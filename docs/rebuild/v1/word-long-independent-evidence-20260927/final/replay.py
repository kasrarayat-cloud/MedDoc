"""Replay existing artifacts once on the frozen repaired engine; add exact author boundary."""
import hashlib
import importlib.util
import json
import sys
from pathlib import Path
from unittest.mock import patch
from docx import Document
HERE=Path(__file__).resolve().parent
OLD=HERE.parent
sys.path.insert(0,'/opt/meddoc-pro/.worktrees/v1-word-long-questions/rebuild')
spec=importlib.util.spec_from_file_location('old_word_probe',OLD/'probe.py')
helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)
spec=importlib.util.spec_from_file_location('documents.word_long_repaired',HERE/'word-v3.py')
w=importlib.util.module_from_spec(spec);sys.modules[spec.name]=w;spec.loader.exec_module(w)
assert hashlib.sha256((HERE/'word-v3.py').read_bytes()).hexdigest()=='3f2ca5136a876d153b045e5c8dd37b8b8b24b303833263966664c63290a3b856'
prior=json.loads((OLD/'result.json').read_text())
declared=json.loads((OLD/'predeclared.json').read_text())
(HERE/'predeclared.json').write_text(json.dumps({'retained':declared,'additional':{
    'name':'author-source-boundary','heading':'Report','question':'Please provide the initial','expected':'hold'}},indent=2)+'\n')
rows=[]
with patch('socket.socket.connect',side_effect=AssertionError('Network denied')):
    for old,control in zip(prior['cases'],declared):
        assert old['name']==control['name']
        raw=(OLD/(old['name']+'.docx')).read_bytes()
        assert hashlib.sha256(raw).hexdigest()==old['sha256']
        try:
            if control['action']=='inspect':result=w.inspect(raw)
            else:
                slot='word/document.xml#b/tbl[0]/tr[0]/tc[1]/p[1]' if control['action']=='nested' else 'word/document.xml#b/tbl[0]/tr[0]/tc[0]'
                result={'fields':w.locate(raw,[{'kind':'long_text','label':'What medications are taken?','slot':slot}]),'notes':[]}
        except Exception as exc:result={'error_type':type(exc).__name__,'error':str(exc)}
        rows.append({'name':old['name'],'original_sha256':old['sha256'],'expected':old['expected'],'source':old['source'],'finished':result})
    doc=Document();doc.add_heading('Report',1);doc.add_paragraph('Please provide the initial');helper.box(doc)
    raw=helper.encode(doc);(HERE/'author-source-boundary.docx').write_bytes(raw)
    rows.append({'name':'author-source-boundary','original_sha256':hashlib.sha256(raw).hexdigest(),'expected':'hold',
                 'finished':w.inspect(raw)})
result={'source_commit':'cf04596ed841064bb883af579e64c444cd462a08','engine_sha256':hashlib.sha256((HERE/'word-v3.py').read_bytes()).hexdigest(),'cases':rows}
(HERE/'result.json').write_text(json.dumps(result,indent=2)+'\n')
for row in rows:
    value=row['finished']
    print(row['name'], 'field_count=',len(value.get('fields',[])), 'context=',[f['context'] for f in value.get('fields',[])],
          'notes=',value.get('notes',[]), 'error=',value.get('error'))
