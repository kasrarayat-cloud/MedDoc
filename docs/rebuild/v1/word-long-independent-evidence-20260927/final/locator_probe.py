"""Finite resolved-locator ownership check with real pure engine, no database/provider."""
import copy
import hashlib
import io
import json
import os
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

ROOT=Path('/opt/meddoc-pro/.worktrees/v1-word-long-questions/rebuild')
HERE=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT))
os.environ.setdefault('DJANGO_SETTINGS_MODULE','quillmd.settings')
import django
django.setup()
from docx import Document
from core import mapping_assistance as m, mapping_wire as wire
from documents import word_v3 as w


def main():
    doc=Document();doc.add_heading('Actual source section',1);doc.add_table(rows=1,cols=2)
    out=io.BytesIO();doc.save(out);raw=out.getvalue()
    path=HERE/'locator-source.docx';path.write_bytes(raw)
    digest=hashlib.sha256(raw).hexdigest()
    owner='word/document.xml#b/tbl[0]/tr[0]/tc[0]'
    slots=[owner,owner+'/p[0]','word/document.xml#b/tbl[0]/tr[0]/tc[1]']
    label='Complete original question'
    original=w.locate(raw,[{'kind':'text','label':label,'slot':owner}])[0]
    context='Retained complete prior DONOR ONLY context '
    variants=[('exact',None,None,0),('fingerprint','fingerprint','0'*64,0),
              ('write-target','target',{'type':'cell_replace'},0),
              ('descendant',None,None,1),('foreign',None,None,2)]
    (HERE/'locator-predeclared.json').write_text(json.dumps(variants,indent=2)+'\n')
    results=[]
    def locate_fields(local,targets):
        assert local==path
        return {'sha256':digest,'fields':w.locate(raw,targets)}
    with patch('socket.socket.connect',side_effect=AssertionError('Network denied')),patch.object(m,'private_path',return_value=path),patch.object(m,'locate_fields',side_effect=locate_fields):
        for name,key,value,sid in variants:
            prior=copy.deepcopy(original);prior['context']=context
            if key:prior['locator'][key]=value
            version=SimpleNamespace(fields=[prior],template=SimpleNamespace(sha256=digest,original_path='synthetic'))
            payload={'source_sha256':digest,'slots':[{'slot':slot,'text':'','heading':'Actual source section','row_label':'','options':[]} for slot in slots],
                     'fields':[{'key':prior['key'],'label':prior['label'],'kind':prior['kind'],'options':[],'slots':[owner]}]}
            payload['text_refs']=wire.catalog(payload)
            compiled={'request':{'messages':[{'content':json.dumps(payload)}]},'request_sha256':'synthetic-local-request','profile_sha256':'synthetic-local-profile'}
            body={'f':{'k':[],'r':[0],'n':[],'u':[]},'s':{'c':[],'n':[0,1,2],'u':[]},
                  'x':[[0,sid,0,payload['text_refs'].index([0,4,0]),0,len(label),-1]],'e':False}
            try:
                result=m.validate_proposal(version,compiled,body)
                rows={'fields':result['fields'],'unresolved':result['review']['unresolved']}
            except Exception as exc:
                rows={'error_type':type(exc).__name__,'error':str(exc)}
            results.append({'case':name,'prior':prior,'output':rows})
    result={'engine_sha256':hashlib.sha256(Path(w.__file__).read_bytes()).hexdigest(),
            'mapping_sha256':hashlib.sha256(Path(m.__file__).read_bytes()).hexdigest(), 'cases':results}
    (HERE/'locator-result.json').write_text(json.dumps(result,indent=2)+'\n')
    for row in results:
        print(row['case'],json.dumps(row['output']))


if __name__=='__main__':main()
