"""Bounded independent pure-engine controls; no original corpus or provider."""
import hashlib
import importlib.util
import io
import json
import sys
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
sys.path.insert(0, '/opt/meddoc-pro/.worktrees/v1-word-long-questions/rebuild')


def load(name, filename):
    spec = importlib.util.spec_from_file_location(name, HERE/filename)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


def box(doc):
    table = doc.add_table(rows=1, cols=1)
    borders = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single'); borders.append(el)
    table.cell(0,0)._tc.get_or_add_tcPr().append(borders)
    return table


def encode(doc):
    out=io.BytesIO(); doc.save(out); return out.getvalue()


def controls():
    for name, question, expected in (
        ('adjective-positive','Please describe the changes since the initial report?', 'map'),
        ('plain-sign','Please sign the initial report here?', 'hold'),
        ('countersign','Please countersign the initial report?', 'hold'),
        ('initial-imperative','Please initial the initial report?', 'hold'),
        ('signature-noun','Please provide your signature on the initial report?', 'hold'),
        ('astral-1000','What ' + '\U0001f9ea'*994 + '?','map'),
        ('astral-1001','What ' + '\U0001f9ea'*995 + '?','hold'),
    ):
        doc=Document();doc.add_paragraph(question);box(doc)
        yield name,doc,'inspect',expected,question
    for name, heading, expected in (
        ('root-manual-donor', 'Clinical section ' + 'continued '*16 + 'DONOR ONLY','context'),
        ('manual-1000','C'*990+'DONOR ONLY','context'),
        ('manual-1001','C'*991+'DONOR ONLY','error'),
    ):
        doc=Document();doc.add_heading(heading,1);doc.add_paragraph('What medications are taken?');box(doc)
        yield name,doc,'locate',expected,heading
    doc=Document();p=doc.add_paragraph('DONOR ONLY');ps=OxmlElement('w:pStyle');ps.set(qn('w:val'),'MissingSourceStyle');p._p.get_or_add_pPr().append(ps)
    box(doc)
    yield 'manual-unknown-scope',doc,'locate','error',''
    doc=Document();table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).paragraphs[0].text='DONOR ONLY'
    table.cell(0,0).paragraphs[0].style='Heading 1'
    table.cell(0,1).paragraphs[0].text='PATIENT ONLY'
    table.cell(0,1).paragraphs[0].style='Heading 1'
    table.cell(0,1).add_paragraph()
    yield 'manual-local-sibling-scope',doc,'nested','local','PATIENT ONLY'


def main():
    cases=list(controls())
    (HERE/'predeclared.json').write_text(json.dumps([{'name':n,'action':a,'expected':e,'source':s} for n,_,a,e,s in cases],indent=2)+'\n')
    engines={'base':load('documents.word_long_base','base-word-v3.py'),
             'finished':load('documents.word_long_finished','finished-word-v3.py')}
    rows=[]
    with patch('socket.socket.connect',side_effect=AssertionError('Network denied')):
        for name,doc,action,expected,source in cases:
            raw=encode(doc);(HERE/(name+'.docx')).write_bytes(raw)
            row={'name':name,'expected':expected,'source':source,'sha256':hashlib.sha256(raw).hexdigest(),'engines':{}}
            for key,w in engines.items():
                try:
                    if action=='inspect':
                        result=w.inspect(raw)
                    else:
                        slot='word/document.xml#b/tbl[0]/tr[0]/tc[1]/p[1]' if action=='nested' else 'word/document.xml#b/tbl[0]/tr[0]/tc[0]'
                        result={'fields':w.locate(raw,[{'kind':'long_text','label':'What medications are taken?','slot':slot}]),'notes':[]}
                    row['engines'][key]=result
                except Exception as exc:
                    row['engines'][key]={'error_type':type(exc).__name__,'error':str(exc)}
            rows.append(row)
            finished=row['engines']['finished']
            print(name, 'fields=',[(f['label'],f['context'],f['locator']['slot']) for f in finished.get('fields',[])], 'notes=',finished.get('notes',[]),'error=',finished.get('error'))
    (HERE/'result.json').write_text(json.dumps({'finished_commit':'90c5aa7bfbab8e24722beb1189eee6043bd73ac3',
        'engine_sha256':hashlib.sha256((HERE/'finished-word-v3.py').read_bytes()).hexdigest(),'cases':rows},indent=2)+'\n')


if __name__=='__main__':main()
