import io,json,hashlib,sys,os
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,os.environ.get("TASK_ROOT","/opt/meddoc-pro/.worktrees/v1-release/rebuild"))
from docx import Document
from documents import word_v3 as w
heading="Clinical section " + "continued "*16 + "DONOR ONLY"
doc=Document();doc.add_heading(heading,level=1);doc.add_paragraph("What medications are taken?");doc.add_table(rows=1,cols=1)
buf=io.BytesIO();doc.save(buf)
with patch("socket.socket.connect",side_effect=AssertionError("Network denied")):
 fields=w.locate(buf.getvalue(),[{"kind":"long_text","label":"What medications are taken?","slot":"word/document.xml#b/tbl[0]/tr[0]/tc[0]"}])
print(json.dumps({"engine_sha256":hashlib.sha256(Path(w.__file__).read_bytes()).hexdigest(),"heading":heading,"fields":fields},indent=2))
