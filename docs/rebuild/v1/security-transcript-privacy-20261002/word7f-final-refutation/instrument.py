import os, sys, io, json, hashlib, socket, time, copy, ast, subprocess
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace

ROOT = Path('/opt/meddoc-pro/.worktrees/v1-sol-word-coverage')
SCRATCH = Path('/opt/meddoc-pro/rebuild/.runtime/v1-sol-continuation-20261002/security')
COMMIT = '7f3b003d4b1478cad086ad88d8734166547b6498'
assert subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip() == COMMIT
assert not subprocess.check_output(['git','status','--porcelain'],cwd=ROOT,text=True).strip()
paths = ['rebuild/documents/' + n for n in ('word_v3.py', 'engine.py', 'service.py', 'import_selection.py', 'source_policy.py')]
paths += ['rebuild/core/' + n for n in ('import_binding.py', 'v3_mapping.py', 'views.py', 'jobs.py', 'migrations/0035_import_manifest.py')]
witness = SCRATCH / 'refute-word7f-final-source'
witness.mkdir(mode=0o700)
files = []
for rel in paths:
    raw = (ROOT / rel).read_bytes()
    assert raw == subprocess.check_output(['git', 'show', COMMIT + ':' + rel], cwd=ROOT)
    out = witness / rel
    out.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
    fd = os.open(out, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600)
    with os.fdopen(fd, 'wb') as f: f.write(raw)
    files.append({'path': rel, 'sha256': hashlib.sha256(raw).hexdigest()})
excluded = 'rebuild/core/tests/test_source_coverage_workflow.py'
diff = subprocess.check_output(['git', 'diff', '--', excluded], cwd=ROOT)
(witness / 'excluded-fixture.diff').write_bytes(diff)
manifest = {'commit': COMMIT, 'contract': 'clean final source; ten exact committed product components',
    'files': files, 'measured_fixture': {'path': excluded, 'sha256': hashlib.sha256((ROOT/excluded).read_bytes()).hexdigest(), 'diff_sha256': hashlib.sha256(diff).hexdigest()}}
(witness / 'manifest.json').write_text(json.dumps(manifest, indent=2))
def verify():
    assert subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip() == COMMIT
    for row in files: assert hashlib.sha256((ROOT/row['path']).read_bytes()).hexdigest() == row['sha256']
verify()
out = SCRATCH / 'refute-word7f-final-cases'
out.mkdir(mode=0o700)
sys.path.insert(0, str(ROOT/'rebuild'))
os.environ.update(DJANGO_SETTINGS_MODULE='quillmd.settings', QUILLMD_ENV_MANAGED='1', QUILLMD_SECRET_KEY='synthetic-read-only',
    QUILLMD_DB_PASSWORD='synthetic-unusable', QUILLMD_DB_PORT='1', QUILLMD_DATA_MODE='synthetic', QUILLMD_AI_MODE='fixture', QUILLMD_STATE_DIR=str(out))
import django
django.setup()
from docx import Document
from documents.service import inspect_template, engine_fingerprint
from documents import word_v3
from documents.import_selection import digest, ENGINE_VERSIONS
from documents.test_source_coverage import data
from core import import_binding, v3_mapping
from django.db.backends.base.base import BaseDatabaseWrapper
from django.core.exceptions import ValidationError

def retain(name, raw):
    p = out/name
    fd = os.open(p, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600)
    with os.fdopen(fd, 'wb') as f: f.write(raw)
    return p

cases = [('date_tab_separators', 'Date: [DD\t/\tMM\t/\tYYYY]'),
    ('date_nbsp_separators', 'Date: [DD\u00a0/\u00a0MM/YYYY]'),
    ('date_line_separator', 'Date: [DD\u2028/\u2028MM/YYYY]'),
    ('date_newline_control', 'Date: [DD\n/ MM/YYYY]'),
    ('spaced_nondate_menu', 'Response: [yes / no]'),
    ('prefixed_spaced_menu', 'Response: [type one / type two]'),
    ('unclosed_before_placeholder', 'Site: [left/right; Comments: [enter notes]'),
    ('unclosed_before_date', 'Site: [left/right; Date: [DD/MM/YYYY]'),
    ('unclosed_plain_control', 'Location: [upper/lower'),
    ('closed_paired_control', 'Site: [left/right]; Date: [DD/MM/YYYY]')]

started = time.monotonic()
results = []
quarantine = []
historical = []
with patch.object(socket.socket, 'connect', side_effect=AssertionError('Network forbidden')), patch.object(BaseDatabaseWrapper, 'ensure_connection', side_effect=AssertionError('DB forbidden')):
    for name, text in cases:
        doc = Document()
        doc.add_paragraph(text)
        doc.add_paragraph('Reason: ______')
        source = retain(name+'-original.docx', data(doc))
        for engine in ('legacy', 'v3', 'auto'):
            assert time.monotonic()-started < 120
            row = {'name': name, 'source_text': text, 'requested_engine': engine, 'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest()}
            try: r = inspect_template(source, engine=engine)
            except Exception as e:
                row['inspection_refusal'] = {'type': type(e).__name__, 'code': getattr(e, 'code', None), 'message': str(e)}
                results.append(row)
                continue
            row.update(fields=r['fields'], manifest=r['import_manifest'])
            version = SimpleNamespace(template=SimpleNamespace(sha256=r['sha256'], original_path=source.name, warnings=r.get('warnings', [])),
                contract=r.get('contract', 'legacy'), fields=r['fields'], import_manifest=r['import_manifest'], corrections={})
            try: import_binding.verify(version); row['admission_binding'] = 'allowed'
            except ValidationError as e: row['admission_binding'] = 'refused'; row['binding_refusal'] = str(e)
            if name.startswith('date_') and engine=='v3':
                candidate=next((f for f in r['fields'] if f['label']!='Reason'), None)
                if candidate:
                    try:
                        filled=word_v3.fill(source.read_bytes(), r['fields'], {candidate['key']:'2026-10-03'})
                        row['actual_export_paragraphs']=[p.text for p in Document(io.BytesIO(filled)).paragraphs]
                        row['export_sha256']=hashlib.sha256(filled).hexdigest()
                        retain(name+'-actual-export.docx',filled)
                    except Exception as e: row['fill_refusal']=str(e)
            results.append(row)
    for source_hash, declaration in word_v3.KNOWN_LAYOUT_FAILURES.items():
        for engine in ('legacy', 'v3'):
            fields=[]
            bound={'source_sha256':source_hash,'engine':engine,'engine_version':ENGINE_VERSIONS[engine],
                'engine_code_sha256':engine_fingerprint(engine,source_sha256=source_hash),'mapping_sha256':digest(fields),
                'initial_mapping_sha256':digest(fields),'coverage':{}}
            version=SimpleNamespace(template=SimpleNamespace(sha256=source_hash, original_path='unused-original.docx', warnings=[]),
                contract=engine,fields=fields,import_manifest=bound,corrections={})
            try:import_binding.verify(version);outcome='allowed'
            except ValidationError as e:outcome={'refused':str(e)}
            quarantine.append({'case':declaration['case'],'engine':engine,'coverage_stripped_outcome':outcome,
                'scope':'direct source-hash admission-binding control; no original/corpus/browser qualification'})
        version.import_manifest={};version.contract='legacy'
        historical.append({'case':declaration['case'],'empty_historical_manifest_legacy_admission':list(v3_mapping.admission_refusals(version)),
            'scope':'pure historical guard observation; no actual stored legacy version or browser admission witness'})
verify()
tree=ast.parse((ROOT/'rebuild/core/views.py').read_text())
before=ast.parse(subprocess.check_output(['git','show','09644640759200813c40c3246b92299999a64aca:rebuild/core/views.py'],cwd=ROOT,text=True))
def export_ast(t):return ast.dump(next(n for n in t.body if isinstance(n,(ast.FunctionDef,ast.AsyncFunctionDef)) and n.name=='export'),include_attributes=False)
report={'source_manifest':manifest,'cases':results,'stripped_quarantine_controls':quarantine,'historical_guard_observations':historical,
    'stored_export_route_AST_unchanged_since096':export_ast(tree)==export_ast(before),'provider_requests':0,'database_connections':0,
    'elapsed_seconds':round(time.monotonic()-started,3)}
p=SCRATCH/'refute-word7f-final.json'
fd=os.open(p,os.O_WRONLY|os.O_CREAT|os.O_EXCL|os.O_NOFOLLOW,0o600)
with os.fdopen(fd,'w') as f:json.dump(report,f,indent=2)
print(json.dumps({'commit':COMMIT,'elapsed_seconds':report['elapsed_seconds'],'inspections':len(results),
    'cases':[{'name':r['name'],'engine':r['requested_engine'],'fields':[(f['label'],f['kind']) for f in r.get('fields',[])],
        'issues':[i['code'] for i in r.get('manifest',{}).get('coverage',{}).get('selected_engine_issues',[])],
        'advanced_unresolved':len(r.get('manifest',{}).get('coverage',{}).get('advanced_unresolved_areas',[])),
        'block':r.get('manifest',{}).get('coverage',{}).get('block_generation'),'binding':r.get('admission_binding'),
        'export':r.get('actual_export_paragraphs'),'refusal':r.get('inspection_refusal')} for r in results],
    'quarantine':quarantine,'historical':historical,'stored_export_AST_unchanged':report['stored_export_route_AST_unchanged_since096']},indent=2))
