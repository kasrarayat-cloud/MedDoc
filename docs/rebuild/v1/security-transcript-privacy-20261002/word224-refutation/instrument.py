import os,sys,io,json,hashlib,socket,time
from pathlib import Path
from unittest.mock import patch
ROOT=Path('/opt/meddoc-pro/.worktrees/v1-sol-word-coverage')
SCRATCH=Path('/opt/meddoc-pro/rebuild/.runtime/v1-sol-continuation-20261002/security')
manifest=json.loads((SCRATCH/'refute-word224-source/manifest.json').read_text())
def verify():
    for r in manifest['files']:
        assert hashlib.sha256((ROOT/r['path']).read_bytes()).hexdigest()==r['sha256']
verify()
sys.path.insert(0,str(ROOT/'rebuild'))
os.environ.update(DJANGO_SETTINGS_MODULE='quillmd.settings',QUILLMD_ENV_MANAGED='1',
    QUILLMD_SECRET_KEY='synthetic-read-only',QUILLMD_DB_PASSWORD='synthetic-unusable',
    QUILLMD_DB_PORT='1',QUILLMD_DATA_MODE='synthetic',QUILLMD_AI_MODE='fixture')
import django
django.setup()
from docx import Document
from documents.service import inspect_template
from documents import word_v3
from documents.test_source_coverage import data
from django.db.backends.base.base import BaseDatabaseWrapper
out=SCRATCH/'refute-word224-cases'
out.mkdir(mode=0o700)
inputs=[('apostrophe_alternatives',"Role: [patient's/carer's]"),
    ('parenthesized_alternatives','Role: [patient/carer (specify)]'),
    ('digit_alternatives','Severity: [mild/moderate/grade 3]'),
    ('punctuated_alternatives','Consent: [yes/no?]'),
    ('positive_date','Date: [dd/mm/yyyy]'),
    ('source_list_nocolon','Gender\na) Male\nb) Female'),
    ('source_list_period','Gender.\na) Male\nb) Female'),
    ('source_list_instruction','Discharge instructions\na) Call clinic\nb) Keep wound dry')]
rows=[];started=time.monotonic()
with patch.object(socket.socket,'connect',side_effect=AssertionError('Network forbidden')), \
    patch.object(BaseDatabaseWrapper,'ensure_connection',side_effect=AssertionError('DB forbidden')):
    for name,text in inputs:
        assert time.monotonic()-started<90
        doc=Document()
        for paragraph in text.split('\n'):doc.add_paragraph(paragraph)
        doc.add_paragraph('Reason: ______')
        original=data(doc);source=out/(name+'-original.docx')
        fd=os.open(source,os.O_WRONLY|os.O_CREAT|os.O_EXCL|os.O_NOFOLLOW,0o600)
        with os.fdopen(fd,'wb') as f:f.write(original)
        r=inspect_template(source,engine='v3')
        row={'name':name,'source_text':text,'source_sha256':hashlib.sha256(original).hexdigest(),
            'fields':r['fields'],'coverage':r['import_manifest']['coverage']}
        candidate=next((f for f in r['fields'] if f['label']!='Reason'),None)
        if candidate and name!='positive_date':
            value='SYNTHETIC UNRESTRICTED REPLACEMENT'
            try:
                filled=word_v3.fill(original,r['fields'],{candidate['key']:value})
                dest=out/(name+'-actual-export.docx')
                fd=os.open(dest,os.O_WRONLY|os.O_CREAT|os.O_EXCL|os.O_NOFOLLOW,0o600)
                with os.fdopen(fd,'wb') as f:f.write(filled)
                row['actual_written_paragraphs']=[p.text for p in Document(io.BytesIO(filled)).paragraphs]
                row['export_sha256']=hashlib.sha256(filled).hexdigest()
            except Exception as e:row['fill_refusal']={'type':type(e).__name__,'message':str(e)}
        rows.append(row)
verify()
p=SCRATCH/'refute-word224.json';fd=os.open(p,os.O_WRONLY|os.O_CREAT|os.O_EXCL|os.O_NOFOLLOW,0o600)
result={'source_manifest':manifest,'cases':rows,'provider_requests':0,'database_connections':0,
    'elapsed_seconds':round(time.monotonic()-started,3)}
with os.fdopen(fd,'w') as f:json.dump(result,f,indent=2)
print(json.dumps({'commit':manifest['commit'],'elapsed_seconds':result['elapsed_seconds'],
    'cases':[{'name':r['name'],'source':r['source_text'],'fields':[(f['label'],f['kind']) for f in r['fields']],
        'issues':[i['code'] for i in r['coverage'].get('selected_engine_issues',[])],
        'block_generation':r['coverage']['block_generation'],'automatic_complete':r['coverage']['automatic_complete'],
        'actual_written_paragraphs':r.get('actual_written_paragraphs'),'fill_refusal':r.get('fill_refusal')}
        for r in rows]},indent=2))
