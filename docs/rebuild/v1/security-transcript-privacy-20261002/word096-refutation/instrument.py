import os,sys,io,json,hashlib,socket,time,copy
from pathlib import Path
from unittest.mock import patch
ROOT=Path('/opt/meddoc-pro/.worktrees/v1-sol-word-coverage')
SCRATCH=Path('/opt/meddoc-pro/rebuild/.runtime/v1-sol-continuation-20261002/security')
manifest=json.loads((SCRATCH/'refute-word096-source/manifest.json').read_text())
def verify():
    for r in manifest['files']:
        assert hashlib.sha256((ROOT/r['path']).read_bytes()).hexdigest()==r['sha256']
verify()
sys.path.insert(0,str(ROOT/'rebuild'))
os.environ.update(DJANGO_SETTINGS_MODULE='quillmd.settings',QUILLMD_ENV_MANAGED='1',QUILLMD_SECRET_KEY='synthetic-read-only',QUILLMD_DB_PASSWORD='synthetic-unusable',QUILLMD_DB_PORT='1',QUILLMD_DATA_MODE='synthetic',QUILLMD_AI_MODE='fixture')
import django;django.setup()
from docx import Document
from documents.service import inspect_template
from documents import word_v3
from documents.test_source_coverage import data
from django.db.backends.base.base import BaseDatabaseWrapper
out=SCRATCH/'refute-word096-cases';out.mkdir(mode=0o700)
def retain(name,raw):
    p=out/name;fd=os.open(p,os.O_WRONLY|os.O_CREAT|os.O_EXCL|os.O_NOFOLLOW,0o600)
    with os.fdopen(fd,'wb') as f:f.write(raw)
    return hashlib.sha256(raw).hexdigest()
cases=[('ordinary_split_runs',['Role: [patient','/','carer]'],'normal'),
('punctuation_split_runs',["Role: [patient's",'/',"carer's]"],'normal'),
('newline_inside_alternative',['Role: [patient/\ncarer]'],'normal'),
('tab_inside_alternative',['Role: [patient/\tcarer]'],'normal'),
('trailing_empty_alternative',['Role: [patient/carer/]'],'normal'),
('middle_empty_alternative',['Role: [patient//carer]'],'normal'),
('type_label_alternatives',['Procedure type: [type A/type B]'],'normal'),
('state_label_alternatives',['Jurisdiction: [state/local]'],'normal'),
('action_alternatives',['Action: [add/remove]'],'normal'),
('directive_positive',['Plan: [enter left/right]'],'normal'),
('date_positive',['Date: [dd/mm/yyyy]'],'normal'),
('signature_forged',['Patient signature: ______'],'forge'),
('witness_forged',['Witness signature: ______'],'forge'),
('initials_forged',['Nurse initials: ______'],'forge'),
('patient_policy_forged',['Completed by the patient:','Comments: ______'],'policy'),
('ordinary_initial',['Initial assessment: ______'],'ordinary')]
started=time.monotonic();rows=[]
with patch.object(socket.socket,'connect',side_effect=AssertionError('Network forbidden')),patch.object(BaseDatabaseWrapper,'ensure_connection',side_effect=AssertionError('DB forbidden')):
    for name,segments,mode in cases:
        assert time.monotonic()-started<110
        doc=Document()
        if mode=='policy':
            for text in segments:doc.add_paragraph(text)
            doc.add_heading('Clinician',1)
        else:
            p=doc.add_paragraph()
            for i,text in enumerate(segments):
                rr=p.add_run(text);rr.bold=bool(i%2)
        doc.add_paragraph('Reason: ______');raw=data(doc)
        source=out/(name+'-original.docx');sourcehash=retain(source.name,raw)
        try:r=inspect_template(source,engine='v3')
        except Exception as e:
            rows.append({'name':name,'inspect_refusal':{'type':type(e).__name__,'message':str(e)}});continue
        row={'name':name,'source_paragraphs':[p.text for p in doc.paragraphs],'source_sha256':sourcehash,'fields':r['fields'],'coverage':r['import_manifest']['coverage']}
        candidate=next((f for f in r['fields'] if f['label']!='Reason'),None)
        if candidate:
            proposed=copy.deepcopy(candidate)
            if mode=='forge':
                proposed.update(kind='text',label='Ordinary renamed answer',write_protected=False)
                proposed.pop('review_reason',None)
            if mode=='policy':proposed.pop('answer_policy',None);proposed['label']='Ordinary renamed answer'
            try:
                filled=word_v3.fill(raw,[proposed],{proposed['key']:'SYNTHETIC UNRESTRICTED REPLACEMENT'})
                row['actual_written_paragraphs']=[p.text for p in Document(io.BytesIO(filled)).paragraphs]
                row['export_sha256']=retain(name+'-actual-export.docx',filled)
            except Exception as e:row['fill_refusal']={'type':type(e).__name__,'code':getattr(e,'code',None),'message':str(e)}
        rows.append(row)
verify()
result={'source_manifest':manifest,'cases':rows,'provider_requests':0,'database_connections':0,'elapsed_seconds':round(time.monotonic()-started,3)}
p=SCRATCH/'refute-word096.json';fd=os.open(p,os.O_WRONLY|os.O_CREAT|os.O_EXCL|os.O_NOFOLLOW,0o600)
with os.fdopen(fd,'w') as f:json.dump(result,f,indent=2)
print(json.dumps({'commit':manifest['commit'],'elapsed_seconds':result['elapsed_seconds'],'cases':[{'name':r['name'],'source':r.get('source_paragraphs'),'fields':[(f['label'],f['kind']) for f in r.get('fields',[])],'issues':[i['code'] for i in r.get('coverage',{}).get('selected_engine_issues',[])],'block_generation':r.get('coverage',{}).get('block_generation'),'actual_written_paragraphs':r.get('actual_written_paragraphs'),'fill_refusal':r.get('fill_refusal'),'inspect_refusal':r.get('inspect_refusal')} for r in rows]},indent=2))
