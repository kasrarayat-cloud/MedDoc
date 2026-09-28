import hashlib, io, json, os, socket, sys
from pathlib import Path
ROOT=Path('/opt/meddoc-pro/.worktrees/v1-document-relations')
sys.path.insert(0,str(ROOT/'rebuild'))
os.environ['QUILLMD_ENV_MANAGED']='1'
socket.socket.connect=lambda *a,**k: (_ for _ in ()).throw(RuntimeError('network forbidden'))
from docx import Document
from docx.oxml import OxmlElement
from documents import word_v3 as w
from documents.test_question_boxes import data
NAMES=['source_policy.py','test_source_policy.py','word_v3.py']
def hashes(): return {n:hashlib.sha256((ROOT/'rebuild/documents'/n).read_bytes()).hexdigest() for n in NAMES}
def pair(doc,label='What progress was observed?'):
 t=doc.add_table(rows=1,cols=2);t.cell(0,0).text=label;return t
results=[]
def check(name,doc,expect_patient=False,expect_attestation=False,fill_value=None,mutate=None,expect_code=None):
 try:
  fs=w.inspect(data(doc))['fields']; f=next(f for f in fs if f['locator']['target']['type']=='cell')
  policy=f.get('answer_policy'); outcome={'patient':isinstance(policy,dict) and policy.get('source_author')=='patient','attestation':isinstance(policy,dict) and policy.get('attestation')=='document_review','policy':policy,'kind':f['kind'],'label':f['label'],'write_protected':f.get('write_protected')}
  matched=outcome['patient']==expect_patient and outcome['attestation']==expect_attestation
  if fill_value is not None:
   if mutate:mutate()
   try:w.fill(data(doc),fs,{f['key']:fill_value});code='accepted'
   except w.V3Failure as e:code=e.code
   outcome['fill']=code;matched=matched and code==expect_code
  results.append({'case':name,'expected':{'patient':expect_patient,'attestation':expect_attestation,'fill':expect_code},'observed':outcome,'matched':matched})
 except Exception as e:results.append({'case':name,'error':type(e).__name__+': '+str(e),'matched':False})
before=hashes()
for role in ['For the patient:','To be completed by the patient:','Completed by patient:']:
 d=Document();d.add_paragraph(role);pair(d);check('explicit '+role,d,True)
d=Document();d.add_paragraph('Patient details:');pair(d);check('patient fact does not assign role',d)
d=Document();d.add_paragraph('For the patient:');d.add_paragraph('For the clinician:');pair(d);check('role boundary',d)
d=Document();d.add_paragraph('For the patient:');d.add_heading('Clinical assessment',1);pair(d);check('heading boundary',d)
d=Document();d.add_paragraph('For the patient:');p=d.add_paragraph();p._p.get_or_add_pPr().append(OxmlElement('w:sectPr'));pair(d);check('distinct section boundary',d)
d=Document();p=d.add_paragraph('For the patient:');p._p.get_or_add_pPr().append(OxmlElement('w:sectPr'));pair(d);check('role paragraph itself ends section',d)
d=Document();d.add_paragraph('For the patient:');d.add_paragraph('For example: describe changes since your last visit.');pair(d);check('ordinary example is not a new role',d,True)
d=Document();d.add_paragraph('For the patient:');pair(d,'Patient signature');check('signature protected independently',d,True,fill_value='Signed',expect_code='protected_signing_area')
question='Please confirm that the patient has had sight of this report and/or had an opportunity to add their comments, if appropriate or necessary.'
d=Document();pair(d,question);check('full attestation source',d,False,True)
d=Document();pair(d,'Confirm the patient report of chest pain was reviewed by the clinician.');check('clinician review of symptom report',d)
d=Document();p=d.add_paragraph('For the patient:');pair(d);check('unchanged source accepted',d,True,fill_value='Synthetic progress',expect_code='accepted')
d=Document();p=d.add_paragraph('For the patient:');pair(d);check('changed source refused',d,True,fill_value='Synthetic progress',mutate=lambda:setattr(p,'text','For the clinician:'),expect_code='mapping_mismatch')
d=Document();t=pair(d,question);check('changed attestation refused',d,False,True,fill_value='Yes',mutate=lambda:setattr(t.cell(0,0),'text','What symptoms did the patient report?'),expect_code='mapping_mismatch')
print(json.dumps({'source_before':before,'source_after':hashes(),'cases':results},indent=2))
