import hashlib, io, json, os, subprocess, sys, tempfile
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
ROOT=Path('/opt/meddoc-pro/.worktrees/v1-document-relations')
TMP=Path('/tmp/claude-1000/v1-disposition-integration-20260928/source-policy');TMP.mkdir(parents=True,exist_ok=True)
rows=[]
for case in ('positive','section','example','symptom'):
 d=Document()
 if case!='symptom':
  p=d.add_paragraph('For the patient:')
  if case=='section':p._p.get_or_add_pPr().append(OxmlElement('w:sectPr'))
  if case=='example':d.add_paragraph('For example: describe changes since your last visit.')
 t=d.add_table(rows=1,cols=2);t.cell(0,0).text='Confirm the patient report of chest pain was reviewed by the clinician.' if case=='symptom' else 'What progress was observed?'
 with tempfile.TemporaryDirectory(prefix=case+'-',dir=TMP) as directory:
  p=Path(directory);d.save(p/'input.docx');(p/'request.json').write_text(json.dumps({'operation':'inspect','format':'docx','engine':'v3'}))
  r=subprocess.run([sys.executable,'-I','-B',str(ROOT/'rebuild/documents/worker.py'),str(p)],env={'PATH':'/usr/bin:/bin','LANG':'C.UTF-8','TZ':'UTC'},cwd=p,capture_output=True,text=True,timeout=30)
  out=json.loads(r.stdout);rows.append({'case':case,'returncode':r.returncode,'stderr':r.stderr,'sandbox':out.get('sandbox'),'error':out.get('error'),'fields':[{'key':f['key'],'label':f['label'],'answer_policy':f.get('answer_policy')} for f in out.get('fields',[])]})
print(json.dumps(rows,indent=2))
