"""Pure synthetic refutation of frozen document-relations diff; no DB or external I/O."""
from pathlib import Path
import copy
import hashlib
import io
import json
import socket
import subprocess
import sys
import types

ROOT = Path('/opt/meddoc-pro/.worktrees/v1-document-relations')
sys.path.insert(0, str(ROOT / 'rebuild'))
def denied(*a, **kw): raise AssertionError('No network')
socket.socket.connect = denied
socket.create_connection = denied
from docx import Document
from documents import word_v3 as w
from documents.test_question_boxes import box, data
from documents.test_document_relations import link, numbered

source = (ROOT / 'rebuild/documents/word_v3.py').read_bytes()
old_source = subprocess.check_output(['git', '-C', str(ROOT), 'show', '9f2bdad03dfc135befceae12a600bb0f86c7f7be:rebuild/documents/word_v3.py'])
baseline = types.ModuleType('documents.independent_relation_baseline')
baseline.__package__ = 'documents'
sys.modules[baseline.__name__] = baseline
exec(compile(old_source, '9f2bdad0:word_v3.py', 'exec'), baseline.__dict__)


def conditional_choice():
    doc = Document()
    question = 'If pain is present, select one severity:'
    doc.add_paragraph(question)
    numbered(doc, ('Mild', 'Severe'), 'lowerLetter')
    doc.add_paragraph('For help, please refer to the guidance at:')
    link(doc.add_paragraph())
    original = data(doc)
    old = baseline.inspect(original)['fields']; new = w.inspect(original)['fields']
    return {'input': question + ' / a. Mild / b. Severe / For help, please refer to the guidance at: / fixed link',
            'baseline_fields': [{'label': f['label'], 'options': f.get('options')} for f in old],
            'new_fields': [{'label': f['label'], 'options': f.get('options')} for f in new],
            'demonstrated_wrong_output': bool([f for f in old if f['label'] == question.rstrip(':') and f.get('options') == ['Mild', 'Severe']]) and not any(f.get('options') == ['Mild', 'Severe'] for f in new)}


def moved_reference_destination():
    doc = Document(); doc.add_heading('History', 1)
    doc.add_paragraph('What symptoms were observed?'); first = box(doc)
    follow = doc.add_paragraph('In relation to the above question, what did the patient say?'); box(doc)
    original = data(doc); fields = w.inspect(original)['fields']
    answer = next(f for f in fields if f.get('source_questions'))
    # Insert an identical physical box without moving either referenced paragraph's indexed path.
    follow._p.addprevious(copy.deepcopy(first._tbl))
    changed = data(doc)
    current_fields = w.inspect(changed)['fields']
    try:
        output = w.fill(changed, fields, {answer['key']: 'SYNTHETIC patient response'})
        values = [t.cell(0, 0).text for t in Document(io.BytesIO(output)).tables]
        return {'mapping_target': answer['locator']['slot'], 'relation': answer['source_questions'],
                'fresh_relation_count': sum(bool(f.get('source_questions')) for f in current_fields),
                'physical_box_values_after_fill': values,
                'demonstrated_wrong_output': values == ['', 'SYNTHETIC patient response', '']}
    except w.V3Failure as exc:
        return {'refused': exc.code, 'demonstrated_wrong_output': False}


def other_author_box():
    doc = Document(); doc.add_heading('Parent comments', 1)
    statement = 'This section is for the parent to provide comments:'
    doc.add_paragraph(statement)
    numbered(doc, ('No comments', 'Other (please specify)'))
    box(doc)
    original = data(doc); fields = w.inspect(original)['fields']
    target = next((f for f in fields if f['locator']['target']['type'] == 'cell'), None)
    if not target: return {'demonstrated_wrong_output': False, 'fields': fields}
    try:
        output = w.fill(original, fields, {target['key']: 'SYNTHETIC parent agrees'})
        value = Document(io.BytesIO(output)).tables[0].cell(0, 0).text
        return {'statement_matches_authorship': bool(w.Detector.question_box_authorship(statement)),
                'field': target, 'exported_author_box': value,
                'demonstrated_wrong_output': not target.get('write_protected') and value == 'SYNTHETIC parent agrees'}
    except w.V3Failure as exc:
        return {'refused': exc.code, 'demonstrated_wrong_output': False}


def filled_author_box():
    doc = Document(); doc.add_heading('Manager comments', 1)
    doc.add_paragraph('This section is for the line manager to add comments.'); table = box(doc)
    fields = w.inspect(data(doc))['fields']; field = fields[0]
    table.cell(0, 0).text = 'SYNTHETIC manager statement already present'
    changed = data(doc)
    field.pop('write_protected', None); field.pop('review_reason', None)
    index = w.SlotIndex(w.open_package(changed))
    slot = index.by_id[field['locator']['part'] + '#' + field['locator']['slot']]
    field['locator']['fingerprint'] = w.fingerprint(slot.el)
    field['locator']['target'] = {'type': 'cell_replace', 'all': True}
    try:
        output = w.fill(changed, fields, {field['key']: 'SYNTHETIC forged replacement'})
        value = Document(io.BytesIO(output)).tables[0].cell(0, 0).text
        return {'input': 'A nonempty author-only comments box and forged unprotected current-anchor mapping',
                'exported_author_box': value, 'demonstrated_wrong_output': value == 'SYNTHETIC forged replacement'}
    except w.V3Failure as exc:
        return {'refused': exc.code, 'demonstrated_wrong_output': False}


results = {f.__name__: f() for f in [conditional_choice, moved_reference_destination, other_author_box, filled_author_box]}
assert source == (ROOT / 'rebuild/documents/word_v3.py').read_bytes()
print(json.dumps({'source_sha256': hashlib.sha256(source).hexdigest(), 'source_unchanged': True,
                  'pure_synthetic': True, 'cases': results}, indent=2))
