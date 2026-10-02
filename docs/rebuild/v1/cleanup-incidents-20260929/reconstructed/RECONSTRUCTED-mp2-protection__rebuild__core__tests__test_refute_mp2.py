"""Refutation probes (scratch only). Synthetic BLANK forms; no provider construction."""
import io
import json
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from core import mapping_assistance as mapping, v3_mapping
from core.security import private_path
from core.tests.test_mapping_places import PlaceRule, saved, bordered_box, W, WP, A, WPS
from documents.service import fill_template, assistance_places


def pict_run():
    run = OxmlElement('w:r'); run.append(OxmlElement('w:pict')); return run


class Probe(PlaceRule):
    def show(self, name, version, compiled, payload, result, change):
        adopted = self.adopted(result)
        print('\n=== %s' % name)
        print('held places:', [h for h in assistance_places(private_path(version.template.original_path))['held']])
        print('adopted:', [(f['label'], f['kind'], f['locator']['slot'], f['locator']['target'], f.get('write_protected'))
                           for f in adopted])
        print('answerable adopted:', [f['label'] for f in v3_mapping.answerable(adopted)])
        print('rejected:', result['review']['rejected_changes'])
        print('unresolved:', result['review']['unresolved'])
        return adopted

    def fill_text(self, version, fields, answers):
        out = private_path('outputs/%s.docx' % id(fields)); out.parent.mkdir(parents=True, exist_ok=True)
        fill_template(private_path(version.template.original_path), fields, answers, out)
        doc = Document(io.BytesIO(out.read_bytes()))
        return [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]

    def run_addition(self, name, data, slot, kind, label, roles=(0, 1, 2), drop=()):
        version = self.version(name, data, drop=drop)
        compiled, payload = self.compile(version)
        print('\n--- %s slots:' % name, [(r['slot'], r['text'], r['heading'], r.get('row_label')) for r in payload['slots']])
        print('engine fields:', [(f['label'], f['kind'], f['locator']['slot'], f.get('write_protected')) for f in version.fields])
        change = self.addition(payload, slot, kind, label, roles)
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        adopted = self.show(name, version, compiled, payload, result, change)
        return version, result, adopted

    # P1: a SIGNING question box (the F1 shape) whose one-cell box also holds a picture (a printed signature line).
    def test_p1_signing_box_with_picture(self):
        doc = Document(); doc.add_paragraph('Participant ID: __________')
        doc.add_paragraph('Consultant review', style='Heading 2'); doc.add_paragraph('Signed by the consultant:')
        box = bordered_box(doc); box.cell(0, 0).paragraphs[0]._p.append(pict_run())
        version, result, adopted = self.run_addition('p1', saved(doc), 'b/tbl[0]/tr[0]/tc[0]/p[0]', 'long_text',
                                                     'Consultant review', roles=(1,))
        if adopted:
            print('FILLED:', self.fill_text(version, result['fields'], {adopted[0]['key']: 'SYNTHETIC-ANSWER'}))

    # P1b: same signing box holding printed guidance text instead of a picture.
    def test_p1b_signing_box_with_guidance(self):
        doc = Document(); doc.add_paragraph('Participant ID: __________')
        doc.add_paragraph('Consultant review', style='Heading 2'); doc.add_paragraph('Signed by the consultant:')
        box = bordered_box(doc); box.cell(0, 0).text = 'Use black ink'
        version, result, adopted = self.run_addition('p1b', saved(doc), 'b/tbl[0]/tr[0]/tc[0]/p[0]', 'text',
                                                     'Use black ink', roles=(0,))
        if adopted:
            print('FILLED:', self.fill_text(version, result['fields'], {adopted[0]['key']: 'SYNTHETIC-ANSWER'}))

    # P3: a signature COLUMN cell holding a picture (the engine makes no area there).
    def test_p3_signature_column_with_picture(self):
        doc = Document(); doc.add_paragraph('Clinic record', style='Heading 2')
        table = doc.add_table(rows=2, cols=2); table.cell(0, 0).text = 'Name'; table.cell(0, 1).text = 'Signature'
        table.cell(1, 1).paragraphs[0]._p.append(pict_run())
        version, result, adopted = self.run_addition('p3', saved(doc), 'b/tbl[0]/tr[1]/tc[1]/p[0]', 'text',
                                                     'Clinic record', roles=(1,))
        if adopted:
            print('FILLED:', self.fill_text(version, result['fields'], {adopted[0]['key']: 'SYNTHETIC-ANSWER'}))

    # P4: a typed-underscore signature line captioned BELOW.
    def test_p4_underscore_line_captioned_below(self):
        doc = Document(); doc.add_paragraph('Consent', style='Heading 2')
        doc.add_paragraph('I agree to the procedure described above.')
        doc.add_paragraph('______________________________')
        doc.add_paragraph('Patient signature')
        version, result, adopted = self.run_addition('p4', saved(doc), 'b/p[2]', 'text', 'Consent', roles=(1,))
        if adopted:
            print('FILLED:', self.fill_text(version, result['fields'], {adopted[0]['key']: 'SYNTHETIC-ANSWER'}))

    # P5: a paragraph framed through its STYLE (w:framePr in styles.xml), not its own pPr.
    def test_p5_style_frame(self):
        doc = Document(); doc.add_paragraph('Plan: ______________')
        style = doc.styles.add_style('OfficeFrame', 1)
        ppr = style.element.get_or_add_pPr()
        frame = OxmlElement('w:framePr')
        for n, v in (('w', '2000'), ('hAnchor', 'page'), ('vAnchor', 'page'), ('x', '8000'), ('y', '14000')):
            frame.set(qn('w:' + n), v)
        ppr.insert(0, frame)
        doc.add_paragraph('Office code: ______________', style='OfficeFrame')
        version, result, adopted = self.run_addition('p5', saved(doc), 'b/p[1]', 'text', 'Office code', drop=('Office code',))


def append_body(doc, xml):
    body = doc.element.body
    body.insert(len(body) - 1, etree.fromstring(xml))


class Probe2(Probe):
    # P6a: a content-LOCKED inline plain-text control (Word's write-protected content).
    def test_p6a_locked_inline_control(self):
        doc = Document(); doc.add_paragraph('Plan: ______________')
        append_body(doc, f'<w:p xmlns:w="{W}"><w:r><w:t xml:space="preserve">Office use: </w:t></w:r><w:sdt><w:sdtPr>'
                         '<w:alias w:val="Office use"/><w:tag w:val="office"/><w:lock w:val="contentLocked"/><w:text/></w:sdtPr>'
                         '<w:sdtContent><w:r><w:t xml:space="preserve"> </w:t></w:r></w:sdtContent></w:sdt></w:p>')
        version = self.version('p6a', saved(doc))
        compiled, payload = self.compile(version)
        print('\n--- p6a slots:', [(r['slot'], r['kind'] if 'kind' in r else '', r['text'], r['heading']) for r in payload['slots']])
        print('engine fields:', [(f['label'], f['kind'], f['locator']['slot'], f['locator']['target'], f.get('write_protected')) for f in version.fields])
        for slot, label in (('b/p[1]/sdt[0]', 'Office use'), ('b/p[1]', 'Office use')):
            sid = next((n for n, r in enumerate(payload['slots']) if r['slot'].endswith('#' + slot)), None)
            if sid is None:
                print('no slot', slot); continue
            dropped = self.version('p6a-drop', saved(doc), drop=[f['label'] for f in version.fields if 'Office' in f['label']])
            compiled2, payload2 = self.compile(dropped)
            try:
                change = self.addition(payload2, slot, 'text', label)
            except AssertionError as exc:
                print('cannot quote', exc); continue
            result = mapping.validate_proposal(dropped, compiled2, self.respond(payload2, [change]))
            adopted = self.show('p6a ' + slot, dropped, compiled2, payload2, result, change)
            if adopted:
                try:
                    print('FILLED:', self.fill_text(dropped, result['fields'], {adopted[0]['key']: 'SYNTHETIC-ANSWER'}))
                except Exception as exc:
                    print('FILL REFUSED:', exc)

    # P6b: a block-level control with sdtContentLocked around a printed answer line.
    def test_p6b_locked_block_control(self):
        doc = Document(); doc.add_paragraph('Plan: ______________')
        append_body(doc, f'<w:sdt xmlns:w="{W}"><w:sdtPr><w:alias w:val="Locked statement"/><w:lock w:val="sdtContentLocked"/>'
                         '</w:sdtPr><w:sdtContent><w:p><w:r><w:t>Consultant statement: ______________</w:t></w:r></w:p>'
                         '</w:sdtContent></w:sdt>')
        version = self.version('p6b', saved(doc))
        compiled, payload = self.compile(version)
        print('\n--- p6b slots:', [(r['slot'], r['text'], r['heading']) for r in payload['slots']])
        print('engine fields:', [(f['label'], f['kind'], f['locator']['slot'], f['locator']['target'], f.get('write_protected')) for f in version.fields])
        dropped = self.version('p6b-drop', saved(doc), drop=[f['label'] for f in version.fields if 'Consultant' in f['label']])
        compiled2, payload2 = self.compile(dropped)
        for row in payload2['slots']:
            if 'sdt' in row['slot']:
                slot = row['slot'].split('#', 1)[1]
                try:
                    change = self.addition(payload2, slot, 'text', 'Consultant statement')
                except AssertionError as exc:
                    print('cannot quote', exc); continue
                result = mapping.validate_proposal(dropped, compiled2, self.respond(payload2, [change]))
                adopted = self.show('p6b ' + slot, dropped, compiled2, payload2, result, change)
                if adopted:
                    try:
                        print('FILLED:', self.fill_text(dropped, result['fields'], {adopted[0]['key']: 'SYNTHETIC-ANSWER'}))
                    except Exception as exc:
                        print('FILL REFUSED:', exc)


import zipfile, re as _re


class Probe3(Probe):
    def fill_xml(self, version, fields, answers):
        out = private_path('outputs/x%s.docx' % id(fields)); out.parent.mkdir(parents=True, exist_ok=True)
        fill_template(private_path(version.template.original_path), fields, answers, out)
        return zipfile.ZipFile(io.BytesIO(out.read_bytes())).read('word/document.xml').decode()

    def test_p6b_xml(self):
        doc = Document(); doc.add_paragraph('Plan: ______________')
        append_body(doc, f'<w:sdt xmlns:w="{W}"><w:sdtPr><w:alias w:val="Locked statement"/><w:lock w:val="sdtContentLocked"/>'
                         '</w:sdtPr><w:sdtContent><w:p><w:r><w:t>Consultant statement: ______________</w:t></w:r></w:p>'
                         '</w:sdtContent></w:sdt>')
        version = self.version('p6bx', saved(doc))
        print('\nengine fields p6bx:', [f['label'] for f in version.fields])
        compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/sdt[0]/p[0]', 'text', 'Consultant statement')
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        adopted = self.adopted(result)
        print('p6bx adopted:', [(f['label'], f['locator']['slot'], f['locator']['target']) for f in adopted],
              'unresolved:', result['review']['unresolved'])
        xml = self.fill_xml(version, result['fields'], {adopted[0]['key']: 'SYNTHETIC-ANSWER'})
        m = _re.search(r'<w:sdt>.*?</w:sdt>', xml, _re.S)
        print('LOCKED SDT AFTER FILL:', m.group(0) if m else None)

    def test_p7_consent_box_with_picture(self):
        doc = Document(); doc.add_paragraph('Participant ID: __________')
        doc.add_paragraph('Procedure consent', style='Heading 2')
        doc.add_paragraph('Record that the patient consents to proceed:')
        box = bordered_box(doc); box.cell(0, 0).paragraphs[0]._p.append(pict_run())
        version, result, adopted = self.run_addition('p7', saved(doc), 'b/tbl[0]/tr[0]/tc[0]/p[0]', 'long_text',
                                                     'Procedure consent', roles=(1,))
        # Control: the SAME box without the picture is held (F1).
        doc2 = Document(); doc2.add_paragraph('Participant ID: __________')
        doc2.add_paragraph('Procedure consent', style='Heading 2')
        doc2.add_paragraph('Record that the patient consents to proceed:')
        bordered_box(doc2)
        self.run_addition('p7-control', saved(doc2), 'b/tbl[0]/tr[0]/tc[0]/p[0]', 'long_text', 'Procedure consent', roles=(1,))

    def test_p1_control_without_picture(self):
        doc = Document(); doc.add_paragraph('Participant ID: __________')
        doc.add_paragraph('Consultant review', style='Heading 2'); doc.add_paragraph('Signed by the consultant:')
        bordered_box(doc)
        self.run_addition('p1-control', saved(doc), 'b/tbl[0]/tr[0]/tc[0]/p[0]', 'long_text', 'Consultant review', roles=(1,))

    def test_p3_control_without_picture(self):
        doc = Document(); doc.add_paragraph('Clinic record', style='Heading 2')
        table = doc.add_table(rows=2, cols=2); table.cell(0, 0).text = 'Name'; table.cell(0, 1).text = 'Signature'
        self.run_addition('p3-control', saved(doc), 'b/tbl[0]/tr[1]/tc[1]/p[0]', 'text', 'Clinic record', roles=(1,),
                          drop=('Signature',))

    def test_p8_countersigned_vocabulary(self):
        doc = Document(); doc.add_paragraph('Plan: ______________'); doc.add_paragraph('Countersigned by: ______________')
        self.run_addition('p8', saved(doc), 'b/p[1]', 'text', 'Countersigned by', drop=('Countersigned by',))
