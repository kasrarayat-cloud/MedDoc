"""F1-export-text: export_checks on the product's own legacy fill, with only the Findings answer paragraph replaced."""
import io, os, re, sys, tempfile, zipfile, shutil
from pathlib import Path
TREE = Path(sys.argv[1])
sys.path.insert(0, str(TREE / 'rebuild'))
from ops import sonnet5_qa as sq
from documents import fill_template, inspect_template
from docx import Document

work = Path(tempfile.mkdtemp(prefix='rr-export-', dir=os.environ.get('SCRATCH')))
form = sq.build_form()
original = work / 'form.docx'; original.write_bytes(form)
fields = inspect_template(original)['fields']
print('fields', [(f['key'], f['label']) for f in fields])
answers = {f['key']: sq.EXPECTED[f['label']] for f in fields}
out = work / 'out.docx'
fill_template(original, fields, answers, out)
export = out.read_bytes()
print('baseline export_checks:', sq.export_checks(form, export, sq.EXPECTED))

z = zipfile.ZipFile(io.BytesIO(export))
xml = z.read('word/document.xml').decode()
root_tag = re.search(r'<w:document[^>]*>', xml).group(0)
print('mc declared:', 'xmlns:mc=' in root_tag, '| w14 declared:', 'xmlns:w14=' in root_tag, '| Ignorable:', re.search(r'mc:Ignorable="[^"]*"', root_tag).group(0) if 'mc:Ignorable' in root_tag else None)
# the Findings answer paragraph, as the fill wrote it
rows = re.findall(r'<w:tr\b.*?</w:tr>', xml, re.S)
findings_row = [r for r in rows if '>Findings<' in r][0]
cells = re.findall(r'<w:tc>.*?</w:tc>|<w:tc\b.*?</w:tc>', findings_row, re.S)
answer_cell = cells[1]
paras = re.findall(r'<w:p\b.*?</w:p>|<w:p/>', answer_cell, re.S)
print('findings answer cell paragraphs:', len(paras)); print('as filled:', paras[0][:400])
assert len(paras) == 1

VARIANTS = {
 'A_alternate_content': '<w:p><w:r><mc:AlternateContent><mc:Choice Requires="w14"></mc:Choice><mc:Fallback><w:t xml:space="preserve">No </w:t></mc:Fallback></mc:AlternateContent></w:r><w:r><w:t>deformity documented.</w:t></w:r></w:p>',
 'B_tracked_deletion': '<w:p><w:del w:id="91" w:author="x" w:date="2026-09-29T00:00:00Z"><w:r><w:t xml:space="preserve">No </w:t></w:r></w:del><w:r><w:t>deformity documented.</w:t></w:r></w:p>',
}
for name, para in VARIANTS.items():
    new_cell = answer_cell.replace(paras[0], para)
    assert answer_cell.count(paras[0]) == 1 and new_cell != answer_cell
    new_xml = xml.replace(answer_cell, new_cell); assert xml.count(answer_cell) == 1
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as w:
        for info in z.infolist():
            w.writestr(info, new_xml.encode() if info.filename == 'word/document.xml' else z.read(info.filename))
    variant = buf.getvalue()
    vz = zipfile.ZipFile(io.BytesIO(variant))
    same = all(vz.read(n) == z.read(n) for n in z.namelist() if n != 'word/document.xml') and sorted(vz.namelist()) == sorted(z.namelist())
    problems = sq.export_checks(form, variant, sq.EXPECTED)
    doc = Document(io.BytesIO(variant))
    shown = doc.tables[0].rows[1].cells[1].text
    from lxml import etree
    cell_el = doc.tables[0].rows[1].cells[1]._tc
    print(f'{name}: other parts byte-identical={same} | export_checks -> {problems} | _cell_text={sq._cell_text(cell_el)!r} | python-docx cell.text={shown!r}')
shutil.rmtree(work)
