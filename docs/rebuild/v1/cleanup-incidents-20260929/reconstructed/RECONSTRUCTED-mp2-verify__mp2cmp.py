"""Same U1/U2 inputs run against whichever tree is on sys.path (HEAD 48aa5e8d or d40c8689). Synthetic blank forms only.
No provider construction; non-loopback sockets refused."""
import copy
import io
import json
import socket
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

_real_connect = socket.socket.connect


def _guard(self, address):
    host = address[0] if isinstance(address, tuple) else address
    if isinstance(host, str) and host not in ('127.0.0.1', '::1', 'localhost') and not host.startswith('/'):
        raise AssertionError(f'non-loopback socket refused: {host!r}')
    return _real_connect(self, address)


socket.socket.connect = _guard

from django.test import SimpleTestCase, override_settings
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from core import generation, mapping_assistance as mapping, mapping_wire as wire, v3_mapping
from core.security import private_path
from core.tests.test_mapping_assistance import SETTINGS
from documents.service import inspect_template, fill_template, DocumentError

BODY = 'word/document.xml#'


def saved(doc):
    out = io.BytesIO(); doc.save(out); return out.getvalue()


def bordered_box(doc):
    table = doc.add_table(rows=1, cols=1)
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); borders.append(edge)
    table.cell(0, 0)._tc.get_or_add_tcPr().append(borders)
    return table


def box_form(heading, question):
    doc = Document(); doc.add_paragraph('Participant ID: __________')
    doc.add_paragraph(heading, style='Heading 2'); doc.add_paragraph(question); bordered_box(doc)
    return saved(doc)


def lines_form(second):
    doc = Document(); doc.add_paragraph('Plan: ______________'); doc.add_paragraph(second)
    return saved(doc)


@override_settings(**SETTINGS)
class Cmp(SimpleTestCase):
    def setUp(self):
        deny = patch('anthropic.Anthropic', side_effect=AssertionError('Unmocked provider construction is forbidden.'))
        deny.start(); self.addCleanup(deny.stop)
        self.directory = tempfile.TemporaryDirectory(prefix='mp2-cmp-'); self.addCleanup(self.directory.cleanup)
        override = override_settings(MEDIA_ROOT=Path(self.directory.name)); override.enable(); self.addCleanup(override.disable)
        self.config = generation.configuration()

    def version(self, name, data, drop=()):
        path = private_path(f'originals/{name}.docx'); path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(data)
        info = inspect_template(path, engine='v3')
        fields = [f for f in copy.deepcopy(info['fields']) if f['label'] not in drop]
        return SimpleNamespace(template=SimpleNamespace(format='docx', original_path=f'originals/{name}.docx', sha256=info['sha256']),
                               contract='v3', document_backend={}, fields=fields), info['fields']

    def ref(self, payload, sid, text, roles):
        for position, (owner, role, index) in enumerate(payload['text_refs']):
            if owner != sid or role not in roles:
                continue
            source = (payload['slots'][sid]['options'][index] if role == 3 else
                      payload['slots'][sid][('text', 'heading', 'row_label')[role]])
            if text in source:
                return [position, source.index(text), source.index(text) + len(text)]
        raise AssertionError((sid, text, roles))

    def run_one(self, name, data, slot, kind, label, label_roles, options=(), drop=(), replace_label=None):
        version, engine = self.version(name, data, drop)
        compiled = mapping.compile_request(version, self.config)
        payload = json.loads(compiled['request']['messages'][0]['content'])
        sid = next(n for n, row in enumerate(payload['slots']) if row['slot'] == BODY + slot)
        replace = []
        if replace_label:
            fid = next(n for n, f in enumerate(payload['fields']) if f['label'] == replace_label)
            replace = [fid]
            change = [fid, sid, wire.KINDS.index(kind), payload['text_refs'].index([fid, 4, 0]), 0, len(label), -1]
        else:
            change = [-1, sid, wire.KINDS.index(kind), *self.ref(payload, sid, label, label_roles), -1]
        for option in options:
            change.extend(self.ref(payload, sid, option, (3, 0)))
        mapped = {s for f in payload['fields'] for s in f['slots']}
        body = {'f': {'k': [n for n in range(len(payload['fields'])) if n not in replace], 'r': replace, 'n': [], 'u': []},
                's': {'c': [n for n, row in enumerate(payload['slots']) if v3_mapping.is_mapped(row['slot'], mapped)],
                      'n': [n for n, row in enumerate(payload['slots']) if not v3_mapping.is_mapped(row['slot'], mapped)], 'u': []},
                'x': [change], 'e': False}
        result = mapping.validate_proposal(version, compiled, body)
        adopted = [f for f in result['fields'] if f.get('origin') == 'model_assisted']
        out = {'engine_fields': [(f['label'], f['kind'], f.get('options'), sorted(v3_mapping.mapped_slots([f]))) for f in engine],
               'adopted': [(f['label'], f['kind'], f.get('options'), f['locator']['slot'], f['locator']['target']) for f in adopted],
               'rejected': [r['reason'] for r in result['review']['rejected_changes']],
               'unresolved': [r['reason'] for r in result['review']['unresolved']]}
        if adopted:
            answers = {f['key']: (f['options'][0] if f['kind'] == 'choice' else 'SYNTHETIC-ANSWER') for f in adopted}
            try:
                fill_template(private_path(version.template.original_path), result['fields'], answers,
                              Path(self.directory.name) / f'{name}.docx')
                out['fill'] = 'ok'
            except DocumentError as exc:
                out['fill'] = f'refused: {exc} ({exc.code})'
        print('\nCMP ' + name + ' ' + json.dumps(out, default=str))

    def test_all(self):
        self.run_one('U1-circle', lines_form('Visit type (circle one): Initial / Follow-up'), 'b/p[1]', 'choice',
                     'Visit type', (0,), options=('Initial', 'Follow-up'))
        self.run_one('U1-circle-text', lines_form('Visit type (circle one): Initial / Follow-up'), 'b/p[1]', 'text',
                     'Visit type', (0,))
        self.run_one('U1-circle-control-New', lines_form('Visit type (circle one): New / Follow-up'), 'b/p[1]', 'choice',
                     'Visit type', (0,), options=('New', 'Follow-up'))
        self.run_one('U1-checkbox-choice', lines_form('Visit type: ☐ Initial ☐ Follow-up'), 'b/p[1]', 'choice',
                     'Visit type', (0,), options=('Initial', 'Follow-up'), drop=('Visit type',))
        self.run_one('U1-checkbox-text', lines_form('Visit type: ☐ Initial ☐ Follow-up'), 'b/p[1]', 'text',
                     'Visit type', (0,), drop=('Visit type',))
        self.run_one('U1-checkbox-choice-glyphs', lines_form('Visit type: ☐ Initial ☐ Follow-up'), 'b/p[1]/glyphs', 'choice',
                     'Visit type', (0,), options=('Initial', 'Follow-up'), drop=('Visit type',))
        self.run_one('U1-checkbox-control-New', lines_form('Visit type: ☐ New ☐ Follow-up'), 'b/p[1]', 'choice',
                     'Visit type', (0,), options=('New', 'Follow-up'), drop=('Visit type',))
        self.run_one('U1-nominal-replacement', lines_form('Is this the initial presentation? ☐ Yes ☐ No'),
                     'b/p[1]/glyphs', 'choice', 'Is this the initial presentation?', (0,), options=('Yes', 'No'),
                     replace_label='Is this the initial presentation?')
        self.run_one('U1-initial-assessment-text', lines_form('Initial assessment: ______________'), 'b/p[1]', 'text',
                     'Initial assessment', (0,), drop=('Initial assessment',))
        slot = 'b/tbl[0]/tr[0]/tc[0]'
        self.run_one('U2-initial-question', box_form('Clinical history', 'Describe the initial presentation?'), slot,
                     'long_text', 'Clinical history', (1,), drop=('Describe the initial presentation?',))
        self.run_one('U2-question-control', box_form('Clinical history', 'Describe the presentation?'), slot,
                     'long_text', 'Clinical history', (1,), drop=('Describe the presentation?',))
        self.run_one('U2-initial-heading', box_form('Initial Assessment', 'Please describe the presenting complaint?'), slot,
                     'long_text', 'Initial Assessment', (1,), drop=('Please describe the presenting complaint?',))
        self.run_one('U2-heading-control', box_form('Wound review', 'Please describe the presenting complaint?'), slot,
                     'long_text', 'Wound review', (1,), drop=('Please describe the presenting complaint?',))
        # Without drop: does the ENGINE make a field for these boxes itself?
        self.run_one('U2-initial-question-nodrop', box_form('Clinical history', 'Describe the initial presentation?'), slot,
                     'long_text', 'Clinical history', (1,))
        self.run_one('U2-initial-heading-nodrop', box_form('Initial Assessment', 'Please describe the presenting complaint?'), slot,
                     'long_text', 'Initial Assessment', (1,))
