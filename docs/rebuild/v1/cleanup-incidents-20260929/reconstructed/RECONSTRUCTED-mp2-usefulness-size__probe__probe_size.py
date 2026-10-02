import copy, io, json, os, tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, override_settings
from docx import Document
from core import generation, mapping_assistance as mapping
from core.security import private_path
from core.tests.test_mapping_assistance import SETTINGS
from documents.service import inspect_template
OUT = os.environ.get('PROBE_OUT')

def big_form(count, length=130, table_rows=0):
    doc = Document(); doc.add_paragraph('Synthetic blank intake form')
    words = ('please read each instruction carefully before completing the section below and return the form to the '
             'front desk when finished ').split()
    for k in range(count):
        text = f'Instruction {k}: '; i = k
        while len(text) < length:
            text += words[i % len(words)] + ' '; i += 1
        doc.add_paragraph(text.strip())
    if table_rows:
        t = doc.add_table(rows=table_rows, cols=2)
        for r in range(table_rows): t.cell(r, 0).text = f'Item {r}'
    for label in ('Reason for visit: ____________________', 'Comments: ____________________', 'Plan: ____________________'):
        doc.add_paragraph(label)
    out = io.BytesIO(); doc.save(out); return out.getvalue()

@override_settings(**SETTINGS)
class Size(SimpleTestCase):
    def setUp(self):
        deny = patch('anthropic.Anthropic', side_effect=AssertionError('no provider')); deny.start(); self.addCleanup(deny.stop)
        self.directory = tempfile.TemporaryDirectory(prefix='probe-size-'); self.addCleanup(self.directory.cleanup)
        o = override_settings(MEDIA_ROOT=Path(self.directory.name)); o.enable(); self.addCleanup(o.disable)
        self.config = generation.configuration()

    def test_sizes(self):
        rows = []
        for count, length, table_rows in [(750,130,0),(770,130,0),(772,130,0),(774,130,0),(776,130,0),(778,130,0),(780,130,0),(600,130,150),(500,130,300),(300,130,500)]:
            data = big_form(count, length, table_rows)
            name = f'b{count}-{length}-{table_rows}'
            path = private_path(f'originals/{name}.docx'); path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(data)
            info = inspect_template(path, engine='v3')
            version = SimpleNamespace(template=SimpleNamespace(format='docx', original_path=f'originals/{name}.docx', sha256=info['sha256']),
                                      contract='v3', document_backend={}, fields=info['fields'])
            try:
                c = mapping.compile_request(version, self.config)
                rows.append({'form': name, 'size': len(mapping.canonical(c['request']).encode()), 'presentation': c.get('presentation'), 'bound': c['input_token_bound']})
            except ValidationError as e:
                rows.append({'form': name, 'refused': ' '.join(e.messages)})
        with open(OUT, 'w') as fh: json.dump(rows, fh, indent=0)
