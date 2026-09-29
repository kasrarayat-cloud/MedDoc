"""Refutation probe (scratch only): ordinary-question model changes, adopted or held, per tree."""
import copy, io, json, os, tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, override_settings
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from core import generation, mapping_assistance as mapping, mapping_wire as wire, v3_mapping
from core.security import private_path
from core.tests.test_mapping_assistance import SETTINGS
from documents.service import inspect_template

BODY = 'word/document.xml#'
OUT = os.environ.get('PROBE_OUT')


def saved(doc):
    out = io.BytesIO(); doc.save(out); return out.getvalue()


def bordered_box(doc):
    table = doc.add_table(rows=1, cols=1)
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); borders.append(edge)
    table.cell(0, 0)._tc.get_or_add_tcPr().append(borders)
    return table


@override_settings(**SETTINGS)
class Probe(SimpleTestCase):
    def setUp(self):
        deny = patch('anthropic.Anthropic', side_effect=AssertionError('no provider'))
        deny.start(); self.addCleanup(deny.stop)
        self.directory = tempfile.TemporaryDirectory(prefix='probe-')
        self.addCleanup(self.directory.cleanup)
        override = override_settings(MEDIA_ROOT=Path(self.directory.name)); override.enable(); self.addCleanup(override.disable)
        self.config = generation.configuration()
        self.results = []

    def version(self, name, data, drop=()):
        path = private_path(f'originals/{name}.docx'); path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(data)
        info = inspect_template(path, engine='v3')
        fields = [f for f in copy.deepcopy(info['fields']) if f['label'] not in drop]
        return SimpleNamespace(template=SimpleNamespace(format='docx', original_path=f'originals/{name}.docx', sha256=info['sha256']),
                               contract='v3', document_backend={}, fields=fields), info

    def compile(self, version):
        compiled = mapping.compile_request(version, self.config)
        return compiled, json.loads(compiled['request']['messages'][0]['content'])

    @staticmethod
    def sid(payload, slot):
        return next(n for n, row in enumerate(payload['slots']) if row['slot'] == BODY + slot or row['slot'] == slot)

    @staticmethod
    def quote(payload, sid, text, roles=(0, 1, 2)):
        for position, (owner, role, _) in enumerate(payload['text_refs']):
            if owner == sid and role in roles:
                source = payload['slots'][sid][('text', 'heading', 'row_label')[role]]
                if text in source:
                    return [position, source.index(text), source.index(text) + len(text)]
        raise AssertionError(f'{text!r} not at slot {sid}: {payload["slots"][sid]}')

    def respond(self, payload, changes, replace=()):
        mapped = {slot for f in payload['fields'] for slot in f['slots']}
        return {'f': {'k': [n for n in range(len(payload['fields'])) if n not in replace], 'r': list(replace), 'n': [], 'u': []},
                's': {'c': [n for n, row in enumerate(payload['slots']) if v3_mapping.is_mapped(row['slot'], mapped)],
                      'n': [n for n, row in enumerate(payload['slots']) if not v3_mapping.is_mapped(row['slot'], mapped)], 'u': []},
                'x': [list(c) for c in changes], 'e': False}

    def change(self, payload, fid, slot, kind, label, roles=(0, 1, 2), options=(), option_roles=(0, 1, 2)):
        sid = self.sid(payload, slot)
        if label is None:
            ref = [payload['text_refs'].index([fid, 4, 0]), 0, len(payload['fields'][fid]['label'])]
        else:
            ref = self.quote(payload, sid, label, roles)
        extra = []
        for option in options:
            extra += self.quote(payload, sid, option, option_roles)
        return [fid, sid, wire.KINDS.index(kind), *ref, -1, *extra]

    def run_case(self, name, version, compiled, payload, change, replace=()):
        try:
            result = mapping.validate_proposal(version, compiled, self.respond(payload, [change], replace=replace))
            adopted = [(f['label'], f['kind'], f.get('options')) for f in result['fields'] if f.get('origin') == 'model_assisted']
            reasons = [r['reason'] for r in result['review']['rejected_changes']]
            outcome = {'case': name, 'adopted': adopted, 'held': reasons}
        except ValidationError as exc:
            outcome = {'case': name, 'refused': ' '.join(exc.messages)}
        self.results.append(outcome)
        return outcome

    def tearDown(self):
        with open(OUT, 'a') as fh:
            for row in self.results:
                fh.write(json.dumps(row) + '\n')

    # --- cases ---------------------------------------------------------------------------------------------
    def test_choice_with_initial_option(self):
        doc = Document(); doc.add_paragraph('Plan: ______________')
        doc.add_paragraph('Visit type: ☐ Initial ☐ Follow-up')
        data = saved(doc)
        version, info = self.version('visit', data)
        self.results.append({'case': 'visit/engine_fields', 'fields': [(f['label'], f['kind'], f.get('options'), f['locator']['slot'], f['locator']['target'].get('type')) for f in info['fields']]})
        compiled, payload = self.compile(version)
        self.results.append({'case': 'visit/slots', 'slots': [(r['slot'], r['text']) for r in payload['slots']]})
        # Replacement of the engine's own choice field, keeping its label, at its own paragraph.
        prior = next((f for f in version.fields if f['label'].startswith('Visit type')), None)
        if prior is not None:
            fid = next(n for n, f in enumerate(payload['fields']) if f['key'] == prior['key'])
            for kind, options in (('choice', ('Initial', 'Follow-up')),):
                ch = self.change(payload, fid, 'b/p[1]', kind, None, options=options)
                self.run_case(f'visit/replace-{kind}', version, compiled, payload, ch, replace=[fid])
        # The engine's own field missing: an addition of the same question.
        version2, _ = self.version('visit-miss', data, drop=(prior['label'],) if prior else ())
        compiled2, payload2 = self.compile(version2)
        ch = self.change(payload2, -1, 'b/p[1]', 'choice', 'Visit type', options=('Initial', 'Follow-up'))
        self.run_case('visit/add-choice', version2, compiled2, payload2, ch)
        ch = self.change(payload2, -1, 'b/p[1]', 'text', 'Visit type')
        self.run_case('visit/add-text', version2, compiled2, payload2, ch)

    def test_yes_no_initial_presentation(self):
        doc = Document(); doc.add_paragraph('Plan: ______________')
        doc.add_paragraph('Is this the initial presentation? ☐ Yes ☐ No')
        data = saved(doc)
        version, info = self.version('initpres', data)
        self.results.append({'case': 'initpres/engine_fields', 'fields': [(f['label'], f['kind'], f.get('options'), f['locator']['slot']) for f in info['fields']]})
        compiled, payload = self.compile(version)
        prior = next((f for f in version.fields if 'initial' in f['label'].lower()), None)
        if prior is not None:
            fid = next(n for n, f in enumerate(payload['fields']) if f['key'] == prior['key'])
            ch = self.change(payload, fid, 'b/p[1]', prior['kind'] if prior['kind'] in wire.KINDS else 'choice', None,
                             options=tuple(prior.get('options') or ()))
            self.run_case('initpres/replace-own', version, compiled, payload, ch, replace=[fid])

    def test_question_box_under_initial_heading(self):
        for heading in ('Initial Assessment', 'Wound review'):
            doc = Document(); doc.add_paragraph('Participant ID: __________')
            doc.add_paragraph(heading, style='Heading 2'); doc.add_paragraph('Please describe the presenting complaint?')
            bordered_box(doc)
            data = saved(doc)
            version, info = self.version('box-' + heading.split()[0], data, drop=('Please describe the presenting complaint?',))
            self.results.append({'case': f'box[{heading}]/engine_fields', 'fields': [(f['label'], f['kind'], f['locator']['slot']) for f in info['fields']]})
            compiled, payload = self.compile(version)
            ch = self.change(payload, -1, 'b/tbl[0]/tr[0]/tc[0]', 'long_text', heading, roles=(1,))
            self.run_case(f'box[{heading}]/add-long_text', version, compiled, payload, ch)

    def test_paragraph_with_date_and_initials_parts(self):
        doc = Document(); doc.add_paragraph('Plan: ______________')
        doc.add_paragraph('Pain score: ______\tDate: ______\tNurse initials: ______')
        data = saved(doc)
        version, info = self.version('vitals', data)
        self.results.append({'case': 'vitals/engine_fields', 'fields': [(f['label'], f['kind'], f['locator']['slot'], f['locator']['target']) for f in info['fields']]})


class Natural(Probe):
    """Natural engine misses (no drop=): the engine makes no field, the model adds one."""

    def test_circle_one_choice_line(self):
        for text, options in (('Visit type (circle one): Initial / Follow-up', ('Initial', 'Follow-up')),
                              ('Visit type (circle one): New / Follow-up', ('New', 'Follow-up'))):
            doc = Document(); doc.add_paragraph('Plan: ______________'); doc.add_paragraph(text)
            data = saved(doc)
            version, info = self.version('circle-' + options[0], data)
            self.results.append({'case': f'circle[{text}]/engine_fields', 'fields': [(f['label'], f['kind'], f.get('options')) for f in info['fields']]})
            compiled, payload = self.compile(version)
            ch = self.change(payload, -1, 'b/p[1]', 'choice', 'Visit type', options=options)
            out = self.run_case(f'circle[{text}]/add-choice', version, compiled, payload, ch)

    def test_question_box_asking_about_initial(self):
        for question in ('Describe the initial presentation?', 'Describe the presentation?'):
            doc = Document(); doc.add_paragraph('Participant ID: __________')
            doc.add_paragraph('Clinical history', style='Heading 2'); doc.add_paragraph(question)
            bordered_box(doc)
            data = saved(doc)
            version, info = self.version('qbox-' + str(len(question)), data)
            self.results.append({'case': f'qbox[{question}]/engine_fields', 'fields': [(f['label'], f['kind']) for f in info['fields']], 'notes': [n for n in info.get('notes', []) if 'Unassigned' in n]})
            if any(f['label'] == question for f in version.fields):
                version.fields = [f for f in version.fields if f['label'] != question]
            compiled, payload = self.compile(version)
            ch = self.change(payload, -1, 'b/tbl[0]/tr[0]/tc[0]', 'long_text', 'Clinical history', roles=(1,))
            self.run_case(f'qbox[{question}]/add-long_text', version, compiled, payload, ch)

    def test_replacement_target_types(self):
        doc = Document(); doc.add_paragraph('Plan: ______________')
        doc.add_paragraph('Visit type: ☐ Initial ☐ Follow-up')
        data = saved(doc)
        version, info = self.version('visit-r', data)
        compiled, payload = self.compile(version)
        prior = next(f for f in version.fields if f['label'].startswith('Visit type'))
        fid = next(n for n, f in enumerate(payload['fields']) if f['key'] == prior['key'])
        ch = self.change(payload, fid, 'b/p[1]', 'choice', None, options=('Initial', 'Follow-up'))
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [ch], replace=[fid]))
        self.results.append({'case': 'visit-r/adopted-locators', 'fields': [(f['label'], f['kind'], f['locator']['slot'], f['locator']['target']) for f in result['fields'] if f.get('origin') == 'model_assisted'], 'held': [r['reason'] for r in result['review']['rejected_changes']]})


class NaturalText(Probe):
    def test_text_question_mentioning_initial(self):
        for text, label in (('Presenting complaint at initial visit:', 'Presenting complaint at initial visit'),
                            ('Presenting complaint at first visit:', 'Presenting complaint at first visit')):
            doc = Document(); doc.add_paragraph('Plan: ______________'); doc.add_paragraph(text)
            data = saved(doc)
            version, info = self.version('txt-' + str(len(text)), data)
            self.results.append({'case': f'text[{text}]/engine_fields', 'fields': [(f['label'], f['kind']) for f in info['fields']]})
            compiled, payload = self.compile(version)
            ch = self.change(payload, -1, 'b/p[1]', 'long_text', label)
            self.run_case(f'text[{text}]/add-long_text', version, compiled, payload, ch)
