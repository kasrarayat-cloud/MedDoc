"""Scratch-only independent reproduction of RR2-M1 (mapping adoption COMMIT answer lost after it committed)."""
import io
import json
import tempfile
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.test import TransactionTestCase, override_settings
from django.utils import timezone
from docx import Document

from core import generation, jobs, mapping_assistance as mapping, mapping_jobs
from core.billing import welcome_credit
from core.jobs import deliverable
from core.management.commands.audit_billing import audit_report
from core.models import (AuditEvent, Department, Job, LedgerEntry, Membership, Organization, PlatformCost,
                         ProviderReceipt, Reservation, Template, TemplateVersion, Usage, User, Wallet)
from core.security import private_path
from core.tests.test_mapping_assistance import SETTINGS
from core.tests.test_settlement_outcome_unknown import CommitAnswerLost
from documents.service import inspect_template


def report(tag, **values):
    print(f'\n[RR2-M1 {tag}] ' + json.dumps({k: str(v) for k, v in values.items()}, sort_keys=True), flush=True)


@override_settings(**SETTINGS)
class MappingAdoptionCommitAnswerLost(TransactionTestCase):

    def setUp(self):
        deny = patch('anthropic.Anthropic', side_effect=AssertionError('unmocked provider construction'))
        deny.start(); self.addCleanup(deny.stop)
        self.directory = tempfile.TemporaryDirectory(prefix='rr2-m1-')
        self.addCleanup(self.directory.cleanup)
        override = override_settings(MEDIA_ROOT=Path(self.directory.name)); override.enable(); self.addCleanup(override.disable)
        self.user = User.objects.create_user(username='rr2m1', email='rr2m1@synthetic.invalid', verified_at=timezone.now())
        org = Organization.objects.create(name='Synthetic rr2 m1')
        self.department = Department.objects.create(organization=org, name='Synthetic rr2 m1')
        Membership.objects.create(user=self.user, department=self.department, role='clinician')
        welcome_credit(self.user)
        doc = Document(); table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Plan'; table.cell(1, 0).text = 'History'
        doc.add_paragraph('Instructions: ' + 'Safe blank source wording. ' * 20)
        doc.add_paragraph('Signature: ____________')
        buffer = io.BytesIO(); doc.save(buffer)
        path = private_path('originals/mapping.docx'); path.parent.mkdir(parents=True); path.write_bytes(buffer.getvalue())
        info = inspect_template(path, engine='v3')
        self.template = Template.objects.create(department=self.department, name='Synthetic rr2 m1', format='docx',
            original_name='mapping.docx', original_path='originals/mapping.docx', sha256=info['sha256'], created_by=self.user)
        self.version = TemplateVersion.objects.create(template=self.template, number=1, fields=info['fields'],
                                                      contract='v3', created_by=self.user)

    def response(self, body):
        data = {'model': 'claude-sonnet-5', 'stop_reason': 'end_turn', 'content': [{'type': 'text', 'text': json.dumps(body)}]}
        return SimpleNamespace(model='claude-sonnet-5', stop_reason='end_turn', content=[SimpleNamespace(**data['content'][0])],
            id='msg_synthetic_rr2', _request_id='req_synthetic_rr2',
            usage=SimpleNamespace(input_tokens=100, output_tokens=30, cache_read_input_tokens=0, cache_creation_input_tokens=0),
            model_dump=lambda **kwargs: data)

    def run_live_mapping(self, landed):
        config = generation.configuration()
        compiled = mapping.compile_request(self.version, config)
        body = mapping_jobs.send(compiled, {**config, 'mode': 'fixture'})[0]
        job = mapping_jobs.enqueue(self.version, self.user, config, compiled['request_sha256'])
        key_before = job.request_key
        wallets_before = sorted((w.kind, w.balance, w.reserved) for w in Wallet.objects.filter(user=self.user))
        harness = CommitAnswerLost(landed)
        real_validate = mapping.validate_proposal

        def validate(*args, **kwargs):
            result = real_validate(*args, **kwargs)
            harness.arm()  # the next outermost COMMIT is the adoption transaction (mapping_jobs.py:132-143)
            return result
        with harness.patched(), patch('anthropic.Anthropic') as client, \
                patch('core.mapping_assistance.validate_proposal', side_effect=validate):
            client.return_value.messages.create.return_value = self.response(body)
            self.assertTrue(jobs.run_one())
            harness.join(self)
        job.refresh_from_db()
        return job, harness, key_before, wallets_before, client

    def observe(self, job):
        versions = list(TemplateVersion.objects.filter(template=self.template).order_by('number'))
        usage = Usage.objects.filter(job=job).first()
        return dict(
            state=job.state, error=job.error, request_key=job.request_key,
            versions=[(v.number, (v.assistance or {}).get('job_id')) for v in versions],
            review_ready=AuditEvent.objects.filter(action='mapping.review_ready', target=str(job.pk)).count(),
            mapping_failed=AuditEvent.objects.filter(action='mapping.failed', target=str(job.pk)).count(),
            usage_absorbed=AuditEvent.objects.filter(action='usage.absorbed', target=str(job.pk)).count(),
            error_after_commit=AuditEvent.objects.filter(action='job.error_after_commit', target=str(job.pk)).count(),
            charged=usage.charged_cost if usage else None, actual=usage.actual_cost if usage else None,
            usage_ledger=sum(LedgerEntry.objects.filter(job=job, kind='usage').values_list('amount', flat=True)),
            absorbed_ledger=sum(LedgerEntry.objects.filter(job=job, kind='absorbed').values_list('amount', flat=True)),
            platform_absorbed=list(PlatformCost.objects.filter(job=job, kind='absorbed_failure').values_list('amount', flat=True)),
            receipt=ProviderReceipt.objects.filter(job=job).count(),
            holds=sorted(Reservation.objects.filter(job=job).values_list('state', flat=True)),
            deliverable=deliverable(job), audit_all_checks_pass=audit_report()['all_checks_pass'])

    @override_settings(AI_MODE='live', QUILLMD_MAPPING_ASSISTANCE_LIVE=True)
    def test_landed(self):
        job, harness, key_before, wallets_before, client = self.run_live_mapping(landed=True)
        obs = self.observe(job)
        report('landed', fired=harness.fired, waiter_seen=harness.waiter_seen, key_before=key_before,
               provider_calls=client.return_value.messages.create.call_count,
               wallets_before=wallets_before,
               wallets_after=sorted((w.kind, w.balance, w.reserved) for w in Wallet.objects.filter(user=self.user)), **obs)
        self.assertTrue(harness.fired)
        self.assertTrue(harness.waiter_seen)
        # The adoption COMMITTED: v2 exists, bound to this job, and review_ready was audited.
        self.assertEqual(obs['versions'], [(1, None), (2, str(job.pk))])
        self.assertEqual(obs['review_ready'], 1)
        # The defect as claimed: the committed success is overwritten as a platform failure and the charge reversed.
        self.assertEqual(job.state, 'failed')
        self.assertEqual(job.error, 'platform_failure: Mapping output could not be adopted.')
        self.assertIn(':failed:', job.request_key)
        self.assertEqual(obs['charged'], Decimal('0.00075'))
        self.assertEqual(obs['absorbed_ledger'], Decimal('0.00075'))
        self.assertEqual(obs['platform_absorbed'], [Decimal('0.0005')])
        self.assertEqual(obs['mapping_failed'], 1)
        self.assertFalse(obs['deliverable'])
        self.assertTrue(obs['audit_all_checks_pass'])

    @override_settings(AI_MODE='live', QUILLMD_MAPPING_ASSISTANCE_LIVE=True)
    def test_not_landed_control(self):
        """Control: the adoption did NOT commit. Platform failure absorbed is the right record here."""
        job, harness, key_before, wallets_before, client = self.run_live_mapping(landed=False)
        obs = self.observe(job)
        report('not_landed', fired=harness.fired, waiter_seen=harness.waiter_seen, **obs)
        self.assertTrue(harness.fired)
        self.assertEqual(obs['versions'], [(1, None)])
        self.assertEqual(obs['review_ready'], 0)
        self.assertEqual(job.state, 'failed')
