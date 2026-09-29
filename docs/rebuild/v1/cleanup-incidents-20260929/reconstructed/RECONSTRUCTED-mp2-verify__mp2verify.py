"""Independent reproduction of R1/R2/R3/U1/U2 against the v1-mapping-partial worktree. Synthetic blank forms only.
No provider construction; non-loopback sockets refused."""
import copy
import io
import json
import socket
import zipfile
from pathlib import Path
from unittest.mock import patch

_real_connect = socket.socket.connect


def _guard(self, address):
    host = address[0] if isinstance(address, tuple) else address
    if isinstance(host, str) and host not in ('127.0.0.1', '::1', 'localhost') and not host.startswith('/'):
        raise AssertionError(f'non-loopback socket refused: {host!r}')
    return _real_connect(self, address)


socket.socket.connect = _guard

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from core import mapping_assistance as mapping, v3_mapping
from core.security import private_path
from core.tests.test_mapping_places import PlaceRule, bordered_box, saved
from documents.service import DocumentError, fill_template
from documents import assistance as doc_assistance

OUT = Path('/tmp/claude-1000/-opt-meddoc-pro/64d93580-8414-4553-b17d-ed04b0550ae1/scratchpad/mp2-verify/out')
OUT.mkdir(parents=True, exist_ok=True)
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def pict_run():
    run = OxmlElement('w:r'); run.append(OxmlElement('w:pict')); return run


def box_form_with(heading, question, fill=None):
    """box_form, but the box's cell holds `fill`: None (empty), 'pict' (one run with an empty w:pict), or text."""
    doc = Document(); doc.add_paragraph('Participant ID: __________')
    doc.add_paragraph(heading, style='Heading 2'); doc.add_paragraph(question)
    cell = bordered_box(doc).cell(0, 0)
    if fill == 'pict':
        cell.paragraphs[0]._p.append(pict_run())
    elif fill:
        cell.paragraphs[0].text = fill
    return saved(doc)


def column_form(pict):
    doc = Document(); doc.add_paragraph('Clinic record', style='Heading 2')
    table = doc.add_table(rows=2, cols=2); table.cell(0, 0).text = 'Name'; table.cell(0, 1).text = 'Signature'
    if pict:
        table.cell(1, 1).paragraphs[0]._p.append(pict_run())
    return saved(doc)


def locked_sdt_form(lock=True):
    doc = Document(); doc.add_paragraph('Plan: ______________')
    body = doc.element.body
    lock_xml = '<w:lock w:val="sdtContentLocked"/>' if lock else ''
    sdt = etree.fromstring(
        f'<w:sdt xmlns:w="{W}"><w:sdtPr><w:alias w:val="Locked statement"/><w:tag w:val="locked"/>{lock_xml}</w:sdtPr>'
        '<w:sdtContent><w:p><w:r><w:t xml:space="preserve">Consultant statement: ______________</w:t></w:r></w:p>'
        '</w:sdtContent></w:sdt>')
    body.insert(len(body) - 1, sdt)   # before sectPr
    return saved(doc)


def lines_form(second):
    doc = Document(); doc.add_paragraph('Plan: ______________'); doc.add_paragraph(second)
    return saved(doc)


class Repro(PlaceRule):
    maxDiff = None
    record = {}

    def run_change(self, name, data, slot, kind, label, roles=(0, 1, 2), options=(), drop=(), fields=None):
        version = self.version(name, data, drop=drop)
        compiled, payload = self.compile(version)
        places = doc_assistance.places(data)
        change = self.addition(payload, slot, kind, label, roles)
        sid = self.sid(payload, slot)
        for option in options:
            change.extend(self.option_ref(payload, sid, option))
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        adopted = self.adopted(result)
        out = {'engine_fields': [(f['label'], f['kind'], sorted(v3_mapping.mapped_slots([f]))) for f in version.fields],
               'held': places['held'], 'anchored': places['anchored'],
               'adopted': [{'label': f['label'], 'kind': f['kind'], 'options': f.get('options'),
                            'locator': f.get('locator'), 'write_protected': f.get('write_protected'),
                            'answerable': f['key'] in [a['key'] for a in v3_mapping.answerable(result['fields'])]}
                           for f in adopted],
               'rejected': [r['reason'] for r in result['review']['rejected_changes']],
               'unresolved': [(r['slot'], r['reason']) for r in result['review']['unresolved']]}
        if adopted:
            out['fill'] = self.fill(name, version, result)
        self.record[name] = out
        print('\n=== ' + name + '\n' + json.dumps(out, indent=1, default=str))
        return out

    @staticmethod
    def option_ref(payload, sid, option):
        for position, (owner, role, index) in enumerate(payload['text_refs']):
            if owner == sid and role == 3 and payload['slots'][sid]['options'][index] == option:
                return [position, 0, len(option)]
        for position, (owner, role, _) in enumerate(payload['text_refs']):
            if owner == sid and role == 0:
                source = payload['slots'][sid]['text']
                if option in source:
                    return [position, source.index(option), source.index(option) + len(option)]
        raise AssertionError(option)

    def fill(self, name, version, result):
        adopted = self.adopted(result)
        answers = {}
        for f in adopted:
            answers[f['key']] = f['options'][0] if f['kind'] == 'choice' else 'SYNTHETIC-ANSWER'
        output = Path(self.directory.name) / f'{name}-out.docx'
        try:
            fill_template(private_path(version.template.original_path), result['fields'], answers, output)
        except DocumentError as exc:
            return {'refused': str(exc), 'code': exc.code}
        xml = zipfile.ZipFile(output).read('word/document.xml').decode()
        root = etree.fromstring(xml.encode())
        hits = []
        for t in root.iter(f'{{{W}}}t'):
            if 'SYNTHETIC-ANSWER' in (t.text or '') or (answers and any(a in (t.text or '') for a in answers.values() if a != 'SYNTHETIC-ANSWER')):
                anc = [a.tag.split('}')[1] for a in t.iterancestors() if a.tag.split('}')[1] in ('tc', 'tbl', 'sdt', 'sdtContent', 'p', 'body')]
                paragraph = next(a for a in t.iterancestors() if a.tag == f'{{{W}}}p')
                hits.append({'ancestors': anc, 'paragraph_text': ''.join(x.text or '' for x in paragraph.iter(f'{{{W}}}t'))})
        sdt_lock = [etree.tostring(s, encoding='unicode')[:600] for s in root.iter(f'{{{W}}}sdt')]
        return {'hits': hits, 'sdt_xml': sdt_lock}

    # R1 ------------------------------------------------------------------------------------------------------------
    def test_R1(self):
        slot = 'b/tbl[0]/tr[0]/tc[0]/p[0]'
        self.run_change('R1-P1-pict', box_form_with('Consultant review', 'Signed by the consultant:', 'pict'),
                        slot, 'long_text', 'Consultant review', roles=(1,))
        self.run_change('R1-P1-control-empty', box_form_with('Consultant review', 'Signed by the consultant:', None),
                        slot, 'long_text', 'Consultant review', roles=(1,))
        self.run_change('R1-P1b-ink', box_form_with('Consultant review', 'Signed by the consultant:', 'Use black ink'),
                        slot, 'text', 'Use black ink', roles=(0,))
        self.run_change('R1-P7-consent-pict', box_form_with('Consultant review', 'Record that the patient consents to proceed:', 'pict'),
                        slot, 'long_text', 'Consultant review', roles=(1,))
        self.run_change('R1-P7-control-empty', box_form_with('Consultant review', 'Record that the patient consents to proceed:', None),
                        slot, 'long_text', 'Consultant review', roles=(1,))
        self.run_change('R1-P7b-consent-word-pict', box_form_with('Consultant review', 'I confirm the patient gave consent to proceed:', 'pict'),
                        slot, 'long_text', 'Consultant review', roles=(1,))
        self.run_change('R1-P7b-control-empty', box_form_with('Consultant review', 'I confirm the patient gave consent to proceed:', None),
                        slot, 'long_text', 'Consultant review', roles=(1,))

    # R2 ------------------------------------------------------------------------------------------------------------
    def test_R2(self):
        self.run_change('R2-pict', column_form(True), 'b/tbl[0]/tr[1]/tc[1]/p[0]', 'text', 'Clinic record', roles=(1,))
        self.run_change('R2-control-empty', column_form(False), 'b/tbl[0]/tr[1]/tc[1]/p[0]', 'text', 'Clinic record',
                        roles=(1,), drop=('Signature',))

    # R3 ------------------------------------------------------------------------------------------------------------
    def test_R3(self):
        self.run_change('R3a-locked-p', locked_sdt_form(), 'b/sdt[0]/p[0]', 'text', 'Consultant statement', roles=(0,))
        self.run_change('R3b-locked-sdt', locked_sdt_form(), 'b/sdt[0]', 'text', 'Consultant statement', roles=(0,))

    # U1 ------------------------------------------------------------------------------------------------------------
    def test_U1(self):
        self.run_change('U1-circle', lines_form('Visit type (circle one): Initial / Follow-up'), 'b/p[1]', 'choice',
                        'Visit type', roles=(0,), options=('Initial', 'Follow-up'))
        self.run_change('U1-circle-control-new', lines_form('Visit type (circle one): New / Follow-up'), 'b/p[1]', 'choice',
                        'Visit type', roles=(0,), options=('New', 'Follow-up'))
        self.run_change('U1-checkbox-choice', lines_form('Visit type: ☐ Initial ☐ Follow-up'), 'b/p[1]', 'choice',
                        'Visit type', roles=(0,), options=('Initial', 'Follow-up'), drop=('Visit type',))
        self.run_change('U1-checkbox-text', lines_form('Visit type: ☐ Initial ☐ Follow-up'), 'b/p[1]', 'text',
                        'Visit type', roles=(0,), drop=('Visit type',))

    def test_U1_replacement(self):
        data = lines_form('Is this the initial presentation? ☐ Yes ☐ No')
        version = self.version('U1-repl', data)
        compiled, payload = self.compile(version)
        print('\n=== U1-repl engine fields', [(f['label'], f['kind'], f.get('options'), sorted(v3_mapping.mapped_slots([f]))) for f in version.fields])
        prior = next(f for f in version.fields if f['kind'] == 'choice')
        slot = next(iter(sorted(v3_mapping.mapped_slots([prior]))))
        slot = slot.split('#', 1)[1]
        change = self.replacement(payload, prior['key'], slot, prior['kind'])
        sid = self.sid(payload, slot)
        for option in prior['options']:
            change.extend(self.option_ref(payload, sid, option))
        result = mapping.validate_proposal(version, compiled,
                                           self.respond(payload, [change], replace=[self.field_id(payload, prior['key'])]))
        print(json.dumps({'adopted': [(f['label'], f['kind'], f.get('options')) for f in self.adopted(result)],
                          'rejected': [r['reason'] for r in result['review']['rejected_changes']],
                          'held': doc_assistance.places(data)['held']}, indent=1))
        # The engine's own fill of its own field on this line (it is answerable, not a signing area).
        answers = {prior['key']: prior['options'][0]}
        output = Path(self.directory.name) / 'u1-engine-out.docx'
        fill_template(private_path(version.template.original_path), version.fields, answers, output)
        print('engine answerable:', [f['key'] for f in v3_mapping.answerable(version.fields)] == [prior['key']] or
              [f['label'] for f in v3_mapping.answerable(version.fields)])

    # U2 ------------------------------------------------------------------------------------------------------------
    def test_U2(self):
        slot = 'b/tbl[0]/tr[0]/tc[0]'
        self.run_change('U2-initial-question', box_form_with('Clinical history', 'Describe the initial presentation?'),
                        slot, 'long_text', 'Clinical history', roles=(1,), drop=('Describe the initial presentation?',))
        self.run_change('U2-initial-question-control', box_form_with('Clinical history', 'Describe the presentation?'),
                        slot, 'long_text', 'Clinical history', roles=(1,), drop=('Describe the presentation?',))
        self.run_change('U2-initial-heading', box_form_with('Initial Assessment', 'Please describe the presenting complaint?'),
                        slot, 'long_text', 'Initial Assessment', roles=(1,), drop=('Please describe the presenting complaint?',))
        self.run_change('U2-initial-heading-control', box_form_with('Wound review', 'Please describe the presenting complaint?'),
                        slot, 'long_text', 'Wound review', roles=(1,), drop=('Please describe the presenting complaint?',))


def column_form_text(text):
    doc = Document(); doc.add_paragraph('Clinic record', style='Heading 2')
    table = doc.add_table(rows=2, cols=2); table.cell(0, 0).text = 'Name'; table.cell(0, 1).text = 'Signature'
    table.cell(1, 1).text = text
    return saved(doc)


class Extra(Repro):
    def test_extra(self):
        # R1 printed guidance in the box, labelled from the heading.
        self.run_change('R1-P1c-ink-heading-label', box_form_with('Consultant review', 'Signed by the consultant:', 'Use black ink'),
                        'b/tbl[0]/tr[0]/tc[0]/p[0]', 'long_text', 'Consultant review', roles=(1,))
        # R2 with printed guidance in the Signature column cell.
        self.run_change('R2-ink', column_form_text('Use black ink'), 'b/tbl[0]/tr[1]/tc[1]/p[0]', 'text', 'Use black ink', roles=(0,))
        # R3: does the engine map the UNLOCKED control? (is lock read by detection at all)
        self.run_change('R3-unlocked-p', locked_sdt_form(lock=False), 'b/sdt[0]/p[0]', 'text', 'Consultant statement', roles=(0,))

    def test_R3_lock_in_output(self):
        version = self.version('R3lock', locked_sdt_form())
        compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/sdt[0]/p[0]', 'text', 'Consultant statement', (0,))
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        f = self.adopted(result)[0]
        output = Path(self.directory.name) / 'r3lock.docx'
        fill_template(private_path(version.template.original_path), result['fields'], {f['key']: 'SYNTHETIC-ANSWER'}, output)
        root = etree.fromstring(zipfile.ZipFile(output).read('word/document.xml'))
        for sdt in root.iter(f'{{{W}}}sdt'):
            pr = sdt.find(f'{{{W}}}sdtPr')
            print('\nR3 output sdtPr children:', [(c.tag.split('}')[1], dict((k.split('}')[1], v) for k, v in c.attrib.items())) for c in pr])
            print('R3 output sdtContent text:', ''.join(t.text or '' for t in sdt.iter(f'{{{W}}}t')))
        # The manual clinician path (locate_fields, as v3_add_field uses) gives the same locator.
        from documents.service import locate_fields
        manual = locate_fields(private_path(version.template.original_path), [{'slot': 'word/document.xml#b/sdt[0]/p[0]', 'kind': 'text', 'label': 'Consultant statement'}])
        print('R3 manual locate:', manual['fields'][0]['locator']['slot'], manual['fields'][0]['locator']['target'])
