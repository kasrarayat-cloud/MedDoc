"""Bounded retained signing counterexamples; synthetic only, no provider."""
import json
from pathlib import Path
from docx import Document
from documents import word_v3 as w
from documents.test_question_boxes import box, data
here=Path(__file__).parent
rows=[]
for heading,label in [('Report','Please provide the initial'),('', 'Please countersign the initial report?'),('', 'Please describe the changes since the initial report?')]:
    doc=Document()
    if heading: doc.add_paragraph(heading, style='Heading 1')
    doc.add_paragraph(label);box(doc)
    rows.append({'heading':heading,'label':label,'result':w.inspect(data(doc))})
print(json.dumps(rows,indent=2))
assert not rows[0]['result']['fields'] and not rows[1]['result']['fields']
assert len(rows[2]['result']['fields'])==1
