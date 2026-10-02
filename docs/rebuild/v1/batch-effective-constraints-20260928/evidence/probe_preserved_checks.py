from pathlib import Path
import os,sys,json,hashlib,io
TARGET=Path('/opt/meddoc-pro/.worktrees/v1-batch-effective-constraints')
sys.path.insert(0,str(TARGET/'rebuild'))
os.environ.update(QUILLMD_ENV_MANAGED='1',QUILLMD_SECRET_KEY='synthetic-batch-refutation',QUILLMD_DB_HOST='127.0.0.1',QUILLMD_DB_PORT='1',QUILLMD_DB_NAME='no_database',QUILLMD_DB_USER='no_database',QUILLMD_DB_PASSWORD='synthetic',QUILLMD_AI_MODE='fixture',DJANGO_SETTINGS_MODULE='quillmd.settings')
import django;django.setup()
from core import generation_batch as b,v3_mapping
from documents import word_v3 as w
from docx import Document
from types import SimpleNamespace
R=Path(__file__).parent
D=Document();D.add_paragraph('Select at most one active group.')
t=D.add_table(rows=103,cols=2)
for i,label in enumerate(('Group A','Group B')):
 t.cell(i,0).text=label;t.cell(i,1).text=('☒ Active ☐ Inactive' if i==0 else '☐ Active ☐ Inactive')
for i in range(2,103):t.cell(i,0).text=f'Note {i:03d}'
rawio=io.BytesIO();D.save(rawio);raw=rawio.getvalue();(R/'prechecked-original.docx').write_bytes(raw)
info=w.inspect(raw);fs=info['fields'];selected={f['label']:f for f in fs if f['label'] in ('Group A','Group B')};assert len(selected)==2,fs
rules={'contract':b.CONSTRAINTS,'source_sha256':hashlib.sha256(raw).hexdigest(),'fields_sha256':b.digest(fs),'coverage':'complete','provenance':'reviewed_assisted','evidence_sha256':'1'*64,'instruction_references':[{'id':'rule','part':'word/document.xml','slot':'b/p[0]','sha256':hashlib.sha256(b'Select at most one active group.').hexdigest()}],'rules':[{'kind':'max_active_groups','source_id':'rule','maximum':1,'groups':[[{'key':selected[label]['key'],'values':['Active']}] for label in ('Group A','Group B')]}]}
version=SimpleNamespace(fields=fs,template=SimpleNamespace(sha256=rules['source_sha256']),generation_constraints=rules)
answers={f['key']:{'text':'','status':'missing','reviewed':True,'left_blank':True} for f in fs}
answers[selected['Group B']['key']]={'text':'Active','status':'supported','reviewed':True,'edited_by':'synthetic-clinician'}
b.check_answers(version,answers)
normalized=v3_mapping.fill_answers(fs,answers);out=w.fill(raw,fs,normalized);(R/'prechecked-export.docx').write_bytes(out)
physical={}
for label,f in selected.items():
 slot=w.SlotIndex(w.open_package(out)).by_id[f['locator']['part']+'#'+f['locator']['slot']]
 # Read physical paragraph text directly rather than trusting constraint/field state readers.
 physical[label]=w.paragraph_text(slot.el)[0] if slot.kind=='p' else w.Detector(w.open_package(out)).text(slot.el)
result={'source_commit':'b1c3b95cca1a807f31eb5e3244258e8373bf9751','generation_batch_sha256':hashlib.sha256((TARGET/'rebuild/core/generation_batch.py').read_bytes()).hexdigest(),'field_count':len(fs),'batch_needed':b.needed(fs),'fields':selected,'constraints':rules,'answers':{k:answers[f['key']] for k,f in selected.items()},'constraint_validation':'accepted','physical_export':physical,'input_sha256':hashlib.sha256(raw).hexdigest(),'export_sha256':hashlib.sha256(out).hexdigest()}
(R/'prechecked-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False,indent=2))
