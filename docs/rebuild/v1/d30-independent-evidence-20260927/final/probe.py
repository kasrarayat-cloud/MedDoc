"""Fixed independent D30 followup: retained eight artifacts + twelve local controls."""
import hashlib
import importlib.util
import io
import json
import sys
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
OLD = HERE.parent
sys.path.insert(0, '/opt/meddoc-pro/.worktrees/v1-word-d30/rebuild')


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


helpers = load('d30_original_probe_helpers', OLD / 'probe.py')


def set_ref(parent, tag, value):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    el.set(qn('w:val'), value)


def paragraph_case(mode):
    doc = Document()
    paragraph = doc.add_paragraph('What symptoms changed?')
    if mode == 'unknown':
        set_ref(paragraph._p.get_or_add_pPr(), 'w:pStyle', 'AbsentParagraph')
    else:
        style = doc.styles.add_style('Proof Paragraph', WD_STYLE_TYPE.PARAGRAPH)
        paragraph.style = style
        set_ref(style.element, 'w:basedOn', style.style_id if mode == 'cycle' else 'AbsentParent')
    helpers.box(doc)
    return doc


def table_case(mode):
    doc = Document()
    doc.add_paragraph('What symptoms changed?')
    table = helpers.box(doc, direct=False)
    if mode == 'unknown':
        set_ref(table._tbl.tblPr, 'w:tblStyle', 'AbsentTable')
    else:
        style = doc.styles.add_style('Proof Table', WD_STYLE_TYPE.TABLE)
        table.style = style
        set_ref(style.element, 'w:basedOn', style.style_id if mode == 'cycle' else 'AbsentParent')
    helpers.borders(table._tbl.tblPr, 'w:tblBorders')
    return doc


def controls():
    for mode in ('unknown', 'cycle', 'missing-parent'):
        yield 'paragraph-' + mode, paragraph_case(mode), 'hold', ''
    for mode in ('unknown', 'cycle', 'missing-parent'):
        yield 'table-' + mode, table_case(mode), 'hold', ''
    for override in (False, True):
        doc = Document()
        doc.add_paragraph('What symptoms changed?')
        style = doc.styles.add_style('Break Spacer', WD_STYLE_TYPE.PARAGRAPH)
        style.paragraph_format.page_break_before = True
        p = doc.add_paragraph('', style)
        if override:
            p.paragraph_format.page_break_before = False
        helpers.box(doc)
        yield 'inherited-break' + ('-overridden-false' if override else ''), doc, 'map' if override else 'hold', ''
    doc = Document()
    parent = doc.styles.add_style('Inherited Major', WD_STYLE_TYPE.PARAGRAPH)
    parent.base_style = doc.styles['Heading 1']
    child = doc.styles.add_style('Inherited Local', WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = doc.styles['Heading 2']
    doc.add_paragraph('Part B', parent)
    doc.add_paragraph('Donor', child)
    doc.add_paragraph('Medication history', child)
    doc.add_paragraph('What medications are currently taken?')
    helpers.box(doc)
    yield 'full-inherited-scope', doc, 'map', 'Part B > Donor > Medication history'
    for child_nil in (False, True):
        doc = Document()
        doc.add_paragraph('What symptoms changed?')
        table = helpers.box(doc, direct=False)
        parent = doc.styles.add_style('Simple Parent Box', WD_STYLE_TYPE.TABLE)
        props = OxmlElement('w:tblPr')
        helpers.borders(props, 'w:tblBorders')
        parent.element.append(props)
        child = doc.styles.add_style('Simple Child Box', WD_STYLE_TYPE.TABLE)
        child.base_style = parent
        table.style = child
        if child_nil:
            props = OxmlElement('w:tblPr')
            helpers.borders(props, 'w:tblBorders', 'nil')
            child.element.append(props)
        yield 'simple-ancestor-borders' + ('-child-nil' if child_nil else ''), doc, 'hold' if child_nil else 'map', ''
    doc = Document()
    doc.add_paragraph('Which medications are currently taken?')
    doc.add_paragraph()
    doc.add_paragraph('Please provide their doses and frequencies.')
    helpers.box(doc)
    yield 'split-question-one-spacer', doc, 'hold', ''


def reduced(result):
    return {'fields': [{'label': f['label'], 'context': f['context'],
                       'slot': f['locator']['slot'], 'target': f['locator']['target']}
                      for f in result['fields']], 'notes': result['notes']}


def main():
    source = HERE / 'word-v3.py'
    assert hashlib.sha256(source.read_bytes()).hexdigest() == 'd4d82856ca066c07b537f3c3d28f9a10b1b3bcbf441d516e6ba636699ea8bc77'
    original = json.loads((OLD / 'result.json').read_text())['cases']
    prepared = list(controls())
    (HERE / 'predeclared.json').write_text(json.dumps({
        'original_artifacts': [{'name': r['name'], 'sha256': r['original_sha256'], 'expected': r['expected']} for r in original],
        'controls': [{'name': n, 'expected': expected, 'context': context} for n, _, expected, context in prepared]
    }, indent=2) + '\n')
    engines = {'previous': load('documents.word_v3_d30_previous', OLD / 'candidate-word-v3.py'),
               'finished': load('documents.word_v3_d30_finished', source)}
    records = []
    with patch('socket.socket.connect', side_effect=AssertionError('Network denied')):
        for old in original:
            raw = (OLD / (old['name'] + '.docx')).read_bytes()
            assert hashlib.sha256(raw).hexdigest() == old['original_sha256']
            outputs = {name: reduced(engine.inspect(raw)) for name, engine in engines.items()}
            assert outputs['previous'] == old['engines']['candidate']
            records.append({'name': old['name'], 'kind': 'retained-original', 'sha256': old['original_sha256'], 'outputs': outputs})
        for name, doc, expected, context in prepared:
            raw = helpers.encode(doc)
            (HERE / (name + '.docx')).write_bytes(raw)
            outputs = {key: reduced(engine.inspect(raw)) for key, engine in engines.items()}
            finished = outputs['finished']
            passed = bool(not finished['fields'] and finished['notes']) if expected == 'hold' else bool(
                len(finished['fields']) == 1 and not finished['notes'] and finished['fields'][0]['context'] == context)
            records.append({'name': name, 'kind': 'new-local-control', 'sha256': hashlib.sha256(raw).hexdigest(),
                            'expected': expected, 'expected_context': context, 'pass': passed, 'outputs': outputs})
    result = {'finished_commit': 'f170f5a7e10d557679d3fb6000948f7bd5df5e16',
              'finished_source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'cases': records}
    (HERE / 'result.json').write_text(json.dumps(result, indent=2) + '\n')
    for row in records:
        print(row['name'], json.dumps(row['outputs']['finished']), 'PASS=' + str(row.get('pass', 'inspect')))


if __name__ == '__main__':
    main()
