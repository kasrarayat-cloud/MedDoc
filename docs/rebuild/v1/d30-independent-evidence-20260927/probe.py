"""Finite independent synthetic structural refutation; no corpus, provider or database."""
import hashlib
import importlib.util
import io
import json
import sys
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
sys.path.insert(0, '/opt/meddoc-pro/.worktrees/v1-word-d30/rebuild')


def load(name, filename):
    spec = importlib.util.spec_from_file_location(name, HERE / filename)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


def borders(parent, tag, value='single'):
    b = OxmlElement(tag)
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side); e.set(qn('w:val'), value); b.append(e)
    parent.append(b)


def box(doc, direct=True):
    t = doc.add_table(rows=1, cols=1)
    if direct: borders(t.cell(0, 0)._tc.get_or_add_tcPr(), 'w:tcBorders')
    return t


def encode(doc):
    out = io.BytesIO(); doc.save(out); return out.getvalue()


def cases():
    d = Document(); d.add_paragraph('Donor', 'Heading 1'); d.add_paragraph('What medications are currently taken?'); box(d)
    yield 'direct-donor-heading', d, 'Context must retain Donor.'
    d = Document(); style = d.styles.add_style('Donor Section', WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = d.styles['Heading 1']; d.add_paragraph('Donor', style)
    d.add_paragraph('What medications are currently taken?'); box(d)
    yield 'inherited-donor-heading', d, 'Equivalent inherited Heading 1 context must retain Donor or visibly refuse.'
    d = Document(); d.add_paragraph('Which medications are currently taken?')
    d.add_paragraph('Please provide their doses and frequencies.'); box(d)
    yield 'two-paragraph-question', d, 'Cannot silently assign only the anaphoric second part; retain complete question context or warn.'
    d = Document(); d.add_paragraph('What symptoms changed?'); box(d)
    yield 'plain-positive', d, 'One complete question maps to its adjacent box.'
    d = Document(); d.add_paragraph('What symptoms changed?'); t = box(d, False); t.style = 'Table Grid'
    borders(t.cell(0, 0)._tc.get_or_add_tcPr(), 'w:tcBorders', 'nil')
    yield 'cell-nil', d, 'Explicit cell nil overrides style and must refuse.'
    d = Document(); d.add_paragraph('What symptoms changed?'); t = box(d, False)
    parent = d.styles.add_style('Base Box', WD_STYLE_TYPE.TABLE)
    cond = OxmlElement('w:tblStylePr'); cond.set(qn('w:type'), 'firstRow')
    tcpr = OxmlElement('w:tcPr'); borders(tcpr, 'w:tcBorders', 'nil'); cond.append(tcpr); parent.element.append(cond)
    child = d.styles.add_style('Derived Box', WD_STYLE_TYPE.TABLE); child.base_style = parent
    pr = OxmlElement('w:tblPr'); borders(pr, 'w:tblBorders'); child.element.append(pr); t.style = child
    look = t._tbl.tblPr.find(qn('w:tblLook')); look.set(qn('w:firstRow'), '1')
    yield 'ancestor-conditional-nil', d, 'Conditional nil exists in the style ancestry; automatic explicit-border proof must not ignore it.'
    d = Document(); d.add_paragraph('Patient name: __________. What symptoms changed?'); box(d)
    yield 'inline-blank', d, 'Preserve the original name field and warn about the additional box.'
    d = Document(); d.add_paragraph('What symptoms changed?'); p = d.add_paragraph()
    p.paragraph_format.page_break_before = True; box(d)
    yield 'page-break-spacer', d, 'A physical page-break paragraph is not a plain empty spacer; assess refusal.'


def main():
    prepared = list(cases())
    (HERE / 'predeclared.json').write_text(json.dumps([{'name': n, 'expected': x} for n, _, x in prepared], indent=2) + '\n')
    engines = {'base': load('documents.word_v3_independent_base', 'base-word-v3.py'),
               'candidate': load('documents.word_v3_independent_candidate', 'candidate-word-v3.py')}
    records = []
    with patch('socket.socket.connect', side_effect=AssertionError('Network denied')):
        for name, doc, expected in prepared:
            raw = encode(doc); (HERE / (name + '.docx')).write_bytes(raw)
            row = {'name': name, 'expected': expected, 'original_sha256': hashlib.sha256(raw).hexdigest(), 'engines': {}}
            for key, w in engines.items():
                result = w.inspect(raw)
                row['engines'][key] = {'fields': [{'label': f['label'], 'context': f['context'],
                    'slot': f['locator']['slot'], 'target': f['locator']['target']} for f in result['fields']],
                    'notes': result['notes']}
                if name == 'plain-positive' and key == 'candidate':
                    f = result['fields'][0]
                    exported = w.fill(raw, result['fields'], {f['key']: 'SYNTHETIC exact answer.'})
                    (HERE / 'positive-export.docx').write_bytes(exported)
                    with ZipFile(io.BytesIO(raw)) as a, ZipFile(io.BytesIO(exported)) as b:
                        row['unchanged_parts'] = all(a.read(p) == b.read(p) for p in a.namelist() if p != 'word/document.xml')
                    row['export_sha256'] = hashlib.sha256(exported).hexdigest()
            records.append(row)
    result = {'candidate_sha256': hashlib.sha256((HERE / 'candidate-word-v3.py').read_bytes()).hexdigest(),
              'base_sha256': hashlib.sha256((HERE / 'base-word-v3.py').read_bytes()).hexdigest(), 'cases': records}
    (HERE / 'result.json').write_text(json.dumps(result, indent=2) + '\n')
    print(json.dumps(result, indent=2))


if __name__ == '__main__': main()
