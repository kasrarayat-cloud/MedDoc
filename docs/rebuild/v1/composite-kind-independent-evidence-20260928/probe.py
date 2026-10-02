from pathlib import Path
import hashlib,importlib.util,io,json,socket,subprocess,sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).parent;T=Path('/opt/meddoc-pro/.worktrees/v1-workflow-candidate')
raw=(T/'rebuild/documents/word_v3.py').read_bytes();assert hashlib.sha256(raw).hexdigest()=='21399290ec826f9c0c2b9e419b725a93008f39da565558ded076372990af0bc6'
(R/'candidate.py').write_bytes(raw)
(R/'baseline.py').write_bytes(subprocess.check_output(['git','show','ab41d7107f8ccb34bf64b034fdb28446fc7478b4:rebuild/documents/word_v3.py'],cwd=T))
socket.socket=lambda *a,**kw: (_ for _ in ()).throw(AssertionError('network forbidden'))
def load(key):
 spec=importlib.util.spec_from_file_location('documents._independent_'+key,R/(key+'.py'));m=importlib.util.module_from_spec(spec);sys.modules[spec.name]=m;spec.loader.exec_module(m);return m
base=load('baseline');candidate=load('candidate')
def data(doc):
 b=io.BytesIO();doc.save(b);return b.getvalue()
cases=[('Date of birth','date','date'),('D.O.B.','date','date'),('Time (24 h)','time','time'),('Date and description of restrictions','date','long_text'),('Date and details of examination','date','long_text'),('Time and reason for attendance','time','long_text'),('Signature and comments','signature','signature'),('Initials and details','initials','initials'),('Patient name','text','text'),('Comments','long_text','long_text')]
results=[]
for label,old,new in cases:
 doc=Document();t=doc.add_table(rows=1,cols=2);t.cell(0,0).text=label;raw=data(doc)
 a=base.inspect(raw);b=candidate.inspect(raw);assert len(a['fields'])==len(b['fields'])==1
 af,bf=a['fields'][0],b['fields'][0];assert af['kind']==old and bf['kind']==new,(label,af,bf)
 assert {k:v for k,v in af.items() if k!='kind'}=={k:v for k,v in bf.items() if k!='kind'}
 value='February 2, 2026; retain the complete documented restriction and time 09:30.'
 if new not in ('signature','initials'):
  ea=base.fill(raw,a['fields'],{af['key']:value});eb=candidate.fill(raw,b['fields'],{bf['key']:value});assert ea==eb
 results.append({'label':label,'base_kind':old,'candidate_kind':new,'all_other_metadata_identical':True,'export_equal_when_applicable':True})
for alias in ('Date and summary of findings','Time and history details'):
 doc=Document();sdt=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr');sdt.append(pr);tag=OxmlElement('w:alias');tag.set(qn('w:val'),alias);pr.append(tag);pr.append(OxmlElement('w:date'));content=OxmlElement('w:sdtContent');sdt.append(content);run=OxmlElement('w:r');content.append(run);text=OxmlElement('w:t');text.text='Click here to enter a date.';run.append(text);doc.add_paragraph()._p.append(sdt);raw=data(doc)
 a=base.inspect(raw);b=candidate.inspect(raw);assert a==b;assert a['fields'][0]['kind']=='date'
 results.append({'native_alias':alias,'base_equals_candidate':True,'kind':'date'})
p=R/'result.json';assert not p.exists();p.write_text(json.dumps({'candidate_sha256':hashlib.sha256((R/'candidate.py').read_bytes()).hexdigest(),'baseline_sha256':hashlib.sha256((R/'baseline.py').read_bytes()).hexdigest(),'cases':results,'finding':'NOTHING FOUND within bounded claims'},indent=2));print(len(results),'controls PASS')
