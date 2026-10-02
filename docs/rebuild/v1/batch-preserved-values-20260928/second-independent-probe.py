"""Independent frozen-diff probes. Documents remain in memory; no DB or network."""
import copy
import hashlib
import io
import json
import socket
import sys
from pathlib import Path
from types import SimpleNamespace

SOURCE = Path('/opt/meddoc-pro/.worktrees/v1-batch-preserved-values/rebuild')
sys.path.insert(0, str(SOURCE))
def deny(*args, **kwargs):
    raise AssertionError('Network and DB connections forbidden by independent review')
socket.socket.connect = deny
socket.create_connection = deny
from django.conf import settings
settings.configure(SECRET_KEY='synthetic-only', INSTALLED_APPS=[])
from core import generation_batch as b, v3_mapping
from documents import word_v3 as w
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.simpletypes import ST_OnOff

def el(name, **attrs):
    node = OxmlElement(name)
    for key, value in attrs.items(): node.set(qn('w:' + key), value)
    return node

def run(parent, child):
    node = el('w:r'); node.append(child); parent.append(node)
    return node

def text(parent, value):
    child = el('w:t'); child.text = value; return run(parent, child)

def legacy(p, value, mode='normal', default=''):
    begin = el('w:fldChar', fldCharType='begin'); ff = el('w:ffData')
    ff.append(el('w:name', val='Priority')); inp = el('w:textInput')
    inp.append(el('w:type', val='regular'))
    if default: inp.append(el('w:default', val=default))
    ff.append(inp); begin.append(ff); run(p._p, begin)
    instr = el('w:instrText'); instr.text = ' FORMTEXT '; run(p._p, instr)
    if mode != 'no_separator': run(p._p, el('w:fldChar', fldCharType='separate'))
    if mode == 'hyperlink':
        link=el('w:hyperlink',anchor='somewhere'); text(link,value); p._p.append(link)
    elif mode == 'nested':
        run(p._p, el('w:fldChar', fldCharType='begin'))
        ins=el('w:instrText'); ins.text=' DOCVARIABLE value '; run(p._p,ins)
        run(p._p,el('w:fldChar',fldCharType='separate')); text(p._p,value)
        run(p._p,el('w:fldChar',fldCharType='end'))
    else: text(p._p,value)
    run(p._p,el('w:fldChar',fldCharType='end'))

def native(p, value, *, scalar=False, marker=None, missing_content=False):
    control=el('w:sdt'); pr=el('w:sdtPr'); control.append(pr)
    if scalar: pr.append(el('w:date'))
    else:
        dd=el('w:dropDownList')
        for option in ['1','2','3']: dd.append(el('w:listItem',displayText=option,value=option))
        pr.append(dd)
    if marker is not None:
        flag=el('w:showingPlcHdr')
        if marker != 'present': flag.set(qn('w:val'), marker)
        pr.append(flag)
    if not missing_content:
        content=el('w:sdtContent'); text(content,value); control.append(content)
    p._p.append(control)

INSTRUCTION='Each priority must be unique and in 1, 2, 3.'

def execute(name, maker, first='', second='', expected='REFUSED', copies=False):
    doc=Document(); doc.add_paragraph(INSTRUCTION)
    maker(doc)
    for i in range(101): doc.add_paragraph(f'Ordinary {i:03}: ____________')
    buffer=io.BytesIO(); doc.save(buffer); raw=buffer.getvalue()
    fields=w.inspect(raw)['fields']
    selected=[f for f in fields if f['locator']['target']['type'] in ('ff_text','sdt_text','sdt_choice')]
    assert len(selected)==2,(name,selected)
    selected.sort(key=lambda f: f['label'])
    m={'contract':b.CONSTRAINTS,'source_sha256':hashlib.sha256(raw).hexdigest(),
       'fields_sha256':b.digest(fields),'coverage':'complete','provenance':'reviewed_assisted',
       'evidence_sha256':'1'*64,'instruction_references':[{'id':'rule','part':'word/document.xml','slot':'b/p[0]',
       'sha256':hashlib.sha256(INSTRUCTION.encode()).hexdigest()}],
       'rules':[{'kind':'unique_values','source_id':'rule','keys':[f['key'] for f in selected],'allowed_values':['1','2','3']}]}
    version=SimpleNamespace(fields=fields,template=SimpleNamespace(sha256=m['source_sha256']),generation_constraints=m)
    b.constraints(version)
    answers={f['key']:{'text':'','reviewed':True,'status':'left_blank'} for f in fields}
    for f,answer in zip(selected,(first,second)): answers[f['key']].update(text=answer,status='supported' if answer else 'left_blank')
    reader=lambda controls,writes:w.effective_values(raw,controls,writes)
    try:
        b.validate_constraints(m,answers,fields,read_options=reader)
        outcome='ACCEPTED'; reason=''
    except Exception as exc: outcome='REFUSED'; reason=type(exc).__name__+': '+str(exc)
    try:
        effective={k:sorted(v) for k,v in b.effective_values(m,fields,answers,reader).items()}
    except Exception as exc: effective={'error':type(exc).__name__+': '+str(exc)}
    try:
        filled=w.fill(raw,fields,v3_mapping.fill_answers(fields,answers)); package=w.open_package(filled)
        parts={f['locator']['part'] for f in selected}
        parts.update(x['part'] for f in selected for x in f['locator']['target'].get('copies',[]))
        exported={part:[w.text_of(p) for p in package.tree(part).iter(qn('w:p'))][:3] for part in sorted(parts)}
        export_sha=hashlib.sha256(filled).hexdigest()
    except Exception as exc: exported={'error':type(exc).__name__+': '+str(exc)};export_sha=None
    return {'name':name,'expected':expected,'outcome':outcome,'match':outcome==expected,'reason':reason,
            'field_count':len(fields),'effective_values':effective,'exported':exported,'constrained_fields':selected,
            'input_sha256':hashlib.sha256(raw).hexdigest(),'export_sha256':export_sha}

def pair(doc, kind='ff', a='1', z='', **kw):
    for label,value in [('Priority A: ',a),('Priority B: ',z)]:
        p=doc.add_paragraph(label)
        if kind=='ff': legacy(p,value,**(kw if label.startswith('Priority A') else {}))
        else: native(p,value,scalar=kind=='scalar',**(kw if label.startswith('Priority A') else {}))

rows=[]
rows.append(execute('F-B3-original',lambda d:pair(d),second='1'))
rows.append(execute('F-B3-explicit-resolve',lambda d:pair(d),first='2',second='1',expected='ACCEPTED'))
rows.append(execute('F-B4-original',lambda d:pair(d,'dropdown',a='4',z='2')))
rows.append(execute('F-B4-explicit-resolve',lambda d:pair(d,'dropdown',a='4',z='2'),first='3',expected='ACCEPTED'))
for mode in ('hyperlink','nested','no_separator'):
    rows.append(execute('legacy-'+mode,lambda d,mode=mode:pair(d,mode=mode),second='1'))
rows.append(execute('legacy-default-only',lambda d:pair(d,a='',z='2',default='1')))
rows.append(execute('legacy-genuinely-empty',lambda d:pair(d,a='',z='2'),expected='ACCEPTED'))
for kind in ('scalar','dropdown'):
    for flag in ('0','false','off','bogus'):
        rows.append(execute(kind+'-placeholder-'+flag,lambda d,kind=kind,flag=flag:pair(d,kind,a='1',z='2',marker=flag),second='1'))
    rows.append(execute(kind+'-real-placeholder',lambda d,kind=kind:pair(d,kind,a='Choose an item.',z='2',marker='present'),expected='ACCEPTED'))
    rows.append(execute(kind+'-missing-content',lambda d,kind=kind:pair(d,kind,a='',z='2',missing_content=True)))
    rows.append(execute(kind+'-known-value',lambda d,kind=kind:pair(d,kind,a='1',z='2'),expected='ACCEPTED'))
    rows.append(execute(kind+'-known-duplicate',lambda d,kind=kind:pair(d,kind,a='1',z='2'),second='1'))
    for pair_values in [('1','1'),('1','2')]:
        def maker(doc,kind=kind,pair_values=pair_values):
            sec=doc.sections[0];sec.different_first_page_header_footer=True
            for header,value in zip((sec.header,sec.first_page_header),pair_values):
                p=header.paragraphs[0];p.text='Priority A: ';native(p,value,scalar=kind=='scalar')
            native(doc.add_paragraph('Priority B: '),'3',scalar=kind=='scalar')
        rows.append(execute(kind+'-replicas-'+''.join(pair_values),maker,expected='ACCEPTED' if pair_values[0]==pair_values[1] else 'REFUSED'))
        rows.append(execute(kind+'-replicas-explicit-'+''.join(pair_values),maker,first='2',expected='ACCEPTED'))

result={'frozen_head':'952c7e2100e451882c1fd837980ab0212dc4f586',
        'source_hashes':{p:hashlib.sha256((SOURCE/p).read_bytes()).hexdigest() for p in ['core/generation_batch.py','documents/word_v3.py','core/v3_mapping.py']},
        'on_off_semantics':{s:ST_OnOff.convert_from_xml(s) for s in ['0','false','off','1','true','on']},
        'network_database_provider_calls':0,'cases':rows,'unexpected':[r['name'] for r in rows if not r['match']]}
print(json.dumps(result,indent=2,ensure_ascii=False))
