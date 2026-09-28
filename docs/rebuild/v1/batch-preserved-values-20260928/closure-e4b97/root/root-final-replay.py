import hashlib
import io
import json
import socket
import sys
from pathlib import Path
from types import SimpleNamespace

SOURCE = Path('/opt/meddoc-pro/.worktrees/v1-batch-preserved-values/rebuild')
sys.path.insert(0, str(SOURCE))
def deny(*args, **kwargs):
    raise AssertionError('No network or database connections allowed')
socket.socket.connect = deny
socket.create_connection = deny
from django.conf import settings
settings.configure(SECRET_KEY='synthetic-only', INSTALLED_APPS=[])
from core import generation_batch as b, v3_mapping
from documents import word_v3 as w
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def el(name, **attributes):
    node = OxmlElement(name)
    for key, value in attributes.items():
        node.set(qn('w:' + key), value)
    return node

def run(p, child):
    item = el('w:r'); item.append(child); p._p.append(item)

def legacy_text(p, value):
    begin = el('w:fldChar', fldCharType='begin')
    ff = el('w:ffData'); ff.append(el('w:name', val='Priority'))
    ff.append(el('w:enabled'))
    inp = el('w:textInput'); inp.append(el('w:type', val='regular'))
    inp.append(el('w:default', val=value)); ff.append(inp)
    begin.append(ff); run(p, begin)
    instruction = el('w:instrText'); instruction.text=' FORMTEXT '; run(p,instruction)
    run(p, el('w:fldChar', fldCharType='separate'))
    value_node=el('w:t'); value_node.text=value; run(p,value_node)
    run(p, el('w:fldChar', fldCharType='end'))

def dropdown(p, value):
    sdt=el('w:sdt'); pr=el('w:sdtPr'); dd=el('w:dropDownList')
    for option in ['1','2','3']:
        dd.append(el('w:listItem', displayText=option, value=option))
    pr.append(dd); sdt.append(pr)
    content=el('w:sdtContent'); r=el('w:r'); t=el('w:t'); t.text=value
    r.append(t); content.append(r); sdt.append(content); p._p.append(sdt)

def xml_paragraphs(data):
    root=w.open_package(data).tree('word/document.xml')
    return [''.join(p.itertext()) for p in root.iter(qn('w:p'))]

def probe(kind):
    doc=Document(); instruction='Each priority must be unique and in 1, 2, 3.'
    doc.add_paragraph(instruction)
    if kind=='preserved_scalar':
        legacy_text(doc.add_paragraph('Priority A: '),'1')
        legacy_text(doc.add_paragraph('Priority B: '),'')
    else:
        dropdown(doc.add_paragraph('Priority A: '),'4')
        dropdown(doc.add_paragraph('Priority B: '),'2')
    for i in range(101):doc.add_paragraph(f'Ordinary {i:03}: ____________')
    output=io.BytesIO();doc.save(output);raw=output.getvalue()
    fields=w.inspect(raw)['fields']
    selected=[f for f in fields if f['locator']['target']['type']==('ff_text' if kind=='preserved_scalar' else 'sdt_choice')]
    assert len(selected)==2, selected
    m={'contract':b.CONSTRAINTS,'source_sha256':hashlib.sha256(raw).hexdigest(),
       'fields_sha256':b.digest(fields),'coverage':'complete','provenance':'reviewed_assisted',
       'evidence_sha256':'1'*64,'instruction_references':[{'id':'rule','part':'word/document.xml','slot':'b/p[0]',
       'sha256':hashlib.sha256(instruction.encode()).hexdigest()}],
       'rules':[{'kind':'unique_values','source_id':'rule','keys':[f['key'] for f in selected],'allowed_values':['1','2','3']}]}
    version=SimpleNamespace(fields=fields,template=SimpleNamespace(sha256=m['source_sha256']),generation_constraints=m)
    b.constraints(version)
    answers={f['key']:{'text':'','reviewed':True,'status':'left_blank'} for f in fields}
    if kind=='preserved_scalar':answers[selected[1]['key']].update(text='1',status='supported')
    reader=lambda controls,writes:w.effective_options(raw,controls,writes)
    b.validate_constraints(m,answers,fields,read_options=reader)
    filled=w.fill(raw,fields,v3_mapping.fill_answers(fields,answers))
    effective=b.effective_values(m,fields,answers,reader)
    return {'case':kind,'field_count':len(fields),'batch_needed':b.needed(fields),
            'constrained_fields':selected,'constraint_validation':'ACCEPTED',
            'effective_values':{k:sorted(v) for k,v in effective.items()},
            'source_paragraphs':xml_paragraphs(raw)[:3], 'export_paragraphs':xml_paragraphs(filled)[:3],
            'input_sha256':hashlib.sha256(raw).hexdigest(),'export_sha256':hashlib.sha256(filled).hexdigest()}


from django.core.exceptions import ValidationError
results=[]
for case in ('preserved_scalar','unknown_dropdown'):
    try:
        value=probe(case)
        results.append({'case':case,'refused':False,'unexpected_result':value})
    except ValidationError as exc:
        results.append({'case':case,'refused':True,'reason':str(exc)})
print(json.dumps({'source_commit':'90297ea971528b0b90f68a0e770ea89088672761',
 'source_sha256':{p:hashlib.sha256((SOURCE/p).read_bytes()).hexdigest() for p in ['core/generation_batch.py','documents/word_v3.py']},'results':results},indent=2))
assert all(result['refused'] for result in results)
