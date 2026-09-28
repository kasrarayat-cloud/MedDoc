import hashlib
import io
import json
import pathlib
import sys

source = pathlib.Path(sys.argv[1]).resolve()
outdir = pathlib.Path(sys.argv[2]).resolve()
outdir.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, str(source / 'rebuild'))
from documents import word_v3 as w
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE


def run(text):
    r = OxmlElement('w:r')
    t = OxmlElement('w:t'); t.text = text; r.append(t)
    return r


def hyperlink(d, text=None):
    el = OxmlElement('w:hyperlink')
    el.set(qn('r:id'), d.part.relate_to('mailto:forms@example.invalid',
        RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    if text is not None:
        el.append(run(text))
    return el


def sdt(*, placeholder=False):
    el = OxmlElement('w:sdt'); pr = OxmlElement('w:sdtPr')
    alias = OxmlElement('w:alias'); alias.set(qn('w:val'), 'Clinical report'); pr.append(alias)
    if placeholder:
        pr.append(OxmlElement('w:showingPlcHdr'))
    el.append(pr); content = OxmlElement('w:sdtContent'); el.append(content)
    return el, content


def fixture(case):
    d = Document()
    if case == 'block-placeholder-fixed-link':
        el, content = sdt(placeholder=True)
        p = OxmlElement('w:p'); p.append(run('Enter report. Contact: '))
        p.append(hyperlink(d, 'forms@example.invalid')); content.append(p)
        d._element.body.insert(0, el)
    elif case == 'inline-linked-marker':
        p = d.add_paragraph('Name: '); el, content = sdt()
        content.append(hyperlink(d, 'Enter name')); p._p.append(el)
    elif case == 'hyperlink-around-sdt':
        p = d.add_paragraph('Name: '); link = hyperlink(d)
        el, content = sdt(placeholder=True); content.append(run('Enter name'))
        link.append(el); p._p.append(link)
    elif case == 'unlinked-sdt-positive':
        p = d.add_paragraph('Name: '); el, content = sdt(placeholder=True)
        content.append(run('Enter name')); p._p.append(el)
    elif case == 'linked-checkbox':
        p = d.add_paragraph(); p._p.append(hyperlink(d, '☐ Send copies'))
    elif case == 'separate-link-and-sdt-positive':
        p = d.add_paragraph(); p._p.append(hyperlink(d, 'Clinic instructions'))
        p.add_run(' Name: '); el, content = sdt(placeholder=True)
        content.append(run('Enter name')); p._p.append(el)
    else:
        raise ValueError(case)
    b = io.BytesIO(); d.save(b); return b.getvalue()


def links(raw):
    p = w.open_package(raw); tree = p.tree('word/document.xml')
    return [{'text': ''.join(h.itertext()), 'visible_text': ''.join(x.text or '' for x in h.iter(w.T)),
             'xml': w.etree.tostring(h).decode()} for h in tree.iter(w.q('w:hyperlink'))]


results = {'source': str(source), 'word_v3_sha256': hashlib.sha256((source/'rebuild/documents/word_v3.py').read_bytes()).hexdigest(), 'cases': []}
for case in ('block-placeholder-fixed-link', 'inline-linked-marker', 'hyperlink-around-sdt',
             'unlinked-sdt-positive', 'linked-checkbox', 'separate-link-and-sdt-positive'):
    raw = fixture(case); (outdir/(case+'.docx')).write_bytes(raw)
    result = {'case': case, 'before_links': links(raw), 'original_sha256': hashlib.sha256(raw).hexdigest()}
    try:
        info = w.inspect(raw)
        result.update(fields=info['fields'], notes=info['notes'])
        answers = {f['key']: (True if f['kind'] == 'checkbox' else 'Synthetic patient answer') for f in info['fields']}
        result['answers'] = answers
        filled = w.fill(raw, info['fields'], answers)
        (outdir/(case+'-filled.docx')).write_bytes(filled)
        result['after_links'] = links(filled)
        result['output_sha256'] = hashlib.sha256(filled).hexdigest()
        result['hyperlinks_changed'] = result['before_links'] != result['after_links']
    except Exception as exc:
        result['exception'] = {'type': type(exc).__name__, 'message': str(exc), 'code': getattr(exc, 'code', None)}
    results['cases'].append(result)
(outdir/'results.json').write_text(json.dumps(results, indent=2)+'\n')
print(json.dumps(results, indent=2))
