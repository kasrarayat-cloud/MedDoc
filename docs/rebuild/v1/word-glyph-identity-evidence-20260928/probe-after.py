from pathlib import Path
import io,json,hashlib
from docx import Document
from documents import word_v3 as w
from documents.test_word_v3 import _sdt as sdt,ff_field
R=Path(__file__).parent/'after';R.mkdir(exist_ok=False)
out=[]
for control in ('glyph','sdt','ff'):
 for after in (False,True):
  doc=Document();p=doc.add_paragraph('Choices? ')
  if control=='glyph':p.add_run('☐ Before: _____')
  elif control=='sdt':p._p.append(sdt('checkbox'));p.add_run(' Before: _____')
  else:ff_field(p,'checkbox');p.add_run(' Before: _____')
  if after:
   p.add_run(' ')
   if control=='glyph':p.add_run('☐ After')
   elif control=='sdt':p._p.append(sdt('checkbox'));p.add_run(' After')
   else:ff_field(p,'checkbox');p.add_run(' After')
  stream=io.BytesIO();doc.save(stream);raw=stream.getvalue();key=f'{control}-{after}'
  (R/(key+'.docx')).write_bytes(raw)
  detector=w.Detector(w.open_package(raw));slot=next(s for s in detector.index.slots if s.kind=='p');text=detector.ptext(slot.el)[0];start=text.index('_____')
  refs=w.place_box_refs(detector.index,slot)
  checkbox={'key':'c','kind':'multi_choice','label':'Choices','options':['Before','After'] if after else ['Before'],'locator':{'contract':'v3','part':slot.part,'slot':slot.path,'fingerprint':w.fingerprint(slot.el),'target':{'type':'checkboxes','boxes':refs}}}
  write={'key':'t','kind':'text','label':'Before details','locator':{'contract':'v3','part':slot.part,'slot':slot.path,'fingerprint':w.fingerprint(slot.el),'target':{'type':'span','start':start,'end':start+5,'shape':'blank'}}}
  try:
   filled=w.fill(raw,[checkbox,write],{'c':['Before','After'] if after else ['Before'],'t':'literal ☐ symbol retained'})
   (R/(key+'-filled.docx')).write_bytes(filled);new=w.SlotIndex(w.open_package(filled));states=w.option_states(new,checkbox)
   out.append({'case':key,'text':w.paragraph_text(new.by_id[slot.id].el)[0],'readback':states,'expected_printed_after_checked':after,'result':'exported'})
  except Exception as e:out.append({'case':key,'error':str(e),'type':type(e).__name__})
(R/'result.json').write_text(json.dumps({'engine_sha256':hashlib.sha256(Path(w.__file__).read_bytes()).hexdigest(),'constructed_disjoint_fields_not_upload_test':True,'cases':out},indent=2,ensure_ascii=False));print(json.dumps(out,ensure_ascii=False,indent=2))
