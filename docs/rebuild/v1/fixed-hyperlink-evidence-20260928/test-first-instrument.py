"""Fixed hyperlink text is not a replaceable answer hint."""
import io
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

from . import word_v3 as w


def link(paragraph, text):
    element = OxmlElement('w:hyperlink')
    element.set(qn('r:id'), paragraph.part.relate_to('mailto:forms@example.invalid',
                RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1')
    props.append(color); run.append(props)
    value = OxmlElement('w:t'); value.text = text; run.append(value)
    element.append(run); paragraph._p.append(element)


def data(document):
    out = io.BytesIO(); document.save(out); return out.getvalue()


class HyperlinkDestinations(unittest.TestCase):
    def test_colored_contact_is_never_an_answer_hint_in_common_parts(self):
        for position in ('body', 'table', 'header', 'footer'):
            with self.subTest(position=position):
                d = Document()
                if position == 'body': p = d.add_paragraph('Send the completed form to: ')
                elif position == 'table':
                    p = d.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
                    p.add_run('Send the completed form to: ')
                else:
                    p = getattr(d.sections[0], position).paragraphs[0]
                    p.add_run('Send the completed form to: ')
                link(p, 'forms@example.invalid')
                self.assertEqual(w.inspect(data(d))['fields'], [])

    def test_explicit_answer_marker_inside_link_is_visible_and_unassigned(self):
        for text in ('________', 'Enter name', '{{name}}'):
            with self.subTest(text=text):
                d = Document(); p = d.add_paragraph('Name: '); link(p, text)
                info = w.inspect(data(d))
                self.assertEqual(info['fields'], [])
                self.assertTrue(any('hyperlink' in message.lower() for message in info['notes']))

    def test_normal_blank_after_link_keeps_text_and_relationship(self):
        d = Document(); p = d.add_paragraph(); link(p, 'Clinic instructions')
        p.add_run(' Patient name: ________')
        raw = data(d); info = w.inspect(raw)
        self.assertEqual(len(info['fields']), 1)
        output = w.fill(raw, info['fields'], {info['fields'][0]['key']: 'Synthetic Example'})
        before = w.open_package(raw); after = w.open_package(output)
        self.assertEqual(before.read('word/_rels/document.xml.rels'), after.read('word/_rels/document.xml.rels'))
        for package in (before, after):
            self.assertEqual([n.text for h in package.tree('word/document.xml').iter(w.q('w:hyperlink'))
                              for n in h.iter(w.T)], ['Clinic instructions'])
        self.assertIn('Synthetic Example', w.Detector(after).text(after.tree('word/document.xml'), own=False))

    def test_old_mapping_cannot_replace_fixed_hyperlink(self):
        d = Document(); p = d.add_paragraph('Send the completed form to: ')
        link(p, 'forms@example.invalid'); raw = data(d)
        index = w.SlotIndex(w.open_package(raw)); slot = index.by_id['word/document.xml#b/p[0]']
        text, _ = w.paragraph_text(slot.el); start = text.index('forms@')
        field = {'key': 'old', 'kind': 'text', 'label': 'Send the completed form to',
                 'locator': {'contract': 'v3', 'part': slot.part, 'slot': slot.path,
                             'fingerprint': w.fingerprint(slot.el),
                             'target': {'type': 'span', 'start': start, 'end': len(text), 'shape': 'hint'}}}
        with self.assertRaisesRegex(w.WordError, 'hyperlink'):
            w.fill(raw, [field], {'old': 'Synthetic incorrect destination'})

    def test_insertion_in_or_at_end_of_hyperlink_refuses(self):
        for offset in (3, len('Clinic instructions')):
            with self.subTest(offset=offset):
                d = Document(); p = d.add_paragraph(); link(p, 'Clinic instructions')
                with self.assertRaisesRegex(w.WordError, 'hyperlink'):
                    w.replace_range(p._p, offset, offset, 'Synthetic answer')

    def test_unlinked_colored_explicit_hint_remains_fillable(self):
        d = Document(); p = d.add_paragraph('Name: '); run = p.add_run('Enter name')
        from docx.shared import RGBColor
        run.font.color.rgb = RGBColor.from_string('0563C1')
        raw = data(d); info = w.inspect(raw)
        self.assertEqual(len(info['fields']), 1)
        w.fill(raw, info['fields'], {info['fields'][0]['key']: 'Synthetic Example'})
