from pathlib import Path
import importlib.util,sys,json,subprocess,hashlib,io
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path('/opt/meddoc-pro/.worktrees/v1-notes-workflows/rebuild/.runtime/word-paired-independent-20260927')
BASE='4caa34431a60e6ce3409ecadbbb6989b50659d7b';FINAL='a921dfc312acbd6b9ad08b198720a08253481301'
modules={}
for name,ref in [('base',BASE),('candidate',FINAL)]:
 raw=subprocess.check_output(['git','show',ref+':rebuild/documents/word_v3.py'],cwd='/opt/meddoc-pro')
 p=ROOT/(name+'-word_v3.py');p.write_bytes(raw)
 spec=importlib.util.spec_from_file_location('documents.independent_'+name,p)
 mod=importlib.util.module_from_spec(spec);sys.modules[spec.name]=mod;spec.loader.exec_module(mod);modules[name]=mod

def make(kind):
 doc=Document();table=doc.add_table(rows=1,cols=2)
 cell=table.cell(0,0);cell.text='Past medical history:';blank=cell.add_paragraph()
 if kind=='paragraph_shading':
  shd=OxmlElement('w:shd');shd.set(qn('w:val'),'clear');shd.set(qn('w:fill'),'D9D9D9');blank._p.get_or_add_pPr().append(shd)
 elif kind=='paragraph_highlight':
  run=blank.add_run('        ');shade=OxmlElement('w:highlight');shade.set(qn('w:val'),'yellow');run._r.get_or_add_rPr().append(shade)
 elif kind=='numbered':blank.style='List Number'
 elif kind=='bullet':blank.style='List Bullet'
 elif kind=='paragraph_shading_inherited':
  style=doc.styles.add_style('Shaded Answer',WD_STYLE_TYPE.PARAGRAPH)
  shd=OxmlElement('w:shd');shd.set(qn('w:val'),'clear');shd.set(qn('w:fill'),'D9D9D9');style.element.get_or_add_pPr().append(shd);blank.style=style
 elif kind=='underlined':blank.add_run('         ').underline=True
 stream=io.BytesIO();doc.save(stream);return stream.getvalue()
result={'base':BASE,'candidate':FINAL,'source_hashes':{name:hashlib.sha256((ROOT/(name+'-word_v3.py')).read_bytes()).hexdigest() for name in modules},'cases':[]}
for kind in ['plain','paragraph_shading','paragraph_highlight','numbered','bullet','paragraph_shading_inherited','underlined']:
 source=make(kind);(ROOT/(kind+'.docx')).write_bytes(source);row={'kind':kind,'source_sha256':hashlib.sha256(source).hexdigest()}
 for name,mod in modules.items():
  out=mod.inspect(source)
  row[name]={'fields':[{'label':f['label'],'slot':f['locator']['slot'],'target':f['locator']['target']} for f in out['fields']], 'warnings':out.get('warnings'),'notes':out.get('notes')}
 result['cases'].append(row)
(ROOT/'probe.json').write_text(json.dumps(result,indent=2))
for row in result['cases']:print(row['kind'],len(row['base']['fields']),'->',len(row['candidate']['fields']),row['candidate']['warnings'])
