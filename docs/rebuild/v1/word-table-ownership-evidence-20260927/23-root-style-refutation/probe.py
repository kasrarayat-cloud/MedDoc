import sys,json,hashlib
from pathlib import Path
sys.path.insert(0,'/opt/meddoc-pro/.worktrees/v1-word-workflows/rebuild')
from documents import word_v3 as w
from documents.test_question_boxes import data
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
out=[]
for stylekind in ['character_underline','inherited_character_underline','unknown_character','default_table_underline','simple_spacing']:
 d=Document();t=d.add_table(rows=1,cols=2);t.cell(0,0).text='Date of Birth';p=t.cell(0,0).add_paragraph();r=p.add_run('        ')
 if 'character' in stylekind:
  st=d.styles.add_style('ClinicalBlank',WD_STYLE_TYPE.CHARACTER)
  if stylekind=='inherited_character_underline':
   base=d.styles.add_style('ClinicalBlankBase',WD_STYLE_TYPE.CHARACTER);base.font.underline=True;st.base_style=base
  elif stylekind=='character_underline':st.font.underline=True
  r.style=st
  if stylekind=='unknown_character':r._r.rPr.find(qn('w:rStyle')).set(qn('w:val'),'UnknownBlank')
 if stylekind=='default_table_underline':
  st=d.styles.add_style('DefaultClinicalTable',WD_STYLE_TYPE.TABLE);st.font.underline=True;st.element.set(qn('w:default'),'1')
 raw=data(d);result=w.inspect(raw)
 pth=Path('/opt/meddoc-pro/.worktrees/v1-retention-enforcement/rebuild/.runtime/word-spacing-root-refutation/date-label')/(stylekind+'.docx');pth.write_bytes(raw)
 out.append({'case':stylekind,'fields':[{'slot':f['locator']['slot'],'target':f['locator']['target']['type'],'label':f['label']} for f in result['fields']],'notes':result.get('notes'),'original_sha256':hashlib.sha256(raw).hexdigest()})
print(json.dumps(out,indent=2))
