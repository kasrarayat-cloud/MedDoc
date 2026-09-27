import sys,io,json,hashlib
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
ROOT=Path('/opt/meddoc-pro/.worktrees/v1-word-d19/rebuild');sys.path.insert(0,str(ROOT))
from docx import Document
from lxml import etree
from documents import word_v3 as w
from documents.test_signing_areas import signing_form
OUT=Path(__file__).resolve().parent

def save(doc):
    out=io.BytesIO();doc.save(out);return out.getvalue()
def parts(data):
    with ZipFile(io.BytesIO(data)) as z:return {n:z.read(n) for n in z.namelist()}
def signing_bands(container):
    p=container.add_paragraph()
    for i in range(3):
        if i:p.add_run('\t')
        p.add_run('\t\t\t').underline=True
    container.add_paragraph('Name of Participant\t\tDate\t\tSignature')
def result(name,data,field,answer):
    try:
        output=w.fill(data,[field],{field['key']:answer})
        row={'case':name,'outcome':'EXPORTED','bytes_changed':data!=output,'parts_changed':[k for k,v in parts(data).items() if parts(output)[k]!=v]}
    except w.V3Failure as exc:row={'case':name,'outcome':'REFUSED','code':exc.code,'message':str(exc)}
    return row
results=[]
data=signing_form(table=False);fields=w.inspect(data)['fields'];safe=next(f for f in fields if not f.get('write_protected'));protected=next(f for f in fields if f['label']=='Name of Participant')
for mode in ('flags-stripped','extra','copies','paragraph-options','checkbox-destination'):
    f=deepcopy(safe)
    if mode=='flags-stripped':
        f=deepcopy(protected);f.pop('write_protected',None);f.pop('review_reason',None);f['kind']='text';f['label']='Ordinary value'
    if mode=='extra':f['locator']['target']={'type':'paragraph'};f['locator']['extra']=[{k:protected['locator'][k] for k in ('part','slot','fingerprint')}]
    if mode=='copies':f['locator']['target']['copies']=[{k:protected['locator'][k] for k in ('part','slot','fingerprint')}]
    if mode=='paragraph-options':f.update(kind='choice',options=['Yes']);f['locator']['target']={'type':'paragraph_options','paragraphs':[protected['locator']['slot']],'printed':['Yes']}
    if mode=='checkbox-destination':f.update(kind='checkbox',options=['Checked']);f['locator']['target']={'type':'checkboxes','boxes':[{'kind':'glyph','part':protected['locator']['part'],'slot':protected['locator']['slot'],'glyph':0}]}
    results.append(result(mode,data,f,'Yes' if mode=='paragraph-options' else ('X\nY' if mode=='extra' else 'SYNTHETIC')))
    assert results[-1]['outcome']=='REFUSED' and results[-1]['code']=='protected_signing_area',results[-1]
# Alias to the protected caption rather than a band, through the real manual-location entrypoint.
idx=w.SlotIndex(w.open_package(data)); caption=next(s for s in idx.slots if s.kind=='p' and 'Name of Participant' in w.paragraph_text(s.el)[0])
manual=w.locate(data,[{'slot':caption.id,'label':'Synthetic caption alias','kind':'text'}])[0]
results.append(result('manual-caption-alias',data,manual,'SYNTHETIC'))
assert results[-1]['code']=='protected_signing_area'
# No floating boxes: ordinary short/long answers work and signing nodes are unchanged.
for name,value in [('ordinary-short','SYNTHETIC-1'),('ordinary-long','SYNTHETIC '*80)]:
    out=w.fill(data,fields,{safe['key']:value});before=etree.fromstring(parts(data)['word/document.xml']);after=etree.fromstring(parts(out)['word/document.xml'])
    assert [etree.tostring(x) for x in before.find(w.q('w:body'))][1:]==[etree.tostring(x) for x in after.find(w.q('w:body'))][1:]
    results.append({'case':name,'outcome':'EXPORTED','signing_regions_unchanged':True})
# Independently build protected captions in header and nested cell structures.
for shape in ('header','nested-table'):
    doc=Document();doc.add_paragraph('Record number: __________')
    if shape=='header':container=doc.sections[0].header
    else:container=doc.add_table(rows=1,cols=1).cell(0,0).add_table(rows=1,cols=1).cell(0,0)
    signing_bands(container);body=save(doc);fs=w.inspect(body)['fields'];hold=next(f for f in fs if f['label']=='Name of Participant');ordinary=next(f for f in fs if f['label']=='Record number')
    forged=deepcopy(ordinary);forged['locator']['target']['copies']=[{k:hold['locator'][k] for k in ('part','slot','fingerprint')}]
    results.append(result(shape+'-copy',body,forged,'SYNTHETIC'))
    assert results[-1]['code']=='protected_signing_area'
    if shape=='nested-table':
        table=hold['locator']['slot'].split('/tr[')[0];f=deepcopy(ordinary);f['locator']['target']={'type':'table_rows','table':table,'rows':[0],'columns':[0]}
        results.append(result('table-ancestor-alias',body,f,[['SYNTHETIC']]))
        assert results[-1]['code']=='protected_signing_area'
# A source-derived floating refusal survives removed flags/fields; missing export remains unchanged.
floating=signing_form();fs=w.inspect(floating)['fields'];f=deepcopy(next(f for f in fs if not f.get('write_protected')))
assert w.FLOATING_INITIALS_REFUSAL in w.inspect(floating)['notes']
for name,value in [('floating-short','SYNTHETIC'),('floating-long','SYNTHETIC '*80)]:
    results.append(result(name,floating,f,value));assert results[-1]['code']=='unqualified_floating_initials'
blank=w.fill(floating,[f],{});assert parts(blank)==parts(floating)
results.append({'case':'floating-blank','outcome':'EXPORTED','all_package_parts_unchanged':True})
print(json.dumps({'source_hashes':{name:hashlib.sha256((ROOT/name).read_bytes()).hexdigest() for name in ('documents/word_v3.py','core/v3_mapping.py','core/generation.py','core/views.py')},'cases':results},indent=2))
