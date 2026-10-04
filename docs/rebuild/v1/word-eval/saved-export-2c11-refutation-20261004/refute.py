"""Independent source-only adversarial inputs for finished2c11 verification."""
import copy
import hashlib
import io
import json
import socket
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from docx import Document

S=Path('/opt/meddoc-pro/.worktrees/v1-word-refute-saved-export-2c11e951')
assert subprocess.check_output(['git','rev-parse','HEAD'],cwd=S).decode().strip()=='2c11e951cc2e864fc387d77ff784ea33b21f2d14'
assert not subprocess.check_output(['git','status','--porcelain'],cwd=S).strip()
sys.path.insert(0,str(S/'rebuild'))
socket.socket.connect=lambda *a,**k:(_ for _ in ()).throw(AssertionError('offline'))
from documents import word_v3 as w
from core import v3_mapping
from ops import sonnet5_stages as st
B=Path(__file__).parent/'two-paragraph-final';B.mkdir(mode=0o700)
sha=lambda b:hashlib.sha256(b).hexdigest()
doc=Document();table=doc.add_table(rows=2,cols=2)
table.cell(0,0).text='Findings';table.cell(1,0).text='Plan'
doc.add_paragraph('Fixed synthetic neighboring text.')
stream=io.BytesIO();doc.save(stream);original=stream.getvalue()
fields=w.inspect(original)['fields'];assert len(fields)==2
saved={fields[0]['key']:{'status':'supported','text':'Alpha result.\nBeta follow-up.','reviewed':True},
       fields[1]['key']:{'status':'left_blank','text':'','reviewed':True}}
honest=w.fill(original,fields,v3_mapping.fill_answers(fields,saved))
# Independent valid representation of the same complete visible two-line
# scalar: Word paragraph boundaries, rather than a run's line-break node.
# This is the exact frozen false-negative shape the new branch handles.
stream=io.BytesIO()
with ZipFile(io.BytesIO(honest)) as source,ZipFile(stream,'w') as output:
    output.comment=source.comment
    for item in source.infolist():
        raw=source.read(item.filename)
        if item.filename=='word/document.xml':
            root=E.fromstring(raw)
            cell=root.find(w.q('w:body')+'/'+w.q('w:tbl')+'/'+w.q('w:tr')).findall(w.q('w:tc'))[1]
            for p in cell.findall(w.q('w:p')):cell.remove(p)
            for line in saved[fields[0]['key']]['text'].split('\n'):
                p=E.SubElement(cell,w.q('w:p'));r=E.SubElement(p,w.q('w:r'));t=E.SubElement(r,w.q('w:t'));t.text=line
            raw=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
        output.writestr(item,raw)
honest=stream.getvalue()
(B/'original.docx').write_bytes(original);(B/'fields.json').write_text(json.dumps(fields,indent=2)+'\n')
(B/'saved.json').write_text(json.dumps(saved,indent=2)+'\n');(B/'honest.docx').write_bytes(honest)
positive,reading=st.saved_export_verification(original,honest,fields,saved)
assert positive['workflow_success'] and reading['passed']
assert not positive['frozen_export_verification']['workflow_success']
key=fields[0]['key'];cases=[]
for name,hidden in [('hidden-all-answer-runs',True),('visible-answer-runs',False)]:
    result=io.BytesIO()
    with ZipFile(io.BytesIO(honest)) as src,ZipFile(result,'w') as out:
        out.comment=src.comment
        for item in src.infolist():
            raw=src.read(item.filename)
            if item.filename=='word/document.xml':
                root=E.fromstring(raw);index=w.SlotIndex(w.open_package(honest))
                slot=index.by_id[fields[0]['locator']['part']+'#'+fields[0]['locator']['slot']]
                source_path=slot.el.getroottree().getpath(slot.el)
                destination=root.xpath(source_path,namespaces=root.nsmap)[0]
                runs=list(destination.iter(w.q('w:r')));assert runs
                for run in runs:
                    properties=run.find(w.q('w:rPr'))
                    if properties is None:properties=E.Element(w.q('w:rPr'));run.insert(0,properties)
                    vanish=E.SubElement(properties,w.q('w:vanish'));vanish.set(w.q('w:val'),'1' if hidden else '0')
                raw=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
            out.writestr(item,raw)
    exported=result.getvalue();target=B/name;target.mkdir(mode=0o700)
    (target/'export.docx').write_bytes(exported)
    corrected,independent=st.saved_export_verification(original,exported,fields,saved)
    (target/'verification.json').write_text(json.dumps({'corrected':corrected,'read_back':independent},indent=2)+'\n')
    case={'case':name,'export_sha256':sha(exported),'all_answer_runs_explicitly_hidden':hidden,
          'frozen_pass':corrected['frozen_export_verification']['workflow_success'],
          'corrected_pass':corrected['workflow_success'],'read_back_pass':independent['passed'],
          'corrections':list(corrected['visible_break_readback_corrections']),
          'runs':len(runs),'exact_read_text':independent['fields'][key]['read']==saved[key]['text']}
    cases.append(case)
record={'commit':'2c11e951cc2e864fc387d77ff784ea33b21f2d14','source_sha256':sha((S/'rebuild/ops/sonnet5_stages.py').read_bytes()),
        'original_sha256':sha(original),'honest_sha256':sha(honest),'cases':cases,
        'finding':'Explicitly hidden multiline answer is accepted by the new correction. Both answer runs carry w:rPr/w:vanish w:val=1; exact text readers count their XML text despite the hidden formatting. Frozen criterion was false and new workflow_success true.',
        'FILE_LINES':['rebuild/ops/sonnet5_stages.py:604–610','rebuild/ops/sonnet5_stages.py:614–616','rebuild/documents/word_v3.py:10421–10429'],
        'network_DB_render_provider_calls':0,'native_Word':'not executed; hidden-formatting property directly established in valid OOXML',
        'source_clean_after':not subprocess.check_output(['git','status','--porcelain'],cwd=S).strip()}
(B/'result.json').write_text(json.dumps(record,indent=2)+'\n')
print(json.dumps(cases))
raise SystemExit(1 if cases[0]['corrected_pass'] else 0)
