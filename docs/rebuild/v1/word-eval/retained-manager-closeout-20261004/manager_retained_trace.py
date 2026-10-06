"""The one retained manager source on S2, source-only; no new corpus review."""
import hashlib
import io
import json
import socket
import sys
from pathlib import Path
from docx import Document
S=Path('/opt/meddoc-pro/.worktrees/v1-qualified-candidate-481b65ad')
sys.path.insert(0,str(S/'rebuild'))
socket.socket.connect=lambda *a,**k:(_ for _ in ()).throw(AssertionError('offline'))
from documents import word_v3 as w, engine
B=Path(__file__).parent
W=B.parent
P=W/'core-word-9c15c45952f9-1791076014635128282/qa/word-artifacts/manager-reserved-statement'
raw=(P/'original.docx').read_bytes()
out=B/'manager-retained-S2';out.mkdir(mode=0o700)
(out/'original.docx').write_bytes(raw)
legacy_doc,paragraphs,cells,legacy_fields,_=engine.docx_inspect(out/'original.docx')
roles=engine.legacy_roles(legacy_fields,legacy_fields,source=(raw,legacy_doc,paragraphs,cells))
assert roles['legacy_unresolved_actions']
assert all(r['printed_label']=='This section is for the line manager to add comments'
           and r['place'][0]=='word/document.xml#b/tbl[0]/tr[0]/tc[1]' for r in roles['legacy_unresolved_actions'])
info=w.inspect(raw);manager,plan=info['fields']
assert manager.get('write_protected') and not plan.get('write_protected')
ordinary=w.fill(raw,info['fields'],{plan['key']:'Rest.'})
(out/'ordinary-export.docx').write_bytes(ordinary)
doc=Document(io.BytesIO(ordinary))
assert doc.tables[0].cell(0,1).text=='' and doc.tables[0].cell(1,1).text=='Rest.'
forged=json.loads(json.dumps(manager));forged['label']='Clinical value';forged['kind']='text'
forged.pop('write_protected',None);forged.pop('review_reason',None)
try:
    wrong=w.fill(raw,[forged],{forged['key']:'Dr Synthetic Finalizer'})
except Exception as e:
    refusal={'type':type(e).__name__,'code':getattr(e,'code',None),'message':str(e)}
    assert refusal['code']=='protected_signing_area',refusal
else:
    (out/'unsafe-export.docx').write_bytes(wrong)
    raise AssertionError('Physical manager source was written')
sha=lambda b:hashlib.sha256(b).hexdigest()
record={'source_commit':'481b65adb8ffa97053b694f0de3ba4dbaac72d0a','retained_original':str(P/'original.docx'),
        'original_sha256':sha(raw),'fresh_fields':info['fields'],'legacy_fields':legacy_fields,
        'legacy_role_report':roles,'forged_field':forged,'refusal':refusal,
        'ordinary_export_sha256':sha(ordinary),'manager_answer_unchanged':True,'Plan_answer_complete':True,
        'scope':'One exact retained application manager original, current physical/legacy source guard; no DB, generation, corpus, provider, or native Word claim.'}
(out/'result.json').write_text(json.dumps(record,indent=2,default=str)+'\n')
print(json.dumps({'manager_protected':True,'legacy_unresolved_actions':len(roles['legacy_unresolved_actions']),
                  'ordinary_Plan_written':True,'forged_refusal':refusal['code']}))
