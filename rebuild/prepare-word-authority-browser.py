"""Build the SYNTHETIC blank legacy form and the expected values for word-authority-browser.mjs.

Run with the candidate venv python (python-docx), no Django and no network:
    <venv>/bin/python -B prepare-word-authority-browser.py
Writes (new files only, mode 0600, directory 0700):
    .runtime/word-authority-form.docx   a BLANK two-column legacy form: a heading, an instruction paragraph and
                                        eight question rows with empty answer cells (no patient data)
    .runtime/word-authority-input.json  labels, the synthetic notes, the expected ordinary values, the signing
                                        rows that must stay untouched, and the original's sha256
Known failure L01 (legacy default path): 'Patient signature' and 'Initials' are ordinary text fields for the
legacy inspector, so the labelled note lines fill them. The browser script asserts they are NOT filled.
The original eight-row case is retained as a refused, incomplete human-consent mapping. A separately
named seven-row control omits only that consent row and exercises the original signing/export assertions.
"""
import hashlib
import json
import os
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
OUT = HERE / '.runtime'
HEADING = 'Synthetic Legacy Intake Form'
INSTRUCTION = 'Use black ink.'
ROWS = ['Patient name', 'Chief complaint', 'Smoking status', 'Patient signature', 'Initials',
        'Consent form signed', 'Initial assessment', 'Date of visit']
NOTES = {'Patient name': 'Synthetic Example', 'Chief complaint': 'Cough for three days.',
         'Smoking status': 'Never smoker', 'Patient signature': 'SYNTHETIC-SIGNATURE-LEAK', 'Initials': 'SX',
         'Consent form signed': 'Yes', 'Initial assessment': 'Stable.', 'Date of visit': '2026-09-30'}
SIGNING = ['Patient signature', 'Initials']

def prepare_case(out, stem, rows, purpose):
    """Create a new blank original and hash-bound manifest, never replace either."""
    form = out / (stem + '-form.docx')
    manifest = out / (stem + '-input.json')
    if form.exists() or manifest.exists():
        raise FileExistsError('Refusing to replace a prepared Word authority case.')
    document = Document()
    document.core_properties.author = 'QuillMD synthetic fixture'
    document.core_properties.title = 'Synthetic legacy signing-row form'
    document.add_heading(HEADING, level=1)
    document.add_paragraph(INSTRUCTION)
    table = document.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    for label in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = ''
    with open(form, 'xb') as stream:
        document.save(stream)
    data = {
        'synthetic_only': True, 'purpose': purpose, 'form': form.name,
        'original_sha256': hashlib.sha256(form.read_bytes()).hexdigest(),
        'heading': HEADING, 'instruction': INSTRUCTION, 'rows': rows,
        'notes': '\n'.join(f'{label}: {NOTES[label]}' for label in rows),
        'ordinary': {label: NOTES[label] for label in rows if label not in SIGNING},
        'signing': {label: NOTES[label] for label in SIGNING},
    }
    with open(manifest, 'x') as stream:
        json.dump(data, stream, indent=2)
    return data


def prepare(out=OUT):
    os.umask(0o077)
    out.mkdir(mode=0o700, parents=True, exist_ok=True)
    stems = ('word-authority', 'word-authority-usable-control')
    # Check the entire output set before creating anything, including retained old eight-row cases.
    if any((out / (stem + suffix)).exists() for stem in stems
           for suffix in ('-form.docx', '-input.json')):
        raise FileExistsError('Refusing to replace prepared Word authority evidence.')
    original = prepare_case(out, stems[0], ROWS, 'retained-eight-row-human-consent-refusal')
    control = prepare_case(out, stems[1], [row for row in ROWS if row != 'Consent form signed'],
                           'separate-seven-row-signing-export-positive-control')
    return original, control


if __name__ == '__main__':
    prepare()
    print('Prepared retained eight-row refusal and separate seven-row signing/export control.')
