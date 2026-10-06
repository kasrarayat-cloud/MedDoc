"""Read persisted SOAP/custom examples and their actual served downloads.

Default is preparation only. Execute through the new demo's normal protected
launcher after its actual browser/manual clinical workflows. This does not
create, generate, review, finalize, charge, seed or edit a clinical record.
Authentication creates ordinary sessions/audit events. Fixture configuration
never proves a live model response, provider charge or human microphone test.
"""
import argparse
from collections import Counter
from hashlib import md5, sha256
from html import unescape
from io import BytesIO
import json
import os
from pathlib import Path
import re
import socket
import stat
import sys
import tarfile
import time
import uuid
from zipfile import ZipFile

from ops import verify_qa_authorization as auth
from ops import v1_private_install_preparation as preparation

LABELS = {'soap': ('Subjective', 'Objective', 'Assessment', 'Plan'),
          'custom': ('History', 'Findings', 'Plan')}
SCRATCH = auth.TASK / 'storage-resume-20261003T215613Z/operations'
# Exact static inputs from the pinned django-otp 1.7.3 dependency. These are
# public assets, not a permission to admit other untracked runtime content.
STATIC_DEPENDENCY_ASSETS = {
    'django_otp/css/login.css': ('django_otp/static/django_otp/css/login.css',
        'c46b11a02eb0bd7835fa26862a84efc9bf79bdc6661343796bae317eaf56de39'),
    'otp_totp/css/qrcode.css': ('django_otp/plugins/otp_totp/static/otp_totp/css/qrcode.css',
        '59b3b4964fe158beac514820ff47c1177c4105c972191ecb796d2c439660d757'),
    'otp_totp/js/qrcode.js': ('django_otp/plugins/otp_totp/static/otp_totp/js/qrcode.js',
        '943b25309bc5a4ad175bc09ad34347ed3148708d899d20588c7e4d8bd00bc32e'),
}


def digest(raw):
    return sha256(raw).hexdigest()


def validate_health(health, release, surface):
    auth.require(health == {'status': 'ok', 'data_mode': 'synthetic-only',
                           'surface': surface, 'release': release}, 'served_health_identity_differs')


def read_artifact(path, expected, maximum=15 * 1024 * 1024):
    """Read one original/result/export; never follow an arbitrary private link."""
    auth.require(re.fullmatch('[a-f0-9]{64}', expected or ''), 'missing_artifact_hash')
    raw = auth.protected_file(path)
    auth.require(len(raw) <= maximum and digest(raw) == expected, 'artifact_bytes_changed')
    return raw


def field_summary(answers):
    return dict(sorted(Counter(row.get('status', 'unrecorded') for row in answers.values()).items()))


def fresh_mfa_delay(last_counter, now):
    """Do not replay a code consumed by the preceding browser acceptance run."""
    auth.require(type(last_counter) is int and isinstance(now, (int, float)), 'invalid_authenticator_counter')
    delay = max(0, (last_counter + 1) * 30 - now + 0.2)
    auth.require(delay <= 32, 'authenticator_clock_or_counter_requires_review')
    return delay


def validate_example(record, original, first, exported, kind):
    """Validate only the two declared ordinary positive fixtures, preserving holds.

    The reader supplies stored data; this function does not independently prove
    those rows were created through HTTP. The separate route/browser receipts
    remain required, as do native Word and human microphone acceptance.
    """
    auth.require(kind in LABELS, 'unsupported_example_kind')
    fields = record['fields']; keys = [field['key'] for field in fields]
    auth.require(tuple(field['label'] for field in fields) == LABELS[kind]
        and len(set(keys)) == len(keys) and all(field.get('locator', {}).get('contract') == 'v3'
                                               for field in fields), 'positive_fixture_mapping_differs')
    auth.require(record['contract'] == 'v3' and record['mapping_confirmed']
        and record['source_sha256'] == record['original_sha256'] == digest(original), 'original_mapping_binding_failed')
    auth.require(record['status'] == 'finalized' and record['finalized_by'] == record['doctor']
        and bool(record['finalized_at']), 'clinician_finalization_missing')
    auth.require(record['department'] == record['patient_department'] == record['template_department']
        and record['job_encounter'] == record['encounter'] and record['job_user'] == record['doctor']
        and record['job_state'] == 'succeeded' and record['job_kind'] == 'generation', 'persisted_scope_or_job_failed')
    config = record['config']; profile = config.get('request_profile', {})
    auth.require(config.get('mode') == 'fixture' and config.get('provider') == 'anthropic'
        and config.get('model') == preparation.MODEL
        and config.get('ordinary_field_contract') == preparation.CONTRACT
        and profile.get('profile') == 'claude-sonnet-5-5-ordinary-v1'
        and profile.get('max_output_tokens') == 20000 and 'invalid' not in profile,
        'persisted_fixture_profile_failed')
    answers = record['answers']
    auth.require(set(answers) == set(keys) and isinstance(first, dict)
        and set(first) == set(keys), 'answer_catalogue_differs')
    useful = 0
    corrections = []
    blanks = []
    for key in keys:
        before = first[key]; after = answers[key]
        auth.require(isinstance(before, dict) and isinstance(after, dict), 'malformed_answer_record')
        if before.get('status') == 'supported' and before.get('text', '').strip():
            evidence = before.get('evidence') or []
            auth.require(bool(evidence), 'supported_answer_without_source')
            for span in evidence:
                start, end = span.get('start'), span.get('end')
                auth.require(type(start) is int and type(end) is int and 0 <= start < end <= len(record['notes'])
                    and record['notes'][start:end] == span.get('quote'), 'source_quote_differs')
            useful += 1
        auth.require(after.get('reviewed') is True and after.get('reviewed_by') == record['doctor'],
                     'explicit_clinician_review_missing')
        if after.get('status') == 'left_blank':
            auth.require(not after.get('text') and (after.get('blank_decision') or {}).get('by') == record['doctor'],
                         'reviewed_blank_decision_missing')
            blanks.append(key)
        if after.get('status') != 'left_blank' and (after.get('text') != before.get('text') or after.get('status') == 'edited'):
            corrections.append(key)
    auth.require(useful > 0, 'zero_useful_first_pass_answers')
    usage = record['usage']
    auth.require(usage['simulated'] is True and usage['actual_cost'] == usage['charged_cost'] == '0'
        and record['receipt_matches_usage'] and not record['held_reservations'], 'fixture_usage_or_reservation_failed')
    auth.require(digest(exported) == record['export_sha256'], 'final_download_hash_differs')
    from docx import Document
    before_doc, after_doc = Document(BytesIO(original)), Document(BytesIO(exported))
    expected = [(field['label'], answers[field['key']].get('text', '')) for field in fields]
    auth.require(len(before_doc.tables) == len(after_doc.tables) == 1
        and [(row.cells[0].text, row.cells[1].text) for row in before_doc.tables[0].rows]
            == [(label, '') for label in LABELS[kind]]
        and [(row.cells[0].text, row.cells[1].text) for row in after_doc.tables[0].rows] == expected
        and [paragraph.text for paragraph in before_doc.paragraphs] == [paragraph.text for paragraph in after_doc.paragraphs],
        'positive_fixture_export_fidelity_failed')
    with ZipFile(BytesIO(original)) as before_zip, ZipFile(BytesIO(exported)) as after_zip:
        auth.require(set(before_zip.namelist()) == set(after_zip.namelist())
            and all(before_zip.read(name) == after_zip.read(name) for name in before_zip.namelist()
                    if name != 'word/document.xml'), 'neighbor_package_bytes_changed')
    return {'first_pass_statuses': field_summary(first), 'final_statuses': field_summary(answers),
            'useful_first_pass_answers': useful, 'clinician_corrections': corrections,
            'reviewed_blank_keys': blanks, 'supported_layout': 'declared two-column growing answer cells',
            'native_word_inspection': False, 'human_microphone_acceptance': False,
            'live_model_qualification': False, 'invoice_reconciliation': False}


def compiled_static_file(path, root):
    """Read a bounded public collectstatic output without following links."""
    auth.require(path.is_relative_to(root) and not any(p.is_symlink() for p in (path, *path.parents)),
                 'compiled_static_link_refused')
    before = path.lstat()
    auth.require(stat.S_ISREG(before.st_mode) and before.st_uid == os.getuid()
        and before.st_nlink == 1 and stat.S_IMODE(before.st_mode) in (0o600, 0o644)
        and before.st_size <= 2 * 1024 * 1024, 'compiled_static_file_metadata_differs')
    descriptor = os.open(path, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK)
    with os.fdopen(descriptor, 'rb') as stream:
        raw = stream.read(2 * 1024 * 1024 + 1)
        after = os.fstat(stream.fileno())
    current = path.lstat()
    identity = ('st_dev', 'st_ino', 'st_uid', 'st_mode', 'st_nlink', 'st_size', 'st_mtime_ns', 'st_ctime_ns')
    auth.require(len(raw) == before.st_size and all(getattr(before, key) == getattr(value, key)
        for value in (after, current) for key in identity), 'compiled_static_file_changed')
    return raw


def compiled_static_identity(root, archived):
    """Admit only genuine current collectstatic assets, hashes and compression.

    The five current CSS/JS inputs need no URL/import rewriting. If that changes,
    qualification must extend this strict byte comparison; it never guesses at
    generated source or accepts an arbitrary static tree.
    """
    from importlib.metadata import distribution
    from whitenoise.compress import Compressor
    generated = root / 'staticfiles'
    auth.require(not any(name.startswith('rebuild/staticfiles/') for name in archived),
                 'compiled_static_is_not_archived_source')
    requirements = (root / 'requirements.txt').read_text()
    for name, version in (('Django', '5.2.17'), ('django-otp', '1.7.3'), ('whitenoise', '6.12.0')):
        auth.require(re.search(r'^' + re.escape(name) + '==' + re.escape(version) + r'$', requirements, re.M)
            and distribution(name).version == version, 'compiled_static_dependency_version_differs')
    app_names = {name.removeprefix('rebuild/static/') for name in archived if name.startswith('rebuild/static/')}
    auth.require(app_names == {'app.css', 'app.js'}, 'compiled_static_source_catalogue_differs')
    inputs = {name: (root / 'static' / name).read_bytes() for name in app_names}
    dependency = distribution('django-otp')
    recorded = {str(path): path for path in dependency.files or []}
    for name, (relative, expected) in STATIC_DEPENDENCY_ASSETS.items():
        import base64
        path = recorded.get(relative)
        auth.require(path is not None and path.hash and path.hash.mode == 'sha256',
                     'compiled_static_dependency_record_missing')
        raw = Path(dependency.locate_file(path)).read_bytes()
        auth.require(digest(raw) == expected and base64.urlsafe_b64encode(sha256(raw).digest()).decode().rstrip('=')
            == path.hash.value, 'compiled_static_dependency_asset_differs')
        inputs[name] = raw
    expected_files, paths = {}, {}
    compressor = Compressor(quiet=True)
    for name, raw in inputs.items():
        path = Path(name)
        hashed = str(path.with_name(path.stem + '.' + md5(raw, usedforsecurity=False).hexdigest()[:12] + path.suffix))
        paths[name] = hashed
        for output in (name, hashed):
            expected_files[output] = raw
            if compressor.use_brotli:
                compressed = compressor.compress_brotli(raw)
                auth.require(len(compressed) <= len(raw) * 0.95, 'compiled_static_brotli_not_qualified')
                expected_files[output + '.br'] = compressed
            compressed = compressor.compress_gzip(raw)
            auth.require(len(compressed) <= len(raw) * 0.95, 'compiled_static_gzip_not_qualified')
            expected_files[output + '.gz'] = compressed
    manifest = {'paths': paths, 'version': '1.1',
        'hash': md5(json.dumps(sorted(paths.items())).encode(), usedforsecurity=False).hexdigest()[:12]}
    manifest_raw = compiled_static_file(generated / 'staticfiles.json', generated)
    try:
        actual_manifest = json.loads(manifest_raw)
    except (ValueError, UnicodeDecodeError):
        raise RuntimeError('compiled_static_manifest_invalid') from None
    auth.require(actual_manifest == manifest and json.dumps(actual_manifest).encode() == manifest_raw,
                 'compiled_static_manifest_differs')
    # Finder traversal order may change JSON key order. Its complete logical
    # mapping/hash and the normal serializer are checked independently above.
    expected_files['staticfiles.json'] = manifest_raw
    actual_files, actual_dirs = {}, set()
    for directory, names, files in os.walk(generated, followlinks=False):
        parent = Path(directory)
        info = parent.lstat()
        auth.require(stat.S_ISDIR(info.st_mode) and not parent.is_symlink() and info.st_uid == os.getuid()
            and stat.S_IMODE(info.st_mode) in (0o700, 0o755), 'compiled_static_directory_metadata_differs')
        auth.require(all(not (parent / name).is_symlink() for name in names), 'compiled_static_directory_link_refused')
        actual_dirs.add(str(parent.relative_to(generated)))
        auth.require(len(actual_dirs) <= 16 and len(actual_files) + len(files) <= 32,
                     'compiled_static_catalogue_bound')
        for name in files:
            path = parent / name
            actual_files[str(path.relative_to(generated))] = compiled_static_file(path, generated)
    expected_dirs = {'.'} | {str(parent) for name in expected_files for parent in Path(name).parents}
    auth.require(actual_dirs == expected_dirs and actual_files == expected_files, 'compiled_static_inventory_or_bytes_differ')
    return {name: digest(raw) for name, raw in sorted(actual_files.items())}


def archive_identity(demo, expected):
    auth.require(re.fullmatch('[a-f0-9]{40}', expected or ''), 'invalid_release_identity')
    auth.require(demo == demo.resolve() and demo.parent == auth.TASK
        and demo.name == 'demo-' + expected[:8] and expected[:8] != '789c0372', 'new_demo_scope_required')
    prepared = json.loads(auth.protected_file(demo / 'prepared.json'))
    auth.require(prepared['commit'] == expected and prepared['root'] == str(demo)
        and prepared['database'] == 'quillmd_qa_int_sol_' + expected[:8]
        and prepared['ports'] == [9160, 9161] and prepared['ai_mode'] == 'fixture'
        and prepared['field_relations'] is True and prepared['role_suggestions'] is False,
        'prepared_identity_differs')
    raw = auth.protected_file(demo / 'source.tar')
    auth.require(digest(raw) == prepared['source_tar_sha256'], 'archive_hash_differs')
    seen = {}
    with tarfile.open(fileobj=BytesIO(raw)) as archive:
        auth.require(archive.pax_headers.get('comment') == expected, 'archive_release_differs')
        for member in archive.getmembers():
            path = Path(member.name)
            auth.require(not path.is_absolute() and '..' not in path.parts and path.parts[0] == 'rebuild'
                and (member.isfile() or member.isdir()), 'archive_boundary_failed')
            if not member.isfile():
                continue
            actual = demo / 'src' / path
            info = actual.lstat()
            auth.require(member.name not in seen and not any(item.is_symlink() for item in (actual, *actual.parents))
                and stat.S_ISREG(info.st_mode) and info.st_nlink == 1 and info.st_uid == os.getuid()
                and stat.S_IMODE(info.st_mode) in (0o600, 0o700), 'source_file_changed_or_linked')
            seen[member.name] = digest(archive.extractfile(member).read())
            auth.require(digest(actual.read_bytes()) == seen[member.name], 'served_source_bytes_changed')
    auth.require(seen == prepared['runtime_source_sha256'], 'complete_archive_inventory_differs')
    actual_files = set()
    root = demo / 'src/rebuild'
    if os.path.lexists(root / 'staticfiles'):
        compiled_static_identity(root, seen)
    for directory, names, files in os.walk(root, followlinks=False):
        # Compiled public assets were independently verified above. The existing
        # seed/browser writes only the separate top-level private state.
        if Path(directory) == root:
            names[:] = [name for name in names if name not in ('.runtime', 'staticfiles')]
        auth.require(all(not (Path(directory) / name).is_symlink() for name in names), 'untracked_source_directory_linked')
        for name in files:
            actual_files.add(str((Path(directory) / name).relative_to(demo / 'src')))
    auth.require(actual_files == set(seen), 'untracked_runtime_source_refused')
    return prepared


def main():
    os.umask(0o077)
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--demo-dir', type=Path); parser.add_argument('--expected-commit')
    for name in ('doctor-pid', 'admin-pid', 'worker-pid'):
        parser.add_argument('--' + name, type=int)
    for name in ('soap-encounter', 'soap-job', 'custom-encounter', 'custom-job'):
        parser.add_argument('--' + name, type=uuid.UUID)
    parser.add_argument('--run-id'); parser.add_argument('--execute', action='store_true')
    args = parser.parse_args()
    if not args.execute:
        print(json.dumps({'prepared_only': True, 'clinical_mutations': 0, 'provider_requests': 0,
            'max_requests': auth.MAX_REQUESTS, 'deadline_seconds': auth.DEADLINE_SECONDS,
            'scope': 'Two actual finalized fixture examples; ordinary login/session/audit events only'})); return
    auth.require(all((args.demo_dir, args.expected_commit, args.doctor_pid, args.admin_pid, args.worker_pid,
        args.soap_encounter, args.soap_job, args.custom_encounter, args.custom_job, args.run_id)), 'complete_execution_identity_required')
    auth.require(re.fullmatch('[a-z0-9][a-z0-9-]{1,55}', args.run_id), 'distinct_run_identity_required')
    demo = args.demo_dir; prepared = archive_identity(demo, args.expected_commit)
    source = demo / 'src/rebuild'; database = prepared['database']
    auth.require(Path.cwd() == source and not os.path.lexists(source / '.env'), 'normal_demo_launcher_required')
    processes = [auth.process_identity(args.doctor_pid, source, args.expected_commit, database, 'doctor', 9160),
        auth.process_identity(args.admin_pid, source, args.expected_commit, database, 'admin', 9161),
        auth.process_identity(args.worker_pid, source, args.expected_commit, database, 'doctor')]
    connect = socket.socket.connect
    def local(sock, address):
        auth.require(isinstance(address, tuple) and address[0] == '127.0.0.1' and address[1] in (55731, 9160, 9161),
                     'out_of_scope_network_connection')
        return connect(sock, address)
    socket.socket.connect = local
    sys.path.insert(0, str(source)); os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'quillmd.settings')
    import django
    django.setup()
    from django.conf import settings
    from django.db import connection
    from django_otp.oath import totp
    from django_otp.plugins.otp_totp.models import TOTPDevice
    from core import models as m
    from core.security import private_path
    from core.soap_document import original_bytes
    auth.require(settings.SYNTHETIC_ONLY and settings.AI_MODE == 'fixture' and settings.RELEASE_SHA == args.expected_commit
        and settings.V3_FIELD_RELATIONS is True and settings.AI_MODEL == preparation.MODEL
        and settings.QUILLMD_WORD_ENGINE == 'auto'
        and settings.DATABASES['default']['NAME'] == database
        and settings.DATABASES['default']['HOST'] == '127.0.0.1'
        and str(settings.DATABASES['default']['PORT']) == '55731', 'launcher_profile_differs')
    with connection.cursor() as cursor:
        cursor.execute("SET lock_timeout = '3s'"); cursor.execute("SET statement_timeout = '8s'")
        cursor.execute('SET default_transaction_read_only = on')
    out = SCRATCH / ('served-persisted-' + args.run_id)
    auth.require(not out.exists(), 'retained_run_identity_not_reusable'); out.mkdir(mode=0o700)
    served = auth.Served(); started = time.monotonic()
    result = {'commit': args.expected_commit, 'database': database, 'processes': processes,
        'complete': False, 'clinical_mutations': 0, 'provider_requests': 0, 'examples': {},
        'application_configuration': {'ai_mode': settings.AI_MODE, 'model': settings.AI_MODEL,
            'word_engine': settings.QUILLMD_WORD_ENGINE, 'field_relations': settings.V3_FIELD_RELATIONS,
            'live_model_qualification': False}, 'served_health': {}}
    try:
        accounts = json.loads(auth.protected_file(source / '.runtime/private-access.json'))
        auth.require(all(accounts[name]['email'].endswith('@synthetic.quillmd.invalid')
            for name in ('doctor', 'admin', 'ultrasound')), 'synthetic_seed_identity_required')
        doctor = m.User.objects.get(email=accounts['doctor']['email'])
        admin = m.User.objects.get(email=accounts['admin']['email'])
        foreign = m.User.objects.get(email=accounts['ultrasound']['email'])
        auth.require(admin.is_staff and not m.Membership.objects.filter(user=admin, active=True).exists(), 'financial_only_admin_required')
        doctor_client = served.login(accounts['doctor'], 9160)
        foreign_client = served.login(accounts['ultrasound'], 9160)
        # The existing served authorization gate owns the password-only
        # negative. Repeating it here would throttle the same seed account.
        # Its/browser's successful code also cannot be replayed in this run.
        device = TOTPDevice.objects.get(user=admin, confirmed=True)
        auth.require(device.step == 30 and device.t0 == 0 and device.digits == 6
            and device.key == accounts['admin']['totp_key_hex'], 'seed_authenticator_identity_differs')
        pause = fresh_mfa_delay(device.last_t, time.time())
        if pause:
            time.sleep(pause)
        admin_client = served.login(accounts['admin'], 9161,
            otp=str(totp(bytes.fromhex(accounts['admin']['totp_key_hex']), step=30, t0=0, digits=6)).zfill(6))
        for surface, port in (('doctor', 9160), ('admin', 9161)):
            health = json.loads(served.expect(served.client(), port, '/health/', 200, label=surface + '_health'))
            validate_health(health, args.expected_commit, surface)
            result['served_health'][surface] = health
        financial_pages = [served.expect(admin_client, 9161, '/console/', 200, label='financial_admin_positive'),
            served.expect(admin_client, 9161, '/console/services/', 200, label='configured_services_positive')]
        served.expect(doctor_client, 9160, '/console/', 404, label='doctor_console_denied')
        for kind in LABELS:
            auth.require(time.monotonic() - started < auth.DEADLINE_SECONDS, 'persisted_read_deadline')
            encounter_id, job_id = getattr(args, kind + '_encounter'), getattr(args, kind + '_job')
            e = m.Encounter.objects.select_related('template_version__template', 'patient').get(pk=encounter_id)
            job = m.Job.objects.get(pk=job_id); version = e.template_version; template = version.template
            auth.require(e.created_by_id == doctor.pk and m.Membership.objects.filter(user=doctor, department=e.department,
                active=True, role__in=('owner', 'clinician')).exists()
                and not m.Membership.objects.filter(user=foreign, department=e.department, active=True).exists(), 'positive_foreign_scope_missing')
            original = read_artifact(private_path(template.original_path), template.sha256)
            if kind == 'soap':
                auth.require(original == original_bytes(), 'application_soap_original_differs')
            raw = read_artifact(private_path(job.result_path), job.result_sha256, 2 * 1024 * 1024)
            saved = json.loads(raw); auth.require(set(saved) == {'answers'}, 'fixture_original_answers_required')
            exported = read_artifact(private_path(e.export_path), e.export_sha256)
            usage = m.Usage.objects.get(job=job); receipt = m.ProviderReceipt.objects.get(job=job)
            columns = ('provider', 'provider_request_id', 'provider_resource_id', 'actual_cost', 'units', 'rates', 'simulated')
            record = {'encounter': str(e.pk), 'doctor': str(doctor.pk), 'notes': e.notes,
                'department': str(e.department_id), 'patient_department': str(e.patient.department_id),
                'template_department': str(template.department_id), 'template': str(template.pk),
                'mapping': str(version.pk), 'mapping_number': version.number, 'fields': version.fields,
                'contract': version.contract, 'mapping_confirmed': version.confirmed is not None,
                'source_sha256': version.import_manifest.get('source_sha256'), 'original_sha256': template.sha256,
                'status': e.status, 'answers': e.answers, 'finalized_by': str(e.finalized_by_id),
                'finalized_at': e.finalized_at.isoformat() if e.finalized_at else None,
                'export_sha256': e.export_sha256, 'job': str(job.pk), 'job_encounter': str(job.encounter_id),
                'job_user': str(job.user_id), 'job_kind': job.kind, 'job_state': job.state, 'job_revision': job.revision,
                'config': job.config, 'usage': {'simulated': usage.simulated, 'actual_cost': str(usage.actual_cost.normalize()),
                    'charged_cost': str(usage.charged_cost.normalize())},
                'receipt_matches_usage': all(getattr(receipt, name) == getattr(usage, name) for name in columns),
                'held_reservations': m.Reservation.objects.filter(job=job, state='held').count(),
                'attempts': list(m.Job.objects.filter(encounter=e).order_by('created_at').values('id', 'kind', 'state', 'provider'))}
            def financial_fingerprint():
                data = {'usage': list(m.Usage.objects.filter(job=job).values()),
                        'receipt': list(m.ProviderReceipt.objects.filter(job=job).values()),
                        'reservations': list(m.Reservation.objects.filter(job=job).order_by('id').values())}
                return digest(json.dumps(data, sort_keys=True, default=str).encode())
            before_finance = financial_fingerprint()
            before_manifest = json.dumps(version.import_manifest, sort_keys=True)
            before_fingerprint = digest(json.dumps(record, sort_keys=True, default=str).encode())
            # Preserve the first actual input/output before any usefulness or
            # fidelity assertion. A failure must retain the original case.
            for name, value in ((kind + '-original.docx', original), (kind + '-first-result.json', raw),
                (kind + '-final-export.docx', exported), (kind + '-persisted-record.json', record)):
                auth.write_new(out / name, value)
            summary = validate_example(record, original, saved['answers'], exported, kind)
            served.expect(doctor_client, 9160, f'/encounters/{e.pk}/', 200, label=kind + '_doctor_positive')
            actual = served.expect(doctor_client, 9160, f'/encounters/{e.pk}/export/', 200, label=kind + '_served_download')
            auth.require(actual == exported, 'served_download_differs_from_final_file')
            for route in (f'/encounters/{e.pk}/', f'/encounters/{e.pk}/export/',
                          f'/templates/{template.pk}/original/'):
                served.expect(foreign_client, 9160, route, 404, label=kind + '_foreign_department_denied')
                served.expect(admin_client, 9161, route, 404, label=kind + '_admin_clinical_denied')
            page = served.expect(admin_client, 9161, f'/console/jobs/{job.pk}/', 200, label=kind + '_admin_charge_positive')
            for body in (*financial_pages, page):
                text = unescape(body.decode(errors='replace'))
                clinical = [e.notes, *(row.get('text', '') for row in e.answers.values())]
                auth.require(all(value not in text for value in clinical if len(value) >= 8), 'clinical_text_on_financial_page')
            e.refresh_from_db(); job.refresh_from_db(); version.refresh_from_db(); usage.refresh_from_db(); receipt.refresh_from_db()
            auth.require(e.answers == record['answers'] and e.notes == record['notes'] and e.status == record['status']
                and version.fields == record['fields'] and json.dumps(version.import_manifest, sort_keys=True) == before_manifest
                and job.config == record['config'] and job.state == record['job_state']
                and financial_fingerprint() == before_finance
                and read_artifact(private_path(template.original_path), template.sha256) == original
                and read_artifact(private_path(e.export_path), e.export_sha256) == exported
                and all(getattr(receipt, name) == getattr(usage, name) for name in columns)
                and m.Reservation.objects.filter(job=job, state='held').count() == record['held_reservations'], 'denied_reads_changed_persisted_state')
            result['examples'][kind] = {**summary, 'encounter': str(e.pk), 'job': str(job.pk),
                'persisted_record_sha256': before_fingerprint, 'first_result_sha256': job.result_sha256,
                'original_sha256': template.sha256, 'export_sha256': e.export_sha256,
                'financial_fingerprint_before_and_after': before_finance}
        auth.require(archive_identity(demo, args.expected_commit) == prepared, 'source_changed_during_verification')
        auth.require(processes == [auth.process_identity(args.doctor_pid, source, args.expected_commit, database, 'doctor', 9160),
            auth.process_identity(args.admin_pid, source, args.expected_commit, database, 'admin', 9161),
            auth.process_identity(args.worker_pid, source, args.expected_commit, database, 'doctor')], 'process_changed_during_verification')
        result.update(complete=True, requests=served.requests, elapsed_seconds=time.monotonic() - started)
    except BaseException as error:
        import traceback
        auth.write_new(out / 'original-failure.txt', traceback.format_exc().encode())
        result.update(complete=False, failure_kind=type(error).__name__, requests=served.requests,
                      elapsed_seconds=time.monotonic() - started)
        raise RuntimeError('Fixture persisted verification failed; preserve protected original evidence.') from None
    finally:
        auth.write_new(out / 'http-records.json', served.records)
        auth.write_new(out / 'result.json', result)
        connection.close()
    print(json.dumps(result, sort_keys=True, default=str))


if __name__ == '__main__':
    main()
