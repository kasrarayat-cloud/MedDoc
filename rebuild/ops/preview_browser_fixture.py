"""Synthetic fixture control for preview browser tests; refuses the real DB.

Sensitive test values travel in protected files and captured stdin/stdout pipes,
never process arguments or browser reports. This is not inbox-delivery evidence.
"""
import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import secrets
import stat
import sys


def require(value, message):
    if not value:
        raise RuntimeError(message)


def private_write(path, payload):
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600)
    with os.fdopen(fd, 'w') as stream:
        json.dump(payload, stream, indent=2)
        stream.write('\n')


def context():
    import django
    django.setup()
    from django.conf import settings
    from django.db import connection
    require(settings.SETTINGS_MODULE == 'ops.preview_browser_settings'
            and connection.settings_dict['NAME'].startswith('quillmd_preview_browser_')
            and connection.settings_dict['PORT'] == '55439'
            and connection.settings_dict['HOST'] == '127.0.0.1'
            and connection.settings_dict['USER'] == 'quillmd_ci'
            and settings.AI_MODE == 'fixture' and settings.QUILLMD_EMAIL_MODE == 'file'
            and not settings.RESEND_API_KEY and not settings.STRIPE_SECRET_KEY
            and settings.SYNTHETIC_ONLY, 'Only the dedicated synthetic preview browser database is permitted.')
    state = Path(os.environ['QUILLMD_PREVIEW_BROWSER_STATE'])
    meta = state.lstat()
    require(state.resolve() == state and '.runtime' in state.parts
            and stat.S_ISDIR(meta.st_mode) and stat.S_IMODE(meta.st_mode) == 0o700
            and meta.st_uid == os.getuid(), 'Browser state must be a canonical operator-owned0700 runtime subdirectory.')
    return settings, state


def prepare(settings, state):
    from core.models import User, Organization, Department, Membership, Patient, Template, TemplateVersion, Encounter
    from core.billing import welcome_credit
    from django.utils import timezone
    from docx import Document
    require(User.objects.count() == 0, 'Prepare only a new empty isolated fixture database.')
    control = {'origin': settings.QUILLMD_BROWSER_TEST_ORIGIN,
               'account_email_origin': settings.QUILLMD_APP_BASE_URL, 'state': str(state),
               'mail_evidence': 'isolated file outbox; no provider request or inbox receipt',
               'owner': {'email': 'preview-owner@synthetic.quillmd.invalid', 'password': secrets.token_urlsafe(28)},
               'invitees': {}, 'fixtures': {}}
    for browser in ('chrome', 'edge'):
        control['invitees'][browser] = {'email': f'preview-{browser}@synthetic.quillmd.invalid',
                                      'password': secrets.token_urlsafe(28), 'replacement': secrets.token_urlsafe(28)}
    owner = User.objects.create_user(username=control['owner']['email'], email=control['owner']['email'],
                                    password=control['owner']['password'], verified_at=timezone.now(),
                                    first_name='Synthetic', last_name='Preview Owner')
    welcome_credit(owner)
    org = Organization.objects.create(name='Synthetic preview browser organization')
    other = Organization.objects.create(name='Synthetic other browser organization')
    for kind, organization, name in [('own', org, 'X-ray'), ('ultrasound', org, 'Ultrasound'), ('other', other, 'X-ray')]:
        department = Department.objects.create(organization=organization, name=name)
        if kind == 'own':
            Membership.objects.create(user=owner, department=department, role='owner')
        patient = Patient.objects.create(department=department, display_name='SYNTHETIC PREVIEW ' + kind.upper(), reference='PREVIEW-' + kind, created_by=owner)
        original = f'originals/{kind}.docx'
        exported = f'exports/{kind}.docx'
        (settings.MEDIA_ROOT / 'originals').mkdir(parents=True, exist_ok=True, mode=0o700)
        (settings.MEDIA_ROOT / 'exports').mkdir(parents=True, exist_ok=True, mode=0o700)
        document = Document()
        document.add_paragraph('Synthetic preview access fixture. No clinical workflow or provider call is represented.')
        document.add_paragraph('History: synthetic browser test.')
        document.save(settings.MEDIA_ROOT / original)
        blob = (settings.MEDIA_ROOT / original).read_bytes()
        (settings.MEDIA_ROOT / original).chmod(0o600)
        (settings.MEDIA_ROOT / exported).write_bytes(blob)
        (settings.MEDIA_ROOT / exported).chmod(0o600)
        digest = hashlib.sha256(blob).hexdigest()
        template = Template.objects.create(department=department, name='Synthetic preview ' + kind, format='docx',
                                           original_name=kind + '.docx', original_path=original, sha256=digest, created_by=owner)
        version = TemplateVersion.objects.create(template=template, number=1, fields=[], created_by=owner)
        encounter = Encounter.objects.create(department=department, patient=patient, template_version=version, created_by=owner,
                                             service_date=timezone.localdate(), notes='Synthetic access fixture only.', status='finalized',
                                             finalized_by=owner, finalized_at=timezone.now(), export_path=exported, export_sha256=digest)
        control['fixtures'][kind] = {'department_id': str(department.pk), 'patient_id': str(patient.pk),
                                     'template_id': str(template.pk), 'encounter_id': str(encounter.pk), 'file_sha256': digest}
    private_write(state / 'control.json', control)
    return {'prepared': True, 'synthetic_only': True, 'provider_calls': 0, 'mail_sends': 0}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--action', required=True, choices=['prepare', 'mail', 'state', 'reset-limits'])
    args = parser.parse_args()
    payload = json.loads(sys.stdin.read() or '{}')
    settings, state = context()
    if args.action == 'prepare':
        result = prepare(settings, state)
    else:
        from core.models import User, Membership, Wallet, LedgerEntry, AuditEvent, LoginAttempt
        control = json.loads((state / 'control.json').read_text())
        if args.action == 'mail':
            actor = payload['browser']
            require(actor in control['invitees'], 'Unknown isolated browser recipient.')
            purpose = payload['purpose']
            require(purpose in ('invitation', 'recovery'), 'Only the two account-flow messages are inspected.')
            recipient = control['invitees'][actor]['email']
            route = 'invite' if purpose == 'invitation' else 'recovery/confirm'
            from email import policy
            from email.parser import BytesParser
            links = []
            for path in sorted(settings.EMAIL_FILE_PATH.glob('*.log'), key=lambda item: item.stat().st_mtime_ns):
                for raw in path.read_bytes().split(b'\n' + b'-' * 79 + b'\n'):
                    message = BytesParser(policy=policy.default).parsebytes(raw.strip())
                    if str(message.get('To', '')).strip() != recipient:
                        continue
                    content = message.get_content()
                    match = re.search(re.escape(control['account_email_origin']) + '/' + route + r'/[A-Za-z0-9_/-]+/', content)
                    if match:
                        links.append(match.group(0))
            require(len(links) == 1, 'Expected exactly one isolated account email for this browser and purpose.')
            # Account email links use the real canonical origin. Browser
            # navigation maps only that known path to this isolated listener;
            # neither a delivery receipt nor public-route proof is inferred.
            result = {'url': control['origin'] + links[0][len(control['account_email_origin']):],
                      'canonical_account_url': links[0], 'source': 'isolated-private-file-outbox',
                      'isolated_listener_mapping': True, 'provider_calls': 0, 'inbox_receipt': False}
        elif args.action == 'state':
            actor = payload['browser']
            require(actor in control['invitees'], 'Unknown isolated browser recipient.')
            user = User.objects.filter(email=control['invitees'][actor]['email']).first()
            owner = User.objects.get(email=control['owner']['email'])
            result = {'exists': bool(user), 'owner_password_sha256': hashlib.sha256(owner.password.encode()).hexdigest(),
                      'provider_jobs': 0, 'mail_mode': settings.QUILLMD_EMAIL_MODE}
            if user:
                result.update(user_id=str(user.pk), verified=bool(user.verified_at), is_staff=user.is_staff,
                              password_sha256=hashlib.sha256(user.password.encode()).hexdigest(),
                              memberships=[{'department_id': str(row.department_id), 'role': row.role, 'active': row.active} for row in Membership.objects.filter(user=user)],
                              wallets=[{'kind': row.kind, 'balance': str(row.balance)} for row in Wallet.objects.filter(user=user)],
                              welcome_grants=LedgerEntry.objects.filter(wallet__user=user, kind='welcome').count(),
                              recovery_events=AuditEvent.objects.filter(actor=user, action='auth.password_recovered').count())
            from core.models import Job, Usage, ProviderReceipt
            result['provider_jobs'] = Job.objects.count()
            result['provider_usage'] = Usage.objects.count()
            result['provider_receipts'] = ProviderReceipt.objects.count()
        else:
            # Each ordinary browser gets an independent limiter fixture. The
            # dedicated limiter unit tests establish actual throttling behavior.
            deleted, _ = LoginAttempt.objects.all().delete()
            result = {'isolated_rate_limit_fixture_rows_cleared': deleted}
    print(json.dumps(result))


if __name__ == '__main__':
    try:
        main()
    except Exception:
        print(json.dumps({'error': 'Isolated preview fixture operation refused; sensitive details suppressed.'}))
        raise SystemExit(1)
