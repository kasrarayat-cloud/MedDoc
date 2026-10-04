"""Source-only fault controls for the future served reader; no DB/provider/service."""
import contextlib
from io import BytesIO, StringIO
import json
from pathlib import Path
import tarfile
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from documents import word_v3
from ops import verify_fixture_persisted as probe


class PersistedFixtureProbeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        document = Document(); document.add_heading('Synthetic SOAP positive', 0)
        table = document.add_table(rows=0, cols=2)
        for label in probe.LABELS['soap']:
            row = table.add_row(); row.cells[0].text = label
        stream = BytesIO(); document.save(stream); cls.original = stream.getvalue()
        cls.fields = word_v3.inspect(cls.original)['fields']

    def example(self):
        first, final = {}, {}
        notes = 'Subjective: cough for two days.\nObjective: lungs clear.'
        for index, field in enumerate(self.fields):
            key = field['key']
            if index < 2:
                quote = 'cough for two days.' if index == 0 else 'lungs clear.'
                start = notes.index(quote)
                before = {'status': 'supported', 'text': quote,
                          'evidence': [{'start': start, 'end': start + len(quote), 'quote': quote}]}
                after = {**before, 'reviewed': True, 'reviewed_by': 'synthetic-doctor'}
            else:
                before = {'status': 'missing' if index == 2 else 'conflict', 'text': '', 'evidence': []}
                after = {**before, 'status': 'left_blank', 'reviewed': True, 'reviewed_by': 'synthetic-doctor',
                         'blank_decision': {'by': 'synthetic-doctor', 'from_status': before['status']}}
            first[key], final[key] = before, after
        exported = word_v3.fill(self.original, self.fields, {key: row['text'] for key, row in final.items()})
        record = {'fields': self.fields, 'contract': 'v3', 'mapping_confirmed': True,
            'source_sha256': probe.digest(self.original), 'original_sha256': probe.digest(self.original),
            'status': 'finalized', 'finalized_by': 'synthetic-doctor', 'doctor': 'synthetic-doctor',
            'finalized_at': '2026-10-04T00:00:00Z', 'department': 'own', 'patient_department': 'own',
            'template_department': 'own', 'job_encounter': 'own-encounter', 'encounter': 'own-encounter',
            'job_user': 'synthetic-doctor', 'job_state': 'succeeded', 'job_kind': 'generation',
            'config': {'mode': 'fixture', 'provider': 'anthropic', 'model': probe.preparation.MODEL,
                'ordinary_field_contract': probe.preparation.CONTRACT,
                'request_profile': {'profile': 'claude-sonnet-5-5-ordinary-v1', 'max_output_tokens': 20000}},
            'answers': final, 'notes': notes, 'usage': {'simulated': True, 'actual_cost': '0', 'charged_cost': '0'},
            'receipt_matches_usage': True, 'held_reservations': 0, 'export_sha256': probe.digest(exported)}
        return record, first, exported

    def test_actual_inspected_bound_original_and_export_report_holds_and_reviewed_blanks_separately(self):
        record, first, exported = self.example()
        result = probe.validate_example(record, self.original, first, exported, 'soap')
        self.assertEqual(result['first_pass_statuses'], {'conflict': 1, 'missing': 1, 'supported': 2})
        self.assertEqual(result['useful_first_pass_answers'], 2)
        self.assertEqual(len(result['reviewed_blank_keys']), 2)
        self.assertEqual(result['clinician_corrections'], [])
        self.assertFalse(result['live_model_qualification'])
        self.assertFalse(result['native_word_inspection'])

    def test_a_manual_answer_does_not_become_first_pass_correctness(self):
        record, first, _ = self.example(); key = self.fields[2]['key']
        record['answers'][key].update(text='Clinician entered assessment.', status='edited')
        record['answers'][key].pop('blank_decision')
        exported = word_v3.fill(self.original, self.fields, {key: row['text'] for key, row in record['answers'].items()})
        record['export_sha256'] = probe.digest(exported)
        result = probe.validate_example(record, self.original, first, exported, 'soap')
        self.assertEqual(result['useful_first_pass_answers'], 2)
        self.assertEqual(result['clinician_corrections'], [key])
        self.assertEqual(len(result['reviewed_blank_keys']), 1)

    def test_missing_review_blank_identity_source_quote_or_zero_usefulness_refuse(self):
        mutations = (
            lambda record, first: record['answers'][self.fields[0]['key']].update(reviewed=False),
            lambda record, first: record['answers'][self.fields[2]['key']]['blank_decision'].update(by='other-doctor'),
            lambda record, first: first[self.fields[0]['key']]['evidence'][0].update(quote='invented finding'),
            lambda record, first: [row.update(status='missing', text='') for row in first.values()],
        )
        for change in mutations:
            record, first, exported = self.example(); change(record, first)
            with self.subTest(change=change), self.assertRaises(RuntimeError):
                probe.validate_example(record, self.original, first, exported, 'soap')

    def test_live_wrong_profile_foreign_scope_or_false_fixture_finance_refuse(self):
        mutations = (
            lambda record: record['config'].update(mode='live'),
            lambda record: record['config']['request_profile'].update(max_output_tokens=8000),
            lambda record: record['config'].update(model='claude-sonnet-4-6'),
            lambda record: record.update(patient_department='foreign'),
            lambda record: record['usage'].update(simulated=False),
            lambda record: record['usage'].update(charged_cost='0.000000000001'),
            lambda record: record.update(receipt_matches_usage=False),
            lambda record: record.update(held_reservations=1),
        )
        for change in mutations:
            record, first, exported = self.example(); change(record)
            with self.subTest(change=change), self.assertRaises(RuntimeError):
                probe.validate_example(record, self.original, first, exported, 'soap')

    def test_hash_matching_export_that_changes_neighbor_package_is_still_refused(self):
        record, first, exported = self.example(); stream = BytesIO()
        with ZipFile(BytesIO(exported)) as original, ZipFile(stream, 'w', ZIP_DEFLATED) as altered:
            for item in original.infolist():
                altered.writestr(item, b'<wrong/>' if item.filename == 'docProps/core.xml' else original.read(item.filename))
        changed = stream.getvalue(); record['export_sha256'] = probe.digest(changed)
        with self.assertRaisesRegex(RuntimeError, 'neighbor_package_bytes_changed'):
            probe.validate_example(record, self.original, first, changed, 'soap')

    def archive(self, directory):
        target = 'a' * 40; demo = directory / ('demo-' + target[:8]); demo.mkdir(mode=0o700)
        (demo / 'src').mkdir(mode=0o700); source = demo / 'src/rebuild'; source.mkdir(mode=0o700)
        (source / 'one.py').write_bytes(b'print("synthetic immutable source")\n')
        (source / 'one.py').chmod(0o600)
        stream = BytesIO()
        with tarfile.open(fileobj=stream, mode='w', format=tarfile.PAX_FORMAT,
                          pax_headers={'comment': target}) as archive:
            archive.add(source / 'one.py', arcname='rebuild/one.py')
        raw = stream.getvalue()
        prepared = {'commit': target, 'root': str(demo), 'database': 'quillmd_qa_int_sol_' + target[:8],
            'ports': [9160, 9161], 'ai_mode': 'fixture', 'field_relations': True, 'role_suggestions': False,
            'source_tar_sha256': probe.digest(raw),
            'runtime_source_sha256': {'rebuild/one.py': probe.digest((source / 'one.py').read_bytes())}}
        for path, value in ((demo / 'source.tar', raw), (demo / 'prepared.json', json.dumps(prepared).encode())):
            path.write_bytes(value); path.chmod(0o600)
        return target, demo, source, prepared

    def test_full_archive_identity_refuses_stale_source_missing_inventory_and_old_demo(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary); directory.chmod(0o700)
            target, demo, source, prepared = self.archive(directory)
            with patch.object(probe.auth, 'TASK', directory):
                self.assertEqual(probe.archive_identity(demo, target), prepared)
                (source / 'foreign.py').write_bytes(b'untracked source')
                with self.assertRaisesRegex(RuntimeError, 'untracked_runtime_source_refused'):
                    probe.archive_identity(demo, target)
                (source / 'foreign.py').unlink()
                changed = dict(prepared, runtime_source_sha256={})
                (demo / 'prepared.json').write_text(json.dumps(changed))
                with self.assertRaisesRegex(RuntimeError, 'complete_archive_inventory_differs'):
                    probe.archive_identity(demo, target)
                (demo / 'prepared.json').write_text(json.dumps(prepared))
                (source / 'one.py').write_bytes(b'stale source')
                with self.assertRaisesRegex(RuntimeError, 'served_source_bytes_changed'):
                    probe.archive_identity(demo, target)
                with self.assertRaisesRegex(RuntimeError, 'new_demo_scope_required'):
                    probe.archive_identity(directory / 'demo-789c0372', '789c0372' + '0' * 32)

    def test_default_cli_does_not_open_credentials_or_start_database_network_services(self):
        output = StringIO()
        with patch.object(sys_module(), 'argv', ['verify_fixture_persisted']), \
                patch.object(probe.auth, 'protected_file', side_effect=AssertionError('unexpected private read')), \
                patch.object(probe.socket.socket, 'connect', side_effect=AssertionError('unexpected network')), \
                contextlib.redirect_stdout(output):
            probe.main()
        result = json.loads(output.getvalue())
        self.assertTrue(result['prepared_only']); self.assertEqual(result['clinical_mutations'], 0)
        self.assertEqual(result['provider_requests'], 0)

    def test_preceding_browser_mfa_code_is_not_replayed_or_waited_for_without_bound(self):
        self.assertAlmostEqual(probe.fresh_mfa_delay(100, 3005), 25.2)
        self.assertEqual(probe.fresh_mfa_delay(99, 3005), 0)
        with self.assertRaisesRegex(RuntimeError, 'authenticator_clock_or_counter_requires_review'):
            probe.fresh_mfa_delay(110, 3005)

    def test_actual_synthetic_health_labels_match_both_mounted_surfaces(self):
        for surface in ('doctor', 'admin'):
            health = {'status': 'ok', 'data_mode': 'synthetic-only', 'surface': surface, 'release': 'a' * 40}
            self.assertIsNone(probe.validate_health(health, 'a' * 40, surface))

    def test_old_label_real_mode_stale_release_foreign_host_and_extra_identity_are_refused(self):
        valid = {'status': 'ok', 'data_mode': 'synthetic-only', 'surface': 'doctor', 'release': 'a' * 40}
        for change in ({'data_mode': 'synthetic'}, {'data_mode': 'real'}, {'data_mode': 'live'},
                       {'release': 'b' * 40}, {'surface': 'admin'}, {'extra': 'not the bound health contract'}):
            with self.subTest(change=change), self.assertRaisesRegex(RuntimeError, 'served_health_identity_differs'):
                probe.validate_health({**valid, **change}, 'a' * 40, 'doctor')


def sys_module():
    import sys
    return sys


if __name__ == '__main__':
    unittest.main()
