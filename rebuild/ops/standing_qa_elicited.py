"""Fixed synthetic elicited-report qualification through standing QA machinery.

No import reads credentials, creates a database, or contacts a provider. The
original three CI fixtures stay offline; these independently authored forms
measure the same proof transport using fresh, explicit clinical questions.
"""
import datetime
import hashlib
import io
import json
import os
from pathlib import Path
import zipfile

try:
    from . import provider_campaign as pc, sonnet5_stages as stages
except ImportError:
    import provider_campaign as pc, sonnet5_stages as stages

STAGE = 'elicited_qualification'
PRICING = 'ordinary_notes_field_relations'
MANIFEST = Path(__file__).with_name('fixtures') / 'elicited-qualification-v1.json'
MANIFEST_SHA256 = '73b118daa070b5ae65d150e3c4e394982caccc18b771fc3533de748a85706111'
FORM_LABELS = {
    'elicited-single-v1': ('Patient response about relationships',),
    'elicited-seven-v1': ('Housing suitability', 'Monitoring relationships',
        'Patient response about relationships', 'Clinical diagnosis', 'Medication administration',
        'Physical illness affecting attendance', 'Describe the past clinician encounter in which the patient was seen.'),
}
CASE_IDS = ('E01', 'E02', 'E03', 'E04', 'E05')
require = pc.require


def build(form):
    """Deterministic blank form, authored without reading a development original."""
    from docx import Document
    doc = Document()
    doc.core_properties.title = 'Synthetic clinical reporting worksheet'
    doc.core_properties.author = doc.core_properties.last_modified_by = 'QuillMD synthetic engineering'
    doc.core_properties.created = doc.core_properties.modified = datetime.datetime(2026, 10, 4)
    doc.core_properties.comments = 'Project-authored synthetic fixture; no patient data or signing action.'
    doc.add_heading('Synthetic clinical reporting worksheet', 0)
    table = doc.add_table(rows=len(FORM_LABELS[form]), cols=2)
    table.style = 'Table Grid'
    for row, label in zip(table.rows, FORM_LABELS[form]):
        row.cells[0].text = label
        row.cells[1].text = ''
    raw = io.BytesIO(); doc.save(raw)
    output = io.BytesIO()
    with zipfile.ZipFile(raw) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            item = zipfile.ZipInfo(name, date_time=(2026, 10, 4, 0, 0, 0))
            item.compress_type = zipfile.ZIP_DEFLATED
            item.create_system = 3; item.external_attr = 0o600 << 16
            target.writestr(item, source.read(name))
    return output.getvalue()


def frozen_inputs():
    raw = MANIFEST.read_bytes()
    root = Path(__file__).resolve().parents[2]
    pc.committed_file(MANIFEST, root, 'elicited qualification manifest')
    pc.committed_file(Path(__file__).resolve(), root, 'elicited qualification runner')
    require(MANIFEST_SHA256 is not None and hashlib.sha256(raw).hexdigest() == MANIFEST_SHA256,
            'The elicited qualification manifest is unreviewed or changed.')
    manifest = json.loads(raw)
    require(manifest['stage'] == STAGE and manifest['provider_upload_allowed'] is True
            and manifest['independently_authored'] is True and manifest['corpus_qualification'] is False,
            'The elicited fixture rights or scope changed.')
    require(tuple(case['id'] for case in manifest['cases']) == CASE_IDS
            and set(manifest['forms']) == set(FORM_LABELS), 'The complete fixed elicited plan changed.')
    for form, spec in manifest['forms'].items():
        require(hashlib.sha256(build(form)).hexdigest() == spec['sha256'], 'A synthetic original changed.')
        require([f['label'] for f in spec['fields']] == list(FORM_LABELS[form]),
                'The synthetic physical destination catalogue changed.')
    for case in manifest['cases']:
        require(hashlib.sha256(case['notes'].encode()).hexdigest() == case['notes_sha256'],
                'A declared synthetic source changed.')
    return {'elicited': manifest, 'freeze': {'manifest_sha256': MANIFEST_SHA256,
            'frozen_files': {str(MANIFEST): MANIFEST_SHA256},
            'previously_exposed': True, 'exact_original_CI_fixtures_provider_eligible': False}}


def plan_sends(inputs=None):
    exact = frozen_inputs()
    inputs = exact if inputs is None else inputs
    require(inputs == exact, 'Only the complete fixed elicited qualification inputs are eligible.')
    return [stages.Send(i, 'elicited:' + case['id'], 'generation', case['form'], 'field_relations',
                        STAGE, case['id'], PRICING) for i, case in enumerate(inputs['elicited']['cases'], 1)]


def send_inputs(send, inputs):
    require(send in plan_sends(inputs), 'No arbitrary elicited send is eligible.')
    case = next(c for c in inputs['elicited']['cases'] if c['id'] == send.case)
    return {'form_sha256': inputs['elicited']['forms'][case['form']]['sha256'],
            'notes_sha256': case['notes_sha256']}


def build_forms(target, inputs):
    require(inputs == frozen_inputs(), 'The form builder requires the exact reviewed freeze.')
    target = pc.private_directory(target)
    result = {}
    for form in FORM_LABELS:
        raw = build(form); path = target / (form + '.docx')
        if os.path.lexists(path):
            require(not path.is_symlink() and path.read_bytes() == raw, 'A stored synthetic original changed.')
        else: pc.write_new(path, raw, raw=True)
        result[form] = raw
    return result


def mapping_admission(original_sha, fields, manifest):
    from documents.import_selection import digest
    coverage = manifest.get('coverage', {})
    require(manifest.get('engine') == 'v3' and manifest.get('source_sha256') == original_sha
            and manifest.get('mapping_sha256') == manifest.get('initial_mapping_sha256') == digest(fields),
            'The source and initial immutable mapping do not match this original.')
    require(coverage.get('block_generation') is False
            and not coverage.get('selected_engine_issues')
            and all(coverage.get(key) == [] for key in
                    ('legacy_omissions', 'legacy_unmapped', 'v3_uncovered_legacy_places')),
            'An omitted or unsupported destination cannot qualify the elicited workflow.')
    return coverage


def price_sends(sends, inputs, state_dir, declaration, limits):
    from django.test import override_settings
    from core import generation as g
    from documents.service import inspect_template
    build_forms(Path(state_dir) / 'forms', inputs)
    priced = {}
    with override_settings(**stages.steps.PATH_FLAGS['field_relations']):
        for send in sends:
            send_inputs(send, inputs)
            case = next(c for c in inputs['elicited']['cases'] if c['id'] == send.case)
            fields = inspect_template(Path(state_dir) / 'forms' / (send.form + '.docx'), engine='v3')['fields']
            require(fields == inputs['elicited']['forms'][send.form]['fields'], 'The current detected mapping changed.')
            config = g.configuration(fields)
            require(config['model'] == 'claude-sonnet-5-5' and config.get('ordinary_field_contract')
                    == 'ordinary-field-relations-v4' and config['request_profile']['max_output_tokens'] == 20000,
                    'The elicited send requires the qualified Sonnet 5.5 ordinary-v4 profile.')
            require(config['pricing_version'] == declaration.pricing_version(PRICING)
                    and stages.within_bounds(config['request_profile'], declaration.expect(limits)),
                    'The elicited quote/profile differs from the prospectively declared configuration.')
            maximum = g.maximum_cost(case['notes'], fields, config, context_labels=[], requirements=[])
            priced[send.id] = {'maximum_cost': maximum, 'bound': pc.attempt_bound(maximum),
                'pricing_version': config['pricing_version'], 'max_output_tokens': 20000,
                'timeout_seconds': config['request_profile']['timeout_seconds']}
    return priced


def usefulness(case, fields, answers):
    """Exact first-pass text/attribution; corrections never become automatic credit."""
    expected = case['desired_supported']; held = case['required_holds']
    require(set(expected) | set(held) == {f['label'] for f in fields} and not set(expected) & set(held),
            'The declared first-pass expectations do not cover the whole physical catalogue.')
    result = {}
    for field in fields:
        row = answers.get(field['key'], {}); label = field['label']
        if label in expected:
            result[field['key']] = ('supported_correct' if row.get('status') == 'supported'
                and row.get('text') == expected[label] else 'unsupported_accepted'
                if row.get('status') == 'supported' or row.get('text') else 'over_held')
        else:
            result[field['key']] = ('held_correct' if row.get('status') in ('conflict','missing')
                and not row.get('text') else 'unsupported_accepted')
    return result


class ElicitedRun(stages.Run):
    """Reuse the normal job/send/receipt/audit/review/export chain without substitution."""
    def notes_send(self, send):
        from django.test import override_settings
        require(self.stage == STAGE, 'The elicited worker requires its separately declared stage.')
        send_inputs(send, self.inputs)
        case = next(c for c in self.inputs['elicited']['cases'] if c['id'] == send.case)
        form = self.inputs['elicited']['forms'][case['form']]
        with override_settings(**stages.steps.PATH_FLAGS['field_relations']):
            template, version = self.upload(self.forms[send.form],
                f'SYNTHETIC ELICITED {send.case} attempt {len(self.ledger.rows()) + 1}', 'elicited.docx')
            require(template.sha256 == form['sha256'] and version.contract == 'v3'
                    and version.fields == form['fields'], 'The actual immutable original/mapping changed.')
            coverage = mapping_admission(template.sha256, version.fields, version.import_manifest)
            response = stages.steps.confirm_mapping(self.world.client, template.pk, version.pk)
            version.refresh_from_db()
            require(response.status_code == 302 and version.confirmed, 'Normal mapping confirmation was refused.')
            def judge(send, version, answers, exported, original, **kwargs):
                outcomes = usefulness(case, version.fields, answers)
                return outcomes, [], {'elicited_first_pass': {'fields': outcomes,
                    'original_CI_fixtures_offline_only': True, 'manual_corrections': 0}}
            return self.generation_send(send, template, version, case['notes'], extra={
                'case': send.case, 'elicited_fixture_manifest_sha256': MANIFEST_SHA256,
                'coverage_before_confirmation': coverage}, judge=judge)
