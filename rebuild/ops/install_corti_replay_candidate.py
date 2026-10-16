"""Seed only the separate replay candidate. Never contacts a provider.

Invoke after migrations using ops.corti_candidate_settings. The protected access
file is written exclusively; this script cannot reset an existing account/run.
"""
import argparse
from datetime import date
import hashlib
import json
from pathlib import Path
import secrets
from uuid import UUID

import django
django.setup()

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from docx import Document
from core import coding_service, testing_authority
from core.management.commands.import_coding_reference import prepare_reference, save_reference
from core.management.commands.import_terminology import save_release
from core.models import User, Organization, Department, Membership, Patient, Template, TemplateVersion, Encounter, CodingContext, CodeSet
from core.security import private_path, audit
from documents import inspect_template


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--manifest',required=True)
    parser.add_argument('--sha256',required=True)
    args=parser.parse_args()
    if (settings.SETTINGS_MODULE != 'ops.corti_candidate_settings' or settings.CORTI_CODING_MODE != 'replay'
            or not settings.DATABASES['default']['NAME'].startswith('test_') or not settings.SYNTHETIC_ONLY):
        raise RuntimeError('Only the isolated synthetic replay installation is eligible.')
    access=settings.STATE_DIR/'clinician-access.json'
    if access.exists() or User.objects.exists() or CodeSet.objects.exists():
        raise RuntimeError('Candidate already has records; refusing to replace any history.')
    prepared=prepare_reference(args.manifest,args.sha256)
    metadata=dict(prepared['provenance']['manifest']['release'])
    metadata['id']=UUID(metadata['id'])
    for key in ('effective_from','effective_until'):metadata[key]=date.fromisoformat(metadata[key])
    save_release(prepared['reference']['tabular']['codes'],metadata)
    reference,_=save_reference(prepared)
    note=settings.CORTI_REPLAY_FIXTURE['request']['context'][0]['text']
    password=secrets.token_urlsafe(24)
    with transaction.atomic():
        actor=User.objects.create_user(id='2247e403-a830-4807-9a34-6903b95a0f50',
            username='corti-review@synthetic.quillmd.invalid',email='corti-review@synthetic.quillmd.invalid',
            password=password,first_name='Synthetic',last_name='Reviewer',verified_at=timezone.now(),is_staff=False)
        org=Organization.objects.create(name='Synthetic Corti candidate')
        department=Department.objects.create(id=testing_authority.ACTORS[str(actor.pk)],organization=org,name='Synthetic GI')
        Membership.objects.create(user=actor,department=department,role='clinician')
        patient=Patient.objects.create(department=department,display_name='Synthetic Taylor Reed',reference='CORTI-REPLAY',created_by=actor)
        document=Document()
        document.sections[0].header.paragraphs[0].text='QuillMD · Synthetic replay candidate'
        document.add_heading('Synthetic gastroenterology office note',0)
        document.add_paragraph('Fictional encounter · service date 16 October 2026')
        document.add_paragraph('Clinical note: {{Clinical note}}')
        path=private_path('originals/corti-replay-blank.docx');path.parent.mkdir(mode=0o700,parents=True,exist_ok=True)
        document.save(path);path.chmod(0o600)
        inspected=inspect_template(path)
        template=Template.objects.create(department=department,created_by=actor,name='Corti synthetic encounter replay',format='docx',
            original_name=path.name,original_path='originals/'+path.name,sha256=hashlib.sha256(path.read_bytes()).hexdigest())
        version=TemplateVersion.objects.create(template=template,number=1,fields=inspected['fields'],created_by=actor)
        key=inspected['fields'][0]['key']
        encounter=Encounter.objects.create(department=department,patient=patient,template_version=version,
            created_by=actor,service_date=date(2026,10,16),notes=note,status='review',revision=1,
            answers={key:{'text':note,'status':'edited','reviewed':False,'evidence':[{'start':0,'end':len(note),'quote':note}]}})
        CodingContext.objects.create(encounter=encounter,updated_by=actor,
            data={**coding_service.CONTEXT_DEFAULTS,'setting':'outpatient_office','patient_status':'established','em_basis':'time'})
        audit(actor,'synthetic.corti_replay_seeded',encounter.pk,department=department,
            fixture_sha256=settings.CORTI_REPLAY_FIXTURE['response_sha256'],provider_calls=0,
            source='Retained fictional encounter; manual saved-answer fixture, not new generation.')
    with access.open('x') as handle:
        access.chmod(0o600)
        json.dump({'email':actor.email,'password':password},handle)
    receipt={'mode':'synthetic_retained_response_replay','encounter_id':str(encounter.pk),
        'clinician_url':f'http://127.0.0.1:9146/encounters/{encounter.pk}/',
        'coding_url':f'http://127.0.0.1:9146/encounters/{encounter.pk}/?tab=coding',
        'role':'clinician','is_staff':False,'provider_calls':0,'reference_version':reference.release.version,
        'reference_sha256':reference.release.sha256,'reference_digest':reference.digest,
        'original_verdict':settings.CORTI_REPLAY_FIXTURE['original_verdict']}
    path=settings.STATE_DIR/'installation.json';path.write_text(json.dumps(receipt,indent=2)+'\n');path.chmod(0o600)
    print(json.dumps(receipt))


if __name__=='__main__':main()
