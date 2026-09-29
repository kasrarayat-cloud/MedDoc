#!/usr/bin/env python3
"""The isolated QA world, and the model_check stage, of the V1 staged synthetic Sonnet 5 campaign.

The contract is the committed declaration docs/foundation/mutations/20260928-quillmd-v1-sonnet5-staged-campaign.md:
its limits, model, request settings, pricing versions, admission and stop rules, credential file, QA-only route,
synthetic-only inputs and retained evidence. Nothing here widens it. This file implements model_check and the ONE door
every stage goes through (outer_run -> the scrubbed request process -> inner); the mapping and ordinary_notes stages
are ops/sonnet5_stages.py, run through the same door.

One model_check attempt, in the request process (a fresh process with the scrubbed environment, re-executed by `run`):
  the QA world is seeded through the application's own models and views (a synthetic organization, doctor, patient, the
  doctor's one welcome credit, and the FIXED legacy-contract form uploaded through the real upload view); the FIXED
  three-line synthetic note and the form are hashed and written to the attempt's evidence BEFORE any send; the real
  authenticated views create the encounter, show the quote and take the Generate request (which quotes, enqueues and
  holds); the job's own maximum cost is read and admitted against the ledger; the attempt is written; the real worker
  function `core.jobs.run_one` runs ONCE with the outbound request captured before it leaves; the outcome is written; the
  stop conditions are checked (each adopted answer's evidence on its own question's notes line, containing the answer);
  the review FORM's answer inputs are read, every answer is saved reviewed, the document is finalized and its
  authenticated export is verified against the original (all of each answer cell's text; no hidden text or nested
  table); the billing audit runs; the evidence is scanned for the key (a tree it cannot read whole is a campaign stop);
  and a stage_passed record is written ONLY if every check held — otherwise a stage_failed record with the reasons.
  Nothing is ever re-sent.

Commands (from the candidate checkout's rebuild/):
  plan                     the fixed case, its hashes, the declared price and its maximum cost; no database, no request
  plan --stage mapping|ordinary_notes
                           the stage's fixed order, each send's maximum and bound, the stage worst case against the
                           remaining stage and total limits (the ledger read without its lock); no credential, network
                           or database
  run --stage model_check|mapping|ordinary_notes
                           the stage (refused unless the declaration, credential and ledger all allow it)
  status                   the ledger against every declared limit; read-only
  recover                  records any attempt left without an outcome as interrupted_unknown and fails its stage
  reopen --stage <id> --diagnosis <path>
                           reopens a FAILED stage after offline repair (provider_campaign.Campaign.reopen); nothing is sent
Exit status: 0 PASS (or a clean plan/status/recover/reopen), 3 FAIL (stage_failed written), 2 REFUSED (nothing sent).
"""
import argparse
import copy
import hashlib
import html.parser
import io
import json
import os
import re
import secrets
import signal
import stat
import sys
import time
import zipfile
from decimal import Decimal
from pathlib import Path

try:
    from . import provider_campaign as pc
except ImportError:                                  # (run as a script: ops/ is on sys.path)
    import provider_campaign as pc

RUNNER = Path(__file__).resolve()
REBUILD = RUNNER.parents[1]
REPOSITORY = REBUILD.parent
DECLARATION = REPOSITORY / 'docs' / 'foundation' / 'mutations' / '20260928-quillmd-v1-sonnet5-staged-campaign.md'
LIVE_ROOT = pc.LIVE_ROOT
STAGE = 'model_check'
EXPECTED_MODEL = 'claude-sonnet-5'
#: The one synthetic QA database server (container v1-sonnet5-qa-pg), and the only database-name prefix accepted.
DATABASE = {'name': 'quillmd_qa_sonnet5', 'user': 'quillmd_qa', 'host': '127.0.0.1', 'port': '55731'}
DATABASE_PREFIX = 'quillmd_qa_'
DATABASE_PASSWORD_FILE = LIVE_ROOT / 'qa-db' / 'password'
#: Stated by the declaration in prose, not in a structured line: its pricing versions were computed "under these
#: settings (and QUILLMD_DEEPGRAM_RATE_PER_MINUTE=0.0043, no discount)". The pricing-version check refuses any drift.
PRICE_SETTINGS = {'QUILLMD_DEEPGRAM_RATE_PER_MINUTE': '0.0043', 'QUILLMD_DEEPGRAM_RATE_PER_SECOND': '',
                  'QUILLMD_PROVIDER_DISCOUNT_RATE': '0'}
#: The model check runs the plain Sonnet 5 generation profile: every contract flag that changes the request is off.
PLAIN_PROFILE_FLAGS = {'QUILLMD_WM_NARRATIVE': '0', 'QUILLMD_V3_HISTORY_CONTEXT': '0', 'QUILLMD_V3_GENERATION_BATCHES': '0',
                       'QUILLMD_V3_FIELD_RELATIONS': '0', 'QUILLMD_MAPPING_ASSISTANCE': '0',
                       'QUILLMD_MAPPING_ASSISTANCE_LIVE': '0', 'QUILLMD_WORD_ENGINE': 'legacy'}
#: A stop: output at or above 70% of the profile's output limit.
OUTPUT_SHARE_LIMIT = Decimal('0.70')
STATIC_STORAGES = {'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
                   'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}}

# ── the FIXED case ───────────────────────────────────────────────────────────────────────────────────────────────
FORM_NAME = 'Synthetic Sonnet 5 model check'
FORM_FILENAME = 'sonnet5-model-check.docx'
FORM_TITLE = 'Synthetic clinical assessment'
FORM_LABELS = ('History', 'Findings', 'Plan')
#: The form's bytes, fixed: python-docx's output with every zip entry's timestamp and attributes normalized.
FORM_SHA256 = 'e0a3bdf4e98f07612d94572d877bacd91aab60dfd8f80e672b0d30d36c7ffeae'
NOTES = ('History: Synthetic patient reports left wrist discomfort after a fall.\n'
         'Findings: No deformity documented.\n'
         'Plan: Follow-up assessment documented.')
NOTES_SHA256 = hashlib.sha256(NOTES.encode()).hexdigest()
#: What a correct draft adopts for each question: the complete documented value after its own label, nothing else.
EXPECTED = {'History': 'Synthetic patient reports left wrist discomfort after a fall.',
            'Findings': 'No deformity documented.',
            'Plan': 'Follow-up assessment documented.'}
DOCTOR_EMAIL = 'sonnet5-qa-doctor@synthetic.quillmd.invalid'
PATIENT_REFERENCE = 'SONNET5-QA-001'
LOGIN = 'django.test.Client.force_login (the synthetic QA doctor; no password is used)'
#: The engine files whose hashes every attempt's provenance records.
ENGINE_FILES = ('core/jobs.py', 'core/generation.py', 'core/billing.py', 'core/views.py', 'core/security.py',
                'core/models.py', 'core/middleware.py', 'core/management/commands/audit_billing.py',
                'documents/service.py', 'documents/engine.py', 'documents/worker.py', 'documents/layout.py',
                'quillmd/settings.py', 'quillmd/urls.py', 'ops/provider_campaign.py', 'ops/sonnet5_qa.py')
#: The files a reopening must show were repaired (A2), repository-relative: every engine file above and the stage
#: runner with the frozen diagnostic graders and walk steps it imports. Every new attempt records their sha256.
STAGE_RUNNER_FILES = ('rebuild/ops/sonnet5_stages.py',
                      'docs/rebuild/v1/sonnet5-campaign/diagnostics/mapping/grade_app_path.py',
                      'docs/rebuild/v1/sonnet5-campaign/diagnostics/mapping/measure_baseline.py',
                      'docs/rebuild/v1/sonnet5-campaign/diagnostics/notes/app_path_steps.py',
                      'docs/rebuild/v1/sonnet5-campaign/diagnostics/notes/offline_replay.py')
WATCHED_FILES = tuple('rebuild/' + name for name in ENGINE_FILES) + STAGE_RUNNER_FILES
EXIT = {'PASS': 0, 'FAIL': 3, 'REFUSED': 2}
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

require = pc.require
Refused = pc.Refused


def build_form():
    """The fixed legacy-contract form (a History / Findings / Plan question table, like seed_private's 'Clinical
    assessment'), as deterministic bytes."""
    from docx import Document
    document = Document()
    document.add_heading(FORM_TITLE, 0)
    document.add_paragraph('Private test fixture · no patient data')
    table = document.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    for label in FORM_LABELS:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = ''
    produced = io.BytesIO()
    document.save(produced)
    source = zipfile.ZipFile(io.BytesIO(produced.getvalue()))
    fixed = io.BytesIO()
    with zipfile.ZipFile(fixed, 'w') as archive:
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            archive.writestr(entry, source.read(info.filename))
    return fixed.getvalue()


def env_file_present(rebuild=REBUILD):
    """Whether the checkout carries rebuild/.env. Answered from the directory entry alone: nothing is opened."""
    return os.path.lexists(os.path.join(str(rebuild), '.env'))


def sha256(data):
    return hashlib.sha256(data if isinstance(data, bytes) else data.encode()).hexdigest()


def private_evidence(root, number):
    """The attempt's evidence directory, new. A directory left by a run that died BEFORE its attempt record was written
    is set aside, not reused: nothing can have been sent under a number with no attempt record."""
    base = pc.private_directory(Path(root) / 'evidence')
    target = base / f'{number:03d}'
    if target.exists():
        target.rename(base / f'orphan-{number:03d}-{time.time_ns()}')
    target.mkdir(mode=0o700)
    return target


# ── the ledger's open ends ───────────────────────────────────────────────────────────────────────────────────────

def settle_open_stages(campaign):
    """Every attempt with no outcome is recorded `interrupted_unknown`; every stage with an attempt no decision covers
    (an open stage's attempts since its last reopening) is FAILED — the run ended before it was decided, so it is never
    repeated in place: offline diagnosis, then `reopen`. Needs the lock."""
    ledger = campaign.ledger
    recovered = pc.recover_interrupted(ledger)
    rows = ledger.rows()
    outcomes = {attempt['number']: outcome for attempt, outcome in rows}
    failed = []
    for stage, numbers in sorted(ledger.undecided(rows).items()):
        reasons = [f'attempt {number} ended {outcomes[number]["class"]}' + (' (recovered after an interruption)'
                   if outcomes[number].get('recovered') else '') for number in numbers]
        ledger.write_stage(stage, 'failed', {'attempts': numbers, 'reasons': [
            'The run ended before this stage was decided; it is stopped for offline diagnosis and is admitted again only '
            'after `reopen` with a committed diagnosis and repair.'] + reasons})
        failed.append(stage)
    return {'recovered': recovered, 'failed_stages': failed}


# ── preconditions ────────────────────────────────────────────────────────────────────────────────────────────────

def preconditions(declaration, stage, root, *, database_prefix=DATABASE_PREFIX, rebuild=REBUILD, environ=None,
                  ledger=None):
    """Asserted before any enqueue. Returns what was checked, for the provenance manifest."""
    from django.conf import settings
    from django.db import connection
    from core import generation
    environ = os.environ if environ is None else environ
    require(not env_file_present(rebuild), 'The candidate checkout has a rebuild/.env; the campaign never runs beside one.')
    database = connection.settings_dict
    require(str(database['NAME']).startswith(database_prefix),
            f'The database must be a {database_prefix}* synthetic QA database.')
    require((str(database['HOST']), str(database['PORT'])) == (DATABASE['host'], DATABASE['port']),
            'The database must be the QA server on 127.0.0.1:55731.')
    root = Path(root).resolve()
    state, media = Path(settings.STATE_DIR).resolve(), Path(settings.MEDIA_ROOT).resolve()
    require(state.is_relative_to(root) and state != root and media.is_relative_to(state),
            'STATE_DIR (and its files) must be under the campaign runtime directory.')
    require(settings.SYNTHETIC_ONLY is True and settings.QUILLMD_DATA_MODE == 'synthetic', 'Synthetic-only mode is required.')
    require(settings.AI_MODE == 'live', 'The QA process must be in live AI mode.')
    require(settings.AI_PROVIDER == declaration.provider == 'anthropic', 'The provider must be anthropic.')
    require(settings.AI_MODEL == declaration.model == EXPECTED_MODEL, f'The model must be {EXPECTED_MODEL}.')
    require((settings.AI_INPUT_RATE, settings.AI_OUTPUT_RATE, settings.AI_CACHED_RATE)
            == (declaration.rates.input, declaration.rates.output, declaration.rates.cached),
            'The configured rates are not the declared rates.')
    require((settings.AI_THINKING, settings.AI_EFFORT, settings.AI_INFERENCE_GEO)
            == (declaration.thinking, declaration.effort, declaration.inference_geo),
            'The configured thinking, effort or inference geography is not the declared one.')
    require(not (settings.WM_NARRATIVE or settings.V3_FIELD_RELATIONS or settings.V3_GENERATION_BATCHES
                 or settings.V3_HISTORY_CONTEXT), 'The model check runs the plain generation profile only.')
    config = generation.configuration()
    require(config['pricing_version'] == declaration.pricing_version(stage),
            f'The application prices {config["pricing_version"]}, not the declared {stage} version.')
    require(config['request_profile'].get('profile') == 'claude-sonnet-5-v1' and 'invalid' not in config['request_profile'],
            'The request profile is not the Sonnet 5 generation profile.')
    key = pc.load_credential(declaration.credential_file, declaration)
    require(pc.constant_time_equal(key, environ.get('ANTHROPIC_API_KEY', '')),
            'The request process does not carry the declared credential.')
    pc.scrubbed_now(environ)
    # (a copy of the key anywhere under the campaign root stops the campaign; found now, it stops it before a send
    # rather than after one)
    stray = pc.leaked(root, key, exclude=(declaration.credential_file,))
    require(not stray, 'The credential is already present in retained files: ' + ', '.join(stray))
    if ledger is not None:
        rows = ledger.rows()
        require(all(outcome is not None for _, outcome in rows), 'The ledger is not reconciled: an attempt has no outcome.')
        require(not ledger.undecided(rows), 'The ledger is not reconciled: a stage with attempts has no decision.')
    return {'database': {'name': database['NAME'], 'host': database['HOST'], 'port': str(database['PORT'])},
            'state_dir': str(state), 'data_mode': settings.QUILLMD_DATA_MODE, 'ai_mode': settings.AI_MODE,
            'provider': settings.AI_PROVIDER, 'model': settings.AI_MODEL,
            'rates': {'input': settings.AI_INPUT_RATE, 'output': settings.AI_OUTPUT_RATE, 'cached': settings.AI_CACHED_RATE},
            'thinking': settings.AI_THINKING, 'effort': settings.AI_EFFORT, 'inference_geo': settings.AI_INFERENCE_GEO,
            'pricing_version': config['pricing_version'], 'request_profile': config['request_profile'],
            'word_engine': settings.QUILLMD_WORD_ENGINE,
            'price_settings': {name: os.environ.get(name, '') for name in PRICE_SETTINGS}}


# ── the synthetic world ──────────────────────────────────────────────────────────────────────────────────────────

class World:
    def __init__(self, user, department, patient, template, version, client):
        self.user, self.department, self.patient = user, department, patient
        self.template, self.version, self.client = template, version, client


class Account:
    def __init__(self, user, department, patient, client):
        self.user, self.department, self.patient, self.client = user, department, patient, client


def seed_account(root):
    """The synthetic organization, doctor, patient and welcome credit (idempotent), through the application's own models
    and its welcome-credit grant, and a client logged in as the doctor. The doctor's access record is written under the
    campaign runtime directory — never BASE_DIR/.runtime/private-access.json. Every stage uses this one account."""
    from django.db import transaction
    from django.test import Client
    from django.utils import timezone
    from core.billing import welcome_credit
    from core.models import Department, Membership, Organization, Patient, User
    from core.security import audit
    root = Path(root)
    user = User.objects.filter(username=DOCTOR_EMAIL).first()
    if user is None:
        access = pc.private_directory(root / 'access')
        password = secrets.token_urlsafe(24)
        with transaction.atomic():
            user = User.objects.create_user(username=DOCTOR_EMAIL, email=DOCTOR_EMAIL, password=password,
                                            first_name='Synthetic', last_name='Sonnet 5 QA Doctor',
                                            verified_at=timezone.now())
            organization = Organization.objects.create(name='Synthetic Sonnet 5 QA Clinic')
            department = Department.objects.create(organization=organization, name='Synthetic assessment')
            Membership.objects.create(user=user, department=department, role='owner')
            Patient.objects.create(department=department, display_name='SYNTHETIC SONNET 5 QA PATIENT',
                                   reference=PATIENT_REFERENCE, created_by=user)
            audit(user, 'synthetic.account_seeded', user.pk)
            pc.write_new(access / 'qa-doctor.json', {'email': DOCTOR_EMAIL, 'password': password, 'user_id': str(user.pk),
                                                     'login_used_by_runner': LOGIN, 'created_at': pc.now()})
    welcome_credit(user)            # (the application's own one-time grant, keyed to the account: idempotent)
    membership = Membership.objects.get(user=user, active=True)
    department = membership.department
    patient = Patient.objects.get(department=department, reference=PATIENT_REFERENCE)
    client = Client()
    client.force_login(user)
    return Account(user, department, patient, client)


def seed_world(root):
    """The synthetic account (seed_account) and the FIXED legacy-contract form uploaded through the real upload view."""
    from django.core.files.uploadedfile import SimpleUploadedFile
    from core.models import Template
    account = seed_account(root)
    user, department, patient, client = account.user, account.department, account.patient, account.client
    form = build_form()
    require(sha256(form) == FORM_SHA256, 'The fixed form did not build to its recorded bytes.')
    template = Template.objects.filter(department=department, sha256=FORM_SHA256).order_by('created_at').first()
    if template is None:
        response = client.post('/templates/upload/', {
            'name': FORM_NAME, 'department': str(department.pk), 'synthetic_only': 'on',
            'file': SimpleUploadedFile(FORM_FILENAME, form,
                                       content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        require(response.status_code == 302, f'The form upload was refused ({response.status_code}).')
        template = Template.objects.get(department=department, sha256=FORM_SHA256)
    version = template.versions.order_by('-number').first()
    require(version is not None and [f['label'] for f in version.fields] == list(FORM_LABELS)
            and not version.review_requirements and not version.document_backend and not version.contract
            and not version.needs_confirmation, 'The uploaded form is not the fixed legacy-contract form.')
    return World(user, department, patient, template, version, client)


def cancel_stale_jobs(world):
    """A queued job no run owns (a run that died between Generate and its attempt record) is cancelled through the
    product before anything else: it never ran, so nothing was sent, and the worker takes the OLDEST queued job."""
    from core.jobs import cancel
    from core.models import Job
    require(not Job.objects.filter(state='running').exists(), 'A job is running in the QA database; nothing is started.')
    cancelled = []
    for job in Job.objects.filter(state='queued').order_by('created_at'):
        cancel(job, job.user)
        cancelled.append(str(job.pk))
    return cancelled


def page_quote(client, encounter):
    """The Generate form exactly as the page shows it: the signed quote, the revision and the maximum charge."""
    import html
    from django.core import signing
    page = client.get(f'/encounters/{encounter.pk}/')
    require(page.status_code == 200, f'The encounter page failed ({page.status_code}).')
    text = page.content.decode()
    form = re.search(r'<form[^>]*data-generate-form[^>]*>(.*?)</form>', text, re.S)
    require(form, 'The page offers no Generate form.')
    token = re.search(r'name="usage_quote" value="([^"]*)"', form.group(1))
    revision = re.search(r'name="revision" value="(\d+)"', form.group(1))
    shown = re.search(r'Reserve up to \$([0-9.]+)\.', text)
    require(token and revision and shown, 'The page shows no complete quote.')
    token = html.unescape(token.group(1))
    quoted = signing.loads(token, salt='usage-quote')['config']
    return token, revision.group(1), Decimal(shown.group(1)), quoted


# ── evidence ─────────────────────────────────────────────────────────────────────────────────────────────────────

def rows_of(model, **filters):
    return [{key: value for key, value in row.items()} for row in model.objects.filter(**filters).values()]


def job_records(job):
    from core.models import Job, LedgerEntry, PlatformCost, ProviderReceipt, Reservation, Usage, Wallet
    return {'job': rows_of(Job, pk=job.pk), 'usage': rows_of(Usage, job=job),
            'provider_receipts': rows_of(ProviderReceipt, job=job), 'reservations': rows_of(Reservation, job=job),
            'ledger_entries': rows_of(LedgerEntry, job=job), 'platform_costs': rows_of(PlatformCost, job=job),
            'wallets': rows_of(Wallet, pk__in=[r.wallet_id for r in Reservation.objects.filter(job=job)])}


def raw_result(job):
    """(bytes, parsed) of the job's stored provider result file, or (None, None)."""
    from core.security import private_path
    if not job.result_path:
        return None, None
    data = private_path(job.result_path).read_bytes()
    require(hashlib.sha256(data).hexdigest() == job.result_sha256, 'The stored provider result does not match its hash.')
    return data, json.loads(data)


def usage_from_receipt(units):
    """The provider-reported usage a receipt recorded, in the Messages API's own names."""
    return {'input_tokens': units.get('input_tokens'), 'output_tokens': units.get('output_tokens'),
            'cache_read_input_tokens': units.get('cached_input_tokens', 0),
            'cache_creation_input_tokens': units.get('cache_creation_tokens', 0),
            'cache_creation': ({'ephemeral_5m_input_tokens': units['cache_creation_5m_tokens'],
                                'ephemeral_1h_input_tokens': units['cache_creation_1h_tokens']}
                               if 'cache_creation_5m_tokens' in units else None),
            'inference_geo': units.get('inference_geo')}


def classify(job, receipt, usage_row, captured):
    """One attempt's outcome class from the product's own records."""
    from core.jobs import NOT_PROCESSED
    if receipt is not None and usage_row is not None and job.state in ('succeeded', 'failed'):
        return 'completed'                  # the provider answered and it was settled: usage is known
    if job.state == 'failed' and (job.error or '').startswith(NOT_PROCESSED) and receipt is None and not captured:
        return 'not_sent'                   # the product proved no request, and none was captured
    if job.state == 'reconciliation':
        return 'reconciliation'
    return 'uncertain'


def question_line(notes, label):
    """(start, end) of the ONE notes line that documents `label` (`Label: value`), end excluding its line break; None
    when the notes do not have exactly one. Lines are split as str.splitlines does, so no line separator hides."""
    spans, offset = [], 0
    for line in notes.splitlines(keepends=True):
        body = line.splitlines()[0] if line.splitlines() else ''
        if body.startswith(label + ':'):
            spans.append((offset, offset + len(body)))
        offset += len(line)
    return spans[0] if len(spans) == 1 else None


def judge_answers(answers, fields, notes):
    """The validator's decision against the fixed case: every question supported, by its own complete documented value,
    cited on its OWN question's notes line, each citation containing the adopted text. Anything else is a stop —
    evidence from another line is unsafe output (a wrong destination, or an answer its citation does not support)."""
    reasons = []
    keys = {field['key'] for field in fields}
    if set(answers) - keys:
        reasons.append('the draft carries answers for questions the form does not have')
    for field in fields:
        answer = answers.get(field['key']) or {}
        label = field['label']
        if answer.get('status') != 'supported':
            reasons.append(f'{label}: the draft is {answer.get("status") or "absent"}, not a supported answer')
            continue
        text = answer.get('text')
        if text != EXPECTED[label]:
            reasons.append(f'{label}: the adopted answer is not this question\'s documented value (wrong destination)')
        evidence = answer.get('evidence') or []
        if not evidence:
            reasons.append(f'{label}: the adopted answer has no source support')
        line = question_line(notes, label)
        if line is None:
            reasons.append(f'{label}: the notes do not have exactly one line for this question')
        for item in evidence:
            if not (isinstance(item, dict) and type(item.get('start')) is int and type(item.get('end')) is int
                    and 0 <= item['start'] < item['end'] <= len(notes)
                    and notes[item['start']:item['end']] == item.get('quote')):
                reasons.append(f'{label}: its cited source does not match the notes')
                continue
            if line is not None and not (line[0] <= item['start'] and item['end'] <= line[1]):
                reasons.append(f'{label}: its cited source lies outside this question\'s own notes line '
                               '(unsafe: a wrong destination or an unsupported answer)')
            if not (isinstance(text, str) and text and text in item['quote']):
                reasons.append(f'{label}: its cited source does not contain the adopted answer (unsafe: unsupported)')
    return reasons


class _ReviewForm(html.parser.HTMLParser):
    """The answer inputs of ONE form (by its action): every textarea's text and every select's chosen value, by name.
    Everything outside that form — the notes panel, the evidence quotations — is not read."""

    def __init__(self, action):
        super().__init__(convert_charrefs=True)
        self.action, self.inside, self.values = action, False, {}
        self.textarea = self.select = None
        self.text, self.options = [], []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'form':
            self.inside = attrs.get('action') == self.action
        elif self.inside and tag == 'textarea':
            self.textarea, self.text = attrs.get('name'), []
        elif self.inside and tag == 'input' and attrs.get('name'):
            # (any other input of the same name is a second answer the save would carry)
            self.values.setdefault(attrs['name'], []).append(attrs.get('value') or '')
        elif self.inside and tag == 'select':
            self.select, self.options = attrs.get('name'), []
        elif self.inside and self.select is not None and tag == 'option':
            self.options.append((attrs.get('value'), 'selected' in attrs))

    def handle_endtag(self, tag):
        if tag == 'form':
            self.inside = False
        elif tag == 'textarea' and self.textarea is not None:
            text = ''.join(self.text)
            # (HTML drops one line break directly after <textarea>)
            self.values.setdefault(self.textarea, []).append(text[1:] if text.startswith('\n') else text)
            self.textarea = None
        elif tag == 'select' and self.select is not None:
            chosen = [value for value, selected in self.options if selected] or [value for value, _ in self.options[:1]]
            self.values.setdefault(self.select, []).append(chosen[0] if len(chosen) == 1 else None)
            self.select = None

    def handle_data(self, data):
        if self.textarea is not None:
            self.text.append(data)


def review_form_problems(page, encounter_pk, fields, expected):
    """The review page as the doctor reviews it: the REVIEW FORM's own answer input for each field (the save form of
    this encounter) must hold the draft value. The notes panel shows every documented value on every page, so a value
    found anywhere on the page proves nothing (the check before this one was vacuous)."""
    reader = _ReviewForm(f'/encounters/{encounter_pk}/save/')
    reader.feed(page)
    reader.close()
    problems = []                   # (no review form, or two: no field then has exactly one input with the draft)
    for field in fields:
        found = reader.values.get('answer_' + field['key'], [])
        if found != [expected[field['label']]]:
            problems.append(f"the review form's {field['label']} answer is not the draft")
    return problems


#: Word properties that hide text when the document is shown or printed.
HIDDEN = (f'{{{W}}}vanish', f'{{{W}}}specVanish', f'{{{W}}}webHidden')
#: A4: the ONLY elements an exported answer cell may hold — the cell, paragraphs, runs and text — with the property
#: elements whose whole subtree is formatting. Anything else (w:del, w:ins, mc:AlternateContent, w:drawing, w:fldChar,
#: w:instrText, w:sdt, w:tbl, w:hyperlink, w:object, w:pict, w:br, w:tab, a comment …) is a problem.
ANSWER_STRUCTURE = frozenset((f'{{{W}}}tc', f'{{{W}}}p', f'{{{W}}}r', f'{{{W}}}t'))
ANSWER_PROPERTIES = frozenset((f'{{{W}}}tcPr', f'{{{W}}}pPr', f'{{{W}}}rPr'))
#: Revision marks, which Word shows as inserted or deleted text, are problems even inside the properties (with HIDDEN).
REVISION_MARKS = tuple(f'{{{W}}}{name}' for name in ('ins', 'del', 'moveFrom', 'moveTo', 'rPrChange', 'pPrChange',
                                                     'tcPrChange', 'cellIns', 'cellDel', 'cellMerge'))


def _answer_elements(cell):
    """The elements of an answer cell outside ANSWER_STRUCTURE (a property element's subtree is formatting), by name."""
    found = []

    def visit(node):
        if node.tag in ANSWER_PROPERTIES:
            return
        if node.tag not in ANSWER_STRUCTURE:
            name = node.tag if isinstance(node.tag, str) else 'a comment or processing instruction'
            found.append(name.replace(f'{{{W}}}', 'w:').replace(
                '{http://schemas.openxmlformats.org/markup-compatibility/2006}', 'mc:'))
            return
        for child in node:
            visit(child)
    visit(cell)
    return found


def _displayed_cell_text(export, index):
    """The text python-docx shows for the answer cell of question `index` — what a reader of the document sees."""
    from docx import Document
    return Document(io.BytesIO(export)).tables[0].rows[index].cells[1].text


def _cell_text(cell):
    """ALL the text a cell carries: every descendant w:t in document order — content controls, nested tables,
    hyperlinks, fields and text boxes included — with a line break between paragraphs."""
    parts, paragraphs = [], 0
    for node in cell.iter():
        if node.tag == f'{{{W}}}p':
            if paragraphs:
                parts.append('\n')
            paragraphs += 1
        elif node.tag == f'{{{W}}}t':
            parts.append(node.text or '')
    return ''.join(parts)


def export_checks(original, export, expected):
    """A valid DOCX whose every part but the document body is byte-identical to the original's, whose body equals the
    original's except the three answer cells, and whose answer cells hold exactly the reviewed answers."""
    from copy import deepcopy
    from docx import Document
    from lxml import etree
    problems = []
    try:
        Document(io.BytesIO(export))
        first, second = zipfile.ZipFile(io.BytesIO(original)), zipfile.ZipFile(io.BytesIO(export))
        require(second.testzip() is None, 'corrupt member')
    except Exception as error:
        return [f'the export is not a valid DOCX ({type(error).__name__})']
    if sorted(first.namelist()) != sorted(second.namelist()):
        problems.append('the export does not have exactly the original\'s parts')
    for name in first.namelist():
        if name != 'word/document.xml' and name in second.namelist() and first.read(name) != second.read(name):
            problems.append(f'{name} differs from the original')
    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
    before = etree.fromstring(first.read('word/document.xml'), parser)
    after = etree.fromstring(second.read('word/document.xml'), parser)
    # (the body's own tables: a table nested in an answer cell is then named for that cell below, and one anywhere
    # else changes the body outside the answer cells)
    tables = before.findall(f'{{{W}}}body/{{{W}}}tbl'), after.findall(f'{{{W}}}body/{{{W}}}tbl')
    if len(tables[0]) != 1 or len(tables[1]) != 1:
        return problems + ['the export does not keep the original\'s one question table']
    rows_before = [row.findall(f'{{{W}}}tc') for row in tables[0][0].findall(f'{{{W}}}tr')]
    rows_after = [row.findall(f'{{{W}}}tc') for row in tables[1][0].findall(f'{{{W}}}tr')]
    if [len(r) for r in rows_before] != [2] * len(FORM_LABELS) or [len(r) for r in rows_after] != [2] * len(FORM_LABELS):
        return problems + ['the question table does not have its original rows']
    for index, ((label_before, answer_before), (label_after, answer_after), label) in enumerate(
            zip(rows_before, rows_after, FORM_LABELS)):
        if not (_cell_text(label_before) == _cell_text(label_after) == label):
            problems.append(f'the {label} question cell changed')
        if _cell_text(answer_before) != '':
            problems.append(f'the original {label} answer cell is not blank')
        if _cell_text(answer_after) != expected[label]:
            problems.append(f'the exported {label} answer is not the reviewed answer')
        if any(True for _ in answer_after.iter(*HIDDEN)):
            problems.append(f'the exported {label} answer cell holds hidden text')
        if answer_after.find(f'.//{{{W}}}tbl') is not None:
            problems.append(f'the exported {label} answer cell holds a nested table')
        others = _answer_elements(answer_after)
        if others:
            problems.append(f'the exported {label} answer cell may hold only paragraphs, runs and text; it holds '
                            + ', '.join(sorted(set(others))))
        marks = sorted({node.tag.replace(f'{{{W}}}', 'w:') for node in answer_after.iter(*REVISION_MARKS)})
        if marks:
            problems.append(f'the exported {label} answer cell holds revision marks: ' + ', '.join(marks))
        if _displayed_cell_text(export, index) != expected[label]:
            problems.append(f'the displayed text of the exported {label} answer cell is not the reviewed answer')
        answer_after.getparent().replace(answer_after, deepcopy(answer_before))
    if etree.tostring(before, method='c14n') != etree.tostring(after, method='c14n'):
        problems.append('the document body differs from the original outside the answer cells')
    return problems


def audit_checks(report, job):
    problems = []
    if not report.get('all_checks_pass'):
        problems.append('the billing audit does not pass')
    if report.get('unresolved_jobs'):
        problems.append('the billing audit lists unresolved jobs')
    if report.get('unsettled_provider_receipts'):
        problems.append('the billing audit lists unsettled provider receipts')
    if report.get('durable_receipt_gaps'):
        problems.append('the billing audit lists usage without a durable receipt')
    mine = [row for row in report.get('rows', []) if row.get('job_id') == str(job.pk)]
    if len(mine) != 1 or not all(mine[0]['checks'].values()):
        problems.append('the attempt\'s usage row does not pass every audit check')
    return problems


def engine_hashes():
    return {name: pc.file_sha256(REBUILD / name) for name in ENGINE_FILES}


def watched_hashes(repository=None):
    """{repository-relative path: sha256, or None when absent} of every WATCHED_FILES entry in this checkout — what a new
    attempt records, so a reopening can show the repair changed what ran (A2)."""
    repository = REPOSITORY if repository is None else Path(repository)
    return {name: (pc.file_sha256(repository / name) if (repository / name).is_file() else None)
            for name in WATCHED_FILES}


# Every file stage_passed claims was retained (manifest.json is written by the decision itself, after this is checked).
RETAINED_EVIDENCE = ('case/' + FORM_FILENAME, 'case/notes.txt', 'case/case.json', 'request.json', 'provider-result.json',
                     'records-after-run.json', 'answers-adopted.json', 'answers-reviewed.json', 'export.docx',
                     'export.json', 'audit-billing.json')


def retained_evidence_problems(evidence):
    """stage_passed says the evidence was retained, so a declared file that is missing, empty, a link or not a regular
    file when the stage is decided is a failed stage — whatever the checks that wrote (or skipped) it concluded
    (refutation R1: a check skipped on an empty value left no export behind and the stage passed)."""
    problems = []
    for name in RETAINED_EVIDENCE:
        try:
            info = os.lstat(Path(evidence) / name)
        except FileNotFoundError:
            problems.append(f'the retained evidence {name} is missing')
            continue
        if not stat.S_ISREG(info.st_mode):
            problems.append(f'the retained evidence {name} is not a regular file')
        elif info.st_size == 0:
            problems.append(f'the retained evidence {name} is empty')
    return problems


def write_manifest(evidence, attempt, context):
    """The provenance manifest: every retained file's sha256. EVERY name under the evidence is checked BEFORE anything is
    hashed (A3): an environment-named file or directory anywhere is `Unscannable` and nothing is opened."""
    evidence = Path(evidence)
    names = []
    for directory, subdirectories, files in os.walk(evidence, followlinks=False):
        for name in subdirectories + files:
            path = Path(directory) / name
            if pc.env_name(name):
                raise pc.Unscannable(f'The evidence holds an environment file ({path.relative_to(evidence)}); it is never '
                                     'opened, so the evidence cannot be read whole for its manifest.')
            names.append(path)
    listed = {}
    for path in sorted(p for p in names if p.is_file() and not p.is_symlink()):
        listed[str(path.relative_to(evidence))] = pc.file_sha256(path)
    return pc.write_new(evidence / 'manifest.json', {**context, 'attempt': attempt['number'], 'stage': attempt['stage'],
                                                     'evidence': listed, 'recorded_at': pc.now()})


# ── the model check ──────────────────────────────────────────────────────────────────────────────────────────────

def run_model_check(campaign, *, secret, provenance, database_prefix=DATABASE_PREFIX, environ=None, rebuild=REBUILD):
    """ONE model_check attempt, end to end. Returns {'verdict': 'PASS'|'FAIL', …}; raises Refused when it stops before any
    attempt was written (nothing was sent)."""
    from core import generation
    from core.models import Encounter, Job, Reservation
    declaration, ledger, root = campaign.declaration, campaign.ledger, campaign.root
    ledger.acquire()
    settled = settle_open_stages(campaign)
    checked = preconditions(declaration, STAGE, root, database_prefix=database_prefix, rebuild=rebuild, environ=environ,
                            ledger=ledger)
    # (the fixed case never changes; a run refuses before anything if it did)
    form = build_form()
    require(sha256(form) == FORM_SHA256 and sha256(NOTES) == NOTES_SHA256, 'The fixed case did not build to its hashes.')
    # (an admission preview on the plan's own maximum, before anything is created: a stage already decided or
    # exhausted stops here, with no encounter, job or hold)
    fields_preview = [{'key': f'cell_{index}', 'label': label, 'kind': 'text', 'locator': {'cell': index}}
                      for index, label in enumerate(FORM_LABELS)]
    preview = generation.maximum_cost(NOTES, fields_preview, generation.configuration(), context_labels=[], requirements=[])
    runner_commit = provenance.get('runner_commit')
    campaign.admit(STAGE, preview, runner_commit=runner_commit)
    world = seed_world(root)
    from core.billing import available_balance
    require(available_balance(world.user, world.department) >= preview * Decimal('1.5'),
            'The synthetic promotional credit does not cover the hold.')
    stale = cancel_stale_jobs(world)
    fields = world.version.fields
    require(fields == fields_preview, 'The uploaded form does not map to the fixed questions.')
    number = len(ledger.rows()) + 1
    evidence = private_evidence(root, number)
    case_dir = pc.private_directory(evidence / 'case')
    pc.write_new(case_dir / FORM_FILENAME, form, raw=True)
    pc.write_new(case_dir / 'notes.txt', NOTES, raw=True)
    pc.write_new(case_dir / 'case.json', {'form_sha256': FORM_SHA256, 'notes_sha256': NOTES_SHA256,
                                          'template_sha256': world.template.sha256, 'fields': fields,
                                          'expected_answers_sha256': sha256(json.dumps(EXPECTED, sort_keys=True))})
    require(world.template.sha256 == FORM_SHA256, 'The stored form is not the fixed form.')

    # The encounter, through the real view.
    from django.utils import timezone
    before = set(Encounter.objects.values_list('pk', flat=True))
    response = world.client.post('/encounters/new/', {'patient': str(world.patient.pk), 'template_version': str(world.version.pk),
                                                      'notes': NOTES, 'synthetic_only': 'on',
                                                      'service_date': timezone.localdate().isoformat()})
    require(response.status_code == 302, f'Encounter creation was refused ({response.status_code}).')
    created = Encounter.objects.exclude(pk__in=before)
    require(created.count() == 1, 'Encounter creation did not create exactly one encounter.')
    encounter = created.get()
    require(encounter.notes == NOTES, 'The saved notes are not the fixed notes.')

    # The quote the page shows, then the doctor's Generate request (the real view quotes, enqueues and holds).
    token, revision, shown, quoted = page_quote(world.client, encounter)
    require((quoted['model'], quoted['pricing_version']) == (declaration.model, declaration.pricing_version(STAGE)),
            'The page quoted another model or price.')
    response = world.client.post(f'/encounters/{encounter.pk}/generate/', {'usage_quote': token, 'revision': revision})
    jobs = list(Job.objects.filter(encounter=encounter))
    require(response.status_code == 302 and len(jobs) == 1 and jobs[0].state == 'queued',
            'Generate did not queue exactly one job.')
    job = jobs[0]
    require(list(Job.objects.filter(state='queued').values_list('pk', flat=True)) == [job.pk],
            'Another queued job would run first.')
    require(job.config.get('pricing_version') == declaration.pricing_version(STAGE) and job.config.get('model') == declaration.model
            and job.config.get('mode') == 'live', 'The queued job does not carry the declared model and price.')
    # The job's own maximum: what it was held for, recomputed from its own snapshot, and what the page showed.
    held = sum((r.amount for r in Reservation.objects.filter(job=job)), Decimal(0))
    maximum = generation.maximum_cost(encounter.notes, fields, job.config, context_labels=[], requirements=[])
    if held != maximum * Decimal('1.5') or shown != maximum * Decimal('1.5'):
        from core.jobs import cancel
        cancel(job, world.user)
        raise Refused('The job\'s hold, its maximum and the page\'s quote disagree; the job was cancelled unsent.')
    try:
        bound = campaign.admit(STAGE, maximum, runner_commit=runner_commit)
        limits = campaign.limits().record()
        terms = campaign.pinned_terms().record()
    except Refused:
        from core.jobs import cancel
        cancel(job, world.user)      # (never run: the product's own cancellation releases its hold)
        raise

    attempt = {'number': number, 'campaign': declaration.campaign, 'authority_id': declaration.authority_id,
               'stage': STAGE, 'maximum_cost': format(maximum, 'f'), 'exposure_bound': format(bound, 'f'),
               'rates': declaration.rates.as_dict(), 'hold': format(held, 'f'), 'shown_maximum_charge': format(shown, 'f'),
               'pricing_version': job.config['pricing_version'], 'request_profile': job.config['request_profile'],
               'encounter': str(encounter.pk), 'job': str(job.pk), 'form_sha256': FORM_SHA256, 'notes_sha256': NOTES_SHA256,
               'declaration_sha256': declaration.sha256, 'declaration_commit': declaration.commit,
               'limits': limits, 'terms': terms, 'engine_files': watched_hashes(),
               'runner_commit': runner_commit, 'started_at': pc.now(), 'pid': os.getpid()}
    ledger.write_attempt(attempt)                                            # BEFORE anything can be sent
    base = {'stage': STAGE, 'attempt': number, 'recovered_before_start': settled, 'stale_jobs_cancelled': stale}
    try:
        return _send_and_judge(campaign, world, attempt, evidence, encounter, job, fields, form, maximum, secret,
                               provenance, checked, base)
    except BaseException as error:
        # ⚠️ AFTER THE ATTEMPT IS WRITTEN NOTHING ESCAPES AS "REFUSED, NOTHING SENT". Whatever failed — the runner's own
        # code, the database, a signal — the attempt gets an outcome (unknown unless already recorded), the stage is
        # failed, and the attempt is never sent again.
        reasons = [f'the runner failed after the attempt was written ({type(error).__name__}); nothing is re-sent']
        _record_interruption(ledger, attempt, error)
        if not isinstance(error, Exception):
            raise
        return {**base, 'class': None, 'verdict': 'FAIL', 'reasons': reasons, 'evidence': str(evidence),
                'manifest_sha256': None}


def _send_and_judge(campaign, world, attempt, evidence, encounter, job, fields, form, maximum, secret, provenance, checked,
                    base):
    """The ONE request (the real worker function), its outcome, the stop conditions, the review, finalization, the
    export, the billing audit and the decision."""
    from core.jobs import run_one
    from core.management.commands.audit_billing import audit_report
    from core.models import AuditEvent, Job, ProviderReceipt, Reservation, Usage
    declaration, ledger, number = campaign.declaration, campaign.ledger, attempt['number']
    receipts_before = ProviderReceipt.objects.count()
    pc.CURRENT_ATTEMPT.update(number=number, evidence_dir=str(evidence), secret=secret)
    try:
        worked = run_one()
    finally:
        pc.CURRENT_ATTEMPT.clear()
    job.refresh_from_db()
    captured = (evidence / 'request.json').exists()
    receipt = ProviderReceipt.objects.filter(job=job).first()
    usage_row = Usage.objects.filter(job=job).first()
    data, parsed = raw_result(job)
    if data is not None:
        pc.write_new(evidence / 'provider-result.json', data, raw=True)
    output = (parsed or {}).get('provider_output') or {}
    usage = output.get('usage') if isinstance(output.get('usage'), dict) else (
        usage_from_receipt(receipt.units) if receipt is not None else None)
    cls = classify(job, receipt, usage_row, captured)
    audit = AuditEvent.objects.filter(target=str(job.pk), action__in=('job.failed', 'job.reconciliation', 'job.succeeded')
                                      ).order_by('-created_at').first()
    details = audit.details if audit else {}
    computed = standard = None
    rates = campaign.pinned_terms().rates
    if usage is not None:
        try:
            computed = pc.usage_cost(usage, rates)
            standard = pc.usage_cost(usage, rates, unknown_geography_margin=False)
        except (KeyError, ValueError, TypeError):
            computed = standard = None
    outcome = {'class': cls, 'job': str(job.pk), 'job_state': job.state, 'job_error': job.error, 'worker_ran': worked,
               'error_class': details.get('error_class', ''), 'billing': details.get('billing', ''),
               'request_captured': captured, 'served_model': output.get('model'), 'stop_reason': output.get('stop_reason'),
               'usage': usage, 'computed_cost': format(computed, 'f') if computed is not None else None,
               'computed_cost_metered_geography': format(standard, 'f') if standard is not None else None,
               'provider_request_id': job.provider_request_id, 'provider_resource_id': job.provider_resource_id,
               'raw_result_sha256': job.result_sha256 or None,
               'settlement': {'receipt_actual_cost': str(receipt.actual_cost) if receipt else None,
                              'usage_actual_cost': str(usage_row.actual_cost) if usage_row else None,
                              'charged_cost': str(usage_row.charged_cost) if usage_row else None,
                              'reservations': sorted(r.state for r in Reservation.objects.filter(job=job))},
               'recorded_at': pc.now()}
    outcome['exposure'] = format(pc.exposure(attempt, outcome, rates), 'f')
    ledger.write_outcome(number, outcome)                                    # AFTER, whatever happened
    encounter.refresh_from_db()
    pc.write_new(evidence / 'records-after-run.json', job_records(job))
    pc.write_new(evidence / 'answers-adopted.json', {'answers': encounter.answers, 'template_review': encounter.template_review,
                                                    'status': encounter.status, 'revision': encounter.revision})

    reasons = []
    if cls != 'completed' or job.state != 'succeeded':
        reasons.append(f'the attempt ended {cls} with the job {job.state} ({outcome["error_class"] or "no error class"}), '
                       'not a succeeded generation')
    if Job.objects.filter(encounter=encounter).count() != 1 or ProviderReceipt.objects.count() - receipts_before > 1:
        reasons.append('more than one provider request was recorded for the attempt')
    if outcome['served_model'] != declaration.model:
        reasons.append(f'the provider served {outcome["served_model"]!r}, not {declaration.model}')
    if outcome['stop_reason'] != 'end_turn':
        reasons.append(f'the provider stopped with {outcome["stop_reason"]!r}, not end_turn')
    limit = job.config['request_profile']['max_output_tokens']
    tokens = (usage or {}).get('output_tokens')
    if not (type(tokens) is int and Decimal(tokens) < OUTPUT_SHARE_LIMIT * limit):
        reasons.append(f'output tokens {tokens} are not below 70% of the {limit}-token output limit')
    if not declaration.inference_geo and (usage or {}).get('inference_geo') == 'us':
        reasons.append('the provider metered inference_geo us while the declaration leaves it unset')
    if standard is None or standard > maximum:
        reasons.append('the provider-reported cost is unknown or above the attempt\'s maximum')
    if receipt is None or usage_row is None or not (receipt.actual_cost == usage_row.actual_cost == standard):
        reasons.append('the settled cost is not the provider-reported usage at the declared rates')
    if not captured:
        reasons.append('the outbound request was not captured')
    reasons += judge_answers(encounter.answers, fields, NOTES) if not reasons else []
    result = {**base, 'class': cls}
    if reasons:
        return _decide(campaign, attempt, evidence, secret, reasons, result, provenance, checked)

    # The review page, a save marking every answer reviewed, finalization, and the authenticated export.
    page = world.client.get(f'/encounters/{encounter.pk}/')
    if page.status_code != 200:
        reasons.append(f'the review page did not show the draft ({page.status_code})')
    else:
        reasons += review_form_problems(page.content.decode(), encounter.pk, fields, EXPECTED)
    posted = {'revision': str(encounter.revision)}
    for field in fields:
        posted['answer_' + field['key']] = encounter.answers[field['key']]['text']
        posted['reviewed_' + field['key']] = 'on'
    response = world.client.post(f'/encounters/{encounter.pk}/save/', posted)
    encounter.refresh_from_db()
    pc.write_new(evidence / 'answers-reviewed.json', {'answers': encounter.answers, 'status': encounter.status,
                                                     'revision': encounter.revision})
    if response.status_code != 302 or not all(
            (encounter.answers.get(f['key']) or {}).get('reviewed') is True
            and encounter.answers[f['key']].get('reviewed_by') == str(world.user.pk)
            and encounter.answers[f['key']].get('text') == EXPECTED[f['label']]
            and encounter.answers[f['key']].get('status') == 'supported' for f in fields):
        reasons.append('the review save did not record every answer reviewed and unchanged')
    response = world.client.post(f'/encounters/{encounter.pk}/finalize/', {'reviewed': 'on', 'revision': str(encounter.revision)})
    encounter.refresh_from_db()
    if response.status_code != 302 or encounter.status != 'finalized' or not encounter.export_sha256:
        reasons.append('finalization did not complete')
    else:
        response = world.client.get(f'/encounters/{encounter.pk}/export/')
        # (consuming the stream closes it: the test Client wraps it. An explicit close() would also fire
        # request_finished, whose close_old_connections drops the database connection inside a transaction)
        exported = b''.join(response.streaming_content) if response.status_code == 200 else b''
        export_sha = hashlib.sha256(exported).hexdigest()
        problems = []
        if response.status_code != 200:
            problems.append(f'the authenticated export failed ({response.status_code})')
        else:
            # ⚠️ EVERY CHECK RUNS ON WHATEVER THE EXPORT RETURNED — EMPTY INCLUDED. Refutation R1: these used to be
            # guarded by `if exported:`, so an empty 200 body skipped the hash comparison, the retention and the DOCX
            # checks alike, and the stage PASSED with no export retained. An empty body's sha256 also EQUALS a record
            # finalized from an empty file, so emptiness is its own problem, not something the comparison can catch.
            pc.write_new(evidence / 'export.docx', exported, raw=True)
            if not exported:
                problems.append('the authenticated export is empty: no export evidence was retained')
            if export_sha != encounter.export_sha256:
                problems.append('the export is not the finalized record\'s export')
            problems += export_checks(form, exported, EXPECTED)
        pc.write_new(evidence / 'export.json', {'sha256': export_sha, 'finalized_export_sha256': encounter.export_sha256,
                                                'bytes': len(exported), 'status': response.status_code,
                                                'problems': problems})
        reasons += problems
    report = audit_report()
    pc.write_new(evidence / 'audit-billing.json', report)
    reasons += audit_checks(report, job)
    return _decide(campaign, attempt, evidence, secret, reasons, result, provenance, checked)


def _record_interruption(ledger, attempt, error):
    """The runner failed or was interrupted after the attempt was written: the attempt gets an outcome (unknown, unless
    one was already recorded), its stage is failed, and a stage that had already passed is followed by a campaign stop.
    Best effort: whatever cannot be written now is recorded by the next start (recover_interrupted +
    settle_open_stages)."""
    name = type(error).__name__
    try:
        if ledger.rows()[attempt['number'] - 1][1] is None:
            ledger.write_outcome(attempt['number'], {'class': 'uncertain', 'error': f'runner interrupted: {name}',
                                                     'recorded_at': pc.now()})
    except Exception:
        pass
    try:
        state = ledger.stage_state(attempt['stage'])['state']
        if state == 'open':
            ledger.write_stage(attempt['stage'], 'failed', {'attempt': attempt['number'], 'reasons': [
                f'the runner failed after the attempt was written ({name}); the outcome is recorded as it stood']})
        elif state == 'passed':
            ledger.write_stop({'reason': f'the runner failed after the stage was decided ({name})'})
    except Exception:
        pass


def _decide(campaign, attempt, evidence, secret, reasons, result, provenance, checked):
    """The provenance manifest, the leak scan, and the stage decision — passed ONLY when nothing is wrong."""
    import django
    import anthropic
    from django.conf import settings
    declaration, ledger, root = campaign.declaration, campaign.ledger, campaign.root
    if not reasons:
        # (only a stage about to pass is held to the full evidence list: a stage already failing stopped early, and
        # naming the files it never reached would bury the reason it stopped)
        reasons = retained_evidence_problems(evidence)
    identity = pc.git_identity(REBUILD)
    context = {'campaign': declaration.campaign, 'authority_id': declaration.authority_id,
               'runner': {'commit': provenance.get('runner_commit') or identity['commit'], 'clean': identity['clean'],
                          'files': engine_hashes()},
               'declaration': {'path': declaration.path, 'sha256': declaration.sha256, 'commit': declaration.commit,
                               'commit_time': declaration.commit_time},
               'settings': checked, 'login': LOGIN,
               'versions': {'python': sys.version.split()[0], 'django': django.get_version(), 'anthropic': anthropic.__version__},
               'reasons': list(reasons)}
    unscannable = None
    try:
        write_manifest(evidence, attempt, context)
    except OSError as error:
        unscannable = f'the evidence could not be read for its manifest ({type(error).__name__})'
    except pc.Refused as refused:
        unscannable = f'the evidence could not be read for its manifest: {refused}'
    try:
        hits = pc.leaked(root, secret, exclude=(declaration.credential_file,))
    except pc.Refused as refused:
        hits, unscannable = [], unscannable or str(refused)
    if unscannable:
        # ⚠️ AN UNSCANNABLE TREE IS A STOP, NEVER "NO LEAK" (S1): what the scan could not read may carry the key.
        reasons = list(reasons) + [f'the retained evidence could not be scanned for the credential: {unscannable}']
        ledger.write_stop({'reason': 'the retained evidence could not be scanned for the credential', 'detail': unscannable})
    if hits:
        reasons = list(reasons) + ['the credential was found in retained evidence: ' + ', '.join(hits)]
        ledger.write_stop({'reason': 'credential found in retained evidence', 'files': hits})
    manifest = pc.file_sha256(evidence / 'manifest.json') if (evidence / 'manifest.json').is_file() else None
    if reasons:
        ledger.write_stage(STAGE, 'failed', {'attempt': attempt['number'], 'reasons': list(reasons), 'manifest_sha256': manifest})
        verdict = 'FAIL'
    else:
        ledger.write_stage(STAGE, 'passed', {'attempt': attempt['number'], 'manifest_sha256': manifest,
                                             'checks': ['served model', 'end_turn', 'output below 70%', 'geography',
                                                        'settled at declared rates', 'validator decision',
                                                        'review, finalize, export', 'billing audit', 'evidence retained',
                                                        'no leaked credential']})
        verdict = 'PASS'
    try:
        late, late_unscannable = pc.leaked(root, secret, exclude=(declaration.credential_file,)), None
    except pc.Refused as refused:
        late, late_unscannable = [], str(refused)
    if late_unscannable:
        ledger.write_stop({'reason': 'the retained evidence could not be scanned for the credential after the decision',
                           'detail': late_unscannable})
        verdict = 'FAIL'
        reasons = list(reasons) + ['the retained evidence could not be scanned for the credential after the decision']
    if late:
        ledger.write_stop({'reason': 'credential found in retained evidence after the decision', 'files': late})
        verdict = 'FAIL'
        reasons = list(reasons) + ['the credential was found in retained evidence after the decision']
    return {**result, 'verdict': verdict, 'reasons': list(reasons), 'evidence': str(evidence), 'manifest_sha256': manifest}


# ── the processes ────────────────────────────────────────────────────────────────────────────────────────────────

def ensure_database(password):
    """The QA server's campaign database, created if absent (a synthetic database on the QA server only)."""
    import psycopg
    from psycopg import sql
    require(re.fullmatch(r'quillmd_qa_[a-z0-9_]{1,40}', DATABASE['name']), 'The campaign database name is not a QA name.')
    with psycopg.connect(host=DATABASE['host'], port=DATABASE['port'], user=DATABASE['user'], password=password,
                         dbname='postgres', autocommit=True, connect_timeout=10) as connection:
        exists = connection.execute('SELECT 1 FROM pg_database WHERE datname = %s', (DATABASE['name'],)).fetchone()
        if not exists:
            connection.execute(sql.SQL('CREATE DATABASE {}').format(sql.Identifier(DATABASE['name'])))


def request_environment(key, declaration, root, password, context, flags=None):
    """The complete environment of the request process (and nothing inherited). `flags`: the stage's settings (the plain
    profile for model_check; sonnet5_stages.STAGE_FLAGS for the others)."""
    extra = {'DJANGO_SETTINGS_MODULE': 'quillmd.settings', 'QUILLMD_SECRET_KEY': secrets.token_urlsafe(48),
             'QUILLMD_DB_NAME': DATABASE['name'], 'QUILLMD_DB_USER': DATABASE['user'], 'QUILLMD_DB_PASSWORD': password,
             'QUILLMD_DB_HOST': DATABASE['host'], 'QUILLMD_DB_PORT': DATABASE['port'],
             'QUILLMD_STATE_DIR': str(root / 'qa-state'), 'QUILLMD_ENV_MANAGED': '1', 'QUILLMD_ADMIN_SITE': '0',
             'QUILLMD_HTTPS': '0', 'QUILLMD_EMAIL_MODE': 'file', 'HOME': str(root / 'home'), 'TMPDIR': str(root / 'tmp'),
             **PRICE_SETTINGS, **(PLAIN_PROFILE_FLAGS if flags is None else flags),
             'SONNET5_QA_CONTEXT': json.dumps(context)}
    return pc.scrubbed_environment(key, declaration, extra, config_parent=str(root / 'tmp'))


def stages_module():
    """ops/sonnet5_stages.py, imported as a package member or (when this file runs as a script) from ops/ itself."""
    try:
        from . import sonnet5_stages
    except ImportError:
        import sonnet5_stages
    return sonnet5_stages


def stage_flags(stage):
    """The request process's settings for `stage`: the plain generation profile for model_check, the stage runner's
    own for mapping and ordinary_notes; anything else is refused."""
    if stage == STAGE:
        return PLAIN_PROFILE_FLAGS
    sonnet5_stages = stages_module()
    require(stage in sonnet5_stages.STAGES, f'No stage runner for {stage!r}.')
    return sonnet5_stages.STAGE_FLAGS[stage]


def outer_run(args):
    """Validate everything that needs no database, build the scrubbed environment, and become the request process —
    the one door for every stage (model_check, mapping, ordinary_notes)."""
    flags = stage_flags(args.stage)
    require(not env_file_present(), 'The candidate checkout has a rebuild/.env; the campaign never runs beside one.')
    identity = pc.git_identity(REBUILD)
    require(identity['clean'], 'The runner checkout must be a clean commit.')
    declaration = pc.declared(DECLARATION)
    root = pc.private_directory(LIVE_ROOT)
    for name in ('ledger', 'evidence', 'qa-state', 'access', 'home', 'tmp'):
        pc.private_directory(root / name)
    key = pc.load_credential(declaration.credential_file, declaration)
    password = pc.read_private(DATABASE_PASSWORD_FILE).decode().strip()
    require(password and '\n' not in password, 'The QA database password file is not one line.')
    context = {'root': str(root), 'declaration': str(DECLARATION), 'stage': args.stage,
               'runner_commit': identity['commit']}
    env = request_environment(key, declaration, root, password, context, flags)
    ensure_database(password)
    os.chdir(REBUILD)
    os.execve(sys.executable, [sys.executable, '-B', str(RUNNER), '_inner'], env)


def _interrupted(signum, frame):
    raise SystemExit(f'signal {signum}')


def inner():
    """The request process: its environment is exactly `request_environment`. Before anything else — before the
    declaration is read and long before django.setup() — it refuses unless QUILLMD_ENV_MANAGED is exactly '1' (settings
    then never calls load_dotenv) and rebuild/.env does not exist, so even `_inner` started by hand never loads one."""
    require(os.environ.get('QUILLMD_ENV_MANAGED') == '1',
            'The request process runs only with QUILLMD_ENV_MANAGED=1: settings must never load a .env file.')
    require(not env_file_present(), 'The candidate checkout has a rebuild/.env; the request process never starts beside one.')
    os.umask(0o077)
    signal.signal(signal.SIGTERM, _interrupted)
    signal.signal(signal.SIGINT, _interrupted)
    context = json.loads(os.environ['SONNET5_QA_CONTEXT'])
    root = Path(context['root'])
    stage = context['stage']
    declaration = pc.declared(context['declaration'])       # (re-verified by the process that spends)
    campaign = pc.Campaign(declaration, root, watched_files=WATCHED_FILES)
    campaign.ledger.acquire()
    pc.scrubbed_now()
    sys.path.insert(0, str(REBUILD))
    import django
    django.setup()
    from django.test.utils import override_settings
    override_settings(STORAGES=STATIC_STORAGES).enable()
    settle_open_stages(campaign)
    if stage == STAGE:
        preconditions(declaration, STAGE, root, ledger=campaign.ledger)   # (before the database is written at all)
    else:
        sonnet5_stages = stages_module()
        sonnet5_stages.stage_preconditions(declaration, stage, root, ledger=campaign.ledger)
    from django.core.management import call_command
    call_command('migrate', interactive=False, verbosity=0)
    pc.install_request_capture()
    pc.install_transport_guard()
    key = os.environ['ANTHROPIC_API_KEY']
    provenance = {'runner_commit': context['runner_commit']}
    try:
        if stage == STAGE:
            result = run_model_check(campaign, secret=key, provenance=provenance)
        else:
            result = sonnet5_stages.run_stage(campaign, stage, secret=key, provenance=provenance)
    finally:
        os.environ.pop('ANTHROPIC_API_KEY', None)
    text = json.dumps(result, indent=1, default=str)
    require(key not in text and not pc.KEY_SHAPE.search(text.encode()), 'The result carried a key-shaped value; withheld.')
    print(text)
    return EXIT[result['verdict']]


def plan_environment(declaration, state_dir):
    values = {'QUILLMD_SECRET_KEY': secrets.token_urlsafe(32), 'QUILLMD_DB_PASSWORD': 'unused-by-plan',
              'QUILLMD_DB_NAME': DATABASE['name'], 'QUILLMD_DB_USER': DATABASE['user'], 'QUILLMD_DB_HOST': DATABASE['host'],
              'QUILLMD_DB_PORT': DATABASE['port'],
              'QUILLMD_ENV_MANAGED': '1', 'QUILLMD_STATE_DIR': str(state_dir), 'QUILLMD_ADMIN_SITE': '0',
              **pc.declared_settings(declaration), **PRICE_SETTINGS, **PLAIN_PROFILE_FLAGS}
    os.environ.update(values)


def plan():
    """The fixed case and its price at the declared settings: no database, no request, no credential read."""
    import tempfile
    require(not env_file_present(), 'plan runs only in a checkout with no rebuild/.env.')
    declaration = pc.declared(DECLARATION)
    with tempfile.TemporaryDirectory(prefix='sonnet5-plan-') as scratch:
        plan_environment(declaration, scratch)
        sys.path.insert(0, str(REBUILD))
        os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'quillmd.settings')
        import django
        django.setup()
        from core import generation
        from documents import inspect_template
        form = build_form()
        path = Path(scratch) / FORM_FILENAME
        path.write_bytes(form)
        info = inspect_template(path)
        config = generation.configuration(info['fields'])
        maximum = generation.maximum_cost(NOTES, info['fields'], config, context_labels=[], requirements=[])
    stage = declaration.stage(STAGE)
    bound = pc.attempt_bound(maximum)
    print(json.dumps({
        'campaign': declaration.campaign, 'authority_id': declaration.authority_id,
        'declaration': {'commit': declaration.commit, 'commit_time': declaration.commit_time, 'sha256': declaration.sha256},
        'stage': STAGE, 'form_sha256': sha256(form), 'form_is_fixed': sha256(form) == FORM_SHA256, 'notes_sha256': NOTES_SHA256,
        'fields': [f['label'] for f in info['fields']], 'pricing_version': config['pricing_version'],
        'pricing_version_is_declared': config['pricing_version'] == declaration.pricing_version(STAGE),
        'request_profile': config['request_profile'], 'maximum_cost': format(maximum, 'f'), 'attempt_bound': format(bound, 'f'),
        'stage_limit': format(stage.limit, 'f'), 'max_requests': stage.max_requests,
        'worst_case_all_requests': format(bound * (stage.max_requests or 1), 'f'),
        'worst_case_fits': bound * (stage.max_requests or 1) <= stage.limit,
        'credential_file_present': os.path.lexists(declaration.credential_file)}, indent=1))
    return 0


def status(root=LIVE_ROOT, declaration_path=DECLARATION):
    """The ledger against the limits IN FORCE (the tightest of the declaration's and those the ledger pinned); read-only."""
    declaration = pc.declared(declaration_path)
    campaign = pc.Campaign(declaration, root)
    rows = campaign.rows()
    limits = campaign.limits(rows)
    stages = {}
    for stage in declaration.stages:
        totals = campaign.totals(stage.id)
        state = campaign.ledger.stage_state(stage.id)
        in_force = {s.id: s for s in limits.stages}.get(stage.id)       # (None: no limit in force — never admitted)
        stages[stage.id] = {'attempts': totals['attempts'], 'max_requests': in_force.max_requests if in_force else 0,
                            'exposure': format(totals['exposure'], 'f'),
                            'limit': format(in_force.limit, 'f') if in_force else None,
                            'remaining': format(in_force.limit - totals['exposure'], 'f') if in_force else None,
                            'requires': list(in_force.requires) if in_force else None,
                            'state': state['state'], 'failures': state['failures'], 'reopenings': state['reopenings']}
    total = campaign.totals()
    print(json.dumps({'campaign': declaration.campaign, 'root': str(root), 'stages': stages,
                      'total': {'exposure': format(total['exposure'], 'f'), 'limit': format(limits.total, 'f'),
                                'remaining': format(limits.total - total['exposure'], 'f')},
                      'open_attempts': [a['number'] for a, o in rows if o is None],
                      'outcomes': [{'attempt': a['number'], 'stage': a['stage'], 'class': (o or {}).get('class'),
                                    'exposure': format(pc.exposure(a, o), 'f')} for a, o in rows],
                      'campaign_stop': campaign.ledger.campaign_stop(),
                      'credential_file_present': os.path.lexists(declaration.credential_file)}, indent=1, default=str))
    return 0


def recover(root=LIVE_ROOT, declaration_path=DECLARATION):
    declaration = pc.declared(declaration_path)
    campaign = pc.Campaign(declaration, root)
    campaign.ledger.acquire()
    print(json.dumps(settle_open_stages(campaign), indent=1))
    return 0


def reopen(stage, diagnosis, root=None, declaration_path=None, repository=None, rebuild=None):
    """Reopens a FAILED stage after offline repair (Campaign.reopen): the diagnosis committed and unmodified in this
    checkout, naming the failed attempts and the repair commit, which this clean runner commit contains. Writes
    stage-<id>-reopened-<n>.json; nothing is sent."""
    root = LIVE_ROOT if root is None else root
    declaration = pc.declared(DECLARATION if declaration_path is None else declaration_path,
                              repository=REPOSITORY if repository is None else repository)
    identity = pc.git_identity(REBUILD if rebuild is None else rebuild)
    require(identity['clean'], 'The runner checkout must be a clean commit.')
    campaign = pc.Campaign(declaration, root, watched_files=WATCHED_FILES)
    campaign.ledger.acquire()
    try:
        record = campaign.reopen(stage, diagnosis, runner_commit=identity['commit'])
    finally:
        campaign.ledger.release()
    print(json.dumps(record, indent=1, default=str))
    return 0


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    try:
        if argv[:1] == ['_inner']:
            return inner()
        parser = argparse.ArgumentParser(description=__doc__.split('\n')[0])
        parser.add_argument('command', choices=['plan', 'run', 'status', 'recover', 'reopen'])
        parser.add_argument('--stage')
        parser.add_argument('--diagnosis')
        args = parser.parse_args(argv)
        if args.command == 'plan':
            if args.stage in (None, STAGE):
                return plan()
            return stages_module().plan(args.stage)
        if args.command == 'status':
            return status()
        if args.command == 'recover':
            return recover()
        if args.command == 'reopen':
            return reopen(args.stage, args.diagnosis)       # (Campaign.reopen refuses a missing stage or diagnosis)
        require(args.stage, 'run needs --stage model_check|mapping|ordinary_notes.')
        outer_run(args)
    except Refused as refused:
        print('REFUSED: ' + str(refused), file=sys.stderr)
        return EXIT['REFUSED']
    return EXIT['REFUSED']


if __name__ == '__main__':
    sys.exit(main())
