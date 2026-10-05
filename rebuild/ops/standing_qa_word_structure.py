"""Three prospectively frozen project-authored Word cases, through normal QA.

No import reads a credential, creates application state or sends a request.
Canonical elicited/notes/corpus inputs retain their original identities.
"""
import datetime
import hashlib
import io
import json
import os
from pathlib import Path
import re
import zipfile

try:
    from . import provider_campaign as pc, sonnet5_stages as stages, standing_qa_elicited as elicited
except ImportError:
    import provider_campaign as pc, sonnet5_stages as stages, standing_qa_elicited as elicited

STAGE = 'word_structure_qualification'
PRICING = 'ordinary_notes_field_relations'
CASE_IDS = ('WQ01', 'WQ02', 'WQ03')
FORM_IDS = ('synthetic-divided-date-v1', 'synthetic-nominal-insertion-v1', 'synthetic-caption-catalogue-v1')
MANIFEST = Path(__file__).with_name('fixtures') / 'word-structure-qualification-v1.json'
MANIFEST_SHA256 = 'a6886adaaee377acf6365ff6a4bb04cd2b6639c8fdc7cdc7430346a2ad5bb0fb'
require = pc.require


def calendar_identity(value):
    """Independent finite calendar reader, never the document writer's formatter."""
    try:
        if re.fullmatch(r'[0-9]{4}-[0-9]{2}-[0-9]{2}', value):
            return datetime.date.fromisoformat(value).isoformat()
        slash = re.fullmatch(r'([0-9]{2})/([0-9]{2})/([0-9]{4})', value)
        if slash:
            a, b, year = map(int, slash.groups()); dates = set()
            for day, month in ((a, b), (b, a)):
                try:dates.add(datetime.date(year, month, day).isoformat())
                except ValueError:pass
            return next(iter(dates)) if len(dates) == 1 else None
        month = re.fullmatch(r'([0-9]{2})/(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)/([0-9]{4})', value)
        if month:
            number = ('Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec').index(month[2]) + 1
            return datetime.date(int(month[3]), number, int(month[1])).isoformat()
    except (TypeError, ValueError):pass
    return None


def usefulness(case, fields, answers):
    result = elicited.usefulness(case, fields, answers)
    for field in fields:
        label = field['label']; allowed = case.get('accepted_date_values', {}).get(label)
        if allowed is None:continue
        require(field['kind'] == 'date' and label in case['desired_supported'], 'Only a declared ordinary date has alternatives.')
        expected = calendar_identity(case['desired_supported'][label])
        require(expected and all(calendar_identity(value) == expected for value in allowed),
                'The prospectively declared date representations disagree on a real calendar date.')
        row = answers.get(field['key'], {})
        if row.get('status') == 'supported' and row.get('text') in allowed:
            result[field['key']] = 'supported_correct'
    return result


def build(form):
    """Deterministic blank originals; no development original is opened."""
    from docx import Document
    from docx.shared import Inches
    require(form in FORM_IDS, 'No arbitrary Word qualification original is eligible.')
    doc = Document()
    doc.core_properties.title = 'Project-authored synthetic Word qualification'
    doc.core_properties.author = doc.core_properties.last_modified_by = 'QuillMD synthetic engineering'
    doc.core_properties.created = doc.core_properties.modified = datetime.datetime(2026, 10, 5)
    doc.core_properties.comments = 'Authored synthetic blank fixture; no patient or original corpus data.'
    if form == FORM_IDS[0]:
        doc.add_paragraph('Visit date:'); paragraph = doc.add_paragraph()
        for n in range(9):
            if n in (2, 5):paragraph.add_run(' / ')
            paragraph.add_run('   ').underline = True
            if n + 1 < 9 and n + 1 not in (2, 5):paragraph.add_run(' ')
        paragraph.add_run('.'); doc.add_paragraph('\t'.join('ddmmmyyyy'))
    elif form == FORM_IDS[1]:
        doc.add_paragraph('We assessed the case and determined that the {insert category} treatment should end.')
    else:
        for labels in (('Patient name', 'Visit date', 'Clinical notes'), ('Witness name', 'Signature', 'Date')):
            table = doc.add_table(rows=2, cols=3)
            for n, label in enumerate(labels):table.cell(1, n).text = label
            for column in table.columns:column.width = Inches(2)
            for row in table.rows:
                for cell in row.cells:cell.width = Inches(2)
    for label in (('Plan', 'Allergies', 'Physician signature') if form != FORM_IDS[2] else ('Allergies',)):
        doc.add_paragraph(label + ': ____________')
    raw = io.BytesIO(); doc.save(raw); output = io.BytesIO()
    with zipfile.ZipFile(raw) as source, zipfile.ZipFile(output, 'w') as target:
        for name in sorted(source.namelist()):
            item = zipfile.ZipInfo(name, date_time=(2026, 10, 5, 0, 0, 0))
            item.compress_type = zipfile.ZIP_DEFLATED; item.create_system = 3; item.external_attr = 0o600 << 16
            target.writestr(item, source.read(name))
    return output.getvalue()


def issue_identity(issue):
    return {key: issue[key] for key in ('code', 'part', 'slot', 'start', 'end') if key in issue}


def frozen_inputs():
    raw = MANIFEST.read_bytes(); root = Path(__file__).resolve().parents[2]
    pc.committed_file(MANIFEST, root, 'Word structure qualification manifest')
    pc.committed_file(Path(__file__).resolve(), root, 'Word structure qualification runner')
    require(hashlib.sha256(raw).hexdigest() == MANIFEST_SHA256, 'The Word qualification freeze changed.')
    manifest = json.loads(raw)
    require(manifest['stage'] == STAGE and manifest['provider_upload_allowed'] is True
            and manifest['independently_authored'] is True and manifest['corpus_qualification'] is False,
            'The Word qualification rights or scope changed.')
    require(tuple(case['id'] for case in manifest['cases']) == CASE_IDS
            and tuple(manifest['forms']) == FORM_IDS, 'The exact three-case Word plan changed.')
    for form, spec in manifest['forms'].items():
        require(hashlib.sha256(build(form)).hexdigest() == spec['sha256'], 'A synthetic blank original changed.')
    for case in manifest['cases']:
        require(hashlib.sha256(case['notes'].encode()).hexdigest() == case['notes_sha256'],
                'A declared synthetic clinical source changed.')
    return {'word_structure': manifest, 'freeze': {'manifest_sha256': MANIFEST_SHA256,
            'frozen_files': {str(MANIFEST): MANIFEST_SHA256}, 'previously_exposed': True,
            'original_corpus_provider_eligible': False}}


def plan_sends(inputs=None):
    exact = frozen_inputs(); inputs = exact if inputs is None else inputs
    require(inputs == exact, 'Only the complete fixed three-case Word freeze is eligible.')
    return [stages.Send(n, 'word-structure:' + case['id'], 'generation', case['form'], 'field_relations',
                        STAGE, case['id'], PRICING)
            for n, case in enumerate(inputs['word_structure']['cases'], 1)]


def send_inputs(send, inputs):
    require(send in plan_sends(inputs), 'No arbitrary Word qualification send is eligible.')
    case = next(c for c in inputs['word_structure']['cases'] if c['id'] == send.case)
    return {'form_sha256': inputs['word_structure']['forms'][send.form]['sha256'],
            'notes_sha256': case['notes_sha256']}


def build_forms(target, inputs):
    require(inputs == frozen_inputs(), 'Only the exact reviewed Word originals may be built.')
    target = pc.private_directory(target); forms = {}
    for form in FORM_IDS:
        raw = build(form); path = target / (form + '.docx')
        if os.path.lexists(path):
            require(not path.is_symlink() and path.read_bytes() == raw, 'A stored synthetic original changed.')
        else:pc.write_new(path, raw, raw=True)
        forms[form] = raw
    return forms


def mapping_admission(original_sha, fields, manifest, spec):
    """Only prospectively named human-event exclusions; never remove a field."""
    from documents.import_selection import digest
    coverage = manifest.get('coverage', {})
    require(manifest.get('engine') == 'v3' and manifest.get('source_sha256') == original_sha == spec['sha256']
            and fields == spec['fields'] and manifest.get('mapping_sha256')
            == manifest.get('initial_mapping_sha256') == digest(fields), 'The exact original/mapping binding changed.')
    require(not coverage.get('comparison_unresolved') and not coverage.get('layout_quarantine')
            and coverage.get('block_generation') is bool(spec['source_issues'])
            and [issue_identity(row) for row in coverage.get('selected_engine_issues', [])] == spec['source_issues'],
            'The current unsupported source differs from the prospectively declared human-event exclusions.')
    return coverage


def price_sends(sends, inputs, state_dir, declaration, limits):
    from django.test import override_settings
    from core import generation as g
    from documents.service import inspect_template
    build_forms(Path(state_dir) / 'forms', inputs); priced = {}
    with override_settings(**stages.steps.PATH_FLAGS['field_relations']):
        for send in sends:
            send_inputs(send, inputs)
            case = next(c for c in inputs['word_structure']['cases'] if c['id'] == send.case)
            fields = inspect_template(Path(state_dir) / 'forms' / (send.form + '.docx'), engine='v3')['fields']
            require(fields == inputs['word_structure']['forms'][send.form]['fields'], 'The detected catalogue changed.')
            config = g.configuration(fields)
            require(config['model'] == 'claude-sonnet-5-5' and config.get('ordinary_field_contract')
                    == 'ordinary-field-relations-v4' and config['request_profile']['max_output_tokens'] == 20000,
                    'The Word send requires the qualified Sonnet5.5 ordinary-v4 profile.')
            require(config['pricing_version'] == declaration.pricing_version(PRICING)
                    and stages.within_bounds(config['request_profile'], declaration.expect(limits)),
                    'The Word quote/profile differs from its prospective declaration.')
            maximum = g.maximum_cost(case['notes'], fields, config, context_labels=[], requirements=[])
            priced[send.id] = {'maximum_cost': maximum, 'bound': pc.attempt_bound(maximum),
                'pricing_version': config['pricing_version'], 'max_output_tokens': 20000,
                'timeout_seconds': config['request_profile']['timeout_seconds']}
    return priced


class WordStructureRun(stages.Run):
    """Normal owned upload, mapping review, quote, worker, capture and export."""
    def notes_send(self, send):
        from copy import deepcopy
        from django.test import override_settings
        from core import import_binding
        from core.models import TemplateVersion
        require(self.stage == STAGE, 'The Word worker requires its separately declared stage.')
        send_inputs(send, self.inputs)
        case = next(c for c in self.inputs['word_structure']['cases'] if c['id'] == send.case)
        spec = self.inputs['word_structure']['forms'][send.form]
        with override_settings(**stages.steps.PATH_FLAGS['field_relations']):
            template, version = self.upload(self.forms[send.form],
                f'SYNTHETIC WORD STRUCTURE {send.case} attempt {len(self.ledger.rows()) + 1}', 'synthetic.docx')
            require(version.contract == 'v3', 'The normal v3 import was not selected.')
            coverage = mapping_admission(template.sha256, version.fields, version.import_manifest, spec)
            review = None
            if spec['source_issues']:
                old_fields, old_manifest = deepcopy(version.fields), deepcopy(version.import_manifest)
                issues = coverage['selected_engine_issues']
                response = self.world.client.post(f'/templates/{template.pk}/', {
                    'action': 'v3_coverage', 'mapping_version': str(version.pk),
                    **{'exclude_' + issue['id']: 'on' for issue in issues}})
                require(response.status_code == 302, 'The explicit source-exclusion review was refused; nothing sent.')
                new = TemplateVersion.objects.filter(template=template).order_by('-number').first()
                version.refresh_from_db()
                require(new.pk != version.pk and version.fields == new.fields == old_fields
                        and version.import_manifest == new.import_manifest == old_manifest
                        and import_binding.coverage_reviewed(new), 'Source review did not retain the exact mapping.')
                reviews = new.corrections.get('coverage_review', [])
                require(len(reviews) == len(issues) and {r['id'] for r in reviews} == {r['id'] for r in issues}
                        and all(r.get('disposition') == 'leave_unchanged' and r.get('source_sha256') == template.sha256
                                and r.get('reviewed_by') == str(self.world.user.pk) for r in reviews),
                        'The declared human-event exclusions were not explicitly reviewed by this owner.')
                review = {'original_version': str(version.pk), 'reviewed_version': str(new.pk),
                          'coverage_review': reviews, 'partial_coverage': True, 'automatic_complete': False}
                version = new
            response = stages.steps.confirm_mapping(self.world.client, template.pk, version.pk)
            version.refresh_from_db()
            require(response.status_code == 302 and version.confirmed, 'Normal mapping confirmation was refused.')
            def judge(send, version, answers, exported, original, **kwargs):
                outcomes = usefulness(case, version.fields, answers)
                return outcomes, [], {'word_structure_first_pass': {'fields': outcomes,
                    'manual_answer_corrections': 0, 'corpus_qualification': False}}
            return self.generation_send(send, template, version, case['notes'], extra={
                'case': send.case, 'word_structure_fixture_manifest_sha256': MANIFEST_SHA256,
                'coverage_before_confirmation': coverage, 'source_exclusion_review': review}, judge=judge)
