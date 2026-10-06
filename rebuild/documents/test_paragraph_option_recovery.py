"""Explicit recovery of complete printed lists; never an automatic corpus pass."""
from copy import deepcopy
import io
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import word_v3 as w


def saved(doc):
    stream = io.BytesIO(); doc.save(stream); return stream.getvalue()


def source(options=('Own flat', 'With family', 'With partner', 'Other')):
    doc = Document(); doc.add_heading('Accommodation', level=2)
    doc.add_paragraph('Which accommodation arrangements apply:')
    for option in options:
        p = doc.add_paragraph(option, style='List Bullet')
        numbering = OxmlElement('w:numPr')
        for tag, value in (('w:ilvl', '0'), ('w:numId', '1')):
            node = OxmlElement(tag); node.set(qn('w:val'), value); numbering.append(node)
        p._p.get_or_add_pPr().append(numbering)
    doc.add_paragraph('Other details: ______')
    doc.sections[0].header.paragraphs[0].text = 'Synthetic clinic'
    return saved(doc)


def located(raw, kind='multi_choice', slot='word/document.xml#b/p[1]'):
    return w.locate(raw, [{'slot': slot, 'kind': kind, 'label': 'Accommodation arrangements'}])[0]


class ParagraphOptionRecovery(unittest.TestCase):
    def test_explicit_one_and_several_choices_bind_every_original_option(self):
        raw = source(); info = w.inspect(raw)
        self.assertFalse(any(f['locator']['target']['type'] == 'paragraph_options' for f in info['fields']))
        for kind, value, expected in (('choice', 'With partner', [False, False, True, False]),
                                      ('multi_choice', ['Own flat', 'With partner'], [True, False, True, False])):
            with self.subTest(kind=kind):
                field = located(raw, kind); self.assertEqual(field['kind'], kind)
                self.assertEqual(field['origin'], 'manual')
                self.assertEqual(field['options'], ['Own flat', 'With family', 'With partner', 'Other'])
                self.assertEqual(len(field['locator']['extra']), 3)
                self.assertEqual(field['source_questions'][0]['cardinality'], 'one' if kind == 'choice' else 'many')
                answers = {field['key']: value}; out = w.fill(raw, [field], answers)
                self.assertEqual(w.effective_options(raw, [field], answers)[field['key']], expected)
                before, after = Document(io.BytesIO(raw)), Document(io.BytesIO(out))
                for number, marked in enumerate(expected, 2):
                    self.assertEqual(after.paragraphs[number].text,
                                     ('☒ ' if marked else '') + before.paragraphs[number].text)
                    self.assertEqual(before.paragraphs[number].style.name, after.paragraphs[number].style.name)
                self.assertEqual(before.paragraphs[1].text, after.paragraphs[1].text)
                self.assertEqual(before.paragraphs[6].text, after.paragraphs[6].text)
                with ZipFile(io.BytesIO(raw)) as a, ZipFile(io.BytesIO(out)) as b:
                    self.assertEqual(a.namelist(), b.namelist())
                    for part in a.namelist():
                        if part != 'word/document.xml': self.assertEqual(a.read(part), b.read(part))
                self.assertEqual(w.fill(raw, [field], {}), raw)
                self.assertEqual(w.fill(raw, [field], answers), out)

    def test_only_the_question_or_first_option_can_choose_the_whole_group(self):
        raw = source()
        self.assertEqual(located(raw)['locator'], located(raw, slot='word/document.xml#b/p[2]')['locator'])
        for slot in ('word/document.xml#b/p[3]', 'word/document.xml#b/p[6]'):
            with self.subTest(slot=slot), self.assertRaises(w.V3Failure): located(raw, slot=slot)

    def test_hidden_control_or_long_list_member_cannot_silently_drop_a_suffix(self):
        for variant in ('hidden', 'blank', 'long', 'checked', 'duplicate', 'level', 'unknown_numbering', 'checkbox_marker'):
            with self.subTest(variant=variant):
                doc = Document(io.BytesIO(source()))
                p = doc.paragraphs[4]
                if variant == 'hidden': p.runs[0].font.hidden = True
                elif variant == 'blank': p.text = '______'
                elif variant == 'long': p.text = 'Long option ' * 30
                elif variant == 'checked': p.text = '☒ With partner'
                elif variant == 'duplicate': p.text = 'Own flat'
                elif variant in ('level', 'unknown_numbering'):
                    props = p._p.get_or_add_pPr(); numbering = OxmlElement('w:numPr')
                    level = OxmlElement('w:ilvl'); level.set(qn('w:val'), '1' if variant == 'level' else '0'); numbering.append(level)
                    identifier = OxmlElement('w:numId'); identifier.set(qn('w:val'), '1' if variant == 'level' else '999'); numbering.append(identifier)
                    props.remove(props.find(qn('w:numPr'))); props.append(numbering)
                else:
                    numbering = doc.part.numbering_part.element
                    abstract = next(a for a in numbering.findall(qn('w:abstractNum')) if a.get(qn('w:abstractNumId')) == '8')
                    abstract.find('.//' + qn('w:lvlText')).set(qn('w:val'), '☐')
                with self.assertRaises(w.V3Failure): located(saved(doc))

    def test_modified_group_metadata_cannot_move_or_relabel_physical_options(self):
        raw = source(); original = located(raw)
        for variant in ('options', 'destination', 'descriptor', 'cardinality', 'definitions', 'removed_relation', 'removed_origin'):
            with self.subTest(variant=variant):
                field = deepcopy(original)
                if variant == 'options': field['options'][0] = 'Invented option'
                elif variant == 'destination': field['locator']['target']['paragraphs'][-1] = 'b/p[6]'
                elif variant == 'descriptor': field['source_questions'][0]['descriptors'][1]['text'] = 'Invented option'
                elif variant == 'cardinality': field['kind'] = 'choice'
                elif variant == 'definitions': field['source_questions'][0]['definitions']['word/numbering.xml'] = 'changed'
                elif variant == 'removed_relation': field.pop('source_questions')
                else: field.pop('source_questions'); field.pop('origin')
                with self.assertRaises(w.V3Failure): w.fill(raw, [field], {field['key']: field['options'][0]})

    def test_unknown_answer_and_multiple_single_choices_refuse_before_export(self):
        raw = source()
        for kind, value in (('choice', ['Own flat', 'With partner']), ('multi_choice', ['Unknown'])):
            field = located(raw, kind)
            with self.subTest(kind=kind), self.assertRaises(w.V3Failure): w.fill(raw, [field], {field['key']: value})
