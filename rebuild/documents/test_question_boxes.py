"""Synthetic structural controls for a printed question followed by a separate answer box."""
import io
import unittest
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from . import word_v3 as w


def box(doc, *, bordered=True):
    table = doc.add_table(rows=1, cols=1)
    if bordered:
        borders = OxmlElement('w:tcBorders')
        for side in ('top', 'left', 'bottom', 'right'):
            edge = OxmlElement('w:' + side)
            edge.set(qn('w:val'), 'single')
            borders.append(edge)
        table.cell(0, 0)._tc.get_or_add_tcPr().append(borders)
    return table


def data(doc):
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


class QuestionBoxes(unittest.TestCase):
    def test_mapped_instruction_cell_owns_its_plain_row_or_column_label(self):
        from docx.shared import RGBColor
        for stacked in (False, True):
            for suffix in ('.', '?'):
                with self.subTest(stacked=stacked, suffix=suffix):
                    doc = Document(); table = doc.add_table(rows=2 if stacked else 1, cols=1 if stacked else 2)
                    table.cell(0, 0).text = 'Please give examples of symptoms experienced:'
                    answer = table.cell(1, 0) if stacked else table.cell(0, 1)
                    answer.paragraphs[0].add_run('Describe the frequency and duration' + suffix).font.color.rgb = RGBColor.from_string('0000FF')
                    original = data(doc); fields = w.inspect(original)['fields']
                    self.assertEqual(len(fields), 1)
                    self.assertEqual(fields[0]['locator']['target']['type'], 'cell_append')
                    result = w.fill(original, fields, {fields[0]['key']: 'Twice a week for one hour.'})
                    before = w.open_package(original); after = w.open_package(result)
                    a = w.SlotIndex(before).by_id['word/document.xml#b/tbl[0]/tr[0]/tc[0]'].el
                    b = w.SlotIndex(after).by_id['word/document.xml#b/tbl[0]/tr[0]/tc[0]'].el
                    self.assertEqual(etree.tostring(a), etree.tostring(b))

    def test_mapped_cell_does_not_remove_an_explicit_header_answer(self):
        from docx.shared import RGBColor
        for stacked in (False, True):
            doc = Document(); table = doc.add_table(rows=2 if stacked else 1, cols=1 if stacked else 2)
            table.cell(0, 0).text = 'Observation: __________'
            answer = table.cell(1, 0) if stacked else table.cell(0, 1)
            answer.paragraphs[0].add_run('Describe its duration?').font.color.rgb = RGBColor.from_string('0000FF')
            fields = w.inspect(data(doc))['fields']
            self.assertTrue(any(f['locator']['target']['type'] == 'span' for f in fields))

    @staticmethod
    def layout_borders(table, color='FFFFFF'):
        borders = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); edge.set(qn('w:color'), color)
            borders.append(edge)
        table._tbl.tblPr.append(borders)

    def test_presentation_heading_table_spacing_is_not_an_answer(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        for variant in ('presentation', 'boxed', 'explicit_blank', 'question', 'shaded', 'unknown_style', 'theme_border', 'duplicate_properties', 'marked_spacing'):
            with self.subTest(variant=variant):
                doc = Document(); table = doc.add_table(rows=1, cols=2)
                self.layout_borders(table, '000000' if variant == 'boxed' else 'FFFFFF')
                for i, text in enumerate(('Regional service', 'Assessment form')):
                    para = table.cell(0, 0).paragraphs[0] if not i else table.cell(0, 0).add_paragraph()
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER; para.add_run(text).bold = True
                table.cell(0, 0).add_paragraph()
                if variant == 'explicit_blank': table.cell(0, 1).text = 'Identifier: ______'
                if variant == 'question': table.cell(0, 0).paragraphs[1].text = 'What symptoms are present?'
                if variant == 'shaded':
                    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), '000000')
                    table.cell(0, 1)._tc.get_or_add_tcPr().append(shd)
                if variant == 'unknown_style':
                    ref = OxmlElement('w:tblStyle'); ref.set(qn('w:val'), 'MissingTableStyle'); table._tbl.tblPr.append(ref)
                if variant == 'theme_border':
                    table._tbl.tblPr.find(qn('w:tblBorders')).find(qn('w:insideV')).set(qn('w:themeColor'), 'accent1')
                if variant == 'duplicate_properties': table.cell(0, 1)._tc.append(OxmlElement('w:tcPr'))
                if variant == 'marked_spacing': table.cell(0, 0).paragraphs[-1].add_run('        ').underline = True
                fields = w.inspect(data(doc))['fields']
                self.assertEqual(bool(fields), variant != 'presentation')

    def test_merge_continuation_tail_keeps_only_the_owned_upper_answer(self):
        from docx.shared import RGBColor
        for variant in ('tail', 'boxed', 'explicit_blank', 'unmerged_neighbor'):
            with self.subTest(variant=variant):
                doc = Document(); table = doc.add_table(rows=2, cols=2)
                self.layout_borders(table)
                table.cell(0, 0).text = 'Date'
                if variant != 'unmerged_neighbor': table.cell(0, 0).merge(table.cell(1, 0))
                else: table.cell(1, 0).text = 'Further date'
                table.cell(0, 1).paragraphs[0].add_run('Enter the date of submission.').font.color.rgb = RGBColor.from_string('0000FF')
                for row in (0, 1):
                    borders = OxmlElement('w:tcBorders')
                    for side in (('top', 'left', 'bottom', 'right') if row == 0 or variant == 'boxed' else ('top',)):
                        edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); edge.set(qn('w:color'), '000000'); borders.append(edge)
                    table.cell(row, 1)._tc.get_or_add_tcPr().append(borders)
                if variant == 'explicit_blank': table.cell(1, 1).text = '________'
                fields = w.inspect(data(doc))['fields']
                lower = [f for f in fields if 'tr[1]/tc[1]' in f['locator']['slot']]
                self.assertEqual(bool(lower), variant != 'tail')
                self.assertTrue(any('tr[0]/tc[1]' in f['locator']['slot'] for f in fields))

    def test_paired_cell_label_spacing_is_not_a_second_answer(self):
        for layout in ('ordinary', 'merged', 'multicolumn', 'nested'):
            with self.subTest(layout=layout):
                doc = Document()
                owner = doc if layout != 'nested' else doc.add_table(rows=1, cols=1).cell(0, 0)
                table = owner.add_table(rows=1, cols=4 if layout == 'multicolumn' else 3 if layout == 'merged' else 2)
                label = table.cell(0, 0)
                if layout == 'merged': label = label.merge(table.cell(0, 1))
                label.text = 'Date of Birth'
                label.add_paragraph()
                label.add_paragraph()
                if layout == 'multicolumn':
                    table.cell(0, 2).text = 'Date of Visit'
                    table.cell(0, 2).add_paragraph()
                result = w.inspect(data(doc))
                self.assertFalse([f for f in result['fields'] if f['locator']['target']['type'] == 'paragraph'])
                self.assertEqual(len([f for f in result['fields'] if f['locator']['target']['type'] == 'cell']),
                                 2 if layout == 'multicolumn' else 1)

    def test_paired_cell_keeps_explicit_and_independent_answer_lines(self):
        for marker in ('underline', 'border', 'inherited_border', 'table_border', 'default_border',
                       'unknown_style', 'standalone'):
            with self.subTest(marker=marker):
                doc = Document()
                table = doc.add_table(rows=1, cols=1 if marker == 'standalone' else 2)
                label = table.cell(0, 0)
                label.text = 'Date of Birth'
                blank = label.add_paragraph()
                if marker == 'underline': blank.add_run('     ').underline = True
                elif marker in ('border', 'inherited_border', 'table_border', 'default_border'):
                    border = OxmlElement('w:pBdr');edge = OxmlElement('w:bottom')
                    edge.set(qn('w:val'), 'single');border.append(edge)
                    if marker == 'border': blank._p.get_or_add_pPr().append(border)
                    elif marker == 'inherited_border':
                        style = doc.styles.add_style('Answer Line', WD_STYLE_TYPE.PARAGRAPH)
                        style.element.get_or_add_pPr().append(border);blank.style = style
                    elif marker == 'table_border':
                        style = doc.styles.add_style('Answer Table', WD_STYLE_TYPE.TABLE)
                        condition = OxmlElement('w:tblStylePr');condition.set(qn('w:type'), 'firstCol')
                        props = OxmlElement('w:pPr');props.append(border);condition.append(props)
                        style.element.append(condition);table.style = style
                    else:
                        defaults = doc.styles.element.find(qn('w:docDefaults'))
                        props = defaults.find(qn('w:pPrDefault')).find(qn('w:pPr'))
                        props.append(border)
                elif marker == 'unknown_style':
                    ref = OxmlElement('w:pStyle');ref.set(qn('w:val'), 'UnavailableAnswerStyle')
                    blank._p.get_or_add_pPr().append(ref)
                result = w.inspect(data(doc))
                self.assertTrue([f for f in result['fields'] if f['locator']['target']['type'] == 'paragraph'])

    def test_paired_spacing_requires_resolved_run_and_table_styles(self):
        cases = ('character', 'inherited_character', 'unknown_character', 'cyclic_character',
                 'default_character', 'default_table', 'inherited_default_table',
                 'unknown_default_table_parent', 'multiple_default_tables', 'explicit_table_override',
                 'explicit_nil_run_override', 'plain_character')
        for case in cases:
            with self.subTest(case=case):
                doc = Document()
                table = doc.add_table(rows=1, cols=2)
                table.cell(0, 0).text = 'Date of Birth'
                paragraph = table.cell(0, 0).add_paragraph()
                run = paragraph.add_run('        ')
                if 'table' in case:
                    kind = WD_STYLE_TYPE.TABLE
                    style = doc.styles.add_style('Answer Table', kind)
                    if case != 'multiple_default_tables':
                        for candidate in doc.styles.element:
                            if candidate.get(qn('w:type')) == 'table':
                                candidate.attrib.pop(qn('w:default'), None)
                    style.element.set(qn('w:default'), '1')
                    if case in ('inherited_default_table', 'unknown_default_table_parent'):
                        parent = doc.styles.add_style('Answer Table Parent', kind)
                        parent.font.underline = True
                        style.base_style = parent
                        if case == 'unknown_default_table_parent':
                            style.element.find(qn('w:basedOn')).set(qn('w:val'), 'MissingTableParent')
                    else:
                        style.font.underline = True
                    if case == 'explicit_table_override':
                        table.style = doc.styles.add_style('Plain Explicit Table', kind)
                else:
                    style = doc.styles.add_style('Answer Character', WD_STYLE_TYPE.CHARACTER)
                    if case == 'inherited_character':
                        parent = doc.styles.add_style('Answer Character Parent', WD_STYLE_TYPE.CHARACTER)
                        parent.font.underline = True
                        style.base_style = parent
                    elif case != 'plain_character':
                        style.font.underline = True
                    if case == 'default_character':
                        for candidate in doc.styles.element:
                            if candidate.get(qn('w:type')) == 'character':
                                candidate.attrib.pop(qn('w:default'), None)
                        style.element.set(qn('w:default'), '1')
                    else:
                        run.style = style
                    if case == 'unknown_character':
                        run._r.rPr.find(qn('w:rStyle')).set(qn('w:val'), 'MissingAnswerCharacter')
                    if case == 'cyclic_character':
                        style.base_style = style
                    if case == 'explicit_nil_run_override':
                        run.underline = False  # Conservative retention is permitted; never erase a marked area.
                result = w.inspect(data(doc))
                paragraphs = [f for f in result['fields'] if f['locator']['target']['type'] == 'paragraph']
                self.assertEqual(bool(paragraphs), case not in ('plain_character', 'explicit_table_override'))
                self.assertEqual(len([f for f in result['fields'] if f['locator']['target']['type'] == 'cell']), 1)

    def test_paired_spacing_keeps_shading_lists_and_run_answer_marks(self):
        for marker in ('shading', 'inherited_shading', 'highlight', 'number', 'bullet',
                       'run_border', 'strike', 'double_strike', 'leader_tab'):
            with self.subTest(marker=marker):
                doc = Document()
                table = doc.add_table(rows=1, cols=2)
                table.cell(0, 0).text = 'Past medical history:'
                blank = table.cell(0, 0).add_paragraph()
                if marker in ('shading', 'inherited_shading'):
                    shade = OxmlElement('w:shd');shade.set(qn('w:fill'), 'D9D9D9')
                    if marker == 'shading': blank._p.get_or_add_pPr().append(shade)
                    else:
                        style = doc.styles.add_style('Shaded Answer', WD_STYLE_TYPE.PARAGRAPH)
                        style.element.get_or_add_pPr().append(shade);blank.style = style
                elif marker in ('number', 'bullet'):
                    blank.style = 'List Number' if marker == 'number' else 'List Bullet'
                else:
                    run = blank.add_run('        ')
                    if marker == 'leader_tab': run.add_tab()
                    else:
                        tag = {'highlight': 'highlight', 'run_border': 'bdr',
                               'strike': 'strike', 'double_strike': 'dstrike'}[marker]
                        mark = OxmlElement('w:' + tag)
                        mark.set(qn('w:val'), 'yellow' if marker == 'highlight' else 'single' if marker == 'run_border' else '1')
                        run._r.get_or_add_rPr().append(mark)
                result = w.inspect(data(doc))
                self.assertEqual(len(result['fields']), 2)
                self.assertTrue(any(f['locator']['target']['type'] == 'paragraph' for f in result['fields']))

    def test_paired_cell_spacing_does_not_claim_controls_or_textboxes(self):
        from .test_word_v3 import _sdt, ff_field
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Date of Birth'
        table.cell(0, 0).add_paragraph()
        p = doc.add_paragraph('Laterality: ')
        p._p.append(_sdt('dropdown', options=['Choose an item.', 'Left', 'Right']))
        p = doc.add_paragraph('Reviewed ')
        ff_field(p, 'checkbox', checked=True)
        p = doc.add_paragraph()
        pict = OxmlElement('w:pict')
        shape = etree.SubElement(pict, '{urn:schemas-microsoft-com:vml}shape')
        shape.set('style', 'width:200pt;height:70pt')
        textbox = etree.SubElement(shape, '{urn:schemas-microsoft-com:vml}textbox')
        content = OxmlElement('w:txbxContent');textbox.append(content)
        inner = OxmlElement('w:p');content.append(inner)
        run = OxmlElement('w:r');inner.append(run)
        text = OxmlElement('w:t');text.text = 'Textbox comments: ______';run.append(text)
        p.add_run()._r.append(pict)
        original = data(doc)
        result = w.inspect(original)
        self.assertEqual(len([f for f in result['fields'] if f['locator']['target']['type'] == 'sdt_choice']), 1)
        self.assertTrue(any('/txbx[' in f['locator']['slot'] and f['label'] == 'Textbox comments' for f in result['fields']))
        self.assertTrue(any(f['locator']['target']['type'] == 'checkboxes' for f in result['fields']))
        target = next(f for f in result['fields'] if f['locator']['slot'] == 'b/tbl[0]/tr[0]/tc[1]')
        exported = w.fill(original, result['fields'], {target['key']: 'SYNTHETIC date'})
        before = w.SlotIndex(w.open_package(original));after = w.SlotIndex(w.open_package(exported))
        for slot in before.roots:
            if slot.kind != 'tbl':
                self.assertEqual(etree.tostring(slot.el), etree.tostring(after.by_id[slot.id].el))

    def test_complete_long_question_and_section_are_retained(self):
        doc = Document()
        context = 'Clinical section ' + 'complete context ' * 25
        question = 'What changed? ' + 'Keep every qualifier. ' * 25
        doc.add_paragraph(context.strip(), style='Heading 1')
        doc.add_paragraph(question.strip())
        box(doc)
        result = w.inspect(data(doc))
        self.assertEqual(len(result['fields']), 1)
        self.assertEqual(result['fields'][0]['label'], question.strip())
        self.assertEqual(result['fields'][0]['context'], context.strip())

    def test_full_source_limits_keep_boundary_and_refuse_after_scope_expansion(self):
        for length in (1000, 1001):
            for what in ('label', 'context'):
                with self.subTest(length=length, what=what):
                    doc = Document()
                    context = 'C' * length if what == 'context' else 'Clinical history'
                    question = 'What ' + 'q' * (length - 6) + '?' if what == 'label' else 'What changed?'
                    doc.add_paragraph(context, style='Heading 1')
                    doc.add_paragraph(question)
                    box(doc)
                    result = w.inspect(data(doc))
                    if length == 1000:
                        self.assertEqual(len(result['fields']), 1)
                        self.assertEqual(result['fields'][0]['label'], question)
                        self.assertEqual(result['fields'][0]['context'], context)
                    else:
                        self.assertEqual(result['fields'], [])
                        self.assertTrue(any(question in n and 'limit' in n for n in result['notes']))
        doc = Document()
        long_context = 'Clinical ' + 'x' * 981
        doc.add_paragraph(long_context, style='Heading 1')
        for _ in range(2):
            doc.add_paragraph('What changed?')
            box(doc)
        result = w.inspect(data(doc))
        self.assertEqual(result['fields'], [])
        self.assertEqual(sum('nothing was shortened' in n for n in result['notes']), 2)

    def test_manual_derived_context_is_complete_or_visibly_refused(self):
        for length in (600, 1001):
            doc = Document()
            context = 'Clinical ' + 'c' * (length - 9)
            doc.add_paragraph(context, style='Heading 1')
            box(doc)
            target = {'slot': 'word/document.xml#b/tbl[0]/tr[0]/tc[0]', 'kind': 'long_text', 'label': 'Comments'}
            with self.subTest(length=length):
                if length == 600:
                    self.assertEqual(w.locate(data(doc), [target])[0]['context'], context)
                else:
                    with self.assertRaises(w.V3Failure): w.locate(data(doc), [target])

    def test_manual_root_verified_long_donor_scope_is_not_lost(self):
        doc = Document()
        context = 'Clinical section ' + 'continued ' * 16 + 'DONOR ONLY'
        doc.add_paragraph(context, style='Heading 1')
        doc.add_paragraph('What medications are currently taken?')
        box(doc)
        target = {'slot': 'word/document.xml#b/tbl[0]/tr[0]/tc[0]', 'kind': 'long_text', 'label': 'Medications'}
        self.assertEqual(w.locate(data(doc), [target])[0]['context'], context)

    def test_manual_long_metadata_never_slices_and_oversize_refuses(self):
        doc = Document()
        box(doc)
        source = data(doc)
        target = {'slot': 'word/document.xml#b/tbl[0]/tr[0]/tc[0]',
                  'kind': 'long_text', 'label': 'Question ' + 'q' * 491,
                  'context': 'Section ' + 'c' * 592}
        field = w.locate(source, [target])[0]
        self.assertEqual(field['label'], target['label'])
        self.assertEqual(field['context'], target['context'])
        for name in ('label', 'context'):
            with self.subTest(name=name), self.assertRaises(w.V3Failure):
                w.locate(source, [{**target, name: 'x' * 1001}])

    def test_initial_report_adjective_is_not_a_signing_instruction(self):
        for phrase, accepted in (
            ('Please describe the changes since the initial report?', True),
            ('Please describe changes in the case of the initial report?', True),
            ('Please notarize the initial report?', False),
            ('Please describe the changes since the initial report, then initial here?', False),
            ('Please describe the initial report and provide initials?', False),
            ('Please describe the initial report and provide signature?', False),
            ('Please provide Initial?', False), ('Please provide Initials?', False),
            ('Please sign the initial report here?', False),
            ('Please countersign the initial report?', False),
            ('Please counter-sign the initial report?', False),
            ('Has the initial report been countersigned?', False),
            ('Please undersign the initial report?', False),
            ('Please co-sign the initial report?', False),
            ('Please re-sign the initial report?', False),
            ('Please provide signoff on the initial report?', False),
            ('Please describe signing the initial report?', False),
            ('Please countersign since the initial report?', False),
            ('Please initial the report here?', False), ('Please initial here?', False)):
            with self.subTest(phrase=phrase):
                doc = Document()
                doc.add_paragraph(phrase)
                box(doc)
                result = w.inspect(data(doc))
                fields = [f for f in result['fields'] if f['locator']['target']['type'] == 'cell']
                self.assertEqual(len(fields), int(accepted))
                if accepted:
                    self.assertEqual(fields[0]['label'], phrase)
                    exported = w.fill(data(doc), fields, {fields[0]['key']: 'SYNTHETIC'})
                    self.assertEqual(w.text_of(w.SlotIndex(w.open_package(exported)).by_id['word/document.xml#b/tbl[0]/tr[0]/tc[0]'].el), 'SYNTHETIC')

    def test_initial_adjective_cannot_cross_question_section_boundary(self):
        for heading, question in (('Report', 'Please provide the initial'),
                                  ('Initial', 'Please describe the initial report?'),
                                  ('Please countersign the initial report', 'What changed?')):
            with self.subTest(heading=heading, question=question):
                doc = Document()
                doc.add_paragraph(heading, style='Heading 1')
                doc.add_paragraph(question)
                box(doc)
                result = w.inspect(data(doc))
                self.assertEqual(result['fields'], [])
                self.assertTrue(any('Signing or consent' in note for note in result['notes']))

    def test_inherited_heading_scope_and_unknown_or_cyclic_ancestry(self):
        for ancestor in ('Heading 1', 'Unknown', 'cycle'):
            with self.subTest(ancestor=ancestor):
                doc = Document()
                style = doc.styles.add_style('Donor Section', WD_STYLE_TYPE.PARAGRAPH)
                style.base_style = doc.styles['Heading 1']
                if ancestor != 'Heading 1':
                    style.element.find(qn('w:basedOn')).set(qn('w:val'), 'DonorSection' if ancestor == 'cycle' else ancestor)
                doc.add_paragraph('Donor', style)
                doc.add_paragraph('What medications are currently taken?')
                box(doc)
                result = w.inspect(data(doc))
                if ancestor == 'Heading 1':
                    self.assertEqual(result['fields'][0]['context'], 'Donor')
                else:
                    self.assertEqual(result['fields'], [])
                    self.assertTrue(any('style ancestry' in note for note in result['notes']))

    def test_multi_paragraph_question_keeps_both_sources_in_review(self):
        for blanks in (0, 1, 2):
            with self.subTest(blanks=blanks):
                doc = Document()
                first = 'Which medications are currently taken?'
                last = 'Please provide their doses and frequencies.'
                doc.add_paragraph(first)
                for _ in range(blanks):
                    doc.add_paragraph()
                doc.add_paragraph(last)
                box(doc)
                result = w.inspect(data(doc))
                self.assertEqual(result['fields'], [])
                self.assertTrue(any(first in note and last in note for note in result['notes']))

    def test_page_break_before_spacer_is_not_plain_empty_space(self):
        for inherited in (False, True):
            with self.subTest(inherited=inherited):
                doc = Document()
                doc.add_paragraph('What symptoms changed?')
                p = doc.add_paragraph()
                if inherited:
                    style = doc.styles.add_style('New page', WD_STYLE_TYPE.PARAGRAPH)
                    style.paragraph_format.page_break_before = True
                    p.style = style
                else:
                    p.paragraph_format.page_break_before = True
                box(doc)
                result = w.inspect(data(doc))
                self.assertEqual(result['fields'], [])
                self.assertTrue(any('page break' in note for note in result['notes']))

    def test_conditional_border_in_ancestor_style_is_not_claimed_as_explicit(self):
        doc = Document()
        doc.add_paragraph('What symptoms changed?')
        table = box(doc, bordered=False)
        parent = doc.styles.add_style('Base Box', WD_STYLE_TYPE.TABLE)
        conditional = OxmlElement('w:tblStylePr')
        conditional.set(qn('w:type'), 'firstRow')
        parent.element.append(conditional)
        child = doc.styles.add_style('Child Box', WD_STYLE_TYPE.TABLE)
        child.base_style = parent
        props = OxmlElement('w:tblPr')
        borders = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right'):
            edge = OxmlElement('w:' + side)
            edge.set(qn('w:val'), 'single')
            borders.append(edge)
        props.append(borders)
        child.element.append(props)
        table.style = child
        result = w.inspect(data(doc))
        self.assertEqual(result['fields'], [])
        self.assertTrue(any('border' in note for note in result['notes']))

    def test_question_with_its_own_inline_answer_does_not_lose_that_field(self):
        doc = Document()
        doc.add_paragraph('Patient name: __________. What symptoms changed?')
        box(doc)
        result = w.inspect(data(doc))
        self.assertEqual([(f['label'], f['locator']['target']['type']) for f in result['fields']],
                         [('Patient name', 'span')])
        self.assertTrue(any('another answer area' in note for note in result['notes']))

    def test_explicit_style_borders_and_cell_nil_override(self):
        for nil in (False, True):
            with self.subTest(nil=nil):
                doc = Document()
                doc.add_paragraph('What changed?')
                table = box(doc, bordered=False)
                table.style = 'Table Grid'
                if nil:
                    borders = OxmlElement('w:tcBorders')
                    edge = OxmlElement('w:left')
                    edge.set(qn('w:val'), 'nil')
                    borders.append(edge)
                    table.cell(0, 0)._tc.get_or_add_tcPr().append(borders)
                result = w.inspect(data(doc))
                self.assertEqual(len(result['fields']), 0 if nil else 1)
                self.assertEqual(bool(result['notes']), nil)

    def test_adjacent_and_spaced_questions_use_the_box_with_complete_scope(self):
        for spaces in (0, 1, 2):
            with self.subTest(spaces=spaces):
                doc = Document()
                for text in ('Part B', 'Social Supervisor Report', 'Supervision'):
                    doc.add_paragraph().add_run(text).bold = True
                question = 'When did you last see the patient? Give the date and location.'
                doc.add_paragraph(question)
                for _ in range(spaces):
                    doc.add_paragraph()
                box(doc)
                fields = w.inspect(data(doc))['fields']
                self.assertEqual(len(fields), 1)
                field = fields[0]
                self.assertEqual(field['label'], question)
                self.assertEqual(field['context'], 'Part B > Social Supervisor Report > Supervision')
                self.assertEqual(field['kind'], 'long_text')
                self.assertEqual(field['locator']['slot'], 'b/tbl[0]/tr[0]/tc[0]')

    def test_ambiguous_boxes_have_a_visible_source_location(self):
        cases = ('heading', 'option', 'signature', 'reference', 'long', 'intervening', 'three-spaces',
                 'section-break', 'unbordered')
        for case in cases:
            with self.subTest(case=case):
                doc = Document()
                question = {'heading': 'Clinical history', 'option': 'Other (please specify)',
                            'signature': 'Please provide the clinician signature.',
                            'reference': 'In relation to the above question, what changed?',
                            'long': 'What happened? ' + 'Keep every qualifier. ' * 50}.get(case, 'What changed?')
                p = doc.add_paragraph(question)
                if case == 'option':
                    p.style = 'List Bullet'
                if case == 'intervening':
                    doc.add_table(rows=1, cols=2)
                if case == 'three-spaces':
                    for _ in range(3):
                        doc.add_paragraph()
                if case == 'section-break':
                    p._p.get_or_add_pPr().append(OxmlElement('w:sectPr'))
                target = box(doc, bordered=case != 'unbordered')
                result = w.inspect(data(doc))
                self.assertFalse([f for f in result['fields'] if f['locator']['slot'].endswith('/tc[0]')])
                self.assertTrue(any('Unassigned answer box' in x and 'word/document.xml#b/tbl[' in x
                                    for x in result['notes']), result)
                if case == 'long':
                    self.assertTrue(any(question.strip() in x for x in result['notes']))

    def test_empty_ordinary_long_preserve_other_document_parts_and_neighbours(self):
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = 'Synthetic clinical branding'
        doc.add_paragraph('What symptoms changed?')
        box(doc)
        doc.add_paragraph('Printed consent and signature area — leave unchanged')
        original = data(doc)
        fields = w.inspect(original)['fields']
        field = next(f for f in fields if f['label'] == 'What symptoms changed?')
        for answer in ('', 'SYNTHETIC symptoms improved.', 'SYNTHETIC long answer. ' * 150):
            with self.subTest(length=len(answer)):
                output = w.fill(original, fields, {field['key']: answer})
                with ZipFile(io.BytesIO(original)) as before, ZipFile(io.BytesIO(output)) as after:
                    self.assertEqual(before.namelist(), after.namelist())
                    for part in before.namelist():
                        if part != 'word/document.xml':
                            self.assertEqual(before.read(part), after.read(part))
                    old = etree.fromstring(before.read('word/document.xml')).find(w.q('w:body'))
                    new = etree.fromstring(after.read('word/document.xml')).find(w.q('w:body'))
                    self.assertEqual(etree.tostring(old[0]), etree.tostring(new[0]))
                    self.assertEqual(etree.tostring(old[2]), etree.tostring(new[2]))
                    self.assertEqual(w.text_of(new[1]).strip(), answer.strip())

    def test_part_change_keeps_distinct_supervisors_and_does_not_reuse_a_question(self):
        doc = Document()
        for part, role in (('Part B', 'Social Supervisor'), ('Part C', 'Clinical Supervisor')):
            for text in (part, role, 'Supervision'):
                doc.add_paragraph().add_run(text).bold = True
            doc.add_paragraph('When did you last see the patient?')
            box(doc)
            box(doc)
        result = w.inspect(data(doc))
        fields = [f for f in result['fields'] if f['locator']['target']['type'] == 'cell']
        self.assertEqual(len(fields), 2)
        self.assertIn('Part B > Social Supervisor > Supervision', fields[0]['context'])
        self.assertIn('Part C > Clinical Supervisor > Supervision', fields[1]['context'])
        self.assertEqual(len([n for n in result['notes'] if 'Unassigned answer box' in n]), 2)


class NestedCheckboxOwnership(unittest.TestCase):
    def form(self, *, heading='', subsection=False, spacer='plain', mixed=False, write_in=False):
        doc = Document()
        if heading: doc.add_paragraph(heading, style='Heading 1')
        outer = doc.add_table(rows=1, cols=1).cell(0, 0)
        outer.text = 'Which symptoms occur? Tick all that apply.'
        if subsection: outer.add_paragraph('During exertion:')
        blank = outer.add_paragraph()
        if spacer == 'underline': blank.add_run('       ').underline = True
        if spacer == 'other_question': blank.text = 'Which symptoms occurred last year?'
        inner = outer.add_table(rows=1, cols=2)
        inner.cell(0, 0).text = '☐ Tightness'
        inner.cell(0, 1).text = ('☐ Weakness [☐ Right ☐ Left]' if mixed else
                                '☐ Other: ______' if write_in else '☒ Nausea')
        return doc

    def test_full_nested_question_and_cardinality_preserve_scope_and_checked_state(self):
        for heading, subsection in (('', False), ('Donor', False), ('Current visit', True)):
            with self.subTest(heading=heading, subsection=subsection):
                original = data(self.form(heading=heading, subsection=subsection))
                result = w.inspect(original)
                fields = [f for f in result['fields'] if f['locator']['target']['type'] == 'checkboxes']
                self.assertEqual(len(fields), 1)
                target = fields[0]
                self.assertEqual(target['kind'], 'multi_choice')
                self.assertEqual(target['label'], 'During exertion:' if subsection else
                                 'Which symptoms occur? Tick all that apply.')
                if heading: self.assertIn(heading, target['context'])
                if subsection: self.assertIn('Which symptoms occur? Tick all that apply.', target['context'])
                self.assertFalse([f for f in result['fields'] if f['locator']['target']['type'] == 'paragraph'])
                self.assertEqual(target['printed_ticked'], [1])
                blank = w.fill(original, result['fields'], {})
                self.assertEqual(w.option_states(w.SlotIndex(w.open_package(blank)), target), [False, True])
                filled = w.fill(original, result['fields'], {target['key']: ['Tightness', 'Nausea']})
                self.assertEqual(w.option_states(w.SlotIndex(w.open_package(filled)), target), [True, True])
                before = w.SlotIndex(w.open_package(original)); after = w.SlotIndex(w.open_package(filled))
                for slot in before.slots:
                    if slot.kind == 'p' and '/tbl[0]/tr[0]/tc[0]/tbl[0]/' not in slot.path:
                        self.assertEqual(etree.tostring(slot.el), etree.tostring(after.by_id[slot.id].el))

    def test_nested_marked_spacer_and_competing_questions_are_not_silently_claimed(self):
        for spacer in ('underline', 'other_question'):
            with self.subTest(spacer=spacer):
                result = w.inspect(data(self.form(spacer=spacer)))
                self.assertFalse([f for f in result['fields'] if f['kind'] == 'multi_choice'])
                self.assertTrue(any('Nested option' in note for note in result['notes']))

    def test_nested_source_cannot_lose_local_subject_or_page_boundary(self):
        for case in ('foreign_prefix', 'unpaired_subsection', 'page_break', 'long_question'):
            with self.subTest(case=case):
                doc = self.form()
                owner = doc.tables[0].cell(0, 0)
                if case == 'foreign_prefix':
                    source = OxmlElement('w:p');run = OxmlElement('w:r');text = OxmlElement('w:t')
                    text.text = 'Donor';run.append(text);source.append(run)
                    owner.paragraphs[0]._p.addprevious(source)
                elif case == 'unpaired_subsection':
                    owner.paragraphs[1].text = 'Donor:'
                    source = OxmlElement('w:p');run = OxmlElement('w:r');text = OxmlElement('w:t')
                    text.text = 'During exertion:';run.append(text);source.append(run)
                    owner.paragraphs[1]._p.addnext(source)
                elif case == 'page_break':
                    owner.paragraphs[1]._p.get_or_add_pPr().append(OxmlElement('w:pageBreakBefore'))
                else:
                    owner.paragraphs[0].text = 'Which symptoms occur? Tick all that apply. ' + 'Preserve this qualifier. ' * 20
                result = w.inspect(data(doc))
                groups = [f for f in result['fields'] if f['locator']['target']['type'] == 'checkboxes']
                if case == 'long_question':
                    self.assertEqual(len(groups), 1)
                    self.assertEqual(groups[0]['label'], owner.paragraphs[0].text.strip())
                else:
                    self.assertFalse(groups)
                    self.assertTrue(any('Nested option' in note for note in result['notes']))
                    if case == 'foreign_prefix': self.assertTrue(any('Donor' in note for note in result['notes']))

    def test_nested_subordinate_boxes_refuse_and_write_in_is_explicitly_visible(self):
        mixed = w.inspect(data(self.form(mixed=True)))
        self.assertFalse([f for f in mixed['fields'] if f['locator']['target']['type'] == 'checkboxes'])
        self.assertTrue(any('subordinate' in note and 'Right' in note and 'Left' in note for note in mixed['notes']))
        write_in = w.inspect(data(self.form(write_in=True)))
        self.assertTrue(any(f['label'] == 'Other:' and f['locator']['target']['type'] == 'span'
                            for f in write_in['fields']))

    def test_single_nested_option_keeps_its_complete_question_and_subsection(self):
        for subsection in (False, True):
            with self.subTest(subsection=subsection):
                doc = Document()
                doc.add_paragraph('Current visit', style='Heading 1')
                outer = doc.add_table(rows=1, cols=1).cell(0, 0)
                outer.text = 'Which finding concerns the donor? Tick all that apply.'
                if subsection: outer.add_paragraph('During exertion:')
                outer.add_paragraph()
                outer.add_table(rows=1, cols=1).cell(0, 0).text = '☐ Other'
                result = w.inspect(data(doc))
                group = next(f for f in result['fields'] if f['kind'] == 'checkbox')
                self.assertEqual(group['label'], 'Other')
                self.assertIn(outer.paragraphs[0].text, group['context'])
                self.assertIn('Current visit', group['context'])
                if subsection: self.assertIn('During exertion:', group['context'])
                self.assertEqual(group['options'], ['Other'])

    def test_nested_grid_cannot_drop_internal_headers_or_governing_signing_context(self):
        for case in ('headers', 'consent', 'signature'):
            with self.subTest(case=case):
                doc = Document()
                if case != 'headers':
                    doc.add_paragraph('Consent' if case == 'consent' else 'Signed declaration', style='Heading 1')
                outer = doc.add_table(rows=1, cols=1).cell(0, 0)
                outer.text = 'Which procedures are authorized? Tick all that apply.'
                outer.add_paragraph()
                grid = outer.add_table(rows=2 if case == 'headers' else 1, cols=2)
                if case == 'headers':
                    grid.cell(0, 0).text = 'Donor';grid.cell(0, 1).text = 'Recipient'
                grid.cell(len(grid.rows)-1, 0).text = '☐ Abnormal' if case == 'headers' else '☐ Blood sampling'
                grid.cell(len(grid.rows)-1, 1).text = '☐ Abnormal' if case == 'headers' else '☐ Tissue storage'
                result = w.inspect(data(doc))
                self.assertFalse([f for f in result['fields'] if f['locator']['target']['type'] == 'checkboxes'])
                self.assertTrue(any('Nested option' in note for note in result['notes']))
                if case == 'headers':
                    self.assertTrue(any('Donor' in note and 'Recipient' in note for note in result['notes']))


class NestedWriteIns(unittest.TestCase):
    def form(self, *, control='glyph', checked=False, blank='underscores', label='Additional detail:',
             heading='Current visit', subsection=True, suffix=''):
        from .test_word_v3 import checkbox_control, ff_field
        doc = NestedCheckboxOwnership().form(heading=heading, subsection=subsection)
        paragraph = doc.tables[0].cell(0, 0).tables[0].cell(0, 1).paragraphs[0]
        paragraph.clear()
        if control == 'glyph': paragraph.add_run('☒' if checked else '☐')
        elif control == 'sdt': paragraph._p.append(checkbox_control(checked=checked))
        else: ff_field(paragraph, 'checkbox', checked=checked)
        paragraph.add_run(' ' + label + ' ')
        paragraph.add_run('_____' if blank == 'underscores' else '\t').underline = blank == 'tab'
        if suffix: paragraph.add_run(suffix)
        return doc

    def test_disjoint_blank_and_choice_preserve_source_and_independent_states(self):
        for control in ('glyph', 'sdt', 'ff'):
            for blank in ('underscores', 'tab'):
                with self.subTest(control=control, blank=blank):
                    raw = data(self.form(control=control, checked=True, blank=blank))
                    result = w.inspect(raw)
                    spans = [f for f in result['fields'] if f['locator']['target']['type'] == 'span']
                    self.assertEqual(len(spans), 1)
                    field = spans[0]
                    self.assertEqual(field['label'], 'Additional detail:')
                    self.assertEqual(field['context'],
                        'Current visit > Which symptoms occur? Tick all that apply. > During exertion:')
                    group = next(f for f in result['fields'] if f['kind'] == 'multi_choice')
                    self.assertEqual(group['printed_ticked'], [1])
                    self.assertFalse(result['notes'])
                    with ZipFile(io.BytesIO(raw)) as before, ZipFile(io.BytesIO(w.fill(raw, result['fields'], {}))) as after:
                        self.assertEqual({n: before.read(n) for n in before.namelist()},
                                         {n: after.read(n) for n in after.namelist()})
                    for value in ('Ordinary detail.', 'Literal ☐ and ☒ remain text.\n' + 'Retained long detail. ' * 45):
                        output = w.fill(raw, result['fields'], {field['key']: value})
                        index = w.SlotIndex(w.open_package(output));slot = index.by_id['word/document.xml#' + field['locator']['slot']]
                        self.assertEqual(w.paragraph_text(slot.el)[0],
                                         (w.CHECKED if control == 'glyph' else w.BOX) + ' Additional detail: ' + value)
                        # Read the first physical source control, never an inserted glyph ordinal.
                        self.assertTrue(w.box_state(index, group['locator']['target']['boxes'][-1]))
                        if blank == 'tab':
                            inserted = next(node for node in slot.el.iter(w.T) if value.split('\n')[0] in (node.text or ''))
                            self.assertIsNotNone(inserted.getparent().find(w.q('w:rPr') + '/' + w.q('w:u')))
                    output = w.fill(raw, result['fields'], {group['key']: ['Tightness']})
                    index = w.SlotIndex(w.open_package(output));slot = index.by_id['word/document.xml#' + field['locator']['slot']]
                    self.assertIn('_____' if blank == 'underscores' else '\t', w.paragraph_text(slot.el)[0])
                    self.assertFalse(w.box_state(index, group['locator']['target']['boxes'][-1]))

    def test_ambiguous_blank_shapes_remain_visible(self):
        for suffix in (' then _____', ' after review', ' ☐ Separate'):
            with self.subTest(suffix=suffix):
                result = w.inspect(data(self.form(suffix=suffix)))
                self.assertFalse([f for f in result['fields'] if f['locator']['target']['type'] == 'span'])
                self.assertTrue(any('Printed' in note and suffix.strip() in note for note in result['notes']))

    def test_signing_scope_and_complete_context_limits_are_not_relaxed(self):
        for heading, label in (('Consent', 'Additional detail:'), ('Signed declaration', 'Additional detail:'),
                               ('Current visit', 'Initial here:'), ('Current visit', 'Signature:')):
            with self.subTest(heading=heading, label=label):
                result = w.inspect(data(self.form(heading=heading, label=label)))
                spans = [f for f in result['fields'] if f['locator']['target']['type'] == 'span'
                         and f['kind'] not in ('signature', 'initials')]
                self.assertFalse(spans)
        doc = self.form(heading='', subsection=False)
        question = 'Which ' + 'q' * 993 + '?'
        self.assertEqual(len(question), 1000)
        doc.tables[0].cell(0, 0).paragraphs[0].text = question
        result = w.inspect(data(doc))
        field = next(f for f in result['fields'] if f['locator']['target']['type'] == 'span')
        self.assertEqual(field['context'], question)
        doc.tables[0].cell(0, 0).paragraphs[0].text += 'x'
        result = w.inspect(data(doc))
        self.assertFalse([f for f in result['fields'] if f['locator']['target']['type'] == 'span'])
        self.assertTrue(any('exceeds' in note for note in result['notes']))

    def test_changed_source_identity_refuses_write_in(self):
        raw = data(self.form(control='ff'))
        result = w.inspect(raw)
        field = next(f for f in result['fields'] if f['locator']['target']['type'] == 'span')
        field['locator']['fingerprint'] = '0' * 64
        with self.assertRaises(w.V3Failure):
            w.fill(raw, result['fields'], {field['key']: 'Synthetic detail.'})


class TabColumnOwnership(unittest.TestCase):
    """Source proposals only. Rendered ownership is qualified by the layout service."""

    def proposal(self, raw):
        detector = w.Detector(w.open_package(raw), layout_proposal=True)
        areas, notes = detector.detect()
        fields = [w.area_to_field(i + 1, area) for i, area in enumerate(areas)]
        for field in fields: w.annotate_printed(detector.index, field)
        return {'fields': fields, 'notes': notes}

    def proposed_fill(self, raw, fields, answers):
        # Private preparer path tests writer mechanics, never rendered meaning.
        return w.fill(raw, fields, answers, retained_layout=[])

    def form(self, *, control='glyph', inherited=False, problem='', many=False, count=2):
        from .test_word_v3 import checkbox_control, ff_field
        doc = Document();doc.add_paragraph('Current visit', style='Heading 1')
        cell = doc.add_table(rows=1, cols=2).cell(0, 0)
        cell.text = 'Examination; tick all that apply.' if many else 'Examination'
        options = ['Normal', 'Abnormal', 'Unavailable'][:count]
        header = cell.add_paragraph('\t' + '\t'.join(options))
        row = cell.add_paragraph('Finding')
        for i in range(count):
            row.add_run('\t' + str(8-i) + ' ')
            if control == 'glyph':row.add_run('☐')
            elif control == 'sdt':row._p.append(checkbox_control())
            else:ff_field(row, 'checkbox', name='C'+str(i))
        def stops(ppr, shift=0):
            tabs = OxmlElement('w:tabs')
            for i in range(count):
                tab = OxmlElement('w:tab');tab.set(qn('w:val'), 'center')
                tab.set(qn('w:pos'), str(1800 + i*900 + shift));tabs.append(tab)
            ppr.append(tabs)
        if inherited:
            style = doc.styles.add_style('SyntheticColumns', WD_STYLE_TYPE.PARAGRAPH)
            stops(style.element.get_or_add_pPr())
            header.style = style;row.style = style
        else:
            stops(header._p.get_or_add_pPr());stops(row._p.get_or_add_pPr())
        if problem == 'mismatch':
            row._p.get_or_add_pPr().remove(row._p.get_or_add_pPr().find(qn('w:tabs')))
            stops(row._p.get_or_add_pPr(), 20)
        elif problem == 'unknown_style':
            row._p.get_or_add_pPr().insert(0,OxmlElement('w:pStyle'))
            row._p.get_or_add_pPr().find(qn('w:pStyle')).set(qn('w:val'),'MissingColumns')
        elif problem == 'cleared':
            tabs=OxmlElement('w:tabs');tab=OxmlElement('w:tab');tab.set(qn('w:val'),'clear')
            tab.set(qn('w:pos'),'1800');tabs.append(tab);row._p.get_or_add_pPr().append(tabs)
        elif problem == 'indent':
            ind=OxmlElement('w:ind');ind.set(qn('w:left'),'720');row._p.get_or_add_pPr().append(ind)
        elif problem == 'duplicate':header.text='\tNormal\tNormal'
        elif problem == 'missing':header._p.getparent().remove(header._p)
        elif problem == 'intervening':header.insert_paragraph_before('Donor only')
        elif problem == 'signing':cell.paragraphs[0].text='Signed declaration'
        return doc

    def test_source_proposals_preserve_literal_options_and_writer_controls(self):
        for control in ('glyph','sdt','ff'):
            for inherited in (False,True):
                for count in (1,2,3):
                    with self.subTest(control=control,inherited=inherited,count=count):
                        raw=data(self.form(control=control,inherited=inherited,count=count))
                        result=self.proposal(raw)
                        fields=[f for f in result['fields'] if f['locator']['target']['type']=='checkboxes']
                        self.assertEqual(len(fields),1)
                        field=fields[0]
                        self.assertEqual(field['label'],'Finding')
                        self.assertEqual(field['options'],['Normal','Abnormal','Unavailable'][:count])
                        self.assertIn('Examination',field['context']);self.assertIn('Current visit',field['context'])
                        self.assertFalse(result['notes'])
                        chosen=True if count==1 else field['options'][-1]
                        out=self.proposed_fill(raw,result['fields'],{field['key']:chosen})
                        self.assertEqual(w.option_states(w.SlotIndex(w.open_package(out)),field),[False]*(count-1)+[True])
                        before=w.SlotIndex(w.open_package(raw));after=w.SlotIndex(w.open_package(out))
                        for slot in before.slots:
                            if slot.kind=='p' and not w.place_box_refs(before,slot):
                                self.assertEqual(etree.tostring(slot.el),etree.tostring(after.by_id[slot.id].el))

    def test_unknown_or_conflicting_geometry_and_source_are_visible(self):
        for problem in ('mismatch','unknown_style','cleared','indent','duplicate','missing','intervening','signing'):
            with self.subTest(problem=problem):
                result=self.proposal(data(self.form(problem=problem,inherited=problem=='cleared')))
                self.assertFalse([f for f in result['fields'] if f['locator']['target']['type']=='checkboxes'])
                self.assertTrue(any('Tab-column' in n and 'Finding' in n for n in result['notes']))

    def test_literal_cardinality_is_preserved_and_missing_states_unchanged(self):
        raw=data(self.form(many=True));result=self.proposal(raw)
        field=next(f for f in result['fields'] if f['locator']['target']['type']=='checkboxes')
        self.assertEqual(field['kind'],'multi_choice')
        out=self.proposed_fill(raw,result['fields'],{field['key']:field['options']})
        self.assertEqual(w.option_states(w.SlotIndex(w.open_package(out)),field),[True,True])
        with ZipFile(io.BytesIO(raw)) as old,ZipFile(io.BytesIO(self.proposed_fill(raw,result['fields'],{}))) as new:
            self.assertEqual({n:old.read(n) for n in old.namelist()},{n:new.read(n) for n in new.namelist()})

    def test_inherited_stops_clear_override_and_same_indent_are_resolved(self):
        doc=self.form(inherited=True)
        cell=doc.tables[0].cell(0,0)
        for paragraph in cell.paragraphs[1:]:
            ppr=paragraph._p.get_or_add_pPr()
            tabs=OxmlElement('w:tabs')
            for position,kind in ((1800,'clear'),(2000,'center')):
                tab=OxmlElement('w:tab');tab.set(qn('w:val'),kind);tab.set(qn('w:pos'),str(position));tabs.append(tab)
            ppr.append(tabs)
            ind=OxmlElement('w:ind');ind.set(qn('w:left'),'120');ppr.append(ind)
        result=self.proposal(data(doc));field=next(f for f in result['fields'] if f['kind']=='choice')
        self.assertEqual((field['label'],field['options']),('Finding',['Normal','Abnormal']))
        self.assertFalse(result['notes'])

    def test_marked_column_heading_is_preserved_without_adjacent_subject(self):
        for heading in ('Donor','Consent'):
            with self.subTest(heading=heading):
                doc=self.form();table=doc.tables[0];row=table.add_row()._tr
                row.getparent().remove(row);table._tbl.insert(list(table._tbl).index(table.rows[0]._tr),row)
                header=table.rows[0];header.cells[0].text=heading;header.cells[1].text='Recipient'
                header._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                result=self.proposal(data(doc));fields=[f for f in result['fields'] if f['locator']['target']['type']=='checkboxes']
                if heading=='Consent':
                    self.assertFalse(fields);self.assertTrue(any('protected' in n for n in result['notes']))
                else:
                    self.assertEqual(len(fields),1);self.assertIn('Donor',fields[0]['context'])
                    self.assertNotIn('Recipient',fields[0]['context'])

    def test_direct_numeric_option_words_are_not_reinterpreted_as_column_codes(self):
        doc=Document();doc.add_table(rows=1,cols=1).cell(0,0).text='Score?\t☐ 1\t☐ 2'
        result=self.proposal(data(doc));field=next(f for f in result['fields'] if f['kind']=='choice')
        self.assertEqual(field['options'],['1','2'])
        self.assertFalse(any('Tab-column' in note for note in result['notes']))

    def test_unused_tab_stops_cannot_hide_other_inherited_spacer_markers(self):
        doc=self.form();header=doc.tables[0].cell(0,0).paragraphs[1]
        spacer=header.insert_paragraph_before()
        style=doc.styles.add_style('MarkedBlank',WD_STYLE_TYPE.CHARACTER)
        style.element.get_or_add_rPr().append(OxmlElement('w:u'))
        spacer.add_run().style=style
        result=self.proposal(data(doc))
        self.assertFalse([f for f in result['fields'] if f['locator']['target']['type']=='checkboxes'])
        self.assertTrue(any('Tab-column' in note for note in result['notes']))

    def test_governing_cardinality_and_conflicting_scopes(self):
        for conflict in (False, True):
            with self.subTest(conflict=conflict):
                doc=self.form()
                doc.paragraphs[0].text='Symptoms; tick all that apply.'
                cell=doc.tables[0].cell(0,0);cell.paragraphs[1].text='\tPain\tSwelling'
                if conflict:cell.paragraphs[0].text='Examination; choose one.'
                raw=data(doc);result=self.proposal(raw)
                fields=[f for f in result['fields'] if f['locator']['target']['type']=='checkboxes']
                if conflict:
                    self.assertFalse(fields);self.assertTrue(any('cardinality' in n for n in result['notes']))
                else:
                    self.assertEqual(fields[0]['kind'],'multi_choice')
                    out=self.proposed_fill(raw,result['fields'],{fields[0]['key']:['Pain','Swelling']})
                    self.assertEqual(w.option_states(w.SlotIndex(w.open_package(out)),fields[0]),[True,True])
        doc=self.form();doc.paragraphs[0].text='Symptoms; select up to 1.'
        result=self.proposal(data(doc))
        self.assertFalse([f for f in result['fields'] if f['locator']['target']['type']=='checkboxes'])
        self.assertTrue(any('unrepresented selection limit' in n for n in result['notes']))

    def test_hidden_or_unknown_source_visibility_never_supplies_columns(self):
        for source in ('header','title','row','heading','marked_header'):
            for mode in ('direct','paragraph','character','unknown_character'):
                with self.subTest(source=source,mode=mode):
                    doc=self.form();table=doc.tables[0];cell=table.cell(0,0)
                    para={'header':cell.paragraphs[1],'title':cell.paragraphs[0],
                          'row':cell.paragraphs[2],'heading':doc.paragraphs[0]}.get(source)
                    if source=='marked_header':
                        row=table.add_row()._tr;row.getparent().remove(row)
                        table._tbl.insert(list(table._tbl).index(table.rows[0]._tr),row)
                        row=table.rows[0];row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                        row.cells[0].text='Donor';para=row.cells[0].paragraphs[0]
                    if mode=='direct':para.runs[0]._r.get_or_add_rPr().append(OxmlElement('w:vanish'))
                    elif mode=='unknown_character':
                        ref=OxmlElement('w:rStyle');ref.set(qn('w:val'),'MissingVisibility')
                        para.runs[0]._r.get_or_add_rPr().append(ref)
                    else:
                        kind=WD_STYLE_TYPE.PARAGRAPH if mode=='paragraph' else WD_STYLE_TYPE.CHARACTER
                        parent=doc.styles.add_style('HiddenParent',kind)
                        parent.element.get_or_add_rPr().append(OxmlElement('w:webHidden'))
                        child=doc.styles.add_style('HiddenChild',kind);child.base_style=parent
                        if mode=='paragraph':
                            if source=='heading':parent.base_style=doc.styles['Heading 1']
                            para.style=child
                        else:para.runs[0].style=child
                    result=self.proposal(data(doc))
                    self.assertFalse([f for f in result['fields'] if f['locator']['target']['type']=='checkboxes'])
                    self.assertTrue(any('Tab-column' in n for n in result['notes']))

    def test_visibility_defaults_and_explicit_off_have_bounded_meaning(self):
        for mode in ('doc_default','table_default','character_default','paragraph_default',
                     'table_inherited','character_cycle','explicit_off','ancestor_on_direct_off'):
            with self.subTest(mode=mode):
                doc=self.form();header=doc.tables[0].cell(0,0).paragraphs[1]
                if mode=='doc_default':
                    holder=doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
                elif mode.endswith('_default'):
                    kind=mode.split('_')[0]
                    style=next(s for s in doc.styles.element if s.get(qn('w:type'))==kind
                               and s.get(qn('w:default'))=='1')
                    holder=style.get_or_add_rPr()
                elif mode=='table_inherited':
                    parent=doc.styles.add_style('HiddenTableParent',WD_STYLE_TYPE.TABLE)
                    child=doc.styles.add_style('HiddenTableChild',WD_STYLE_TYPE.TABLE);child.base_style=parent
                    doc.tables[0].style=child;holder=parent.element.get_or_add_rPr()
                elif mode in ('character_cycle','ancestor_on_direct_off'):
                    style=doc.styles.add_style('VisibilityAncestor',WD_STYLE_TYPE.CHARACTER)
                    header.runs[0].style=style;holder=style.element.get_or_add_rPr()
                    if mode=='character_cycle':style.base_style=style
                    else:
                        off=OxmlElement('w:vanish');off.set(qn('w:val'),'0')
                        header.runs[0]._r.get_or_add_rPr().append(off)
                else:holder=header.runs[0]._r.get_or_add_rPr()
                marker=OxmlElement('w:vanish')
                if mode=='explicit_off':marker.set(qn('w:val'),'0')
                holder.append(marker)
                result=self.proposal(data(doc))
                fields=[f for f in result['fields'] if f['locator']['target']['type']=='checkboxes']
                self.assertEqual(bool(fields),mode=='explicit_off')
                if mode!='explicit_off':self.assertTrue(any('Tab-column' in n for n in result['notes']))
