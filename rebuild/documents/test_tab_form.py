"""Owner-supplied blank form acceptance; output is temporary and synthetic only."""

from copy import deepcopy
import hashlib
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from .service import DocumentError, fill_template, inspect_template


ORIGINAL = Path(__file__).with_name("__fixtures__") / "Colonoscopy_Operative_Template.docx"
SOURCE_SHA256 = "8ae1e2b0ebeec317c038b8bfeed6a0824055f44d689c039ab450404593f2815a"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
# Fixed expectations derived independently from the received original's XML.
# A logical field can occupy several physical answer tabs; these are not 28 questions.
EXPECTED = [
    ("PATIENT NAME", [(2, 1)]),
    ("PATIENT DOB", [(2, 4)]),
    ("PATIENT ID #", [(3, 1)]),
    ("DATE OF PROCEDURE", [(3, 4)]),
    ("SURGEON / ENDOSCOPIST", [(4, 1)]),
    ("PROCEDURE", [(5, 1)]),
    ("ANESTHESIA", [(6, 1)]),
    ("BOWEL PREPARATION / BBPS", [(6, 4)]),
    ("EXTENT OF EXAMINATION", [(7, 1)]),
    ("PREOPERATIVE DIAGNOSIS / INDICATION", [(9, 0)]),
    ("POSTOPERATIVE FINDINGS — 1", [(11, 1)]),
    ("POSTOPERATIVE FINDINGS — 2", [(12, 1)]),
    ("POSTOPERATIVE FINDINGS — 3", [(13, 1)]),
    ("CONSENT", [(15, 0), (16, 0)]),
    ("DESCRIPTION OF PROCEDURE", [(18, 0), (19, 0), (20, 0), (21, 0)]),
    ("SPECIMENS", [(22, 1)]),
    ("ESTIMATED BLOOD LOSS", [(23, 1)]),
    ("COMPLICATIONS", [(23, 4)]),
    ("RECOMMENDATIONS — 1", [(25, 1)]),
    ("RECOMMENDATIONS — 2", [(26, 1)]),
    ("RECOMMENDATIONS — 3", [(27, 1)]),
    ("Electronically signed by", [(29, 1)]),
    ("DD (dictated)", [(30, 1)]),
    ("DT (transcribed)", [(30, 4)]),
]


def document_xml(path):
    with ZipFile(path) as package:
        return etree.fromstring(package.read("word/document.xml"))


def body_paragraphs(path):
    return document_xml(path).findall("./w:body/w:p", NS)


def visible_text(element):
    return "".join(element.xpath(".//w:t/text()", namespaces=NS))


class UnderlinedTabFormTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="quillmd-tab-acceptance-")
        self.directory = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def fields(self):
        return inspect_template(ORIGINAL)["fields"]

    def assert_failure(self, operation, *args, codes=None):
        with self.assertRaises(DocumentError) as result:
            operation(*args)
        if codes:
            self.assertIn(result.exception.code, codes)
        return result.exception

    def assert_original_and_parts_preserved(self, output):
        self.assertEqual(hashlib.sha256(ORIGINAL.read_bytes()).hexdigest(), SOURCE_SHA256)
        with ZipFile(ORIGINAL) as original, ZipFile(output) as filled:
            self.assertEqual(original.namelist(), filled.namelist())
            for name in original.namelist():
                if name != "word/document.xml":
                    self.assertEqual(original.read(name), filled.read(name), name)
        original_paragraphs, filled_paragraphs = body_paragraphs(ORIGINAL), body_paragraphs(output)
        self.assertEqual(len(filled_paragraphs), 31)
        for index, original_paragraph in enumerate(original_paragraphs):
            filled_paragraph = filled_paragraphs[index]
            # Preserve every label in place and retain its original run style.
            original_label_runs = [r for r in original_paragraph.findall("w:r", NS) if r.findall("w:t", NS)]
            filled_label_runs = [r for r in filled_paragraph.findall("w:r", NS) if r.findall("w:t", NS)]
            cursor = 0
            for run in original_label_runs:
                text = visible_text(run)
                while cursor < len(filled_label_runs) and visible_text(filled_label_runs[cursor]) != text:
                    cursor += 1
                self.assertLess(cursor, len(filled_label_runs), (index, text))
                old, new = run.find("w:rPr", NS), filled_label_runs[cursor].find("w:rPr", NS)
                self.assertEqual(etree.tostring(old) if old is not None else None,
                                 etree.tostring(new) if new is not None else None, (index, text))
                cursor += 1
            # Existing tab-stop definitions, paragraph rules, and section page setup survive.
            for tag in ("tabs", "pBdr"):
                old, new = original_paragraph.find(f"w:pPr/w:{tag}", NS), filled_paragraph.find(f"w:pPr/w:{tag}", NS)
                self.assertEqual(etree.tostring(old) if old is not None else None,
                                 etree.tostring(new) if new is not None else None, (index, tag))
        self.assertEqual(etree.tostring(document_xml(ORIGINAL).find("./w:body/w:sectPr", NS)),
                         etree.tostring(document_xml(output).find("./w:body/w:sectPr", NS)))

    def test_original_inventory_and_exact_reviewable_mapping(self):
        self.assertEqual(ORIGINAL.stat().st_size, 38393)
        self.assertEqual(hashlib.sha256(ORIGINAL.read_bytes()).hexdigest(), SOURCE_SHA256)
        paragraphs = body_paragraphs(ORIGINAL)
        self.assertEqual(len(paragraphs), 31)
        actual_answers, separators = [], []
        for paragraph_index, paragraph in enumerate(paragraphs):
            for run_index, run in enumerate(paragraph.findall("w:r", NS)):
                for child_index, child in enumerate(run):
                    if child.tag != qn("w:tab"):
                        continue
                    underline = run.find("w:rPr/w:u", NS)
                    destination = actual_answers if underline is not None and underline.get(qn("w:val")) == "single" else separators
                    destination.append((paragraph_index, run_index, child_index))
        expected_answers = [(p, r, 1) for _, positions in EXPECTED for p, r in positions]
        self.assertEqual(actual_answers, expected_answers)
        self.assertEqual(separators, [(p, 2, 1) for p in (2, 3, 6, 23, 30)])
        fields = self.fields()
        self.assertEqual([field["label"] for field in fields], [label for label, _ in EXPECTED])
        self.assertEqual(len({field["key"] for field in fields}), 24)
        mapped = []
        for field, (label, positions) in zip(fields, EXPECTED):
            locator = field["locator"]["tab_form"]
            self.assertEqual(locator["version"], 1)
            self.assertEqual(locator["source_sha256"], SOURCE_SHA256)
            anchors = locator["anchors"]
            self.assertEqual([(a["paragraph"], a["run"], a["child"]) for a in anchors], [(p, r, 1) for p, r in positions], label)
            self.assertTrue(all(a["part"] == "word/document.xml" for a in anchors))
            mapped.extend((a["paragraph"], a["run"], a["child"]) for a in anchors)
        self.assertEqual(mapped, actual_answers)
        self.assertEqual(len(mapped), len(set(mapped)))
        self.assertEqual(fields, self.fields(), "Repeated inspection must be stable")

    def test_concise_answers_occupy_independent_original_regions_once(self):
        fields = self.fields()
        answers = {field["key"]: f"Synthetic answer {index:02d}" for index, field in enumerate(fields)}
        output = self.directory / "concise.docx"
        result = fill_template(ORIGINAL, fields, answers, output)
        self.assertEqual(result["filled"], 24)
        self.assertEqual(result["missing"], [])
        paragraphs = body_paragraphs(output)
        whole = visible_text(document_xml(output))
        for field, (_, positions) in zip(fields, EXPECTED):
            answer = answers[field["key"]]
            self.assertEqual(whole.count(answer), 1, field["label"])
            region = "".join(visible_text(paragraphs[p]) for p, _ in positions)
            self.assertIn(answer, region, field["label"])
        self.assertEqual(len(document_xml(output).xpath(".//w:r/w:tab", namespaces=NS)), 9,
                         "Five separators and four unused physical continuation lines stay intact")
        # Paired labels must still separate their answers in the same paragraph.
        for left, right, paragraph in ((0, 1, 2), (2, 3, 3), (6, 7, 6), (16, 17, 23), (22, 23, 30)):
            text = visible_text(paragraphs[paragraph])
            left_answer, right_answer = answers[fields[left]["key"]], answers[fields[right]["key"]]
            self.assertLess(text.index(left_answer), text.index(EXPECTED[right][0]))
            self.assertLess(text.index(EXPECTED[right][0]), text.index(right_answer))
        self.assert_original_and_parts_preserved(output)

    def test_missing_consent_signature_and_dates_remain_missing_without_invention(self):
        fields = self.fields()
        output = self.directory / "missing.docx"
        result = fill_template(ORIGINAL, fields, {}, output)
        self.assertEqual(result["filled"], 0)
        self.assertEqual(set(result["missing"]), {field["key"] for field in fields})
        self.assertEqual(visible_text(document_xml(ORIGINAL)), visible_text(document_xml(output)))
        self.assertEqual(len(document_xml(output).xpath(".//w:r[w:rPr/w:u[@w:val='single']]/w:tab", namespaces=NS)), 28)
        self.assert_original_and_parts_preserved(output)

    def test_explicit_continuation_lines_use_existing_regions_and_expand_the_last(self):
        fields = self.fields()
        consent = ["Synthetic consent first", "Synthetic consent second", "Synthetic consent overflow"]
        procedure = [f"Synthetic procedure line {i}" for i in range(6)]
        answers = {fields[13]["key"]: "\n".join(consent), fields[14]["key"]: "\n".join(procedure)}
        output = self.directory / "continuations.docx"
        result = fill_template(ORIGINAL, fields, answers, output)
        self.assertEqual(result["filled"], 2)
        paragraphs = body_paragraphs(output)
        self.assertEqual(visible_text(paragraphs[15]), consent[0])
        self.assertEqual(visible_text(paragraphs[16]), "".join(consent[1:]))
        for offset in range(3):
            self.assertEqual(visible_text(paragraphs[18 + offset]), procedure[offset])
        self.assertEqual(visible_text(paragraphs[21]), "".join(procedure[3:]))
        whole = visible_text(document_xml(output))
        for line in consent + procedure:
            self.assertEqual(whole.count(line), 1, line)
        self.assertEqual(len(paragraphs[16].xpath(".//w:br", namespaces=NS)), 1)
        self.assertEqual(len(paragraphs[21].xpath(".//w:br", namespaces=NS)), 2)
        self.assert_original_and_parts_preserved(output)

    def test_long_answers_flow_in_their_actual_regions_without_duplication(self):
        fields = self.fields()
        answers = {field["key"]: f"Short synthetic {index:02d}" for index, field in enumerate(fields)}
        for index, count in ((0, 80), (13, 160), (14, 450), (16, 100), (22, 80)):
            answers[fields[index]["key"]] = " ".join(f"synthetic{index}word{n}" for n in range(count))
        output = self.directory / "long.docx"
        fill_template(ORIGINAL, fields, answers, output)
        paragraphs = body_paragraphs(output)
        whole = visible_text(document_xml(output))
        for field, (_, positions) in zip(fields, EXPECTED):
            answer = answers[field["key"]]
            self.assertEqual(whole.count(answer), 1, field["label"])
            region = "".join(visible_text(paragraphs[p]) for p, _ in positions)
            self.assertIn(answer, region, field["label"])
        for left, right, paragraph in ((0, 1, 2), (16, 17, 23), (22, 23, 30)):
            text = visible_text(paragraphs[paragraph])
            self.assertLess(text.index(answers[fields[left]["key"]]), text.index(EXPECTED[right][0]))
            self.assertLess(text.index(EXPECTED[right][0]), text.index(answers[fields[right]["key"]]))
        self.assert_original_and_parts_preserved(output)

    def test_mapping_change_and_same_locations_with_changed_source_are_refused(self):
        fields = self.fields()
        output = self.directory / "forbidden.docx"
        modified = deepcopy(fields)
        modified[0]["locator"]["tab_form"]["anchors"][0]["run"] = 4
        self.assert_failure(fill_template, ORIGINAL, modified, {}, output, codes={"mapping_mismatch", "stale_mapping"})
        changed_source = self.directory / "changed-source.docx"
        with ZipFile(ORIGINAL) as incoming, ZipFile(changed_source, "w") as outgoing:
            for member in incoming.infolist():
                value = incoming.read(member.filename)
                if member.filename == "docProps/core.xml":
                    value += b"\n"
                outgoing.writestr(member, value)
        self.assert_failure(fill_template, changed_source, fields, {}, output, codes={"mapping_mismatch", "stale_mapping"})
        self.assertFalse(output.exists())

    def test_orphan_answer_tab_does_not_become_an_invented_clinical_question(self):
        path = self.directory / "orphan.docx"
        document = Document()
        document.add_paragraph().add_run("\t").underline = True
        document.save(path)
        self.assert_failure(inspect_template, path)

    def test_decorative_underlining_rule_and_separator_tabs_are_not_answers(self):
        path = self.directory / "decoration.docx"
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Decorative heading").underline = True
        paragraph.add_run("\tseparator")
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        border.append(bottom)
        paragraph._p.get_or_add_pPr().append(border)
        document.save(path)
        self.assert_failure(inspect_template, path, codes={"no_fields"})

    def test_prior_placeholder_contract_is_not_silently_extended(self):
        path = self.directory / "prior-placeholder.docx"
        document = Document()
        document.add_paragraph("{{Existing question}}")
        document.add_paragraph("Later ordinary label: ").add_run("\t").underline = True
        document.save(path)
        fields = inspect_template(path)["fields"]
        self.assertEqual(len(fields), 1)
        self.assertEqual(fields[0]["label"], "Existing question")
        self.assertNotIn("tab_form", fields[0]["locator"])
