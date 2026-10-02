"""Printed instructions remain evidence when an answer is entered beside a question."""
import copy
import json
import tempfile
from pathlib import Path
import unittest

from docx import Document
from docx.shared import RGBColor
from lxml import etree

from . import word_v3 as w
from .test_question_boxes import data


def legacy_text(*, name='', label='', limit='1', points='2', result=' ', default=' ', help_text='Press TAB to continue.'):
    from .test_word_v3 import ff_text_field
    doc = Document(); p = doc.add_paragraph(label); ff_text_field(p, name=name)
    begin = next(r for r in p._p.findall(w.R) if r.find(w.FLDCHAR) is not None)
    ff = begin.find(w.FLDCHAR).find(w.q('w:ffData'))
    ti = ff.find(w.q('w:textInput'))
    if limit is not None: ti.append(w._el('w:maxLength', **{'w:val':limit}))
    ti.append(w._el('w:default', **{'w:val':default}))
    if help_text is not None: ff.append(w._el('w:statusText', **{'w:type':'text','w:val':help_text}))
    end = p._p.findall(w.R)[-1]
    run = w._el('w:r'); text = w._el('w:t'); text.text = result; run.append(text); end.addprevious(run)
    for r in p._p.findall(w.R):
        pr = w._el('w:rPr')
        if points is not None:
            for tag in ('w:sz','w:szCs'): pr.append(w._el(tag, **{'w:val':points}))
        r.insert(0,pr)
    return doc


class LegacyTextBounds(unittest.TestCase):
    def mapped(self, **kwargs):
        raw = data(legacy_text(**kwargs)); fields = w.inspect(raw)['fields']
        self.assertEqual(len(fields), 1)
        return raw, fields

    def result(self, raw, field):
        index = w.SlotIndex(w.open_package(raw)); loc = field['locator']
        runs, _ = w._field_result_runs(index.by_id[loc['part']+'#'+loc['slot']].el)
        return ''.join(w.text_of(r) for r in runs)

    def test_native_limit_is_enforced_from_source_for_old_and_forged_metadata(self):
        raw, fields = self.mapped(name='Code', label='Code: ', points='22')
        self.assertEqual(fields[0]['max_chars'], 1)
        for supplied in (1, None, 99999):
            altered = copy.deepcopy(fields)
            if supplied is None: altered[0].pop('max_chars')
            else: altered[0]['max_chars'] = supplied
            before = copy.deepcopy(altered)
            self.assertEqual(self.result(w.fill(raw, altered, {altered[0]['key']:'Ω'}), altered[0]), 'Ω')
            with self.assertRaises(w.V3Failure) as failure: w.fill(raw, altered, {altered[0]['key']:'ΩΩ'})
            self.assertEqual(failure.exception.code, 'answer_too_long')
            self.assertEqual(altered, before)

    def test_omitted_limit_is_unlimited_but_explicit_zero_and_malformed_caps_are_protected(self):
        raw, fields = self.mapped(name='Details', label='Details: ', points='22', limit=None)
        self.assertNotIn('max_chars', fields[0]); answer = 'A complete observation. ' * 40
        self.assertEqual(self.result(w.fill(raw, fields, {fields[0]['key']:answer}), fields[0]), answer)
        for limit in ('0','-1','32768','bad',''):
            with self.subTest(limit=limit):
                raw, fields = self.mapped(name='Code', limit=limit)
                self.assertTrue(fields[0]['write_protected'])
                self.assertEqual(w.fill(raw, fields, {}), raw)
                fields[0].pop('write_protected'); fields[0].pop('review_reason')
                with self.assertRaises(w.V3Failure) as failure: w.fill(raw, fields, {fields[0]['key']:'A'})
                self.assertEqual(failure.exception.code, 'protected_source_guidance')

    def test_populated_overlength_original_is_preserved_and_never_truncated_on_load(self):
        raw, fields = self.mapped(name='Code', label='Code: ', points='22', result='ORIGINAL', default='ORIGINAL')
        self.assertEqual(w.fill(raw, fields, {}), raw)
        self.assertEqual(self.result(raw, fields[0]), 'ORIGINAL')
        with self.assertRaises(w.V3Failure): w.fill(raw, fields, {fields[0]['key']:'LONGER ORIGINAL'})
        self.assertEqual(self.result(w.fill(raw, fields, {fields[0]['key']:'A'}), fields[0]), 'A')

    def test_micro_guidance_is_visible_protected_metadata_and_keeps_full_help(self):
        help_text = 'Press TAB to continue.\nKeep this full source guidance Ω.'
        raw, fields = self.mapped(help_text=help_text)
        field = fields[0]; self.assertTrue(field['write_protected']); self.assertEqual(field['max_chars'], 1)
        self.assertIn('help', field['review_reason'])
        self.assertEqual(field['source_instructions'][0]['text'], help_text)
        self.assertEqual(w.fill(raw, fields, {}), raw)
        for metadata in ('new', 'removed', 'false'):
            forged = copy.deepcopy(fields)
            if metadata != 'new':
                for key in ('write_protected','review_reason','max_chars','source_instructions'): forged[0].pop(key,None)
                if metadata == 'false': forged[0]['write_protected'] = False
            with self.assertRaises(w.V3Failure) as failure: w.fill(raw, forged, {field['key']:'A'})
            self.assertEqual(failure.exception.code, 'protected_source_guidance')

    def test_micro_source_protection_cannot_be_bypassed_by_replacing_its_parent(self):
        raw, fields = self.mapped(); field = fields[0]
        p = w.SlotIndex(w.open_package(raw)).by_id['word/document.xml#b/p[0]']
        forged = {'key':'old', 'kind':'text', 'label':'Prior mapping', 'locator':{
            'contract':'v3','part':p.part,'slot':p.path,'fingerprint':w.fingerprint(p.el),'target':{'type':'paragraph'}}}
        with self.assertRaises(w.V3Failure) as failure: w.fill(raw, [forged], {'old':'A'})
        self.assertEqual(failure.exception.code, 'protected_source_guidance')
        self.assertEqual(w.fill(raw, [forged], {}), raw)

    def test_real_named_visible_populated_and_mixed_style_inputs_are_not_micro_guidance(self):
        for variation in ({'name':'Code'}, {'label':'Code: '}, {'points':'22'}, {'points':None},
                          {'result':'A'}, {'default':'A'}, {'help_text':None}, {'limit':None}):
            with self.subTest(variation=variation):
                raw, fields = self.mapped(**variation)
                self.assertFalse(fields[0].get('write_protected'))
                self.assertEqual(self.result(w.fill(raw, fields, {fields[0]['key']:'B'}), fields[0]), 'B')

    def test_direct_writer_checks_native_cap_and_guidance_before_mutation(self):
        for kwargs, answer, code in (({'name':'Code','points':'22'},'AB','answer_too_long'),
                                     ({},'A','protected_source_guidance')):
            raw, fields = self.mapped(**kwargs)
            slot = w.SlotIndex(w.open_package(raw)).by_id['word/document.xml#b/p[0]/ff[0]']
            before = etree.tostring(slot.el.getparent())
            with self.assertRaises(w.V3Failure) as failure: w.write_ff_text(slot.el,answer)
            self.assertEqual(failure.exception.code,code)
            self.assertEqual(etree.tostring(slot.el.getparent()),before)

    def test_repeated_headers_use_the_tightest_native_limit(self):
        doc = Document(); doc.settings.odd_and_even_pages_header_footer = True
        for header, limit in zip((doc.sections[0].header, doc.sections[0].even_page_header), ('2','1')):
            p = legacy_text(name='Code',label='Code: ',points='22',limit=limit).paragraphs[0]._p
            header._element.replace(header.paragraphs[0]._p,p)
        raw = data(doc); fields = w.inspect(raw)['fields']; self.assertEqual(len(fields),1)
        self.assertEqual(fields[0]['max_chars'],1)
        self.assertEqual(len(fields[0]['locator']['target']['copies']),1)
        fields[0].pop('max_chars')
        with self.assertRaises(w.V3Failure) as failure: w.fill(raw,fields,{fields[0]['key']:'AB'})
        self.assertEqual(failure.exception.code,'answer_too_long')
        output = w.fill(raw,fields,{fields[0]['key']:'A'})
        index = w.SlotIndex(w.open_package(output))
        for slot in [s for s in index.slots if s.kind=='ff']:
            self.assertEqual(''.join(w.text_of(r) for r in w._field_result_runs(slot.el)[0]),'A')

    def test_visible_paired_cell_prompt_retains_a_real_one_character_input(self):
        doc = Document(); table = doc.add_table(rows=1,cols=2); table.cell(0,0).text='Code'
        table.cell(0,1)._tc.replace(table.cell(0,1).paragraphs[0]._p,legacy_text().paragraphs[0]._p)
        raw = data(doc); fields = w.inspect(raw)['fields']; self.assertEqual(len(fields),1)
        self.assertFalse(fields[0].get('write_protected')); self.assertEqual(fields[0]['max_chars'],1)
        self.assertEqual(self.result(w.fill(raw,fields,{fields[0]['key']:'A'}),fields[0]),'A')

    def test_sandbox_json_roundtrip_keeps_protection_and_original_help_xml(self):
        from documents import inspect_template, fill_template, DocumentError
        raw = data(legacy_text())
        with tempfile.TemporaryDirectory() as directory:
            original = Path(directory)/'original.docx'; final = Path(directory)/'final.docx'; original.write_bytes(raw)
            fields = json.loads(json.dumps(inspect_template(original,engine='v3')['fields']))
            self.assertTrue(fields[0]['write_protected'])
            fill_template(original,fields,{},final)
            self.assertEqual(final.read_bytes(),raw)
            fields[0].pop('write_protected'); fields[0].pop('review_reason')
            with self.assertRaises(DocumentError): fill_template(original,fields,{fields[0]['key']:'A'},final)
            self.assertEqual(original.read_bytes(),raw)


def form():
    doc = Document(); table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'What information has the patient reviewed?'
    for index, text in enumerate(('Do not share third-party information with the patient.',
                                   'The patient must have an opportunity to add comments.')):
        p = table.cell(0, 1).paragraphs[0] if index == 0 else table.cell(0, 1).add_paragraph()
        p.add_run(text).font.color.rgb = RGBColor.from_string('808080')
    return data(doc)


class GuidancePreservation(unittest.TestCase):
    def test_linked_guidance_is_kept_beside_its_answer_and_relationship(self):
        from .test_document_relations import link
        doc = Document(); table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Please confirm the patient category and management level:'
        p = table.cell(0, 1).paragraphs[0]
        p.add_run('See the published guidance. ').font.color.rgb = RGBColor.from_string('0070C0')
        link(p)
        for r in p._p.findall('.//' + w.R):
            props = r.find(w.q('w:rPr'))
            if props is None: props = etree.SubElement(r, w.q('w:rPr'))
            color = etree.SubElement(props, w.q('w:color')); color.set(w.q('w:val'), '0070C0')
        original = data(doc); fields = w.inspect(original)['fields']
        field = next(f for f in fields if f['locator']['slot'] == 'b/tbl[0]/tr[0]/tc[1]')
        before = w.open_package(original)
        final = w.fill(original, fields, {field['key']: 'The patient has category two management.'})
        after = w.open_package(final)
        self.assertEqual(before.zip.read('word/_rels/document.xml.rels'), after.zip.read('word/_rels/document.xml.rels'))
        before_p = w.SlotIndex(before).by_id['word/document.xml#b/tbl[0]/tr[0]/tc[1]'].el.find(w.q('w:p'))
        after_p = w.SlotIndex(after).by_id['word/document.xml#b/tbl[0]/tr[0]/tc[1]'].el.find(w.q('w:p'))
        self.assertEqual(etree.tostring(before_p), etree.tostring(after_p))
        self.assertIn('Published guidance', field['source_instructions'][0]['text'])

    def test_reference_examples_are_not_an_answer_column(self):
        for title in ('The following are examples of common accommodations.',
                      'Examples of treatment plans for reference', 'For reference: possible conditions'):
            doc = Document(); table = doc.add_table(rows=2, cols=1)
            table.cell(0, 0).text = title
            table.cell(1, 0).paragraphs[0].add_run('Attend scheduled appointments at [specify clinic].').font.color.rgb = RGBColor.from_string('0070C0')
            original = data(doc)
            self.assertEqual(w.inspect(original)['fields'], [], title)
            anchor = w.SlotIndex(w.open_package(original)).by_id['word/document.xml#b/tbl[0]/tr[1]/tc[0]']
            old = w.area_to_field(1, w.Area('text', title, anchor, {'type': 'cell_replace', 'all': True}))
            with self.assertRaises(w.V3Failure):
                w.fill(original, [old], {old['key']: 'The patient must attend weekly.'})

    def test_request_for_examples_still_has_an_answer_destination(self):
        doc = Document(); table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text = 'Please give examples of symptoms experienced:'
        table.cell(1, 0).paragraphs[0].add_run('Describe the frequency and duration.').font.color.rgb = RGBColor.from_string('0070C0')
        fields = w.inspect(data(doc))['fields']
        self.assertEqual(len(fields), 1)
        self.assertEqual(fields[0]['locator']['slot'], 'b/tbl[0]/tr[1]/tc[0]')

    def test_repeated_header_instructions_all_reach_the_mapping(self):
        from docx.shared import Inches
        doc = Document(); doc.settings.odd_and_even_pages_header_footer = True
        instructions = ['Do not disclose information about the mother.',
                        'Do not disclose information about the father.']
        for header, text in zip((doc.sections[0].header, doc.sections[0].even_page_header), instructions):
            table = header.add_table(rows=1, cols=2, width=Inches(6))
            table.cell(0, 0).text = 'What progress have you made?'
            table.cell(0, 1).paragraphs[0].add_run(text).font.color.rgb = RGBColor.from_string('808080')
        original = data(doc); fields = w.inspect(original)['fields']
        self.assertEqual(len(fields), 1)
        self.assertEqual([item['text'] for item in fields[0]['source_instructions']], instructions)
        changed = copy.deepcopy(fields); changed[0]['source_instructions'].pop()
        with self.assertRaises(w.V3Failure): w.fill(original, changed, {changed[0]['key']: 'A complete observation.'})

    def test_equal_header_labels_with_different_subjects_are_separate_answers(self):
        from docx.shared import Inches
        doc = Document(); doc.settings.odd_and_even_pages_header_footer = True
        for header, role in zip((doc.sections[0].header, doc.sections[0].even_page_header), ('Donor', 'Recipient')):
            header.paragraphs[0].text = role; header.paragraphs[0].style = 'Heading 1'
            table = header.add_table(rows=1, cols=2, width=Inches(6))
            table.cell(0, 0).text = 'Name'
            table.cell(0, 1).paragraphs[0].add_run('Enter the full name.').font.color.rgb = RGBColor.from_string('808080')
        original = data(doc); fields = w.inspect(original)['fields']
        self.assertEqual(len(fields), 2)
        self.assertEqual(len({field['context'] for field in fields}), 2)
        target = next(f for f in fields if 'Donor' in f['context'])
        final = w.fill(original, fields, {target['key']: 'SYNTHETIC DONOR'})
        before = w.open_package(original); after = w.open_package(final)
        self.assertEqual(before.zip.read('word/header2.xml'), after.zip.read('word/header2.xml'))

    def test_new_mapping_preserves_all_instruction_paragraphs_and_appends_full_answer(self):
        original = form(); fields = w.inspect(original)['fields']; field = fields[0]
        self.assertEqual(field['locator']['target']['type'], 'cell_append')
        self.assertEqual(len(field['source_instructions']), 1)
        self.assertIn('Do not share third-party information', field['source_instructions'][0]['text'])
        for answer in ('The patient reviewed the report.', ('A complete reviewed sentence. ' * 150).strip()):
            with self.subTest(long=len(answer) > 100):
                result = w.fill(original, fields, {field['key']: answer})
                before = w.SlotIndex(w.open_package(original)).by_id['word/document.xml#b/tbl[0]/tr[0]/tc[1]'].el
                after = w.SlotIndex(w.open_package(result)).by_id['word/document.xml#b/tbl[0]/tr[0]/tc[1]'].el
                self.assertEqual([etree.tostring(p) for p in before.findall(w.q('w:p'))],
                                 [etree.tostring(p) for p in after.findall(w.q('w:p'))][:-1])
                self.assertEqual(w.text_of(after.findall(w.q('w:p'))[-1]), answer)
                self.assertIsNone(w.effective_values(original, fields, {field['key']: answer})[field['key']])

    def test_old_whole_cell_mapping_is_not_rewritten_and_cannot_delete_printed_guidance(self):
        original = form(); fields = w.inspect(original)['fields']
        fields[0]['locator']['target'] = {'type': 'cell_replace', 'all': True}
        fields[0].pop('source_instructions', None)
        old = copy.deepcopy(fields); answer = 'The patient reviewed the report.'
        result = w.fill(original, fields, {fields[0]['key']: answer})
        text = w.text_of(w.SlotIndex(w.open_package(result)).by_id['word/document.xml#b/tbl[0]/tr[0]/tc[1]'].el)
        self.assertIn('Do not share third-party information with the patient.', text)
        self.assertIn(answer, text)
        self.assertEqual(fields, old)
        self.assertIsNone(w.effective_values(original, fields, {fields[0]['key']: answer})[fields[0]['key']])

    def test_instruction_identity_is_checked_and_unwritten_constrained_state_is_unknown(self):
        original = form(); fields = w.inspect(original)['fields']; key = fields[0]['key']
        self.assertIsNone(w.effective_values(original, fields, {})[key])
        fields[0]['source_instructions'][0]['text'] = 'The patient has already consented.'
        with self.assertRaises(w.V3Failure) as failure:
            w.fill(original, fields, {key: 'The patient reviewed the report.'})
        self.assertEqual(failure.exception.code, 'mapping_mismatch')

    def test_preserved_coloured_values_are_not_erased_from_constrained_state(self):
        for color in ('808080', '0000FF'):
            for option in (False, True):
                doc = Document(); table = doc.add_table(rows=2 if option else 1, cols=2)
                if option:
                    table.cell(0, 0).text = 'Question'; table.cell(0, 1).text = 'Yes / No'
                row = 1 if option else 0
                table.cell(row, 0).text = 'Do you smoke?' if option else 'Priority'
                table.cell(row, 1).paragraphs[0].add_run('Yes' if option else '1').font.color.rgb = RGBColor.from_string(color)
                original = data(doc); inspected = w.inspect(original)['fields']
                self.assertEqual(len(inspected), 1)
                for legacy in (False, True):
                    with self.subTest(color=color, option=option, legacy=legacy):
                        fields = copy.deepcopy(inspected); field = fields[0]
                        if legacy:
                            field['locator']['target'].update(type='cell_replace', all=True)
                            field.pop('source_instructions')
                        expected = [None, None] if option else None
                        for answers in ({}, {field['key']: 'No' if option else '2'}):
                            self.assertEqual(w.effective_values(original, fields, answers)[field['key']], expected)
                        output = w.fill(original, fields, {field['key']: 'No' if option else '2'})
                        cell = w.SlotIndex(w.open_package(output)).by_id['word/document.xml#b/tbl[0]/tr[%d]/tc[1]' % row].el
                        self.assertEqual([w.text_of(p) for p in cell.findall(w.q('w:p'))], ['Yes', 'No'] if option else ['1', '2'])

    def test_complete_instruction_bound_survives_sandbox_json_and_export(self):
        from documents import inspect_template, fill_template, DocumentError
        text = 'Preserved guidance: ' + 'x' * (w.MAX_SOURCE_INSTRUCTION_CHARS - len('Preserved guidance: ') - 1) + 'Ω'
        doc = Document(); table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'What observations were made?'
        table.cell(0, 1).paragraphs[0].add_run(text).font.color.rgb = RGBColor.from_string('808080')
        original = data(doc)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'original.docx'; path.write_bytes(original)
            mapped = inspect_template(path, engine='v3')
            fields = json.loads(json.dumps(mapped['fields']))
            self.assertEqual(fields[0]['source_instructions'][0]['text'], text)
            self.assertEqual(len(text), 20000)
            export = Path(directory) / 'export.docx'; answer = 'The observed finding is recorded in full.'
            fill_template(path, fields, {fields[0]['key']: answer}, export)
            cell = w.SlotIndex(w.open_package(export.read_bytes())).by_id['word/document.xml#b/tbl[0]/tr[0]/tc[1]'].el
            self.assertEqual([w.text_of(p) for p in cell.findall(w.q('w:p'))], [text, answer])
            color = cell.findall(w.q('w:p'))[-1].find('.//' + w.q('w:color'))
            self.assertEqual(color.get(w.q('w:val')), '000000')
            self.assertEqual(path.read_bytes(), original)
            table.cell(0, 1).paragraphs[0].runs[0].text += 'x'
            path.write_bytes(data(doc))
            with self.assertRaises(DocumentError) as failure:
                inspect_template(path, engine='v3')
            self.assertEqual(failure.exception.code, 'mapping_limit')

    def test_missing_or_changed_instruction_anchor_refuses_export(self):
        original = form(); first = w.inspect(original)['fields']
        for mutate in (lambda f: f.pop('source_instructions'),
                       lambda f: f['source_instructions'][0].update(slot='b/p[0]'),
                       lambda f: f['source_instructions'][0].update(fingerprint='f' * 64)):
            fields = copy.deepcopy(first); mutate(fields[0])
            with self.assertRaises(w.V3Failure) as failure:
                w.fill(original, fields, {fields[0]['key']: 'A complete observation.'})
            self.assertEqual(failure.exception.code, 'mapping_mismatch')
