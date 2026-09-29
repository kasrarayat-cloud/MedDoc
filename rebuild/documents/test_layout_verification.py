"""Source/transport/PDF refutations; synthetic boundaries are never renderer evidence."""
import copy
import hashlib
import io
import json
import os
from pathlib import Path
import socket
import struct
import tempfile
import threading
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from pypdf import PdfWriter
from pypdf.generic import ArrayObject, DictionaryObject, FloatObject, NameObject, NumberObject, DecodedStreamObject

from . import layout_source as source, word_v3 as w
from .layout_pdf_worker import verify
from .layout_transport import receive, send, verify_document
from .service import DocumentError
from .test_question_boxes import data


def tab_form(**kwargs):
    from .test_question_boxes import TabColumnOwnership
    return TabColumnOwnership().form(**kwargs)


def synthetic_verified(raw, mapping=None, answers=None, readings=None):
    """Mock only the external renderer in application contract tests. (`readings`: the parent's syntax readings the
    real verifier forwards to its isolated preparer; in-process the same pinned parser reads the same clauses.)"""
    prepared, actual, _ = source.prepare(raw, mapping=mapping, answers=answers)
    fields = prepared['fields']
    for field in fields:
        if field.get('layout_candidate'):
            field['layout_verified'] = {'synthetic_test_only': True}
    return {'fields': fields, 'warnings': prepared['warnings'], 'receipt': {'synthetic_test_only': True},
            'states': prepared['states']}, actual


class LayoutSource(unittest.TestCase):
    def setUp(self):
        self.raw = data(tab_form())

    def test_tab_stops_only_propose_and_json_receipts_cannot_authorize_writes(self):
        inspected = w.inspect(self.raw)
        self.assertEqual(inspected['layout_pending'], 1)
        self.assertFalse(inspected['fields'])
        self.assertTrue(any('layout' in note.lower() for note in inspected['notes']))
        proposal, actual, instrumented = source.prepare(self.raw)
        field = proposal['fields'][0]
        field['layout_verified'] = {'verified': True, 'source_sha256': source.digest(self.raw)}
        for fields in (proposal['fields'], []):
            with self.assertRaises(w.V3Failure) as caught:
                w.fill(self.raw, fields, {field['key']: 'Normal'})
            self.assertEqual(caught.exception.code, 'layout_verification_required')
        self.assertEqual(actual, self.raw)
        self.assertEqual(instrumented, self.raw)
        self.assertEqual(proposal['manifest']['contract'], 'word-tab-containment-v2')
        with ZipFile(io.BytesIO(self.raw)) as before, ZipFile(io.BytesIO(w.fill(self.raw, [field], {}))) as after:
            self.assertEqual({n: before.read(n) for n in before.namelist()},
                             {n: after.read(n) for n in after.namelist()})

    def test_mapping_changes_and_partial_constraint_reads(self):
        prepared, _, _ = source.prepare(self.raw)
        for key, value in [('options', ['Abnormal', 'Normal']), ('label', 'Other question'), ('context', 'Another subject')]:
            changed = copy.deepcopy(prepared['fields']); changed[0][key] = value
            with self.assertRaises(w.V3Failure):
                source.prepare(self.raw, mapping=changed, answers={changed[0]['key']: 'Normal'})
        field = prepared['fields'][0]
        with self.assertRaises(w.V3Failure) as caught:
            w.effective_values(self.raw, [field], {})
        self.assertEqual(caught.exception.code, 'layout_verification_required')
        final, actual, _ = source.prepare(self.raw, mapping=[field], answers={field['key']: 'Normal'})
        self.assertEqual(final['states'][field['key']], [True, False])
        self.assertEqual(w.option_states(w.SlotIndex(w.open_package(actual)), field), [True, False])
        # Word's original container can carry metadata that a no-op ZIP rewrite
        # changes even while every part stays identical.
        container = io.BytesIO(self.raw)
        with ZipFile(container, 'a') as archive: archive.comment = b'Original package metadata'
        raw = container.getvalue()
        self.assertEqual(w.fill(raw, [field], {field['key']:''}), raw)
        self.assertEqual(source.prepare(raw, mapping=[field], answers={field['key']:''})[1], raw)

    def test_instrumentation_preserves_output_and_all_package_parts(self):
        for control in ('glyph', 'sdt', 'ff'):
            with self.subTest(control=control):
                raw = data(tab_form(control=control))
                prepared, actual, marked = source.prepare(raw)
                self.assertEqual(actual, raw)
                names = [a['name'] for a in prepared['manifest']['anchors']]
                self.assertEqual(names, [])
                self.assertEqual(marked, raw)
                with ZipFile(io.BytesIO(raw)) as before, ZipFile(io.BytesIO(marked)) as after:
                    self.assertEqual(before.namelist(), after.namelist())
                    for name in before.namelist():
                        if name != 'word/document.xml': self.assertEqual(before.read(name), after.read(name))

    def test_layout_preparation_retains_unrelated_guidance_and_patient_authority(self):
        from docx.shared import RGBColor
        doc = tab_form()
        doc.add_paragraph('For the patient:')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'What progress have you made?'
        instruction = 'Describe your own experience; do not sign this answer for another person.'
        table.cell(0, 1).paragraphs[0].add_run(instruction).font.color.rgb = RGBColor.from_string('808080')
        raw = data(doc)
        ordinary = next(f for f in w.inspect(raw)['fields'] if f['locator']['target']['type'] == 'cell_append')
        prepared, _, _ = source.prepare(raw)
        field = next(f for f in prepared['fields'] if f['locator'] == ordinary['locator'])
        self.assertEqual(field['source_instructions'], ordinary['source_instructions'])
        self.assertEqual(field['answer_policy'], ordinary['answer_policy'])
        self.assertEqual(field['answer_policy']['source_author'], 'patient')
        answer = 'The patient reports weekly supervised contact.'
        _, exported, _ = source.prepare(raw, mapping=prepared['fields'], answers={field['key']: answer})
        cell = w.SlotIndex(w.open_package(exported)).by_id['word/document.xml#b/tbl[1]/tr[0]/tc[1]'].el
        self.assertEqual([w.text_of(p) for p in cell.findall(w.q('w:p'))], [instruction, answer])


class LayoutPdf(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name)
        self.anchors = [{'name': 'QMD' + str(i), 'candidate': 'row1', 'ordinal': 0,
                         'role': role, 'edge': edge} for i, (role, edge) in enumerate(
                             [('header', 'start'), ('header', 'end'), ('row', 'start'), ('row', 'end')])]
        self.manifest = {'anchors': self.anchors, 'candidates': [{'id': 'row1', 'source': {'options': ['Yes']}}]}
        (self.path / 'manifest.json').write_text(json.dumps(self.manifest))

    def pdf(self, name, *, marked=False, shift=0, missing=False, duplicate=False, graphics=False, outside=False):
        writer = PdfWriter(); page = writer.add_blank_page(width=400, height=400)
        stream = DecodedStreamObject(); stream.set_data(b'0 0 m 10 10 l S' if graphics else b'')
        page[NameObject('/Contents')] = writer._add_object(stream)
        if marked:
            pairs = []
            for i, anchor in enumerate(self.anchors):
                if missing and i == 3: continue
                x = [100, 140, 110 + shift, 130 + shift][i]
                if outside and i == 0: x = -1
                dest = ArrayObject([page.indirect_reference, NameObject('/XYZ'), FloatObject(x),
                                    FloatObject(300 if i < 2 else 270), NumberObject(0)])
                pairs.extend([NameObject('/' + anchor['name']), dest])
            if duplicate: pairs.extend(pairs[:2])
            writer._root_object[NameObject('/Names')] = DictionaryObject({NameObject('/Dests'):
                DictionaryObject({NameObject('/Names'): ArrayObject(pairs)})})
        writer.write(self.path / name)

    def test_complete_aligned_spans_pass_and_wrong_column_fails(self):
        self.pdf('document.pdf'); self.pdf('instrumented.pdf', marked=True)
        self.assertTrue(verify(self.path)['passed'])
        self.pdf('instrumented.pdf', marked=True, shift=40)
        self.assertFalse(verify(self.path)['passed'])

    def test_missing_duplicate_outside_and_changed_graphics_refuse(self):
        self.pdf('document.pdf')
        for change in ('missing', 'duplicate', 'outside', 'graphics'):
            with self.subTest(change=change):
                self.pdf('instrumented.pdf', marked=True, **{change: True})
                with self.assertRaises(ValueError): verify(self.path)


class LayoutTransport(unittest.TestCase):
    def test_frame_bounds_and_roundtrip(self):
        left, right = socket.socketpair()
        with left, right:
            send(left, {'checked': [1, 2]}); self.assertEqual(receive(right), {'checked': [1, 2]})
            left.sendall(struct.pack('!I', 48 * 1024 * 1024 + 1))
            with self.assertRaises(ValueError): receive(right)

    def test_unavailable_service_and_untrusted_socket_owner_refuse(self):
        with patch('django.conf.settings', SimpleNamespace(QUILLMD_LAYOUT_SOCKET='', QUILLMD_LAYOUT_SERVICE_UID='')):
            with self.assertRaises(DocumentError): verify_document(b'original')
        with tempfile.TemporaryDirectory() as temp:
            path = str(Path(temp) / 'service.sock')
            with socket.socket(socket.AF_UNIX) as server:
                server.bind(path)
                with patch('django.conf.settings', SimpleNamespace(QUILLMD_LAYOUT_SOCKET=path,
                        QUILLMD_LAYOUT_SERVICE_UID=os.getuid() + 1)):
                    with self.assertRaises(DocumentError): verify_document(b'original')

    def test_live_peer_cannot_return_a_receipt_for_another_original(self):
        with tempfile.TemporaryDirectory() as temp:
            path = str(Path(temp) / 'service.sock')
            with socket.socket(socket.AF_UNIX) as server:
                server.bind(path); server.listen(1)
                def answer():
                    connection, _ = server.accept()
                    with connection:
                        receive(connection)
                        send(connection, {'verified': True, 'receipt': {'source_sha256': 'wrong-source'}})
                thread = threading.Thread(target=answer); thread.start()
                try:
                    with patch('django.conf.settings', SimpleNamespace(QUILLMD_LAYOUT_SOCKET=path,
                            QUILLMD_LAYOUT_SERVICE_UID=os.getuid())):
                        with self.assertRaises(DocumentError): verify_document(b'original')
                finally: thread.join(5)
                self.assertFalse(thread.is_alive())


class LayoutHeldProposals(unittest.TestCase):
    """The service publishes only measured destinations and names every individually held proposal."""

    def setUp(self):
        from .layout_server import retain_verified
        self.retain = retain_verified
        row = lambda n: {'part': 'word/document.xml', 'slot': 'b/tbl[0]/tr[0]/tc[0]/p[%d]' % n}
        self.manifest = {'candidates': [
            {'id': 'required', 'source': {'row': row(1), 'label': 'Hair', 'context': 'Head'}},
            {'id': 'held', 'optional': True, 'untouched': True,
             'source': {'row': row(3), 'label': 'Supple', 'context': 'Neck'}},
            {'id': 'measured', 'optional': True, 'untouched': True,
             'source': {'row': row(5), 'label': 'Masses', 'context': 'Neck'}}]}
        self.fields = [{'key': 'v3_1', 'layout_candidate': 'required'}, {'key': 'v3_2', 'layout_candidate': 'held'},
                       {'key': 'v3_3', 'layout_candidate': 'measured'}, {'key': 'v3_4'}]

    def receipt(self, identity):
        return {'candidate': identity}

    def test_held_proposals_are_removed_and_reported_with_their_printed_source(self):
        fields, warnings = self.retain(copy.deepcopy(self.fields), ['existing'], self.manifest,
                                       {'held': [{'candidate': 'held', 'reason': 'outside'}]}, self.receipt)
        self.assertEqual([f['key'] for f in fields], ['v3_1', 'v3_3', 'v3_4'])
        self.assertEqual([f.get('layout_verified') for f in fields],
                         [{'candidate': 'required'}, {'candidate': 'measured'}, None])
        self.assertEqual(warnings[0], 'existing')
        self.assertEqual(len(warnings), 2)
        self.assertIn('b/tbl[0]/tr[0]/tc[0]/p[3]', warnings[1]); self.assertIn('Supple', warnings[1])
        self.assertIn('Neck', warnings[1])

    def test_nothing_else_can_be_held(self):
        for held in ([{'candidate': 'required'}], [{'candidate': 'unknown'}], [{'candidate': 'held'}, {'candidate': 'held'}],
                     [{}], ['held']):
            with self.subTest(held=held), self.assertRaises(ValueError):
                self.retain(copy.deepcopy(self.fields), [], self.manifest, {'held': held}, self.receipt)
        touched = copy.deepcopy(self.manifest); touched['candidates'][1]['untouched'] = False
        with self.assertRaises(ValueError):
            self.retain(copy.deepcopy(self.fields), [], touched, {'held': [{'candidate': 'held'}]}, self.receipt)
