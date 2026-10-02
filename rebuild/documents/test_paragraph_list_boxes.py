"""Synthetic question/list/blank-box ownership; no corpus or model qualification."""
import copy
import io
import json
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from . import word_v3 as w
from .test_document_relations import link, numbered
from .test_question_boxes import box, data

QUESTION = 'What type of residence does the patient occupy:'
VALUES = ('Own home', 'With family', 'Other (please specify)')


def form(*, kind='bullet', question=QUESTION, values=VALUES):
    doc = Document(); doc.add_heading('Accommodation', 1); doc.add_paragraph(question)
    numbered(doc, values, kind)
    box(doc)
    return doc


def rewrite(raw, part, change):
    out = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as before, ZipFile(out, 'w', ZIP_DEFLATED) as after:
        for name in before.namelist():
            value = before.read(name)
            after.writestr(name, change(value) if name == part else value)
    return out.getvalue()


class ParagraphListBoxes(unittest.TestCase):
    def owned(self, raw):
        result = w.inspect(raw)
        fields = [f for f in result['fields'] if f['locator']['target']['type'] == 'cell']
        self.assertEqual(len(fields), 1)
        field = fields[0]
        self.assertEqual(field['label'], QUESTION)
        self.assertEqual(field['kind'], 'long_text')
        self.assertNotIn('options', field)
        self.assertEqual(field['source_questions'][0]['cardinality'], 'unspecified')
        self.assertEqual(result['layout_pending'], 0)
        return field

    def test_whole_question_owns_box_instead_of_last_other_descriptor(self):
        self.owned(data(form()))

    def test_no_other_item_needed_and_lettered_descriptors_do_not_become_choices(self):
        for kind in ('bullet', 'lowerLetter', 'decimal'):
            with self.subTest(kind=kind):
                raw = data(form(kind=kind, values=VALUES[:2])); field = self.owned(raw)
                self.assertEqual(len(w.inspect(raw)['fields']), 1)
                self.assertEqual([x['text'] for x in field['source_questions'][0]['descriptors']], list(VALUES[:2]))

    def test_blank_single_and_combined_answer_change_only_answer_cell(self):
        doc = form(); doc.paragraphs[2].runs[0].bold = True
        link(doc.paragraphs[3], ' guidance Ω')
        original = data(doc); field = self.owned(original); fields = [field]
        original_mapping = copy.deepcopy(fields)
        for answer in ('', 'The patient resides in a Care Home.', 'The patient lives in their own home with family.'):
            with self.subTest(answer=answer):
                filled = w.fill(original, fields, {field['key']: answer})
                before = w.SlotIndex(w.open_package(original)); after = w.SlotIndex(w.open_package(filled))
                key = field['locator']['part'] + '#' + field['locator']['slot']
                self.assertEqual(w.text_of(after.by_id[key].el), answer)
                # Replace only the filled cell with its original; every body node,
                # property, numbered bullet, relationship and other part must agree.
                replacement = after.by_id[key].el
                replacement.getparent().replace(replacement, copy.deepcopy(before.by_id[key].el))
                self.assertEqual(etree.tostring(before.package.tree('word/document.xml')),
                                 etree.tostring(after.package.tree('word/document.xml')))
                with ZipFile(io.BytesIO(original)) as a, ZipFile(io.BytesIO(filled)) as b:
                    self.assertEqual(a.namelist(), b.namelist())
                    for name in a.namelist():
                        if name != 'word/document.xml': self.assertEqual(a.read(name), b.read(name), name)
                self.assertEqual(fields, original_mapping)

    def test_numbering_instance_and_page_break_do_not_split_one_question(self):
        doc = form()
        numbering = doc.part.numbering_part.element
        existing = next(n for n in numbering.findall(qn('w:num')) if n.get(qn('w:numId')) == '10')
        other = copy.deepcopy(existing); other.set(qn('w:numId'), '11'); numbering.append(other)
        item = doc.paragraphs[3]
        item._p.find('.//' + qn('w:numId')).set(qn('w:val'), '11')
        page = OxmlElement('w:br'); page.set(qn('w:type'), 'page'); item.runs[0]._r.insert(0, page)
        field = self.owned(data(doc))
        self.assertEqual([x['numbering'][0] for x in field['source_questions'][0]['descriptors']], ['10', '11', '10'])

    def test_complete_relation_survives_json_and_refuses_changed_source_or_definitions(self):
        doc = form(); link(doc.paragraphs[3]); original = data(doc)
        field = json.loads(json.dumps(self.owned(original)))
        for part, old, new in (
                ('word/document.xml', b'Own home', b'Care Home'),
                ('word/document.xml', b'What type', b'Which type'),
                ('word/document.xml', b'Accommodation', b'Signed consent'),
                ('word/numbering.xml', b'w:val="bullet"', b'w:val="decimal"'),
                ('word/_rels/document.xml.rels', b'https://example.test/guidance', b'https://example.test/different')):
            with self.subTest(part=part, old=old):
                changed = rewrite(original, part, lambda value: value.replace(old, new))
                self.assertNotEqual(changed, original)
                with self.assertRaises(w.V3Failure) as failure:
                    w.fill(changed, [field], {field['key']: 'The patient lives with family.'})
                self.assertEqual(failure.exception.code, 'mapping_mismatch')

    def test_relation_rejects_incomplete_forged_source_and_changed_physical_owner(self):
        original = data(form()); field = self.owned(original)
        for change in (lambda r: r['descriptors'].pop(), lambda r: r.update(cardinality='one'),
                       lambda r: r['descriptors'][0].update(text='Invented descriptor'),
                       lambda r: r['definitions'].clear()):
            changed = copy.deepcopy(field); change(changed['source_questions'][0])
            with self.assertRaises(w.V3Failure) as failure:
                w.fill(original, [changed], {field['key']: 'Explicit synthetic answer.'})
            self.assertEqual(failure.exception.code, 'mapping_mismatch')
        doc = form(); new_question = doc.add_paragraph('What unrelated event occurred?')
        doc.tables[0]._tbl.addprevious(new_question._p)
        with self.assertRaises(w.V3Failure) as failure:
            w.fill(data(doc), [field], {field['key']: 'Explicit synthetic answer.'})
        self.assertEqual(failure.exception.code, 'mapping_mismatch')

    def test_explicit_selection_controls_are_not_reclassified_as_free_text(self):
        for instruction in ('Choose one residence:', 'Select all that apply:', 'Which residence (one only):'):
            with self.subTest(instruction=instruction):
                raw = data(form(kind='lowerLetter', question=instruction))
                fields = w.inspect(raw)['fields']
                self.assertFalse([f for f in fields if f.get('source_questions')])
                choice = next(f for f in fields if f['locator']['target']['type'] == 'paragraph_options')
                self.assertEqual(choice['options'], list(VALUES))

    def test_boundaries_and_uncertain_source_never_claim_the_list_box(self):
        for case in ('heading', 'section', 'other_question', 'line_answer', 'hidden', 'control',
                     'unknown_numbering', 'mixed_depth', 'unbordered', 'merged', 'filled', 'too_long'):
            with self.subTest(case=case):
                doc = form()
                if case in ('heading', 'other_question'):
                    p = doc.add_heading('Separate section', 1) if case == 'heading' else doc.add_paragraph('What unrelated event occurred?')
                    doc.tables[0]._tbl.addprevious(p._p)
                elif case == 'section':
                    doc.paragraphs[3]._p.get_or_add_pPr().append(OxmlElement('w:sectPr'))
                elif case == 'line_answer': doc.paragraphs[3].add_run(' ______')
                elif case == 'hidden': doc.paragraphs[3].runs[0].font.hidden = True
                elif case == 'control': doc.paragraphs[3].add_run(' ☐')
                elif case == 'unknown_numbering': doc.paragraphs[3]._p.find('.//' + qn('w:numId')).set(qn('w:val'), '999')
                elif case == 'mixed_depth':
                    depth = OxmlElement('w:ilvl'); depth.set(qn('w:val'), '1'); doc.paragraphs[3]._p.find('.//' + qn('w:numPr')).append(depth)
                elif case == 'unbordered':
                    tc = doc.tables[0].cell(0, 0)._tc.tcPr; tc.remove(tc.find(qn('w:tcBorders')))
                elif case == 'merged':
                    span = OxmlElement('w:gridSpan'); span.set(qn('w:val'), '2'); doc.tables[0].cell(0, 0)._tc.tcPr.append(span)
                elif case == 'filled': doc.tables[0].cell(0, 0).text = 'Existing answer.'
                elif case == 'too_long': doc.paragraphs[3].text = 'x' * (w.MAX_MAPPING_CONTEXT + 1)
                self.assertFalse([f for f in w.inspect(data(doc))['fields']
                                  if any(r.get('relation') == 'paragraph_list_answer' for r in f.get('source_questions', []))])

    def test_legacy_other_box_mapping_is_consumed_without_rewriting(self):
        original = data(form()); field = self.owned(original)
        field.pop('source_questions'); field['label'] = 'Other (please specify)'
        field['context'] = field['context'].replace('Printed descriptors:', 'Printed options:')
        old = copy.deepcopy(field); answer = 'The patient lives with family.'
        output = w.fill(original, [field], {field['key']: answer})
        self.assertEqual(field, old)
        self.assertEqual(Document(io.BytesIO(output)).tables[0].cell(0, 0).text, answer)

    def test_destination_hidden_directly_or_by_style_is_not_an_answer_field(self):
        for case in ('direct', 'paragraph', 'inherited', 'table'):
            with self.subTest(case=case):
                doc = form(); paragraph = doc.tables[0].cell(0, 0).paragraphs[0]
                if case == 'direct':
                    paragraph.add_run('').font.hidden = True
                else:
                    kind = WD_STYLE_TYPE.TABLE if case == 'table' else WD_STYLE_TYPE.PARAGRAPH
                    style = doc.styles.add_style('HiddenAnswer', kind); style.font.hidden = True
                    if case == 'inherited':
                        child = doc.styles.add_style('AnswerChild', kind); child.base_style = style; style = child
                    if case == 'table': doc.tables[0].style = style
                    else: paragraph.style = style
                self.assertFalse(w.inspect(data(doc))['fields'])

    def test_visible_destination_and_later_style_change_are_rechecked_at_export(self):
        doc = form(); style = doc.styles.add_style('HiddenAnswer', WD_STYLE_TYPE.PARAGRAPH)
        style.font.hidden = False; doc.tables[0].cell(0, 0).paragraphs[0].style = style
        raw = data(doc); field = self.owned(raw); answer = 'The patient lives with family.'
        output = w.fill(raw, [field], {field['key']: answer})
        index = w.SlotIndex(w.open_package(output)); detector = w.Detector(index.package)
        cell = index.by_id[field['locator']['part'] + '#' + field['locator']['slot']]
        self.assertEqual(w.text_of(cell.el), answer)
        self.assertTrue(all(detector.tab_source_visible(p) for p in cell.children if p.kind == 'p'))
        changed = rewrite(raw, 'word/styles.xml', lambda b: b.replace(b'<w:vanish w:val="0"/>', b'<w:vanish/>'))
        self.assertNotEqual(changed, raw)
        with self.assertRaises(w.V3Failure) as failure:
            w.fill(changed, [field], {field['key']: answer})
        self.assertEqual(failure.exception.code, 'mapping_mismatch')

    def test_preceding_selection_or_other_only_instruction_does_not_own_whole_answer(self):
        for text in ('Tick one residence below. Use the box only to specify Other.',
                     'Use the box only to specify Other.', 'Use the box for Other only.',
                     'Choose all that apply below.'):
            with self.subTest(instruction=text):
                doc = form(); p = doc.add_paragraph(text); doc.paragraphs[1]._p.addprevious(p._p)
                self.assertFalse([f for f in w.inspect(data(doc))['fields']
                                  if any(r.get('relation') == 'paragraph_list_answer' for r in f.get('source_questions', []))])

    def test_complete_local_instruction_is_preserved_and_bound_at_export(self):
        doc = form(); instruction = 'Use clinical observations recorded during this period.'
        p = doc.add_paragraph(instruction); doc.paragraphs[1]._p.addprevious(p._p)
        raw = data(doc); field = self.owned(raw); relation = field['source_questions'][0]
        self.assertIn(instruction, field['context'])
        self.assertEqual([i['text'] for i in relation['instructions']], [instruction])
        answer = 'The patient lives with family.'
        self.assertEqual(Document(io.BytesIO(w.fill(raw, [field], {field['key']: answer}))).tables[0].cell(0, 0).text, answer)
        changed = rewrite(raw, 'word/document.xml', lambda b: b.replace(instruction.encode(), b'Use the box only to specify Other.'))
        with self.assertRaises(w.V3Failure) as failure:
            w.fill(changed, [field], {field['key']: answer})
        self.assertEqual(failure.exception.code, 'mapping_mismatch')
        incomplete = copy.deepcopy(field); incomplete['source_questions'][0].pop('instructions')
        with self.assertRaises(w.V3Failure) as failure:
            w.fill(raw, [incomplete], {field['key']: answer})
        self.assertEqual(failure.exception.code, 'mapping_mismatch')

    def test_numbering_checkbox_markers_including_override_and_symbol_font_refuse(self):
        for case in ('literal', 'checked', 'wingdings', 'override'):
            with self.subTest(marker=case):
                doc = form(); numbering = doc.part.numbering_part.element
                level = next(a for a in numbering.findall(qn('w:abstractNum'))
                             if a.get(qn('w:abstractNumId')) == '90').find(qn('w:lvl'))
                if case == 'override':
                    number = next(n for n in numbering.findall(qn('w:num')) if n.get(qn('w:numId')) == '10')
                    override = OxmlElement('w:lvlOverride'); override.set(qn('w:ilvl'), '0')
                    level = copy.deepcopy(level); override.append(level); number.append(override)
                marker = OxmlElement('w:lvlText'); marker.set(qn('w:val'), {'checked': '☒', 'wingdings': 'o'}.get(case, '☐'))
                level.append(marker)
                if case == 'wingdings':
                    props = OxmlElement('w:rPr'); fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'), 'Wingdings')
                    props.append(fonts); level.append(props)
                self.assertFalse([f for f in w.inspect(data(doc))['fields']
                                  if any(r.get('relation') == 'paragraph_list_answer' for r in f.get('source_questions', []))])

    def test_plain_courier_bullet_remains_a_descriptor(self):
        doc = form(); level = next(a for a in doc.part.numbering_part.element.findall(qn('w:abstractNum'))
                                  if a.get(qn('w:abstractNumId')) == '90').find(qn('w:lvl'))
        marker = OxmlElement('w:lvlText'); marker.set(qn('w:val'), 'o'); level.append(marker)
        props = OxmlElement('w:rPr'); fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'), 'Courier New')
        props.append(fonts); level.append(props)
        self.owned(data(doc))

    def test_unresolved_references_cannot_claim_a_complete_list_question(self):
        for reference in ('above question', 'previous answer', 'following section', 'as above', 'same as'):
            with self.subTest(reference=reference):
                raw = data(form(question='What residence applies in relation to the ' + reference + ':'))
                self.assertFalse(w.inspect(raw)['fields'])
