"""Same physical table columns appoint captions below ordinary/signing blanks."""
import io
import unittest
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches

from . import word_v3 as word


def form(*, conflict=False, merged=False):
    document = Document()
    table = document.add_table(rows=3 if conflict else 2, cols=5)
    widths = (2.0, .15, 2.0, .15, 1.2)
    for n, width in enumerate(widths):
        table.columns[n].width = Inches(width)
        for row in table.rows:row.cells[n].width = Inches(width)
    at = 1 if conflict else 0
    for n, text in ((0, 'Participant name'), (2, 'Signature'), (4, 'Date')):
        table.cell(at + 1, n).text = text
        if conflict:table.cell(0, n).text = 'Hospital number'
    if merged:table.cell(at, 0).merge(table.cell(at, 1))
    document.add_paragraph('Clinical findings: ______')
    output = io.BytesIO(); document.save(output)
    return output.getvalue(), at


def owned_fields(data):
    return [f for f in word.inspect(data)['fields']
            if any(r.get('relation') == 'caption_below_cell' for r in f.get('source_questions', []))]


class CaptionedBlankRows(unittest.TestCase):
    def test_same_columns_bind_name_signature_date_and_never_spacer_checkboxes(self):
        data, _ = form(); fields = owned_fields(data)
        self.assertEqual([(f['label'], f['kind'], f['locator']['slot']) for f in fields],
                         [('Participant name', 'text', 'b/tbl[0]/tr[0]/tc[0]'),
                          ('Signature', 'signature', 'b/tbl[0]/tr[0]/tc[2]'),
                          ('Date', 'date', 'b/tbl[0]/tr[0]/tc[4]')])
        catalog = word.inspect(data)['fields']
        self.assertFalse(any(f['kind'] == 'checkbox' for f in catalog))
        signing = next(f for f in fields if f['kind'] == 'signature')
        with self.assertRaises(word.V3Failure):
            word.fill(data, catalog, {signing['key']: 'Dr Synthetic'})

    def test_ordinary_name_date_and_clinical_writes_preserve_every_caption_and_package_part(self):
        data, _ = form(); catalog = word.inspect(data)['fields']
        answers = {f['key']: {'Participant name': 'Synthetic Person', 'Date': '2026-10-05',
                              'Clinical findings': 'No focal swelling.'}[f['label']]
                   for f in catalog if f['label'] != 'Signature'}
        output = word.fill(data, catalog, answers)
        original = word.SlotIndex(word.open_package(data)); final = word.SlotIndex(word.open_package(output))
        for n in range(5):
            identity = 'word/document.xml#b/tbl[0]/tr[1]/tc[%d]' % n
            self.assertEqual(word.fingerprint(original.by_id[identity].el), word.fingerprint(final.by_id[identity].el))
        for n in (1, 2, 3):
            identity = 'word/document.xml#b/tbl[0]/tr[0]/tc[%d]' % n
            self.assertEqual(word.fingerprint(original.by_id[identity].el), word.fingerprint(final.by_id[identity].el))
        with ZipFile(io.BytesIO(data)) as a, ZipFile(io.BytesIO(output)) as b:
            self.assertEqual(a.namelist(), b.namelist())
            for name in a.namelist():
                if name != 'word/document.xml':self.assertEqual(a.read(name), b.read(name))

    def test_competing_caption_or_inconsistent_columns_refuse_new_ownership(self):
        for options in ({'conflict': True}, {'merged': True}):
            with self.subTest(options=options):
                data, _ = form(**options)
                self.assertEqual(owned_fields(data), [])

    def test_old_or_forged_spacer_box_cannot_write_fixed_caption_spacing(self):
        data, _ = form(); index = word.SlotIndex(word.open_package(data))
        for path in ('b/tbl[0]/tr[0]/tc[1]', 'b/tbl[0]/tr[1]/tc[1]'):
            with self.subTest(path=path):
                slot = index.by_id['word/document.xml#' + path]
                field = {'key': 'old_box', 'kind': 'checkbox', 'label': 'Signature', 'context': '',
                         'options': ['Signature'], 'locator': {'contract': 'v3', 'part': slot.part,
                            'slot': slot.path, 'fingerprint': word.fingerprint(slot.el),
                            'target': {'type': 'checkboxes', 'boxes': [{'kind': 'cell', 'part': slot.part, 'slot': slot.path}]}}}
                with self.assertRaises(word.V3Failure):word.fill(data, [field], {'old_box': 'checked'})

    def test_normal_response_grid_keeps_original_choices(self):
        document = Document(); table = document.add_table(rows=3, cols=3)
        for n, text in enumerate(('Finding', 'Yes', 'No')):table.cell(0, n).text = text
        table.cell(1, 0).text = 'Swelling'; table.cell(2, 0).text = 'Tenderness'
        output = io.BytesIO(); document.save(output); data = output.getvalue()
        fields = word.inspect(data)['fields']
        self.assertEqual(owned_fields(data), [])
        self.assertEqual([f['kind'] for f in fields], ['choice', 'choice'])
        self.assertTrue(all(f['options'] == ['Yes', 'No'] for f in fields))

    def test_previous_block_caption_cannot_migrate_to_the_next_actor_block(self):
        data, _ = form(); document = Document(io.BytesIO(data)); table = document.tables[0]
        for _ in range(2):table.add_row()
        for row in table.rows[2:]:
            for n, cell in enumerate(row.cells):cell.width = table.columns[n].width
        for n, text in ((0, 'Researcher name'), (2, 'Signature'), (4, 'Date')):
            table.cell(3, n).text = text
        stream = io.BytesIO(); document.save(stream); fields = owned_fields(stream.getvalue())
        self.assertEqual([(f['label'], f['kind']) for f in fields],
                         [('Participant name', 'text'), ('Signature', 'signature'), ('Date', 'date'),
                          ('Researcher name', 'text'), ('Signature', 'signature'), ('Date', 'date')])
