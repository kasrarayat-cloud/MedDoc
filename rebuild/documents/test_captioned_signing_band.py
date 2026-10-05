"""Source-bound signing bands; pure XML/parser controls, no app resources."""
from copy import deepcopy
import io
import unittest

from docx import Document

from . import word_v3 as word


def source(captions='Signature of Recorder\tDate\tTime', *, same_line=False, gap=False):
    doc = Document(); doc.add_paragraph('Printed Name of Recorder:')
    p = doc.add_paragraph()
    p.add_run('\t\t').underline = True; p.add_run('   ')
    p.add_run('________'); p.add_run('   ')
    p.add_run('\t\t').underline = True
    if same_line:p.add_run('  ' + captions)
    else:
        if gap:doc.add_paragraph()
        doc.add_paragraph(captions)
    doc.add_paragraph('Clinical findings: ______')
    out = io.BytesIO(); doc.save(out)
    return out.getvalue()


def fields(data):
    return word.inspect(data)['fields']


class CaptionedSigningBand(unittest.TestCase):
    def test_following_and_same_line_signing_captions_hold_every_owned_fragment(self):
        for same_line in (False, True):
            with self.subTest(same_line=same_line):
                data = source(same_line=same_line)
                own = [f for f in fields(data) if f['locator']['slot'] == 'b/p[1]']
                self.assertGreaterEqual(len(own), 2)
                self.assertTrue(all(f.get('write_protected') for f in own))
                for field in own:
                    with self.assertRaises(word.V3Failure):
                        word.fill(data, [field], {field['key']: 'Dr Synthetic'})
                self.assertEqual(word.fill(data, fields(data), {}), data)

    def test_removed_protection_or_changed_kind_cannot_authorize_the_actual_band(self):
        data = source()
        field = next(f for f in fields(data) if f['locator']['slot'] == 'b/p[1]')
        changed = deepcopy(field); changed.update(kind='text', label='Ordinary clinical value')
        for key in ('write_protected', 'review_reason', 'source_questions'):
            changed.pop(key, None)
        with self.assertRaises(word.V3Failure):
            word.fill(data, [changed], {changed['key']: 'Dr Synthetic'})

    def test_unrelated_clinical_answer_keeps_its_own_role_and_untouched_band(self):
        data = source(); catalog = fields(data)
        field = next(f for f in catalog if f['label'] == 'Clinical findings')
        self.assertFalse(field.get('write_protected'))
        output = word.fill(data, catalog, {field['key']: 'No focal swelling.'})
        old = word.SlotIndex(word.open_package(data)); new = word.SlotIndex(word.open_package(output))
        for path in ('b/p[0]', 'b/p[1]', 'b/p[2]'):
            self.assertEqual(word.fingerprint(old.by_id['word/document.xml#' + path].el),
                             word.fingerprint(new.by_id['word/document.xml#' + path].el))

    def test_ordinary_date_and_identity_captions_do_not_become_signing_bands(self):
        for captions in ('Hospital number\tCollection date\tCollection time',
                         "Printed Name of Witness to participant's signature\tRelationship to participant"):
            with self.subTest(captions=captions):
                data = source(captions); detector = word.Detector(word.open_package(data))
                slot = detector.index.by_id['word/document.xml#b/p[1]']
                self.assertIsNone(detector.captioned_signing_band(slot))

    def test_foreign_nonadjacent_caption_does_not_supply_band_ownership(self):
        data = source(gap=True); detector = word.Detector(word.open_package(data))
        slot = detector.index.by_id['word/document.xml#b/p[1]']
        self.assertIsNone(detector.captioned_signing_band(slot))

    def test_actor_order_and_parenthetical_captions_do_not_release_the_band(self):
        for captions in ('Recorder signature\tDate\tTime', 'Signature (Recorder)\tDate\tTime',
                         'Participant signature\tDate', 'Witness initials\tDate', 'Name\tSign below\tDate'):
            with self.subTest(captions=captions):
                raw = source(captions); own = [f for f in fields(raw) if f['locator']['slot'] == 'b/p[1]']
                self.assertTrue(own and all(f.get('write_protected') for f in own))
                for field in own:
                    forged = deepcopy(field); forged.pop('write_protected', None); forged['kind'] = 'text'
                    with self.assertRaises(word.V3Failure):
                        word.fill(raw, [forged], {field['key']: 'Dr Synthetic'})
