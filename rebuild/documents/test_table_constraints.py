"""Offline source-bound constraints for newly prepared short table forms."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document

from .engine import DocumentFailure, answer_constraint_error, constrained_answer_text, docx_inspect, validate_answer_constraints
from .table_constraints import with_table_constraints


class TableConstraintTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def form(self, labels, columns=2):
        document = Document()
        table = document.add_table(rows=len(labels) if columns == 2 else 1, cols=columns)
        if columns == 2:
            for row, label in zip(table.rows, labels):
                row.cells[0].text = label
        else:
            for index, label in enumerate(labels):
                table.rows[0].cells[index * 2].text = label
        path = self.root / 'form.docx'
        document.save(path)
        fields = docx_inspect(path)[3]
        return path, fields

    def test_actual_temperature_and_choice_cells_reject_mismatched_answers(self):
        path, original = self.form(['Temperature (°C)', 'Fever (Yes / No)'])
        frozen = deepcopy(original)
        fields = with_table_constraints(path, original)
        self.assertEqual(original, frozen)
        self.assertEqual([field['answer_constraints'] for field in fields], [
            {'version': 1, 'data_type': 'number', 'unit': '°C'},
            {'version': 1, 'data_type': 'choice', 'options': ['Yes', 'No']}])
        for bad in ['98.6 °F', 'normal', '37.2 kg']:
            answers = {fields[0]['key']: bad, fields[1]['key']: 'No'}
            with self.subTest(bad=bad), self.assertRaises(DocumentFailure) as refused:
                validate_answer_constraints(fields, answers)
            self.assertEqual(refused.exception.code, 'invalid_answer_constraint')
        for bad in ['Sometimes', 'Yes / No']:
            self.assertIsNotNone(answer_constraint_error(fields[1], bad))
        valid = {fields[0]['key']: '37.2 °C', fields[1]['key']: 'No'}
        validate_answer_constraints(fields, valid)
        self.assertEqual(constrained_answer_text(fields[0], valid[fields[0]['key']]), '37.2')
        self.assertEqual(with_table_constraints(path, fields), fields)

    def test_only_printed_date_and_numeric_type_declarations_are_enforced(self):
        path, fields = self.form(['Visit date (YYYY-MM-DD)', 'Count (integer)', 'Result [number]'])
        fields = with_table_constraints(path, fields)
        self.assertEqual([field['answer_constraints'] for field in fields], [
            {'version': 1, 'data_type': 'date', 'format': 'YYYY-MM-DD'},
            {'version': 1, 'data_type': 'integer'}, {'version': 1, 'data_type': 'number'}])
        for field, bad in zip(fields, ['2026-02-30', '2.5', 'normal']):
            self.assertIsNotNone(answer_constraint_error(field, bad))
        for field, good in zip(fields, ['2026-10-09', '-2', '1.5']):
            self.assertIsNone(answer_constraint_error(field, good))

    def test_nonterminal_and_square_bracket_declarations_reject_invalid_answers(self):
        cases = [
            ('Temperature (°C), record fasting result', '98.6 °F', '37.2 °C'),
            ('Fever (Yes / No), select one', 'Maybe', 'No'),
            ('Visit date (YYYY-MM-DD), enter visit', '2026-02-30', '2026-10-09'),
            ('Count (integer), include total', '2.5', '2'),
            ('Temperature [°C]', '98.6 °F', '37.2 °C'),
            ('Fever [Yes / No], select one', 'Maybe', 'Yes'),
            ('Visit date [YYYY-MM-DD], enter visit', '2026-02-30', '2026-10-09'),
            ('Count [integer], include total', '2.5', '-2'),
        ]
        path, original = self.form([label for label, _, _ in cases])
        original_bytes = path.read_bytes()
        derived = with_table_constraints(path, original)
        for field, (label, bad, good) in zip(derived, cases):
            with self.subTest(label=label):
                self.assertEqual(field['label'], label)
                self.assertIsNotNone(answer_constraint_error(field, bad))
                self.assertIsNone(answer_constraint_error(field, good))
        self.assertEqual(path.read_bytes(), original_bytes)

    def test_conflicting_nonterminal_declarations_remain_uncertain(self):
        labels = ['Temperature (°C), record result °F', 'Temperature [°C] [°F]',
                  'Count (integer), result (number)', 'Visit (YYYY-MM-DD) [DD/MM/YYYY]',
                  'Assessment, record mg']
        path, fields = self.form(labels)
        self.assertTrue(all('answer_constraints' not in field
                            for field in with_table_constraints(path, fields)))

    def test_ambiguous_unknown_and_plural_declarations_do_not_acquire_constraints(self):
        labels = ['Date of birth', 'Location (Hand / Foot)', 'Weight (kg or lb)',
                  'Medication(s)', 'Elapsed (s)', 'Dose (mg/kg)', 'Weight (kg) (g)']
        path, fields = self.form(labels)
        derived = with_table_constraints(path, fields)
        self.assertTrue(all('answer_constraints' not in field for field in derived))

    def test_multiple_qualified_pairs_bind_each_constraint_to_its_own_question(self):
        path, fields = self.form(['Weight (kg)', 'Fever (Yes / No)'], columns=4)
        result = with_table_constraints(path, fields)
        self.assertEqual(result[0]['locator'], {'cell': 0})
        self.assertEqual(result[1]['locator'], {'cell': 1})
        self.assertEqual(result[0]['answer_constraints']['unit'], 'kg')
        self.assertEqual(result[1]['answer_constraints']['options'], ['Yes', 'No'])
        self.assertIsNotNone(answer_constraint_error(result[0], 'No'))
        self.assertIsNotNone(answer_constraint_error(result[1], '70 kg'))

    def test_unqualified_row_cannot_borrow_a_unit_or_answer_association(self):
        document = Document()
        ambiguous = document.add_table(rows=1, cols=3)
        ambiguous.rows[0].cells[0].text = 'Estimated blood loss'
        ambiguous.rows[0].cells[2].text = 'mL'
        supported = document.add_table(rows=1, cols=2)
        supported.rows[0].cells[0].text = 'Assessment'
        path = self.root / 'form.docx'
        document.save(path)
        fields = docx_inspect(path)[3]
        self.assertEqual([field['label'] for field in fields], ['Assessment'])
        self.assertEqual(with_table_constraints(path, fields), fields)
        forged = deepcopy(fields)
        forged[0]['label'] = 'Estimated blood loss (mL)'
        self.assertNotIn('answer_constraints', with_table_constraints(path, forged)[0])

    def test_doctor_display_label_does_not_supply_or_remove_source_constraints(self):
        path, fields = self.form(['Temperature (°C)'])
        renamed = deepcopy(fields)
        renamed[0]['label'] = 'Temperature in Fahrenheit'
        result = with_table_constraints(path, renamed)
        self.assertEqual(result[0]['label'], 'Temperature in Fahrenheit')
        self.assertEqual(result[0]['answer_constraints']['unit'], '°C')
        forged = deepcopy(result)
        forged[0]['answer_constraints']['unit'] = '°F'
        with self.assertRaises(DocumentFailure) as refused:
            with_table_constraints(path, forged)
        self.assertEqual(refused.exception.code, 'content_changed')

    def test_signing_and_attestation_cells_get_no_new_constraints_or_destinations(self):
        path, fields = self.form(['Patient signature (number)', 'Physician signature (kg)',
                                  'Attestation (Yes / No)', 'Consent (Yes / No)'])
        before = deepcopy(fields)
        self.assertEqual(with_table_constraints(path, fields), before)
        self.assertTrue(all('answer_constraints' not in field for field in fields))
        # It preserves the inspector's existing destination policy; it never
        # inserts an anchor into a signing area or changes the source document.
        self.assertEqual(docx_inspect(path)[3], before)

    def test_missing_duplicate_or_forged_cell_is_refused(self):
        path, fields = self.form(['Temperature (°C)', 'Fever (Yes / No)'])
        wrong_location = deepcopy(fields)
        wrong_location[0]['locator']['cell'] = 1
        boolean_cell = deepcopy(fields)
        boolean_cell[0]['locator']['cell'] = False
        for forged in [fields[:-1], [fields[0], fields[0]], wrong_location, boolean_cell]:
            with self.subTest(forged=forged), self.assertRaises(DocumentFailure):
                with_table_constraints(path, forged)

    def test_original_zip_and_mixed_placeholder_mapping_remain_unchanged(self):
        path, fields = self.form(['Temperature (°C)'])
        document = Document(path)
        document.add_paragraph('Visit date: {{Visit date}} YYYY-MM-DD')
        document.save(path)
        fields = docx_inspect(path)[3]
        before = path.read_bytes()
        with ZipFile(path) as package:
            parts = {name: package.read(name) for name in package.namelist()}
        result = with_table_constraints(path, fields)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), hashlib.sha256(before).hexdigest())
        with ZipFile(path) as package:
            self.assertEqual({name: package.read(name) for name in package.namelist()}, parts)
        self.assertIn('answer_constraints', result[0])
        self.assertEqual(result[1], fields[1])

    def test_actual_top_level_helper_call_after_document_isolation(self):
        path, _ = self.form(['Temperature (°C), record result', 'Fever [Yes / No], select one'])
        original = path.read_bytes()
        modules = str(Path(__file__).resolve().parent)
        # The workers preload their complete dependency graph before Landlock.
        # A plain top-level import check cannot prove the function's lazy
        # dependency is also available when it actually processes a document.
        program = '\n'.join([
            'import json, sys', 'from pathlib import Path',
            'sys.path.insert(0, ' + repr(modules) + ')',
            'from engine import docx_inspect', 'from sandbox import isolate',
            'import preparation', 'from table_constraints import with_table_constraints',
            'from preparation_roles import role_schema',
            'workspace = Path(' + repr(str(self.root)) + ')', 'isolate(workspace)',
            'path = workspace / "form.docx"',
            'fields = with_table_constraints(path, docx_inspect(path)[3])',
            'assert fields[0]["answer_constraints"]["unit"] == "°C"',
            'assert fields[1]["answer_constraints"]["options"] == ["Yes", "No"]',
            'assert role_schema("checker")["additionalProperties"] is False',
            'print(json.dumps({"processed_under_isolation": True, "fields": len(fields)}))',
        ])
        result = subprocess.run([sys.executable, '-I', '-B', '-c', program],
                                cwd=self.root, env={'PATH': '/usr/bin:/bin', 'LANG': 'C.UTF-8', 'TZ': 'UTC'},
                                capture_output=True, text=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout), {'processed_under_isolation': True, 'fields': 2})
        self.assertEqual(path.read_bytes(), original)
