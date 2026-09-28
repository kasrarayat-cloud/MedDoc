"""Printed value columns and coded choices are distinct answer structures."""
import unittest

from docx import Document

from . import word_v3 as w, layout_source
from .test_question_boxes import data
from . import test_question_boxes as question_fixtures


def grid(headers=('Left', 'Right'), *, boxes=False, instruction='Pulses'):
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    table = outer.cell(0, 0).add_table(rows=4, cols=3)
    table.cell(0, 0).text = instruction
    for col, text in enumerate(headers, 1): table.cell(0, col).text = text
    for row, label in enumerate(('Radial', 'Femoral', 'Tibial'), 1):
        table.cell(row, 0).text = label
        if boxes:
            for col in (1, 2): table.cell(row, col).text = '☐'
    return doc


class StructuralColumns(unittest.TestCase):
    def test_blank_side_columns_keep_independent_values_and_fixed_headers(self):
        raw = data(grid()); fields = w.inspect(raw)['fields']
        self.assertEqual(len(fields), 6)
        self.assertTrue(all(f['kind'] == 'text' for f in fields))
        self.assertEqual({f['label'] for f in fields},
                         {f'{row} — {side}' for row in ('Radial', 'Femoral', 'Tibial') for side in ('Left', 'Right')})
        answers = {f['key']: str(i + 1) for i, f in enumerate(fields)}
        result = w.fill(raw, fields, answers)
        before = w.SlotIndex(w.open_package(raw)); after = w.SlotIndex(w.open_package(result))
        for f in fields:
            ref = f['locator']['part'] + '#' + f['locator']['slot']
            self.assertEqual(w.text_of(after.by_id[ref].el), answers[f['key']])
        from lxml import etree
        for col in range(3):
            header = f'word/document.xml#b/tbl[0]/tr[0]/tc[0]/tbl[0]/tr[0]/tc[{col}]'
            self.assertEqual(etree.tostring(before.by_id[header].el), etree.tostring(after.by_id[header].el))
        self.assertEqual(w.fill(raw, fields, {}), raw)

    def test_true_option_grids_keep_choices_and_never_map_the_header(self):
        for headers, boxes, instruction in [(('Yes', 'No'), False, 'Finding'),
                                            (('Left', 'Right'), True, 'Affected side'),
                                            (('Left', 'Right'), False, 'Select one side')]:
            with self.subTest(headers=headers, boxes=boxes, instruction=instruction):
                raw = data(grid(headers, boxes=boxes, instruction=instruction))
                fields = w.inspect(raw)['fields']
                self.assertEqual(len(fields), 3)
                self.assertTrue(all(f['kind'] == 'choice' for f in fields))
                self.assertTrue(all(f['options'] == list(headers) for f in fields))
                self.assertFalse(any('/tr[0]/tc[1]' in f['locator']['slot'].split('/tbl[0]/')[-1] for f in fields))

    def test_summary_title_owns_the_single_unlabelled_coded_row(self):
        for control in ('glyph', 'sdt', 'ff'):
            with self.subTest(control=control):
                doc = question_fixtures.TabColumnOwnership().form(control=control, count=3)
                doc.tables[0].cell(0, 1).text = 'Fixed details'
                cell = doc.tables[0].cell(0, 0)
                cell.paragraphs[0].text = 'Examination summary'
                row = cell.paragraphs[-1]
                row.runs[0].text = ''
                raw = data(doc)
                proposal, _, _ = layout_source.prepare(raw)
                self.assertEqual(len(proposal['fields']), 1)
                field = proposal['fields'][0]
                self.assertEqual(field['label'], 'Examination summary')
                self.assertEqual(field['options'], ['Normal', 'Abnormal', 'Unavailable'])
                self.assertEqual(w.inspect(raw)['fields'], [])
                _, result, _ = layout_source.prepare(raw, mapping=proposal['fields'], answers={field['key']: 'Abnormal'})
                self.assertEqual(w.option_states(w.SlotIndex(w.open_package(result)), field), [False, True, False])

    def test_unresolved_coded_row_never_becomes_numeric_options(self):
        for row in ('Finding      1 ☐    2 ☐', 'Finding\t\t1 ☐\t2 ☐'):
            doc = Document(); cell = doc.add_table(rows=1, cols=1).cell(0, 0)
            cell.text = 'Examination'; cell.add_paragraph(row)
            result = w.inspect(data(doc))
            self.assertFalse(any(f.get('options') for f in result['fields']))
            self.assertTrue(any('coded' in note.lower() for note in result['notes']))
