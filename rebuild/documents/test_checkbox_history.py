"""Bounded native-checkbox event lists; synthetic examples are not corpus qualification."""
import copy
import io
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from lxml import etree

from . import word_v3 as w
from .test_question_boxes import data
from .test_word_v3 import checkbox_control, ff_field, _sdt


def form(question='Treatments received', options=('Physiotherapy', 'Occupational therapy', 'Hydrotherapy'),
         *, control='ff', instruction='', heading='Clinical history', separate=False, after='', bold_after=False,
         table_before='', table_after=''):
    doc = Document(); doc.add_heading(heading, 1)
    if table_before: doc.add_table(rows=1, cols=1).cell(0, 0).text = table_before
    if instruction: doc.add_paragraph(instruction)
    if separate: doc.add_paragraph(question + ':')
    for i, option in enumerate(options):
        p = doc.add_paragraph(question + ': ' if i == 0 and not separate else '')
        if control == 'ff': ff_field(p, 'checkbox', name='Answer' + str(i))
        elif control == 'sdt': p._p.append(checkbox_control())
        else: p.add_run('☐')
        p.add_run(' ' + option)
    if after:
        doc.add_paragraph('')
        doc.add_paragraph().add_run(after).bold = bold_after
    if table_after: doc.add_table(rows=1, cols=1).cell(0, 0).text = table_after
    return doc


def boxed_question(doc, question, options, *, before=None):
    """A separate printed question that owns its own native boxes inside a one-cell table (the D31 page-1 shape)."""
    table = doc.add_table(rows=1, cols=1); cell = table.cell(0, 0)
    cell.paragraphs[0].text = question
    for i, option in enumerate(options):
        p = cell.add_paragraph(); ff_field(p, 'checkbox', name='Own' + str(i)); p.add_run(' ' + option)
    if before is not None:
        before._p.addprevious(table._tbl)
    return table


def rewrite(raw, old, new):
    out = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as a, ZipFile(out, 'w', ZIP_DEFLATED) as b:
        for name in a.namelist():
            value = a.read(name)
            b.writestr(name, value.replace(old.encode(), new.encode()) if name == 'word/document.xml' else value)
    return out.getvalue()


class CheckboxHistory(unittest.TestCase):
    def field(self, raw):
        fields = [f for f in w.inspect(raw)['fields'] if f['locator']['target']['type'] == 'checkboxes']
        self.assertEqual(len(fields), 1)
        return fields[0]

    def test_distinct_treatment_and_course_histories_select_multiple_native_controls(self):
        for question, options in (
                ('Treatments received', ('Physiotherapy', 'Occupational therapy', 'Hydrotherapy')),
                ('Courses attended', ('First aid', 'Food hygiene', 'Literacy'))):
            for control in ('ff', 'sdt'):
                with self.subTest(question=question, control=control):
                    raw = data(form(question, options, control=control)); field = self.field(raw)
                    self.assertEqual(field['kind'], 'multi_choice')
                    self.assertEqual(field['options'], list(options))
                    original = copy.deepcopy(field)
                    for chosen in ([], [options[1]], [options[0], options[2]]):
                        output = w.fill(raw, [field], {field['key']: chosen})
                        index = w.SlotIndex(w.open_package(output))
                        self.assertEqual([w.box_state(index, b) for b in field['locator']['target']['boxes']],
                                         [o in chosen for o in options])
                        before = w.SlotIndex(w.open_package(raw))
                        for box in field['locator']['target']['boxes']:
                            ident = box['part'] + '#' + box['slot']
                            node = index.by_id[ident].el
                            node.getparent().replace(node, copy.deepcopy(before.by_id[ident].el))
                        self.assertEqual(etree.tostring(before.package.tree('word/document.xml')),
                                         etree.tostring(index.package.tree('word/document.xml')))
                    self.assertEqual(field, original)

    def test_separate_complete_owner_and_local_instruction_are_bound(self):
        raw = data(form(separate=True, instruction='Record events during the reporting period.'))
        field = self.field(raw)
        self.assertEqual(field['kind'], 'multi_choice')
        self.assertIn('Record events during the reporting period.', field['context'])
        self.assertEqual(field['source_questions'][0]['question'], 'Treatments received')
        self.assertGreaterEqual(len(field['source_questions'][0]['sources']), 5)

    def test_explicit_exclusive_binary_and_single_event_scopes_keep_one_choice(self):
        for args in (
                {'instruction': 'Tick one.'}, {'heading': 'Select one'},
                {'question': 'Most recent treatment received'}, {'question': 'Current treatment received'},
                {'question': 'Single treatment received'}, {'question': 'Latest treatment received'},
                {'options': ('Yes', 'No')}, {'options': ('Mild', 'Moderate', 'Severe')},
                {'options': ('Physiotherapy', 'Yes', 'No')},
                {'options': ('Physiotherapy', 'None')}):
            with self.subTest(args=args):
                raw = data(form(**args)); field = self.field(raw)
                self.assertEqual(field['kind'], 'choice')
                with self.assertRaises(w.V3Failure):
                    w.fill(raw, [field], {field['key']: field['options'][:2]})

    def test_conflicting_or_unrepresented_limits_do_not_authorize_history_selections(self):
        for instruction in ('Select up to two.', 'Choose one. Tick all that apply.', 'At least two are required.'):
            with self.subTest(instruction=instruction):
                result = w.inspect(data(form(instruction=instruction)))
                self.assertFalse([f for f in result['fields'] if f['locator']['target']['type'] == 'checkboxes'])
                self.assertTrue(any('selection' in n for n in result['notes']))

    def test_only_one_owner_and_just_one_command_retain_exclusive_selection(self):
        for args in ({'question': 'Only one treatment received'}, {'instruction': 'Select just one.'}):
            with self.subTest(args=args):
                raw = data(form(**args)); field = self.field(raw)
                self.assertEqual(field['kind'], 'choice')
                output = w.fill(raw, [field], {field['key']: field['options'][0]})
                index = w.SlotIndex(w.open_package(output))
                self.assertEqual([w.box_state(index, b) for b in field['locator']['target']['boxes']], [True, False, False])
                with self.assertRaises(w.V3Failure):
                    w.fill(raw, [field], {field['key']: field['options'][:2]})

    def test_unknown_or_negative_selection_instructions_never_grant_multiple(self):
        for instruction in ('Do not select multiple options.', 'Do not tick all that apply.',
                            'Select one or more.', 'Never choose several options.'):
            with self.subTest(instruction=instruction):
                result = w.inspect(data(form(instruction=instruction)))
                self.assertFalse([f for f in result['fields'] if f['locator']['target']['type'] == 'checkboxes'])
                self.assertTrue(any('selection' in n for n in result['notes']))

    def test_complete_positive_all_that_apply_command_remains_useful(self):
        raw = data(form(instruction='Please select all that apply.')); field = self.field(raw)
        self.assertEqual(field['kind'], 'multi_choice')
        output = w.fill(raw, [field], {field['key']: field['options'][:2]})
        index = w.SlotIndex(w.open_package(output))
        self.assertEqual([w.box_state(index, b) for b in field['locator']['target']['boxes']], [True, True, False])

    def test_prior_unrelated_tick_one_table_does_not_own_history(self):
        # D31's shape: the tick-one instruction belongs to a separate question
        # that prints and owns its own boxes, so it does not govern the list.
        doc = form()
        boxed_question(doc, 'Type of application (please tick one):', ('Conditional', 'Absolute'),
                       before=doc.paragraphs[1])
        raw = data(doc)
        fields = [f for f in w.inspect(raw)['fields'] if f['locator']['target']['type'] == 'checkboxes']
        own = next(f for f in fields if f['label'].startswith('Type of application'))
        field = next(f for f in fields if f['label'] == 'Treatments received')
        self.assertEqual((own['kind'], field['kind']), ('choice', 'multi_choice'))
        out = w.fill(raw, [field], {field['key']: field['options'][:2]})
        index = w.SlotIndex(w.open_package(out))
        self.assertEqual([w.box_state(index, b) for b in field['locator']['target']['boxes']], [True, True, False])
        self.assertEqual(Document(io.BytesIO(out)).tables[0].cell(0, 0).paragraphs[0].text,
                         'Type of application (please tick one):')
        with self.assertRaises(w.V3Failure):
            w.fill(raw, [own], {own['key']: own['options']})

    def assertUnresolved(self, raw, note='selection'):
        result = w.inspect(raw)
        self.assertFalse([f for f in result['fields'] if f['locator']['target']['type'] == 'checkboxes'
                          and f['kind'] == 'multi_choice'])
        self.assertFalse([f for f in result['fields'] if f.get('source_questions')])
        self.assertTrue(any(note in n for n in result['notes']))

    def test_adjacent_selection_instruction_outside_the_owner_is_never_skipped(self):
        # A printed "tick one" beside the list, but outside the paragraphs the
        # rule reads, must never let the list accept several selections.
        for args in ({'after': 'Tick one only.'}, {'after': 'Please tick one box.'},
                     {'after': 'Select one', 'bold_after': True}, {'after': 'Tick all that apply.'},
                     {'table_after': 'Tick one only.'}, {'table_before': 'Please tick one box only.'},
                     {'table_before': 'Type of application (please tick one):'},
                     # A sample glyph printed inside an instruction is not a separate question's own box.
                     {'table_before': 'Mark ☒ one box only.'}, {'after': 'Tick ☐ one only.'}):
            with self.subTest(args=args):
                self.assertUnresolved(data(form(**args)))

    def test_form_wide_or_floating_selection_instruction_is_never_skipped(self):
        # Printed guidance whose reach is the whole form (a preamble, a page
        # header) or whose position is unknown (a floating text box) is never
        # read as part of the local relation, so it cannot be ignored either.
        textbox = ('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                   'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape style="position:absolute;'
                   'margin-left:300pt;width:120pt;height:30pt"><v:textbox><w:txbxContent><w:p><w:r><w:t>%s</w:t>'
                   '</w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>')
        preamble = form(); preamble.paragraphs[0]._p.addprevious(
            copy.deepcopy(preamble.add_paragraph('Please tick only one box for each question.')._p))
        preamble.paragraphs[-1]._p.getparent().remove(preamble.paragraphs[-1]._p)
        header = form(); header.sections[0].header.paragraphs[0].text = 'Tick one box per question.'
        in_heading = form(); in_heading.paragraphs[0]._p.append(parse_xml(textbox % 'Tick one only'))
        elsewhere = form(heading='Clinical history')
        early = elsewhere.paragraphs[0]._p
        intro = copy.deepcopy(early); early.addprevious(intro)
        note = elsewhere.add_paragraph('Details below.'); note._p.append(parse_xml(textbox % 'Tick one only'))
        early.addprevious(note._p)
        framed = form(); frame = framed.add_paragraph('Tick one only')
        frame._p.get_or_add_pPr().append(parse_xml(
            '<w:framePr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'w:w="2400" w:hAnchor="page" w:vAnchor="page" w:x="8000" w:y="3000"/>'))
        framed.paragraphs[0]._p.addprevious(frame._p)
        for name, doc in (('preamble', preamble), ('page header', header), ('text box in heading', in_heading),
                          ('text box elsewhere', elsewhere), ('positioned frame elsewhere', framed)):
            with self.subTest(case=name):
                self.assertUnresolved(data(doc))

    def test_form_wide_text_without_a_selection_command_keeps_the_history_useful(self):
        textbox = ('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                   'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape style="position:absolute;'
                   'margin-left:300pt;width:120pt;height:30pt"><v:textbox><w:txbxContent><w:p><w:r><w:t>%s</w:t>'
                   '</w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>')
        doc = form(); doc.sections[0].header.paragraphs[0].text = 'Please answer all questions on this form.'
        grid = doc.add_paragraph('Tick one box per row in the table below.')
        doc.paragraphs[0]._p.addprevious(copy.deepcopy(grid._p)); grid._p.getparent().remove(grid._p)
        doc.paragraphs[0]._p.append(parse_xml(textbox % 'Office use'))
        raw = data(doc); field = self.field(raw)
        self.assertEqual(field['kind'], 'multi_choice')
        output = w.fill(raw, [field], {field['key']: field['options'][:2]})
        index = w.SlotIndex(w.open_package(output))
        self.assertEqual([w.box_state(index, b) for b in field['locator']['target']['boxes']], [True, True, False])

    def test_a_selection_instruction_pointing_at_the_list_from_further_away_is_not_skipped(self):
        above = form()
        pointer = above.add_table(rows=1, cols=1); pointer.cell(0, 0).text = 'Please tick one box in the list below.'
        name = above.add_table(rows=1, cols=2); name.cell(0, 0).text = 'Name:'
        above.paragraphs[1]._p.addprevious(pointer._tbl); above.paragraphs[1]._p.addprevious(name._tbl)
        below = form(table_after='Other: give details.'); below.add_paragraph('For the list above, tick one only.')
        for name, doc in (('pointer two blocks above', above), ('pointer two blocks below', below)):
            with self.subTest(case=name):
                self.assertUnresolved(data(doc))

    def test_neighbouring_content_without_selection_wording_keeps_the_history_useful(self):
        plain = form(after='Record further details in the box below.')
        tables = form(table_before='Other: give the dates and destination.',
                      table_after='Other leave: give the dates and destination.')
        following = form(); boxed_question(following, 'Type of application (please tick one):', ('Conditional', 'Absolute'))
        for name, doc in (('plain paragraph after', plain), ('text tables around', tables),
                          ('next question owns its tick-one boxes', following)):
            with self.subTest(case=name):
                raw = data(doc)
                fields = [f for f in w.inspect(raw)['fields'] if f['locator']['target']['type'] == 'checkboxes']
                field = next(f for f in fields if f['label'] == 'Treatments received')
                self.assertEqual(field['kind'], 'multi_choice')
                self.assertTrue(all(f['kind'] == 'choice' for f in fields if f is not field))
                out = w.fill(raw, [field], {field['key']: field['options'][1:]})
                index = w.SlotIndex(w.open_package(out))
                self.assertEqual([w.box_state(index, b) for b in field['locator']['target']['boxes']], [False, True, True])

    def test_ordinary_complete_single_commands_remain_useful_single_choices(self):
        for instruction in ('Tick one box.', 'Tick one only.', 'Choose one of the following:', '(Tick one)',
                            'Select a single answer.', 'Pick one.', 'Tick 1 box.', 'Please select ONE option only.',
                            'Record the treatment. Tick one.', 'Pick 1.', 'Indicate 1 answer.'):
            with self.subTest(instruction=instruction):
                raw = data(form(instruction=instruction)); field = self.field(raw)
                self.assertEqual(field['kind'], 'choice')
                self.assertIn(instruction, field['context'])
                output = w.fill(raw, [field], {field['key']: field['options'][2]})
                index = w.SlotIndex(w.open_package(output))
                self.assertEqual([w.box_state(index, b) for b in field['locator']['target']['boxes']], [False, False, True])
                with self.assertRaises(w.V3Failure):
                    w.fill(raw, [field], {field['key']: field['options'][:2]})

    def test_a_numbered_part_heading_is_not_a_quantity(self):
        for heading in ('Part One: Leave history', 'Section One'):
            with self.subTest(heading=heading):
                raw = data(form(heading=heading)); field = self.field(raw)
                self.assertEqual(field['kind'], 'multi_choice')
                self.assertIn(heading, field['context'])
        self.assertUnresolved(data(form(heading='Part One: select one or more')))

    def test_ordinary_complete_multiple_commands_remain_useful(self):
        for instruction in ('Tick as many as apply.', 'Record the treatments. Tick all that apply.',
                            '(Tick all that apply)', 'All that apply.', 'Select any that apply.'):
            with self.subTest(instruction=instruction):
                raw = data(form(instruction=instruction)); field = self.field(raw)
                self.assertEqual(field['kind'], 'multi_choice')
                output = w.fill(raw, [field], {field['key']: [field['options'][0], field['options'][2]]})
                index = w.SlotIndex(w.open_package(output))
                self.assertEqual([w.box_state(index, b) for b in field['locator']['target']['boxes']], [True, False, True])

    def test_single_or_limited_wording_outside_the_command_grammar_never_grants_multiple(self):
        for instruction in ('Put an X in the box that applies.', 'Put a cross in one box.',
                            'Record either physiotherapy or hydrotherapy.', 'These are mutually exclusive.',
                            'Pick two.', 'Tick one. Do not tick more than one.', 'Max 2.', 'Choose the best option.'):
            with self.subTest(instruction=instruction):
                self.assertUnresolved(data(form(instruction=instruction)))

    def test_owner_quantity_other_than_exactly_one_is_unresolved(self):
        for question in ('One or more treatments received', 'One or two treatments received',
                         'One to three treatments received', 'One and only one treatment received'):
            with self.subTest(question=question):
                self.assertUnresolved(data(form(question=question)))

    def test_interrogative_or_conditional_owner_is_not_an_event_history(self):
        for question, options in (('Has a pregnancy test been completed', ('Yes - urine', 'Yes - serum', 'Declined')),
                                  ('Were any courses attended', ('First aid', 'Literacy', 'Food hygiene')),
                                  ('If yes, treatments received', ('Physiotherapy', 'Hydrotherapy'))):
            with self.subTest(question=question):
                raw = data(form(question, options)); field = self.field(raw)
                self.assertEqual(field['kind'], 'choice')
                self.assertNotIn('source_questions', field)
                with self.assertRaises(w.V3Failure):
                    w.fill(raw, [field], {field['key']: field['options'][:2]})

    def test_negative_or_exclusive_answer_options_keep_one_choice(self):
        for options in (('Physiotherapy', 'Hydrotherapy', 'None of the above'), ('Physiotherapy', 'No treatment'),
                        ('Physiotherapy', 'Yes - other'), ('Physiotherapy', 'Not taken')):
            with self.subTest(options=options):
                raw = data(form(options=options)); field = self.field(raw)
                self.assertEqual(field['kind'], 'choice')
                with self.assertRaises(w.V3Failure):
                    w.fill(raw, [field], {field['key']: field['options'][:2]})

    def test_referenced_owner_or_option_is_unresolved(self):
        doc = form(); note = OxmlElement('w:r'); ref = OxmlElement('w:footnoteReference')
        ref.set(qn('w:id'), '1'); note.append(ref); doc.paragraphs[1].runs[0]._r.addnext(note)
        cases = [('footnote', data(doc)), ('owner asterisk', data(form(question='Treatments received*'))),
                 ('option dagger', data(form(options=('Physiotherapy', 'Hydrotherapy†'))))]
        for name, raw in cases:
            with self.subTest(case=name):
                self.assertUnresolved(raw, note='native-checkbox history source is unresolved')

    def test_positioned_source_paragraph_is_unresolved(self):
        # A frame prints the paragraph somewhere else on the page: it is not
        # proven to be this list's adjacent instruction.
        doc = form(instruction='Record events during the reporting period.')
        doc.paragraphs[1]._p.get_or_add_pPr().append(parse_xml(
            '<w:framePr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'w:w="2400" w:hAnchor="page" w:vAnchor="page" w:x="8000" w:y="3000"/>'))
        self.assertUnresolved(data(doc), note='native-checkbox history source is unresolved')

    def test_explicit_single_selection_variants_never_accept_multiple_outside_history(self):
        # The ordinary checkbox question path (not the history rule).
        for question in ('Treatment(s) - select just one:', 'Treatment(s) - select exactly one:',
                         'Treatment(s) - pick one:', 'Which apply? Select a single option:',
                         'Treatment(s) - indicate only one:'):
            for control in ('glyph', 'ff'):
                with self.subTest(question=question, control=control):
                    doc = Document(); p = doc.add_paragraph(question + ' ')
                    for i, option in enumerate(('Physiotherapy', 'Hydrotherapy', 'Literacy')):
                        if control == 'ff': ff_field(p, 'checkbox', name='G' + str(i))
                        else: p.add_run('☐')
                        p.add_run(' ' + option + '   ')
                    raw = data(doc); field = self.field(raw)
                    self.assertEqual(field['kind'], 'choice')
                    with self.assertRaises(w.V3Failure):
                        w.fill(raw, [field], {field['key']: field['options'][:2]})
        for question in ('Treatment(s):', 'Treatment(s) - tick all that apply:'):
            with self.subTest(question=question):
                doc = Document(); p = doc.add_paragraph(question + ' ')
                for option in ('Physiotherapy', 'Hydrotherapy', 'Literacy'):
                    p.add_run('☐'); p.add_run(' ' + option + '   ')
                raw = data(doc); field = self.field(raw)
                self.assertEqual(field['kind'], 'multi_choice')
                output = w.fill(raw, [field], {field['key']: field['options'][:2]})
                index = w.SlotIndex(w.open_package(output))
                self.assertEqual([w.box_state(index, b) for b in field['locator']['target']['boxes']], [True, True, False])

    def test_d31_shaped_leave_history_accepts_two_leave_selections(self):
        # Synthetic mirror of D31's A15 layout (no corpus content): a long
        # section instruction box, question/answer tables, then the leave list
        # as two tab-separated native boxes per paragraph with spacer lines.
        doc = Document(); doc.add_heading('Conditional Discharge Proposal', 1)
        boxed_question(doc, 'Type of discharge applied for (please tick one):', ('Conditional', 'Absolute'))
        doc.add_table(rows=1, cols=2).rows[0].cells[1].text = 'Please see the list at annex B and select accordingly'
        doc.add_table(rows=1, cols=2).rows[0].cells[1].text = 'Responsible Clinician (including employing organisation)'
        doc.add_paragraph(''); doc.add_paragraph('')
        rows = (('Compassionate (day)', 'Compassionate (overnight)'), ('Medical (day)', 'Medical (overnight)'),
                ('Long Term Escorted Leave of Absence',))
        for n, row in enumerate(rows):
            p = doc.add_paragraph('Community leave taken:\t' if n == 0 else '')
            for k, option in enumerate(row):
                ff_field(p, 'checkbox', name='Leave%d%d' % (n, k)); p.add_run(' ' + option + ('    \t        \t' if k == 0 and len(row) > 1 else ''))
            doc.add_paragraph('')
        doc.add_table(rows=1, cols=2).rows[0].cells[1].text = 'Leave taken within the hospital grounds; leave to attend Court;'
        raw = data(doc)
        fields = [f for f in w.inspect(raw)['fields'] if f['locator']['target']['type'] == 'checkboxes']
        leave = next(f for f in fields if f['label'] == 'Community leave taken')
        discharge = next(f for f in fields if f['label'].startswith('Type of discharge'))
        self.assertEqual((leave['kind'], discharge['kind']), ('multi_choice', 'choice'))
        self.assertEqual(len(leave['options']), 5)
        chosen = ['Compassionate (day)', 'Compassionate (overnight)']
        output = w.fill(raw, [leave], {leave['key']: chosen})
        index = w.SlotIndex(w.open_package(output))
        self.assertEqual([w.box_state(index, b) for b in leave['locator']['target']['boxes']],
                         [o in chosen for o in leave['options']])
        with self.assertRaises(w.V3Failure):
            w.fill(raw, [discharge], {discharge['key']: discharge['options']})

    def test_full_source_and_mapping_identity_are_revalidated_at_export(self):
        raw = data(form(instruction='Record events during the reporting period.')); field = self.field(raw)
        self.assertEqual(field['kind'], 'multi_choice')
        for old, new in (('Treatments received', 'Most recent treatment received'),
                         ('Physiotherapy', 'Changed option'), ('Clinical history', 'Select one'),
                         ('Record events during the reporting period.', 'Tick one.')):
            with self.subTest(source=old):
                with self.assertRaises(w.V3Failure) as failure:
                    w.fill(rewrite(raw, old, new), [field], {field['key']: field['options'][:2]})
                self.assertEqual(failure.exception.code, 'mapping_mismatch')
        for alter in (lambda f: f.update(label='Invented question'), lambda f: f.update(context='Invented context'),
                      lambda f: f['options'].reverse(), lambda f: f.update(kind='choice'),
                      lambda f: f['source_questions'][0]['sources'].pop()):
            changed = copy.deepcopy(field); alter(changed)
            with self.assertRaises(w.V3Failure) as failure:
                w.fill(raw, [changed], {changed['key']: changed['options'][0]})
            self.assertEqual(failure.exception.code, 'mapping_mismatch')

    def test_unrelated_checkbox_and_dropdown_defaults_do_not_change(self):
        for args in ({'question': 'Preferred treatment'}, {'control': 'glyph'}):
            with self.subTest(args=args):
                field = self.field(data(form(**args)))
                self.assertEqual(field['kind'], 'choice')
                self.assertNotIn('source_questions', field)
        doc = Document(); p = doc.add_paragraph('Treatments received: ')
        p._p.append(_sdt('dropdown', options=('Physiotherapy', 'Hydrotherapy')))
        field = w.inspect(data(doc))['fields'][0]
        self.assertEqual(field['kind'], 'choice')

    def test_legacy_mapping_keeps_recorded_kind_and_source_metadata(self):
        raw = data(form()); field = self.field(raw)
        field['kind'] = 'choice'; field.pop('source_questions', None)
        previous = copy.deepcopy(field)
        output = w.fill(raw, [field], {field['key']: 'Physiotherapy'})
        index = w.SlotIndex(w.open_package(output))
        self.assertEqual([w.box_state(index, b) for b in field['locator']['target']['boxes']], [True, False, False])
        self.assertEqual(previous, field)
        with self.assertRaises(w.V3Failure):
            w.fill(raw, [field], {field['key']: field['options'][:2]})


LEAVE = ('Day', 'Overnight', 'Weekend')
TEXTBOX = ('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape style="position:absolute;'
           'margin-left:300pt;width:120pt;height:30pt"><v:textbox><w:txbxContent><w:p><w:r><w:t>%s</w:t>'
           '</w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>')


def native_list(doc, question='Community leave taken', options=LEAVE):
    """A native FORMCHECKBOX list, one option per paragraph, the owner on the first."""
    for i, option in enumerate(options):
        p = doc.add_paragraph(question + ': ' if i == 0 else '')
        ff_field(p, 'checkbox', name='L%d' % i); p.add_run(' ' + option)
    return doc


def inline_question(question, options, control):
    """An ordinary checkbox question: the question and its boxes in one paragraph (glyph or native)."""
    doc = Document(); p = doc.add_paragraph(question + ' ')
    for i, option in enumerate(options):
        if control == 'ff': ff_field(p, 'checkbox', name='G%d' % i)
        else: p.add_run('☐')
        p.add_run(' ' + option + '   ')
    return doc


class SelectionLimitRefutations(unittest.TestCase):
    """Refuter inputs against the native-checkbox history and the ordinary one-only reading (c002b620)."""

    def fields(self, raw, label):
        return [f for f in w.inspect(raw)['fields']
                if f['locator']['target']['type'] == 'checkboxes' and f['label'].startswith(label)]

    def ticks(self, raw, field, chosen):
        output = w.fill(raw, [field], {field['key']: chosen})
        index = w.SlotIndex(w.open_package(output))
        return [w.box_state(index, b) for b in field['locator']['target']['boxes']]

    def assertNeverTwo(self, raw, label='Community leave taken'):
        """A limited instruction may govern the list: it stays one choice or is held, and never takes two."""
        fields = self.fields(raw, label)
        self.assertLessEqual(len(fields), 1)
        for field in fields:
            self.assertEqual(field['kind'], 'choice')
            with self.assertRaises(w.V3Failure):
                w.fill(raw, [field], {field['key']: field['options'][:2]})
        self.assertTrue(any('selection' in n for n in w.inspect(raw)['notes']))

    def assertOneChoice(self, raw, label, chosen, expected):
        fields = self.fields(raw, label)
        self.assertEqual(len(fields), 1)
        field = fields[0]
        self.assertEqual(field['kind'], 'choice')
        self.assertEqual(self.ticks(raw, field, chosen), expected)
        with self.assertRaises(w.V3Failure):
            w.fill(raw, [field], {field['key']: field['options'][:2]})
        return field

    def assertMulti(self, raw, label, expected=(True, True, False)):
        fields = self.fields(raw, label)
        self.assertEqual(len(fields), 1)
        field = fields[0]
        self.assertEqual(field['kind'], 'multi_choice')
        self.assertEqual(self.ticks(raw, field, field['options'][:2]), list(expected))
        return field

    # A-F1: a tick-one printed higher up in the section, in a parent heading or below the list.
    def test_tick_one_in_the_section_or_a_parent_heading_never_grants_two(self):
        cases = []
        for instruction in ('Tick one box only for the leave question.', 'Please tick one box only.', 'For leave, tick one only.'):
            doc = Document(); doc.add_heading('Leave', 1); doc.add_paragraph(instruction)
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = 'Report on leave:'; table.cell(0, 1).text = 'Give the dates and destination.'
            doc.add_paragraph(''); cases.append(('above a table: ' + instruction, native_list(doc)))
        doc = Document(); doc.add_heading('Leave', 1); doc.add_paragraph('Please tick one box only in this section.')
        doc.add_heading('Community leave', 2); cases.append(('parent section', native_list(doc)))
        doc = Document(); doc.add_heading('Leave', 1); doc.add_heading('Community leave', 2)
        doc.add_paragraph('Tick one box only.'); doc.add_heading('Details', 3); cases.append(('under a sub-heading', native_list(doc)))
        doc = Document()
        for level, heading in enumerate(('Discharge report', 'Clinical history', 'Leave (tick one only)', 'Community leave')):
            doc.add_heading(heading, level + 1)
        cases.append(('heading level 3', native_list(doc)))
        doc = Document(); doc.add_paragraph().add_run('Please tick one box only').bold = True
        doc.add_paragraph().add_run('Community leave').bold = True; cases.append(('bold lines', native_list(doc)))
        doc = Document(); doc.add_heading('Leave', 1); native_list(doc); doc.add_paragraph('')
        doc.add_paragraph('Dates of leave: see record.'); doc.add_paragraph('Tick one box only.'); cases.append(('below', doc))
        for name, doc in cases:
            with self.subTest(case=name):
                self.assertNeverTwo(data(doc))

    # A-F2: singular-box instructions directly above the list.
    def test_singular_box_instructions_never_grant_two(self):
        for instruction in ('Place an X in the box that best describes the leave.', 'Place an X in the most appropriate box.',
                            'Enter X in the correct box.', 'Place an X in the relevant box.',
                            'Use the box that best describes the leave.', '✓ the relevant box.', 'Place a ✓ in the applicable box.'):
            with self.subTest(instruction=instruction):
                self.assertNeverTwo(data(form('Community leave taken', LEAVE, instruction=instruction)))

    # A-F3: a plain tick-one in the page header or the preamble.
    def test_plain_tick_one_in_a_page_header_or_preamble_never_grants_two(self):
        header = Document(); header.add_heading('Leave', 1); native_list(header)
        header.sections[0].header.paragraphs[0].text = 'Tick one box only.'
        preamble = Document(); preamble.add_paragraph('Please use black ink and tick one box only.')
        preamble.add_heading('Clinical history', 1); native_list(preamble)
        for name, doc in (('page header', header), ('preamble', preamble)):
            with self.subTest(case=name):
                self.assertNeverTwo(data(doc))

    # C-F1: a numeral quantity; C-F2: a pointer across a heading.
    def test_numeral_quantities_and_pointers_across_a_heading_never_grant_two(self):
        for instruction in ('Put an X in 1 box.', 'Not more than 1 treatment may be recorded.', 'Do not record more than 1.',
                            'Record 1 treatment.', 'Just 1.', '1 answer.', 'Record 2 treatments.'):
            with self.subTest(instruction=instruction):
                self.assertNeverTwo(data(form(instruction=instruction)), 'Treatments received')
        with self.subTest(owner='1 treatment received'):
            raw = data(form(question='1 treatment received'))
            self.assertOneChoice(raw, '1 treatment received', 'Hydrotherapy', [False, False, True])
        above = Document(); above.add_paragraph('In the list below, tick one only.'); above.add_heading('Clinical history', 1)
        native_list(above, 'Treatments received', ('Physiotherapy', 'Occupational therapy', 'Hydrotherapy'))
        below = Document(); below.add_heading('Clinical history', 1)
        native_list(below, 'Treatments received', ('Physiotherapy', 'Occupational therapy', 'Hydrotherapy'))
        below.add_heading('Notes', 1); below.add_paragraph('In the list above, tick one only.')
        for name, doc in (('pointer above a heading', above), ('pointer below a heading', below)):
            with self.subTest(case=name):
                self.assertNeverTwo(data(doc), 'Treatments received')

    # A-F4 and B-F5: one level, count or total.
    def test_single_level_count_or_total_keeps_one_choice(self):
        for question, options in (('Highest level of education completed', ('Primary', 'Secondary', 'University')),
                                  ('Number of therapy sessions attended', ('0-2', '3-5', '6 or more')),
                                  ('Total days of leave taken', ('Under 7 days', '7 to 14 days', 'Over 14 days'))):
            with self.subTest(question=question):
                doc = Document(); doc.add_heading('Clinical history', 1); native_list(doc, question, options)
                self.assertOneChoice(data(doc), question, options[1], [False, True, False])
        doc = Document(); doc.add_heading('Immunisation', 1); native_list(doc, 'Doses received', ('One', 'Two', 'Three'))
        self.assertOneChoice(data(doc), 'Doses received', 'Two', [False, True, False])

    # A-F5 and A-F8: ranges and alternatives are multi-select, not one-only.
    def test_ranges_and_alternatives_stay_multiple(self):
        symptoms = ('Headache', 'Nausea', 'Dizziness')
        for question, options in (('Symptoms - tick 1-3 that apply:', symptoms), ('Symptoms - tick 1–3 that apply:', symptoms),
                                  ('Symptoms (tick all that apply; more than a single answer may be ticked):', symptoms),
                                  ('Symptoms - pick one, some or all that apply:', symptoms),
                                  ('Side(s) affected - indicate one/both:', ('Left', 'Right')),
                                  ('Symptoms - select 1+ option(s):', symptoms),
                                  ('Symptoms - tick one or more that apply:', symptoms),
                                  ('Symptoms - select one or more option(s):', symptoms),
                                  ('Symptoms - choose one or more that apply:', symptoms)):
            for control in ('glyph', 'ff'):
                with self.subTest(question=question, control=control):
                    raw = data(inline_question(question, options, control))
                    self.assertMulti(raw, question.split(' ')[0], [True, True, False][:len(options)])

    # A-F6 and C-F3: explicit one-only wording on an ordinary "(s)" question.
    def test_explicit_one_only_wording_keeps_ordinary_questions_single(self):
        options = ('Physiotherapy', 'Hydrotherapy', 'Literacy')
        for question in ('Treatment(s) (one box only):', 'Treatment(s) - only one may be ticked:',
                         'Treatment(s) - only one answer allowed:', 'Treatment(s) - tick the ONE main treatment:',
                         'Treatment(s) - select the single best option:', 'Treatment(s) - no more than one:',
                         'Treatment(s) - maximum one:', 'Treatment(s) - select just one and give details:',
                         'Treatment(s) - pick one and give details:', 'Treatment(s) - indicate one to be reviewed:',
                         'Treatment(s) - pick one - see notes:', 'Treatment(s) - pick 1 / see notes:'):
            for control in ('glyph', 'ff'):
                with self.subTest(question=question, control=control):
                    self.assertOneChoice(data(inline_question(question, options, control)), 'Treatment(s)',
                                         'Hydrotherapy', [False, True, False])

    # A-F7: common printed multi-select commands above a native history.
    def test_common_all_that_apply_commands_map_multiple(self):
        for instruction in ('Tick all boxes that apply.', 'Tick all applicable boxes.', 'Tick as many boxes as apply.',
                            'Tick all that are applicable.', 'Please tick all relevant boxes.', 'You may tick more than one box.',
                            'Tick any/all that apply.'):
            with self.subTest(instruction=instruction):
                self.assertMulti(data(form('Community leave taken', LEAVE, instruction=instruction)), 'Community leave taken')

    # B-F1, B-F2, B-F3, B-F4: exclusive Yes/No answers keep one choice whatever is printed around them.
    def test_exclusive_yes_no_event_questions_keep_one_choice(self):
        cases = []
        for heading in ('Safety checks', 'Health check', 'Tumour markers', 'Treatment options', 'Multiple sclerosis review',
                        'One-stop clinic', 'Day one review', 'Selection criteria', 'Optional section'):
            doc = Document(); doc.add_heading(heading, 1); cases.append(('heading ' + heading, native_list(doc, 'Bloods taken', ('Yes', 'No'))))
        for instruction in ('Please tick the relevant box.', 'Please tick the appropriate box.', 'Please tick as appropriate.',
                            'Tick Yes or No.', 'Please select Yes or No.', 'Please mark with an X.',
                            'Please cross the appropriate box.', 'Please tick.', 'Tick if applicable.',
                            'Please tick one box for each item.'):
            doc = Document(); doc.add_heading('Discharge', 1); doc.add_paragraph(instruction)
            cases.append(('instruction ' + instruction, native_list(doc, 'Bloods taken', ('Yes', 'No'))))
        for sentence in ('Use one font throughout the document.', 'Participants of this study will receive payment in the '
                         'form of a check or electronic funds transfer (EFT).'):
            doc = Document(); doc.add_heading('Specimens', 1); native_list(doc, 'Bloods taken', ('Yes', 'No'))
            doc.add_heading('Notes', 1); doc.add_paragraph(sentence); cases.append(('form-wide ' + sentence, doc))
        for sentence in ('If yes, mark the site on the body map below.', 'If no, please indicate the reason in the box below.'):
            doc = Document(); doc.add_heading('Specimens', 1); native_list(doc, 'Bloods taken', ('Yes', 'No'))
            doc.add_paragraph(sentence); cases.append(('follow-up ' + sentence, doc))
        doc = Document(); doc.add_heading('Specimens', 1); native_list(doc, 'Swab taken', ('Yes', 'No'))
        native_list(doc, 'Bloods taken', ('Yes', 'No')); cases.append(('after a separate yes/no question', doc))
        for line in ('Date of procedure: ______________', 'Date: ____/____/______'):
            doc = Document(); doc.add_heading('Specimens', 1); doc.add_paragraph(line)
            p = doc.add_paragraph('Bloods taken: '); ff_field(p, 'checkbox', name='S0'); p.add_run(' Yes ')
            ff_field(p, 'checkbox', name='S1'); p.add_run(' No'); cases.append(('under ' + line, doc))
        for name, doc in cases:
            with self.subTest(case=name):
                self.assertOneChoice(data(doc), 'Bloods taken', 'Yes', [True, False])

    # B-F1, B-F2, B-F3: words and sentences that are not selection guidance keep a genuine history multiple.
    def test_non_selection_words_and_separate_questions_keep_the_history_multiple(self):
        treatments = ('Physiotherapy', 'Occupational therapy', 'Hydrotherapy')
        cases = []
        for heading in ('Tumour markers', 'Treatment options', 'Multiple sclerosis review', 'One-stop clinic', 'Day one review',
                        'Selection criteria', 'Optional section', 'Pre-operative checklist'):
            doc = Document(); doc.add_heading(heading, 1); cases.append(('heading ' + heading, native_list(doc, 'Treatments received', treatments)))
        for sentence in ('Use one font throughout the document.',
                         'IF MORE THAN ONE PAGE IS USED, EACH PAGE (VAF 10-1086) MUST BE CONSECUTIVELY NUMBERED.',
                         'Participants of this study will receive payment in the form of a check or electronic funds transfer (EFT).',
                         'Even if you are completing the checklist manually, keep the document.',
                         'For each question, please indicate how often you had this feeling.'):
            doc = Document(); doc.add_heading('Specimens', 1); native_list(doc, 'Treatments received', treatments)
            doc.add_heading('Notes', 1); doc.add_paragraph(sentence); cases.append(('form-wide ' + sentence[:30], doc))
        doc = Document(); doc.add_heading('Specimens', 1); native_list(doc, 'Treatments received', treatments)
        doc.add_heading('Notes', 1); doc.add_paragraph('Record any other notes.'); doc.add_section()
        doc.add_paragraph('Office')._p.append(parse_xml(TEXTBOX % 'FOR OFFICE USE ONLY'))
        cases.append(('office text box in a later section', doc))
        doc = Document(); doc.add_heading('Specimens', 1); native_list(doc, 'Treatments received', treatments)
        doc.sections[0].footer.paragraphs[0].text = 'This form is for clinical use only.'; cases.append(('footer', doc))
        doc = Document(); doc.add_heading('Clinical history', 1); doc.add_paragraph('Date of referral: ______________')
        cases.append(('under a date line', native_list(doc, 'Treatments received', treatments)))
        doc = Document(); doc.add_heading('Clinical history', 1); native_list(doc, 'Bloods taken', ('Yes', 'No'))
        cases.append(('after a separate yes/no question', native_list(doc, 'Treatments received', treatments)))
        for name, doc in cases:
            with self.subTest(case=name):
                self.assertMulti(data(doc), 'Treatments received')

    # B-F6: an all-that-apply list that also names a primary item is never narrowed to one.
    def test_all_that_apply_with_a_primary_designation_is_not_narrowed_to_one(self):
        options = ('Pain', 'Swelling', 'Fever', 'Rash')
        for question in ('Symptoms (tick all that apply and indicate one as the main symptom):',
                         'Symptoms (tick all that apply; pick one as primary)'):
            native = Document(); native.add_heading('Presentation', 1); p = native.add_paragraph(question + ' ')
            for i, option in enumerate(options):
                ff_field(p, 'checkbox', name='S%d' % i); p.add_run(' ' + option + ' ')
            glyph = Document(); glyph.add_heading('Presentation', 1); glyph.add_paragraph(question)
            for option in options:
                glyph.add_paragraph('☐ ' + option)
            for name, doc in (('native', native), ('glyph', glyph)):
                with self.subTest(question=question, control=name):
                    raw = data(doc)
                    fields = self.fields(raw, 'Symptoms')
                    self.assertFalse([f for f in fields if f['kind'] == 'choice'])
                    for field in fields:
                        self.assertEqual(self.ticks(raw, field, ['Pain', 'Fever']), [True, False, True, False])
                    if not fields:
                        self.assertTrue(any('cardinality instructions conflict' in n for n in w.inspect(raw)['notes']))

    # Controls for the repair itself: what it must not hold, narrow or widen.
    def test_repair_controls_keep_ordinary_answers_useful(self):
        treatments = ('Physiotherapy', 'Occupational therapy', 'Hydrotherapy')
        multi = []
        doc = Document(); doc.add_heading('Clinical history', 1)
        glyph = doc.add_paragraph('Sex (tick one): ☐ Male ☐ Female')
        multi.append(('a separate glyph question above owns its tick-one', native_list(doc, 'Treatments received', treatments)))
        doc = Document(); doc.add_heading('Demographics', 1); doc.add_paragraph('Sex (tick one): ☐ Male ☐ Female')
        doc.add_heading('Clinical history', 1)
        multi.append(('a glyph question in another section owns its tick-one', native_list(doc, 'Treatments received', treatments)))
        doc = Document(); doc.add_paragraph('You may tick more than one box for each question.'); doc.add_heading('Clinical history', 1)
        multi.append(('open count in the preamble', native_list(doc, 'Treatments received', treatments)))
        for preamble in ('Section 1: tick all boxes that apply.', 'Part One - tick all boxes that apply.'):
            doc = Document(); doc.add_paragraph(preamble); doc.add_heading('Clinical history', 1)
            multi.append(('a section number in the preamble: ' + preamble, native_list(doc, 'Treatments received', treatments)))
        multi.append(('dated instruction', form(instruction='Record treatments received within 6 weeks of discharge.')))
        multi.append(('clinical option name', form(options=('Fragile X', 'Karyotype', 'Microarray'), question='Tests completed')))
        doc = Document(); doc.add_heading('Clinical history', 1); native_list(doc, 'Treatments received', treatments)
        doc.add_paragraph('Give the date of the change at least 14 days before the move taking place.')
        doc.add_heading('Other', 1); doc.add_paragraph('Please indicate at least 14 days in advance.')
        multi.append(('measured amounts are not limits', doc))
        for name, doc in multi:
            with self.subTest(case=name):
                raw = data(doc)
                label = 'Tests completed' if 'option' in name else 'Treatments received'
                self.assertMulti(raw, label)
        self.assertEqual(glyph.text, 'Sex (tick one): ☐ Male ☐ Female')
        # A label paragraph directly above a separate question's own native boxes owns its tick-one.
        doc = Document(); doc.add_heading('Application', 1); doc.add_paragraph('Type of application (please tick one):')
        for i, option in enumerate(('Conditional', 'Absolute')):
            p = doc.add_paragraph(); ff_field(p, 'checkbox', name='T%d' % i); p.add_run(' ' + option)
        doc.add_heading('Clinical history', 1); native_list(doc, 'Treatments received', treatments)
        raw = data(doc)
        self.assertMulti(raw, 'Treatments received')
        self.assertOneChoice(raw, 'Type of application', 'Absolute', [False, True])
        # One-answer and singular-box phrases printed elsewhere in the form may govern the list.
        for sentence in ('Use the box that best describes the treatment.', 'One answer only.', '✓ one only.',
                         'Put X against one.', 'Tick 1.'):
            with self.subTest(elsewhere=sentence):
                doc = Document(); doc.add_paragraph(sentence); doc.add_heading('Clinical history', 1)
                self.assertNeverTwo(data(native_list(doc, 'Treatments received', treatments)), 'Treatments received')
        # A footnote and a positioned frame print away from the list: never owned by the block they sit beside.
        raw = data(native_list(Document(), 'Treatments received', treatments))
        out = io.BytesIO()
        with ZipFile(io.BytesIO(raw)) as a, ZipFile(out, 'w', ZIP_DEFLATED) as b:
            for name in a.namelist():
                b.writestr(name, a.read(name))
            b.writestr('word/footnotes.xml', '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
                       '2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>Tick one box only.</w:t></w:r></w:p></w:footnote>'
                       '</w:footnotes>')
        self.assertNeverTwo(out.getvalue(), 'Treatments received')
        doc = Document(); doc.add_heading('Application', 1); frame = doc.add_paragraph('Leave type (tick one only):')
        frame._p.get_or_add_pPr().append(parse_xml(
            '<w:framePr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'w:w="2400" w:hAnchor="page" w:vAnchor="page" w:x="8000" w:y="3000"/>'))
        p = doc.add_paragraph('Type: '); ff_field(p, 'checkbox', name='T0'); p.add_run(' A '); ff_field(p, 'checkbox', name='T1')
        p.add_run(' B'); doc.add_heading('Clinical history', 1)
        self.assertNeverTwo(data(native_list(doc, 'Treatments received', treatments)), 'Treatments received')
        # A glyph question owns its tick-one in a table row or as a label above its option lines; a legend with
        # sample glyphs does not.
        doc = Document(); doc.add_heading('Application', 1); row = doc.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = 'Type of application (please tick one):'; row.cells[1].text = '☐ Conditional ☐ Absolute'
        doc.add_heading('Clinical history', 1); self.assertMulti(data(native_list(doc, 'Treatments received', treatments)),
                                                                 'Treatments received')
        doc = Document(); doc.add_heading('Application', 1); doc.add_paragraph('Type of application (please tick one):')
        doc.add_paragraph('☐ Conditional'); doc.add_paragraph('☐ Absolute'); doc.add_heading('Clinical history', 1)
        self.assertMulti(data(native_list(doc, 'Treatments received', treatments)), 'Treatments received')
        for legend in ('Mark ☒ one box only, leave the others ☐.', 'Legend: mark ☒ for the one answer, ☐ for the others.'):
            with self.subTest(legend=legend):
                doc = Document(); doc.add_paragraph(legend); doc.add_heading('Clinical history', 1)
                self.assertNeverTwo(data(native_list(doc, 'Treatments received', treatments)), 'Treatments received')
        # An answer label is not a one-only limit on the question it belongs to.
        raw = data(inline_question('Symptom(s) (tick all that apply):', ('Pain in only one joint', 'Swelling', 'Fever'), 'ff'))
        self.assertMulti(raw, 'Symptom(s)')
        # A standalone tick-one above a separate question and the list is not that question's label.
        for instruction in ('Tick one box only.', 'Please tick one box only:'):
            with self.subTest(standalone=instruction):
                doc = Document(); doc.add_heading('Leave', 1); doc.add_paragraph(instruction)
                p = doc.add_paragraph('Escorted: '); ff_field(p, 'checkbox', name='E0'); p.add_run(' Yes ')
                ff_field(p, 'checkbox', name='E1'); p.add_run(' No')
                self.assertNeverTwo(data(native_list(doc, 'Treatments received', treatments)), 'Treatments received')
        # A one-only sentence owned by a separate question stops being owned when it names a wider reach.
        for label in ('Please tick one box for each question in this section:',
                      'Leave type (tick one box only for each question below):'):
            with self.subTest(label=label):
                doc = Document(); doc.add_heading('Clinical history', 1)
                boxed_question(doc, label, ('Yes', 'No'))
                doc.add_paragraph(''); native_list(doc, 'Treatments received', treatments)
                self.assertNeverTwo(data(doc), 'Treatments received')
        # Ordinary questions: a count that is an amount, and "more than one" in a question, are not selection limits.
        options = ('Physiotherapy', 'Hydrotherapy', 'Literacy')
        for question in ('Symptom(s) in the one week before admission:', 'Symptom(s) lasting no more than 1 day:'):
            with self.subTest(question=question):
                self.assertMulti(data(inline_question(question, options, 'ff')), 'Symptom(s)')
        raw = data(inline_question('Do you have more than one type of headache?', ('No', 'Yes - describe the worst'), 'ff'))
        self.assertOneChoice(raw, 'Do you have more than one type', 'No', [True, False])


def yes_no(doc, label='Escorted: ', tail=None):
    """A separate Yes/No question that owns its native boxes; `tail` is printed after a line break in its paragraph."""
    p = doc.add_paragraph(label); ff_field(p, 'checkbox', name='E0'); p.add_run(' Yes ')
    ff_field(p, 'checkbox', name='E1'); p.add_run(' No')
    if tail is not None:
        p.add_run().add_break(); p.add_run(tail)
    return p


def with_part(raw, name, xml):
    out = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as a, ZipFile(out, 'w', ZIP_DEFLATED) as b:
        for item in a.namelist():
            b.writestr(item, a.read(item))
        b.writestr(name, xml)
    return out.getvalue()


def d31_shape(label):
    """The D31-shaped synthetic form (as in test_d31_shaped_leave_history_accepts_two_leave_selections) with the
    page-1 boxed label spelled `label`; like the real cell, an empty line separates the label from its boxes."""
    doc = Document(); doc.add_heading('Conditional Discharge Proposal', 1)
    cell = boxed_question(doc, label, ('Conditional', 'Absolute')).cell(0, 0)
    cell.paragraphs[0]._p.addnext(copy.deepcopy(cell.add_paragraph('')._p)); cell._tc.remove(cell.paragraphs[-1]._p)
    doc.add_table(rows=1, cols=2).rows[0].cells[1].text = 'Please see the list at annex B and select accordingly'
    doc.add_table(rows=1, cols=2).rows[0].cells[1].text = 'Responsible Clinician (including employing organisation)'
    doc.add_paragraph(''); doc.add_paragraph('')
    rows = (('Compassionate (day)', 'Compassionate (overnight)'), ('Medical (day)', 'Medical (overnight)'),
            ('Long Term Escorted Leave of Absence',))
    for n, row in enumerate(rows):
        p = doc.add_paragraph('Community leave taken:\t' if n == 0 else '')
        for k, option in enumerate(row):
            ff_field(p, 'checkbox', name='Leave%d%d' % (n, k))
            p.add_run(' ' + option + ('    \t        \t' if k == 0 and len(row) > 1 else ''))
        doc.add_paragraph('')
    doc.add_table(rows=1, cols=2).rows[0].cells[1].text = 'Leave taken within the hospital grounds; leave to attend Court;'
    return doc


TREATMENTS = ('Physiotherapy', 'Occupational therapy', 'Hydrotherapy')


class SelectionLimitRound2(unittest.TestCase):
    """Refuter inputs F1-F9 against 233ac9a1 (the native-checkbox history and the ordinary one-only reading)."""
    fields, ticks = SelectionLimitRefutations.fields, SelectionLimitRefutations.ticks
    assertNeverTwo = SelectionLimitRefutations.assertNeverTwo
    assertOneChoice = SelectionLimitRefutations.assertOneChoice
    assertMulti = SelectionLimitRefutations.assertMulti

    @staticmethod
    def preamble(sentence, question='Treatments received', options=TREATMENTS):
        doc = Document(); doc.add_paragraph(sentence); doc.add_heading('Clinical history', 1)
        return native_list(doc, question, options)

    @staticmethod
    def header(sentence, question='Treatments received', options=TREATMENTS):
        doc = Document(); doc.add_heading('Clinical history', 1); native_list(doc, question, options)
        doc.sections[0].header.paragraphs[0].text = sentence
        return doc

    @staticmethod
    def floating(sentence):
        doc = Document(); doc.add_heading('Leave', 1)._p.append(parse_xml(TEXTBOX % sentence))
        return native_list(doc)

    # F1: a negated "more than one" is a limit, and floating selection wording is never skipped.
    def test_f1_negated_open_counts_and_floating_selection_wording_never_grant_two(self):
        for sentence in ('Please do not tick more than one box for each question.', 'Never tick more than 1 box per question.',
                         'No more than one box may be ticked.', 'Do not select more than one option.',
                         'Do not select multiple options.', 'No more than one answer per question.',
                         'Never select several options.', 'You cannot tick more than one box.',
                         "Don't tick more than one box.", 'Don’t tick more than one box.', 'Do not tick more than one.',
                         *('No more than one %s per question.' % noun for noun in (
                             'box', 'answer', 'option', 'response', 'choice', 'category'))):
            for where, build in (('preamble', self.preamble), ('page header', self.header)):
                with self.subTest(where=where, sentence=sentence):
                    self.assertNeverTwo(data(build(sentence)), 'Treatments received')
        for sentence in ('Do not tick more than one box.', 'Choose the leave type that applies.'):
            with self.subTest(text_box=sentence):
                self.assertNeverTwo(data(self.floating(sentence)))
        doc = Document(); frame = doc.add_paragraph('Choose the leave type that applies.'); doc.add_heading('Leave', 1)
        frame._p.get_or_add_pPr().append(parse_xml(
            '<w:framePr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'w:w="2400" w:hAnchor="page" w:vAnchor="page" w:x="8000" w:y="3000"/>'))
        with self.subTest(frame='Choose the leave type that applies.'):
            self.assertNeverTwo(data(native_list(doc)))
        # Controls: permission for several, and an attribution line in an office text box, stay multiple.
        self.assertMulti(data(self.preamble('You may tick more than one box for each question.')), 'Treatments received')
        for office in ('Checked by: ____________', 'Marked by: ____________', 'Page X of Y'):
            with self.subTest(office=office):
                self.assertMulti(data(self.floating(office)), 'Community leave taken')

    # F2: a scale key or a warning after "tick one" is not a range.
    def test_f2_tick_one_followed_by_a_scale_key_or_a_warning_stays_one_choice(self):
        scale, symptoms = ('1', '2', '3', '4', '5'), ('Pain', 'Swelling', 'Fever')
        for question, options, chosen, expected in (
                ('How strongly does each apply? Tick one – 1 (never) to 5 (always):', scale, '4',
                 [False, False, False, True, False]),
                ('How much do these statements apply to you? Circle one - 1 = not at all, 5 = completely:', scale, '4',
                 [False, False, False, True, False]),
                ('Pain score(s) (tick one, 1 = mild, 5 = worst):', scale, '4', [False, False, False, True, False]),
                ('Symptom(s) (tick one, all others must be left blank):', symptoms, 'Swelling', [False, True, False]),
                ('Main symptom(s) - tick one — two ticks will void the answer:', symptoms, 'Swelling', [False, True, False])):
            for control in ('glyph', 'ff'):
                with self.subTest(question=question, control=control):
                    self.assertOneChoice(data(inline_question(question, options, control)), question.split(' ')[0],
                                         chosen, expected)
        # Controls: a dash range that ends the phrase, or is printed unspaced before the answers, is still a range.
        for question in ('Symptom(s) - tick 1–3:', 'Symptom(s) (tick 1-3):', 'Symptom(s) - tick 1-3 of the following:',
                         'Symptom(s) - tick 1-3 which apply:', 'Symptom(s) - tick 1-3; see notes:',
                         'Symptom(s) - tick 1-3, see notes:', 'Symptom(s) - tick 1-3. See notes:',
                         'Symptom(s) - pick one, two or three:', 'Symptom(s) - tick 1-3 boxes:', 'Symptom(s) - tick 1-3 options:',
                         'Symptom(s) - tick 1-3 answers:', 'Symptom(s) - tick 1-3 items:', 'Symptom(s) - tick 1-3 responses:',
                         'Symptom(s) - tick 1-3 categories:', 'Symptom(s) - tick 1-3 choices:'):
            for control in ('glyph', 'ff'):
                with self.subTest(range=question, control=control):
                    self.assertMulti(data(inline_question(question, symptoms, control)), 'Symptom(s)')

    # F3: a note, or an instruction on a line of its own, is not a separate question's label.
    def test_f3_a_standalone_instruction_is_not_owned_by_the_question_below_it(self):
        cases = []
        for note in ('Please note: only one box may be ticked.', 'Note: tick one box only.', 'Instructions: tick one box only.',
                     'NB: tick one box only.', 'Important: please tick only one box.', 'Guidance: one answer only.',
                     'Leave type - tick one', 'Leave type (tick one):'):
            for question, options in (('Community leave taken', LEAVE), ('Treatments received', TREATMENTS)):
                doc = Document(); doc.add_heading('Leave', 1); doc.add_paragraph(note); yes_no(doc)
                cases.append((note + ' / ' + question, question, native_list(doc, question, options)))
        doc = Document(); doc.add_heading('Leave', 1); yes_no(doc, tail='Leave type (tick one box only):')
        cases.append(('after a line break', 'Community leave taken', native_list(doc)))
        doc = Document(); doc.add_heading('Leave', 1)
        p = doc.add_paragraph('Leave type (tick one box only):'); p.add_run().add_break(); p.add_run('Escorted: ')
        ff_field(p, 'checkbox', name='E0'); p.add_run(' Yes '); ff_field(p, 'checkbox', name='E1'); p.add_run(' No')
        cases.append(('before a line break', 'Community leave taken', native_list(doc)))
        doc = Document(); doc.add_heading('Leave', 1); doc.add_paragraph('Community leave: only one box may be ticked.')
        boxed_question(doc, 'Escorted?', ('Yes', 'No')); doc.add_paragraph('')
        cases.append(('above a boxed question', 'Community leave taken', native_list(doc)))
        # A question's own tick-one label with a second limit printed after its boxes.
        doc = Document(); doc.add_heading('Leave', 1); yes_no(doc, 'Escorted (tick one): ', tail='Leave type (tick one box only):')
        cases.append(('second limit after the boxes', 'Community leave taken', native_list(doc)))
        # Only a paragraph's first sentence can be a label; selection wording after the boxes is not a label command;
        # a label needs a question name.
        doc = Document(); doc.add_heading('Clinical history', 1)
        doc.add_paragraph('Sex (tick one): ☐ Male ☐ Female. Tick one box only for every question.')
        cases.append(('second sentence after a label', 'Treatments received', native_list(doc, 'Treatments received', TREATMENTS)))
        doc = Document(); doc.add_heading('Clinical history', 1)
        doc.add_paragraph('Select treatment plan: ☐ Standard ☐ Intensive (only one)')
        cases.append(('limit after the boxes', 'Treatments received', native_list(doc, 'Treatments received', TREATMENTS)))
        doc = Document(); doc.add_heading('Clinical history', 1); boxed_question(doc, '(tick one):', ('Yes', 'No'))
        doc.add_paragraph('')
        cases.append(('label without a name', 'Treatments received', native_list(doc, 'Treatments received', TREATMENTS)))
        # A notice in a cell beside its row's boxes is not a label either (only a cell holding its own boxes is).
        doc = Document(); doc.add_heading('Clinical history', 1); row = doc.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = 'Note: tick one box only.'; cell = row.cells[1].paragraphs[0]
        for i, option in enumerate(('Yes', 'No')):
            ff_field(cell, 'checkbox', name='R%d' % i); cell.add_run(' ' + option + ' ')
        doc.add_paragraph('')
        cases.append(('notice beside a row of boxes', 'Treatments received', native_list(doc, 'Treatments received', TREATMENTS)))
        # A body notice is not a label even directly above unlabelled box lines; a bracketed command after a colon is.
        doc = Document(); doc.add_heading('Leave', 1); doc.add_paragraph('Note: tick one box only.')
        for i, option in enumerate(('Escorted', 'Unescorted')):
            p = doc.add_paragraph(); ff_field(p, 'checkbox', name='E%d' % i); p.add_run(' ' + option)
        doc.add_heading('Clinical history', 1)
        cases.append(('body notice above box lines', 'Treatments received', native_list(doc, 'Treatments received', TREATMENTS)))
        # A positioned frame prints elsewhere on the page: never the label of the box lines that follow it in the file.
        doc = Document(); doc.add_heading('Leave', 1); frame = doc.add_paragraph('Leave type (tick one only):')
        frame._p.get_or_add_pPr().append(parse_xml(
            '<w:framePr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'w:w="2400" w:hAnchor="page" w:vAnchor="page" w:x="8000" w:y="3000"/>'))
        for i, option in enumerate(('Escorted', 'Unescorted')):
            p = doc.add_paragraph(); ff_field(p, 'checkbox', name='E%d' % i); p.add_run(' ' + option)
        doc.add_heading('Clinical history', 1)
        cases.append(('framed label above box lines', 'Treatments received', native_list(doc, 'Treatments received', TREATMENTS)))
        # A tick-one printed after a box is that answer's text, not the label of the box lines below it.
        doc = Document(); doc.add_heading('Leave', 1)
        for i, option in enumerate(('Escort required (tick one only):', 'Yes', 'No')):
            p = doc.add_paragraph(); ff_field(p, 'checkbox', name='E%d' % i); p.add_run(' ' + option)
        doc.add_heading('Clinical history', 1)
        cases.append(('tick-one after a box', 'Treatments received', native_list(doc, 'Treatments received', TREATMENTS)))
        for name, label, doc in cases:
            with self.subTest(case=name):
                self.assertNeverTwo(data(doc), label)

    # F4: a form-wide one-box rule that also says where the boxes are printed still limits.
    def test_f4_form_wide_one_box_rules_naming_a_column_or_table_still_limit(self):
        for sentence in ('Please answer each question by ticking one box in the column provided.',
                         'Please tick one box only for each question, in the column on the right.',
                         'For every question tick only one box in the answer column.',
                         'Please tick one box only (see table of codes on page 4).',
                         'Tick one box per row in the table below, and one box only for every other question.'):
            for where, build in (('preamble', self.preamble), ('page header', self.header)):
                with self.subTest(where=where, sentence=sentence):
                    self.assertNeverTwo(data(build(sentence)), 'Treatments received')
        for sentence in ('Tick one box per row in the table below.', 'In the table below, tick one box per row.',
                         'In the grid below, tick one box per row.', 'Mark one box in each row of the grid.',
                         'Select only one box per row.', 'Tick one box only in each row of the table.'):
            with self.subTest(row_only=sentence):
                self.assertMulti(data(self.preamble(sentence)), 'Treatments received')

    # F5: a label that covers both parts, or names the list's subject, is not owned by the Yes/No question.
    def test_f5_a_label_reaching_both_parts_governs_the_list(self):
        for label in ('Leave and treatment (tick one box for leave and one for treatment):',
                      'Leave and treatment (tick one box only in both parts):',
                      'Escort and treatment (tick one box only in these two lines):',
                      'Escorted (tick one box only throughout this form):'):
            with self.subTest(label=label):
                doc = Document(); doc.add_heading('Leave', 1); doc.add_paragraph(label); yes_no(doc)
                self.assertNeverTwo(data(native_list(doc, 'Treatments received', TREATMENTS)), 'Treatments received')
        # The same in a boxed question that owns its boxes, and a label whose name is itself a limit.
        for label in ('Escorted (tick one box only throughout this form):',
                      'Please tick one box for each question in this section (one only):'):
            with self.subTest(boxed=label):
                doc = Document(); doc.add_heading('Clinical history', 1); boxed_question(doc, label, ('Yes', 'No'))
                doc.add_paragraph('')
                self.assertNeverTwo(data(native_list(doc, 'Treatments received', TREATMENTS)), 'Treatments received')

    # F6: a singular "the answer/response/category that best ..." directly above the list.
    def test_f6_singular_best_answer_phrases_above_the_list_never_grant_two(self):
        for instruction in ('Record the answer that best describes the leave taken.', 'Enter the response that best fits.',
                            'Use the category which best matches the leave.'):
            with self.subTest(instruction=instruction):
                self.assertNeverTwo(data(form('Community leave taken', LEAVE, instruction=instruction)))

    # F7: D31's A15 stays multiple whatever the spelling of the separate boxed question's own tick-one label.
    def test_f7_d31_label_spellings_keep_the_leave_history_multiple(self):
        chosen = ['Compassionate (day)', 'Compassionate (overnight)']
        for label in ('Type of discharge applied for (please tick one):', 'Type of discharge applied for (please tick one)',
                      'Type of discharge applied for - please tick one',
                      'Type of discharge applied for (please tick one of the following):',
                      'Type of discharge applied for (please tick one - see page 3):',
                      'Type of discharge applied for (tick one, see notes):',
                      'Type of discharge applied for (tick one – see overleaf):',
                      'Type of discharge applied for (tick one — see page 2):', 'Type of discharge applied for (tick 1):',
                      'Type of discharge applied for (s41) (please tick one):',
                      'Type of discharge applied for: (please tick one)', 'Type of discharge applied for: please tick one',
                      'Type of discharge applied for [please tick one]:', 'Type of discharge applied for: [please tick one]',
                      *('Type of discharge applied for (please %s one):' % verb for verb in (
                          'check', 'choose', 'select', 'mark', 'circle', 'pick', 'indicate', 'cross')),
                      *('Type of discharge applied for (tick one %s):' % noun for noun in (
                          'option', 'answer', 'response', 'item', 'box', 'category', 'choice')),
                      'Type of discharge applied for (tick only one):', 'Type of discharge applied for (tick just one):',
                      'Type of discharge applied for (tick exactly one):', 'Type of discharge applied for (select a single option):',
                      'Type of discharge applied for (tick one box only):', 'Type of discharge applied for (one only):',
                      'Type of discharge applied for (one box only):', 'Type of discharge applied for (only one):',
                      'Type of discharge applied for (only one box):'):
            with self.subTest(label=label):
                raw = data(d31_shape(label))
                leave = self.fields(raw, 'Community leave taken')[0]
                discharge = self.fields(raw, 'Type of discharge')[0]
                self.assertEqual((leave['kind'], discharge['kind']), ('multi_choice', 'choice'))
                self.assertEqual(self.ticks(raw, leave, chosen), [o in chosen for o in leave['options']])
                with self.assertRaises(w.V3Failure):
                    w.fill(raw, [discharge], {discharge['key']: discharge['options']})
        doc = Document(); doc.add_heading('Clinical history', 1)
        row = doc.add_table(rows=1, cols=2).rows[0]; row.cells[0].text = 'Smoking status (tick one)'
        cell = row.cells[1].paragraphs[0]
        for i, option in enumerate(('Never', 'Ex', 'Current')):
            ff_field(cell, 'checkbox', name='S%d' % i); cell.add_run(' ' + option + ' ')
        doc.add_paragraph(''); raw = data(native_list(doc, 'Treatments received', TREATMENTS))
        self.assertMulti(raw, 'Treatments received')
        # A glyph question in a table cell owns its own tick-one; so does a bracketed command after a body label's colon.
        doc = Document(); doc.add_heading('Clinical history', 1)
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Sex (tick one): ☐ Male ☐ Female'; doc.add_paragraph('')
        self.assertMulti(data(native_list(doc, 'Treatments received', TREATMENTS)), 'Treatments received')
        doc = Document(); doc.add_heading('Leave', 1); doc.add_paragraph('Leave type: (tick one)')
        for i, option in enumerate(('Escorted', 'Unescorted')):
            p = doc.add_paragraph(); ff_field(p, 'checkbox', name='E%d' % i); p.add_run(' ' + option)
        doc.add_heading('Clinical history', 1)
        self.assertMulti(data(native_list(doc, 'Treatments received', TREATMENTS)), 'Treatments received')
        # A question name that is a selection verb but not a limit still owns its own tick-one.
        doc = Document(); doc.add_heading('Clinical history', 1)
        boxed_question(doc, 'Select a treatment plan (tick one):', ('Standard', 'Intensive')); doc.add_paragraph('')
        raw = data(native_list(doc, 'Treatments received', TREATMENTS))
        self.assertMulti(raw, 'Treatments received')
        self.assertOneChoice(raw, 'Select a treatment plan', 'Intensive', [False, True])

    # F8: wording that permits several boxes is not a one-answer limit; the singular wording still is.
    def test_f8_plural_box_wording_elsewhere_keeps_the_history_multiple(self):
        for sentence in ('Please tick the appropriate box(es).', 'Tick the relevant box or boxes.', 'Tick the relevant box/boxes.',
                         *('Tick the relevant %s(s).' % n for n in ('answer', 'option', 'response', 'category', 'square')),
                         *('Tick the relevant %s or %s.' % (n, p) for n, p in (
                             ('answer', 'answers'), ('option', 'options'), ('response', 'responses'),
                             ('category', 'categories'), ('square', 'squares'))),
                         *('Tick the relevant %s/%s.' % (n, p) for n, p in (
                             ('answer', 'answers'), ('option', 'options'), ('response', 'responses'),
                             ('category', 'categories'), ('square', 'squares')))):
            with self.subTest(sentence=sentence):
                self.assertMulti(data(self.preamble(sentence)), 'Treatments received')
        for sentence in ('Please tick the relevant box.', 'Tick the relevant answer.', 'Tick the relevant option or tick N/A.'):
            with self.subTest(singular=sentence):
                self.assertNeverTwo(data(self.preamble(sentence)), 'Treatments received')

    # F9: a failing control for every guard the refuter could remove without a failure.
    def test_f9_every_selection_guard_has_a_control(self):
        raw = data(native_list(Document(), 'Treatments received', TREATMENTS))
        self.assertNeverTwo(with_part(raw, 'word/endnotes.xml', '<w:endnotes xmlns:w="http://schemas.openxmlformats.org/'
                                      'wordprocessingml/2006/main"><w:endnote w:id="1"><w:p><w:r><w:t>Tick one box only.'
                                      '</w:t></w:r></w:p></w:endnote></w:endnotes>'), 'Treatments received')
        for control in ('glyph', 'ff'):
            with self.subTest(at_most=control):
                self.assertOneChoice(data(inline_question('Treatment(s) - at most one:', ('Physiotherapy', 'Hydrotherapy',
                                     'Literacy'), control)), 'Treatment(s)', 'Hydrotherapy', [False, True, False])
        doc = form(table_before='Tick whichever applies.')
        SelectionLimitRound2.assertUnresolved(self, data(doc))
        for options in (('Under 7 days', '7 to 14 days', 'Over 14 days'), ('Less than a week', 'More than a week'),
                        ('Fewer than three', 'Greater than three'), ('At least weekly', 'At most monthly'),
                        ('Up to a week', 'Between one and two weeks', 'Within a month'), ('Above target', 'Below target')):
            with self.subTest(options=options):
                doc = Document(); doc.add_heading('Clinical history', 1); native_list(doc, 'Leave taken', options)
                self.assertOneChoice(data(doc), 'Leave taken', options[1], [o == options[1] for o in options])
        for question, options in (('Level of care received', ('Independent', 'Supervised', 'Assisted')),
                                  ('Duration of therapy received', ('Short course', 'Standard course', 'Extended course')),
                                  ('Length of treatment received', ('Short course', 'Standard course', 'Extended course')),
                                  ('Frequency of sessions attended', ('Daily', 'Weekly', 'Monthly')),
                                  ('Amount of leave taken', ('Little', 'Some', 'Extensive'))):
            with self.subTest(question=question):
                doc = Document(); doc.add_heading('Clinical history', 1); native_list(doc, question, options)
                self.assertOneChoice(data(doc), question, options[1], [False, True, False])

    def assertUnresolved(self, raw, note='selection'):
        CheckboxHistory.assertUnresolved(self, raw, note)
