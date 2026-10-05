"""Pure original-clause/physical-slot controls, without application resources."""
from copy import deepcopy
import io
import unittest
from zipfile import ZipFile

from docx import Document

from . import word_v3 as word


def fixture(printed):
    doc = Document(); doc.add_paragraph(printed); doc.add_paragraph('Plan: ____________')
    doc.sections[0].header.paragraphs[0].text = 'Fixed letterhead'
    out = io.BytesIO(); doc.save(out)
    return out.getvalue()


def insertion(raw):
    return next(f for f in word.inspect(raw)['fields']
                if f['locator']['target'].get('shape') == 'insertion')


class NominalInsertionRole(unittest.TestCase):
    def test_historical_decision_keeps_actor_and_only_fills_nominal_argument(self):
        for actor, report, decision, argument, noun in [
                ('We', 'reviewed', 'decided', 'type', 'services'),
                ('I', 'considered', 'determined', 'kind', 'care'),
                ('Our team', 'assessed', 'found', 'category', 'treatment')]:
            printed = f'{actor} {report} the case and {decision} that the {{insert {argument}}} {noun} should end.'
            with self.subTest(printed=printed):
                raw = fixture(printed); field = insertion(raw)
                self.assertFalse(field.get('write_protected'))
                final = word.fill(raw, [field], {field['key']: 'skilled nursing'})
                old, new = word.SlotIndex(word.open_package(raw)), word.SlotIndex(word.open_package(final))
                self.assertEqual(word.text_of(new.by_id['word/document.xml#b/p[0]'].el),
                                 printed.replace('{insert ' + argument + '}', 'skilled nursing'))
                self.assertEqual(word.fingerprint(old.by_id['word/document.xml#b/p[1]'].el),
                                 word.fingerprint(new.by_id['word/document.xml#b/p[1]'].el))
                with ZipFile(io.BytesIO(raw)) as before, ZipFile(io.BytesIO(final)) as after:
                    for name in before.namelist():
                        if name != 'word/document.xml':self.assertEqual(before.read(name), after.read(name))

    def test_permission_signing_author_and_document_frames_retain_holds(self):
        for printed in (
                'We reviewed this document and decided that the {insert type} services should end.',
                'We reviewed consent and decided that the {insert type} services should end.',
                'We reviewed the case and decided that I authorize {insert type} treatment.',
                'We reviewed the case and decided that the physician approves {insert type} treatment.',
                'We reviewed the case and decided that the {insert type} signature should be added.',
                'Manager comments: We reviewed the case and decided that the {insert type} services should end.',
                'I certify the case and decided that the {insert type} services should end.'):
            with self.subTest(printed=printed):
                raw = fixture(printed); field = insertion(raw)
                self.assertTrue(field.get('write_protected'))
                forged = deepcopy(field); forged.pop('write_protected', None); forged['kind'] = 'text'
                with self.assertRaises(word.V3Failure):
                    word.fill(raw, [forged], {forged['key']: 'skilled nursing'})

    def test_unknown_current_negated_or_non_nominal_act_is_not_new_authority(self):
        for printed in (
                'We review the case and decide that the {insert type} services should end.',
                'We did not review the case and decided that the {insert type} services should end.',
                'We reviewed the case and did not decide that the {insert type} services should end.',
                'We reviewed the case and decided that {insert type} should end services.',
                'We reviewed the case and decided that the {insert type} ended services.',
                'We reviewed the case and decided that {insert name} should end services.',
                'We reviewed the case and decided that the {insert type} services should end; I approve this form.'):
            with self.subTest(printed=printed):
                raw = fixture(printed); field = insertion(raw)
                self.assertTrue(field.get('write_protected'))
                with self.assertRaises(word.V3Failure):
                    word.fill(raw, [field], {field['key']: 'Synthetic answer'})

    def test_source_relation_and_target_cannot_be_reassigned_by_review(self):
        raw = fixture('We reviewed the case and decided that the {insert type} services should end.')
        original = insertion(raw)
        for mutation in ('request', 'question', 'fingerprint', 'target'):
            with self.subTest(mutation=mutation), self.assertRaises(word.V3Failure):
                field = deepcopy(original)
                if mutation == 'target':field['locator']['target']['start'] -= 1
                elif mutation == 'fingerprint':field['source_questions'][0]['source']['fingerprint'] = '0' * 64
                else:field['source_questions'][0][mutation] = 'Ordinary category'
                word.fill(raw, [field], {field['key']: 'skilled nursing'})
