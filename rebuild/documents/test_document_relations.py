"""Source ownership controls; synthetic documents are not corpus qualification."""
import copy
import io
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

from . import word_v3 as w
from .test_question_boxes import box, data


def link(paragraph, text='Published guidance', address='https://example.test/guidance'):
    node = OxmlElement('w:hyperlink')
    node.set(qn('r:id'), paragraph.part.relate_to(address, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); value = OxmlElement('w:t'); value.text = text
    run.append(value); node.append(run); paragraph._p.append(node)


def numbered(doc, values, kind='bullet'):
    abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), '90')
    level = OxmlElement('w:lvl'); level.set(qn('w:ilvl'), '0')
    fmt = OxmlElement('w:numFmt'); fmt.set(qn('w:val'), kind)
    level.append(fmt); abstract.append(level); doc.part.numbering_part.element.append(abstract)
    number = OxmlElement('w:num'); number.set(qn('w:numId'), '10')
    ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), '90'); number.append(ref)
    doc.part.numbering_part.element.append(number)
    for value in values:
        p = doc.add_paragraph(value)
        props = OxmlElement('w:numPr'); num = OxmlElement('w:numId'); num.set(qn('w:val'), '10')
        props.append(num); p._p.get_or_add_pPr().append(props)


class DocumentRelations(unittest.TestCase):
    def test_terminal_write_request_owns_its_plain_blank_before_a_neighboring_question(self):
        doc = Document(); table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Examination'; table.cell(0, 1).text = 'Comments'
        cell = table.cell(1, 0); cell.text = '19. Other (specify)'; cell.add_paragraph()
        cell.add_paragraph('Mental status: existing neighboring entry')
        raw = data(doc); field = next(f for f in w.inspect(raw)['fields'] if f['locator']['slot'].endswith('tc[0]/p[1]'))
        self.assertEqual(field['source_questions'][0]['question'], '19. Other (specify)')
        final = Document(io.BytesIO(w.fill(raw, [field], {field['key']: 'Skin finding'})))
        self.assertEqual([p.text for p in final.tables[0].cell(1, 0).paragraphs],
                         ['19. Other (specify)', 'Skin finding', 'Mental status: existing neighboring entry'])
        self.assertEqual(final.tables[0].cell(1, 1)._tc.xml, table.cell(1, 1)._tc.xml)

    def test_sole_terminal_write_request_appends_below_caption_without_replacing_it(self):
        for directive in ('specify', 'please describe', 'explain'):
            with self.subTest(directive=directive):
                doc = Document(); table = doc.add_table(rows=2, cols=2)
                table.cell(0, 0).text = 'Examination'; table.cell(0, 1).text = 'Comments'
                caption = 'Other (' + directive + ')'; cell = table.cell(1, 0); cell.text = caption
                raw = data(doc); field = next(f for f in w.inspect(raw)['fields']
                                            if f['locator']['slot'] == 'b/tbl[0]/tr[1]/tc[0]')
                self.assertEqual(field['locator']['target'], {'type': 'cell_append'})
                self.assertEqual(field['kind'], 'long_text')
                answer = ('Complete synthetic finding. ' * 100).strip()
                final = Document(io.BytesIO(w.fill(raw, [field], {field['key']: answer})))
                self.assertEqual([p.text for p in final.tables[0].cell(1, 0).paragraphs], [caption, answer])
                self.assertEqual(final.tables[0].cell(1, 0).paragraphs[0]._p.xml, cell.paragraphs[0]._p.xml)
                self.assertEqual(final.tables[0].cell(1, 1)._tc.xml, table.cell(1, 1)._tc.xml)
                for variant in ('relation', 'replace'):
                    altered = copy.deepcopy(field)
                    if variant == 'relation': altered.pop('source_questions')
                    else: altered['locator']['target'] = {'type': 'cell_replace'}
                    with self.assertRaises(w.V3Failure): w.fill(raw, [altered], {field['key']: answer})

    def test_terminal_request_does_not_release_human_authoring_or_signing(self):
        for caption in ('Patient signature (specify)', '19. Patient signature (specify)', 'Consent (specify)',
                        'This section is for Alex Reed to add comments (specify)'):
            with self.subTest(caption=caption):
                doc = Document(); cell = doc.add_table(rows=1, cols=1).cell(0, 0); cell.text = caption
                raw = data(doc); fields = w.inspect(raw)['fields']; self.assertTrue(fields)
                field = copy.deepcopy(fields[0]); field['label'] = 'Clinical answer'; field['kind'] = 'text'
                field.pop('write_protected', None); field.pop('review_reason', None)
                with self.assertRaises(w.V3Failure): w.fill(raw, [field], {field['key']: 'Dr Synthetic Finalizer'})

    def test_ordinary_signed_receipt_and_initial_reading_requests_still_take_complete_text(self):
        for caption in ('Signed receipt (specify)', 'Initial reading (specify)'):
            for below in (False, True):
                with self.subTest(caption=caption, below=below):
                    doc = Document(); cell = doc.add_table(rows=1, cols=1).cell(0, 0); cell.text = caption
                    if below:
                        cell.add_paragraph(); cell.add_paragraph('Unchanged neighboring question')
                    raw = data(doc); field = next(f for f in w.inspect(raw)['fields'] if f.get('source_questions'))
                    self.assertIn(field['kind'], ('text', 'long_text')); self.assertFalse(field.get('write_protected'))
                    final = Document(io.BytesIO(w.fill(raw, [field], {field['key']: 'Reviewed ordinary value'})))
                    self.assertIn('Reviewed ordinary value', final.tables[0].cell(0, 0).text)

    def test_independent_request_before_coded_columns_keeps_its_blank_with_unused_inherited_tabs(self):
        from docx.enum.style import WD_STYLE_TYPE
        from docx.shared import Inches
        from .test_word_v3 import ff_field
        doc = Document(); style = doc.styles.add_style('Synthetic tabbed spacing', WD_STYLE_TYPE.PARAGRAPH)
        style.paragraph_format.tab_stops.add_tab_stop(Inches(1))
        cell = doc.add_table(rows=1, cols=1).cell(0, 0); cell.text = '19. Other (specify)'
        blank = cell.add_paragraph(); blank.style = style
        cell.add_paragraph('Mental status\tNormal\tAbnormal')
        row = cell.add_paragraph('\t1 '); ff_field(row, 'checkbox', name='Check1')
        row.add_run('\t2 '); ff_field(row, 'checkbox', name='Check2')
        raw = data(doc); info = w.inspect(raw)
        field = next(f for f in info['fields'] if f['locator']['slot'] == 'b/tbl[0]/tr[0]/tc[0]/p[1]')
        self.assertEqual(field['source_questions'][0]['question'], '19. Other (specify)')
        self.assertTrue(info['coverage_issues'])
        self.assertFalse([f for f in info['fields'] if f['label'] == 'Mental status'])

    def test_existing_filled_hidden_or_status_qualified_cells_are_not_empty_requested_answers(self):
        for variant in ('filled', 'hidden', 'status', 'trailing', 'unknown_style'):
            with self.subTest(variant=variant):
                doc = Document(); cell = doc.add_table(rows=1, cols=1).cell(0, 0)
                cell.text = {'status': 'Other (specified)', 'trailing': 'Other (specify) recorded'}.get(variant, 'Other (specify)')
                if variant == 'filled': cell.add_paragraph('Existing entry')
                elif variant == 'hidden': cell.paragraphs[0].runs[0].font.hidden = True
                elif variant == 'unknown_style':
                    node = OxmlElement('w:pStyle'); node.set(qn('w:val'), 'UnknownStyle'); cell.paragraphs[0]._p.get_or_add_pPr().append(node)
                raw = data(doc)
                self.assertFalse([f for f in w.inspect(raw)['fields'] if any(
                    relation.get('relation') == 'parenthetical_write_request' for relation in f.get('source_questions', []))])

    def test_a_rows_other_option_cannot_become_an_unconditional_write_request(self):
        for below in (False, True):
            doc = Document(); row = doc.add_table(rows=1, cols=4).rows[0]
            for cell, text in zip(row.cells, ('Residence', 'Own home', 'With family', 'Other (specify)')):
                cell.text = text
            if below: row.cells[-1].add_paragraph()
            self.assertFalse([f for f in w.inspect(data(doc))['fields'] if any(
                relation.get('relation') == 'parenthetical_write_request' for relation in f.get('source_questions', []))])

    def test_heading_with_own_digit_destinations_cannot_claim_following_spacers(self):
        doc = Document(); cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        p = cell.paragraphs[0]; p.style = 'Heading 3'
        for label in ('Waist Circumference: ', 'in.  Hip Circumference: '):
            p.add_run(label)
            for _ in range(3):
                run = p.add_run(); symbol = OxmlElement('w:sym')
                symbol.set(qn('w:font'), 'Symbol'); symbol.set(qn('w:char'), 'F0F0')
                run._r.append(symbol)
        p.add_run('in.'); cell.add_paragraph(); cell.add_paragraph()
        raw = data(doc); info = w.inspect(raw); fields = info['fields']
        self.assertEqual([f['label'] for f in fields], ['Waist Circumference', 'Hip Circumference'])
        self.assertTrue(all(f['locator']['slot'] == 'b/tbl[0]/tr[0]/tc[0]/p[0]' for f in fields))
        self.assertEqual([(f['locator']['target']['start'], f['locator']['target']['end']) for f in fields], [(21, 24), (48, 51)])
        out = Document(io.BytesIO(w.fill(raw, fields, {fields[0]['key']: '032', fields[1]['key']: '040'})))
        self.assertEqual(out.tables[0].cell(0, 0).paragraphs[0].text, 'Waist Circumference: 032in.  Hip Circumference: 040in.')
        self.assertTrue(all(not p.text for p in out.tables[0].cell(0, 0).paragraphs[1:]))
        self.assertEqual(w.fill(raw, fields, {}), raw)

    def test_all_recognized_inline_blanks_keep_their_own_heading_destination(self):
        for placeholder in ('|__|__|__|', '{{Findings}}', '[Describe findings]'):
            with self.subTest(placeholder=placeholder):
                doc = Document(); cell = doc.add_table(rows=1, cols=1).cell(0, 0)
                p = cell.paragraphs[0]; p.style = 'Heading 3'; p.text = 'Findings: ' + placeholder
                cell.add_paragraph(); cell.add_paragraph(); raw = data(doc); fields = w.inspect(raw)['fields']
                self.assertEqual(len(fields), 1)
                self.assertEqual(fields[0]['locator']['slot'], 'b/tbl[0]/tr[0]/tc[0]/p[0]')
                self.assertEqual(fields[0]['locator']['target']['type'], 'span')
        # An unfilled underline uses the same precise blank reader. A plain
        # label without its own blank still owns its directly following box.
        for underline in (True, False):
            doc = Document(); cell = doc.add_table(rows=1, cols=1).cell(0, 0)
            p = cell.paragraphs[0]; p.style = 'Heading 3'; p.text = 'Findings:'
            if underline: p.add_run('   ').underline = True
            cell.add_paragraph(); cell.add_paragraph(); fields = w.inspect(data(doc))['fields']
            self.assertEqual(len(fields), 1)
            self.assertEqual(fields[0]['locator']['slot'], 'b/tbl[0]/tr[0]/tc[0]/p[' + ('0' if underline else '1') + ']')

    def test_previous_question_has_complete_local_source(self):
        doc = Document(); doc.add_heading('Relationships', 1)
        first = 'How is the reporting condition monitored outside self-report?'
        follow = 'In relation to the above question, what did the patient say?'
        doc.add_paragraph(first); box(doc); doc.add_paragraph(follow); box(doc)
        fields = w.inspect(data(doc))['fields']
        answer = next(f for f in fields if f['label'] == follow)
        self.assertIn(first, answer['context'])
        self.assertIn('Relationships', answer['context'])
        self.assertEqual(answer['locator']['slot'], 'b/tbl[1]/tr[0]/tc[0]')
        self.assertEqual(answer['source_questions'][0]['question'], first)
        source = data(doc)
        w.fill(source, fields, {answer['key']: 'Synthetic response.'})
        doc.paragraphs[1].text = 'A different clinical question?'
        with self.assertRaises(w.V3Failure) as failure:
            w.fill(data(doc), fields, {answer['key']: 'Synthetic response.'})
        self.assertEqual(failure.exception.code, 'mapping_mismatch')

    def test_reference_does_not_cross_boundaries_or_borrow_an_answer(self):
        for boundary in ('heading', 'intervening', 'unsigned', 'answer', 'following', 'unmapped'):
            with self.subTest(boundary=boundary):
                doc = Document(); doc.add_heading('History', 1)
                doc.add_paragraph('Please provide the signature?' if boundary == 'unsigned' else 'What changed?')
                box(doc, bordered=boundary != 'unmapped')
                if boundary == 'heading': doc.add_heading('Plan', 1)
                if boundary == 'intervening': doc.add_paragraph('Printed intervening instruction.')
                reference = {'answer': 'above answer', 'following': 'following question'}.get(boundary, 'above question')
                doc.add_paragraph('In relation to the ' + reference + ', what happened?'); box(doc)
                fields = w.inspect(data(doc))['fields']
                self.assertFalse([f for f in fields if f['locator']['slot'] == 'b/tbl[1]/tr[0]/tc[0]'])

    def test_fixed_link_introductions_are_not_answer_fields(self):
        for introduction in ('For questions please contact:', 'Please refer to the application forms at:'):
            with self.subTest(introduction=introduction):
                doc = Document(); doc.add_paragraph(introduction); doc.add_paragraph(); link(doc.add_paragraph())
                doc.add_paragraph('Patient contact: __________')
                original = data(doc); fields = w.inspect(original)['fields']
                self.assertEqual([f['label'] for f in fields], ['Patient contact'])
                exported = w.fill(original, fields, {fields[0]['key']: 'Synthetic person'})
                self.assertIn(b'https://example.test/guidance', w.Package(exported).zip.read('word/_rels/document.xml.rels'))

    def test_conditional_guidance_list_is_not_a_choice(self):
        doc = Document(); doc.add_paragraph('If you need any of the following:')
        numbered(doc, ('Change of conditions', 'Request supporting evidence'), 'lowerLetter')
        doc.add_paragraph('Then please refer to the guidance at:'); link(doc.add_paragraph())
        self.assertEqual(w.inspect(data(doc))['fields'], [])

    def test_section_heading_does_not_add_a_phantom_answer(self):
        doc = Document(); doc.add_heading('Annex: Patient comments:', 1)
        doc.add_paragraph('What would you like recorded?'); box(doc)
        self.assertEqual([f['label'] for f in w.inspect(data(doc))['fields']], ['What would you like recorded?'])

    def test_heading_with_an_explicit_blank_still_has_an_answer(self):
        doc = Document(); doc.add_heading('Patient name: __________', 1)
        self.assertEqual([f['label'] for f in w.inspect(data(doc))['fields']], ['Patient name'])

    def test_list_box_has_full_question_without_guessing_list_cardinality(self):
        doc = Document(); doc.add_heading('Accommodation', 1)
        doc.add_paragraph('What type of residence does the patient occupy:')
        numbered(doc, ('Own home', 'With family', 'Other (please specify)'))
        for _ in range(3): doc.add_paragraph()
        box(doc)
        fields = w.inspect(data(doc))['fields']
        field = next(f for f in fields if f['locator']['target']['type'] == 'cell')
        self.assertEqual(field['label'], 'What type of residence does the patient occupy:')
        self.assertIn('What type of residence does the patient occupy:', field['context'])
        self.assertIn('Own home | With family | Other (please specify)', field['context'])
        for answer in ('', 'Synthetic supported residence.', 'Synthetic long response. ' * 100):
            filled = w.fill(data(doc), fields, {field['key']: answer})
            self.assertEqual(Document(__import__('io').BytesIO(filled)).tables[0].cell(0, 0).text, answer)

    def test_other_box_does_not_cross_an_intervening_boundary(self):
        for boundary in ('heading', 'table', 'question', 'consent'):
            with self.subTest(boundary=boundary):
                doc = Document(); doc.add_paragraph('Consent choices:' if boundary == 'consent' else 'Residence types:')
                numbered(doc, ('Own home', 'Other (please specify)'))
                if boundary == 'heading': doc.add_heading('Unrelated', 1)
                elif boundary == 'table': doc.add_table(rows=1, cols=2)
                elif boundary == 'question': doc.add_paragraph('New unrelated instruction.')
                table = box(doc); result = w.inspect(data(doc))
                cells = [f for f in result['fields'] if f['locator']['target']['type'] == 'cell']
                self.assertFalse([f for f in cells if f['label'].lower().startswith('other')])

    def test_third_party_author_area_is_visible_and_cannot_be_forged_writable(self):
        doc = Document(); doc.add_heading('Manager comments', 1)
        statement = 'This section is for the line manager of the clinician to add any comments or concerns.'
        doc.add_paragraph(statement); box(doc)
        original = data(doc); fields = w.inspect(original)['fields']
        self.assertEqual(len(fields), 1)
        field = fields[0]; self.assertTrue(field['write_protected'])
        self.assertEqual(field['label'], statement)
        field.pop('write_protected'); field.pop('review_reason')
        with self.assertRaises(w.V3Failure) as failure:
            w.fill(original, fields, {field['key']: 'The manager agrees.'})
        self.assertEqual(failure.exception.code, 'protected_signing_area')

    def test_explicit_conditional_selection_keeps_its_choices_beside_help(self):
        for verb in ('select', 'choose', 'check', 'tick', 'mark', 'circle', 'indicate'):
            with self.subTest(verb=verb):
                doc = Document(); doc.add_paragraph('If pain is present, ' + verb + ' one severity:')
                numbered(doc, ('Mild', 'Severe'), 'lowerLetter')
                doc.add_paragraph('For help, please refer to the guidance at:'); link(doc.add_paragraph())
                fields = w.inspect(data(doc))['fields']
                self.assertEqual(len(fields), 1)
                self.assertEqual(fields[0]['options'], ['Mild', 'Severe'])

    def test_reference_rechecks_physical_destination_after_identical_box_insertion(self):
        doc = Document(); doc.add_heading('History', 1)
        doc.add_paragraph('What symptoms were observed?'); first = box(doc)
        follow = doc.add_paragraph('In relation to the above question, what did the patient say?'); box(doc)
        fields = w.inspect(data(doc))['fields']; field = next(f for f in fields if f.get('source_questions'))
        follow._p.addprevious(copy.deepcopy(first._tbl))
        with self.assertRaises(w.V3Failure) as failure:
            w.fill(data(doc), fields, {field['key']: 'Synthetic response.'})
        self.assertEqual(failure.exception.code, 'mapping_mismatch')

    def test_other_and_already_filled_author_boxes_remain_protected(self):
        for other in (False, True):
            for filled in (False, True):
                with self.subTest(other=other, filled=filled):
                    doc = Document(); doc.add_heading('Parent comments', 1)
                    doc.add_paragraph('This section is for the parent to provide comments:')
                    if other: numbered(doc, ('No comments', 'Other (please specify)'))
                    table = box(doc); fields = w.inspect(data(doc))['fields']
                    field = next(f for f in fields if f['locator']['target']['type'] == 'cell')
                    self.assertTrue(field['write_protected'])
                    if filled: table.cell(0, 0).text = 'Existing synthetic parent statement.'
                    original = data(doc); idx = w.SlotIndex(w.open_package(original))
                    slot = idx.by_id[field['locator']['part'] + '#' + field['locator']['slot']]
                    field.pop('write_protected'); field.pop('review_reason')
                    field['locator']['fingerprint'] = w.fingerprint(slot.el)
                    field['locator']['target'] = {'type': 'cell_replace', 'all': True}
                    with self.assertRaises(w.V3Failure) as failure:
                        w.fill(original, fields, {field['key']: 'Invented parent statement.'})
                    self.assertEqual(failure.exception.code, 'protected_signing_area')
