"""Multi-column question/answer rows, and tables combined with underlined blanks.

Two documented gaps, both previously visible only as a warning:

  * a question/answer row wider than two columns was never mapped, so a four-column vitals
    grid or a `Label | Answer | Unit` row produced `no_fields` and the form could not be used;
  * the tab adapter ran ONLY when nothing else matched, so a form holding both a two-column
    table and a labelled underlined section mapped the table and discarded the section — while
    the same section on its own imported perfectly.

Both are now supported where placement is unambiguous, and both stay unsupported where it is
not. The ambiguous cases are as much the subject of this file as the supported ones: a row
whose blanks outnumber its labels, and a row whose first cell is empty, must NOT be guessed.

Fixtures are authored here so the suite is self-contained. Synthetic only.
"""

from pathlib import Path
import re
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from .engine import CONTRACT_V1, CONTRACT_V2, docx_inspect, multi_column_answers
from .service import DocumentError, fill_template, inspect_template

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
OWNER_FORM = Path(__file__).with_name("__fixtures__") / "Colonoscopy_Operative_Template.docx"


def underlined_tab(paragraph):
    run = paragraph.add_run()
    properties = OxmlElement("w:rPr")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run._r.append(properties)
    run._r.append(OxmlElement("w:tab"))
    return run


def grid(document, rows, columns):
    table = document.add_table(rows=len(rows), cols=columns)
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.rows[r].cells[c].text = value
    return table


def labelled_blank(document, label):
    paragraph = document.add_paragraph()
    paragraph.add_run(label)
    underlined_tab(paragraph)
    return paragraph


def texts(path):
    """Every w:t across body, headers and footers, as a LIST.

    Never join these with '' to test containment: that fuses text across element boundaries
    and invents tokens the document does not contain.
    """
    found = []
    with ZipFile(path) as package:
        for name in package.namelist():
            if re.fullmatch(r"word/(document|header\d*|footer\d*)\.xml", name):
                found += etree.fromstring(package.read(name)).xpath("//w:t/text()", namespaces=NS)
    return found


class Base(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)

    def save(self, document, name):
        path = Path(self.directory.name) / name
        document.save(path)
        return str(path)

    def labels(self, path):
        return [f["label"] for f in inspect_template(path)["fields"]]

    def fill(self, path, name, answers=None, fields=None):
        fields = fields if fields is not None else inspect_template(path)["fields"]
        answers = answers if answers is not None else {f["key"]: f"Answer for {f['label']}" for f in fields}
        out = str(Path(self.directory.name) / name)
        summary = fill_template(path, fields, answers, out)
        return out, summary, answers


class MultiColumnRule(Base):
    """The rule itself, before any document is involved."""

    def cells_of(self, values):
        document = Document()
        grid(document, [values], len(values))
        return document.element.body.iter(qn("w:tr")).__next__().findall(qn("w:tc"))

    def test_strict_pairs_are_offered(self):
        offered = multi_column_answers(self.cells_of(["Height", "", "Weight", ""]))
        self.assertEqual([label for label, _ in offered], ["Height", "Weight"])

    def test_a_part_filled_pairs_row_is_refused_because_data_wears_the_same_shape(self):
        """⚠️ THIS TEST ASSERTED THE OPPOSITE, AND THE REFUTATION PASS WAS RIGHT.

        It read `Procedure | Colonoscopy | Consent | ____` as a pairs grid with a preprinted
        value and offered `Consent`. But `Aspirin | 81 mg | Daily | ____` is THE SAME SHAPE,
        and there the question offered was `Daily` — a frequency — with its answer written into
        the Route column of a medication row. Measured, both of them.

        Nothing in the row tells the two apart, so neither is mapped. The row is reported.
        """
        for row in (["Procedure", "Colonoscopy", "Consent", ""],
                    ["Aspirin", "81 mg", "Daily", ""],
                    ["Date", "Finding", "Action", ""]):
            with self.subTest(row=row):
                self.assertEqual(multi_column_answers(self.cells_of(row)), [])

    def test_no_odd_width_row_is_mapped_however_tidy_it_looks(self):
        """⚠️ ALSO REVERSED. `Label | blank | Unit` was mapped as one question.

        `Estimated blood loss | ____ | mL` cannot be told from `Sodium | 140 | ____`, where the
        answer lands in the REFERENCE RANGE column beside a result that already reads 140, or
        from `Diabetes | ____ | X`, free text into the Yes column of a history grid whose No
        column is ticked. All three are "one blank after a non-empty first cell".
        """
        for row in (["Estimated blood loss", "", "mL"], ["Estimated blood loss", "mL", ""],
                    ["Sodium", "140", ""], ["Diabetes", "", "X"]):
            with self.subTest(row=row):
                self.assertEqual(multi_column_answers(self.cells_of(row)), [])

    def test_a_cell_holding_a_ticked_symbol_is_not_an_empty_answer_cell(self):
        """`cell_text` reads only `w:t`, so a Wingdings tick reads as empty and gets erased."""
        cells = self.cells_of(["Consent obtained", "", "Sedation", ""])
        run = cells[1].xpath(".//w:p")[0].makeelement(qn("w:r"), {})
        symbol = OxmlElement("w:sym")
        symbol.set(qn("w:font"), "Wingdings")
        symbol.set(qn("w:char"), "F0FE")
        run.append(symbol)
        cells[1].xpath(".//w:p")[0].append(run)
        self.assertEqual(multi_column_answers(cells), [])
        # Positive control: without the symbol the identical row IS a pairs grid.
        self.assertEqual([label for label, _ in
                          multi_column_answers(self.cells_of(["Consent obtained", "", "Sedation", ""]))],
                         ["Consent obtained", "Sedation"])

    def test_two_blanks_after_one_label_are_refused(self):
        self.assertEqual(multi_column_answers(self.cells_of(["Allergies", "", ""])), [])

    def test_a_blank_in_a_label_position_is_refused(self):
        self.assertEqual(multi_column_answers(self.cells_of(["", "", "Clinician", ""])), [])

    def test_a_fully_populated_row_offers_nothing(self):
        self.assertEqual(multi_column_answers(self.cells_of(["Drug", "Dose", "Route"])), [])

    def test_two_columns_are_left_to_the_existing_rule(self):
        self.assertEqual(multi_column_answers(self.cells_of(["Indication", ""])), [])

    def test_an_overlong_label_cell_is_refused(self):
        self.assertEqual(multi_column_answers(self.cells_of(["x" * 81, "", "Weight", ""])), [])


class MultiColumnImport(Base):
    def test_four_column_pairs_map_in_reading_order(self):
        document = Document()
        document.add_paragraph("SYNTHETIC VITALS FORM")
        grid(document, [["Height", "", "Weight", ""], ["Pulse", "", "Blood pressure", ""]], 4)
        path = self.save(document, "vitals.docx")
        self.assertEqual(self.labels(path), ["Height", "Weight", "Pulse", "Blood pressure"])

    def test_six_columns_work_the_same_way(self):
        document = Document()
        grid(document, [["Date", "", "Time", "", "Room", ""]], 6)
        self.assertEqual(self.labels(self.save(document, "six.docx")), ["Date", "Time", "Room"])

    def test_a_header_row_is_not_mistaken_for_questions(self):
        document = Document()
        grid(document, [["Measure", "Value", "Measure", "Value"], ["Systolic", "", "Diastolic", ""]], 4)
        self.assertEqual(self.labels(self.save(document, "header.docx")), ["Systolic", "Diastolic"])

    def test_label_answer_unit_is_reported_not_mapped(self):
        """An accepted narrowing: `value | blank | unit` reads exactly like a result table row."""
        document = Document()
        grid(document, [["Estimated blood loss", "", "mL"], ["Procedure duration", "", "minutes"]], 3)
        with self.assertRaises(DocumentError) as raised:
            inspect_template(self.save(document, "units.docx"))
        self.assertEqual(raised.exception.code, "no_fields")
        self.assertIn("not two columns", str(raised.exception))

    def test_an_ambiguous_row_is_refused_and_the_refusal_explains_itself(self):
        document = Document()
        grid(document, [["Allergies", "", ""], ["Medications", "", ""]], 3)
        with self.assertRaises(DocumentError) as raised:
            inspect_template(self.save(document, "ambiguous.docx"))
        self.assertEqual(raised.exception.code, "no_fields")
        self.assertIn("not two columns", str(raised.exception))

    def test_a_merged_row_is_still_reported_not_mapped(self):
        document = Document()
        table = grid(document, [["Patient number", "", "Weight", ""]], 4)
        properties = table.rows[0].cells[1]._tc.get_or_add_tcPr()
        merge = OxmlElement("w:vMerge")
        merge.set(qn("w:val"), "restart")
        properties.append(merge)
        document.add_paragraph("PLAN: {{Plan}}")
        result = inspect_template(self.save(document, "merged.docx"))
        self.assertEqual([f["label"] for f in result["fields"]], ["Plan"])
        self.assertTrue(any("merged" in w for w in result["warnings"]),
                        "a merge must stay reported, not become a multi-column guess")


class MixedStructures(Base):
    def build(self, name, *, placeholder=False, multi=False, ambiguous=False):
        document = Document()
        if multi:
            grid(document, [["Height", "", "Weight", ""]], 4)
        else:
            grid(document, [["Referring clinician", ""]], 2)
        if placeholder:
            document.add_paragraph("ASSESSMENT: {{Assessment}}")
        document.add_heading("PROCEDURE DETAIL", level=1)
        if ambiguous:
            paragraph = document.add_paragraph()
            paragraph.add_run("MEASUREMENTS: ")
            underlined_tab(paragraph)
            underlined_tab(paragraph)
        else:
            labelled_blank(document, "EXTENT: ")
        return self.save(document, name)

    def test_a_table_and_an_underlined_section_now_map_together(self):
        path = self.build("mixed.docx")
        self.assertEqual(self.labels(path), ["Referring clinician", "EXTENT"])

    def test_table_placeholder_and_underlined_blank_all_three(self):
        path = self.build("three.docx", placeholder=True)
        self.assertEqual(self.labels(path), ["Referring clinician", "Assessment", "EXTENT"])

    def test_a_multi_column_grid_composes_with_an_underlined_blank(self):
        path = self.build("multi-mixed.docx", multi=True)
        self.assertEqual(self.labels(path), ["Height", "Weight", "EXTENT"])

    def test_an_ambiguous_underlined_area_leaves_the_rest_of_the_form_working(self):
        """NEVER WORSE THAN BEFORE.

        inspect_tabs refuses adjacent blanks without separate labels, and is right to. Letting
        that refusal escape would turn a form that imports today into one that does not, so it
        falls back to the note it replaced and the table still maps.
        """
        path = self.build("ambiguous-mixed.docx", ambiguous=True)
        result = inspect_template(path)
        self.assertEqual([f["label"] for f in result["fields"]], ["Referring clinician"])
        self.assertTrue(any("could NOT be placed unambiguously" in w for w in result["warnings"]),
                        result["warnings"])

    def test_a_mixed_mapping_fills_every_shape_without_losing_the_others(self):
        path = self.build("mixed-fill.docx", placeholder=True)
        out, summary, answers = self.fill(path, "mixed-filled.docx")
        joined = " ".join(texts(out))
        self.assertEqual(summary["filled"], 3)
        for answer in answers.values():
            self.assertIn(answer, joined)
            self.assertEqual(joined.count(answer), 1, "an answer must be inserted exactly once")
        for static in ("Referring clinician", "PROCEDURE DETAIL", "EXTENT: "):
            self.assertIn(static, joined, "original text must survive the fill")
        self.assertIn("ASSESSMENT: ", joined, "text around a replaced placeholder must survive")


class FillBehaviour(Base):
    def vitals(self):
        document = Document()
        document.add_paragraph("SYNTHETIC VITALS FORM")
        grid(document, [["Height", "", "Weight", ""]], 4)
        return self.save(document, "fill-vitals.docx")

    def test_normal_answers_land_in_their_own_cells(self):
        path = self.vitals()
        out, summary, answers = self.fill(path, "normal.docx")
        joined = " ".join(texts(out))
        self.assertEqual(summary["filled"], 2)
        self.assertEqual(texts(out).count("Answer for Height"), 1)
        for static in ("SYNTHETIC VITALS FORM", "Height", "Weight"):
            self.assertIn(static, joined)

    def test_missing_answers_stay_blank_and_are_named(self):
        path = self.vitals()
        fields = inspect_template(path)["fields"]
        out, summary, _ = self.fill(path, "missing.docx", answers={f["key"]: "" for f in fields})
        self.assertEqual(summary["filled"], 0)
        self.assertEqual(sorted(summary["missing"]), sorted(f["key"] for f in fields))
        self.assertIn("Height", texts(out))

    def test_a_long_answer_is_kept_whole_and_not_duplicated(self):
        path = self.vitals()
        fields = inspect_template(path)["fields"]
        long_answer = "Fictional long answer. " * 60
        out, summary, _ = self.fill(path, "long.docx",
                                    answers={fields[0]["key"]: long_answer, fields[1]["key"]: ""})
        joined = " ".join(texts(out))
        self.assertEqual(summary["filled"], 1)
        self.assertIn(long_answer, joined)
        self.assertEqual(joined.count(long_answer), 1)
        self.assertIn("Weight", joined, "a neighbouring label must not be overwritten by expansion")


class ContractCompatibility(Base):
    """A mapping saved before this change must still export, unchanged.

    `validate_mapping` compares `fields != mapping` exactly, so a document that now yields MORE
    fields would refuse its own stored mapping. Filling re-reads under the legacy contract when
    the current one does not match, and exact equality is the proof that it is the right one.
    """

    def mixed(self):
        document = Document()
        grid(document, [["Referring clinician", ""]], 2)
        document.add_heading("PROCEDURE DETAIL", level=1)
        labelled_blank(document, "EXTENT: ")
        return self.save(document, "contract.docx")

    def test_the_two_contracts_see_the_document_differently(self):
        path = self.mixed()
        self.assertEqual([f["label"] for f in docx_inspect(path, contract=CONTRACT_V1)[3]],
                         ["Referring clinician"])
        self.assertEqual([f["label"] for f in docx_inspect(path, contract=CONTRACT_V2)[3]],
                         ["Referring clinician", "EXTENT"])

    def test_a_legacy_mapping_still_fills(self):
        path = self.mixed()
        legacy = docx_inspect(path, contract=CONTRACT_V1)[3]
        answers = {f["key"]: "Legacy answer" for f in legacy}
        out = str(Path(self.directory.name) / "legacy.docx")
        summary = fill_template(path, legacy, answers, out)
        self.assertEqual(summary["filled"], 1)
        self.assertIn("Legacy answer", " ".join(texts(out)))

    def test_a_mapping_that_matches_neither_contract_is_still_refused(self):
        """The fallback must not become a way to fill with any mapping at all."""
        path = self.mixed()
        bogus = [{"key": "cell_0", "label": "Invented", "kind": "text", "locator": {"cell": 0}},
                 {"key": "cell_9", "label": "Also invented", "kind": "text", "locator": {"cell": 9}}]
        with self.assertRaises(DocumentError) as raised:
            fill_template(path, bogus, {}, str(Path(self.directory.name) / "bogus.docx"))
        self.assertEqual(raised.exception.code, "mapping_mismatch")


class OwnerFormUnchanged(unittest.TestCase):
    """The shipped colonoscopy form must be untouched by any of this."""

    def test_shape_is_exactly_as_before(self):
        result = inspect_template(str(OWNER_FORM))
        fields = result["fields"]
        self.assertEqual(len(fields), 24)
        self.assertEqual(sum(len(f["locator"]["tab_form"]["anchors"]) for f in fields), 28)
        self.assertTrue(all("tab_form" in f["locator"] for f in fields))
        self.assertEqual([w for w in result["warnings"] if "NOT mapped" in w], [],
                         "a clean tab form must collect no unmapped-area note")

    def test_it_still_fills_through_the_unchanged_tab_writer(self):
        with tempfile.TemporaryDirectory() as directory:
            fields = inspect_template(str(OWNER_FORM))["fields"]
            answers = {f["key"]: f"Fictional {i}" for i, f in enumerate(fields)}
            out = str(Path(directory) / "owner.docx")
            summary = fill_template(str(OWNER_FORM), fields, answers, out)
            self.assertEqual(summary["filled"], 24)
            self.assertEqual(summary["missing"], [])


if __name__ == "__main__":
    unittest.main()


class LegacyMappingSurvivesAContractOnlyRefusal(Base):
    """⚠️ THE FALLBACK EXISTED AND COULD NOT BE REACHED.

    `process` inspects under the CURRENT contract before it can compare structures, so a
    refusal the current contract invented — a fixed-height answer cell, or a field count the
    wider contract pushes past the limit — escaped first. An encounter whose every field a
    clinician had already reviewed became un-finalizable and un-exportable, refused by name
    over a table row the doctor never mapped. Measured on both triggers.
    """

    def exact_height_pairs(self, name="exact-height.docx"):
        document = Document()
        two = document.add_table(rows=1, cols=2)
        two.rows[0].cells[0].text = "Allergies"
        wide = grid(document, [["Height", "", "Weight", ""]], 4)
        properties = wide.rows[0]._tr.get_or_add_trPr()
        height = OxmlElement("w:trHeight")
        height.set(qn("w:val"), "340")
        height.set(qn("w:hRule"), "exact")
        properties.append(height)
        return self.save(document, name)

    def stored_legacy_mapping(self, path):
        return docx_inspect(Path(path), contract=CONTRACT_V1)[3]

    def test_the_current_contract_really_does_refuse_this_document(self):
        """The precondition. Without it every assertion below could pass for the wrong reason."""
        with self.assertRaises(DocumentError) as raised:
            inspect_template(self.exact_height_pairs())
        self.assertEqual(raised.exception.code, "overflow_risk")

    def test_a_mapping_made_before_that_refusal_existed_still_fills(self):
        path = self.exact_height_pairs()
        stored = self.stored_legacy_mapping(path)
        self.assertEqual([f["label"] for f in stored], ["Allergies"])
        out, summary, _ = self.fill(path, "legacy-out.docx",
                                    answers={stored[0]["key"]: "Penicillin - anaphylaxis"},
                                    fields=stored)
        self.assertEqual(summary["filled"], 1)
        with ZipFile(out) as package:
            body = package.read("word/document.xml").decode()
        self.assertIn("Penicillin - anaphylaxis", body)

    def test_a_fresh_upload_of_the_same_document_is_still_refused(self):
        """The fallback is a compatibility path, not a way to weaken the current contract."""
        with self.assertRaises(DocumentError) as raised:
            inspect_template(self.exact_height_pairs())
        self.assertEqual(raised.exception.code, "overflow_risk")

    def test_a_mapping_that_does_not_match_the_legacy_inspection_is_refused(self):
        path = self.exact_height_pairs()
        wrong = [{**f, "locator": {"cell": 99}} for f in self.stored_legacy_mapping(path)]
        with self.assertRaises(DocumentError) as raised:
            self.fill(path, "wrong-out.docx", answers={}, fields=wrong)
        self.assertEqual(raised.exception.code, "overflow_risk")
