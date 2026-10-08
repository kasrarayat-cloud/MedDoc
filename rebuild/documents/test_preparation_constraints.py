"""Synthetic counterexamples for newly reviewed underscore answer constraints."""
from copy import deepcopy
import hashlib
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile
from docx import Document

from .engine import answer_constraint_error, docx_inspect, structure
from .preparation import candidate, constraint_from_context, inventory
from .service import DocumentError, fill_template


class PreparationConstraintTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def prepare(self, paragraphs):
        original = self.root / 'original.docx'
        doc = Document()
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
        doc.save(original)
        value = inventory(original)
        prepared = self.root / 'prepared.docx'
        result = candidate(original, prepared, value, [anchor['id'] for anchor in value['anchors']])
        return original, prepared, value, result['prepared_fields']

    def test_exact_original_underscore_constraints_and_signature_are_preserved(self):
        original, prepared, value, fields = self.prepare([
            'Choose one listed option for Fever. Preserve °C. Patient signature must remain blank.',
            'Patient name: ____________________', 'Date of birth: __________',
            'Temperature: ______ °C', 'Fever: Yes / No ______', 'Patient signature: ____________________'])
        self.assertEqual([field['answer_constraints'] for field in fields], [
            {'version': 1, 'data_type': 'text'}, {'version': 1, 'data_type': 'text'},
            {'version': 1, 'data_type': 'number', 'unit': '°C'},
            {'version': 1, 'data_type': 'choice', 'options': ['Yes', 'No']}])
        self.assertEqual([field['kind'] for field in fields], ['text'] * 4)
        with ZipFile(original) as before, ZipFile(prepared) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(before.read(name), after.read(name))
        visible = '\n'.join(paragraph.text for paragraph in Document(prepared).paragraphs)
        self.assertIn('Fever: Yes / No {{Fever: Yes / No}}', visible)
        self.assertIn('Patient signature: ____________________', visible)
        self.assertTrue(all(not question.get('answer_constraints') for question in value['questions']
                            if question['protected_signing']))

    def test_incorrect_choice_wrong_unit_and_nonnumeric_value_cannot_export(self):
        original, prepared, _, fields = self.prepare(['Temperature: ____ °C', 'Fever: ____ (Yes / No)'])
        source_hash = hashlib.sha256(original.read_bytes()).hexdigest()
        for index, bad in [(0, '98.6 °F'), (0, 'normal'), (0, 'NaN'), (1, 'Sometimes'), (1, 'Yes / No')]:
            answers = {fields[0]['key']: '37.2 °C', fields[1]['key']: 'Yes'}
            answers[fields[index]['key']] = bad
            output = self.root / f'bad-{index}-{len(bad)}.docx'
            with self.subTest(bad=bad), self.assertRaises(DocumentError) as refused:
                fill_template(prepared, fields, answers, output)
            self.assertEqual(refused.exception.code, 'invalid_answer_constraint')
            self.assertFalse(output.exists())
        self.assertEqual(hashlib.sha256(original.read_bytes()).hexdigest(), source_hash)

    def test_correct_measurement_writes_printed_unit_exactly_once_and_reuses_mapping(self):
        _, prepared, _, fields = self.prepare(['Temperature: ____ °C', 'Fever: ____ (Yes / No)'])
        before = deepcopy(fields)
        for name, temperature in [('typed', '37.2 °C'), ('reviewed-voice', '37.2')]:
            output = self.root / (name + '.docx')
            fill_template(prepared, fields, {fields[0]['key']: temperature, fields[1]['key']: 'No'}, output)
            visible = '\n'.join(paragraph.text for paragraph in Document(output).paragraphs)
            self.assertIn('Temperature: 37.2 °C', visible)
            self.assertEqual(visible.count('°C'), 1)
            self.assertIn('Fever: No (Yes / No)', visible)
            self.assertEqual(fields, before)

    def test_only_explicit_date_format_and_integer_declaration_are_enforced(self):
        _, prepared, _, fields = self.prepare(['Date of birth: ____', 'Visit date: ____ YYYY-MM-DD',
                                               'Count: ____ (integer)'])
        self.assertEqual(fields[0]['answer_constraints']['data_type'], 'text')
        self.assertIsNone(answer_constraint_error(fields[0], 'format not established by original'))
        for bad in ['2026-02-30', '08/10/2026', '2026-10-8']:
            self.assertIsNotNone(answer_constraint_error(fields[1], bad))
        self.assertIsNone(answer_constraint_error(fields[1], '2026-10-08'))
        self.assertIsNotNone(answer_constraint_error(fields[2], '2.5'))
        self.assertIsNone(answer_constraint_error(fields[2], '-2'))
        for field in fields:
            self.assertIsNone(answer_constraint_error(field, ''))
        valid = {fields[0]['key']: 'no format declared', fields[1]['key']: '2026-10-08', fields[2]['key']: '2'}
        fill_template(prepared, fields, valid, self.root / 'date-integer.docx')
        invalid = {**valid, fields[1]['key']: '2026-02-30'}
        with self.assertRaises(DocumentError) as refused:
            fill_template(prepared, fields, invalid, self.root / 'invalid-date.docx')
        self.assertEqual(refused.exception.code, 'invalid_answer_constraint')

    def test_malformed_constraint_schema_is_rejected_even_for_a_blank_answer(self):
        for metadata in [{'version': True, 'data_type': 'text'}, {'version': 1, 'data_type': []},
                         {'version': 1, 'data_type': 'choice'}, {'version': 1, 'data_type': 'date'},
                         {'version': 1, 'data_type': 'number', 'unit': 'invented'},
                         {'version': 1, 'data_type': 'text', 'options': ['Yes', 'No']}]:
            with self.subTest(metadata=metadata):
                self.assertIsNotNone(answer_constraint_error({'label': 'Question', 'answer_constraints': metadata}, ''))

    def test_arbitrary_slash_text_and_combined_declarations_remain_uncertain(self):
        for prefix, suffix in [('Location: Hand / Foot ', ''), ('Date of birth: ', ''),
                               ('Pressure: ', 'mmHg or kPa'), ('Count: ', 'as clinically appropriate'),
                               ('Question: Yes / No ', '°C'), ('Location (Hand / Foot): ', ''),
                               ('Weight (kg): ', 'g')]:
            with self.subTest(prefix=prefix, suffix=suffix):
                self.assertEqual(constraint_from_context(prefix, suffix), {'version': 1, 'data_type': 'text'})
        self.assertEqual(constraint_from_context('Grade: choose one: Mild / Moderate / Severe ', ''),
                         {'version': 1, 'data_type': 'choice', 'options': ['Mild', 'Moderate', 'Severe']})

    def test_final_parenthetical_unit_and_explicit_choice_are_bound_before_blank(self):
        _, prepared, _, fields = self.prepare(['Weight (kg): ____', 'Fever (Yes / No): ____',
                                               'Grade (choose one: Mild / Moderate / Severe): ____'])
        self.assertEqual(fields[0]['answer_constraints'], {'version': 1, 'data_type': 'number', 'unit': 'kg'})
        self.assertEqual(fields[1]['answer_constraints'], {'version': 1, 'data_type': 'choice', 'options': ['Yes', 'No']})
        self.assertEqual(fields[2]['answer_constraints'], {'version': 1, 'data_type': 'choice',
                                                        'options': ['Mild', 'Moderate', 'Severe']})
        answers = {fields[0]['key']: '70 kg', fields[1]['key']: 'No', fields[2]['key']: 'Mild'}
        output = self.root / 'parenthetical-export.docx'
        fill_template(prepared, fields, answers, output)
        self.assertEqual(Document(output).paragraphs[0].text, 'Weight (kg): 70')
        self.assertIsNotNone(answer_constraint_error(fields[0], '154 lb'))
        self.assertIsNotNone(answer_constraint_error(fields[1], 'Unknown'))

    def test_parenthesized_plural_marker_never_becomes_a_seconds_constraint(self):
        _, prepared, _, fields = self.prepare(['Medication(s): ____', 'Allergy(s): ____',
                                               'Diagnosis (s): ____', 'Duration (s): ____'])
        self.assertTrue(all(field['answer_constraints'] == {'version': 1, 'data_type': 'text'} for field in fields))
        answers = {field['key']: value for field, value in zip(fields,
                   ['Synthetic Drug X', 'Synthetic Drug Y', 'Synthetic condition', 'Not specified'])}
        output = self.root / 'plural-export.docx'
        fill_template(prepared, fields, answers, output)
        visible = [paragraph.text for paragraph in Document(output).paragraphs]
        self.assertEqual(visible, ['Medication(s): Synthetic Drug X', 'Allergy(s): Synthetic Drug Y',
                                   'Diagnosis (s): Synthetic condition', 'Duration (s): Not specified'])

    def test_explicit_seconds_remain_enforced_without_guessing_from_a_label(self):
        _, prepared, _, fields = self.prepare(['Elapsed: ____ s', 'Elapsed (unit: s): ____'])
        self.assertTrue(all(field['answer_constraints'] == {'version': 1, 'data_type': 'number', 'unit': 's'}
                            for field in fields))
        output = self.root / 'seconds-export.docx'
        fill_template(prepared, fields, {field['key']: '2 s' for field in fields}, output)
        self.assertEqual([paragraph.text for paragraph in Document(output).paragraphs],
                         ['Elapsed: 2 s', 'Elapsed (unit: s): 2'])
        self.assertIsNotNone(answer_constraint_error(fields[0], 'Synthetic Drug X'))

    def test_constraint_forgery_and_placement_drift_are_rejected(self):
        _, prepared, _, fields = self.prepare(['Temperature: ____ °C', 'Fever: ____ (Yes / No)'])
        for mutation in ['unit', 'options', 'locator']:
            forged = deepcopy(fields)
            if mutation == 'unit':
                forged[0]['answer_constraints']['unit'] = '°F'
            elif mutation == 'options':
                forged[1]['answer_constraints']['options'].append('Sometimes')
            else:
                forged[0]['locator']['paragraph'] = 1
            with self.subTest(mutation=mutation), self.assertRaises(DocumentError) as refused:
                fill_template(prepared, forged, {forged[0]['key']: '', forged[1]['key']: ''},
                              self.root / (mutation + '.docx'))
            self.assertEqual(refused.exception.code, 'mapping_mismatch')

    def test_old_approved_mapping_semantics_and_new_version_are_distinct(self):
        _, prepared, _, fields = self.prepare(['Temperature: ____ °C'])
        legacy = docx_inspect(prepared)[3]
        legacy_snapshot = deepcopy(legacy)
        self.assertNotEqual(structure(fields), structure(legacy))
        # A retained old seal continues to mean exactly its original text map.
        old_output = self.root / 'old-approved-export.docx'
        fill_template(prepared, legacy, {legacy[0]['key']: 'old free text'}, old_output)
        self.assertEqual(legacy, legacy_snapshot)
        self.assertIn('old free text', Document(old_output).paragraphs[0].text)
        with self.assertRaises(DocumentError):
            fill_template(prepared, fields, {fields[0]['key']: 'old free text'}, self.root / 'new-invalid.docx')
