"""Local selected objects remain distinct from referenced condition data."""
from copy import deepcopy
import io
import unittest
from zipfile import ZipFile

from docx import Document
from . import word_v3 as w
from . import test_source_coverage as coverage


class NamedLocalObjectScope(unittest.TestCase):
    positive = ('Modality – Select Home if received, as defined in Quantity above. '
                'Select Work when not received. Select Remote when unclear.')

    def form(self, instruction=None, **kwargs):
        return coverage.NamedSelectionInstructions().form(instruction or self.positive, **kwargs)

    def group(self, info):
        return next(f for f in info['fields']if f['label']=='Modality')

    def test_complete_selected_objects_are_local_and_all_original_ink_is_preserved(self):
        for native in (False,True):
            for reference in ('Quantity above','‘Quantity’ above','"Recorded amount" below'):
                with self.subTest(native=native,reference=reference):
                    instruction=self.positive.replace('Quantity above',reference)
                    raw=self.form(instruction,native=native);info=w.inspect(raw);field=self.group(info)
                    self.assertEqual((field['kind'],field['options']),('choice',['Home','Work','Remote']))
                    relation=next(r for r in field['source_questions']if r['relation']=='named_selection_instruction')
                    self.assertEqual(relation['selected_objects'],['Home','Work','Remote'])
                    self.assertEqual(relation['instructions'],[instruction.split(' – ',1)[1]])
                    output=w.fill(raw,info['fields'],{field['key']:'Work'})
                    after=w.SlotIndex(w.open_package(output))
                    self.assertEqual(w.constrained_option_states(after,field),[False,True,False])
                    before_doc,after_doc=Document(io.BytesIO(raw)),Document(io.BytesIO(output))
                    self.assertEqual(before_doc.tables[0].cell(0,1)._tc.xml,after_doc.tables[0].cell(0,1)._tc.xml)
                    self.assertEqual([p._p.xml for p in before_doc.paragraphs],[p._p.xml for p in after_doc.paragraphs])
                    with ZipFile(io.BytesIO(raw))as before,ZipFile(io.BytesIO(output))as after:
                        for name in before.namelist():
                            if name!='word/document.xml':self.assertEqual(before.read(name),after.read(name))
                    self.assertEqual(w.fill(raw,info['fields'],{}),raw)

    def test_printed_skip_qualifiers_remain_exact_answer_names_and_unselected_controls(self):
        instruction=('Modality – Select ‘Yes’ if received, as defined in ‘Quantity’ above. '
                     'If the record is complete, select ‘Yes.’ Select ‘No’ when absent (e.g., no record). '
                     'Select ‘unknown’ when unclear.')
        raw=self.form(instruction,native=True);doc=Document(io.BytesIO(raw))
        for p,option in zip(doc.tables[0].cell(0,0).paragraphs[1:],('Yes','No (skip to #12)','Unknown (skip to end)')):
            p.runs[-1].text=' '+option
        raw=coverage.data(doc);info=w.inspect(raw);field=self.group(info)
        self.assertEqual(field['options'],['Yes','No (skip to #12)','Unknown (skip to end)'])
        with self.assertRaises(w.V3Failure):w.fill(raw,info['fields'],{field['key']:'No'})
        output=w.fill(raw,info['fields'],{field['key']:'No (skip to #12)'})
        self.assertEqual(w.constrained_option_states(w.SlotIndex(w.open_package(output)),field),[False,True,False])
        before,after=Document(io.BytesIO(raw)),Document(io.BytesIO(output))
        self.assertEqual([p.text for p in before.tables[0].cell(0,0).paragraphs],
                         [p.text for p in after.tables[0].cell(0,0).paragraphs])
        self.assertEqual(before.paragraphs[-1]._p.xml,after.paragraphs[-1]._p.xml)
        self.assertEqual(w.fill(raw,info['fields'],{}),raw)

    def test_selection_reach_unknown_objects_and_malformed_framing_remain_held(self):
        for instruction in ('Modality – Select options from the question above.',
                'Modality – Select Home for every question. Select Work when absent. Select Remote when unclear.',
                self.positive.replace('if received, as defined in Quantity above','if required for the question above'),
                self.positive.replace('Select Home','Select Other'),
                self.positive.replace('Quantity above','Quantity above and all following questions'),
                self.positive.replace('Select Home','Do not select Home'),
                self.positive.replace('Quantity above','(Quantity above'),
                self.positive.replace('Home if','‘Home if')):
            with self.subTest(instruction=instruction):
                raw=self.form(instruction);info=w.inspect(raw)
                self.assertFalse([f for f in info['fields']if f['label']=='Modality'])
                self.assertTrue(any(i['code']=='named_selection_source'for i in info['coverage_issues']))
                self.assertEqual(w.fill(raw,info['fields'],{}),raw)

    def test_competing_many_hidden_duplicate_and_extended_questions_do_not_gain_scope(self):
        for config in ({'extra_instruction':'Modality – Select all relevant options.'},
                       {'duplicate':True},{'hidden':True},
                       {'instruction':self.positive.replace('Modality –','Modality details –')}):
            with self.subTest(config=config):
                raw=self.form(**config);info=w.inspect(raw)
                own=[f for f in info['fields']if f['label']=='Modality']
                self.assertFalse(any(any(r.get('selected_objects')for r in f.get('source_questions',[]))for f in own))
                self.assertFalse(any(f['kind']=='multi_choice'for f in own))
                self.assertEqual(w.fill(raw,info['fields'],{}),raw)

    def test_human_instructions_keep_actual_source_roles_even_with_a_forged_map(self):
        valid=self.form();fields=w.inspect(valid)['fields'];field=self.group({'fields':fields})
        for act in ('I approve this document','This area is for Alex Reed to enter the answer',
                    'Witness must write the following','Alex Reed is to write his comments'):
            with self.subTest(act=act):
                raw=self.form(self.positive+';'+act+':');info=w.inspect(raw)
                self.assertFalse([f for f in info['fields']if f['label']=='Modality'and not f.get('write_protected')])
                forged=deepcopy(field);index=w.SlotIndex(w.open_package(raw))
                for ref in [forged['locator'],*forged['locator'].get('extra',[]),*forged['locator']['target']['boxes']]:
                    slot=index.by_id[ref['part']+'#'+ref['slot']];ref['fingerprint']=w.fingerprint(slot.el)
                with self.assertRaises(w.V3Failure):w.fill(raw,[forged],{forged['key']:'Home'})
                self.assertEqual(w.fill(raw,info['fields'],{}),raw)

    def test_original_command_and_physical_metadata_cannot_be_removed_or_rewritten(self):
        raw=self.form();info=w.inspect(raw);field=self.group(info)
        for mutation in ('selected-objects','instructions','source','relationship','context','kind','options','boxes'):
            with self.subTest(mutation=mutation):
                forged=deepcopy(field)
                relation=next(r for r in forged['source_questions']if r['relation']=='named_selection_instruction')
                if mutation=='selected-objects':relation['selected_objects'].reverse()
                elif mutation=='instructions':relation['instructions']=['Select all relevant options']
                elif mutation=='source':relation['sources'][-1]['fingerprint']='0'*64
                elif mutation=='relationship':forged.pop('source_questions')
                elif mutation=='context':forged['context']='Another question'
                elif mutation=='kind':forged['kind']='multi_choice'
                elif mutation=='options':forged['options'].reverse()
                else:forged['locator']['target']['boxes'].reverse()
                with self.assertRaises(w.V3Failure):w.fill(raw,[forged],{forged['key']:'Home'})

    def test_scope_proof_does_not_supply_a_cardinality_command(self):
        raw=self.form();det=w.Detector(w.open_package(raw));det.detect()
        field=self.group(w.inspect(raw));owners,entries=det.named_guidance_index();owner=owners['modality'][0]
        result=det.named_selection_source('Modality',{id(owner.el)},field['options'])
        self.assertEqual(result['commands'],[])
        self.assertEqual(result['selected_objects'],['Home','Work','Remote'])
        self.assertIsNone(det.named_local_object_scope([(owner,'Select Home if defined above.')],field['options']))


class NamedConditionSourceActs(unittest.TestCase):
    suffix = 'Select Work when not received. Select Remote when unclear.'

    def test_complete_nested_appointed_and_ambiguous_response_clauses_stay_held(self):
        clauses = (
            'Select Home if applicable (this area is for Alex Reed to enter the answer), as defined in Quantity above.',
            'Select Home if applicable (Alex Reed must complete this answer), as defined in Quantity above.',
            'Select Home if applicable, as defined in Quantity above, and Alex Reed is to write his comments in this answer area.',
            'Select Home if received, as defined in Quantity above. Alex Reed is to write his comments here.',
            'Select Home if the manager signs this question, as defined in Quantity above.',
            'Select Home if applicable (context (I approve this document)), as defined in Quantity above.',
            'Select Home if applicable (Dr. Avery is required to enter this answer), as defined in Quantity above.',
            'Select Home if applicable (Alex Reed must not complete this answer), as defined in Quantity above.',
        )
        for native in (False, True):
            for clause in clauses:
                with self.subTest(native=native, clause=clause):
                    instruction = 'Modality – ' + clause + ' ' + self.suffix
                    self.assertTrue(w.Detector.named_instruction_source_action(instruction))
                    raw = coverage.NamedSelectionInstructions().form(instruction, native=native)
                    info = w.inspect(raw)
                    self.assertFalse([f for f in info['fields'] if f['label'] == 'Modality' and not f.get('write_protected')])
                    self.assertTrue(any('human' in x['detail'].casefold() for x in info['coverage_issues']))
                    self.assertEqual(w.fill(raw, info['fields'], {}), raw)

    def test_clinical_recorded_conditions_remain_local_with_identical_ink(self):
        for condition in ('the manager recorded the observation',
                          'the patient has clinical signs',
                          'the patient had initial symptoms',
                          'the signed referral was received',
                          'applicable (the clinician recorded initial symptoms)',
                          'applicable (the manager recorded the observation)'):
            for native in (False, True):
                with self.subTest(condition=condition, native=native):
                    instruction = 'Modality – Select Home if ' + condition + ', as defined in Quantity above. ' + self.suffix
                    self.assertFalse(w.Detector.named_instruction_source_action(instruction))
                    raw = coverage.NamedSelectionInstructions().form(instruction, native=native)
                    info = w.inspect(raw); field = next(f for f in info['fields'] if f['label'] == 'Modality')
                    output = w.fill(raw, info['fields'], {field['key']: 'Work'})
                    self.assertEqual(w.constrained_option_states(w.SlotIndex(w.open_package(output)), field), [False, True, False])
                    before, after = Document(io.BytesIO(raw)), Document(io.BytesIO(output))
                    self.assertEqual([p._p.xml for p in before.paragraphs], [p._p.xml for p in after.paragraphs])
                    self.assertEqual(before.tables[0].cell(0,1)._tc.xml, after.tables[0].cell(0,1)._tc.xml)
                    self.assertEqual(w.fill(raw, info['fields'], {}), raw)

    def test_existing_general_authorship_reader_still_requires_the_writing_relation(self):
        for text in ('Alex Reed must write this answer', 'Alex Reed should enter the response',
                     'This area is for Alex Reed to record his comments'):
            self.assertTrue(w.Detector.question_box_authorship(text))
        for text in ('Alex Reed must complete this answer', 'the manager recorded the observation',
                     'This section is for the clinician to record initial symptoms',
                     'This section is for the clinician to record signed document receipt'):
            self.assertFalse(w.Detector.question_box_authorship(text))
        self.assertTrue(w.Detector.prospective_response('Alex Reed must complete this answer', located=True))
        self.assertFalse(w.Detector.prospective_response('The patient must start the medication', located=True))


class NamedBalancedSourceQualifiers(unittest.TestCase):
    action = 'this area is for Alex Reed to enter the answer'
    suffix = ', as defined in Quantity above. Select Work when not received. Select Remote when unclear.'

    def test_complete_bracketed_quoted_and_nested_authorship_is_not_a_scope_escape(self):
        for wrapper in ('[{}]', '{{{}}}', '“{}”', '‘{}’', '"{}"', "'{}'", '([{}])', '“context ({})”', '“context {{{}}}”'):
            for native in (False, True):
                with self.subTest(wrapper=wrapper, native=native):
                    instruction = 'Modality – Select Home if applicable ' + wrapper.format(self.action) + self.suffix
                    self.assertTrue(w.Detector.named_instruction_source_action(instruction))
                    raw = coverage.NamedSelectionInstructions().form(instruction, native=native); info = w.inspect(raw)
                    self.assertFalse([f for f in info['fields'] if f['label'] == 'Modality' and not f.get('write_protected')])
                    self.assertEqual(w.fill(raw, info['fields'], {}), raw)

    def test_ordinary_balanced_qualifiers_and_quoted_definition_names_remain_usable(self):
        for qualifier in ('[the manager has recorded the observation]', '{the manager recorded the observation}', '“the manager recorded the observation”',
                          '[the patient had several initial symptoms]', '“context (the signed referral was received)”'):
            instruction = 'Modality – Select Home if applicable ' + qualifier + self.suffix
            raw = coverage.NamedSelectionInstructions().form(instruction); info = w.inspect(raw)
            field = next(f for f in info['fields'] if f['label'] == 'Modality')
            output = w.fill(raw, info['fields'], {field['key']: 'Work'})
            self.assertEqual(w.constrained_option_states(w.SlotIndex(w.open_package(output)), field), [False, True, False])
            before, after = Document(io.BytesIO(raw)), Document(io.BytesIO(output))
            self.assertEqual(before.paragraphs[-1]._p.xml, after.paragraphs[-1]._p.xml)
            self.assertEqual(w.fill(raw, info['fields'], {}), raw)

    def test_short_bracket_placeholders_do_not_become_ordinary_instruction_ink(self):
        # The established source reader interprets these short bracketed noun
        # phrases as answer placeholders. A new scope proof cannot erase that
        # separate answer destination merely because its words look clinical.
        for qualifier in ('[the manager recorded the observation]', '[the patient had initial symptoms]'):
            instruction = 'Modality – Select Home if applicable ' + qualifier + self.suffix
            self.assertFalse(w.Detector.named_instruction_source_action(instruction))
            self.assertTrue(list(w.BRACKET.finditer(qualifier)))
            raw = coverage.NamedSelectionInstructions().form(instruction); info = w.inspect(raw)
            self.assertFalse([f for f in info['fields'] if f['label'] == 'Modality'])
            self.assertTrue(any(x['code'] == 'named_selection_source' for x in info['coverage_issues']))
            self.assertEqual(w.fill(raw, info['fields'], {}), raw)

    def test_malformed_qualifiers_cannot_prove_a_complete_named_scope(self):
        for qualifier in ('[the manager recorded the observation)', '(the manager recorded the observation]',
                          '[the manager recorded the observation', '“the manager recorded the observation',
                          '{the manager recorded the observation]', '{the manager recorded the observation'):
            instruction = 'Modality – Select Home if applicable ' + qualifier + self.suffix
            self.assertIsNone(w.Detector.named_instruction_clauses(instruction))
            raw = coverage.NamedSelectionInstructions().form(instruction); info = w.inspect(raw)
            self.assertFalse([f for f in info['fields'] if f['label'] == 'Modality'])
            self.assertTrue(any(x['code'] == 'named_selection_source' for x in info['coverage_issues']))
            self.assertEqual(w.fill(raw, info['fields'], {}), raw)
