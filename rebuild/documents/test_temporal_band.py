"""Pure physical ownership controls; no application, DB, network or renderer."""
from copy import deepcopy
import io
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree

from . import word_v3 as word


def band_form(caption='Date of assessment:', pattern='ddmmmyyyy', *, populated=False,
              separators=(2, 5), foreign_guide=False, hidden=False):
    document = Document()
    document.add_paragraph(caption)
    paragraph = document.add_paragraph()
    for n, _ in enumerate(pattern):
        if n in separators:
            paragraph.add_run(' / ')
        if n == 2:
            start = OxmlElement('w:bookmarkStart')
            start.set(word.q('w:id'), '17'); start.set(word.q('w:name'), 'month')
            paragraph._p.append(start)
        run = paragraph.add_run('Z' if populated and n == 1 else '   ')
        run.underline = True
        if hidden:
            run.font.hidden = True
        if n == 4:
            end = OxmlElement('w:bookmarkEnd'); end.set(word.q('w:id'), '17')
            paragraph._p.append(end)
        if n + 1 < len(pattern) and n + 1 not in separators:
            paragraph.add_run(' ')
    paragraph.add_run('.')
    if foreign_guide:
        document.add_table(rows=1, cols=1).cell(0, 0).text = pattern
    else:
        document.add_paragraph('\t'.join(pattern))
    document.add_paragraph('Plan: ____________')
    document.add_paragraph('Decision: ☐ Yes ☐ No')
    output = io.BytesIO(); document.save(output)
    return output.getvalue()


def detected(data):
    return word.inspect(data)['fields']


def bands(data):
    return [f for f in detected(data) if f['locator']['target'].get('shape') == 'temporal_band']


class TemporalBandOwnershipTests(unittest.TestCase):
    def test_one_field_owns_caption_complete_band_and_guide(self):
        data = band_form()
        fields = detected(data)
        dates = [f for f in fields if f['kind'] == 'date']
        self.assertEqual(len(dates), 1)
        self.assertEqual(dates[0]['label'], 'Date of assessment')
        self.assertEqual(dates[0]['locator']['slot'], 'b/p[1]')
        self.assertEqual(len(dates[0]['locator']['target']['parts']), 9)
        self.assertFalse(any(f['label'] == '/' for f in fields))
        self.assertFalse(any(f['locator']['slot'] in ('b/p[0]', 'b/p[2]') for f in fields))

    def test_surgical_writes_preserve_package_neighbours_separators_and_bookmarks(self):
        data = band_form(); field = bands(data)[0]
        final = word.fill(data, [field], {field['key']: '2026-10-05'})
        with ZipFile(io.BytesIO(data)) as before, ZipFile(io.BytesIO(final)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(before.read(name), after.read(name))
        original = word.SlotIndex(word.open_package(data)); written = word.SlotIndex(word.open_package(final))
        for identity, slot in original.by_id.items():
            if slot.kind == 'p' and identity != 'word/document.xml#b/p[1]':
                self.assertEqual(word.fingerprint(slot.el), word.fingerprint(written.by_id[identity].el))
        old = original.by_id['word/document.xml#b/p[1]'].el
        new = written.by_id['word/document.xml#b/p[1]'].el
        self.assertEqual(word.paragraph_text(new)[0], '0 5 / O c t / 2 0 2 6.')
        for tag in ('w:pPr', 'w:bookmarkStart', 'w:bookmarkEnd'):
            self.assertEqual([etree.tostring(n) for n in old.findall(word.q(tag))],
                             [etree.tostring(n) for n in new.findall(word.q(tag))])
        self.assertEqual([etree.tostring(n) for n in old.findall(word.R) if '/' in word.text_of(n)],
                         [etree.tostring(n) for n in new.findall(word.R) if '/' in word.text_of(n)])
        self.assertEqual(word.fill(data, [field], {}), data)

    def test_formats_follow_their_own_guide_and_calendar(self):
        for pattern, separators, expected in [('ddmmyyyy', (2, 4), '0 5 / 1 0 / 2 0 2 6.'),
                ('mmddyyyy', (2, 4), '1 0 / 0 5 / 2 0 2 6.'), ('yyyymmdd', (4, 6), '2 0 2 6 / 1 0 / 0 5.')]:
            with self.subTest(pattern=pattern):
                data = band_form(pattern=pattern, separators=separators); field = bands(data)[0]
                output = word.fill(data, [field], {field['key']: '2026-10-05'})
                index = word.SlotIndex(word.open_package(output))
                self.assertEqual(word.paragraph_text(index.by_id['word/document.xml#b/p[1]'].el)[0], expected)
        data = band_form(); field = bands(data)[0]
        for value in ('2026-02-30', '05/gar/2026', '123456789', '2026-10-05 extra', '5/Oct/2026'):
            with self.subTest(value=value), self.assertRaises(word.V3Failure):
                word.fill(data, [field], {field['key']: value})

    def test_readback_requires_every_original_separator_and_returns_one_date(self):
        data = band_form(); field = bands(data)[0]
        final = word.fill(data, [field], {field['key']: '2026-10-05'})
        original = word.SlotIndex(word.open_package(data))
        index = word.SlotIndex(word.open_package(final))
        self.assertEqual(word.constrained_scalar_value(original, index, field, True), '05/Oct/2026')
        self.assertEqual(word.constrained_scalar_value(original, original, field, False), '')
        paragraph = index.by_id['word/document.xml#b/p[1]'].el
        slash = next(node for node in paragraph.iter(word.T) if '/' in (node.text or ''))
        slash.text = slash.text.replace('/', '-')
        self.assertIsNone(word.constrained_scalar_value(original, index, field, True))

    def test_normal_effective_values_reads_iso_and_printed_date_inputs(self):
        data = band_form(); field = bands(data)[0]; key = field['key']
        for answer in ('2026-10-05', '05/Oct/2026'):
            with self.subTest(answer=answer):
                self.assertEqual(word.effective_values(data, [field], {key: answer}),
                                 {key: '05/Oct/2026'})
        self.assertEqual(word.effective_values(data, [field], {}), {key: ''})

    def test_normal_effective_values_rejects_mutated_written_band(self):
        data = band_form(); field = bands(data)[0]; key = field['key']
        answers = {key: '2026-10-05'}
        final = word.fill(data, [field], answers)
        for mutation in ('separator', 'calendar', 'hidden', 'guide', 'caption'):
            with self.subTest(mutation=mutation):
                package = word.open_package(final)
                index = word.SlotIndex(package)
                paragraph = index.by_id['word/document.xml#b/p[1]'].el
                if mutation == 'separator':
                    node = next(n for n in paragraph.iter(word.T) if '/' in (n.text or ''))
                    node.text = node.text.replace('/', '-')
                elif mutation == 'calendar':
                    characters = [n for n in paragraph.iter(word.T)
                                  if (n.text or '').strip() and '/' not in n.text and n.text != '.']
                    for node, value in zip(characters, '30Feb2026'):node.text = value
                elif mutation == 'hidden':
                    run = next(n for n in paragraph.findall(word.R)
                               if word.text_of(n).strip() and '/' not in word.text_of(n))
                    properties = run.find(word.q('w:rPr'))
                    if properties is None:
                        properties = OxmlElement('w:rPr'); run.insert(0, properties)
                    properties.append(OxmlElement('w:vanish'))
                else:
                    path = 'b/p[2]' if mutation == 'guide' else 'b/p[0]'
                    node = next(index.by_id['word/document.xml#' + path].el.iter(word.T))
                    node.text = 'mm' if mutation == 'guide' else 'Date of witness approval:'
                written = package.write({'word/document.xml': package.tree('word/document.xml')})
                self.assertIsNone(word.effective_values(data, [field], answers, _written=written)[key])

    def test_stage_readback_compares_only_this_source_format_to_its_logical_date(self):
        from ops import sonnet5_stages as stages
        data = band_form(); field = bands(data)[0]; key = field['key']
        for answer in ('2026-10-05', '05/Oct/2026'):
            with self.subTest(answer=answer):
                final = word.fill(data, [field], {key: answer})
                saved = {key: {'text': answer, 'status': 'supported', 'reviewed': True}}
                self.assertTrue(stages.export_read_back(data, final, [field], saved)['passed'])
        final = word.fill(data, [field], {key: '2026-10-05'})
        for answer in ('2026-10-06', '2026-02-30', '05/gar/2026'):
            with self.subTest(answer=answer):
                saved = {key: {'text': answer, 'status': 'supported', 'reviewed': True}}
                self.assertFalse(stages.export_read_back(data, final, [field], saved)['passed'])

    def test_foreign_missing_inconsistent_hidden_or_populated_source_does_not_grant_band(self):
        for options in ({'foreign_guide': True}, {'pattern': 'ddmmyy'}, {'separators': (1, 5)},
                        {'hidden': True}, {'populated': True}):
            with self.subTest(options=options):
                self.assertEqual(bands(band_form(**options)), [])

    def test_signing_date_stays_protected(self):
        data = band_form(caption='Date of witness approval:'); field = bands(data)[0]
        self.assertTrue(field['write_protected'])
        self.assertEqual(word.fill(data, [field], {}), data)
        with self.assertRaises(word.V3Failure):
            word.fill(data, [field], {field['key']: '2026-10-05'})

    def test_source_bound_band_cannot_move_or_drop_its_guide(self):
        data = band_form(); original = bands(data)[0]
        for change in ('guide', 'part', 'format'):
            field = deepcopy(original)
            if change == 'guide':field['locator']['target']['guide']['slot'] = 'b/p[3]'
            elif change == 'part':field['locator']['target']['parts'][0][1] += 1
            else:field['locator']['target']['format'] = 'mmddyyyy'
            with self.subTest(change=change), self.assertRaises(word.V3Failure):
                word.fill(data, [field], {field['key']: '2026-10-05'})

    def test_saved_scalar_locator_keeps_its_original_physical_write(self):
        data = band_form(); index = word.SlotIndex(word.open_package(data))
        slot = index.by_id['word/document.xml#b/p[1]']
        saved = {'key': 'v3_4', 'label': 'Date of assessment', 'kind': 'date', 'context': '',
                 'locator': {'contract': 'v3', 'part': slot.part, 'slot': slot.path,
                             'fingerprint': word.fingerprint(slot.el),
                             'target': {'type': 'span', 'shape': 'utab', 'start': 0, 'end': 3}}}
        output = word.fill(data, [saved], {'v3_4': '5'})
        text = word.paragraph_text(word.SlotIndex(word.open_package(output)).by_id[slot.id].el)[0]
        self.assertEqual(text, '5' + word.paragraph_text(slot.el)[0][3:])
        self.assertEqual(saved['locator']['target']['end'], 3)
