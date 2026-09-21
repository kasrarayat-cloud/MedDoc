"""An answer area that is not mapped must be NAMED, never silently dropped.

Four shapes previously left the importer through a bare `continue` or through mapping
precedence. In every case the import SUCCEEDED, the doctor saw a plausible field list at
mapping review, and the rest of the form was gone with nothing said:

  * a question/answer row that is not two columns          (three- and four-column tables)
  * a row whose cells are merged                           (vMerge / gridSpan / nested table)
  * a blank answer cell holding a line break, image or drawing
  * labelled underlined blanks in a form that ALSO has a table or placeholder mapping,
    which the tab adapter never sees because it runs only when nothing else matched

None of these become mapped here — enabling ambiguous shapes is exactly what the product
should not do. What changes is that each is reported as a review note with a supported
alternative, and that a wholesale refusal explains which shape it found.

Fixtures are authored in this file and their expectations stated before the parser runs.
Synthetic only; no patient data, no network, no database.
"""

from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .service import DocumentError, inspect_template


def underlined_tab(paragraph):
    """A single-underlined tab-only run: the one shape the tab adapter treats as a blank."""
    run = paragraph.add_run()
    properties = OxmlElement("w:rPr")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run._r.append(properties)
    run._r.append(OxmlElement("w:tab"))
    return run


def qa_table(document, rows, columns):
    table = document.add_table(rows=len(rows), cols=columns)
    for index, values in enumerate(rows):
        for column, value in enumerate(values):
            table.rows[index].cells[column].text = value
    return table


class UnmappedAreasAreReported(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)

    def save(self, document, name):
        path = Path(self.directory.name) / name
        document.save(path)
        return str(path)

    def inspect(self, document, name):
        return inspect_template(self.save(document, name))

    def notes(self, result):
        return [w for w in result["warnings"] if "NOT mapped" in w]

    # ---------- the shapes that must now be reported ----------

    def test_four_column_pairs_are_reported_not_dropped(self):
        """Expected independently: 2 label/blank pairs per row, 2 rows -> 4 answer areas."""
        document = Document()
        qa_table(document, [["Height", "", "Weight", ""], ["Pulse", "", "BP", ""]], 4)
        qa_table(document, [["Referring clinician", ""]], 2)
        result = self.inspect(document, "four-column.docx")

        self.assertEqual([f["label"] for f in result["fields"]], ["Referring clinician"])
        notes = self.notes(result)
        self.assertTrue(any("not two columns" in n for n in notes), notes)
        self.assertTrue(any("4 table row" not in n for n in notes))   # 2 rows skipped, not 4
        self.assertTrue(any(n.startswith("2 table row(s)") for n in notes), notes)
        self.assertTrue(any("{{Question name}}" in n for n in notes),
                        f"A note must offer a supported alternative, literally. Got: {notes}")

    def test_merged_answer_cells_are_reported_not_dropped(self):
        """Expected independently: 3 questions; 2 of them share a vertically merged answer."""
        document = Document()
        table = qa_table(document, [["Patient number", ""], ["Medications", ""], ["Dose", ""]], 2)
        for index, value in ((1, "restart"), (2, "continue")):
            properties = table.rows[index].cells[1]._tc.get_or_add_tcPr()
            merge = OxmlElement("w:vMerge")
            merge.set(qn("w:val"), value)
            properties.append(merge)
        result = self.inspect(document, "merged.docx")

        self.assertEqual([f["label"] for f in result["fields"]], ["Patient number"])
        self.assertTrue(any("merged" in n for n in self.notes(result)), self.notes(result))

    def test_answer_cell_with_a_line_break_is_reported_not_dropped(self):
        """Expected independently: 2 questions; one answer cell carries a <w:br>."""
        document = Document()
        table = qa_table(document, [["Chief complaint", ""], ["Plan", ""]], 2)
        paragraph = table.rows[1].cells[1].paragraphs[0]
        paragraph.add_run()._r.append(OxmlElement("w:br"))
        result = self.inspect(document, "line-break.docx")

        self.assertEqual([f["label"] for f in result["fields"]], ["Chief complaint"])
        self.assertTrue(any("line break" in n for n in self.notes(result)), self.notes(result))

    def test_underlined_section_beside_a_table_is_reported_not_dropped(self):
        """Precedence, made visible.

        The CONTROL is the point: the identical underlined section ALONE imports fine, so the
        silence in the combined document was not the section being unsupported — it was the
        tab adapter never being reached.
        """
        alone = Document()
        alone.add_heading("FINDINGS", level=1)
        paragraph = alone.add_paragraph()
        paragraph.add_run("EXTENT: ")
        underlined_tab(paragraph)
        control = self.inspect(alone, "underline-alone.docx")
        self.assertTrue(control["fields"], "CONTROL: the underlined section alone must map.")
        self.assertTrue(any("tab_form" in f["locator"] for f in control["fields"]))

        combined = Document()
        qa_table(combined, [["Referring clinician", ""]], 2)
        combined.add_heading("FINDINGS", level=1)
        paragraph = combined.add_paragraph()
        paragraph.add_run("EXTENT: ")
        underlined_tab(paragraph)
        result = self.inspect(combined, "underline-with-table.docx")

        self.assertEqual([f["label"] for f in result["fields"]], ["Referring clinician"])
        notes = self.notes(result)
        self.assertTrue(any("underlined answer blanks" in n for n in notes), notes)
        self.assertTrue(any("ONE method" in n for n in notes), notes)

    def test_a_wholesale_refusal_explains_which_shape_it_found(self):
        """A form of visible blanks must not be told it contains nothing."""
        document = Document()
        qa_table(document, [["Allergies", "", ""], ["Medications", "", ""]], 3)
        with self.assertRaises(DocumentError) as raised:
            self.inspect(document, "three-column-only.docx")
        self.assertEqual(raised.exception.code, "no_fields")
        message = str(raised.exception)
        self.assertIn("not two columns", message)
        self.assertIn("{{Question name}}", message)

    # ---------- and the silence that must be preserved ----------

    def test_an_ordinary_supported_form_collects_no_notes(self):
        """The guard against advisory noise: a clean form must stay clean."""
        document = Document()
        document.add_paragraph("PATIENT NAME: {{Patient name}}")
        qa_table(document, [["Indication", ""], ["Plan", ""]], 2)
        result = self.inspect(document, "clean.docx")

        self.assertEqual(len(result["fields"]), 3)
        self.assertEqual(self.notes(result), [],
                         "A fully supported form must not be given unmapped-area notes.")

    def test_a_filled_data_table_is_not_mistaken_for_unmapped_answers(self):
        """A three-column table with no blank cell is data, not a skipped question."""
        document = Document()
        qa_table(document, [["Drug", "Dose", "Route"], ["Aspirin", "81 mg", "Oral"]], 3)
        document.add_paragraph("PLAN: {{Plan}}")
        result = self.inspect(document, "data-table.docx")

        self.assertEqual([f["label"] for f in result["fields"]], ["Plan"])
        self.assertEqual(self.notes(result), [],
                         "A fully populated data table must not be reported as a skipped answer area.")


if __name__ == "__main__":
    unittest.main()
