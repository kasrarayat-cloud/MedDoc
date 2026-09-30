"""Word engine precision: signing structures and non-answer text, on synthetic python-docx shapes built here.

Every shape mirrors a STRUCTURE (a placeholder naming the signing act, a caption ruled along its top, an attestation
line, a phrase printed underlined, a bracketed aside inside a sentence ...), never a particular form or its wording.
Each rule has a positive control that must keep today's answer place.
"""
import io
import unittest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import word_v3 as word


def build(*lines, style=None):
    """A document whose body paragraphs are `lines`; a line may be a callable(doc) that adds its own paragraph."""
    doc = Document()
    for line in lines:
        if callable(line):
            line(doc)
        else:
            doc.add_paragraph(line)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def fields_of(data):
    return word.inspect(data)["fields"]


def at(fields, n):
    return [f for f in fields if f["locator"]["slot"] == "b/p[%d]" % n]


def top_border(element):
    border = OxmlElement("w:pBdr")
    edge = OxmlElement("w:top")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), "4")
    border.append(edge)
    element.append(border)


def caption(text, *, how="direct"):
    def add(doc):
        p = doc.add_paragraph()
        for n, part in enumerate(text.split("\t")):
            if n:
                p.add_run("\t")
            p.add_run(part)
        if how == "direct":
            top_border(p._p.get_or_add_pPr())
        elif how == "style":
            name = "Signing Rule"
            try:
                style = doc.styles[name]
            except KeyError:
                style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                top_border(style.element.get_or_add_pPr())
            p.style = style
    return add


def underlined(*parts, hint=()):
    """A paragraph of runs; parts are (text, underline) pairs; indexes in `hint` are written in guidance ink."""
    def add(doc):
        p = doc.add_paragraph()
        for n, (text, under) in enumerate(parts):
            r = p.add_run(text)
            r.underline = under
            if n in hint:
                color = OxmlElement("w:color")
                color.set(qn("w:val"), "0070C0")
                r._r.get_or_add_rPr().append(color)
    return add


class SigningPlaceholders(unittest.TestCase):
    """K7: a placeholder whose own words are the signing act is a signature place."""

    def test_sign_token_is_a_signature_place_and_ends_the_area(self):
        data = build("Yours sincerely,", "[Sign]", "[Name of signatory]", "[Title of signatory]")
        fields = fields_of(data)
        sign = at(fields, 1)
        self.assertEqual([f["kind"] for f in sign], ["signature"])
        self.assertNotIn("extra", sign[0]["locator"])
        self.assertEqual(sign[0]["locator"]["target"]["type"], "span")
        # the placeholders below it are their own place, not lines of the signature
        self.assertTrue(at(fields, 2))
        self.assertTrue(all(f["kind"] not in ("signature", "initials") for f in at(fields, 2)))

    def test_signing_token_variants_and_controls(self):
        cases = [("[Signature]", "signature"), ("<e-signature>", "signature"), ("{{Sign here}}", "signature"),
                 ("[Initials]", "initials"), ("[Signature of investigator]", "signature"),
                 ("[Initial assessment]", None), ("[Name of signatory]", None), ("[Date of signature]", None),
                 ("[Sign-off date]", None)]
        for token, kind in cases:
            with self.subTest(token=token):
                fields = at(fields_of(build("Closing line.", token)), 1)
                self.assertEqual(len(fields), 1)
                if kind in ("signature", "initials"):
                    self.assertEqual(fields[0]["kind"], kind)
                else:
                    self.assertNotIn(fields[0]["kind"], ("signature", "initials"))

    def test_signing_token_in_a_placeholder_cell(self):
        def table(doc):
            t = doc.add_table(rows=2, cols=2)
            t.cell(0, 0).text, t.cell(0, 1).text = "Clinician", "[Sign]"
            t.cell(1, 0).text, t.cell(1, 1).text = "Clinic", "[Name of clinic]"
        fields = fields_of(build(table))
        self.assertEqual([f["kind"] for f in fields], ["signature", "text"])

    def test_sign_token_inside_a_line_keeps_its_neighbours(self):
        fields = at(fields_of(build("Yours faithfully [Sign] on [date of letter] at [Name of clinic]")), 0)
        self.assertEqual([f["kind"] for f in fields], ["signature", "date", "text"])
        data = build("Yours faithfully [Sign] on [date of letter] at [Name of clinic]")
        allf = fields_of(data)
        answers = {f["key"]: "03/14/2026" if f["kind"] == "date" else "Synthetic Clinic"
                   for f in allf if f["kind"] not in ("signature", "initials")}
        out = word.fill(data, allf, answers)
        self.assertIn(b"Synthetic Clinic", _document_xml(out))
        with self.assertRaises(word.V3Failure):
            word.fill(data, allf, {[f for f in allf if f["kind"] == "signature"][0]["key"]: "X"})


def _document_xml(data):
    from zipfile import ZipFile
    return ZipFile(io.BytesIO(data)).read("word/document.xml")


class DisplacedSigningLine(unittest.TestCase):
    """K7: 'Signature:' / spacing / a drawn line - the drawn line is the signature place."""

    def test_drawn_line_after_spacing_is_the_signature(self):
        data = build("Patient's signature:", "", "______________________________", "Date: ____________")
        fields = fields_of(data)
        self.assertEqual(at(fields, 1), [])
        sign = at(fields, 2)
        self.assertEqual([f["kind"] for f in sign], ["signature"])
        self.assertEqual((sign[0]["locator"]["target"]["start"], sign[0]["locator"]["target"]["end"]), (0, 30))
        self.assertEqual([f["kind"] for f in at(fields, 3)], ["date"])
        self.assertEqual(sum(f["kind"] == "signature" for f in fields), 1)

    def test_two_spacing_lines_and_initials(self):
        fields = fields_of(build("Initials of the witness:", "", "", "__________"))
        self.assertEqual([f["kind"] for f in at(fields, 3)], ["initials"])
        self.assertEqual(at(fields, 1) + at(fields, 2), [])

    def test_line_under_a_label_line_naming_a_signature_and_a_date(self):
        fields = fields_of(build("Recorder Signature:\t\t\tDate:", ""))
        self.assertTrue(fields)
        self.assertTrue(all(f["kind"] == "signature" for f in fields if f["locator"]["slot"] == "b/p[1]"))
        fields = fields_of(build("Date of signature:", ""))
        self.assertEqual([f["kind"] for f in at(fields, 1)], ["date"])

    def test_controls_keep_their_lines(self):
        for label in ("Name:", "Signature date:", "Date:"):
            with self.subTest(label=label):
                fields = fields_of(build(label, "", "______________________________"))
                self.assertFalse([f for f in fields if f["kind"] in ("signature", "initials")])
        # a line holding more than the drawn blank is not the displaced signing line
        fields = fields_of(build("Patient's signature:", "", "________ Date: ________"))
        self.assertFalse([f for f in at(fields, 2) if f["kind"] == "signature"])


class CaptionRuledSigningLine(unittest.TestCase):
    """K7: an empty line over a caption ruled along its top ('Signature Line' style)."""

    def form(self, how="direct", captions="Signature of Participant\t\t\tPrint Name of Participant\t\t\tDate"):
        return build("Please sign below.", "", caption(captions, how=how), "",
                     caption("Signature of Person Obtaining Consent\t\t\tPrint Name\t\t\tDate", how=how))

    def test_signature_on_the_line_above_each_ruled_caption(self):
        for how in ("direct", "style"):
            with self.subTest(how=how):
                result = word.inspect(self.form(how))
                fields = result["fields"]
                self.assertEqual([(f["kind"], f["label"]) for f in at(fields, 1)],
                                 [("signature", "Signature of Participant")])
                self.assertEqual([(f["kind"], f["label"]) for f in at(fields, 3)],
                                 [("signature", "Signature of Person Obtaining Consent")])
                self.assertEqual(at(fields, 2) + at(fields, 4), [])
                self.assertFalse([f for f in fields if f["kind"] in ("date", "text", "long_text")])
                self.assertTrue(any("completed at signing" in n for n in result["notes"]))

    def test_initials_caption(self):
        fields = fields_of(build("", caption("Initials of Participant\t\tDate")))
        self.assertEqual([f["kind"] for f in at(fields, 0)], ["initials"])

    def test_controls(self):
        # no rule along the caption: not this structure
        fields = fields_of(self.form(how="none"))
        self.assertFalse([f for f in at(fields, 1) if f["kind"] == "signature"])
        # a ruled caption with no signing caption
        fields = fields_of(build("", caption("Print Name\t\t\tDate")))
        self.assertFalse([f for f in fields if f["kind"] in ("signature", "initials")])
        # the line above holds text: not an empty writing line
        fields = fields_of(build("Witnessed", caption("Signature of Witness\t\tDate")))
        self.assertFalse(at(fields, 0))

        # a ruled section heading is not a caption
        def heading(doc):
            p = doc.add_heading("Signature", 2)
            top_border(p._p.get_or_add_pPr())
        self.assertFalse(fields_of(build("", heading)))


class ESignatureAttestation(unittest.TestCase):
    """K7: 'Electronically signed on <date> by:' - the place after 'by' is the signing act."""

    def attestation(self, doc):
        p = doc.add_paragraph()
        p.add_run("\t")
        for text in ("Electronically Signed On ", "<<insert date and time>>", " ", "By:"):
            p.add_run(text).underline = True

    def test_attestation_by_is_a_signature_owning_the_signer_line(self):
        data = build(self.attestation, "\tJane Q. Example, MD")
        fields = fields_of(data)
        line = at(fields, 0)
        self.assertEqual(sorted(f["kind"] for f in line), ["date", "signature"])
        sign = [f for f in line if f["kind"] == "signature"][0]
        self.assertEqual([x["slot"] for x in sign["locator"]["extra"]], ["b/p[1]"])
        date = [f for f in line if f["kind"] == "date"][0]
        text = word.paragraph_text(word.open_package(data).tree("word/document.xml").find(".//" + word.q("w:p")))[0]
        self.assertEqual(text[date["locator"]["target"]["start"]:date["locator"]["target"]["end"]],
                         "<<insert date and time>>")
        out = word.fill(data, fields, {date["key"]: "03/14/2026"})
        xml = _document_xml(out)
        self.assertIn(b"03/14/2026", xml)
        self.assertNotIn(b"&lt;03/14/2026", xml)
        with self.assertRaises(word.V3Failure):
            word.fill(data, fields, {sign["key"]: "Jane Q. Example"})

    def test_controls(self):
        fields = fields_of(build(self.attestation, "Comments"))
        self.assertFalse([f for f in fields if f["locator"].get("extra")])
        fields = fields_of(build("Reviewed by:", "Jane Q. Example, MD"))
        self.assertFalse([f for f in fields if f["kind"] == "signature"])
        self.assertFalse([f for f in fields if f["locator"].get("extra")])
        fields = fields_of(build("Electronically signed by: ______________"))
        self.assertEqual([f["kind"] for f in fields], ["signature"])


class SegmentedLabels(unittest.TestCase):
    """K7: an 'Initials:' label printed after another label's answer gap is its own place."""

    def test_initials_after_circled_options(self):
        data = build("Circle Yes or No:      Yes      No      Initials:          Date: ")
        fields = fields_of(data)
        kinds = {f["label"]: f["kind"] for f in fields}
        self.assertEqual(kinds.get("Initials"), "initials")
        self.assertEqual(kinds.get("Date"), "date")
        self.assertFalse([f for f in fields if f["kind"] == "date" and "Initials" in f["label"]])
        spans = [(f["locator"]["target"]["start"], f["locator"]["target"]["end"]) for f in fields]
        self.assertEqual(len(set(spans)), len(spans))

    def test_name_signature_date_on_one_line_stay_separately_answerable(self):
        data = build("Name: ____________   Signature: ____________   Date: ____________")
        fields = fields_of(data)
        self.assertEqual([f["kind"] for f in fields], ["text", "signature", "date"])
        name, sign, date = fields
        out = word.fill(data, fields, {name["key"]: "Jane Example", date["key"]: "03/14/2026"})
        self.assertIn(b"Jane Example", _document_xml(out))
        with self.assertRaises(word.V3Failure):
            word.fill(data, fields, {sign["key"]: "Jane Example"})


class DrawnLeadersAndDates(unittest.TestCase):
    """A dotted leader typed as mixed dots and ellipses, and a drawn dd/mm/yyyy, are each one place."""

    def test_signature_leader_and_drawn_date(self):
        text = "Signature:……………........................…………\t\t\t\tDate: _____/_____/_____"
        data = build(text)
        fields = fields_of(data)
        self.assertEqual([(f["kind"], f["label"]) for f in fields], [("signature", "Signature"), ("date", "Date")])
        sign, date = fields
        self.assertEqual(text[sign["locator"]["target"]["start"]:sign["locator"]["target"]["end"]],
                         "……………........................…………")
        self.assertEqual(text[date["locator"]["target"]["start"]:date["locator"]["target"]["end"]], "_____/_____/_____")
        out = word.fill(data, fields, {date["key"]: "03/14/2026"})
        filled = word.paragraph_text(word.open_package(out).tree("word/document.xml").find(".//" + word.q("w:p")))[0]
        self.assertEqual(filled, text.replace("_____/_____/_____", "03/14/2026"))

    def test_a_date_shaped_blank_after_a_signing_label_is_signed(self):
        fields = fields_of(build("Sign here: __/__/____"))
        self.assertEqual([f["kind"] for f in fields], ["signature"])

    def test_two_groups_are_not_a_date(self):
        fields = fields_of(build("Blood pressure: ____/____ mmHg"))
        self.assertFalse([f for f in fields if f["kind"] == "date"])


class UnderlinedPhrase(unittest.TestCase):
    """K11: a space printed underlined inside an underlined phrase is emphasis, not a drawn blank."""

    def test_underlined_phrase_has_no_answer_place(self):
        data = build(underlined(("Rate how you felt ", False), ("during", True), (" ", True),
                                ("the past 30 days", True), (".", False)))
        self.assertEqual(fields_of(data), [])

    def test_controls(self):
        fields = fields_of(build(underlined(("Name:", False), ("\t\t\t", True))))
        self.assertEqual(len(fields), 1)
        # a guidance-ink neighbour: a hint in the answer position keeps its place
        fields = fields_of(build(underlined(("Completed by: ", False), ("Print", True), (" ", True), ("Name", True),
                                            hint=(1, 3))))
        self.assertTrue([f for f in fields if f["locator"]["target"].get("shape") == "utab"])
        # three or more underlined spaces are a drawn blank
        fields = fields_of(build(underlined(("Signed", True), ("      ", True), ("here", True))))
        self.assertEqual(len(fields), 1)



def listed(doc, items, *, num, fmt="decimal"):
    """List paragraphs [(text, level)] of a list of its own (numId `num`, every level printed as `fmt`)."""
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(100 + int(num)))
    for lvl in range(3):
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), str(lvl))
        f = OxmlElement("w:numFmt")
        f.set(qn("w:val"), fmt)
        level.append(f)
        abstract.append(level)
    numbering.append(abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(100 + int(num)))
    number.append(ref)
    numbering.append(number)
    for text, lvl in items:
        p = doc.add_paragraph(text)
        props = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(lvl))
        nid = OxmlElement("w:numId")
        nid.set(qn("w:val"), str(num))
        props.append(ilvl)
        props.append(nid)
        p._p.get_or_add_pPr().append(props)


def spans_of(data):
    body = word.open_package(data).tree("word/document.xml")
    texts = [word.paragraph_text(p)[0] for p in body.iter(word.q("w:p"))]
    out = []
    for f in fields_of(data):
        t = f["locator"]["target"]
        if t.get("type") == "span":
            idx = word.SlotIndex(word.open_package(data)).by_id[f["locator"]["part"] + "#" + f["locator"]["slot"]]
            out.append(word.paragraph_text(idx.el)[0][t["start"]:t["end"]])
    return out


class AuthorsBracketedText(unittest.TestCase):
    """K10: a bracketed or angled phrase inside running prose that is the author's text, not a place to write."""

    def test_conditions_examples_and_instruction_groups_are_printed(self):
        for text in ("Participation is voluntary. <for hospital sites> You were invited because you stayed.",
                     "<if applicable> You would take part in the study for [X weeks/months].",
                     "Samples <if collected><coded data and coded samples> will be kept.",
                     "[IF YES] What ethnic group do you belong to?",
                     "You will get (*insert name of drug [e.g., tablets] for [number of days] days*).",
                     "Data may come from a [insert device here, e.g. a fitness tracker].",
                     "Your rights are described in <add local regulation (e.g. a Code of Conduct) as applicable>.",
                     "Contact the ethics committee <remove if not applicable> first."):
            with self.subTest(text=text):
                result = word.inspect(build(text))
                self.assertEqual(spans_of(build(text)), [])
                self.assertTrue(any("kept as printed" in n for n in result["notes"]))

    def test_instruction_group_alone(self):
        self.assertEqual(spans_of(build("(* You will receive [name of drug] each day *)")), [])
        self.assertEqual(len(spans_of(build("You will receive [name of drug] each day."))), 1)

    def test_note_alone_on_its_line_is_not_the_next_lines_label(self):
        fields = fields_of(build("<if applicable>", "[Name of the second parent]"))
        self.assertEqual(at(fields, 0), [])
        self.assertEqual([f["label"] for f in at(fields, 1)], ["Name of the second parent"])

    def test_wording_for_another_respondent_is_printed(self):
        for text in ("I get [My child gets] along with friends.", "What is your [your child\u2019s] month of birth?",
                     "The people helping me [my child] stuck with me [us].", "I felt I had [my child had] support.",
                     "The help I [we] wanted [for my child] came."):
            with self.subTest(text=text):
                self.assertEqual(spans_of(build(text)), [])

    def test_bracketed_heading_and_printed_notice(self):
        self.assertEqual(fields_of(build("[CAREGIVER RESPONSE:]", "[This page intentionally left blank]")), [])

    def test_header_prose_placeholders(self):
        for header, n in (("This project <insert project acronym> has received funding.", 0),
                          ("Clinic letter for [patient name] dated [date]", 2)):
            with self.subTest(header=header):
                doc = Document()
                doc.sections[0].header.paragraphs[0].text = header
                doc.add_paragraph("Visit notes")
                out = io.BytesIO()
                doc.save(out)
                self.assertEqual(len(fields_of(out.getvalue())), n)

    def test_grid_row_aside(self):
        def grid(doc, box):
            t = doc.add_table(rows=1, cols=2)
            t.cell(0, 0).text = "prescription opioids (oxycodone [OxyContin], hydrocodone)"
            t.cell(0, 1).text = "\u2610 Yes \u2610 No" if box else "Used in the past month"
        with_box = fields_of(build(lambda doc: grid(doc, True)))
        self.assertFalse([f for f in with_box if f["locator"]["target"].get("shape") == "bracket"])
        without = fields_of(build(lambda doc: grid(doc, False)))
        self.assertTrue([f for f in without if f["locator"]["target"].get("shape") == "bracket"])

    def test_places_to_write_stay(self):
        cases = [("I understand that [Name of applicant] wishes to apply.", 1),
                 ("The applicant confirmed that [he/she] attended.", 1),
                 ("Dear [Insert name of addressee],", 1),
                 ("[Your Address]", 1), ("[Your E-mail (if applicable)]", 1),
                 ("Other caregivers you may call are: [Name], at [Telephone number].", 2),
                 ("The purpose of this treatment is [insert purpose].", 1),
                 ("It is approved to treat [X condition].", 1),
                 ("Name of applicant: [Name of applicant]", 1),
                 ("I, [your name], agree to take part in this study.", 1),
                 ("We [your practice name] confirm the referral.", 1),
                 ("[IF YES] Date of diagnosis: [date]", 1),
                 ("The visit took place on [Date of visit, e.g. DD/MM/YYYY].", 1)]
        for text, n in cases:
            with self.subTest(text=text):
                self.assertEqual(len(spans_of(build(text))), n)


class LeadInsAndHeadings(unittest.TestCase):
    """K11: colon-ended text that introduces printed content is not a label with an answer space."""

    def test_sentence_lead_in_before_content(self):
        for lead in ("Some of the side effects that researchers know about are:",
                     "If you choose to take part in this study, you will need to:",
                     "I have been asked to give consent for my child to take part:",
                     "Please read the following before you sign:"):
            with self.subTest(lead=lead):
                fields = fields_of(build(lead, "Nausea and vomiting may happen after each dose."))
                self.assertEqual(at(fields, 0), [])
                fields = fields_of(build(lead, "", "Name of subject: ____________"))
                self.assertEqual(at(fields, 0), [])

    def test_questions_requests_and_open_labels_stay(self):
        for label, after in (("Have you ever been diagnosed with cancer:", "Nausea and vomiting are common."),
                             ("If yes, please specify the reason:", "Nausea and vomiting are common."),
                             ("Medications you are currently taking:", ""),
                             ("Medications you are currently taking:", "______________________"),
                             ("Name:", "Nausea and vomiting are common.")):
            with self.subTest(label=label, after=after):
                fields = fields_of(build(label, after) if after else build(label))
                self.assertEqual(len(at(fields, 0)), 1)

    def test_heading_over_typed_bullets(self):
        fields = fields_of(build("Pre-treatment:", "- Dental status:  ", "- Tumour size:  "))
        self.assertEqual(at(fields, 0), [])
        self.assertEqual(len(at(fields, 1)) + len(at(fields, 2)), 2)
        # an item that is a label itself keeps its answer space
        fields = fields_of(build("- Evolution:  ", "  - Stabilization:  ", "  - Regression:  "))
        self.assertEqual([len(at(fields, n)) for n in range(3)], [1, 1, 1])

    def test_list_heading_over_deeper_prose_items(self):
        data = build(lambda doc: listed(doc, [("Side effects of study groups:", 0),
                                              ("For single-arm studies, list every side effect.", 1),
                                              ("For multiple-arm studies, use a table.", 1)], num=21))
        self.assertEqual(at(fields_of(data), 0), [])
        # items that are labels themselves, and label items of one list level, keep their answer spaces
        data = build(lambda doc: listed(doc, [("Height:", 0), ("Weight:", 0), ("Pulse:", 0)], num=22))
        self.assertEqual([len(at(fields_of(data), n)) for n in range(3)], [1, 1, 1])
        data = build(lambda doc: listed(doc, [("Vital signs:", 0)], num=25),
                     lambda doc: listed(doc, [("Weight: ", 0), ("Pulse: ", 0)], num=26))
        self.assertEqual([len(at(fields_of(data), n)) for n in range(3)], [1, 1, 1])

    def test_unclosed_bracketed_instruction(self):
        self.assertEqual(fields_of(build("[FOR ALL REASSESSMENTS:", "Record the date of the review.")), [])

    def test_column_header_over_answers_and_printed_values(self):
        def answers(doc):
            t = doc.add_table(rows=2, cols=2)
            t.cell(0, 0).text = "Visit"
            t.cell(0, 1).text = "Admission Date:"
            t.cell(1, 0).text = "First visit"
        fields = fields_of(build(answers))
        self.assertFalse([f for f in fields if f["locator"]["slot"].startswith("b/tbl[0]/tr[0]")])
        self.assertTrue([f for f in fields if f["locator"]["slot"] == "b/tbl[0]/tr[1]/tc[1]"])

        def printed(doc):
            t = doc.add_table(rows=3, cols=2)
            t.cell(0, 0).text, t.cell(0, 1).text = "Version", "Summary of Revisions Made:"
            t.cell(1, 0).text, t.cell(1, 1).text = "1.0", "First approved version"
            t.cell(2, 0).text, t.cell(2, 1).text = "2.0", "Removed page numbers"
        self.assertEqual(fields_of(build(printed)), [])

        def named_below(doc):
            # the answer cell under the header is two rows down; the row between answers with boxes
            t = doc.add_table(rows=3, cols=2)
            t.cell(0, 0).text, t.cell(0, 1).text = "Unit", "Admission Date:"
            t.cell(1, 0).text, t.cell(1, 1).text = "Intensive care", "\u2610 Yes \u2610 No"
            t.cell(2, 0).text = "Step-down"
        fields = fields_of(build(named_below))
        self.assertFalse([f for f in fields if f["locator"]["slot"].startswith("b/tbl[0]/tr[0]")])
        self.assertEqual([f["label"] for f in fields if f["locator"]["slot"] == "b/tbl[0]/tr[2]/tc[1]"],
                         ["Step-down \u2014 Admission Date"])

        def one_printed_row(doc):
            # one printed line under a label is not a column of printed values: the label keeps its answer space
            t = doc.add_table(rows=2, cols=1)
            t.cell(0, 0).text, t.cell(1, 0).text = "Allergies:", "Please list all known allergies"
        self.assertEqual([f["label"] for f in fields_of(build(one_printed_row))], ["Allergies"])

        def one_row(doc):
            t = doc.add_table(rows=1, cols=3)
            for n, text in enumerate(("Investigator Name:", "Protocol:", "Site Name:")):
                t.cell(0, n).text = text
        self.assertEqual(len(fields_of(build(one_row))), 3)

    def test_section_captions_over_options_signing_rows_and_grids(self):
        def options(doc):
            t = doc.add_table(rows=2, cols=3)
            a = t.cell(0, 0).merge(t.cell(0, 2))
            a.text = "Comorbidities:"
            for n, text in enumerate(("Diabetes", "Hypertension", "Asthma")):
                t.cell(1, n).text = text
        self.assertEqual(fields_of(build(options)), [])

        def signing(doc):
            t = doc.add_table(rows=3, cols=3)
            a = t.cell(0, 0).merge(t.cell(0, 2))
            a.text = "Investigator:"
            for n, text in enumerate(("Signature of Investigator", "Print Name", "Date")):
                t.cell(2, n).text = text
        fields = fields_of(build(signing))
        self.assertFalse([f for f in fields if f["locator"]["slot"].startswith("b/tbl[0]/tr[0]")])
        self.assertEqual([(f["kind"], f["locator"]["slot"]) for f in fields if f["kind"] == "signature"],
                         [("signature", "b/tbl[0]/tr[1]/tc[0]")])

        def grid(doc):
            t = doc.add_table(rows=4, cols=3)
            for n, text in enumerate(("Medication", "Dose", "Ongoing:")):
                t.cell(0, n).text = text
        fields = fields_of(build(grid))
        self.assertFalse([f for f in fields if f["locator"]["target"]["type"] == "span"])

    def test_signer_row_heading(self):
        def row(doc, cells):
            t = doc.add_table(rows=1, cols=len(cells))
            for n, text in enumerate(cells):
                t.cell(0, n).text = text
        fields = fields_of(build(lambda doc: row(doc, ("Completed by:", "Name:", "Signature:", "Date:"))))
        self.assertEqual([f["kind"] for f in fields], ["text", "signature", "date"])
        fields = fields_of(build(lambda doc: row(doc, ("Completed by:", "Protocol:"))))
        self.assertEqual([f["label"] for f in fields], ["Completed by", "Protocol"])

    def test_instruction_list_is_not_an_option_list(self):
        data = build(lambda doc: doc.add_paragraph("For the lay title:"),
                     lambda doc: listed(doc, [("Provide a brief title of the study in lay language.", 0),
                                              ("Use general terms.", 0)], num=23, fmt="lowerLetter"))
        self.assertFalse([f for f in fields_of(data) if f["kind"] in ("choice", "multi_choice")])
        data = build(lambda doc: doc.add_paragraph("Which arm was treated:"),
                     lambda doc: listed(doc, [("Left", 0), ("Right", 0)], num=24, fmt="lowerLetter"))
        self.assertEqual([f["kind"] for f in fields_of(data)], ["choice"])
        data = build(lambda doc: doc.add_paragraph("Reason for referral:"),
                     lambda doc: listed(doc, [("Chest pain.", 0), ("Breathlessness.", 0)], num=27, fmt="lowerLetter"))
        self.assertEqual([f["kind"] for f in fields_of(data)], ["choice"])


class TableHeadingsAndAnswerLines(unittest.TestCase):
    """K5 (subset) and K15."""

    def test_signing_row_captions_under_a_section_caption(self):
        def table(doc, caption_line):
            t = doc.add_table(rows=3, cols=3)
            a = t.cell(0, 0).merge(t.cell(0, 2))
            a.text = "Assent of Minor"
            for n, text in enumerate(("Signature of Minor", "Print Name of Minor", "Date")):
                t.cell(2, n).text = text
            if caption_line:
                t.cell(2, 2).add_paragraph()
        fields = fields_of(build(lambda doc: table(doc, False)))
        # the signature cell is a signature place; its neighbours keep their own (cell-keyed) answer places
        self.assertEqual([(f["kind"], f["label"], f["locator"]["slot"]) for f in fields],
                         [("signature", "Signature of Minor", "b/tbl[0]/tr[1]/tc[0]"),
                          ("text", "Print Name of Minor", "b/tbl[0]/tr[1]/tc[1]"), ("date", "Date", "b/tbl[0]/tr[1]/tc[2]")])
        # a caption's own empty line is spacing: never the date answer of the signing row
        fields = fields_of(build(lambda doc: table(doc, True)))
        self.assertFalse([f for f in fields if f["locator"]["slot"].startswith("b/tbl[0]/tr[2]")])

    def hint(self, cell, text):
        cell.text = ""
        r = cell.paragraphs[0].add_run(text)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0070C0")
        r._r.get_or_add_rPr().append(color)

    def test_heading_cell_of_a_repeated_header_row(self):
        def table(doc):
            t = doc.add_table(rows=4, cols=3)
            for r, side in ((0, "** Right Leg **"), (2, "** Left Leg **")):
                self.hint(t.cell(r, 0), side)
                t.cell(r, 1).text, t.cell(r, 2).text = "DIAMETER (mm)", "REFLUX TIME (ms)"
            t.cell(1, 0).text, t.cell(3, 0).text = "GSV at the knee", "GSV at the knee"
        fields = fields_of(build(table))
        self.assertFalse([f for f in fields if f["locator"]["slot"].startswith("b/tbl[0]/tr[2]/tc[0]")])
        self.assertEqual(len([f for f in fields if "/tr[3]/" in f["locator"]["slot"] + "/"]), 2)

    def instruction(self, empty_line):
        def table(doc):
            t = doc.add_table(rows=2, cols=2)
            t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Findings"
            t.cell(1, 0).text = "Left knee"
            self.hint(t.cell(1, 1), "If swollen, describe here")
            if empty_line:
                t.cell(1, 1).add_paragraph()
        return build(table)

    def test_cell_append_uses_the_empty_answer_line(self):
        for empty_line, paragraphs in ((True, 2), (False, 2)):
            with self.subTest(empty_line=empty_line):
                data = self.instruction(empty_line)
                fields = fields_of(data)
                appended = [f for f in fields if f["locator"]["target"]["type"] == "cell_append"]
                self.assertEqual(len(appended), 1)
                out = word.fill(data, fields, {appended[0]["key"]: "Mild effusion"})
                cell = word.SlotIndex(word.open_package(out)).by_id["word/document.xml#b/tbl[0]/tr[1]/tc[1]"].el
                paras = cell.findall(word.q("w:p"))
                self.assertEqual(len(paras), paragraphs)
                self.assertEqual(word.text_of(paras[-1]), "Mild effusion")
                self.assertEqual(word.text_of(paras[0]), "If swollen, describe here")


if __name__ == "__main__":
    unittest.main()
