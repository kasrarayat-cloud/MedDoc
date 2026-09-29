"""The ONE structure-only authority over signing, initials, identity and attestation places (word_v3.held_places).

Synthetic BLANK forms only; no network, no provider. The pinned parser's readings are used in-process (as the trusted
parent supplies them to the isolated parser); `NoReadings` proves the holding direction when a clause is unread.
"""
import io
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import word_v3 as w
from .service import DocumentError, fill_template, inspect_template, locate_fields

BODY = "word/document.xml#"


def saved(doc):
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def bordered_box(doc, text=None, pict=False):
    table = doc.add_table(rows=1, cols=1)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:val"), "single")
        borders.append(edge)
    table.cell(0, 0)._tc.get_or_add_tcPr().append(borders)
    if text:
        table.cell(0, 0).text = text
    if pict:
        run = OxmlElement("w:r")
        run.append(OxmlElement("w:pict"))
        table.cell(0, 0).paragraphs[0]._p.append(run)
    return table


def question_box_form(question, variant="empty", heading="Consultant review"):
    """R1: a heading, a printed question and its own one-cell bordered box, empty / holding a drawing / printing ink."""
    doc = Document()
    doc.add_paragraph("Participant ID: __________")
    doc.add_paragraph(heading, style="Heading 2")
    doc.add_paragraph(question)
    bordered_box(doc, text="Use black ink" if variant == "ink" else None, pict=variant == "pict")
    return saved(doc)


def column_form(variant="empty", header="Signature"):
    """R2: a 'Name | <header>' table whose second cell is empty / holds a drawing / prints ink guidance."""
    doc = Document()
    doc.add_paragraph("Clinic record", style="Heading 2")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = header
    if variant == "ink":
        table.cell(1, 1).text = "Use black ink"
    if variant == "pict":
        run = OxmlElement("w:r")
        run.append(OxmlElement("w:pict"))
        table.cell(1, 1).paragraphs[0]._p.append(run)
    return saved(doc)


def lines_form(*lines):
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    return saved(doc)


def admin_grid():
    doc = Document()
    doc.add_paragraph("Medicines given", style="Heading 2")
    table = doc.add_table(rows=3, cols=4)
    for column, header in enumerate(("Drug", "Dose", "Time", "Signature")):
        table.cell(0, column).text = header
    return saved(doc)


def held(data):
    package = w.open_package(data)
    return {slot[len(BODY):] if slot.startswith(BODY) else slot: entry for slot, entry in w.held_places(package).items()}


def text_field(slot, target=None, kind="text", label="Answer", key="probe"):
    part, path = (slot.split("#", 1) if "#" in slot else ("word/document.xml", slot))
    return {"key": key, "label": label, "kind": kind, "context": "",
            "locator": {"contract": w.CONTRACT, "part": part, "slot": path, "target": target or {"type": "paragraph"}}}


class NoReadings:
    """Run the authority with an EMPTY readings map (what the isolated parser has when the parent reads nothing)."""

    def __enter__(self):
        self.saved = w.READINGS
        w.use_readings({})

    def __exit__(self, *exc):
        w.use_readings(self.saved)


class QuestionBoxesWhateverTheyHold(unittest.TestCase):
    """R1: the question and section are read BEFORE and whatever the box holds; the box is held in every variant."""

    def test_signing_and_consent_question_boxes_are_held_empty_drawing_or_guidance(self):
        for question in ("Signed by the consultant:", "I confirm the patient gave consent to proceed:",
                         "Record that the patient consents to proceed:", "Patient consented: yes",
                         "Has the report been countersigned?", "Countersignature:"):
            for variant in ("empty", "pict", "ink"):
                with self.subTest(question=question, variant=variant):
                    entry = held(question_box_form(question, variant)).get("b/tbl[0]/tr[0]/tc[0]")
                    self.assertIsNotNone(entry)
                    self.assertIn(entry.code, ("consent", "declaration", "attestation", "signing"))
                    self.assertIsNone(entry.ranges)

    def test_a_signing_section_holds_the_box_under_an_ordinary_question(self):
        entry = held(question_box_form("What was discussed?", "ink", heading="Consultant signature")).get("b/tbl[0]/tr[0]/tc[0]")
        self.assertEqual(entry.code, "consent")

    def test_control_an_ordinary_question_box_is_not_held_in_any_variant(self):
        for variant in ("empty", "pict", "ink"):
            with self.subTest(variant=variant):
                self.assertNotIn("b/tbl[0]/tr[0]/tc[0]", held(question_box_form("Please describe the wound?", variant,
                                                                                 heading="Wound review")))


class ColumnsRowsAndBoxes(unittest.TestCase):
    """R2: a cell under a signing column header is held whatever it holds; an admin grid holds only that column."""

    def test_a_signature_column_cell_is_held_empty_drawing_or_guidance(self):
        for header in ("Signature", "Countersigned", "Countersignature", "Signatory", "Initials", "Sign"):
            for variant in ("empty", "pict", "ink"):
                with self.subTest(header=header, variant=variant):
                    places = held(column_form(variant, header))
                    self.assertEqual(places["b/tbl[0]/tr[1]/tc[1]"].code, "signing")
                    self.assertNotIn("b/tbl[0]/tr[1]/tc[0]", places)

    def test_an_administration_grid_holds_only_its_signature_column(self):
        places = held(admin_grid())
        cells = {slot for slot in places if slot.startswith("b/tbl[0]/")}
        self.assertEqual({slot.split("/")[3] for slot in cells if "/p[" not in slot}, {"tc[3]"})
        field = next(f for f in w.inspect(admin_grid())["fields"] if f["kind"] == "table_rows")
        self.assertNotIn("write_protected", field)      # the drug, dose and time columns stay answerable
        package = w.open_package(admin_grid())
        source = w.Detector(package, layout_proposal=True)
        source.detect()
        authority = w.held_places(package, source)
        self.assertIsNone(w.held_violation(source.index, authority, field))
        self.assertEqual(w.held_table_columns(source.index, authority, field), [3])
        w.fill(admin_grid(), [field], {field["key"]: [["Drug A", "5 mg", "08:00", ""]]})
        with self.assertRaises(w.V3Failure) as refused:
            w.fill(admin_grid(), [field], {field["key"]: [["Drug A", "5 mg", "08:00", "SYNTHETIC-ANSWER"]]})
        self.assertEqual(refused.exception.code, "held_place")

    def test_an_initials_row_holds_its_box_whatever_the_box_prints(self):
        for printed in (None, "Initial here", "Use black ink"):
            with self.subTest(printed=printed):
                doc = Document()
                table = doc.add_table(rows=1, cols=2)
                table.cell(0, 0).text = "Clinician initials"
                if printed:
                    table.cell(0, 1).text = printed
                self.assertEqual(held(saved(doc))["b/tbl[0]/tr[0]/tc[1]"].code, "signing")

    def test_a_signature_box_captioned_inside_holds_its_empty_writing_lines(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Signature of examining clinician"
        table.cell(0, 0).add_paragraph("")
        places = held(saved(doc))
        self.assertEqual(places["b/tbl[0]/tr[0]/tc[0]"].code, "signing")
        index = w.SlotIndex(w.open_package(saved(doc)))
        self.assertTrue(w.held_touches(index, {BODY + k: v for k, v in places.items()},
                                       text_field("b/tbl[0]/tr[0]/tc[0]/p[1]")))

    def test_guidance_that_mentions_signing_is_not_a_signing_caption(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "What progress have you made?"
        table.cell(0, 1).text = "Describe your own experience; do not sign this answer for another person."
        self.assertNotIn("b/tbl[0]/tr[0]/tc[1]", held(saved(doc)))
        # A sentence that mentions a signed form is held where it is printed, but it is no caption of a signing box:
        # the answer line below it in the same cell stays an ordinary answer place.
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Describe the procedure performed; the signed consent form is filed with the notes."
        table.cell(0, 0).add_paragraph("")
        places = held(saved(doc))
        self.assertIn("b/tbl[0]/tr[0]/tc[0]/p[0]", places)
        self.assertNotIn("b/tbl[0]/tr[0]/tc[0]", places)
        index = w.SlotIndex(w.open_package(saved(doc)))
        self.assertIsNone(w.held_violation(index, {BODY + k: v for k, v in places.items()},
                                           text_field("b/tbl[0]/tr[0]/tc[0]/p[1]", kind="long_text")))

    def test_a_drawn_box_under_an_initials_instruction_and_a_signing_control_alias_are_held(self):
        from .test_signing_areas import signing_form
        entry = held(signing_form())["b/tbl[0]/tr[0]/tc[1]"]
        self.assertEqual(entry.code, "signing")
        self.assertTrue(entry.evidence.startswith("drawn box under"), entry.evidence)
        doc = Document()
        paragraph = doc.add_paragraph("Reviewer: ")
        sdt = OxmlElement("w:sdt")
        properties = OxmlElement("w:sdtPr")
        alias = OxmlElement("w:alias")
        alias.set(qn("w:val"), "Consultant signature")
        properties.append(alias)
        sdt.append(properties)
        content = OxmlElement("w:sdtContent")
        run = OxmlElement("w:r")
        run.append(OxmlElement("w:t"))
        content.append(run)
        sdt.append(content)
        paragraph._p.append(sdt)
        self.assertEqual(held(saved(doc))["b/p[0]/sdt[0]"].code, "signing")

    def test_a_locked_content_control_is_held_as_protected(self):
        doc = Document()
        paragraph = doc.add_paragraph("Statement: ")
        sdt = OxmlElement("w:sdt")
        properties = OxmlElement("w:sdtPr")
        lock = OxmlElement("w:lock")
        lock.set(qn("w:val"), "sdtContentLocked")
        properties.append(lock)
        sdt.append(properties)
        content = OxmlElement("w:sdtContent")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Printed statement"
        run.append(text)
        content.append(run)
        sdt.append(content)
        paragraph._p.append(sdt)
        self.assertEqual(held(saved(doc))["b/p[0]/sdt[0]"].code, "protected")


class InflectionsAndCaptions(unittest.TestCase):
    """The existing readers' inflection gaps, closed without new vocabulary: consent/consents/consented/consenting,
    'signed' inside countersigned, countersignature, signatory, 'Sign:'/'Sign below' as a caption, the adjective
    'initial', and one blank captioned '... signature and date'."""

    def kinds(self, *lines):
        return [(f["label"], f["kind"], bool(f.get("write_protected"))) for f in w.inspect(lines_form(*lines))["fields"]]

    def test_signing_captions_read_every_inflection(self):
        for line, kind in (("Countersigned by: ______", "signature"), ("Countersignature: ______", "signature"),
                           ("Signatory: ______", "signature"), ("Sign: ______", "signature"),
                           ("Sign below: ______", "signature"), ("Co-signed by: ______", "signature"),
                           ("Consultant signature and date: ______", "signature"),
                           ("Initials: ______", "initials"), ("Initial here: ______", "initials"),
                           ("Initialled by: ______", "initials")):
            with self.subTest(line=line):
                self.assertEqual(self.kinds(line)[0][1], kind)
                self.assertEqual(held(lines_form(line))["b/p[0]"].code, "signing")

    def test_clinical_signs_and_the_adjective_initial_stay_ordinary(self):
        for line, kind in (("Vital signs: ______", "text"), ("Signs and symptoms: ______", "text"),
                           ("Sign of infection: ______", "text"), ("Initial assessment: ______", "text"),
                           ("Initial presentation: ______", "text"), ("Initial dose: ______", "text"),
                           ("Signature date: ______", "date"), ("Date signed: ______", "date"),
                           ("Consent obtained: ______", "text"), ("Type I diabetes: ______", "text")):
            with self.subTest(line=line):
                self.assertEqual(self.kinds(line), [(clean(line), kind, False)])
                self.assertNotIn("b/p[0]", held(lines_form(line)))

    def test_question_box_signing_reads_every_consent_inflection(self):
        for text in ("Record that the patient consents to proceed:", "Patient consented: yes", "Consenting adult?",
                     "Consent given?"):
            with self.subTest(text=text):
                self.assertTrue(w.Detector.question_box_signing(text))
        self.assertFalse(w.Detector.question_box_signing("Vital signs recorded?"))

    def test_an_unread_clause_is_read_as_holding(self):
        with NoReadings():
            self.assertEqual(w.kind_for_label("Initial assessment"), "initials")
            self.assertEqual(w.kind_for_label("Sign of infection"), "signature")
            self.assertTrue(w.first_person("Type I diabetes"))
            self.assertTrue(w.embedded_consent("Consent obtained"))
            self.assertEqual(w.kind_for_label("Vital signs"), "text")        # a plural is never a candidate
        self.assertEqual(w.kind_for_label("Initial assessment"), "text")


def clean(line):
    return w.clean_label(line.split("_")[0])


class DeclarationsAndAttestations(unittest.TestCase):
    """A first-person declaration (and the block it heads) is held; an embedded consent is an unresolved attestation; a
    factual third-person field stays ordinary."""

    def test_first_person_statements_and_their_heading_are_held(self):
        doc = Document()
        doc.add_paragraph("Declaration", style="Heading 2")
        doc.add_paragraph("I declare the person examined agreed to this examination: ______")
        doc.add_paragraph("☐ Yes ☐ No")
        doc.add_paragraph("Findings", style="Heading 2")
        doc.add_paragraph("Plan: ______")
        places = held(saved(doc))
        self.assertEqual(places["b/p[0]"].code, "declaration")
        self.assertEqual(places["b/p[1]"].code, "declaration")
        self.assertEqual(places["b/p[2]"].code, "declaration")     # its own answer line
        self.assertNotIn("b/p[3]", places)
        self.assertNotIn("b/p[4]", places)
        # The engine reads the answer line as the declaration's Yes/No choice: a new inspection write-protects it.
        fields = w.inspect(saved(doc))["fields"]
        choice = next(f for f in fields if f["locator"]["slot"] == "b/p[2]/glyphs")
        self.assertEqual((choice["kind"], choice["write_protected"]), ("choice", True))
        self.assertFalse(next(f for f in fields if f["label"] == "Plan").get("write_protected"))

    def test_a_declaration_row_holds_its_box_on_either_side(self):
        for statement_first in (True, False):
            with self.subTest(statement_first=statement_first):
                doc = Document()
                table = doc.add_table(rows=1, cols=2)
                statement, box = (0, 1) if statement_first else (1, 0)
                table.cell(0, statement).text = "I confirm the risks were explained"
                table.cell(0, box).text = "☐"
                self.assertEqual(held(saved(doc))["b/tbl[0]/tr[0]/tc[%d]" % box].code, "declaration")

    def test_an_embedded_consent_is_an_unresolved_attestation_and_a_fact_is_ordinary(self):
        places = held(lines_form("Record that the patient consents to proceed: ______", "Consent obtained: ______",
                                 "Patient consented: ______"))
        self.assertEqual(places["b/p[0]"].code, "attestation")
        self.assertNotIn("b/p[1]", places)
        self.assertNotIn("b/p[2]", places)


class RangedCaptionLines(unittest.TestCase):
    def test_only_the_signature_part_of_a_mixed_caption_line_is_held(self):
        data = lines_form("Name: ______ Signature: ______ Date: ______")
        entry = held(data)["b/p[0]"]
        self.assertEqual(entry.code, "signing")
        self.assertEqual(len(entry.ranges), 1)
        fields = w.inspect(data)["fields"]
        self.assertEqual([(f["label"], f["kind"], bool(f.get("write_protected"))) for f in fields],
                         [("Name", "text", False), ("Signature", "signature", False), ("Date", "date", False)])
        name, _, date = fields
        written = w.fill(data, fields, {name["key"]: "SYNTHETIC NAME", date["key"]: "2026-09-29"})
        self.assertIn("SYNTHETIC NAME", w.text_of(w.SlotIndex(w.open_package(written)).by_id[BODY + "b/p[0]"].el))
        # A text field at the signature part itself is refused by fill (the stored-mapping protection).
        start, end = entry.ranges[0]
        forged = {**name, "key": "forged", "locator": {**name["locator"], "target": {"type": "span", "start": start + 12,
                                                                                         "end": start + 12, "shape": "after_label"}}}
        with self.assertRaises(w.V3Failure) as refused:
            w.fill(data, [forged], {"forged": "SYNTHETIC-ANSWER"})
        self.assertEqual(refused.exception.code, "held_place")

    def test_a_signed_question_with_boxes_holds_a_written_answer_but_not_a_box_selection(self):
        entry = held(lines_form("Consent form signed? ☐ Yes ☐ No"))["b/p[0]"]
        self.assertTrue(entry.selection)
        package = w.open_package(lines_form("Consent form signed? ☐ Yes ☐ No"))
        index = w.SlotIndex(package)
        authority = {BODY + "b/p[0]": entry}
        box = {"key": "box", "label": "Yes", "kind": "checkbox", "options": ["Yes"],
               "locator": {"contract": w.CONTRACT, "part": "word/document.xml", "slot": "b/p[0]",
                           "target": {"type": "checkboxes", "boxes": [{"kind": "glyph", "part": "word/document.xml",
                                                                        "slot": "b/p[0]/glyphs", "glyph": 0}]}}}
        self.assertIsNone(w.held_violation(index, authority, box))
        self.assertIsNotNone(w.held_violation(index, authority, text_field("b/p[0]")))


class OneAuthorityEveryWriter(unittest.TestCase):
    """Fill, inspect and locate consult the SAME authority; it reads only the original's structure."""

    def test_fill_refuses_a_stored_answerable_field_on_every_held_shape(self):
        cases = [(question_box_form("Signed by the consultant:", variant), "b/tbl[0]/tr[0]/tc[0]", {"type": "cell"})
                 for variant in ("empty", "pict", "ink")]
        cases += [(column_form(variant), "b/tbl[0]/tr[1]/tc[1]", {"type": "cell"}) for variant in ("empty", "pict", "ink")]
        cases += [(lines_form("Signature: ______"), "b/p[0]", {"type": "span", "start": 10, "end": 10, "shape": "after_label"})]
        for data, slot, target in cases:
            with self.subTest(slot=slot):
                index = w.SlotIndex(w.open_package(data))
                field = text_field(slot, target)
                field["locator"]["fingerprint"] = w.fingerprint(index.by_id[BODY + slot].el)
                with self.assertRaises(w.V3Failure) as refused:
                    w.fill(data, [field], {"probe": "SYNTHETIC-ANSWER"})
                self.assertEqual(refused.exception.code, "held_place")

    def test_locate_refuses_a_non_signing_kind_on_a_held_place_and_allows_a_signing_kind(self):
        data = column_form("pict")
        with self.assertRaises(w.V3Failure) as refused:
            w.locate(data, [{"slot": BODY + "b/tbl[0]/tr[1]/tc[1]/p[0]", "kind": "text", "label": "Clinic record"}])
        self.assertEqual(refused.exception.code, "held_place")
        located = w.locate(data, [{"slot": BODY + "b/tbl[0]/tr[1]/tc[1]/p[0]", "kind": "signature", "label": "Signature"}])
        self.assertEqual(located[0]["kind"], "signature")

    def test_the_authority_never_depends_on_emptiness_or_on_an_engine_area(self):
        for question in ("Signed by the consultant:", "I confirm the patient gave consent to proceed:"):
            places = [held(question_box_form(question, variant)).get("b/tbl[0]/tr[0]/tc[0]")
                      for variant in ("empty", "pict", "ink")]
            self.assertEqual(len({(p.code, p.ranges) for p in places}), 1, question)
        places = [held(column_form(variant))["b/tbl[0]/tr[1]/tc[1]"] for variant in ("empty", "pict", "ink")]
        self.assertEqual(len({(p.code, p.ranges) for p in places}), 1)


class IsolatedParserExchange(unittest.TestCase):
    """Through the service: the isolated parser asks for its clause readings, the parent reads them with the pinned
    parser, and the result equals the in-process reading - bound to the original's sha256."""

    def test_the_sandboxed_authority_equals_the_in_process_one(self):
        import hashlib
        import tempfile
        from pathlib import Path
        from .service import held_places
        data = lines_form("Initial assessment: ______", "Initial here: ______", "I confirm the plan: ______",
                          "Vital signs: ______", "Sign below: ______")
        with tempfile.TemporaryDirectory(prefix="held-places-") as directory:
            path = Path(directory) / "form.docx"
            path.write_bytes(data)
            report = held_places(path)
            self.assertEqual(report["sha256"], hashlib.sha256(data).hexdigest())
            self.assertEqual({row["slot"]: row["code"] for row in report["held"]},
                             {BODY + k: v.code for k, v in held(data).items()})
            inspected = inspect_template(path, engine="v3")
            self.assertEqual([(f["kind"], bool(f.get("write_protected"))) for f in inspected["fields"]],
                             [(f["kind"], bool(f.get("write_protected"))) for f in w.inspect(data)["fields"]])
            with self.assertRaises(DocumentError) as refused:
                locate_fields(path, [{"slot": BODY + "b/p[4]", "kind": "text", "label": "Sign below"}])
            self.assertEqual(refused.exception.code, "held_place")
            field = text_field("b/p[2]", {"type": "span", "start": 18, "end": 18, "shape": "after_label"})
            field["locator"]["fingerprint"] = w.fingerprint(w.SlotIndex(w.open_package(data)).by_id[BODY + "b/p[2]"].el)
            with self.assertRaises(DocumentError) as refused:
                fill_template(path, [field], {"probe": "SYNTHETIC-ANSWER"}, Path(directory) / "out.docx")
            self.assertEqual(refused.exception.code, "held_place")
            self.assertFalse((Path(directory) / "out.docx").exists())


if __name__ == "__main__":
    unittest.main()
