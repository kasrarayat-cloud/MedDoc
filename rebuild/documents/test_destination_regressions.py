"""Fixed physical-destination controls from the retained a620ebaf failures.

Pure synthetic OOXML tests: no database, provider, renderer or native Word.
Whole-document qualification remains a separate workflow gate.
"""
from copy import deepcopy
import io
import unittest
from zipfile import ZipFile
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import engine, word_v3


def saved(doc):
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


class DestinationRegressions(unittest.TestCase):
    def assert_held(self, raw):
        info = word_v3.inspect(raw)
        ordinary = next(f for f in info['fields'] if f['label'] == 'Clinical reason')
        targets = [f for f in info['fields'] if f['key'] != ordinary['key']]
        self.assertTrue(targets)
        for field in targets:
            forged = deepcopy(field)
            forged['kind'] = 'checkbox' if field['kind'] == 'checkbox' else 'text'
            forged.pop('write_protected', None)
            forged.pop('review_reason', None)
            with self.subTest(destination=field['label']):
                self.assertIn(forged['key'], word_v3.destination_roles(raw, [forged])['violations'])
                with self.assertRaises(word_v3.V3Failure):
                    word_v3.fill(raw, [forged], {forged['key']: 'checked' if forged['kind'] == 'checkbox'
                                              else 'Synthetic finalizer'})
        self.assertFalse(ordinary.get('write_protected'))
        exported = word_v3.fill(raw, info['fields'], {ordinary['key']: 'Synthetic reason'})
        self.assertIn('Clinical reason: Synthetic reason',
                      [p.text for p in Document(io.BytesIO(exported)).paragraphs])
        return info, exported

    def test_retained_actor_modal_failures_protect_actor_action_checkbox_and_table(self):
        # These are the seven exact unsafe retained inputs, never unseen cases.
        for source in ('Person ______ will confirm this form: ______',
                       'Person ______ will attest this form: ______',
                       'Person ______ shall verify this document: ______',
                       'Person ______ cannot verify this document: ______',
                       'Person ______ would approve this form: ______',
                       '☐ Person ______ will confirm this form: ______'):
            with self.subTest(source=source):
                doc = Document(); doc.add_paragraph(source); doc.add_paragraph('Clinical reason: ______')
                _, exported = self.assert_held(saved(doc))
                self.assertEqual(Document(io.BytesIO(exported)).paragraphs[0].text, source)
        doc = Document(); table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Person will confirm this form'; table.cell(0, 1).text = '______'
        doc.add_paragraph('Clinical reason: ______')
        _, exported = self.assert_held(saved(doc))
        self.assertEqual(Document(io.BytesIO(exported)).tables[0].cell(0, 1).text, '______')

    def test_genuine_signing_roles_never_receive_finalizer_identity(self):
        for source in ('Clinician signature: ______', 'Witness signature and date: ______',
                       'Patient signature: ______', 'Manager approval: ______',
                       '☐ I grant permission for this procedure: ______',
                       '☐ I observed the patient signing: ______'):
            with self.subTest(source=source):
                doc = Document(); doc.add_paragraph(source); doc.add_paragraph('Clinical reason: ______')
                _, exported = self.assert_held(saved(doc))
                self.assertEqual(Document(io.BytesIO(exported)).paragraphs[0].text, source)

    def test_retained_independent_caption_after_two_spaces_or_linebreak_is_ordinary(self):
        for separator in ('  ', '\n', '\t\t', '    ', '. '):
            with self.subTest(separator=repr(separator)):
                source = 'I approve this form: ______' + separator + 'Clinical reason: ______'
                doc = Document(); doc.add_paragraph(source); raw = saved(doc)
                info = word_v3.inspect(raw)
                ordinary = next(f for f in info['fields'] if f['label'] == 'Clinical reason')
                self.assertFalse(ordinary.get('write_protected'))
                self.assertEqual(word_v3.destination_roles(raw, [ordinary])['violations'], {})
                signer = info['fields'][0]
                self.assertTrue(signer.get('write_protected'))
                exported = word_v3.fill(raw, info['fields'], {ordinary['key']: 'Synthetic reason'})
                self.assertEqual(Document(io.BytesIO(exported)).paragraphs[0].text,
                                 source.removesuffix('______') + 'Synthetic reason')
        # A blank actor followed by a modal/action remains one human clause.
        for separator in ('  ', '\n'):
            doc = Document(); doc.add_paragraph('Person ______' + separator + 'will confirm this form: ______')
            doc.add_paragraph('Clinical reason: ______')
            self.assert_held(saved(doc))

    def test_ordinary_initial_signed_quoted_clinical_controls_remain_writable(self):
        for caption in ('Initial symptoms occurred at', 'Initial diagnosis',
                        'Signed consent was received on', 'Signed consent was recorded on',
                        'The patient has received the signed document on',
                        'The patient will receive this document on',
                        'The patient has received the signed prescription on',
                        'Consent was obtained by', 'Person will review the medication dose',
                        '"I observed a sign of dehydration"', '"Initial pain score"'):
            for shape in ('paragraph', 'table'):
                with self.subTest(caption=caption, shape=shape):
                    doc = Document()
                    if shape == 'paragraph': doc.add_paragraph(caption + ': ______')
                    else:
                        table = doc.add_table(rows=1, cols=2)
                        table.cell(0, 0).text = caption; table.cell(0, 1).text = '______'
                    raw = saved(doc); info = word_v3.inspect(raw); target = info['fields'][0]
                    self.assertFalse(target.get('write_protected'))
                    self.assertNotIn(target['kind'], ('signature', 'initials'))
                    self.assertEqual(word_v3.destination_roles(raw, [target])['violations'], {})
                    exported = Document(io.BytesIO(word_v3.fill(raw, info['fields'],
                                                                {target['key']: 'Synthetic value'})))
                    self.assertEqual(exported.tables[0].cell(0, 1).text if shape == 'table'
                                     else exported.paragraphs[0].text,
                                     'Synthetic value' if shape == 'table' else caption + ': Synthetic value')

    def test_precise_clinical_write_preserves_original_checks_link_and_package_parts(self):
        doc = Document(); p = doc.add_paragraph('☒ Existing selection')
        link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), 'retained_reference')
        run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'Reference'
        run.append(text); link.append(run); p._p.append(link)
        doc.add_paragraph('I approve this form: ______  Clinical reason: ______')
        doc.add_paragraph('Neighbor: ______'); raw = saved(doc)
        info = word_v3.inspect(raw)
        target = next(f for f in info['fields'] if f['label'] == 'Clinical reason')
        out = word_v3.fill(raw, info['fields'], {target['key']: 'Synthetic reason'})
        with ZipFile(io.BytesIO(raw)) as before, ZipFile(io.BytesIO(out)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for part in before.namelist():
                if part != 'word/document.xml': self.assertEqual(before.read(part), after.read(part))
        actual = Document(io.BytesIO(out))
        self.assertEqual(actual.paragraphs[0].text, Document(io.BytesIO(raw)).paragraphs[0].text)
        self.assertEqual(actual.paragraphs[2].text, 'Neighbor: ______')
        self.assertEqual(actual.paragraphs[0]._p.find(qn('w:hyperlink')).get(qn('w:anchor')), 'retained_reference')

    def test_counted_answer_restriction_is_not_a_person_granting_permission(self):
        for caption in ('Treatment(s) - only one answer allowed:',
                        'Treatment(s) - one option permitted:'):
            with self.subTest(caption=caption):
                doc = Document(); doc.add_paragraph(caption + ' ☐ Therapy ☐ Exercise')
                raw = saved(doc); info = word_v3.inspect(raw); field = info['fields'][0]
                self.assertFalse(field.get('write_protected'))
                out = Document(io.BytesIO(word_v3.fill(raw, [field], {field['key']: 'Exercise'})))
                self.assertIn('☒ Exercise', out.paragraphs[0].text)
        for source in ('☐ I allow only one answer', '☐ Person allowed this procedure',
                       '☐ Person allowed one answer to this form'):
            with self.subTest(source=source):
                doc = Document(); doc.add_paragraph(source); doc.add_paragraph('Clinical reason: ______')
                self.assert_held(saved(doc))

    def test_reserved_authorship_is_physical_and_does_not_depend_on_a_job_title_list(self):
        for author in ('the line manager', 'the radiology supervisor', 'Dr. Synthetic'):
            for shape in ('table', 'paragraph'):
                with self.subTest(author=author, shape=shape):
                    source = 'This section is for ' + author + ' to add comments'
                    doc = Document()
                    if shape == 'table':
                        table = doc.add_table(rows=1, cols=2)
                        table.cell(0, 0).text = source
                    else:
                        doc.add_paragraph(source + ': ______')
                    doc.add_paragraph('Clinical reason: ______')
                    self.assert_held(saved(doc))
        # An authorship role in a neighboring question grants no authority and
        # cannot hold a separately printed clinical answer cell.
        doc = Document(); table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'This section is for the line manager to add comments'
        table.cell(1, 0).text = 'Clinical reason'; raw = saved(doc)
        info = word_v3.inspect(raw)
        ordinary = next(f for f in info['fields'] if f['label'] == 'Clinical reason')
        self.assertFalse(ordinary.get('write_protected'))
        out = Document(io.BytesIO(word_v3.fill(raw, info['fields'], {ordinary['key']: 'Rest'})))
        self.assertEqual(out.tables[0].cell(0, 1).text, '')
        self.assertEqual(out.tables[0].cell(1, 1).text, 'Rest')

    def test_retained_refutation_ownership_frames_refuse_precise_manager_writes(self):
        # Six exposed 2de63efe failures: three ownership frames in two physical
        # layouts. They are retained controls, never independent holdout cases.
        for source in ('This section is for the line manager to write comments',
                       'This section is reserved for the line manager to add comments',
                       'Comments to be added by the line manager'):
            for shape in ('table', 'paragraph'):
                with self.subTest(source=source, shape=shape):
                    doc = Document()
                    if shape == 'table':
                        table = doc.add_table(rows=2, cols=2)
                        table.cell(0, 0).text = source
                        table.cell(1, 0).text = 'Clinical reason'
                    else:
                        doc.add_paragraph(source + ': ______  Clinical reason: ______')
                    raw = saved(doc); info = word_v3.inspect(raw)
                    ordinary = next(f for f in info['fields'] if f['label'] == 'Clinical reason')
                    assigned = next(f for f in info['fields'] if f['key'] != ordinary['key'])
                    with self.assertRaises(word_v3.V3Failure):
                        word_v3.fill(raw, info['fields'], {assigned['key']: 'Synthetic finalizer'})
                    forged = deepcopy(assigned)
                    forged['kind'] = 'text'; forged['label'] = 'Unrelated answer'
                    forged.pop('write_protected', None); forged.pop('review_reason', None)
                    self.assertIn(forged['key'], word_v3.destination_roles(raw, [forged])['violations'])
                    with self.assertRaises(word_v3.V3Failure):
                        word_v3.fill(raw, [forged], {forged['key']: 'Synthetic finalizer'})
                    self.assertFalse(ordinary.get('write_protected'))
                    self.assertEqual(word_v3.destination_roles(raw, [ordinary])['violations'], {})
                    out = Document(io.BytesIO(word_v3.fill(raw, info['fields'], {ordinary['key']: 'Rest'})))
                    if shape == 'table':
                        self.assertEqual(out.tables[0].cell(0, 1).text, '')
                        self.assertEqual(out.tables[0].cell(1, 1).text, 'Rest')
                    else:
                        self.assertEqual(out.paragraphs[0].text, source + ': ______  Clinical reason: Rest')

    def test_complete_ownership_grammar_is_shared_by_checkbox_and_comment_areas(self):
        for source in ('This section is for the radiology supervisor to dictate comments',
                       'This section is reserved for Dr. Synthetic to supply a statement',
                       'Concerns to be recorded by the nominated advocate'):
            with self.subTest(source=source):
                doc = Document(); doc.add_paragraph('☐ ' + source)
                doc.add_paragraph('Clinical reason: ______')
                _, out = self.assert_held(saved(doc))
                self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text, '☐ ' + source)

    def test_authorship_receipts_and_clinical_recording_remain_ordinary(self):
        for caption in ('Line manager comments received on',
                        'This section is for the clinician to record initial symptoms',
                        'Clinical comments about the line manager',
                        'Comments received by the line manager on',
                        'This section is for the clinician to record signed document receipt'):
            with self.subTest(caption=caption):
                doc = Document(); table = doc.add_table(rows=1, cols=2)
                table.cell(0, 0).text = caption; raw = saved(doc)
                info = word_v3.inspect(raw); field = info['fields'][0]
                self.assertFalse(field.get('write_protected'))
                self.assertEqual(word_v3.destination_roles(raw, [field])['violations'], {})
                out = Document(io.BytesIO(word_v3.fill(raw, [field], {field['key']: 'Synthetic value'})))
                self.assertEqual(out.tables[0].cell(0, 1).text, 'Synthetic value')
        doc = Document()
        doc.add_paragraph('This section is for Dr. Synthetic to add comments. Clinical reason: ______')
        raw = saved(doc); info = word_v3.inspect(raw); field = info['fields'][0]
        self.assertFalse(field.get('write_protected'))
        self.assertEqual(word_v3.destination_roles(raw, [field])['violations'], {})


class LegacyPhysicalSources(unittest.TestCase):
    def report(self, labels):
        doc = Document(); table = doc.add_table(rows=len(labels), cols=2)
        for number, label in enumerate(labels): table.cell(number, 0).text = label
        raw = saved(doc)
        with TemporaryDirectory(prefix='quillmd-legacy-physical-') as directory:
            path = Path(directory) / 'input.docx'; path.write_bytes(raw)
            document, paragraphs, cells, fields, _ = engine.docx_inspect(path)
            result = engine.legacy_roles(fields, fields, source=(raw, document, paragraphs, cells))
        return raw, fields, result

    def test_retained_ambiguous_consent_and_manager_statements_are_bound_to_original_extent(self):
        for label in ('Consent form signed', 'Care plan signed and date of next review',
                      'This section is for the line manager to add comments'):
            with self.subTest(label=label):
                raw, fields, result = self.report([label, 'Plan'])
                actions = result['legacy_unresolved_actions']
                self.assertTrue(actions)
                self.assertTrue(all(row['printed_label'] == label for row in actions))
                self.assertTrue(all(row['place'][0] == 'word/document.xml#b/tbl[0]/tr[0]/tc[1]' for row in actions))
                self.assertEqual(Document(io.BytesIO(raw)).tables[0].cell(0, 1).text, '')

    def test_ordinary_initial_signed_receipt_and_known_human_signer_controls_keep_their_roles(self):
        raw, fields, result = self.report(['Initial diagnosis', 'Initial GCS', 'Signed consent was received on',
                                          'The patient has received the signed document on', 'Electronically signed by',
                                          'Signature (participant)', 'Nurse initials', 'Plan'])
        self.assertEqual(result['legacy_unresolved_actions'], [])
        roles = result['legacy_roles']
        for field in fields[:4] + fields[-1:]: self.assertEqual(roles[field['key']]['role'], '')
        for field in fields[4:7]: self.assertTrue(roles[field['key']]['role'])

    def test_direct_legacy_export_refuses_unresolved_original_without_writing_output(self):
        raw, fields, _ = self.report(['Consent form signed', 'Plan'])
        with TemporaryDirectory(prefix='quillmd-legacy-export-') as directory:
            workspace = Path(directory); (workspace / 'input.docx').write_bytes(raw)
            with self.assertRaises(engine.DocumentFailure) as refusal:
                engine.process(workspace, {'operation': 'fill', 'format': 'docx', 'engine': 'legacy',
                                          'mapping': fields, 'answers': {fields[0]['key']: 'Yes'}})
            self.assertEqual(refusal.exception.code, 'unresolved_legacy_roles')
            self.assertFalse((workspace / 'output.docx').exists())
            self.assertEqual((workspace / 'input.docx').read_bytes(), raw)
