"""Reproducible synthetic format qualification; no provider calls or retained fills.

Run from rebuild/: python -m documents.qualification --report .runtime/document-qualification.json
Only bounded metadata is retained. The accepted blank WM IOV form is optional.
"""
import argparse
import hashlib
import json
from pathlib import Path
import tempfile

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.shared import Inches
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from .service import DocumentError, fill_template, inspect_template


def run(report_path, wm_iov=None):
    results = []
    with tempfile.TemporaryDirectory(prefix="quillmd-qualification-") as directory:
        root = Path(directory)

        def qualify(name, original, expected, answers=None, verify=None):
            original_hash = hashlib.sha256(original.read_bytes()).hexdigest()
            row = {"case": name, "source_sha256": original_hash, "expected": expected}
            output = root / (name + "-filled" + original.suffix)
            try:
                inspection = inspect_template(original)
                row.update(fields=len(inspection["fields"]), sandbox=inspection["sandbox"])
                if answers is not None:
                    values = {field["key"]: answers for field in inspection["fields"]}
                    fill_template(original, inspection["fields"], values, output)
                    if verify:
                        verify(output, values)
                    row["complete_answer_characters"] = len(answers)
                row["observed"] = "supported"
            except DocumentError as error:
                row.update(observed=error.code, message=str(error), no_partial_export=not output.exists())
            row["original_unchanged"] = hashlib.sha256(original.read_bytes()).hexdigest() == original_hash
            row["passed"] = row["observed"] == expected and row["original_unchanged"]
            results.append(row)

        def word(name, kind):
            document = Document()
            document.add_heading("Synthetic qualification form", 0)
            if kind == "styled":
                paragraph = document.add_paragraph("Question: ")
                paragraph.runs[0].bold = True
                paragraph.add_run("{{His").italic = True
                paragraph.add_run("tory}}")
                document.sections[0].header.paragraphs[0].text = "Synthetic {{Visit date}}"
                document.sections[0].footer.paragraphs[0].text = "{{Plan}}"
            elif kind == "sdt":
                node = OxmlElement("w:sdt")
                content = OxmlElement("w:sdtContent")
                content.append(OxmlElement("w:p")); node.append(content)
                document._element.body.append(node)
            elif kind == "textbox":
                document.add_paragraph()._p.append(OxmlElement("w:txbxContent"))
            elif kind == "tracked":
                document.add_paragraph()._p.append(OxmlElement("w:ins"))
            elif kind == "legacy":
                document.add_paragraph()._p.append(OxmlElement("w:fldSimple"))
            else:
                table = document.add_table(rows=1, cols=2)
                table.cell(0, 0).text = "History"
                if kind == "fixed":
                    table.rows[0].height = Inches(.1)
                    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
                if kind == "merged":
                    table.cell(0, 0).merge(table.cell(0, 1))
                if kind == "nested":
                    table.cell(0, 1).add_table(rows=1, cols=2)
            location = root / (name + ".docx")
            document.save(location)
            return location

        def pdf(name, kind="text", width=450, height=90):
            location = root / (name + ".pdf")
            document = canvas.Canvas(str(location))
            document.drawString(30, 780, "Synthetic question: History")
            if kind == "text":
                document.acroForm.textfield(name="History", x=30, y=650, width=width,
                    height=height, fontName="Helvetica", fontSize=10, maxlen=20000,
                    fieldFlags="multiline")
            elif kind == "checkbox":
                document.acroForm.checkbox(name="History", x=30, y=650)
            elif kind == "dropdown":
                document.acroForm.choice(name="History", x=30, y=650,
                    options=["Yes", "No"], value="Yes")
            elif kind == "scan":
                # A bounded, generated image models a scan without using patient content.
                from PIL import Image, ImageDraw
                picture = Image.new("RGB", (500, 100), "white")
                ImageDraw.Draw(picture).text((10, 10), "Synthetic scanned question and blank answer", fill="black")
                png = root / "synthetic-scan.png"; picture.save(png)
                document.drawImage(str(png), 30, 650, width=450, height=90)
            document.showPage(); document.save()
            return location

        def verify_word(output, values):
            assert Document(output).tables[0].cell(0, 1).text == next(iter(values.values()))

        def verify_pdf(output, values):
            reader = PdfReader(output)
            assert reader.get_form_text_fields() == values
            assert "Synthetic question: History" in reader.pages[0].extract_text()
            assert "/AP" in reader.pages[0]["/Annots"][0].get_object()

        qualify("docx-question-table", word("table", "table"), "supported", "Synthetic documented finding.", verify_word)
        qualify("docx-styled-body-header-footer", word("styled", "styled"), "supported", "Synthetic text")
        qualify("docx-long-19000-characters", word("long", "table"), "supported", "A" * 19000, verify_word)
        for kind, expected in [("sdt", "unsupported_structure"), ("textbox", "unsupported_structure"),
                               ("tracked", "unsupported_structure"), ("legacy", "unsupported_structure"),
                               ("fixed", "overflow_risk"), ("merged", "no_fields"), ("nested", "no_fields")]:
            qualify("docx-" + kind, word(kind, kind), expected)
        qualify("pdf-multiline-text", pdf("text"), "supported", "Synthetic first line.\nSynthetic second line.", verify_pdf)
        qualify("pdf-answer-overflow", pdf("narrow", width=80), "answer_overflow", "W" * 40)
        qualify("pdf-unicode", pdf("unicode"), "unsupported_characters", "Synthetic αβγ")
        qualify("pdf-checkbox", pdf("checkbox", "checkbox"), "unsupported_field")
        qualify("pdf-dropdown", pdf("dropdown", "dropdown"), "unsupported_field")
        qualify("pdf-flat-text", pdf("flat", "flat"), "flat_pdf")
        qualify("pdf-image-scan", pdf("scan", "scan"), "flat_pdf")
        dangerous = root / "javascript.pdf"
        writer = PdfWriter(clone_from=pdf("javascript-base")); writer.add_js("app.alert('synthetic')"); writer.write(dangerous)
        qualify("pdf-javascript", dangerous, "unsafe_upload")
        if wm_iov:
            qualify("accepted-blank-wm-iov", Path(wm_iov), "unsupported_structure")
    report = {"synthetic_only": True, "provider_calls": 0, "scratch_removed": True,
              "cases": results, "passed": all(row["passed"] for row in results)}
    report_path = Path(report_path)
    report_path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    report_path.write_text(json.dumps(report, indent=2) + "\n")
    report_path.chmod(0o600)
    print(json.dumps({"cases": len(results), "passed": sum(row["passed"] for row in results),
                      "all_passed": report["passed"]}))
    return report["passed"]


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", required=True)
    parser.add_argument("--wm-iov")
    args = parser.parse_args()
    raise SystemExit(0 if run(args.report, args.wm_iov) else 1)
