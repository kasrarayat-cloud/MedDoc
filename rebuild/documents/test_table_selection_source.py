"""Pure source/ownership controls for a captioned original selection table."""
from copy import deepcopy
import io
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import word_v3 as word


def data(document):
    out = io.BytesIO(); document.save(out); return out.getvalue()


def flat(document, caption='Options (Select all that apply):', cells=2):
    label = document.add_paragraph(caption)
    table = document.add_table(rows=1, cols=cells)
    for i, cell in enumerate(table.rows[0].cells):
        cell.text = '☐ Option %s-A' % i
        cell.add_paragraph('☐ Option %s-B' % i)
    return label, table


def matrix(document, instruction='Locations (Select all that apply):', title='Region Table'):
    caption = document.add_paragraph(instruction)
    document.add_paragraph('')
    title_p = document.add_paragraph(title)
    table = document.add_table(rows=3, cols=3)
    for cell, text in zip(table.rows[0].cells, ('Region', 'Right', 'Left')):
        cell.text = text
    for r, row in enumerate(table.rows[1:]):
        row.cells[0].text = 'Location %s' % r
        for cell in row.cells[1:]:
            cell.text = ''
            cell.add_paragraph('☐ Acute  ☐ Chronic  ☐ Uncertain')
    return caption, title_p, table


def related(fields):
    return [f for f in fields if any(r.get('relation') == 'table_selection_instruction'
                                   for r in f.get('source_questions', []))]


def old_fields(source):
    detector = word.Detector(word.open_package(source), _table_instructions=False)
    areas, _ = detector.detect()
    return [word.area_to_field(i, a) for i, a in enumerate(areas, 1)
            if a.target.get('type') == 'checkboxes']


def native(p, label, checked=False):
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    ff = OxmlElement('w:ffData'); ff.append(OxmlElement('w:enabled'))
    checkbox = OxmlElement('w:checkBox'); state = OxmlElement('w:default')
    state.set(qn('w:val'), '1' if checked else '0'); checkbox.append(state); ff.append(checkbox); begin.append(ff)
    p.add_run()._r.append(begin)
    instruction = OxmlElement('w:instrText'); instruction.text = ' FORMCHECKBOX '
    p.add_run()._r.append(instruction)
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end'); p.add_run()._r.append(end)
    p.add_run(' ' + label)


class TableSelectionSource(unittest.TestCase):
    def test_native_parallel_list_reads_and_writes_every_original_option(self):
        doc = Document(); doc.add_paragraph('Sequences (Select all that apply):')
        table = doc.add_table(rows=1, cols=2)
        for c, cell in enumerate(table.rows[0].cells):
            cell.text = ''
            for i in range(2):
                native(cell.add_paragraph(), 'Sequence %s-%s' % (c, i), checked=c == 1 and i == 0)
        source = data(doc); field = related(word.inspect(source)['fields'])[0]
        self.assertEqual(field['kind'], 'multi_choice'); self.assertEqual(len(field['options']), 4)
        final = word.fill(source, [field], {field['key']: [field['options'][0], field['options'][3]]})
        index = word.SlotIndex(word.open_package(final))
        self.assertEqual([word.box_state(index, b) for b in field['locator']['target']['boxes']],
                         [True, False, False, True])
        before = word.open_package(source).tree('word/document.xml')
        after = word.open_package(final).tree('word/document.xml')
        for tree in (before, after):
            for node in list(tree.iter(word.q('w:checked'))): node.getparent().remove(node)
        self.assertEqual(word.fingerprint(before), word.fingerprint(after))

    def test_complete_parallel_cells_share_all_options_and_preserve_original_package(self):
        doc = Document(); doc.add_paragraph('Fixed introduction'); flat(doc)
        doc.add_paragraph('Untargeted selection: ☒ Fixed  ☐ Other')
        source = data(doc); fields = related(word.inspect(source)['fields'])
        self.assertEqual(len(fields), 1)
        field = fields[0]; self.assertEqual(field['kind'], 'multi_choice')
        self.assertEqual(field['options'], ['Option 0-A', 'Option 0-B', 'Option 1-A', 'Option 1-B'])
        self.assertEqual(word.fill(source, fields, {}), source)
        final = word.fill(source, fields, {field['key']: ['Option 0-A', 'Option 1-B']})
        index = word.SlotIndex(word.open_package(final))
        self.assertEqual([word.box_state(index, b) for b in field['locator']['target']['boxes']],
                         [True, False, False, True])
        before = word.SlotIndex(word.open_package(source))
        for slot in before.slots:
            if slot.kind == 'p' and not slot.path.startswith('b/tbl[0]/'):
                self.assertEqual(word.fingerprint(slot.el), word.fingerprint(index.by_id[slot.id].el))
        with ZipFile(io.BytesIO(source)) as original, ZipFile(io.BytesIO(final)) as exported:
            self.assertEqual(original.namelist(), exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(original.read(name), exported.read(name), name)

    def test_single_caption_retains_one_selection_over_both_cells(self):
        doc = Document(); flat(doc, 'Options (choose one):')
        source = data(doc); field = related(word.inspect(source)['fields'])[0]
        self.assertEqual(field['kind'], 'choice'); self.assertEqual(len(field['options']), 4)
        with self.assertRaises(word.V3Failure):
            word.fill(source, [field], {field['key']: ['Option 0-A', 'Option 1-B']})
        final = word.fill(source, [field], {field['key']: 'Option 1-B'})
        index = word.SlotIndex(word.open_package(final))
        self.assertEqual([word.box_state(index, b) for b in field['locator']['target']['boxes']],
                         [False, False, False, True])

    def test_header_matched_title_scopes_instruction_to_only_that_matrix(self):
        doc = Document(); matrix(doc)
        doc.add_paragraph('Region Table')
        second = doc.add_table(rows=2, cols=2)
        second.cell(0, 0).text = 'Region'; second.cell(0, 1).text = 'Right'
        second.cell(1, 0).text = 'Later'; second.cell(1, 1).text = '☐ Acute  ☐ Chronic'
        fields = word.inspect(data(doc))['fields']; scoped = related(fields)
        self.assertEqual(len(scoped), 4)
        self.assertTrue(all(f['kind'] == 'multi_choice' for f in scoped))
        self.assertTrue(all(f['locator']['slot'].startswith('b/tbl[0]/') for f in scoped))
        self.assertFalse(any(f in scoped for f in fields if f['locator']['slot'].startswith('b/tbl[1]/')))

    def test_mismatched_title_intervening_question_and_wider_instruction_do_not_supply_scope(self):
        for mode in ('wrong_title', 'question', 'wider', 'no_command', 'one_over_rows'):
            with self.subTest(mode=mode):
                doc = Document()
                if mode == 'wrong_title':
                    matrix(doc, title='Other Table')
                elif mode == 'one_over_rows':
                    matrix(doc, instruction='Locations (choose one):')
                else:
                    instruction = ('Options for each question (Select all that apply):' if mode == 'wider'
                                   else 'Options:' if mode == 'no_command' else 'Options (Select all that apply):')
                    label, table = flat(doc, instruction)
                    if mode == 'question':
                        paragraph = OxmlElement('w:p'); run = OxmlElement('w:r'); text = OxmlElement('w:t')
                        text.text = 'Other question: ______'; run.append(text); paragraph.append(run)
                        table._tbl.addprevious(paragraph)
                self.assertFalse(related(word.inspect(data(doc))['fields']))

    def test_unsafe_caption_title_and_header_cannot_supply_new_scope(self):
        for mode in ('caption_link', 'caption_hidden', 'title_drawing', 'header_native'):
            with self.subTest(mode=mode):
                doc = Document(); caption, title, table = matrix(doc)
                if mode == 'caption_link':
                    link = OxmlElement('w:hyperlink'); run = caption._p.find(qn('w:r'))
                    caption._p.remove(run); link.append(run); caption._p.append(link)
                elif mode == 'caption_hidden':
                    caption.runs[0].font.hidden = True
                elif mode == 'title_drawing':
                    title.add_run()._r.append(OxmlElement('w:drawing'))
                else:
                    table.cell(0, 0).paragraphs[0].add_run()._r.append(OxmlElement('w:fldChar'))
                self.assertFalse(related(word.inspect(data(doc))['fields']))

    def test_local_cardinality_conflict_is_held_and_independent_questions_never_merge(self):
        doc = Document(); _, table = flat(doc)
        table.cell(0, 0).add_paragraph('Choose one only.')
        info = word.inspect(data(doc))
        self.assertFalse(any(f['kind'] == 'multi_choice' for f in related(info['fields'])))
        self.assertTrue(info['coverage_issues'])
        source = data(doc); prior = old_fields(source)
        with self.assertRaises(word.V3Failure):
            word.fill(source, prior, {f['key']: f['options'][0] for f in prior})
        doc = Document(); _, table = flat(doc)
        table.cell(0, 0).paragraphs[0].text = 'First question: ☐ Option A  ☐ Option B'
        table.cell(0, 1).paragraphs[0].text = 'Second question: ☐ Option C  ☐ Option D'
        fields = related(word.inspect(data(doc))['fields'])
        self.assertTrue(fields)
        self.assertTrue(all(len({b['slot'].split('/tc[')[1].split(']')[0]
                                for b in f['locator']['target']['boxes']}) == 1 for f in fields))

    def test_forged_or_stripped_instruction_reordered_boxes_and_duplicate_writers_refuse(self):
        doc = Document(); flat(doc); source = data(doc); field = related(word.inspect(source)['fields'])[0]
        for change in ('strip', 'kind', 'options', 'boxes', 'source'):
            with self.subTest(change=change):
                altered = deepcopy(field)
                if change == 'strip': altered.pop('source_questions')
                elif change == 'kind': altered['kind'] = 'choice'
                elif change == 'options': altered['options'].reverse()
                elif change == 'boxes': altered['locator']['target']['boxes'].reverse()
                else: altered['source_questions'][0]['sources'][0]['fingerprint'] = '0' * 64
                with self.assertRaises(word.V3Failure):
                    word.fill(source, [altered], {altered['key']: 'Option 0-A'})
        duplicate = deepcopy(field); duplicate['key'] = 'duplicate'
        with self.assertRaises(word.V3Failure):
            word.fill(source, [field, duplicate], {field['key']: ['Option 0-A'], 'duplicate': ['Option 1-B']})

    def test_exact_prior_one_answer_mappings_preserve_their_controls_without_widening_source(self):
        doc = Document(); flat(doc); source = data(doc); fields = old_fields(source)
        self.assertEqual(len(fields), 2)
        answers = {f['key']: f['options'][0] for f in fields}
        final = word.fill(source, fields, answers)
        index = word.SlotIndex(word.open_package(final))
        for field in fields:
            self.assertEqual([word.box_state(index, b) for b in field['locator']['target']['boxes']], [True, False])
        forged = deepcopy(fields[0]); forged['options'].reverse()
        with self.assertRaises(word.V3Failure):
            word.fill(source, [forged], {forged['key']: forged['options'][0]})
        doc = Document(); flat(doc, 'Options (choose one):'); source = data(doc); fields = old_fields(source)
        with self.assertRaises(word.V3Failure):
            word.fill(source, fields, {f['key']: f['options'][0] for f in fields})
        final = word.fill(source, [fields[0]], {fields[0]['key']: fields[0]['options'][0]})
        index = word.SlotIndex(word.open_package(final))
        self.assertTrue(word.box_state(index, fields[0]['locator']['target']['boxes'][0]))
        # An untargeted pre-existing check in the other cell is fixed source;
        # the legacy subset cannot silently untick it or add a second selection.
        doc = Document(); _, table = flat(doc, 'Options (choose one):')
        table.cell(0, 1).paragraphs[0].text = '☒ Option 1-A'
        source = data(doc); fields = old_fields(source)
        with self.assertRaises(word.V3Failure):
            word.fill(source, [fields[0]], {fields[0]['key']: fields[0]['options'][0]})

    def test_human_owned_caption_cannot_be_released_by_ordinary_label(self):
        doc = Document(); flat(doc, 'I certify this statement (Select all that apply):')
        source = data(doc); fields = related(word.inspect(source)['fields'])
        self.assertTrue(fields)
        for field in fields:
            altered = deepcopy(field); altered['label'] = 'Clinical options'; altered.pop('write_protected', None)
            with self.assertRaises(word.V3Failure):
                word.fill(source, [altered], {altered['key']: ['Option 0-A']})


if __name__ == '__main__':
    unittest.main()
