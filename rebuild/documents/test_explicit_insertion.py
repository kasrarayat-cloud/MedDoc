"""Pure regressions for source-owned publisher insertions and narrative bands."""
from copy import deepcopy
import io
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import word_v3 as word


def data(document):
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def paragraphs(raw):
    package = word.open_package(raw)
    return list(package.tree('word/document.xml').find(word.q('w:body')).iterchildren(word.q('w:p')))


def band(document, question, count=3):
    document.add_paragraph(question, style='Body Text')
    for _ in range(count):
        document.add_paragraph('', style='Body Text')


class ExplicitInsertion(unittest.TestCase):
    def test_inline_request_writes_only_its_original_span_and_retains_package(self):
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = 'Fixed letterhead'
        paragraph = doc.add_paragraph('Treatment category: ')
        paragraph.add_run('{insert type}').italic = True
        paragraph.add_run(' services are covered.').bold = True
        doc.add_paragraph('Existing selection: ☒ Yes  ☐ No')
        source = data(doc)
        field = next(f for f in word.inspect(source)['fields']
                     if f['locator']['target'].get('shape') == 'insertion')
        final = word.fill(source, [field], {field['key']: 'skilled nursing'})
        self.assertEqual(word.text_of(paragraphs(final)[0]),
                         'Treatment category: skilled nursing services are covered.')
        self.assertEqual(word.fingerprint(paragraphs(source)[1]), word.fingerprint(paragraphs(final)[1]))
        self.assertEqual(word.fingerprint(paragraphs(source)[0].find(word.q('w:pPr')))
                         if paragraphs(source)[0].find(word.q('w:pPr')) is not None else None,
                         word.fingerprint(paragraphs(final)[0].find(word.q('w:pPr')))
                         if paragraphs(final)[0].find(word.q('w:pPr')) is not None else None)
        with ZipFile(io.BytesIO(source)) as before, ZipFile(io.BytesIO(final)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(before.read(name), after.read(name), name)

    def test_whole_heading_is_an_explicit_insert_not_a_clinic_identity_guess(self):
        doc = Document()
        doc.add_paragraph('Insert contact information here', style='Heading 1')
        source = data(doc)
        fields = word.inspect(source)['fields']
        self.assertEqual(len(fields), 1)
        self.assertEqual(fields[0]['kind'], 'long_text')
        self.assertEqual(word.fill(source, fields, {}), source)
        final = word.fill(source, fields, {fields[0]['key']: 'Synthetic practice\n555-0100'})
        self.assertEqual(word.text_of(paragraphs(final)[0]), 'Synthetic practice\n555-0100')
        self.assertEqual(word.fingerprint(paragraphs(source)[0].find(word.q('w:pPr'))),
                         word.fingerprint(paragraphs(final)[0].find(word.q('w:pPr'))))

    def test_descriptor_slash_has_one_scalar_identifier_but_menus_remain_unresolved(self):
        for request in ('provider/plan toll-free telephone number', 'facility/department postal address'):
            with self.subTest(request=request):
                doc = Document(); doc.add_paragraph('Call {insert ' + request + '}.')
                source = data(doc); fields = word.inspect(source)['fields']
                self.assertEqual(len(fields), 1)
                self.assertIn('555-0100', word.text_of(paragraphs(word.fill(source, fields, {fields[0]['key']: '555-0100'}))[0]))
        for request in ('left/right', 'aspirin/warfarin', 'one or two', 'select a category'):
            with self.subTest(request=request):
                doc = Document(); doc.add_paragraph('Value {insert ' + request + '}.')
                info = word.inspect(data(doc))
                self.assertFalse(any(f['locator']['target'].get('shape') == 'insertion' for f in info['fields']))
                self.assertTrue(any(i['code'] == 'unresolved_insertion_source' for i in info['coverage_issues']))

    def test_unknown_braces_and_unbalanced_requests_are_not_new_scalar_answers(self):
        for printed in ('Fixed {clinical text}', '{insert type', '{insert {type}}', '{insert ' + 'x' * 170 + '}'):
            with self.subTest(printed=printed):
                doc = Document(); doc.add_paragraph(printed)
                info = word.inspect(data(doc))
                self.assertFalse(any(f['locator']['target'].get('shape') == 'insertion' for f in info['fields']))
                if 'insert' in printed:
                    self.assertTrue(any(i['code'] == 'unresolved_insertion_source' for i in info['coverage_issues']))

    def test_hidden_linked_native_and_revision_sources_are_not_plain_insertions(self):
        for marker in ('hidden', 'hyperlink', 'sdt', 'ins'):
            with self.subTest(marker=marker):
                doc = Document(); p = doc.add_paragraph()
                if marker == 'hidden':
                    p.add_run('{insert type}').font.hidden = True
                else:
                    wrapper = OxmlElement('w:' + marker)
                    run = OxmlElement('w:r'); text = OxmlElement('w:t')
                    text.text = '{insert type}'; run.append(text); wrapper.append(run); p._p.append(wrapper)
                if marker == 'ins':
                    with self.assertRaises(word.V3Failure) as refused:
                        word.inspect(data(doc))
                    self.assertEqual(refused.exception.code, 'tracked_changes')
                    continue
                info = word.inspect(data(doc))
                self.assertFalse(any(f['locator']['target'].get('shape') == 'insertion' for f in info['fields']))

    def test_removed_or_moved_source_relation_and_duplicate_writer_refuse(self):
        doc = Document(); doc.add_paragraph('Category: {insert type} services.')
        source = data(doc); field = word.inspect(source)['fields'][0]
        for mutation in ('drop', 'move', 'kind', 'request', 'duplicate'):
            with self.subTest(mutation=mutation), self.assertRaises(word.V3Failure) as refused:
                edited = deepcopy(field); fields = [edited]
                if mutation == 'drop': edited.pop('source_questions')
                elif mutation == 'move': edited['locator']['target'].update(start=0, end=8)
                elif mutation == 'kind': edited['kind'] = 'date'
                elif mutation == 'request': edited['source_questions'][0]['request'] = '{insert identity}'
                else:
                    second = deepcopy(field); second['key'] = 'duplicate'; fields.append(second)
                word.fill(source, fields, {f['key']: 'Synthetic answer' for f in fields})
            self.assertEqual(refused.exception.code, 'mapping_mismatch')

    def test_signing_and_other_author_assertions_stay_human_owned(self):
        for printed in ('I certify this report: {insert name}.', 'Manager signature: {insert name}.',
                        'I approve this document with {insert name}.'):
            with self.subTest(printed=printed):
                doc = Document(); doc.add_paragraph(printed)
                source = data(doc); fields = word.inspect(source)['fields']
                with self.assertRaises(word.V3Failure):
                    word.fill(source, fields, {f['key']: 'Dr Synthetic' for f in fields})

    def test_complete_long_question_owns_only_same_style_plain_blank_band(self):
        doc = Document()
        question = '• Detailed explanation of the clinical observations and relevant evidence used to determine the proposed ongoing treatment recommendations:'
        band(doc, question, count=4)
        doc.add_paragraph('', style='Normal')
        doc.add_paragraph('Fixed following instructions.')
        source = data(doc); fields = word.inspect(source)['fields']
        field = next(f for f in fields if any(r['relation'] == 'narrative_empty_band' for r in f.get('source_questions', [])))
        self.assertEqual(field['source_questions'][0]['question'], question)
        self.assertEqual([r['slot'] for r in field['locator']['extra']], ['b/p[2]', 'b/p[3]', 'b/p[4]'])
        final = word.fill(source, [field], {field['key']: 'Synthetic evidence\nSecond line'})
        old, new = paragraphs(source), paragraphs(final)
        for index in (0, 5, 6):
            self.assertEqual(word.fingerprint(old[index]), word.fingerprint(new[index]), index)
        self.assertIn('Synthetic evidence', word.text_of(new[1]))

    def test_real_inline_tab_does_not_appoint_the_following_spacer(self):
        doc = Document(); doc.add_paragraph('Date:\t'); doc.add_paragraph('')
        fields = word.inspect(data(doc))['fields']
        self.assertEqual(len(fields), 1)
        self.assertEqual(fields[0]['locator']['slot'], 'b/p[0]')
        self.assertEqual(fields[0]['locator']['target']['shape'], 'after_label')

    def test_consecutive_named_placeholders_have_independent_physical_writers(self):
        doc = Document(); cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = '[Addressee name]'
        cell.add_paragraph('[Address]')
        cell.add_paragraph('[Your contact telephone number]')
        cell.add_paragraph('{{Email address}}')
        source = data(doc); fields = word.inspect(source)['fields']
        self.assertEqual(len(fields), 4)
        self.assertEqual([f['locator']['slot'] for f in fields],
                         ['b/tbl[0]/tr[0]/tc[0]/p[%d]' % i for i in range(4)])
        self.assertTrue(all(not f['locator'].get('extra') for f in fields))
        answers = {f['key']: 'Synthetic value %d' % i for i, f in enumerate(fields)}
        final = word.fill(source, fields, answers)
        index = word.SlotIndex(word.open_package(final))
        for field in fields:
            self.assertEqual(word.text_of(index.by_id[field['locator']['part'] + '#' + field['locator']['slot']].el),
                             answers[field['key']])

    def test_anonymous_drawn_lines_still_continue_one_answer(self):
        doc = Document(); doc.add_paragraph('Clinical observations')
        doc.add_paragraph('__________'); doc.add_paragraph('__________')
        source = data(doc); fields = word.inspect(source)['fields']
        self.assertEqual(len(fields), 1)
        self.assertEqual(fields[0]['kind'], 'long_text')
        self.assertEqual([f['slot'] for f in fields[0]['locator']['extra']], ['b/p[2]'])
        final = word.fill(source, fields, {fields[0]['key']: 'Synthetic first line\nSecond line'})
        self.assertIn('Synthetic first line', word.text_of(paragraphs(final)[1]))

    def test_named_placeholder_keeps_its_preceding_signing_caption_role(self):
        doc = Document(); doc.add_paragraph('[Manager signature]'); doc.add_paragraph('[Name]')
        source = data(doc); fields = word.inspect(source)['fields']
        self.assertEqual(len(fields), 1)
        self.assertEqual(fields[0]['kind'], 'signature')
        with self.assertRaises(word.V3Failure):
            word.fill(source, fields, {fields[0]['key']: 'Dr Synthetic'})

    def test_narrative_caption_cannot_release_signature_or_patient_authorship(self):
        doc = Document(); doc.add_paragraph('Completed by the patient:')
        band(doc, '• Detailed account of the observations:')
        source = data(doc); fields = word.inspect(source)['fields']
        field = next(f for f in fields if f['locator']['target']['type'] == 'paragraph')
        self.assertEqual(field['answer_policy']['source_author'], 'patient')
        stripped = deepcopy(field); stripped.pop('answer_policy')
        with self.assertRaises(word.V3Failure):
            word.fill(source, [stripped], {stripped['key']: 'Clinician invented statement'})
        doc = Document(); band(doc, '• Manager signature:')
        source = data(doc); fields = word.inspect(source)['fields']
        with self.assertRaises(word.V3Failure):
            word.fill(source, fields, {f['key']: 'Dr Synthetic' for f in fields})


if __name__ == '__main__':
    unittest.main()
