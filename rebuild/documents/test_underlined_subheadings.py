"""Fixed underlined group captions and their separately printed answer lines."""
from copy import deepcopy
from io import BytesIO
import unittest
from zipfile import ZipFile

from docx import Document
from . import word_v3 as w


def form(caption='Before treatment:', *, gap='', children=('- First observation:  ', '- Second observation:  '),
         in_cell=False, mutate=None):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = 'Fixed letterhead'
    doc.sections[0].footer.paragraphs[0].text = 'Fixed footer'
    if in_cell:
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 1).text = 'Fixed neighboring entry'
        owner = table.cell(0, 0)
        heading = owner.paragraphs[0]
    else:
        owner = doc
        heading = doc.add_paragraph()
    heading.add_run(caption + gap).underline = True
    lines = [owner.add_paragraph(text) for text in children]
    doc.add_paragraph('Fixed receipt: already received')
    doc.add_paragraph('Selection retained: ☒ Present ☐ Absent')
    if mutate: mutate(doc, heading, lines)
    stream = BytesIO(); doc.save(stream); return stream.getvalue()


def slots(data):
    package = w.open_package(data)
    detector = w.Detector(package)
    detector.detect()
    return package, detector


def old_caption_field(original):
    _, detector = slots(original)
    heading = next(s for s in detector.index.slots if s.kind == 'p'
                   and w.text_of(s.el).startswith('Before treatment:'))
    text = w.paragraph_text(heading.el)[0]
    return dict(key='old_caption', label='Before treatment', kind='text',
                locator=dict(contract='v3', part=heading.part, slot=heading.path,
                             fingerprint=w.fingerprint(heading.el),
                             target=dict(type='span', start=len(text), end=len(text), shape='after_label')))


class UnderlinedSubheadings(unittest.TestCase):
    def test_complete_heading_is_fixed_and_subordinate_answers_keep_group_context(self):
        for cell in (False, True):
            with self.subTest(cell=cell):
                original = form(in_cell=cell)
                mapping = w.inspect(original)
                fields = {f['label']: f for f in mapping['fields']}
                self.assertNotIn('Before treatment', fields)
                for label in ('First observation', 'Second observation'):
                    self.assertIn('Before treatment', fields[label]['context'])
                values = {fields['First observation']['key']: 'First complete answer',
                          fields['Second observation']['key']: 'Second complete answer\nA second line'}
                exported = w.fill(original, mapping['fields'], values)
                old, detector = slots(original); new, _ = slots(exported)
                heading = next(s for s in detector.index.slots if s.id in detector.fixed_subheading_slots)
                newer = w.SlotIndex(new).by_id[heading.id]
                self.assertEqual(w.fingerprint(heading.el), w.fingerprint(newer.el))
                body_text = w.text_of(new.tree('word/document.xml'))
                self.assertIn('First complete answer', body_text)
                self.assertIn('Second complete answer', body_text)
                self.assertIn('A second line', body_text)
                self.assertIn('Fixed neighboring entry' if cell else 'Fixed receipt: already received', body_text)
                self.assertIn('☒ Present ☐ Absent', body_text)
                with ZipFile(BytesIO(original)) as a, ZipFile(BytesIO(exported)) as b:
                    for name in a.namelist():
                        if name != 'word/document.xml': self.assertEqual(a.read(name), b.read(name))
                self.assertEqual(w.fill(original, mapping['fields'], {}), original)

    def test_historical_and_forged_caption_writers_refuse_before_export(self):
        for cell in (False, True):
            original = form(in_cell=cell); historical = old_caption_field(original)
            for label in ('Before treatment', 'Ordinary value', 'Signed receipt'):
                field = deepcopy(historical); field['label'] = label
                with self.subTest(cell=cell, label=label):
                    with self.assertRaisesRegex(w.V3Failure, 'fixed underlined group heading.*Reinspect'):
                        w.fill(original, [field], {'old_caption': 'Synthetic heading answer'})
                    self.assertEqual(w.fill(original, [field], {}), original)

    def test_source_fixed_caption_cannot_be_added_as_a_new_manual_answer(self):
        for cell in (False, True):
            original = form(in_cell=cell)
            _, detector = slots(original)
            heading = next(s for s in detector.index.slots if s.id in detector.fixed_subheading_slots)
            for kind in ('text', 'long_text'):
                with self.subTest(cell=cell, kind=kind):
                    with self.assertRaisesRegex(w.V3Failure, 'fixed underlined group heading.*Nothing was added') as denied:
                        w.locate(original, [dict(slot=heading.id, kind=kind, label='Synthetic ordinary answer')])
                    self.assertEqual(denied.exception.code, 'invalid_mapping')

    def test_manual_subordinate_and_genuine_parent_answers_remain_usable(self):
        for original, prefix in ((form(), '- First observation:'), (form(gap='  '), 'Before treatment:')):
            _, detector = slots(original)
            slot = next(s for s in detector.index.slots if s.kind == 'p' and w.text_of(s.el).startswith(prefix))
            fields = w.locate(original, [dict(slot=slot.id, kind='text', label='Clinician reviewed value')])
            exported = w.fill(original, fields, {fields[0]['key']: 'Explicit manual value'})
            self.assertIn('Explicit manual value', w.text_of(w.open_package(exported).tree('word/document.xml')))

    def test_genuine_underlined_fields_and_parent_answers_are_not_excluded(self):
        for kwargs in ({'children': ()}, {'children': ('- First observation:  ',)},
                       {'gap': '  '}, {'gap': '\t'}, {'gap': ' ____'},
                       {'children': ('', '- First observation:  ', '- Second observation:  ')},
                       {'children': ('- First observation:  ', '- First observation:  ')}):
            with self.subTest(kwargs=kwargs):
                original = form(**kwargs); _, detector = slots(original)
                self.assertFalse(detector.fixed_subheading_slots)
                mapping = w.inspect(original)
                field = next(f for f in mapping['fields'] if f['label'] == 'Before treatment')
                exported = w.fill(original, mapping['fields'], {field['key']: 'Actual parent answer'})
                self.assertIn('Actual parent answer', w.text_of(w.open_package(exported).tree('word/document.xml')))

    def test_mixed_underlining_hidden_child_or_control_never_proves_a_fixed_heading(self):
        def mixed(doc, heading, lines):
            heading.runs[0].text = 'Before '
            heading.add_run('treatment:')
        def hidden(doc, heading, lines): lines[1].runs[0].font.hidden = True
        def control(doc, heading, lines): lines[1].add_run(' ☐')
        for change in (mixed, hidden, control):
            with self.subTest(change=change.__name__):
                _, detector = slots(form(mutate=change))
                self.assertFalse(detector.fixed_subheading_slots)

    def test_scope_cannot_cross_a_table_or_a_new_question(self):
        def boundary(doc, heading, lines):
            table = doc.add_table(rows=1, cols=1); table.cell(0, 0).text = 'Separate question:  '
            heading._p.addnext(table._tbl)
        for kwargs in ({'mutate': boundary},
                       {'children': ('A separate question:  ', '- First observation:  ', '- Second observation:  ')}):
            with self.subTest(kwargs=kwargs):
                _, detector = slots(form(**kwargs))
                self.assertFalse(detector.fixed_subheading_slots)

    def test_complete_human_caption_keeps_its_existing_source_protection(self):
        for caption in ('Date Alex Reed should sign:', 'Witness attestation:',
                        'This area is for Alex Reed to enter the answer:'):
            with self.subTest(caption=caption):
                original = form(caption=caption); _, detector = slots(original)
                self.assertFalse(detector.fixed_subheading_slots)
                mapping = w.inspect(original)
                fields = [f for f in mapping['fields'] if f['locator']['slot'] == 'b/p[0]']
                self.assertTrue(fields)
                self.assertTrue(all(f.get('write_protected') or f['kind'] in ('signature', 'initials') for f in fields))
                for field in fields:
                    with self.assertRaises(w.V3Failure):
                        w.fill(original, mapping['fields'], {field['key']: 'Dr Synthetic Finalizer'})
