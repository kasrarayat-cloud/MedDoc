"""Complete original-source ruled narrative detection and historical recovery."""
from copy import deepcopy
import io
import unittest
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from . import word_v3 as w


def saved(doc):
    stream = io.BytesIO(); doc.save(stream); return stream.getvalue()


def source(lines=14):
    doc = Document(); doc.add_heading('Clinical impressions', level=2)
    doc.add_paragraph('Impression: Overall impression of this patient')
    doc.add_paragraph('[Describe mental status, physical condition and differential diagnoses.]')
    doc.add_paragraph()
    for number in range(lines):
        p = doc.add_paragraph(); p.paragraph_format.tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT,
                                                                        WD_TAB_LEADER.LINES)
        p.add_run().add_tab()
        if number != lines - 1:
            p = doc.add_paragraph(); p.paragraph_format.tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT,
                                                                            WD_TAB_LEADER.LINES)
    doc.add_paragraph('Neighbor: ______')
    doc.sections[0].header.paragraphs[0].text = 'Synthetic clinic'
    doc.sections[0].footer.paragraphs[0].text = 'Preserved footer'
    return saved(doc)


def located(raw, slot='word/document.xml#b/p[4]'):
    return w.locate(raw, [{'slot': slot, 'kind': 'long_text', 'label': 'Impression'}])[0]


class RuledParagraphRecovery(unittest.TestCase):
    def test_explicit_review_binds_every_line_spacer_question_and_fixed_guidance(self):
        raw = source(); field = located(raw)
        self.assertEqual(field['origin'], 'manual')
        self.assertEqual(field['locator']['target'], {'type': 'ruled_paragraphs'})
        self.assertEqual(len(field['locator']['extra']), 13)
        relation = field['source_questions'][0]
        self.assertEqual(relation['relation'], 'manual_ruled_paragraph_answer')
        self.assertEqual(len(relation['descriptors']), 27)
        self.assertEqual(relation['source']['slot'], 'b/p[1]')
        self.assertEqual({ref['slot'] for ref in relation['instructions']}, {'b/p[2]', 'b/p[3]'})
        self.assertIn('differential diagnoses', field['context'])
        self.assertEqual(w.fill(raw, [field], {}), raw)
        automatic = [f for f in w.inspect(raw)['fields'] if f['locator']['target']['type'] == 'ruled_paragraphs']
        self.assertEqual(len(automatic), 1)
        self.assertNotIn('origin', automatic[0])
        self.assertEqual(automatic[0]['locator'], field['locator'])
        self.assertEqual(automatic[0]['source_questions'], field['source_questions'])

    def test_long_answer_is_complete_once_and_original_structure_and_other_parts_are_preserved(self):
        raw = source(); field = located(raw)
        answer = ('\n'.join('Synthetic sentence %02d, complete clinical review.' % n for n in range(19)) + '\n' + 'Long detail ' * 180).rstrip()
        out = w.fill(raw, [field], {field['key']: answer})
        self.assertEqual(w.effective_values(raw, [field], {field['key']: answer})[field['key']], answer)
        original, final = w.SlotIndex(w.open_package(raw)), w.SlotIndex(w.open_package(out))
        writes = {field['locator']['slot']} | {ref['slot'] for ref in field['locator']['extra']}
        self.assertEqual(len(original.slots), len(final.slots))
        for slot in original.slots:
            actual = final.by_id[slot.id]
            if slot.part != 'word/document.xml' or slot.path not in writes:
                self.assertEqual(w.fingerprint(slot.el), w.fingerprint(actual.el))
            else:
                self.assertEqual(w.fingerprint(slot.el.find(qn('w:pPr'))), w.fingerprint(actual.el.find(qn('w:pPr'))))
                self.assertEqual(len(slot.el.findall('.//' + qn('w:tab'))), len(actual.el.findall('.//' + qn('w:tab'))))
        with ZipFile(io.BytesIO(raw)) as before, ZipFile(io.BytesIO(out)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for part in before.namelist():
                if part != 'word/document.xml': self.assertEqual(before.read(part), after.read(part))
        self.assertEqual(w.fill(raw, [field], {field['key']: answer}), out)

    def test_middle_line_cannot_create_a_partial_suffix_and_inconsistent_late_lines_refuse_whole_group(self):
        with self.assertRaises(w.V3Failure): located(source(), slot='word/document.xml#b/p[6]')
        for variant in ('filled', 'hidden', 'leader', 'stop', 'unknown_style', 'native_control', 'break', 'numbering'):
            with self.subTest(variant=variant):
                doc = Document(io.BytesIO(source())); p = doc.paragraphs[26]
                if variant == 'filled': p.add_run('Existing entry')
                elif variant == 'hidden': p.runs[0].font.hidden = True
                elif variant in ('leader', 'stop'):
                    tab = p._p.find('.//' + qn('w:tab'))
                    tab.set(qn('w:leader' if variant == 'leader' else 'w:pos'), 'dot' if variant == 'leader' else '7000')
                elif variant == 'unknown_style':
                    node = OxmlElement('w:pStyle'); node.set(qn('w:val'), 'UnknownStyle'); p._p.get_or_add_pPr().append(node)
                elif variant == 'native_control':
                    from .test_word_v3 import _sdt
                    p._p.append(_sdt('rich', placeholder='Click or tap here to enter text.'))
                elif variant == 'break': p.add_run().add_break()
                else:
                    node = OxmlElement('w:numPr'); p._p.get_or_add_pPr().append(node)
                with self.assertRaises(w.V3Failure): located(saved(doc))

    def test_mapping_cannot_strip_ownership_or_drop_or_move_a_physical_line(self):
        raw = source(); original = located(raw)
        for variant in ('missing_relation', 'missing_origin', 'missing_extra', 'move_extra', 'descriptor', 'guidance', 'definitions', 'kind'):
            with self.subTest(variant=variant):
                field = deepcopy(original)
                if variant == 'missing_relation': field.pop('source_questions')
                elif variant == 'missing_origin': field.pop('origin'); field.pop('source_questions')
                elif variant == 'missing_extra': field['locator']['extra'].pop()
                elif variant == 'move_extra': field['locator']['extra'][-1] = {'part': 'word/document.xml', 'slot': 'b/p[31]', 'fingerprint': field['locator']['fingerprint']}
                elif variant == 'descriptor': field['source_questions'][0]['descriptors'].pop()
                elif variant == 'guidance': field['source_questions'][0]['instructions'].pop()
                elif variant == 'definitions': field['source_questions'][0]['definitions']['word/styles.xml'] = 'changed'
                else: field['kind'] = 'text'
                with self.assertRaises(w.V3Failure): w.fill(raw, [field], {field['key']: 'Synthetic answer'})

    def test_actual_signing_caption_cannot_be_relabelled_as_a_ruled_clinical_answer(self):
        doc = Document(io.BytesIO(source())); doc.paragraphs[1].text = 'Patient signature:'
        with self.assertRaises(w.V3Failure): located(saved(doc))

    def test_automatic_complete_answer_and_historical_manual_answer_export_identically(self):
        raw = source()
        automatic, = [f for f in w.inspect(raw)['fields'] if f['locator']['target']['type'] == 'ruled_paragraphs']
        manual = located(raw)
        answer = '\n'.join('Synthetic complete sentence %02d.' % n for n in range(21))
        automatic_export = w.fill(raw, [automatic], {automatic['key']: answer})
        manual_export = w.fill(raw, [manual], {manual['key']: answer})
        self.assertEqual(automatic_export, manual_export)
        self.assertEqual(w.effective_values(raw, [automatic], {automatic['key']: answer})[automatic['key']], answer)
        self.assertEqual(w.fill(raw, [automatic], {}), raw)

    def test_complete_human_clause_guidance_heading_and_continuous_act_refuse_automatic_and_forged_old_mapping(self):
        original = source()
        old = located(original)
        for variant in ('caption', 'guidance', 'heading', 'continuous_act', 'assigned_author', 'approval_clause'):
            with self.subTest(variant=variant):
                doc = Document(io.BytesIO(original))
                if variant == 'caption': doc.paragraphs[1].text = 'Date Alex Reed should sign:'
                elif variant == 'guidance': doc.paragraphs[2].text = '[Witness must enter the following:]'
                elif variant == 'heading': doc.paragraphs[0].text = 'Witness attestation'
                elif variant == 'continuous_act': doc.paragraphs[0].text = 'I certify that I witnessed the patient sign this consent form.'
                elif variant == 'assigned_author': doc.paragraphs[1].text = 'This area is for Dr. Alex Reed to enter the answer:'
                else: doc.paragraphs[2].text = '[Describe findings;I approve this document.]'
                raw = saved(doc)
                mapping = w.inspect(raw)
                self.assertFalse(any(f['locator']['target']['type'] == 'ruled_paragraphs' for f in mapping['fields']))
                self.assertTrue(any(i['code'] == 'unmapped_answer' for i in mapping['coverage_issues']))
                with self.assertRaises(w.V3Failure): located(raw)
                with self.assertRaises(w.V3Failure): w.fill(raw, [old], {old['key']: 'Dr Synthetic Finalizer'})

    def test_automatic_partial_group_is_excluded_and_old_reordered_or_partial_writer_refuses(self):
        original = source()
        automatic, = [f for f in w.inspect(original)['fields'] if f['locator']['target']['type'] == 'ruled_paragraphs']
        for variant in ('late_stop', 'late_ink', 'different_neighbor', 'reordered', 'partial'):
            with self.subTest(variant=variant):
                field = deepcopy(automatic)
                if variant in ('reordered', 'partial'):
                    if variant == 'reordered': field['locator']['extra'].reverse()
                    else: field['locator']['extra'].pop()
                    with self.assertRaises(w.V3Failure): w.fill(original, [field], {field['key']: 'Synthetic answer'})
                    continue
                doc = Document(io.BytesIO(original))
                if variant == 'late_stop': doc.paragraphs[26]._p.find('.//' + qn('w:tab')).set(qn('w:pos'), '7000')
                elif variant == 'late_ink': doc.paragraphs[26].add_run('Existing entry')
                else:
                    doc.paragraphs[16].text = 'Other question: ______'
                raw = saved(doc)
                mapping = w.inspect(raw)
                self.assertFalse(any(f['locator']['target']['type'] == 'ruled_paragraphs' for f in mapping['fields']))
                self.assertTrue(any(i['code'] == 'unmapped_answer' for i in mapping['coverage_issues']))
                with self.assertRaises(w.V3Failure): w.fill(raw, [field], {field['key']: 'Synthetic answer'})

    def test_ordinary_initial_signed_receipt_and_clinical_question_have_useful_complete_destinations(self):
        for caption, guidance in (
                ('Initial observations:', '[Describe the current clinical findings.]'),
                ('Signed referral receipt:', '[Signed report received on Tuesday.]'),
                ('Review of systems:', '[Describe symptoms; the patient should remain hydrated.]')):
            with self.subTest(caption=caption):
                doc = Document(io.BytesIO(source()))
                doc.paragraphs[1].text = caption
                doc.paragraphs[2].text = guidance
                raw = saved(doc)
                field, = [f for f in w.inspect(raw)['fields'] if f['locator']['target']['type'] == 'ruled_paragraphs']
                answer = 'Synthetic complete clinical response.\nReviewed second sentence.'
                out = w.fill(raw, [field], {field['key']: answer})
                self.assertEqual(w.effective_values(raw, [field], {field['key']: answer})[field['key']], answer)
                self.assertEqual(w.fill(raw, [field], {}), raw)
                before, after = w.SlotIndex(w.open_package(raw)), w.SlotIndex(w.open_package(out))
                writes = {field['locator']['slot'], *(x['slot'] for x in field['locator']['extra'])}
                for slot in before.slots:
                    if slot.part != 'word/document.xml' or slot.path not in writes:
                        self.assertEqual(w.fingerprint(slot.el), w.fingerprint(after.by_id[slot.id].el))

    def test_refreshing_all_physical_fingerprints_cannot_relabel_a_human_source_as_clinical(self):
        ordinary = located(source())
        for target, text in ((1, 'Date Alex Reed should sign:'),
                             (2, '[Witness should write the following:]'),
                             (2, '[Describe findings;I approve this document.]')):
            with self.subTest(target=target, text=text):
                doc = Document(io.BytesIO(source())); doc.paragraphs[target].text = text
                raw = saved(doc); index = w.SlotIndex(w.open_package(raw))
                field = deepcopy(ordinary)
                refs = [field['locator'], *field['locator']['extra']]
                relation = field['source_questions'][0]
                refs += [relation['source'], *relation['descriptors'], *relation['instructions']]
                for ref in refs:
                    slot = index.by_id[ref['part'] + '#' + ref['slot']]
                    ref['fingerprint'] = w.fingerprint(slot.el)
                    if 'text' in ref: ref['text'] = w.paragraph_text(slot.el)[0]
                relation['question'] = w.paragraph_text(index.by_id['word/document.xml#b/p[1]'].el)[0].strip()
                with self.assertRaises(w.V3Failure): w.fill(raw, [field], {field['key']: 'Dr Synthetic Finalizer'})
