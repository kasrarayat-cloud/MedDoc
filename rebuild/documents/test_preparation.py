"""Blank-only finite placement, preservation and model-schema counterexamples."""
from copy import deepcopy
import hashlib
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile
from docx import Document
from lxml import etree
from .engine import DocumentFailure, docx_inspect, W_NS
from .preparation import inventory,candidate,validate_role,model_inventory,validate_candidate_fields,digest

class PreparationDocumentTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)

    def make(self,paragraphs=(),table=None):
        d=Document()
        for text in paragraphs:d.add_paragraph(text)
        if table:
            t=d.add_table(rows=len(table),cols=len(table[0]))
            for row,values in zip(t.rows,table):
                for cell,value in zip(row.cells,values):cell.text=value
        p=self.root/'original.docx';d.save(p);return p

    def role(self,inv,role):
        field,value=('decision','map') if role=='proposer' else ('verdict','pass')
        return {'inventory_sha256':inv['inventory_sha256'],'anchors':[
            {'anchor_id':a['id'],'question_id':a['question_id'],field:value} for a in inv['anchors']],
            'stories':[{'story_id':s['id'],'content_sha256':s['content_sha256'],'verdict':'retain'}
                       for s in model_inventory(inv)['stories']]}

    def test_underscores_prepare_without_changing_questions_units_options_or_signing(self):
        p=self.make(['Instructions: preserve °C and Yes / No.','Temperature: ______ °C',
                     'Fever: ______ (Yes / No)','Patient signature: __________'])
        inv=inventory(p);self.assertFalse(inv['issues'])
        before=p.read_bytes();q=self.root/'candidate.docx'
        result=candidate(p,q,inv,[a['id'] for a in inv['anchors']])
        fields=docx_inspect(q)[3];validate_candidate_fields(inv,fields)
        self.assertEqual([f['label'] for f in fields],['Temperature','Fever'])
        self.assertEqual(len(result['changes']),2);self.assertEqual(p.read_bytes(),before)
        with ZipFile(p) as old,ZipFile(q) as new:
            for name in old.namelist():
                if name!='word/document.xml':self.assertEqual(old.read(name),new.read(name))
            text=new.read('word/document.xml').decode()
            self.assertIn('Patient signature: __________',text);self.assertIn('°C',text);self.assertIn('Yes / No',text)

    def test_clinical_sign_words_keep_answer_fields_and_true_signatures_stay_blank(self):
        clinical=['Vital signs','Signs and symptoms','Significant findings','Design assessment','Assignment']
        protected=['Patient signature','Clinician signature','Signed by','Consent','Attestation','Certification',
                   'Sign here','Please sign below','Sign and date','Clinician sign']
        p=self.make([label+': ______' for label in clinical+protected]);inv=inventory(p)
        self.assertFalse(inv['issues'])
        labels={q['id']:q['text'] for q in inv['questions']}
        self.assertEqual([labels[a['question_id']] for a in inv['anchors'] if a['kind']=='literal_blank'],clinical)
        self.assertEqual(len([a for a in inv['anchors'] if a['kind']=='protected_blank']),len(protected))
        q=self.root/'candidate.docx';candidate(p,q,inv,[a['id'] for a in inv['anchors']])
        fields=docx_inspect(q)[3];validate_candidate_fields(inv,fields)
        self.assertEqual([f['label'] for f in fields],clinical)
        text='\n'.join(paragraph.text for paragraph in Document(q).paragraphs)
        for label in protected:self.assertIn(label+': ______',text)

    def test_unrecognized_singular_sign_requires_assisted_clarification(self):
        for label in ['Sign','Physician to sign','Consultant should sign','A sign of concern']:
            for table in [False,True]:
                p=self.make(['Assessment: ____',label+': ____'] if not table else ['Assessment: ____'],
                            [[label,'']] if table else None)
                with self.subTest(table=table,label=label):
                    self.assertIn('guided_signing',{i['code'] for i in inventory(p)['issues']})

    def test_ambiguous_cells_require_explicit_finite_doctor_placement(self):
        p=self.make(['Clinician signature: ______'],[['Temperature','','°C'],['Assessment','','']])
        inv=inventory(p);self.assertFalse(inv['issues'])
        choose=[a for a in inv['anchors'] if a['kind']=='choose_cell'];self.assertEqual(len(choose),2)
        with self.assertRaises(DocumentFailure):candidate(p,self.root/'no.docx',inv,[a['id'] for a in inv['anchors']])
        placement={a['id']:a['choices'][-1]['id'] for a in choose}
        q=self.root/'yes.docx';proof=candidate(p,q,inv,[a['id'] for a in inv['anchors']],placement)
        self.assertEqual(proof['reviewed_multi_column_rows'],[{'part':'word/document.xml','row':1}])
        fields=docx_inspect(q)[3];validate_candidate_fields(inv,fields)
        self.assertEqual({f['label'] for f in fields},{'Temperature','Assessment'})
        self.assertEqual(Document(q).tables[0].cell(1,1).text,'')
        self.assertEqual(Document(q).tables[0].cell(1,2).text,'{{Assessment}}')

    def test_unknown_or_signing_cell_placement_refuses(self):
        p=self.make(['Patient signature: ____'],[['Assessment','','']]);inv=inventory(p)
        choose=next(a for a in inv['anchors'] if a['kind']=='choose_cell')
        for bad in ['c0','c99','../../file','w:p']:
            with self.subTest(bad=bad),self.assertRaises(DocumentFailure):
                candidate(p,self.root/'out.docx',inv,[a['id'] for a in inv['anchors']],{choose['id']:bad})

    def test_uncovered_wide_rows_cannot_hide_behind_another_manual_placement(self):
        for row in [['','Plan',''],['x'*121,'','°C'],['x'*121,'',''],['Patient signature','','']]:
            p=self.make(['Name: {{Name}}'],[['Assessment','',''],row]);inv=inventory(p)
            with self.subTest(row=row):
                self.assertTrue(inv['issues'])
                choices={a['id']:a['choices'][0]['id'] for a in inv['anchors'] if a['kind']=='choose_cell'}
                with self.assertRaises(DocumentFailure):
                    candidate(p,self.root/'bad.docx',inv,[a['id'] for a in inv['anchors']],choices)

    def test_neighbouring_other_question_is_unresolved(self):
        p=self.make(['Name: ______'],[['Assessment','','Plan']]);inv=inventory(p)
        self.assertTrue(inv['issues'])

    def test_split_blank_and_merged_layout_are_never_green(self):
        d=Document();p=d.add_paragraph('Name: ');p.add_run('__');p.add_run('__')
        t=d.add_table(rows=1,cols=3);t.cell(0,0).text='Assessment';t.cell(0,1).merge(t.cell(0,2))
        original=self.root/'original.docx';d.save(original);inv=inventory(original)
        self.assertTrue(inv['issues'])

    def test_contiguous_split_blank_with_long_second_run_remains_unresolved(self):
        d=Document();p=d.add_paragraph('Patient name: ');p.add_run('__');p.add_run('____')
        original=self.root/'original.docx';d.save(original);inv=inventory(original)
        self.assertTrue(inv['issues']);self.assertFalse(any(a['kind']=='literal_blank' for a in inv['anchors']))

    def test_pre_and_post_root_processing_instructions_survive_candidate(self):
        p=self.make(['Assessment: ____']);modified=self.root/'with-pi.docx'
        with ZipFile(p) as old,ZipFile(modified,'w') as new:
            for entry in old.infolist():
                raw=old.read(entry)
                if entry.filename=='word/document.xml':
                    tree=etree.fromstring(raw);tree.addprevious(etree.ProcessingInstruction('quill-preserve','retained-layout-note'))
                    tree.addnext(etree.ProcessingInstruction('quill-tail','retained-tail-note'))
                    raw=etree.tostring(tree.getroottree(),xml_declaration=True,encoding='UTF-8')
                new.writestr(entry,raw)
        inv=inventory(modified);q=self.root/'candidate.docx';candidate(modified,q,inv,[a['id'] for a in inv['anchors']])
        with ZipFile(q) as package:
            raw=package.read('word/document.xml')
            self.assertIn(b'<?quill-preserve retained-layout-note?>',raw)
            self.assertIn(b'<?quill-tail retained-tail-note?>',raw)

    def test_missing_duplicate_swapped_extra_and_code_results_refuse(self):
        p=self.make(['Name: ____','Temperature: ____ °C']);inv=inventory(p)
        for role in ['proposer','checker']:
            base=self.role(inv,role);validate_role(inv,base,role)
            bad=[]
            a=deepcopy(base);a['anchors'].pop();bad.append(a)
            a=deepcopy(base);a['anchors'][1]=a['anchors'][0];bad.append(a)
            a=deepcopy(base);a['anchors'][0]['question_id']='q2';bad.append(a)
            a=deepcopy(base);a['anchors'][0]['xpath']='//w:p';bad.append(a)
            a=deepcopy(base);a['inventory_sha256']='0'*64;bad.append(a)
            a=deepcopy(base);a['stories'].pop();bad.append(a)
            a=deepcopy(base);a['stories'][0]['content_sha256']='0'*64;bad.append(a)
            for value in bad:
                with self.subTest(role=role,value=value),self.assertRaises(DocumentFailure):validate_role(inv,value,role)

    def test_model_never_receives_parts_paths_xml_or_locators(self):
        inv=inventory(self.make(['Assessment: ____']))
        data=model_inventory(inv)
        self.assertNotIn('parts',data)
        for anchor in data['anchors']:self.assertNotIn('locator',anchor);self.assertNotIn('field_key',anchor)
        for question in data['questions']:self.assertNotIn('story',question)
        inv['issues']=[{'code':'unresolved_blank','message':'Check this blank','story':'word/document.xml'}]
        self.assertNotIn('story',model_inventory(inv)['issues'][0])

    def test_changed_original_or_self_rehashed_locator_cannot_create_candidate(self):
        p=self.make(['Assessment: ____']);inv=inventory(p)
        value=deepcopy(inv);value['anchors'][0]['locator']['paragraph']=999
        value['inventory_sha256']=digest({k:v for k,v in value.items() if k!='inventory_sha256'})
        with self.assertRaises(DocumentFailure):candidate(p,self.root/'bad.docx',value,['a1'])
        d=Document(p);d.add_paragraph('New question: ____');d.save(p)
        with self.assertRaises(DocumentFailure):candidate(p,self.root/'changed.docx',inv,['a1'])

    def test_existing_mapping_candidate_is_byte_identical(self):
        p=self.make(table=[['Assessment',''],['Electronically signed by','']]);inv=inventory(p)
        q=self.root/'candidate.docx';result=candidate(p,q,inv,[a['id'] for a in inv['anchors']])
        self.assertEqual(q.read_bytes(),p.read_bytes());self.assertFalse(result['changes'])

    def test_forged_advanced_mapping_cannot_bypass_original_pin(self):
        p=self.make(['Name: ____']);sha=hashlib.sha256(p.read_bytes()).hexdigest()
        with self.assertRaises(DocumentFailure):inventory(p,{'sha256':sha,'backend':'wm-iov-foss-v2','fields':[]})

    def test_networkless_worker_rejects_forged_wm_mapping_by_original_pin(self):
        from .preparation_service import inspect_preparation
        from .service import DocumentError
        p=self.make(['Name: ____']);sha=hashlib.sha256(p.read_bytes()).hexdigest()
        with self.assertRaises(DocumentError) as refused:
            inspect_preparation(p,{'sha256':sha,'backend':'wm-iov-foss-v2','fields':[]})
        self.assertEqual(refused.exception.code,'stale_mapping')

    def test_active_fields_external_links_and_altchunks_still_refuse(self):
        p=self.make(['Name: ____'])
        for tag in ['fldSimple','altChunk','instrText']:
            bad=self.root/(tag+'.docx')
            with ZipFile(p) as old,ZipFile(bad,'w') as out:
                for entry in old.infolist():
                    raw=old.read(entry)
                    if entry.filename=='word/document.xml':
                        tree=etree.fromstring(raw);etree.SubElement(tree.find('{'+W_NS+'}body'),'{'+W_NS+'}'+tag)
                        raw=etree.tostring(tree)
                    out.writestr(entry,raw)
            with self.subTest(tag=tag),self.assertRaises(DocumentFailure):inventory(bad)

    def test_candidate_field_coverage_refuses_lost_question(self):
        inv=inventory(self.make(['Name: ____','Assessment: ____']))
        with self.assertRaises(DocumentFailure):validate_candidate_fields(inv,[{'label':'Name'}])
