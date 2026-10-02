"""Descriptive free-text questions retain narrative typing beside date/time words."""
import io
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import word_v3 as w


def inspect(document):
    stream = io.BytesIO()
    document.save(stream)
    return w.inspect(stream.getvalue())['fields']


class CompositeKinds(unittest.TestCase):
    def test_descriptive_free_text_and_signing_boundaries(self):
        cases = [('Date of Birth', 'date'), ('Date of assessment', 'date'),
                 ('Time of assessment', 'time'),
                 ('Date of Order and details of restrictions', 'long_text'),
                 ('Time of onset and description of symptoms', 'long_text'),
                 ('Date and summary of previous assessment', 'long_text'),
                 ('Signature comments', 'signature'), ('Initials and comments', 'initials'),
                 ('Patient name', 'text')]
        for label, expected in cases:
            with self.subTest(label=label):
                document = Document()
                table = document.add_table(rows=1, cols=2)
                table.cell(0, 0).text = label
                fields = inspect(document)
                self.assertEqual(len(fields), 1)
                self.assertEqual(fields[0]['kind'], expected)
                self.assertEqual(fields[0]['label'], label)

    def test_native_date_remains_structurally_typed(self):
        document = Document()
        control = OxmlElement('w:sdt')
        properties = OxmlElement('w:sdtPr')
        control.append(properties)
        alias = OxmlElement('w:alias')
        alias.set(qn('w:val'), 'Date and assessment details')
        properties.append(alias)
        properties.append(OxmlElement('w:date'))
        content = OxmlElement('w:sdtContent')
        control.append(content)
        run = OxmlElement('w:r')
        content.append(run)
        text = OxmlElement('w:t')
        text.text = 'Click here to enter a date.'
        run.append(text)
        document.add_paragraph()._p.append(control)
        fields = inspect(document)
        self.assertEqual(len(fields), 1)
        self.assertEqual(fields[0]['kind'], 'date')
