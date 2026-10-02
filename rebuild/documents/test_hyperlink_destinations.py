"""Fixed hyperlink text is not a replaceable answer hint."""
import io
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

from . import word_v3 as w
from .test_word_v3 import _sdt


def link(paragraph, text):
    element = OxmlElement('w:hyperlink')
    element.set(qn('r:id'), paragraph.part.relate_to('mailto:forms@example.invalid',
                RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1')
    props.append(color); run.append(props)
    value = OxmlElement('w:t'); value.text = text; run.append(value)
    element.append(run); paragraph._p.append(element)


def data(document):
    out = io.BytesIO(); document.save(out); return out.getvalue()


class FixedReferenceCoverage(unittest.TestCase):
    def test_literal_mailto_reference_is_preserved_without_inventing_an_answer(self):
        doc=Document();p=doc.add_paragraph('Submit the completed report by email to: ')
        link(p,'forms@example.invalid');doc.add_paragraph('Reason for visit: ______')
        result=w.inspect(data(doc))
        self.assertEqual([f['label'] for f in result['fields']],['Reason for visit'])
        self.assertEqual(result['coverage_issues'],[])
        self.assertTrue(any('hyperlink' in n.lower() for n in result['notes']))
        # The target alone cannot make an actual linked placeholder a reference.
        doc=Document();p=doc.add_paragraph('Answer: ');link(p,'Enter answer')
        doc.add_paragraph('Reason for visit: ______');held=w.inspect(data(doc))
        self.assertTrue(any(i['code']=='fixed_link_answer' for i in held['coverage_issues']))


def linked_control(document, kind, block=False):
    paragraph = document.add_paragraph('Answer: ')
    control = _sdt(kind, options=('Yes', 'No'), placeholder='Enter answer')
    properties = control.find(qn('w:sdtPr'))
    if kind == 'date':
        properties.append(OxmlElement('w:date'))
    properties.append(OxmlElement('w:showingPlcHdr'))
    content = control.find(qn('w:sdtContent'))
    content.clear()
    link(paragraph, 'Fixed clinic instructions')
    hyperlink = paragraph._p[-1]
    if block:
        paragraph.text = 'Instructions'
        inner = OxmlElement('w:p')
        inner.append(hyperlink)
        content.append(inner)
        document.element.body.insert(len(document.element.body) - 1, control)
    else:
        content.append(hyperlink)
        paragraph._p.append(control)
    return control


class HyperlinkDestinations(unittest.TestCase):
    def test_colored_contact_is_never_an_answer_hint_in_common_parts(self):
        for position in ('body', 'table', 'header', 'footer'):
            with self.subTest(position=position):
                d = Document()
                if position == 'body': p = d.add_paragraph('Send the completed form to: ')
                elif position == 'table':
                    p = d.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
                    p.add_run('Send the completed form to: ')
                else:
                    p = getattr(d.sections[0], position).paragraphs[0]
                    p.add_run('Send the completed form to: ')
                link(p, 'forms@example.invalid')
                self.assertEqual(w.inspect(data(d))['fields'], [])

    def test_explicit_answer_marker_inside_link_is_visible_and_unassigned(self):
        for text in ('________', 'Enter name', '{{name}}'):
            with self.subTest(text=text):
                d = Document(); p = d.add_paragraph('Name: '); link(p, text)
                info = w.inspect(data(d))
                self.assertEqual(info['fields'], [])
                self.assertTrue(any('hyperlink' in message.lower() for message in info['notes']))

    def test_link_only_value_cell_is_not_replacement_guidance(self):
        for suffix in ('', ' . '):
            with self.subTest(suffix=suffix):
                d = Document(); table = d.add_table(rows=1, cols=2)
                table.cell(0, 0).text = 'Send the completed form to'
                paragraph = table.cell(0, 1).paragraphs[0]
                link(paragraph, 'forms@example.invalid')
                from docx.shared import RGBColor
                paragraph.add_run(suffix).font.color.rgb = RGBColor.from_string('0563C1')
                self.assertEqual(w.inspect(data(d))['fields'], [])

    def test_old_whole_cell_mapping_cannot_remove_later_link(self):
        d = Document(); cell = d.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = 'Enter report'; link(cell.add_paragraph(), 'Clinic instructions')
        with self.assertRaisesRegex(w.V3Failure, 'hyperlink'):
            w.write_cell(cell._tc, 'Synthetic report', replace=True, whole=True)

    def test_normal_blank_after_link_keeps_text_and_relationship(self):
        d = Document(); p = d.add_paragraph(); link(p, 'Clinic instructions')
        p.add_run(' Patient name: ________')
        raw = data(d); info = w.inspect(raw)
        self.assertEqual(len(info['fields']), 1)
        output = w.fill(raw, info['fields'], {info['fields'][0]['key']: 'Synthetic Example'})
        before = w.open_package(raw); after = w.open_package(output)
        self.assertEqual(before.zip.read('word/_rels/document.xml.rels'), after.zip.read('word/_rels/document.xml.rels'))
        for package in (before, after):
            self.assertEqual([n.text for h in package.tree('word/document.xml').iter(w.q('w:hyperlink'))
                              for n in h.iter(w.T)], ['Clinic instructions'])
        self.assertIn('Synthetic Example', w.Detector(after).text(after.tree('word/document.xml'), own=False))

    def test_old_mapping_cannot_replace_fixed_hyperlink(self):
        d = Document(); p = d.add_paragraph('Send the completed form to: ')
        link(p, 'forms@example.invalid'); raw = data(d)
        index = w.SlotIndex(w.open_package(raw)); slot = index.by_id['word/document.xml#b/p[0]']
        text, _ = w.paragraph_text(slot.el); start = text.index('forms@')
        field = {'key': 'old', 'kind': 'text', 'label': 'Send the completed form to',
                 'locator': {'contract': 'v3', 'part': slot.part, 'slot': slot.path,
                             'fingerprint': w.fingerprint(slot.el),
                             'target': {'type': 'span', 'start': start, 'end': len(text), 'shape': 'hint'}}}
        with self.assertRaisesRegex(w.V3Failure, 'hyperlink'):
            w.fill(raw, [field], {'old': 'Synthetic incorrect destination'})

    def test_insertion_in_or_at_end_of_hyperlink_refuses(self):
        for offset in (3, len('Clinic instructions')):
            with self.subTest(offset=offset):
                d = Document(); p = d.add_paragraph(); link(p, 'Clinic instructions')
                with self.assertRaisesRegex(w.V3Failure, 'hyperlink'):
                    w.replace_range(p._p, offset, offset, 'Synthetic answer')

    def test_unlinked_colored_explicit_hint_remains_fillable(self):
        d = Document(); p = d.add_paragraph('Name: '); run = p.add_run('Enter name')
        from docx.shared import RGBColor
        run.font.color.rgb = RGBColor.from_string('0563C1')
        raw = data(d); info = w.inspect(raw)
        self.assertEqual(len(info['fields']), 1)
        w.fill(raw, info['fields'], {info['fields'][0]['key']: 'Synthetic Example'})

    def test_unlinked_blue_guidance_cell_remains_fillable(self):
        from docx.shared import RGBColor
        d = Document(); table = d.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Clinical report'
        run = table.cell(0, 1).paragraphs[0].add_run('Enter the clinical report here')
        run.font.color.rgb = RGBColor.from_string('0563C1')
        raw = data(d); info = w.inspect(raw)
        self.assertEqual(len(info['fields']), 1)
        self.assertEqual(info['fields'][0]['locator']['target']['type'], 'cell_replace')
        w.fill(raw, info['fields'], {info['fields'][0]['key']: 'Synthetic report.'})

    def test_link_bearing_content_controls_are_visible_but_not_whole_answers(self):
        for kind in ('rich', 'dropdown', 'date'):
            for block in (False, True):
                with self.subTest(kind=kind, block=block):
                    d = Document(); linked_control(d, kind, block)
                    info = w.inspect(data(d))
                    self.assertEqual(info['fields'], [])
                    self.assertTrue(any('hyperlink' in note.lower() and 'word/document.xml#' in note
                                        for note in info['notes']))

    def test_historical_control_mapping_cannot_remove_link(self):
        for kind in ('rich', 'dropdown', 'date'):
            for block in (False, True):
                with self.subTest(kind=kind, block=block):
                    d = Document(); linked_control(d, kind, block)
                    raw = data(d); index = w.SlotIndex(w.open_package(raw))
                    slot = next(s for s in index.slots if s.kind == 'sdt')
                    field = {'key': 'old', 'kind': 'choice' if kind == 'dropdown' else 'text',
                             'label': 'Answer', 'options': ['Yes', 'No'],
                             'locator': {'contract': 'v3', 'part': slot.part, 'slot': slot.path,
                                         'fingerprint': w.fingerprint(slot.el),
                                         'target': {'type': 'sdt_choice' if kind == 'dropdown' else 'sdt_text'}}}
                    with self.assertRaises(w.V3Failure) as rejected:
                        w.fill(raw, [field], {'old': 'Yes'})
                    self.assertEqual(rejected.exception.code, 'protected_hyperlink')
                    self.assertEqual(w.fill(raw, [field], {}), raw)

    def test_content_control_refusal_precedes_any_source_mutation(self):
        from lxml import etree
        d = Document(); control = linked_control(d, 'rich', block=True)
        before = etree.tostring(control)
        with self.assertRaises(w.V3Failure) as rejected:
            w.write_sdt_text(control, 'Synthetic answer')
        self.assertEqual(rejected.exception.code, 'protected_hyperlink')
        self.assertEqual(etree.tostring(control), before)

    def test_linked_checkbox_is_unassigned_and_historical_write_refuses(self):
        for position in ('body', 'table'):
            with self.subTest(position=position):
                d = Document()
                paragraph = (d.add_paragraph() if position == 'body' else
                             d.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0])
                link(paragraph, '☐ Send copies')
                info = w.inspect(data(d))
                self.assertEqual(info['fields'], [])
                self.assertTrue(any('hyperlink' in note.lower() for note in info['notes']))
                for writer in (w.set_glyph, lambda p, n, value: w.pin_glyph(p, n)):
                    with self.assertRaises(w.V3Failure) as rejected:
                        writer(paragraph._p, 0, True)
                    self.assertEqual(rejected.exception.code, 'protected_hyperlink')

    def test_checkbox_control_cannot_clear_or_change_linked_content(self):
        from lxml import etree
        for mode in ('fixed-contact', 'linked-glyph'):
            for position in ('body', 'table'):
                with self.subTest(mode=mode, position=position):
                    d = Document()
                    paragraph = (d.add_paragraph() if position == 'body' else
                                 d.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0])
                    control = _sdt('checkbox')
                    content = control.find(qn('w:sdtContent'))
                    if mode == 'linked-glyph':
                        content.clear()
                    link(paragraph, '☐' if mode == 'linked-glyph' else 'forms@example.invalid')
                    content.append(paragraph._p[-1])
                    paragraph._p.append(control); paragraph.add_run(' Send copies')
                    info = w.inspect(data(d))
                    self.assertEqual(info['fields'], [])
                    self.assertTrue(any('hyperlink' in note.lower() for note in info['notes']))
                    before = etree.tostring(control)
                    with self.assertRaises(w.V3Failure) as rejected:
                        w.set_sdt_checkbox(control, True)
                    self.assertEqual(rejected.exception.code, 'protected_hyperlink')
                    self.assertEqual(etree.tostring(control), before)

    def test_checkbox_beside_linked_option_text_keeps_link_and_remains_fillable(self):
        for control in (False, True):
            with self.subTest(control=control):
                d = Document(); paragraph = d.add_paragraph()
                if control:
                    paragraph._p.append(_sdt('checkbox'))
                else:
                    paragraph.add_run('☐ ')
                link(paragraph, 'Send copies')
                raw = data(d); info = w.inspect(raw)
                self.assertEqual(len(info['fields']), 1)
                output = w.fill(raw, info['fields'], {info['fields'][0]['key']: True})
                before, after = w.open_package(raw), w.open_package(output)
                self.assertEqual([w.etree.tostring(h) for h in before.tree('word/document.xml').iter(w.q('w:hyperlink'))],
                                 [w.etree.tostring(h) for h in after.tree('word/document.xml').iter(w.q('w:hyperlink'))])

    def test_separate_blank_in_link_bearing_block_control_still_fills(self):
        d = Document(); control = linked_control(d, 'rich', block=True)
        paragraph = d.add_paragraph('Patient name: ________')
        control.find(qn('w:sdtContent')).append(paragraph._p)
        raw = data(d); info = w.inspect(raw)
        self.assertEqual(len(info['fields']), 1)
        self.assertEqual(info['fields'][0]['locator']['target']['type'], 'span')
        output = w.fill(raw, info['fields'], {info['fields'][0]['key']: 'Synthetic Example'})
        before, after = w.open_package(raw), w.open_package(output)
        self.assertEqual(before.zip.read('word/_rels/document.xml.rels'),
                         after.zip.read('word/_rels/document.xml.rels'))
        self.assertEqual([n.text for h in after.tree('word/document.xml').iter(w.q('w:hyperlink'))
                          for n in h.iter(w.T)], ['Fixed clinic instructions'])
        self.assertIn('Synthetic Example', w.Detector(after).text(after.tree('word/document.xml'), own=False))
