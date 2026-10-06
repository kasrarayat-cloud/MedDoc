"""Physical mixed-line destinations and ordinary clinician recording controls."""
import io
from copy import deepcopy
from zipfile import ZipFile
import unittest
from docx import Document
from docx.oxml import OxmlElement
from . import word_v3 as w


def saved(doc):
    out=io.BytesIO();doc.save(out);return out.getvalue()


class OperatorRecording(unittest.TestCase):
    def form(self, caption):
        doc=Document();table=doc.add_table(rows=2,cols=2)
        table.cell(0,0).text=caption;table.cell(1,0).text='Clinical reason'
        table.cell(1,1).text='Original neighbor entry'
        doc.sections[0].header.paragraphs[0].text='Original header'
        doc.sections[0].footer.paragraphs[0].text='Original footer'
        return saved(doc)

    def test_complete_current_operator_recording_ordinary_argument_preserves_original(self):
        for argument in ('initial symptoms','initial physical observations','signed document receipt',
                         'signed referral received and recorded'):
            with self.subTest(argument=argument):
                caption='This section is for the clinician to record '+argument
                raw=self.form(caption);info=w.inspect(raw);field=next(f for f in info['fields']if f['locator']['slot']=='b/tbl[0]/tr[0]/tc[1]')
                self.assertFalse(field.get('write_protected'))
                self.assertEqual(w.destination_roles(raw,[field])['violations'],{})
                out=w.fill(raw,info['fields'],{field['key']:'Synthetic recorded observation'})
                doc=Document(io.BytesIO(out));self.assertEqual(doc.tables[0].cell(0,0).text,caption)
                self.assertEqual(doc.tables[0].cell(0,1).text,'Synthetic recorded observation')
                self.assertEqual(doc.tables[0].cell(1,1).text,'Original neighbor entry')
                with ZipFile(io.BytesIO(raw))as before,ZipFile(io.BytesIO(out))as after:
                    self.assertEqual(before.comment,after.comment)
                    for name in before.namelist():
                        if name!='word/document.xml':self.assertEqual(before.read(name),after.read(name),name)

    def test_named_author_and_complete_operator_actions_never_gain_permission(self):
        for caption in ('This section is for Alex Reed to record initial symptoms',
                        'This section is for Dr. Synthetic to record symptoms',
                        'This section is for the clinician to write his comments',
                        'This section is for the clinician to record initial symptoms and I approve this document',
                        'This section is for the clinician to record initial patient signature',
                        'This section is reserved for the clinician to record initial symptoms',
                        'This section is reserved for the line manager to supply comments',
                        'Date and time Alex Reed is to write his comments:'):
            with self.subTest(caption=caption):
                raw=self.form(caption);info=w.inspect(raw);field=next(f for f in info['fields']if f['locator']['slot']=='b/tbl[0]/tr[0]/tc[1]')
                self.assertTrue(field.get('write_protected'))
                with self.assertRaises(w.V3Failure):w.fill(raw,info['fields'],{field['key']:'Dr Synthetic Finalizer'})
                forged=deepcopy(field);forged.update(label='Clinical value',kind='text')
                for key in ('write_protected','destination_roles','answer_policy'):forged.pop(key,None)
                with self.assertRaises(w.V3Failure):w.fill(raw,[forged],{forged['key']:'Dr Synthetic Finalizer'})
                self.assertEqual(w.fill(raw,info['fields'],{}),raw)


class PairedColonGaps(unittest.TestCase):
    def form(self, text, *, body=False):
        doc=Document();doc.sections[0].header.paragraphs[0].text='Original header'
        doc.sections[0].footer.paragraphs[0].text='Original footer'
        if body:p=doc.add_paragraph(text)
        else:
            # The retained mixed line occupies a single-cell table, not a
            # caption cell paired with a separate answer column.
            table=doc.add_table(rows=1,cols=1);table.cell(0,0).text=text
            p=table.cell(0,0).paragraphs[0]
        doc.add_paragraph('Unanswered neighbor: ______')
        doc.add_paragraph('Existing selection: ☒ Recorded ☐ Missing')
        return doc,p

    def own(self, info, body=False):
        slot='b/p[0]'if body else'b/tbl[0]/tr[0]/tc[0]/p[0]'
        return [f for f in info['fields']if f['locator']['slot']==slot]

    def test_two_complete_captions_have_distinct_physical_spans_and_complete_answers(self):
        for body in (False,True):
            for text in ('5. Number of completed sessions:        6. Average duration of sessions:    minutes',
                         'First reading:        Second reading:    units'):
                with self.subTest(body=body,text=text):
                    doc,p=self.form(text,body=body);raw=saved(doc);info=w.inspect(raw);fields=self.own(info,body)
                    self.assertEqual(len(fields),2)
                    targets=[f['locator']['target']for f in fields]
                    self.assertLessEqual(targets[0]['end'],targets[1]['start'])
                    self.assertTrue(all(text[t['start']:t['end']].isspace()for t in targets))
                    self.assertEqual(w.fill(raw,info['fields'],{}),raw)
                    answers={fields[0]['key']:'12',fields[1]['key']:'45'}
                    out=w.fill(raw,info['fields'],answers);before=w.SlotIndex(w.open_package(raw));after=w.SlotIndex(w.open_package(out))
                    for f,value in zip(fields,('12','45')):
                        self.assertEqual(w.constrained_scalar_value(before,after,f,True),value)
                    emitted=w.paragraph_text(after.by_id[fields[0]['locator']['part']+'#'+fields[0]['locator']['slot']].el)[0]
                    self.assertEqual(emitted.replace(' 12','',1).replace(' 45','',1),text)
                    for slot in before.slots:
                        if slot.kind=='p'and slot.id!=fields[0]['locator']['part']+'#'+fields[0]['locator']['slot']:
                            self.assertEqual(w.etree.tostring(slot.el),w.etree.tostring(after.by_id[slot.id].el))
                    with ZipFile(io.BytesIO(raw))as a,ZipFile(io.BytesIO(out))as b:
                        for name in a.namelist():
                            if name!='word/document.xml':self.assertEqual(a.read(name),b.read(name),name)

    def test_printed_question_ink_inside_an_existing_blank_is_not_a_new_caption(self):
        doc,p=self.form('First reading:    ');p.add_run('[Second reading:    units]')
        raw=saved(doc);fields=self.own(w.inspect(raw))
        self.assertFalse([f for f in fields if f['locator']['target'].get('shape')=='after_label'
                          and f['label']=='Second reading'])
        self.assertEqual(w.fill(raw,w.inspect(raw)['fields'],{}),raw)

    def test_choices_sentences_and_ambiguous_unit_prefix_are_not_promoted_by_spaces(self):
        for text in ('Symptoms:    ☐ Cough ☐ Fever','First reading:    normal result. Other observations continue.',
                     'First reading:    kg Second reading:    units'):
            with self.subTest(text=text):
                doc,p=self.form(text);raw=saved(doc);info=w.inspect(raw);fields=self.own(info)
                if text.startswith('Symptoms'):
                    self.assertFalse([f for f in fields if f['locator']['target'].get('type')=='span'])
                elif 'kg Second' in text:self.assertEqual(len(fields),1)
                else:self.assertEqual(fields,[])
                self.assertEqual(w.fill(raw,info['fields'],{}),raw)

    def test_repeated_captions_keep_positions_and_do_not_copy_one_answer(self):
        doc,p=self.form('Reading:        Reading:    units');raw=saved(doc);info=w.inspect(raw);fields=self.own(info)
        self.assertEqual(len(fields),2);self.assertEqual([f['label']for f in fields],['Reading','Reading'])
        self.assertNotEqual(fields[0]['context'],fields[1]['context'])
        self.assertTrue(all('occurrence' in f['context']for f in fields))
        out=w.fill(raw,info['fields'],{fields[0]['key']:'12'})
        text=Document(io.BytesIO(out)).tables[0].cell(0,0).text
        self.assertEqual(text,'Reading: 12        Reading:    units')

    def test_actual_human_caption_stays_protected_beside_plain_gaps(self):
        for caption in ('Alex Reed should sign','Manager is required to write this response','I approve this document'):
            with self.subTest(caption=caption):
                doc,p=self.form('First reading:        '+caption+':    minutes')
                raw=saved(doc);info=w.inspect(raw);fields=self.own(info)
                actual=next(f for f in fields if f['label']==caption)
                self.assertTrue(actual.get('write_protected'))
                with self.assertRaises(w.V3Failure):w.fill(raw,info['fields'],{actual['key']:'Dr Synthetic Finalizer'})
                forged=deepcopy(actual);forged.update(label='Ordinary value',kind='text')
                for key in ('write_protected','destination_roles','answer_policy'):forged.pop(key,None)
                with self.assertRaises(w.V3Failure):w.fill(raw,[forged],{forged['key']:'Dr Synthetic Finalizer'})
                self.assertEqual(w.fill(raw,info['fields'],{}),raw)

    def test_duplicate_or_overlapping_writer_refuses_before_any_answer_is_lost(self):
        doc,p=self.form('First reading:        Second reading:    units');raw=saved(doc)
        info=w.inspect(raw);fields=self.own(info)
        for mode in ('duplicate-answered','duplicate-blank','overlap','whole-paragraph','renamed-caption','changed-kind'):
            with self.subTest(mode=mode):
                extra=deepcopy(fields[0]);extra['key']='additional_writer'
                if mode=='overlap':extra['locator']['target']['end']+=5
                elif mode=='whole-paragraph':extra['locator']['target']={'type':'paragraph'}
                elif mode=='renamed-caption':extra['label']='New caption'
                elif mode=='changed-kind':extra['kind']='date'
                mapping=info['fields']+[extra]
                answers={fields[0]['key']:'12',fields[1]['key']:'45'}
                if mode!='duplicate-blank':answers[extra['key']]='99'
                with self.assertRaises(w.V3Failure)as error:w.fill(raw,mapping,answers)
                self.assertEqual(error.exception.code,'mapping_mismatch')
                self.assertEqual(w.fill(raw,mapping,{}),raw)

    def test_composed_readback_retains_all_fixed_ink_and_refuses_ambiguity_and_hidden_values(self):
        doc,p=self.form('First reading:        Second reading:    units');raw=saved(doc)
        info=w.inspect(raw);fields=self.own(info);answers={fields[0]['key']:'12',fields[1]['key']:'45'}
        output=w.fill(raw,info['fields'],answers);before=w.SlotIndex(w.open_package(raw))
        expected=w.effective_values(raw,info['fields'],answers,_written=output)
        self.assertEqual([expected[f['key']]for f in fields],['12','45'])
        for text in ('First reading: 12        Third reading: 45    units',
                     'First reading: 12        Second reading: 45    kilograms',
                     'First reading: 12        Second reading: 45',
                     'Second reading: 45        First reading: 12    units',
                     'First reading: 12        Second reading: not-an-answer        Second reading: 45    units'):
            with self.subTest(text=text):
                changed=Document(io.BytesIO(output));changed.tables[0].cell(0,0).text=text
                after=w.SlotIndex(w.open_package(saved(changed)))
                self.assertTrue(all(w.constrained_scalar_value(before,after,f,True)is None for f in fields))
        for mode in ('run','paragraph','style'):
            with self.subTest(hidden=mode):
                changed=Document(io.BytesIO(output));paragraph=changed.tables[0].cell(0,0).paragraphs[0]
                if mode=='run':paragraph.runs[1].font.hidden=True
                elif mode=='style':paragraph.style.font.hidden=True
                else:
                    pr=paragraph._p.get_or_add_pPr();rpr=OxmlElement('w:rPr');rpr.append(OxmlElement('w:vanish'));pr.append(rpr)
                after=w.SlotIndex(w.open_package(saved(changed)))
                self.assertTrue(all(w.constrained_scalar_value(before,after,f,True)is None for f in fields))
        for mode in ('fingerprint','target','label','kind','context','source-relationship'):
            with self.subTest(forgery=mode):
                forged=deepcopy(fields[0])
                if mode=='fingerprint':forged['locator']['fingerprint']='0'*64
                elif mode=='target':forged['locator']['target']['end']+=1
                elif mode=='source-relationship':forged['source_questions']=[{'relation':'unproved'}]
                else:forged[mode]='Changed'
                after=w.SlotIndex(w.open_package(output))
                self.assertIsNone(w.constrained_scalar_value(before,after,forged,True))

    def test_multiline_values_read_back_and_neighbor_source_cannot_become_an_answer(self):
        doc,p=self.form('First reading:        Second reading:    units');raw=saved(doc)
        info=w.inspect(raw);fields=self.own(info);answers={fields[0]['key']:'first\ncomplete\tvalue',fields[1]['key']:'second\nfull value'}
        output=w.fill(raw,info['fields'],answers)
        actual=w.effective_values(raw,info['fields'],answers,_written=output)
        self.assertEqual([actual[f['key']]for f in fields],list(answers.values()))
        # Original-source identities remain separate in the same process.
        other_doc,p=self.form('Different caption:        Second reading:    units');other=saved(other_doc)
        with self.assertRaises(w.V3Failure):w.fill(other,info['fields'],answers)
        original=w.SlotIndex(w.open_package(other));final=w.SlotIndex(w.open_package(output))
        self.assertIsNone(w.constrained_scalar_value(original,final,fields[0],True))

    def test_existing_single_answer_mapping_remains_valid_and_does_not_copy_the_new_answer(self):
        doc,p=self.form('First reading:        Second reading:    units');raw=saved(doc)
        fields=self.own(w.inspect(raw));old=[deepcopy(fields[0])]
        out=w.fill(raw,old,{old[0]['key']:'12'})
        before=w.SlotIndex(w.open_package(raw));after=w.SlotIndex(w.open_package(out))
        self.assertEqual(w.constrained_scalar_value(before,after,old[0],True),'12')
        self.assertEqual(w.constrained_scalar_value(before,after,fields[1],False),'')
        self.assertEqual(Document(io.BytesIO(out)).tables[0].cell(0,0).text,
                         'First reading: 12        Second reading:    units')

    def test_two_cell_caption_answer_layout_is_not_claimed_as_the_single_cell_pair(self):
        doc=Document();table=doc.add_table(rows=1,cols=2)
        table.cell(0,0).text='First reading:        Second reading:    units'
        table.cell(0,1).text='Original adjacent entry';raw=saved(doc);info=w.inspect(raw)
        self.assertFalse(self.own(info))
        self.assertEqual(w.fill(raw,info['fields'],{}),raw)

    def test_three_complete_captions_and_explicit_false_visibility_remain_usable(self):
        doc,p=self.form('First reading:        Second reading:        Third reading:')
        raw=saved(doc);info=w.inspect(raw);fields=self.own(info)
        self.assertEqual(len(fields),3)
        answers={f['key']:str(i)for i,f in enumerate(fields,1)}
        output=w.fill(raw,info['fields'],answers)
        for value in ('0','false','off'):
            changed=Document(io.BytesIO(output));paragraph=changed.tables[0].cell(0,0).paragraphs[0]
            for run in paragraph.runs:
                mark=OxmlElement('w:vanish');mark.set(w.q('w:val'),value);run._r.get_or_add_rPr().append(mark)
            read=w.effective_values(raw,info['fields'],answers,_written=saved(changed))
            self.assertEqual([read[f['key']]for f in fields],['1','2','3'])
