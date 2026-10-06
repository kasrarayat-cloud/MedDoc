"""Ordinary below-caption bands and requests, and printed sources that are never separate destinations.

Pure parser/writer controls on project-authored synthetic documents; no corpus content, no app resources.
The merged-hierarchy cases were written by the D08 specialist (2026-10-05) and are kept verbatim; only the
source claim they pin was integrated, not its new reading rules.
"""
import io
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import word_v3 as w
from .test_word_v3 import ff_field

BODY = 'word/document.xml#'


def data(doc):
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def body(paragraphs):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    return data(doc)


class BelowCaptionBranches(unittest.TestCase):
    """Bulleted bands and parenthetical requests stay ordinary answers: the signing-caption hold reads the plain
    below-caption branch only."""

    def outcomes(self, raw):
        fields = w.inspect(raw)['fields']
        result = {}
        for field in fields:
            if field['label'] in ('Clinical summary', 'Plan'):
                continue
            try:
                w.fill(raw, fields, {field['key']: '14/03/2026'})
                result[field['label']] = 'written'
            except w.V3Failure:
                result[field['label']] = 'refused'
        return result

    def test_ordinary_bulleted_bands_and_requests_stay_writable(self):
        for paragraphs in (['Clinical summary: ____', '• Detailed explanation of the clinical observations:',
                            '', '', '', 'Plan: ____'],
                           ['Clinical summary: ____', 'Other allergy (please specify):', '', 'Plan: ____']):
            with self.subTest(caption=paragraphs[1]):
                outcome = self.outcomes(body(paragraphs))
                self.assertTrue(outcome)
                self.assertEqual(set(outcome.values()), {'written'}, outcome)


def merged_hierarchy(*, header='Type of unit: (Alpha or Beta)', family='ff'):
    """A column question over two-row merged items: a colon parent with indented subtypes and a peer below."""
    doc = Document()
    doc.add_paragraph('Unit stays')
    table = doc.add_table(rows=5, cols=3)
    for n, text in enumerate(('Stay', header, 'Remarks')):
        table.cell(0, n).text = text
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))

    def box(p, text):
        if family == 'ff':
            ff_field(p, 'checkbox')
        else:
            p.add_run('☐')
        p.add_run(text)
    for start in (1, 3):
        table.cell(start, 0).text = 'First stay:' if start == 1 else 'Second stay:'
        first, last = table.cell(start, 0), table.cell(start + 1, 0)
        first.merge(last)
        for extra in first.paragraphs[1:]:
            first._tc.remove(extra._p)
        cell = table.cell(start, 1)
        cell.text = ''
        box(cell.paragraphs[0], ' Alpha:')
        for text in (' Child one', ' Child two', ' Child three'):
            p = cell.add_paragraph()
            box(p, text)
            ind = OxmlElement('w:ind')
            ind.set(qn('w:left'), '300')
            p._p.get_or_add_pPr().append(ind)
        cell = table.cell(start + 1, 1)
        cell.text = ''
        box(cell.paragraphs[0], ' Beta')
        table.cell(start, 2).text = 'Preserved remark'
        table.cell(start + 1, 2).text = 'Preserved remark continuation'
    return data(doc)


class MergedHierarchySources(unittest.TestCase):
    def test_column_question_is_never_a_separate_destination(self):
        for family in ('ff', 'glyph'):
            with self.subTest(family=family):
                raw = merged_hierarchy(family=family)
                fields = w.inspect(raw)['fields']
                groups = [f for f in fields if any(r.get('relation') == 'merged_checkbox_hierarchy'
                                                   for r in f.get('source_questions', []))]
                self.assertEqual(len(groups), 4)
                self.assertEqual(groups[0]['label'], 'Type of unit: (Alpha or Beta)')
                # No writer of any kind on the column question, the item labels or the merge tails.
                for path in ('b/tbl[0]/tr[0]/tc[1]', 'b/tbl[0]/tr[1]/tc[0]', 'b/tbl[0]/tr[2]/tc[0]',
                             'b/tbl[0]/tr[3]/tc[0]', 'b/tbl[0]/tr[4]/tc[0]'):
                    self.assertEqual([f for f in fields if f['locator']['slot'] == path
                                      or f['locator']['slot'].startswith(path + '/')], [], path)
                # Answering EVERY field of the fresh mapping is accepted by its own verifier.
                root, child = [f for f in groups if '/tr[1]/' in f['locator']['slot']]
                answers = {f['key']: 'Synthetic remark' for f in fields if f not in groups
                           and f['kind'] not in ('choice', 'multi_choice', 'checkbox', 'signature', 'initials')}
                answers.update({root['key']: 'Alpha', child['key']: 'Child two'})
                out = w.fill(raw, fields, answers)
                index = w.SlotIndex(w.open_package(out))
                self.assertEqual(w.constrained_option_states(index, root), [True, False])
                self.assertEqual(w.constrained_option_states(index, child), [False, True, False])
                source = w.SlotIndex(w.open_package(raw))
                header = BODY + 'b/tbl[0]/tr[0]/tc[1]'
                self.assertEqual(w.fingerprint(source.by_id[header].el), w.fingerprint(index.by_id[header].el))

    def test_contradicting_parent_and_subtype_answers_still_refuse(self):
        raw = merged_hierarchy()
        fields = w.inspect(raw)['fields']
        groups = [f for f in fields if any(r.get('relation') == 'merged_checkbox_hierarchy'
                                           for r in f.get('source_questions', []))]
        root, child = [f for f in groups if '/tr[1]/' in f['locator']['slot']]
        with self.assertRaises(w.V3Failure) as error:
            w.fill(raw, fields, {root['key']: 'Beta', child['key']: 'Child two'})
        self.assertEqual(error.exception.code, 'conditional_answer_unselected')

    def test_a_stored_flat_writer_on_the_column_question_is_still_refused(self):
        raw = merged_hierarchy()
        fields = w.inspect(raw)['fields']
        index = w.SlotIndex(w.open_package(raw))
        cell = index.by_id[BODY + 'b/tbl[0]/tr[0]/tc[1]']
        forged = {'key': 'former_flat', 'kind': 'text', 'label': 'Stay', 'context': '', 'options': [],
                  'locator': {'contract': 'v3', 'part': 'word/document.xml', 'slot': 'b/tbl[0]/tr[0]/tc[1]',
                              'fingerprint': w.fingerprint(cell.el), 'target': {'type': 'cell_replace'}}}
        with self.assertRaises(w.V3Failure) as error:
            w.fill(raw, fields + [forged], {'former_flat': 'Synthetic value'})
        self.assertEqual(error.exception.code, 'mapping_mismatch')



if __name__ == '__main__':
    unittest.main()
