"""Complete physical date/time hints beside clock controls stay separate writers."""
import io
from copy import deepcopy
from zipfile import ZipFile
import unittest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from . import word_v3 as w
from .test_word_v3 import ff_field, checkbox_control


def data(doc):
    out=io.BytesIO();doc.save(out);return out.getvalue()


def form(caption='Date and Time Collected:',family='glyph',checked=False):
    d=Document();d.sections[0].header.paragraphs[0].text='Synthetic header'
    d.sections[0].footer.paragraphs[0].text='Synthetic footer'
    t=d.add_table(rows=2,cols=3)
    for i,text in enumerate(('Participant',caption,'Accession Number:')):t.cell(0,i).text=text
    t.cell(1,0).text='Original participant entry';t.cell(1,2).text='Original accession entry'
    c=t.cell(1,1);c.text='MM/DD/YYYY';c.add_paragraph(':(HH:MM)')
    if family:
        p=c.add_paragraph()
        for i,text in enumerate(('AM','PM','24-hr clock')):
            if family=='glyph':p.add_run(('☒' if checked and i==0 else '☐')+' '+text+' ')
            elif family=='ff':ff_field(p,'checkbox',checked=checked and i==0);p.add_run(' '+text+' ')
            else:p._p.append(checkbox_control(checked and i==0));p.add_run(' '+text+' ')
    d.add_paragraph('Neighbor findings: ______');return d


def pairs(result):
    return [f for f in result['fields'] if any(r.get('relation')=='temporal_format_hint' for r in f.get('source_questions',[]))]


class TemporalFormatPairs(unittest.TestCase):
    def test_native_and_glyph_clocks_existing_entries_and_package_parts_are_preserved(self):
        for family in ('glyph','ff','sdt'):
            with self.subTest(family=family):
                raw=data(form(family=family,checked=True));r=w.inspect(raw);fields=pairs(r)
                self.assertEqual([f['kind'] for f in fields],['date','time'])
                self.assertEqual([f['locator']['slot'] for f in fields],['b/tbl[0]/tr[1]/tc[1]/p[0]','b/tbl[0]/tr[1]/tc[1]/p[1]'])
                clock=next(f for f in r['fields'] if f.get('options')==['AM','PM','24-hr clock'])
                oldindex=w.SlotIndex(w.open_package(raw));oldclock=oldindex.by_id['word/document.xml#b/tbl[0]/tr[1]/tc[1]/p[2]']
                out=w.fill(raw,r['fields'],{fields[0]['key']:'10/04/2026',fields[1]['key']:'14:35'})
                current=Document(io.BytesIO(out));t=current.tables[0]
                self.assertEqual([p.text for p in t.cell(1,1).paragraphs[:2]],['10/04/2026','14:35'])
                self.assertEqual(t.cell(0,1).text,'Date and Time Collected:')
                self.assertEqual(t.cell(1,0).text,'Original participant entry');self.assertEqual(t.cell(1,2).text,'Original accession entry')
                self.assertEqual(current.paragraphs[0].text,'Neighbor findings: ______')
                idx=w.SlotIndex(w.open_package(out));self.assertEqual(w.etree.tostring(idx.by_id[oldclock.id].el),w.etree.tostring(oldclock.el))
                self.assertEqual(w.constrained_option_states(idx,clock),[True,False,False])
                with ZipFile(io.BytesIO(raw))as a,ZipFile(io.BytesIO(out))as b:
                    self.assertEqual(a.namelist(),b.namelist())
                    for n in a.namelist():
                        if n!='word/document.xml':self.assertEqual(a.read(n),b.read(n),n)
                self.assertEqual(w.fill(raw,r['fields'],{}),raw)

    def test_renamed_ordinary_complete_caption_needs_no_template_subject_noun(self):
        for caption in ('Specimen date and time:','Date and time recorded:','Date and time signed referral received:','Date and time initial physical observations:'):
            with self.subTest(caption=caption):
                raw=data(form(caption,family=None));f=pairs(w.inspect(raw));self.assertEqual(len(f),2)
                self.assertTrue(all(x['source_questions'][0]['question']==caption for x in f))
                out=w.fill(raw,f,{f[0]['key']:'10/04/2026',f[1]['key']:'14:35'})
                self.assertEqual(Document(io.BytesIO(out)).tables[0].cell(0,1).text,caption)

    def test_human_signing_and_author_captions_refuse_physical_date_forgery(self):
        for caption in ('Date and time Alex Reed should sign:','Date and time Avery is required to sign:',
                        'Witness signature date and time:','Date and time I approve this document:',
                        'Date and time this section is for Alex Reed to write comments:'):
            with self.subTest(caption=caption):
                raw=data(form(caption));r=w.inspect(raw)
                for x in pairs(r):
                    self.assertTrue(x.get('write_protected'))
                    with self.assertRaises(w.V3Failure):w.fill(raw,r['fields'],{x['key']:'10/04/2026'})
                    f=deepcopy(x);f.update(label='Clinical value',kind='text')
                    for k in ('write_protected','destination_roles','answer_policy'):f.pop(k,None)
                    with self.assertRaises(w.V3Failure):w.fill(raw,[f],{f['key']:'Dr Synthetic Finalizer'})
                s=w.SlotIndex(w.open_package(raw)).by_id['word/document.xml#b/tbl[0]/tr[1]/tc[1]/p[0]']
                f={'key':'forged','kind':'text','label':'Clinical value','context':'','locator':{'part':s.part,'slot':s.path,
                   'fingerprint':w.fingerprint(s.el),'contract':'v3','target':{'type':'span','start':0,'end':10,'shape':'format_hint'}}}
                with self.assertRaises(w.V3Failure):w.fill(raw,[f],{'forged':'Dr Synthetic Finalizer'})

    def test_compound_caption_keeps_partial_coverage_and_complete_human_clauses_protected(self):
        for caption in ('Date and time collected; I approve this document:',
                        'Date and time collected.I approve this document:',
                        'Date and time collected. I approve this document:',
                        'Date and time collected: Witness name:',
                        'Date and time collected: I approve this document:',
                        'Date and time collected:\nI approve this document:'):
            with self.subTest(caption=caption):
                raw=data(form(caption));r=w.inspect(raw);self.assertEqual(pairs(r),[])
                issues=[x for x in r['coverage_issues'] if x['slot'].startswith('b/tbl[0]/tr[1]/tc[1]/p[')]
                # The newline caption cannot qualify even the immediate column
                # header and remains outside this compiler's finite source scope.
                if '\n' not in caption:self.assertTrue(issues)
                if 'I approve' not in caption:continue
                index=w.SlotIndex(w.open_package(raw))
                for i,end in ((0,10),(1,8)):
                    slot=index.by_id['word/document.xml#b/tbl[0]/tr[1]/tc[1]/p[%d]'%i]
                    f={'key':'forged','kind':'text','label':'Clinical value','context':'',
                       'locator':{'contract':'v3','part':slot.part,'slot':slot.path,'fingerprint':w.fingerprint(slot.el),
                                  'target':{'type':'span','start':0,'end':end,'shape':'format_hint'}}}
                    with self.assertRaises(w.V3Failure):w.fill(raw,[f],{'forged':'Dr Synthetic Finalizer'})

    def test_partial_reordered_and_extra_questions_are_not_complete_pairs(self):
        for date,time,tail in [('MM/DD/YYYY extra',':(HH:MM)',None),('MM/DD/YYYY','HH:MM pending',None),
                              (':(HH:MM)','MM/DD/YYYY',None),('MM/DD/YYYY',':(HH:MM)','Symptoms: ______'),
                              ('MM/DD/YYYY',':(HH:MM)','☐ Yes ☐ No')]:
            with self.subTest(date=date,time=time,tail=tail):
                d=form(family=None);c=d.tables[0].cell(1,1);c.paragraphs[0].text=date;c.paragraphs[1].text=time
                if tail:c.add_paragraph(tail)
                self.assertEqual(pairs(w.inspect(data(d))),[])

    def test_neighbor_merged_competing_or_unqualified_captions_cannot_own_hints(self):
        for mode in ('neighbor','merge','second-caption','nonterminal','wrong-kind'):
            with self.subTest(mode=mode):
                d=form();t=d.tables[0]
                if mode=='neighbor':t.cell(0,1).text='Other result:';t.cell(0,0).text='Date and Time Collected:'
                elif mode=='merge':t.cell(0,1).merge(t.cell(0,2)).text='Date and Time Collected:'
                elif mode=='second-caption':t.cell(0,1).add_paragraph('Date and time received:')
                elif mode=='nonterminal':t.cell(0,1).text='Date and Time Collected for reference'
                else:t.cell(0,1).text='Accession Number:'
                self.assertEqual(pairs(w.inspect(data(d))),[])

    def test_hidden_or_linked_ink_is_not_appointed_by_format_text(self):
        for which in ('caption','date','time','clock'):
            for style in (False,True):
                with self.subTest(which=which,style=style):
                    d=form();c=d.tables[0].cell(1,1)
                    p=d.tables[0].cell(0,1).paragraphs[0] if which=='caption' else c.paragraphs[{'date':0,'time':1,'clock':2}[which]]
                    if style:
                        from docx.enum.style import WD_STYLE_TYPE
                        s=d.styles.add_style('HiddenTemporal',WD_STYLE_TYPE.CHARACTER);s.font.hidden=True;p.runs[0].style=s
                    else:p.runs[0].font.hidden=True
                    self.assertEqual(pairs(w.inspect(data(d))),[])
        d=form();p=d.tables[0].cell(1,1).paragraphs[0];link=OxmlElement('w:hyperlink');link.set(qn('r:id'),'rIdUnresolved')
        for r in list(p._p.findall(qn('w:r'))):link.append(r)
        p._p.append(link);self.assertEqual(pairs(w.inspect(data(d))),[])

    def test_metadata_removed_relationship_partial_span_and_whole_parent_forgery_refuse(self):
        raw=data(form());date=pairs(w.inspect(raw))[0]
        for mutation in ('kind','label','context','remove-relation','format','source','span','boundary-insert','whole-cell'):
            with self.subTest(mutation=mutation):
                f=deepcopy(date)
                if mutation in ('kind','label','context'):f[mutation]='text' if mutation=='kind' else 'Forged'
                elif mutation=='remove-relation':f.pop('source_questions')
                elif mutation=='format':f['source_questions'][0]['format']='YYYY/MM/DD'
                elif mutation=='source':f['source_questions'][0]['sources'][0]=f['source_questions'][0]['sources'][1]
                elif mutation=='span':f['locator']['target']['end']-=1
                elif mutation=='boundary-insert':f['locator']['target'].update(start=10,end=10,shape='after_label')
                else:
                    s=w.SlotIndex(w.open_package(raw)).by_id['word/document.xml#b/tbl[0]/tr[1]/tc[1]']
                    f['locator'].update(slot=s.path,fingerprint=w.fingerprint(s.el),target={'type':'cell_replace'})
                with self.assertRaises(w.V3Failure):w.fill(raw,[f],{f['key']:'Synthetic wrong answer'})

    def test_repeated_temporal_paragraph_writer_never_overwrites_an_answer(self):
        raw=data(form());r=w.inspect(raw);date,time=pairs(r)
        for source_key,shared_object in ((False,False),(True,False),(True,True)):
            duplicate=date if shared_object else deepcopy(date)
            if not source_key:duplicate['key']='second-date-writer'
            fields=[date,duplicate,time]
            cases=[{date['key']:'04/10/2026'}]
            if not source_key:cases.append({date['key']:'04/10/2026',duplicate['key']:'05/10/2026 EXTRA'})
            for answers in cases:
                with self.subTest(source_key=source_key,shared_object=shared_object,answers=answers),self.assertRaises(w.V3Failure)as error:
                    w.fill(raw,fields,answers)
                self.assertEqual(error.exception.code,'mapping_mismatch')
            self.assertEqual(w.fill(raw,fields,{}),raw)
        out=w.fill(raw,r['fields'],{date['key']:'04/10/2026',time['key']:'14:35'})
        self.assertEqual([p.text for p in Document(io.BytesIO(out)).tables[0].cell(1,1).paragraphs[:2]],['04/10/2026','14:35'])

    def test_different_original_cannot_reuse_the_old_pair(self):
        first=data(form());f=pairs(w.inspect(first));second=data(form('Date and time received:'))
        with self.assertRaises(w.V3Failure):w.fill(second,f,{f[0]['key']:'10/04/2026'})
        current=pairs(w.inspect(second));self.assertEqual(len(current),2)
        out=w.fill(second,current,{current[0]['key']:'10/04/2026',current[1]['key']:'14:35'})
        self.assertEqual(Document(io.BytesIO(out)).tables[0].cell(1,1).paragraphs[0].text,'10/04/2026')
