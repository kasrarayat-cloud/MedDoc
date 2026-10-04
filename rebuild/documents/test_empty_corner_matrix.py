"""Original-source ownership of an empty corner above a checkbox response matrix."""
from copy import deepcopy
from io import BytesIO
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from . import word_v3 as w


def native_box(p, kind, checked=False):
    if kind == 'glyph':
        p.add_run('☒' if checked else '☐')
    elif kind == 'ff':
        begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
        data = OxmlElement('w:ffData'); data.append(OxmlElement('w:enabled'))
        cb = OxmlElement('w:checkBox'); state = OxmlElement('w:default')
        state.set(qn('w:val'), '1' if checked else '0'); cb.append(state); data.append(cb); begin.append(data)
        p.add_run()._r.append(begin)
        instruction = OxmlElement('w:instrText'); instruction.text = ' FORMCHECKBOX '
        p.add_run()._r.append(instruction)
        end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end'); p.add_run()._r.append(end)
    else:
        sdt = OxmlElement('w:sdt'); pr = OxmlElement('w:sdtPr'); sdt.append(pr)
        cb = etree.SubElement(pr, w.q('w14:checkbox'))
        state = etree.SubElement(cb, w.q('w14:checked')); state.set(w.q('w14:val'), '1' if checked else '0')
        content = OxmlElement('w:sdtContent'); sdt.append(content)
        run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = '☒' if checked else '☐'
        run.append(text); content.append(run); p._p.append(sdt)


def form(kind='glyph', labels=('Observation alpha', 'Observation beta', 'Observation gamma'),
         headers=('Yes', 'No', 'Refused'), mutate=None):
    doc = Document(); doc.add_paragraph('Synthetic response inventory')
    doc.sections[0].header.paragraphs[0].text = 'Fixed letterhead'
    doc.sections[0].footer.paragraphs[0].text = 'Fixed footer'
    table = doc.add_table(rows=len(labels) + 1, cols=4)
    table.cell(0, 0).text = '\u00a0'
    for col, label in enumerate(headers, 1): table.cell(0, col).text = label
    for row, label in enumerate(labels, 1):
        table.cell(row, 0).text = label
        for col in (1, 2, 3):
            native_box(table.cell(row, col).paragraphs[0], kind, checked=row == 2 and col == 2)
    doc.add_paragraph('Fixed neighboring history')
    if mutate: mutate(doc, table)
    data = BytesIO(); doc.save(data); return data.getvalue()


def by_label(mapping): return {field['label']: field for field in mapping['fields']}


class EmptyCornerMatrix(unittest.TestCase):
    def test_complete_original_axis_yields_each_row_and_never_a_header_answer(self):
        for kind in ('glyph', 'ff', 'sdt'):
            with self.subTest(kind=kind):
                original = form(kind); mapping = w.inspect(original); fields = by_label(mapping)
                self.assertEqual(set(fields), {'Observation alpha', 'Observation beta', 'Observation gamma'})
                for field in fields.values():
                    self.assertEqual((field['kind'], field['options']), ('choice', ['Yes', 'No', 'Refused']))
                    self.assertNotIn('/tr[0]/', field['locator']['slot'])
                answer = fields['Observation alpha']
                exported = w.fill(original, mapping['fields'], {answer['key']: 'Refused'})
                for label, field in fields.items():
                    self.assertEqual(w.constrained_option_states(w.SlotIndex(w.open_package(exported)), field),
                                     [False, False, True] if label == 'Observation alpha' else [False, True, False] if label == 'Observation beta' else [False, False, False])
                with ZipFile(BytesIO(original)) as old, ZipFile(BytesIO(exported)) as new:
                    for part in old.namelist():
                        if part != 'word/document.xml': self.assertEqual(old.read(part), new.read(part))
                root = w.open_package(exported).tree('word/document.xml')
                header = root.find(w.q('w:body')).find(w.q('w:tbl')).find(w.q('w:tr'))
                self.assertEqual([w.text_of(cell) for cell in w.row_cells_of(header)[1:]], ['Yes', 'No', 'Refused'])
                self.assertEqual(w.fill(original, mapping['fields'], {}), original)

    def test_historical_or_forged_header_writer_refuses_with_original_unchanged(self):
        original = form(); mapping = w.inspect(original)
        package = w.open_package(original); detector = w.Detector(package)
        slot = detector.index.by_id['word/document.xml#b/tbl[0]/tr[0]/tc[0]']
        header = dict(key='historic_header', kind='text', label='Observation alpha', context='Synthetic response inventory',
                      locator=dict(part=slot.part, slot=slot.path, target=dict(type='cell'),
                                   contract='v3', fingerprint=w.fingerprint(slot.el)))
        for fields in ([header], [*mapping['fields'], header]):
            with self.subTest(fields=len(fields)):
                with self.assertRaisesRegex(w.V3Failure, 'fixed response-column header spacing.*Reinspect'):
                    w.fill(original, fields, {'historic_header': 'Yes'})
                self.assertEqual(w.fill(original, fields, {}), original)
        self.assertTrue(any('fixed spacing' in note and 'reinspect' in note for note in mapping['notes']))

    def test_axis_needs_complete_unique_visible_physical_bands(self):
        def hide_header(doc, table): table.cell(0, 2).paragraphs[0].runs[0].font.hidden = True
        def hidden_box(doc, table): table.cell(1, 1).paragraphs[0].runs[0].font.hidden = True
        def merged_header(doc, table): table.cell(0, 1).merge(table.cell(0, 2))
        def merged_body(doc, table): table.cell(1, 1).merge(table.cell(2, 1))
        def extra_control(doc, table): table.cell(1, 1).paragraphs[0].add_run(' ☐')
        def extra_unprinted_control(doc, table): native_box(table.cell(1, 1).paragraphs[0], 'ff')
        def marked_corner(doc, table): table.cell(0, 0).paragraphs[0].runs[0].underline = True
        variants = [('one row', dict(labels=('Observation alpha',))),
                    ('duplicate headers', dict(headers=('Yes', 'Yes', 'Refused'))),
                    ('independent fields', dict(headers=('Name', 'Date', 'Signature'))),
                    *[(fn.__name__, dict(mutate=fn)) for fn in (hide_header, hidden_box, merged_header, merged_body,
                                                              extra_control, extra_unprinted_control, marked_corner)]]
        for name, kwargs in variants:
            with self.subTest(name=name):
                original = form(**kwargs); detector = w.Detector(w.open_package(original))
                table = detector.index.by_id['word/document.xml#b/tbl[0]']
                self.assertIsNone(detector.response_matrix_corner(table, 0))
                self.assertEqual(w.fill(original, w.inspect(original)['fields'], {}), original)

    def test_human_row_remains_protected_and_ordinary_receipts_stay_usable(self):
        for label in ('Witness must write the following:', 'This area is for Alex Reed to enter the answer:',
                      'Date Alex Reed should sign:', 'I approve this document:'):
            with self.subTest(label=label):
                original = form(labels=(label, 'Ordinary observation', 'Neighbor observation'))
                mapping = w.inspect(original); human = by_label(mapping).get(label.rstrip(':')) or by_label(mapping).get(label)
                self.assertIsNotNone(human)
                with self.assertRaises(w.V3Failure): w.fill(original, mapping['fields'], {human['key']: 'Yes'})
                forged = deepcopy(human); forged['label'] = 'Ordinary observation'; forged.pop('write_protected', None)
                with self.assertRaises(w.V3Failure): w.fill(original, [forged], {forged['key']: 'Yes'})
        for label in ('Initial observations', 'Signed referral received', 'Review of systems'):
            with self.subTest(label=label):
                original = form(labels=(label, 'Second observation', 'Third observation'))
                mapping = w.inspect(original); field = by_label(mapping)[label]
                exported = w.fill(original, mapping['fields'], {field['key']: 'No'})
                self.assertEqual(w.constrained_option_states(w.SlotIndex(w.open_package(exported)), field), [False, True, False])

    def test_all_that_apply_remains_held_and_does_not_promote_the_row(self):
        def many(doc, table): doc.paragraphs[0].text = 'Select all that apply.'
        original = form(mutate=many); mapping = w.inspect(original)
        self.assertFalse(any(f['label'].startswith('Observation') and f['kind'] == 'multi_choice' for f in mapping['fields']))
        self.assertTrue(any(issue['code'] == 'unmapped_answer' for issue in mapping['coverage_issues']))
        self.assertEqual(w.fill(original, mapping['fields'], {}), original)

    def test_unqualified_response_bands_cannot_fallback_to_a_header_answer(self):
        def duplicate_native(doc, table):
            for row in (1, 2, 3):
                for col in (1, 2, 3): native_box(table.cell(row, col).paragraphs[0], 'ff')
        def hidden_controls(doc, table):
            for row in (1, 2, 3):
                for col in (1, 2, 3): table.cell(row, col).paragraphs[0].runs[0].font.hidden = True
        def merged_controls(doc, table):
            table.cell(1, 1).merge(table.cell(2, 1))
        def merged_headers(doc, table): table.cell(0, 1).merge(table.cell(0, 2))
        def hidden_headers(doc, table): table.cell(0, 1).paragraphs[0].runs[0].font.hidden = True
        def competing_headers(doc, table): table.cell(0, 2).text = 'Yes'
        for mutate in (duplicate_native, hidden_controls, merged_controls, merged_headers, hidden_headers, competing_headers):
            with self.subTest(mutate=mutate.__name__):
                original = form(mutate=mutate); mapping = w.inspect(original)
                self.assertFalse(any(field['locator']['slot'] == 'b/tbl[0]/tr[0]/tc[0]' for field in mapping['fields']))
                detector = w.Detector(w.open_package(original))
                table = detector.index.by_id['word/document.xml#b/tbl[0]']
                self.assertIsNone(detector.response_matrix_corner(table, 0))
                self.assertIsNotNone(detector.response_matrix_header_corner(table, 0))
                self.assertTrue(any(issue['code'] == 'unmapped_response_matrix' for issue in mapping['coverage_issues']))
                corner = detector.index.by_id['word/document.xml#b/tbl[0]/tr[0]/tc[0]']
                forged = dict(key='old_header', kind='text', label='Observation alpha', context='',
                              locator=dict(part=corner.part, slot=corner.path, target=dict(type='cell'),
                                           contract='v3', fingerprint=w.fingerprint(corner.el)))
                with self.assertRaisesRegex(w.V3Failure, 'fixed response-column header spacing.*Reinspect'):
                    w.fill(original, [forged], {'old_header': 'Yes'})
                self.assertEqual(w.fill(original, mapping['fields'], {}), original)

    def test_drawn_or_explicit_answer_corner_is_not_quarantined_as_spacing(self):
        def drawn(doc, table): table.cell(0, 0).text = '___________'
        def explicit(doc, table): table.cell(0, 0).text = 'Reference: ___________'
        def underlined(doc, table): table.cell(0, 0).paragraphs[0].runs[0].underline = True
        for mutate in (drawn, explicit, underlined):
            with self.subTest(mutate=mutate.__name__):
                original = form(mutate=mutate); detector = w.Detector(w.open_package(original))
                self.assertIsNone(detector.response_matrix_header_corner(detector.index.by_id['word/document.xml#b/tbl[0]'], 0))


if __name__ == '__main__': unittest.main()
