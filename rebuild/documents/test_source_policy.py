"""Only explicit source roles govern authorship; patient facts do not imply consent."""
import unittest

from docx import Document
from docx.oxml import OxmlElement
from . import word_v3 as w
from .test_question_boxes import data


def pair(doc, text):
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = text
    return table


class SourcePolicy(unittest.TestCase):
    def test_explicit_patient_scope_is_carried_until_the_next_role(self):
        for heading in ('For the patient:', 'To be completed by the patient:', 'Completed by patient:'):
            with self.subTest(heading=heading):
                doc = Document(); doc.add_heading('Patient comments', 1); doc.add_paragraph(heading)
                pair(doc, 'What progress have you made?'); pair(doc, 'What else would you like to say?')
                doc.add_paragraph('For the clinician:'); pair(doc, 'What progress was observed?')
                fields = w.inspect(data(doc))['fields']
                first, second, last = [f for f in fields if f['locator']['target']['type'] == 'cell']
                self.assertEqual(first['answer_policy']['source_author'], 'patient')
                self.assertEqual(first['answer_policy'], second['answer_policy'])
                self.assertEqual(first['answer_policy']['source']['text'], heading)
                self.assertNotIn('answer_policy', last)

    def test_patient_mentions_and_heading_changes_do_not_imply_patient_authorship(self):
        for heading in ('Patient details:', 'Patient comments', 'For the patient:'):
            with self.subTest(heading=heading):
                doc = Document(); doc.add_paragraph(heading)
                if heading == 'For the patient:': doc.add_heading('Clinician findings', 1)
                pair(doc, 'What symptoms did the patient report?')
                fields = w.inspect(data(doc))['fields']
                field = next(f for f in fields if f['locator']['target']['type'] == 'cell')
                self.assertNotIn('answer_policy', field)

    def test_explicit_review_attestation_keeps_the_complete_source(self):
        question = ('Please confirm that the patient has had sight of this report and/or '
                    'had an opportunity to add their comments, if appropriate or necessary.')
        doc = Document(); pair(doc, question); pair(doc, 'What did the patient report today?')
        fields = w.inspect(data(doc))['fields']
        first = fields[0]
        self.assertEqual(first['answer_policy']['attestation'], 'document_review')
        self.assertEqual(first['answer_policy']['attestation_source']['text'], question)
        self.assertNotIn('answer_policy', fields[1])

    def test_changed_author_source_cannot_keep_its_old_policy_at_export(self):
        doc = Document(); role = doc.add_paragraph('For the patient:')
        pair(doc, 'What progress have you made?')
        fields = w.inspect(data(doc))['fields']
        field = next(f for f in fields if f.get('answer_policy'))
        w.fill(data(doc), fields, {field['key']: 'The patient says progress is good.'})
        role.text = 'For the clinician:'
        with self.assertRaises(w.V3Failure) as failure:
            w.fill(data(doc), fields, {field['key']: 'The patient says progress is good.'})
        self.assertEqual(failure.exception.code, 'mapping_mismatch')

    def test_role_on_section_boundary_does_not_leak_into_next_section(self):
        doc = Document(); role = doc.add_paragraph('For the patient:')
        role._p.get_or_add_pPr().append(OxmlElement('w:sectPr'))
        pair(doc, 'What symptoms were observed?')
        field = next(f for f in w.inspect(data(doc))['fields'] if f['locator']['target']['type'] == 'cell')
        self.assertNotIn('answer_policy', field)

    def test_patient_example_instruction_keeps_patient_scope(self):
        doc = Document(); doc.add_paragraph('For the patient:')
        doc.add_paragraph('For example: describe changes since your last visit.')
        pair(doc, 'What progress have you made?')
        field = next(f for f in w.inspect(data(doc))['fields'] if f['locator']['target']['type'] == 'cell')
        self.assertEqual(field['answer_policy']['source_author'], 'patient')

    def test_example_marker_case_does_not_change_patient_authorship(self):
        for heading in ('For EXAMPLE:', 'FOR EXAMPLE:', 'For instance:', 'For INSTANCE:'):
            with self.subTest(heading=heading):
                doc = Document(); doc.add_paragraph('For the patient:'); doc.add_paragraph(heading)
                pair(doc, 'What progress have you made?')
                field = next(f for f in w.inspect(data(doc))['fields'] if f['locator']['target']['type'] == 'cell')
                self.assertEqual(field.get('answer_policy', {}).get('source_author'), 'patient')

    def test_clinician_review_of_symptom_report_is_not_patient_attestation(self):
        doc = Document(); pair(doc, 'Confirm the patient report of chest pain was reviewed by the clinician.')
        self.assertNotIn('answer_policy', w.inspect(data(doc))['fields'][0])

    def test_literal_document_objects_keep_the_full_conjunctive_review_instruction(self):
        for noun in ('report', 'form', 'document', 'application'):
            with self.subTest(noun=noun):
                question = ('Please confirm that the patient has had sight of this ' + noun +
                            ' and had an opportunity to add their comments')
                doc = Document(); pair(doc, question)
                field = w.inspect(data(doc))['fields'][0]
                self.assertEqual(field['answer_policy']['attestation'], 'document_review')
                self.assertEqual(field['answer_policy']['attestation_source']['text'], question)
        doc = Document(); pair(doc, 'Confirm the patient reports that the application of cream was reviewed by the clinician.')
        self.assertNotIn('answer_policy', w.inspect(data(doc))['fields'][0])

    def test_explicit_acronym_or_named_completion_role_ends_patient_authorship(self):
        for role in ('For the RC:', 'For RC:', 'For QA:', 'To be completed by Alex:', 'Completed by the reviewer:'):
            with self.subTest(role=role):
                doc = Document(); doc.add_paragraph('For the patient:'); pair(doc, 'What would you like to say?')
                doc.add_paragraph(role); pair(doc, 'What was confirmed?')
                fields = [f for f in w.inspect(data(doc))['fields'] if f['locator']['target']['type'] == 'cell']
                self.assertIn('answer_policy', fields[0])
                self.assertNotIn('answer_policy', fields[1])
