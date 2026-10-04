"""Explicit source-bound recovery; automatic clinical caption association stays unchanged."""
from copy import deepcopy
import io
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import word_v3 as w


def saved(doc):
    stream = io.BytesIO(); doc.save(stream); return stream.getvalue()


def source(band='▯▯▯/▯▯▯', caption='Blood pressure'):
    doc = Document(); doc.add_heading('Supine measurements', level=2)
    doc.add_paragraph(caption); doc.add_paragraph(band)
    doc.add_paragraph('\tSystolic\tdiastolic'); doc.add_paragraph('Neighbor: ______')
    doc.sections[0].header.paragraphs[0].text = 'Synthetic clinic'
    doc.sections[0].footer.paragraphs[0].text = 'Preserved footer'
    return saved(doc)


def located(raw, number, label='Reviewed pressure', kind='text', slot='word/document.xml#b/p[2]'):
    return w.locate(raw, [{'slot': slot, 'kind': kind, 'label': label, 'numeric_part': number}])[0]


class NumericPartRecovery(unittest.TestCase):
    def test_explicit_review_divides_complete_band_without_changing_automatic_detection(self):
        raw = source(); before = w.inspect(raw)['fields']
        self.assertFalse(any(field['locator']['target'].get('shape') == 'numeric_part' for field in before))
        parts = next(row['numeric_parts'] for row in w.outline(raw)['slots'] if row['slot'].endswith('#b/p[2]'))
        self.assertEqual(parts, [{'number': 1, 'digits': 3}, {'number': 2, 'digits': 3}])
        for number, bounds in ((1, (0, 3)), (2, (4, 7))):
            field = located(raw, number)
            self.assertEqual(field['origin'], 'manual'); self.assertEqual(field['max_chars'], 3)
            self.assertEqual(field['locator']['target'], {'type': 'span', 'start': bounds[0], 'end': bounds[1],
                                                        'shape': 'numeric_part', 'number': number})
            self.assertEqual(field['source_questions'][0]['source']['fingerprint'], field['locator']['fingerprint'])
            self.assertEqual(w.fill(raw, [field], {}), raw)
        self.assertEqual(w.inspect(raw)['fields'], before)

    def test_both_parts_and_variable_lengths_preserve_separators_structure_and_other_parts(self):
        raw = source(); fields = [located(raw, 1, 'Supine systolic'), located(raw, 2, 'Supine diastolic')]
        answers = {fields[0]['key']: '120', fields[1]['key']: '80'}
        output = w.fill(raw, fields, answers)
        self.assertEqual(Document(io.BytesIO(output)).paragraphs[2].text, '120/80')
        self.assertEqual(w.effective_values(raw, fields, answers), answers)
        before, after = w.SlotIndex(w.open_package(raw)), w.SlotIndex(w.open_package(output))
        for slot in before.slots:
            if slot.id != 'word/document.xml#b/p[2]':
                self.assertEqual(w.fingerprint(slot.el), w.fingerprint(after.by_id[slot.id].el), slot.id)
        with ZipFile(io.BytesIO(raw)) as a, ZipFile(io.BytesIO(output)) as b:
            self.assertEqual(a.namelist(), b.namelist())
            for name in a.namelist():
                if name != 'word/document.xml': self.assertEqual(a.read(name), b.read(name))
        answers = {fields[1]['key']: '075'}
        self.assertEqual(Document(io.BytesIO(w.fill(raw, fields, answers))).paragraphs[2].text, '▯▯▯/075')
        self.assertEqual(w.effective_values(raw, fields, answers), {fields[0]['key']: '', fields[1]['key']: '075'})
        raw = source(' ▯▯ / ▯▯ / ▯▯▯▯ ')
        fields = [located(raw, n, 'Reviewed part %d' % n) for n in (1, 2, 3)]
        answers = {field['key']: answer for field, answer in zip(fields, ('4', '09', '2026'))}
        self.assertEqual(Document(io.BytesIO(w.fill(raw, fields, answers))).paragraphs[2].text, ' 4 / 09 / 2026 ')
        self.assertEqual(w.effective_values(raw, fields, answers), answers)
        self.assertIsNone(w.numeric_part_value('▯▯▯/▯▯▯', '120:80', 1))
        self.assertIsNone(w.numeric_part_value('▯▯▯/▯▯▯', '120/80 and 90', 2))

    def test_invalid_filled_hidden_linked_control_or_mixed_bands_refuse_whole_group(self):
        for band in ('▯▯▯', '120/▯▯▯', '▯▯▯/80', '▯▯▯/▯▯▯ mmHg', '▯▯▯\n/▯▯▯', '▯▯▯//▯▯▯'):
            with self.subTest(band=band), self.assertRaises(w.V3Failure): located(source(band), 1)
        for variant in ('hidden', 'hyperlink', 'control', 'break', 'numbering', 'unknown_style'):
            doc = Document(io.BytesIO(source())); paragraph = doc.paragraphs[2]
            if variant == 'hidden': paragraph.runs[0].font.hidden = True
            elif variant == 'hyperlink':
                link = OxmlElement('w:hyperlink'); paragraph._p.append(link); link.append(paragraph.runs[0]._r)
            elif variant == 'control':
                from .test_word_v3 import _sdt
                paragraph._p.append(_sdt('rich'))
            elif variant == 'break': paragraph.add_run().add_break()
            elif variant == 'numbering': paragraph._p.get_or_add_pPr().append(OxmlElement('w:numPr'))
            else:
                style = OxmlElement('w:pStyle'); style.set(qn('w:val'), 'UnknownStyle'); paragraph._p.get_or_add_pPr().append(style)
            with self.subTest(variant=variant), self.assertRaises(w.V3Failure): located(saved(doc), 1)
        for number in (0, 3, -1, True, '1', 1.0):
            with self.subTest(number=number), self.assertRaises(w.V3Failure): located(source(), number)
        for kind in ('long_text', 'date', 'checkbox'):
            with self.subTest(kind=kind), self.assertRaises(w.V3Failure): located(source(), 1, kind=kind)

    def test_oversized_non_digit_or_forged_answers_refuse_atomically(self):
        raw = source(); original = located(raw, 1)
        for answer in ('1234', '12/80', '12.5', '-12', '1 2', '１２', 'Dr Synthetic Finalizer'):
            with self.subTest(answer=answer), self.assertRaises(w.V3Failure): w.fill(raw, [original], {original['key']: answer})
        for variant in ('relation', 'shape_relation', 'bounds', 'number', 'count', 'capacity', 'copy', 'source', 'kind'):
            field = deepcopy(original)
            if variant in ('relation', 'shape_relation'):
                field.pop('source_questions'); field.pop('origin')
                if variant == 'shape_relation': field['locator']['target']['shape'] = 'segmented'
            elif variant == 'bounds': field['locator']['target']['end'] = 7
            elif variant == 'number': field['locator']['target']['number'] = 2
            elif variant == 'count': field['source_questions'][0]['parts'] = 3
            elif variant == 'capacity': field.pop('max_chars')
            elif variant == 'copy': field['locator']['target']['copies'] = [deepcopy(field['locator'])]
            elif variant == 'source': field['source_questions'][0]['source']['fingerprint'] = 'changed'
            else: field['kind'] = 'long_text'
            with self.subTest(variant=variant), self.assertRaises(w.V3Failure): w.fill(raw, [field], {field['key']: '120'})

    def test_duplicate_and_whole_band_overlap_refuse_while_original_mapping_still_exports(self):
        raw = source(); original = located(raw, 1); second = located(raw, 1, 'Repeated part')
        with self.assertRaises(w.V3Failure): w.fill(raw, [original, second], {original['key']: '120'})
        old = next(field for field in w.inspect(raw)['fields'] if field['locator']['slot'] == 'b/p[2]')
        with self.assertRaises(w.V3Failure): w.fill(raw, [original, old], {original['key']: '120'})
        self.assertEqual(Document(io.BytesIO(w.fill(raw, [old], {old['key']: '120/80'}))).paragraphs[2].text, '120/80')

    def test_actual_signing_source_stays_protected_and_ordinary_labels_still_work(self):
        for caption in ('Manager signature:', 'Witness signature:', 'Consent obtained:', 'I approve this document:'):
            with self.subTest(caption=caption), self.assertRaises(w.V3Failure):
                located(source(caption=caption), 1, 'Ordinary clinical value')
        for caption in ('Initial reading', 'Signed receipt status', 'Review of systems'):
            with self.subTest(caption=caption):
                raw = source(caption=caption); field = located(raw, 1, 'Ordinary clinical value')
                self.assertEqual(Document(io.BytesIO(w.fill(raw, [field], {field['key']: '120'}))).paragraphs[2].text,
                                 '120/▯▯▯')

    def test_old_whole_band_and_forged_partial_writes_obey_physical_continuity(self):
        for table in (False, True):
            for spacers in (0, 1, 3):
                for boundary in ('', 'clinical', 'heading', 'marked'):
                    doc = Document()
                    host = doc.add_table(rows=2, cols=1).cell(0, 0) if table else doc
                    host.add_paragraph('Manager signature:')
                    if table: host = doc.tables[0].cell(1, 0)
                    for _ in range(spacers): host.add_paragraph()
                    if boundary == 'clinical': host.add_paragraph('Blood pressure')
                    elif boundary == 'heading': host.add_paragraph('Clinical history', style='Heading 2')
                    elif boundary == 'marked': host.add_paragraph().add_run().add_break()
                    band = host.add_paragraph('▯▯▯/▯▯▯'); raw = saved(doc)
                    index = w.SlotIndex(w.open_package(raw))
                    slot = next(s for s in index.slots if s.kind == 'p' and w.paragraph_text(s.el)[0] == '▯▯▯/▯▯▯')
                    # A marked blank cannot cancel a real table-column signer
                    # header. Only a substantive clinical caption or heading
                    # supplies a different source for the numeric area there.
                    held = boundary == '' or table and boundary == 'marked'
                    with self.subTest(table=table, spacers=spacers, boundary=boundary):
                        if held:
                            with self.assertRaises(w.V3Failure): located(raw, 1, slot=slot.id)
                            forged = {'key': 'forged', 'label': 'Clinical value', 'kind': 'text', 'locator':
                                      {'contract': 'v3', 'part': slot.part, 'slot': slot.path,
                                       'fingerprint': w.fingerprint(slot.el),
                                       'target': {'type': 'span', 'start': 0, 'end': 7, 'shape': 'segmented'}}}
                            with self.assertRaises(w.V3Failure): w.fill(raw, [forged], {'forged': '120/80'})
                        else:
                            field = located(raw, 1, slot=slot.id)
                            self.assertIn('120/▯▯▯', w.text_of(w.open_package(w.fill(raw, [field], {field['key']: '120'})).tree('word/document.xml')))
