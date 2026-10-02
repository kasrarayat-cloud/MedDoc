"""Contract v3 through the sandboxed worker, on synthetic blank forms built here (no patient data)."""

from copy import deepcopy
import io
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from .service import DocumentError, fill_template, inspect_template, locate_fields, outline_template

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"


def _sdt(kind, *, options=(), placeholder="Choose an item.", nested=None):
    sdt = OxmlElement("w:sdt")
    pr = OxmlElement("w:sdtPr")
    sdt.append(pr)
    if kind == "dropdown":
        dd = OxmlElement("w:dropDownList")
        for o in options:
            li = OxmlElement("w:listItem")
            li.set(qn("w:displayText"), o)
            li.set(qn("w:value"), o)
            dd.append(li)
        pr.append(dd)
        pr.append(OxmlElement("w:showingPlcHdr"))
    elif kind == "checkbox":
        cb = etree.SubElement(pr, "{%s}checkbox" % W14)
        etree.SubElement(cb, "{%s}checked" % W14).set("{%s}val" % W14, "0")
    content = OxmlElement("w:sdtContent")
    sdt.append(content)
    if nested is not None:
        content.append(nested)
    else:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "☐" if kind == "checkbox" else placeholder
        r.append(t)
        content.append(r)
    return sdt


def ff_field(p, kind, options=(), checked=False, name="F"):
    """A legacy form field (FORMDROPDOWN or FORMCHECKBOX) appended to python-docx paragraph `p`."""
    def run(child):
        r = OxmlElement("w:r")
        r.append(child)
        p._p.append(r)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    data = OxmlElement("w:ffData")
    n = OxmlElement("w:name")
    n.set(qn("w:val"), name)
    data.append(n)
    data.append(OxmlElement("w:enabled"))
    if kind == "dropdown":
        dd = OxmlElement("w:ddList")
        for o in options:
            e = OxmlElement("w:listEntry")
            e.set(qn("w:val"), o)
            dd.append(e)
        data.append(dd)
    else:
        cb = OxmlElement("w:checkBox")
        cb.append(OxmlElement("w:sizeAuto"))
        d = OxmlElement("w:default")
        d.set(qn("w:val"), "1" if checked else "0")
        cb.append(d)
        data.append(cb)
    begin.append(data)
    run(begin)
    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = " FORMDROPDOWN " if kind == "dropdown" else " FORMCHECKBOX "
    run(instr)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run(end)


def ff_result(path):
    with ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    return [r.get("{%s}val" % W) for r in root.iter("{%s}result" % W)]


def build_form(path, *, nested_wrapper=False):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Patient name: ____________"
    doc.add_paragraph("Synthetic clinic intake")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Referring physician"
    table.cell(1, 0).text = "Reason for visit"
    table.cell(2, 0).text = "Allergies"
    doc.add_paragraph("Date of visit: ____________")
    doc.add_paragraph("Smoker?  ☐ Yes  ☐ No")
    p = doc.add_paragraph("Laterality: ")
    p._p.append(_sdt("dropdown", options=["Choose an item.", "Left", "Right", "Bilateral"]))
    doc.add_paragraph("Clinician signature: ____________")
    if nested_wrapper:
        inner_p = OxmlElement("w:p")
        inner_r = OxmlElement("w:r")
        inner_t = OxmlElement("w:t")
        inner_t.text = "Side: "
        inner_r.append(inner_t)
        inner_p.append(inner_r)
        inner_p.append(_sdt("dropdown", options=["Left", "Right"]))
        wrapper = _sdt("rich", nested=inner_p)
        doc.element.body.insert(len(doc.element.body) - 1, wrapper)
    doc.save(path)


def fields_by_label(fields):
    return {f["label"]: f for f in fields}


def part_text(path, part="word/document.xml"):
    with ZipFile(path) as z:
        root = etree.fromstring(z.read(part))
    return "".join(t.text or "" for t in root.iter("{%s}t" % W))


class ContractV3(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.original = self.dir / "form.docx"
        build_form(self.original)

    def tearDown(self):
        self.tmp.cleanup()

    def test_inspect_finds_every_structure_through_the_sandbox(self):
        info = inspect_template(self.original, engine="v3")
        self.assertEqual(info.get("contract"), "v3")
        by = fields_by_label(info["fields"])
        self.assertEqual(by["Referring physician"]["locator"]["target"]["type"], "cell")
        self.assertEqual(by["Date of visit"]["kind"], "date")
        self.assertEqual(by["Smoker?"]["kind"], "choice")
        self.assertEqual(by["Smoker?"]["options"], ["Yes", "No"])
        self.assertEqual(by["Laterality"]["options"], ["Left", "Right", "Bilateral"])
        self.assertEqual(by["Clinician signature"]["kind"], "signature")
        self.assertIn("word/header1.xml", {f["locator"]["part"] for f in info["fields"]})

    def test_fill_writes_answers_and_copies_every_other_part(self):
        info = inspect_template(self.original, engine="v3")
        by = fields_by_label(info["fields"])
        answers = {by["Referring physician"]["key"]: "Dr Example", by["Date of visit"]["key"]: "03/14/2026",
                   by["Smoker?"]["key"]: "No", by["Laterality"]["key"]: "Left", by["Patient name"]["key"]: "Synthetic Person"}
        out = self.dir / "out.docx"
        result = fill_template(self.original, info["fields"], answers, out)
        self.assertEqual(result["filled"], 5)
        body = part_text(out)
        self.assertIn("Dr Example", body)
        self.assertIn("03/14/2026", body)
        self.assertIn("Left", body)
        self.assertIn("Synthetic Person", part_text(out, "word/header1.xml"))
        self.assertIn("☐ Yes", body)             # the unchosen option stays printed
        self.assertIn("☒", body)                 # the chosen box is ticked
        with ZipFile(self.original) as a, ZipFile(out) as b:
            changed = {"word/document.xml", "word/header1.xml"}
            self.assertEqual([i.filename for i in a.infolist()], [i.filename for i in b.infolist()])
            for name in a.namelist():
                if name not in changed:
                    self.assertEqual(a.read(name), b.read(name), name)

    def test_signature_is_never_filled_and_a_moved_anchor_is_stale(self):
        info = inspect_template(self.original, engine="v3")
        by = fields_by_label(info["fields"])
        with self.assertRaises(DocumentError) as sig:
            fill_template(self.original, info["fields"], {by["Clinician signature"]["key"]: "Dr X"}, self.dir / "s.docx")
        self.assertEqual(sig.exception.code, "invalid_answer")
        tampered = deepcopy(info["fields"])
        for f in tampered:
            if f["label"] == "Referring physician":
                f["locator"]["fingerprint"] = "0" * 64
        with self.assertRaises(DocumentError) as stale:
            fill_template(self.original, tampered, {by["Referring physician"]["key"]: "x"}, self.dir / "t.docx")
        self.assertEqual(stale.exception.code, "stale_mapping")

    def test_an_option_not_printed_is_refused(self):
        info = inspect_template(self.original, engine="v3")
        by = fields_by_label(info["fields"])
        with self.assertRaises(DocumentError) as bad:
            fill_template(self.original, info["fields"], {by["Smoker?"]["key"]: "Former"}, self.dir / "o.docx")
        self.assertEqual(bad.exception.code, "invalid_answer")

    def test_a_control_wrapping_other_controls_is_read_by_its_parts_and_never_overwritten(self):
        wrapped = self.dir / "wrapped.docx"
        build_form(wrapped, nested_wrapper=True)
        info = inspect_template(wrapped, engine="v3")
        side = [f for f in info["fields"] if f["label"] == "Side"]
        self.assertEqual(len(side), 1)
        self.assertEqual(side[0]["options"], ["Left", "Right"])
        out = self.dir / "wrapped-out.docx"
        fill_template(wrapped, info["fields"], {side[0]["key"]: "Right"}, out)
        with ZipFile(wrapped) as a, ZipFile(out) as b:
            count = lambda z: etree.fromstring(z.read("word/document.xml")).findall(".//{%s}sdt" % W)
            self.assertEqual(len(count(a)), len(count(b)))

    def test_manual_field_is_located_recorded_as_manual_and_filled(self):
        outline = outline_template(self.original)
        slot = next(r for r in outline["slots"] if r["kind"] == "p" and r["preview"] == "Synthetic clinic intake")
        added = locate_fields(self.original, [{"slot": slot["slot"], "kind": "text", "label": "Clinic"}])["fields"]
        self.assertEqual(added[0]["origin"], "manual")
        out = self.dir / "manual.docx"
        fill_template(self.original, added, {added[0]["key"]: "North"}, out)
        self.assertIn("Synthetic clinic intake North", part_text(out))
        with self.assertRaises(DocumentError):
            cell = next(r for r in outline["slots"] if r["kind"] == "tc" and r["preview"] == "Allergies")
            locate_fields(self.original, [{"slot": cell["slot"], "kind": "text", "label": "Overwrite"}])

    def test_legacy_engine_is_unchanged_for_new_inspections_when_selected(self):
        try:
            info = inspect_template(self.original, engine="legacy")
        except DocumentError:
            return                               # the legacy contract refuses content controls, as before
        self.assertNotEqual(info.get("contract"), "v3")


def docx_bytes(doc):
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def knee_form(path):
    """F1: dropdown options that differ only by an abbreviation a list-marker strip would remove."""
    doc = Document()
    doc.add_paragraph("Synthetic knee injection note")
    p = doc.add_paragraph("Injection site: ")
    p._p.append(_sdt("dropdown", options=["Choose an item.", "Rt. knee", "Lt. knee", "Bil. knees"]))
    Path(path).write_bytes(docx_bytes(doc))


class OptionReader(unittest.TestCase):
    """F1: the engine chooses an option exactly as the review reads it: the stored label, trimmed, case only."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_the_reviewed_option_is_the_exported_option(self):
        knee_form(self.dir / "knee.docx")
        info = inspect_template(self.dir / "knee.docx", engine="v3")
        field = next(f for f in info["fields"] if f["kind"] == "choice")
        self.assertEqual(field["options"], ["Rt. knee", "Lt. knee", "Bil. knees"])
        for answer, shown in (("Lt. knee", "Injection site: Lt. knee"), (" lt. KNEE ", "Injection site: Lt. knee"),
                              ("Bil. knees", "Injection site: Bil. knees")):
            out = self.dir / ("out-%d.docx" % len(list(self.dir.iterdir())))
            fill_template(self.dir / "knee.docx", info["fields"], {field["key"]: answer}, out)
            self.assertIn(shown, part_text(out), answer)
        with self.assertRaises(DocumentError) as refused:
            fill_template(self.dir / "knee.docx", info["fields"], {field["key"]: "knee"}, self.dir / "never.docx")
        self.assertEqual(refused.exception.code, "invalid_answer")

    def test_options_that_differ_only_in_punctuation_stay_distinct(self):
        doc = Document()
        doc.add_paragraph("Synthetic follow-up")
        p = doc.add_paragraph("Imaging reviewed: ")
        p._p.append(_sdt("dropdown", options=["Choose an item.", "Yes", "Yes*", "No"]))
        (self.dir / "star.docx").write_bytes(docx_bytes(doc))
        info = inspect_template(self.dir / "star.docx", engine="v3")
        field = next(f for f in info["fields"] if f["kind"] == "choice")
        self.assertEqual(field["options"], ["Yes", "Yes*", "No"])
        fill_template(self.dir / "star.docx", info["fields"], {field["key"]: "Yes*"}, self.dir / "star-out.docx")
        self.assertIn("Imaging reviewed: Yes*", part_text(self.dir / "star-out.docx"))

    def test_list_markers_are_digits_single_letters_and_roman_numerals_only(self):
        from documents.word_v3 import clean_label
        kept = ["Rt. knee effusion", "Lt. knee", "Bil. knees", "Dr. Example", "No. of lesions", "Mr. Synthetic",
                "Hx. reviewed", "Pt. consent", "(c) Route"]     # R4: a single letter in parentheses is kept (fail closed)
        for text in kept:
            self.assertEqual(clean_label(text), text)
        stripped = {"1. Name": "Name", "12) Plan": "Plan", "a) Yes": "Yes", "B. No": "No", "(3) Dose": "Dose",
                    "c) Route": "Route", "iv. Findings": "Findings", "(xii) Impression": "Impression",
                    "XXX. Item": "Item", "vii) Item": "Item"}
        for text, label in stripped.items():
            self.assertEqual(clean_label(text), label, text)

    def test_laterality_survives_in_detected_and_typed_labels(self):
        doc = Document()
        doc.add_paragraph("Synthetic joint exam")
        doc.add_paragraph("Rt. knee effusion: ____________")
        doc.add_paragraph("Lt. knee effusion: ____________")
        doc.add_paragraph("")
        (self.dir / "joint.docx").write_bytes(docx_bytes(doc))
        info = inspect_template(self.dir / "joint.docx", engine="v3")
        self.assertEqual([f["label"] for f in info["fields"]], ["Rt. knee effusion", "Lt. knee effusion"])
        added = locate_fields(self.dir / "joint.docx", [{"slot": "word/document.xml#b/p[3]", "kind": "text",
                                                          "label": "Rt. knee aspirate volume"}])["fields"]
        self.assertEqual(added[0]["label"], "Rt. knee aspirate volume")



def repeated_options_form(path):
    """F7: one question printing the same option twice (boxes, a dropdown, printed words)."""
    doc = Document()
    doc.add_paragraph("Synthetic swelling check")
    doc.add_paragraph("Swelling:  ☐ Left  ☐ Right  ☐ Left")
    p = doc.add_paragraph("Answer given: ")
    p._p.append(_sdt("dropdown", options=["Choose an item.", "Yes", "No", "Yes"]))
    doc.add_paragraph("Was the dressing changed?")
    doc.add_paragraph("YES\t\tNO\t\tYES")
    Path(path).write_bytes(docx_bytes(doc))


class RepeatedOptions(unittest.TestCase):
    """F7: every printed option can be chosen, even one printed twice; fill maps by position."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.form = self.dir / "repeated.docx"
        repeated_options_form(self.form)
        self.fields = inspect_template(self.form, engine="v3")["fields"]
        self.by = fields_by_label(self.fields)

    def tearDown(self):
        self.tmp.cleanup()

    def test_repeated_options_get_print_order_suffixes(self):
        self.assertEqual(self.by["Swelling"]["options"], ["Left", "Right", "Left (2)"])
        self.assertEqual(self.by["Answer given"]["options"], ["Yes", "No", "Yes (2)"])
        self.assertEqual(self.by["Was the dressing changed?"]["options"], ["YES", "NO", "YES (2)"])
        self.assertEqual(self.by["Swelling"]["locator"]["target"]["printed"], ["Left", "Right", "Left"])

    def test_the_second_printing_is_the_one_ticked(self):
        out = self.dir / "second.docx"
        fill_template(self.form, self.fields, {self.by["Swelling"]["key"]: "Left (2)",
                                               self.by["Answer given"]["key"]: "Yes (2)",
                                               self.by["Was the dressing changed?"]["key"]: "YES (2)"}, out)
        body = part_text(out)
        self.assertIn("Swelling:  ☐ Left  ☐ Right  ☒ Left", body)
        self.assertIn("Answer given: Yes", body)
        self.assertNotIn("Yes (2)", body, "the printed option is written, never the suffix")
        self.assertIn("YESNO☒ YES", body, "tabs are not text: the words read YES, NO, YES")

    def test_a_legacy_dropdown_selects_the_entry_at_its_printing(self):
        doc = Document()
        doc.add_paragraph("Synthetic legacy form")
        ff_field(doc.add_paragraph("Answer given: "), "dropdown", ["Yes", "No", "Yes"])
        form = self.dir / "legacy-dd.docx"
        form.write_bytes(docx_bytes(doc))
        fields = inspect_template(form, engine="v3")["fields"]
        self.assertEqual(fields[0]["options"], ["Yes", "No", "Yes (2)"])
        for answer, index in (("Yes (2)", "2"), ("Yes", "0"), ("No", "1")):
            out = self.dir / ("dd-%s.docx" % index)
            fill_template(form, fields, {fields[0]["key"]: answer}, out)
            self.assertEqual(ff_result(out), [index], answer)

    def test_the_first_printing_stays_the_first(self):
        out = self.dir / "first.docx"
        fill_template(self.form, self.fields, {self.by["Swelling"]["key"]: "Left",
                                               self.by["Was the dressing changed?"]["key"]: "YES"}, out)
        body = part_text(out)
        self.assertIn("Swelling:  ☒ Left  ☐ Right  ☐ Left", body)
        self.assertIn("☒ YESNOYES", body)
        self.assertEqual(body.count("☒"), 2)



def ticked_sdt_checkbox():
    sdt = _sdt("checkbox")
    sdt.find(qn("w:sdtPr")).find("{%s}checkbox" % W14).find("{%s}checked" % W14).set("{%s}val" % W14, "1")
    sdt.find(qn("w:sdtContent")).find(qn("w:r")).find(qn("w:t")).text = "☒"
    return sdt


def prechecked_form(path):
    """F3: a blank form printing some boxes ticked - a checkbox control, a glyph, a legacy checkbox, a choice, a
    check-all-that-apply line, a radio pair, and a box grid with one box ticked."""
    doc = Document()
    doc.add_paragraph("Synthetic procedure checklist")
    doc.add_paragraph("Consent obtained ")._p.append(ticked_sdt_checkbox())
    doc.add_paragraph("History reviewed today.")
    doc.add_paragraph("Hypertension ☒")
    doc.add_paragraph("Medication list reviewed.")
    doc.add_paragraph("Smoker?  ☐ Yes  ☒ No")
    doc.add_paragraph("Symptoms (check all that apply): ☒ Pain ☐ Swelling ☒ Redness")
    ff_field(doc.add_paragraph("Band applied "), "checkbox", checked=True, name="Band")
    doc.add_paragraph("Medication list reviewed again.")
    doc.add_paragraph("Laterality:  ○ Left  ◉ Right")
    grid = doc.add_table(rows=3, cols=3)
    for c, h in enumerate(["Symptom", "Yes", "No"]):
        grid.cell(0, c).text = h
    grid.cell(1, 0).text = "Fever"
    grid.cell(2, 0).text = "Cough"
    for r, c, mark in ((1, 1, "☐"), (1, 2, "☒"), (2, 1, "☐"), (2, 2, "☐")):
        grid.cell(r, c).text = mark
    Path(path).write_bytes(docx_bytes(doc))


def paragraphs(path, part="word/document.xml"):
    with ZipFile(path) as z:
        root = etree.fromstring(z.read(part))
    return ["".join(t.text or "" for t in p.iter("{%s}t" % W)) for p in root.iter("{%s}p" % W)]


def control_states(path):
    with ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    sdt = [c.get("{%s}val" % W14) for c in root.iter("{%s}checked" % W14)]
    ff = [c.get("{%s}val" % W) for c in root.iter("{%s}checked" % W)]
    return sdt, ff


class Unticking(unittest.TestCase):
    """F3: filling sets every box of an answered field - chosen ticked, every other unticked; unanswered stays."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.form = self.dir / "prechecked.docx"
        prechecked_form(self.form)
        self.fields = inspect_template(self.form, engine="v3")["fields"]
        self.by = fields_by_label(self.fields)

    def tearDown(self):
        self.tmp.cleanup()

    def fill(self, answers, name="out.docx"):
        out = self.dir / name
        fill_template(self.form, self.fields, {self.by[label]["key"]: value for label, value in answers.items()}, out)
        return out

    def test_the_blank_forms_printed_ticks_are_recorded(self):
        printed = {label: f.get("printed_ticked") for label, f in self.by.items()}
        self.assertEqual(printed["Consent obtained"], [0])
        self.assertEqual(printed["Hypertension"], [0])
        self.assertEqual(printed["Smoker?"], [1])
        self.assertEqual(printed["Symptoms (check all that apply)"], [0, 2])
        self.assertEqual(printed["Band applied"], [0])
        self.assertEqual(printed["Laterality"], [1])
        self.assertEqual(printed["Fever"], [1])
        self.assertIsNone(printed["Cough"], "nothing printed ticked, nothing recorded")

    def test_a_checkbox_answered_unticked_is_exported_unticked(self):
        out = self.fill({"Consent obtained": False, "Hypertension": False, "Band applied": False})
        text = paragraphs(out)
        self.assertIn("Consent obtained ☐", text)
        self.assertIn("Hypertension ☐", text)
        self.assertEqual(control_states(out), (["0"], ["0"]))

    def test_a_choice_ticks_the_chosen_and_unticks_the_others(self):
        out = self.fill({"Smoker?": "Yes", "Symptoms (check all that apply)": ["Swelling"], "Laterality": "Left",
                         "Fever": "Yes"})
        text = paragraphs(out)
        self.assertIn("Smoker?  ☒ Yes  ☐ No", text)
        self.assertIn("Symptoms (check all that apply): ☐ Pain ☒ Swelling ☐ Redness", text)
        self.assertIn("Laterality:  ◉ Left  ○ Right", text)
        row = text[text.index("Fever"):text.index("Fever") + 3]
        self.assertEqual(row, ["Fever", "☒", "☐"])

    def test_an_unanswered_field_stays_as_printed(self):
        out = self.fill({"Hypertension": True})
        text = paragraphs(out)
        self.assertIn("Smoker?  ☐ Yes  ☒ No", text)
        self.assertIn("Symptoms (check all that apply): ☒ Pain ☐ Swelling ☒ Redness", text)
        self.assertIn("Consent obtained ☒", text)
        self.assertEqual(control_states(out), (["1"], []))

    def test_a_written_cell_mark_is_removed_when_the_other_option_is_chosen(self):
        from documents import word_v3
        doc = Document()
        grid = doc.add_table(rows=1, cols=3)
        grid.cell(0, 0).text = "Fever"
        grid.cell(0, 2).text = "X"
        yes, no = grid.cell(0, 1)._tc, grid.cell(0, 2)._tc
        self.assertEqual(word_v3.text_of(no).strip(), "X")
        self.assertTrue(word_v3.set_cell_box(yes, True))
        self.assertTrue(word_v3.set_cell_box(no, False))
        self.assertEqual((word_v3.text_of(yes).strip(), word_v3.text_of(no).strip()), ("X", ""))
        self.assertFalse(word_v3.set_cell_box(no, False), "an empty cell has nothing to untick")


class ExportReadBack(unittest.TestCase):
    """F3 belt and braces: an export whose ticked boxes are not exactly the chosen ones is refused."""

    def test_a_writer_that_leaves_a_box_as_printed_refuses_the_export(self):
        from unittest.mock import patch
        from documents import word_v3
        with tempfile.TemporaryDirectory() as tmp:
            form = Path(tmp) / "prechecked.docx"
            prechecked_form(form)
            data = form.read_bytes()
            fields = word_v3.inspect(data)["fields"]
            smoker = fields_by_label(fields)["Smoker?"]
            self.assertIn("☒ Yes", "".join(paragraphs_bytes(word_v3.fill(data, fields, {smoker["key"]: "Yes"}))))
            real = word_v3.set_glyph
            with patch.object(word_v3, "set_glyph", side_effect=lambda p, n, ticked: real(p, n, ticked) if ticked else False):
                with self.assertRaises(word_v3.V3Failure) as refused:
                    word_v3.fill(data, fields, {smoker["key"]: "Yes"})
            self.assertEqual(refused.exception.code, "fill_mismatch")
            self.assertIn('"Smoker?"', str(refused.exception))


def paragraphs_bytes(data):
    root = etree.fromstring(ZipFile(io.BytesIO(data)).read("word/document.xml"))
    return ["".join(t.text or "" for t in p.iter("{%s}t" % W)) for p in root.iter("{%s}p" % W)]


class BoxNumbers(unittest.TestCase):
    """F2/F6: a manually added checkbox is the box the clinician named, and a box the place lacks is refused."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        doc = Document()
        doc.add_paragraph("Synthetic smoking status")
        doc.add_paragraph("Smoker?  ☐ Yes  ☐ No")
        doc.add_paragraph("Diabetes ")._p.append(_sdt("checkbox"))
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Severity"
        table.cell(0, 1).text = "☐ Mild ☐ Severe"
        self.form = self.dir / "boxes.docx"
        self.form.write_bytes(docx_bytes(doc))
        self.places = {row["slot"]: row for row in outline_template(self.form)["slots"]}

    def tearDown(self):
        self.tmp.cleanup()

    def add(self, slot, box, label="Added box"):
        return locate_fields(self.form, [{"slot": slot, "kind": "checkbox", "label": label, "glyph": box - 1}])["fields"]

    def test_the_outline_counts_each_places_boxes(self):
        self.assertEqual(self.places["word/document.xml#b/p[1]"]["boxes"], 2)
        self.assertEqual(self.places["word/document.xml#b/p[1]/glyphs"]["boxes"], 2)
        self.assertEqual(self.places["word/document.xml#b/p[2]/sdt[0]"]["boxes"], 1)
        self.assertEqual(self.places["word/document.xml#b/tbl[0]/tr[0]/tc[1]"]["boxes"], 2)
        self.assertEqual(self.places["word/document.xml#b/p[0]"]["boxes"], 0)

    def test_a_paragraph_place_takes_the_named_box(self):
        added = self.add("word/document.xml#b/p[1]", 2, "Non-smoker")
        self.assertEqual(added[0]["locator"]["target"]["boxes"][0]["glyph"], 1)
        out = self.dir / "second-box.docx"
        fill_template(self.form, added, {added[0]["key"]: True}, out)
        self.assertIn("Smoker?  ☐ Yes  ☒ No", part_text(out))

    def test_a_cell_place_takes_the_named_box(self):
        try:
            added = self.add("word/document.xml#b/tbl[0]/tr[0]/tc[1]", 2, "Severe")
        except DocumentError as refused:
            self.fail("a cell holding printed boxes was refused as a place: %s" % refused)
        out = self.dir / "cell-box.docx"
        fill_template(self.form, added, {added[0]["key"]: True}, out)
        self.assertIn("☐ Mild ☒ Severe", part_text(out))

    def test_a_box_the_place_does_not_have_is_refused(self):
        for slot, box, count in (("word/document.xml#b/p[1]", 3, "2 printed boxes"),
                                 ("word/document.xml#b/p[1]/glyphs", 7, "2 printed boxes"),
                                 ("word/document.xml#b/p[2]/sdt[0]", 2, "1 printed box;"),
                                 ("word/document.xml#b/tbl[0]/tr[0]/tc[1]", 3, "2 printed boxes"),
                                 ("word/document.xml#b/p[1]", 0, "2 printed boxes")):
            with self.assertRaises(DocumentError) as refused:
                self.add(slot, box)
            self.assertEqual(refused.exception.code, "invalid_mapping", slot)
            self.assertIn(count, str(refused.exception), slot)


def laterality_form(path):
    """Options that differ only by a side marker, on every detector path that builds options: a box line with
    '(L) knee' / '(R) knee', one with 'L knee' / 'R knee', a dropdown, an option grid's column headers, a column
    checklist's row labels, and narrow tick cells beside printed options."""
    from docx.shared import Twips
    doc = Document()
    doc.add_paragraph("Synthetic joint procedure")
    doc.add_paragraph("Injected joint:  ☐ (L) knee  ☐ (R) knee")
    doc.add_paragraph("Aspirated joint:  ☐ L knee  ☐ R knee")
    doc.add_paragraph("Imaged joint: ")._p.append(_sdt("dropdown", options=["Choose an item.", "(L) knee", "(R) knee"]))
    grid = doc.add_table(rows=4, cols=3)
    for c, h in enumerate(["Finding", "(L) knee", "(R) knee"]):
        grid.cell(0, c).text = h
    for r, name in enumerate(["Effusion", "Warmth", "Crepitus"], 1):
        grid.cell(r, 0).text = name
    doc.add_paragraph("Dressing site")
    ticks = doc.add_table(rows=2, cols=2)
    for r, name in enumerate(["(L) knee", "(R) knee"]):
        ticks.cell(r, 0).width = Twips(300)
        ticks.cell(r, 1).text = name
    # Several answers need the printed command (the structural selection rule): the checklist's own instruction,
    # printed under its own heading (an instruction printed between two questions belongs to neither, R4-F5).
    doc.add_heading("Examination", 2)
    doc.add_paragraph("Joints examined - tick all that apply")
    check = doc.add_table(rows=4, cols=2)
    check.cell(0, 0).text = "Joint"
    check.cell(0, 1).text = "Examined"
    for r, name in enumerate(["(L) knee", "(R) knee", "(L) hip"], 1):
        check.cell(r, 0).text = name
        check.cell(r, 1).text = "☐"
    Path(path).write_bytes(docx_bytes(doc))


class OptionsKeepTheirPrintedText(unittest.TestCase):
    """An option label is the printed option text, trimmed only; the list-marker strip is for question labels."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.form = self.dir / "laterality.docx"
        laterality_form(self.form)
        self.fields = inspect_template(self.form, engine="v3")["fields"]
        self.by = fields_by_label(self.fields)

    def tearDown(self):
        self.tmp.cleanup()

    def test_every_option_path_keeps_the_side(self):
        sides = ["(L) knee", "(R) knee"]
        self.assertEqual(self.by["Injected joint"]["options"], sides)                  # box line
        self.assertEqual(self.by["Aspirated joint"]["options"], ["L knee", "R knee"])  # box line, bare letters
        self.assertEqual(self.by["Imaged joint"]["options"], sides)                    # dropdown
        for finding in ("Effusion", "Warmth", "Crepitus"):                             # option grid headers
            self.assertIn(finding, self.by, "each grid row is one choice between the two sides")
            self.assertEqual((self.by[finding]["kind"], self.by[finding]["options"]), ("choice", sides), finding)
        self.assertEqual(self.by["Examined"]["options"], sides + ["(L) hip"])          # column checklist rows
        ticks = [f for f in self.fields if f["label"] == "Choose"]                     # narrow tick cells
        self.assertEqual([f["options"] for f in ticks], [sides])
        for f in self.fields:
            self.assertNotIn("knee (2)", f.get("options") or [], f["label"])

    def test_each_side_is_exported_in_its_own_box(self):
        out = self.dir / "sides.docx"
        answers = {"Injected joint": "(R) knee", "Aspirated joint": "L knee", "Imaged joint": "(L) knee",
                   "Effusion": "(R) knee", "Examined": ["(R) knee", "(L) hip"]}
        fill_template(self.form, self.fields, {self.by[k]["key"]: v for k, v in answers.items()}, out)
        text = [line for line in paragraphs(out)]
        self.assertIn("Injected joint:  ☐ (L) knee  ☒ (R) knee", text)
        self.assertIn("Aspirated joint:  ☒ L knee  ☐ R knee", text)
        self.assertIn("Imaged joint: (L) knee", text)
        at = text.index("Effusion")
        self.assertEqual(text[at:at + 3], ["Effusion", "", "X"])
        at = text.index("Joint")
        self.assertEqual(text[at + 2:at + 8], ["(L) knee", "☐", "(R) knee", "☒", "(L) hip", "☒"])

    def test_the_other_option_paths_keep_the_printed_text(self):
        doc = Document()
        doc.add_paragraph("Synthetic joint log")
        site = doc.add_table(rows=2, cols=2)                         # a slash header over an answer cell
        site.cell(0, 0).text = "Joint"
        site.cell(0, 1).text = "L. / R."
        site.cell(1, 0).text = "Knee"
        doc.add_paragraph("Doses given")
        log = doc.add_table(rows=3, cols=3)                          # a log table's column headers
        for c, h in enumerate(["Drug", "(L) dose", "(R) dose"]):
            log.cell(0, c).text = h
        ff_field(doc.add_paragraph("Splinted side: "), "dropdown", ["(L) knee ", " (R) knee"])
        form = self.dir / "paths.docx"
        form.write_bytes(docx_bytes(doc))
        fields = inspect_template(form, engine="v3")["fields"]
        by = fields_by_label(fields)
        self.assertIn("Knee", by, "'L. / R.' over the answer cell makes it a choice")
        self.assertEqual((by["Knee"]["kind"], by["Knee"]["options"]), ("choice", ["L.", "R."]))
        self.assertEqual(next(f for f in fields if f["kind"] == "table_rows")["options"], ["Drug", "(L) dose", "(R) dose"])
        self.assertEqual(by["Splinted side"]["options"], ["(L) knee", "(R) knee"])
        out = self.dir / "paths-out.docx"
        fill_template(form, fields, {by["Knee"]["key"]: "R.", by["Splinted side"]["key"]: "(R) knee"}, out)
        self.assertEqual(ff_result(out), ["1"])
        text = paragraphs(out)
        self.assertEqual(text[text.index("Knee"):text.index("Knee") + 2], ["Knee", "R."])
        outline = outline_template(form)["slots"]
        place = next(r for r in outline if r["kind"] == "p" and r["preview"] == "Synthetic joint log")
        added = locate_fields(form, [{"slot": place["slot"], "kind": "choice", "label": "Side",
                                      "options": [" (L) knee", "(R)  knee "]}])["fields"]
        self.assertEqual(added[0]["options"], ["(L) knee", "(R) knee"])

    def test_option_cells_and_lettered_items_keep_the_side(self):
        doc = Document()
        doc.add_paragraph("Synthetic procedure")
        row = doc.add_table(rows=1, cols=3)                          # one box and its option per cell
        row.cell(0, 0).text = "Injected side"
        row.cell(0, 1).text = "☐ (L) knee"
        row.cell(0, 2).text = "☐ (R) knee"
        doc.add_paragraph("Which knee was examined?")                # lettered list items under a question
        numbering = doc.part.numbering_part.element
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), "90")
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        for tag, val in (("w:start", "1"), ("w:numFmt", "lowerLetter"), ("w:lvlText", "%1)")):
            x = OxmlElement(tag)
            x.set(qn("w:val"), val)
            lvl.append(x)
        abstract.append(lvl)
        numbering.insert(0, abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), "90")
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), "90")
        num.append(ref)
        numbering.append(num)
        for side in ("(L) knee", "(R) knee"):
            ppr = doc.add_paragraph(side)._p.get_or_add_pPr()
            numpr = OxmlElement("w:numPr")
            for tag, val in (("w:ilvl", "0"), ("w:numId", "90")):
                x = OxmlElement(tag)
                x.set(qn("w:val"), val)
                numpr.append(x)
            ppr.append(numpr)
        form = self.dir / "cells-and-items.docx"
        form.write_bytes(docx_bytes(doc))
        fields = inspect_template(form, engine="v3")["fields"]
        by = fields_by_label(fields)
        self.assertEqual(by["Injected side"]["options"], ["(L) knee", "(R) knee"])
        self.assertEqual((by["Which knee was examined?"]["locator"]["target"]["type"], by["Which knee was examined?"]["options"]),
                         ("paragraph_options", ["(L) knee", "(R) knee"]))
        out = self.dir / "cells-and-items-out.docx"
        fill_template(form, fields, {by["Injected side"]["key"]: "(R) knee", by["Which knee was examined?"]["key"]: "(L) knee"}, out)
        text = paragraphs(out)
        self.assertIn("☐ (L) knee", text)
        self.assertIn("☒ (R) knee", text)
        self.assertIn("☒ (L) knee", text)

    def test_question_labels_keep_the_marker_rule_and_options_do_not(self):
        from documents.word_v3 import clean_label, option_text
        self.assertEqual(clean_label("a) Pain"), "Pain")
        self.assertEqual(option_text("a) Pain"), "a) Pain")
        self.assertEqual(option_text("(a) Pain"), "(a) Pain")
        self.assertEqual(option_text("  (L)\tknee  "), "(L) knee")
        self.assertEqual(option_text("Other: ______"), "Other", "the write-in blank is not the option")
        self.assertEqual(option_text("b) Left knee -"), "b) Left knee")


def checkbox_control(checked=False, *, states=True, flag=True):
    """A checkbox content control as Word writes it: the w14:checked flag, the checkedState/uncheckedState glyphs
    (MS Gothic U+2612 / U+2610) and the glyph the state names as its content."""
    sdt = _sdt("checkbox")
    cb = sdt.find(qn("w:sdtPr")).find("{%s}checkbox" % W14)
    cb.find("{%s}checked" % W14).set("{%s}val" % W14, "1" if checked else "0")
    if not flag:
        cb.remove(cb.find("{%s}checked" % W14))
    if states:
        for tag, code in (("checkedState", "2612"), ("uncheckedState", "2610")):
            state = etree.SubElement(cb, "{%s}%s" % (W14, tag))
            state.set("{%s}val" % W14, code)
            state.set("{%s}font" % W14, "MS Gothic")
    sdt.find(qn("w:sdtContent")).find(qn("w:r")).find(qn("w:t")).text = "☒" if checked else "☐"
    return sdt


def text_run(text):
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    return r


def controls_shown(data):
    """[(glyph shown, w14:checked flag)] for every checkbox content control of an exported package, in order."""
    root = etree.fromstring(ZipFile(io.BytesIO(data)).read("word/document.xml"))
    out = []
    for sdt in root.iter(qn("w:sdt")):
        cb = sdt.find(".//{%s}checkbox" % W14)
        if cb is None:
            continue
        flag = cb.find("{%s}checked" % W14)
        out.append(("".join(t.text or "" for t in sdt.find(qn("w:sdtContent")).iter(qn("w:t"))),
                    flag.get("{%s}val" % W14) if flag is not None else None))
    return out


class ControlReadBack(unittest.TestCase):
    """R3: a checkbox control reads as ticked or unticked only when its flag AND the glyph it displays agree."""

    def anticoagulated(self):
        doc = Document()
        doc.add_paragraph("Synthetic intake")
        p = doc.add_paragraph("Anticoagulated? ")
        p._p.append(checkbox_control(checked=True))
        p._p.append(text_run(" Yes  "))
        p._p.append(checkbox_control())
        p._p.append(text_run(" No"))
        return docx_bytes(doc)

    def test_a_writer_that_leaves_the_ticked_glyph_is_refused(self):
        from unittest.mock import patch
        from documents import word_v3
        data = self.anticoagulated()
        fields = word_v3.inspect(data)["fields"]
        self.assertEqual((fields[0]["label"], fields[0]["options"]), ("Anticoagulated?", ["Yes", "No"]))
        real = word_v3.set_sdt_checkbox

        def regressed(sdt, ticked):             # the refuter's writer: the flag set, the printed glyph left
            if ticked:
                return real(sdt, True)
            cb = sdt.find(word_v3.q("w:sdtPr")).find(word_v3.q("w14:checkbox"))
            cb.find(word_v3.q("w14:checked")).set(word_v3.q("w14:val"), "0")
            return True
        with patch.object(word_v3, "set_sdt_checkbox", side_effect=regressed):
            with self.assertRaises(word_v3.V3Failure) as refused:
                word_v3.fill(data, fields, {fields[0]["key"]: "No"})
        self.assertEqual(refused.exception.code, "fill_mismatch")

    def test_unticking_a_control_shows_its_unchecked_glyph(self):
        from documents import word_v3
        data = self.anticoagulated()
        fields = word_v3.inspect(data)["fields"]
        self.assertEqual(fields[0].get("printed_ticked"), [0])
        out = word_v3.fill(data, fields, {fields[0]["key"]: "No"})
        self.assertEqual(controls_shown(out), [("☐", "0"), ("☒", "1")])

    def test_an_added_flag_keeps_the_schema_order(self):
        """R5: w14:checked precedes w14:checkedState and w14:uncheckedState (CT_SdtCheckbox is a sequence)."""
        from documents import word_v3
        doc = Document()
        doc.add_paragraph("Synthetic intake")
        p = doc.add_paragraph("Smoker? ")
        p._p.append(checkbox_control(flag=False))      # an unticked control written without its optional flag
        p._p.append(text_run(" Yes  "))
        p._p.append(checkbox_control(flag=False))
        p._p.append(text_run(" No"))
        data = docx_bytes(doc)
        fields = word_v3.inspect(data)["fields"]
        out = word_v3.fill(data, fields, {fields[0]["key"]: "Yes"})
        root = etree.fromstring(ZipFile(io.BytesIO(out)).read("word/document.xml"))
        orders = [[etree.QName(c).localname for c in cb] for cb in root.iter("{%s}checkbox" % W14)]
        self.assertEqual(orders, [["checked", "checkedState", "uncheckedState"]] * 2)
        self.assertEqual(controls_shown(out), [("☒", "1"), ("☐", "0")])


def control_checklist_bytes():
    """The refuter's p1c input: a check-all-that-apply column whose boxes are checkbox controls, Warfarin printed
    ticked."""
    doc = Document()
    doc.add_paragraph("Synthetic medication reconciliation")
    table = doc.add_table(rows=4, cols=2)
    table.cell(0, 0).text = "Medication"
    table.cell(0, 1).text = "Continue (check all that apply)"
    for r, name in ((1, "Warfarin"), (2, "Metformin"), (3, "Aspirin")):
        table.cell(r, 0).text = name
        table.cell(r, 1).paragraphs[0]._p.append(checkbox_control(checked=(r == 1)))
    return docx_bytes(doc)


def control_grid_bytes():
    """The refuter's p1 input: a Yes/No grid whose answer cells are checkbox controls, Chest pain printed Yes."""
    doc = Document()
    doc.add_paragraph("Synthetic review of systems")
    table = doc.add_table(rows=3, cols=3)
    for c, h in enumerate(["Symptom", "Yes", "No"]):
        table.cell(0, c).text = h
    table.cell(1, 0).text = "Chest pain"
    table.cell(2, 0).text = "Dyspnea"
    for r in (1, 2):
        for c in (1, 2):
            table.cell(r, c).paragraphs[0]._p.append(checkbox_control(checked=(r == 1 and c == 1)))
    return docx_bytes(doc)


class ControlsInBoxCells(unittest.TestCase):
    """R2: a table cell whose box is a checkbox control is read and written as that control."""

    def test_a_control_cell_printed_ticked_is_recorded_and_kept_when_unanswered(self):
        from documents import word_v3
        data = control_checklist_bytes()
        fields = word_v3.inspect(data)["fields"]
        field = next(f for f in fields if f["kind"] == "multi_choice")
        self.assertEqual(field["options"], ["Warfarin", "Metformin", "Aspirin"])
        self.assertEqual(field.get("printed_ticked"), [0], "Warfarin is printed ticked")
        self.assertEqual(controls_shown(word_v3.fill(data, fields, {})), [("☒", "1"), ("☐", "0"), ("☐", "0")])

    def test_an_answer_sets_every_control_of_the_checklist(self):
        from documents import word_v3
        data = control_checklist_bytes()
        fields = word_v3.inspect(data)["fields"]
        field = next(f for f in fields if f["kind"] == "multi_choice")
        out = word_v3.fill(data, fields, {field["key"]: ["Metformin"]})
        self.assertEqual(controls_shown(out), [("☐", "0"), ("☒", "1"), ("☐", "0")])
        out = word_v3.fill(data, fields, {field["key"]: ["Aspirin", "Warfarin"]})
        self.assertEqual(controls_shown(out), [("☒", "1"), ("☐", "0"), ("☒", "1")])

    def test_a_grid_of_controls_takes_the_other_option(self):
        from documents import word_v3
        data = control_grid_bytes()
        fields = word_v3.inspect(data)["fields"]
        chest = fields_by_label(fields)["Chest pain"]
        self.assertEqual((chest["kind"], chest["options"], chest.get("printed_ticked")), ("choice", ["Yes", "No"], [0]))
        out = word_v3.fill(data, fields, {chest["key"]: "No"})
        self.assertEqual(controls_shown(out), [("☐", "0"), ("☒", "1"), ("☐", "0"), ("☐", "0")])

    def test_a_cell_holding_a_control_beside_another_box_is_unread_and_refused(self):
        from documents import word_v3
        doc = Document()
        doc.add_paragraph("Synthetic review of systems")
        table = doc.add_table(rows=2, cols=3)
        for c, h in enumerate(["Symptom", "Yes", "No"]):
            table.cell(0, c).text = h
        table.cell(1, 0).text = "Cough"
        table.cell(1, 1).text = "☐"                   # a printed box, and a legacy checkbox printed ticked beside it
        ff_field(table.cell(1, 1).paragraphs[0], "checkbox", checked=True, name="CoughYes")
        table.cell(1, 2).text = "☐"
        data = docx_bytes(doc)
        fields = word_v3.inspect(data)["fields"]
        cough = fields_by_label(fields)["Cough"]
        self.assertEqual((cough["options"], cough.get("printed_ticked"), cough.get("printed_unread")),
                         (["Yes", "No"], None, [0]), "never read as unticked")
        with self.assertRaises(word_v3.V3Failure) as refused:
            word_v3.fill(data, fields, {cough["key"]: "No"})
        self.assertEqual(refused.exception.code, "unsupported_layout")
        self.assertIn("checkbox control beside another box", str(refused.exception))


def mixed_side_bytes(control="sdt"):
    """The refuter's p2 input: one printed line mixing a checkbox control between two printed glyph boxes."""
    doc = Document()
    doc.add_paragraph("Synthetic knee note")
    p = doc.add_paragraph("Side: ☐ Right ")
    if control == "sdt":
        p._p.append(checkbox_control())
    else:
        ff_field(p, "checkbox", name="Left")
    p._p.append(text_run(" Left ☐ Bilateral"))
    return docx_bytes(doc)


def mixed_cell_bytes():
    """D03's shape (DEV corpus, place b/tbl[0]/tr[17]/tc[0]/p[4]): printed boxes and two checkbox controls on one
    line in a table cell, and a cell mixing a printed box, a control and a printed box."""
    doc = Document()
    doc.add_paragraph("Synthetic social history")
    table = doc.add_table(rows=2, cols=1)
    p = table.cell(0, 0).paragraphs[0]
    p.add_run("☐ Alcohol beverages: ☐ Wine [")
    p._p.append(checkbox_control())
    p._p.append(text_run(" Red "))
    p._p.append(checkbox_control())
    p._p.append(text_run(" White ], ☐ Other:__________"))
    q = table.cell(1, 0).paragraphs[0]
    q.add_run("☐ Mild ")
    q._p.append(checkbox_control())
    q._p.append(text_run(" Moderate ☐ Severe"))
    return docx_bytes(doc)


def glyphs_shown(data, contains):
    """The printed text (glyphs, controls as their glyph) of the exported paragraph containing `contains`."""
    return next(line for line in paragraphs_bytes(data) if contains in line)


def side_questions_bytes():
    """The refuter's p6 input: questions whose only difference is a side marker, in table rows and in paragraphs."""
    doc = Document()
    doc.add_paragraph("Synthetic knee examination")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Finding"
    table.cell(0, 1).text = "Present?"
    table.cell(1, 0).text = "(L) Effusion"
    table.cell(1, 1).text = "☐ Yes ☐ No"
    table.cell(2, 0).text = "(R) Effusion"
    table.cell(2, 1).text = "☐ Yes ☐ No"
    doc.add_paragraph("L. Tenderness? ☐ Yes ☐ No")
    doc.add_paragraph("R. Tenderness? ☐ Yes ☐ No")
    return docx_bytes(doc)


class SideMarkersInQuestions(unittest.TestCase):
    """R4: a side written as a list marker - (L) (R) L. R. - stays in the question label; so does any single letter in
    parentheses (fail closed). '(L) Effusion' and '(R) Effusion' both became 'Effusion'."""

    def test_the_side_stays_in_every_question_label(self):
        from documents import word_v3
        fields = word_v3.inspect(side_questions_bytes())["fields"]
        self.assertEqual([(f["label"], f["kind"], f["options"]) for f in fields],
                         [("(L) Effusion", "choice", ["Yes", "No"]), ("(R) Effusion", "choice", ["Yes", "No"]),
                          ("L. Tenderness?", "choice", ["Yes", "No"]), ("R. Tenderness?", "choice", ["Yes", "No"])])

    def test_side_letters_and_single_letters_in_parentheses_are_kept(self):
        from documents.word_v3 import clean_label
        for text in ("(L) Effusion", "(R) Effusion", "(l) knee", "(r) knee", "L. Tenderness?", "R. Tenderness?",
                     "L) knee", "R) knee", "l. knee", "r) knee", "(a) Pain", "(B) Hip", "(i) Item", "(x) Item"):
            self.assertEqual(clean_label(text), text)
        for text, label in (("a) Yes", "Yes"), ("B. No", "No"), ("(3) Dose", "Dose"), ("12) Plan", "Plan"),
                            ("iv. Findings", "Findings"), ("(xii) Impression", "Impression"), ("(ii) Item", "Item")):
            self.assertEqual(clean_label(text), label, text)


class LegacyDropdownDefault(unittest.TestCase):
    """R6: a legacy dropdown with no w:result shows its w:default entry (ECMA-376 ddList), not the first."""

    def test_the_default_entry_is_the_one_printed(self):
        from documents import word_v3
        doc = Document()
        doc.add_paragraph("Synthetic injection note")
        p = doc.add_paragraph("Side: ")
        ff_field(p, "dropdown", options=["Left", "Right", "Bilateral"], name="Side")
        default = OxmlElement("w:default")
        default.set(qn("w:val"), "1")
        p._p.find(".//" + qn("w:ddList")).insert(0, default)
        data = docx_bytes(doc)
        fields = word_v3.inspect(data)["fields"]
        self.assertEqual((fields[0]["options"], fields[0].get("printed_ticked")), (["Left", "Right", "Bilateral"], [1]))
        for answer, shown in (("Left", "0"), ("Bilateral", "2")):
            with tempfile.TemporaryDirectory() as tmp:
                Path(tmp, "out.docx").write_bytes(word_v3.fill(data, fields, {fields[0]["key"]: answer}))
                self.assertEqual(ff_result(Path(tmp, "out.docx")), [shown], answer)


def ff_text_field(p, name="Consent"):
    """A legacy text form field (FORMTEXT) appended to python-docx paragraph `p`, with an empty result."""
    for kind, child in (("begin", None), ("instr", None), ("separate", None), ("end", None)):
        r = OxmlElement("w:r")
        if kind == "instr":
            node = OxmlElement("w:instrText")
            node.text = " FORMTEXT "
        else:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                data = OxmlElement("w:ffData")
                n = OxmlElement("w:name")
                n.set(qn("w:val"), name)
                data.append(n)
                data.append(OxmlElement("w:enabled"))
                data.append(OxmlElement("w:textInput"))
                node.append(data)
        r.append(node)
        p._p.append(r)


class CheckboxNeedsABox(unittest.TestCase):
    """R7: a manual checkbox on a text content control or a legacy text field is refused when it is added, instead
    of being accepted, confirmed, and refused at every export."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        doc = Document()
        doc.add_paragraph("Synthetic consent")
        doc.add_paragraph("Consent obtained: ")._p.append(_sdt("rich", placeholder="Click or tap here to enter text."))
        ff_text_field(doc.add_paragraph("Witness: "), name="Witness")
        self.form = self.dir / "consent.docx"
        self.form.write_bytes(docx_bytes(doc))
        self.places = {row["slot"]: row for row in outline_template(self.form)["slots"]}

    def tearDown(self):
        self.tmp.cleanup()

    def test_a_checkbox_on_a_text_control_or_text_field_is_refused(self):
        for slot in ("word/document.xml#b/p[1]/sdt[0]", "word/document.xml#b/p[2]/ff[0]"):
            self.assertIn(slot, self.places)
            with self.assertRaises(DocumentError) as refused:
                locate_fields(self.form, [{"slot": slot, "kind": "checkbox", "label": "Consent obtained", "glyph": 0}])
            self.assertEqual(refused.exception.code, "invalid_mapping", slot)
            self.assertIn("is not a checkbox", str(refused.exception), slot)

    def test_the_same_places_still_take_text(self):
        for n, slot in enumerate(("word/document.xml#b/p[1]/sdt[0]", "word/document.xml#b/p[2]/ff[0]")):
            field = locate_fields(self.form, [{"slot": slot, "kind": "text", "label": "Consent note"}])["fields"][0]
            out = self.dir / ("text-%d.docx" % n)
            fill_template(self.form, [field], {field["key"]: "Signed by synthetic patient"}, out)
            self.assertIn("Signed by synthetic patient", part_text(out))


class SurvivingMutantPins(unittest.TestCase):
    """Guards the round-2 mutation batch showed no test would miss: M01, M13 and M16."""

    def test_a_repeat_suffix_never_takes_a_printed_options_words(self):
        """M01: 'Left' printed again after a printed 'Left (2)' is 'Left (3)', or two options share one label and
        neither can be chosen."""
        from documents import word_v3
        self.assertEqual(word_v3.unique_options(["Left", "Left (2)", "Left"]), ["Left", "Left (2)", "Left (3)"])
        doc = Document()
        doc.add_paragraph("Synthetic swelling check")
        doc.add_paragraph("Swelling:  ☐ Left  ☐ Left (2)  ☐ Left")
        data = docx_bytes(doc)
        fields = word_v3.inspect(data)["fields"]
        self.assertEqual(fields[0]["options"], ["Left", "Left (2)", "Left (3)"])
        for answer, line in (("Left (3)", "Swelling:  ☐ Left  ☐ Left (2)  ☒ Left"),
                             ("Left (2)", "Swelling:  ☐ Left  ☒ Left (2)  ☐ Left")):
            self.assertIn(line, paragraphs_bytes(word_v3.fill(data, fields, {fields[0]["key"]: answer})), answer)

    def test_a_manual_checkbox_records_its_printed_tick(self):
        """M13: a box added by hand on a box the blank form prints ticked says so on the page."""
        with tempfile.TemporaryDirectory() as tmp:
            form = Path(tmp) / "prechecked.docx"
            prechecked_form(form)
            place = next(r for r in outline_template(form)["slots"] if r["kind"] == "p" and r["preview"] == "Hypertension ☒")
            added = locate_fields(form, [{"slot": place["slot"], "kind": "checkbox", "label": "Hypertension", "glyph": 0}])
            self.assertEqual(added["fields"][0].get("printed_ticked"), [0])

    def test_a_dropdown_or_option_cell_showing_other_words_is_refused(self):
        """M16: the read-back of a dropdown control and of an option-text cell compares the words shown."""
        from unittest.mock import patch
        from documents import word_v3
        doc = Document()
        doc.add_paragraph("Synthetic joint log")
        doc.add_paragraph("Laterality: ")._p.append(_sdt("dropdown", options=["Choose an item.", "Left", "Right"]))
        site = doc.add_table(rows=2, cols=2)
        site.cell(0, 0).text = "Joint"
        site.cell(0, 1).text = "L. / R."
        site.cell(1, 0).text = "Knee"
        data = docx_bytes(doc)
        by = fields_by_label(word_v3.inspect(data)["fields"])
        fields = [by["Laterality"], by["Knee"]]
        self.assertEqual((by["Knee"]["locator"]["target"].get("write"), by["Knee"]["options"]), ("option_text", ["L.", "R."]))
        self.assertIn("Laterality: Left", paragraphs_bytes(word_v3.fill(data, fields, {by["Laterality"]["key"]: "Left",
                                                                                          by["Knee"]["key"]: "L."})))
        real_sdt, real_cell = word_v3.write_sdt_text, word_v3.write_cell
        for name, wrong, answers in (
                ("write_sdt_text", lambda sdt, answer: real_sdt(sdt, "Right"), {by["Laterality"]["key"]: "Left"}),
                ("write_cell", lambda tc, answer, *a: real_cell(tc, "R.", *a), {by["Knee"]["key"]: "L."})):
            with patch.object(word_v3, name, side_effect=wrong):
                with self.assertRaises(word_v3.V3Failure) as refused:
                    word_v3.fill(data, fields, answers)
            self.assertEqual(refused.exception.code, "fill_mismatch", name)


def fever_cell_bytes():
    """A cell 'Yes' printed box and a legacy checkbox 'No': the cell's text shows the printed box only."""
    doc = Document()
    doc.add_paragraph("Synthetic review of systems")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Fever"
    cell = table.cell(0, 1).paragraphs[0]
    cell.add_run("☐ Yes ")
    ff_field(cell, "checkbox", name="FeverNo")
    cell._p.append(text_run(" No"))
    return docx_bytes(doc)


class MixedBoxNumbers(unittest.TestCase):
    """R1: the box number a clinician gives counts EXACTLY the boxes the place's preview shows, in its order - printed
    glyph boxes and checkbox controls alike - and the export ticks exactly that box, whatever its kind."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def form(self, data, name):
        path = self.dir / name
        path.write_bytes(data)
        return path, {row["slot"]: row for row in outline_template(path)["slots"]}

    def add_and_fill(self, path, slot, box, label):
        field = locate_fields(path, [{"slot": slot, "kind": "checkbox", "label": label, "glyph": box - 1}])["fields"][0]
        out = self.dir / ("%s-%d.docx" % (label.replace(" ", "-"), box))
        fill_template(path, [field], {field["key"]: True}, out)
        return field, out.read_bytes()

    def test_a_line_mixing_a_control_counts_every_box_it_shows(self):
        path, places = self.form(mixed_side_bytes(), "side.docx")
        for slot in ("word/document.xml#b/p[1]", "word/document.xml#b/p[1]/glyphs"):
            self.assertEqual(places[slot]["preview"], "Side: ☐ Right ☐ Left ☐ Bilateral")
            self.assertEqual(places[slot]["boxes"], 3, slot)
        _, out = self.add_and_fill(path, "word/document.xml#b/p[1]", 2, "Left")
        self.assertEqual(controls_shown(out), [("☒", "1")], "box 2 is the control printed before 'Left'")
        self.assertEqual(glyphs_shown(out, "Right"), "Side: ☐ Right ☒ Left ☐ Bilateral")
        _, out = self.add_and_fill(path, "word/document.xml#b/p[1]/glyphs", 3, "Bilateral")
        self.assertEqual((controls_shown(out), glyphs_shown(out, "Right")), ([("☐", "0")], "Side: ☐ Right ☐ Left ☒ Bilateral"))
        _, out = self.add_and_fill(path, "word/document.xml#b/p[1]", 1, "Right")
        self.assertEqual((controls_shown(out), glyphs_shown(out, "Right")), ([("☐", "0")], "Side: ☒ Right ☐ Left ☐ Bilateral"))
        with self.assertRaises(DocumentError) as refused:
            self.add_and_fill(path, "word/document.xml#b/p[1]", 4, "Beyond")
        self.assertIn("That place has 3 printed boxes", str(refused.exception))

    def test_a_legacy_checkbox_on_the_line_is_numbered_where_it_shows(self):
        path, places = self.form(mixed_side_bytes("ff"), "side-ff.docx")
        self.assertEqual((places["word/document.xml#b/p[1]"]["preview"], places["word/document.xml#b/p[1]"]["boxes"]),
                         ("Side: ☐ Right ☐ Left ☐ Bilateral", 3))
        _, out = self.add_and_fill(path, "word/document.xml#b/p[1]", 2, "Left")
        self.assertEqual(control_states(self.dir / "Left-2.docx"), ([], ["1"]))
        self.assertEqual(glyphs_shown(out, "Right"), "Side: ☐ Right  Left ☐ Bilateral")

    def test_d03s_red_wine_is_the_third_box_shown(self):
        path, places = self.form(mixed_cell_bytes(), "social.docx")
        line = "word/document.xml#b/tbl[0]/tr[0]/tc[0]/p[0]"
        self.assertEqual(places[line]["preview"], "☐ Alcohol beverages: ☐ Wine [☐ Red ☐ White ], ☐ Other:__________")
        self.assertEqual(places[line]["boxes"], 5)
        _, out = self.add_and_fill(path, line, 3, "Red wine")
        self.assertEqual(controls_shown(out)[:2], [("☒", "1"), ("☐", "0")], "box 3 is the control before 'Red'")
        self.assertIn("☐ Other", glyphs_shown(out, "Alcohol"))
        _, out = self.add_and_fill(path, line, 5, "Other drink")
        self.assertEqual(controls_shown(out)[:2], [("☐", "0"), ("☐", "0")])
        self.assertIn("☒ Other", glyphs_shown(out, "Alcohol"))

    def test_a_cell_counts_its_controls_as_its_preview_shows_them(self):
        path, places = self.form(mixed_cell_bytes(), "severity.docx")
        cell = "word/document.xml#b/tbl[0]/tr[1]/tc[0]"
        self.assertEqual((places[cell]["preview"], places[cell]["boxes"]), ("☐ Mild ☐ Moderate ☐ Severe", 3))
        _, out = self.add_and_fill(path, cell, 2, "Moderate")
        self.assertEqual(controls_shown(out)[2:], [("☒", "1")])
        self.assertEqual(glyphs_shown(out, "Mild"), "☐ Mild ☒ Moderate ☐ Severe")
        _, out = self.add_and_fill(path, cell, 3, "Severe")
        self.assertEqual((controls_shown(out)[2:], glyphs_shown(out, "Mild")), ([("☐", "0")], "☐ Mild ☐ Moderate ☒ Severe"))

    def test_a_cell_whose_preview_does_not_show_its_boxes_refuses_a_number(self):
        path, places = self.form(fever_cell_bytes(), "fever.docx")
        tc, line = "word/document.xml#b/tbl[0]/tr[0]/tc[1]", "word/document.xml#b/tbl[0]/tr[0]/tc[1]/p[0]"
        self.assertEqual((places[tc]["preview"], places[tc]["boxes"], places[tc].get("unnumbered_boxes")), ("☐ Yes No", 0, True))
        with self.assertRaises(DocumentError) as refused:
            self.add_and_fill(path, tc, 2, "No fever")
        self.assertEqual(refused.exception.code, "invalid_mapping")
        self.assertIn("choose the line inside the cell, or the checkbox itself", str(refused.exception))
        self.assertEqual((places[line]["preview"], places[line]["boxes"]), ("☐ Yes ☐ No", 2))
        _, out = self.add_and_fill(path, line, 2, "No fever")
        self.assertEqual(control_states(self.dir / "No-fever-2.docx"), ([], ["1"]))

    def test_a_cell_whose_preview_shows_boxes_nested_inside_it_refuses_a_number(self):
        doc = Document()
        doc.add_paragraph("Synthetic severity")
        outer = doc.add_table(rows=1, cols=3)
        outer.cell(0, 0).paragraphs[0].text = "☐ Mild ☐ Severe"          # its own line, and a nested table's box
        outer.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "☐ Other"
        control = checkbox_control()                                   # a control showing no box, and a nested box
        control.find(qn("w:sdtContent")).find(qn("w:r")).find(qn("w:t")).text = "x"
        outer.cell(0, 1).paragraphs[0]._p.append(control)
        outer.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = "☐"
        ff_field(outer.cell(0, 2).paragraphs[0], "checkbox", name="Hidden")   # a legacy checkbox the preview hides,
        outer.cell(0, 2).add_table(rows=1, cols=1).cell(0, 0).text = "☐"   # and a nested box: 1 shown, 1 of its own
        path, places = self.form(docx_bytes(doc), "nested.docx")
        for tc, preview in (("word/document.xml#b/tbl[0]/tr[0]/tc[0]", "☐ Mild ☐ Severe☐ Other"),
                            ("word/document.xml#b/tbl[0]/tr[0]/tc[1]", "x☐"),
                            ("word/document.xml#b/tbl[0]/tr[0]/tc[2]", "☐")):
            self.assertEqual((places[tc]["preview"], places[tc]["boxes"], places[tc].get("unnumbered_boxes")),
                             (preview, 0, True), tc)
            with self.assertRaises(DocumentError) as refused:
                self.add_and_fill(path, tc, 1, "Box one")
            self.assertEqual(refused.exception.code, "invalid_mapping", tc)
        self.assertEqual(places["word/document.xml#b/tbl[0]/tr[0]/tc[0]/p[0]"]["boxes"], 2)


def wrapper_control(inner, kind="richText"):
    """An inline content control (a rich text or a group control) wrapping `inner`, as Word nests them."""
    sdt = OxmlElement("w:sdt")
    pr = OxmlElement("w:sdtPr")
    pr.append(OxmlElement("w:" + kind))
    sdt.append(pr)
    content = OxmlElement("w:sdtContent")
    content.append(inner)
    sdt.append(content)
    return sdt


def own_controls_shown(data):
    """[(glyph shown, flag)] for every checkbox content control of an export - the control itself, never a control that
    only wraps one - in document order."""
    root = etree.fromstring(ZipFile(io.BytesIO(data)).read("word/document.xml"))
    out = []
    for sdt in root.iter(qn("w:sdt")):
        cb = sdt.find(qn("w:sdtPr") + "/{%s}checkbox" % W14)
        if cb is None:
            continue
        flag = cb.find("{%s}checked" % W14)
        out.append(("".join(t.text or "" for t in sdt.find(qn("w:sdtContent")).iter(qn("w:t"))),
                    flag.get("{%s}val" % W14) if flag is not None else None))
    return out


def wrapped_grid_bytes(kind="richText"):
    """The refuter's F1 grid: Chest pain's Yes cell holds a checkbox control printed ticked INSIDE another control."""
    doc = Document()
    doc.add_paragraph("Synthetic review of systems")
    table = doc.add_table(rows=3, cols=3)
    for c, h in enumerate(["Symptom", "Yes", "No"]):
        table.cell(0, c).text = h
    table.cell(1, 0).text = "Chest pain"
    table.cell(2, 0).text = "Dyspnea"
    table.cell(1, 1).paragraphs[0]._p.append(wrapper_control(checkbox_control(checked=True), kind))
    for r, c in ((1, 2), (2, 1), (2, 2)):
        table.cell(r, c).paragraphs[0]._p.append(checkbox_control())
    return docx_bytes(doc)


def wrapped_checklist_bytes():
    """The refuter's F1 check-all column: Warfarin's box, printed ticked, is a checkbox control inside a control."""
    doc = Document()
    doc.add_paragraph("Synthetic medication reconciliation")
    table = doc.add_table(rows=4, cols=2)
    table.cell(0, 0).text = "Medication"
    table.cell(0, 1).text = "Continue (check all that apply)"
    for r, name in ((1, "Warfarin"), (2, "Metformin"), (3, "Aspirin")):
        table.cell(r, 0).text = name
        box = checkbox_control(checked=(r == 1))
        table.cell(r, 1).paragraphs[0]._p.append(wrapper_control(box) if r == 1 else box)
    return docx_bytes(doc)


def wrapped_line_bytes(legacy=False):
    """The refuter's F1 line: 'Smoker? [wrapped ☒] Yes ☐ No' (or its Yes a legacy checkbox inside a control)."""
    doc = Document()
    doc.add_paragraph("Synthetic intake")
    p = doc.add_paragraph("Smoker? ")
    if legacy:
        inner = Document().add_paragraph()
        ff_field(inner, "checkbox", checked=True, name="SmokerYes")
        wrapper = wrapper_control(inner._p[0])
        for run in list(inner._p)[1:]:
            wrapper.find(qn("w:sdtContent")).append(run)
    else:
        wrapper = wrapper_control(checkbox_control(checked=True))
    p._p.append(wrapper)
    p._p.append(text_run(" Yes ☐ No"))
    return docx_bytes(doc)


class NestedCheckboxControls(unittest.TestCase):
    """Round 3 F1: a checkbox control nested inside another control is read and written as that control, end to end -
    detection, the printed state, the writer and the read-back - or its place is refused; never invisible."""

    def test_a_wrapped_control_in_a_grid_cell_is_read_and_written_as_that_control(self):
        from documents import word_v3
        for kind in ("richText", "group"):
            data = wrapped_grid_bytes(kind)
            fields = word_v3.inspect(data)["fields"]
            chest = fields_by_label(fields)["Chest pain"]
            self.assertEqual((chest["kind"], chest["options"], chest.get("printed_ticked"), chest.get("printed_unread")),
                             ("choice", ["Yes", "No"], [0], None), kind)
            self.assertEqual(own_controls_shown(word_v3.fill(data, fields, {chest["key"]: "No"})),
                             [("☐", "0"), ("☒", "1"), ("☐", "0"), ("☐", "0")], kind)
            self.assertEqual(own_controls_shown(word_v3.fill(data, fields, {}))[:2], [("☒", "1"), ("☐", "0")],
                             "left blank it stays as printed - and the page was told it is printed ticked")

    def test_a_wrapped_control_in_a_checklist_is_unticked_when_not_chosen(self):
        from documents import word_v3
        data = wrapped_checklist_bytes()
        fields = word_v3.inspect(data)["fields"]
        field = next(f for f in fields if f["kind"] == "multi_choice")
        self.assertEqual((field["options"], field.get("printed_ticked")), (["Warfarin", "Metformin", "Aspirin"], [0]))
        self.assertEqual(own_controls_shown(word_v3.fill(data, fields, {field["key"]: ["Metformin"]})),
                         [("☐", "0"), ("☒", "1"), ("☐", "0")])

    def test_a_wrapped_control_on_a_line_is_one_of_the_questions_options(self):
        from documents import word_v3
        data = wrapped_line_bytes()
        fields = word_v3.inspect(data)["fields"]
        self.assertEqual([(f["label"], f["kind"], f["options"], f.get("printed_ticked")) for f in fields],
                         [("Smoker?", "choice", ["Yes", "No"], [0])])
        out = word_v3.fill(data, fields, {fields[0]["key"]: "No"})
        self.assertEqual(own_controls_shown(out), [("☐", "0")])
        self.assertEqual(glyphs_shown(out, "Smoker"), "Smoker? ☐ Yes ☒ No")
        places = {row["slot"]: row for row in word_v3.outline(data)["slots"]}
        self.assertEqual((places["word/document.xml#b/p[1]"]["preview"], places["word/document.xml#b/p[1]"]["boxes"]),
                         ("Smoker? ☐ Yes ☐ No", 2))

    def test_a_wrapped_legacy_checkbox_is_read_and_written_too(self):
        from documents import word_v3
        data = wrapped_line_bytes(legacy=True)
        fields = word_v3.inspect(data)["fields"]
        self.assertEqual([(f["label"], f["options"], f.get("printed_ticked")) for f in fields],
                         [("Smoker?", ["Yes", "No"], [0])])
        out = word_v3.fill(data, fields, {fields[0]["key"]: "No"})
        root = etree.fromstring(ZipFile(io.BytesIO(out)).read("word/document.xml"))
        self.assertEqual([c.get(qn("w:val")) for c in root.iter(qn("w:checked"))], ["0"])
        self.assertIn("☒ No", glyphs_shown(out, "Smoker"))

    def test_the_read_back_sees_a_control_the_writer_missed(self):
        """Belt and braces: a reader that loses a nested control (the round-3 shape) can never pass an export."""
        from unittest.mock import patch
        from documents import word_v3
        data = wrapped_grid_bytes()
        fields = word_v3.inspect(data)["fields"]
        chest = fields_by_label(fields)["Chest pain"]
        real = word_v3.box_pieces

        def blind(p):                              # the round-3 reader: a control inside a control is not seen
            return [(kind, node, n) for kind, node, n in real(p) if node is None or node.getparent().tag != qn("w:sdtContent")]
        with patch.object(word_v3, "box_pieces", side_effect=blind):
            with self.assertRaises(word_v3.V3Failure) as refused:
                word_v3.fill(data, fields, {chest["key"]: "No"})
        self.assertIn(refused.exception.code, ("fill_mismatch", "unsupported_layout"))

    def test_a_control_that_cannot_be_read_in_place_refuses_its_line(self):
        """A checkbox control inside a checkbox control: the line shows one box where the package holds two, so the
        line is not read as a question, its boxes are not numbered, and no box on it is ever called unticked."""
        from documents import word_v3
        doc = Document()
        doc.add_paragraph("Synthetic intake")
        p = doc.add_paragraph("Allergy? ")
        outer = checkbox_control()
        outer.find(qn("w:sdtContent")).append(checkbox_control(checked=True))
        p._p.append(outer)
        p._p.append(text_run(" Yes ☐ No"))
        data = docx_bytes(doc)
        info = word_v3.inspect(data)
        self.assertEqual([f["label"] for f in info["fields"] if f["kind"] in ("choice", "checkbox")], [])
        self.assertTrue(any("cannot read in place" in note for note in info["notes"]), info["notes"])
        places = {row["slot"]: row for row in word_v3.outline(data)["slots"]}
        self.assertEqual((places["word/document.xml#b/p[1]"]["boxes"], places["word/document.xml#b/p[1]"].get("unnumbered_boxes")),
                         (0, True))
        index = word_v3.SlotIndex(word_v3.open_package(data))
        glyph = {"kind": "glyph", "part": "word/document.xml", "slot": "b/p[1]/glyphs", "glyph": 0}
        self.assertIsNone(word_v3.box_state(index, glyph))
        with self.assertRaises(word_v3.V3Failure) as refused:
            word_v3.locate(data, [{"slot": "word/document.xml#b/p[1]", "kind": "checkbox", "label": "No", "glyph": 1}])
        self.assertEqual(refused.exception.code, "invalid_mapping")


def wrap_in_control(el):
    """Wrap a table cell (w:tr/w:sdt/w:tc) or a whole row (w:tbl/w:sdt/w:tr) in a content control, in place."""
    sdt = OxmlElement("w:sdt")
    sdt.append(OxmlElement("w:sdtPr"))
    content = OxmlElement("w:sdtContent")
    sdt.append(content)
    el.getparent().replace(el, sdt)
    content.append(el)
    return sdt


def yes_no_grid(doc, rows=("Chest pain", "Dyspnea")):
    table = doc.add_table(rows=len(rows) + 1, cols=3)
    for c, h in enumerate(["Symptom", "Yes", "No"]):
        table.cell(0, c).text = h
    for r, label in enumerate(rows, 1):
        table.cell(r, 0).text = label
        table.cell(r, 1).text = "☐"
        table.cell(r, 2).text = "☐"
    return table


class WrappedCellsAndRows(unittest.TestCase):
    """Round 3 P1: a content control wrapping a whole table cell (w:tr/w:sdt/w:tc) or a whole row (w:tbl/w:sdt/w:tr)
    is looked through, so every cell keeps its column; the table_rows writer counts columns as the index does."""

    def test_a_wrapped_yes_cell_keeps_every_later_cell_in_its_column(self):
        from documents import word_v3
        doc = Document()
        doc.add_paragraph("Synthetic review of systems")
        table = yes_no_grid(doc)
        wrap_in_control(table.cell(1, 1)._tc)
        data = docx_bytes(doc)
        fields = word_v3.inspect(data)["fields"]
        chest = fields_by_label(fields)["Chest pain"]
        self.assertEqual((chest["kind"], chest["options"]), ("choice", ["Yes", "No"]))
        rows = paragraphs_bytes(word_v3.fill(data, fields, {chest["key"]: "No"}))
        self.assertEqual(rows[rows.index("Chest pain"):rows.index("Chest pain") + 3], ["Chest pain", "☐", "☒"])
        rows = paragraphs_bytes(word_v3.fill(data, fields, {chest["key"]: "Yes"}))
        self.assertEqual(rows[rows.index("Chest pain"):rows.index("Chest pain") + 3], ["Chest pain", "☒", "☐"])
        places = {row["slot"]: row for row in word_v3.outline(data)["slots"]}
        self.assertIn("word/document.xml#b/tbl[0]/tr[1]/tc[1]", places, "the wrapped cell is a place of its own")
        self.assertFalse([f for f in fields if f["locator"]["target"].get("type") == "sdt_text"],
                         "the control around the cell is not an answer area")

    def test_a_wrapped_answer_cell_takes_its_answer(self):
        from documents import word_v3
        doc = Document()
        doc.add_paragraph("Synthetic intake")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Allergies"
        table.cell(1, 0).text = "Plan"
        wrap_in_control(table.cell(0, 1)._tc)
        data = docx_bytes(doc)
        fields = word_v3.inspect(data)["fields"]
        self.assertEqual([(f["label"], f["locator"]["target"]["type"]) for f in fields], [("Allergies", "cell"), ("Plan", "cell")])
        out = word_v3.fill(data, fields, {fields[0]["key"]: "None documented", fields[1]["key"]: "Review"})
        self.assertEqual(paragraphs_bytes(out)[1:], ["Allergies", "None documented", "Plan", "Review"])

    def test_a_wrapped_row_is_read_in_its_place(self):
        from documents import word_v3
        doc = Document()
        doc.add_paragraph("Synthetic review of systems")
        table = yes_no_grid(doc, rows=("Chest pain", "Dyspnea", "Cough"))
        wrap_in_control(table.rows[1]._tr)
        data = docx_bytes(doc)
        fields = word_v3.inspect(data)["fields"]
        by = fields_by_label(fields)
        self.assertEqual({label: by[label]["options"] for label in ("Chest pain", "Dyspnea", "Cough")},
                         {"Chest pain": ["Yes", "No"], "Dyspnea": ["Yes", "No"], "Cough": ["Yes", "No"]})
        rows = paragraphs_bytes(word_v3.fill(data, fields, {by["Chest pain"]["key"]: "No", by["Cough"]["key"]: "Yes"}))
        self.assertEqual(rows[rows.index("Chest pain"):], ["Chest pain", "☐", "☒", "Dyspnea", "☐", "☐", "Cough", "☒", "☐"])

    def medication_log(self, wrapped):
        doc = Document()
        doc.add_paragraph("Synthetic medications")
        table = doc.add_table(rows=3, cols=3)
        for c, h in enumerate(["Medication", "Dose", "Route"]):
            table.cell(0, c).text = h
        for r in (1, 2):
            tr = table.rows[r]._tr
            before = OxmlElement("w:gridBefore")
            before.set(qn("w:val"), "1")
            tr.get_or_add_trPr().append(before)
            tr.remove(tr.findall(qn("w:tc"))[0])
        if wrapped:
            wrap_in_control(table.rows[2]._tr)
        return docx_bytes(doc)

    def test_the_table_rows_writer_counts_columns_as_the_index_does(self):
        """Found repairing P1: log rows that skip a grid column (w:gridBefore) - the writer put the dose under Route and
        dropped the route."""
        from documents import word_v3
        for wrapped in (False, True):
            data = self.medication_log(wrapped)
            field = next(f for f in word_v3.inspect(data)["fields"] if f["kind"] == "table_rows")
            self.assertEqual((field["options"], field["locator"]["target"]["rows"]), (["Dose", "Route"], [1, 2]), wrapped)
            out = word_v3.fill(data, [field], {field["key"]: [["10 mg", "oral"], ["5 mg", "iv"], ["1 g", "topical"]]})
            self.assertEqual(paragraphs_bytes(out)[4:], ["10 mg", "oral", "5 mg", "iv", "1 g", "topical"], wrapped)


def block_wrapped_grid_bytes():
    """The refuter's P2 input: a Yes/No grid of checkbox controls whose Chest pain Yes control sits in a paragraph
    wrapped by a block-level content control inside its cell (w:tc/w:sdt/w:p)."""
    doc = Document()
    doc.add_paragraph("Synthetic review of systems")
    table = doc.add_table(rows=3, cols=3)
    for c, h in enumerate(["Symptom", "Yes", "No"]):
        table.cell(0, c).text = h
    table.cell(1, 0).text = "Chest pain"
    table.cell(2, 0).text = "Dyspnea"
    for r in (1, 2):
        for c in (1, 2):
            table.cell(r, c).paragraphs[0]._p.append(checkbox_control())
    wrap_in_control(table.cell(1, 1).paragraphs[0]._p)
    return docx_bytes(doc)


class LoneBoxInAnOptionRow(unittest.TestCase):
    """Round 3 P2: a box in a paragraph wrapped by a block-level control is its cell's box, and a lone box is never
    labelled with its row's question while that row holds another box the detector did not give it."""

    def test_a_block_wrapped_yes_control_is_the_rows_yes(self):
        from documents import word_v3
        data = block_wrapped_grid_bytes()
        fields = word_v3.inspect(data)["fields"]
        self.assertNotIn("Checkbox", [f["label"] for f in fields])
        chest = fields_by_label(fields)["Chest pain"]
        self.assertEqual((chest["kind"], chest["options"]), ("choice", ["Yes", "No"]))
        self.assertEqual(own_controls_shown(word_v3.fill(data, fields, {chest["key"]: "Yes"})),
                         [("☒", "1"), ("☐", "0"), ("☐", "0"), ("☐", "0")])
        self.assertEqual(own_controls_shown(word_v3.fill(data, fields, {chest["key"]: "No"})),
                         [("☐", "0"), ("☒", "1"), ("☐", "0"), ("☐", "0")])
        places = {row["slot"]: row for row in word_v3.outline(data)["slots"]}
        self.assertEqual(places["word/document.xml#b/tbl[0]/tr[1]/tc[1]"]["boxes"], 1)

    def assert_no_lone_row_question(self, data, label):
        from documents import word_v3
        info = word_v3.inspect(data)
        lone = [f for f in info["fields"] if f["kind"] == "checkbox" and f["label"] == label]
        self.assertEqual(lone, [], "a lone box took the row's question although the row prints another box")
        self.assertTrue(any("another box" in note for note in info["notes"]), info["notes"])

    def test_a_lone_no_box_beside_a_yes_box_it_was_not_given_is_refused(self):
        doc = Document()
        doc.add_paragraph("Synthetic review of systems")
        table = doc.add_table(rows=2, cols=3)
        for c, h in enumerate(["Symptom", "Yes", "No"]):
            table.cell(0, c).text = h
        table.cell(1, 0).text = "Chest pain"
        table.cell(1, 1).text = "☐ Yes"
        table.cell(1, 2).text = "☐"
        self.assert_no_lone_row_question(docx_bytes(doc), "Chest pain")

    def test_a_lone_no_box_beside_an_unreadable_yes_is_refused(self):
        doc = Document()
        doc.add_paragraph("Synthetic review of systems")
        table = doc.add_table(rows=2, cols=3)
        for c, h in enumerate(["Symptom", "Yes", "No"]):
            table.cell(0, c).text = h
        table.cell(1, 0).text = "Chest pain"
        outer = checkbox_control()
        outer.find(qn("w:sdtContent")).append(checkbox_control(checked=True))
        table.cell(1, 1).paragraphs[0]._p.append(outer)
        table.cell(1, 2).text = "☐"
        self.assert_no_lone_row_question(docx_bytes(doc), "Chest pain")

    def test_a_row_with_one_box_per_question_keeps_both(self):
        from documents import word_v3
        doc = Document()
        doc.add_paragraph("Synthetic history")
        table = doc.add_table(rows=1, cols=4)
        for c, text in enumerate(["Fever", "☐", "Consent obtained", "☐"]):
            table.cell(0, c).text = text
        fields = word_v3.inspect(docx_bytes(doc))["fields"]
        self.assertEqual([(f["label"], f["kind"]) for f in fields], [("Fever", "checkbox"), ("Consent obtained", "checkbox")])


def sided_grid_bytes(side="Left", *, heading=False, after=False):
    doc = Document()
    doc.add_paragraph("Synthetic joint examination")
    if heading:
        doc.add_paragraph(side + " knee:")
        table = yes_no_grid(doc, ("Effusion",))
    else:
        table = doc.add_table(rows=2, cols=4)
        headers = ["Finding", "Yes", "No", "Side"] if after else ["Side", "Finding", "Yes", "No"]
        values = ["Effusion", "☐", "☐", side] if after else [side, "Effusion", "☐", "☐"]
        for row, vals in enumerate((headers, values)):
            for col, value in enumerate(vals):
                table.cell(row, col).text = value
    return docx_bytes(doc)


class RoundThreeRemainingDefects(unittest.TestCase):
    def test_a_row_side_governs_multiple_details_until_the_next_printed_side(self):
        from documents import word_v3
        doc = Document()
        table = doc.add_table(rows=1, cols=6)
        for column, text in enumerate(("Right lung", "Effusion: ____", "Consolidation: ____",
                                       "Left lung", "Effusion: ____", "Consolidation: ____")):
            table.cell(0, column).text = text
        fields = word_v3.inspect(docx_bytes(doc))["fields"]
        self.assertEqual([f["label"] for f in fields], ["Right lung — Effusion", "Right lung — Consolidation",
                                                     "Left lung — Effusion", "Left lung — Consolidation"])

    def test_paired_side_cells_govern_only_their_own_answer(self):
        from documents import word_v3
        for right, left in (("Right", "Left"), ("Right knee", "Left knee")):
            doc = Document()
            table = doc.add_table(rows=1, cols=4)
            for column, text in enumerate((right, "Effusion: __________", left, "Effusion: __________")):
                table.cell(0, column).text = text
            data = docx_bytes(doc)
            fields = word_v3.inspect(data)["fields"]
            self.assertEqual([f["label"] for f in fields], [right + " — Effusion", left + " — Effusion"])
            exported = Document(io.BytesIO(word_v3.fill(data, fields, {fields[0]["key"]: "Absent", fields[1]["key"]: "Present"})))
            self.assertEqual([c.text for c in exported.tables[0].rows[0].cells],
                             [right, "Effusion: Absent", left, "Effusion: Present"])

    def test_a_separate_side_cell_qualifies_the_question_even_after_the_boxes(self):
        from documents import word_v3
        for side in ("L", "R", "Left", "Right", "Bilateral"):
            for after in (False, True):
                with self.subTest(side=side, after=after):
                    data = sided_grid_bytes(side, after=after)
                    field = next(f for f in word_v3.inspect(data)["fields"] if "Effusion" in f["label"])
                    self.assertIn(side, field["label"])
                    out = paragraphs_bytes(word_v3.fill(data, [field], {field["key"]: "Yes"}))
                    at = out.index("Effusion")
                    self.assertEqual(out[at + 1:at + 3], ["☒", "☐"])

    def test_a_plain_side_heading_qualifies_its_questions(self):
        from documents import word_v3
        field = next(f for f in word_v3.inspect(sided_grid_bytes(heading=True))["fields"] if "Effusion" in f["label"])
        self.assertIn("Left knee", field["label"])

    def test_units_and_questionnaire_phrases_are_not_sides(self):
        from documents.word_v3 import side_phrase
        for text in ("per L", "l. other – specify", "l. I got the help I wanted", "The services were right for me"):
            self.assertFalse(side_phrase(text), text)
        for text in ("L", "R", "(L)", "Left knee", "Right side", "Costophrenical pathologies (right)"):
            self.assertTrue(side_phrase(text), text)

    def test_a_heading_in_another_table_cell_does_not_govern_the_next_row(self):
        from documents import word_v3
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).paragraphs[0].add_run("Right leg").bold = True
        table.cell(0, 1).paragraphs[0].add_run("Left leg").bold = True
        table.cell(1, 0).text = "Notes: __________"
        table.cell(1, 1).text = "Notes: __________"
        fields = word_v3.inspect(docx_bytes(doc))["fields"]
        labels = [f["label"] for f in fields]
        self.assertIn("Right leg — Notes", labels)
        self.assertIn("Left leg — Notes", labels)
        self.assertFalse(any("Left leg — Right leg" in label or "Right leg — Left leg" in label for label in labels))

    def test_a_repeated_side_header_in_one_table_resets_the_side(self):
        from documents import word_v3
        doc = Document()
        table = doc.add_table(rows=4, cols=2)
        for row, side in ((0, "Right leg"), (2, "Left leg")):
            table.cell(row, 0).text = side
            table.cell(row, 1).text = "Measurement"
            table.cell(row + 1, 0).text = "Vein diameter"
        fields = word_v3.inspect(docx_bytes(doc))["fields"]
        questions = [f for f in fields if "Vein diameter" in f["label"]]
        self.assertEqual([f["label"] for f in questions],
                         ["Right leg — Vein diameter — Measurement", "Left leg — Vein diameter — Measurement"])

    def test_duplicate_context_uses_readable_line_text(self):
        from documents import word_v3
        doc = Document()
        doc.add_paragraph("Synthetic symptoms")
        doc.add_paragraph("Pain? ☐ Yes ☐ No")
        doc.add_paragraph("Pain? ☐ Yes ☐ No")
        fields = word_v3.inspect(docx_bytes(doc))["fields"]
        self.assertEqual(len(fields), 2)
        for field in fields:
            self.assertIn("Pain?", field["context"])
            self.assertNotRegex(field["context"], r"(?:b|hf)/|/glyphs|p\[|sdt\[")
        self.assertNotEqual(fields[0]["context"], fields[1]["context"])

    def test_text_fields_are_not_drawn_as_printed_boxes(self):
        from documents import word_v3
        for legacy in (False, True):
            doc = Document()
            p = doc.add_paragraph("Side: ")
            if legacy:
                ff_text_field(p, "Side")
            else:
                p._p.append(_sdt("rich", placeholder=" "))
            p.add_run(" Left ☐ Right")
            data = docx_bytes(doc)
            place = next(r for r in word_v3.outline(data)["slots"] if r["kind"] == "p")
            self.assertEqual(place["preview"], "Side: [text field] Left ☐ Right")
            self.assertEqual(place["boxes"], 1)
            field = word_v3.locate(data, [{"slot": place["slot"], "kind": "checkbox", "label": "Right", "box": 1}])[0]
            self.assertIn("Left ☒ Right", "".join(paragraphs_bytes(word_v3.fill(data, [field], {field["key"]: True}))))

    def test_an_unaddressable_checkbox_in_a_hyperlink_cannot_be_mapped_by_number(self):
        from documents import word_v3
        doc = Document()
        p = doc.add_paragraph("Synthetic symptom: ")
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), "local")
        link.append(checkbox_control())
        p._p.append(link)
        p.add_run(" Yes")
        data = docx_bytes(doc)
        with self.assertRaises(word_v3.V3Failure):
            word_v3.locate(data, [{"slot": "word/document.xml#b/p[0]", "kind": "checkbox", "label": "Symptom", "box": 1}])

    def test_private_use_symbol_codes_match_the_control_state(self):
        from documents import word_v3
        control = checkbox_control(checked=True)
        state = control.find(qn("w:sdtPr") + "/{%s}checkbox/{%s}checkedState" % (W14, W14))
        state.set("{%s}val" % W14, "F0FE")
        content = control.find(qn("w:sdtContent"))
        content.find(qn("w:r") + "/" + qn("w:t")).text = "\u00fe"
        self.assertIs(word_v3.sdt_checkbox_state(control), True)

    def test_a_control_printing_multiple_glyphs_is_unread(self):
        from documents import word_v3
        control = checkbox_control(checked=True)
        control.find(qn("w:sdtContent") + "/" + qn("w:r") + "/" + qn("w:t")).text = "☒☐"
        self.assertIsNone(word_v3.sdt_checkbox_state(control))

    def test_a_wrapped_paragraphs_box_is_read_by_the_cell_writer(self):
        from documents import word_v3
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        control = checkbox_control(checked=True)
        cell.paragraphs[0]._p.append(control)
        wrap_in_control(cell.paragraphs[0]._p)
        self.assertEqual(word_v3.cell_boxes(cell._tc), [("sdt", control)])
        word_v3.set_cell_box(cell._tc, False)
        self.assertIs(word_v3.sdt_checkbox_state(control), False)


if __name__ == "__main__":
    unittest.main()


class SourceBoundGlyphWrites(unittest.TestCase):
    """Inserted checkbox-looking text is not an original printed control."""
    def form(self, before='glyph', after='glyph', *, local_namespace=False):
        from . import word_v3 as w
        doc = Document();p = doc.add_paragraph('Choices? ')
        def control(kind, label):
            if kind == 'glyph': p.add_run('☐')
            elif kind == 'sdt': p._p.append(checkbox_control())
            else: ff_field(p, 'checkbox', name=label)
            p.add_run(' ' + label)
        if before: control(before, 'Before');p.add_run(': ')
        p.add_run('_____')
        if after: p.add_run(' ');control(after, 'After')
        if local_namespace:
            text = w.text_of(p._p)
            p._p[:] = []
            run = etree.SubElement(p._p, qn('w:r'), nsmap={'x': W})
            etree.SubElement(run, qn('w:t')).text = text
        stream = io.BytesIO();doc.save(stream);raw = stream.getvalue()
        index = w.SlotIndex(w.open_package(raw));slot = next(s for s in index.slots if s.kind == 'p')
        start = w.paragraph_text(slot.el)[0].index('_____')
        locator = {'contract': 'v3', 'part': slot.part, 'slot': slot.path, 'fingerprint': w.fingerprint(slot.el)}
        choice = {'key': 'c', 'kind': 'multi_choice', 'label': 'Choices',
                  'options': (['Before'] if before else []) + (['After'] if after else []),
                  'locator': dict(locator, target={'type': 'checkboxes', 'boxes': w.place_box_refs(index, slot)})}
        text = {'key': 't', 'kind': 'text', 'label': 'Description',
                'locator': dict(locator, target={'type': 'span', 'start': start, 'end': start + 5, 'shape': 'blank'})}
        return raw, [choice, text]

    def test_inserted_box_character_never_becomes_the_printed_after_box(self):
        from . import word_v3 as w
        for before, after in (('glyph', 'glyph'), (None, 'glyph'), ('glyph', None)):
            with self.subTest(before=before, after=after):
                raw, fields = self.form(before, after)
                answer = 'literal ☐ symbol retained'
                result = w.fill(raw, fields, {'c': fields[0]['options'], 't': answer})
                root = etree.fromstring(ZipFile(io.BytesIO(result)).read('word/document.xml'))
                physical = ''.join(root.itertext())
                expected = 'Choices? ' + ('☒ Before: ' if before else '') + answer + (' ☒ After' if after else '')
                self.assertEqual(physical, expected)

    def test_native_controls_and_missing_states_remain_independent_of_inserted_boxes(self):
        from . import word_v3 as w
        for before, after in (('sdt', 'glyph'), ('glyph', 'sdt'), ('ff', 'glyph'), ('glyph', 'ff'),
                              ('sdt', 'sdt'), ('ff', 'ff')):
            with self.subTest(before=before, after=after):
                raw, fields = self.form(before, after)
                missing = w.fill(raw, fields, {})
                with ZipFile(io.BytesIO(raw)) as old, ZipFile(io.BytesIO(missing)) as new:
                    self.assertEqual({n: old.read(n) for n in old.namelist()},
                                     {n: new.read(n) for n in new.namelist()})
                answer = 'ordinary ☐ text\n' + 'long retained words ' * 80
                result = w.fill(raw, fields, {'c': ['After'], 't': answer})
                root = etree.fromstring(ZipFile(io.BytesIO(result)).read('word/document.xml'))
                text = ''.join(root.itertext())
                self.assertIn('ordinary ☐ text', text)
                self.assertIn('long retained words ' * 80, text)
                if before == 'glyph': self.assertIn('☐ Before:', text)
                if after == 'glyph': self.assertTrue(text.endswith('☒ After'))
                sdt_states = [state for _, state in controls_shown(result)]
                self.assertEqual(sdt_states, (['0'] if before == 'sdt' else []) + (['1'] if after == 'sdt' else []))
                ff_states = root.xpath('//w:checkBox/w:checked/@w:val', namespaces={'w': W})
                self.assertEqual(ff_states, (['0'] if before == 'ff' else []) + (['1'] if after == 'ff' else []))

    def test_deleted_original_box_and_stale_source_reference_are_refused(self):
        from . import word_v3 as w
        raw, fields = self.form()
        changed = deepcopy(fields)
        changed[1]['locator']['target'].update(start=0, end=len(w.paragraph_text(
            next(s.el for s in w.SlotIndex(w.open_package(raw)).slots if s.kind == 'p'))[0]))
        with self.assertRaises(w.V3Failure):
            w.fill(raw, changed, {'c': ['Before', 'After'], 't': '☒ fabricated boxes ☒'})
        changed = deepcopy(fields);changed[0]['locator']['fingerprint'] = '0' * 64
        with self.assertRaises(w.V3Failure) as error:
            w.fill(raw, changed, {'c': ['Before', 'After'], 't': 'ordinary'})
        self.assertEqual(error.exception.code, 'mapping_mismatch')

    def test_final_readback_reads_original_leaf_not_inserted_mark(self):
        from . import word_v3 as w
        from unittest.mock import patch
        raw, fields = self.form()
        real_write = w.Package.write
        def corrupt(package, changed):
            output = real_write(package, changed)
            other = w.open_package(output);root = other.tree('word/document.xml')
            texts = list(root.iter(w.T))
            original_after = [node for node in texts if node.text == '☒'][-1]
            original_after.text = '☐'
            inserted = next(node for node in texts if 'literal ☐ symbol' in (node.text or ''))
            inserted.text = inserted.text.replace('☐', '☒')
            return real_write(other, {'word/document.xml': root})
        with patch.object(w.Package, 'write', corrupt):
            with self.assertRaises(w.V3Failure) as error:
                w.fill(raw, fields, {'c': ['Before', 'After'], 't': 'literal ☐ symbol retained'})
        self.assertEqual(error.exception.code, 'fill_mismatch')

    def test_readback_accepts_local_namespace_aliases_for_original_runs(self):
        from . import word_v3 as w
        raw, fields = self.form(local_namespace=True)
        result = w.fill(raw, fields, {'c': ['Before', 'After'], 't': 'literal ☐ symbol retained'})
        root = etree.fromstring(ZipFile(io.BytesIO(result)).read('word/document.xml'))
        self.assertEqual(''.join(root.itertext()), 'Choices? ☒ Before: literal ☐ symbol retained ☒ After')


# ─── the engine switch (EC3 / EC4 / EC5): every legacy-acceptable shape at least as good on v3 ────────────────────

OWNER_FORM = Path(__file__).with_name("__fixtures__") / "Colonoscopy_Operative_Template.docx"


def _underlined_tab(paragraph):
    """'Label:' + an underlined tab, the legacy tab-line shape (python-docx paragraph)."""
    run = paragraph.add_run("\t")
    run.font.underline = True


def merged_answer_bytes(*, label_merged=False, second_label="Allergies"):
    """SF03: a section row, then a vertically merged answer cell beside two questions (or beside one question whose
    label cell is merged the same way), then a plain row."""
    doc = Document()
    doc.add_heading("Synthetic merged form", 1)
    table = doc.add_table(rows=4, cols=2)
    top = table.cell(0, 0).merge(table.cell(0, 1))
    top.text = "History"
    table.cell(1, 0).text = "Reason for visit"
    table.cell(2, 0).text = "Current medications"
    if label_merged:
        table.cell(2, 0).merge(table.cell(3, 0))
    elif second_label:
        table.cell(3, 0).text = second_label
    table.cell(2, 1).merge(table.cell(3, 1))
    plain = doc.add_table(rows=1, cols=2)
    plain.cell(0, 0).text = "Plan"
    return docx_bytes(doc)


def instruction_cell_bytes(*, header=False, grey_header=False):
    """SF15: a two-column table whose FIRST answer cell prints a grey instruction; with header=True a printed header
    row ('Question | Answer') above (in grey ink with grey_header=True), the control for a real column header."""
    from docx.shared import RGBColor
    doc = Document()
    doc.add_heading("Synthetic instruction form", 1)
    rows = 3 if header else 2
    table = doc.add_table(rows=rows, cols=2)
    first = 1 if header else 0
    if header:
        for col, text in enumerate(("Question", "Answer")):
            run = table.cell(0, col).paragraphs[0].add_run(text)
            if grey_header:
                run.font.color.rgb = RGBColor.from_string("808080")
    table.cell(first, 0).text = "Reason for visit"
    table.cell(first + 1, 0).text = "Plan"
    table.cell(first, 1).paragraphs[0].add_run("Describe the main complaint.").font.color.rgb = RGBColor.from_string("808080")
    return docx_bytes(doc)


def numbered_slots_bytes(*, heading=True, marks=("1.", "2.", "3.")):
    """D02's numbered slots: under two section headings, lines '1.' '2.' '3.' each with an underlined tab."""
    doc = Document()
    for section in ("POSTOPERATIVE FINDINGS", "RECOMMENDATIONS"):
        if heading:
            doc.add_heading(section, 2)
        else:
            doc.add_paragraph(section + " (plain text, not a heading)")
        for mark in marks:
            _underlined_tab(doc.add_paragraph(mark))
    return docx_bytes(doc)


class EngineSwitchShapes(unittest.TestCase):
    """The blocking differences the legacy-vs-v3 comparison measured (EC3, EC4, EC5), through the sandboxed worker."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def inspect(self, data, name="form.docx"):
        path = self.dir / name
        path.write_bytes(data)
        return inspect_template(path, engine="v3"), path

    # EC3 ─ a vertically merged answer cell shared by two questions is never given to one of them silently
    def test_ec3_a_merged_answer_cell_shared_by_two_questions_is_mapped_for_neither_and_said(self):
        info, _ = self.inspect(merged_answer_bytes())
        slots = {f["locator"]["slot"]: f["label"] for f in info["fields"]}
        self.assertNotIn("b/tbl[0]/tr[2]/tc[1]", slots, "the shared merged cell was given to one question")
        self.assertNotIn("Current medications", slots.values())
        self.assertNotIn("Allergies", slots.values())
        notes = [n for n in info["warnings"] if "merged" in n]
        self.assertEqual(len(notes), 1, info["warnings"])
        self.assertIn("Current medications", notes[0])
        self.assertIn("Allergies", notes[0])
        # Everything else on the form is still mapped.
        self.assertEqual(sorted(slots.values()), ["Plan", "Reason for visit"])

    def test_ec3_control_a_merged_answer_cell_of_one_question_is_still_mapped(self):
        for kwargs in ({"label_merged": True}, {"second_label": ""}):
            with self.subTest(**kwargs):
                info, _ = self.inspect(merged_answer_bytes(**kwargs))
                by = fields_by_label(info["fields"])
                self.assertIn("Current medications", by)
                self.assertEqual(by["Current medications"]["locator"]["slot"], "b/tbl[0]/tr[2]/tc[1]")
                self.assertFalse([n for n in info["warnings"] if "merged" in n], info["warnings"])

    # EC4 ─ an instruction printed in another row's answer cell is never this row's column header
    def test_ec4_another_rows_instruction_is_not_this_questions_label_or_context(self):
        info, path = self.inspect(instruction_cell_bytes())
        by = fields_by_label(info["fields"])
        self.assertIn("Plan", by, [f["label"] for f in info["fields"]])
        self.assertNotIn("Describe", by["Plan"]["context"])
        # The instruction cell is still the first question's own answer place (the gain over legacy is kept).
        self.assertEqual(by["Reason for visit"]["locator"]["target"]["type"], "cell_append")
        out = self.dir / "out.docx"
        fill_template(path, info["fields"], {by["Plan"]["key"]: "Rest.", by["Reason for visit"]["key"]: "Knee."}, out)
        self.assertIn("Rest.", part_text(out))

    def test_ec4_control_a_printed_header_row_still_names_the_column(self):
        for grey in (False, True):
            with self.subTest(grey_header=grey):
                info, _ = self.inspect(instruction_cell_bytes(header=True, grey_header=grey))
                labels = [f["label"] for f in info["fields"]]
                self.assertIn("Plan — Answer", labels)

    # EC5 ─ numbered slots carry their section, so two sections' '1.' are different questions
    def test_ec5_numbered_slots_carry_their_section(self):
        info, _ = self.inspect(numbered_slots_bytes())
        labels = [f["label"] for f in info["fields"]]
        self.assertEqual(labels, ["POSTOPERATIVE FINDINGS — %d" % n for n in (1, 2, 3)]
                         + ["RECOMMENDATIONS — %d" % n for n in (1, 2, 3)])
        self.assertEqual(len(set(labels)), len(labels))

    def test_ec5_the_owner_colonoscopy_form_has_unique_labels_as_on_legacy(self):
        info = inspect_template(OWNER_FORM, engine="v3")
        labels = [f["label"] for f in info["fields"]]
        self.assertEqual(len(set(labels)), len(labels), labels)
        for section in ("POSTOPERATIVE FINDINGS", "RECOMMENDATIONS"):
            for n in (1, 2, 3):
                self.assertIn("%s — %d" % (section, n), labels)

    def test_ec5_control_without_a_section_heading_the_marker_is_kept(self):
        info, _ = self.inspect(numbered_slots_bytes(heading=False))
        self.assertFalse([f for f in info["fields"] if "—" in f["label"]], [f["label"] for f in info["fields"]])


# ─── the refutation of 8e00df9e (refute-wec): EC3 / EC4 / EC5 repairs and the false greens it found ────────────────

def _shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _grey(paragraph, text):
    from docx.shared import RGBColor
    paragraph.add_run(text).font.color.rgb = RGBColor.from_string("808080")


def _tab_line(container, mark, bold=False):
    paragraph = container.add_paragraph()
    paragraph.add_run(mark).bold = bold or None
    _underlined_tab(paragraph)
    return paragraph


class RefutationRepairs(unittest.TestCase):
    """Each finding of the refutation (M1-M10, F02, F04, F05, F07-F09) as a pinned input: a numbered slot never carries
    another question's label, a merged answer cell beside two printed questions is mapped for none of them whatever
    fills it, and another row's instruction is never a column header."""

    def inspect(self, data):
        from . import word_v3 as w
        return w.inspect(data)

    def labels(self, data):
        return [f["label"] for f in self.inspect(data)["fields"]]

    # M1 ─ a bold question label is not a section: numbered slots carry only a Word heading of their own container
    def test_m1_a_bold_question_label_never_names_the_numbered_slots_under_another_question(self):
        doc = Document()
        doc.add_paragraph().add_run("Known allergies").bold = True
        _underlined_tab(doc.add_paragraph())
        doc.add_paragraph("Current medications:")
        for mark in ("1.", "2.", "3."):
            _tab_line(doc, mark)
        labels = self.labels(docx_bytes(doc))
        self.assertFalse([label for label in labels if label.startswith("Known allergies —")], labels)
        self.assertEqual([label for label in labels if label.endswith(("1.", "2.", "3."))], ["1.", "2.", "3."])

    def test_m1_a_bold_label_cell_never_names_the_numbered_slots_of_the_rows_below(self):
        doc = Document()
        table = doc.add_table(rows=4, cols=2)
        table.cell(0, 0).paragraphs[0].add_run("Diagnosis").bold = True
        table.cell(1, 0).text = "Plan:"
        for row, mark in ((2, "1."), (3, "2.")):
            paragraph = table.cell(row, 1).paragraphs[0]
            paragraph.add_run(mark)
            _underlined_tab(paragraph)
        labels = self.labels(docx_bytes(doc))
        self.assertFalse([label for label in labels if "Diagnosis —" in label], labels)

    def test_m1_control_a_word_heading_directly_above_names_the_slots_a_question_line_between_does_not(self):
        """Round 3 (conservative): a section only when unambiguous. A Word heading directly above the slots names them; a
        plain question line between ('Current medications:') leaves the bare marker (declared cost: legacy names them
        by the section)."""
        doc = Document()
        doc.add_heading("History", 2)
        for mark in ("1.", "2."):
            _tab_line(doc, mark)
        doc.add_heading("Medicines", 2)
        doc.add_paragraph("Current medications:")
        for mark in ("1.", "2."):
            _tab_line(doc, mark)
        labels = [label for label in self.labels(docx_bytes(doc)) if label.endswith(("1", "2", "1.", "2."))]
        self.assertEqual(labels, ["History — 1", "History — 2", "1.", "2."])

    def test_m8_bold_markers_take_their_section_not_the_previous_marker(self):
        for marks in (("1.", "2.", "3."), ("(1)", "(2)", "(3)")):
            with self.subTest(marks=marks):
                doc = Document()
                for section in ("POSTOPERATIVE FINDINGS", "RECOMMENDATIONS"):
                    doc.add_heading(section, 2)
                    for mark in marks:
                        _tab_line(doc, mark, bold=True)
                self.assertEqual(self.labels(docx_bytes(doc)), ["POSTOPERATIVE FINDINGS — %d" % n for n in (1, 2, 3)]
                                 + ["RECOMMENDATIONS — %d" % n for n in (1, 2, 3)])

    # M9 ─ a heading in another table cell is not the section of this cell's slots
    def test_m9_a_column_heading_to_the_right_never_names_the_left_columns_slots(self):
        for style in (None, "Heading 2"):
            with self.subTest(style=style):
                doc = Document()
                table = doc.add_table(rows=2, cols=2)
                for col, text in enumerate(("FINDINGS", "SPECIMENS")):
                    paragraph = table.cell(0, col).paragraphs[0]
                    paragraph.add_run(text).bold = True
                    if style:
                        paragraph.style = doc.styles[style]
                for col, marks in ((0, ("1.", "2.")), (1, ("3.", "4."))):
                    cell = table.cell(1, col)
                    for n, mark in enumerate(marks):
                        paragraph = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
                        paragraph.add_run(mark)
                        _underlined_tab(paragraph)
                fields = self.inspect(docx_bytes(doc))["fields"]
                left = [f["label"] for f in fields if f["locator"]["slot"].startswith("b/tbl[0]/tr[1]/tc[0]")]
                self.assertEqual(len(left), 2, [f["label"] for f in fields])
                self.assertFalse([label for label in left if "SPECIMENS" in label], left)

    # F07 / F08 / F09 ─ the claimed '2)' and '(3)' markers
    def test_f07_f08_f09_parenthesised_markers_carry_their_section_and_number(self):
        for marks in (("(1)", "(2)", "(3)"), ("1)", "2)", "3)")):
            with self.subTest(marks=marks):
                self.assertEqual(self.labels(numbered_slots_bytes(marks=marks)),
                                 ["POSTOPERATIVE FINDINGS — %d" % n for n in (1, 2, 3)]
                                 + ["RECOMMENDATIONS — %d" % n for n in (1, 2, 3)])

    # M2 / M3 ─ two rows that each print a question share the merged cell, whatever their last label cell says
    def merged_rows(self, rows, *, extra=None, header=None):
        """`rows`: [[left cells...]] beside one answer column merged over all of them (under an optional `header`
        row); then a separate Plan table."""
        doc = Document()
        width = max(len(r) for r in rows) + 1
        top = 1 if header else 0
        table = doc.add_table(rows=len(rows) + top, cols=width)
        for c, text in enumerate(header or ()):
            table.cell(0, c).text = text
        for r, cells in enumerate(rows):
            for c, text in enumerate(cells):
                if text:
                    table.cell(r + top, c).text = text
        merged = table.cell(top, width - 1).merge(table.cell(top + len(rows) - 1, width - 1))
        if extra:
            extra(merged)
        plan = doc.add_table(rows=1, cols=2)
        plan.cell(0, 0).text = "Plan"
        return docx_bytes(doc)

    def assert_shared(self, data, *, names=()):
        info = self.inspect(data)
        # Only the separate Plan table's own answer cell is mapped. (Its label may carry the header of the table it
        # continues, 'Plan — Effusion findings': pre-existing on b01e9cfa, not EC3.)
        self.assertEqual([f["locator"]["slot"] for f in info["fields"]], ["b/tbl[1]/tr[0]/tc[1]"],
                         [(f["label"], f["locator"]["slot"]) for f in info["fields"]])
        notes = [n for n in info["notes"] if "merged" in n]
        self.assertEqual(len(notes), 1, info["notes"])
        for name in names:
            self.assertIn(name, notes[0])

    def test_m2_rows_with_the_same_last_label_but_different_items_share_the_cell(self):
        self.assert_shared(self.merged_rows([["Right eye", "Visual acuity"], ["Left eye", "Visual acuity"]]),
                           names=("Right eye", "Left eye"))
        self.assert_shared(self.merged_rows([["Pre-procedure", "Blood pressure:"], ["Post-procedure", "Blood pressure"]]),
                           names=("Pre-procedure", "Post-procedure"))

    def test_m3_option_words_beside_a_merged_cell_are_still_two_questions(self):
        self.assert_shared(self.merged_rows([["Left"], ["Right"]], header=("Knee", "Effusion findings")),
                           names=("Left", "Right"))

    # F04 / M10 ─ whatever fills the shared cell (a drawn blank, an underlined tab, a content control), no rule maps it
    def test_f04_m10_a_shared_merged_cell_is_mapped_for_none_whatever_fills_it(self):
        def underscores(cell):
            cell.paragraphs[0].add_run("____________________")

        def tab(cell):
            _underlined_tab(cell.paragraphs[0])

        def control(cell):
            cell.paragraphs[0]._p.append(_sdt("text", placeholder="Click here to enter text."))
        for name, fill in (("underscores", underscores), ("tab", tab), ("control", control)):
            with self.subTest(fill=name):
                data = self.merged_rows([["Current medications"], ["Allergies"]], extra=fill)
                info = self.inspect(data)
                inside = [f for f in info["fields"] if f["locator"]["slot"].startswith("b/tbl[0]/tr[0]/tc[1]")]
                self.assertEqual(inside, [], [(f["label"], f["locator"]["slot"]) for f in inside])
                self.assert_shared(data, names=("Current medications", "Allergies"))

    # F02 ─ the walk stops at the end of its own merge
    def test_f02_a_second_merge_lower_in_the_column_is_its_own_cell(self):
        doc = Document()
        table = doc.add_table(rows=6, cols=2)
        for row, text in ((0, "Current medications"), (2, "Plan"), (3, "Allergies"), (4, "Social history"), (5, "Follow-up")):
            table.cell(row, 0).text = text
        table.cell(0, 1).merge(table.cell(1, 1))
        table.cell(3, 1).merge(table.cell(4, 1))
        info = self.inspect(docx_bytes(doc))
        by = fields_by_label(info["fields"])
        self.assertEqual(by["Current medications"]["locator"]["slot"], "b/tbl[0]/tr[0]/tc[1]")
        self.assertNotIn("Allergies", by)
        notes = [n for n in info["notes"] if "merged" in n]
        self.assertEqual(len(notes), 1, info["notes"])
        self.assertIn("Social history", notes[0])
        self.assertNotIn("Current medications", notes[0])

    # M4 ─ a row whose answer cell continues a merged answer is a question row, never the next row's header
    def test_m5_an_instruction_behind_a_merged_question_cell_is_not_the_next_rows_header(self):
        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "Reason for visit"
        table.cell(0, 0).merge(table.cell(1, 0))
        _grey(table.cell(0, 1).paragraphs[0], "Describe the main complaint.")
        _grey(table.cell(1, 1).paragraphs[0], "Include when it started.")
        table.cell(2, 0).text = "Plan"
        self.assertIn("Plan", self.labels(docx_bytes(doc)))

    def test_f05_a_header_row_ending_in_an_ordinary_cell_keeps_naming_the_columns(self):
        doc = Document()
        table = doc.add_table(rows=3, cols=3)
        table.cell(0, 0).text = "Medication"
        _grey(table.cell(0, 1).paragraphs[0], "Name and strength")
        table.cell(0, 2).text = "Dose"
        table.cell(1, 0).text = "Morning"
        table.cell(2, 0).text = "Evening"
        from . import word_v3 as w
        data = docx_bytes(doc)
        detector = w.Detector(w.open_package(data))
        self.assertEqual(detector.row_kind(next(s for s in detector.index.slots if s.kind == "tbl"), 0), "label")
        # (The grey header cell itself is also mapped as 'Medication' by rule (d): pre-existing on b01e9cfa, not EC4.)
        labels = [label for label in self.labels(data) if label.startswith(("Morning", "Evening"))]
        self.assertEqual(labels, ["Morning — Name and strength", "Morning — Dose", "Evening — Name and strength",
                                  "Evening — Dose"])

    # M7 ─ white header text on a dark fill is ordinary ink; a repeated header row is a header row
    def test_m7_a_dark_header_row_keeps_its_log_table(self):
        from docx.shared import RGBColor
        for header_flag in (True, False):
            with self.subTest(tblHeader=header_flag):
                doc = Document()
                table = doc.add_table(rows=4, cols=2)
                for col, text in enumerate(("Medication", "Dose and frequency")):
                    cell = table.cell(0, col)
                    _shade(cell, "1F3864")
                    run = cell.paragraphs[0].add_run(text)
                    run.bold = True
                    if col == 1:
                        run.font.color.rgb = RGBColor.from_string("FFFFFF")
                if header_flag:
                    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                fields = self.inspect(docx_bytes(doc))["fields"]
                self.assertIn("table_rows", [f["kind"] for f in fields], [(f["label"], f["kind"]) for f in fields])
                # (rule (d) also maps the white header cell as 'Medication', cell_append: pre-existing on b01e9cfa)

    def test_m7_white_column_headers_on_a_dark_fill_still_name_the_rows_below(self):
        from docx.shared import RGBColor
        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        for col, text in enumerate(("Lesion", "Diameter (mm)")):
            cell = table.cell(0, col)
            _shade(cell, "1F3864")
            run = cell.paragraphs[0].add_run(text)
            run.bold = True
            if col == 1:
                run.font.color.rgb = RGBColor.from_string("FFFFFF")
        table.cell(1, 0).text = "Lesion A"
        table.cell(2, 0).text = "Lesion B"
        labels = [label for label in self.labels(docx_bytes(doc)) if label.startswith("Lesion ")]
        self.assertEqual(labels, ["Lesion A — Diameter (mm)", "Lesion B — Diameter (mm)"])

    # The controls the repairs must keep (each measured on a dev form; synthetic twins here)
    def test_ec4_control_a_header_row_with_a_guidance_cell_that_is_no_questions_own_stays_a_header(self):
        """'Drug(s) under Investigation: | IND No.: | Study (Indication)' (the last in guidance ink) over a row of guidance
        answer cells is a header row: its cells name the columns below."""
        from . import word_v3 as w
        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        table.cell(0, 0).text = "Drug under investigation:"
        table.cell(0, 1).text = "IND No.:"
        for row, col, text in ((0, 2, "Study (Indication)"), (1, 0, "Enter drug"), (1, 1, "Enter IND No.")):
            _grey(table.cell(row, col).paragraphs[0], text)
        detector = w.Detector(w.open_package(docx_bytes(doc)))
        tslot = next(s for s in detector.index.slots if s.kind == "tbl")
        self.assertEqual(detector.row_kind(tslot, 0), "label")
        self.assertEqual(detector.column_header(tslot, 1, 1), "IND No.:")

    def test_m4_control_a_second_header_row_under_merged_header_cells_keeps_naming_the_columns(self):
        """'Taken | Completed by' with 'Taken' (and an empty corner cell) merged down over a second header row
        'Initial | Date': the second row continues printed header cells, so it stays a header row, and the 'Initial'
        column keeps its initials areas and their signing role."""
        from . import word_v3 as w
        doc = Document()
        table = doc.add_table(rows=3, cols=4)
        table.cell(0, 0).text = "Date taken"
        table.cell(0, 1).merge(table.cell(0, 2)).text = "Completed by"
        table.cell(0, 0).merge(table.cell(1, 0))
        table.cell(0, 3).merge(table.cell(1, 3))
        table.cell(1, 1).text = "Initial"
        table.cell(1, 2).text = "Date"
        data = docx_bytes(doc)
        fields = w.inspect(data)["fields"]
        kinds = {f["locator"]["slot"]: f["kind"] for f in fields}
        self.assertEqual(kinds.get("b/tbl[0]/tr[2]/tc[1]"), "initials", [(f["label"], f["kind"], f["locator"]["slot"]) for f in fields])
        self.assertIn("signing_area", {r["source"] for r in w.destination_roles(data, fields)["roles"]})

    def test_ec3_declared_cost_a_question_printed_inside_a_shared_merged_cell_is_withdrawn_too(self):
        """Round 3: no exemption. A merged cell that prints its own question ('Primary event (diagnosis): ____') beside
        rows of other questions is withdrawn with the note (declared cost, the D35 shape)."""
        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        for row, text in enumerate(("Onset date: ____", "Onset time: ____", "Occurred at: ____")):
            table.cell(row, 0).text = text
        table.cell(0, 1).merge(table.cell(2, 1)).text = "Primary event (diagnosis): ____________"
        info = self.inspect(docx_bytes(doc))
        self.assertNotIn("Primary event (diagnosis)", [f["label"] for f in info["fields"]])
        self.assertEqual(len([n for n in info["notes"] if "merged" in n]), 1, info["notes"])

    def test_ec5_a_slot_line_styled_as_a_heading_is_never_the_next_slots_section(self):
        """Legacy reads a heading only on a line without an answer tab; so does section_heading."""
        doc = Document()
        for section in ("FINDINGS", "PLAN"):
            doc.add_heading(section, 2)
            for mark in ("1.", "2."):
                _tab_line(doc, mark).style = doc.styles["Heading 3"]
        self.assertEqual(self.labels(docx_bytes(doc)), ["FINDINGS — 1", "FINDINGS — 2", "PLAN — 1", "PLAN — 2"])

    def test_ec3_control_a_signing_area_in_a_shared_merged_cell_keeps_its_role(self):
        """A merged blank beside two signature rows stays a signing place (never withdrawn: its role depends on it)."""
        from . import word_v3 as w
        data = self.merged_rows([["Patient signature"], ["Witness signature"]])
        fields = w.inspect(data)["fields"]
        self.assertIn("signature", [f["kind"] for f in fields], [(f["label"], f["kind"]) for f in fields])
        self.assertIn("signing_area", {r["source"] for r in w.destination_roles(data, fields)["roles"]})

    def test_m4_control_a_sub_header_row_beside_a_merged_printed_header_stays_a_header(self):
        """'Completed by' over 'Initial | Date', beside a printed header 'Site' merged down over both header rows: the
        second row continues a PRINTED cell (to the right of its own labels), so it stays a header row and the
        'Initial' column keeps its initials areas."""
        from . import word_v3 as w
        doc = Document()
        table = doc.add_table(rows=3, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Completed by"
        table.cell(0, 2).text = "Site"
        table.cell(0, 2).merge(table.cell(1, 2))
        table.cell(1, 0).text = "Initial"
        table.cell(1, 1).text = "Date"
        data = docx_bytes(doc)
        fields = w.inspect(data)["fields"]
        kinds = {f["locator"]["slot"]: f["kind"] for f in fields}
        self.assertEqual(kinds.get("b/tbl[0]/tr[2]/tc[0]"), "initials", [(f["label"], f["kind"], f["locator"]["slot"]) for f in fields])

    def test_ec5_a_section_by_outline_level_is_a_word_heading_as_on_legacy(self):
        """documents.tab_form._heading also reads an outline level on the paragraph, or in its style's ancestry."""
        from docx.enum.style import WD_STYLE_TYPE
        for how in ("paragraph", "style"):
            with self.subTest(how=how):
                doc = Document()
                section = doc.add_paragraph("FINDINGS")
                level = OxmlElement("w:outlineLvl")
                level.set(qn("w:val"), "1")
                if how == "paragraph":
                    section._p.get_or_add_pPr().append(level)
                else:
                    style = doc.styles.add_style("Section Title", WD_STYLE_TYPE.PARAGRAPH)
                    style.element.get_or_add_pPr().append(level)
                    section.style = style
                for mark in ("1.", "2."):
                    _tab_line(doc, mark)
                self.assertEqual(self.labels(docx_bytes(doc)), ["FINDINGS — 1", "FINDINGS — 2"])

    def test_declared_residual_ec1_the_finalizers_own_signer_row_stays_hand_signed_on_v3(self):
        """EC1 withdrawn (refutation of 358cb5b0): on contract v3 the owner colonoscopy form's 'Electronically signed by'
        line is a Signature area signed by hand, never filled - not even with the finalizer's own name, which the legacy
        engine accepts there (core.tests.test_destination_roles test_pc14). Declared residual, owner-routed."""
        info = inspect_template(OWNER_FORM, engine="v3")
        signer = fields_by_label(info["fields"])["Electronically signed by"]
        self.assertEqual(signer["kind"], "signature")
        self.assertNotIn("answer_policy", signer)
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(DocumentError) as refused:
                fill_template(OWNER_FORM, info["fields"], {signer["key"]: "Synthetic Alpha Clinician"},
                              Path(directory) / "out.docx")
        self.assertEqual(refused.exception.code, "invalid_answer")


# ─── round 3: the refutation of c1cae13d (refute-wec-r2) - readers removed, not refined ───────────────────────────────

def _body_append(doc, xml):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    element = parse_xml(xml.replace("NS", nsdecls("w")))
    doc.element.body.insert(len(doc.element.body) - 1, element)
    return element


class RoundThree(unittest.TestCase):
    """Each round-2 finding as a pinned input, against the round-3 rules: EC3 withdraws every answer area in a shared
    merged cell (no exemption); EC4 keeps only the question+guidance pairing (white ink never guidance); EC5 names a
    numbered slot only by an unambiguous section, else keeps the bare marker."""
    inspect = RefutationRepairs.inspect
    labels = RefutationRepairs.labels
    merged_rows = RefutationRepairs.merged_rows
    assert_shared = RefutationRepairs.assert_shared

    def shared_guidance(self, rows, guidance):
        def fill(cell):
            _grey(cell.paragraphs[0], guidance)
        return self.merged_rows(rows, extra=fill)

    # WEC2-m-1 / R2F-01 HIGH ─ guidance in a shared cell never keeps it for the first question
    def test_r2f01_guidance_mentioning_the_first_question_never_keeps_the_shared_cell(self):
        for rows, guidance in (([["Current medications"], ["Allergies"]], "List current medications and allergies."),
                               ([["Pain"], ["Swelling"]], "Note any painful movement or swelling."),
                               ([["Rate"], ["Rhythm"]], "Record separately."),
                               ([["Plan"], ["Follow-up"]], "Describe the plan and the follow-up interval."),
                               ([["Right eye", "Visual acuity"], ["Left eye", "Visual acuity"]],
                                "Record the visual acuity of each eye.")):
            with self.subTest(guidance=guidance):
                self.assert_shared(self.shared_guidance(rows, guidance))

    # WEC2-m-2 ─ a continued label is read as rule (d) reads it
    def test_m2_a_continued_label_counts_as_its_question(self):
        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "Current medications"
        table.cell(0, 0).merge(table.cell(1, 0))
        table.cell(0, 1).text = "Name, dose and frequency of each:"
        table.cell(1, 1).merge(table.cell(2, 1))
        table.cell(2, 0).text = "Allergies"
        plan = doc.add_table(rows=1, cols=2)
        plan.cell(0, 0).text = "Plan"
        info = self.inspect(docx_bytes(doc))
        self.assertFalse([f for f in info["fields"] if f["locator"]["slot"].startswith("b/tbl[0]/tr[1]/tc[1]")],
                         [(f["label"], f["locator"]["slot"]) for f in info["fields"]])
        notes = [n for n in info["notes"] if "merged" in n]
        self.assertEqual(len(notes), 1, info["notes"])
        self.assertIn("Current medications", notes[0])
        self.assertIn("Allergies", notes[0])

    # WEC2-m-3 ─ a shared cell printing its own question is withdrawn, never side-qualified from one row
    def test_m3_a_question_printed_in_a_cell_shared_by_two_sides_is_withdrawn(self):
        for content in ("Effusion: ____________", "Effusion:  ☐ Yes   ☐ No"):
            with self.subTest(content=content):
                def fill(cell, content=content):
                    cell.paragraphs[0].add_run(content)
                data = self.merged_rows([["Right knee"], ["Left knee"]], extra=fill)
                labels = [f["label"] for f in self.inspect(data)["fields"]]
                self.assertFalse([label for label in labels if "knee" in label], labels)
                self.assert_shared(data, names=("Right knee", "Left knee"))

    # WEC2-m-5 ─ a shared cell whose only area is a signing area still gets its note; the role is kept
    def test_m5_a_signing_only_shared_cell_is_noted_and_keeps_its_role(self):
        from . import word_v3 as w
        data = self.merged_rows([["Clinician signature"], ["Comments to the GP"]])
        info = w.inspect(data)
        self.assertEqual(len([n for n in info["notes"] if "merged" in n]), 1, info["notes"])
        self.assertIn("signing_area", {r["source"] for r in w.destination_roles(data, info["fields"])["roles"]})

    # R2F-10 FA ─ the row count reads the cells LEFT of the merge, as rule (d) reads a row's question (WEC2-m-6, a side
    # printed to the right, is recorded, not changed)
    def test_r2f10_fa_a_column_to_the_right_of_a_tall_answer_cell_is_no_second_question(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        table.cell(0, 0).text = "Current medications"
        table.cell(0, 2).text = "Reviewed"
        table.cell(1, 2).text = "Date"
        table.cell(0, 1).merge(table.cell(1, 1))
        info = self.inspect(docx_bytes(doc))
        self.assertIn("b/tbl[0]/tr[0]/tc[1]", [f["locator"]["slot"] for f in info["fields"]], [(f["label"], f["locator"]["slot"]) for f in info["fields"]])
        self.assertFalse([n for n in info["notes"] if "merged" in n], info["notes"])

    # R2F-05 / R2F-10 FG ─ initials and protected areas in a shared cell are kept with their roles
    def test_r2f05_initials_in_a_shared_cell_keep_their_signing_role(self):
        from . import word_v3 as w
        data = self.merged_rows([["Nurse initials"], ["Doctor initials"]])
        fields = w.inspect(data)["fields"]
        self.assertIn("initials", [f["kind"] for f in fields])
        self.assertIn("signing_area", {r["source"] for r in w.destination_roles(data, fields)["roles"]})

    def test_r2f10_fg_a_locked_control_in_a_shared_cell_keeps_its_protected_role(self):
        from . import word_v3 as w

        def locked(cell):
            sdt = _sdt("text", placeholder="Click here to enter text.")
            lock = OxmlElement("w:lock")
            lock.set(qn("w:val"), "sdtContentLocked")
            sdt.find(qn("w:sdtPr")).append(lock)
            cell.paragraphs[0]._p.append(sdt)
        data = self.merged_rows([["Current medications"], ["Allergies"]], extra=locked)
        fields = w.inspect(data)["fields"]
        sources = {r["source"] for r in w.destination_roles(data, fields)["roles"]}
        self.assertIn("protected_area", sources, sources)
        # the protected field itself is kept in the mapping (shown as protected), not withdrawn
        self.assertTrue([f for f in fields if f.get("write_protected") and f["locator"]["slot"].startswith("b/tbl[0]/tr[0]/tc[1]")],
                        [(f["label"], f["locator"]["slot"], f.get("write_protected")) for f in fields])

    # R2F-10 FB ─ a blank-text row ('dd/mm/yyyy') prints no question
    def test_r2f10_fb_a_date_format_row_is_no_second_question(self):
        info = self.inspect(self.merged_rows([["Date of injury"], ["dd/mm/yyyy"]]))
        self.assertIn("Date of injury", [f["label"] for f in info["fields"]])

    # R2F-11 FW ─ one note per cell, however many areas it held
    def test_r2f11_fw_one_note_per_shared_cell(self):
        def two(cell):
            cell.paragraphs[0]._p.append(_sdt("text", placeholder="Click here to enter text."))
            cell.add_paragraph()._p.append(_sdt("text", placeholder="Click here to enter text."))
        info = self.inspect(self.merged_rows([["Current medications"], ["Allergies"]], extra=two))
        self.assertEqual(len([n for n in info["notes"] if "merged" in n]), 1, info["notes"])

    # P-1 ─ a withdrawal never moves a neighbouring signing role over the withdrawn place
    def test_p1_a_withdrawn_control_stays_an_ordinary_place_beside_a_signature(self):
        def control_and_signature(cell):
            paragraph = cell.paragraphs[0]
            paragraph._p.append(_sdt("text", placeholder="Click here to enter text."))
            paragraph.add_run("   Signature: ______________")
        data = self.merged_rows([["Current medications"], ["Allergies"]], extra=control_and_signature)
        path = Path(tempfile.mkdtemp()) / "p1.docx"
        path.write_bytes(data)
        slot = next(r["slot"] for r in outline_template(path)["slots"] if r["kind"] == "sdt")
        added = locate_fields(path, [{"slot": slot, "kind": "text", "label": "Current medications"}])["fields"]
        self.assertEqual(len(added), 1)
        path.unlink()

    # WR2H-1 HIGH ─ one-letter question cells are questions (no numbering column reader)
    def test_wr2h1_one_letter_question_rows_keep_their_own_questions(self):
        doc = Document()
        table = doc.add_table(rows=4, cols=2)
        for row, (letter, prompt) in enumerate((("S", "Presenting complaint in the patient's words."),
                                                ("O", "Examination findings."), ("A", "Assessment."), ("P", None))):
            table.cell(row, 0).text = letter
            if prompt:
                _grey(table.cell(row, 1).paragraphs[0], prompt)
        labels = self.labels(docx_bytes(doc))
        self.assertEqual(labels, ["S", "O", "A", "P"])

    def test_wr2h6_a_numbered_header_row_stays_a_header_row(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        table.cell(0, 0).text = "1."
        table.cell(0, 1).text = "Drug under investigation:"
        _grey(table.cell(0, 2).paragraphs[0], "Study (Indication)")
        from . import word_v3 as w
        detector = w.Detector(w.open_package(docx_bytes(doc)))
        self.assertEqual(detector.row_kind(next(s for s in detector.index.slots if s.kind == "tbl"), 0), "label")

    # WR2H-2 HIGH ─ a heading the slots cannot be sure of (in a table, a control, a text box) leaves the bare marker
    def test_wr2h2_a_heading_in_a_table_control_or_text_box_never_lets_an_earlier_section_name_the_slots(self):
        def banner(doc):
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).paragraphs[0].style = doc.styles["Heading 2"]
            table.cell(0, 0).paragraphs[0].add_run("PLAN")

        def control(doc):
            _body_append(doc, '<w:sdt NS><w:sdtPr><w:lock w:val="sdtContentLocked"/></w:sdtPr><w:sdtContent><w:p><w:pPr>'
                              '<w:pStyle w:val="Heading2"/></w:pPr><w:r><w:t>PLAN</w:t></w:r></w:p></w:sdtContent></w:sdt>')

        def text_box(doc):
            _body_append(doc, '<w:p NS xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict><v:shape style="width:200pt;'
                              'height:30pt"><v:textbox><w:txbxContent><w:p><w:pPr><w:pStyle w:val="Heading2"/></w:pPr>'
                              '<w:r><w:t>PLAN</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>')
        for name, place in (("table", banner), ("control", control), ("text box", text_box)):
            with self.subTest(heading_in=name):
                doc = Document()
                doc.add_heading("HISTORY", 2)
                place(doc)
                for mark in ("1.", "2."):
                    _tab_line(doc, mark)
                labels = self.labels(docx_bytes(doc))
                self.assertFalse([label for label in labels if "HISTORY" in label], labels)
                self.assertEqual([label for label in labels if label in ("1.", "2.")], ["1.", "2."], labels)

    # WR2H-3 ─ a repeated header row that is a question + instruction row is a question row
    def test_wr2h3_a_tblheader_question_row_is_not_the_next_rows_header(self):
        data = instruction_cell_bytes()
        from docx import Document as Load
        doc = Load(io.BytesIO(data))
        doc.tables[0].rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        labels = self.labels(docx_bytes(doc))
        self.assertIn("Plan", labels)
        self.assertFalse([label for label in labels if "Describe" in label], labels)

    # WR2H-4 / R2F-02 ─ white header ink never pairs as an instruction, whatever and wherever the fill
    def test_wr2h4_r2f02_white_header_ink_keeps_the_log_table_on_any_fill(self):
        from docx.shared import RGBColor
        for fill, style in (("5B9BD5", None), ("ED7D31", None), ("4472C4", None), (None, "Medium Shading 1 Accent 1")):
            with self.subTest(fill=fill, style=style):
                doc = Document()
                table = doc.add_table(rows=4, cols=2)
                if style:
                    table.style = doc.styles[style]
                for col, text in enumerate(("Medication", "Dose and frequency")):
                    cell = table.cell(0, col)
                    if fill:
                        _shade(cell, fill)
                    run = cell.paragraphs[0].add_run(text)
                    run.bold = True
                    if col == 1:
                        run.font.color.rgb = RGBColor.from_string("FFFFFF")
                fields = self.inspect(docx_bytes(doc))["fields"]
                self.assertIn("table_rows", [f["kind"] for f in fields], [(f["label"], f["kind"]) for f in fields])

    def test_white_ink_by_theme_or_near_white_never_pairs_as_an_instruction(self):
        from docx.shared import RGBColor
        for ink in ("theme", "F2F2F2"):
            with self.subTest(ink=ink):
                doc = Document()
                table = doc.add_table(rows=4, cols=2)
                for col, text in enumerate(("Medication", "Dose and frequency")):
                    cell = table.cell(0, col)
                    _shade(cell, "1F3864")
                    run = cell.paragraphs[0].add_run(text)
                    if col == 1:
                        if ink == "theme":
                            color = OxmlElement("w:color")
                            color.set(qn("w:val"), "FFFFFF")
                            color.set(qn("w:themeColor"), "background1")
                            run._r.get_or_add_rPr().append(color)
                        else:
                            run.font.color.rgb = RGBColor.from_string(ink)
                fields = self.inspect(docx_bytes(doc))["fields"]
                self.assertIn("table_rows", [f["kind"] for f in fields], [(f["label"], f["kind"]) for f in fields])

    def test_a_signing_area_kept_in_a_cell_beside_two_sides_is_never_side_qualified(self):
        def signature(cell):
            cell.paragraphs[0].add_run("Signature: ____________")
        data = self.merged_rows([["Right knee"], ["Left knee"]], extra=signature)
        fields = self.inspect(data)["fields"]
        kept = [f for f in fields if f["kind"] == "signature" and f["locator"]["slot"].startswith("b/tbl[0]/")]
        self.assertEqual([f["label"] for f in kept], ["Signature"], [(f["label"], f["kind"], f["locator"]["slot"]) for f in fields])

    def test_a_heading_inside_the_slots_own_table_cell_leaves_the_bare_marker(self):
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].style = doc.styles["Heading 2"]
        cell.paragraphs[0].add_run("FINDINGS")
        for mark in ("1.", "2."):
            paragraph = cell.add_paragraph()
            paragraph.add_run(mark)
            _underlined_tab(paragraph)
        self.assertFalse([label for label in self.labels(docx_bytes(doc)) if "—" in label])

    def test_wr2h4_grey_guidance_on_a_dark_row_is_still_guidance(self):
        from docx.shared import RGBColor
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        for col in (0, 1):
            _shade(table.cell(0, col), "1F3864")
        table.cell(0, 0).paragraphs[0].add_run("Reason for visit").font.color.rgb = RGBColor.from_string("FFFFFF")
        table.cell(0, 1).paragraphs[0].add_run("Describe the main complaint.").font.color.rgb = RGBColor.from_string("BFBFBF")
        table.cell(1, 0).text = "Plan"
        labels = self.labels(docx_bytes(doc))
        self.assertIn("Plan", labels)
        self.assertIn("Reason for visit", labels)

    # WR2H-5 ─ an empty label box merged beside a log's header never drops the log (M4 is dropped)
    def test_wr2h5_an_empty_merged_label_box_keeps_the_log_table(self):
        doc = Document()
        doc.add_heading("Synthetic drug chart", 2)
        table = doc.add_table(rows=6, cols=3)
        table.cell(0, 0).text = "Ward"
        table.cell(1, 0).text = "Consultant"
        table.cell(2, 0).text = "Medication"
        table.cell(2, 1).text = "Dose"
        table.cell(0, 2).merge(table.cell(2, 2))
        fields = self.inspect(docx_bytes(doc))["fields"]
        self.assertIn("table_rows", [f["kind"] for f in fields], [(f["label"], f["kind"]) for f in fields])

    # WR2H-7 / R2F-11 FV / FU ─ the outline level Word reads decides; the default style never makes a heading
    def test_wr2h7_a_style_that_sets_body_text_is_no_heading(self):
        from docx.enum.style import WD_STYLE_TYPE
        doc = Document()
        doc.add_heading("HISTORY", 2)
        style = doc.styles.add_style("QuestionLabel", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Heading 3"]
        level = OxmlElement("w:outlineLvl")
        level.set(qn("w:val"), "9")
        style.element.get_or_add_pPr().append(level)
        doc.add_paragraph("Known allergies", style=style)
        for mark in ("1.", "2."):
            _tab_line(doc, mark)
        labels = self.labels(docx_bytes(doc))
        self.assertFalse([label for label in labels if "—" in label], labels)

    def test_a_paragraph_outline_level_of_body_text_overrides_a_heading_style(self):
        doc = Document()
        heading = doc.add_heading("FINDINGS", 2)
        level = OxmlElement("w:outlineLvl")
        level.set(qn("w:val"), "9")
        heading._p.get_or_add_pPr().append(level)
        for mark in ("1.", "2."):
            _tab_line(doc, mark)
        self.assertFalse([label for label in self.labels(docx_bytes(doc)) if "—" in label])

    def test_r2f11_fv_a_default_style_with_an_outline_level_makes_no_heading(self):
        doc = Document()
        level = OxmlElement("w:outlineLvl")
        level.set(qn("w:val"), "0")
        doc.styles["Normal"].element.get_or_add_pPr().append(level)
        doc.add_paragraph("FINDINGS")
        for mark in ("1.", "2."):
            _tab_line(doc, mark)
        self.assertFalse([label for label in self.labels(docx_bytes(doc)) if "—" in label])

    def test_r2f11_fu_a_lower_case_heading_style_id_is_a_heading(self):
        from docx.enum.style import WD_STYLE_TYPE
        doc = Document()
        style = doc.styles.add_style("heading9x", WD_STYLE_TYPE.PARAGRAPH)
        style.element.set(qn("w:styleId"), "heading2")
        doc.add_paragraph("FINDINGS", style=style)
        for mark in ("1.", "2."):
            _tab_line(doc, mark)
        self.assertEqual(self.labels(docx_bytes(doc)), ["FINDINGS — 1", "FINDINGS — 2"])

    # R2F-04 / R2F-07 ─ a heading line that holds an answer (a tab, a drawn blank) is a question: bare markers
    def test_r2f04_r2f07_a_heading_line_holding_an_answer_never_names_the_slots(self):
        for answer in ("tab", "underscores"):
            with self.subTest(answer=answer):
                doc = Document()
                doc.add_heading("FINDINGS", 2)
                line = doc.add_paragraph("Diagnosis: " + ("____________" if answer == "underscores" else ""),
                                         style=doc.styles["Heading 3"])
                if answer == "tab":
                    _underlined_tab(line)
                for mark in ("1.", "2."):
                    _tab_line(doc, mark)
                labels = self.labels(docx_bytes(doc))
                self.assertFalse([label for label in labels if label.endswith(("— 1", "— 2"))], labels)

    # R2F-06 ─ slots in a table cell take no body heading
    def test_r2f06_slots_in_a_table_cell_never_take_the_body_heading_before_the_table(self):
        doc = Document()
        doc.add_heading("HISTORY", 2)
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Current medications"
        for n, mark in enumerate(("1.", "2.")):
            paragraph = table.cell(0, 1).paragraphs[0] if n == 0 else table.cell(0, 1).add_paragraph()
            paragraph.add_run(mark)
            _underlined_tab(paragraph)
        self.assertFalse([label for label in self.labels(docx_bytes(doc)) if "HISTORY" in label])

    # R2F-10 FT ─ empty paragraphs between a heading and its slots are no ambiguity
    def test_r2f10_ft_an_empty_heading_paragraph_is_skipped(self):
        doc = Document()
        doc.add_heading("FINDINGS", 2)
        doc.add_heading("", 2)
        for mark in ("1.", "2."):
            _tab_line(doc, mark)
        self.assertEqual(self.labels(docx_bytes(doc)), ["FINDINGS — 1", "FINDINGS — 2"])

    # R2F-12 ─ M5: a merged-down question cell with a box, or a continuation with no restart above, is the question
    def test_r2f12_a_boxed_or_unrestarted_merged_question_cell_is_still_the_rows_question(self):
        for variant in ("box", "no_restart"):
            with self.subTest(variant=variant):
                doc = Document()
                table = doc.add_table(rows=3, cols=2)
                table.cell(0, 0).text = "☐ Smoker" if variant == "box" else "Reason for visit"
                if variant == "box":
                    table.cell(0, 0).merge(table.cell(1, 0))
                else:
                    vmerge = OxmlElement("w:vMerge")
                    table.cell(1, 0)._tc.get_or_add_tcPr().append(vmerge)
                _grey(table.cell(0, 1).paragraphs[0], "Packs per day.")
                _grey(table.cell(1, 1).paragraphs[0], "Years smoked.")
                table.cell(2, 0).text = "Plan"
                labels = self.labels(docx_bytes(doc))
                self.assertIn("Plan", labels)
                self.assertFalse([label for label in labels if "Years smoked" in label], labels)
