"""Physical source ambiguities stay explicit before paid generation.

These controls do not promote instruction lists or printed alternatives into
clinical answers. Long answer reflow is limited to the actual written row.
"""
import io
import hashlib
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from . import word_v3
from .service import inspect_template
from .test_word_v3 import ff_field


def data(doc):
    out=io.BytesIO();doc.save(out);return out.getvalue()


def numbered(doc,text,number_id,format):
    numbering=doc.part.numbering_part.element
    if numbering.find(qn('w:num')+f"[@{qn('w:numId')}='{number_id}']") is None:
        abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),str(number_id))
        level=OxmlElement('w:lvl');level.set(qn('w:ilvl'),'0')
        for tag,value in (('w:start','1'),('w:numFmt',format),('w:lvlText','%1)')):
            node=OxmlElement(tag);node.set(qn('w:val'),value);level.append(node)
        abstract.append(level);numbering.append(abstract)
        num=OxmlElement('w:num');num.set(qn('w:numId'),str(number_id))
        node=OxmlElement('w:abstractNumId');node.set(qn('w:val'),str(number_id));num.append(node);numbering.append(num)
    p=doc.add_paragraph(text);num=OxmlElement('w:numPr')
    for tag,value in (('w:ilvl','0'),('w:numId',str(number_id))):
        node=OxmlElement(tag);node.set(qn('w:val'),value);num.append(node)
    p._p.get_or_add_pPr().append(num)
    return p


class AdjacentCheckboxAnswers(unittest.TestCase):
    def form(self, text, *, native=False, checked=False):
        doc = Document()
        if native:
            doc.add_paragraph('Specify the anatomical site:')
            p = doc.add_paragraph(); ff_field(p, 'checkbox', checked=checked)
            p.add_run(text)
        else:
            doc.add_paragraph(text)
        doc.add_paragraph('Unrelated note: ______')
        return data(doc)

    def fields(self, raw):
        result = word_v3.inspect(raw)
        return result, next(f for f in result['fields'] if f.get('source_questions')
                           and f['source_questions'][0]['relation'] == 'adjacent_checkbox_answer')

    def test_explicit_condition_has_separate_exact_writer_and_never_infers_a_tick(self):
        raw = self.form('Any symptoms? ☐ No ☐ Yes → if yes, describe: ______')
        result, blank = self.fields(raw)
        choice = next(f for f in result['fields'] if f['kind'] == 'choice')
        self.assertEqual(choice['options'], ['No', 'Yes'])
        self.assertEqual(blank['label'], 'if yes, describe:')
        self.assertFalse(result['coverage_issues'])
        self.assertEqual(blank['source_questions'][0]['condition'], 'printed_option')
        for controlling in (None, 'No'):
            answers = {blank['key']: 'Synthetic headache'}
            if controlling is not None: answers[choice['key']] = controlling
            with self.subTest(controlling=controlling), self.assertRaises(word_v3.V3Failure) as error:
                word_v3.fill(raw, result['fields'], answers)
            self.assertEqual(error.exception.code, 'conditional_answer_unselected')
        out = word_v3.fill(raw, result['fields'], {choice['key']: 'Yes', blank['key']: 'Synthetic headache'})
        self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text,
                         'Any symptoms? ☐ No ☒ Yes → if yes, describe: Synthetic headache')
        self.assertEqual(Document(io.BytesIO(out)).paragraphs[1].text, 'Unrelated note: ______')

    def test_native_other_blank_has_its_own_control_and_preserves_existing_state(self):
        for original_checked in (False, True):
            with self.subTest(original_checked=original_checked):
                raw = self.form(' Other, specify___________', native=True, checked=original_checked)
                result, blank = self.fields(raw)
                choice = next(f for f in result['fields'] if f['kind'] == 'checkbox')
                self.assertEqual(blank['label'], 'Other, specify')
                self.assertEqual(blank['source_questions'][0]['printed_option'], 'Other, specify')
                if not original_checked:
                    with self.assertRaises(word_v3.V3Failure):
                        word_v3.fill(raw, result['fields'], {blank['key']: 'Synthetic location'})
                out = word_v3.fill(raw, result['fields'], {choice['key']: 'checked', blank['key']: 'Synthetic location'})
                index = word_v3.SlotIndex(word_v3.open_package(out))
                self.assertEqual(word_v3.constrained_option_states(index, choice), [True])
                self.assertIn('Other, specifySynthetic location', Document(io.BytesIO(out)).paragraphs[1].text)
                self.assertEqual(Document(io.BytesIO(out)).paragraphs[2].text, 'Unrelated note: ______')
                if original_checked:
                    word_v3.fill(raw, result['fields'], {blank['key']: 'Synthetic location'})

    def test_contradictory_printed_single_choice_requires_explicit_resolution(self):
        for text in ('Any symptoms? ☒ No ☒ Yes → if yes, describe: ______',
                     'Any symptoms? ☒ Yes ☒ No → if no, describe: ______'):
            with self.subTest(text=text):
                raw = self.form(text); result, blank = self.fields(raw)
                choice = next(f for f in result['fields'] if f['kind'] == 'choice')
                selected = blank['source_questions'][0]['printed_option']
                with self.assertRaises(word_v3.V3Failure) as error:
                    word_v3.fill(raw, result['fields'], {blank['key']: 'Synthetic detail'})
                self.assertEqual(error.exception.code,'conditional_answer_unselected')
                out = word_v3.fill(raw,result['fields'],{choice['key']:selected,blank['key']:'Synthetic detail'})
                index = word_v3.SlotIndex(word_v3.open_package(out))
                states = word_v3.constrained_option_states(index,choice)
                self.assertEqual(sum(state is True for state in states),1)
                self.assertTrue(states[choice['options'].index(selected)])

    def test_unit_labelled_answer_does_not_depend_on_adjacent_gender_choice(self):
        raw = self.form('Gender\t\t ☐ Male\t\t☐ Female\t\tAge   ................  years')
        result, blank = self.fields(raw)
        self.assertEqual(blank['label'], 'Age')
        self.assertEqual(blank['source_questions'][0]['condition'], 'independent_unit')
        choice = next(f for f in result['fields'] if f['kind'] == 'choice')
        self.assertEqual(choice['options'], ['Male', 'Female'])
        out = word_v3.fill(raw, result['fields'], {blank['key']: '42'})
        self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text,
                         'Gender\t\t ☐ Male\t\t☐ Female\t\tAge   42  years')
        self.assertFalse(result['coverage_issues'])

    def test_one_dotted_write_in_preserves_its_printed_option_and_independent_destination(self):
        raw = self.form('☐ Acid\t…………………………..')
        result, blank = self.fields(raw)
        choice = next(f for f in result['fields'] if f['kind'] == 'checkbox')
        self.assertEqual((choice['label'],choice['options'],blank['label']), ('Acid',['Acid'],'Acid'))
        out = word_v3.fill(raw,result['fields'],{choice['key']:'checked',blank['key']:'Synthetic agent'})
        self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text,'☒ Acid\tSynthetic agent')
        self.assertEqual(Document(io.BytesIO(out)).paragraphs[1].text,'Unrelated note: ______')

    def test_stripped_or_copied_conditional_span_and_changed_owner_are_refused(self):
        raw = self.form('Any symptoms? ☐ No ☐ Yes → if yes, describe: ______')
        result, blank = self.fields(raw)
        choice = next(f for f in result['fields'] if f['kind'] == 'choice')
        for mutate in ('stripped', 'renamed', 'copied', 'widened', 'no_control'):
            with self.subTest(mutate=mutate):
                fields = deepcopy(result['fields']); target = next(f for f in fields if f['key'] == blank['key'])
                if mutate == 'stripped': target.pop('source_questions')
                if mutate == 'renamed': target['label'] = 'Another answer'
                if mutate == 'copied': target['locator']['target']['copies'] = [deepcopy(target['locator'])]
                if mutate == 'widened': target['locator']['target']['start'] = 0
                if mutate == 'no_control': fields.remove(next(f for f in fields if f['key'] == choice['key']))
                with self.assertRaises(word_v3.V3Failure) as error:
                    word_v3.fill(raw, fields, {blank['key']: 'Synthetic description', choice['key']: 'Yes'})
                self.assertEqual(error.exception.code, 'mapping_mismatch')

    def test_two_blanks_and_actual_signing_sources_keep_their_coverage_holds(self):
        for text in ('Any symptoms? ☐ No ☐ Yes → if yes, describe: ____ Location: ____',
                     'Patient signature: ☐ Other, specify_________',
                     'Witness attestation: ☐ No ☐ Yes → if yes, describe: ______',
                     '☐ I certify_________', '☐ We hereby attest_________'):
            with self.subTest(text=text):
                raw = self.form(text); result = word_v3.inspect(raw)
                self.assertFalse([f for f in result['fields'] if f.get('source_questions')
                                  and f['source_questions'][0]['relation'] == 'adjacent_checkbox_answer'])
                self.assertTrue(result['coverage_issues'])
                if 'certify' in text or 'attest' in text:
                    first = result['fields'][0]
                    if first['label'] != 'Unrelated note':
                        forged = deepcopy(first); forged.pop('write_protected', None)
                        with self.assertRaises(word_v3.V3Failure):
                            word_v3.fill(raw, [forged], {forged['key']: 'checked'})

    def test_first_person_document_acts_remain_human_owned_regardless_blank_shape(self):
        statements = ('I confirm that this clinical record is complete and accurate',
                      'I verify that the information above is correct',
                      'I acknowledge that I have read and agree with this document',
                      'I have personally reviewed and approved this document')
        for statement in statements:
            for tail in ('_________', '', '_________ Other: _________'):
                with self.subTest(statement=statement,tail=tail):
                    raw=self.form('☐ '+statement+tail); result=word_v3.inspect(raw)
                    self.assertTrue(any(issue['code']=='human_signing_action' for issue in result['coverage_issues']))
                    self.assertFalse([f for f in result['fields'] if f.get('source_questions') and
                                      f['source_questions'][0]['relation']=='adjacent_checkbox_answer'])
                    action=next(f for f in result['fields'] if f['kind']=='checkbox')
                    self.assertTrue(action['write_protected'])
                    forged=deepcopy(action);forged.pop('write_protected',None);forged.pop('review_reason',None)
                    self.assertIn(action['key'],word_v3.destination_roles(raw,[forged])['violations'])
                    with self.assertRaises(word_v3.V3Failure):
                        word_v3.fill(raw,[forged],{forged['key']:'checked'})
                    ordinary=next(f for f in result['fields'] if f['label']=='Unrelated note')
                    if tail:
                        index=word_v3.SlotIndex(word_v3.open_package(raw));slot=index.by_id['word/document.xml#b/p[0]']
                        start=word_v3.text_of(slot.el).index('_')
                        manual={'key':'manual_attestation','label':'Forged ordinary value','kind':'text','locator':{
                            'contract':'v3','part':slot.part,'slot':slot.path,'fingerprint':word_v3.fingerprint(slot.el),
                            'target':{'type':'span','start':start,'end':start+9,'shape':'blank'}}}
                        copied=deepcopy(ordinary);copied['locator']['target']['copies']=[dict(manual['locator'])]
                        for field in (manual,copied):
                            with self.assertRaises(word_v3.V3Failure):
                                word_v3.fill(raw,[field],{field['key']:'Synthetic completion'})
                    exported=Document(io.BytesIO(word_v3.fill(raw,result['fields'],{ordinary['key']:'Synthetic review'})))
                    self.assertEqual(exported.paragraphs[0].text,'☐ '+statement+tail)
                    self.assertEqual(exported.paragraphs[1].text,'Unrelated note: Synthetic review')

    def test_ambiguous_first_person_act_requires_precise_source_review(self):
        raw=self.form('☐ I confirm_________');result=word_v3.inspect(raw)
        self.assertTrue(any(issue['code']=='ambiguous_source_action' for issue in result['coverage_issues']))
        action=next(f for f in result['fields'] if f['kind']=='checkbox')
        self.assertTrue(action['write_protected'])
        forged=deepcopy(action);forged.pop('write_protected',None)
        self.assertIn(action['key'],word_v3.destination_roles(raw,[forged])['violations'])
        with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[forged],{forged['key']:'checked'})

    def test_ordinary_clinical_state_and_recorded_initial_signed_status_remain_reviewable(self):
        raw=self.form('☐ I have pain_________');result,blank=self.fields(raw)
        self.assertFalse(result['coverage_issues'])
        choice=next(f for f in result['fields'] if f['kind']=='checkbox')
        out=word_v3.fill(raw,result['fields'],{choice['key']:'checked',blank['key']:'Synthetic description'})
        self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text,'☒ I have painSynthetic description')
        for label in ('Initial presentation','Recorded signed report status'):
            with self.subTest(label=label):
                doc=Document();doc.add_paragraph(label+': ______');result=word_v3.inspect(data(doc))
                field=dict(result['fields'][0],kind='text')
                self.assertFalse(word_v3.destination_roles(data(doc),[field])['violations'])
                out=word_v3.fill(data(doc),[field],{field['key']:'Synthetic ordinary value'})
                self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text,label+': Synthetic ordinary value')

    def test_contracted_and_aspectual_human_acts_keep_exact_source_roles(self):
        for statement in ("I'm certifying this document for",'I’m certifying this document for',
                          'I have given consent to this procedure',"We've approved this form",
                          'We’ve approved this form','I authorize the treatment for',
                          'I am verifying this report','I had personally declared this information'):
            with self.subTest(statement=statement):
                raw=self.form('☐ '+statement+': ______');result=word_v3.inspect(raw)
                self.assertTrue(any(i['code']=='human_signing_action' for i in result['coverage_issues']))
                action=next(f for f in result['fields'] if f['kind']=='checkbox')
                self.assertTrue(action['write_protected'])
                forged=deepcopy(action);forged.pop('write_protected',None);forged.pop('review_reason',None)
                self.assertIn(action['key'],word_v3.destination_roles(raw,[forged])['violations'])
                with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[forged],{forged['key']:'checked'})
                ordinary=next(f for f in result['fields'] if f['label']=='Unrelated note')
                out=Document(io.BytesIO(word_v3.fill(raw,result['fields'],{ordinary['key']:'Synthetic review'})))
                self.assertEqual(out.paragraphs[0].text,'☐ '+statement+': ______')

    def test_physically_governing_document_acts_protect_the_response_without_signer_inference(self):
        for sources in (('I confirm this form is accurate.','☐ Yes, completed by: ______'),
                        ('Witness statement: I saw the patient sign.','☐ Yes, completed by: ______'),
                        ('Approval of this report: ☐ I do so on: ______',),
                        ('I confirm this form is accurate.','☐ I do so on: ______')):
            with self.subTest(sources=sources):
                doc=Document()
                for source in sources:doc.add_paragraph(source)
                doc.add_paragraph('Unrelated note: ______');raw=data(doc);result=word_v3.inspect(raw)
                self.assertTrue(any(i['code']=='human_signing_action' for i in result['coverage_issues']))
                action=next(f for f in result['fields'] if f['kind']=='checkbox')
                self.assertTrue(action['write_protected'])
                forged=deepcopy(action);forged.pop('write_protected',None);forged.pop('review_reason',None)
                with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[forged],{forged['key']:'checked'})
                ordinary=next(f for f in result['fields'] if f['label']=='Unrelated note')
                index=word_v3.SlotIndex(word_v3.open_package(raw));slot=index.by_id[f'word/document.xml#b/p[{len(sources)-1}]']
                start=word_v3.text_of(slot.el).index('_')
                manual={'key':'manual_action','label':'Forged completion','kind':'text','locator':{
                    'contract':'v3','part':slot.part,'slot':slot.path,'fingerprint':word_v3.fingerprint(slot.el),
                    'target':{'type':'span','start':start,'end':start+6,'shape':'blank'}}}
                with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[manual],{manual['key']:'Synthetic completion'})
                out=Document(io.BytesIO(word_v3.fill(raw,result['fields'],{ordinary['key']:'Synthetic review'})))
                self.assertEqual([p.text for p in out.paragraphs[:-1]],list(sources))

    def test_recorded_status_and_independent_question_or_cell_do_not_inherit_another_act(self):
        for source in ('Consent was obtained by: ______','Signed consent was received on: ______',
                       'Initial symptoms occurred at: ______','I take warfarin at: ______'):
            with self.subTest(source=source):
                raw=self.form('☐ '+source);result=word_v3.inspect(raw)
                self.assertFalse([i for i in result['coverage_issues'] if i['code'] in
                                  ('human_signing_action','ambiguous_source_action')])
                state=next(f for f in result['fields'] if f['kind']=='checkbox')
                self.assertFalse(state.get('write_protected'))
                out=word_v3.fill(raw,result['fields'],{state['key']:'checked'})
                self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text,'☒ '+source)
        for boundary in ('independent_question','table','different_cell','independent_personal_state'):
            with self.subTest(boundary=boundary):
                doc=Document();doc.add_paragraph('I confirm this form is accurate.')
                text='Any symptoms? ☐ No ☐ Yes → if yes, describe: ______'
                if boundary=='table':doc.add_table(rows=1,cols=1).cell(0,0).text='Fixed unrelated block'
                if boundary=='different_cell':
                    table=doc.add_table(rows=1,cols=2)
                    table.cell(0,0).text='I confirm this form is accurate.'
                    table.cell(0,1).text='☐ Other, specify_________'
                else:doc.add_paragraph('☐ I have pain_________' if boundary=='independent_personal_state' else text)
                doc.add_paragraph('Unrelated note: ______');result=word_v3.inspect(data(doc))
                self.assertFalse([i for i in result['coverage_issues'] if i['code'] in
                                  ('human_signing_action','ambiguous_source_action')])


    def test_nominal_action_places_are_held_without_enumerating_signer_names(self):
        for label in ('Manager approval','Department lead approval','Approval by manager',
                      'Approval from department lead','Review and approval','Verification of this report'):
            for tail in ('','_________'):
                with self.subTest(label=label,tail=tail):
                    source=label+': ☐ Approved'+tail;raw=self.form(source);result=word_v3.inspect(raw)
                    self.assertTrue(any(i['code'] in ('human_signing_action','ambiguous_source_action')
                                        for i in result['coverage_issues']))
                    action=next(f for f in result['fields'] if f['kind']=='checkbox')
                    self.assertTrue(action['write_protected'])
                    forged=deepcopy(action);forged.pop('write_protected',None);forged.pop('review_reason',None)
                    self.assertIn(action['key'],word_v3.destination_roles(raw,[forged])['violations'])
                    with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[forged],{forged['key']:'checked'})
                    ordinary=next(f for f in result['fields'] if f['label']=='Unrelated note')
                    out=word_v3.fill(raw,result['fields'],{ordinary['key']:'Synthetic review'})
                    self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text,source)
        source='Approval status: ☐ Recorded';raw=self.form(source);result=word_v3.inspect(raw)
        self.assertFalse([i for i in result['coverage_issues'] if i['code'] in
                          ('human_signing_action','ambiguous_source_action')])
        state=next(f for f in result['fields'] if f['kind']=='checkbox')
        self.assertFalse(state.get('write_protected'))
        out=word_v3.fill(raw,result['fields'],{state['key']:'checked'})
        self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text,'Approval status: ☒ Recorded')

    def test_noncheckbox_document_actions_are_source_owned_per_destination(self):
        for source in ('I confirm that this clinical record is complete and accurate:',
                       'Manager approval:', 'Approval by manager:'):
            for shape in ('paragraph', 'table'):
                with self.subTest(source=source,shape=shape):
                    doc=Document()
                    if shape=='table':
                        table=doc.add_table(rows=2,cols=2)
                        table.cell(0,0).text=source;table.cell(1,0).text='Reason for visit:'
                    else:
                        doc.add_paragraph(source+' ______    Reason for visit: ______')
                    buf=io.BytesIO();doc.save(buf);raw=buf.getvalue();result=word_v3.inspect(raw)
                    held=[f for f in result['fields'] if f.get('write_protected')]
                    self.assertEqual(len(held),1)
                    self.assertTrue(any(i['code'] in ('human_signing_action','ambiguous_source_action')
                                        for i in result['coverage_issues']))
                    forged=deepcopy(held[0]);forged.pop('write_protected');forged.pop('review_reason')
                    self.assertIn(forged['key'],word_v3.destination_roles(raw,[forged])['violations'])
                    with self.assertRaises(word_v3.V3Failure):
                        word_v3.fill(raw,[forged],{forged['key']:'SYNTHETIC UNAUTHORIZED ACT'})
                    ordinary=next(f for f in result['fields'] if f['label']=='Reason for visit')
                    self.assertFalse(ordinary.get('write_protected'))
                    out=word_v3.fill(raw,result['fields'],{ordinary['key']:'Synthetic review'})
                    filled=Document(io.BytesIO(out))
                    if shape=='table':
                        self.assertEqual(filled.tables[0].cell(0,0).text,source)
                        self.assertEqual(filled.tables[0].cell(0,1).text,'')
                        self.assertEqual(filled.tables[0].cell(1,1).text,'Synthetic review')
                    else:
                        self.assertEqual(filled.paragraphs[0].text,source+' ______    Reason for visit: Synthetic review')

    def test_recorded_status_and_clinical_review_sources_are_independent_of_other_actions(self):
        doc=Document();table=doc.add_table(rows=4,cols=2)
        labels=('Manager approval:', 'Initial presentation:', 'Recorded signed report status:', 'Approval status:')
        for row,label in enumerate(labels):table.cell(row,0).text=label
        buf=io.BytesIO();doc.save(buf);raw=buf.getvalue();result=word_v3.inspect(raw)
        self.assertEqual(sum(bool(f.get('write_protected')) for f in result['fields']),1)
        # The existing initial/signed lexical kinds still require the explicit
        # ordinary-field mapping decision. This source-role repair does not
        # globally release those kinds or count that correction as automatic.
        reviewed=deepcopy(result['fields'])
        for field in reviewed:
            if not field.get('write_protected'):field['kind']='text'
        answers={f['key']:'Synthetic ordinary status' for f in reviewed if not f.get('write_protected')}
        self.assertEqual(word_v3.destination_roles(raw,[f for f in reviewed if not f.get('write_protected')])['violations'],{})
        out=word_v3.fill(raw,reviewed,answers);filled=Document(io.BytesIO(out))
        self.assertEqual(filled.tables[0].cell(0,1).text,'')
        for row in range(1,4):self.assertEqual(filled.tables[0].cell(row,1).text,'Synthetic ordinary status')

    def source_act_forms(self,statement):
        for shape in ('checkbox','paragraph','table'):
            doc=Document()
            if shape=='table':
                table=doc.add_table(rows=1,cols=2);table.cell(0,0).text=statement+':'
            else:doc.add_paragraph(('☐ ' if shape=='checkbox' else '')+statement+' ______')
            doc.add_paragraph('Unrelated note: ______');buffer=io.BytesIO();doc.save(buffer)
            yield shape,buffer.getvalue()

    def assert_source_act_stays_human_owned(self,statement):
        for shape,raw in self.source_act_forms(statement):
            with self.subTest(statement=statement,shape=shape):
                result=word_v3.inspect(raw);held=[f for f in result['fields'] if f.get('write_protected')]
                self.assertGreaterEqual(len(held),1)
                self.assertTrue(any(i['code'] in ('human_signing_action','ambiguous_source_action')
                                    for i in result['coverage_issues']))
                for original in held:
                    field=deepcopy(original);field.pop('write_protected');field.pop('review_reason')
                    if field['kind'] in ('signature','initials'):field['kind']='text'
                    self.assertIn(field['key'],word_v3.destination_roles(raw,[field])['violations'])
                    with self.assertRaises(word_v3.V3Failure):
                        word_v3.fill(raw,[field],{field['key']:'checked' if field['kind']=='checkbox' else 'Synthetic completion'})
                ordinary=next(f for f in result['fields'] if f['label']=='Unrelated note')
                self.assertFalse(ordinary.get('write_protected'))
                before=Document(io.BytesIO(raw));after=Document(io.BytesIO(word_v3.fill(raw,result['fields'],
                                                        {ordinary['key']:'Synthetic review'})))
                self.assertEqual([p.text for p in after.paragraphs[:-1]],[p.text for p in before.paragraphs[:-1]])
                if shape=='table':
                    self.assertEqual([c.text for c in after.tables[0].rows[0].cells],
                                     [c.text for c in before.tables[0].rows[0].cells])

    def test_checkbox_action_condition_cannot_supply_signing_authority(self):
        for statement in ('By checking this box, I agree to this procedure',
                          'By checking this box I agree to this procedure',
                          'By checking this box you authorize this treatment',
                          'By ticking this box; we certify that this record is complete'):
            self.assert_source_act_stays_human_owned(statement)

    def test_documentary_nominal_action_does_not_become_an_ordinary_option(self):
        for statement in ('Confirmation that this record is complete and accurate',
                          'Certification of this clinical record','Certification of treatment completion'):
            self.assert_source_act_stays_human_owned(statement)

    def test_documentary_imperative_cannot_be_completed_by_generation(self):
        for statement in ('Confirm this form is complete','Please verify the information above is correct',
                          'Authorize this treatment','Consent to this procedure'):
            self.assert_source_act_stays_human_owned(statement)

    def test_third_person_action_never_infers_a_signer_and_ordinary_observations_stay_useful(self):
        for statement in ('The signer certifies that this record is accurate',
                          'The department lead confirms that this form is accurate',
                          'The department lead authorizes treatment'):
            self.assert_source_act_stays_human_owned(statement)
        for statement in ('Confirm left knee swelling','The laboratory results confirm no infection',
                          'Recorded approval status','Initial presentation','Recorded signed report status'):
            for shape,raw in self.source_act_forms(statement):
                with self.subTest(statement=statement,shape=shape):
                    result=word_v3.inspect(raw)
                    self.assertFalse([i for i in result['coverage_issues'] if i['code'] in
                                      ('human_signing_action','ambiguous_source_action')])
                    fields=deepcopy(result['fields'])
                    for field in fields:
                        if field['kind'] in ('signature','initials'):field['kind']='text'
                    target=next(f for f in fields if f['kind']=='checkbox') if shape=='checkbox' else \
                           next(f for f in fields if f['label']!='Unrelated note')
                    self.assertFalse(target.get('write_protected'))
                    self.assertEqual(word_v3.destination_roles(raw,[target])['violations'],{})
                    word_v3.fill(raw,fields,{target['key']:'checked' if target['kind']=='checkbox'
                                            else 'Synthetic ordinary status'})

    def test_permission_and_witness_arguments_remain_human_owned_in_each_physical_shape(self):
        for statement in ('I grant permission for this procedure',
                          'I observed the patient signing',
                          'The patient signed in my presence',
                          'I permit the examination to proceed',
                          'We have given our permission for treatment',
                          'I saw the patient sign',
                          'I witnessed their signature'):
            self.assert_source_act_stays_human_owned(statement)

    def test_colonless_question_cells_keep_their_action_role_without_named_signer_inference(self):
        for caption in ('Manager approval','Witness attestation','Patient consent',
                        'Manager authorization','Witness certification','Approval by department lead'):
            for suffix in ('','.',':'):
                with self.subTest(caption=caption,suffix=suffix):
                    doc=Document();table=doc.add_table(rows=2,cols=2)
                    table.cell(0,0).text=caption+suffix;table.cell(0,1).text='______'
                    table.cell(1,0).text='Initial diagnosis';table.cell(1,1).text='______'
                    doc.add_paragraph('Unrelated note: ______');raw=data(doc);result=word_v3.inspect(raw)
                    held=[f for f in result['fields'] if f.get('write_protected')]
                    self.assertEqual(len(held),1)
                    self.assertTrue(any(i['code'] in ('human_signing_action','ambiguous_source_action')
                                        for i in result['coverage_issues']))
                    forged=deepcopy(held[0]);forged.pop('write_protected');forged.pop('review_reason')
                    self.assertIn(forged['key'],word_v3.destination_roles(raw,[forged])['violations'])
                    with self.assertRaises(word_v3.V3Failure):
                        word_v3.fill(raw,[forged],{forged['key']:'SYNTHETIC UNAUTHORIZED ACT'})
                    ordinary=next(f for f in result['fields'] if f['label']=='Initial diagnosis')
                    self.assertEqual(ordinary['kind'],'text')
                    out=Document(io.BytesIO(word_v3.fill(raw,result['fields'],
                                                        {ordinary['key']:'Synthetic diagnosis'})))
                    self.assertEqual(out.tables[0].cell(0,0).text,caption+suffix)
                    self.assertEqual(out.tables[0].cell(0,1).text,'______')
                    self.assertEqual(out.tables[0].cell(1,1).text,'Synthetic diagnosis')

    def test_incidental_initial_and_recorded_signed_arguments_are_ordinary_new_mapping_fields(self):
        for caption in ('Initial symptoms occurred at','Initial diagnosis',
                        'Signed consent was received on','Signed consent received on',
                        'Recorded signed report status'):
            for shape in ('paragraph','table'):
                with self.subTest(caption=caption,shape=shape):
                    doc=Document()
                    if shape=='table':
                        table=doc.add_table(rows=1,cols=2)
                        table.cell(0,0).text=caption;table.cell(0,1).text='______'
                    else:doc.add_paragraph(caption+': ______')
                    doc.add_paragraph('Clinician signature: ______');raw=data(doc);result=word_v3.inspect(raw)
                    ordinary=next(f for f in result['fields'] if f['label']==caption)
                    self.assertEqual(ordinary['kind'],'text');self.assertFalse(ordinary.get('write_protected'))
                    self.assertEqual(word_v3.destination_roles(raw,[ordinary])['violations'],{})
                    out=Document(io.BytesIO(word_v3.fill(raw,result['fields'],
                                                        {ordinary['key']:'Synthetic recorded value'})))
                    if shape=='table':self.assertEqual(out.tables[0].cell(0,1).text,'Synthetic recorded value')
                    else:self.assertEqual(out.paragraphs[0].text,caption+': Synthetic recorded value')
                    self.assertIn('Clinician signature: ______',[p.text for p in out.paragraphs])
                    signer=next(f for f in result['fields'] if f['label']=='Clinician signature')
                    signer=deepcopy(signer);signer['kind']='text';signer.pop('write_protected',None)
                    with self.assertRaises(word_v3.V3Failure):
                        word_v3.fill(raw,[signer],{signer['key']:'Synthetic finalizer name'})

    def test_witness_words_and_permission_nouns_do_not_replace_ordinary_clinical_arguments(self):
        for statement in ('I observed the patient walking','I saw ankle swelling',
                          'I observed Tinel signs','Recorded permission status',
                          'Consent was obtained by','Signed consent was received on'):
            for shape,raw in self.source_act_forms(statement):
                with self.subTest(statement=statement,shape=shape):
                    result=word_v3.inspect(raw)
                    self.assertFalse([i for i in result['coverage_issues'] if i['code'] in
                                      ('human_signing_action','ambiguous_source_action')])
                    fields=deepcopy(result['fields'])
                    target=next(f for f in fields if f['kind']=='checkbox') if shape=='checkbox' else \
                           next(f for f in fields if f['label']!='Unrelated note')
                    self.assertFalse(target.get('write_protected'))
                    self.assertEqual(word_v3.destination_roles(raw,[target])['violations'],{})
                    word_v3.fill(raw,fields,{target['key']:'checked' if target['kind']=='checkbox'
                                            else 'Synthetic recorded value'})

    def test_personal_signing_event_arguments_remain_held_without_a_witness_verb_whitelist(self):
        for statement in ('I saw the patient sign the consent form',
                          'I watched the patient sign the consent form',
                          'I witnessed the patient sign this consent form',
                          'I was present when the patient signed',
                          'I was present during the signing',
                          'I stood beside the patient when she signed'):
            self.assert_source_act_stays_human_owned(statement)

    def test_personal_permission_arguments_require_role_review_with_unfamiliar_or_negative_acts(self):
        for statement in ('I provide informed consent for this procedure',
                          'I withdraw my consent to this procedure',
                          'We gave written permission for surgery',
                          'I revoke my authorization for this treatment'):
            self.assert_source_act_stays_human_owned(statement)

    def test_printed_speaker_appositions_never_establish_signing_identity(self):
        for statement in ('I, the witness, saw the patient sign',
                          'I, the manager, approve this document',
                          'We, the undersigned, verify the above information',
                          'I (patient) allow this procedure',
                          'I — the clinician — certify this record',
                          'I [the clinician] certify this record'):
            self.assert_source_act_stays_human_owned(statement)
        for statement in ('I, the patient, have ankle swelling','I (patient) take warfarin',
                          'We, the clinical team, observed ankle swelling'):
            for shape,raw in self.source_act_forms(statement):
                with self.subTest(ordinary=statement,shape=shape):
                    result=word_v3.inspect(raw)
                    self.assertFalse([i for i in result['coverage_issues'] if i['code'] in
                                      ('human_signing_action','ambiguous_source_action')])
                    field=next(f for f in result['fields'] if f['label']!='Unrelated note')
                    self.assertEqual(word_v3.destination_roles(raw,[field])['violations'],{})

    def test_grammatical_number_does_not_release_action_captions_or_document_objects(self):
        for statement in ('Witness declarations','Declarations that these records are accurate',
                          'The reviewers approve these documents','I verify the above records'):
            self.assert_source_act_stays_human_owned(statement)

    def test_capacity_phrases_and_third_person_witness_events_never_supply_authority(self):
        for statement in ('The witness saw the patient sign',
                          'The witness observed the patient signing this form',
                          'I as the parent allow this procedure',
                          'As the witness, I saw the patient sign',
                          'I—the witness—saw the patient sign',
                          'The patient’s witness saw her sign',
                          'By checking this box, I, the patient, authorise the treatment',
                          'As manager, I approve this document'):
            self.assert_source_act_stays_human_owned(statement)
        for statement in ('As the patient, I have ankle swelling',
                          'I as the permit holder have ankle swelling',
                          'The witness observed the patient walking',
                          'The clinician saw Tinel signs'):
            for shape,raw in self.source_act_forms(statement):
                with self.subTest(ordinary=statement,shape=shape):
                    result=word_v3.inspect(raw)
                    self.assertFalse([i for i in result['coverage_issues'] if i['code'] in
                                      ('human_signing_action','ambiguous_source_action')])
                    fields=result['fields'];target=next(f for f in fields if f['kind']=='checkbox') if shape=='checkbox' else \
                                                 next(f for f in fields if f['label']!='Unrelated note')
                    self.assertEqual(word_v3.destination_roles(raw,[target])['violations'],{})
                    word_v3.fill(raw,fields,{target['key']:'checked' if target['kind']=='checkbox'
                                            else 'Synthetic ordinary status'})

    def test_name_blank_inside_attestation_subject_is_not_a_finalizer_identity_destination(self):
        for source in ('I [enter full name] certify that this record is accurate: ______    Reason for visit: ______',
                       'I (initials ______) certify that this record is accurate: ______    Reason for visit: ______',
                       'I — [enter full name] — certify that this record is accurate: ______    Reason for visit: ______',
                       'As [enter full name], I certify that this record is accurate: ______    Reason for visit: ______',
                       'I as [enter full name] allow this procedure: ______    Reason for visit: ______',
                       'In my capacity as [enter full name], I certify this record: ______    Reason for visit: ______',
                       'Acting as [enter full name], I approve this form: ______    Reason for visit: ______',
                       'As [enter full name] I certify this record: ______    Reason for visit: ______',
                       'I ______ certify this record: ______    Reason for visit: ______',
                       'I as ______ allow this procedure: ______    Reason for visit: ______'):
            with self.subTest(source=source):
                doc=Document();doc.add_paragraph(source);raw=data(doc);result=word_v3.inspect(raw)
                fields=result['fields'];ordinary=next(f for f in fields if f['label']=='Reason for visit')
                identity=fields[0];self.assertTrue(identity['write_protected'])
                forged=deepcopy(identity);forged['kind']='text';forged.pop('write_protected');forged.pop('review_reason')
                self.assertIn(forged['key'],word_v3.destination_roles(raw,[forged])['violations'])
                with self.assertRaises(word_v3.V3Failure):
                    word_v3.fill(raw,[forged],{forged['key']:'Synthetic finalizer name'})
                self.assertFalse(ordinary.get('write_protected'))
                out=word_v3.fill(raw,fields,{ordinary['key']:'Synthetic reason'})
                self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text,
                                 source.removesuffix('______')+'Synthetic reason')
        source='I [enter full name] have ankle swelling: ______    Reason for visit: ______'
        doc=Document();doc.add_paragraph(source);raw=data(doc);fields=word_v3.inspect(raw)['fields']
        identity=fields[0];self.assertFalse(identity.get('write_protected'))
        self.assertEqual(word_v3.destination_roles(raw,[identity])['violations'],{})
        out=word_v3.fill(raw,fields,{identity['key']:'Synthetic patient name'})
        self.assertIn('I Synthetic patient name have ankle swelling',Document(io.BytesIO(out)).paragraphs[0].text)

    def test_signing_events_keep_unfamiliar_objects_and_introductory_capacities_source_held(self):
        for statement in ('I saw the patient sign it',
                          'I witnessed the patient sign below',
                          'The nurse saw the patient sign it',
                          'In my capacity as witness, I confirm this form',
                          'As a witness I confirm this form',
                          'Acting as witness, I acknowledge this statement',
                          'Witness observation of the patient signing'):
            self.assert_source_act_stays_human_owned(statement)
        # A literal "sign here" is already a physical signature kind. Its
        # checkbox still cannot tick even if its ordinary blank is withdrawn.
        raw=self.form('☐ I watched the patient sign here: ______')
        result=word_v3.inspect(raw);field=deepcopy(next(f for f in result['fields'] if f.get('write_protected')))
        field.pop('write_protected');field.pop('review_reason');field['kind']='checkbox'
        self.assertIn(field['key'],word_v3.destination_roles(raw,[field])['violations'])
        with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[field],{field['key']:'checked'})
        for statement in ('Acting as clinician, I saw ankle swelling',
                          'In my capacity as clinician, I observed signs of dehydration',
                          'As the patient I have ankle swelling',
                          'I observed a sign of infection',
                          'The clinician observed a sign of dehydration'):
            for shape,raw in self.source_act_forms(statement):
                with self.subTest(ordinary=statement,shape=shape):
                    result=word_v3.inspect(raw)
                    self.assertFalse([i for i in result['coverage_issues'] if i['code'] in
                                      ('human_signing_action','ambiguous_source_action')])
                    target=next(f for f in result['fields'] if f['kind']=='checkbox') if shape=='checkbox' else \
                           next(f for f in result['fields'] if f['label']!='Unrelated note')
                    self.assertEqual(word_v3.destination_roles(raw,[target])['violations'],{})
                    word_v3.fill(raw,result['fields'],{target['key']:'checked' if target['kind']=='checkbox'
                                                    else 'Synthetic clinical finding'})

    def test_all_actor_and_response_destinations_bind_to_the_exact_human_act_clause(self):
        for human,separator in (('Person ______ saw the patient sign: ______','    '),
                                ('Name ______ saw the patient sign: ______','\t\t'),
                                ('The ______ approves this document: ______','; '),
                                ('The person named ______ permits this procedure: ______','. '),
                                ('I [enter first name] [enter last name] approve this report: ______','    '),
                                ('I, Dr. ______, certify this record: ______','; ')):
            with self.subTest(human=human,separator=separator):
                source=human+separator+'Initial diagnosis: ______'
                doc=Document();doc.add_paragraph(source);raw=data(doc);result=word_v3.inspect(raw)
                ordinary=next(f for f in result['fields'] if f['label']=='Initial diagnosis')
                self.assertEqual(ordinary['kind'],'text');self.assertFalse(ordinary.get('write_protected'))
                acts=[f for f in result['fields'] if f['key']!=ordinary['key']]
                self.assertGreaterEqual(len(acts),2)
                for field in acts:
                    self.assertTrue(field.get('write_protected'),field)
                    forged=deepcopy(field);forged['kind']='text';forged.pop('write_protected');forged.pop('review_reason')
                    self.assertIn(field['key'],word_v3.destination_roles(raw,[forged])['violations'])
                    with self.assertRaises(word_v3.V3Failure):
                        word_v3.fill(raw,[forged],{field['key']:'Synthetic finalizer name'})
                out=word_v3.fill(raw,result['fields'],{ordinary['key']:'Synthetic ordinary diagnosis'})
                self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text,
                                 source.removesuffix('______')+'Synthetic ordinary diagnosis')

    def test_quoted_and_conditional_document_acts_do_not_authorize_the_actor(self):
        for statement in ('"I approve this document"','“I approve this document”',
                          "'We consent to this procedure'",'‘We consent to this procedure’',
                          'If required, I authorise this treatment'):
            self.assert_source_act_stays_human_owned(statement)
        for statement in ('"I have ankle swelling"','If required, I record ankle swelling',
                          '“The clinician observed clinical signs of dehydration”'):
            for shape,raw in self.source_act_forms(statement):
                with self.subTest(ordinary=statement,shape=shape):
                    result=word_v3.inspect(raw)
                    self.assertFalse([i for i in result['coverage_issues'] if i['code'] in
                                      ('human_signing_action','ambiguous_source_action')])


class SourceCoverage(unittest.TestCase):
    def inspect(self,doc):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'original.docx';path.write_bytes(data(doc))
            return inspect_template(path,engine='v3')

    def test_unpunctuated_numbered_lists_require_exact_review_without_invented_choices(self):
        for title,items in (('Gender',('Male','Female','Other')),
                            ('Discharge instructions',('Call clinic','Keep wound dry'))):
            with self.subTest(title=title):
                doc=Document();numbered(doc,title,90,'decimal')
                for item in items:numbered(doc,item,91,'lowerLetter')
                doc.add_paragraph('Reason for visit: ______')
                result=self.inspect(doc)
                self.assertEqual([f['label'] for f in result['fields']],['Reason for visit'])
                issues=result['import_manifest']['coverage']['selected_engine_issues']
                self.assertEqual(len(issues),1+len(items))
                self.assertEqual({i['code'] for i in issues},{'unclassified_response_list'})
                self.assertEqual({i['slot'] for i in issues},{f'b/p[{i}]' for i in range(1+len(items))})
                self.assertTrue(result['import_manifest']['coverage']['block_generation'])

    def test_explicit_question_keeps_qualified_lettered_choices(self):
        doc=Document();doc.add_paragraph('Which knee was examined?')
        for item in ('Left knee','Right knee'):numbered(doc,item,91,'lowerLetter')
        result=self.inspect(doc)
        choice=next(f for f in result['fields'] if f['label']=='Which knee was examined?')
        self.assertEqual((choice['kind'],choice['options']),('choice',['Left knee','Right knee']))
        self.assertFalse(result['import_manifest']['coverage'].get('selected_engine_issues'))

    def test_bare_factor_list_is_held_without_creating_blanks(self):
        doc=Document();doc.add_paragraph('History of associated factors:')
        doc.add_paragraph('a) Hypertension');doc.add_paragraph('b) Peripheral arterial disease')
        doc.add_paragraph('Reason for visit: ______')
        result=self.inspect(doc)
        self.assertEqual([f['label'] for f in result['fields']],['Reason for visit'])
        issues=result['import_manifest']['coverage']['selected_engine_issues']
        self.assertEqual([i['slot'] for i in issues],['b/p[0]','b/p[1]','b/p[2]'])

    def test_literal_lists_do_not_depend_on_parent_punctuation_or_casing(self):
        for title in ('Gender','Gender.','Gender:','Instructions'):
            for prefix in ('a)','A)'):
                with self.subTest(title=title,prefix=prefix):
                    doc=Document();doc.add_paragraph(title)
                    doc.add_paragraph(prefix+' Male');doc.add_paragraph(('b)' if prefix=='a)' else 'B)')+' Female')
                    doc.add_paragraph('Reason: ______');result=self.inspect(doc)
                    self.assertEqual([f['label'] for f in result['fields']],['Reason'])
                    self.assertEqual(len(result['import_manifest']['coverage']['selected_engine_issues']),3)
                    self.assertTrue(result['import_manifest']['coverage']['block_generation'])

    def test_list_does_not_bridge_a_table_or_unrelated_numbering(self):
        doc=Document();numbered(doc,'Instructions',90,'decimal');numbered(doc,'One item',91,'lowerLetter')
        doc.add_table(rows=1,cols=1).cell(0,0).text='Fixed boundary'
        numbered(doc,'Separate item',91,'lowerLetter');doc.add_paragraph('Reason: ______')
        result=self.inspect(doc)
        self.assertFalse([i for i in result['import_manifest']['coverage'].get('selected_engine_issues',[])
                          if i['code']=='unclassified_response_list'])

    def test_printed_pronoun_alternatives_are_exact_unwritable_source_ranges(self):
        doc=Document();doc.add_paragraph('The care received was right for [me/us].')
        doc.add_paragraph('Reason: ______')
        result=self.inspect(doc);self.assertEqual([f['label'] for f in result['fields']],['Reason'])
        issue=result['import_manifest']['coverage']['selected_engine_issues'][0]
        self.assertEqual(issue['code'],'printed_bracket_alternatives')
        text=doc.paragraphs[0].text;self.assertEqual(text[issue['start']:issue['end']],'[me/us]')
        self.assertEqual(issue['slot'],'b/p[0]')
        output=word_v3.fill(data(doc),result['fields'],{result['fields'][0]['key']:'Synthetic review'})
        self.assertEqual(Document(io.BytesIO(output)).paragraphs[0].text,text)

    def test_alternative_parts_cannot_escape_through_punctuation_unicode_or_digits(self):
        for alternatives in ('left/right-hand','me/us/other-person',"patient's/carer's",
                             'patient/carer (specify)','mild/moderate/grade 3','α/β','左/右',
                             'patient’s/carer’s','1/2/3','type A/type B','state/local','add/remove',
                             'patient/carer/','patient//carer','patient/\ncarer'):
            with self.subTest(alternatives=alternatives):
                doc=Document();doc.add_paragraph('Selection: ['+alternatives+']');doc.add_paragraph('Reason: ______')
                result=self.inspect(doc)
                self.assertEqual([f['label'] for f in result['fields']],['Reason'])
                issue=result['import_manifest']['coverage']['selected_engine_issues'][0]
                self.assertEqual(issue['code'],'printed_bracket_alternatives')
                self.assertEqual(doc.paragraphs[0].text[issue['start']:issue['end']],'['+alternatives+']')
                self.assertTrue(result['import_manifest']['coverage']['block_generation'])

    def test_ordinary_placeholders_and_dates_remain_usable(self):
        doc=Document();doc.add_paragraph('Name: [patient name]')
        doc.add_paragraph('Date: [dd/mm/yyyy]');doc.add_paragraph('Plan: [enter clinical plan]')
        doc.add_paragraph('Printed abbreviation [SDU] and [This page is intentionally left blank]')
        result=self.inspect(doc)
        self.assertFalse(result['import_manifest']['coverage'].get('selected_engine_issues'))
        self.assertEqual({f['label'] for f in result['fields']},{'Name patient name','Date dd/mm/yyyy','Plan enter clinical plan'})

    def test_spaced_date_formats_fill_while_spaced_literal_alternatives_stay_held(self):
        for hint in ('DD / MM / YYYY','dd\t/\tmm / yy','MM\u00a0/\u00a0DD / YYYY'):
            with self.subTest(hint=hint):
                doc=Document();doc.add_paragraph('Date: ['+hint+']');result=self.inspect(doc)
                self.assertFalse(result['import_manifest']['coverage'].get('selected_engine_issues'))
                field=result['fields'][0];self.assertEqual(field['kind'],'date')
                out=word_v3.fill(data(doc),result['fields'],{field['key']:'03/14/2026'})
                self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text,'Date: 03/14/2026')
        doc=Document();doc.add_paragraph('Site: [left / right]');doc.add_paragraph('Reason: ______')
        result=self.inspect(doc);self.assertTrue(result['import_manifest']['coverage']['block_generation'])
        self.assertEqual([f['label'] for f in result['fields']],['Reason'])

    def test_unclosed_slash_source_is_reviewable_without_inventing_an_answer(self):
        doc=Document();doc.add_paragraph('Site: [left/right');doc.add_paragraph('Reason: ______')
        result=self.inspect(doc);self.assertEqual([f['label'] for f in result['fields']],['Reason'])
        issue=result['import_manifest']['coverage']['selected_engine_issues'][0]
        self.assertEqual(issue['code'],'unbalanced_bracketed_source')
        self.assertEqual((issue['start'],issue['end']),(6,len(doc.paragraphs[0].text)))
        out=word_v3.fill(data(doc),result['fields'],{result['fields'][0]['key']:'Synthetic review'})
        self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text,'Site: [left/right')

    def test_unclosed_slash_before_another_placeholder_keeps_exact_source_and_later_writer(self):
        for suffix,answer in (('Comments: [enter notes]','Synthetic notes'),
                              ('Date: [DD/MM/YYYY]','03/14/2026')):
            with self.subTest(suffix=suffix):
                text='Site: [left/right; '+suffix
                doc=Document();doc.add_paragraph(text);doc.add_paragraph('Reason: ______')
                result=self.inspect(doc);issues=result['import_manifest']['coverage']['selected_engine_issues']
                self.assertEqual(len(issues),1);issue=issues[0]
                self.assertEqual(issue['code'],'unbalanced_bracketed_source')
                later=text.index('[',text.index('[')+1)
                self.assertEqual((issue['start'],issue['end']),(6,later))
                self.assertEqual(text[issue['start']:issue['end']],'[left/right; '+suffix.split('[')[0])
                self.assertTrue(result['import_manifest']['coverage']['block_generation'])
                field=next(f for f in result['fields'] if f['locator']['slot']=='b/p[0]')
                target=field['locator']['target'];self.assertGreaterEqual(target['start'],issue['end'])
                out=word_v3.fill(data(doc),result['fields'],{field['key']:answer})
                exported=Document(io.BytesIO(out))
                self.assertEqual(exported.paragraphs[0].text,text[:later]+answer)
                self.assertEqual(exported.paragraphs[1].text,'Reason: ______')

    def test_closed_alternatives_do_not_consume_the_later_placeholder_or_date(self):
        for suffix,answer in (('Comments: [enter notes]','Synthetic notes'),
                              ('Date: [DD/MM/YYYY]','03/14/2026')):
            with self.subTest(suffix=suffix):
                text='Site: [left/right]; '+suffix
                doc=Document();doc.add_paragraph(text);doc.add_paragraph('Reason: ______')
                result=self.inspect(doc);issues=result['import_manifest']['coverage']['selected_engine_issues']
                self.assertEqual(len(issues),1);issue=issues[0]
                self.assertEqual(issue['code'],'printed_bracket_alternatives')
                self.assertEqual(text[issue['start']:issue['end']],'[left/right]')
                field=next(f for f in result['fields'] if f['locator']['slot']=='b/p[0]')
                out=word_v3.fill(data(doc),result['fields'],{field['key']:answer})
                self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text,'Site: [left/right]; '+suffix.split('[')[0]+answer)

    def inner_tabs(self):
        doc=Document();p=doc.add_paragraph('Instruction: [identify services before')
        run=p.add_run('\t');run.font.underline=True;p.add_run(' his/her')
        run=p.add_run('\t');run.font.underline=True;p.add_run(' last review]')
        doc.add_paragraph('Reason: ______')
        return doc

    def test_all_blank_shapes_inside_unresolved_brackets_stay_unchanged(self):
        doc=self.inner_tabs();result=self.inspect(doc)
        self.assertEqual([f['label'] for f in result['fields']],['Reason'])
        issue=result['import_manifest']['coverage']['selected_engine_issues'][0]
        self.assertEqual(issue['code'],'printed_bracket_alternatives')
        self.assertTrue(result['import_manifest']['coverage']['block_generation'])
        field=result['fields'][0];out=word_v3.fill(data(doc),result['fields'],{field['key']:'Synthetic adjacent answer'})
        export=Document(io.BytesIO(out));self.assertEqual(export.paragraphs[0].text,doc.paragraphs[0].text)
        self.assertEqual(export.paragraphs[1].text,'Reason: Synthetic adjacent answer')
        self.assertEqual(etree.tostring(export.paragraphs[0]._p),etree.tostring(doc.paragraphs[0]._p))

    def test_original_unresolved_range_refuses_forged_span_and_copy_writers(self):
        doc=self.inner_tabs();raw=data(doc);result=self.inspect(doc)
        issue=result['import_manifest']['coverage']['selected_engine_issues'][0]
        package=word_v3.open_package(raw);index=word_v3.SlotIndex(package);slot=index.by_id['word/document.xml#b/p[0]']
        text,_=word_v3.paragraph_text(slot.el);offset=text.index('\t')
        forged={'key':'manual_source','label':'Forged source blank','kind':'text','locator':{
            'contract':'v3','part':slot.part,'slot':slot.path,'fingerprint':word_v3.fingerprint(slot.el),
            'target':{'type':'span','start':offset,'end':offset+1,'shape':'utab'}}}
        from copy import deepcopy
        copy=deepcopy(result['fields'][0]);copy['locator']['target']['copies']=[{
            'part':slot.part,'slot':slot.path,'fingerprint':word_v3.fingerprint(slot.el)}]
        for field in (forged,copy):
            with self.subTest(field=field['key']):
                with self.assertRaises(word_v3.V3Failure) as failure:
                    word_v3.fill(raw,[field],{field['key']:'SYNTHETIC FORGED WRITE'})
                self.assertEqual(failure.exception.code,'unresolved_source_range')

    def test_long_bracketed_source_is_not_silently_discarded_at_grammar_limit(self):
        for length in (121,2100):
            with self.subTest(length=length):
                doc=Document();doc.add_paragraph('['+'complete this authoring instruction '+'x'*length+']')
                doc.add_paragraph('Reason: ______');result=self.inspect(doc)
                self.assertEqual([f['label'] for f in result['fields']],['Reason'])
                issue=result['import_manifest']['coverage']['selected_engine_issues'][0]
                self.assertEqual(issue['code'],'unmapped_bracketed_source')
                self.assertEqual((issue['start'],issue['end']),(0,len(doc.paragraphs[0].text)))

    def test_inline_writer_expands_only_written_exact_row_and_preserves_neighbor_nodes(self):
        doc=Document();table=doc.add_table(rows=2,cols=1)
        table.cell(0,0).text='Finding: ______';table.cell(1,0).text='Fixed neighboring content'
        for row in table.rows:
            height=OxmlElement('w:trHeight');height.set(qn('w:hRule'),'exact');height.set(qn('w:val'),'180')
            row._tr.get_or_add_trPr().append(height)
        raw=data(doc);result=self.inspect(doc)
        field=next(f for f in result['fields'] if f['label']=='Finding')
        self.assertEqual(field['locator']['target']['type'],'span')
        answer='SYNTHETIC '+'long review '*5
        with self.assertRaises(word_v3.V3Failure):
            word_v3.fill(raw,result['fields'],{field['key']:'SYNTHETIC '+'long review '*60})
        out=word_v3.fill(raw,result['fields'],{field['key']:answer})
        original=Document(io.BytesIO(raw));export=Document(io.BytesIO(out))
        self.assertEqual(export.tables[0].rows[0]._tr.find('.//'+qn('w:trHeight')).get(qn('w:hRule')),'atLeast')
        self.assertEqual(etree.tostring(export.tables[0].rows[1]._tr),etree.tostring(original.tables[0].rows[1]._tr))
        self.assertEqual(export.tables[0].cell(0,0).text,'Finding: '+answer)
        with ZipFile(io.BytesIO(raw)) as before,ZipFile(io.BytesIO(out)) as after:
            self.assertEqual(before.namelist(),after.namelist())
            for name in before.namelist():
                if name!='word/document.xml':self.assertEqual(before.read(name),after.read(name),name)

    def fixed_header(self,doc):
        p=doc.sections[0].header.paragraphs[0];p.text='Patient name: ______'
        pict=OxmlElement('w:pict');shape=etree.SubElement(pict,'{urn:schemas-microsoft-com:vml}shape')
        shape.set('style','position:absolute;width:72pt;height:12pt')
        box=etree.SubElement(shape,'{urn:schemas-microsoft-com:vml}textbox')
        content=etree.SubElement(box,qn('w:txbxContent'));pp=etree.SubElement(content,qn('w:p'))
        rr=etree.SubElement(pp,qn('w:r'));etree.SubElement(rr,qn('w:t')).text='FIXED ADJACENT LABEL'
        p.add_run()._r.append(pict)

    def test_fixed_header_geometry_is_explicit_before_generation_and_forged_fill_refuses(self):
        doc=Document();self.fixed_header(doc);doc.add_paragraph('Reason: ______');result=self.inspect(doc)
        self.assertTrue(result['import_manifest']['coverage']['block_generation'])
        issues=result['import_manifest']['coverage']['selected_engine_issues']
        self.assertTrue(any(i['code']=='unqualified_fixed_container' and i['part']=='word/header1.xml' for i in issues))
        field=next(f for f in result['fields'] if f['locator']['part']=='word/header1.xml')
        self.assertTrue(field['write_protected'])
        forged=dict(field,write_protected=False)
        with self.assertRaises(word_v3.V3Failure):word_v3.fill(data(doc),[forged],{forged['key']:'SYNTHETIC replacement'})
        ordinary=next(f for f in result['fields'] if f['label']=='Reason')
        self.assertFalse(ordinary.get('write_protected'))
        out=word_v3.fill(data(doc),result['fields'],{ordinary['key']:'Synthetic ordinary review'})
        self.assertEqual(Document(io.BytesIO(out)).sections[0].header.paragraphs[0].text,'Patient name: ______')

    def test_positioned_table_also_holds_neighbor_growth_but_plain_document_stays_answerable(self):
        doc=Document();floating=doc.add_table(rows=1,cols=2);floating.cell(0,0).text='Finding'
        pos=OxmlElement('w:tblpPr');pos.set(qn('w:vertAnchor'),'text');pos.set(qn('w:tblpY'),'300')
        floating._tbl.tblPr.append(pos)
        regular=doc.add_table(rows=1,cols=2);regular.cell(0,0).text='Plan'
        result=self.inspect(doc);by={f['label']:f for f in result['fields']}
        self.assertTrue(by['Finding']['write_protected']);self.assertTrue(by['Plan']['write_protected'])
        self.assertTrue(result['import_manifest']['coverage']['block_generation'])
        forged=dict(by['Finding'],write_protected=False)
        with self.assertRaises(word_v3.V3Failure):word_v3.fill(data(doc),[forged],{forged['key']:'Synthetic replacement'})
        plan=dict(by['Plan'],write_protected=False)
        with self.assertRaises(word_v3.V3Failure):word_v3.fill(data(doc),[plan],{plan['key']:'Synthetic plan'})
        out=word_v3.fill(data(doc),result['fields'],{})
        parsed=Document(io.BytesIO(out));self.assertEqual(parsed.tables[0].cell(0,1).text,'')
        self.assertEqual(parsed.tables[1].cell(0,1).text,'')
        plain=Document();table=plain.add_table(rows=1,cols=2);table.cell(0,0).text='Plan'
        good=self.inspect(plain);field=good['fields'][0]
        self.assertFalse(good['import_manifest']['coverage'].get('selected_engine_issues'))
        result=word_v3.fill(data(plain),good['fields'],{field['key']:'Synthetic plan'})
        self.assertEqual(Document(io.BytesIO(result)).tables[0].cell(0,1).text,'Synthetic plan')

    def test_cell_append_preserves_original_paragraph_without_unqualified_wrap_repair(self):
        doc=Document();cell=doc.add_table(rows=1,cols=1).cell(0,0);cell.text='FIXED SOURCE GUIDANCE'
        old=etree.tostring(cell.paragraphs[0]._p)
        word_v3.write_cell(cell._tc,'SYNTHETIC synthetic',append=True)
        self.assertEqual(etree.tostring(cell.paragraphs[0]._p),old)
        self.assertIsNone(cell.paragraphs[1]._p.find('./'+qn('w:pPr')+'/'+qn('w:wordWrap')))
        self.assertEqual(cell.paragraphs[1].text,'SYNTHETIC synthetic')

    def test_exact_measured_source_quarantine_cannot_be_cleared_by_forged_writer(self):
        doc=Document();doc.add_paragraph('Reason: ______');raw=data(doc)
        baseline=word_v3.inspect(raw);anchor=baseline['fields'][0]['locator']
        declaration={'case':'synthetic negative control','failed_anchors':[anchor],
                     'evidence':['retained synthetic failure control']}
        with patch.dict(word_v3.KNOWN_LAYOUT_FAILURES,{hashlib.sha256(raw).hexdigest():declaration}):
            result=word_v3.inspect(raw)
            self.assertEqual(result['layout_quarantine']['source_sha256'],hashlib.sha256(raw).hexdigest())
            self.assertTrue(result['fields'][0]['write_protected'])
            forged=dict(result['fields'][0],write_protected=False)
            with self.assertRaises(word_v3.V3Failure) as refusal:
                word_v3.fill(raw,[forged],{forged['key']:'Synthetic replacement'})
            self.assertEqual(refusal.exception.code,'unqualified_layout')
            different=Document();different.add_paragraph('Plan: ______')
            self.assertIsNone(word_v3.inspect(data(different))['layout_quarantine'])


class TerminalChoiceWriteIns(unittest.TestCase):
    def form(self, *, control='glyph', checked=False, caption='Other, specify:', nested=False, question='Current symptoms (choose one):'):
        from .test_word_v3 import checkbox_control
        doc=Document()
        if nested:
            outer=doc.add_table(rows=1,cols=1).cell(0,0);outer.text='Clinical history'
            table=outer.add_table(rows=1,cols=2);owner=table.cell(0,0);owner.text=question
            neighbor=table.cell(0,1);neighbor.text='Past symptoms (choose one):'
            neighbor.add_paragraph('☐ Yes');neighbor.add_paragraph('☐ No')
            paragraphs=[owner.add_paragraph(),owner.add_paragraph()]
        else:
            doc.add_paragraph(question);paragraphs=[doc.add_paragraph(),doc.add_paragraph()]
        for index,p in enumerate(paragraphs):
            value=index==1 and checked
            if control=='ff':ff_field(p,'checkbox',checked=value,name='terminal'+str(index))
            elif control=='sdt':p._p.append(checkbox_control(checked=value))
            else:p.add_run('☒' if value else '☐')
            p.add_run(' '+('Yes' if index==0 else caption))
        doc.add_paragraph('Unrelated notes: ______')
        return data(doc)

    def mapped(self, raw):
        result=word_v3.inspect(raw)
        blanks=[f for f in result['fields']if f.get('source_questions')and f['source_questions'][0]['relation']=='adjacent_checkbox_answer'
                and f['locator']['target'].get('shape')=='after_label']
        self.assertEqual(len(blanks),1,'A terminal source-labelled Other answer was omitted')
        blank=blanks[0];owner=next(f for f in result['fields']if f['locator']['target']==blank['source_questions'][0]['owner_target'])
        return result,blank,owner

    def test_terminal_caption_keeps_control_and_complete_long_answer_separate(self):
        for control,nested in (('glyph',False),('sdt',False),('ff',False),('ff',True)):
            with self.subTest(control=control,nested=nested):
                raw=self.form(control=control,nested=nested);result,blank,owner=self.mapped(raw)
                self.assertEqual(owner['options'],['Yes','Other, specify'])
                self.assertEqual(blank['label'],'Other, specify:')
                self.assertEqual(word_v3.fill(raw,result['fields'],{}),raw)
                answer='Synthetic retained detail.\n'+'Complete long detail; '*70
                out=word_v3.fill(raw,result['fields'],{owner['key']:'Other, specify',blank['key']:answer})
                before=word_v3.SlotIndex(word_v3.open_package(raw));after=word_v3.SlotIndex(word_v3.open_package(out))
                self.assertEqual(word_v3.constrained_scalar_value(before,after,blank,True),answer.strip())
                # Readback normalizes surrounding whitespace; the physical
                # destination must nevertheless retain the entire supplied text.
                identity=blank['locator']['part']+'#'+blank['locator']['slot']
                self.assertTrue(word_v3.paragraph_text(after.by_id[identity].el)[0].endswith(answer))
                self.assertEqual(word_v3.constrained_option_states(after,owner),[False,True])
                self.assertEqual(Document(io.BytesIO(out)).paragraphs[-1].text,'Unrelated notes: ______')
                with ZipFile(io.BytesIO(raw))as old,ZipFile(io.BytesIO(out))as new:
                    self.assertTrue(all(old.read(n)==new.read(n)for n in old.namelist()if n!='word/document.xml'))

    def test_selected_original_is_usable_but_no_tick_or_resolution_is_invented(self):
        for checked in (False,True):
            raw=self.form(control='ff',checked=checked,caption='Other known:');result,blank,owner=self.mapped(raw)
            if not checked:
                with self.assertRaises(word_v3.V3Failure)as error:word_v3.fill(raw,result['fields'],{blank['key']:'Synthetic cause'})
                self.assertEqual(error.exception.code,'conditional_answer_unselected')
            else:
                out=word_v3.fill(raw,result['fields'],{blank['key']:'Synthetic cause'})
                self.assertEqual(word_v3.constrained_option_states(word_v3.SlotIndex(word_v3.open_package(out)),owner),[False,True])
            with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,result['fields'],{owner['key']:'Yes',blank['key']:'Synthetic cause'})

    def test_stale_removed_copied_changed_or_unowned_terminal_mapping_is_refused(self):
        raw=self.form();result,blank,owner=self.mapped(raw)
        for change in ('relation','renamed','copied','widened','no_owner','wrong_owner'):
            with self.subTest(change=change):
                fields=deepcopy(result['fields']);field=next(f for f in fields if f['key']==blank['key'])
                if change=='relation':field.pop('source_questions')
                elif change=='renamed':field['label']='Some other answer'
                elif change=='copied':field['locator']['target']['copies']=[deepcopy(field['locator'])]
                elif change=='widened':field['locator']['target']['start']=0
                elif change=='wrong_owner':field['source_questions'][0]['owner_label']='Other question'
                else:fields.remove(next(f for f in fields if f['key']==owner['key']))
                with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,fields,{owner['key']:'Other, specify',blank['key']:'Synthetic detail'})
        with self.assertRaises(word_v3.V3Failure):word_v3.fill(self.form(caption='Other known:'),result['fields'],{owner['key']:'Other, specify',blank['key']:'Synthetic detail'})

    def test_glyph_readback_is_source_owned_and_keeps_answer_glyphs_literal(self):
        raw=self.form();result,blank,owner=self.mapped(raw)
        answer='Clinical detail contains a literal ☐ and ☒; neither is an option.'
        out=word_v3.fill(raw,result['fields'],{owner['key']:'Other, specify',blank['key']:answer})
        before=word_v3.SlotIndex(word_v3.open_package(raw));after=word_v3.SlotIndex(word_v3.open_package(out))
        self.assertEqual(word_v3.constrained_scalar_value(before,after,blank,True),answer)
        self.assertEqual(word_v3.effective_values(raw,result['fields'],{owner['key']:'Other, specify',blank['key']:answer})[blank['key']],answer)
        for change in ('caption','source','owner','boxes','fingerprint'):
            with self.subTest(change=change):
                forged=deepcopy(blank)
                if change=='caption':forged['label']='Some other clinical answer'
                elif change=='source':forged['source_questions'][0]['source_text']='Some other printed source'
                elif change=='owner':forged['source_questions'][0]['owner_label']='Some other question'
                elif change=='boxes':forged['source_questions'][0]['owner_target']['boxes'][-1]['glyph']=999
                else:forged['locator']['fingerprint']='not-this-original'
                self.assertIsNone(word_v3.constrained_scalar_value(before,after,forged,True))
        identity=blank['locator']['part']+'#'+blank['locator']['slot']
        for change in ('caption','unknown_box'):
            changed=word_v3.SlotIndex(word_v3.open_package(out))
            paragraph=changed.by_id[identity].el
            for node in paragraph.iter(word_v3.q('w:t')):
                if change=='caption' and node.text and 'Other, specify:' in node.text:
                    node.text=node.text.replace('Other, specify:','Another caption:')
                elif change=='unknown_box' and node.text=='☒':node.text='X'
            self.assertIsNone(word_v3.constrained_scalar_value(before,changed,blank,True))
        different=word_v3.SlotIndex(word_v3.open_package(self.form(caption='Other known:')))
        self.assertIsNone(word_v3.constrained_scalar_value(different,after,blank,True))

    def test_nested_write_in_does_not_authorize_caption_cell_or_signing_neighbor(self):
        raw=self.form(control='ff',nested=True);result,blank,owner=self.mapped(raw)
        identity=blank['locator']['part']+'#'+blank['locator']['slot']
        index=word_v3.SlotIndex(word_v3.open_package(raw));cell=index.by_id[identity].parent
        for target in ({'type':'cell_append'},{'type':'cell_replace','all':True}):
            forged=deepcopy(blank);forged['label']='Ordinary narrative';forged.pop('source_questions')
            forged['locator'].update(slot=cell.path,fingerprint=word_v3.fingerprint(cell.el),target=target)
            with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[*result['fields'],{**forged,'key':'forged'}],{'forged':'Dr Synthetic Finalizer'})
        doc=Document(io.BytesIO(raw));doc.add_paragraph('I certify that I witnessed the patient sign this form.')
        doc.add_paragraph('Witness name: ______');protected=data(doc)
        mapped,blank,owner=self.mapped(protected)
        out=word_v3.fill(protected,mapped['fields'],{owner['key']:'Other, specify',blank['key']:'Synthetic detail'})
        self.assertEqual(Document(io.BytesIO(out)).paragraphs[-2].text,doc.paragraphs[-2].text)
        self.assertEqual(Document(io.BytesIO(out)).paragraphs[-1].text,'Witness name: ______')
        signature=next(f for f in mapped['fields']if f['label'].startswith('Witness name'))
        signature=deepcopy(signature);signature.update(kind='text',label='Ordinary narrative');signature.pop('write_protected',None);signature.pop('destination_roles',None)
        with self.assertRaises(word_v3.V3Failure):word_v3.fill(protected,[signature],{signature['key']:'Dr Synthetic Finalizer'})

    def test_prospective_author_and_reserved_area_do_not_gain_a_conditional_writer(self):
        for question in ('Witness must write the following:','Witness should write the following:',
                'Witness must enter the following:','Manager is required to write this response:',
                'Alex Reed cannot write this response:','Dr. Alex Reed should write this response:',
                'Alex Reed shall write this response:',
                'This area is for Alex Reed to enter the answer:',
                'This section is for Dr. Alex Reed to record symptoms:',
                'This box is reserved for another person to supply a statement:',
                'Answer to be entered by Alex Reed:'):
            with self.subTest(question=question):
                raw=self.form(caption='Other known:',question=question);result=word_v3.inspect(raw)
                self.assertFalse([f for f in result['fields']if f.get('source_questions')and f['source_questions'][0]['relation']=='adjacent_checkbox_answer'])
                choices=[f for f in result['fields']if f.get('options')and 'Other known' in f['options']]
                self.assertTrue(choices)
                for choice in choices:
                    forged=deepcopy(choice);forged.pop('write_protected',None);forged.pop('destination_roles',None)
                    with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[forged],{forged['key']:'Other known'})
                self.assertEqual(word_v3.fill(raw,result['fields'],{}),raw)
        for question in ('Review of systems:','Initial reading:','Signed consent receipt:',
                'Medication must include the following:','Patient must avoid the following:'):
            raw=self.form(caption='Other known:',question=question);result,blank,owner=self.mapped(raw)
            out=word_v3.fill(raw,result['fields'],{owner['key']:'Other known',blank['key']:'Synthetic observed detail'})
            self.assertEqual(word_v3.effective_values(raw,result['fields'],{owner['key']:'Other known',blank['key']:'Synthetic observed detail'})[blank['key']],'Synthetic observed detail')

    def test_unresolved_tail_and_human_source_do_not_gain_a_writer(self):
        for caption,question in (('Other details: More instructions','Current symptoms:'),('Other','Current symptoms:'),
                ('Other: _____ Location: _____','Current symptoms:'),('Other (I certify this document):','Current symptoms:'),
                ('Other, specify:','Witness attestation:'),('Other, specify:','1. I approve this document:'),
                ('Other, specify:','This section is for Alex Reed to write comments:')):
            with self.subTest(caption=caption,question=question):
                raw=self.form(caption=caption,question=question);result=word_v3.inspect(raw)
                self.assertFalse([f for f in result['fields']if f.get('source_questions')and f['source_questions'][0]['relation']=='adjacent_checkbox_answer'])
                self.assertEqual(word_v3.fill(raw,result['fields'],{}),raw)


class PrintedOptionBranchInstructions(unittest.TestCase):
    def form(self, option='Yes → If yes, focus on the following questions.', *, native=False):
        doc=Document();p=doc.add_paragraph('Symptoms present? ')
        if native:
            from .test_word_v3 import checkbox_control
            if native=='sdt':p._p.append(checkbox_control())
            else:ff_field(p,'checkbox',name='branch0')
            p.add_run(' No ')
            if native=='sdt':p._p.append(checkbox_control())
            else:ff_field(p,'checkbox',name='branch1')
            p.add_run(' '+option)
        else:p.add_run('☐ No ☐ '+option)
        doc.add_paragraph('Details: ______');doc.add_paragraph('Independent question? ☐ No ☐ Yes')
        return data(doc)

    def group(self, raw):
        fields=word_v3.inspect(raw)['fields'];return fields,next(f for f in fields if f['label']=='Symptoms present?')

    def test_branch_caption_and_instruction_have_separate_bound_meanings(self):
        for native in (False,True,'sdt'):
            raw=self.form(native=native);fields,choice=self.group(raw)
            self.assertEqual(choice['options'],['No','Yes'])
            self.assertEqual(word_v3.printed_options(choice),['No','Yes → If yes, focus on the following questions.'])
            self.assertIn('focus on the following questions.',choice['context'])
            out=word_v3.fill(raw,fields,{choice['key']:'Yes'})
            index=word_v3.SlotIndex(word_v3.open_package(out))
            self.assertEqual(word_v3.constrained_option_states(index,choice),[False,True])
            expected=Document(io.BytesIO(raw)).paragraphs[0].text
            actual=Document(io.BytesIO(out)).paragraphs[0].text
            self.assertEqual(actual.replace('☒','☐'),expected)
            self.assertEqual(Document(io.BytesIO(out)).paragraphs[1].text,'Details: ______')
            self.assertEqual(Document(io.BytesIO(out)).paragraphs[2].text,'Independent question? ☐ No ☐ Yes')
            self.assertEqual(word_v3.fill(raw,fields,{}),raw)

    def test_terminal_conditional_caption_keeps_a_separate_following_answer(self):
        raw=self.form('Yes → if yes, please expand:');fields,choice=self.group(raw)
        self.assertEqual(choice['options'],['No','Yes'])
        details=next(f for f in fields if f['label']=='Details')
        out=word_v3.fill(raw,fields,{choice['key']:'Yes',details['key']:'Synthetic complete detail'})
        doc=Document(io.BytesIO(out))
        self.assertEqual(doc.paragraphs[0].text,'Symptoms present? ☐ No ☒ Yes → if yes, please expand:')
        self.assertEqual(doc.paragraphs[1].text,'Details: Synthetic complete detail')
        self.assertEqual(doc.paragraphs[2].text,'Independent question? ☐ No ☐ Yes')

    def test_no_global_yes_alias_or_cross_question_boundary_is_created(self):
        for caption in ('Yes if yes, focus on the following questions.',
                'Yes → If no, focus on the following questions.',
                'Yes → If yes, explain: _____',
                'Yes → If yes, location? _____',
                'Yes → If yes, focus here → then answer another question.',
                'Yes → If yes,'):
            raw=self.form(caption);fields,choice=self.group(raw)
            if caption in ('Yes → If yes, explain: _____','Yes → If yes, location? _____'):
                # The established drawn conditional writer is a separate contract.
                continue
            self.assertNotEqual(choice['options'],['No','Yes'])
            with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,fields,{choice['key']:'Yes'})

    def test_changed_removed_or_forged_branch_mapping_cannot_select_another_box(self):
        raw=self.form();fields,choice=self.group(raw)
        for change in ('relation','context','options','printed','boxes'):
            forged=deepcopy(choice)
            if change=='relation':forged.pop('source_questions')
            elif change=='context':forged['context']='Some other printed context'
            elif change=='options':forged['options']=['Yes','No']
            elif change=='printed':forged['locator']['target']['printed']=['No','Yes']
            else:forged['locator']['target']['boxes'].reverse()
            with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[forged],{forged['key']:'Yes'})
        with self.assertRaises(word_v3.V3Failure):word_v3.fill(self.form('Yes → If yes, answer the following questions.'),fields,{choice['key']:'Yes'})

    def test_existing_selection_branch_retains_state_and_ordinary_history_arguments(self):
        for instruction in ('review the dose.','record the initial reading.',
                'document the signed referral received.'):
            raw=self.form('Yes → If yes, '+instruction);fields,choice=self.group(raw)
            self.assertEqual(choice['options'],['No','Yes'])
            self.assertFalse(choice.get('write_protected'))
            out=word_v3.fill(raw,fields,{choice['key']:'No'})
            self.assertEqual(word_v3.constrained_option_states(word_v3.SlotIndex(word_v3.open_package(out)),choice),[True,False])
            self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text.replace('☒','☐'),Document(io.BytesIO(raw)).paragraphs[0].text)
        raw=self.form();fields,choice=self.group(raw)
        out=word_v3.fill(raw,fields,{choice['key']:'No'})
        remapped=word_v3.inspect(out)['fields'];current=next(f for f in remapped if f['label']=='Symptoms present?')
        self.assertEqual(word_v3.fill(out,remapped,{}),out)
        next_output=word_v3.fill(out,remapped,{current['key']:'Yes'})
        self.assertEqual(word_v3.constrained_option_states(word_v3.SlotIndex(word_v3.open_package(next_output)),current),[False,True])

    def test_original_full_option_mapping_remains_usable_without_a_new_alias(self):
        raw=self.form();fields,choice=self.group(raw);relation=choice['source_questions'][-1]
        legacy=deepcopy(choice);legacy['options']=relation['printed'];legacy['context']=relation['original_context']
        legacy['locator']['target']=relation['original_target'];legacy.pop('source_questions')
        out=word_v3.fill(raw,[legacy],{legacy['key']:legacy['options'][-1]})
        self.assertEqual(word_v3.constrained_option_states(word_v3.SlotIndex(word_v3.open_package(out)),legacy),[False,True])
        with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[legacy],{legacy['key']:'Yes'})
        for change in ('caption','context','box'):
            forged=deepcopy(legacy)
            if change=='caption':forged['label']='Other clinical question'
            elif change=='context':forged['context']='Some other source'
            else:forged['locator']['target']['boxes'].reverse()
            with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[forged],{forged['key']:forged['options'][-1]})

    def test_human_branch_act_is_not_released_as_an_ordinary_yes(self):
        for native in (False,True,'sdt'):
            for instruction in ('I approve this document.','I certify that I witnessed the patient sign this consent form.',
                    'the line manager will review this form.','focus here. I approve this document.',
                    'focus here;I approve this document.','focus here.\nI approve this document.'):
                with self.subTest(native=native,instruction=instruction):
                    raw=self.form('Yes → If yes, '+instruction,native=native);fields,choice=self.group(raw)
                    self.assertTrue(choice.get('write_protected'))
                    forged=deepcopy(choice);forged.pop('write_protected',None);forged.pop('destination_roles',None);forged['label']='Symptoms present';forged['options']=['No','Yes']
                    with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[forged],{forged['key']:'Yes'})
                    self.assertEqual(word_v3.fill(raw,fields,{}),raw)

    def test_original_nested_relationship_and_full_option_remain_bound(self):
        from .test_question_boxes import IndependentNestedQuestionCells
        fixture=IndependentNestedQuestionCells()
        for control in ('glyph','ff'):
            raw=fixture.form(control=control)
            doc=Document(io.BytesIO(raw));cell=doc.tables[0].cell(0,0).tables[0].cell(0,0)
            cell.paragraphs[1].add_run(' → If yes, review the dose.')
            raw=data(doc);fields=word_v3.inspect(raw)['fields']
            choice=next(f for f in fields if f['label']=='1. Current symptoms (choose one):')
            relation=next(r for r in choice['source_questions'] if r['relation']=='option_branch_instruction')
            legacy=deepcopy(choice);legacy['options']=relation['printed'];legacy['context']=relation['original_context']
            legacy['locator']['target']=relation['original_target'];legacy['source_questions']=relation['original_relations']
            self.assertEqual([r['relation']for r in legacy['source_questions']],['nested_cell_options'])
            full=legacy['options'][0];out=word_v3.fill(raw,[legacy],{legacy['key']:full})
            self.assertEqual(word_v3.constrained_option_states(word_v3.SlotIndex(word_v3.open_package(out)),legacy),[True,False])
            self.assertEqual(Document(io.BytesIO(out)).tables[0].cell(0,0).tables[0].cell(0,1)._tc.xml,
                             Document(io.BytesIO(raw)).tables[0].cell(0,0).tables[0].cell(0,1)._tc.xml)
            with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[legacy],{legacy['key']:'Yes'})
            for change in ('relation','source','context','fingerprint','extra','target','kind','caption','options'):
                forged=deepcopy(legacy)
                if change=='relation':forged.pop('source_questions')
                elif change=='source':forged['source_questions'][0]['sources'][0]['fingerprint']='0'*64
                elif change=='context':forged['context']='Other source'
                elif change=='fingerprint':forged['locator']['fingerprint']='0'*64
                elif change=='extra':forged['locator']['extra']=[]
                elif change=='target':forged['locator']['target']['boxes'].reverse()
                elif change=='kind':forged['kind']='multi_choice'
                elif change=='caption':forged['label']='Different question'
                else:forged['options'].reverse()
                with self.assertRaises(word_v3.V3Failure):word_v3.fill(raw,[forged],{forged['key']:full})
            cell.paragraphs[1].add_run(';I approve this document.');human=data(doc)
            human_fields=word_v3.inspect(human)['fields'];candidate=next(f for f in human_fields if f['label']==choice['label'])
            candidate.pop('write_protected',None);candidate.pop('destination_roles',None)
            with self.assertRaises(word_v3.V3Failure):word_v3.fill(human,[candidate],{candidate['key']:candidate['options'][0]})
            self.assertEqual(word_v3.fill(raw,fields,{}),raw)
