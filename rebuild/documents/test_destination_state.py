"""roles-2: ONE three-state destination decision (ordinary / protected / uncertain), read from the ORIGINAL's printed
label AND structure, on both engines.

Synthetic blank forms only, built in memory. Ids: REG = the 18 clinical labels f5fe720f no longer generated; HOLD = the
17 must-hold labels; L1:/L2:/L3: = the 28 recorded residual shapes (docs/rebuild/v1/destination-roles-20260930/
residuals.json); PC = positive controls (an ordinary field beside each refused or held place still fills).
"""
import io
import unittest
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import word_v3 as w

SYN = 'SYNTHETIC-ANSWER'
BODY = 'word/document.xml#'
LINE = '______________________'

#: The 18 labels f5fe720f no longer generated, and how roles-2 classifies each (state, basis).
REG18 = {
    'Troponin (initial)': ('ordinary', 'measure_head'),
    'Lactate (initial)': ('ordinary', 'measure_head'),
    'Dose (initial)': ('ordinary', 'measure_head'),
    'Troponin initial': ('ordinary', 'measure_head'),
    'Troponin — Initial (ng/L)': ('ordinary', 'measure_unit'),
    'Initial (baseline) weight': ('ordinary', 'attributive'),
    'Initial and ongoing management': ('ordinary', 'attributive'),
    'Visit type (initial/follow-up)': ('ordinary', 'sequence_option'),
    'Initial / repeat': ('ordinary', 'sequence_option'),
    'Gene expression signature': ('ordinary', 'scientific_signature'),
    'Consent form signed by patient': ('ordinary', 'document_signed_by'),
    'Drug chart signed by prescriber': ('ordinary', 'document_signed_by'),
    'Patient signature obtained': ('ordinary', 'signature_state'),
    'Witness signature obtained': ('ordinary', 'signature_state'),
    # Not restored, classified honestly: the label alone cannot tell a sign-off status or the time of a signing from
    # the place the consultant (or the patient on admission) signs - uncertain: never generated, the clinician types it.
    'Signed off by consultant': ('uncertain', 'sign_off'),
    'Signed on admission': ('uncertain', 'signed_when'),
    # A truly signing place: 'Initial' alone is the imperative / the initials box (the same words as the must-hold
    # 'Initial:'); 'Patient initials' is the patient's own initials (the same grammar as 'Nurse initials').
    'Initial': ('protected', 'bare_initial'),
    'Patient initials': ('protected', 'initials_noun'),
}
HOLD17 = ('Patient signature', 'Initials', 'Witness signature', 'Consultant signature', 'Signature of examining clinician',
          "Participant's signature", 'Nurse initials', 'Initial here', 'Initial each page', 'Signature and date',
          'Signature (patient)', 'Signed by (witness)', 'Nurse (initial)', 'Initial:', 'Staff initial', 'Nurse initial',
          'Initials — Nurse')


def data(doc):
    out = io.BytesIO(); doc.save(out); return out.getvalue()


def xml(blob):
    return zipfile.ZipFile(io.BytesIO(blob)).read('word/document.xml').decode()


def tabbed(p, text):
    parts = text.split('\t')
    for i, part in enumerate(parts):
        if part:
            p.add_run(part)
        if i < len(parts) - 1:
            p.add_run()._r.append(OxmlElement('w:tab'))
    return p


def lines(*texts, heading='Clinic record'):
    doc = Document()
    if heading:
        doc.add_heading(heading, level=2)
    for text in texts:
        tabbed(doc.add_paragraph(), text)
    doc.add_paragraph('Plan: ______')
    return data(doc)


def bordered(table):
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); borders.append(edge)
    table.cell(0, 0)._tc.get_or_add_tcPr().append(borders)
    return table


def box_after(question, inner=(), heading='Consultant review'):
    doc = Document(); doc.add_heading(heading, level=2)
    doc.add_paragraph(question)
    cell = bordered(doc.add_table(rows=1, cols=1)).cell(0, 0)
    for i, text in enumerate(inner):
        tabbed(cell.paragraphs[0] if i == 0 else cell.add_paragraph(), text)
    doc.add_paragraph('Plan: ______')
    return data(doc)


def two_col(rows, br=False):
    doc = Document(); doc.add_heading('Clinic record', level=2)
    table = doc.add_table(rows=len(rows), cols=2); table.style = 'Table Grid'
    for i, label in enumerate(rows):
        cell = table.cell(i, 0)
        if '|' in label:
            first, second = label.split('|', 1)
            if br:
                p = cell.paragraphs[0]; p.add_run(first); p.add_run()._r.append(OxmlElement('w:br')); p.add_run(second)
            else:
                cell.paragraphs[0].add_run(first); cell.add_paragraph(second)
        else:
            cell.paragraphs[0].add_run(label)
    return data(doc)


def control_line(caption, alias):
    doc = Document(); doc.add_heading('Clinic record', level=2)
    p = doc.add_paragraph(caption)
    sdt = OxmlElement('w:sdt'); pr = OxmlElement('w:sdtPr')
    al = OxmlElement('w:alias'); al.set(qn('w:val'), alias); pr.append(al); pr.append(OxmlElement('w:text'))
    pr.append(OxmlElement('w:showingPlcHdr')); sdt.append(pr)
    content = OxmlElement('w:sdtContent'); r = OxmlElement('w:r'); t = OxmlElement('w:t')
    t.text = 'Click or tap here to enter text.'; r.append(t); content.append(r); sdt.append(content)
    p._p.append(sdt)
    doc.add_paragraph('Plan: ______')
    return data(doc)


def report(blob):
    fields = w.inspect(blob)['fields']
    return fields, w.destination_roles(blob, fields)


def generated(field, rep):
    """Whether generation asks this field (v3): an answer kind, not write-protected, no role violation, not held."""
    return (field['kind'] in ('text', 'long_text', 'date', 'time', 'choice', 'multi_choice', 'checkbox')
            and not field.get('write_protected') and field['key'] not in rep['violations']
            and field['key'] not in rep['held'])


def by_label(fields):
    return {f['label']: f for f in fields}


def manual(blob, slot, kind='text', label='Manual'):
    try:
        return w.locate(blob, [{'slot': BODY + slot, 'kind': kind, 'label': label}])[0], None
    except w.V3Failure as exc:
        return None, exc.code


class TheEighteenAndTheSeventeen(unittest.TestCase):
    """(a) each of the 18 labels: classified, and generated and exported where clinical; (b) the 17 held."""

    def test_each_of_the_18_is_classified(self):
        for label, (state, basis) in REG18.items():
            with self.subTest(label=label):
                reading = w.destination_state(label)
                self.assertEqual((reading['state'], reading['basis']), (state, basis))
                self.assertEqual(reading['reading'], w.ROLE_READING)

    def test_the_clinical_ones_are_generated_and_filled_on_v3(self):
        for label, (state, _) in REG18.items():
            with self.subTest(label=label):
                blob = lines(label + ': ______')
                fields, rep = report(blob)
                field = fields[0]
                if state == 'ordinary':
                    self.assertNotIn(field['kind'], ('signature', 'initials'))
                    self.assertTrue(generated(field, rep))
                    self.assertIn(SYN, xml(w.fill(blob, fields, {field['key']: SYN})))
                    self.assertEqual(w.legacy_role_kind(label), '')
                else:
                    self.assertIn(field['kind'], ('signature', 'initials'))
                    self.assertFalse(generated(field, rep))
                    self.assertTrue(w.legacy_role_kind(label))

    def test_the_17_must_hold_labels_are_protected_on_both_engines(self):
        for label in HOLD17:
            with self.subTest(label=label):
                self.assertEqual(w.destination_state(label)['state'], 'protected')
                self.assertIn(w.legacy_role_kind(label), ('signature', 'initials'))
                blob = lines(label + ' ______' if label.endswith(':') else label + ': ______')
                fields, rep = report(blob)
                self.assertIn(fields[0]['kind'], ('signature', 'initials'))
                self.assertTrue(any(r['class'] == 'signing' for r in rep['roles']))
                self.assertEqual(manual(blob, 'b/p[1]')[1], 'destination_role')

    def test_an_uncertain_place_is_never_generated_and_takes_a_clinicians_text_field(self):
        for label in ('Signed off by consultant', 'Signed on admission', 'Sign', 'Reviewed and signed',
                      'Wristband (initial)'):
            with self.subTest(label=label):
                blob = lines(label + ': ______')
                fields, rep = report(blob)
                self.assertEqual(w.destination_state(label)['state'], 'uncertain')
                self.assertIn(fields[0]['kind'], ('signature', 'initials'))
                self.assertFalse(generated(fields[0], rep))
                self.assertIn(fields[0]['key'], rep['uncertain_places'])
                added, code = manual(blob, 'b/p[1]')
                self.assertIsNone(code)
                rest = fields[1:] + [added]
                rep = w.destination_roles(blob, rest)
                self.assertIn(added['key'], rep['held'])           # held: never generated
                self.assertNotIn(added['key'], rep['violations'])  # not refused: the clinician may type it
                self.assertIn('TYPED', xml(w.fill(blob, rest, {added['key']: 'TYPED'})))


class Structure(unittest.TestCase):
    """The same decision from structure: a sequence word beside a bare 'Initial', a unit in the row."""

    def test_a_bare_initial_beside_a_sequence_word_is_a_qualifier(self):
        blob = two_col(['Initial', 'Repeat', 'Plan'])
        fields, rep = report(blob)
        initial = by_label(fields)['Initial']
        self.assertEqual(initial['kind'], 'text')
        self.assertTrue(generated(initial, rep))
        self.assertEqual(w.legacy_role_kind('Initial', siblings=['Repeat', 'Plan']), '')
        # alone, it is the initials place
        fields, rep = report(two_col(['Initial', 'Plan']))
        self.assertEqual(by_label(fields)['Initial']['kind'], 'initials')
        self.assertEqual(w.legacy_role_kind('Initial', siblings=['Plan']), 'initials')

    def test_a_lab_table_reads_its_row_label_and_column_header_each_as_itself(self):
        """'Troponin (initial)' in a results table with a 'Result' column and a units column is a measurement (the
        engine's joined label 'Troponin (initial) — Result' closes no qualifier); 'Nurse (initial)' and 'Nurse
        initials' rows, and an 'Initials' column, stay held; an 'Initial' column beside 'Repeat' is a qualifier."""
        def lab(rows, headers):
            doc = Document(); table = doc.add_table(rows=len(rows) + 1, cols=len(headers)); table.style = 'Table Grid'
            for c, h in enumerate(headers):
                table.cell(0, c).text = h
            for r, row in enumerate(rows, start=1):
                for c, v in enumerate(row):
                    table.cell(r, c).text = v
            doc.add_paragraph('Plan: ______')
            return data(doc)
        fields, rep = report(lab((('Troponin (initial)', '', 'ng/L'), ('Fibre initial', '', 'mg'), ('Nurse initials', '', ''),
                                  ('Nurse (initial)', '', '')), ('Test', 'Result', 'Units')))
        by = by_label(fields)
        self.assertTrue(generated(by['Troponin (initial) — Result'], rep))
        self.assertTrue(generated(by['Fibre initial — Result'], rep))
        self.assertFalse(generated(by['Nurse initials — Result'], rep))
        self.assertFalse(generated(by['Nurse (initial) — Result'], rep))
        fields, rep = report(lab((('Troponin', '', ''),), ('Analyte', 'Initial', 'Repeat')))
        self.assertTrue(generated(by_label(fields)['Troponin — Initial'], rep))
        fields, rep = report(lab((('Nurse', '', ''),), ('Staff', 'Initials', 'Date')))
        self.assertFalse(generated(by_label(fields)['Nurse — Initials'], rep))
        self.assertTrue(generated(by_label(fields)['Nurse — Date'], rep))

    def test_a_person_beside_a_sequence_word_stays_held(self):
        self.assertEqual(w.legacy_role_kind('Nurse initial', siblings=['Repeat']), 'initials')

    def test_an_unknown_head_is_uncertain_and_a_unit_makes_it_a_measurement(self):
        self.assertEqual(w.destination_state('Wristband (initial)')['state'], 'uncertain')
        self.assertEqual(w.destination_state('Phlebotomist (initial)')['state'], 'protected')     # a person's head
        self.assertEqual(w.destination_state('Fibre initial (mg)')['basis'], 'measure_unit')
        self.assertEqual(w.destination_state('Fibre initial', unit=True)['state'], 'ordinary')
        # a person never becomes a measurement, unit or not
        self.assertEqual(w.destination_state('Nurse (initial) (mg)')['state'], 'protected')


class Residuals(unittest.TestCase):
    """(d) the 28 recorded residual shapes: every unsafe write refused or held, each with an ordinary positive control."""

    def assert_held(self, blob, labels_held, controls=('Plan',)):
        fields, rep = report(blob)
        for f in fields:
            if f['label'] in controls:
                self.assertTrue(generated(f, rep), f['label'])
                self.assertIn(SYN, xml(w.fill(blob, fields, {f['key']: SYN})))
        held = [f for f in fields if not generated(f, rep)]
        self.assertEqual(len([f for f in fields if f['label'] not in controls]), len(held))
        self.assertEqual(sorted(f['label'] for f in held), sorted(labels_held))
        return fields, rep

    def test_l1_rr1_01_ordinary_areas_inside_a_signing_box_are_its_signing_places(self):
        blob = box_after('Signed by the consultant:', [LINE, 'Signature, date: ______', 'Dr ' + LINE + ' (consultant)',
                                                       LINE + '    Date: ______'])
        fields, rep = report(blob)
        held = [f for f in fields if not generated(f, rep)]
        self.assertEqual(len(held), 4)
        date = [f for f in fields if f['label'] == 'Date']
        self.assertEqual(len(date), 1)                       # PC: the date beside the signature line
        self.assertTrue(generated(date[0], rep))
        self.assertIn('2026-09-30', xml(w.fill(blob, fields, {date[0]['key']: '2026-09-30'})))

    def test_l1_rr1_02_l2_rr3_01_a_signature_joined_to_a_date_is_one_signing_place(self):
        for label in ('Signature, date', 'Signature - date', 'Signature, date/time', 'Patient signature, date',
                      'Print name, sign and date', 'Patient to sign and date', 'Signed by (name, date)',
                      'Electronically signed by (name, date)'):
            with self.subTest(label=label):
                self.assertEqual(w.destination_state(label)['state'], 'protected')
                self.assert_held(lines(label + ': ______'), [label])

    def test_l1_rr1_03_l1_rr1_04_l1_rr1_09_l1_rr1_11_l2_rr3_02_l1_f4_signing_labels(self):
        for label in ('Initial 1', 'Initial 2', 'Initial x2', 'Initial on each page', 'Initial against each entry',
                      'Sign in box', 'To be signed', 'Signed consultant', 'Signed patient', 'Consultant signed',
                      'Registrar signed', 'Patient signed', 'Initial in box', 'Sign below', 'Sign and print name',
                      'Sign & stamp', 'Please sign', 'Given by (sign)', 'Countersignature', 'Initialled by',
                      'Nurse sign', 'Print name and sign', 'Name and sign', 'Countersign', 'Witness to sign',
                      'Nurse 2 sign', 'Patient signs here', 'Consultant sign-off', 'Time out signed by'):
            with self.subTest(label=label):
                self.assertEqual(w.destination_state(label)['state'], 'protected')
                self.assertTrue(w.legacy_role_kind(label))
                blob = lines(label + ': ______')
                self.assert_held(blob, [label])
                self.assertEqual(manual(blob, 'b/p[1]')[1], 'destination_role')
        for label in ('Sign', 'Reviewed and signed'):
            with self.subTest(label=label):
                self.assertEqual(w.destination_state(label)['state'], 'uncertain')
                self.assert_held(lines(label + ': ______'), [label])

    def test_l1_rr1_06_a_writing_line_under_or_over_a_signing_caption(self):
        cases = (
            (('Consultant signature:', '', LINE), 'b/p[3]'),
            ((LINE, 'Consultant signature'), 'b/p[1]'),
            (('Consultant signature:', LINE + ' (use black ink)'), 'b/p[2]'),
            (('Signed by the consultant:', 'Dr ' + LINE), 'b/p[2]'),
        )
        for texts, slot in cases:
            with self.subTest(texts=texts):
                blob = lines(*texts)
                fields, rep = report(blob)
                self.assertFalse(any(generated(f, rep) for f in fields if f['label'] != 'Plan'))
                self.assertEqual(manual(blob, slot)[1], 'destination_role')
        # (d) the date on the signature line stays an ordinary answer (PC02)
        blob = lines('Consultant signature:', LINE + '    Date: ______')
        fields, rep = report(blob)
        date = [f for f in fields if f['label'] == 'Date' and generated(f, rep)]
        self.assertEqual(len(date), 1)
        self.assertIn('2026-09-30', xml(w.fill(blob, fields, {date[0]['key']: '2026-09-30'})))

    def test_l1_f1_captions_side_by_side_under_lines(self):
        blob = lines(LINE + '\t\t' + LINE, 'Patient signature\t\t\tDate')
        fields, rep = report(blob)
        self.assertFalse(any(generated(f, rep) for f in fields if f['label'] != 'Plan'))

    def test_l1_f2_a_declaration_above_captions_under_lines(self):
        for texts, protected, uncertain in (
                (('Please complete in block capitals.', LINE, 'Signature of clinician', LINE, 'Name and grade'),
                 'b/p[2]', 'b/p[4]'),
                (('I consent to this procedure.', LINE, 'Patient signature', LINE, 'Witness signature'), 'b/p[2]', None),
                (('I confirm the above information is correct.', LINE, 'Patient signature', LINE, 'Print name'),
                 'b/p[2]', 'b/p[4]')):
            with self.subTest(texts=texts[0]):
                blob = lines(*texts)
                fields, rep = report(blob)
                self.assertFalse(any(generated(f, rep) for f in fields if f['label'] != 'Plan'))
                self.assertEqual(manual(blob, protected)[1], 'destination_role')
                if uncertain:
                    # the line between a signing caption and its own ordinary caption: uncertain (a typed name fills)
                    added, code = manual(blob, uncertain)
                    self.assertIsNone(code)
                    self.assertIn(added['key'], w.destination_roles(blob, [added])['held'])

    def test_l1_f3_l1_rr1_05_a_label_cell_wrapped_over_two_paragraphs(self):
        for rows, br in ((['Consultant|signature', 'Plan'], False), (['Nurse|initials', 'Plan'], False),
                         (['Consultant|signature', 'Plan'], True), (['Patient|signature', 'Plan'], False)):
            with self.subTest(rows=rows, br=br):
                blob = two_col(rows, br=br)
                fields, rep = report(blob)
                first = fields[0]
                self.assertIn(first['kind'], ('signature', 'initials'))
                self.assertFalse(generated(first, rep))
                self.assertTrue(generated(fields[1], rep))

    def test_l1_f5_a_content_control_beside_a_signing_caption(self):
        for alias in ('Answer', 'Text1', 'Clinician'):
            with self.subTest(alias=alias):
                blob = control_line('Consultant signature: ', alias)
                fields, rep = report(blob)
                self.assertEqual(fields[0]['kind'], 'signature')
                self.assertFalse(generated(fields[0], rep))
        blob = control_line('Hospital number: ', 'Answer')       # PC: an ordinary caption keeps its control
        fields, rep = report(blob)
        self.assertTrue(generated(fields[0], rep))

    def test_l1_f6_a_countersignature_or_signatures_box(self):
        for question in ('Please provide your countersignature', 'Countersignature?', 'Signatures?',
                         'Please provide signatures below'):
            with self.subTest(question=question):
                blob = box_after(question, heading='Review')
                fields, rep = report(blob)
                self.assertEqual(fields[0]['kind'], 'signature')
                self.assertEqual(manual(blob, 'b/tbl[0]/tr[0]/tc[0]', 'long_text')[1], 'destination_role')
        blob = box_after('Please describe the presenting complaint?', heading='Review')   # PC: an ordinary box
        fields, rep = report(blob)
        self.assertTrue(generated(fields[0], rep))

    def test_l1_f7_person_captions_inside_a_signatures_box(self):
        for inner in (['Patient: ______________', 'Parent/guardian: ______________'],
                      ['Consultant: ______________', 'Registrar: ______________']):
            with self.subTest(inner=inner):
                fields, rep = report(box_after('Signatures', inner, heading='Consent'))
                self.assertFalse(any(generated(f, rep) for f in fields if f['label'] != 'Plan'))
        # PC (R2-01): an ordinary question inside a consent frame keeps its answer
        fields, rep = report(box_after('Signatures', ['Patient name: ______________', 'Date: ______'],
                                       heading='Consent'))
        self.assertTrue(all(generated(f, rep) for f in fields))

    def test_l1_f8_guidance_between_a_caption_and_its_line_leaves_the_line_held(self):
        for guidance in ('Use black ink.', '(use black ink)'):
            with self.subTest(guidance=guidance):
                fields, rep = report(lines('Consultant signature', guidance, LINE))
                self.assertFalse(any(generated(f, rep) for f in fields if f['label'] != 'Plan'))

    def test_l1_f9_a_drawn_or_wide_line_under_a_signing_caption_refuses_a_manual_add(self):
        doc = Document(); doc.add_heading('Clinic record', level=2)
        p = doc.add_paragraph()
        bdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom'); bottom.set(qn('w:val'), 'single'); bdr.append(bottom)
        p._p.get_or_add_pPr().append(bdr)
        doc.add_paragraph('Consultant signature'); doc.add_paragraph('Plan: ______')
        self.assertEqual(manual(data(doc), 'b/p[1]')[1], 'destination_role')
        for rule in ('＿＿＿＿＿＿＿＿＿＿＿', '‗‗‗‗‗‗‗‗'):
            with self.subTest(rule=rule):
                self.assertEqual(manual(lines('Consultant signature', rule), 'b/p[2]')[1], 'destination_role')

    def test_l1_rr1_07_a_signing_box_read_through_a_bookmark_a_control_empties_or_a_layout_cell(self):
        def form(kind):
            doc = Document(); doc.add_heading('Consultant review', level=2)
            if kind == 'nested':
                cell = doc.add_table(rows=1, cols=1).cell(0, 0)
                cell.paragraphs[0].add_run('Signed by the consultant:')
                bordered(cell.add_table(rows=1, cols=1))
                slot = 'b/tbl[0]/tr[0]/tc[0]/tbl[0]/tr[0]/tc[0]'
            else:
                p = doc.add_paragraph('Signed by the consultant:')
                if kind == 'sdt':
                    sdt = OxmlElement('w:sdt'); sdt.append(OxmlElement('w:sdtPr')); content = OxmlElement('w:sdtContent')
                    p._p.addprevious(sdt); content.append(p._p); sdt.append(content)
                for _ in range(3 if kind == 'empties' else 0):
                    doc.add_paragraph('')
                table = bordered(doc.add_table(rows=1, cols=1))
                if kind == 'bookmark':
                    mark = OxmlElement('w:bookmarkStart'); mark.set(qn('w:id'), '0'); mark.set(qn('w:name'), 'x')
                    table._tbl.addprevious(mark)
                slot = 'b/tbl[0]/tr[0]/tc[0]'
            doc.add_paragraph('Plan: ______')
            return data(doc), slot
        for kind in ('nested', 'bookmark', 'sdt', 'empties'):
            with self.subTest(kind=kind):
                blob, slot = form(kind)
                self.assertEqual(manual(blob, slot, 'long_text')[1], 'destination_role')
                fields, rep = report(blob)
                plan = by_label(fields)['Plan']
                self.assertIn(SYN, xml(w.fill(blob, fields, {plan['key']: SYN})))


class RealFormShapes(unittest.TestCase):
    """(d) the development-form exposures (c/word-census exposure.txt), rebuilt synthetically: each unsafe write is
    refused or held, each beside an ordinary field that still fills."""

    def test_d15_a_signature_line_read_as_three_areas_is_one_signing_run_and_its_date_fills(self):
        blob = lines('Signature:……………........................…………\t\t\t\tDate: _____/_____/_____')
        fields, rep = report(blob)
        run = [f for f in fields if f['locator']['target'].get('start', 99) < 50 and f['label'] != 'Plan']
        self.assertEqual(len(run), 3)
        self.assertFalse(any(generated(f, rep) for f in run))
        date = next(f for f in fields if f['kind'] == 'date')                 # PC02
        self.assertTrue(generated(date, rep))
        self.assertIn('30', xml(w.fill(blob, fields, {date['key']: '30'})))

    def test_d01_an_electronic_signature_attestation_keeps_its_signer_held(self):
        blob = lines('Electronically Signed On <<insert date and time>> By:')
        fields, rep = report(blob)
        by = next(f for f in fields if f['locator']['target']['start'] > 50)
        self.assertEqual(by['kind'], 'signature')
        date = next(f for f in fields if f['kind'] == 'date')                 # the attested date stays an answer
        self.assertTrue(generated(date, rep))

    def test_d17_a_sign_placeholder_is_held_and_a_name_placeholder_fills(self):
        doc = Document(); doc.add_heading('Letter', level=1); doc.add_paragraph('Dear [Patient name],')
        doc.add_paragraph('Yours sincerely,'); doc.add_paragraph('[Sign]'); doc.add_paragraph('Plan: ______')
        blob = data(doc)
        fields, rep = report(blob)
        sign = next(f for f in fields if f['locator']['slot'] == 'b/p[3]')
        self.assertIn(sign['key'], rep['uncertain_places'])
        self.assertFalse(generated(sign, rep))
        name = next(f for f in fields if f['locator']['slot'] == 'b/p[1]')
        self.assertTrue(generated(name, rep))

    def test_d15_an_authors_note_about_the_wording_is_no_answer(self):
        fields, rep = report(lines('You will be asked to give samples <if POS collects genetic data> during the visit.'))
        self.assertFalse(generated(fields[0], rep))
        self.assertIn(fields[0]['key'], rep['uncertain_places'])
        fields, rep = report(lines('Dear <Patient name>,'))                    # PC: a placeholder naming its answer
        self.assertTrue(generated(fields[0], rep))

    def test_d29_an_agent_heading_of_a_signing_row_is_held_and_the_name_fills(self):
        doc = Document(); doc.add_heading('Review', level=1)
        table = doc.add_table(rows=1, cols=4); table.style = 'Table Grid'
        for i, text in enumerate(('Completed by:', 'Name:', 'Signature:', 'Date:')):
            table.cell(0, i).text = text
        doc.add_paragraph('Plan: ______')
        fields, rep = report(data(doc))
        by = by_label(fields)
        self.assertFalse(generated(by['Completed by'], rep))
        self.assertTrue(generated(by['Name'], rep))
        self.assertTrue(generated(by['Date'], rep))

    def test_d36_a_caption_line_of_several_segments_holds_its_whole_writing_line(self):
        fields, rep = report(lines('', 'Signature of Participant\t\t\tPrint Name of Participant\t\t\tDate', '',
                                   heading='Signatures'))
        self.assertFalse(generated(fields[0], rep))

    def test_d36_an_initials_segment_with_no_area_refuses_a_manual_add_and_the_date_fills(self):
        text = 'Circle Yes or No:     Yes     No     Initials:          Date:  '
        blob = lines(text, heading='Signatures')
        fields, rep = report(blob)
        date = next(f for f in fields if f['kind'] == 'date')
        self.assertTrue(generated(date, rep))
        at = text.index('Initials:') + len('Initials:') + 2
        roles = [r for r in rep['roles'] if r['source'] == 'governed_line']
        self.assertTrue(any(r['class'] == 'signing' and r['start'] <= at < r['end'] for r in roles))

    def test_a_first_person_consent_declaration_is_consent_wording_and_a_consent_question_is_ordinary(self):
        fields, rep = report(lines('I have been asked to give consent for my child to participate in this medical study:',
                                   'Consent obtained? ______'))
        declaration, question = fields[0], fields[1]
        self.assertFalse(generated(declaration, rep))
        self.assertIn(declaration['key'], rep['uncertain_places'])
        self.assertTrue(generated(question, rep))
        for label in ('I agree to take part in this study', 'I hereby consent to the procedure'):
            self.assertEqual(w.destination_state(label)['state'], 'uncertain')
        for label in ("I declare that I have taken note of the patient's withdrawal of consent", 'Consent obtained?',
                      'I understand that the applicant was examined', 'I want to take part in this study'):
            self.assertEqual(w.destination_state(label)['state'], 'ordinary', label)


class PositiveControls(unittest.TestCase):
    """Ordinary answers beside the new readings keep generating (refuter batteries' over-hold shapes)."""

    def test_a_square_bracket_hint_inside_a_question_is_the_forms_own_hint(self):
        for text in ('Ethnicity: White / Other [If other, please specify]', 'Have you had surgery? [If yes, give details]',
                     'Mobility aid: [e.g. walking stick, frame]', 'Referral details: [if any]'):
            with self.subTest(text=text):
                fields, rep = report(lines(text))
                self.assertTrue(generated(fields[0], rep))

    def test_options_beside_a_signature_caption_are_no_part_of_it(self):
        blob = lines('Reviewed: ☐ Yes ☐ No    Signature: ______')
        fields, rep = report(blob)
        reviewed = by_label(fields)['Reviewed']
        self.assertEqual(reviewed['kind'], 'choice')
        self.assertTrue(generated(reviewed, rep))

    def test_a_sentence_before_a_caption_is_not_its_caption(self):
        fields, rep = report(lines('Report digitally signed. Transcribed by: ________'))
        self.assertTrue(generated(by_label(fields)['Transcribed by'], rep))


class ConsentChoices(unittest.TestCase):
    """A factual consent question is ordinary; the signer's own consent choice is a consent ACTION (uncertain: never
    ticked from the notes, the clinician may tick it). Measured on development consent forms (D15, D33, D34, D37)."""

    def test_the_signers_consent_choice_is_held_and_a_consent_question_is_asked(self):
        for text in ('I agree that my samples may be kept in a biobank:  ☐ Yes   ☐ No',
                     'I give consent to store my coded data:  ☐ Yes   ☐ No',
                     '<if POS collecting genetic data>I give consent to share my data:  ☐ Yes  ☐ No',
                     '(~Include if appropriate. ~) I agree that my doctor may contact me:  ☐ Yes  ☐ No'):
            with self.subTest(text=text):
                blob = lines(text)
                fields, rep = report(blob)
                choice = next(f for f in fields if f['kind'] == 'choice')
                self.assertIn(choice['key'], rep['held'])
                self.assertNotIn(choice['key'], rep['violations'])
                w.fill(blob, fields, {choice['key']: 'Yes'})                    # the clinician may tick it
        fields, rep = report(lines('Consent obtained?  ☐ Yes  ☐ No'))
        self.assertTrue(generated(fields[0], rep))

    def test_printed_yes_no_words_under_a_consent_sentence_and_a_consent_question_asked_of_the_signer(self):
        for question in ('I agree that my samples and related health information may be kept in a biobank.',
                         'Will you allow your samples to be stored for future research?',
                         'Do you agree to be contacted about future studies?'):
            with self.subTest(question=question):
                fields, rep = report(lines(question, '\t YES\t\tNO', heading='Optional studies'))
                choice = next(f for f in fields if f['kind'] == 'choice')
                self.assertIn(choice['key'], rep['held'])
        for label in ('Do you allow pets at home?', 'Did the patient consent?', 'Can you allow time for the test?'):
            self.assertEqual(w.destination_state(label)['state'], 'ordinary', label)


class LegacyEngineRoles(unittest.TestCase):
    """The legacy 'roles' operation reads the label cell's printed words (paragraphs and line breaks as spaces), its
    stacked captions one by one, and the labels beside it in its table."""

    def roles(self, blob):
        import tempfile
        from pathlib import Path
        from . import engine
        with tempfile.TemporaryDirectory() as tmp:
            ws = Path(tmp); (ws / 'input.docx').write_bytes(blob)
            found = engine.process(ws, {'operation': 'inspect', 'format': 'docx', 'engine': 'legacy'})
            roles = engine.process(ws, {'operation': 'roles', 'format': 'docx', 'engine': 'legacy',
                                        'mapping': found['fields'], 'answers': {}})
        return {f['label']: roles['legacy_roles'][f['key']] for f in found['fields']}

    def test_l1_rr1_05_a_wrapped_label_cell_is_read_as_its_words(self):
        for br in (False, True):
            with self.subTest(br=br):
                rows = self.roles(two_col(['Consultant|signature', 'Nurse|initials', 'Plan'], br=br))
                self.assertEqual(rows['Consultantsignature']['printed_label'], 'Consultant signature')
                self.assertEqual((rows['Consultantsignature']['state'], rows['Nurseinitials']['state']),
                                 ('protected', 'protected'))
                self.assertEqual(rows['Plan']['state'], 'ordinary')

    def test_l1_f1_two_captions_stacked_in_one_label_cell_are_uncertain(self):
        for rows_in in (['Signature|Date', 'Plan'], ['Date of|signature', 'Plan']):
            with self.subTest(rows=rows_in):
                rows = self.roles(two_col(rows_in))
                held = next(v for k, v in rows.items() if k != 'Plan')
                self.assertEqual((held['state'], held['basis']), ('uncertain', 'stacked_caption'))
                fields, rep = report(two_col(rows_in))           # the same reading on contract v3
                self.assertIn(fields[0]['key'], rep['uncertain_places'])
                self.assertFalse(generated(fields[0], rep))

    def test_two_captions_a_tab_separates_in_one_label_cell_are_uncertain(self):
        """'Signature of patient<tab>Date' prints two captions over one answer cell. Read as joined words it is the date
        question 'Signature of patient date', which released the legacy row f5fe720f held (its runs joined to
        'patientDate') and let contract v3 generate a date there. Each caption is read alone, as stacked lines are."""
        def tab_col(rows):
            doc = Document(); doc.add_heading('Clinic record', level=2)
            table = doc.add_table(rows=len(rows), cols=2); table.style = 'Table Grid'
            for i, label in enumerate(rows):
                tabbed(table.cell(i, 0).paragraphs[0], label)
            return data(doc)
        for label in ('Signature of patient\t\t\tDate', 'Signature\tDate', 'Consultant signature\tDate signed',
                      'Date\tSignature'):
            with self.subTest(label=label):
                blob = tab_col([label, 'Plan'])
                held = next(v for k, v in self.roles(blob).items() if k != 'Plan')
                self.assertEqual((held['state'], held['basis']), ('uncertain', 'stacked_caption'))
                fields, rep = report(blob)                       # the same reading on contract v3
                self.assertIn(fields[0]['key'], rep['uncertain_places'])
                self.assertFalse(generated(fields[0], rep))
        # positive controls: tab-separated captions with no signing place, and a label read whole on evidence
        for label in ('Date\tTime', 'Troponin\tng/L', 'Consent form signed by patient', 'Patient signature obtained'):
            with self.subTest(control=label):
                blob = tab_col([label, 'Plan'])
                row = next(v for k, v in self.roles(blob).items() if k != 'Plan')
                self.assertEqual(row['state'], 'ordinary')
                fields, rep = report(blob)
                self.assertTrue(generated(fields[0], rep))
        # a governed writing line keeps its own caption: the blank after 'Date:' is the date's
        fields, rep = report(lines('Signature of patient\t\t\tDate: ______'))
        self.assertEqual(fields[0]['kind'], 'date')

    def test_a_sequence_word_beside_a_bare_initial_in_the_same_table(self):
        rows = self.roles(two_col(['Initial', 'Repeat', 'Plan']))
        self.assertEqual((rows['Initial']['state'], rows['Initial']['basis']), ('ordinary', 'sequence_sibling'))
        self.assertEqual(self.roles(two_col(['Initial', 'Plan']))['Initial']['state'], 'protected')


class Versioned(unittest.TestCase):
    def test_the_report_names_its_reading(self):
        blob = lines('Name: ______')
        self.assertEqual(w.destination_roles(blob, w.inspect(blob)['fields'])['reading'], w.ROLE_READING)
        self.assertEqual(w.legacy_destination('Name')['reading'], w.ROLE_READING)


if __name__ == '__main__':
    unittest.main()
