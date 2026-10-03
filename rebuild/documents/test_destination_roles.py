"""Destination roles: whether a place may receive an answer at all, read from the ORIGINAL's existing structure.

Synthetic blank forms only, built in memory. Red-first ids (RF..) and positive controls (PC..) are the destination-
authority matrix's (v1-product-20260930 a1 mapping-matrix redfirst.json).
"""
import io
import re
import unittest
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import word_v3 as w

SYN = 'SYNTHETIC-ANSWER'
BODY = 'word/document.xml#'


def data(doc):
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def xml(blob):
    return zipfile.ZipFile(io.BytesIO(blob)).read('word/document.xml').decode()


def box(doc, text=None, pict=False):
    table = doc.add_table(rows=1, cols=1)
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        edge = OxmlElement('w:' + side)
        edge.set(qn('w:val'), 'single')
        borders.append(edge)
    table.cell(0, 0)._tc.get_or_add_tcPr().append(borders)
    if text:
        table.cell(0, 0).paragraphs[0].add_run(text)
    if pict:
        table.cell(0, 0).paragraphs[0].add_run()._r.append(OxmlElement('w:pict'))
    return table


def grid(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            if text:
                table.cell(i, j).paragraphs[0].add_run(text)
    return table


def line_form(text):
    doc = Document()
    doc.add_heading('Clinic record', level=2)
    doc.add_paragraph(text)
    doc.add_paragraph('Plan: ______')
    return data(doc)


def signing_box_form(fill=None):
    doc = Document()
    doc.add_heading('Consultant review', level=2)
    doc.add_paragraph('Signed by the consultant:')
    box(doc, text=fill if fill not in (None, 'pict') else None, pict=fill == 'pict')
    doc.add_paragraph('Plan: ______')
    return data(doc)


def refusal(blob, fields, key, value=SYN):
    try:
        w.fill(blob, fields, {key: value})
    except w.V3Failure as exc:
        return exc.code
    return None


def forged(blob, slot, kind='text', label='Forged'):
    """What a stored mapping from before this change could hold: the candidate's own locate result for `slot`, built
    by bypassing only the role check (the placement is exactly what the candidate located)."""
    roles = w.destination_roles_of
    try:
        w.destination_roles_of = lambda package: []
        return w.locate(blob, [{'slot': BODY + slot, 'kind': kind, 'label': label}])[0]
    finally:
        w.destination_roles_of = roles


class SpanAndSegment(unittest.TestCase):
    def test_pc01_pc02_an_ordinary_blank_beside_a_signature_on_one_line_stays_answerable(self):
        for text in ('Name: ______    Signature: ______', 'Date of visit: ______    Signature: ______'):
            with self.subTest(text=text):
                blob = line_form(text)
                fields = w.inspect(blob)['fields']
                first = fields[0]
                self.assertNotEqual(first['kind'], 'signature')
                self.assertEqual(w.destination_roles(blob, fields)['violations'], {})
                out = w.fill(blob, fields, {first['key']: '2026-09-30' if first['kind'] == 'date' else SYN})
                self.assertIn('2026-09-30' if first['kind'] == 'date' else SYN, xml(out))

    def test_rf06_a_signature_owns_its_line_segment_up_to_the_next_engine_area(self):
        """'Signature: ______ (use black ink)': an insertion after the printed guidance does not touch the signature's
        span (11-17) but is inside the segment the signature owns; the candidate wrote '... (use black ink) SYN'."""
        blob = line_form('Signature: ______ (use black ink)')
        roles = [r for r in w.destination_roles(blob)['roles'] if r['class'] == 'signing']
        self.assertEqual([(r['place'], r['start'], r['end']) for r in roles], [(BODY + 'b/p[1]', 0, 34)])
        with self.assertRaises(w.V3Failure) as refused:
            w.locate(blob, [{'slot': BODY + 'b/p[1]', 'kind': 'text', 'label': 'Signature'}])
        self.assertEqual(refused.exception.code, 'destination_role')
        stored = forged(blob, 'b/p[1]', label='Signature')
        self.assertEqual(stored['locator']['target'], {'type': 'span', 'start': 33, 'end': 33, 'shape': 'after_label'})
        fields = [f for f in w.inspect(blob)['fields'] if f['kind'] != 'signature'] + [stored]
        self.assertEqual(refusal(blob, fields, stored['key']), 'protected_signing_area')
        self.assertIn(stored['key'], w.destination_roles(blob, fields)['violations'])

    def test_the_segment_ends_where_the_next_engine_area_starts(self):
        blob = line_form('Signature: ______    Name: ______')
        fields = w.inspect(blob)['fields']
        name = next(f for f in fields if f['label'] == 'Name')
        role = next(r for r in w.destination_roles(blob, fields)['roles'] if r['class'] == 'signing')
        self.assertEqual((role['start'], role['end']), (0, name['locator']['target']['start']))
        self.assertIsNone(refusal(blob, fields, name['key']))


class RemoveThenAdd(unittest.TestCase):
    """RF01-RF05 (X02): remove the engine's signing field, then add text at the same place. Candidate: WROTE."""

    CASES = (('RF01', 'Signature: ______________', 'b/p[1]'),
             ('RF03', "Participant's signature: ______", 'b/p[1]'),
             ('RF04', 'Initials: ______', 'b/p[1]'))

    def test_rf01_rf03_rf04_lines(self):
        for rid, text, slot in self.CASES:
            with self.subTest(rid=rid):
                blob = line_form(text)
                engine = w.inspect(blob)['fields']
                self.assertIn(engine[0]['kind'], ('signature', 'initials'))
                with self.assertRaises(w.V3Failure) as refused:
                    w.locate(blob, [{'slot': BODY + slot, 'kind': 'text', 'label': text.split(':')[0]}])
                self.assertEqual(refused.exception.code, 'destination_role')
                # The same place as a Signature or Initials field is allowed: it is signed by hand, never filled.
                own = w.locate(blob, [{'slot': BODY + slot, 'kind': engine[0]['kind'], 'label': 'Signed here'}])[0]
                self.assertEqual(own['kind'], engine[0]['kind'])
                stored = forged(blob, slot)
                mapping = engine[1:] + [stored]
                self.assertEqual(refusal(blob, mapping, stored['key']), 'protected_signing_area')
                plan = next(f for f in engine if f['label'] == 'Plan')
                self.assertIsNone(refusal(blob, mapping, plan['key']))   # PC16: the rest of the form still fills

    def test_rf02_signature_column_cell(self):
        doc = Document(); doc.add_heading('Clinic record', level=2); grid(doc, [['Name', 'Signature'], ['', '']])
        blob = data(doc)
        engine = w.inspect(blob)['fields']
        self.assertEqual([f['kind'] for f in engine], ['text', 'signature'])
        with self.assertRaises(w.V3Failure):
            w.locate(blob, [{'slot': BODY + 'b/tbl[0]/tr[1]/tc[1]', 'kind': 'text', 'label': 'Signature'}])
        stored = forged(blob, 'b/tbl[0]/tr[1]/tc[1]')
        self.assertEqual(refusal(blob, [engine[0], stored], stored['key']), 'protected_signing_area')
        self.assertIsNone(refusal(blob, [engine[0], stored], engine[0]['key']))

    def test_rf05_declaration_signature_line(self):
        doc = Document(); doc.add_heading('Declaration', level=2)
        doc.add_paragraph('I confirm the above is correct.'); doc.add_paragraph('Signature of examining clinician: ______')
        blob = data(doc)
        with self.assertRaises(w.V3Failure):
            w.locate(blob, [{'slot': BODY + 'b/p[2]', 'kind': 'text', 'label': 'Signature of examining clinician'}])
        stored = forged(blob, 'b/p[2]')
        self.assertEqual(refusal(blob, [stored], stored['key']), 'protected_signing_area')


class SigningQuestionBox(unittest.TestCase):
    def test_rf18_a_signing_box_is_a_role_whatever_it_prints(self):
        """The question_box_signing reading, taken BEFORE the empty-cell check: empty, a drawing, or guidance ink."""
        for fill in (None, 'pict', 'Use black ink'):
            with self.subTest(fill=fill):
                blob = signing_box_form(fill)
                roles = w.destination_roles(blob)['roles']
                self.assertIn(('signing', 'question_box_signing', BODY + 'b/tbl[0]/tr[0]/tc[0]', None, None),
                              [(r['class'], r['source'], r['place'], r['start'], r['end']) for r in roles])
                for slot in ('b/tbl[0]/tr[0]/tc[0]', 'b/tbl[0]/tr[0]/tc[0]/p[0]'):
                    try:
                        w.locate(blob, [{'slot': BODY + slot, 'kind': 'long_text', 'label': 'Consultant review'}])
                    except w.V3Failure as exc:
                        self.assertIn(exc.code, ('destination_role', 'invalid_mapping'))
                    else:
                        self.fail('a long text answer was located in a signing box')
                stored = forged(blob, 'b/tbl[0]/tr[0]/tc[0]/p[0]', kind='long_text')
                self.assertEqual(refusal(blob, [stored], stored['key']), 'protected_signing_area')

    def test_notes_and_areas_of_the_box_are_unchanged(self):
        """The reading adds a role only: the engine's own fields and notes are exactly what they were."""
        blob = signing_box_form()
        found = w.inspect(blob)
        self.assertEqual([f['label'] for f in found['fields']], ['Plan'])
        self.assertTrue(any('Signing or consent context needs explicit review' in n for n in found['notes']))

    def test_an_ordinary_question_box_is_not_a_role(self):
        doc = Document(); doc.add_heading('History', level=2); doc.add_paragraph('Describe the presenting complaint?')
        box(doc)
        blob = data(doc)
        fields = w.inspect(blob)['fields']
        self.assertEqual([f['kind'] for f in fields], ['long_text'])
        self.assertEqual(w.destination_roles(blob)['roles'], [])
        self.assertIsNone(refusal(blob, fields, fields[0]['key']))


class ProtectedAndPolicy(unittest.TestCase):
    def test_pc11_the_protected_signing_block_stays_refused(self):
        doc = Document(); doc.add_paragraph('Discharge letter'); doc.add_paragraph('Plan: ______')
        p = doc.add_paragraph()
        for n in range(2):
            if n:
                p.add_run().add_tab()
            r = p.add_run(); r.underline = True; r._r.append(OxmlElement('w:tab'))
        cap = doc.add_paragraph(); cap.add_run('Signature'); cap.add_run().add_tab(); cap.add_run('Date')
        blob = data(doc)
        roles = {(r['class'], r['source'], r['place']) for r in w.destination_roles(blob)['roles']}
        self.assertIn(('protected', 'protected_area', BODY + 'b/p[2]'), roles)
        self.assertIn(('protected', 'protected_region', BODY + 'b/p[3]'), roles)
        with self.assertRaises(w.V3Failure):
            w.locate(blob, [{'slot': BODY + 'b/p[2]', 'kind': 'text', 'label': 'Signature'}])

    def test_a_manual_field_in_a_patient_section_carries_the_printed_author_role(self):
        doc = Document(); doc.add_paragraph('Pre-operative questionnaire'); doc.add_paragraph('Completed by the patient:')
        doc.add_paragraph('Allergies'); doc.add_paragraph('I smoke: ☐ Yes ☐ No')
        blob = data(doc)
        added = w.locate(blob, [{'slot': BODY + 'b/p[2]', 'kind': 'text', 'label': 'Allergies'}])[0]
        self.assertEqual(added['answer_policy']['source_author'], 'patient')
        self.assertEqual(w.destination_roles(blob, [added])['violations'], {})
        # The same field WITHOUT the printed requirement (a stored mapping from before) is refused at export.
        stripped = {k: v for k, v in added.items() if k != 'answer_policy'}
        self.assertEqual(w.destination_roles(blob, [stripped])['violations'][stripped['key']]['class'], 'policy')
        self.assertEqual(refusal(blob, [stripped], stripped['key'], 'The patient reports none.'), 'mapping_mismatch')
        # PC12: the engine's own patient-policy selection still fills.
        engine = w.inspect(blob)['fields']
        smoke = next(f for f in engine if f['label'] == 'I smoke')
        self.assertIsNone(refusal(blob, engine, smoke['key'], 'Yes'))

    def test_the_attestation_cell_is_a_policy_role(self):
        question = ('Please confirm that the patient has had sight of this report and/or '
                    'had an opportunity to add their comments, if appropriate or necessary.')
        doc = Document(); t = doc.add_table(rows=1, cols=2); t.cell(0, 0).text = question
        blob = data(doc)
        engine = w.inspect(blob)['fields']
        self.assertEqual(engine[0]['answer_policy']['attestation'], 'document_review')
        roles = [r for r in w.destination_roles(blob, engine)['roles'] if r['source'] == 'attestation']
        self.assertEqual([r['place'] for r in roles], [BODY + 'b/tbl[0]/tr[0]/tc[1]'])
        self.assertEqual(w.destination_roles(blob, engine)['violations'], {})
        self.assertTrue(engine[0]['write_protected'])
        relabelled = {**{k: v for k, v in engine[0].items()
                        if k not in ('answer_policy', 'write_protected', 'review_reason')}, 'label': 'Comments'}
        # The retained original now also binds an explicit human action. A
        # protection marker means the unchanged mapping cannot write at all;
        # stripping it still cannot convert that place to ordinary comments.
        self.assertEqual(w.destination_roles(blob, [relabelled])['violations'][relabelled['key']]['class'], 'protected')
        self.assertEqual(refusal(blob, [relabelled], relabelled['key'], 'Yes'), 'protected_signing_area')

    def test_a_locked_content_control_is_protected(self):
        doc = Document(); p = doc.add_paragraph('Hospital number: ')
        sdt = OxmlElement('w:sdt'); pr = OxmlElement('w:sdtPr'); lock = OxmlElement('w:lock')
        lock.set(qn('w:val'), 'sdtContentLocked'); pr.append(lock); pr.append(OxmlElement('w:text')); sdt.append(pr)
        content = OxmlElement('w:sdtContent'); r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = ' '
        r.append(t); content.append(r); sdt.append(content); p._p.append(sdt)
        doc.add_paragraph('Plan: ______')
        blob = data(doc)
        engine = w.inspect(blob)['fields']
        locked = next(f for f in engine if f['locator']['target']['type'] == 'sdt_text')
        self.assertTrue(locked.get('write_protected'))
        self.assertIn(('protected', 'locked_control'),
                      {(r['class'], r['source']) for r in w.destination_roles(blob, engine)['roles']})
        self.assertEqual(w.destination_roles(blob, engine)['violations'], {})


class LegacyRoleKind(unittest.TestCase):
    """The legacy default path's reader is the engine's own kind rule with its two measured over-reads removed."""

    CHECKLIST_SIGNED = ['Consent form signed', 'Consent signed', 'Drug chart signed', 'Discharge summary signed',
                        'DNACPR form signed', 'Prescription signed', 'Care plan signed', 'Theatre checklist signed',
                        'Time signed', 'Signed consent in notes', 'Sign in', 'Sign out', 'Sign-off',
                        'Date consent signed']
    INITIAL_READINGS = ['Initial assessment', 'Initial dose', 'Initial reading', 'Initial GCS', 'Initial observations',
                        'Initial impression', 'Initial plan', 'Initial presentation', 'Initial management',
                        'Initial lactate', 'Troponin — Initial']
    SIGNING = ['Signature', 'Patient signature', 'Witness signature', 'Consultant signature', 'Signed by',
               'Signature of examining clinician', "Participant's signature", 'First doctor signature',
               'Signed off by consultant', 'Sign here', 'Initials', 'Clinician initials', 'Initial the box',
               'Initial:']

    def test_checklist_states_and_adjectives_are_not_roles(self):
        for label in self.CHECKLIST_SIGNED + self.INITIAL_READINGS:
            with self.subTest(label=label):
                self.assertEqual(w.legacy_role_kind(label), '')

    def test_signatures_and_initials_are_roles(self):
        for label in self.SIGNING:
            with self.subTest(label=label):
                self.assertIn(w.legacy_role_kind(label), ('signature', 'initials'))

    def test_every_legacy_match_is_a_kind_rule_match(self):
        """Restricting, never widening: each reading is one the engine's kind rule already gives."""
        for label in self.CHECKLIST_SIGNED + self.INITIAL_READINGS + self.SIGNING + ['Signature date', 'Babinski sign']:
            with self.subTest(label=label):
                kind = w.legacy_role_kind(label)
                if kind:
                    self.assertEqual(w.Detector.kind_for(None, label), kind)



# ─── refutation round 1 (e555540d): each test is red on e555540d and names the finding it closes ─────────────────
def inline_control(p, kind='text', lock=None, alias=None, placeholder='Click or tap here to enter text.'):
    sdt = OxmlElement('w:sdt'); pr = OxmlElement('w:sdtPr')
    if alias:
        a = OxmlElement('w:alias'); a.set(qn('w:val'), alias); pr.append(a)
    if lock:
        lk = OxmlElement('w:lock'); lk.set(qn('w:val'), lock); pr.append(lk)
    pr.append(OxmlElement('w:showingPlcHdr')); pr.append(OxmlElement('w:' + kind)); sdt.append(pr)
    content = OxmlElement('w:sdtContent'); r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = placeholder
    r.append(t); content.append(r); sdt.append(content); p._p.append(sdt)
    return sdt


def formtext(p):
    def run(child):
        r = OxmlElement('w:r'); r.append(child); p._p.append(r)
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    data_ = OxmlElement('w:ffData'); name = OxmlElement('w:name'); name.set(qn('w:val'), 'Text1'); data_.append(name)
    data_.append(OxmlElement('w:enabled')); data_.append(OxmlElement('w:textInput')); begin.append(data_); run(begin)
    instr = OxmlElement('w:instrText'); instr.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    instr.text = ' FORMTEXT '; run(instr)
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate'); run(sep)
    t = OxmlElement('w:t'); t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); t.text = '     '; run(t)
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end'); run(end)


def part_xml(blob, part):
    return zipfile.ZipFile(io.BytesIO(blob)).read(part).decode()


def located(blob, slot, kind='text', label='Answer'):
    try:
        return w.locate(blob, [{'slot': slot, 'kind': kind, 'label': label}])[0], None
    except w.V3Failure as exc:
        return None, exc.code


class SigningPlaceReader(unittest.TestCase):
    """ONE reader of whether a printed label names a signing place (legacy role; v3 area role). Fails closed."""

    def test_l1f1_l1f5_every_refuted_signing_label_is_a_place(self):
        """L1:F1 (HIGH): e555540d needed 'by' after 'signed', so 'Signed (Consultant)' and eight more were not roles and
        generation filled them. L1:F5: 'Initial here', 'Signed for and on behalf of', 'Countersigned by consultant', a
        signature-and-date place and a non-breaking space leaked on the legacy path."""
        for label, kind in (('Signed (Consultant)', 'signature'), ('Signed (print name)', 'signature'),
                            ('Electronically signed (print name)', 'signature'), ('Signed (GP)', 'signature'),
                            ('Signed (nurse)', 'signature'), ('Signed (name)', 'signature'),
                            ('Digitally signed (name)', 'signature'), ('E-signed (name)', 'signature'),
                            ('Signed (registrar)', 'signature'), ('Signed (the consultant)', 'signature'),
                            ('Signed (patient)', 'signature'), ('Signed for and on behalf of', 'signature'),
                            ('Countersigned by consultant', 'signature'), ('Signature and date', 'signature'),
                            ('Patient signature / date', 'signature'), ('Sign & date', 'signature'),
                            ('Please sign\u00a0here', 'signature'), ('Signed:', 'signature'),
                            ('Initial here', 'initials'), ('Initial each page', 'initials'),
                            ('Initial and date', 'initials')):
            with self.subTest(label=label):
                self.assertEqual(w.legacy_role_kind(label), kind)
                self.assertEqual(w.signing_place(label), kind)

    def test_r2_06_a_number_after_initial_is_the_adjective(self):
        """R2-06: 'Initial 12-lead ECG' and 'Initial 4AT score' were held as initials (a digit is never an imperative's
        object)."""
        for label in ('Initial 12-lead ECG', 'Initial 4AT score', 'Initial GCS', 'Initial assessment'):
            with self.subTest(label=label):
                self.assertEqual(w.legacy_role_kind(label), '')

    def test_declared_residuals_stay_held(self):
        """Fail closed: the same wording prints a real signing place ('Note signed by attending physician', 'Patient
        signature', 'Nurse (initial)', 'Initial:'). Held means never generated; a typed value still exports."""
        for label in ('Consent form signed by patient', 'Drug chart signed by prescriber', 'Patient signature obtained',
                      'Witness signature obtained', 'Troponin (initial)', 'Lactate (initial)', 'Initial / repeat',
                      'Initial'):
            with self.subTest(label=label):
                self.assertNotEqual(w.signing_place(label), '')

    def test_date_questions_about_a_signature_stay_date_questions(self):
        for label in ('Signature date', 'Date of signature', 'Date signed', 'Date consent signed'):
            with self.subTest(label=label):
                self.assertEqual(w.signing_place(label), '')
                self.assertEqual(w.Detector.kind_for(None, label), 'date')


class SignatureAndDate(unittest.TestCase):
    def test_l1f2_a_signature_and_date_place_is_a_signing_place(self):
        """L1:F2 (HIGH): the date word removed the signature reading, so 'Signature and date: ____' became a DATE field
        and the export wrote 'Signature and date: Dr Example 30/09/2026'."""
        for text in ('Signature and date: ______________', 'Patient signature / date: ______________'):
            with self.subTest(text=text):
                blob = line_form(text)
                engine = w.inspect(blob)['fields']
                self.assertEqual(engine[0]['kind'], 'signature')
                stored = {**engine[0], 'kind': 'date', 'key': 'stored_date'}   # a mapping stored before this repair
                self.assertEqual(refusal(blob, [stored] + engine[1:], 'stored_date', '30/09/2026'),
                                 'protected_signing_area')
                self.assertIsNone(refusal(blob, engine, engine[1]['key'], 'Rest.'))

    def test_pc04_a_signature_date_stays_a_date(self):
        blob = line_form('Signature date: ______')
        engine = w.inspect(blob)['fields']
        self.assertEqual(engine[0]['kind'], 'date')
        self.assertIn('2026-09-30', xml(w.fill(blob, engine, {engine[0]['key']: '2026-09-30'})))


class InlineSigningControls(unittest.TestCase):
    """L1:F3: 'Consultant signature: ' + an inline control + ' (use black ink)'. The control's area owned only itself, so
    a manual or stored field after the guidance wrote 'Consultant signature: ... (use black ink) Dr Example'."""

    def form(self, kind):
        doc = Document(); doc.add_heading('Clinic record', level=2)
        p = doc.add_paragraph('Consultant signature: ')
        if kind == 'sdt':
            inline_control(p)
        elif kind == 'picture':
            inline_control(p, kind='picture', placeholder='')
        else:
            formtext(p)
        p.add_run(' (use black ink)')
        doc.add_paragraph('Plan: ______')
        return data(doc)

    def test_l1f3_an_inline_signing_control_owns_its_line(self):
        for kind in ('sdt', 'ff', 'picture'):
            with self.subTest(kind=kind):
                blob = self.form(kind)
                field, code = located(blob, BODY + 'b/p[1]', label='Consultant signature')
                self.assertIsNone(field)
                self.assertEqual(code, 'destination_role')
                stored = forged(blob, 'b/p[1]', label='Consultant signature')
                engine = [f for f in w.inspect(blob)['fields'] if f['kind'] not in ('signature', 'initials')]
                self.assertEqual(refusal(blob, engine + [stored], stored['key']), 'protected_signing_area')
                plan = next(f for f in engine if f['label'] == 'Plan')
                self.assertIsNone(refusal(blob, engine + [stored], plan['key'], 'Rest.'))

    def test_a_line_beside_an_inline_control_keeps_its_own_answer(self):
        """The control's segment ends at the next engine area: 'Name' on the same line stays answerable."""
        doc = Document(); doc.add_heading('Clinic record', level=2)
        p = doc.add_paragraph('Consultant signature: '); inline_control(p); p.add_run('    Name: ______')
        blob = data(doc)
        engine = w.inspect(blob)['fields']
        name = next(f for f in engine if f['label'] == 'Name')
        self.assertEqual(w.destination_roles(blob, engine)['violations'], {})
        self.assertIn(SYN, xml(w.fill(blob, engine, {name['key']: SYN})))


class SigningBoxQuestionShapes(unittest.TestCase):
    """L1:F4: the box reading stopped at a question paragraph holding a line break or a drawing, and never read a
    header or footer box, so a long answer was written into the consultant's signing box."""

    def body_form(self, variant):
        doc = Document(); doc.add_heading('Consultant review', level=2)
        p = doc.add_paragraph('Signed by the consultant:')
        if variant == 'br':
            p.add_run().add_break(); p.add_run('(use black ink)')
        else:
            p.add_run()._r.append(OxmlElement('w:drawing'))
        box(doc); doc.add_paragraph('Plan: ______')
        return data(doc)

    def test_l1f4_a_question_with_a_break_or_drawing_still_makes_a_signing_box(self):
        for variant in ('br', 'drawing'):
            with self.subTest(variant=variant):
                blob = self.body_form(variant)
                self.assertIn(('signing', 'question_box_signing', BODY + 'b/tbl[0]/tr[0]/tc[0]'),
                              {(r['class'], r['source'], r['place']) for r in w.destination_roles(blob)['roles']})
                field, code = located(blob, BODY + 'b/tbl[0]/tr[0]/tc[0]', 'long_text', 'Consultant review')
                self.assertEqual(code, 'destination_role')
                stored = forged(blob, 'b/tbl[0]/tr[0]/tc[0]', 'long_text')
                self.assertEqual(refusal(blob, [stored], stored['key']), 'protected_signing_area')

    def test_l1f4_a_footer_signing_box(self):
        doc = Document(); doc.add_heading('Consultant review', level=2); doc.add_paragraph('Plan: ______')
        footer = doc.sections[0].footer
        footer.paragraphs[0].add_run('Signed by the consultant:')
        from docx.shared import Inches
        table = footer.add_table(rows=1, cols=1, width=Inches(5))
        borders = OxmlElement('w:tcBorders')
        for side in ('top', 'left', 'bottom', 'right'):
            edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); borders.append(edge)
        table.cell(0, 0)._tc.get_or_add_tcPr().append(borders)
        blob = data(doc)
        slot = 'word/footer1.xml#hf/tbl[0]/tr[0]/tc[0]'
        self.assertIn(('signing', 'question_box_signing', slot),
                      {(r['class'], r['source'], r['place']) for r in w.destination_roles(blob)['roles']})
        field, code = located(blob, slot, 'long_text', 'Consultant review')
        self.assertEqual(code, 'destination_role')


class SigningCaptionAndLine(unittest.TestCase):
    """L1:F6: a caption over its writing line is one signing place."""

    def test_l1f6_the_writing_line_under_a_signature_caption(self):
        """(a) 'Consultant signature:' then '____': the engine maps the underline as an ordinary 'Answer' field, and fill
        wrote 'Dr Example' on it. Now that field sits on a signing place (confirmation lists it for removal)."""
        doc = Document(); doc.add_heading('Clinic record', level=2)
        doc.add_paragraph('Consultant signature:'); doc.add_paragraph('____________________'); doc.add_paragraph('Plan: ______')
        blob = data(doc)
        engine = w.inspect(blob)['fields']
        answer = next(f for f in engine if f['locator']['slot'] == 'b/p[2]')
        self.assertEqual(answer['kind'], 'text')
        self.assertIn(answer['key'], w.destination_roles(blob, engine)['violations'])
        self.assertEqual(refusal(blob, engine, answer['key'], 'Dr Example'), 'protected_signing_area')
        plan = next(f for f in engine if f['label'] == 'Plan')
        self.assertIsNone(refusal(blob, [f for f in engine if f is not answer], plan['key'], 'Rest.'))

    def test_l1f6_the_caption_of_a_vml_signature_line(self):
        """(b) the engine's signature area is the VML line paragraph; its caption paragraph took a manual text answer."""
        from docx.oxml import parse_xml
        doc = Document(); doc.add_heading('Clinic record', level=2); doc.add_paragraph('Consultant signature:')
        doc.add_paragraph().add_run()._r.append(parse_xml(
            '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
            '<v:shape id="SigLine1" style="width:192pt;height:96pt"><o:signatureline v:ext="edit" '
            'id="{00000000-0000-0000-0000-000000000001}" issignatureline="t"/></v:shape></w:pict>'))
        doc.add_paragraph('Plan: ______')
        blob = data(doc)
        self.assertEqual([f['kind'] for f in w.inspect(blob)['fields']][:1], ['signature'])
        field, code = located(blob, BODY + 'b/p[1]', label='Consultant signature')
        self.assertEqual(code, 'destination_role')
        stored = forged(blob, 'b/p[1]', label='Consultant signature')
        self.assertEqual(refusal(blob, [stored], stored['key']), 'protected_signing_area')


class SigningColumns(unittest.TestCase):
    def log(self, headers):
        doc = Document(); doc.add_heading('Medication log', level=2)
        t = doc.add_table(rows=3, cols=len(headers)); t.style = 'Table Grid'
        for c, h in enumerate(headers):
            t.cell(0, c).text = h
        doc.add_paragraph('Plan: ______')
        return data(doc)

    def test_l1f7_declared_residual_a_typed_initials_column_fills_as_on_43fcf214(self):
        """L1:F7 (pre-existing on 43fcf214) is a declared residual again: a log table's column carries no role, because
        reading one from its printed header refused ordinary clinical columns (refutation of 258650f2, L1:R1-03). Only
        the clinician's typed rows reach it (table_rows is never generated)."""
        blob = self.log(('Drug', 'Dose', 'Nurse initials'))
        engine = w.inspect(blob)['fields']
        rows = next(f for f in engine if f['kind'] == 'table_rows')
        self.assertEqual(w.destination_roles(blob, engine)['violations'], {})
        out = xml(w.fill(blob, engine, {rows['key']: [['Paracetamol', '1 g', 'JE'], ['Ibuprofen', '400 mg', '']]}))
        self.assertIn('Paracetamol', out)
        self.assertIn('JE', out)

    def test_pc10_an_admin_grid_without_a_signing_column_is_unchanged(self):
        blob = self.log(('Drug', 'Dose', 'Initial dose'))
        engine = w.inspect(blob)['fields']
        rows = next(f for f in engine if f['kind'] == 'table_rows')
        self.assertIn('500 mg', xml(w.fill(blob, engine, {rows['key']: [['Paracetamol', '1 g', '500 mg']]})))


class FramesAndNeighbours(unittest.TestCase):
    def frame(self, h1, inside, prompt=None):
        doc = Document()
        if h1:
            doc.add_heading(h1, level=1)
        if prompt:
            doc.add_paragraph(prompt)
        else:
            doc.add_heading('Clinical details', level=2)
        table = box(doc)
        cell = table.cell(0, 0)
        for n, text in enumerate(inside):
            (cell.paragraphs[0] if n == 0 else cell.add_paragraph()).add_run(text)
        doc.add_heading('Management', level=2); doc.add_paragraph('Plan: ______')
        return data(doc)

    def test_r2_01_a_frame_holding_ordinary_questions_is_not_a_signing_box(self):
        """R2-01 (HIGH): the box reading made a whole frame a signing place from its heading ('Initial referral',
        'Consent') or its prompt ('to be signed by the referring clinician'), and every field inside was refused."""
        inside = ('Presenting complaint: ______', 'Current medication: ______', 'Allergies: ______')
        for h1, prompt in (('Initial referral', None), ('Consent', None),
                           ('Referral', 'Referral (to be signed by the referring clinician)')):
            with self.subTest(h1=h1, prompt=prompt):
                blob = self.frame(h1, inside, prompt)
                engine = w.inspect(blob)['fields']
                self.assertEqual(w.destination_roles(blob, engine)['violations'], {})
                answers = {f['key']: 'Chest pain' for f in engine if f['label'] == 'Presenting complaint'}
                self.assertIn('Presenting complaint: Chest pain', re.sub(r'<[^>]+>', '', xml(w.fill(blob, engine, answers))))

    def test_the_rest_of_a_signing_frame_stays_the_frames(self):
        """Excepting the engine's areas inside a signing box never opens the box's other paragraphs: 'Signed by the
        consultant:' over a box holding 'Name: ____' and a guidance line - Name fills, the guidance line takes nothing."""
        doc = Document(); doc.add_heading('Consultant review', level=2); doc.add_paragraph('Signed by the consultant:')
        cell = box(doc).cell(0, 0)
        cell.paragraphs[0].add_run('Name: ______'); cell.add_paragraph().add_run('Use black ink')
        doc.add_paragraph('Plan: ______')
        blob = data(doc)
        engine = w.inspect(blob)['fields']
        name = next(f for f in engine if f['label'] == 'Name')
        self.assertIsNone(refusal(blob, engine, name['key']))
        self.assertEqual(located(blob, BODY + 'b/tbl[0]/tr[0]/tc[0]/p[1]', 'long_text', 'Consultant review')[1],
                         'destination_role')
        stored = forged(blob, 'b/tbl[0]/tr[0]/tc[0]/p[1]', 'long_text')
        self.assertEqual(refusal(blob, engine + [stored], stored['key']), 'protected_signing_area')

    def test_r2_01_a_signature_inside_a_frame_keeps_its_own_role(self):
        blob = self.frame('Consent', ('Patient name: ______', 'Signature: ______'))
        engine = w.inspect(blob)['fields']
        name = next(f for f in engine if f['label'] == 'Patient name')
        self.assertIsNone(refusal(blob, engine, name['key']))
        field, code = located(blob, BODY + 'b/tbl[0]/tr[0]/tc[0]/p[1]', label='Signature')
        self.assertEqual(code, 'destination_role')

    def test_r2_02_an_after_label_neighbour_keeps_its_answer(self):
        """R2-02 (HIGH): 'Printed name:    Signature: ______' - the signature's segment started at Printed name's
        insertion point instead of the end of its printed extent, so Printed name was refused."""
        for text, label in (('Printed name:    Signature: ______', 'Printed name'), ('Name:    Signature: ______', 'Name'),
                            ('Date:    Signature: ______', 'Date'), ('GMC number:    Signature: ______', 'GMC number'),
                            ('Printed name:\tSignature: ______', 'Printed name'),
                            ('Consultant:    Signature: ______    Date: ______', 'Consultant')):
            with self.subTest(text=text):
                blob = line_form(text)
                engine = w.inspect(blob)['fields']
                field = next(f for f in engine if f['label'] == label)
                self.assertEqual(w.destination_roles(blob, engine)['violations'], {})
                value = '2026-09-30' if field['kind'] == 'date' else SYN
                self.assertIn(value, xml(w.fill(blob, engine, {field['key']: value})))
        # The guidance after a signature is still the signature's (RF06).
        blob = line_form('Printed name:    Signature: ______ (use black ink)')
        self.assertEqual(located(blob, BODY + 'b/p[1]', label='Signature')[1], 'destination_role')


class ReadingsReleased(unittest.TestCase):
    def question_box(self, heading, question):
        doc = Document(); doc.add_heading(heading, level=2); doc.add_paragraph(question); box(doc)
        doc.add_paragraph('Plan: ______')
        return data(doc)

    def test_r2_03_an_adjective_initial_no_longer_makes_a_signing_box(self):
        """R2-03: the engine declines these boxes (its own reading is unchanged), but the role also refused the
        clinician's manual add, the only recovery path."""
        for heading, question in (('History', 'Describe the initial presentation?'),
                                  ('Initial assessment', 'What was the presenting complaint?'),
                                  ('Treatment', 'What initial treatment was given?')):
            with self.subTest(question=question):
                blob = self.question_box(heading, question)
                self.assertTrue(any('Signing or consent context needs explicit review' in n
                                    for n in w.inspect(blob)['notes']))
                self.assertEqual(w.destination_roles(blob)['roles'], [])
                field, code = located(blob, BODY + 'b/tbl[0]/tr[0]/tc[0]', 'long_text', question)
                self.assertIsNone(code)

    def test_a_signing_verb_still_governs_an_initial_report(self):
        for question in ('Please countersign the initial report?', 'Please sign the initial report here?',
                         'Please initial here?', 'Signed by the consultant:'):
            with self.subTest(question=question):
                blob = self.question_box('Review', question)
                self.assertEqual(located(blob, BODY + 'b/tbl[0]/tr[0]/tc[0]', 'long_text', 'Review')[1], 'destination_role')

    def test_declared_residual_a_consent_question_box_stays_a_signing_box(self):
        blob = self.question_box('Capacity', 'Does the patient have capacity to consent to treatment?')
        self.assertEqual(located(blob, BODY + 'b/tbl[0]/tr[0]/tc[0]', 'long_text', 'Capacity')[1], 'destination_role')

    def test_r2_04_a_clinical_reading_the_kind_rule_types_initials_is_not_a_role(self):
        """Retain the seven historical controls with current physical-source behavior.

        Initial clinical readings now import as ordinary text automatically.
        The unresolved consent action remains held; it cannot be manually
        re-added as a clinical answer. Historical expectations are retained in
        the continuation's pre-resume replay, rather than weakening the source.
        """
        for label in ('Initial GCS', 'Initial assessment', 'Initial dose', 'Initial ECG', 'Initial lactate',
                      'Consent form signed', 'Time signed'):
            with self.subTest(label=label):
                blob = line_form(label + ': ______')
                engine = w.inspect(blob)['fields']
                self.assertEqual([r for r in w.destination_roles(blob)['roles'] if r['class'] == 'signing'], [])
                if label == 'Consent form signed':
                    self.assertTrue(engine[0].get('write_protected'))
                    self.assertEqual(located(blob, BODY + 'b/p[1]', label=label)[1], 'destination_role')
                    self.assertEqual(refusal(blob, engine, engine[0]['key'], 'Yes'), 'protected_signing_area')
                    continue
                if label.startswith('Initial '):
                    self.assertEqual(engine[0]['kind'], 'text')
                    self.assertFalse(engine[0].get('write_protected'))
                    self.assertIn('SYNVAL', xml(w.fill(blob, engine, {engine[0]['key']: 'SYNVAL'})))
                added, code = located(blob, BODY + 'b/p[1]', label=label)
                self.assertIsNone(code)
                self.assertIn('SYNVAL', xml(w.fill(blob, engine[1:] + [added], {added['key']: 'SYNVAL'})))

    def test_r2_05_an_unlocked_control_inside_a_locked_frame_is_answerable(self):
        """R2-05: a contentLocked rich-text frame around 'Hospital number: ' + an unlocked plain-text control: the
        frame's lock refused the control the form's author left unlocked for the answer."""
        for lock in ('contentLocked', 'sdtContentLocked'):
            with self.subTest(lock=lock):
                doc = Document(); doc.add_heading('Clinic record', level=2)
                frame = OxmlElement('w:sdt'); pr = OxmlElement('w:sdtPr'); lk = OxmlElement('w:lock')
                lk.set(qn('w:val'), lock); pr.append(lk); pr.append(OxmlElement('w:richText')); frame.append(pr)
                content = OxmlElement('w:sdtContent'); frame.append(content)
                p = doc.add_paragraph('Hospital number: '); inline_control(p, alias='Hospital number')
                content.append(p._p)
                doc.element.body.find(qn('w:sectPr')).addprevious(frame)
                doc.add_paragraph('Plan: ______')
                blob = data(doc)
                engine = w.inspect(blob)['fields']
                number = next(f for f in engine if f['label'] == 'Hospital number')
                self.assertEqual(w.destination_roles(blob, engine)['violations'], {})
                self.assertIn('SYN-123', xml(w.fill(blob, engine, {number['key']: 'SYN-123'})))
                # The frame's own printed text is still protected.
                stored = forged(blob, 'b/sdt[0]/p[0]', label='Hospital number')
                self.assertEqual(refusal(blob, [stored], stored['key']), 'protected_signing_area')


if __name__ == '__main__':
    unittest.main()


# ─── salvage of a32243eb (b2): usefulness on the legacy default path, the dash-qualified reading ───────────────────
class DashQualifiedInitial(unittest.TestCase):
    """PC07 / PC18: 43fcf214 generated 'Troponin — Initial' and a32243eb held it as an initials place. The closing
    'initial' after a dash separator is released (a narrowing of INITIALS_PLACE, no new signing word); every other
    closing or bare 'initial' stays a place."""

    def test_the_dash_qualified_reading_is_released(self):
        for label in ('Troponin — Initial', 'Troponin – Initial', 'Troponin - Initial', 'Lactate — initial:',
                      'Troponin—Initial'):
            with self.subTest(label=label):
                self.assertEqual(w.signing_place(label), '')
                self.assertEqual(w.legacy_role_kind(label), '')

    def test_every_other_initial_place_stays_held(self):
        for label in ('Initial', 'Initial:', 'Initials', 'Nurse initial', 'Nurse (initial)', 'Troponin (initial)',
                      'Troponin — Initials', 'Initial — Troponin', 'Initial here', 'Troponin-Initial',
                      'Lactate — initial here', 'Nurse initials — Initial'):
            with self.subTest(label=label):
                self.assertEqual(w.signing_place(label), 'initials')

    def test_the_declared_cost_a_person_before_the_dash_is_released_as_on_43fcf214(self):
        # Declared, not hidden: the reader cannot tell 'Nurse' from 'Troponin'; 43fcf214 generated both.
        self.assertEqual(w.signing_place('Nurse — Initial'), '')

    def test_the_v3_kind_rule_is_unchanged(self):
        # On contract v3 the area stays an initials field (never generated); only its role is released.
        self.assertEqual(w.Detector.kind_for(None, 'Troponin — Initial'), 'initials')


# ─── b2 repair round (refutation of 258650f2): regressions vs 43fcf214 removed by dropping or narrowing their cause ───
def log_table(headers, group=None, rows=3):
    doc = Document(); doc.add_paragraph('Synthetic results log')
    extra = 1 if group else 0
    t = doc.add_table(rows=rows + 1 + extra, cols=len(headers)); t.style = 'Table Grid'
    if group:
        text, start, end = group
        for c in range(len(headers)):
            if c < start or c > end:
                t.cell(0, c).text = headers[c]
        merged = t.cell(0, start).merge(t.cell(0, end)); merged.text = text
        for c in range(len(headers)):
            if start <= c <= end:
                t.cell(1, c).text = headers[c]
            else:
                t.cell(0, c).merge(t.cell(1, c))
    else:
        for c, h in enumerate(headers):
            t.cell(0, c).text = h
    return data(doc)


class RepairDatePairs(unittest.TestCase):
    """L1:R1-04 / L1:R1-05: the signature-and-date pair counts only with 'signed' opening the label (a 'signed' after
    other words is the checklist state signing_place already releases), and the initials-and-date pair only when it
    closes the label. Both are narrowings of a32243eb's pair arms: fewer labels become signing places."""

    RESTORED_KINDS = {'Date signed / date received': 'date', 'Date of diagnosis / initial stage': 'date',
                      'Care plan signed and date of next review': 'long_text', 'Patient initials and date of birth': 'date',
                      'Date of referral / initial contact': 'date', 'Consent form signed and dated': 'signature'}

    def test_the_kind_rule_is_43fcf214s_for_a_pair_that_is_not_one_place(self):
        for label, kind in self.RESTORED_KINDS.items():
            with self.subTest(label=label):
                self.assertEqual(w.Detector.kind_for(None, label), kind)
                self.assertEqual(w.signing_place(label), '')
                self.assertEqual(w.legacy_role_kind(label), '')

    def test_a_signature_or_initials_and_date_place_stays_a_signing_place(self):
        for label, kind in (('Signature and date', 'signature'), ('Signature and date:', 'signature'),
                            ('Patient signature / date', 'signature'), ('Sign & date', 'signature'),
                            ('Signed and date', 'signature'), ('Signed / date', 'signature'),
                            ('Co-signed and date', 'signature'), ('Date and signature', 'signature'),
                            ('Participant signature and date', 'signature'), ('Signature + date', 'signature'),
                            ('Initials and date', 'initials'), ('Initial and date', 'initials'), ('Date / initials', 'initials'),
                            ('Nurse initials and date', 'initials'), ('Initials and date (dd/mm/yyyy)', 'initials')):
            with self.subTest(label=label):
                self.assertEqual(w.Detector.kind_for(None, label), kind)
                self.assertEqual(w.signing_place(label), kind)
                self.assertEqual(w.legacy_role_kind(label), kind)


class RepairLegacyRoleIsTheV3Rule(unittest.TestCase):
    """L1:R1-04 and PC05: the legacy role holds a printed label only where the v3 path would (the engine's kind rule
    types it signature or initials AND signing_place names a signing place). 'Tinel sign both wrists' and 'Sign and
    symptom review' are text to the kind rule - generated on v3 and on 43fcf214 - and were held only on legacy."""

    def test_a_label_the_kind_rule_reads_as_text_is_no_legacy_role(self):
        for label in ('Tinel sign both wrists', 'Sign and symptom review'):
            with self.subTest(label=label):
                self.assertEqual(w.Detector.kind_for(None, label), 'text')
                self.assertEqual(w.legacy_role_kind(label), '')

    def test_rf16_rf17_signing_rows_stay_roles(self):
        for label in ('Patient signature', 'Witness signature', "Participant's signature", 'Initials',
                      'Consultant signature', 'Signature of examining clinician', 'Signature', 'Sign here',
                      'Initial here', 'Signed (patient)', 'Signature and date', 'Nurse initials'):
            with self.subTest(label=label):
                self.assertIn(w.legacy_role_kind(label), ('signature', 'initials'))

    def test_the_legacy_role_is_never_wider_than_the_kind_rule(self):
        corpus = (LegacyRoleKind.CHECKLIST_SIGNED + LegacyRoleKind.INITIAL_READINGS + LegacyRoleKind.SIGNING
                  + ['Sign below', 'Sign and print name', 'Sign & stamp', 'Signatures', 'Sign each page',
                     'Tinel sign both wrists', 'Sign and symptom review', 'Gene expression signature', 'Sig­nature'])
        for label in corpus:
            with self.subTest(label=label):
                kind = w.legacy_role_kind(label)
                if kind:
                    self.assertEqual(w.Detector.kind_for(None, w._plain(label)), kind)


class RepairLogColumns(unittest.TestCase):
    """L1:R1-03: a log table's column is no destination role (as on 43fcf214); its rows are typed by the clinician only
    (table_rows is never generated)."""

    CASES = ((['Date', 'Visit type (initial/follow-up)', 'Clinician notes'], None, ['01/09/2026', 'Follow-up', 'Stable']),
             (['Drug', 'Dose (initial)', 'Dose (maintenance)'], None, ['Levothyroxine', '25 micrograms', '50 micrograms']),
             (['Date', 'Initial', 'Repeat'], ('Troponin (ng/L)', 1, 2), ['01/09/2026', '12', '14']),
             (['Date', 'Troponin initial', 'Troponin 3 h'], None, ['01/09/2026', '12', '14']),
             (['Problem', 'Initial and ongoing management'], None, ['Sepsis', 'IV antibiotics']))

    def test_the_typed_rows_fill(self):
        for headers, group, row in self.CASES:
            with self.subTest(headers=headers):
                blob = log_table(headers, group=group)
                engine = w.inspect(blob)['fields']
                rows = next(f for f in engine if f['kind'] == 'table_rows')
                self.assertEqual([r for r in w.destination_roles(blob, engine)['roles'] if r['source'] == 'signing_column'], [])
                out = xml(w.fill(blob, engine, {rows['key']: [row]}))
                for value in row:
                    self.assertIn(value, out)
