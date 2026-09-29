"""Small synthetic controls for the D19 development finding; no corpus qualification."""
import io
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree

from . import word_v3 as word


def signing_form(*, mismatch=False, table=True):
    doc = Document()
    doc.add_paragraph('Participant ID: __________')
    if table:
        doc.add_paragraph('Please initial box')
        t = doc.add_table(rows=3, cols=2)
        for number in range(2):
            t.cell(number, 0).text = 'Consent statement %s' % (number + 1)
            p = t.cell(number, 1).paragraphs[0]._p
            r = OxmlElement('w:r'); p.append(r)
            drawing = OxmlElement('w:drawing'); r.append(drawing)
            ns = {'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
                  'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
                  'wps': 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape'}
            anchor = etree.SubElement(drawing, '{%s}anchor' % ns['wp'], nsmap=ns)
            shape = etree.SubElement(anchor, '{%s}wsp' % ns['wps'])
            prop = etree.SubElement(shape, '{%s}spPr' % ns['wps'])
            etree.SubElement(prop, '{%s}prstGeom' % ns['a'], prst='rect')
    p = doc.add_paragraph()
    for n in range(3):
        if n: p.add_run('\t')
        p.add_run('\t\t\t').underline = True
    doc.add_paragraph('Name of Participant\t\tDate\t\tSignature' if not mismatch else
                      'Name of Participant\t\tSignature')
    doc.add_paragraph()
    out = io.BytesIO(); doc.save(out)
    return out.getvalue()


class SigningAreas(unittest.TestCase):
    def test_captioned_signing_bands_preserve_types_and_actual_spans(self):
        result = word.inspect(signing_form())
        fields = result['fields']
        protected = [f for f in fields if f.get('write_protected')]
        self.assertEqual(len(protected), 5)
        bands = [f for f in protected if f['locator']['target']['type'] == 'span']
        self.assertEqual([(f['label'], f['kind']) for f in bands],
                         [('Name of Participant', 'text'), ('Date', 'date'), ('Signature', 'signature')])
        self.assertEqual([(f['locator']['target']['start'], f['locator']['target']['end']) for f in bands],
                         [(0, 3), (4, 7), (8, 11)])
        self.assertTrue(all('signing block' in f['review_reason'] for f in bands))
        self.assertEqual([f['label'] for f in fields if not f.get('write_protected')], ['Participant ID'])
        self.assertTrue(all(f['locator']['slot'] != 'b/p[4]' for f in fields))

    def test_floating_initials_have_no_guessed_consent_label(self):
        result = word.inspect(signing_form())
        boxes = [f for f in result['fields'] if f['kind'] == 'initials']
        self.assertEqual(len(boxes), 2)
        self.assertTrue(all('unresolved' in f['label'] for f in boxes))
        self.assertTrue(all('Consent statement' not in f['context'] for f in boxes))
        self.assertTrue(all(f['write_protected'] for f in boxes))
        self.assertTrue(any('floating' in note.lower() for note in result['notes']))

    def test_caption_count_mismatch_is_visible_and_not_guessed(self):
        result = word.inspect(signing_form(mismatch=True, table=False))
        held = [f for f in result['fields'] if f.get('write_protected')]
        self.assertEqual(len(held), 3)
        self.assertTrue(all('unresolved' in f['label'] for f in held))
        self.assertTrue(any('captions' in n for n in result['notes']))

    def test_protected_short_long_and_forged_flags_refuse_without_writing(self):
        data = signing_form()
        fields = word.inspect(data)['fields']
        self.assertEqual(sum(bool(f.get('write_protected')) for f in fields), 5)
        for field in [f for f in fields if f.get('write_protected')]:
            for answer in ('AB', 'Synthetic answer ' * 100):
                with self.subTest(label=field['label'], length=len(answer)):
                    with self.assertRaisesRegex(word.V3Failure, 'protected'):
                        word.fill(data, fields, {field['key']: answer})
                    forged = [{**f, 'kind': 'text', 'write_protected': False} for f in fields]
                    with self.assertRaisesRegex(word.V3Failure, 'protected'):
                        word.fill(data, forged, {field['key']: answer})

    def test_supported_answer_and_missing_preserve_all_signing_xml(self):
        data = signing_form(table=False)
        fields = word.inspect(data)['fields']
        key = next(f['key'] for f in fields if f['label'] == 'Participant ID')
        for answers in ({}, {key: 'SYNTHETIC-123'}, {key: 'SYNTHETIC ' * 100}):
            output = word.fill(data, fields, answers)
            with ZipFile(io.BytesIO(data)) as before, ZipFile(io.BytesIO(output)) as after:
                old = etree.fromstring(before.read('word/document.xml'))
                new = etree.fromstring(after.read('word/document.xml'))
                self.assertEqual([etree.tostring(x) for x in old.find(word.q('w:body'))][1:],
                                 [etree.tostring(x) for x in new.find(word.q('w:body'))][1:])

    def test_floating_layout_all_nonempty_exports_refuse_but_missing_is_unchanged(self):
        data = signing_form(); fields = word.inspect(data)['fields']
        self.assertIn(word.FLOATING_INITIALS_REFUSAL, word.inspect(data)['notes'])
        safe = next(f for f in fields if not f.get('write_protected'))
        for answer in ('A', 'SYNTHETIC ' * 100):
            with self.assertRaises(word.V3Failure) as caught:
                word.fill(data, fields, {safe['key']: answer})
            self.assertEqual(caught.exception.code, 'unqualified_floating_initials')
        with ZipFile(io.BytesIO(data)) as before, ZipFile(io.BytesIO(word.fill(data, fields, {}))) as after:
            self.assertTrue(all(before.read(n) == after.read(n) for n in before.namelist()))

    def test_every_primary_writer_and_paragraph_option_secondary_target_is_protected(self):
        from copy import deepcopy
        # No floating boxes: global export containment must not mask this bypass.
        data = signing_form(table=False); fields = word.inspect(data)['fields']
        safe = next(f for f in fields if not f.get('write_protected'))
        name = next(f for f in fields if f['label'] == 'Name of Participant')
        target_kinds = ('span', 'cell', 'cell_replace', 'paragraph', 'sdt_text', 'sdt_choice',
                        'ff_text', 'ff_dropdown', 'checkboxes', 'paragraph_options', 'text_options', 'table_rows')
        for kind in target_kinds:
            forged = deepcopy(safe); forged['locator'] = deepcopy(name['locator'])
            forged['locator']['target']['type'] = kind
            with self.subTest(kind=kind), self.assertRaisesRegex(word.V3Failure, 'protected'):
                word.fill(data, [forged], {forged['key']: 'AB'})
        forged = deepcopy(safe); forged.update(kind='choice', options=['Yes'])
        forged['locator']['target'] = {'type': 'paragraph_options', 'paragraphs': [name['locator']['slot']], 'printed': ['Yes']}
        with self.assertRaisesRegex(word.V3Failure, 'protected'):
            word.fill(data, [forged], {forged['key']: 'Yes'})

    def test_copies_and_descendant_slots_cannot_bypass_protected_source(self):
        from copy import deepcopy
        data = signing_form(); fields = word.inspect(data)['fields']
        safe = next(f for f in fields if not f.get('write_protected'))
        name = next(f for f in fields if f['label'] == 'Name of Participant')
        forged = deepcopy(safe)
        forged['locator']['target']['copies'] = [{k: name['locator'][k] for k in ('part', 'slot', 'fingerprint')}]
        with self.assertRaisesRegex(word.V3Failure, 'protected'):
            word.fill(data, [forged], {forged['key']: 'AB'})
        # An old mapping might target a blank cell below an overlapping shape.
        forged = deepcopy(safe)
        forged['locator'].update(slot='b/tbl[0]/tr[2]/tc[1]', fingerprint=None, target={'type': 'cell'})
        with self.assertRaisesRegex(word.V3Failure, 'protected'):
            word.fill(data, [forged], {forged['key']: 'AB'})
        box = next(f for f in fields if f['kind'] == 'initials')
        forged = deepcopy(safe)
        forged['locator'].update(slot=box['locator']['slot'] + '/p[0]', fingerprint=None)
        with self.assertRaisesRegex(word.V3Failure, 'protected'):
            word.fill(data, [forged], {forged['key']: 'AB'})

    def test_initials_instruction_does_not_cross_a_table_to_an_empty_paragraph(self):
        doc = Document(); doc.add_paragraph('Please initial box')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Printed instructions'
        doc.add_paragraph()
        out = io.BytesIO(); doc.save(out)
        self.assertEqual(word.inspect(out.getvalue())['fields'], [])

    def test_dated_inline_blank_is_a_date(self):
        doc = Document(); doc.add_paragraph('Information sheet dated.................... (version............)')
        out = io.BytesIO(); doc.save(out)
        self.assertEqual([f['kind'] for f in word.inspect(out.getvalue())['fields']], ['date', 'text'])
        doc = Document(); doc.add_paragraph('Dated and signed: __________')
        out = io.BytesIO(); doc.save(out)
        self.assertEqual(word.inspect(out.getvalue())['fields'][0]['kind'], 'signature')


class AssistancePlaces(unittest.TestCase):
    """What the engine itself holds for the model-assisted mapping place rule (documents.assistance.places)."""

    @staticmethod
    def places(data):
        from . import assistance
        result = assistance.places(data)
        return dict(result['held']), set(result['anchored'])

    def test_a_signing_block_caption_and_its_bands_are_held_part_by_part(self):
        held, anchored = self.places(signing_form(table=False))
        # The caption row is three printed labels; its signature part makes it a signing place although it names a date.
        self.assertEqual(word.label_parts('Name of Participant\t\tDate\t\tSignature'), ['Name of Participant', 'Date', 'Signature'])
        self.assertTrue(word.signing_wording('Name of Participant\t\tDate\t\tSignature'))
        self.assertFalse(word.signing_wording('Name of Participant\t\tDate'))
        self.assertEqual(held.get('word/document.xml#b/p[2]'), 'signing')
        self.assertEqual(held.get('word/document.xml#b/p[1]'), 'signing')
        self.assertNotIn('word/document.xml#b/p[0]', held)            # 'Participant ID: ____' is an ordinary answer
        self.assertEqual(anchored, set())

    def test_question_boxes_the_engine_reads_as_signing_or_authorship_are_held(self):
        def box_form(question):
            doc = Document(); doc.add_paragraph('Participant ID: __________'); doc.add_paragraph(question)
            table = doc.add_table(rows=1, cols=1); borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                edge = OxmlElement('w:' + side); edge.set(word.q('w:val'), 'single'); borders.append(edge)
            table.cell(0, 0)._tc.get_or_add_tcPr().append(borders)
            out = io.BytesIO(); doc.save(out); return out.getvalue()
        box = 'word/document.xml#b/tbl[0]/tr[0]/tc[0]'
        held, _ = self.places(box_form('I confirm the risks were explained and the patient gave consent:'))
        self.assertEqual(held.get(box), 'consent')
        held, _ = self.places(box_form('This section is for the patient to add any comments or concerns about their care'))
        self.assertEqual(held.get(box), 'protected')
        held, _ = self.places(box_form('Please describe the wound?'))
        self.assertNotIn(box, held)

    def test_text_boxes_frames_and_headers_are_anchored(self):
        doc = Document(); doc.add_paragraph('Plan: ______')
        frame = doc.add_paragraph('Office code: ______')
        prop = OxmlElement('w:framePr'); prop.set(word.q('w:w'), '2000'); frame._p.get_or_add_pPr().insert(0, prop)
        doc.sections[0].header.paragraphs[0].text = 'Clinic reference: ______'
        out = io.BytesIO(); doc.save(out)
        _, anchored = self.places(out.getvalue())
        self.assertIn('word/document.xml#b/p[1]', anchored)
        self.assertNotIn('word/document.xml#b/p[0]', anchored)
        self.assertTrue(any(slot.startswith('word/header') for slot in anchored))

    def test_fill_and_the_place_rule_read_one_protected_set(self):
        data = signing_form(table=False)
        package = word.open_package(data)
        source = word.Detector(package, layout_proposal=True); source.detect()
        _, protected, _ = word.protected_places(package, source, source.index)
        held, _ = self.places(data)
        self.assertTrue(protected)
        self.assertTrue(set(protected) <= set(held))
