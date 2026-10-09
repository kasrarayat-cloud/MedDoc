"""New prepared versions enforce source metadata without changing old versions."""
from copy import deepcopy
import hashlib
from pathlib import Path
import tempfile
import unittest
from docx import Document

from .engine import docx_inspect, DocumentFailure
from .preparation import inventory, candidate, PROTOCOL, TYPED_PROTOCOL, digest, validate_role
from .preparation_service import inspect_preparation, prepare_candidate
from .preparation_roles import role_rows
from .service import fill_template, DocumentError


class TypedPreparationIntegrationTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.original = self.root / 'original.docx'
        document = Document()
        document.add_paragraph('Keep all instructions. Use the printed units and choices.')
        document.add_paragraph('Visit date: ____ YYYY-MM-DD')
        document.add_paragraph('Count: {{Count}} (integer)')
        table = document.add_table(rows=3, cols=2)
        for row, label in zip(table.rows, ['Temperature (°C)', 'Fever (Yes / No)', 'Clinical assessment']):
            row.cells[0].text = label
        document.add_paragraph('Patient signature: ____')
        document.save(self.original)
        self.original_sha = hashlib.sha256(self.original.read_bytes()).hexdigest()

    def prepare(self, protocol=TYPED_PROTOCOL, isolated=False):
        inspect = inspect_preparation if isolated else inventory
        value = inspect(self.original, protocol=protocol)
        output = self.root / ('isolated.docx' if isolated else protocol + '.docx')
        selected = [anchor['id'] for anchor in value['anchors']]
        result = (prepare_candidate(self.original, value, selected, output) if isolated
                  else candidate(self.original, output, value, selected))
        return value, output, result['prepared_fields']

    def answers(self, fields):
        supplied = {'Visit date': '2026-10-09', 'Count': '2', 'Temperature (°C)': '37.2 °C',
                    'Fever (Yes / No)': 'No', 'Clinical assessment': 'Fictional synthetic demonstration.'}
        return {field['key']: supplied[field['label']] for field in fields}

    def test_actual_isolated_writer_keeps_original_and_binds_new_constraints(self):
        value, output, fields = self.prepare(isolated=True)
        self.assertEqual(value['protocol'], TYPED_PROTOCOL)
        typed = {field['label']: field.get('answer_constraints') for field in fields}
        self.assertEqual(typed['Temperature (°C)'], {'version': 1, 'data_type': 'number', 'unit': '°C'})
        self.assertEqual(typed['Fever (Yes / No)'], {'version': 1, 'data_type': 'choice', 'options': ['Yes', 'No']})
        self.assertEqual(typed['Count'], {'version': 1, 'data_type': 'integer'})
        self.assertIsNone(typed['Clinical assessment'])
        self.assertEqual(hashlib.sha256(self.original.read_bytes()).hexdigest(), self.original_sha)
        self.assertIn('Patient signature: ____', '\n'.join(p.text for p in Document(output).paragraphs))

    def test_wrong_unit_choice_integer_and_date_are_refused_before_export(self):
        _, source, fields = self.prepare()
        keys = {field['label']: field['key'] for field in fields}
        for label, bad in [('Temperature (°C)', '98.6 °F'), ('Fever (Yes / No)', 'Sometimes'),
                           ('Count', '2.5'), ('Visit date', '2026-02-30')]:
            answers = self.answers(fields)
            answers[keys[label]] = bad
            output = self.root / (keys[label] + '-bad.docx')
            with self.subTest(label=label), self.assertRaises(DocumentError) as refusal:
                fill_template(source, fields, answers, output)
            self.assertEqual(refusal.exception.code, 'invalid_answer_constraint')
            self.assertFalse(output.exists())

    def test_corrected_and_reused_mapping_exports_complete_source_preserving_documents(self):
        _, source, fields = self.prepare()
        approved = deepcopy(fields)
        for name in ['typed', 'reviewed-voice']:
            output = self.root / (name + '.docx')
            fill_template(source, fields, self.answers(fields), output)
            document = Document(output)
            self.assertEqual(document.tables[0].cell(0, 0).text, 'Temperature (°C)')
            self.assertEqual(document.tables[0].cell(0, 1).text, '37.2')
            self.assertEqual(document.tables[0].cell(1, 1).text, 'No')
            self.assertIn('Patient signature: ____', '\n'.join(p.text for p in document.paragraphs))
            self.assertEqual(fields, approved)
        self.assertEqual(hashlib.sha256(self.original.read_bytes()).hexdigest(), self.original_sha)

    def test_old_version_keeps_its_original_metadata_and_inventory(self):
        old = inventory(self.original)
        self.assertEqual(old['protocol'], PROTOCOL)
        self.assertFalse(any(field.get('answer_constraints') for field in old['fields']))
        before = deepcopy(old)
        self.prepare()
        self.assertEqual(inventory(self.original), before)
        self.assertFalse(any(field.get('answer_constraints') for field in docx_inspect(self.original)[3]))

    def test_both_roles_must_report_source_truth_even_when_they_agree(self):
        value, _, _ = self.prepare()
        for role in ['proposer', 'checker']:
            response = {'inventory_sha256': value['inventory_sha256'], 'anchors': role_rows(value, role),
                        'stories': []}
            from .preparation import model_inventory
            response['stories'] = [{'story_id': story['id'], 'content_sha256': story['content_sha256'],
                                   'verdict': 'retain'} for story in model_inventory(value)['stories']]
            validate_role(value, response, role)
            row = next(row for row in response['anchors'] if row['metadata']['unit'] == '°C')
            row['metadata']['unit'] = '°F'
            with self.subTest(role=role), self.assertRaises(DocumentFailure):
                validate_role(value, response, role)

    def test_protocol_and_digest_tampering_cannot_reinterpret_original_inventory(self):
        value = inventory(self.original)
        value['protocol'] = TYPED_PROTOCOL
        value['inventory_sha256'] = digest({key: item for key, item in value.items() if key != 'inventory_sha256'})
        with self.assertRaises(DocumentFailure):
            candidate(self.original, self.root / 'forged.docx', value, [a['id'] for a in value['anchors']])

