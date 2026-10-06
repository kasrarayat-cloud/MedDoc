"""Drawn signing bands, signing-caption holds, and the withdrawn paragraph relaxations' counterexamples.

Pure XML/parser controls on project-authored synthetic documents; no corpus content, no app resources.
"""
import collections
import copy
import io
import json
import unittest
from unittest import mock

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from . import word_v3 as word

NAMESPACES = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'wps': 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape',
    'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
    'v': 'urn:schemas-microsoft-com:vml',
}
DECLARED = ' '.join('xmlns:%s="%s"' % item for item in NAMESPACES.items())
BODY = 'word/document.xml#'


def data(document):
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def line_run(number, *, width=2286000, height=0, prst='line', anchored=False, textbox=False, fallback=True,
             drawn_height=0):
    """One drawn line as Word writes it: modern DrawingML with an optional VML fallback."""
    container = 'anchor' if anchored else 'inline'
    attributes = (' behindDoc="0" distT="0" distB="0" distL="0" distR="0" simplePos="0" locked="0"'
                  ' layoutInCell="1" allowOverlap="1" relativeHeight="%d"' % number) if anchored else \
        ' distT="0" distB="0" distL="0" distR="0"'
    position = ('<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset>'
                '</wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
                if anchored else '')
    text = '<wps:txbx><w:txbxContent><w:p><w:r><w:t>x</w:t></w:r></w:p></w:txbxContent></wps:txbx>' if textbox else ''
    drawing = (
        '<w:drawing><wp:{c}{attrs}>{pos}<wp:extent cx="{w}" cy="{h}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '{wrap}<wp:docPr id="{n}" name="Line {n}" descr="blank line"/><wp:cNvGraphicFramePr/><a:graphic>'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp>'
        '<wps:cNvCnPr/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{w}" cy="{dh}"/></a:xfrm>'
        '<a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom><a:ln w="6350"><a:solidFill><a:srgbClr val="000000"/>'
        '</a:solidFill></a:ln></wps:spPr>{text}<wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:{c}>'
        '</w:drawing>').format(c=container, attrs=attributes, pos=position, w=width, h=height, n=number,
                               prst=prst, dh=drawn_height, text=text, wrap='<wp:wrapNone/>' if anchored else '')
    if fallback:
        drawing = ('<mc:AlternateContent><mc:Choice Requires="wps">%s</mc:Choice><mc:Fallback><w:pict>'
                   '<v:line id="Line %d" style="position:absolute" from="0,0" to="180pt,0" strokeweight=".5pt"/>'
                   '</w:pict></mc:Fallback></mc:AlternateContent>') % (drawing, number)
    return etree.fromstring('<w:r %s>%s</w:r>' % (DECLARED, drawing))


def drawn_band(caption='Signature of Patient or Representative\t\t\tDate', *, lines=2, spacer=True, spacer_xml='',
               **shape):
    """Clinical summary / drawn lines / caption below / spacer / ordinary visit date."""
    doc = Document()
    doc.add_paragraph('Clinical summary: ________')
    paragraph = doc.add_paragraph()
    for number in range(lines):
        if number:
            paragraph.add_run('\t\t')
        paragraph._p.append(line_run(number + 1, **shape))
    doc.add_paragraph(caption)
    if spacer:
        blank = doc.add_paragraph()
        if spacer_xml:
            for node in etree.fromstring('<w:wrap %s>%s</w:wrap>' % (DECLARED, spacer_xml)):
                blank._p.append(node)
    doc.add_paragraph('Visit date: ________')
    return data(doc)


def slot_bytes(raw, slot):
    return word.fingerprint(word.SlotIndex(word.open_package(raw)).by_id[BODY + slot].el)


class DrawnSigningBand(unittest.TestCase):
    def test_inline_line_pair_is_one_protected_signing_event(self):
        raw = drawn_band()
        fields = word.inspect(raw)['fields']
        held = [f for f in fields if f.get('write_protected')]
        self.assertEqual([(f['kind'], f['label'], f['locator']['slot'], f['locator']['target']['type']) for f in held],
                         [('signature', 'Signature of Patient or Representative', 'b/p[1]', 'none'),
                          ('date', 'Signature of Patient or Representative Date', 'b/p[1]', 'none')])
        for field in held:
            relation = field['source_questions'][0]
            self.assertEqual(relation['relation'], 'protected_inline_signing_lines')
            self.assertEqual(relation['source']['slot'], 'b/p[2]')
            self.assertEqual(len(relation['drawings']), 2)
        # Neither the caption nor the spacer below it becomes an answer destination.
        self.assertFalse([f for f in fields if f['locator']['slot'] in ('b/p[2]', 'b/p[3]')])
        for field in held:
            with self.assertRaises(word.V3Failure):
                word.fill(raw, fields, {field['key']: '14/03/2026' if field['kind'] == 'date' else 'Pat Example'})

    def test_ordinary_answers_fill_while_lines_fallbacks_caption_and_spacer_stay_byte_identical(self):
        raw = drawn_band()
        fields = word.inspect(raw)['fields']
        visit = next(f for f in fields if f['label'] == 'Visit date')
        summary = next(f for f in fields if f['label'] == 'Clinical summary')
        final = word.fill(raw, fields, {visit['key']: '14/03/2026', summary['key']: 'Synthetic stable review.'})
        for slot in ('b/p[1]', 'b/p[2]', 'b/p[3]'):
            self.assertEqual(slot_bytes(raw, slot), slot_bytes(final, slot), slot)
        index = word.SlotIndex(word.open_package(final))
        self.assertIn('14/03/2026', word.paragraph_text(index.by_id[BODY + 'b/p[4]'].el)[0])
        drawing = index.by_id[BODY + 'b/p[1]'].el
        self.assertEqual(len(drawing.findall('.//{%s}Fallback' % NAMESPACES['mc'])), 2)

    def test_forged_ordinary_date_on_the_signing_spacer_is_refused(self):
        raw = drawn_band()
        fields = word.inspect(raw)['fields']
        spacer = word.SlotIndex(word.open_package(raw)).by_id[BODY + 'b/p[3]']
        forged = {'key': 'v3_99', 'kind': 'date', 'label': 'Signature of Patient or Representative Date', 'context': '',
                  'locator': {'contract': 'v3', 'part': 'word/document.xml', 'slot': 'b/p[3]',
                              'fingerprint': word.fingerprint(spacer.el), 'target': {'type': 'paragraph'}}}
        with self.assertRaises(word.V3Failure):
            word.fill(raw, fields + [forged], {'v3_99': '14/03/2026'})
        with self.assertRaises(word.V3Failure):
            word.fill(raw, [forged], {'v3_99': '14/03/2026'})

    def test_other_drawings_or_no_drawing_are_not_a_drawn_band(self):
        # Only the exact inline line pair is a drawn band; every other shape keeps the engine's existing reading.
        shapes = {'anchored': {'anchored': True}, 'rectangle': {'prst': 'rect'}, 'text box': {'textbox': True},
                  'tall': {'height': 50800}, 'diagonal': {'drawn_height': 100}, 'one line': {'lines': 1},
                  'three lines': {'lines': 3}, 'no drawing': {'lines': 0}}
        for name, options in shapes.items():
            with self.subTest(name):
                fields = word.inspect(drawn_band(**options))['fields']
                self.assertFalse([f for f in fields if any(r.get('relation') == 'protected_inline_signing_lines'
                                                           for r in f.get('source_questions', []))])

    def test_ordinary_captions_under_lines_are_not_signing_events(self):
        raw = drawn_band(caption='Printed name\t\t\tTitle')
        fields = word.inspect(raw)['fields']
        self.assertFalse([f for f in fields if f.get('write_protected')])
        self.assertFalse([f for f in fields if any(r.get('relation') == 'protected_inline_signing_lines'
                                                   for r in f.get('source_questions', []))])

    def test_a_blank_under_the_drawn_caption_stays_held_whatever_its_empty_formatting(self):
        variants = {'tab only': '<w:r><w:tab/></w:r>',
                    'empty underlined run': '<w:r><w:rPr><w:u w:val="single"/></w:rPr></w:r>',
                    'underlined paragraph mark': '<w:pPr><w:rPr><w:u w:val="single"/></w:rPr></w:pPr>'}
        for name, xml in variants.items():
            with self.subTest(name):
                raw = drawn_band(spacer_xml=xml)
                fields = word.inspect(raw)['fields']
                self.assertFalse([f for f in fields if f['locator']['slot'] == 'b/p[3]'])
                blank = word.SlotIndex(word.open_package(raw)).by_id[BODY + 'b/p[3]']
                stored = {'key': 'v3_9', 'kind': 'date', 'label': 'Signature of Patient or Representative Date',
                          'context': '', 'locator': {'contract': 'v3', 'part': 'word/document.xml', 'slot': 'b/p[3]',
                                                     'fingerprint': word.fingerprint(blank.el),
                                                     'target': {'type': 'paragraph'}}}
                with self.assertRaises(word.V3Failure):
                    word.fill(raw, [stored], {'v3_9': '14/03/2026'})
                with self.assertRaises(word.V3Failure):
                    word.locate(raw, [{'kind': 'date', 'label': 'Date', 'slot': BODY + 'b/p[3]'}])

    def test_a_spacer_that_anchors_a_text_box_is_not_part_of_the_band(self):
        # Refutation 2026-10-05 (round 3): an empty-looking spacer that anchors a floating text box was claimed,
        # so the text box's own field was offered and then every export refused.
        box = ('<w:r><w:drawing><wp:anchor behindDoc="0" distT="0" distB="0" distL="0" distR="0" simplePos="0" '
               'locked="0" layoutInCell="1" allowOverlap="1" relativeHeight="9"><wp:simplePos x="0" y="0"/>'
               '<wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
               '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
               '<wp:extent cx="2286000" cy="457200"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapSquare wrapText="bothSides"/>'
               '<wp:docPr id="9" name="Box 9"/><wp:cNvGraphicFramePr/><a:graphic><a:graphicData '
               'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr/>'
               '<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2286000" cy="457200"/></a:xfrm>'
               '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr><wps:txbx><w:txbxContent><w:p><w:r>'
               '<w:t xml:space="preserve">Office use - reference number: ________</w:t></w:r></w:p></w:txbxContent>'
               '</wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>')
        raw = drawn_band(spacer_xml=box)
        fields = word.inspect(raw)['fields']
        band = [f for f in fields if any(r.get('relation') == 'protected_inline_signing_lines'
                                         for r in f.get('source_questions', []))]
        self.assertEqual(len(band), 2)
        writable = [f for f in fields if not f.get('write_protected') and f['kind'] not in ('signature', 'initials')]
        self.assertIn('word/document.xml#b/p[3]/txbx[0]/p[0]',
                      [f['locator']['part'] + '#' + f['locator']['slot'] for f in writable])
        answers = {f['key']: '14/03/2026' if f['kind'] == 'date' else 'REF-123' for f in writable}
        final = word.fill(raw, fields, answers)
        self.assertIn('REF-123', ''.join(word.text_of(n) for n in
                      word.SlotIndex(word.open_package(final)).by_id[BODY + 'b/p[3]'].el.iter(word.TXBX)))

    def test_invisible_range_markers_never_separate_the_band(self):
        # Refutation round 4: a body-level bookmark or permission marker between the lines and their caption hid
        # the band, and the signing date was invented again (also on H17 with one inserted bookmark).
        for marker in ('<w:bookmarkEnd %s w:id="7"/>', '<w:bookmarkStart %s w:id="8" w:name="sig"/>',
                       '<w:permStart %s w:id="9" w:edGrp="everyone"/>', '<w:commentRangeStart %s w:id="3"/>'):
            with self.subTest(marker=marker.split()[0]):
                doc = Document(io.BytesIO(drawn_band()))
                doc.paragraphs[1]._p.addnext(etree.fromstring(marker % DECLARED))
                raw = data(doc)
                fields = word.inspect(raw)['fields']
                band = [f for f in fields if any(r.get('relation') == 'protected_inline_signing_lines'
                                                 for r in f.get('source_questions', []))]
                self.assertEqual(len(band), 2)
                self.assertFalse([f for f in fields if f['kind'] == 'date' and not f.get('write_protected')
                                  and f['locator']['slot'] != 'b/p[4]'])

    def test_a_text_box_or_control_in_the_lines_paragraph_is_never_swallowed(self):
        # Refutation round 4: a VML text box beside the drawn lines was offered as a field and then every export
        # refused; the lines paragraph now holds the lines alone or it is not a band.
        vml = ('<w:r %s><w:pict><v:shape type="#_x0000_t202" style="width:100pt;height:20pt"><v:textbox>'
               '<w:txbxContent><w:p><w:r><w:t xml:space="preserve">Hospital number: ________</w:t></w:r></w:p>'
               '</w:txbxContent></v:textbox></v:shape></w:pict></w:r>' % DECLARED)
        doc = Document(io.BytesIO(drawn_band()))
        doc.paragraphs[1]._p.append(etree.fromstring(vml))
        raw = data(doc)
        fields = word.inspect(raw)['fields']
        self.assertFalse([f for f in fields if any(r.get('relation') == 'protected_inline_signing_lines'
                                                   for r in f.get('source_questions', []))])
        writable = [f for f in fields if not f.get('write_protected') and f['kind'] not in ('signature', 'initials')]
        word.fill(raw, fields, {f['key']: '14/03/2026' if f['kind'] == 'date' else 'Synthetic answer' for f in writable})

    def test_a_clinical_sign_caption_is_not_a_signing_band(self):
        # Rounds 4-5: a clinical 'sign', an adjective 'initial' or an identifier 'initials' is not a signing act.
        for caption in ('Name of sign if positive\t\tSite', 'Troponin (initial)\t\tTime',
                        'Initial and repeat troponin\t\tResult', 'Patient initials\t\tDate of birth'):
            with self.subTest(caption):
                fields = word.inspect(drawn_band(caption=caption))['fields']
                self.assertFalse([f for f in fields if any(r.get('relation') == 'protected_inline_signing_lines'
                                                           for r in f.get('source_questions', []))])

    def test_a_visible_writing_line_below_the_band_is_kept_and_reviewed(self):
        # Refutation round 4: a ruled line under the caption may be another field's line. It stays unwritten and,
        # as in the previous reading, a clinician reviews it before generation. Round 5: a positional-tab leader
        # draws a line too.
        for name, xml in (('underlined tab', '<w:r><w:rPr><w:u w:val="single"/></w:rPr><w:tab/></w:r>'),
                          ('bottom border', '<w:pPr><w:pBdr><w:bottom w:val="single" w:sz="4" w:space="1" '
                                            'w:color="000000"/></w:pBdr></w:pPr>'),
                          ('leader tab', '<w:pPr><w:tabs><w:tab w:val="right" w:leader="underscore" w:pos="9000"/>'
                                         '</w:tabs></w:pPr><w:r><w:tab/></w:r>'),
                          ('positional tab leader', '<w:r><w:ptab w:relativeTo="margin" w:alignment="right" '
                                                    'w:leader="underscore"/></w:r>')):
            with self.subTest(name):
                found = word.inspect(drawn_band(spacer_xml=xml))
                self.assertFalse([f for f in found['fields'] if f['locator']['slot'] == 'b/p[3]'])
                self.assertEqual([i['code'] for i in found['coverage_issues'] if i['slot'] == 'b/p[3]'],
                                 ['ambiguous_signing_band'])
        found = word.inspect(drawn_band())
        self.assertFalse([i for i in found['coverage_issues'] if i['slot'] == 'b/p[3]'])

    def test_a_typed_writing_line_under_the_caption_is_the_bands_reviewed_spacer(self):
        # Refutation round 7: the band claims its caption, so a typed line under it (underscores, dots) lost its
        # label and became a neutral writable 'Answer' that generation filled - the line e0 read as the signature
        # area. It is the band's spacer, kept unwritten and reviewed exactly like a drawn line; a labelled line
        # ('Print name: ____') stays its own field.
        caption = 'Patient signature\t\t\tPrint name'
        for name, line in (('underscores', '_' * 30), ('dots', '.' * 30)):
            with self.subTest(name):
                raw = drawn_band(caption, spacer_xml='<w:r><w:t>%s</w:t></w:r>' % line)
                found = word.inspect(raw)
                self.assertFalse([f for f in found['fields'] if f['locator']['slot'] == 'b/p[3]'])
                self.assertEqual([i['detail'] for i in found['coverage_issues'] if i['slot'] == 'b/p[3]'],
                                 [word.DRAWN_BAND_SPACER_REVIEW + 'Location: ' + BODY + 'b/p[3]'])
                self.assertIn(BODY + 'b/p[3]', word.signing_band_places(raw))
                writable = [f for f in found['fields'] if not f.get('write_protected')]
                final = word.fill(raw, found['fields'], {f['key']: 'Alex Example' for f in writable})
                self.assertEqual(slot_bytes(final, 'b/p[3]'), slot_bytes(raw, 'b/p[3]'))
        raw = drawn_band(caption, spacer_xml='<w:r><w:t xml:space="preserve">Print name: ________</w:t></w:r>')
        found = word.inspect(raw)
        self.assertEqual([(f['label'], bool(f.get('write_protected'))) for f in found['fields']
                          if f['locator']['slot'] == 'b/p[3]'], [('Print name', False)])
        self.assertFalse([i for i in found['coverage_issues'] if i['slot'] == 'b/p[3]'])
        self.assertNotIn(BODY + 'b/p[3]', word.signing_band_places(raw))
        # Refutation round 8: placeholders and bracketed prompts name their own answers; a line of them under the
        # caption stays the form's own fields, never the band's spacer.
        for line, labels in (('{{Next review date}}\t\t{{Review clinic}}', ['Next review date', 'Review clinic']),
                             ('[Insert clinician name]\t[Insert job title]', ['clinician name', 'job title'])):
            with self.subTest(line=line):
                xml = ''.join('<w:r><w:t xml:space="preserve">%s</w:t></w:r>' % part if part else '<w:r><w:tab/></w:r>'
                              for part in line.replace('\t', '\n\t\n').split('\n')
                              for part in ([''] if part == '\t' else [part]) if part or part == '')
                raw = drawn_band(caption, spacer_xml=xml)
                found = word.inspect(raw)
                self.assertEqual([f['label'] for f in found['fields'] if f['locator']['slot'] == 'b/p[3]'], labels)
                self.assertFalse([i for i in found['coverage_issues'] if i['slot'] == 'b/p[3]'])
                self.assertNotIn(BODY + 'b/p[3]', word.signing_band_places(raw))

    def test_a_multi_line_typed_signing_area_is_wholly_the_bands(self):
        # Refutation round 8: only the first typed line became the band's spacer; lines 2+ of a typed signing area
        # (one signature field on e0) lost the caption's label and became neutral writable answers. Round 9: the band
        # takes exactly the caption's own lines - what a band-free reading labels from the caption - so lines under
        # a typed spacer are the band's, while typed lines after an EMPTY spacer (e0 never read them as the
        # caption's) and the form's next own field are not.
        for first in ('typed', 'empty'):
            with self.subTest(first=first):
                doc = Document()
                doc.add_paragraph('Clinical summary: ________')
                lines = doc.add_paragraph()
                lines._p.append(line_run(1))
                lines.add_run('\t\t')
                lines._p.append(line_run(2))
                doc.add_paragraph('Patient signature\t\t\tPrint name')
                doc.add_paragraph('_' * 30 if first == 'typed' else '')
                doc.add_paragraph('_' * 30)
                doc.add_paragraph('.' * 30)
                doc.add_paragraph('Visit date: ________')
                raw = data(doc)
                found = word.inspect(raw)
                band = word.signing_band_places(raw)
                reviewed = sorted(i['slot'] for i in found['coverage_issues'] if i['code'] == 'ambiguous_signing_band')
                if first == 'typed':
                    for place in ('b/p[3]', 'b/p[4]', 'b/p[5]'):
                        self.assertIn(BODY + place, band)
                        self.assertFalse([f for f in found['fields'] if f['locator']['slot'] == place])
                    self.assertEqual(reviewed, ['b/p[3]', 'b/p[4]', 'b/p[5]'])
                else:
                    self.assertIn(BODY + 'b/p[3]', band)
                    self.assertNotIn(BODY + 'b/p[4]', band)
                    self.assertNotIn(BODY + 'b/p[5]', band)
                    self.assertEqual(reviewed, [])
                visit = [f for f in found['fields'] if f['locator']['slot'] == 'b/p[6]']
                self.assertEqual([(f['label'], bool(f.get('write_protected'))) for f in visit], [('Visit date', False)])
                writable = [f for f in found['fields'] if not f.get('write_protected')]
                final = word.fill(raw, found['fields'], {f['key']: 'Alex Example' for f in writable})
                for place in ('b/p[3]', 'b/p[4]', 'b/p[5]') if first == 'typed' else ('b/p[3]',):
                    self.assertEqual(slot_bytes(final, place), slot_bytes(raw, place))

    def test_the_band_takes_exactly_the_captions_own_lines(self):
        # Refutation round 9: a typed-line heuristic both missed lines e0 read as the caption's signing area (grey
        # prompts, segmented or boxed dates, a line plus a date prompt) - leaving them neutral writable answers - and
        # swallowed lines with their own reading (a row captioned below it). The band now takes exactly what a
        # band-free reading labels from its caption.
        grey = '<w:r><w:rPr><w:color w:val="808080"/></w:rPr><w:t xml:space="preserve">%s</w:t></w:r>'
        tabs = '<w:r><w:tab/><w:tab/><w:tab/><w:tab/></w:r>'

        def text(value):
            return '<w:r><w:t xml:space="preserve">%s</w:t></w:r>' % value

        cases = (
            ('Patient signature\t\t\tPrint name', [tabs + grey % 'Enter printed name'], [True]),
            ('Signature of Patient\t\t\tDate', [tabs + text('|__|__|/|__|__|/|__|__|__|__|')], [True]),
            ('Signature of Patient\t\t\tDate', [text('_' * 22) + '<w:r><w:tab/><w:tab/><w:tab/></w:r>' + grey % 'dd/mm/yyyy'],
             [True]),
            ('Signature of patient', [text('_' * 30), text('_' * 30), text('Printed name of patient')], [True, True, False]),
            ('Signature\t\t\tDate\t\t\tTime', [text('_' * 30)], [True]),
            ('Signature of Patient\t\t\tDate', ['', text('_' * 22) + '<w:r><w:tab/><w:tab/></w:r>' + text('_' * 22),
                                                text('Name of interpreter') + '<w:r><w:tab/><w:tab/><w:tab/></w:r>'
                                                + text('Interpreter ID')], [True, False, False]),
        )
        for caption, after, held in cases:
            with self.subTest(caption=caption, after=after[0][:40]):
                doc = Document()
                doc.add_paragraph('Clinical summary: ________')
                lines = doc.add_paragraph()
                for number in range(len([x for x in caption.split('\t') if x.strip()])):
                    if number:
                        lines.add_run('\t\t')
                    lines._p.append(line_run(number + 1, fallback=False))
                doc.add_paragraph(caption)
                for xml in after:
                    paragraph = doc.add_paragraph()
                    for node in list(etree.fromstring('<w:wrap %s>%s</w:wrap>' % (DECLARED, xml))):
                        paragraph._p.append(node)
                doc.add_paragraph('End of record.')
                raw = data(doc)
                band = word.signing_band_places(raw)
                found = word.inspect(raw)
                writable = [f for f in found['fields'] if not f.get('write_protected')]
                final = word.fill(raw, found['fields'], {f['key']: 'Alex Example' for f in writable})
                for number, expected in enumerate(held):
                    place = 'b/p[%d]' % (3 + number)
                    self.assertEqual(BODY + place in band, expected, place)
                    if expected:
                        self.assertFalse([f for f in found['fields'] if f['locator']['slot'] == place])
                        self.assertEqual(slot_bytes(final, place), slot_bytes(raw, place))
                        if after[number]:
                            self.assertIn(place, [i['slot'] for i in found['coverage_issues']
                                                  if i['code'] == 'ambiguous_signing_band'])

    def test_the_band_reads_its_caption_through_side_headings_and_header_copies(self):
        # Refutation round 10: a side heading ('Consent for right knee arthroscopy') qualifies every later label
        # ('<heading> - Signature of patient'), so plain label equality released the caption's typed line as a
        # writable answer; and identical header copies merge into one area whose second copy is a target copy, so the
        # second header's line was released too. Provenance, qualified labels and copies now all count.
        line = '_' * 30
        grey = etree.fromstring('<w:r %s><w:rPr><w:color w:val="808080"/></w:rPr><w:t>Click here to enter text.</w:t>'
                                '</w:r>' % DECLARED)
        for after in ([line], [line, line], ['grey'], ['line and date prompt']):
            with self.subTest(after=after):
                doc = Document()
                doc.add_paragraph('Synthetic consent record')
                doc.add_paragraph().add_run('Consent for right knee arthroscopy').bold = True
                doc.add_paragraph('PATIENT NAME: ')._p.append(
                    etree.fromstring('<w:r %s><w:rPr><w:u w:val="single"/></w:rPr><w:tab/></w:r>' % DECLARED))
                doc.add_paragraph()._p.append(line_run(1, fallback=False))
                if after == ['line and date prompt']:
                    # The caption's columns read onto the line ('<heading> - Signature of patient Date', no source).
                    doc.paragraphs[-1]._p.append(line_run(2, fallback=False))
                    doc.add_paragraph('Signature of patient\t\t\tDate')
                else:
                    doc.add_paragraph('Signature of patient')
                for text in after:
                    if text == 'grey':
                        doc.add_paragraph()._p.append(copy.deepcopy(grey))
                    elif text == 'line and date prompt':
                        prompt = doc.add_paragraph('_' * 22 + '\t\t\t')
                        prompt._p.append(etree.fromstring('<w:r %s><w:rPr><w:color w:val="808080"/></w:rPr>'
                                                          '<w:t>dd/mm/yyyy</w:t></w:r>' % DECLARED))
                    else:
                        doc.add_paragraph(text)
                doc.add_paragraph('End of record.')
                raw = data(doc)
                band = word.signing_band_places(raw)
                found = word.inspect(raw)
                writable = [f for f in found['fields'] if not f.get('write_protected')]
                final = word.fill(raw, found['fields'], {f['key']: 'Alex Example' for f in writable})
                for number in range(len(after)):
                    place = 'b/p[%d]' % (5 + number)
                    self.assertIn(BODY + place, band)
                    self.assertFalse([f for f in found['fields'] if f['locator']['slot'] == place])
                    self.assertIn(place, [i['slot'] for i in found['coverage_issues']])
                    self.assertEqual(slot_bytes(final, place), slot_bytes(raw, place))
        for lines in (1, 2):
            # (A copy's further lines are its own places: the merged area names only the copy's first line.)
            with self.subTest(header_lines=lines):
                doc = Document()
                doc.add_paragraph('PATIENT NAME: ')._p.append(
                    etree.fromstring('<w:r %s><w:rPr><w:u w:val="single"/></w:rPr><w:tab/></w:r>' % DECLARED))
                section = doc.sections[0]
                section.different_first_page_header_footer = True
                for header in (section.header, section.first_page_header):
                    header.is_linked_to_previous = False
                    header.paragraphs[0]._p.append(line_run(1, fallback=False))
                    header.add_paragraph('Signature of patient')
                    for _ in range(lines):
                        header.add_paragraph(line)
                raw = data(doc)
                band = word.signing_band_places(raw)
                for part in ('word/header1.xml', 'word/header2.xml'):
                    for number in range(lines):
                        self.assertIn('%s#hf/p[%d]' % (part, 2 + number), band)
                found = word.inspect(raw)
                self.assertFalse([f for f in found['fields'] if f['locator']['part'] != 'word/document.xml'
                                  and not f.get('write_protected')])

    def test_a_token_under_the_caption_is_held_and_reviewed_yet_mappable_and_legacy_fills_it(self):
        # Refutation rounds 10-11: whose field a lone {{token}} under the caption names cannot be read from its wording -
        # '{{patient_signature}}', '{{Signature1}}' and '{{date_signed}}' passed every signing reader. On v3 it stays the
        # band's, held and reviewed as e0's v3 reading held it (the caption's signature field), but it is not protected:
        # a clinician who finds it is another field ('{{Next review date}}') maps it by hand. A bracket stays protected.
        # On the legacy route a legacy {{token}} in the band no longer holds the whole import (round 11 routing:1): legacy
        # fills it by name, as e0's legacy route did.
        import tempfile
        from pathlib import Path
        from . import engine

        def letter(caption, token):
            doc = Document()
            doc.add_paragraph('Synthetic clinic letter')
            doc.add_paragraph('PATIENT NAME: {{Patient full name}}')
            lines = doc.add_paragraph()
            for number in range(caption.count('\t\t\t') + 1):
                if number:
                    lines.add_run('\t\t')
                lines._p.append(line_run(number + 1, fallback=False))
            doc.add_paragraph(caption)
            doc.add_paragraph(token)
            doc.add_paragraph('End of letter.')
            return data(doc)

        for caption, token in [('Signature of clinician', '{{Clinician name}}'), ('Signature of clinician', '{{Next review date}}'),
                               ('Patient signature', '{{Date signed}}'), ('Patient signature', '{{patient_signature}}'),
                               ('Patient signature', '{{Signature1}}'), ('Patient signature\t\t\tDate', '{{Date}}'),
                               ('Patient signature', '[Patient name]')]:
            with self.subTest(caption=caption, token=token):
                raw = letter(caption, token)
                found = word.inspect(raw)
                self.assertIn(BODY + 'b/p[4]', word.signing_band_places(raw))
                self.assertFalse([f for f in found['fields'] if f['locator']['slot'] == 'b/p[4]'])
                self.assertIn('b/p[4]', [i['slot'] for i in found['coverage_issues'] if i['code'] == 'ambiguous_signing_band'])
                target = [{'kind': 'text', 'label': 'Mapped by hand', 'slot': 'word/document.xml#b/p[4]'}]
                if token.startswith('['):
                    with self.assertRaises(word.V3Failure):      # a bracket line stays protected
                        word.locate(raw, target)
                    continue
                self.assertFalse(word.destination_roles(raw, word.locate(raw, target))['violations'])
                with tempfile.TemporaryDirectory() as tmp:
                    workspace = Path(tmp)
                    (workspace / 'input.docx').write_bytes(raw)
                    result = engine.process(workspace, {'format': 'docx', 'operation': 'inspect', 'engine': 'legacy'})
                    self.assertEqual(result['import_manifest']['engine'], 'legacy')
                    coverage = result['import_manifest']['coverage']
                    self.assertFalse(coverage.get('block_generation'))
                    self.assertFalse(coverage.get('advanced_unresolved_areas'))
                    key = next(f['key'] for f in result['fields'] if f['label'] == token.strip('{}'))
                    engine.process(workspace, {'format': 'docx', 'operation': 'fill', 'mapping': result['fields'],
                                               'answers': {key: 'Synthetic answer'}, 'engine': 'legacy'})
                    self.assertIn('Synthetic answer', Document(workspace / 'output.docx').paragraphs[4].text)

    def test_a_spacer_that_prints_nothing_raises_no_review(self):
        # Refutation round 5: breaks, a plain tab and paragraph-mark formatting print no line; asking a clinician
        # to review them held generation for nothing.
        for name, xml in (('page break', '<w:r><w:br w:type="page"/></w:r>'), ('line break', '<w:r><w:br/></w:r>'),
                          ('plain tab', '<w:r><w:tab/></w:r>'),
                          ('underlined mark', '<w:pPr><w:rPr><w:u w:val="single"/></w:rPr></w:pPr>'),
                          ('highlighted mark', '<w:pPr><w:rPr><w:highlight w:val="yellow"/></w:rPr></w:pPr>'),
                          ('numbering off', '<w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="0"/></w:numPr></w:pPr>')):
            with self.subTest(name):
                found = word.inspect(drawn_band(spacer_xml=xml))
                self.assertFalse([f for f in found['fields'] if f['locator']['slot'] == 'b/p[3]'])
                self.assertFalse([i for i in found['coverage_issues'] if i['slot'] == 'b/p[3]'])

    def test_a_bare_date_caption_under_its_line_is_not_promoted_to_a_signing_event(self):
        raw = drawn_band(caption='Date', lines=1)
        fields = word.inspect(raw)['fields']
        self.assertFalse([f for f in fields if f.get('write_protected')])

    def test_a_drawn_band_inside_a_table_cell_leaves_the_table_log_writable(self):
        # A protected region inside a table makes its table_rows writer refuse, so the drawn band is read at body
        # level only (refutation 2026-10-05). The in-cell paragraphs keep the engine's existing reading.
        doc = Document()
        doc.add_paragraph('Synthetic medication record')
        table = doc.add_table(rows=4, cols=3)
        for n, text in enumerate(('Medication', 'Dose', 'Route')):
            table.cell(0, n).text = text
        cell = table.cell(3, 0).merge(table.cell(3, 2))
        lines = cell.paragraphs[0]
        lines._p.append(line_run(1, fallback=False))
        lines.add_run('\t\t')
        lines._p.append(line_run(2, fallback=False))
        cell.add_paragraph('Signature of Patient or Representative\t\t\tDate')
        cell.add_paragraph('')
        doc.add_paragraph('Neighbor findings: ______')
        raw = data(doc)
        fields = word.inspect(raw)['fields']
        self.assertFalse([f for f in fields if any(r.get('relation') == 'protected_inline_signing_lines'
                                                   for r in f.get('source_questions', []))])
        log = [f for f in fields if f['kind'] == 'table_rows']
        self.assertEqual(len(log), 1)
        final = word.fill(raw, fields, {log[0]['key']: [['Synthetic drug A', '5 mg', 'oral']]})
        self.assertEqual([c.text for c in Document(io.BytesIO(final)).tables[0].rows[1].cells],
                         ['Synthetic drug A', '5 mg', 'oral'])

    def test_the_band_reads_its_lines_without_recording_the_paragraph_readers_issues_or_notes(self):
        # Refutation round 11 (corpus:2): the band asks the paragraph readers whether a line below it holds an answer
        # place, and those readers record coverage issues and notes as they read. A line the band claims is never read
        # by them, so its one review is the band's; a line that ends the band keeps e0's issues and notes, in e0's
        # order (the band-free reading's), with only the band's own note added.
        def band_form(before, below, header=False):
            doc = Document()
            if before:
                doc.add_paragraph(before)
            doc.add_paragraph('Synthetic clinic letter')
            if header:
                section = doc.sections[0].header
                section.is_linked_to_previous = False
                section.paragraphs[0]._p.append(line_run(1, fallback=False))
                section.add_paragraph('Signature of patient')
                section.add_paragraph()._p.append(copy.deepcopy(below))
            else:
                doc.add_paragraph()._p.append(line_run(1, fallback=False))
                doc.add_paragraph('Signature of patient')
                doc.add_paragraph()._p.append(copy.deepcopy(below))
            doc.add_paragraph('End of letter.')
            return data(doc)

        def run(text, grey=False):
            colour = '<w:rPr><w:color w:val="808080"/></w:rPr>' if grey else ''
            return etree.fromstring('<w:r %s>%s<w:t xml:space="preserve">%s</w:t></w:r>' % (DECLARED, colour, text))

        for header, place in ((False, BODY + 'b/p[3]'), (True, 'word/header1.xml#hf/p[2]')):
            with self.subTest(claimed='ambiguous grey insertion', header=header):
                raw = band_form(None, run('Enter name or initials here', grey=True), header)
                part, slot = place.split('#')
                found = word.inspect(raw)
                self.assertIn(place, word.signing_band_places(raw))
                self.assertEqual([i['code'] for i in found['coverage_issues'] if (i['part'], i['slot']) == (part, slot)],
                                 ['ambiguous_signing_band'])
        for before, below in (('Enter clinic or ward here', 'Enter name or initials here'),
                              ('[Ward/Clinic]', '[Patient/Representative]'),
                              (None, '______________ [Patient/Representative]')):
            with self.subTest(ends_the_band=below):
                raw = band_form(before, run(below))
                found = word.inspect(raw)
                self.assertNotIn(BODY + ('b/p[4]' if before else 'b/p[3]'), word.signing_band_places(raw))
                free = word.Detector(word.open_package(raw), _inline_bands=False)
                free_notes = free.detect()[1]
                shown = lambda issues: [(i['code'], i['part'], i['slot'], i['detail']) for i in issues]
                self.assertEqual(shown(found['coverage_issues']), shown(free.coverage_issues))
                self.assertEqual([n for n in found['notes'] if not n.startswith('These printed inline lines')], free_notes)
                self.assertTrue(free.coverage_issues)


def caption_over_blank(caption, style=None):
    doc = Document()
    doc.add_paragraph('Clinical summary: ________')
    doc.add_paragraph(caption, style=style) if style else doc.add_paragraph(caption)
    doc.add_paragraph()
    doc.add_paragraph('Plan: ____________')
    return data(doc)


class DetectionCost(unittest.TestCase):
    """Refutation round 10 (cost): the band-free reading - a whole second detection, run by every detector and every
    band-place reading, four times in one automatic import - took a large band form past the parser's CPU limit where
    e0 imported it. It is now read once per original, and only when a band has a line to judge; the detector's text
    reading has exact fast paths. These pin both structurally, never by timing."""

    def test_the_band_free_reading_runs_once_per_original_and_only_when_a_line_follows(self):
        runs = []
        detect = word.Detector.detect

        def counted(detector):
            if not detector._inline_bands:
                runs.append(detector)
            return detect(detector)

        def band_form(after):
            doc = Document()
            doc.add_paragraph('Synthetic consent record')
            doc.add_paragraph()._p.append(line_run(1, fallback=False))
            doc.add_paragraph('Signature of patient')
            for text in after:
                doc.add_paragraph(text)
            return data(doc)

        word._BAND_FREE_READINGS.clear()
        with mock.patch.object(word.Detector, 'detect', counted):
            nothing_after = band_form([])
            word.inspect(nothing_after)
            word.signing_band_places(nothing_after)
            self.assertEqual(runs, [])                 # nothing follows the caption: the band never asks
            line_after = band_form(['_' * 30, 'End of record.'])
            found = word.inspect(line_after)
            band = word.signing_band_places(line_after)
            word.destination_roles(line_after)
            word.fill(line_after, found['fields'], {})
            self.assertEqual(len(runs), 1)             # one reading, shared by every detection of these bytes
            self.assertIn(BODY + 'b/p[3]', band)       # and it still decides the caption's own line
            self.assertNotIn(BODY + 'b/p[4]', band)
            word.signing_band_places(band_form(['_' * 31, 'End of record.']))
            self.assertEqual(len(runs), 2)             # other bytes have their own reading

    def test_only_a_writing_line_below_a_band_asks_the_band_free_reading(self):
        # Refutation round 11 (cost): on large table-heavy forms the speed-ups saved almost nothing, so the one
        # band-free reading per operation still took forms past the CPU limit where e0 imported them. The reading
        # decides which writing lines are the caption's; printed prose, an empty spacing line or a picture alone below
        # the band is no writing line, and ends the band without it.
        runs = []
        detect = word.Detector.detect

        def counted(detector):
            if not detector._inline_bands:
                runs.append(detector)
            return detect(detector)

        def band_form(after, spacer=True, caption='Signature of patient'):
            doc = Document()
            doc.add_paragraph('Synthetic consent record')
            doc.add_paragraph()._p.append(line_run(1, fallback=False))
            doc.add_paragraph(caption)
            if spacer:
                doc.add_paragraph()
            for text in after:
                if text == 'band':
                    doc.add_paragraph()._p.append(line_run(2, fallback=False))
                    doc.add_paragraph('Signature of witness')
                else:
                    doc.add_paragraph(text)
            return data(doc)

        shapes = [(['End of record.'], True, 0), (['', '', 'End of record.'], True, 0), (['band', 'End.'], True, 0),
                  (['_' * 30], False, 1), (['\u2610 Patient unable to sign'], True, 1), (['{{Date signed}}'], False, 1),
                  (['', 'End.'], True, 0)]
        with mock.patch.object(word.Detector, 'detect', counted):
            for after, spacer, expected in shapes:
                with self.subTest(after=after, spacer=spacer):
                    word._BAND_FREE_READINGS.clear()
                    runs.clear()
                    word.inspect(band_form(after, spacer))
                    self.assertEqual(len(runs), expected)
            # an empty line under a narrative caption ('- Signature of patient:') is a writing line
            word._BAND_FREE_READINGS.clear()
            runs.clear()
            word.inspect(band_form(['', '', 'End.'], True, caption='\u2022 Signature of patient:'))
            self.assertEqual(len(runs), 1)

    def test_the_text_fast_path_reads_exactly_what_the_glyph_map_reads(self):
        # Outside the symbol fonts the glyph map changes only box glyphs and private-use box codes, so a text holding
        # none of them maps to itself - every character of the basic plane is checked; in a symbol font an ASCII letter
        # can draw a box ('o' in Wingdings), and the fast path never applies.
        for font in ('', 'arial'):
            plane = ''.join(chr(code) for code in range(0x10000) if not 0xD800 <= code <= 0xDFFF)
            mapped = [ch for ch in plane if word.map_text(ch, font) != word.map_char(ch, font)]
            self.assertEqual(mapped, [], font)
        samples = [''.join(map(chr, range(0x80))), 'Signature of patient: ____', '\u2610 Yes \u2612 No',
                   'o q r * R S T Q x \xa8 \xa1', ''.join(chr(code) for code in (0xF06F, 0xF0FE, 0x2A, 0x6F)), '',
                   'The patient\u2019s \u201cconsent\u201d \u2014 r\xe9sum\xe9 \u540d\u524d \u25cb Yes \u25c9 No \u2611 \u20dd',
                   '\u25a1 Yes \u274f \u2751 \u2752 \u25a2 \u25a3 \uf0a8 \uf078 \uf052 done']
        for font in ('', 'arial', 'wingdings', 'wingdings 2', 'symbol', 'segoe ui symbol'):
            for text in samples:
                with self.subTest(font=font, text=text[:16]):
                    self.assertEqual(word.map_text(text, font), ''.join(word.map_char(ch, font) for ch in text))
        self.assertEqual(word.map_text('o', 'wingdings'), word.BOX)

    def test_one_detection_pass_reads_each_slot_visibility_once_and_a_later_read_is_fresh(self):
        # Refutation round 11 (cost): on printed-table forms a quarter of every detection re-read the same slots'
        # visibility. A detect() pass changes no XML, so within it each slot is read once; outside a pass every read
        # is fresh again (test_source_style_cache pins that half for a slot whose XML changes).
        doc = Document()
        doc.add_paragraph('Synthetic intake record')
        table = doc.add_table(rows=3, cols=2)
        for row, caption in enumerate(('Name', 'Date of birth', 'Allergies')):
            table.cell(row, 0).text = caption
        doc.add_paragraph('Heart rate:\t\tBlood pressure:\t\t')
        doc.add_paragraph('Weight:\t____\tHeight:\t____')
        doc.add_paragraph('Signature of patient: ______________')
        raw = data(doc)
        reads = []
        visible = word.Detector._tab_source_visible

        def counted(detector, slot):
            reads.append(slot)                  # held, so no two counted slots share an id
            return visible(detector, slot)

        def most_reads(run):
            reads.clear()
            detector = word.Detector(word.open_package(raw))
            with mock.patch.object(word.Detector, '_tab_source_visible', counted):
                run(detector)
            return detector, max(collections.Counter(map(id, reads)).values())

        self.assertGreater(most_reads(lambda d: d._detect())[1], 1)   # the control: without a pass, slots are re-read
        detector, most = most_reads(lambda d: d.detect())
        self.assertEqual(most, 1)
        self.assertIsNone(detector._visible)                           # the memo ends with its pass
        slot = reads[0]
        reads.clear()
        with mock.patch.object(word.Detector, '_tab_source_visible', counted):
            detector.tab_source_visible(slot)
            detector.tab_source_visible(slot)
        self.assertEqual(len(reads), 2)                                # so a later read is fresh
        failing = word.Detector(word.open_package(raw))
        with mock.patch.object(word.Detector, '_detect', side_effect=RuntimeError('synthetic')):
            with self.assertRaises(RuntimeError):
                failing.detect()
        self.assertIsNone(failing._visible)                            # and so does a pass that fails


class BelowCaptionReading(unittest.TestCase):
    """The blank below a printed caption keeps the engine's existing reading: ordinary answers stay writable and a
    source action keeps the guard's own review. A wording-based signing-caption hold was withdrawn after three
    refutation rounds (2026-10-05): every reader over- or under-held, and a hold could change the import route."""

    def below(self, raw):
        return [f for f in word.inspect(raw)['fields'] if f['locator']['slot'] == 'b/p[2]']

    def test_ordinary_date_captions_keep_their_writable_blank(self):
        for caption in ('Date of assessment:', 'Visit date:', 'Date of birth:', 'Date of procedure:',
                        'Date signed report received:', 'Date received signed report:',
                        'Date nurse received the signed report:', 'Review date:',
                        # A clinical 'signature', 'signs' or 'initial' is not a signing act (refutation 2026-10-05).
                        'Gene expression signature result date:', 'Date of vital signs:',
                        'Date of onset of signs and symptoms:', 'Troponin (initial) date:',
                        'Date of initial or repeat test:', 'Date of troponin (initial):',
                        'Signed copy given date:', 'Signed report received date:'):
            with self.subTest(caption):
                raw = caption_over_blank(caption)
                fields = word.inspect(raw)['fields']
                below = [f for f in fields if f['locator']['slot'] == 'b/p[2]']
                self.assertEqual([f.get('write_protected', False) for f in below], [False] * len(below))
                if below:
                    final = word.fill(raw, fields, {below[0]['key']: '14/03/2026'})
                    index = word.SlotIndex(word.open_package(final))
                    self.assertIn('14/03/2026', word.paragraph_text(index.by_id[BODY + 'b/p[2]'].el)[0])


    def test_the_source_action_guard_keeps_its_own_reading_and_review(self):
        # The hold runs after guard_source_actions: an action the guard reads keeps its coverage issue, so the
        # import still needs the clinician's exclusion review before generation (refutation 2026-10-05).
        for caption, code in (('Date of next review:', 'ambiguous_source_action'),
                              ('Date of approval:', 'ambiguous_source_action'),
                              ('Date of consent:', 'human_signing_action')):
            with self.subTest(caption):
                found = word.inspect(caption_over_blank(caption))
                below = [f for f in found['fields'] if f['locator']['slot'] == 'b/p[2]']
                self.assertEqual([bool(f.get('write_protected')) for f in below], [True])
                self.assertEqual([i['code'] for i in found['coverage_issues'] if i['slot'] == 'b/p[2]'], [code])


class DrawnBandRoute(unittest.TestCase):
    """A drawn band is v3-only knowledge: it keeps the import on v3 where e0 did (via its old inferred date)."""

    def form(self, *, band=True):
        doc = Document()
        doc.add_paragraph('Synthetic intake')
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Patient name'
        table.cell(1, 0).text = 'Diagnosis'
        doc.add_paragraph('Please circle [yes/no] where the form asks.')
        if band:
            lines = doc.add_paragraph()
            lines._p.append(line_run(1, fallback=False))
            lines.add_run('\t\t')
            lines._p.append(line_run(2, fallback=False))
            doc.add_paragraph('Signature of Patient or Representative\t\t\tDate')
            doc.add_paragraph()
        doc.add_paragraph('End of form.')
        return data(doc)

    def test_band_keeps_the_reviewable_v3_route(self):
        import tempfile
        from pathlib import Path
        from django.test import override_settings
        from .service import inspect_template
        with tempfile.TemporaryDirectory() as tmp, override_settings(QUILLMD_WORD_AUTO_PREFER_V3=True):
            path = Path(tmp) / 'original.docx'
            path.write_bytes(self.form())
            result = inspect_template(path, engine='auto')
        coverage = result['import_manifest']['coverage']
        self.assertEqual(result['import_manifest']['engine'], 'v3')
        self.assertEqual(len(coverage['v3_only_signing_bands']), 2)
        self.assertTrue(coverage['selected_engine_issues'])
        self.assertFalse(coverage.get('advanced_unresolved_areas'))

    def test_a_legacy_writer_on_the_band_spacer_is_held_by_the_band_not_lost(self):
        # Refutation round 4: an underlined-tab spacer is a legacy writing place; the band protects it, so it must be
        # accounted for rather than "lost", or the import fell to a legacy route whose every export refuses.
        import tempfile
        from pathlib import Path
        from django.test import override_settings
        from .service import inspect_template, fill_template
        doc = Document()
        doc.add_paragraph('Synthetic clinic review form')
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Patient name'
        table.cell(1, 0).text = 'Diagnosis'
        lines = doc.add_paragraph()
        lines._p.append(line_run(1, fallback=False))
        lines.add_run('\t\t')
        lines._p.append(line_run(2, fallback=False))
        caption = doc.add_paragraph('Signature of Patient or Representative\t\t\tDate')
        level = OxmlElement('w:outlineLvl')
        level.set(qn('w:val'), '2')  # the outline level makes legacy read the underlined tab below as an answer place
        caption._p.get_or_add_pPr().append(level)
        spacer = doc.add_paragraph()
        for node in etree.fromstring('<w:wrap %s><w:r><w:rPr><w:u w:val="single"/></w:rPr><w:tab/></w:r></w:wrap>'
                                     % DECLARED):
            spacer._p.append(node)
        doc.add_paragraph('End of form.')
        with tempfile.TemporaryDirectory() as tmp, override_settings(QUILLMD_WORD_AUTO_PREFER_V3=True):
            path = Path(tmp) / 'original.docx'
            path.write_bytes(data(doc))
            result = inspect_template(path, engine='auto')
            self.assertEqual(result['import_manifest']['engine'], 'v3')
            self.assertFalse(result['import_manifest']['coverage']['v3_uncovered_legacy_places'])
            answers = {f['key']: 'Synthetic answer' for f in result['fields']
                       if not f.get('write_protected') and f['kind'] not in ('signature', 'initials')}
            self.assertEqual(len(answers), 2)
            fill_template(path, result['fields'], answers, Path(tmp) / 'out.docx')
            self.assertEqual([c.text for c in Document(Path(tmp) / 'out.docx').tables[0].columns[1].cells],
                             ['Synthetic answer', 'Synthetic answer'])

    def test_a_band_repeated_in_footer_copies_is_held_in_every_copy(self):
        # Refutation round 5: identical bands in the first-page and default footers merge into one field; the
        # routing account must still hold every copy's places, or the import fell to a dead legacy route.
        import tempfile
        from pathlib import Path
        from django.test import override_settings
        from .service import inspect_template, fill_template
        doc = Document()
        doc.add_paragraph('Synthetic clinic review form')
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Patient name'
        table.cell(1, 0).text = 'Diagnosis'
        doc.add_paragraph('End of form.')
        section = doc.sections[0]
        section.different_first_page_header_footer = True
        for footer in (section.first_page_footer, section.footer):
            lines = footer.paragraphs[0]
            lines._p.append(line_run(1, fallback=False))
            lines.add_run('\t\t')
            lines._p.append(line_run(2, fallback=False))
            caption = footer.add_paragraph('Signature of Patient or Representative\t\t\tDate')
            level = OxmlElement('w:outlineLvl')
            level.set(qn('w:val'), '2')
            caption._p.get_or_add_pPr().append(level)
            spacer = footer.add_paragraph()
            for node in etree.fromstring('<w:wrap %s><w:r><w:rPr><w:u w:val="single"/></w:rPr><w:tab/></w:r>'
                                         '</w:wrap>' % DECLARED):
                spacer._p.append(node)
        with tempfile.TemporaryDirectory() as tmp, override_settings(QUILLMD_WORD_AUTO_PREFER_V3=True):
            path = Path(tmp) / 'original.docx'
            path.write_bytes(data(doc))
            result = inspect_template(path, engine='auto')
            self.assertEqual(result['import_manifest']['engine'], 'v3')
            self.assertFalse(result['import_manifest']['coverage']['v3_uncovered_legacy_places'])
            answers = {f['key']: 'Synthetic answer' for f in result['fields']
                       if not f.get('write_protected') and f['kind'] not in ('signature', 'initials')}
            fill_template(path, result['fields'], answers, Path(tmp) / 'out.docx')

    def test_a_form_kept_on_legacy_keeps_exporting_with_a_band(self):
        # Refutation round 5: a form that stays on legacy for another reason (numbered underlined lines that v3
        # does not cover) gained a whole-export refusal from the band's protected spacer. Legacy keeps its previous
        # reading of the band's places; the band's protection belongs to the v3 route.
        import tempfile
        from pathlib import Path
        from . import engine
        underlined = '<w:r %s><w:rPr><w:u w:val="single"/></w:rPr><w:tab/></w:r>' % DECLARED

        def report(band):
            doc = Document()
            doc.add_paragraph('Synthetic endoscopy report')
            name = doc.add_paragraph('PATIENT NAME: ')
            name._p.append(etree.fromstring(underlined))
            heading = doc.add_paragraph('FINDINGS')
            level = OxmlElement('w:outlineLvl')
            level.set(qn('w:val'), '1')
            heading._p.get_or_add_pPr().append(level)
            for number in (1, 2, 3):
                item = doc.add_paragraph('%d. ' % number)
                item._p.append(etree.fromstring(underlined))
            if band:
                lines = doc.add_paragraph()
                lines._p.append(line_run(1, fallback=False))
                lines.add_run('\t\t')
                lines._p.append(line_run(2, fallback=False))
                caption = doc.add_paragraph('Signature of Patient or Representative\t\t\tDate')
                level = OxmlElement('w:outlineLvl')
                level.set(qn('w:val'), '2')
                caption._p.get_or_add_pPr().append(level)
                spacer = doc.add_paragraph()
                spacer._p.append(etree.fromstring(underlined))
            doc.add_paragraph('End of report.')
            return data(doc)

        for band in (False, True):
            with self.subTest(band=band), tempfile.TemporaryDirectory() as tmp:
                workspace = Path(tmp)
                (workspace / 'input.docx').write_bytes(report(band))
                result = engine.process(workspace, {'format': 'docx', 'operation': 'inspect', 'engine': 'auto',
                                                    'prefer_v3_parity': True})
                self.assertEqual(result['import_manifest']['engine'], 'legacy')
                key = next(f['key'] for f in result['fields'] if f['label'].upper().startswith('PATIENT NAME'))
                engine.process(workspace, {'format': 'docx', 'operation': 'fill', 'mapping': result['fields'],
                                           'answers': {key: 'Alex Example'}, 'engine': 'legacy'})
                self.assertIn('Alex Example', Document(workspace / 'output.docx').paragraphs[1].text)

    def test_a_band_review_holds_a_legacy_import_only_where_legacy_would_write_into_the_band(self):
        # Refutation round 6: on the legacy route every v3 coverage issue became an advanced unresolved area that
        # held generation, and a legacy import has no review that clears one. A drawn band whose spacer draws a line
        # therefore made a form e0 exported unusable. Round 7: dropping the review outright let legacy write a
        # generated date into the band's spacer under a heading-styled 'Clinician signature / Date' caption, which
        # e0 held. The review now holds a legacy import exactly when a legacy field without a signing role sits in
        # the band; any other v3 issue (here a typed captioned signing band's) still holds it as on e0.
        import tempfile
        from pathlib import Path
        from . import engine
        underlined = '<w:r %s><w:rPr><w:u w:val="single"/></w:rPr><w:tab/></w:r>' % DECLARED
        spacers = {
            'bottom border': '<w:p %s><w:pPr><w:pBdr><w:bottom w:val="single" w:sz="4" w:space="1" w:color="auto"/>'
                             '</w:pBdr></w:pPr></w:p>' % DECLARED,
            'positional tab leader': '<w:p %s><w:r><w:ptab w:relativeTo="margin" w:alignment="right" '
                                     'w:leader="underscore"/></w:r></w:p>' % DECLARED,
            'leader tab stop': '<w:p %s><w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="9000"/></w:tabs>'
                               '</w:pPr><w:r><w:tab/></w:r></w:p>' % DECLARED,
            'underlined tab': '<w:p %s>%s</w:p>' % (DECLARED, underlined.replace(' ' + DECLARED, '')),
            'typed line': '<w:p %s><w:r><w:t>%s</w:t></w:r></w:p>' % (DECLARED, '_' * 30),
        }

        def report(spacer, *, numbered=True, typed=False, caption='Patient signature\t\t\tPrint name', outlined=False,
                   extra=0):
            doc = Document()
            doc.add_paragraph('Synthetic endoscopy report')
            name = doc.add_paragraph('PATIENT NAME: ')
            name._p.append(etree.fromstring(underlined))
            heading = doc.add_paragraph('FINDINGS')
            level = OxmlElement('w:outlineLvl')
            level.set(qn('w:val'), '1')
            heading._p.get_or_add_pPr().append(level)
            for number in (1, 2, 3) if numbered else ():
                item = doc.add_paragraph('%d. ' % number)
                item._p.append(etree.fromstring(underlined))
            lines = doc.add_paragraph()
            lines._p.append(line_run(1, fallback=False))
            lines.add_run('\t\t')
            lines._p.append(line_run(2, fallback=False))
            heading = doc.add_paragraph(caption)
            if outlined:
                outline = OxmlElement('w:outlineLvl')
                outline.set(qn('w:val'), '2')
                heading._p.get_or_add_pPr().append(outline)
            placeholder = doc.add_paragraph()._p
            placeholder.addnext(etree.fromstring(spacers[spacer]))
            placeholder.getparent().remove(placeholder)
            for _ in range(extra):
                doc.add_paragraph('_' * 30)
            if typed:
                doc.add_paragraph('Printed Name of Recorder:')
                band = doc.add_paragraph()
                band.add_run('\t\t').underline = True
                band.add_run('   ')
                band.add_run('________')
                band.add_run('   ')
                band.add_run('\t\t').underline = True
                doc.add_paragraph('Signature of Recorder\tDate\tTime')
            doc.add_paragraph('End of report.')
            return data(doc)

        def imported(raw, preferred):
            with tempfile.TemporaryDirectory() as tmp:
                workspace = Path(tmp)
                (workspace / 'input.docx').write_bytes(raw)
                result = engine.process(workspace, {'format': 'docx', 'operation': 'inspect', 'engine': preferred,
                                                    'prefer_v3_parity': preferred == 'auto'})
                coverage = result['import_manifest']['coverage']
                key = next(f['key'] for f in result['fields'] if f['label'].upper().startswith('PATIENT NAME'))
                engine.process(workspace, {'format': 'docx', 'operation': 'fill', 'mapping': result['fields'],
                                           'answers': {key: 'Alex Example'}, 'engine': 'legacy'})
                return result['import_manifest']['engine'], coverage, Document(workspace / 'output.docx')

        for spacer in ('bottom border', 'positional tab leader', 'leader tab stop', 'typed line'):
            raw = report(spacer)
            self.assertTrue([issue for issue in word.inspect(raw)['coverage_issues']
                             if issue['code'] == 'ambiguous_signing_band'])
            for preferred in ('auto', 'legacy'):
                with self.subTest(spacer=spacer, preferred=preferred):
                    route, coverage, output = imported(raw, preferred)
                    self.assertEqual(route, 'legacy')
                    self.assertFalse(coverage.get('block_generation'))
                    self.assertNotIn('advanced_unresolved_areas', coverage)
                    self.assertEqual(output.paragraphs[1].text, 'PATIENT NAME: Alex Example')
                    self.assertEqual(output.paragraphs[8].text, Document(io.BytesIO(raw)).paragraphs[8].text)
        # A multi-line typed signing area on a legacy-kept form exports as on e0, every line unchanged.
        raw = report('typed line', extra=2)
        for preferred in ('auto', 'legacy'):
            with self.subTest(spacer='typed lines', preferred=preferred):
                route, coverage, output = imported(raw, preferred)
                self.assertEqual(route, 'legacy')
                self.assertFalse(coverage.get('block_generation'))
                self.assertEqual([p.text for p in output.paragraphs[8:11]],
                                 [p.text for p in Document(io.BytesIO(raw)).paragraphs[8:11]])
        # The settings default engine keeps a band form without lost places on legacy as well.
        _, coverage, _ = imported(report('bottom border', numbered=False), 'legacy')
        self.assertFalse(coverage.get('block_generation'))
        # Any other advanced issue still holds the legacy import: the typed band's own review stays, alone.
        for preferred in ('auto', 'legacy'):
            with self.subTest(control='typed band', preferred=preferred):
                route, coverage, _ = imported(report('bottom border', typed=True), preferred)
                self.assertEqual(route, 'legacy')
                self.assertTrue(coverage['block_generation'])
                held = coverage['advanced_unresolved_areas']
                self.assertEqual([(issue['code'], issue['detail'].startswith('Protected signing block')) for issue in held],
                                 [('ambiguous_signing_band', True)])
        # Under a heading-styled caption legacy reads the underlined spacer as its own field. A role-less one (the
        # band's date) would take a generated date, so the band's review holds the import as e0 held it; a signing-role
        # one is held per field by core (never generated), so the import stays usable.
        for caption, holds in (('Clinician signature\t\t\tDate', True), ('Patient signature\t\t\tPrinted name', False)):
            raw = report('underlined tab', caption=caption, outlined=True)
            for preferred in ('auto', 'legacy'):
                with self.subTest(caption=caption, preferred=preferred):
                    with tempfile.TemporaryDirectory() as tmp:
                        workspace = Path(tmp)
                        (workspace / 'input.docx').write_bytes(raw)
                        result = engine.process(workspace, {'format': 'docx', 'operation': 'inspect', 'engine': preferred,
                                                            'prefer_v3_parity': preferred == 'auto'})
                    coverage = result['import_manifest']['coverage']
                    self.assertEqual(result['import_manifest']['engine'], 'legacy')
                    in_band = [field['label'] for field in result['fields'] if field['label'] == caption]
                    self.assertEqual(in_band, [caption])
                    self.assertEqual(bool(word.legacy_role_kind(caption)), not holds)
                    self.assertEqual(bool(coverage.get('block_generation')), holds)
                    self.assertEqual([issue['detail'].startswith(word.DRAWN_BAND_SPACER_REVIEW)
                                      for issue in coverage.get('advanced_unresolved_areas', [])], [True] if holds else [])
        # When the place readers disagree, nothing says where legacy writes: the band's review keeps holding.
        from unittest import mock
        from . import import_selection
        with tempfile.TemporaryDirectory() as tmp, mock.patch.object(
                import_selection, 'legacy_places', side_effect=ValueError('Original place readers disagree.')):
            workspace = Path(tmp)
            (workspace / 'input.docx').write_bytes(report('bottom border'))
            result = engine.process(workspace, {'format': 'docx', 'operation': 'inspect', 'engine': 'auto',
                                                'prefer_v3_parity': True})
        coverage = result['import_manifest']['coverage']
        self.assertEqual((result['import_manifest']['engine'], coverage.get('comparison_unresolved'),
                          coverage['block_generation']), ('legacy', True, True))
        self.assertTrue(all(issue['detail'].startswith(word.DRAWN_BAND_SPACER_REVIEW)
                            for issue in coverage['advanced_unresolved_areas']))

    def test_a_header_band_holds_a_legacy_import_whatever_the_earlier_fields_span(self):
        # Refutation round 8: legacy places are one per ANCHOR, so a multi-paragraph tab-form field earlier in the
        # form shifted the field/place pairing and a role-less legacy date field in a header band went unseen; the
        # import was admitted and the date written into both header copies. Places are now read per field.
        import tempfile
        from pathlib import Path
        from . import engine
        underlined = '<w:r %s><w:rPr><w:u w:val="single"/></w:rPr><w:tab/></w:r>' % DECLARED
        caption = 'Signature of Patient or Representative\t\t\tDate'

        def report(anchors):
            doc = Document()
            doc.add_paragraph('Synthetic endoscopy report')
            doc.add_paragraph('PATIENT NAME: ')._p.append(etree.fromstring(underlined))
            heading = doc.add_paragraph('COMMENTS')
            outline = OxmlElement('w:outlineLvl')
            outline.set(qn('w:val'), '1')
            heading._p.get_or_add_pPr().append(outline)
            for _ in range(anchors):
                doc.add_paragraph()._p.append(etree.fromstring(underlined))
            doc.add_paragraph('End of report.')
            section = doc.sections[0]
            section.different_first_page_header_footer = True
            for header in (section.header, section.first_page_header):
                header.is_linked_to_previous = False
                lines = header.paragraphs[0]
                lines._p.append(line_run(1, fallback=False))
                lines.add_run('\t\t')
                lines._p.append(line_run(2, fallback=False))
                band = header.add_paragraph(caption)
                level = OxmlElement('w:outlineLvl')
                level.set(qn('w:val'), '2')
                band._p.get_or_add_pPr().append(level)
                header.add_paragraph()._p.append(etree.fromstring(underlined))
            return data(doc)

        for anchors in (1, 2, 3):
            raw = report(anchors)
            for preferred in ('auto', 'legacy'):
                with self.subTest(anchors=anchors, preferred=preferred), tempfile.TemporaryDirectory() as tmp:
                    workspace = Path(tmp)
                    (workspace / 'input.docx').write_bytes(raw)
                    result = engine.process(workspace, {'format': 'docx', 'operation': 'inspect', 'engine': preferred,
                                                        'prefer_v3_parity': preferred == 'auto'})
                    coverage = result['import_manifest']['coverage']
                    route = result['import_manifest']['engine']
                    self.assertTrue(coverage['block_generation'])
                    if route == 'v3':
                        # Every legacy place is kept, so the band selects v3: both spacers are protected there and
                        # their reviews are the clearable selected-engine issues.
                        self.assertEqual((anchors, preferred), (1, 'auto'))
                        issues = coverage['selected_engine_issues']
                        self.assertFalse([f for f in result['fields'] if f['locator']['part'] != BODY[:-1]
                                          and not f.get('write_protected')])
                    else:
                        self.assertIn(caption, [field['label'] for field in result['fields']])
                        issues = coverage['advanced_unresolved_areas']
                    held = [issue['part'] for issue in issues if issue['detail'].startswith(word.DRAWN_BAND_SPACER_REVIEW)]
                    self.assertEqual(sorted(held), ['word/header1.xml', 'word/header2.xml'])
        # Without per-field places nothing says where legacy writes: the reviews keep holding.
        from unittest import mock
        from . import import_selection
        raw = report(2)
        with tempfile.TemporaryDirectory() as tmp, mock.patch.object(
                import_selection, 'legacy_field_places', side_effect=ValueError('Legacy places do not match the fields.')):
            workspace = Path(tmp)
            (workspace / 'input.docx').write_bytes(raw)
            result = engine.process(workspace, {'format': 'docx', 'operation': 'inspect', 'engine': 'legacy'})
        coverage = result['import_manifest']['coverage']
        self.assertTrue(coverage['block_generation'])
        self.assertEqual(len([issue for issue in coverage['advanced_unresolved_areas']
                              if issue['detail'].startswith(word.DRAWN_BAND_SPACER_REVIEW)]), 2)

    def test_legacy_places_are_read_per_field(self):
        from . import import_selection
        fields = [{'locator': {'paragraph': 1}}, {'locator': {'tab_form': {'anchors': [{}, {}, {}]}}},
                  {'locator': {'cell': 0}}]
        places = [('a', 0, 1), ('b', 0, 1), ('c', 0, 1), ('d', 0, 1), ('e', None, None)]
        self.assertEqual(import_selection.legacy_field_places(fields, places),
                         [[('a', 0, 1)], [('b', 0, 1), ('c', 0, 1), ('d', 0, 1)], [('e', None, None)]])
        with self.assertRaises(ValueError):
            import_selection.legacy_field_places(fields, places[:-1])

    def test_band_parts_never_compete_with_the_forms_own_questions(self):
        # Refutation round 4: a bare 'Date' band part shared its question key with the form's own 'Date' and held
        # that ordinary answer in field-relations generation.
        from core.generation import question_key
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Date'
        lines = doc.add_paragraph()
        lines._p.append(line_run(1, fallback=False))
        lines.add_run('\t\t')
        lines._p.append(line_run(2, fallback=False))
        doc.add_paragraph('Signature of Patient\t\tDate')
        doc.add_paragraph()
        fields = word.inspect(data(doc))['fields']
        band = [f for f in fields if any(r.get('relation') == 'protected_inline_signing_lines'
                                         for r in f.get('source_questions', []))]
        self.assertEqual([(f['kind'], f['label']) for f in band],
                         [('signature', 'Signature of Patient'), ('date', 'Signature of Patient Date')])
        ordinary = [f for f in fields if f not in band]
        keys = {question_key(f['label']) for f in band}
        self.assertTrue(ordinary)
        self.assertFalse([f for f in ordinary if question_key(f['label']) in keys])

    def test_without_a_band_the_route_rule_is_unchanged(self):
        from . import import_selection
        self.assertNotIn('protected_inline_signing_lines', json.dumps(word.inspect(self.form(band=False))['fields']))
        legacy = {'fields': [], 'warnings': []}
        self.assertEqual(import_selection.choose(legacy, {'fields': []}, [], None)[0], 'legacy')


class WithdrawnRelaxationsKeepTheirReading(unittest.TestCase):
    """Independent refutation (2026-10-05) showed three paragraph relaxations wrote into section titles, deleted
    line breaks or blocked generation for ordinary forms; they were withdrawn. These are its counterexamples."""

    def test_heading_section_titles_with_field_words_stay_fixed(self):
        for text, style in (('Patient Details:\t\t\t\tNext of Kin Details:', 'Heading 2'),
                            ('Clinical Notes:\t\t\t\tNursing Notes:', 'Heading 2'),
                            ('Patient name:\t\t\t\t\tPatient number:', 'Heading 2')):
            with self.subTest(text):
                doc = Document()
                doc.add_paragraph(text, style=style)
                doc.add_paragraph('Name: ____________________\t\tName: ____________________')
                fields = word.inspect(data(doc))['fields']
                self.assertFalse([f for f in fields if f['locator']['slot'] == 'b/p[0]'])

    def test_a_request_wrapped_by_a_line_break_is_never_written_across_the_break(self):
        doc = Document()
        doc.add_paragraph('Notice to the patient.')
        paragraph = doc.add_paragraph()
        paragraph._p.append(etree.fromstring('<w:r %s><w:t>Enter</w:t><w:br/><w:t xml:space="preserve">the referring '
                                             'clinician here</w:t></w:r>' % DECLARED))
        doc.add_paragraph('Fixed following text.')
        found = word.inspect(data(doc))
        self.assertFalse([f for f in found['fields'] if f['locator']['slot'] == 'b/p[1]'])

    def test_optional_caption_over_irregular_blanks_keeps_its_own_answer_and_generation(self):
        doc = Document()
        doc.add_paragraph('Patient number: ________')
        doc.add_paragraph('Allergies (optional):')
        for _ in range(3):
            doc.add_paragraph().add_run('\t')
        doc.add_paragraph('Visit date: ________')
        raw = data(doc)
        found = word.inspect(raw)
        self.assertFalse([i for i in found['coverage_issues'] if i['code'] == 'unresolved_narrative_band'])
        caption = [f for f in found['fields'] if f['label'].startswith('Allergies')]
        self.assertEqual(len(caption), 1)
        word.fill(raw, found['fields'], {caption[0]['key']: 'Synthetic answer text'})


if __name__ == '__main__':
    unittest.main()
