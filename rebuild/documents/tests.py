"""Synthetic document tests; no patient documents or provider calls."""

from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, TextStringObject
from reportlab.pdfgen import canvas

from .service import DocumentError, fill_template, inspect_template


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def word(self, name="template.docx"):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Clinical question: ").bold = True
        paragraph.add_run("{{Symp").italic = True
        paragraph.add_run("toms}}")
        paragraph.add_run(" — original suffix").underline = True
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Examination findings"
        header = document.sections[0].header.paragraphs[0]
        header.add_run("Synthetic practice · {{Visit date}}")
        path = self.directory / name
        document.save(path)
        return path

    def pdf(self, name="template.pdf", width=300, height=30, flags=""):
        path = self.directory / name
        document = canvas.Canvas(str(path))
        document.drawString(20, 780, "Clinical question: Symptoms")
        document.acroForm.textfield(name="Symptoms", tooltip="Symptoms", x=20, y=730,
                                   width=width, height=height, fontName="Helvetica", fontSize=10,
                                   maxlen=1000, fieldFlags=flags)
        document.showPage()
        document.save()
        return path

    def assert_error(self, code, operation, *args):
        with self.assertRaises(DocumentError) as caught:
            operation(*args)
        self.assertEqual(caught.exception.code, code, str(caught.exception))

    def test_docx_roundtrip_retains_questions_styles_and_original(self):
        original = self.word()
        before = original.read_bytes()
        inspection = inspect_template(original)
        self.assertEqual(inspection["sandbox"], "landlock+seccomp")
        self.assertEqual(len(inspection["fields"]), 3)
        answers = {field["key"]: "Supported synthetic answer αβγ" for field in inspection["fields"]}
        output = self.directory / "filled.docx"
        result = fill_template(original, inspection["fields"], answers, output)
        self.assertEqual(result["filled"], 3)
        self.assertEqual(before, original.read_bytes())
        filled = Document(output)
        self.assertEqual(filled.paragraphs[0].text, "Clinical question: Supported synthetic answer αβγ — original suffix")
        self.assertTrue(filled.paragraphs[0].runs[0].bold)
        self.assertTrue(filled.paragraphs[0].runs[1].italic)
        self.assertTrue(filled.paragraphs[0].runs[-1].underline)
        self.assertEqual(filled.tables[0].cell(0, 0).text, "Examination findings")
        self.assertEqual(filled.tables[0].cell(0, 1).text, "Supported synthetic answer αβγ")
        self.assertIn("Supported synthetic answer αβγ", filled.sections[0].header.paragraphs[0].text)
        self.assertEqual(result["sha256"], hashlib.sha256(before).hexdigest())

    def test_long_docx_answers_are_not_truncated(self):
        original = self.word()
        fields = inspect_template(original)["fields"]
        answer = "Synthetic long answer. " * 1000
        output = self.directory / "long.docx"
        fill_template(original, fields, {field["key"]: answer for field in fields}, output)
        self.assertEqual(Document(output).tables[0].cell(0, 1).text, answer)

    def test_two_placeholders_same_run_and_repeated_label(self):
        original = self.directory / "repeated.docx"
        doc = Document()
        doc.add_paragraph("A: {{Name}} B: {{Name}}!")
        doc.save(original)
        fields = inspect_template(original)["fields"]
        output = self.directory / "result.docx"
        fill_template(original, fields, {fields[0]["key"]: "First", fields[1]["key"]: "Second"}, output)
        self.assertEqual(Document(output).paragraphs[0].text, "A: First B: Second!")

    def test_unknown_or_changed_mapping_and_existing_output_are_refused(self):
        original = self.word()
        fields = inspect_template(original)["fields"]
        output = self.directory / "result.docx"
        bad = deepcopy(fields)
        bad[0]["locator"]["cell"] = 999
        self.assert_error("mapping_mismatch", fill_template, original, bad, {}, output)
        self.assert_error("invalid_mapping", fill_template, original, fields, {"unmapped": "hidden"}, output)
        self.assert_error("immutable_original", fill_template, original, fields, {}, original)
        self.assertFalse(output.exists())

    def test_external_docx_relationship_is_blocked(self):
        original = self.word()
        doc = Document(original)
        doc.part.relate_to("https://invalid.example/leak", RT.HYPERLINK, is_external=True)
        doc.save(original)
        self.assert_error("unsafe_upload", inspect_template, original)

    def test_fixed_height_docx_blocks_hidden_answer(self):
        original = self.word()
        doc = Document(original)
        doc.tables[0].rows[0].height = Inches(0.1)
        doc.tables[0].rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        doc.save(original)
        self.assert_error("overflow_risk", inspect_template, original)

    def test_hidden_run_and_partially_invalid_placeholders_blocked(self):
        original = self.directory / "hidden.docx"
        document = Document()
        document.add_paragraph().add_run("{{Symptoms}}").font.hidden = True
        document.save(original)
        self.assert_error("overflow_risk", inspect_template, original)
        document = Document()
        document.add_paragraph("{{Symptoms}} Also {{malformed%}}")
        document.save(original)
        self.assert_error("invalid_placeholder", inspect_template, original)

    def test_zip_traversal_duplicate_bomb_and_macros_blocked(self):
        base = self.word()
        for name, payload in [("../secret", b"x"), ("word/vbaProject.bin", b"macro"), ("huge.bin", b"x" * 2_000_000), ("word/document.xml", b"duplicate")]:
            with self.subTest(name=name):
                path = self.directory / "bad.docx"
                path.write_bytes(base.read_bytes())
                with ZipFile(path, "a", compression=ZIP_DEFLATED) as package:
                    package.writestr(name, payload)
                self.assert_error("unsafe_upload", inspect_template, path)

    def test_xml_entities_blocked(self):
        original = self.word()
        bad = self.directory / "entity.docx"
        with ZipFile(original) as incoming, ZipFile(bad, "w") as outgoing:
            for name in incoming.namelist():
                value = incoming.read(name)
                if name == "word/document.xml":
                    value = b'<!DOCTYPE foo [<!ENTITY x SYSTEM "file:///etc/passwd">]>' + value
                outgoing.writestr(name, value)
        self.assert_error("unsafe_upload", inspect_template, bad)

    def test_pdf_roundtrip_and_original_bytes(self):
        original = self.pdf()
        before = original.read_bytes()
        fields = inspect_template(original)["fields"]
        output = self.directory / "filled.pdf"
        fill_template(original, fields, {"Symptoms": "Synthetic symptom only"}, output)
        self.assertEqual(PdfReader(output).get_form_text_fields()["Symptoms"], "Synthetic symptom only")
        self.assertEqual(original.read_bytes(), before)
        self.assertIn("Clinical question: Symptoms", PdfReader(output).pages[0].extract_text())
        annotation = PdfReader(output).pages[0]["/Annots"][0].get_object()
        self.assertIn("/AP", annotation)

    def test_pdf_overflow_and_unicode_refused_without_output(self):
        original = self.pdf(width=100)
        fields = inspect_template(original)["fields"]
        output = self.directory / "filled.pdf"
        self.assert_error("answer_overflow", fill_template, original, fields, {"Symptoms": "W" * 30}, output)
        self.assert_error("unsupported_characters", fill_template, original, fields, {"Symptoms": "αβγ"}, output)
        self.assertFalse(output.exists())

    def test_pdf_actions_and_encryption_refused(self):
        original = self.pdf()
        for mode in ("javascript", "link", "encrypted"):
            with self.subTest(mode=mode):
                writer = PdfWriter(clone_from=original)
                if mode == "javascript":
                    writer.add_js("app.alert('synthetic');")
                elif mode == "link":
                    annotation = writer.pages[0]["/Annots"][0].get_object()
                    annotation[NameObject("/A")] = DictionaryObject({NameObject("/S"): NameObject("/URI"), NameObject("/URI"): TextStringObject("https://invalid.example")})
                else:
                    writer.encrypt("synthetic")
                bad = self.directory / "bad.pdf"
                writer.write(bad)
                self.assert_error("unsafe_upload", inspect_template, bad)

    def test_flat_pdf_and_non_document_blocked(self):
        flat = self.directory / "flat.pdf"
        writer = PdfWriter()
        writer.add_blank_page(400, 600)
        writer.write(flat)
        self.assert_error("flat_pdf", inspect_template, flat)
        text = self.directory / "fake.docx"
        text.write_text("Not a document")
        self.assert_error("unsupported_format", inspect_template, text)

    def test_sandbox_denies_network_secret_read_and_process_creation(self):
        # Run the same enforcement in a disposable child; test real kernel behavior.
        secret = self.directory / "secret.txt"
        secret.write_text("synthetic secret")
        workspace = self.directory / "sandbox"
        workspace.mkdir()
        module = Path(__file__).with_name("sandbox.py").resolve()
        program = """
import importlib.util,json,os,socket,sys
spec=importlib.util.spec_from_file_location('sandbox',sys.argv[1])
module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
module.isolate(sys.argv[2])
result=[]
for operation in (lambda:open(sys.argv[3]).read(),lambda:socket.socket(),lambda:os.fork()):
    try: operation();result.append(False)
    except PermissionError: result.append(True)
print(json.dumps(result))
"""
        result = subprocess.run([sys.executable, "-I", "-B", "-c", program, str(module), str(workspace), str(secret)], capture_output=True, timeout=10)
        self.assertEqual(result.returncode, 0, result.stderr.decode())
        self.assertEqual(json.loads(result.stdout), [True, True, True])


if __name__ == "__main__":
    unittest.main()
