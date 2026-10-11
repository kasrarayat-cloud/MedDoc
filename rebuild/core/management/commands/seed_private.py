import json,secrets
from pathlib import Path
from django.conf import settings
from django.core.management.base import BaseCommand,CommandError
from django.utils import timezone
from django_otp.plugins.otp_totp.models import TOTPDevice
from core.models import User,Organization,Department,Membership,Patient,Template,TemplateVersion,Encounter
from core.billing import welcome_credit
from core.security import private_path,audit
from documents import inspect_template
class Command(BaseCommand):
    help='Create synthetic private test accounts. Credentials are written to a mode-0600 local file, never stdout.'
    def handle(self,*args,**options):
        if not settings.SYNTHETIC_ONLY:raise CommandError('Synthetic-only environment required.')
        dest=settings.BASE_DIR/'.runtime'/'private-access.json'
        if dest.exists():self.stdout.write('Private credentials already exist; refusing to replace.');return
        if settings.AI_MODE != 'fixture':raise CommandError('This private CI seed requires fixture AI mode; it never sends paid requests.')
        org=Organization.objects.create(name='Synthetic North Clinic')
        other=Organization.objects.create(name='Synthetic South Clinic')
        xray=Department.objects.create(organization=org,name='X-ray')
        ultrasound=Department.objects.create(organization=org,name='Ultrasound')
        otherdept=Department.objects.create(organization=other,name='X-ray')
        accounts={}
        for name,dep,role,staff in [('doctor',xray,'owner',False),('xray_employee',xray,'employee',False),('ultrasound',ultrasound,'clinician',False),('other_org',otherdept,'owner',False),('admin',None,None,True)]:
            password=secrets.token_urlsafe(24)
            email=f'{name}@synthetic.quillmd.invalid'
            u=User.objects.create_user(username=email,email=email,password=password,first_name='Synthetic',last_name=name.replace('_',' ').title(),verified_at=timezone.now(),is_staff=staff)
            accounts[name]={'email':email,'password':password}
            if dep:Membership.objects.create(user=u,department=dep,role=role)
            if staff:
                d=TOTPDevice.objects.create(user=u,name='Private staging authenticator',confirmed=True)
                accounts[name]['totp_key_hex']=d.key
                accounts[name]['totp_uri']=d.config_url
            welcome_credit(u)
            audit(u,'synthetic.account_seeded',u.pk)
        doctor=User.objects.get(email=accounts['doctor']['email'])
        # Owner has explicit separate membership, never implicit organization-wide patient access.
        for dep,username in [(xray,'doctor'),(ultrasound,'ultrasound'),(otherdept,'other_org')]:
            u=User.objects.get(email=accounts[username]['email'])
            Patient.objects.create(department=dep,display_name='Synthetic Example '+dep.name,reference='SYNTH-001',created_by=u)
        from docx import Document
        relative='originals/'+secrets.token_hex(24)+'.docx';path=private_path(relative);path.parent.mkdir(parents=True,exist_ok=True,mode=0o700)
        doc=Document();doc.add_heading('Synthetic clinical assessment',0);doc.add_paragraph('Private test fixture · no patient data')
        table=doc.add_table(rows=0,cols=2);table.style='Table Grid'
        for label in ['History','Findings','Plan']:
            cells=table.add_row().cells;cells[0].text=label;cells[1].text=''
        doc.save(path);path.chmod(0o600)
        info=inspect_template(path)
        template=Template.objects.create(department=xray,name='Clinical assessment',format=info['format'],original_name='synthetic-assessment.docx',original_path=relative,sha256=info['sha256'],warnings=info['warnings'],created_by=doctor)
        # CI setup uses the supported approval path. It is explicitly simulated,
        # and never serves as evidence of automatic extraction quality.
        from django.test import override_settings
        from core import preparation
        with override_settings(WORD_PREPARATION_ENABLED=True):
            prepared=preparation.create_preparation(template,doctor)
            if not preparation.run_one(preparation_pk=prepared.pk):raise CommandError('Fixture preparation did not run.')
            prepared.refresh_from_db()
            if prepared.state!='review':raise CommandError('Fixture preparation needs review: '+prepared.error_code)
            version=preparation.approve(prepared.pk,doctor,prepared.revision,prepared.candidate_sha256,
                [change['anchor_id'] for change in prepared.changes],True,True,
                [field['key'] for field in prepared.inventory['fields']])
        Encounter.objects.create(department=xray,patient=Patient.objects.get(department=xray),template_version=version,created_by=doctor,service_date=timezone.localdate(),notes='History: Synthetic patient reports left wrist discomfort after a fall.\nFindings: No deformity documented.\nPlan: Follow-up assessment documented.')
        dest.parent.mkdir(exist_ok=True,mode=0o700);dest.write_text(json.dumps(accounts,indent=2));dest.chmod(0o600)
        self.stdout.write(f'Synthetic accounts and blank template ready. Private credentials: {dest}')
