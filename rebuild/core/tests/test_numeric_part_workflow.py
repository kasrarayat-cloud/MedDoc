"""Manual numeric parts through original upload, immutable mapping review and actual export."""
import io

from django.test import TestCase, override_settings
from docx import Document

from core.models import Encounter, MappingConfirmation, TemplateVersion
from core.security import private_path
from documents.test_numeric_part_recovery import source
from .test_word_v3_workflow import V3, V3Fixture


@override_settings(**V3)
class NumericPartWorkflow(V3Fixture, TestCase):
    def setUp(self):
        super().setUp()
        self.numeric_original = source()
        self.template, self.initial = self.uploaded(data=self.numeric_original)
        self.slot = 'word/document.xml#b/p[2]'
        self.whole = next(field for field in self.initial.fields if field['locator']['slot'] == 'b/p[2]')

    def post(self, payload, version=None):
        version = version or self.latest(self.template)
        return self.client.post(f'/templates/{self.template.pk}/', {'mapping_version': version.pk, **payload})

    def add(self, number, label):
        return self.post({'action': 'v3_add', 'slot': self.slot, 'kind': 'text', 'label': label, 'numeric_part': number})

    def remove_whole(self):
        response = self.post({'action': 'v3_save', 'remove_' + self.whole['key']: 'on'})
        self.assertEqual(response.status_code, 302)

    def test_original_upload_remove_whole_add_each_part_review_and_download(self):
        blocked = self.add(1, 'Supine systolic')
        self.assertEqual(blocked.status_code, 400)
        self.assertIn('Remove the whole-band field first', blocked.context['v3_error'])
        self.assertEqual(self.template.versions.count(), 1)
        self.remove_whole()
        self.assertEqual(self.add(1, 'Supine systolic').status_code, 302)
        first = self.latest(self.template)
        page = self.client.get(f'/templates/{self.template.pk}/')
        self.assertTrue(any(row['slot'] == self.slot for row in page.context['places']))
        self.assertContains(page, 'numeric part 1'); self.assertContains(page, 'up to 3 digits')
        self.assertEqual(self.add(2, 'Supine diastolic').status_code, 302)
        final = self.latest(self.template)
        self.assertNotEqual(final.pk, first.pk)
        parts = [field for field in final.fields if field.get('locator', {}).get('target', {}).get('shape') == 'numeric_part']
        self.assertEqual(len(parts), 2); self.assertTrue(all(field['origin'] == 'manual' for field in parts))
        self.assertEqual(self.initial.fields, TemplateVersion.objects.get(pk=self.initial.pk).fields)
        duplicate = self.add(1, 'Repeated systolic')
        self.assertEqual(duplicate.status_code, 400); self.assertEqual(self.latest(self.template).pk, final.pk)
        self.assertFalse(any(row['slot'] == self.slot for row in self.client.get(f'/templates/{self.template.pk}/').context['places']))
        self.confirm(self.template, final)
        self.assertEqual(self.new_encounter(final, notes='Synthetic typed review.').status_code, 302)
        encounter = Encounter.objects.get()
        values = {'Supine systolic': '120', 'Supine diastolic': '80'}
        payload = {'revision': encounter.revision}
        for field in final.fields:
            if field['label'] in values:
                payload['answer_' + field['key']] = values[field['label']]
                payload['reviewed_' + field['key']] = 'on'
            else: payload['blank_' + field['key']] = 'on'
        response = self.client.post(f'/encounters/{encounter.pk}/save/', payload)
        self.assertEqual(response.status_code, 302)
        encounter.refresh_from_db()
        response = self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'}, follow=True)
        encounter.refresh_from_db()
        self.assertEqual(encounter.status, 'finalized', [str(m) for m in response.context['messages']])
        response = self.client.get(f'/encounters/{encounter.pk}/download/')
        output = b''.join(response.streaming_content)
        self.assertEqual(Document(io.BytesIO(output)).paragraphs[2].text, '120/80')
        self.assertEqual(Document(io.BytesIO(output)).paragraphs[3].text, '\tSystolic\tdiastolic')
        self.assertEqual(private_path(self.template.original_path).read_bytes(), self.numeric_original)
        self.assertEqual(MappingConfirmation.objects.get(version=final).confirmed_by, self.user)

    def test_invalid_part_kind_and_stale_requests_do_not_create_versions(self):
        self.remove_whole(); base = self.latest(self.template)
        for number, kind in ((0, 'text'), (3, 'text'), ('bad', 'text'), (1, 'long_text')):
            response = self.post({'action': 'v3_add', 'slot': self.slot, 'kind': kind, 'label': 'Reviewed number', 'numeric_part': number})
            self.assertEqual(response.status_code, 400); self.assertEqual(self.latest(self.template).pk, base.pk)
        self.assertEqual(self.add(1, 'Systolic').status_code, 302)
        current = self.latest(self.template)
        numeric = next(field for field in current.fields if field.get('origin') == 'manual')
        changed = self.post({'action': 'v3_save', 'kind_' + numeric['key']: 'long_text'})
        self.assertEqual(changed.status_code, 400)
        self.assertEqual(self.latest(self.template).pk, current.pk)
        stale = self.post({'action': 'v3_add', 'slot': self.slot, 'kind': 'text', 'label': 'Diastolic', 'numeric_part': 2}, version=base)
        self.assertEqual(stale.status_code, 409)
        self.assertEqual(self.template.versions.count(), 3)
        self.assertEqual(private_path(self.template.original_path).read_bytes(), self.numeric_original)
