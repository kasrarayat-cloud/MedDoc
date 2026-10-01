"""D10's charge-then-refuse (V1 B-W1, 'D10 charged before a refusal at finalization').

A contract v3 form whose tab-column answers need the isolated layout verifier can export NO answer while the verifier
is absent: `word_v3.fill` refuses every non-empty answer set when the original has a tab-column layout candidate, even
one answer in an unrelated text field (documents/word_v3.py fill: `layout_nodes and retained_layout is None and any
answer`), and `layout_transport.verify_document` then refuses with 'layout_unavailable'. The upload already knows: the
fresh inspection reports `layout_pending` and the verifier's refusal. So generation is refused before the quote is used,
before any hold and before any provider call, and the encounter page says why. A failure AFTER a real provider call is
not this gate's business and keeps its true cost.

Synthetic only: a python-docx form with one tab-column row (the D10 shape) and two ordinary text cells. The provider
client is patched; its call count is the evidence that nothing was sent.
"""
import json
from types import SimpleNamespace
from unittest.mock import patch

from django.test import TestCase, override_settings

from core import v3_mapping
from core.models import Encounter, Job, LedgerEntry, ProviderReceipt, Reservation, Usage, Wallet
from documents import layout_transport
from .test_v3_ordinary_notes import LIVE
from .test_word_v3_workflow import V3Fixture

NO_WITNESS = dict(QUILLMD_LAYOUT_SOCKET='', QUILLMD_LAYOUT_SERVICE_UID='')
NOTES = 'Reason for visit: Right knee pain after a fall on the stairs.\n\nPlan: Physiotherapy referral and review in four weeks.'


def mixed_form():
    """D10's shape in miniature: one tab-column row the verifier must measure, two ordinary answer cells."""
    from documents.test_question_boxes import TabColumnOwnership   # imported here: the runner would collect its tests
    doc = TabColumnOwnership().form()
    doc.add_paragraph('Visit details')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Reason for visit'
    table.cell(1, 0).text = 'Plan'
    return _bytes(doc)


def _bytes(doc):
    import io
    out = io.BytesIO(); doc.save(out)
    return out.getvalue()


def ordinary_form():
    from docx import Document
    doc = Document()
    doc.add_paragraph('Visit details')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Reason for visit'
    table.cell(1, 0).text = 'Plan'
    return _bytes(doc)


def provider_reply(**wire):
    """A completed provider-shaped reply selecting, for each question, the source units that name its label."""
    payload = json.loads(wire['messages'][0]['content'])
    units = payload.get('source_units') or []
    body = {'fields': [{'key': q['key'], 'unit_ids': [u['id'] for u in units if q['label'].lower() in u['quote'].lower()],
                        'conflict': False} for q in payload.get('questions') or []]}
    return SimpleNamespace(id='msg_synthetic_gate', _request_id='req_synthetic_gate', stop_reason='end_turn',
                           usage=SimpleNamespace(input_tokens=1000, output_tokens=100, cache_read_input_tokens=0,
                                                 cache_creation_input_tokens=0),
                           content=[SimpleNamespace(type='text', text=json.dumps(body))],
                           model_dump=lambda **kwargs: {'id': 'msg_synthetic_gate'})


def money(user):
    return [(w.kind, w.balance, w.reserved) for w in Wallet.objects.filter(user=user).order_by('kind')]


@override_settings(**dict(LIVE, **NO_WITNESS))
class LayoutGenerationGate(V3Fixture, TestCase):

    def ready(self, original=None, name='SYNTHETIC tab-column visit'):
        template, version = self.uploaded(data=original or mixed_form(), name=name)
        version = self.confirm(template, version)
        self.assertEqual(self.new_encounter(version, notes=NOTES).status_code, 302)
        return template, version, Encounter.objects.latest('created_at')

    def charge_then_refuse(self, encounter, provider):
        """Only reached when a job was queued: complete the old path and report what it cost, as the failure."""
        from core.jobs import run_one
        held = list(Reservation.objects.filter(job__encounter=encounter).values_list('state', 'amount'))
        provider.return_value.messages.create.side_effect = provider_reply
        run_one()
        encounter.refresh_from_db()
        payload = {'revision': encounter.revision}
        for field in encounter.template_version.fields:
            payload['answer_' + field['key']] = (encounter.answers.get(field['key']) or {}).get('text', '')
            payload['reviewed_' + field['key']] = 'on'
        self.client.post(f'/encounters/{encounter.pk}/save/', payload)
        encounter.refresh_from_db()
        final = self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'},
                                 follow=True)
        encounter.refresh_from_db()
        return ('generation was queued: hold %s; provider calls %d; settled %s; charged %s; then finalization: %s -> %s'
                % (held, provider.return_value.messages.create.call_count,
                   list(Reservation.objects.filter(job__encounter=encounter).values_list('state', 'amount')),
                   list(Usage.objects.filter(job__encounter=encounter).values_list('charged_cost', flat=True)),
                   encounter.status, [str(m) for m in final.context['messages']]))

    def test_without_the_verifier_generation_is_refused_before_any_hold_or_provider_call(self):
        template, version, encounter = self.ready()
        self.assertEqual([f['label'] for f in version.fields], ['Reason for visit', 'Plan'])
        before = money(self.user)
        with patch('anthropic.Anthropic') as provider:
            page = self.client.get(f'/encounters/{encounter.pk}/')
            response = self.client.post(f'/encounters/{encounter.pk}/generate/', {'usage_quote': page.context['usage_quote']},
                                        follow=True)
            if Job.objects.filter(encounter=encounter).exists():
                self.fail(self.charge_then_refuse(encounter, provider))
            self.assertFalse(provider.called)
            self.assertFalse(provider.return_value.messages.create.called)
        self.assertEqual([str(m) for m in response.context['messages']], [v3_mapping.GENERATION_LAYOUT_REFUSAL])
        self.assertFalse(Reservation.objects.exists())
        self.assertFalse(Usage.objects.exists())
        self.assertFalse(ProviderReceipt.objects.exists())
        self.assertEqual(money(self.user), before)
        self.assertFalse(LedgerEntry.objects.filter(job__isnull=False).exists())
        encounter.refresh_from_db()
        self.assertEqual((encounter.status, encounter.answers), ('draft', {}))

    def test_the_explanation_is_on_the_mapping_review_and_the_encounter_page_before_generation(self):
        template, version, encounter = self.ready()
        self.assertIn(layout_transport.LAYOUT_UNAVAILABLE, template.warnings)
        self.assertContains(self.client.get(f'/templates/{template.pk}/'), layout_transport.LAYOUT_UNAVAILABLE)
        page = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertEqual(page.context['export_refusal'], layout_transport.LAYOUT_UNAVAILABLE)
        self.assertContains(page, 'data-export-refusal')
        self.assertContains(page, layout_transport.LAYOUT_UNAVAILABLE)

    def test_enqueue_and_the_worker_refuse_before_the_provider_and_release_any_hold(self):
        from core.jobs import enqueue, run_one
        _, _, encounter = self.ready()
        with patch('anthropic.Anthropic') as provider:
            with self.assertRaisesMessage(v3_mapping.ValidationError, v3_mapping.GENERATION_LAYOUT_REFUSAL):
                enqueue(encounter, self.user)
            self.assertFalse(Job.objects.exists())
            self.assertFalse(Reservation.objects.exists())
            # A job queued before this gate existed (or while the verifier was available) meets it in the worker.
            with patch.object(v3_mapping, 'generation_refusal', return_value=''):
                job = enqueue(encounter, self.user)
            self.assertEqual(list(Reservation.objects.filter(job=job).values_list('state', flat=True)), ['held'])
            self.assertTrue(run_one())
            self.assertFalse(provider.return_value.messages.create.called)
        job.refresh_from_db()
        self.assertEqual(job.state, 'failed')
        self.assertIn(v3_mapping.GENERATION_LAYOUT_REFUSAL, job.error)
        self.assertEqual(list(Reservation.objects.filter(job=job).values_list('state', flat=True)), ['released'])
        self.assertFalse(Usage.objects.filter(job=job, charged_cost__gt=0).exists())
        self.assertFalse(ProviderReceipt.objects.exists())

    def test_a_typed_answer_is_refused_at_finalization_with_the_same_words_and_a_blank_export_is_the_original(self):
        template, version, encounter = self.ready()
        payload = {'revision': encounter.revision}
        for field in version.fields:
            payload['answer_' + field['key']] = 'Typed by the clinician.' if field['label'] == 'Plan' else ''
            payload['reviewed_' + field['key']] = 'on'
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/save/', payload).status_code, 302)
        encounter.refresh_from_db()
        final = self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'},
                                 follow=True)
        self.assertIn(layout_transport.LAYOUT_UNAVAILABLE, [str(m) for m in final.context['messages']])
        encounter.refresh_from_db()
        self.assertEqual(encounter.status, 'review')

    def test_old_stored_wording_is_still_recognised_on_the_page(self):
        template, version, encounter = self.ready()
        template.warnings = [w if w != layout_transport.LAYOUT_UNAVAILABLE else layout_transport.LAYOUT_UNAVAILABLE_BEFORE
                             for w in template.warnings]
        template.save(update_fields=['warnings'])
        self.assertEqual(v3_mapping.export_refusal(template), layout_transport.LAYOUT_UNAVAILABLE)


@override_settings(**dict(LIVE, QUILLMD_LAYOUT_SOCKET='/run/synthetic-layout.sock', QUILLMD_LAYOUT_SERVICE_UID='1000'))
class LayoutGenerationGateControls(V3Fixture, TestCase):
    """Positive controls: the gate is the verifier's absence, never 'a form with tab columns' or 'a v3 form'."""

    def setUp(self):
        super().setUp()
        from documents.test_layout_verification import synthetic_verified
        verifier = patch('documents.layout_transport.verify_document', side_effect=synthetic_verified)
        self.verifier = verifier.start(); self.addCleanup(verifier.stop)

    def test_with_a_verifier_that_verifies_generation_is_admitted_and_charged_as_before(self):
        from core.jobs import enqueue, run_one
        template, version = self.uploaded(data=mixed_form(), name='SYNTHETIC verified tab-column visit')
        self.assertTrue(any(f.get('layout_candidate') for f in version.fields))
        self.assertFalse([w for w in template.warnings if 'layout verifier' in w])
        version = self.confirm(template, version)
        self.assertEqual(self.new_encounter(version, notes=NOTES).status_code, 302)
        encounter = Encounter.objects.latest('created_at')
        self.assertEqual(v3_mapping.generation_refusal(version), '')
        self.assertEqual(v3_mapping.export_refusal(template), '')
        with patch('anthropic.Anthropic') as provider:
            provider.return_value.messages.create.side_effect = provider_reply
            job = enqueue(encounter, self.user)
            self.assertTrue(run_one())
            self.assertEqual(provider.return_value.messages.create.call_count, 1)
        job.refresh_from_db()
        self.assertEqual(job.state, 'succeeded')
        self.assertTrue(Usage.objects.get(job=job).charged_cost > 0)

    def test_a_cost_after_a_real_provider_call_is_kept_when_finalization_later_fails(self):
        """The gate is before the hold; it never erases a charge. Verifier present at generation, gone at finalization."""
        from core.jobs import enqueue, run_one
        template, version = self.uploaded(data=mixed_form(), name='SYNTHETIC verifier lost later')
        version = self.confirm(template, version)
        self.assertEqual(self.new_encounter(version, notes=NOTES).status_code, 302)
        encounter = Encounter.objects.latest('created_at')
        with patch('anthropic.Anthropic') as provider:
            provider.return_value.messages.create.side_effect = provider_reply
            job = enqueue(encounter, self.user)
            self.assertTrue(run_one())
        charged = Usage.objects.get(job=job).charged_cost
        self.assertTrue(charged > 0)
        encounter.refresh_from_db()
        payload = {'revision': encounter.revision}
        for field in version.fields:
            payload['answer_' + field['key']] = (encounter.answers.get(field['key']) or {}).get('text', '')
            payload['reviewed_' + field['key']] = 'on'
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/save/', payload).status_code, 302)
        encounter.refresh_from_db()
        from documents.service import DocumentError
        # The verifier is gone by finalization: its own refusal, raised exactly as verify_document raises it.
        self.verifier.side_effect = DocumentError(layout_transport.LAYOUT_UNAVAILABLE, 'layout_unavailable')
        with override_settings(**NO_WITNESS):
            final = self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'},
                                     follow=True)
        self.assertIn(layout_transport.LAYOUT_UNAVAILABLE, [str(m) for m in final.context['messages']])
        self.assertEqual(Usage.objects.get(job=job).charged_cost, charged)
        self.assertEqual(list(Reservation.objects.filter(job=job).values_list('state', flat=True)), ['settled'])

    def test_a_verifier_that_does_not_verify_now_refuses_generation_before_the_hold(self):
        from core.jobs import enqueue
        from documents.service import DocumentError
        template, version = self.uploaded(data=mixed_form(), name='SYNTHETIC unresolved now')
        version = self.confirm(template, version)
        self.assertEqual(self.new_encounter(version, notes=NOTES).status_code, 302)
        encounter = Encounter.objects.latest('created_at')
        self.verifier.side_effect = DocumentError('The rendered form does not establish these tab-column answer '
                                                  'destinations. Review the original.', 'layout_unresolved')
        with patch('anthropic.Anthropic') as provider:
            with self.assertRaisesMessage(v3_mapping.ValidationError, v3_mapping.GENERATION_LAYOUT_REFUSAL):
                enqueue(encounter, self.user)
            self.assertFalse(provider.called)
        self.assertFalse(Job.objects.exists())
        self.assertFalse(Reservation.objects.exists())


@override_settings(**dict(LIVE, **NO_WITNESS))
class OrdinaryFormIsNotGated(V3Fixture, TestCase):
    def test_a_v3_form_with_no_tab_column_answer_generates_without_the_verifier(self):
        from core.jobs import enqueue, run_one
        template, version = self.uploaded(data=ordinary_form(), name='SYNTHETIC ordinary visit')
        self.assertFalse([w for w in template.warnings if 'layout verifier' in w])
        version = self.confirm(template, version)
        self.assertEqual(self.new_encounter(version, notes=NOTES).status_code, 302)
        encounter = Encounter.objects.latest('created_at')
        self.assertEqual(v3_mapping.generation_refusal(version), '')
        self.assertEqual(v3_mapping.export_refusal(template), '')
        with patch('anthropic.Anthropic') as provider:
            provider.return_value.messages.create.side_effect = provider_reply
            job = enqueue(encounter, self.user)
            self.assertTrue(run_one())
            self.assertEqual(provider.return_value.messages.create.call_count, 1)
        job.refresh_from_db()
        self.assertEqual(job.state, 'succeeded')
