"""Older-image writes cannot bypass the exact review of a reinspected legacy form.

These are synthetic database controls, not provider/export qualification. The
blank original and fields come from the real bounded legacy inspector; the
historical Encounter model reproduces what dedac162 can write on this schema.
"""
import copy
import hashlib
import importlib
import io
import json
import os
from pathlib import Path
import tempfile

from django.db import DatabaseError, connection, transaction
from django.db.migrations.exceptions import IrreversibleError
from django.db.migrations.executor import MigrationExecutor
from django.db.migrations.loader import MigrationLoader
from django.db.migrations.recorder import MigrationRecorder
from django.test import TestCase, TransactionTestCase, override_settings
from django.utils import timezone
from docx import Document

from core import import_binding
from core.models import Department, Encounter, Job, MappingConfirmation, Organization, Patient, Reservation, Template, TemplateVersion, User
from documents import inspect_template
from ops import integrity_contract

PREVIOUS = ('core', '0018_servicehealth')
PRIOR = ('core', '0037_transcript_redaction_boundary')
CURRENT = ('core', '0038_reinspection_confirmation_guard')
OLD_GATE = importlib.import_module('core.migrations.0031_previous_release_mapping_writes').GATE
NEW_GATE = importlib.import_module('core.migrations.' + CURRENT[1]).GATE


def encounter_body():
    with connection.cursor() as cursor:
        cursor.execute("SELECT prosrc FROM pg_proc WHERE proname='quillmd_encounter_guard' AND pronamespace='public'::regnamespace")
        return cursor.fetchone()[0]


class ReinspectionFixture:
    def setUp(self):
        super().setUp()
        self.directory = tempfile.TemporaryDirectory(prefix='quillmd-reinspection-db-')
        self.addCleanup(self.directory.cleanup)
        self.user = User.objects.create_user(username='reinspection-db', email='reinspection-db@synthetic.invalid',
                                            verified_at=timezone.now())
        self.organization = Organization.objects.create(name='Synthetic reinspection DB guard')
        self.department = Department.objects.create(organization=self.organization, name='Authorized synthetic department')
        self.patient = Patient.objects.create(department=self.department, display_name='SYNTHETIC GUARD PATIENT',
                                              reference='SYNTHETIC-REINSPECTION-GUARD', created_by=self.user)
        document = Document(); table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Symptoms'
        stream = io.BytesIO(); document.save(stream)
        self.original = stream.getvalue()
        self.path = Path(self.directory.name) / 'blank.docx'; self.path.write_bytes(self.original)
        self.inspection = inspect_template(self.path, engine='legacy')
        self.assertEqual(self.inspection['import_manifest']['engine'], 'legacy')
        self.assertEqual(len(self.inspection['fields']), 1)
        self.original_template, self.original_version = self.version('preserved original')
        self.release_encounter = MigrationLoader(connection).project_state(PREVIOUS).apps.get_model('core', 'Encounter')
        self.archive('actual-inspection.json', inspection=self.inspection,
                     original_sha256=hashlib.sha256(self.original).hexdigest())
        if root := os.environ.get('QUILLMD_WORD_ARTIFACTS'):
            original = Path(root) / 'reinspection-database-guard' / self._testMethodName / 'original.docx'
            with original.open('xb') as stream: stream.write(self.original)
            original.chmod(0o600)

    def version(self, name, *, marker=False, contract=''):
        original_path = Path(self.directory.name) / (name.replace(' ', '-') + '.docx')
        original_path.write_bytes(self.original)
        template = Template.objects.create(department=self.department, name='Synthetic ' + name, format='docx',
            original_name=original_path.name, original_path=str(original_path),
            sha256=hashlib.sha256(self.original).hexdigest(), created_by=self.user)
        manifest = copy.deepcopy(self.inspection['import_manifest'])
        if marker:
            manifest['reinspection'] = {'source_template_id': str(self.original_template.pk),
                'source_version_id': str(self.original_version.pk), 'source_sha256': template.sha256}
        version = TemplateVersion.objects.create(template=template, number=1, fields=self.inspection['fields'],
            contract=contract, import_manifest=manifest, created_by=self.user)
        return template, version

    def legacy_draft(self, version):
        return self.release_encounter.objects.create(department_id=self.department.pk, patient_id=self.patient.pk,
            template_version_id=version.pk, created_by_id=self.user.pk, notes='Synthetic original note.',
            service_date=timezone.localdate(), answers={})

    def put_gate(self, gate):
        with connection.cursor() as cursor: cursor.execute(gate)

    def archive(self, name, **value):
        if not (root := os.environ.get('QUILLMD_WORD_ARTIFACTS')):
            return
        target = Path(root) / 'reinspection-database-guard' / self._testMethodName
        target.mkdir(mode=0o700, parents=True, exist_ok=True)
        data = json.dumps(value, indent=2, sort_keys=True, default=str).encode() + b'\n'
        destination = target / name
        with destination.open('xb') as stream: stream.write(data)
        destination.chmod(0o600)


@override_settings(QUILLMD_WORD_ENGINE='legacy')
class ReinspectionDatabaseGuard(ReinspectionFixture, TestCase):
    def test_original_0031_counterexample_survives_forward_guard_as_a_held_editable_draft(self):
        _, version = self.version('fresh legacy copy', marker=True)
        self.put_gate(OLD_GATE)
        draft = self.legacy_draft(version)
        original = dict(Encounter.objects.filter(pk=draft.pk).values().get())
        self.archive('original-0031-counterexample.json', old_gate=OLD_GATE, import_manifest=version.import_manifest,
            fields=version.fields, original_sha256=hashlib.sha256(self.original).hexdigest(), encounter=original,
            confirmed=False, historical_insert_succeeded=True, job_count=Job.objects.count(), reservation_count=Reservation.objects.count())
        self.put_gate(NEW_GATE)
        self.assertEqual(Encounter.objects.filter(pk=draft.pk).values().get(), original)
        Encounter.objects.filter(pk=draft.pk).update(notes='Synthetic clinician correction.', status='review')
        with self.assertRaisesRegex(DatabaseError, 'cannot be finalized'), transaction.atomic():
            Encounter.objects.filter(pk=draft.pk).update(status='finalized')
        current = Encounter.objects.get(pk=draft.pk)
        self.assertEqual((current.notes, current.status), ('Synthetic clinician correction.', 'review'))
        MappingConfirmation.objects.create(version=version, confirmed_by=self.user)
        Encounter.objects.filter(pk=draft.pk).update(status='finalized', finalized_by=self.user, finalized_at=timezone.now())
        self.assertEqual(Encounter.objects.get(pk=draft.pk).status, 'finalized')
        self.assertEqual(self.path.read_bytes(), self.original)

    def test_previous_image_insert_refuses_unconfirmed_reinspection_and_accepts_its_exact_confirmation(self):
        _, version = self.version('exact confirmed copy', marker=True)
        MappingConfirmation.objects.create(version=self.original_version, confirmed_by=self.user)
        with self.assertRaisesRegex(DatabaseError, 'cannot start a document'), transaction.atomic():
            self.legacy_draft(version)
        self.assertEqual(Encounter.objects.count(), 0)
        self.assertEqual((Job.objects.count(), Reservation.objects.count()), (0, 0))
        MappingConfirmation.objects.create(version=version, confirmed_by=self.user)
        draft = self.legacy_draft(version)
        self.assertEqual(Encounter.objects.get(pk=draft.pk).template_version_id, version.pk)

    def test_existing_legacy_draft_cannot_move_to_an_unconfirmed_reinspection(self):
        draft = self.legacy_draft(self.original_version)
        _, version = self.version('unreviewed move', marker=True)
        with self.assertRaisesRegex(DatabaseError, 'cannot start a document'), transaction.atomic():
            Encounter.objects.filter(pk=draft.pk).update(template_version=version)
        self.assertEqual(Encounter.objects.get(pk=draft.pk).template_version_id, self.original_version.pk)
        MappingConfirmation.objects.create(version=version, confirmed_by=self.user)
        Encounter.objects.filter(pk=draft.pk).update(template_version=version)
        self.assertEqual(Encounter.objects.get(pk=draft.pk).template_version_id, version.pk)

    def test_ordinary_legacy_manifest_and_historical_empty_manifest_keep_existing_behavior(self):
        for version in (self.original_version, TemplateVersion.objects.create(
                template=Template.objects.create(department=self.department, name='Synthetic historical empty import',
                    format='docx', original_name='blank.docx', original_path=str(self.path),
                    sha256=hashlib.sha256(self.original).hexdigest(), created_by=self.user),
                number=1, fields=self.inspection['fields'], created_by=self.user)):
            with self.subTest(manifest=bool(version.import_manifest)):
                draft = self.legacy_draft(version)
                Encounter.objects.filter(pk=draft.pk).update(status='finalized')
                self.assertEqual(Encounter.objects.get(pk=draft.pk).status, 'finalized')

    def test_empty_null_or_false_reinspection_markers_do_not_bypass_database_review(self):
        for number, marker in enumerate(({}, None, False, [], '', 0)):
            template, first = self.version('malformed marker ' + str(number))
            # New template's FIRST manifest must carry the marker: 0035 rejects any subsequent provenance change.
            other = Template.objects.create(department=self.department, name='Synthetic malformed first ' + str(number),
                format='docx', original_name='blank.docx', original_path=str(self.path), sha256=template.sha256, created_by=self.user)
            manifest = {**self.inspection['import_manifest'], 'reinspection': marker}
            version = TemplateVersion.objects.create(template=other, number=1, fields=first.fields,
                import_manifest=manifest, created_by=self.user)
            with self.subTest(marker=marker), self.assertRaisesRegex(DatabaseError, 'cannot start a document'), transaction.atomic():
                self.legacy_draft(version)

    def test_import_provenance_cannot_drop_or_change_reinspection_on_a_correction_version(self):
        template, version = self.version('immutable lineage', marker=True)
        for number, change in enumerate(('drop', 'change'), 2):
            manifest = import_binding.for_mapping(version.import_manifest, version.fields)
            if change == 'drop': manifest.pop('reinspection')
            else: manifest['reinspection']['source_version_id'] = str(version.pk)
            with self.subTest(change=change), self.assertRaisesRegex(DatabaseError, 'keeps its original import choice'), transaction.atomic():
                TemplateVersion.objects.create(template=template, number=number, fields=version.fields,
                    import_manifest=manifest, created_by=self.user)
        self.assertEqual(template.versions.count(), 1)

    def test_department_and_finalized_record_guards_remain_in_force(self):
        other = Department.objects.create(organization=self.organization, name='Foreign synthetic department')
        with self.assertRaisesRegex(DatabaseError, 'Department mismatch'), transaction.atomic():
            Encounter.objects.create(department=other, patient=self.patient, template_version=self.original_version,
                created_by=self.user, notes='Synthetic.', service_date=timezone.localdate())
        draft = self.legacy_draft(self.original_version)
        Encounter.objects.filter(pk=draft.pk).update(status='finalized')
        with self.assertRaisesRegex(DatabaseError, 'Finalized document is immutable'), transaction.atomic():
            Encounter.objects.filter(pk=draft.pk).update(notes='Should never replace a final.')
        self.assertEqual(Encounter.objects.get(pk=draft.pk).notes, 'Synthetic original note.')

    def test_integrity_contract_refuses_the_old_function_even_with_identical_trigger_count(self):
        def assess():
            with connection.cursor() as cursor:
                cursor.execute(integrity_contract.CATALOG); rows = cursor.fetchall()
                cursor.execute(integrity_contract.CONTEXT); context = cursor.fetchall()
            return integrity_contract.assess(rows, context, len(integrity_contract.GUARDS))
        self.assertTrue(assess()['passed'])
        self.put_gate(OLD_GATE)
        stale = assess()
        self.assertEqual((stale['found'], stale['definition_mismatches']), (111, 1))
        self.assertFalse(stale['passed'])
        self.put_gate(NEW_GATE)
        self.assertTrue(assess()['passed'])


@override_settings(QUILLMD_WORD_ENGINE='legacy')
class ReinspectionMigrationReversal(ReinspectionFixture, TransactionTestCase):
    def tearDown(self):
        executor = MigrationExecutor(connection)
        executor.migrate(executor.loader.graph.leaf_nodes())
        super().tearDown()

    def test_any_reinspection_evidence_refuses_reverse_before_guard_or_rows_change(self):
        template, version = self.version('reverse protected copy', marker=True)
        MappingConfirmation.objects.create(version=version, confirmed_by=self.user)
        draft = self.legacy_draft(version)
        before = {'body': encounter_body(), 'applied': set(MigrationRecorder(connection).applied_migrations()),
            'encounters': list(Encounter.objects.values()), 'versions': list(TemplateVersion.objects.values()),
            'confirmations': list(MappingConfirmation.objects.values())}
        self.archive('before-reversal.json', **before)
        with self.assertRaisesRegex(IrreversibleError, '1 mapping version.*reinspection evidence') as refused:
            MigrationExecutor(connection).migrate([PRIOR])
        self.assertEqual(encounter_body(), before['body'])
        self.assertEqual(set(MigrationRecorder(connection).applied_migrations()), before['applied'])
        self.assertEqual(list(Encounter.objects.values()), before['encounters'])
        self.assertEqual(list(TemplateVersion.objects.values()), before['versions'])
        self.assertEqual(list(MappingConfirmation.objects.values()), before['confirmations'])
        for private in (str(template.pk), str(version.pk), str(draft.pk), self.user.email):
            self.assertNotIn(private, str(refused.exception))

    def test_ordinary_legacy_database_can_reverse_only_0038_and_reapply_without_record_change(self):
        draft = self.legacy_draft(self.original_version)
        before = dict(Encounter.objects.filter(pk=draft.pk).values().get())
        MigrationExecutor(connection).migrate([PRIOR])
        self.assertNotIn(CURRENT, MigrationRecorder(connection).applied_migrations())
        self.assertNotIn("import_manifest ? 'reinspection'", encounter_body())
        self.assertEqual(Encounter.objects.filter(pk=draft.pk).values().get(), before)
        MigrationExecutor(connection).migrate([CURRENT])
        self.assertIn("import_manifest ? 'reinspection'", encounter_body())
        self.assertEqual(Encounter.objects.filter(pk=draft.pk).values().get(), before)
