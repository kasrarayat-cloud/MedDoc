"""The unchanged R1-05 notes, affirmative status grammar and governing controls.

These are offline whole-relation/page/finalization tests on a synthetic original.
The retained real-route failure and original hash remain separate evidence.
"""
from io import BytesIO
from unittest.mock import patch

from docx import Document
from django.test import SimpleTestCase, override_settings

from core import generation as g, source_scope, v3_field_relations as r
from documents import word_v3
from .test_v3_field_relations import relation_wire
from .test_v3_supersession_floor import SONNET_55_DEFAULT, raw
from .test_v3_field_relations_contract import RELATIONS


@override_settings(**SONNET_55_DEFAULT)
class CompletedDocumentStatus(SimpleTestCase):
    model = g.SONNET_5_5

    def setUp(self):
        doc = Document()
        doc.add_paragraph('Clinic record')
        doc.add_paragraph('Consent form signed and dated: ______________')
        doc.add_paragraph('Plan: ______________')
        out = BytesIO(); doc.save(out)
        self.original = out.getvalue()
        inspection = word_v3.inspect(self.original)
        self.fields = [f for f in inspection['fields'] if f['label'] == 'Plan']
        self.assertEqual(len(self.fields), 1)
        self.field = self.fields[0]

    def answer(self, notes, pick='Plan: Rest.'):
        return r.validated(raw(notes, self.fields,
            relation_wire(notes, self.fields, {self.field['key']: [pick]}), self.model))[self.field['key']]

    def check(self, notes, expected='supported', pick='Plan: Rest.', text='Rest.'):
        answer = self.answer(notes, pick)
        self.assertEqual((answer['status'], answer['text']),
                         (expected, text if expected == 'supported' else ''), (notes, answer))
        stored = answer if expected == 'supported' else {**answer, 'status': 'supported', 'text': text,
            'evidence': [{'quote': pick, 'start': notes.index(pick), 'end': notes.index(pick) + len(pick)}]}
        page = g.as_currently_governed(g.governance_for(notes, self.fields), notes, {**self.field, **stored})
        self.assertEqual((page['status'], page['text']), (expected, text if expected == 'supported' else ''))
        self.assertEqual(g.governed_decisions(notes, self.fields, {self.field['key']: stored}),
                         [] if expected == 'supported' else ['Plan'])
        return answer

    def test_original_r105_note_has_a_useful_first_answer_and_faithful_export(self):
        notes = 'Consent form signed and dated: Yes\nPlan: Rest.'
        answer = self.check(notes)
        self.assertEqual(answer['evidence'], [{'quote': 'Plan: Rest.', 'start': 35, 'end': 46}])
        filled = word_v3.fill(self.original, self.fields, {self.field['key']: answer['text']})
        texts = [p.text for p in Document(BytesIO(filled)).paragraphs]
        self.assertEqual(texts[:2], ['Clinic record', 'Consent form signed and dated: ______________'])
        self.assertIn('Rest.', texts[2])

    def test_only_affirmative_complete_document_event_grammar_disposes_unfinished_scope(self):
        for source in ('Consent form signed: Yes', 'Document signed and dated: Yes'):
            with self.subTest(source=source):self.check(source + '\nPlan: Rest.')
        for source in ('Consent form signed and dated: No',
                'Consent form signed and dated: not yet', 'Consent form signed and dated: awaiting signature',
                'Consent form will be signed and dated: Yes', 'If consent form signed and dated: Yes',
                'Consent form signed and dated?: Yes', 'Consent form signed and dated: Yes?',
                'Consent form signed and dated?: Yes.', 'Consent form signed and dated: Yes, entered in error',
                'The patient signed and dated: Yes', 'Her father signed and dated: Yes',
                'The manager signed and dated: Yes', 'Consent form signed by her father: Yes'):
            with self.subTest(source=source):self.check(source + '\nPlan: Rest.', 'conflict')

    def test_metadata_does_not_clear_foreign_pronoun_correction_or_patient_opposition(self):
        metadata = 'Consent form signed and dated: Yes\n'
        for prefix, pick, text in (
                ('Her father has ankle pain.\n', 'Plan: Rest.', 'Rest.'),
                ('The doctor underwent colonoscopy.\n', 'Plan: Rest.', 'Rest.'),
                ('The doctor underwent colonoscopy.\n', 'Plan: She will rest.', 'She will rest.'),
                ('', 'Plan: She will rest.', 'She will rest.'),
                ('', 'Plan: I will rest.', 'I will rest.')):
            with self.subTest(prefix=prefix, pick=pick):self.check(prefix + metadata + pick, 'conflict', pick, text)
        for tail in ('\nCorrection: the plan was entered in error.',
                     '\nThe patient will not rest.', '\nIf the patient rests, call the doctor.'):
            with self.subTest(tail=tail):self.check(metadata + 'Plan: Rest.' + tail, 'conflict')

    def test_a_status_never_reanchors_human_scope_but_an_explicit_patient_clause_can(self):
        metadata = 'Consent form signed and dated: Yes\n'
        for lead in ('The doctor underwent colonoscopy.', 'The nurse underwent colonoscopy.',
                     'Alex underwent colonoscopy.', 'Her husband underwent colonoscopy.'):
            with self.subTest(lead=lead):
                self.check(lead + '\n' + metadata + 'Plan: Rest.', 'conflict')
                self.check(lead + '\nThe patient reports ankle pain.\n' + metadata + 'Plan: Rest.')
        self.check('The patient reports ankle pain.\n' + metadata + 'Plan: Rest.')

    def test_unknown_relation_and_signing_destination_gain_no_automatic_permission(self):
        notes = 'Consent form signed and dated: Yes\nPlan: Rest.'
        field = {**self.field, 'kind': 'signature', 'write_protected': True}
        reading = g.governance_for(notes, self.fields)[11]['preceding']
        self.assertFalse(source_scope.completed_document_status_boundary(reading, reading['units'][0],
            reading['units'][1:], field))
        self.assertEqual(g.v3_generation_fields([field])[0], [])
        pick = 'Plan: Rest.'
        stored = {'status': 'supported', 'text': 'Rest.', 'reviewed': True,
                  'evidence': [{'quote': pick, 'start': 35, 'end': 46}]}
        self.assertEqual(g.governed_decisions(notes, self.fields, {self.field['key']: stored}), ['Plan'])

    def test_missing_or_out_of_bounds_syntax_cannot_prove_document_completion(self):
        notes = 'Consent form signed and dated: Yes\nPlan: Rest.'
        reading = g.governance_for(notes, self.fields)[11]['preceding']
        for unavailable in (None, [None, None]):
            with self.subTest(unavailable=unavailable), patch('core.v3_syntax.docs', return_value=unavailable):
                self.assertFalse(source_scope.completed_document_status_boundary(reading,
                    reading['units'][0], reading['units'][1:], self.field))


@override_settings(**RELATIONS)
class CompletedDocumentStatusSonnet5(CompletedDocumentStatus):
    model = g.SONNET_5
