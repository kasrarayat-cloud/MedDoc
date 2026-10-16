"""Offline export boundaries with explicitly synthetic terminology and documents."""
import copy
import hashlib
from contextlib import contextmanager
from datetime import date
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from django.core.exceptions import PermissionDenied, ValidationError
from django.http import Http404
from django.test import RequestFactory, SimpleTestCase, TestCase, override_settings
from django.utils import timezone
from docx import Document
from lxml import etree

from core import coding_service as service
from core.coding_export import W, append_codes, export_status, export_with_codes
from core.models import (AuditEvent, Code, CodeSet, CodingContext, CodingDecision,
                         CodingRelease, CodingRun, Department, Encounter, Membership,
                         Organization, Patient, Template, TemplateVersion, User)


def document_bytes():
    doc = Document()
    doc.add_paragraph('Synthetic finalized encounter — original clinical answer.')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = 'Assessment', 'Synthetic finding.'
    doc.sections[0].header.paragraphs[0].text = 'Synthetic private letterhead'
    doc.sections[0].footer.paragraphs[0].text = 'Retained footer'
    doc.sections[0]._sectPr.find('{' + W + '}cols').set('{' + W + '}num', '2')
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def texts(data):
    with ZipFile(BytesIO(data)) as archive:
        xml = etree.fromstring(archive.read('word/document.xml'))
    return xml.xpath('//w:t/text()', namespaces={'w': W})


class PreservingAppendix(SimpleTestCase):
    def test_only_document_member_changes_and_original_body_prefix_is_exact(self):
        source = document_bytes()
        rows = [{'system': 'ICD-10-CM', 'code': 'Z00.00', 'description': 'Synthetic fixture description'}]
        variant = append_codes(source, rows)
        with ZipFile(BytesIO(source)) as old, ZipFile(BytesIO(variant)) as new:
            self.assertEqual(old.namelist(), new.namelist())
            for name in old.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(old.read(name), new.read(name), name)
            original = old.read('word/document.xml')
            amended = new.read('word/document.xml')
            self.assertTrue(amended.startswith(original[:original.rindex(b'<w:sectPr')]))
            before = etree.fromstring(original).find('{' + W + '}body/{' + W + '}sectPr')
            after = etree.fromstring(amended)
            retained = after.find('{' + W + '}body/{' + W + '}p/{' + W + '}pPr/{' + W + '}sectPr')
            self.assertEqual(etree.tostring(before), etree.tostring(retained).replace(b' xmlns:q="' + W.encode() + b'"', b''))
            final = after.find('{' + W + '}body/{' + W + '}sectPr')
            self.assertIsNone(final.find('{' + W + '}cols'))
            self.assertEqual(final.find('{' + W + '}type').get('{' + W + '}val'), 'nextPage')
        self.assertEqual(texts(variant), texts(source) + ['Diagnosis and procedure codes', 'ICD-10-CM',
                                                        'Z00.00 — Synthetic fixture description'])
        self.assertEqual(append_codes(source, rows), variant)
        self.assertEqual(len(Document(BytesIO(variant)).sections), 2)

    def test_repeat_variant_refused_and_markup_is_escaped(self):
        rows = [{'system': 'ICD-10-CM', 'code': 'Z00.00', 'description': '<Synthetic & escaped>'},
                {'system': 'CPT', 'code': '00000', 'description': 'Fictional test only'}]
        variant = append_codes(document_bytes(), rows)
        self.assertEqual(texts(variant).count('Diagnosis and procedure codes'), 1)
        self.assertIn('Z00.00 — <Synthetic & escaped>', texts(variant))
        self.assertIn('CPT', texts(variant))
        with self.assertRaisesMessage(ValidationError, 'already contains'):
            append_codes(variant, rows)

    def test_malformed_archive_and_section_fail_closed(self):
        with self.assertRaises(ValidationError):
            append_codes(b'not a Word document', [])
        source = BytesIO()
        with ZipFile(source, 'w') as archive:
            archive.writestr('word/document.xml', '<w:document xmlns:w="' + W + '"><w:body><w:p/></w:body></w:document>')
        with self.assertRaisesMessage(ValidationError, 'section boundary'):
            append_codes(source.getvalue(), [])


@override_settings(ADMIN_SITE=False, SYNTHETIC_ONLY=True, CODING_PROVIDER='existing')
class CodingWordExport(TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        override = override_settings(MEDIA_ROOT=Path(temporary.name))
        override.enable(); self.addCleanup(override.disable)
        self.root = Path(temporary.name)
        self.user = User.objects.create_user(username='synthetic-code-export',
            email='synthetic-code-export@synthetic.invalid', verified_at=timezone.now())
        org = Organization.objects.create(name='Synthetic export organization')
        self.department = Department.objects.create(organization=org, name='Synthetic office')
        Membership.objects.create(user=self.user, department=self.department, role='clinician')
        patient = Patient.objects.create(department=self.department, display_name='Synthetic Taylor Export',
                                         reference='SYNTHETIC-EXPORT', created_by=self.user)
        template = Template.objects.create(department=self.department, name='Synthetic Word', format='docx',
            original_name='synthetic.docx', original_path='originals/synthetic.docx', sha256='a' * 64, created_by=self.user)
        version = TemplateVersion.objects.create(template=template, number=1,
            fields=[{'key': 'finding', 'label': 'Finding'}], created_by=self.user)
        self.original = document_bytes()
        self.path = self.root / 'exports' / 'finalized.docx'
        self.path.parent.mkdir(mode=0o700)
        self.path.write_bytes(self.original)
        self.encounter = Encounter.objects.create(department=self.department, patient=patient,
            template_version=version, created_by=self.user, notes='Synthetic finding.',
            answers={'finding': {'text': 'Synthetic finding.'}}, status='finalized', revision=2,
            service_date=date(2026, 10, 16), finalized_at=timezone.now(), finalized_by=self.user,
            export_path='exports/finalized.docx', export_sha256=hashlib.sha256(self.original).hexdigest())
        CodingContext.objects.create(encounter=self.encounter, updated_by=self.user,
            data={**service.CONTEXT_DEFAULTS, 'setting': 'outpatient_office'})
        self.release = CodeSet.objects.create(system='ICD-10-CM', version='SYNTHETIC-EXPORT',
            effective_from=date(2026, 10, 1), effective_until=date(2027, 9, 30),
            source_url='https://synthetic.invalid/explicit-test-reference', sha256='b' * 64)
        CodingRelease.objects.create(release=self.release, digest='c' * 64)
        Code.objects.create(release=self.release, code='Z00.00', description='Synthetic accepted description')
        Code.objects.create(release=self.release, code='Z00.01', description='Synthetic rejected description')
        unavailable = patch('core.coding_service.cpt_status', return_value={'available': False})
        unavailable.start(); self.addCleanup(unavailable.stop)
        inputs, context, releases, key = service.input_snapshot(self.encounter)
        self.result = {'suggestions': [
            {'id': 'accepted', 'system': 'ICD-10-CM', 'code': 'Z00.00', 'description': 'Synthetic accepted description',
             'eligible': True, 'release_id': str(self.release.pk)},
            {'id': 'rejected', 'system': 'ICD-10-CM', 'code': 'Z00.01', 'description': 'Synthetic rejected description',
             'eligible': True, 'release_id': str(self.release.pk)},
            {'id': 'held', 'system': 'CPT', 'code': '00000', 'description': '', 'provider_display': 'UNVERIFIED SYNTHETIC TEXT',
             'eligible': False, 'release_id': ''}], 'questions': [], 'global_blockers': []}
        self.run = CodingRun.objects.create(encounter=self.encounter, user=self.user, input_digest=key,
            inputs=inputs, context=context, releases=releases, state='succeeded', result=copy.deepcopy(self.result),
            selection={'selected_ids': ['accepted'], 'overrides': {}})
        self.decision = self.save_decision()
        # This suite tests export boundaries; the dispatch/terminology engines
        # have separate tests. Returning the value lets forged approvals reach
        # the export's independent held-row and catalog checks.
        validation = patch('core.coding_export.validate_selection', side_effect=lambda _, selections: selections)
        self.validate = validation.start(); self.addCleanup(validation.stop)

    def save_decision(self):
        return CodingDecision.objects.create(run=self.run, user=self.user, input_digest=self.run.input_digest,
            selection_digest=service.selection_digest(self.run), selections=copy.deepcopy(self.run.selection))

    def download(self, user=None, pk=None, query=''):
        request = RequestFactory().get('/synthetic/export-with-codes/' + query)
        request.user = user or self.user
        return export_with_codes(request, pk or self.encounter.pk)

    @contextmanager
    def returned_review(self, run, decision):
        # Finished runs/approvals are immutable in PostgreSQL. Exercise the
        # export's independent defensive checks with returned object copies;
        # never disable those guards or rewrite historical database evidence.
        with patch('core.coding_export.CodingRun.objects.filter') as query, \
                patch('core.coding_export.CodingDecision.objects.filter', return_value=[decision]):
            query.return_value.first.return_value = run
            yield

    def forged_selection(self, ids):
        run, decision = copy.deepcopy(self.run), copy.deepcopy(self.decision)
        run.selection = {'selected_ids': ids, 'overrides': {}}
        decision.selections = copy.deepcopy(run.selection)
        decision.selection_digest = service.selection_digest(run)
        return run, decision

    def test_saved_accepted_only_download_is_repeatable_and_original_is_unchanged(self):
        before = Encounter.objects.filter(pk=self.encounter.pk).values().get()
        response = self.download(query='?selected_ids=held')
        self.assertEqual(response.status_code, 200)
        first = b''.join(response.streaming_content)
        response.file_to_stream.close()
        response = self.download()
        second = b''.join(response.streaming_content)
        response.file_to_stream.close()
        self.assertEqual(first, second)
        self.assertEqual(texts(first), texts(self.original) + ['Diagnosis and procedure codes', 'ICD-10-CM',
                                                             'Z00.00 — Synthetic accepted description'])
        self.assertEqual(self.path.read_bytes(), self.original)
        self.assertEqual(before, Encounter.objects.filter(pk=self.encounter.pk).values().get())
        self.assertEqual(CodingRun.objects.count(), 1)
        self.assertEqual(CodingDecision.objects.count(), 1)
        self.assertEqual(AuditEvent.objects.filter(action='encounter.export_with_codes').count(), 2)
        files = list((self.root / 'exports' / 'coding').rglob('*.docx'))
        self.assertEqual(len(files), 1)
        self.assertEqual(files[0].stat().st_mode & 0o777, 0o600)
        self.validate.assert_called()

    def test_stale_revision_and_patient_reassignment_cannot_export_prior_approval(self):
        for field, value in [('revision', 3), ('notes', 'Changed saved note.'), ('service_date', date(2027, 10, 16))]:
            changed = copy.deepcopy(self.encounter)
            setattr(changed, field, value)
            with patch('core.coding_export.get_object_or_404', return_value=changed):
                self.assertEqual(self.download().status_code, 409, field)
        other = Patient.objects.create(department=self.department, display_name='Synthetic Other', reference='OTHER', created_by=self.user)
        changed = copy.deepcopy(self.encounter)
        changed.patient = other
        with patch('core.coding_export.get_object_or_404', return_value=changed):
            self.assertEqual(self.download().status_code, 409)
        self.assertFalse(AuditEvent.objects.filter(action='encounter.export_with_codes').exists())

    def test_unaccepted_and_forged_held_decisions_are_refused(self):
        for ids in ([], ['held']):
            with self.returned_review(*self.forged_selection(ids)):
                self.assertEqual(self.download().status_code, 409)
        self.assertEqual(self.path.read_bytes(), self.original)

    def test_changed_approval_result_or_selected_description_fails_closed(self):
        changed, decision = copy.deepcopy(self.run), copy.deepcopy(self.decision)
        changed.result['suggestions'][0]['description'] = 'Unreviewed replacement'
        with self.returned_review(changed, decision):
            self.assertEqual(self.download().status_code, 409)
        decision.selection_digest = service.selection_digest(changed)
        with self.returned_review(changed, decision):
            self.assertEqual(self.download().status_code, 409)

    def test_forged_eligible_cpt_still_requires_current_entitlement(self):
        cpt = CodeSet.objects.create(system='CPT', version='SYNTHETIC-CPT-EXPORT',
            effective_from=date(2026, 1, 1), effective_until=date(2026, 12, 31),
            source_url='https://synthetic.invalid/fictional-cpt-reference', sha256='d' * 64)
        CodingRelease.objects.create(release=cpt, digest='e' * 64)
        Code.objects.create(release=cpt, code='00000', description='Fictional test only')
        with patch('core.coding_service.cpt_status', return_value={'available': True, 'release_id': str(cpt.pk)}):
            inputs, context, releases, key = service.input_snapshot(self.encounter)
            result = copy.deepcopy(self.result)
            result['suggestions'][2].update(eligible=True, release_id=str(cpt.pk), description='Fictional test only')
            self.run = CodingRun.objects.create(encounter=self.encounter, user=self.user, input_digest=key,
                inputs=inputs, context=context, releases=releases, state='succeeded', result=result,
                selection={'selected_ids': ['held'], 'overrides': {}})
            self.save_decision()
            with patch('core.coding_export.cpt_status', return_value={'available': False}):
                response = self.download()
            self.assertEqual(response.status_code, 409)
            self.assertIn(b'CPT export remains held', response.content)

    def test_unavailable_reference_and_validator_failure_cannot_use_cached_export(self):
        response = self.download(); self.assertEqual(response.status_code, 200); response.file_to_stream.close()
        self.validate.side_effect = ValidationError('Synthetic terminology no longer eligible')
        self.assertEqual(self.download().status_code, 409)
        self.validate.side_effect = lambda _, selections: selections
        with patch('core.coding_export.Code.objects.select_related') as catalog:
            catalog.return_value.filter.return_value.first.return_value = None
            self.assertEqual(self.download().status_code, 409)

    def test_other_tenant_and_revoked_member_cannot_download(self):
        outsider = User.objects.create_user(username='synthetic-other-export',
            email='synthetic-other-export@synthetic.invalid', verified_at=timezone.now())
        with self.assertRaises(Http404): self.download(user=outsider)
        member = Membership.objects.get(user=self.user, department=self.department)
        member.active = False; member.save(update_fields=['active'])
        with self.assertRaises(Http404): self.download()

    def test_inactive_user_refused_but_readonly_member_can_download(self):
        member = Membership.objects.get(user=self.user, department=self.department)
        member.role = 'employee'; member.save(update_fields=['role'])
        response = self.download(); self.assertEqual(response.status_code, 200); response.file_to_stream.close()
        self.user.is_active = False; self.user.save(update_fields=['is_active'])
        with self.assertRaises(PermissionDenied): self.download()

    def test_pdf_and_original_integrity_failure_leave_original_download_separate(self):
        pdf = copy.deepcopy(self.encounter)
        pdf.template_version.template.format = 'pdf'
        with patch('core.coding_export.get_object_or_404', return_value=pdf):
            self.assertEqual(self.download().status_code, 409)
        self.assertFalse(export_status(pdf)['available'])
        self.path.write_bytes(self.original + b'corrupt')
        self.assertEqual(self.download().status_code, 409)

    def test_no_saved_decision_never_infers_acceptance(self):
        with patch('core.coding_export.CodingDecision.objects.filter', return_value=[]):
            self.assertEqual(self.download().status_code, 409)
            self.assertFalse(export_status(self.encounter)['available'])

    def test_existing_variant_integrity_failure_is_not_overwritten(self):
        response = self.download(); self.assertEqual(response.status_code, 200); response.file_to_stream.close()
        path = next((self.root / 'exports' / 'coding').rglob('*.docx'))
        path.write_bytes(b'preserve failed variant')
        self.assertEqual(self.download().status_code, 409)
        self.assertEqual(path.read_bytes(), b'preserve failed variant')
