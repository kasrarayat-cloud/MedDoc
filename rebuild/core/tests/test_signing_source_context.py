"""Source-owned signing directions cannot hide behind an ordinary inline token."""
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
from pathlib import Path
import tempfile
from types import SimpleNamespace
from unittest.mock import patch
import uuid

from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, TestCase, override_settings
from docx import Document
from docx.oxml import OxmlElement

from core import preparation, word_contract
from core.template_source import approval_binding, physical_source_for_version
from documents import DocumentError
from documents.engine import DocumentFailure, TOKEN, docx_inspect
from documents.preparation import TYPED_PROTOCOL, PROTOCOL, SIGNING_CONTEXT, candidate, digest, inventory, validate_inventory
from documents.preparation_service import inspect_preparation, prepare_candidate
from documents.question_context import CONFIG_KEY as QUESTION_KEY, CONTEXT as QUESTION_CONTEXT
from documents.signing_source_context import CONFIG_KEY, CONTEXT, documented_event_time_request, require_source_protection, source_scopes, validate_binding


class SigningSourceContextTests(SimpleTestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(); self.addCleanup(directory.cleanup)
        self.root = Path(directory.name); self.counter = 0
        settings = override_settings(MEDIA_ROOT=self.root, WORD_PREPARATION_ENABLED=True)
        settings.enable(); self.addCleanup(settings.disable)

    def make(self, paragraphs=(), rows=()):
        document = Document()
        for text in paragraphs: document.add_paragraph(text)
        for values in rows:
            table = document.add_table(rows=1, cols=len(values))
            for cell, text in zip(table.rows[0].cells, values): cell.text = text
        self.counter += 1; path = self.root / f'blank-{self.counter}.docx'; document.save(path)
        return path

    def inspect(self, path, context=CONTEXT):
        return inventory(path, protocol=TYPED_PROTOCOL, signing_context=context,
            question_context=QUESTION_CONTEXT)

    def assert_guided(self, path):
        before = path.read_bytes()
        with self.assertRaises(DocumentFailure) as caught: self.inspect(path)
        self.assertEqual(caught.exception.code, 'guided_signing')
        self.assertEqual(path.read_bytes(), before)

    def test_clinical_signs_remain_ordinary_across_existing_table_and_literal_imports(self):
        labels = ['Record each documented vital sign and weight with its label, value and unit.',
            'Vital signs', 'Vital sign', 'Record the observed clinical sign']
        for label in labels:
            # A comma is valid printed prose; it is not valid inside the
            # unchanged physical placeholder-name grammar. Keep the full
            # requested sentence intact in supported prose/table contexts.
            layouts = ('named', 'inline', 'table', 'literal') if TOKEN.fullmatch('{{' + label + '}}') else ('inline', 'table')
            for layout in layouts:
                with self.subTest(label=label, layout=layout):
                    path = self.make(['{{' + label.rstrip('.') + '}}'] if layout == 'named' else
                        [label + ': {{answer}}'] if layout == 'inline' else
                        [label + ': ______'] if layout == 'literal' else (),
                        [(label, '')] if layout == 'table' else ())
                    value = self.inspect(path)
                    self.assertFalse(value['issues']); self.assertTrue(value['anchors'])
                    self.assertTrue(all(not row['protected_signing'] for row in value['anchors']))

    def test_hidden_signature_attestation_and_mixed_directives_are_held(self):
        labels = ['Patient signature', 'Witness signature', 'Please sign below',
            'I attest that this is true', 'I certify this account',
            'Record each vital sign; please sign below', 'Sign', 'Physician to sign']
        for label in labels:
            with self.subTest(label=label):
                path = self.make([label + ': {{answer}}'])
                old = self.inspect(path, SIGNING_CONTEXT)
                self.assertFalse(old['anchors'][0]['protected_signing'])
                self.assertFalse(old['issues'])
                self.assert_guided(path)

    def test_explicitly_named_and_table_or_literal_signing_paths_are_retained(self):
        for label in ('Patient signature', 'Witness signature', 'Please sign below', 'Attestation'):
            for layout in ('named', 'table', 'literal'):
                with self.subTest(label=label, layout=layout):
                    path = self.make(['{{' + label + '}}'] if layout == 'named' else
                        [label + ': ______'] if layout == 'literal' else (),
                        [(label, '')] if layout == 'table' else ())
                    old = self.inspect(path, SIGNING_CONTEXT); new = self.inspect(path)
                    self.assertTrue(all(row['protected_signing'] for row in new['anchors']))
                    self.assertEqual({key: value for key, value in old.items() if key not in ('signing_context', 'inventory_sha256')},
                        {key: value for key, value in new.items() if key not in ('signing_context', 'inventory_sha256')})

    def test_multiple_anchors_do_not_inherit_another_questions_signing_scope(self):
        for text, flags in [('Patient signature: {{one}}; Vital signs: {{two}}', [True, False]),
                ('Vital sign: {{one}}; Patient signature: {{two}}', [False, True])]:
            with self.subTest(text=text):
                path = self.make([text]); fields = docx_inspect(path)[3]
                scopes = source_scopes(path, fields)
                self.assertEqual([scope['protected_signing'] for scope in scopes], flags)
                self.assertFalse(any(scope['ownership_unresolved'] for scope in scopes))
                self.assert_guided(path)
        path = self.make(['Vital sign: {{one}}; Weight: {{two}}'])
        self.assertFalse(self.inspect(path)['issues'])

    def test_unowned_interstitial_or_trailing_signing_is_guided(self):
        for text in ('{{one}} please sign below; {{two}}',
                'Vital signs: {{one}}; Weight: {{two}} please sign below',
                'Record each vital sign: {{answer}}; please sign below'):
            with self.subTest(text=text): self.assert_guided(self.make([text]))

    def test_inline_answer_cell_uses_its_exact_question_cell(self):
        for label in ('Patient signature', 'Witness signature', 'I attest this is true'):
            with self.subTest(label=label): self.assert_guided(self.make(rows=[(label, '{{answer}}')]))
        path = self.make(rows=[('Vital sign', '{{one}}', 'Patient signature', '{{two}}')])
        scopes = source_scopes(path, docx_inspect(path)[3])
        self.assertEqual([row['protected_signing'] for row in scopes], [False, True])
        self.assert_guided(path)

    def test_marker_only_paragraph_uses_an_adjacent_heading_without_stealing_previous_answer(self):
        self.assert_guided(self.make(['Patient signature:', '{{answer}}']))
        self.assert_guided(self.make(['Witness signature:', '{{answer}}']))
        path = self.make(['Patient signature: {{Patient signature}}', '{{ordinary answer}}'])
        scopes = source_scopes(path, docx_inspect(path)[3])
        self.assertFalse(scopes[1]['protected_signing'])
        self.assertFalse(self.inspect(path)['issues'])

    def test_named_signing_does_not_waive_unowned_additional_instruction(self):
        path = self.make(['{{Patient signature}} please sign below; {{ordinary answer}}'])
        self.assert_guided(path)

    def test_named_signer_cannot_substitute_another_printed_owners_identity(self):
        for printed, token in [('Patient signature', 'Physician signature'),
                ('Physician signature', 'Patient signature'), ('Witness signature', 'Physician signature'),
                ('Physician signature', 'Witness signature'), ('Patient signature', 'Signature'),
                ('Signature of the patient', 'Physician signature'),
                ('Signed by: patient', 'Physician signature'),
                ('Signature: patient', 'Physician signature'),
                ('Signed by - witness', 'Physician signature'),
                ('Signature (patient)', 'Physician signature'),
                ('Signature—physician', 'Patient signature'),
                ('Ｐａｔｉｅｎｔ signature', 'Physician signature'),
                ('Patient sig\u200bnature', 'Physician signature'),
                ('Signed by an unidentified person', 'Physician signature'),
                ('Signed by: an unidentified person', 'Physician signature')]:
            with self.subTest(printed=printed, token=token):
                path = self.make([printed + ': {{' + token + '}}'])
                self.assertFalse(self.inspect(path, SIGNING_CONTEXT)['issues'])
                self.assert_guided(path)
        for printed, token in [('Patient signature', 'Patient signature'),
                ('Witness signature', 'Witness signature'), ('Physician signature', 'Clinician signature'),
                ('Signed by (witness)', 'Witness signature'), ('Please sign below', 'Patient signature'),
                ('Signed by: patient', 'Patient signature'), ('Signature: witness', 'Witness signature'),
                ('Signed by - physician', 'Clinician signature'), ('Signature (patient)', 'Patient signature'),
                ('Do sign below', 'Patient signature')]:
            with self.subTest(printed=printed, token=token):
                value = self.inspect(self.make([printed + ': {{' + token + '}}']))
                self.assertFalse(value['issues']); self.assertTrue(value['anchors'][0]['protected_signing'])

    def test_retrospective_consent_fact_named_path_is_unchanged(self):
        for text in ('What consent was documented? {{answer}}',
                'Describe the consent discussion documented for this visit. {{answer}}'):
            with self.subTest(text=text):
                path = self.make([text]); old = self.inspect(path, SIGNING_CONTEXT); new = self.inspect(path)
                self.assertEqual(new['issues'], old['issues'])
                self.assertEqual(new['anchors'], old['anchors'])

    def test_complete_event_time_questions_are_source_facts_not_authorization_acts(self):
        for label in ('Signature date', 'Date of signature', 'Date of witness signature',
                'Patient signature date', 'Signature (date and time)', 'Time of certification'):
            self.assertEqual(documented_event_time_request(label)['intent'], 'documented_event_time')
            for layout in ('named', 'inline', 'table', 'literal'):
                if layout in ('named', 'literal') and not TOKEN.fullmatch('{{' + label + '}}'):
                    continue
                with self.subTest(label=label, layout=layout):
                    path = self.make(['{{' + label + '}}'] if layout == 'named' else
                        [label + ': {{answer}}'] if layout == 'inline' else
                        [label + ': ______'] if layout == 'literal' else (),
                        [(label, '')] if layout == 'table' else ())
                    new = self.inspect(path)
                    self.assertFalse(new['issues'])
                    self.assertTrue(all(not row['protected_signing'] for row in new['anchors']))
                    if layout == 'literal':
                        original = path.read_bytes(); output = self.root / ('prepared-date-' + path.name)
                        candidate(path, output, new, [anchor['id'] for anchor in new['anchors']])
                        self.assertEqual(path.read_bytes(), original)
                        semantic = word_contract.inventory_for(self.inspect(output))
                    else:
                        semantic = word_contract.inventory_for(new)
                    self.assertTrue(all(not row['protected_signing'] for row in semantic['fields']))
                    if layout != 'inline':
                        self.assertTrue(self.inspect(path, SIGNING_CONTEXT)['anchors'][0]['protected_signing'])

    def test_date_word_does_not_waive_combined_signing_body_or_unknown_event_scope(self):
        for label in ('Signature and date', 'Signature/Date', 'Date and patient signature',
                'Please sign and date', 'I attest this is true; date', 'Certification body and date',
                'Date of patient signature and witness approval', 'Date of unknown signature'):
            with self.subTest(label=label):
                self.assertIsNone(documented_event_time_request(label))
                self.assert_guided(self.make([label + ': {{answer}}']))
        self.assert_guided(self.make(['Patient signature: {{Signature date}}']))
        self.assert_guided(self.make(['Signature date: {{Patient signature}}']))

    def test_signer_objects_consume_complete_known_lists_and_hold_unknown_trailing_material(self):
        for printed, token in [
                ('Signed by patient, witness, and physician', 'Patient and witness signature'),
                ('Signed by patient, or physician', 'Patient signature'),
                ('Signature: patient, witness, and physician', 'Patient and witness signature'),
                ('Patient, witness, and physician signature', 'Patient and witness signature'),
                ('Signed by patient representative', 'Patient signature'),
                ('Signed by patient,', 'Patient signature'),
                ('Signed by patient, witness, and an unidentified person', 'Patient and witness signature'),
                ('Signed by (patient), and physician', 'Patient signature'),
                ('Signed by patient; physician also required', 'Patient signature'),
                ('Sign by patient', 'Physician signature')]:
            with self.subTest(printed=printed, token=token):
                self.assert_guided(self.make([printed + ': {{' + token + '}}']))
        for printed, token in [
                ('Signed by patient, witness, and physician', 'Patient and witness and physician signature'),
                ('Signed by patient, or physician', 'Patient or physician signature'),
                ('Signed by patient, witness and physician', 'Patient and witness and physician signature'),
                ('Signature: (patient, witness, and physician)', 'Patient and witness and physician signature'),
                ('Patient, witness, and physician signature', 'Patient and witness and physician signature'),
                ('Sign by patient', 'Patient signature')]:
            with self.subTest(printed=printed, token=token):
                value = self.inspect(self.make([printed + ': {{' + token + '}}']))
                self.assertFalse(value['issues'])
                self.assertTrue(word_contract.inventory_for(value)['fields'][0]['protected_signing'])

    def test_styled_runs_are_not_word_boundaries_and_hidden_visible_context_is_not_omitted(self):
        path = self.make(); document = Document(path); paragraph = document.add_paragraph()
        paragraph.add_run('Patient sig').bold = True
        paragraph.add_run('nature: '); paragraph.add_run('{{answer}}'); document.save(path)
        self.assert_guided(path)
        path = self.make(); document = Document(path); paragraph = document.add_paragraph()
        wrapper = OxmlElement('w:customXml'); run = OxmlElement('w:r'); text = OxmlElement('w:t')
        text.text = 'Patient signature: '; run.append(text); wrapper.append(run)
        paragraph._p.append(wrapper); paragraph.add_run('{{answer}}'); document.save(path)
        self.assert_guided(path)

    def test_locator_tampering_cannot_change_printed_context_owner(self):
        path = self.make(['Patient signature: {{answer}}']); fields = docx_inspect(path)[3]
        changed = deepcopy(fields); changed[0]['locator']['start'] -= 1
        with self.assertRaises(DocumentFailure): require_source_protection(path, changed, CONTEXT)
        changed = deepcopy(fields); changed[0]['label'] = 'Vital signs'
        with self.assertRaises(DocumentFailure): require_source_protection(path, changed, CONTEXT)

    def test_v1_and_absent_ordered_inventory_are_not_reinterpreted(self):
        path = self.make(['Patient signature: {{answer}}'], [('Vital signs', '')])
        absent = inventory(path, protocol=TYPED_PROTOCOL)
        old = self.inspect(path, SIGNING_CONTEXT)
        self.assertFalse(absent['anchors'][1]['protected_signing'])
        self.assertFalse(old['anchors'][1]['protected_signing'])
        self.assertNotIn('signing_context', absent)
        self.assertEqual(old, self.inspect(path, SIGNING_CONTEXT))
        require_source_protection(path, docx_inspect(path)[3], None)
        require_source_protection(path, docx_inspect(path)[3], SIGNING_CONTEXT)

    def test_unknown_or_incompatible_policy_never_falls_back(self):
        path = self.make(['{{Vital sign}}'])
        for context in ('unknown', '', 1, True):
            with self.subTest(context=context):
                with self.assertRaises(DocumentFailure): self.inspect(path, context)
        with self.assertRaises(DocumentFailure): inventory(path, protocol=PROTOCOL, signing_context=CONTEXT)
        with self.assertRaises(DocumentFailure): inventory(path, known={'source': 'unqualified'},
            protocol=TYPED_PROTOCOL, signing_context=CONTEXT)
        value = self.inspect(path)
        config = {CONFIG_KEY: CONTEXT, QUESTION_KEY: QUESTION_CONTEXT}
        validate_binding(value, config); word_contract.active(value, config)
        for changed, changed_config in [(value, {}), ({key: item for key, item in value.items() if key != 'signing_context'}, config),
                (value, {CONFIG_KEY: SIGNING_CONTEXT}), (value, {CONFIG_KEY: 'unknown'})]:
            with self.subTest(config=changed_config):
                with self.assertRaises(DocumentFailure): validate_binding(changed, changed_config)

    def test_source_and_semantic_signing_protection_must_agree_before_new_role_payload(self):
        for label in ('Attestation', 'Certification', 'Sign', 'Physician to sign'):
            for layout in ('named', 'table'):
                with self.subTest(label=label, layout=layout):
                    path = self.make(['{{' + label + '}}'] if layout == 'named' else (),
                        [(label, '')] if layout == 'table' else ())
                    old = self.inspect(path, SIGNING_CONTEXT); new = self.inspect(path)
                    # Historical admission is unchanged; only the prospective
                    # source-owned policy refuses the cross-contract mismatch.
                    word_contract.inventory_for(old)
                    with self.assertRaises(DocumentFailure) as caught: word_contract.inventory_for(new)
                    self.assertEqual(caught.exception.code, 'guided_signing')
        from documents.semantic_preparation_roles import fixture_role
        from documents.component_contracts import ENCOUNTER_CONTRACT
        for label in ('Patient signature', 'Witness signature', 'Physician signature'):
            with self.subTest(label=label):
                value = self.inspect(self.make(rows=[(label, '')]))
                semantic = word_contract.inventory_for(value)
                self.assertTrue(semantic['fields'][0]['protected_signing'])
                role = fixture_role(value['fields'], semantic, 'proposer',
                    definition_version=4, component_contract=ENCOUNTER_CONTRACT)
                contract = role['fields'][0]['definition']['component_contract']
                self.assertEqual(contract['rendering'], 'protected'); self.assertEqual(contract['components'], [])

    def test_new_queue_rejects_unknown_policy_before_any_creation(self):
        with patch.object(preparation, 'feature_enabled', return_value=True), \
                patch.object(preparation, 'configuration', return_value={'mode': 'fixture', CONFIG_KEY: 'unknown'}), \
                patch('core.preparation_campaign.active_campaign_config', side_effect=lambda config, _: config), \
                patch('core.testing_authority.create_preparation') as create:
            with self.assertRaises(DocumentFailure): preparation.create_preparation(SimpleNamespace(sha256='a' * 64), None)
        create.assert_not_called()

    def version(self, path, context):
        value = self.inspect(path, context)
        key = uuid.uuid4(); actor = uuid.uuid4()
        template = SimpleNamespace(pk=key, sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
            original_path=path.name, format='docx')
        config = {CONFIG_KEY: context, QUESTION_KEY: QUESTION_CONTEXT}
        prep = SimpleNamespace(pk=uuid.uuid4(), template_id=key, template=template, inventory=value,
            config=config, proposal={}, checker={}, corrections={}, revision=0, correction_history=[],
            candidate_path=path.name, candidate_sha256=template.sha256,
            candidate_info={'fields': deepcopy(value['fields'])}, changes=[], state='approved',
            approved_by_id=actor, approved_at=datetime.now(timezone.utc))
        prep.approval_sha256 = approval_binding(prep)
        version = SimpleNamespace(template=template, template_id=key, preparation_id=prep.pk, preparation=prep,
            prepared_path=path.name, prepared_sha256=template.sha256, document_backend={}, review_requirements=[],
            fields=deepcopy(value['fields']))
        return version

    def test_current_reader_binds_new_policy_and_preserves_old_approval_bytes(self):
        path = self.make(rows=[('Vital signs', '')])
        old = self.version(path, SIGNING_CONTEXT)
        prep = old.preparation
        expected = {'template_id': str(prep.template_id), 'original_sha256': prep.template.sha256,
            'inventory_sha256': prep.inventory['inventory_sha256'], 'proposal': {}, 'checker': {},
            'corrections': {}, 'revision': 0, 'correction_history': [],
            'candidate_sha256': prep.candidate_sha256, 'candidate_info': prep.candidate_info,
            'changes': [], QUESTION_KEY: QUESTION_CONTEXT}
        self.assertEqual(approval_binding(prep), digest(expected))
        self.assertEqual(physical_source_for_version(old), path)
        current = self.version(path, CONTEXT)
        self.assertEqual(physical_source_for_version(current), path)
        original_approval = current.preparation.approval_sha256
        current.preparation.config.pop(CONFIG_KEY)
        self.assertNotEqual(approval_binding(current.preparation), original_approval)
        with self.assertRaises(DocumentFailure): physical_source_for_version(current)
        current.preparation.inventory.pop('signing_context')
        current.preparation.inventory['inventory_sha256'] = digest({key: item for key, item in current.preparation.inventory.items()
            if key != 'inventory_sha256'})
        with self.assertRaises(ValidationError): physical_source_for_version(current)

    def test_actual_sandbox_rejects_hidden_signing_and_revalidates_tampered_old_inventory(self):
        path = self.make(['Patient signature: {{answer}}'])
        with self.assertRaises(DocumentError) as caught:
            inspect_preparation(path, protocol=TYPED_PROTOCOL, signing_context=CONTEXT, question_context=QUESTION_CONTEXT)
        self.assertEqual(caught.exception.code, 'guided_signing')
        old = self.inspect(path, SIGNING_CONTEXT); old['signing_context'] = CONTEXT
        old['inventory_sha256'] = digest({key: item for key, item in old.items() if key != 'inventory_sha256'})
        out = self.root / 'refused.docx'
        with self.assertRaises(DocumentError): prepare_candidate(path, old, [a['id'] for a in old['anchors']], out)
        self.assertFalse(out.exists())
        path = self.make(rows=[('Vital sign', '')]); value = self.inspect(path)
        self.assertEqual(inspect_preparation(path, protocol=TYPED_PROTOCOL, signing_context=CONTEXT,
            question_context=QUESTION_CONTEXT), value)
        out = self.root / 'prepared.docx'
        result = prepare_candidate(path, value, [a['id'] for a in value['anchors']], out)
        self.assertEqual(out.read_bytes(), path.read_bytes()); self.assertTrue(result['parts_preserved'])


@override_settings(WORD_PREPARATION_ENABLED=True, ADMIN_SITE=False,
    TESTING_AUTHORITY_MODE='', TESTING_AUTHORITY_ID='', TESTING_AUTHORITY_SHA256='')
class SigningSourceContextDatabaseTests(TestCase):
    def setUp(self):
        from core.tests.test_semantic_preparation import DefinitionPreparationDatabaseTests
        DefinitionPreparationDatabaseTests.setUp(self)

    def test_actual_preparation_is_held_before_provider_or_reservation_and_not_replayed(self):
        from core.models import Template, TemplatePreparation
        from core.security import private_path
        previous_original = (self.root / 'blank.docx').read_bytes()
        relative = 'hidden-signature-original.docx'
        path = private_path(relative)
        document = Document(); document.add_paragraph('Patient signature: {{answer}}'); document.save(path)
        original = path.read_bytes()
        template = Template.objects.create(department=self.department, name='Synthetic hidden signing boundary',
            format='docx', original_name=relative, original_path=relative,
            sha256=hashlib.sha256(original).hexdigest(), created_by=self.user)
        config = {'mode': 'fixture', 'word_preparation_protocol': TYPED_PROTOCOL,
            CONFIG_KEY: CONTEXT, QUESTION_KEY: QUESTION_CONTEXT}
        prep = TemplatePreparation.objects.create(template=template, requested_by=self.user, config=config)
        with patch.object(preparation, 'call_role') as provider:
            self.assertTrue(preparation.run_one(preparation_pk=prep.pk))
            self.assertFalse(preparation.run_one(preparation_pk=prep.pk))
        provider.assert_not_called(); prep.refresh_from_db()
        self.assertEqual(prep.state, 'blocked'); self.assertEqual(prep.error_code, 'guided_signing')
        self.assertEqual(prep.receipts.count(), 0); self.assertEqual(prep.proposal, {})
        self.assertEqual(prep.checker, {}); self.assertEqual(path.read_bytes(), original)
        self.assertEqual((self.root / 'blank.docx').read_bytes(), previous_original)

    def test_ordinary_new_queue_does_not_submit_named_unprotected_authorizations(self):
        from django.db import transaction
        from core.models import Template
        original_fixture = (self.root / 'blank.docx').read_bytes()
        for index, label in enumerate(('Attestation', 'Certification', 'Sign', 'Physician to sign')):
            with self.subTest(label=label), transaction.atomic():
                relative = f'authorization-{index}.docx'; path = self.root / relative
                document = Document(); table = document.add_table(rows=1, cols=2)
                table.cell(0, 0).text = label; document.save(path); original = path.read_bytes()
                template = Template.objects.create(department=self.department, name='Synthetic signing boundary',
                    format='docx', original_name=relative, original_path=relative,
                    sha256=hashlib.sha256(original).hexdigest(), created_by=self.user)
                with patch.object(preparation, 'configuration', return_value={'mode': 'fixture'}), \
                        patch('core.preparation_campaign.active_campaign_config', side_effect=lambda config, _: config):
                    prep = preparation.create_preparation(template, self.user)
                self.assertEqual(prep.config[CONFIG_KEY], CONTEXT)
                self.assertEqual(prep.config[word_contract.CONFIG_KEY], word_contract.GRAPH_PROFILE)
                with patch.object(preparation, 'provider_role') as provider:
                    self.assertTrue(preparation.run_one(preparation_pk=prep.pk))
                    self.assertFalse(preparation.run_one(preparation_pk=prep.pk))
                provider.assert_not_called(); prep.refresh_from_db()
                self.assertEqual(prep.state, 'blocked'); self.assertEqual(prep.error_code, 'guided_signing')
                self.assertEqual(prep.receipts.count(), 0); self.assertEqual(prep.candidate_path, '')
                self.assertFalse(template.versions.exists()); self.assertEqual(path.read_bytes(), original)
                # Each source is an independent ordinary queue at the same
                # unmodified pilot limits, not four consumed authorizations.
                transaction.set_rollback(True)
        self.assertEqual((self.root / 'blank.docx').read_bytes(), original_fixture)
