"""Source-bound display references; historical text and validation stay intact."""
from copy import deepcopy
import json
from pathlib import Path

from django.test import SimpleTestCase
from docx import Document

from core import preparation, preparation_choice_transport as codec, word_contract
from core.security import private_path
from core.tests import test_preparation_unit_policy as unit
from core.tests import test_preparation_compact_components as compact
from core.tests import test_preparation_stream_transport as sdk_tests
from documents import DocumentError
from documents.component_contracts import CONFIG_KEY, CHOICE_CONTRACT
from documents.preparation import TYPED_PROTOCOL, digest, inventory


class PreparationLabelReferenceTests(SimpleTestCase):
    setUp = unit.PreparationUnitPolicyTests.setUp
    canonical = unit.PreparationUnitPolicyTests.canonical
    rows = unit.PreparationUnitPolicyTests.rows
    replace_rows = staticmethod(unit.PreparationUnitPolicyTests.replace_rows)

    def live(self, transport=codec.LABEL_REFERENCE_TRANSPORT):
        return {**self.config, 'mode': 'live', 'word_preparation_protocol': TYPED_PROTOCOL,
                CONFIG_KEY: CHOICE_CONTRACT, preparation.SCHEMA_TRANSPORT_KEY: transport}

    def encode(self, value, inv=None):
        return codec.encode_label_reference_rows(value, (self.inventory if inv is None else inv)['fields'])

    def decode(self, wire, role='proposer', proposal=None, inv=None, evidence=None):
        inv = self.inventory if inv is None else inv
        envelope = preparation.provider_request(inv, role, self.live(), proposal)
        evidence = evidence or {'response_transport': codec.LABEL_REFERENCE_TRANSPORT,
            'request_envelope': envelope, 'request_envelope_sha256': digest(envelope),
            'response': {'stop_reason': 'end_turn', 'content': [
                {'type': 'text', 'text': json.dumps(wire, ensure_ascii=False)}]}}
        return preparation.decode_role_response(inv, role, self.live(), wire, proposal, evidence,
            digest(preparation.role_payload(inv, role, proposal, config=self.live())))

    def labelled(self, question='Patient identification', labels=('Patient name', 'record reference')):
        doc = Document()
        doc.add_table(rows=1, cols=2).cell(0, 0).text = question
        path = Path(self.root) / 'label-reference.docx'
        doc.save(path)
        inv = inventory(path, protocol=TYPED_PROTOCOL)
        value = self.canonical(inv=inv)
        definition = value['answer_definitions']['fields'][0]['definition']
        definition.update(value_type='text', answer_mode='narrative', unit=None, choices=[],
                          patient_name_requirement='none', meaning=' and '.join(labels))
        template = deepcopy(next(codec.components(value)))
        definition['component_contract'].update(rendering='labelled', components=[
            {**deepcopy(template), 'id': 'c' + str(i), 'question_quote': question, 'label': label,
             'meaning': label, 'value_type': 'text', 'unit': None, 'unit_requirement': 'none',
             'choices': [], 'option_quote': None} for i, label in enumerate(labels)])
        word_contract.validate_result(inv, value, 'proposer', self.live())
        return inv, value

    def test_plain_and_both_roles_full_canonical_roundtrip(self):
        proposal = self.canonical()
        for role in ('proposer', 'checker'):
            value = self.canonical(role, proposal)
            before = deepcopy(value)
            wire = self.encode(value)
            self.assertIsNone(self.rows(wire)[0][2])
            decoded = self.decode(wire, role, proposal)
            self.assertEqual(decoded, before)
            self.assertEqual(value, before)
            word_contract.validate_result(self.inventory, decoded, role, self.live())

    def test_generic_heading_keeps_distinct_meaning_labels_without_invention(self):
        inv, value = self.labelled()
        wire = self.encode(value, inv)
        self.assertEqual([row[2][0] for row in self.rows(wire)], ['meaning', 'meaning'])
        self.assertEqual(self.decode(wire, inv=inv), value)
        self.assertEqual({c['question_quote'] for c in codec.components(value)}, {'Patient identification'})
        self.assertEqual([c['label'] for c in codec.components(value)], ['Patient name', 'record reference'])

    def test_question_quote_shortcut_and_original_question_span_are_exact(self):
        inv, value = self.labelled('Recorded alpha and recorded beta', ('Recorded alpha', 'recorded beta'))
        components = list(codec.components(value))
        components[0]['question_quote'] = components[0]['label']
        wire = self.encode(value, inv)
        self.assertEqual(self.rows(wire)[0][2], ['question_quote'])
        self.assertEqual(self.rows(wire)[1][2], ['question', 19, 32])
        self.assertEqual(self.decode(wire, inv=inv), value)

    def test_unicode_code_points_not_bytes_utf16_or_normalization(self):
        for labels in (('😀 alpha', 'βeta'), ('é', 'e\u0301'), ('"alpha"', 'path\\beta'), ('null', 'NULL')):
            inv, value = self.labelled('Identifier pair', labels)
            wire = self.encode(value, inv)
            self.assertEqual(self.decode(wire, inv=inv), value)
            self.assertEqual([c['label'] for c in codec.components(value)], list(labels))
        inv, value = self.labelled('😀 alpha and beta', ('😀 alpha', 'beta'))
        wire = self.encode(value, inv)
        self.assertEqual(self.rows(wire)[1][2], ['question', 12, 16])
        rows = self.rows(wire)
        rows[1][2] = ['question', 13, 17]  # Incorrect UTF-16-style offset.
        with self.assertRaises(DocumentError):
            self.decode(self.replace_rows(wire, rows), inv=inv)

    def test_exact_tag_arity_native_int_bounds_and_word_boundaries(self):
        inv, value = self.labelled('foobar foo and bar', ('foo', 'bar'))
        wire = self.encode(value, inv)
        cases = ['foo', {}, False, [], [None], ['Question_quote'], ['question_quote', 0],
                 ['question'], ['question', 0], ['question', 0, 3, 4], ['other', 0, 3],
                 ['question', True, 3], ['question', 0, False], ['question', 0.0, 3],
                 ['question', '0', 3], ['question', -1, 3], ['question', 0, 1000],
                 ['question', 3, 3], ['question', 4, 3], ['question', 0, 3],
                 ['question', 1, 6], ['meaning', 0, 2], None]
        for reference in cases:
            rows = self.rows(wire)
            rows[0][2] = reference
            with self.subTest(reference=reference), self.assertRaises(DocumentError):
                self.decode(self.replace_rows(deepcopy(wire), rows), inv=inv)

    def test_inventory_and_response_field_keys_are_closed_and_bound_before_lookup(self):
        wire = self.encode(self.canonical())
        fields = self.inventory['fields']
        for wrong in (None, {}, [], fields + [fields[0]], [{**fields[0], 'key': False}]):
            with self.subTest(fields=wrong), self.assertRaises(DocumentError):
                codec.decode(wire, codec.LABEL_REFERENCE_TRANSPORT, fields=wrong)
        for mutate in (lambda rows: rows.append(deepcopy(rows[0])),
                       lambda rows: rows[0].update(field_key='not-in-inventory'),
                       lambda rows: rows[0].update(field_key=False), lambda rows: rows.pop()):
            changed = deepcopy(wire)
            mutate(changed['answer_definitions']['fields'])
            with self.assertRaises(DocumentError):
                codec.decode(changed, codec.LABEL_REFERENCE_TRANSPORT, fields=fields)

    def test_source_text_is_bound_inventory_not_model_supplied_or_another_field(self):
        inv, value = self.labelled('Original alpha and beta', ('alpha', 'beta'))
        wire = self.encode(value, inv)
        rows = self.rows(wire)
        rows[0][2] = ['question', 0, 8]
        forged = self.replace_rows(deepcopy(wire), rows)
        forged['answer_definitions']['fields'][0]['label'] = 'Invented alpha and beta'
        decoded = codec.decode(forged, codec.LABEL_REFERENCE_TRANSPORT, fields=inv['fields'])
        self.assertEqual(next(codec.components(decoded))['label'], 'Original')
        with self.assertRaises(ValueError):
            word_contract.validate_result(inv, decoded, 'proposer', self.live())

    def test_plain_labelled_and_protected_are_not_silently_converted(self):
        plain = self.encode(self.canonical())
        rows = self.rows(plain)
        rows[0][2] = ['question_quote']
        with self.assertRaises(DocumentError):
            self.decode(self.replace_rows(plain, rows))
        inv, value = self.labelled()
        wire = self.encode(value, inv)
        next(codec.contracts(wire))['rendering'] = 'protected'
        with self.assertRaises(DocumentError):
            self.decode(wire, inv=inv)
        # A non-signing destination cannot become protected by sending no rows.
        next(codec.contracts(wire))['components'] = '[]'
        decoded = self.decode(wire, inv=inv)
        with self.assertRaises(ValueError):
            word_contract.validate_result(inv, decoded, 'proposer', self.live())

    def test_encoder_refuses_invalid_labels_and_other_canonical_constraints(self):
        inv, value = self.labelled('service date and total qualifying minutes',
                                   ('service date', 'total qualifying minutes'))
        for mutate in (lambda c: c.update(label='Service date'),
                       lambda c: c.update(label='invented'), lambda c: c.update(label='servic'),
                       lambda c: c.update(required=False), lambda c: c.update(question_quote='absent'),
                       lambda c: c.update(unit='minutes', unit_requirement='documented'),
                       lambda c: c.update(choices=['invented']), lambda c: c['scope'].update(subject='unknown')):
            changed = deepcopy(value)
            mutate(next(codec.components(changed)))
            before = deepcopy(changed)
            with self.assertRaises(DocumentError):
                self.encode(changed, inv)
            self.assertEqual(changed, before)

    def test_duplicate_labels_and_source_quote_checks_still_follow_decoding(self):
        inv, value = self.labelled()
        wire = self.encode(value, inv)
        rows = self.rows(wire)
        rows[1][2] = rows[0][2]
        decoded = self.decode(self.replace_rows(deepcopy(wire), rows), inv=inv)
        with self.assertRaises(ValueError):
            word_contract.validate_result(inv, decoded, 'proposer', self.live())
        rows = self.rows(wire)
        rows[0][1] = 'not in original question'
        decoded = self.decode(self.replace_rows(deepcopy(wire), rows), inv=inv)
        with self.assertRaises(ValueError):
            word_contract.validate_result(inv, decoded, 'proposer', self.live())

    def test_malformed_rows_duplicate_properties_nonfinite_and_cross_version_refuse(self):
        value = self.canonical()
        wire = self.encode(value)
        row = self.rows(wire)[0]
        for rows in (None, {}, '[]', [None], [{}], [row[:-1]], [row + [None]], [json.dumps(row)]):
            with self.assertRaises(DocumentError):
                self.decode(self.replace_rows(deepcopy(wire), rows))
        for encoded in ('[[NaN]]', '[[Infinity]]', '[[{"x":1,"x":2}]]'):
            changed = deepcopy(wire)
            next(codec.contracts(changed))['components'] = encoded
            with self.assertRaises(DocumentError):
                self.decode(changed)
        old = codec.encode_unit_policy_rows(value)
        self.assertEqual(codec.decode(old, codec.UNIT_POLICY_TRANSPORT), value)
        # Plain null has identical row bytes; version authority is the bound request/evidence.
        self.assertEqual(codec.decode(wire, codec.UNIT_POLICY_TRANSPORT), value)
        inv, labelled = self.labelled()
        with self.assertRaises(DocumentError):
            codec.decode(self.encode(labelled, inv), codec.UNIT_POLICY_TRANSPORT)
        # A null plain label has the same representation but still needs v7 inventory.
        with self.assertRaises(DocumentError):
            codec.decode(old, codec.LABEL_REFERENCE_TRANSPORT)

    def test_inventory_and_transport_envelope_changes_cannot_replay_old_evidence(self):
        value = self.canonical()
        wire = self.encode(value)
        envelope = preparation.provider_request(self.inventory, 'proposer', self.live())
        evidence = {'response_transport': codec.LABEL_REFERENCE_TRANSPORT,
                    'request_envelope': envelope, 'request_envelope_sha256': digest(envelope),
                    'response': {'stop_reason': 'end_turn', 'content': [
                        {'type': 'text', 'text': json.dumps(wire)}]}}
        wrong = deepcopy(evidence)
        wrong['response_transport'] = codec.UNIT_POLICY_TRANSPORT
        with self.assertRaises(DocumentError):
            self.decode(wire, evidence=wrong)
        changed = deepcopy(self.inventory)
        changed['fields'][0]['label'] = 'different original question'
        with self.assertRaises((DocumentError, ValueError)):
            self.decode(wire, inv=changed, evidence=evidence)

    def test_full_twenty_sixty_components_choices_units_nulls_and_signing(self):
        inv, parts = compact.compound_form(self.root)
        proposal = compact.compound_canonical(inv, parts, self.live())
        for role in ('proposer', 'checker'):
            value = compact.compound_canonical(inv, parts, self.live(), role, proposal)
            decoded = self.decode(self.encode(value, inv), role, proposal, inv)
            self.assertEqual(decoded, value)
            self.assertEqual(sum(len(c['components']) for c in codec.contracts(decoded)), 60)
            word_contract.validate_result(inv, decoded, role, self.live())

    def test_outer_schema_grammar_and_limits_reasoning_remain_unchanged(self):
        schema = preparation.role_schema('proposer', TYPED_PROTOCOL, self.inventory, self.live())
        old = codec.wire_schema(schema, codec.UNIT_POLICY_TRANSPORT)
        new = codec.wire_schema(schema, codec.LABEL_REFERENCE_TRANSPORT)
        def grammar(value):
            if isinstance(value, dict): return {k: grammar(v) for k, v in value.items() if k != 'description'}
            if isinstance(value, list): return [grammar(v) for v in value]
            return value
        self.assertEqual(grammar(old), grammar(new))
        self.assertIn('label_ref', codec.instructions(codec.LABEL_REFERENCE_TRANSPORT))
        self.assertNotIn('label_ref', codec.instructions(codec.UNIT_POLICY_TRANSPORT))
        inv, _ = compact.compound_form(self.root)
        self.assertEqual(preparation.role_limits(inv, self.live()),
                         preparation.role_limits(inv, self.live(codec.UNIT_POLICY_TRANSPORT)))
        self.assertEqual(preparation.timing_limits(inv, self.live()), (200, 480))


class PreparationLabelReferenceSDKTests(SimpleTestCase):
    setUp = unit.PreparationUnitPolicyTests.setUp
    receipt = sdk_tests.PreparationStreamTransportTests.receipt
    sdk_client = sdk_tests.PreparationStreamTransportTests.sdk_client
    invoke = sdk_tests.PreparationStreamTransportTests.invoke
    trace = sdk_tests.PreparationStreamTransportTests.trace
    canonical = unit.PreparationUnitPolicyTests.canonical

    def live(self, *, choice=False):
        return {**self.config, 'mode': 'live', 'model': preparation.REASONING_MODEL,
                'word_preparation_protocol': TYPED_PROTOCOL, CONFIG_KEY: CHOICE_CONTRACT,
                preparation.SCHEMA_TRANSPORT_KEY: codec.LABEL_REFERENCE_TRANSPORT,
                preparation.REASONING_KEY: preparation.REASONING_PROFILE,
                preparation.stream_transport.CONFIG_KEY: preparation.stream_transport.TRANSPORT}

    def test_real_sdk_both_roles_raw_envelope_cost_and_retained_roundtrip(self):
        proposal = self.canonical()
        for role in ('proposer', 'checker'):
            canonical = self.canonical(role, proposal)
            wire = codec.encode_label_reference_rows(canonical, self.inventory['fields'])
            body = sdk_tests.Body(sdk_tests.events(wire))
            result, receipt, sent, client, factory, reserve = self.invoke(body, role=role, proposal=proposal, choice=True)
            self.assertEqual(result, canonical)
            self.assertEqual(receipt.state, 'succeeded')
            self.assertEqual(len(sent), 1)
            self.assertEqual(sent[0], preparation.provider_request(self.inventory, role, self.live(), proposal))
            self.assertEqual(sent[0]['thinking'], {'type': 'between_tools'})
            self.assertEqual(sent[0]['output_config']['effort'], 'high')
            self.assertEqual(factory.call_args.kwargs['max_retries'], 0)
            self.assertEqual(reserve.call_count, 1)
            self.assertTrue(body.closed)
            self.assertTrue(client.is_closed())
            raw = json.loads(private_path(receipt.response_path).read_bytes())
            self.assertEqual(raw['response_transport'], codec.LABEL_REFERENCE_TRANSPORT)
            self.assertEqual(raw['request_envelope_sha256'], digest(sent[0]))
            self.assertEqual(json.loads(''.join(b.get('text', '') for b in raw['response']['content'])), wire)
            self.assertEqual(preparation.read_retained_role(receipt, 'live'), canonical)
            self.assertIsNotNone(receipt.computed_cost)

    def test_bad_reference_keeps_original_raw_known_cost_and_no_retry(self):
        wire = codec.encode_label_reference_rows(self.canonical(), self.inventory['fields'])
        contract = next(codec.contracts(wire))
        rows = json.loads(contract['components'])
        rows[0][2] = ['question', True, 1000]
        contract['components'] = json.dumps(rows)
        body = sdk_tests.Body(sdk_tests.events(wire))
        result, receipt, sent, client, factory, reserve = self.invoke(body, choice=True)
        self.assertIsInstance(result, DocumentError)
        self.assertEqual(receipt.state, 'failed')
        self.assertIsNotNone(receipt.computed_cost)
        self.assertEqual(len(sent), 1)
        self.assertEqual(reserve.call_count, 1)
        self.assertEqual(factory.call_args.kwargs['max_retries'], 0)
        raw = json.loads(private_path(receipt.response_path).read_bytes())
        self.assertEqual(json.loads(''.join(b.get('text', '') for b in raw['response']['content'])), wire)
        with self.assertRaises(DocumentError):
            preparation.read_retained_role(receipt, 'live')
