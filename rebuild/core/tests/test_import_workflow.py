"""Auto imports through authenticated upload, review, worker, finalization and genuine download.

Synthetic blank documents only; fixture provider, retained first mapping and immutable original.
"""
import copy
import hashlib
import io
import json
import os
from pathlib import Path
from unittest.mock import patch
from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from django.db import DatabaseError, transaction
from docx import Document
from core import import_binding
from core.jobs import enqueue, run_one
from core.models import Encounter, Template, TemplateVersion, Job, Reservation
from core.security import private_path
from .test_word_v3_workflow import V3Fixture, part_text
from .test_workflow import TEST_SETTINGS


def raw(document):
    stream=io.BytesIO();document.save(stream);return stream.getvalue()


def ordinary_form():
    doc=Document();table=doc.add_table(rows=2,cols=2)
    table.cell(0,0).text='Reason for visit';table.cell(1,0).text='Plan'
    doc.add_paragraph('Initial assessment: ____________')
    doc.add_paragraph('Consent form signed: ____________')
    doc.add_paragraph('Nurse initials: ____________')
    doc.add_paragraph('Witness signature: ____________')
    return raw(doc)


def shared_form():
    doc=Document();table=doc.add_table(rows=2,cols=2)
    table.cell(0,0).text='Medications';table.cell(1,0).text='Allergies';table.cell(0,1).merge(table.cell(1,1))
    doc.add_paragraph('Reason for visit: ____________')
    return raw(doc)


@override_settings(**dict(TEST_SETTINGS,QUILLMD_WORD_ENGINE='auto',QUILLMD_ROLE_SUGGESTIONS=False))
class AutoImportWorkflow(V3Fixture,TestCase):
    def setUp(self):
        super().setUp()
        self.template,self.version=self.uploaded(data=ordinary_form())

    def review_human_actions(self):
        issues=self.version.import_manifest['coverage']['selected_engine_issues']
        self.assertTrue(issues)
        self.assertTrue(all(row['code'] in ('human_signing_action','ambiguous_source_action') for row in issues))
        response=self.client.post(f'/templates/{self.template.pk}/',{'action':'v3_coverage',
            'mapping_version':self.version.pk,**{'exclude_'+row['id']:'on' for row in issues}})
        self.assertEqual(response.status_code,302,response.content.decode())
        self.version=self.latest(self.template)

    def correct_roles(self):
        # This ambiguous physical signing caption cannot be reclassified as a
        # clinical Yes answer. Its exact exclusion remains unchanged for the
        # human workflow; the ordinary assessment can change text-like kind.
        self.review_human_actions()
        by=self.by_label(self.version)
        response=self.client.post(f'/templates/{self.template.pk}/',{'action':'v3_save','mapping_version':self.version.pk,
            'kind_'+by['Initial assessment']['key']:'long_text'})
        self.assertEqual(response.status_code,302,response.content.decode())
        self.version=self.latest(self.template)

    def seed(self,engine):
        from django.core.management import call_command
        with override_settings(BASE_DIR=Path(self.directory.name),QUILLMD_WORD_ENGINE=engine):
            call_command('seed_private',stdout=io.StringIO())
        template=Template.objects.get(name='Clinical assessment')
        version=template.versions.get();import_binding.verify(version)
        self.assertEqual(version.import_manifest['source_sha256'],template.sha256)
        self.assertEqual(version.import_manifest['engine'],'v3' if version.contract=='v3' else 'legacy')
        return version

    def test_auto_demo_seed_keeps_qualified_legacy_contract_and_binding(self):
        version=self.seed('auto')
        self.assertNotEqual(version.contract,'v3')
        self.assertEqual(Encounter.objects.filter(template_version=version).count(),1)

    def test_v3_demo_seed_does_not_fabricate_clinician_mapping_confirmation(self):
        version=self.seed('v3')
        self.assertEqual(version.contract,'v3');self.assertIsNone(version.confirmed)
        self.assertEqual(Encounter.objects.filter(template_version=version).count(),0)

    def test_valid_ordinary_parts_route_allows_large_v3_import_without_generic_batch_flag(self):
        from .test_word_v3_workflow import build_questionnaire
        from core import v3_mapping
        eligible=dict(QUILLMD_WORD_ENGINE='v3',V3_FIELD_RELATIONS=True,V3_GENERATION_BATCHES=False,
                      AI_PROVIDER='anthropic',AI_MODEL='claude-sonnet-5-5',AI_INPUT_RATE='2',
                      AI_OUTPUT_RATE='10',AI_CACHED_RATE='0.20',AI_THINKING='between_tools',AI_EFFORT='medium')
        with override_settings(**eligible):
            self.assertEqual(v3_mapping.field_limit(),800)
            template,version=self.uploaded(data=build_questionnaire(101),name='Synthetic ordinary parts import')
            self.assertEqual(len(version.fields),101)
            self.confirm(template,version)
        for invalid in ({'AI_MODEL':'claude-sonnet-4-6'},{'AI_INPUT_RATE':'3'},{'V3_FIELD_RELATIONS':False}):
            with override_settings(**dict(eligible,**invalid)):
                self.assertEqual(v3_mapping.field_limit(),100)

    def test_original_binding_is_immutable_and_copied_to_explicit_review_version(self):
        first=self.version;manifest=copy.deepcopy(first.import_manifest)
        self.assertEqual(first.contract,'v3')
        self.assertEqual(manifest['engine'],'v3')
        self.assertEqual(manifest['source_sha256'],self.template.sha256)
        import_binding.verify(first,allow_unreviewed=True)
        self.correct_roles()
        self.assertEqual(self.version.import_manifest['initial_mapping_sha256'],manifest['initial_mapping_sha256'])
        self.assertEqual(self.version.import_manifest['engine_version'],manifest['engine_version'])
        self.assertNotEqual(self.version.import_manifest['mapping_sha256'],manifest['mapping_sha256'])
        self.assertEqual(self.version.corrections['kind_changed'],
                         {self.by_label(self.version)['Initial assessment']['key']:['text','long_text']})
        self.assertNotIn('ordinary_information_classified',self.version.corrections)
        first.refresh_from_db();self.assertEqual(first.import_manifest,manifest)
        import_binding.verify(self.version)
        with self.assertRaises(DatabaseError),transaction.atomic():
            TemplateVersion.objects.filter(pk=first.pk).update(import_manifest={})
        # A rollback writer omitting the column cannot create an unbound copy
        # of an imported mapping, even though historical templates keep{}default.
        with self.assertRaises(DatabaseError),transaction.atomic():
            TemplateVersion.objects.create(template=self.template,number=3,fields=self.version.fields,
                                           contract='v3',created_by=self.user)

    def test_initial_clinical_kind_recovers_but_ambiguous_signed_form_stays_human_owned(self):
        by=self.by_label(self.version)
        page=self.client.get(f'/templates/{self.template.pk}/')
        rows={r['label']:r for r in page.context['v3_rows']}
        self.assertTrue(rows['Initial assessment']['kind_editable'])
        self.assertEqual(rows['Initial assessment']['chosen_kind'],'text')
        for label in ('Consent form signed','Nurse initials','Witness signature'):
            self.assertFalse(rows[label]['kind_editable'])
        self.assertTrue(self.version.import_manifest['coverage']['block_generation'])
        self.assertEqual(Job.objects.count(),0);self.assertEqual(Reservation.objects.count(),0)
        response=self.client.post(f'/templates/{self.template.pk}/',{'action':'v3_save','mapping_version':self.version.pk,
            'kind_'+by['Consent form signed']['key']:'text','ordinary_'+by['Consent form signed']['key']:'on'})
        self.assertEqual(response.status_code,400)
        self.assertEqual(self.template.versions.count(),1)
        self.review_human_actions()
        self.assertTrue(import_binding.coverage_reviewed(self.version))
        self.assertEqual(import_binding.ordinary_role_fields(self.version),set())
        # Saving an unrelated label keeps the ordinary automatic assessment.
        response=self.client.post(f'/templates/{self.template.pk}/',{'action':'v3_save','mapping_version':self.version.pk,
            'label_'+by['Plan']['key']:'Treatment plan'})
        self.assertEqual(response.status_code,302)
        self.assertEqual(self.by_label(self.latest(self.template))['Initial assessment']['kind'],'text')

    def test_forged_signer_kind_change_is_refused_without_a_new_mapping(self):
        by=self.by_label(self.version)
        for label in ('Nurse initials','Witness signature'):
            response=self.client.post(f'/templates/{self.template.pk}/',{'action':'v3_save','mapping_version':self.version.pk,
                                                                      'kind_'+by[label]['key']:'text'})
            self.assertEqual(response.status_code,400)
        self.assertEqual(self.template.versions.count(),1)

    def test_ambiguous_kind_change_requires_explicit_ordinary_information_classification(self):
        # Existing lexical "initial" metadata with no printed signing action
        # remains an explicit, source-bound clinician mapping decision.
        from .test_destination_roles import legacy_rows
        with override_settings(QUILLMD_WORD_ENGINE='v3'):
            template,version=self.uploaded(data=legacy_rows('Troponin (initial)','Plan'),name='Synthetic mapping ambiguity')
        by=self.by_label(version);key=by['Troponin (initial)']['key']
        self.assertIn(key,import_binding.ordinary_role_fields(version))
        response=self.client.post(f'/templates/{template.pk}/',{'action':'v3_save','mapping_version':version.pk,
            'kind_'+key:'text'})
        self.assertEqual(response.status_code,400)
        self.assertEqual(template.versions.count(),1)
        response=self.client.post(f'/templates/{template.pk}/',{'action':'v3_save','mapping_version':version.pk,
            'kind_'+key:'text','ordinary_'+key:'on'})
        self.assertEqual(response.status_code,302)
        reviewed=self.latest(template)
        self.assertEqual(reviewed.corrections['ordinary_information_classified'],[key])
        self.assertEqual(self.by_label(reviewed)['Troponin (initial)']['kind'],'text')

    def test_typed_fixture_runs_worker_review_finalization_and_download_with_bound_engine(self):
        self.correct_roles();self.confirm(self.template,self.version)
        notes='Reason for visit: Knee pain\n\nPlan: Physiotherapy\n\nInitial assessment: Stable\n\nConsent form signed: Yes'
        self.assertEqual(self.new_encounter(self.version,notes).status_code,302)
        encounter=Encounter.objects.get(template_version=self.version)
        job=enqueue(encounter,self.user);self.assertTrue(run_one());job.refresh_from_db()
        self.assertEqual(job.state,'succeeded',job.error)
        encounter.refresh_from_db();by=self.by_label(self.version)
        for label in ('Consent form signed','Nurse initials','Witness signature'):
            self.assertFalse(encounter.answers.get(by[label]['key'],{}).get('text'))
        # Keep first provider/fixture adoption distinct from explicit clinician corrections.
        first_output=copy.deepcopy(encounter.answers)
        payload={'revision':encounter.revision}
        desired={'Reason for visit':'Knee pain','Plan':'Physiotherapy','Initial assessment':'Stable'}
        for label,value in desired.items():
            key=by[label]['key'];payload['answer_'+key]=value;payload['reviewed_'+key]='on'
        response=self.client.post(f'/encounters/{encounter.pk}/save/',payload)
        self.assertEqual(response.status_code,302,response.content.decode())
        encounter.refresh_from_db()
        with override_settings(QUILLMD_WORD_ENGINE='legacy'):
            response=self.client.post(f'/encounters/{encounter.pk}/finalize/',{'revision':encounter.revision,'reviewed':'on'},follow=True)
        encounter.refresh_from_db()
        self.assertEqual(encounter.status,'finalized',[str(m) for m in response.context['messages']])
        exported=private_path(encounter.export_path).read_bytes();body=part_text(exported)
        for value in desired.values():self.assertIn(value,body)
        self.assertIn('Consent form signed: ____________',body)
        self.assertIn('Nurse initials: ____________',body);self.assertIn('Witness signature: ____________',body)
        self.assertEqual(hashlib.sha256(private_path(self.template.original_path).read_bytes()).hexdigest(),self.template.sha256)
        self.assertEqual(hashlib.sha256(exported).hexdigest(),encounter.export_sha256)
        download=self.client.get(f'/encounters/{encounter.pk}/export/')
        self.assertEqual(download.status_code,200)
        self.assertEqual(b''.join(download.streaming_content),exported)
        self.assertTrue(first_output,'the worker actually returned an independently retained answer result')
        # Explicit opt-in from the protected VPS runner. This preserves the
        # FIRST worker result separately from the clinician-corrected document.
        state=Path(os.environ.get('QUILLMD_STATE_DIR','/unavailable'))
        allowed=Path('/opt/meddoc-pro/rebuild/.runtime/v1-sol-continuation-20261002/word')
        if state.is_relative_to(allowed):
            capture=state.parent/'captures'/'typed-import';capture.mkdir(mode=0o700,parents=True)
            original=private_path(self.template.original_path).read_bytes()
            (capture/'original.docx').write_bytes(original);(capture/'export.docx').write_bytes(exported)
            record={'scope':'synthetic fixture worker HTTP workflow, not live model or native Word acceptance',
                    'source_sha256':self.template.sha256,'export_sha256':encounter.export_sha256,
                    'notes':notes,'first_answers':first_output,'reviewed_answers':encounter.answers,
                    'expected':desired,'mapping':self.version.fields,'import_manifest':self.version.import_manifest,
                    'review_actions':payload,'job_id':str(job.pk),'job_state':job.state}
            (capture/'workflow.json').write_text(json.dumps(record,indent=2,default=str))

    def test_binding_mismatch_refuses_even_a_cached_role_read(self):
        from core import destination_roles
        self.review_human_actions()
        destination_roles.report(self.version)
        bad=copy.deepcopy(self.version.import_manifest);bad['engine']='legacy';self.version.import_manifest=bad
        with self.assertRaises(ValidationError):destination_roles.report(self.version)

    def test_changed_processor_cannot_generate_or_finalize_under_same_engine_name(self):
        from core import destination_roles
        original=copy.deepcopy(self.version.import_manifest)
        self.version.import_manifest['engine_code_sha256']='0'*64
        with self.assertRaises(ValidationError):destination_roles.generation_fields(self.version)
        self.version.import_manifest=original
        # Same selected engine/version and mapping, different installed implementation.
        with patch('documents.service.engine_fingerprint',return_value='1'*64):
            with self.assertRaises(ValidationError):import_binding.verify(self.version)

    def test_incomplete_legacy_route_refuses_before_provider_or_export(self):
        manifest=copy.deepcopy(self.version.import_manifest);manifest['coverage']['block_generation']=True
        self.version.import_manifest=manifest
        with self.assertRaisesMessage(ValidationError,'No generation or charge'):
            import_binding.verify(self.version)
        from core import destination_roles
        with self.assertRaisesMessage(ValidationError,'No generation or charge'):
            destination_roles.generation_fields(self.version)

    def test_real_legacy_enqueue_refuses_omissions_before_job_or_reservation(self):
        with override_settings(QUILLMD_WORD_ENGINE='legacy'):
            template,version=self.uploaded(name='Synthetic legacy omission',data=ordinary_form())
            self.assertNotEqual(version.contract,'v3')
            self.assertTrue(version.import_manifest['coverage']['block_generation'])
            self.assertEqual(self.new_encounter(version,'Reason for visit: Knee pain').status_code,302)
            encounter=Encounter.objects.get(template_version=version)
            with self.assertRaisesMessage(ValidationError,'No generation or charge'):
                enqueue(encounter,self.user)
        self.assertEqual(Job.objects.count(),0)
        self.assertEqual(Reservation.objects.count(),0)

    def test_known_v3_exclusion_blocks_until_exact_clinician_review_and_stays_unwritten(self):
        template,version=self.uploaded(name='Synthetic shared omission',data=shared_form())
        self.assertEqual([f['label'] for f in version.fields],['Reason for visit'])
        issue=version.import_manifest['coverage']['selected_engine_issues'][0]
        with self.assertRaises(ValidationError):import_binding.verify(version)
        blocked=self.client.post(f'/templates/{template.pk}/',{'action':'v3_confirm','mapping_version':version.pk,
                                                             'checked_every_field':'on'})
        self.assertEqual(blocked.status_code,400)
        # General confirmation is not an area-specific exclusion; precise review
        # makes a NEW immutable version and never invents two destinations.
        response=self.client.post(f'/templates/{template.pk}/',{'action':'v3_coverage','mapping_version':version.pk,
                                                               'exclude_'+issue['id']:'on'})
        self.assertEqual(response.status_code,302,response.content.decode())
        reviewed=self.latest(template);self.assertEqual(reviewed.fields,version.fields)
        self.assertNotEqual(reviewed.pk,version.pk);self.assertEqual(reviewed.import_manifest,version.import_manifest)
        import_binding.verify(reviewed)
        self.assertFalse(reviewed.import_manifest['coverage']['automatic_complete'])
        self.confirm(template,reviewed)
        self.assertEqual(self.new_encounter(reviewed,'Reason for visit: Knee pain').status_code,302)
        encounter=Encounter.objects.get(template_version=reviewed);job=enqueue(encounter,self.user)
        self.assertTrue(run_one());job.refresh_from_db();self.assertEqual(job.state,'succeeded',job.error)
        encounter.refresh_from_db();key=reviewed.fields[0]['key']
        response=self.client.post(f'/encounters/{encounter.pk}/save/',{'revision':encounter.revision,
                                      'answer_'+key:'Knee pain','reviewed_'+key:'on'})
        self.assertEqual(response.status_code,302);encounter.refresh_from_db()
        self.client.post(f'/encounters/{encounter.pk}/finalize/',{'revision':encounter.revision,'reviewed':'on'})
        encounter.refresh_from_db();self.assertEqual(encounter.status,'finalized')
        exported=Document(private_path(encounter.export_path))
        self.assertEqual(exported.tables[0].cell(0,1).text,'')
        self.assertEqual(exported.tables[0].cell(1,1).text,'')
        self.assertEqual(exported.tables[0].cell(0,0).text,'Medications')
        self.assertEqual(exported.tables[0].cell(1,0).text,'Allergies')
        self.assertIn('Knee pain',exported.paragraphs[-1].text)

    def test_exclusions_cannot_be_approved_by_manager_or_bypass_writes_via_copies(self):
        template,version=self.uploaded(name='Synthetic exclusion authorization',data=shared_form())
        issue=version.import_manifest['coverage']['selected_engine_issues'][0]
        self.client.force_login(self.manager)
        response=self.client.post(f'/templates/{template.pk}/',{'action':'v3_coverage','mapping_version':version.pk,
                                                               'exclude_'+issue['id']:'on'})
        self.assertEqual(response.status_code,403);self.assertEqual(template.versions.count(),1)
        version.corrections={'coverage_review':[{'id':issue['id'],'disposition':'leave_unchanged',
                'source_sha256':template.sha256,'reviewed_by':str(self.user.pk)}]}
        # The writer's COPY destination participates even when its primary anchor is elsewhere.
        altered=copy.deepcopy(version.fields)
        altered[0]['locator']['target']['copies']=[{'part':issue['part'],'slot':issue['slot']}]
        version.fields=altered
        version.import_manifest=import_binding.for_mapping(version.import_manifest,altered)
        self.assertFalse(import_binding.coverage_reviewed(version))

    def test_exact_write_in_exclusions_allow_adjacent_boxes_but_refuse_overlapping_copy(self):
        doc=Document();doc.add_paragraph('Symptoms (tick all that apply): ☐ Pain ☐ Other: ______ and ______')
        template,version=self.uploaded(name='Synthetic separate write-in exclusions',data=raw(doc))
        issues=version.import_manifest['coverage']['selected_engine_issues']
        self.assertEqual(len(issues),2)
        self.assertEqual(len(version.fields),1)
        response=self.client.post(f'/templates/{template.pk}/',dict(action='v3_coverage',mapping_version=version.pk,
                                          **{'exclude_'+issue['id']:'on' for issue in issues}))
        self.assertEqual(response.status_code,302,response.content.decode())
        reviewed=self.latest(template)
        self.assertTrue(import_binding.coverage_reviewed(reviewed))
        import_binding.verify(reviewed)
        # The checkbox glyphs are adjacent to the excluded blanks. A valid
        # reviewed copy into either blank must nevertheless refuse generation.
        altered=copy.deepcopy(reviewed.fields)
        altered.append(dict(key='manual_copy',label='Unmapped copy',kind='text',locator={
            'part':issues[0]['part'],'slot':issues[0]['slot'],'target':{
                'type':'span','start':issues[0]['start'],'end':issues[0]['end']}}))
        reviewed.fields=altered
        reviewed.import_manifest=import_binding.for_mapping(reviewed.import_manifest,altered)
        self.assertFalse(import_binding.coverage_reviewed(reviewed))


@override_settings(**dict(TEST_SETTINGS,QUILLMD_WORD_ENGINE='auto',QUILLMD_WORD_AUTO_PREFER_V3=True,
    QUILLMD_ROLE_SUGGESTIONS=False,V3_FIELD_RELATIONS=True,V3_GENERATION_BATCHES=False,
    AI_PROVIDER='anthropic',AI_MODEL='claude-sonnet-5-5',AI_INPUT_RATE='2',AI_OUTPUT_RATE='10',
    AI_CACHED_RATE='0.20',AI_THINKING='between_tools',AI_EFFORT='medium',AI_INFERENCE_GEO=''))
class OrdinaryDefaultImport(V3Fixture,TestCase):
    def test_real_new_upload_and_quote_compile_the_frozen_note_on_the_v4_default_route(self):
        from core import generation as g,v3_field_relations as r
        from ops import sonnet5_stages as stages
        frozen=stages.frozen_inputs()
        case=next(c for c in frozen['cases']['cases'] if c['id']=='N01')
        built=stages.build_forms(Path(self.directory.name)/'frozen_forms','ordinary_notes',frozen)
        # The original frozen form uses its existing advanced omission/rejection
        # route. This separately authored simple table proves the new preference;
        # it is supplemental reachability, never a regraded frozen case.
        basic=Document();table=basic.add_table(rows=8,cols=2)
        labels=('Chief complaint','History of present illness','Past medical history','Medications',
                'Allergies','Examination','Assessment','Plan')
        for row,label in zip(table.rows,labels):row.cells[0].text=label
        records=[]
        for name,original in (('frozen N01 form',built[case['form']]),('supplemental basic table',raw(basic))):
            template,version=self.uploaded(data=original,name=name)
            self.assertEqual(version.contract,'v3');self.assertEqual(version.import_manifest['engine'],'v3')
            if name.startswith('supplemental'):
                self.assertTrue(version.import_manifest['coverage']['new_import_v3_preference']['exact_ordinary_parity'])
                self.assertEqual([f['label'] for f in version.fields],list(labels))
            self.confirm(template,version)
            self.assertEqual(self.new_encounter(version,case['notes']).status_code,302)
            encounter=Encounter.objects.get(template_version=version)
            page=self.client.get(f'/encounters/{encounter.pk}/')
            self.assertIsNotNone(page.context['generation_max_charge'])
            config=g.configuration(version.fields)
            self.assertEqual(config['model'],'claude-sonnet-5-5')
            self.assertEqual(config['ordinary_field_contract'],r.CONTRACT)
            compiled=r.compile_request(encounter.notes,version.fields,config)
            wire=compiled['wire'];payload=json.loads(wire['messages'][0]['content'])
            self.assertEqual(payload['contract'],'ordinary-field-relations-v4')
            self.assertEqual(payload['catalogue'],version.fields)
            self.assertEqual(payload['notes'],case['notes'])
            marker=object()
            with patch('core.v3_field_relations.generate',return_value=marker) as dispatch:
                self.assertIs(g.generate(encounter.notes,version.fields,config),marker)
                dispatch.assert_called_once()
            records.append({'scope':'authenticated upload/confirm/encounter/quote and exact request compilation; no provider',
                'form':name,'source_sha256':template.sha256,'mapping':version.fields,'manifest':version.import_manifest,
                'notes_case':'N01','notes_sha256':hashlib.sha256(case['notes'].encode()).hexdigest(),
                'freeze_manifest_sha256':frozen['freeze']['manifest_sha256'],'request':wire,'config':config})
        state=Path(os.environ.get('QUILLMD_STATE_DIR','/unavailable'))
        allowed=Path('/opt/meddoc-pro/rebuild/.runtime/v1-sol-continuation-20261002/word')
        if state.is_relative_to(allowed):
            capture=state.parent/'captures'/('ordinary-default-'+str(encounter.pk))
            capture.mkdir(mode=0o700,parents=True)
            (capture/'route.json').write_text(json.dumps(records,indent=2,default=str))
