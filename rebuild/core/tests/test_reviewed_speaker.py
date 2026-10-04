"""Exact synthetic reviewed SOAP source and independently governed controls."""
import hashlib
import json

from django.test import SimpleTestCase, override_settings

from core import generation as g, v3_field_relations as r
from .test_transcript_context import proof
from .test_v3_supersession_floor import SONNET_55_DEFAULT, raw
from .test_v3_relation_selection_prompt import FIXTURES, SourceHeadingAndSelection
from .test_v3_field_relations import relation_wire


@override_settings(**SONNET_55_DEFAULT)
class ReviewedClinicalSpeaker(SimpleTestCase):
    answers = SourceHeadingAndSelection.answers
    supported = SourceHeadingAndSelection.supported
    held = SourceHeadingAndSelection.held

    def setUp(self):
        self.case = json.loads(FIXTURES.read_text())['cases'][0]
        self.fields = self.case['catalogue']
        # These are the three explicit synthetic review edits retained by the
        # prospective review fixture. This test is not a review-route/provider
        # run and does not substitute the original audio script.
        self.reviewed = self.case['notes'].replace(
            'The earlier six out of ten one entered in error.',
            'The earlier six out of ten was entered in error.').replace(
            'a left ankle strain.', 'a left ankle sprain.').replace(
            'eyes wrapped in cloth', 'ice wrapped in cloth')
        self.assertEqual(hashlib.sha256(self.reviewed.encode()).hexdigest(),
                         'ab251f1b25ff18a0857269e7fc112e7e5dd67021c10b1350d2cab99421b9844d')
        self.assessment = 'The clinician explicitly documents a left ankle sprain.'
        self.directive = ('The clinician explicitly directs rest, ice wrapped in cloth '
                          'for fifteen minutes at a time, and follow-up in three days.')
        self.no_medication = 'No medication is prescribed.'

    def test_complete_reviewed_soap_wire_keeps_three_useful_answers_and_one_explicit_hold(self):
        wire = json.loads(self.case['response_text'])
        for item in wire['fields']:
            item['unit_ids'] = sorted(set(item['context']['answer']) | set(item['context']['qualifier']))
        result = raw(self.reviewed, self.fields, wire, g.SONNET_5_5)
        result['patient_default_context'] = proof(self.reviewed)
        answers = r.validated(result)
        expected = {'Objective': 'Examination mild swelling over the left lateral ankle.\nNo open wound.',
                    'Assessment': self.assessment, 'Plan': self.directive + '\n' + self.no_medication}
        adopted = {}
        for field in self.fields:
            answer = answers[field['key']]
            if field['label'] in expected:
                self.assertEqual((answer['status'], answer['text']), ('supported', expected[field['label']]))
                self.assertEqual('\n'.join(e['quote'] for e in answer['evidence']), expected[field['label']])
                adopted[field['key']] = answer
            else:
                self.assertEqual((answer['status'], answer['text']), ('conflict', ''))
            page = g.as_currently_governed(g.governance_for(self.reviewed, self.fields), self.reviewed, {**field, **answer})
            self.assertEqual((page['status'], page['text']), (answer['status'], answer['text']))
        self.assertEqual(len(adopted), 3)
        self.assertEqual(g.governed_decisions(self.reviewed, self.fields, adopted), [])

    def test_exact_review_and_complete_source_are_prerequisites(self):
        for label, source in (('Assessment', self.assessment), ('Plan', self.directive)):
            with self.subTest(label=label):
                notes = label + '. ' + source
                field, answer = self.answers(notes, label, [source], context=False)
                self.assertEqual((answer['status'], answer['text']), ('conflict', ''))
                result = raw(notes, self.fields, relation_wire(notes, self.fields, {field['key']: [source]}), g.SONNET_5_5)
                result['patient_default_context'] = proof(notes, end=notes.index(source) + 2)
                self.assertEqual(r.validated(result)[field['key']]['status'], 'conflict')
        self.held('Assessment. ' + self.assessment, 'Assessment', [self.assessment], {'Assessment.': 'qualifier'})
        self.held('Plan. ' + self.directive + ' ' + self.no_medication, 'Plan', [self.directive],
                  {self.no_medication: 'qualifier'})

    def test_clinical_speaker_and_foreign_bearer_are_distinct(self):
        for source in ('The clinician has a left ankle sprain.',
                'The clinician explicitly documents her mother\'s left ankle sprain.',
                'Alex documents a left ankle sprain.', 'I document a left ankle sprain.',
                'The patient\'s father documents a left ankle sprain.',
                '"The clinician explicitly documents a left ankle sprain."'):
            with self.subTest(source=source):
                self.held('Assessment. ' + source, 'Assessment', [source])
        for lead in ('Her father underwent colonoscopy.', 'The doctor underwent colonoscopy.'):
            with self.subTest(lead=lead):
                self.held(lead + ' Assessment. ' + self.assessment, 'Assessment', [self.assessment])
        self.supported('The doctor underwent colonoscopy. The patient reports ankle pain. '
                       'Assessment. ' + self.assessment, 'Assessment', [self.assessment])

    def test_clinical_directive_retains_timing_but_no_hypothetical_or_foreign_subject(self):
        self.supported(self.reviewed, 'Plan', [self.directive, self.no_medication])
        for source in ('The clinician may direct rest.', 'The clinician directed rest last year.',
                'If the patient has pain, the clinician directs rest.',
                'The clinician directs rest for her father.',
                'The clinician directs rest for the dog.',
                'The clinician directs that she should rest.',
                'The clinician directs rest and Alex takes warfarin.'):
            with self.subTest(source=source):
                self.held('Plan. ' + source, 'Plan', [source])

    def test_uncertainty_stays_in_the_report_and_source_retraction_still_holds(self):
        source = 'The clinician documents a possible left ankle sprain.'
        self.supported('Assessment. ' + source, 'Assessment', [source])
        for change in ('The report was entered in error.',
                       'The clinician entered the report in error.',
                       'The clinician does not document a left ankle sprain.'):
            with self.subTest(change=change):
                self.held('Assessment. ' + self.assessment + ' ' + change, 'Assessment', [self.assessment])

    def test_a_directive_cannot_establish_a_current_medication_or_erase_a_switch(self):
        from io import BytesIO
        from docx import Document
        from documents import word_v3
        original = Document()
        table = original.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Current medications'
        table.cell(1, 0).text = 'Plan'
        buffer = BytesIO(); original.save(buffer)
        self.fields = word_v3.inspect(buffer.getvalue())['fields']
        directive = 'The clinician explicitly directs warfarin 5 mg daily.'
        self.held('Current medications. ' + directive, 'Current medications', [directive])
        current = 'The patient takes warfarin 5 mg daily.'
        for change in ('The patient switched from warfarin to apixaban.',
                       'Warfarin was stopped.', 'Warfarin was replaced by apixaban.'):
            with self.subTest(change=change):
                self.held(current + ' ' + change + ' Plan. The clinician directs rest.',
                          'Current medications', [current])
