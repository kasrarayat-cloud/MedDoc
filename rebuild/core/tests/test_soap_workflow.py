"""Built-in SOAP uses persisted department scope and the real synthetic document pipeline; zero providers."""
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
from unittest.mock import patch
from zipfile import ZipFile

from django.test import TestCase, override_settings
from django.utils import timezone
from django_otp.plugins.otp_totp.models import TOTPDevice
from docx import Document

from core.billing import welcome_credit
from core.jobs import run_one
from core.models import (AuditEvent, Department, Encounter, Job, Membership, Organization,
                         Patient, Template, TemplateVersion, Usage, User)
from core.security import private_path
from core.soap_views import LABELS, original_bytes
from .test_workflow import TEST_SETTINGS


SOAP_SETTINGS = dict(TEST_SETTINGS, QUILLMD_WORD_ENGINE='v3', AI_PROVIDER='anthropic', AI_MODEL='claude-sonnet-5-5',
                     AI_INPUT_RATE='2', AI_OUTPUT_RATE='10', AI_CACHED_RATE='0.20', AI_THINKING='between_tools',
                     AI_EFFORT='medium', AI_INFERENCE_GEO='', V3_FIELD_RELATIONS=True)


@override_settings(**SOAP_SETTINGS)
class SoapWorkflow(TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(prefix='quillmd-soap-http-')
        self.addCleanup(directory.cleanup)
        self.private = Path(directory.name) / 'private'
        setting = override_settings(MEDIA_ROOT=self.private)
        setting.enable()
        self.addCleanup(setting.disable)
        self.user = User.objects.create_user(username='soap-clinician', email='soap@synthetic.invalid',
                                             verified_at=timezone.now())
        self.organization = Organization.objects.create(name='Synthetic SOAP organization')
        self.department = Department.objects.create(organization=self.organization, name='Synthetic SOAP clinic')
        self.membership = Membership.objects.create(user=self.user, department=self.department, role='clinician')
        self.foreign = Department.objects.create(organization=Organization.objects.create(name='Other synthetic clinic'),
                                                 name='Other department')
        welcome_credit(self.user)
        self.client.force_login(self.user)
        response = self.client.post('/patients/', {'department': self.department.pk, 'display_name': 'SYNTHETIC SOAP PATIENT',
                                                  'reference': 'SYNTHETIC-SOAP', 'synthetic_only': 'on'})
        self.assertEqual(response.status_code, 302)
        self.patient = Patient.objects.get(reference='SYNTHETIC-SOAP')

    def prepare(self):
        response = self.client.post('/templates/soap/', {'department': self.department.pk})
        self.assertEqual(response.status_code, 302, getattr(response, 'context', None))
        template = Template.objects.get(department=self.department)
        self.assertEqual(response.url, f'/templates/{template.pk}/')
        return template, template.versions.order_by('-number').first()

    def confirm(self, template, version):
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_confirm', 'mapping_version': version.pk,
                                                                  'checked_every_field': 'on'})
        self.assertEqual(response.status_code, 302)
        version.refresh_from_db()
        self.assertIsNotNone(version.confirmed)

    def create_document(self, version, notes):
        response = self.client.post('/encounters/new/', {'patient': self.patient.pk, 'template_version': version.pk,
                                                        'notes': notes, 'service_date': timezone.localdate().isoformat(),
                                                        'synthetic_only': 'on'})
        self.assertEqual(response.status_code, 302, getattr(response, 'context', None))
        return Encounter.objects.get(template_version=version)

    def generate(self, encounter):
        response = self.client.post(f'/encounters/{encounter.pk}/generate/')
        self.assertEqual(response.status_code, 302)
        self.assertTrue(run_one())
        encounter.refresh_from_db()
        self.assertEqual(encounter.status, 'review')
        self.assertEqual(Job.objects.get(encounter=encounter).state, 'succeeded')
        usage = Usage.objects.get(job__encounter=encounter)
        self.assertTrue(usage.simulated)
        self.assertEqual((usage.actual_cost, usage.charged_cost), (0, 0))
        for answer in encounter.answers.values():
            self.assertFalse(answer['reviewed'])
            for evidence in answer.get('evidence', []):
                self.assertEqual(encounter.notes[evidence['start']:evidence['end']], evidence['quote'])

    def retain(self, encounter, export=None):
        """Keep separately named synthetic originals/source/config/answers/export before fixture DB teardown."""
        target = os.getenv('QUILLMD_WORD_ARTIFACTS')
        if not target:
            return
        directory = Path(target) / 'soap-http' / self._testMethodName
        directory.mkdir(parents=True, mode=0o700, exist_ok=True)
        version = encounter.template_version
        template = version.template
        (directory / 'original.docx').write_bytes(private_path(template.original_path).read_bytes())
        (directory / 'source-notes.txt').write_text(encounter.notes)
        (directory / 'encounter.json').write_text(json.dumps({
            'synthetic_only': True, 'provider_requests': 0, 'encounter': str(encounter.pk),
            'template': str(template.pk), 'template_sha256': template.sha256,
            'mapping': str(version.pk), 'mapping_number': version.number, 'fields': version.fields,
            'import_manifest': version.import_manifest, 'answers': encounter.answers,
            'job_config': Job.objects.get(encounter=encounter).config,
            'status': encounter.status, 'export_sha256': encounter.export_sha256,
        }, indent=2, default=str) + '\n')
        if export is not None:
            (directory / 'export.docx').write_bytes(export)

    def review(self, encounter, blanks=(), edits=None):
        edits = edits or {}
        payload = {'revision': encounter.revision}
        for field in encounter.template_version.fields:
            key = field['key']
            payload['answer_' + key] = edits.get(field['label'], encounter.answers[key]['text'])
            payload['reviewed_' + key] = 'on'
            if field['label'] in blanks:
                payload['blank_' + key] = 'on'
        response = self.client.post(f'/encounters/{encounter.pk}/save/', payload)
        self.assertEqual(response.status_code, 302)
        encounter.refresh_from_db()
        self.assertTrue(all(answer['reviewed'] for answer in encounter.answers.values()))

    def finalize(self, encounter):
        response = self.client.post(f'/encounters/{encounter.pk}/finalize/', {'reviewed': 'on', 'revision': encounter.revision})
        self.assertEqual(response.status_code, 302)
        encounter.refresh_from_db()
        self.assertEqual(encounter.status, 'finalized')
        response = self.client.get(f'/encounters/{encounter.pk}/export/')
        self.assertEqual(response.status_code, 200)
        data = b''.join(response.streaming_content)
        self.assertEqual(sha256(data).hexdigest(), encounter.export_sha256)
        return data

    def test_blank_original_is_deterministic_and_has_four_growing_destinations(self):
        data = original_bytes()
        self.assertEqual(data, original_bytes())
        document = Document(BytesIO(data))
        self.assertEqual(document.paragraphs[0].text, 'SOAP note')
        self.assertEqual([(row.cells[0].text, row.cells[1].text) for row in document.tables[0].rows],
                         [(label, '') for label in LABELS])
        self.assertTrue(all(row.height is None for row in document.tables[0].rows))

    def test_library_action_prepares_bound_original_and_requires_clinician_mapping_review(self):
        self.assertContains(self.client.get('/'), 'Use SOAP note')
        response = self.client.get('/templates/soap/')
        self.assertContains(response, 'Assessment and Plan')
        self.assertContains(response, 'Review SOAP form')
        self.assertFalse(Template.objects.exists(), 'GET must not save a private file or original')
        template, version = self.prepare()
        self.assertEqual(private_path(template.original_path).read_bytes(), original_bytes())
        self.assertEqual(template.sha256, sha256(original_bytes()).hexdigest())
        self.assertEqual(tuple(field['label'] for field in version.fields), LABELS)
        self.assertEqual(version.contract, 'v3')
        self.assertEqual(version.import_manifest['source_sha256'], template.sha256)
        self.assertIsNone(version.confirmed)
        response = self.client.get(f'/templates/{template.pk}/')
        self.assertContains(response, 'Not yet confirmed')
        response = self.client.post('/encounters/new/', {'patient': self.patient.pk, 'template_version': version.pk,
                                                        'notes': 'Subjective: cough.', 'synthetic_only': 'on'})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Encounter.objects.exists())
        self.confirm(template, version)
        self.assertContains(self.client.get(f'/templates/{template.pk}/'), 'Use this form')

    def test_repeat_start_preserves_original_and_corrected_mapping_version(self):
        template, version = self.prepare()
        initial_path = template.original_path
        subjective = next(field for field in version.fields if field['label'] == 'Subjective')
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_save', 'mapping_version': version.pk,
                                                                 'label_' + subjective['key']: 'Subjective history'})
        self.assertEqual(response.status_code, 302)
        corrected = template.versions.get(number=2)
        self.prepare()
        self.assertEqual((Template.objects.count(), TemplateVersion.objects.count()), (1, 2))
        template.refresh_from_db()
        self.assertEqual(template.original_path, initial_path)
        self.assertEqual(private_path(initial_path).read_bytes(), original_bytes())
        self.assertEqual(template.versions.order_by('-number').first().pk, corrected.pk)
        self.assertEqual(AuditEvent.objects.filter(action='template.upload').count(), 1)

    def test_other_authorized_department_has_its_own_persisted_original_and_mapping(self):
        template, version = self.prepare()
        Membership.objects.create(user=self.user, department=self.foreign, role='owner')
        response = self.client.post('/templates/soap/', {'department': self.foreign.pk})
        self.assertEqual(response.status_code, 302)
        other = Template.objects.get(department=self.foreign)
        self.assertNotEqual(other.pk, template.pk)
        self.assertNotEqual(other.original_path, template.original_path)
        self.assertEqual(other.sha256, template.sha256)
        self.assertNotEqual(other.versions.get().pk, version.pk)

    def test_cross_department_post_and_saved_original_mapping_and_encounter_are_denied(self):
        response = self.client.post('/templates/soap/', {'department': self.foreign.pk})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['form'].errors)
        self.assertFalse(Template.objects.exists())
        self.assertFalse(self.private.exists())
        # Positive control: real objects and a completed clinician confirmation in the authorized department.
        template, version = self.prepare()
        self.confirm(template, version)
        encounter = self.create_document(version, 'Subjective: cough.')
        outsider = User.objects.create_user(username='soap-outsider', email='outsider@synthetic.invalid', verified_at=timezone.now())
        Membership.objects.create(user=outsider, department=self.foreign, role='clinician')
        self.client.force_login(outsider)
        for route in (f'/templates/{template.pk}/', f'/templates/{template.pk}/original/',
                      f'/encounters/{encounter.pk}/', f'/encounters/{encounter.pk}/export/'):
            self.assertEqual(self.client.get(route).status_code, 404, route)
        self.assertFalse(Template.objects.filter(department=self.foreign).exists())

    def test_inactive_membership_and_nonclinical_roles_cannot_prepare_soap(self):
        for role, active in (('clinician', False), ('manager', True), ('employee', True)):
            with self.subTest(role=role, active=active):
                self.membership.role, self.membership.active = role, active
                self.membership.save(update_fields=['role', 'active'])
                response = self.client.post('/templates/soap/', {'department': self.department.pk})
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.context['form'].errors)
                self.assertFalse(Template.objects.exists())
                self.assertFalse(self.private.exists())

    def test_authenticated_mfa_admin_surface_has_no_soap_clinical_access(self):
        self.user.is_staff = True
        self.user.save(update_fields=['is_staff'])
        device = TOTPDevice.objects.create(user=self.user, confirmed=True)
        session = self.client.session
        session['otp_device_id'] = device.persistent_id
        session.save()
        with override_settings(ADMIN_SITE=True):
            self.assertEqual(self.client.get('/console/services/').status_code, 200)
            for method in ('get', 'post'):
                response = getattr(self.client, method)('/templates/soap/', {'department': self.department.pk})
                self.assertEqual(response.status_code, 404)
                self.assertNotIn(b'SYNTHETIC SOAP PATIENT', response.content)
        self.assertFalse(Template.objects.exists())

    def test_corrupt_existing_original_is_refused_without_replacement_or_new_version(self):
        template, version = self.prepare()
        path = private_path(template.original_path)
        path.write_bytes(b'SYNTHETIC preserved corrupt original')
        response = self.client.post('/templates/soap/', {'department': self.department.pk})
        self.assertContains(response, 'saved SOAP original could not be verified')
        self.assertEqual(path.read_bytes(), b'SYNTHETIC preserved corrupt original')
        self.assertEqual((Template.objects.count(), TemplateVersion.objects.count()), (1, 1))

    def test_processor_refusal_does_not_save_partial_original_or_mapping(self):
        from documents import DocumentError
        with patch('documents.inspect_template', side_effect=DocumentError('Synthetic engine refusal')):
            response = self.client.post('/templates/soap/', {'department': self.department.pk})
        self.assertContains(response, 'could not be prepared')
        self.assertFalse(Template.objects.exists())
        self.assertFalse(list(self.private.rglob('*.docx')))

    def test_typed_soap_notes_use_worker_source_review_finalization_and_faithful_export(self):
        template, version = self.prepare()
        self.confirm(template, version)
        notes = ('Subjective: cough for two days.\nObjective: lungs clear.\n'
                 'Assessment: clinician assessed a viral upper respiratory infection.\n'
                 'Plan: rest, hydration and review if symptoms worsen.')
        encounter = self.create_document(version, notes)
        self.generate(encounter)
        self.retain(encounter)
        self.assertEqual({answer['status'] for answer in encounter.answers.values()}, {'supported'}, encounter.answers)
        self.client.post(f'/encounters/{encounter.pk}/finalize/', {'reviewed': 'on', 'revision': encounter.revision})
        encounter.refresh_from_db()
        self.assertEqual(encounter.status, 'review', 'Generated answers still need explicit review')
        self.review(encounter)
        exported = self.finalize(encounter)
        self.retain(encounter, exported)
        document = Document(BytesIO(exported))
        by_label = {field['label']: encounter.answers[field['key']]['text'] for field in version.fields}
        self.assertEqual([(row.cells[0].text, row.cells[1].text) for row in document.tables[0].rows],
                         [(label, by_label[label]) for label in LABELS])
        self.assertEqual(document.paragraphs[0].text, 'SOAP note')
        with ZipFile(BytesIO(original_bytes())) as original, ZipFile(BytesIO(exported)) as completed:
            self.assertEqual(original.namelist(), completed.namelist())
            for name in original.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(original.read(name), completed.read(name), name)
        self.assertEqual(private_path(template.original_path).read_bytes(), original_bytes())
        self.assertEqual(encounter.template_version_id, version.pk)
        self.assertFalse(any(notes in str(event.details) for event in AuditEvent.objects.all()))

    def test_missing_clinician_assessment_and_plan_stay_missing_until_explicit_blank_review(self):
        template, version = self.prepare()
        self.confirm(template, version)
        encounter = self.create_document(version, 'Subjective: cough for two days.\nObjective: lungs clear.')
        self.generate(encounter)
        self.retain(encounter)
        by_label = {field['label']: encounter.answers[field['key']] for field in version.fields}
        for label in ('Assessment', 'Plan'):
            self.assertEqual((by_label[label]['status'], by_label[label]['text']), ('missing', ''))
        self.review(encounter, blanks=('Assessment', 'Plan'))
        exported = self.finalize(encounter)
        self.retain(encounter, exported)
        document = Document(BytesIO(exported))
        self.assertEqual([row.cells[1].text for row in document.tables[0].rows][2:], ['', ''])
        self.assertEqual({encounter.answers[field['key']]['status'] for field in version.fields[2:]}, {'left_blank'})

    def test_clinician_long_answer_survives_export_without_duplicate_or_truncation(self):
        template, version = self.prepare()
        self.confirm(template, version)
        encounter = self.create_document(version, 'Subjective: cough.\nObjective: lungs clear.\nAssessment: viral illness.\nPlan: review.')
        self.generate(encounter)
        long_plan = '\n'.join(f'Synthetic documented follow-up item {index:03d}: review symptoms and hydration.' for index in range(40))
        self.review(encounter, edits={'Plan': long_plan})
        exported = self.finalize(encounter)
        self.retain(encounter, exported)
        document = Document(BytesIO(exported))
        self.assertEqual(document.tables[0].cell(3, 1).text, long_plan)
        self.assertEqual(document.tables[0].cell(3, 0).text, 'Plan')
        self.assertEqual(sum(long_plan == row.cells[1].text for row in document.tables[0].rows), 1)
        self.assertEqual(private_path(template.original_path).read_bytes(), original_bytes())
