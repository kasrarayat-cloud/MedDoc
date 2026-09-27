"""Contract v3 Word forms through the clinical workflow: upload, mapping review, confirmation, answers, export and
fixture-mode generation. Synthetic blank forms are built here; no patient data, no provider request."""

import hashlib
import io
from pathlib import Path
import tempfile
from zipfile import ZipFile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from docx import Document
from lxml import etree

from core.billing import welcome_credit
from core.models import AuditEvent, Department, Encounter, MappingConfirmation, Membership, Organization, Patient, Template, TemplateVersion, User
from core.security import private_path
from documents.test_word_v3 import _sdt
from .test_workflow import TEST_SETTINGS

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
V3 = dict(TEST_SETTINGS, QUILLMD_WORD_ENGINE="v3")


def build_intake_form(path):
    """One synthetic blank form with every v3 kind: header text, table cells, a date line, printed Yes/No boxes, a
    dropdown, check-all-that-apply boxes, a checkbox control, a log table, a signature and initials line."""
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Patient name: ____________"
    doc.add_paragraph("Synthetic clinic intake")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Referring physician"
    table.cell(1, 0).text = "Reason for visit"
    table.cell(2, 0).text = "Allergies"
    doc.add_paragraph("Date of visit: ____________")
    doc.add_paragraph("Smoker?  ☐ Yes  ☐ No")
    p = doc.add_paragraph("Laterality: ")
    p._p.append(_sdt("dropdown", options=["Choose an item.", "Left", "Right", "Bilateral"]))
    doc.add_paragraph("Symptoms (check all that apply): ☐ Pain ☐ Swelling ☐ Redness")
    p = doc.add_paragraph("Diabetes ")
    p._p.append(_sdt("checkbox"))
    doc.add_paragraph("Medications")
    meds = doc.add_table(rows=4, cols=3)
    meds.cell(0, 0).text = "Medication"
    meds.cell(0, 1).text = "Dose"
    meds.cell(0, 2).text = "Frequency"
    doc.add_paragraph("Clinician signature: ____________")
    doc.add_paragraph("Physician initials: ______")
    buffer = io.BytesIO()
    doc.save(buffer)
    Path(path).write_bytes(buffer.getvalue())
    return buffer.getvalue()


def build_questionnaire(rows, signatures=0):
    doc = Document()
    doc.add_paragraph("Synthetic long questionnaire")
    table = doc.add_table(rows=rows, cols=2)
    for i in range(rows):
        table.cell(i, 0).text = "Question %03d" % (i + 1)
    for i in range(signatures):
        doc.add_paragraph("Clinician signature: ____________" if i % 2 == 0 else "Physician initials: ______")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_legacy_table_form():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Allergies"
    table.cell(1, 0).text = "Plan"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def part_text(data, part="word/document.xml"):
    with ZipFile(io.BytesIO(data)) as z:
        root = etree.fromstring(z.read(part))
    return "".join(t.text or "" for t in root.iter("{%s}t" % W))


class V3Fixture:
    """A department with an owner-clinician, a manager and an employee, one synthetic patient, private storage."""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="quillmd-v3-http-")
        self.addCleanup(self.directory.cleanup)
        self.override = override_settings(MEDIA_ROOT=Path(self.directory.name) / "private")
        self.override.enable()
        self.addCleanup(self.override.disable)
        organization = Organization.objects.create(name="Synthetic v3 practice")
        self.department = Department.objects.create(organization=organization, name="Synthetic intake clinic")
        self.user = self.person("v3-clinician", "Alpha", "clinician")
        self.manager = self.person("v3-manager", "Beta", "manager")
        self.employee = self.person("v3-employee", "Gamma", "employee")
        welcome_credit(self.user)
        self.patient = Patient.objects.create(department=self.department, display_name="SYNTHETIC V3 PATIENT",
                                              reference="SYNTHETIC-V3", created_by=self.user)
        self.client.force_login(self.user)
        self.original = build_intake_form(Path(self.directory.name) / "intake.docx")

    def person(self, name, last, role):
        user = User.objects.create_user(username=name, email=f"{name}@synthetic.invalid", first_name="Synthetic",
                                        last_name=last + " Clinician", verified_at=timezone.now())
        Membership.objects.create(user=user, department=self.department, role=role)
        return user

    def upload(self, data=None, name="Synthetic v3 intake"):
        return self.client.post("/templates/upload/", {"name": name, "department": self.department.pk,
                                "synthetic_only": "on", "file": SimpleUploadedFile("intake.docx", data or self.original)})

    def uploaded(self, **kwargs):
        response = self.upload(**kwargs)
        self.assertEqual(response.status_code, 302, getattr(response, "context", None) and response.context["form"].errors)
        template = Template.objects.get(name=kwargs.get("name", "Synthetic v3 intake"))
        return template, template.versions.get(number=1)

    def by_label(self, version):
        return {f["label"]: f for f in version.fields}

    def latest(self, template):
        return template.versions.order_by("-number").first()

    def confirm(self, template, version=None):
        version = version or self.latest(template)
        response = self.client.post(f"/templates/{template.pk}/", {"action": "v3_confirm", "mapping_version": version.pk,
                                                                   "checked_every_field": "on"})
        self.assertEqual(response.status_code, 302)
        return TemplateVersion.objects.get(pk=version.pk)

    def new_encounter(self, version, notes="Referring physician: Dr Example"):
        return self.client.post("/encounters/new/", {"patient": self.patient.pk, "template_version": version.pk,
                                "notes": notes, "service_date": timezone.localdate().isoformat(), "synthetic_only": "on"})


@override_settings(**V3)
class V3Upload(V3Fixture, TestCase):
    def test_upload_follows_the_engine_setting_and_records_the_contract(self):
        template, version = self.uploaded()
        self.assertEqual(version.contract, "v3")
        self.assertTrue(all(f["locator"]["contract"] == "v3" for f in version.fields))
        by = self.by_label(version)
        self.assertEqual(by["Smoker?"]["kind"], "choice")
        self.assertEqual(by["Laterality"]["options"], ["Left", "Right", "Bilateral"])
        self.assertEqual(by["Symptoms (check all that apply)"]["kind"], "multi_choice")
        self.assertEqual(by["Diabetes"]["kind"], "checkbox")
        self.assertEqual({f["kind"] for f in version.fields if f["kind"] in ("signature", "initials", "table_rows")},
                         {"signature", "initials", "table_rows"})
        self.assertIsNone(version.confirmed)
        self.assertEqual(version.document_backend, {}, "a v3 mapping must never look like the qualified WM form")
        event = AuditEvent.objects.get(action="template.upload", target=str(template.pk))
        self.assertEqual(event.details["contract"], "v3")

    @override_settings(QUILLMD_WORD_ENGINE="legacy")
    def test_the_legacy_setting_keeps_the_legacy_contract(self):
        response = self.upload()
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Template.objects.exists())
        legacy = self.upload(data=build_legacy_table_form(), name="Synthetic legacy table")
        self.assertEqual(legacy.status_code, 302)
        version = Template.objects.get(name="Synthetic legacy table").versions.get()
        self.assertEqual(version.contract, "")
        self.assertFalse(version.needs_confirmation)

    def test_more_than_100_answerable_areas_is_refused_whole_with_the_count(self):
        response = self.upload(data=build_questionnaire(101))
        self.assertEqual(response.status_code, 200)
        errors = " ".join(response.context["form"].non_field_errors())
        self.assertIn("101 answer areas", errors)
        self.assertIn("at most 100", errors)
        self.assertIn("nothing was dropped", errors)
        self.assertFalse(Template.objects.exists())
        originals = Path(self.directory.name) / "private" / "originals"
        self.assertEqual(list(originals.glob("*")) if originals.exists() else [], [], "a refused upload is not kept")

    def test_signature_and_initials_areas_do_not_count_toward_the_limit(self):
        template, version = self.uploaded(data=build_questionnaire(100, signatures=2), name="Synthetic 100")
        self.assertEqual(len(version.fields), 102)
        self.assertEqual(sum(1 for f in version.fields if f["kind"] not in ("signature", "initials")), 100)


@override_settings(**V3)
class V3MappingReview(V3Fixture, TestCase):
    def setUp(self):
        super().setUp()
        self.template, self.version = self.uploaded()

    def post(self, data, version=None):
        return self.client.post(f"/templates/{self.template.pk}/", {"mapping_version": (version or self.latest(self.template)).pk, **data})

    def test_the_review_lists_every_field_in_document_order_with_its_details(self):
        response = self.client.get(f"/templates/{self.template.pk}/")
        self.assertEqual(response.status_code, 200)
        rows = response.context["v3_rows"]
        self.assertEqual([row["key"] for row in rows], [f["key"] for f in self.version.fields])
        self.assertEqual([row["number"] for row in rows], list(range(1, len(rows) + 1)))
        by = {row["field"]["label"]: row for row in rows}
        self.assertTrue(by["Referring physician"]["kind_editable"])
        self.assertFalse(by["Laterality"]["kind_editable"])
        self.assertEqual(by["Laterality"]["options"], ["Left", "Right", "Bilateral"])
        self.assertIn("Header", by["Patient name"]["location"])
        self.assertIn("Body", by["Referring physician"]["location"])
        self.assertEqual(by["Patient name"]["max_chars"], 120)
        self.assertEqual({row["origin"] for row in rows}, {"Automatic"})
        body = response.content.decode()
        self.assertIn("Not an answer &mdash; remove", body)
        self.assertIn("Not yet confirmed", body)
        self.assertNotIn("Use this form", body, "an unconfirmed mapping offers no way to start a document")

    def test_the_add_list_excludes_mapped_places_and_is_searchable(self):
        places = self.client.get(f"/templates/{self.template.pk}/").context["places"]
        slots = {row["slot"] for row in places}
        self.assertIn("word/document.xml#b/p[0]", slots)                      # "Synthetic clinic intake"
        for field in self.version.fields:
            anchor = field["locator"]["part"] + "#" + field["locator"]["slot"]
            self.assertNotIn(anchor, slots)
        self.assertNotIn("word/document.xml#b/tbl[0]/tr[0]/tc[1]/p[0]", slots, "inside a mapped cell")
        self.assertNotIn("word/document.xml#b/p[3]", slots, "holds a mapped dropdown")
        found = self.client.get(f"/templates/{self.template.pk}/?find=clinic intake").context["places"]
        self.assertEqual([row["slot"] for row in found], ["word/document.xml#b/p[0]"])

    def test_rename_kind_change_and_remove_make_a_version_with_recorded_corrections(self):
        by = self.by_label(self.version)
        response = self.post({"action": "v3_save", "label_" + by["Allergies"]["key"]: "Known allergies",
                              "kind_" + by["Referring physician"]["key"]: "long_text",
                              "remove_" + by["Reason for visit"]["key"]: "on"})
        self.assertEqual(response.status_code, 302)
        latest = self.latest(self.template)
        self.assertEqual(latest.number, 2)
        self.assertEqual(latest.contract, "v3")
        new = {f["key"]: f for f in latest.fields}
        self.assertEqual(new[by["Allergies"]["key"]]["label"], "Known allergies")
        self.assertEqual(new[by["Referring physician"]["key"]]["kind"], "long_text")
        self.assertNotIn(by["Reason for visit"]["key"], new)
        self.assertEqual(len(latest.fields), len(self.version.fields) - 1)
        self.assertEqual(latest.corrections, {
            "base_version": 1, "added": [],
            "removed": [{"key": by["Reason for visit"]["key"], "label": "Reason for visit", "kind": "long_text", "origin": "automatic"}],
            "kind_changed": {by["Referring physician"]["key"]: ["text", "long_text"]},
            "renamed": [{"key": by["Allergies"]["key"], "from": "Allergies", "to": "Known allergies"}]})
        # Locators are the engine's, untouched; only label and kind change.
        self.assertEqual(new[by["Allergies"]["key"]]["locator"], by["Allergies"]["locator"])
        event = AuditEvent.objects.get(action="template.mapping_version", target=str(latest.pk))
        self.assertEqual(event.details["removed"], [by["Reason for visit"]["key"]])
        self.assertEqual(event.details["renamed"], [by["Allergies"]["key"]])
        self.version.refresh_from_db()
        self.assertEqual(self.version.fields, Template.objects.get(pk=self.template.pk).versions.get(number=1).fields)
        self.assertIsNone(latest.confirmed, "a corrected mapping needs its own confirmation")
        page = self.client.get(f"/templates/{self.template.pk}/").content.decode()
        self.assertIn("Corrections recorded for version 2", page)
        self.assertIn("Renamed: &ldquo;Allergies&rdquo; to &ldquo;Known allergies&rdquo;", page)
        self.assertIn("Removed (not an answer): Reason for visit", page)

    def test_only_text_like_kinds_change_into_one_another(self):
        by = self.by_label(self.version)
        response = self.post({"action": "v3_save", "kind_" + by["Laterality"]["key"]: "text"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("keeps its kind", response.context["v3_error"])
        response = self.post({"action": "v3_save", "kind_" + by["Referring physician"]["key"]: "choice"})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.template.versions.count(), 1)

    def test_a_signature_question_keeps_its_signature_label(self):
        by = self.by_label(self.version)
        response = self.post({"action": "v3_save", "label_" + by["Clinician signature"]["key"]: "Attending"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("keeps a signature label", response.context["v3_error"])
        self.assertEqual(self.template.versions.count(), 1)

    def test_a_missed_field_is_added_as_manual_in_document_order(self):
        response = self.post({"action": "v3_add", "slot": "word/document.xml#b/p[0]", "kind": "text", "label": "Clinic"})
        self.assertEqual(response.status_code, 302)
        latest = self.latest(self.template)
        added = latest.fields[0]
        self.assertEqual((added.get("label"), added.get("kind"), added.get("origin")), ("Clinic", "text", "manual"),
                         "the added paragraph comes first in the document, so it is field 1")
        self.assertTrue(added["key"].startswith("m3_"))
        self.assertEqual(latest.fields[1:], self.version.fields)
        self.assertEqual(latest.corrections, {"base_version": 1, "added": [added["key"]], "removed": [],
                                              "kind_changed": {}, "renamed": []})
        rows = self.client.get(f"/templates/{self.template.pk}/").context["v3_rows"]
        self.assertEqual(rows[0]["origin"], "Manual")

    def test_a_place_that_already_belongs_to_a_field_is_refused(self):
        response = self.post({"action": "v3_add", "slot": "word/document.xml#b/tbl[0]/tr[0]/tc[1]", "kind": "text",
                              "label": "Referrer again"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("already belongs to a field", response.context["v3_error"])
        self.assertEqual(self.template.versions.count(), 1)

    def test_printed_options_need_a_dropdown_place(self):
        response = self.post({"action": "v3_add", "slot": "word/document.xml#b/p[0]", "kind": "choice", "label": "Clinic"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("needs a Word dropdown", response.context["v3_error"])
        self.assertEqual(self.template.versions.count(), 1)

    def test_a_stale_page_saves_nothing(self):
        by = self.by_label(self.version)
        self.assertEqual(self.post({"action": "v3_save", "label_" + by["Allergies"]["key"]: "Known allergies"}).status_code, 302)
        stale = self.post({"action": "v3_save", "label_" + by["Allergies"]["key"]: "Allergy list"}, version=self.version)
        self.assertEqual(stale.status_code, 409)
        added = self.post({"action": "v3_add", "slot": "word/document.xml#b/p[0]", "kind": "text", "label": "Clinic"},
                          version=self.version)
        self.assertEqual(added.status_code, 409)
        self.assertEqual(self.template.versions.count(), 2)
        self.assertEqual(self.by_label(self.latest(self.template))["Known allergies"]["key"], by["Allergies"]["key"])


@override_settings(**V3)
class V3Confirmation(V3Fixture, TestCase):
    def setUp(self):
        super().setUp()
        self.template, self.version = self.uploaded()

    def test_an_unconfirmed_mapping_cannot_start_a_document(self):
        response = self.new_encounter(self.version)
        self.assertEqual(response.status_code, 200)
        errors = " ".join(response.context["form"].non_field_errors())
        self.assertIn("has not been confirmed", errors)
        self.assertFalse(Encounter.objects.exists())

    def test_a_clinician_confirms_and_who_and_when_are_recorded(self):
        before = timezone.now()
        version = self.confirm(self.template)
        self.assertEqual(version.confirmed.confirmed_by, self.user)
        self.assertGreaterEqual(version.confirmed.created_at, before)
        event = AuditEvent.objects.get(action="template.mapping_confirmed", target=str(version.pk))
        self.assertEqual(event.details["number"], 1)
        page = self.client.get(f"/templates/{self.template.pk}/").content.decode()
        self.assertIn("was confirmed by v3-clinician@synthetic.invalid", page)
        self.assertIn("Use this form", page)
        self.assertEqual(self.new_encounter(version).status_code, 302)
        self.assertEqual(Encounter.objects.get().template_version, version)

    def test_confirmation_needs_the_explicit_statement_and_a_clinician(self):
        response = self.client.post(f"/templates/{self.template.pk}/", {"action": "v3_confirm", "mapping_version": self.version.pk})
        self.assertEqual(response.status_code, 400)
        for other in (self.employee, self.manager):
            self.client.force_login(other)
            response = self.client.post(f"/templates/{self.template.pk}/", {"action": "v3_confirm",
                                        "mapping_version": self.version.pk, "checked_every_field": "on"})
            self.assertEqual(response.status_code, 403)
        self.assertFalse(MappingConfirmation.objects.exists())

    def test_a_corrected_mapping_needs_a_new_confirmation(self):
        self.confirm(self.template)
        by = self.by_label(self.version)
        self.client.post(f"/templates/{self.template.pk}/", {"action": "v3_save", "mapping_version": self.version.pk,
                                                             "label_" + by["Allergies"]["key"]: "Known allergies"})
        latest = self.latest(self.template)
        self.assertEqual(latest.number, 2)
        self.assertTrue(latest.needs_confirmation)
        self.assertEqual(self.new_encounter(latest).status_code, 200)
        self.assertEqual(self.new_encounter(self.version).status_code, 302, "the confirmed version stays usable")

    def test_a_mapping_with_no_fields_cannot_be_confirmed(self):
        removal = {"remove_" + f["key"]: "on" for f in self.version.fields}
        self.client.post(f"/templates/{self.template.pk}/", {"action": "v3_save", "mapping_version": self.version.pk, **removal})
        empty = self.latest(self.template)
        self.assertEqual((empty.number, empty.fields), (2, []))
        response = self.client.post(f"/templates/{self.template.pk}/", {"action": "v3_confirm", "mapping_version": empty.pk,
                                                                        "checked_every_field": "on"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("no fields", response.context["v3_error"])
        self.assertFalse(MappingConfirmation.objects.exists())

    def test_the_legacy_label_editor_cannot_bypass_the_recorded_review(self):
        labels = {f["key"]: f["label"] + " renamed" for f in self.version.fields if f["kind"] not in ("signature", "initials")}
        response = self.client.post(f"/templates/{self.template.pk}/", {"mapping_version": self.version.pk,
                                    "labels": __import__("json").dumps(labels)})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.template.versions.count(), 1)

    @override_settings(QUILLMD_WORD_ENGINE="legacy")
    def test_legacy_templates_need_no_confirmation(self):
        legacy = self.upload(data=build_legacy_table_form(), name="Synthetic legacy table")
        self.assertEqual(legacy.status_code, 302)
        version = Template.objects.get(name="Synthetic legacy table").versions.get()
        self.assertEqual(self.new_encounter(version, notes="Allergies: none documented").status_code, 302)
