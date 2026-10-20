"""Preparation SDK failures: private bounded evidence, closed metadata and one send."""
from contextlib import ExitStack, nullcontext
from copy import deepcopy
from decimal import Decimal
import hashlib
import io
import json
import os
from pathlib import Path
import tempfile
from types import SimpleNamespace
from unittest.mock import Mock, patch
import uuid

import httpx
from anthropic import Anthropic, APIResponseValidationError, BadRequestError
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.test import SimpleTestCase, TestCase, override_settings
from docx import Document

from core import preparation, preparation_error_evidence as diagnostics
from core.ai_stage_error_evidence import MAX_BODY_BYTES
from core.security import private_path
from documents.preparation import inventory, TYPED_PROTOCOL


SECRET = 'synthetic-non-key-shaped-private-credential'
CLINICAL = 'synthetic-private-clinical-marker'
MESSAGE = 'The compiled grammar is too large. Simplify the schema.'


def error_response(request, *, request_id='req_offline_error', body=None, status=400):
    return httpx.Response(status, request=request, headers={'request-id': request_id},
        json=body if body is not None else {'type': 'error', 'error': {
            'type': 'invalid_request_error', 'message': MESSAGE + ' ' + CLINICAL + ' ' + SECRET
                + '; Authorization: Bearer another-secret; x-api-key=third-secret'}})


class PreparationErrorEvidenceTests(SimpleTestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(); self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        settings = override_settings(MEDIA_ROOT=self.root, WORD_PREPARATION_ENABLED=True)
        settings.enable(); self.addCleanup(settings.disable)
        doc = Document(); table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Documented symptom history'
        doc.save(self.root / 'blank.docx')
        self.inventory = inventory(self.root / 'blank.docx', protocol=TYPED_PROTOCOL)
        self.config = {'mode': 'live', 'provider': 'anthropic', 'model': 'offline-configured',
            'input_rate': '2', 'output_rate': '10', 'cached_rate': '0.20'}

    def receipt(self):
        receipt = SimpleNamespace(pk=uuid.uuid4(), response_path='', response_sha256='',
            returned_model='', provider_request_id='', provider_resource_id='', computed_cost=None,
            usage={}, rates={}, failure={}, failure_path='', failure_sha256='', simulated=False,
            state='running', error_code='', quote_maximum=Decimal('0.10'))
        receipt.saves = []
        columns = ('state', 'error_code', 'failure', 'provider_request_id', 'failure_path',
            'failure_sha256', 'response_path', 'response_sha256', 'computed_cost')
        receipt.persisted = {key: deepcopy(getattr(receipt, key)) for key in columns}
        receipt.save_failures = set()
        def save(**kw):
            receipt.saves.append(deepcopy({'state': receipt.state, 'failure': receipt.failure,
                'request_id': receipt.provider_request_id, 'update_fields': kw.get('update_fields')}))
            if len(receipt.saves) in receipt.save_failures:
                raise RuntimeError('offline receipt update failure')
            # A single model UPDATE is atomic: columns are published together.
            changed = {key: deepcopy(getattr(receipt, key)) for key in (kw.get('update_fields') or columns)}
            receipt.persisted.update(changed)
        receipt.save = save
        return receipt

    def exercise(self, handler=None, *, strict=False, fault=None, provider_role=None, reserve=None,
            save_failures=()):
        receipt = self.receipt(); sent = []
        receipt.save_failures = set(save_failures)
        def handle(request):
            sent.append(json.loads(request.content))
            return (handler or error_response)(request)
        with ExitStack() as stack:
            http = stack.enter_context(httpx.Client(transport=httpx.MockTransport(handle)))
            client = Anthropic(api_key=SECRET, base_url='https://synthetic-provider.invalid',
                http_client=http, max_retries=0, timeout=60, _strict_response_validation=strict)
            stack.enter_context(patch.dict(os.environ, {'ANTHROPIC_API_KEY': SECRET}))
            stack.enter_context(patch('anthropic.Anthropic', return_value=client))
            stack.enter_context(patch('core.testing_authority.identity', return_value=True))
            reservation = stack.enter_context(patch('core.testing_authority.reserve_role',
                side_effect=reserve, return_value=receipt))
            stack.enter_context(patch('core.preparation_campaign.provider_slot', return_value=nullcontext()))
            stack.enter_context(patch.object(preparation.time, 'monotonic', return_value=100))
            if provider_role is not None:
                stack.enter_context(patch.object(preparation, 'provider_role', side_effect=provider_role))
            if fault is not None: stack.enter_context(fault)
            try:
                preparation.call_role(SimpleNamespace(inventory=self.inventory, config=self.config),
                    'proposer', 280)
            except BaseException as exc:
                return exc, receipt, sent, reservation
        self.fail('Expected an offline error outcome')

    def detail(self, receipt):
        pointer = receipt.failure[diagnostics.POINTER_KEY]
        self.assertEqual(set(pointer), {'contract', 'path', 'sha256'})
        self.assertEqual(pointer['contract'], diagnostics.CONTRACT)
        raw = private_path(pointer['path']).read_bytes()
        self.assertEqual(hashlib.sha256(raw).hexdigest(), pointer['sha256'])
        self.assertEqual(private_path(pointer['path']).stat().st_mode & 0o777, 0o600)
        self.assertLess(len(raw), MAX_BODY_BYTES)
        self.assertNotEqual(pointer['path'], receipt.failure_path)
        return json.loads(raw)['provider_evidence']

    def closed_metadata(self, receipt):
        raw = private_path(receipt.failure_path).read_bytes()
        self.assertEqual(hashlib.sha256(raw).hexdigest(), receipt.failure_sha256)
        self.assertEqual(json.loads(raw), receipt.failure)
        allowed = {'format', 'role', 'stage', 'error_class', 'http_status', 'provider_request_id',
            'provider_error_type', 'provider_stop_reason', 'elapsed_ms', 'response_retained',
            'usage_established', 'automatic_retry', 'provider_invoked', diagnostics.POINTER_KEY}
        self.assertLessEqual(set(receipt.failure), allowed)
        public = json.dumps(receipt.failure) + json.dumps(receipt.saves)
        public += preparation.failure_explanation(SimpleNamespace(state='blocked',
            error_code=receipt.error_code), [receipt])
        for value in (SECRET, CLINICAL, MESSAGE, 'another-secret', 'third-secret'):
            self.assertNotIn(value, public)

    def test_real_sdk_400_retains_private_cause_and_public_request_binding_without_retry(self):
        with patch.object(BadRequestError, '__str__', side_effect=AssertionError('never stringify')):
            error, receipt, sent, _ = self.exercise()
        self.assertIsInstance(error, BadRequestError)
        self.assertEqual(len(sent), 1)
        self.assertEqual(receipt.state, 'unknown'); self.assertIsNone(receipt.computed_cost)
        self.assertEqual(receipt.quote_maximum, Decimal('0.10'))
        self.assertEqual(receipt.provider_request_id, 'req_offline_error')
        self.assertEqual(receipt.failure['http_status'], 400)
        self.assertEqual(receipt.failure['provider_error_type'], 'invalid_request_error')
        self.assertFalse(receipt.failure['automatic_retry'])
        self.assertEqual(receipt.response_path, '')
        detail = self.detail(receipt)
        self.assertEqual(detail['body_source'], 'cached_http_response')
        self.assertIn(MESSAGE, detail['provider_error']['message'])
        self.assertIn(CLINICAL, detail['provider_error']['message'])
        self.assertEqual(detail['provider_request_id'], 'req_offline_error')
        self.assertIsNone(detail['actual_cost']); self.assertEqual(detail['delivery_outcome'], 'unknown')
        for secret in (SECRET, 'another-secret', 'third-secret'):
            self.assertNotIn(secret, json.dumps(detail))
        self.closed_metadata(receipt)

    def test_malformed_and_missing_error_bodies_keep_status_request_id_and_unknown_cost(self):
        for raw, expected in [(b'<html>not a provider JSON error</html>', 'omitted_non_json'),
                (b'[]', 'omitted_non_object'), (b'{"error":null}', 'omitted_unrecognized_error'),
                (b'{"error":{"type":[],"message":17}}', 'sanitized_error_projection')]:
            with self.subTest(raw=raw):
                def handle(request):
                    return httpx.Response(400, content=raw, request=request,
                        headers={'request-id': 'req_malformed'})
                error, receipt, sent, _ = self.exercise(handle)
                self.assertIsInstance(error, BadRequestError); self.assertEqual(len(sent), 1)
                self.assertEqual(self.detail(receipt)['body_retention'], expected)
                self.assertEqual(receipt.failure['provider_error_type'], 'unestablished')
                self.assertEqual(receipt.provider_request_id, 'req_malformed')
                self.assertIsNone(receipt.computed_cost); self.closed_metadata(receipt)

    def test_invalid_success_response_retains_actual_200_without_claiming_rejection_or_free_cost(self):
        def handle(request):
            return httpx.Response(200, request=request, headers={'request-id': 'req_invalid_success'},
                json={'type': 'message', 'content': 'invalid content type', 'usage': None})
        error, receipt, sent, _ = self.exercise(handle, strict=True)
        self.assertIsInstance(error, APIResponseValidationError); self.assertEqual(len(sent), 1)
        self.assertEqual(receipt.state, 'unknown'); self.assertIsNone(receipt.computed_cost)
        self.assertEqual(receipt.failure['http_status'], 200)
        self.assertEqual(receipt.provider_request_id, 'req_invalid_success')
        self.assertEqual(self.detail(receipt)['delivery_outcome'], 'unknown')
        self.closed_metadata(receipt)

    def test_oversize_cached_error_is_hashed_and_omitted_not_silently_truncated(self):
        body = {'error': {'type': 'invalid_request_error', 'message': 'x' * (MAX_BODY_BYTES + 1)}}
        error, receipt, sent, _ = self.exercise(lambda request: error_response(request, body=body))
        detail = self.detail(receipt)
        self.assertEqual(len(sent), 1); self.assertEqual(detail['body_retention'], 'omitted_oversize')
        self.assertGreater(detail['body_bytes'], MAX_BODY_BYTES)
        self.assertEqual(detail['body_sha256'], hashlib.sha256(error.response.content).hexdigest())
        self.assertIsNone(detail['provider_error']); self.closed_metadata(receipt)

    def test_actual_non_key_shaped_credential_cannot_leak_via_error_request_id(self):
        _, receipt, sent, _ = self.exercise(lambda request: error_response(request, request_id=SECRET))
        self.assertEqual(len(sent), 1); self.assertEqual(receipt.provider_request_id, '')
        self.assertEqual(receipt.failure['provider_request_id'], '')
        self.assertIsNone(self.detail(receipt)['provider_request_id'])
        self.closed_metadata(receipt)

    def test_cached_response_supplies_missing_sdk_metadata_without_reading_or_stringifying(self):
        original = Exception()
        original.response = error_response(httpx.Request('POST', 'https://synthetic-provider.invalid'),
            request_id='req_header_only')
        error, receipt, sent, _ = self.exercise(
            provider_role=lambda *a, **k: (_ for _ in ()).throw(original))
        self.assertIs(error, original); self.assertEqual(sent, [])
        self.assertIsNone(receipt.saves[0]['failure']['http_status'])
        self.assertEqual(receipt.saves[0]['request_id'], '')
        self.assertEqual(receipt.failure['http_status'], 400)
        self.assertEqual(receipt.provider_request_id, 'req_header_only')
        self.assertEqual(receipt.failure['provider_error_type'], 'invalid_request_error')
        self.assertEqual(self.detail(receipt)['provider_request_id'], 'req_header_only')
        self.closed_metadata(receipt)

    def test_private_projection_failure_leaves_initial_metadata_durable_and_original_error(self):
        error, receipt, sent, _ = self.exercise(fault=patch.object(diagnostics, 'build',
            side_effect=ValueError('private diagnostic failure')))
        self.assertIsInstance(error, BadRequestError); self.assertEqual(len(sent), 1)
        self.assertNotIn(diagnostics.POINTER_KEY, receipt.failure)
        self.assertEqual(receipt.saves[0]['failure']['http_status'], 400)
        self.assertEqual(receipt.saves[0]['request_id'], 'req_offline_error')
        self.assertIsNone(receipt.computed_cost); self.closed_metadata(receipt)

    def test_private_retention_failure_does_not_erase_metadata_or_replace_provider_exception(self):
        original = preparation.retain_response
        def retain(prep, role, value):
            if value.get('contract') == diagnostics.CONTRACT:
                raise OSError('offline diagnostic write failure')
            return original(prep, role, value)
        error, receipt, sent, _ = self.exercise(fault=patch.object(preparation, 'retain_response', side_effect=retain))
        self.assertIsInstance(error, BadRequestError); self.assertEqual(len(sent), 1)
        self.assertNotIn(diagnostics.POINTER_KEY, receipt.failure)
        self.assertEqual(receipt.failure['http_status'], 400)
        self.assertEqual(receipt.provider_request_id, 'req_offline_error')
        self.closed_metadata(receipt)

    def test_all_retention_failure_still_keeps_initial_closed_row_and_unknown_outcome(self):
        error, receipt, sent, _ = self.exercise(fault=patch.object(preparation, 'retain_response',
            side_effect=OSError('all writes failed')))
        self.assertIsInstance(error, BadRequestError); self.assertEqual(len(sent), 1)
        self.assertEqual(receipt.saves[0]['failure']['http_status'], 400)
        self.assertEqual(receipt.saves[0]['request_id'], 'req_offline_error')
        self.assertEqual(receipt.failure_path, ''); self.assertIsNone(receipt.computed_cost)

    def test_refused_diagnostic_update_keeps_durable_metadata_and_mirror_together(self):
        error, receipt, sent, _ = self.exercise(save_failures={2})
        self.assertIsInstance(error, BadRequestError); self.assertEqual(len(sent), 1)
        self.assertEqual(receipt.state, 'unknown'); self.assertIsNone(receipt.computed_cost)
        self.assertEqual(receipt.failure, receipt.persisted['failure'])
        self.assertNotIn(diagnostics.POINTER_KEY, receipt.failure)
        self.assertEqual(receipt.failure_path, receipt.persisted['failure_path'])
        self.assertEqual(receipt.failure_path, '')
        self.assertEqual(set(receipt.saves[1]['update_fields']),
            {'failure', 'provider_request_id', 'failure_path', 'failure_sha256'})
        self.assertEqual(receipt.persisted['failure']['http_status'], 400)

    def test_primary_metadata_write_failure_does_not_mask_original_or_claim_retention(self):
        error, receipt, sent, _ = self.exercise(save_failures={1})
        self.assertIsInstance(error, BadRequestError); self.assertEqual(len(sent), 1)
        self.assertEqual(receipt.persisted['failure'], {})
        self.assertEqual(receipt.failure_path, ''); self.assertIsNone(receipt.computed_cost)
        self.assertEqual(receipt.quote_maximum, Decimal('0.10'))

    def test_final_timestamp_write_failure_preserves_active_sdk_error(self):
        error, receipt, sent, _ = self.exercise(save_failures={3})
        self.assertIsInstance(error, BadRequestError); self.assertEqual(len(sent), 1)
        self.assertEqual(receipt.failure, receipt.persisted['failure'])
        self.assertEqual(receipt.failure_path, receipt.persisted['failure_path'])
        self.assertIsNone(receipt.computed_cost); self.detail(receipt); self.closed_metadata(receipt)

    def test_final_timestamp_write_failure_is_not_suppressed_after_success(self):
        response, evidence = preparation.provider_role(self.inventory, 'proposer',
            {**self.config, 'mode': 'fixture'}, 1)
        def success(*args, **kw):
            args[-1](evidence)
            return response, evidence
        error, receipt, sent, _ = self.exercise(provider_role=success, save_failures={4})
        self.assertIsInstance(error, RuntimeError); self.assertEqual(sent, [])
        self.assertEqual(str(error), 'offline receipt update failure')
        self.assertEqual(receipt.state, 'succeeded'); self.assertEqual(receipt.failure, {})

    def test_previous_original_response_and_request_binding_are_never_overwritten(self):
        response = error_response(httpx.Request('POST', 'https://synthetic-provider.invalid'))
        error = BadRequestError('unsafe exception text', response=response,
            body={'error': {'type': 'invalid_request_error'}})
        original = {'model': 'offline-model', 'request_id': 'req_original',
            'resource_id': 'msg_original', 'response': {'raw': 'original response bytes'}, 'cost': None}
        def fail(*args, **kw):
            args[-1](original)
            raise error
        caught, receipt, sent, _ = self.exercise(provider_role=fail)
        self.assertIs(caught, error); self.assertEqual(sent, [])
        self.assertEqual(receipt.state, 'failed')
        self.assertEqual(json.loads(private_path(receipt.response_path).read_bytes()), original)
        self.assertEqual(receipt.provider_request_id, 'req_original')
        self.assertEqual(receipt.failure['provider_request_id'], 'req_offline_error')
        self.assertEqual(self.detail(receipt)['provider_request_id'], 'req_offline_error')
        self.closed_metadata(receipt)

    def test_hostile_attributes_and_unread_response_cannot_mask_original_error_or_read_network(self):
        class Hostile(Exception):
            @property
            def code(self): raise SystemExit('unsafe metadata')
            @property
            def body(self): raise KeyboardInterrupt('unsafe body')
            @property
            def status_code(self): raise ValueError('unsafe status')
            @property
            def request_id(self): raise ValueError('unsafe id')
            def __str__(self): raise AssertionError('never stringify')
        original = Hostile()
        error, receipt, _, _ = self.exercise(provider_role=lambda *a, **k: (_ for _ in ()).throw(original))
        self.assertIs(error, original); self.assertEqual(receipt.state, 'unknown')
        self.assertIsNone(receipt.computed_cost); self.closed_metadata(receipt)
        class Unread:
            status_code = 400
            headers = {'request-id': 'req_unread'}
            @property
            def content(self): raise httpx.ResponseNotRead()
            def read(self): raise AssertionError('must not read')
            def json(self): raise AssertionError('must not parse response through API')
        original = Exception(); original.response = Unread()
        private = diagnostics.build(original, provider='anthropic', receipt_id='offline', role='proposer',
            stage='provider_role', metadata={})
        self.assertEqual(private['provider_evidence']['body_retention'], 'unavailable')
        self.assertEqual(private['provider_evidence']['provider_request_id'], 'req_unread')
        class HostileString(str):
            def __hash__(self): raise SystemExit('must not hash hostile metadata')
        original = Exception(); original.body = {'error': {'type': HostileString('invalid_request_error')}}
        metadata = preparation.safe_provider_failure(original, 'proposer', 'provider_role', 1, False, False)
        self.assertEqual(metadata['provider_error_type'], 'unestablished')

    def test_reservation_refusal_never_reaches_provider_or_captures_a_fake_response(self):
        refusal = ValidationError('This preparation role already has a retained attempt; it cannot be resent.')
        error, receipt, sent, _ = self.exercise(reserve=lambda *a, **k: (_ for _ in ()).throw(refusal))
        self.assertIs(error, refusal); self.assertEqual(sent, [])
        self.assertEqual(receipt.failure, {}); self.assertEqual(receipt.saves, [])


@override_settings(WORD_PREPARATION_ENABLED=True, ADMIN_SITE=False,
    TESTING_AUTHORITY_MODE='', TESTING_AUTHORITY_ID='', TESTING_AUTHORITY_SHA256='')
class PreparationErrorEvidenceDatabaseTests(TestCase):
    def setUp(self):
        from core.tests.test_semantic_preparation import DefinitionPreparationDatabaseTests
        DefinitionPreparationDatabaseTests.setUp(self)

    def test_failed_worker_is_not_automatically_rerun_and_original_unknown_receipt_stays_exact(self):
        from core.models import TemplatePreparation
        config = {'mode': 'live', 'provider': 'anthropic', 'model': 'offline-configured',
            'input_rate': '2', 'output_rate': '10', 'cached_rate': '0.20'}
        prep = TemplatePreparation.objects.create(template=self.template, requested_by=self.user, config=config)
        sent = []
        def handle(request):
            sent.append(request.url.path)
            return error_response(request)
        with httpx.Client(transport=httpx.MockTransport(handle)) as http:
            client = Anthropic(api_key=SECRET, base_url='https://synthetic-provider.invalid',
                http_client=http, max_retries=0, timeout=60)
            stdout, stderr = io.StringIO(), io.StringIO()
            with patch.dict(os.environ, {'ANTHROPIC_API_KEY': SECRET}), \
                    patch('anthropic.Anthropic', return_value=client), \
                    patch.object(BadRequestError, '__str__', side_effect=AssertionError('never stringify')), \
                    patch('core.management.commands.worker.signal.signal'), \
                    patch('core.management.commands.worker.mark_stale'), \
                    patch('core.management.commands.worker.reap_abandoned_voice_uploads'), \
                    patch('core.management.commands.worker.run_one', return_value=False), \
                    patch('core.semantic_preparation.run_one', return_value=False), \
                    patch('core.coding_service.run_one', return_value=False):
                call_command('worker', once=True, stdout=stdout, stderr=stderr)
                prep.refresh_from_db(); receipt = prep.receipts.get()
                before = deepcopy((receipt.state, receipt.quote_maximum, receipt.computed_cost,
                    receipt.provider_request_id, receipt.failure, receipt.failure_path, receipt.failure_sha256))
                call_command('worker', once=True, stdout=stdout, stderr=stderr)
        prep.refresh_from_db(); receipt.refresh_from_db()
        self.assertEqual(len(sent), 1); self.assertEqual(prep.state, 'blocked')
        self.assertEqual(prep.proposal, {}); self.assertEqual(prep.checker, {})
        self.assertEqual(receipt.state, 'unknown'); self.assertIsNone(receipt.computed_cost)
        self.assertEqual(receipt.provider_request_id, 'req_offline_error')
        self.assertEqual(before, (receipt.state, receipt.quote_maximum, receipt.computed_cost,
            receipt.provider_request_id, receipt.failure, receipt.failure_path, receipt.failure_sha256))
        self.assertFalse(self.template.versions.filter(preparation=prep).exists())
        self.assertEqual(stdout.getvalue(), ''); self.assertEqual(stderr.getvalue(), '')
