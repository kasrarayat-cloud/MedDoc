"""Complete multi-table forms: bounded planning and failure retention, no network."""
from contextlib import nullcontext
from copy import deepcopy
from decimal import Decimal
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
from unittest.mock import Mock, patch

from django.test import SimpleTestCase
from docx import Document

from core import preparation, word_contract
from documents.preparation import inventory, TYPED_PROTOCOL
from . import test_preparation_output as transport


class CompletePreparationTests(SimpleTestCase):
    receipt = transport.PreparationOutputTests.receipt
    complete_output = transport.PreparationOutputTests.complete_output
    exercise_transport = transport.PreparationOutputTests.exercise_transport

    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.path = Path(temporary.name) / 'complete-form.docx'
        document = Document()
        count = 0
        for rows in (18, 13, 6):
            table = document.add_table(rows=rows, cols=2)
            table.cell(0, 0).text = 'Question'
            table.cell(0, 1).text = 'Answer'
            for row in table.rows[1:]:
                count += 1
                row.cells[0].text = 'Documented consultation detail '+str(count)
        document.save(self.path)
        self.inventory = inventory(self.path, protocol=TYPED_PROTOCOL)
        self.config = {'mode': 'fixture', 'provider': 'anthropic', 'model': 'offline-configured-label',
            'input_rate': '3', 'output_rate': '15', 'cached_rate': '0.3',
            'word_preparation_output_profile': preparation.COMPLETE_OUTPUT_PROFILE,
            preparation.DESTINATION_KEY: preparation.DESTINATION_PROFILE,
            word_contract.CONFIG_KEY: word_contract.PROFILE}

    def test_whole_form_has_one_complete_proposal_and_one_independent_check(self):
        before = deepcopy(self.inventory)
        self.assertEqual(len(self.inventory['fields']), 34)
        maximum, cap = preparation.role_limits(self.inventory, self.config)
        self.assertGreater(maximum, 16000)
        self.assertLessEqual(maximum, 48000)
        self.assertEqual(cap, Decimal('0.10'))
        for role in ('proposer', 'checker'):
            result, receipt, retained, sent, reserve = self.exercise_transport(json.dumps(self.complete_output(role)), role=role)
            self.assertEqual(len(result['anchors']), 34)
            self.assertEqual(len(result['answer_definitions']['fields']), 34)
            self.assertEqual(len(result['stories']), len(self.inventory['stories']))
            self.assertEqual({r['definition']['permitted_destination']['field_key'] for r in result['answer_definitions']['fields']},
                             {f['key'] for f in self.inventory['fields']})
            self.assertEqual(sent[0]['max_tokens'], maximum)
            self.assertEqual(receipt.state, 'succeeded')
            self.assertEqual(len(sent), 1)
            self.assertEqual(len(retained), 1)
            self.assertEqual(reserve.call_args.args[-1], preparation.maximum_cost(sent[0], {**self.config, 'mode': 'live'}, maximum))
        self.assertEqual(self.inventory, before)

    def test_historical_profiles_and_unknown_values_keep_their_boundaries(self):
        legacy = {**self.config, 'word_preparation_output_profile': preparation.OUTPUT_PROFILE}
        with self.assertRaises(preparation.DocumentError) as rejected:
            preparation.role_limits(self.inventory, legacy)
        detail = rejected.exception.processing_limit
        self.assertEqual(detail['allowed'], 16000)
        self.assertEqual(detail['measured'], preparation.role_limits(self.inventory, self.config)[0])
        absent = {k:v for k,v in legacy.items() if k != 'word_preparation_output_profile'}
        self.assertEqual(preparation.role_limits(self.inventory, absent)[0], 2500)
        for value in (None, False, '', 'unknown'):
            with self.subTest(value=value), self.assertRaises(preparation.DocumentError):
                preparation.role_limits(self.inventory, {**self.config, 'word_preparation_output_profile': value})
        # Current profile selection never rewrites the original configuration.
        self.assertEqual(legacy['word_preparation_output_profile'], preparation.OUTPUT_PROFILE)

    def test_output_boundary_is_inclusive_and_next_rounded_step_refuses_before_reservation(self):
        maximum = preparation.role_limits(self.inventory, self.config)[0]
        self.assertEqual(maximum % 500, 0)
        with patch.object(preparation, 'MAX_COMPLETE_OUTPUT', maximum):
            self.assertEqual(preparation.role_limits(self.inventory, self.config)[0], maximum)
        with patch.object(preparation, 'MAX_COMPLETE_OUTPUT', maximum-500), \
                patch('anthropic.Anthropic') as client, patch('core.testing_authority.reserve_role') as reserve:
            with self.assertRaises(preparation.DocumentError) as rejected:
                preparation.call_role(SimpleNamespace(inventory=self.inventory, config=self.config), 'proposer', 999999999)
        self.assertEqual(rejected.exception.processing_limit, {'name': 'response_output_tokens',
            'measured': maximum, 'allowed': maximum-500, 'stage': 'output_planning'})
        client.assert_not_called(); reserve.assert_not_called()

    def test_request_boundary_covers_system_schema_original_and_checker_proposal(self):
        proposer = self.complete_output()
        for role in ('proposer', 'checker'):
            proposal = proposer if role == 'checker' else None
            envelope = preparation.provider_request(self.inventory, role, self.config, proposal)
            size = len(json.dumps(envelope, ensure_ascii=False).encode())
            self.assertLess(size, 512000)
            self.assertGreater(size, len(json.dumps(preparation.role_payload(self.inventory, role, proposal, config=self.config)).encode()))
            with patch.object(preparation, 'MAX_PREPARATION_REQUEST_BYTES', size):
                self.assertEqual(preparation.provider_request(self.inventory, role, self.config, proposal), envelope)
            with patch.object(preparation, 'MAX_PREPARATION_REQUEST_BYTES', size-1), \
                    patch('anthropic.Anthropic') as client, patch('core.testing_authority.reserve_role') as reserve:
                with self.assertRaises(preparation.DocumentError) as rejected:
                    preparation.call_role(SimpleNamespace(inventory=self.inventory, config=self.config), role, 999999999, proposal)
            self.assertEqual(rejected.exception.processing_limit, {'name': 'request_bytes', 'measured': size,
                'allowed': size-1, 'stage': role+'_request'})
            client.assert_not_called(); reserve.assert_not_called()

    def test_larger_forms_keep_bounded_time_and_historical_time_is_unchanged(self):
        self.assertEqual(preparation.timing_limits(self.inventory, self.config), (180, 480))
        self.assertEqual(preparation.timing_limits(self.inventory,
            {**self.config, 'word_preparation_output_profile': preparation.OUTPUT_PROFILE}), (120, 360))

    def test_unknown_usage_and_truncation_keep_paid_attempt_and_no_retry(self):
        complete = json.dumps(self.complete_output())
        for role in ('proposer', 'checker'):
            result, receipt, retained, sent, _ = self.exercise_transport(complete[:-20], 'max_tokens', role=role)
            self.assertEqual(result.code, 'output_limit')
            self.assertEqual(receipt.computed_cost, Decimal('0.0378'))
            self.assertEqual(receipt.state, 'failed')
            self.assertEqual(len(sent), 1)
            self.assertEqual(retained[0]['response']['content'][0]['text'], complete[:-20])
            self.assertFalse(receipt.failure['automatic_retry'])
        result, receipt, _, sent, _ = self.exercise_transport(complete, usage=False)
        self.assertEqual(result.code, 'provider_unknown')
        self.assertIsNone(receipt.computed_cost)
        self.assertEqual(receipt.quote_maximum, Decimal('0.10'))
        self.assertEqual(len(sent), 1)

    def worker(self, break_role, break_part):
        packet = SimpleNamespace(inventory={}, config={**self.config, 'word_preparation_protocol': TYPED_PROTOCOL},
            template=SimpleNamespace(sha256=self.inventory['source_sha256'], department='synthetic-department', original_path='original.docx'),
            requested_by='synthetic-doctor', state='queued', pk='synthetic-prep', save=Mock(), error_code='', proposal={}, checker={})
        manager = Mock()
        manager.select_for_update.return_value.filter.return_value.order_by.return_value.first.side_effect = [packet, None]
        manager.filter.return_value.exists.return_value = False
        receipts = []; retained = []; calls = []
        original_provider = preparation.provider_role
        def reserve(*args):
            receipt = self.receipt(); receipts.append(receipt); return receipt
        def provider(inv, role, config, timeout, proposal, observed, **kwargs):
            calls.append(role)
            response, evidence = original_provider(inv, role, config, timeout, proposal)
            if role == break_role:
                rows = response['answer_definitions']['fields'] if break_part == 'definitions' else response[break_part]
                rows[-1] = deepcopy(rows[0])
            observed(evidence)
            return response, evidence
        with patch.object(preparation, 'feature_enabled', return_value=True), \
                patch.object(preparation.transaction, 'atomic', side_effect=nullcontext), \
                patch.object(preparation.TemplatePreparation, 'objects', manager), \
                patch.object(preparation, 'membership'), patch.object(preparation, 'private_path', return_value=self.path), \
                patch.object(preparation, 'inspect_preparation', return_value=deepcopy(self.inventory)), \
                patch('core.testing_authority.identity', return_value=True), \
                patch('core.testing_authority.reserve_role', side_effect=reserve), \
                patch('core.preparation_campaign.provider_slot', return_value=nullcontext()), \
                patch.object(preparation, 'retain_response', side_effect=lambda *args: (retained.append(deepcopy(args[-1])) or ('retained.json', 'c'*64))), \
                patch.object(preparation, 'provider_role', side_effect=provider), \
                patch.object(preparation, 'build_candidate') as build, patch.object(preparation, 'audit') as audit:
            self.assertTrue(preparation.run_one())
            self.assertFalse(preparation.run_one())
        build.assert_not_called()
        self.assertEqual(packet.state, 'blocked')
        self.assertEqual(packet.error_code, 'incomplete_proposal' if break_part == 'stories' else 'invalid_proposal')
        self.assertEqual(calls, ['proposer'] if break_role == 'proposer' else ['proposer', 'checker'])
        self.assertTrue(all(r.computed_cost == 0 for r in receipts))
        self.assertEqual(receipts[-1].state, 'failed')
        self.assertEqual(receipts[-1].failure['stage'], 'proposal_validation')
        self.assertGreaterEqual(len(retained), len(receipts))
        self.assertEqual(audit.call_count, 1)
        if break_role == 'checker':
            self.assertEqual(receipts[0].state, 'succeeded')
            self.assertEqual(len(packet.proposal['anchors']), 34)

    def test_missing_or_duplicate_coverage_never_produces_candidate_or_resends(self):
        for role in ('proposer', 'checker'):
            for part in ('anchors', 'stories', 'definitions'):
                with self.subTest(role=role, part=part):
                    self.worker(role, part)

    def test_historical_limit_explanation_replays_only_local_planner_and_has_no_content(self):
        packet = SimpleNamespace(state='blocked', error_code='preparation_limit', inventory=self.inventory,
            config={**self.config, 'word_preparation_output_profile': preparation.OUTPUT_PROFILE})
        before = deepcopy(packet.__dict__)
        with patch.object(preparation, 'provider_role') as provider:
            message = preparation.failure_explanation(packet, [])
        self.assertIn('16,000', message)
        self.assertIn('complete-response allowance', message)
        self.assertIn('unchanged original', message)
        self.assertNotIn('Documented consultation detail', message)
        self.assertEqual(packet.__dict__, before)
        provider.assert_not_called()
        for bad in ({'secret': 'do not expose'}, {'name':'request_bytes','measured':True,'allowed':0,'stage':'proposer_request'}):
            self.assertIsNone(preparation.validated_processing_limit(bad))
