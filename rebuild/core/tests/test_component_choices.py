"""Source-owned per-component enums; constructed fixtures, no provider requests."""
from copy import deepcopy
import hashlib
import json
from importlib import import_module
from types import SimpleNamespace as N
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.db import DatabaseError, connection, transaction
from django.template.loader import render_to_string
from django.test import SimpleTestCase, TestCase, override_settings

from documents import component_contracts as cc, semantic_preparation_roles as roles
from core import ai_stages, component_review as editor, source_assertion_generation as generation
from core import semantic_preparation as workflow, word_contract
from core.tests import test_source_assertion_graph as fx
from core.tests import test_component_encounter_scope as scope_fx
from core.tests import test_component_review as review_fx
from core.tests import test_definition_capacity as capacity_fx


def choice_definition(label='Visit mode: office or telephone. Encounter details.', options=('office','telephone'), quote='office or telephone'):
    fields, values = fx.definitions((label,))
    field, definition = fields[0], values['field0']
    definition['component_contract'] = cc.fixture_for(field, definition, cc.CHOICE_CONTRACT)
    part = definition['component_contract']['components'][0]
    part.update(id='mode', question_quote=quote, option_quote=quote, value_type='choice', choices=list(options),
                meaning='The documented visit mode.', scope={'subject':'patient','temporality':'current','polarity':'positive'})
    other = deepcopy(part)
    other.update(id='details', question_quote='Encounter details', option_quote=None, value_type='text', choices=[],
                 meaning='The complete documented encounter details.')
    definition['component_contract']['components'].append(other)
    return field, definition


def choice_case(contract=cc.CHOICE_CONTRACT, *, subject='patient', encounter='current', timing='current', certainty='established', value='office'):
    predicate = ('Her sister documented '+value+' as the visit mode.' if subject == 'other' else
                 'If symptoms persist, a '+value+' visit is planned.' if certainty == 'conditional' else
                 'The documented visit mode was '+value+'.')
    notes='Patient: Synthetic Ada Vale\n'+predicate+'\nThe clinician documented a review.'
    field, definition=choice_definition()
    definition['component_contract']['contract']=contract
    if contract != cc.CHOICE_CONTRACT:
        for part in definition['component_contract']['components']:
            del part['choices'];del part['option_quote']
    facts=[fx.fact(notes,'f0','Synthetic Ada Vale',kind='identity'),
           fx.fact(notes,'f1',value,subject=subject,timing=timing),
           fx.fact(notes,'f2','The clinician documented a review.')]
    assertions=[fx.assertion(notes,'a0',['f0'],'Synthetic Ada Vale',kind='identity',identity_role='patient_name',consumer_role='none'),
        fx.assertion(notes,'a1',['f1'],predicate,kind='context',consumer_role='none',
            subject=subject,encounter=encounter,temporality=timing,certainty=certainty),
        fx.assertion(notes,'a2',['f2'],'The clinician documented a review.',kind='context',consumer_role='none')]
    if subject == 'other':facts[1]['subject']['name_span']=fx.ref(notes,'Her sister')
    if certainty == 'conditional':assertions[1]['qualifier_refs']=[fx.ref(notes,'If symptoms persist')]
    assigned={'fields':[{'key':field['key'],'separator':'; ','components':[
        {'id':identifier,'disposition':'supported','fact_ids':[fact], 'predicate_assertion_ids':[root],'relationship_ids':[]}
        for identifier,fact,root in [('mode','f1','a1'),('details','f2','a2')]]}]}
    return [notes,[field],{field['key']:definition},fx.extraction(notes,facts,assertions),assigned]


def choice_post(prep, field, definition, actor):
    post=review_fx.editor_post(prep,field,definition,actor)
    for part in editor.parts_for(field,definition):
        post['part_'+str(part['index'])+'_choices']=part['choices_text']
    return post


@override_settings(STORAGES={'default':{'BACKEND':'django.core.files.storage.FileSystemStorage'},
                            'staticfiles':{'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class ComponentChoiceTests(SimpleTestCase):
    def test_retained_v1_v2_complete_role_envelopes_have_exact_pre_v3_bytes(self):
        fields=[{'key':'topic','kind':'text','label':'Complete topic','locator':{'cell':0}}]
        inventory=roles.build_inventory(fields,source_version_id='synthetic-legacy-source',source_sha256='a'*64,
                                        original={'questions':['Complete topic']})
        expected={cc.CONTRACT:'9f60ffb0d0efb470203ba49680856d7baf3b90b3f5618e8ef328422122eb3e4d',
                  cc.ENCOUNTER_CONTRACT:'d690b098bd96f1bfe8eb3bf4f82b504fac7c46bdd79e3b16ae0fb0e000d50c5b'}
        for tag,sha in expected.items():
            envelope={role:{'schema':roles.role_schema(role,4,component_contract=tag),
                'instructions':roles.role_instructions(role,4,component_contract=tag),
                'fixture':roles.fixture_role(fields,inventory,role,definition_version=4,component_contract=tag)}
                for role in ('proposer','checker')}
            encoded=json.dumps(envelope,ensure_ascii=False,separators=(',',':')).encode()
            self.assertEqual(hashlib.sha256(encoded).hexdigest(),sha)

    def test_current_word_roles_and_sealed_shared_readers_preserve_own_options(self):
        from core.tests import test_word_answer_contract as word_fx
        from core import preparation,semantic_review
        from core.template_source import approval_binding,source_for_version
        fixture=word_fx.FreshWordContractTests();fixture.setUp();self.addCleanup(fixture.tearDown)
        old=fixture.graph_config(); config={**old,cc.CONFIG_KEY:cc.CHOICE_CONTRACT}
        original=deepcopy((fixture.inv,old))
        proposal=None
        for role in ('proposer','checker'):
            response,evidence=preparation.provider_role(fixture.inv,role,config,1,proposal)
            word_contract.validate_result(fixture.inv,response,role,config)
            self.assertTrue(evidence['simulated'])
            for row in response['answer_definitions']['fields']:
                d=row['definition'];self.assertEqual(d['component_contract']['contract'],cc.CHOICE_CONTRACT)
                for part in d['component_contract']['components']:
                    self.assertEqual(part['choices'],d['choices'] if part['value_type']=='choice' else [])
                    self.assertIsNone(part['option_quote'])
            proposal=response
        version=fixture.version();prep=version.preparation
        prep.config=config
        prep.proposal,_=preparation.provider_role(fixture.inv,'proposer',config,1)
        prep.checker,_=preparation.provider_role(fixture.inv,'checker',config,1,prep.proposal)
        prep.candidate_info.pop(word_contract.INFO_KEY)
        values=word_contract.proposed(prep)
        word_contract.seal(prep,list(values),N(pk='doctor'));prep.approval_sha256=approval_binding(prep)
        self.assertEqual(semantic_review.definitions_for(version),values)
        self.assertEqual(word_contract.require_generation(version),values)
        self.assertEqual(source_for_version(version),fixture.root/'prepared.docx')
        self.assertEqual((fixture.inv,old),original)

    def test_full_graph_reproduces_unfillable_old_choice_and_new_own_options_round_trip(self):
        old=choice_case(cc.ENCOUNTER_CONTRACT)
        old_result,_=fx.run_case(old)
        answer=old_result['answers']['field0']
        self.assertNotEqual(answer['status'],'supported')
        self.assertIn('component_value_type',[r['code'] for r in answer['route_diagnostics']['holds']])
        case=choice_case(); result,calls=fx.run_case(case)
        answer=result['answers']['field0']
        self.assertEqual(answer['status'],'supported')
        self.assertEqual(answer['text'],'office; The clinician documented a review.')
        self.assertEqual(answer['components'][0]['value'],'office')
        self.assertEqual(len(calls),3)
        self.assertEqual(case[2]['field0']['choices'],[])
        self.assertEqual(generation.restore(case[0],result['interpretation']),result['answers'])
        generation.validate_stored_answer(case[0],case[1][0],case[2]['field0'],answer,result['interpretation'])
        altered=deepcopy(result['interpretation'])
        altered['definitions']['field0']['component_contract']['components'][0]['choices'].reverse()
        from core import source_assertion_graph as graph
        altered['sha256']=graph.digest({k:v for k,v in altered.items() if k!='sha256'})
        with self.assertRaises(ValidationError): generation.restore(case[0],altered)

    def test_enum_does_not_waive_subject_encounter_event_time_or_checker_evidence(self):
        for changes in ({'subject':'other'},{'encounter':'past','timing':'past'},{'encounter':'unknown'},
                        {'timing':'future'},{'certainty':'conditional'}):
            case=choice_case(**changes)
            with self.subTest(changes=changes):
                result,_=fx.run_case(case)
                self.assertNotEqual(result['answers']['field0']['status'],'supported')
        case=choice_case()
        def unknown(check,*_):
            check['components'][0].update(verdict='unknown',cause='interpretation',reason='Mode is not established.')
            check['fields'][0].update(verdict='unknown',cause='interpretation',reason='Mode is not established.')
        result,_=fx.run_case(case,unknown)
        self.assertNotEqual(result['answers']['field0']['status'],'supported')

    def test_v3_case_and_whitespace_choose_one_original_value_without_rewriting_literal_text(self):
        for value in ('office','Office','OFFICE','  Office\t'):
            case=choice_case(value=value);result,_=fx.run_case(case)
            answer=result['answers']['field0']
            self.assertEqual(answer['status'],'supported')
            self.assertEqual(answer['components'][0]['text'],value)
            self.assertEqual(answer['components'][0]['value'],'office')
            self.assertEqual(generation.restore(case[0],result['interpretation']),result['answers'])
        case=choice_case(value='Office   visit')
        field,d=choice_definition('Visit mode: office visit or telephone visit. Encounter details.',
                                  ['office visit','telephone visit'],'office visit or telephone visit')
        case[1],case[2]=[field],{field['key']:d}
        result,_=fx.run_case(case)
        self.assertEqual(result['answers']['field0']['components'][0]['value'],'office visit')
        self.assertEqual(result['answers']['field0']['components'][0]['text'],'Office   visit')
        for value in ('clinic','in-office','office!'):
            result,_=fx.run_case(choice_case(value=value))
            self.assertNotEqual(result['answers']['field0']['status'],'supported')
        result,_=fx.run_case(choice_case(value='Office',subject='other'))
        self.assertNotEqual(result['answers']['field0']['status'],'supported')
        old=choice_case(cc.ENCOUNTER_CONTRACT,value='Office');result,_=fx.run_case(old)
        self.assertNotEqual(result['answers']['field0']['status'],'supported')

    def test_complete_ordered_or_lists_and_physical_dropdowns_are_source_bound(self):
        for expression,options in [('office or telephone',['office','telephone']),
            ('red, blue or green',['red','blue','green']),('red, blue, or green',['red','blue','green'])]:
            field,d=choice_definition('Choose '+expression+'. Encounter details.',options,expression)
            self.assertIs(cc.validate(field,d),d['component_contract'])
            self.assertEqual(d['component_contract']['components'][0]['choices'],options)
        # A conjunction between distinct requested parts does not belong to the
        # options owned by the complete final component question.
        field,d=choice_definition('Record visit modality and new or established patient status.',
                                  ['new','established'],'new or established patient status')
        d['component_contract']['components'][0]['option_quote']='new or established'
        d['component_contract']['components'][1]['question_quote']='visit modality'
        self.assertIs(cc.validate(field,d),d['component_contract'])
        comma=deepcopy(field);comma['label']=field['label'].replace(' and new',', new')
        self.assertIs(cc.validate(comma,d),d['component_contract'])
        bad=deepcopy(d);bad['component_contract']['components'][0]['option_quote']='new or established patient status'
        with self.assertRaises(roles.DefinitionError):cc.validate(field,bad)
        field,d=choice_definition()
        field['answer_constraints']={'data_type':'choice','options':['In person','Video, audio','Other / specify']}
        d.update(value_type='choice',answer_mode='atomic',choices=deepcopy(field['answer_constraints']['options']))
        d['component_contract']=cc.fixture_for(field,d,cc.CHOICE_CONTRACT)
        self.assertIs(cc.validate(field,d),d['component_contract'])
        bad=deepcopy(d);bad['component_contract']['components'][0]['choices'].reverse()
        with self.assertRaises(roles.DefinitionError):cc.validate(field,bad)

    def test_partial_reordered_borrowed_and_ambiguous_options_never_become_source_owned(self):
        cases=[('red, blue or green','blue or green',['blue','green']),
            ('red or blue or green','red or blue',['red','blue']),
            ('red; blue or green','blue or green',['blue','green']),
            ('red and blue or green','blue or green',['blue','green']),
            ('red / blue or green','blue or green',['blue','green']),
            ('red, either blue or green','blue or green',['blue','green']),
            ('red or blue','red or blue',['blue','red']),
            ('red or blue','red or blue',['red','green']),
            ('red and blue','red and blue',['red','blue']),
            ('red / blue','red / blue',['red','blue']),
            ('red; blue','red; blue',['red','blue']),
            ('Yes or yes','Yes or yes',['Yes','yes']),
            ('red or  red','red or  red',['red',' red']),
            ('red or blue and green','red or blue',['red','blue']),
            ('red or blue; green','red or blue',['red','blue']),
            ('red or blue and red or blue','red or blue',['red','blue']),
            ('red, blue or green','red, blue or green',['red, blue','green'])]
        for source,quote,options in cases:
            field,d=choice_definition('Choose '+source+'. Encounter details.',options,quote)
            with self.subTest(source=source,options=options),self.assertRaises(roles.DefinitionError):cc.validate(field,d)
        # A noun suffix on the owner cannot hide an original earlier option.
        for separator in (',',';',' /',' and'):
            field,d=choice_definition('Select low'+separator+' medium or high acuity. Encounter details.',
                                      ['medium','high'],'medium or high acuity')
            d['component_contract']['components'][0]['option_quote']='medium or high'
            with self.subTest(separator=separator),self.assertRaises(roles.DefinitionError):
                roles.validate_definitions([field],{field['key']:d})
        field,d=choice_definition('Select low, medium or high acuity. Encounter details.',
                                  ['low','medium','high'],'Select low, medium or high acuity.')
        d['component_contract']['components'][0]['option_quote']='low, medium or high'
        self.assertIs(cc.validate(field,d),d['component_contract'])
        field,d=choice_definition()
        for mutate in (lambda p:p.pop('choices'),lambda p:p.update(choices=[]),lambda p:p.update(choices=['office','office']),
                       lambda p:p.update(value_type='text'),lambda p:p.update(encounter_scope='unknown'),
                       lambda p:p.update(question_quote='fice or telephone')):
            bad=deepcopy(d);mutate(bad['component_contract']['components'][0])
            with self.assertRaises(roles.DefinitionError):cc.validate(field,bad)

    def test_both_independent_preparation_roles_have_explicit_options_and_hierarchy_obligations(self):
        for role in ('proposer','checker'):
            schema=roles.role_schema(role,4,component_contract=cc.CHOICE_CONTRACT)
            part=schema['properties']['fields']['items']['properties']['definition']['properties']['component_contract']['properties']['components']['items']
            self.assertIn('choices',part['required']);self.assertIn('option_quote',part['required'])
            self.assertFalse(part['additionalProperties'])
            text=roles.role_instructions(role,4,component_contract=cc.CHOICE_CONTRACT)
            for words in ('complete exact original options','original order','topic introduction is their context',
                          'independently requested narrative','any real remaining omission, disagreement'):
                self.assertIn(words,text)
        case=choice_case()
        self.assertIn('own approved complete ordered choices',generation.instructions_for(case[2],generation.CHECK_INSTRUCTIONS))

    def test_v3_editor_posts_own_options_with_source_destination_and_review_token_binding(self):
        prep,_,actor=review_fx.ComponentReviewPureTests().make()
        field,d=choice_definition();prep.source_version.fields=[field]
        post=choice_post(prep,field,d,actor)
        self.assertEqual(editor.posted_definition(prep,field,d,post),d)
        self.assertEqual(editor.parts_for(field,d)[-1]['choices'],[])
        for mutate in (lambda p:p.pop('part_0_choices'),lambda p:p.appendlist('part_0_choices','office'),
                       lambda p:p.update(part_0_choices='telephone\noffice')):
            changed=post.copy();mutate(changed)
            with self.subTest(mutate=mutate):
                try: value=editor.posted_definition(prep,field,d,changed)
                except ValidationError:continue
                with self.assertRaises(roles.DefinitionError):roles.validate_definitions([field],{field['key']:value})
        row={'display_label':field['label'],'definition':d,'field':field}
        html=render_to_string('template_component_review.html',{'template':N(pk='00000000-0000-4000-8000-000000000001',name='Synthetic'),
            'preparation':N(pk='00000000-0000-4000-8000-000000000002'),'row':row,'parts':editor.parts_for(field,d),
            'user':N(is_authenticated=False,first_name='',username='synthetic'),
            'editable':True,'component_choices':True,'definition_version':4})
        self.assertIn('name="part_0_choices"',html);self.assertIn('office\ntelephone',html)
        self.assertIn('name="part_0_option_quote"',html)

    def test_v3_nonchoice_inherits_every_encounter_and_literal_qualifier_rule(self):
        for changes,expected in [({},'supported'),({'subject':'other'},None),({'encounter':'unknown'},None),({'polarity':'affirmed'},None)]:
            case=scope_fx.topic_case(**changes)
            d=case[2]['field0'];d['component_contract']['contract']=cc.CHOICE_CONTRACT
            for c in d['component_contract']['components']:c.update(choices=[],option_quote=None)
            result,_=fx.run_case(case)
            if expected:self.assertEqual(result['answers']['field0']['status'],expected)
            else:self.assertNotEqual(result['answers']['field0']['status'],'supported')


@override_settings(WORD_PREPARATION_ENABLED=True,AI_MODE='fixture',TESTING_AUTHORITY_MODE='',TESTING_AUTHORITY_ID='',TESTING_AUTHORITY_SHA256='',
    STORAGES={'default':{'BACKEND':'django.core.files.storage.FileSystemStorage'},
              'staticfiles':{'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class ComponentChoiceCapacityTests(TestCase):
    setUp=capacity_fx.DefinitionCapacityDatabaseTests.setUp
    parent=capacity_fx.DefinitionCapacityDatabaseTests.parent
    reserve=capacity_fx.DefinitionCapacityDatabaseTests.reserve

    def test_real_insert_guard_accepts_only_v2_or_v3_and_preserves_numeric_caps(self):
        ai_stages.assert_installed_guards()
        current=ai_stages.definition_pipeline_config({**capacity_fx.CONFIG,'mode':'fixture','semantic_component_contract':cc.CHOICE_CONTRACT})
        self.assertEqual(ai_stages.definition_limits(current),(32000,240))
        self.assertEqual(current['definition_max_calls'],2);self.assertEqual(current['definition_request_byte_limit'],512000)
        for tag in (cc.ENCOUNTER_CONTRACT,cc.CHOICE_CONTRACT):
            with transaction.atomic():
                self.reserve({**current,'semantic_component_contract':tag});transaction.set_rollback(True)
        variants=[{k:v for k,v in current.items() if k!='semantic_component_contract'}]
        variants += [{**current,'semantic_component_contract':v}for v in (None,False,4,[],{},'unknown',cc.CONTRACT)]
        variants += [{**current,k:v}for k,v in [('definition_max_calls',3),('definition_output_token_limit',32001),
            ('definition_timeout_seconds',241),('definition_request_byte_limit',512001),('definition_output_token_limit','32000'),
            ('semantic_definition_version','4'),('provider','openai'),('model','other')]]
        for bad in variants:
            with self.subTest(config=bad),self.assertRaises(DatabaseError),transaction.atomic():self.reserve(bad)

    def test_exact_single_sql_delta_reverse_then_forward_and_legacy_receipts(self):
        m=import_module('core.migrations.0036_component_choice_capacity')
        previous=import_module('core.migrations.0035_definition_capacity_profile').CAPACITY_SQL
        self.assertEqual(m.PREVIOUS_SQL,previous)
        self.assertEqual(m.CAPACITY_SQL.replace(m.NEW_CONDITION,m.OLD_CONDITION),previous)
        self.assertEqual(m.Migration.operations[0].reverse_sql,previous)
        with transaction.atomic():
            with connection.cursor()as cursor:cursor.execute(m.PREVIOUS_SQL)
            with self.assertRaises(ai_stages.StageNotSent):ai_stages.assert_installed_guards()
            current=ai_stages.definition_pipeline_config({**capacity_fx.CONFIG,'mode':'fixture','semantic_component_contract':cc.CHOICE_CONTRACT})
            with self.assertRaises(DatabaseError),transaction.atomic():self.reserve(current)
            with transaction.atomic():
                self.reserve({**current,'semantic_component_contract':cc.ENCOUNTER_CONTRACT})
                transaction.set_rollback(True)
            with connection.cursor()as cursor:cursor.execute(m.CAPACITY_SQL)
            ai_stages.assert_installed_guards()
            self.reserve(current)

    def test_new_definition_queue_has_v3_but_keeps_prior_requests_immutable(self):
        prep=workflow.create_preparation(self.source,self.user)
        self.assertEqual(prep.config[cc.CONFIG_KEY],cc.CHOICE_CONTRACT)
        self.assertTrue(workflow.run_one());prep.refresh_from_db()
        for response in (prep.proposal,prep.checker):
            self.assertTrue(all(row['definition']['component_contract']['contract']==cc.CHOICE_CONTRACT for row in response['fields']))
        self.assertEqual(prep.ai_stage_receipts.count(),2)

    def test_real_editor_choice_draft_whole_form_approval_render_and_word_fill(self):
        from docx import Document
        from django.urls import reverse
        from documents import inspect_template,fill_template
        from core.models import Template,TemplateVersion
        from core.semantic_review import definitions_for
        from core.template_source import source_for_version
        path=self.root/'choice-blank.docx'
        document=Document();cells=document.add_table(rows=1,cols=2).rows[0].cells
        cells[0].text='Visit mode: office or telephone. Encounter details.';document.save(path)
        info=inspect_template(path)
        template=Template.objects.create(department=self.department,name='Synthetic choice form',format='docx',
            original_name=path.name,original_path=path.name,sha256=hashlib.sha256(path.read_bytes()).hexdigest(),created_by=self.user)
        source=TemplateVersion.objects.create(template=template,number=1,fields=info['fields'],created_by=self.user)
        prep=workflow.create_preparation(source,self.user);self.assertTrue(workflow.run_one());prep.refresh_from_db()
        original=deepcopy((prep.proposal,prep.checker,source.fields,path.read_bytes()))
        field=source.fields[0];current=workflow.corrected_definitions(prep)[field['key']]
        post=choice_post(prep,field,current,self.user)
        controls=dict(part_0_question_quote='Visit mode: office or telephone',part_0_value_type='choice',
            part_0_meaning='The documented visit mode.',part_0_choices='office\ntelephone',part_0_option_quote='office or telephone',
            part_0_polarity='positive',part_0_temporality='current',part_1_add='on',part_1_question_quote='Encounter details',
            part_1_meaning='The complete documented encounter details.',part_1_encounter_scope='current',
            part_1_temporality='current',part_1_polarity='positive')
        for key,value in controls.items():post[key]=value
        self.client.force_login(self.user)
        url=reverse('template_definitions_components',args=[prep.pk,field['key']])
        response=self.client.get(url);self.assertContains(response,'name="part_0_option_quote"')
        response=self.client.post(url,post.urlencode(),content_type='application/x-www-form-urlencoded')
        self.assertEqual(response.status_code,302,response.context['error'] if response.context else response.content[:500]);prep.refresh_from_db()
        values=workflow.corrected_definitions(prep)
        self.assertEqual(values[field['key']]['component_contract']['components'][0]['choices'],['office','telephone'])
        with self.assertRaises(ValidationError):workflow.approve(prep.pk,self.user,prep.revision,prep.inventory['inventory_sha256'],list(values),True)
        workflow.correct(prep.pk,self.user,prep.revision,values,list(values),'Compared each part and every choice with the original source form.')
        prep.refresh_from_db();version=workflow.approve(prep.pk,self.user,prep.revision,prep.inventory['inventory_sha256'],list(values),True)
        self.assertEqual(definitions_for(version),values)
        case=choice_case();case[1]=source.fields;case[2]=values
        case[4]['fields'][0]['key']=field['key']
        for selected,part in zip(case[4]['fields'][0]['components'],values[field['key']]['component_contract']['components']):selected['id']=part['id']
        result,_=fx.run_case(case)
        self.assertEqual(result['answers'][field['key']]['status'],'supported')
        output=self.root/'choice-filled.docx'
        fill_template(source_for_version(version),version.fields,{field['key']:result['answers'][field['key']]['text']},output)
        self.assertIn('office; The clinician documented a review.',Document(output).tables[0].cell(0,1).text)
        prep.refresh_from_db();source.refresh_from_db()
        self.assertEqual((prep.proposal,prep.checker,source.fields,path.read_bytes()),original)

    def test_atomic_physical_choice_canonicalizes_only_equivalent_source_before_exact_word_guard(self):
        from docx import Document
        from documents import inspect_template,fill_template
        from documents.table_constraints import with_table_constraints
        path=self.root/'printed-choice.docx'
        document=Document();cells=document.add_table(rows=1,cols=2).rows[0].cells
        cells[0].text='Visit mode (select one: office / telephone)';document.save(path)
        fields=with_table_constraints(path,inspect_template(path)['fields'])
        self.assertEqual(fields[0]['answer_constraints']['options'],['office','telephone'])
        definitions=roles.default_definitions(fields);d=definitions[fields[0]['key']]
        d.update(version=4,answer_mode='atomic',patient_name_requirement='none',
                 scope={'subject':'patient','temporality':'current','polarity':'positive'})
        d['component_contract']=cc.fixture_for(fields[0],d,cc.CHOICE_CONTRACT)
        def run(value,contract=cc.CHOICE_CONTRACT):
            case=choice_case(value=value);case[1]=fields;case[2]=deepcopy(definitions)
            if contract == cc.ENCOUNTER_CONTRACT:
                c=case[2][fields[0]['key']]['component_contract'];c['contract']=contract
                del c['components'][0]['choices'];del c['components'][0]['option_quote']
            case[4]['fields'][0]['key']=fields[0]['key'];case[4]['fields'][0]['components']=case[4]['fields'][0]['components'][:1]
            case[4]['fields'][0]['components'][0]['id']='answer'
            return case,fx.run_case(case)[0]
        for value in ('Office','OFFICE','  Office\t'):
            case,result=run(value);answer=result['answers'][fields[0]['key']]
            self.assertEqual(answer['status'],'supported');self.assertEqual(answer['text'],'office')
            self.assertEqual(answer['value'],'office');self.assertEqual(answer['components'][0]['value'],'office')
            self.assertEqual(answer['facts'][0]['value_evidence']['quote'],value)
            self.assertTrue(any(value in e['quote'] for e in answer['evidence']))
            self.assertEqual(generation.restore(case[0],result['interpretation']),result['answers'])
        output=self.root/'printed-choice-filled.docx'
        fill_template(path,fields,{fields[0]['key']:answer['text']},output)
        self.assertEqual(Document(output).tables[0].cell(0,1).text,'office')
        with self.assertRaises(ValidationError):generation.validate_value(d,'Office',fields[0])
        for value in ('clinic','office; telephone'):
            _,result=run(value);self.assertNotEqual(result['answers'][fields[0]['key']]['status'],'supported')
        _,result=run('Office',cc.ENCOUNTER_CONTRACT)
        self.assertNotEqual(result['answers'][fields[0]['key']]['status'],'supported')
