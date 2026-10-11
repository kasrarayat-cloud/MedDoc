"""No database/provider calls: source-bound metadata and new approval versions."""
from copy import deepcopy
import hashlib
from pathlib import Path
import tempfile
from types import SimpleNamespace
from unittest.mock import Mock, patch

from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, override_settings
from docx import Document

from core import preparation as prep
from core.template_source import approval_binding, source_for_version
from documents import inspect_template
from documents.preparation import candidate, inventory


class PreparationConstraintApprovalTests(SimpleTestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        override = override_settings(MEDIA_ROOT=self.root, WORD_PREPARATION_ENABLED=True)
        override.enable()
        self.addCleanup(override.disable)
        doc = Document()
        doc.add_paragraph('Temperature: ____ °C')
        doc.add_paragraph('Fever: ____ (Yes / No)')
        doc.add_paragraph('Patient signature: ____')
        self.original = self.root / 'original.docx'
        doc.save(self.original)
        self.inv = inventory(self.original)
        self.prepared = self.root / 'prepared.docx'
        self.result = candidate(self.original, self.prepared, self.inv, [a['id'] for a in self.inv['anchors']])
        self.info = inspect_template(self.prepared)

    def packet(self):
        class TemplateRecord(SimpleNamespace):
            objects = Mock()
        template = TemplateRecord(original_path='original.docx', format='docx', sha256=self.inv['source_sha256'],
                                  department='synthetic-department', versions=Mock())
        template.versions.order_by.return_value.first.return_value = SimpleNamespace(number=3)
        config = {'mode': 'fixture', 'provider': 'anthropic', 'model': 'configured-label'}
        proposal, _ = prep.provider_role(self.inv, 'proposer', config, 1)
        checker, _ = prep.provider_role(self.inv, 'checker', config, 1, proposal)
        packet = SimpleNamespace(pk='new-preparation', template_id='synthetic-template', template=template, config=config,
                                 inventory=self.inv, candidate_sha256=hashlib.sha256(self.prepared.read_bytes()).hexdigest(),
                                 candidate_path='prepared.docx', state='review', revision=0, corrections={},
                                 proposal=proposal, checker=checker, changes=self.result['changes'], save=Mock())
        packet.candidate_info = prep.candidate_inspection(packet, self.info, self.result, {})
        return packet

    def test_candidate_constraints_are_bound_to_original_and_writable_locations(self):
        packet = self.packet()
        self.assertEqual(packet.candidate_info['fields'], self.result['prepared_fields'])
        for change in ['omitted', 'unit', 'options', 'locator']:
            result = deepcopy(self.result)
            if change == 'omitted':
                result.pop('prepared_fields')
            elif change == 'unit':
                result['prepared_fields'][0]['answer_constraints']['unit'] = '°F'
            elif change == 'options':
                result['prepared_fields'][1]['answer_constraints']['options'].append('Sometimes')
            else:
                result['prepared_fields'][0]['locator']['paragraph'] = 9
            with self.subTest(change=change), self.assertRaises(ValueError):
                prep.candidate_inspection(packet, self.info, result, {})

    def test_explicit_approval_creates_next_version_without_mutating_old_mapping(self):
        packet = self.packet()
        old_fields = deepcopy(self.info['fields'])
        old_snapshot = deepcopy(old_fields)
        manager = Mock()
        manager.select_for_update.return_value.select_related.return_value.get.return_value = packet
        versions = Mock()
        actor = SimpleNamespace(pk='synthetic-doctor')
        with patch.object(prep.TemplatePreparation, 'objects', manager), patch.object(prep.TemplateVersion, 'objects', versions), \
                patch.object(prep, 'membership'), patch.object(prep, 'audit'):
            version = prep.approve.__wrapped__(packet.pk, actor, 0, packet.candidate_sha256,
                [row['anchor_id'] for row in packet.changes], True, True)
        self.assertIs(version, versions.create.return_value)
        self.assertEqual(versions.create.call_args.kwargs['number'], 4)
        self.assertEqual(versions.create.call_args.kwargs['fields'], self.result['prepared_fields'])
        self.assertEqual(old_fields, old_snapshot)
        self.assertEqual(packet.approval_sha256, approval_binding(packet))

    def test_missing_explicit_review_cannot_create_version(self):
        packet = self.packet()
        manager = Mock()
        manager.select_for_update.return_value.select_related.return_value.get.return_value = packet
        versions = Mock()
        with patch.object(prep.TemplatePreparation, 'objects', manager), patch.object(prep.TemplateVersion, 'objects', versions), \
                patch.object(prep, 'membership'), patch.object(prep, 'audit'):
            with self.assertRaises(ValidationError):
                prep.approve.__wrapped__(packet.pk, SimpleNamespace(pk='doctor'), 0, packet.candidate_sha256,
                    [row['anchor_id'] for row in packet.changes], False, True)
        versions.create.assert_not_called()
        packet.save.assert_not_called()

    def test_approved_constraint_edit_invalidates_existing_seal(self):
        packet = self.packet()
        packet.state = 'approved'
        packet.approved_by_id = 'synthetic-doctor'
        packet.approved_at = 'recorded-approval'
        packet.approval_sha256 = approval_binding(packet)
        version = SimpleNamespace(template=packet.template, template_id=packet.template_id, preparation=packet,
                                  preparation_id=packet.pk, prepared_path=packet.candidate_path,
                                  prepared_sha256=packet.candidate_sha256, fields=deepcopy(packet.candidate_info['fields']),
                                  document_backend={}, review_requirements=[])
        self.assertEqual(source_for_version(version), self.prepared)
        for key, value in [('unit', '°F'), ('version', True), ('version', 1.0)]:
            version.fields = deepcopy(packet.candidate_info['fields'])
            version.fields[0]['answer_constraints'][key] = value
            with self.subTest(key=key, value=value), self.assertRaises(ValidationError):
                source_for_version(version)
