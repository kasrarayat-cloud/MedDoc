"""Explicit reviewed-transcript actor context, original evidence and clinical counterexamples.

Every model choice is simulated; no provider requests or patient data.
"""
import json
import io
import os
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from uuid import uuid4

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import DatabaseError, transaction
from django.test import SimpleTestCase, TestCase, override_settings

from core import generation as g, transcript_context as tc, v3_field_relations as r
from core.jobs import enqueue, run_one
from core.models import AuditEvent, Encounter, Job, LedgerEntry, Membership, TranscriptReview, Usage
from core.security import private_path
from .test_billing_boundary import BoundaryFixture, deepgram
from .test_billing_policy import STORAGES
from .audio_fixtures import wav_bytes
from .test_v3_field_relations import relation_wire
from .test_v3_supersession_floor import SONNET_55_DEFAULT, form, raw
from .test_word_v3_workflow import V3, V3Fixture, part_text
from docx import Document


def proof(notes, start=0, end=None):
    end = len(notes) if end is None else end
    return tc.token_for(SimpleNamespace(pk=uuid4(), job_id=uuid4(), reviewed_by_id=uuid4(), source_revision=1,
        source_sha256=tc.digest(notes), raw_sha256=tc.digest(notes), result_sha256='0'*64,
        reviewed_sha256=tc.digest(notes[start:end]), span_start=start, span_end=end))


@override_settings(**SONNET_55_DEFAULT)
class PatientDefault(SimpleTestCase):
    def answers(self, notes, picks, context=True, start=0):
        fields = form(*picks)
        keyed = {f['key']: picks[f['label']] for f in fields}
        result = raw(notes, fields, relation_wire(notes, fields, keyed), g.SONNET_5_5)
        if context: result['patient_default_context'] = proof(notes, start=start)
        return fields, r.validated(result)

    def check(self, notes, label, selected, expected='supported', context=True, start=0):
        fields, answers = self.answers(notes, {label: selected}, context=context, start=start)
        answer = answers['f1']
        self.assertEqual(answer['status'], expected, (notes, answer))
        if expected == 'supported':
            self.assertEqual(answer['text'], '\n'.join(selected))
            self.assertEqual([e['quote'] for e in answer['evidence']], selected)
            self.assertEqual(g.governed_decisions(notes, fields, answers), [])
            row = g.as_currently_governed(g.governance_for(notes, fields), notes, {**fields[0], **answer})
            self.assertEqual((row['status'], row['text']), ('supported', answer['text']))
        else:self.assertEqual(answer['text'], '')
        return answer

    def test_subjectless_patient_values_are_useful_with_original_evidence(self):
        for label, source in [('Medications', 'Takes warfarin 5 mg daily.'),
                              ('Symptoms', 'Reports cough.')]:
            with self.subTest(source=source):
                self.check(source, label, [source], 'conflict', context=False)
                self.check(source, label, [source])

    def test_generic_medication_source_remains_an_unresolved_dependency(self):
        source='Takes warfarin 5 mg daily.'
        for label in ['Medications','Current medications']:
            for actor in ['She','The patient']:
                for unknown in ['unknown','UNKNOWN','']:
                    tail=f'{actor} switched from an {unknown} medication to supine positioning.'
                    with self.subTest(label=label,tail=tail):
                        self.check(source+' '+tail,label,[source],'conflict')
        tail='The patient switched from an unknown medication to supine positioning.'
        independent='The patient reports cough.'
        self.check(independent+' '+tail,'Symptoms',[independent])
        dependent='The patient reports cough while taking warfarin.'
        self.check(dependent+' '+tail,'Symptoms',[dependent],'conflict')
        # A concrete old value still governs that value, without treating its
        # separately supported current replacement as discontinued as well.
        change='The patient switched from warfarin medication to apixaban.'
        replacement='The patient takes apixaban 5 mg twice daily.'
        for label in ['Medications','Current medications']:
            self.check(source+' '+change,label,[source],'conflict')
            self.check(replacement+' '+change,label,[replacement])

    def test_foreign_positional_event_does_not_change_patient_medication(self):
        for source in ['Takes warfarin 5 mg daily.','The patient takes warfarin 5 mg daily.']:
            for tail in ['Her father switched from prone to supine positioning.',
                         'Her father switched to the supine position.']:
                with self.subTest(source=source,tail=tail):self.check(source+' '+tail,'Medications',[source])
        # This scopes an unrelated event; it grants no patient actor to a
        # father's actual medication statement or an unresolved drug change.
        foreign='Her father takes warfarin 5 mg daily.'
        self.check(foreign+' The patient switched from prone to supine positioning.','Medications',[foreign],'conflict')

    def test_generic_transition_requires_a_completed_patient_event(self):
        source='Takes warfarin 5 mg daily.'
        for tail in ['Her father switched from an unknown medication to supine positioning.',
                     'Her father previously switched from an unknown medication last year to supine positioning.',
                     'She did not switch from an unknown medication to supine positioning.',
                     'The patient may switch from an unknown medication to supine positioning tomorrow.',
                     'The patient will switch from an unknown medication to supine positioning tomorrow.']:
            for label in ['Medications','Current medications']:
                with self.subTest(tail=tail,label=label):self.check(source+' '+tail,label,[source])

    def test_unidentified_old_medication_does_not_also_discontinue_its_named_target(self):
        change='The patient switched from an unknown medication to apixaban.'
        source='The patient takes warfarin 5 mg daily.'
        target='The patient takes apixaban 5 mg twice daily.'
        for label in ['Medications','Current medications']:
            self.check(source+' '+change,label,[source],'conflict')
            self.check(target+' '+change,label,[target])
            for invalid in ['The patient previously took apixaban last year.',
                            'The patient may take apixaban tomorrow.',
                            'Her father takes apixaban 5 mg twice daily.']:
                with self.subTest(invalid=invalid,label=label):self.check(invalid+' '+change,label,[invalid],'conflict')
            self.check(target+' '+change+' Apixaban was stopped.',label,[target],'conflict')
        moved='The patient switched to apixaban from an unknown medication.'
        for label in ['Medications','Current medications']:
            self.check(source+' '+moved,label,[source],'conflict')
            self.check(target+' '+moved,label,[target])

    def test_medication_question_references_do_not_become_drug_arguments(self):
        source='The patient takes warfarin 5 mg daily.'
        for tail in ['Medications were copied from the prior section.',
                     'Her father switched from the medication section to the plan section.',
                     'Her father switched from an unknown medication to supine positioning. Medications are incorrect.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Medications',[source],'conflict')

    def test_mixed_actor_endings_keep_their_own_patient_and_named_medication(self):
        prefix='Her father switched from prone to supine positioning'
        for source in ['Takes warfarin 5 mg daily.','The patient takes warfarin 5 mg daily.']:
            for label in ['Medications','Current medications']:
                for tail in [' and the patient stopped warfarin.', '; the patient stopped warfarin.',
                             ' and the patient withheld warfarin.',
                             ' and the patient replaced warfarin with apixaban.',
                             ' and the patient stopped all medications.',
                             " and the patient's warfarin was stopped.",
                             ' and she stopped warfarin.']:
                    with self.subTest(source=source,label=label,tail=tail):
                        self.check(source+' '+prefix+tail,label,[source],'conflict')
                for tail in [' and the patient stopped apixaban.',
                             ' and the patient did not stop warfarin.',
                             ' and the patient may stop warfarin tomorrow.',
                             " and the patient stopped her father's warfarin.",
                             ' and the patient reports cough.']:
                    with self.subTest(source=source,label=label,tail=tail):self.check(source+' '+prefix+tail,label,[source])
        target='The patient takes apixaban 5 mg twice daily.'
        self.check(target+' '+prefix+' and the patient stopped warfarin and apixaban.',
            'Current medications',[target],'conflict')
        self.check(target+' '+prefix+' and the patient replaced warfarin with apixaban.',
            'Current medications',[target])
        source='Takes warfarin 5 mg daily.'
        self.check(source+' The patient switched from prone to supine positioning and her father stopped warfarin.',
            'Medications',[source])

    def test_mixed_medication_events_do_not_share_an_actor_or_replacement(self):
        source='The patient takes warfarin 5 mg daily.'
        other='Her father switched from an unknown medication to supine positioning and the patient reports cough.'
        self.check(source+' '+other,'Medications',[source])
        self.check(source+' '+other,'Current medications',[source])
        target='The patient takes apixaban 5 mg twice daily.'
        changes='Her father switched from an unknown medication to apixaban and the patient switched from an unknown medication to warfarin.'
        self.check(target+' '+changes,'Current medications',[target],'conflict')
        self.check(source+' '+changes,'Current medications',[source])
        # The separate first policy diagnostic for an 'unidentified substance'
        # is retained, not certified as a medication contradiction: the noun
        # does not positively name medication. Do not manufacture that class.

    def test_nested_clinical_arguments_retain_the_explicit_actor_and_ended_topic(self):
        prefix='Her father switched from prone to supine positioning and reports that '
        for source in ['Takes warfarin 5 mg daily.','The patient takes warfarin 5 mg daily.']:
            for label in ['Medications','Current medications']:
                for tail in ['the patient stopped warfarin.', 'the patient withheld warfarin.',
                             'the patient replaced warfarin with apixaban.', 'the patient stopped all medications.',
                             'she stopped warfarin.',
                             'warfarin was stopped and the patient reports cough.']:
                    with self.subTest(source=source,label=label,tail=tail):
                        self.check(source+' '+prefix+tail,label,[source],'conflict')
                for tail in ['the patient stopped apixaban.', 'the patient did not stop warfarin.',
                             'the patient may stop warfarin tomorrow.', 'her mother stopped warfarin.',
                             'she stopped apixaban.']:
                    with self.subTest(source=source,label=label,tail=tail):self.check(source+' '+prefix+tail,label,[source])

    def test_bounded_foreign_position_complement_has_no_patient_medication_argument(self):
        source='Takes warfarin 5 mg daily.'
        for label in ['Medications','Current medications']:
            for tail in ['The patient reports that her father switched from prone to supine positioning.',
                         'The patient confirms that her father switched from prone to supine positioning.',
                         'The patient reports that her father had switched from prone to supine positioning.',
                         'The patient reports that her father switched from prone to supine positioning and the patient stopped apixaban.']:
                with self.subTest(label=label,tail=tail):self.check(source+' '+tail,label,[source])
            for tail in ['The patient reports that her father switched from prone to supine positioning and the patient stopped warfarin.',
                         'The patient reports that she switched from an unknown medication to supine positioning.',
                         'The patient reports that someone switched from an unknown medication to supine positioning.']:
                with self.subTest(label=label,tail=tail):self.check(source+' '+tail,label,[source],'conflict')

    def test_owned_patient_denial_is_not_erased_by_a_foreign_outer_actor(self):
        source='Takes warfarin 5 mg daily.'
        for label in ['Medications','Current medications']:
            for tail in ['Her father denies that the patient takes warfarin.',
                         'Her father switched from prone to supine positioning and the patient does not take warfarin.']:
                with self.subTest(label=label,tail=tail):self.check(source+' '+tail,label,[source],'conflict')
            for tail in ['Her father denies that her mother takes warfarin.',
                         'Her father denies that the patient takes apixaban.',
                         'Her father switched from prone to supine positioning and her mother does not take warfarin.',
                         'Her father switched from prone to supine positioning and the patient does not take apixaban.',
                         'Her father switched from prone to supine positioning and the patient did not stop warfarin.',
                         'Her father denies that the patient has warfarin.']:
                with self.subTest(label=label,tail=tail):self.check(source+' '+tail,label,[source])

    def test_family_positional_subject_is_not_replaced_by_an_inferred_patient(self):
        source='Takes warfarin 5 mg daily.'
        for label in ['Medications','Current medications']:
            for tail in ['Her father switched from prone to supine positioning and stopped warfarin.',
                         'Her father switched from prone to supine positioning and replaced warfarin with apixaban.',
                         'The patient reports that her father switched from prone to supine positioning and stopped warfarin.']:
                with self.subTest(label=label,tail=tail):self.check(source+' '+tail,label,[source])

    def test_passive_ended_argument_and_explicit_recipient_keep_their_ownership(self):
        source='Takes warfarin 5 mg daily.'
        prefix='Her father switched from prone to supine positioning and '
        for label in ['Medications','Current medications']:
            for tail in ['warfarin was withheld from her father.',
                         'apixaban was withheld from the patient.',
                         'apixaban was withheld.',
                         'warfarin was not withheld from the patient.']:
                with self.subTest(label=label,tail=tail):self.check(source+' '+prefix+tail,label,[source])
            for tail in ['warfarin was withheld from the patient.',
                         'warfarin was withheld.', 'medication was withheld from the patient.',
                         'warfarin was stopped and the patient reports cough.']:
                with self.subTest(label=label,tail=tail):self.check(source+' '+prefix+tail,label,[source],'conflict')

    def test_bounded_negative_clinical_argument_is_separate_from_patient_medication(self):
        source='Takes warfarin 5 mg daily.'
        prefix='Her father switched from prone to supine positioning and '
        for label in ['Medications','Current medications']:
            self.check(source+' '+prefix+'the patient denies knee pain.',label,[source])
            self.check(source+' '+prefix+'the patient denies knee pain while taking warfarin.',label,[source],'conflict')
            # The existing note-wide modal scope remains an honest unsupported
            # boundary. A base future and an uncertain completed cessation must
            # not be qualified by changing that old scope or its evidence.
            for tail in ['reports that the patient might stop warfarin tomorrow.',
                         'reports that the patient might have stopped warfarin.']:
                with self.subTest(label=label,tail=tail):self.check(source+' '+prefix+tail,label,[source],'conflict')
        pain='The patient reports knee pain.'
        self.check(pain+' '+prefix+'the patient denies knee pain.','Symptoms',[pain],'conflict')
        self.check(source+' '+prefix+'the patient does not take warfarin.','Medications',[source],'conflict')

    def test_nominal_clinical_assertions_gain_only_default_actor_not_drug_status(self):
        for source, label in [('No fever.', 'Examination'),
                              ('Fever.', 'Symptoms'),
                              ('Examination mild swelling over the left lateral ankle.', 'Examination'),
                              ('No open wound.', 'Examination')]:
            with self.subTest(source=source):
                self.check(source, label, [source], 'conflict', context=False)
                self.check(source, label, [source])
        pair='Examination mild swelling over the left lateral ankle. No open wound.'
        self.check(pair, 'Examination', ['Examination mild swelling over the left lateral ankle.', 'No open wound.'])
        self.check('Symptoms.', 'Symptoms', ['Symptoms.'], 'conflict')
        for source, label in [('Warfarin 5 mg daily.', 'Current medications'),
                              ('Stopped warfarin.', 'Current medications'),
                              ('Her father mild cough.', 'Symptoms'),
                              ('Possible ankle strain.', 'Current diagnosis')]:
            with self.subTest(source=source):self.check(source, label, [source], 'conflict')
        for source in ['Family history. No fever.', 'Hypothetical. No fever.']:
            self.check(source, 'Examination', ['No fever.'], 'conflict')

    def test_opposed_anatomical_topics_need_the_same_explicit_source_side(self):
        source='Examination mild swelling over the left lateral ankle.'
        self.check(source+' Denies fever, numbness, and right ankle pain.', 'Examination', [source])
        self.check(source+' The patient has no swelling over the left ankle.', 'Examination', [source], 'conflict')
        self.check(source+' The patient has no swelling.', 'Examination', [source], 'conflict')
        self.check(source+' The patient has no swelling over the right ankle.', 'Examination', [source])
        patient='The patient reports pain in the left knee.'
        self.check(patient+' The patient denies right knee pain.', 'Symptoms', [patient])
        self.check(patient+' The patient denies left knee pain.', 'Symptoms', [patient], 'conflict')
        self.check(patient+' The patient denies knee pain.', 'Symptoms', [patient], 'conflict')
        self.check(patient+' The patient denies right knee pain and has no fever.', 'Symptoms', [patient])

    def test_context_is_exact_and_span_bound_not_a_global_patient_assumption(self):
        source = 'Takes warfarin 5 mg daily.'
        prefix = 'Reports cough.\n\n'
        notes = prefix + source
        self.check(notes, 'Symptoms', ['Reports cough.'], 'conflict', start=len(prefix))
        self.check(notes, 'Medications', [source], start=len(prefix))
        token = proof(notes, start=len(prefix))
        self.assertTrue(tc.verified(notes, token))
        self.assertIsNone(tc.verified(notes+' ', token))
        self.assertIsNone(tc.verified(notes, token[:-1]+'x'))
        self.assertIsNone(tc.verified(notes, proof(notes, end=0)))
        self.assertIsNone(tc.verified(notes, True))

    def test_explicit_other_and_inherited_other_never_gain_patient_subject(self):
        for notes, pick in [('Her father takes warfarin 5 mg daily.', 'Her father takes warfarin 5 mg daily.'),
                            ('Her father reports cough. Takes warfarin 5 mg daily.', 'Takes warfarin 5 mg daily.'),
                            ('The nurse takes warfarin 5 mg daily.', 'The nurse takes warfarin 5 mg daily.'),
                            ('He takes warfarin 5 mg daily.', 'He takes warfarin 5 mg daily.'),
                            ('"Takes warfarin 5 mg daily."', '"Takes warfarin 5 mg daily."')]:
            with self.subTest(notes=notes):self.check(notes, 'Medications', [pick], 'conflict')

    def test_time_uncertainty_denial_and_cessation_keep_their_meaning(self):
        for notes, pick in [('Previously took warfarin 5 mg daily.', 'Previously took warfarin 5 mg daily.'),
                            ('May take warfarin 5 mg daily.', 'May take warfarin 5 mg daily.'),
                            ('Takes warfarin 5 mg daily. Warfarin was stopped.', 'Takes warfarin 5 mg daily.'),
                            ('Takes warfarin 5 mg daily. Warfarin was withheld today.', 'Takes warfarin 5 mg daily.'),
                            ('Takes warfarin 5 mg daily. Warfarin is on hold.', 'Takes warfarin 5 mg daily.'),
                            ('Takes warfarin 5 mg daily. Warfarin was replaced by apixaban.', 'Takes warfarin 5 mg daily.')]:
            with self.subTest(notes=notes):self.check(notes, 'Medications', [pick], 'conflict')
        # A denial stays a denial in the full retained sentence; subject context
        # cannot authorize a model's positive fragment or invented dose.
        notes = 'Denies cough.'
        _, answers = self.answers(notes, {'Symptoms': [notes]})
        self.assertNotEqual(answers['f1'].get('text'), 'cough')
        self.assertTrue(all(e['quote'] == notes for e in answers['f1']['evidence']))

    def test_two_original_voice_safety_controls_remain_blocked(self):
        retained=json.loads((Path(__file__).parent/'voice_first_pass_controls.json').read_text())
        notes=retained['notes']; fields=retained['fields']
        original=raw(notes,fields,retained['first_model_body'],g.SONNET_5_5)
        for context in [False,True]:
            if context:original['patient_default_context']=proof(notes)
            answers=r.validated(original)
            for key in retained['held_controls']:
                with self.subTest(context=context,key=key):
                    self.assertEqual((answers[key]['status'],answers[key]['text']),('conflict',''))
        # The model's exact original correction selection and planning line
        # remain adversarial controls, rather than newly simplified examples.

    def test_original_v4_voice_response_gains_one_useful_answer_without_changing_raw_output(self):
        retained=json.loads((Path(__file__).parent/'voice_first_pass_controls.json').read_text())
        notes=retained['notes']; fields=retained['fields']
        result=raw(notes,fields,retained['first_model_body'],g.SONNET_5_5)
        self.assertEqual(sum(a['status']=='supported' for a in r.validated(result).values()),0)
        result['patient_default_context']=proof(notes)
        answers=r.validated(result)
        self.assertEqual(answers['v3_7']['status'],'supported')
        self.assertEqual(answers['v3_7']['text'],
            'Examination mild swelling over the left lateral ankle.\nNo open wound.')
        self.assertEqual(sum(a['status']=='supported' for a in answers.values()),1)
        self.assertEqual(g.governed_decisions(notes,fields,answers),[])
        for key in retained['held_controls']:self.assertEqual(answers[key]['status'],'conflict')

    def test_literal_prefix_heading_and_side_controls_keep_actual_opposites(self):
        source='Examination mild swelling over the left lateral ankle.'
        for tail in ['The patient has no swelling over the left ankle.',
                     'The patient has no swelling over either ankle.',
                     'The patient has no swelling over both ankles.',
                     'The patient has no swelling over the ankle.',
                     'The patient has no swelling.',
                     'The patient has non-swollen left ankle.',
                     'The patient has swelling over the left ankle. No.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Examination',[source],'conflict')
        self.check(source+' The patient has non-swollen right ankle.','Examination',[source])
        patient='The patient reports pain in the left knee.'
        self.check(patient+' Her father denies left knee pain.','Symptoms',[patient])
        # Retained lexical-denial residual: a negated deny predicate is still
        # conservatively held. It is measured as an unnecessary hold, not
        # proof that the preceding pain statement is actually contradicted.
        self.check(patient+' The patient does not deny left knee pain.','Symptoms',[patient],'conflict')
        self.check(patient+' The patient denies left knee pain.','Symptoms',[patient],'conflict',context=False)
        self.check(patient+' The patient denies left knee pain.','Symptoms',[patient],'conflict',context=True)

    def test_independent_refutation_modifiers_and_exclusive_clinical_topic(self):
        source='Examination mild swelling over the left lateral ankle.'
        for tail in ['The patient has a left ankle without swelling.',
                     'The patient has absent swelling over the left ankle.',
                     'The patient has a swelling-free left ankle.',
                     'The patient has a left ankle free of swelling.',
                     'The patient has an ankle without swelling.',
                     'The patient has no swelling over the left ankle.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Examination',[source],'conflict')
        for tail in ['The patient has a right ankle without swelling.',
                     'The patient has absent swelling over the right ankle.',
                     'The patient has a swelling-free right ankle.',
                     'Her father has no swelling over the left ankle.',
                     'Her father has a left ankle without swelling.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Examination',[source])
        patient='The patient reports pain in the left knee.'
        for tail in ['The patient has a left knee without pain.',
                     'Only her father has left knee pain.']:
            with self.subTest(tail=tail):self.check(patient+' '+tail,'Symptoms',[patient],'conflict')
        for tail in ['Her father has left knee pain.','Only her father has right knee pain.',
                     'Only her father has a cough.','The patient has a right knee without pain.']:
            with self.subTest(tail=tail):self.check(patient+' '+tail,'Symptoms',[patient])
        self.check('The patient takes warfarin. Diet: none.','Medications',['The patient takes warfarin.'])
        self.check('Takes warfarin. None.','Medications',['Takes warfarin.'],'conflict')

    def test_negative_modifier_polarity_and_side_are_local_to_the_named_assertion(self):
        patient='The patient reports pain in the left knee.'
        for tail in ['The patient has no pain relief.',
                     'The patient goes without treatment for pain.',
                     'Her father has a left knee without pain.',
                     'The patient has a right knee without pain.']:
            with self.subTest(tail=tail):self.check(patient+' '+tail,'Symptoms',[patient])
        for tail in ['The patient does not have a left knee without pain.',
                     'The patient has a right knee and a left ankle without pain.',
                     'The patient has knee pain on the right and on the left. The patient has a knee without pain.']:
            with self.subTest(tail=tail):self.check(patient+' '+tail,'Symptoms',[patient],'conflict')
        negative='The patient has a left knee without pain.'
        self.check(negative+' '+patient,'Symptoms',[negative],'conflict')
        self.check(negative+' The patient reports pain in the right knee.','Symptoms',[negative])
        self.check(negative+' Her father reports pain in the left knee.','Symptoms',[negative])
        self.check('Takes warfarin. No diabetes.','Medications',['Takes warfarin.'])

    def test_independent_predicative_refutation_keeps_argument_actor_side_and_polarity(self):
        source='Examination mild swelling over the left lateral ankle.'
        for tail in ['The patient has entirely absent swelling over the left ankle.',
                     'The patient has completely absent swelling over the left ankle.',
                     'Swelling over the left ankle is absent.',
                     "The patient's swelling over the left ankle is absent.",
                     'Swelling is absent over the left ankle.',
                     'The left ankle is free of swelling.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Examination',[source],'conflict')
        for tail in ['The patient has entirely absent swelling over the right ankle.',
                     'Swelling over the right ankle is absent.',
                     'Swelling is absent over the right ankle.',
                     'Swelling over the left ankle is not absent.',
                     "Her father's swelling over the left ankle is absent.",
                     'The patient has completely absent pain over the left ankle.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Examination',[source])
        # A malformed nested cancellation supplies no positive assertion.
        self.check(source+' The patient has not absent swelling over the left ankle.',
                   'Examination',[source],'conflict')
        negative='The patient has a left ankle without swelling.'
        self.check(negative+' The patient has swelling over the right ankle.','Examination',[negative])
        self.check(negative+' The patient has swelling over the left ankle.','Examination',[negative],'conflict')
        self.check(source+" Her father's swelling over the patient's left ankle is absent.",
                   'Examination',[source],'conflict')

    def test_unknown_negative_complement_keeps_only_its_named_clinical_dependency(self):
        source='Examination mild swelling over the left lateral ankle.'
        for tail in ['Swelling over the left ankle was found to be absent.',
                     'There is absent swelling over the left ankle.',
                     'Swelling over the left ankle was considered to be absent.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Examination',[source],'conflict')
        for tail in ['Swelling over the right ankle was found to be absent.',
                     'There is absent swelling over the right ankle.',
                     'Pain over the left ankle was found to be absent.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Examination',[source])
        patient='The patient reports pain in the left knee.'
        self.check(patient+' The patient has no pain relief.','Symptoms',[patient])
        self.check(patient+' The patient goes without treatment for pain.','Symptoms',[patient])

    def test_existing_negative_operator_arguments_need_no_new_word_whitelist(self):
        source='Examination mild swelling over the left lateral ankle.'
        for tail in ['Swelling absent over the left ankle.',
                     'The patient lacks swelling over the left ankle.',
                     'Lack of swelling over the left ankle.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Examination',[source],'conflict')
        for tail in ['The patient lacks swelling over the right ankle.',
                     'Lack of swelling over the right ankle.',
                     'Her father lacks swelling over the left ankle.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Examination',[source])
        self.check(source+' Her father lacks swelling over her left ankle.',
                   'Examination',[source],'conflict')
        patient='The patient reports pain in the left knee.'
        self.check(patient+' The patient lacks pain relief.','Symptoms',[patient])
        # An ambiguous negative relationship cannot itself supply an answer.
        self.check('The patient has swelling absent over the left ankle.',
                   'Examination',['The patient has swelling absent over the left ankle.'],'conflict')

    def test_confirmation_does_not_tune_around_an_unresolved_parser_failure(self):
        # Pinned syntax still reads left as a past verb in this sentence. Retain
        # the first failed result; a clinician can edit and explicitly re-review.
        self.check('Reports left ankle pain.', 'Chief complaint', ['Reports left ankle pain.'], 'conflict')
        self.check('Mild swelling over the left ankle.', 'Examination', ['Mild swelling over the left ankle.'], 'conflict')

    def test_negative_nominalization_complements_and_exclusive_polarity_refutations(self):
        source='Examination mild swelling over the left lateral ankle.'
        for tail in ['The patient has an absence of swelling over the left ankle.',
                     'The patient finds swelling absent over the left ankle.',
                     'The patient reports swelling to be absent over the left ankle.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Examination',[source],'conflict')
        negative='No swelling over the left ankle.'
        for tail in ['Only her father lacks swelling over the left ankle.',
                     'Only her father has no swelling over the left ankle.']:
            with self.subTest(tail=tail):self.check(negative+' '+tail,'Examination',[negative],'conflict')
        for source in ['The patient takes warfarin 5 mg daily.','Takes warfarin 5 mg daily.']:
            self.check(source+' Only her father has no warfarin.','Medications',[source])
        for tail in ['Her father lacks swelling over the left ankle.',
                     'Only her father has no swelling over the right ankle.',
                     'Only her father lacks pain over the left ankle.']:
            with self.subTest(tail=tail):self.check(negative+' '+tail,'Examination',[negative])
        source='Examination mild swelling over the left lateral ankle.'
        for tail in ['The patient has an absence of swelling over the right ankle.',
                     'The patient finds swelling absent over the right ankle.',
                     'The patient reports swelling to be absent over the right ankle.',
                     "Swelling over her father's left ankle was found to be absent."]:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Examination',[source])
        self.check(source+" Swelling over her father's and the patient's left ankle was found to be absent.",
                   'Examination',[source],'conflict')

    def test_negative_operator_argument_and_property_polarity_refutations(self):
        source='Examination mild swelling over the left lateral ankle.'
        for tail in ['The patient reports no evidence of swelling over the left ankle.',
                     'Her father denies that the patient has swelling over the left ankle.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Examination',[source],'conflict')
        for tail in ['The patient reports no evidence of swelling over the right ankle.',
                     'Her father reports no evidence of swelling over the left ankle.',
                     'Her father denies that the patient has swelling over the right ankle.',
                     'Only her father has an absence of swelling over the left ankle.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Examination',[source])
        negative='No swelling over the left ankle.'
        self.check(negative+' Only her father has an absence of swelling over the left ankle.',
                   'Examination',[negative],'conflict')
        self.check(negative+' Only her father has an absence of swelling over the right ankle.',
                   'Examination',[negative])
        # Withdrawn refutation: lack/possession and taking are distinct claims.
        # The patient can possess medication without taking it. Do not widen
        # a guard on an unsupported claim that this is a contradiction.
        medication='The patient takes no warfarin.'
        self.check(medication+' Only her father lacks warfarin.','Current medications',[medication])

    def test_bound_negative_nominal_chain_keeps_the_same_clinical_dependency(self):
        source='Examination mild swelling over the left lateral ankle.'
        for nominal in ['an absence of evidence of swelling',
                        'an absence of evidence of signs of swelling',
                        'an absence of documentation of evidence of signs of swelling']:
            left='The patient reports '+nominal+' over the left ankle.'
            right='The patient reports '+nominal+' over the right ankle.'
            with self.subTest(nominal=nominal):
                self.check(source+' '+left,'Examination',[source],'conflict')
                self.check(source+' '+right,'Examination',[source])
        self.check(source+' Her father reports an absence of evidence of swelling over the left ankle.',
                   'Examination',[source])
        self.check(source+' The patient reports an absence of evidence of pain over the left ankle.',
                   'Examination',[source])

    def test_independent_refutation_ellipses_cannot_retract_a_patient_claim_silently(self):
        source='Takes warfarin 5 mg daily.'
        for tail,label in [('Only her father does.','Medications'),('Not today.','Medications taken today'),
                           ('Except today.','Medications taken today'),('Her father does not.','Medications'),
                           ('Not this week.','Medications taken this week'),
                           ('Only her father actually takes.','Medications'),
                           ('Only her father takes it.','Medications')]:
            with self.subTest(tail=tail):self.check(source+' '+tail,label,[source],'conflict')
        # A modifier before the first claim still cannot be declared independent.
        self.check('Not today. '+source,'Medications taken today',[source],'conflict')

    def test_elliptical_dependency_is_local_and_complete_other_person_claims_are_distinct(self):
        source='Takes warfarin 5 mg daily.'
        for tail in ['Her father takes apixaban.','Her father has diabetes.','Her father reports cough.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Medications',[source])
        # The elliptical substitution points to the intervening cough statement;
        # it does not undo every earlier independent statement in the document.
        self.check(source+' Reports cough. Only her father does.','Medications',[source])
        self.check(source+' Reports cough. Only her father does.','Symptoms',['Reports cough.'],'conflict')

    def test_proper_date_modifiers_cannot_be_declared_independent(self):
        source='Takes warfarin 5 mg daily.'
        for tail in ['Not since Monday.','Not on Fridays.','Only at night.','Not after surgery.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Current medications',[source],'conflict')
        self.check(source+' Reports cough. Not since Monday.','Medications',[source])
        self.check(source+' Reports cough. Not since Monday.','Symptoms',['Reports cough.'],'conflict')

    def test_exclusive_other_actor_changes_only_the_same_predicate_and_claim(self):
        source='Takes warfarin 5 mg daily.'
        for tail in ['Only her father takes warfarin.','Her father alone takes warfarin.',
                     'Her father only takes warfarin.','Her father is the only one to take warfarin.',
                     'Her father is the sole person taking warfarin.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Medications',[source],'conflict')
        for tail in ['Her father takes warfarin.','Her father takes only warfarin.',
                     'Only her father takes apixaban.','Only her father has diabetes.',
                     'Only her father reports cough.','Only her father knows warfarin.',
                     'Her father takes warfarin alone.','Only her father is alone on warfarin.',
                     'Her father only is alone while taking warfarin.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Medications',[source])
        self.check(source+' Reports cough. Only her father takes warfarin.','Medications',[source],'conflict')
        self.check(source+' Reports cough. Only her father takes warfarin.','Symptoms',['Reports cough.'])
        self.check('Reports cough. Only her father has cough.','Symptoms',['Reports cough.'],'conflict')
        side='The patient reports right ankle pain.'
        self.check(side+' Only her father reports left ankle pain.','Symptoms',[side])

    def test_ambiguous_restriction_is_reviewed_without_inventing_drug_status(self):
        source='Takes warfarin 5 mg daily.'
        for tail in ['No longer routine.','No longer daily.','Not regularly.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Current medications',[source],'conflict')
        self.check(source+' No diabetes.','Medications',[source])
        self.check(source+' Reports cough. No longer routine.','Medications',[source])
        self.check(source+' Reports cough. No longer routine.','Symptoms',['Reports cough.'],'conflict')
        on='The patient is on warfarin.'
        self.check(on+' Only her father is alone on warfarin.','Medications',[on])
        self.check(on+' Her father is single with warfarin.','Medications',[on])
        self.check(on+' Her father is the sole person taking warfarin.','Medications',[on],'conflict')

    def test_nominal_classifier_and_exclusive_role_are_affected_claim_dependencies(self):
        source='Takes warfarin 5 mg daily.'
        for tail in ['No longer a daily medication.','Her father is the sole warfarin user.',
                     'Her father is the only patient on warfarin.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Current medications',[source],'conflict')
        for tail in ['Her father is the sole apixaban user.','Her father is the only patient on apixaban.',
                     'Her father knows warfarin.','Her father is an expert on warfarin.',
                     'Only her father knows warfarin.','No diabetes.','No longer a daily symptom.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Medications',[source])
        self.check(source+' Reports cough. No longer a daily medication.','Medications',[source],'conflict')
        self.check(source+' Reports cough. No longer a daily medication.','Symptoms',['Reports cough.'])
        self.check(source+' Reports cough. Her father is the sole warfarin user.','Medications',[source],'conflict')
        self.check(source+' Reports cough. Her father is the sole warfarin user.','Symptoms',['Reports cough.'])

    def test_nominal_temporal_predicate_and_exclusive_negation_have_distinct_scope(self):
        source='Takes warfarin 5 mg daily.'
        for tail in ['No longer the daily medication.','No longer a daily medication.','No longer my daily medication.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Current medications',[source],'conflict')
        for tail in ['Her father is not the only warfarin user.',
                     'Her father is the only person not taking warfarin.',
                     'Only her father is not taking warfarin.','Her father is the only person without warfarin.',
                     'Only her father used to take warfarin.','No diabetes.']:
            with self.subTest(tail=tail):self.check(source+' '+tail,'Medications',[source])
        denied='The patient does not take warfarin.'
        self.check(denied+' Only her father does not take warfarin.','Medications',[denied],'conflict')
        self.check(source+' Reports cough. No longer the daily medication.','Symptoms',['Reports cough.'])
        self.check(source+' Reports cough. No longer the daily medication.','Current medications',[source],'conflict')

    @override_settings(V3_GENERATION_BATCHES=True)
    def test_whole_form_plan_keeps_context_in_request_digest_and_field_records(self):
        from core import generation_batch as b
        from .test_generation_batch import sample, manifest
        notes='Takes warfarin 5 mg daily. Reports cough.'
        e=sample(2,notes)
        e.template_version.fields[0]['label']='Medications'
        e.template_version.fields[1]['label']='Symptoms'
        e.template_version.generation_constraints=manifest(e.template_version.fields)
        config=g.configuration(e.template_version.fields)
        plan=b.plan_for(e,config,patient_context=proof(notes))
        payload=json.loads(b.request(plan,1)['messages'][0]['content'])
        self.assertEqual(payload['patient_default_context']['default_subject'],'encounter_patient')
        self.assertEqual(b.digest(b.request(plan,1)),plan['parts'][0]['request_sha256'])
        selected={e.template_version.fields[0]['key']:['Takes warfarin 5 mg daily.'],
                  e.template_version.fields[1]['key']:['Reports cough.']}
        output={'model':g.SONNET_5_5,'stop_reason':'end_turn','content':[{'type':'text',
                'text':json.dumps(relation_wire(notes,e.template_version.fields,selected))}]}
        answers,_=b.assemble(plan,[output])
        self.assertTrue(all(a['status']=='supported' for a in answers.values()))
        no_context=b.plan_for(e,config)
        self.assertNotEqual(b.digest(plan),b.digest(no_context))
        answers,_=b.assemble(no_context,[output])
        self.assertTrue(all(a['status']=='conflict' for a in answers.values()))
        plan['patient_default_context']='tampered'
        with self.assertRaises(ValidationError):b.request(plan,1)


@override_settings(AI_MODE='fixture', ADMIN_SITE=False, STORAGES=STORAGES)
class ReviewVersion(BoundaryFixture, TestCase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.user)
        self.url = f'/encounters/{self.encounter.pk}/'

    def transcript(self, text='Takes warfarin 5 mg daily.'):
        with patch.dict(os.environ, self.VOICE_ENV):
            quote = self.client.get(self.url).context['usage_quote']
            response = self.client.post(self.url+'voice/', {'audio':SimpleUploadedFile('context.wav',wav_bytes(tag='context')),
                'consent':'on','usage_quote':quote})
            self.assertEqual(response.status_code,302)
            job = Job.objects.get(encounter=self.encounter)
            with patch('core.jobs.httpx.post', return_value=deepgram(text, 'synthetic-context-transcript', duration=1)):
                self.assertTrue(run_one())
        job.refresh_from_db()
        self.assertEqual(job.state, 'succeeded', job.error)
        return job

    def apply(self, job, text='Takes warfarin 5 mg daily.', **changes):
        self.encounter.refresh_from_db()
        return self.client.post(self.url+'notes/', {'revision': self.encounter.revision,
            'transcript_job': str(job.pk), 'corrected_transcript': text,
            'transcript_mode': 'append', 'transcript_reviewed': 'on',
            'transcript_patient_context': 'on', **changes})

    def test_raw_provider_and_clinician_edit_and_context_are_separate_evidence(self):
        job = self.transcript()
        original = self.encounter.notes
        result = private_path(job.result_path).read_bytes()
        corrected = 'Takes warfarin 4 mg daily.'
        self.assertEqual(self.apply(job, corrected).status_code, 302)
        self.encounter.refresh_from_db()
        review = TranscriptReview.objects.get(job=job)
        self.assertEqual(review.reviewed_text, corrected)
        self.assertEqual(review.raw_sha256, tc.digest('Takes warfarin 5 mg daily.'))
        self.assertEqual(review.reviewed_sha256, tc.digest(corrected))
        self.assertEqual(self.encounter.notes, original+'\n\n'+corrected)
        self.assertEqual(private_path(job.result_path).read_bytes(), result)
        self.assertTrue(tc.current(self.encounter))
        self.assertEqual(tc.payload(self.encounter.notes, tc.current(self.encounter))['start'], len(original)+2)
        event = AuditEvent.objects.get(action='transcription.applied', target=str(job.pk))
        self.assertNotIn(corrected, json.dumps(event.details))
        self.assertNotIn(review.context_token, json.dumps(event.details))
        self.assertEqual(Usage.objects.filter(job=job).count(), 1)
        self.assertEqual(LedgerEntry.objects.filter(job=job, kind='usage').count(), 1)
        self.assertEqual(self.apply(job).status_code, 409)

    def test_edit_revert_requires_new_explicit_review_without_retranscription(self):
        job = self.transcript()
        self.assertEqual(self.apply(job, transcript_mode='replace').status_code, 302)
        self.encounter.refresh_from_db(); notes = self.encounter.notes; token = tc.current(self.encounter)
        self.assertEqual(self.client.post(self.url+'notes/', {'revision': self.encounter.revision, 'notes': notes}).status_code, 302)
        self.encounter.refresh_from_db();self.assertEqual(tc.current(self.encounter), token)
        for edited in [notes+'\nReports cough.', notes]:
            self.assertEqual(self.client.post(self.url+'notes/', {'revision': self.encounter.revision, 'notes': edited}).status_code, 302)
            self.encounter.refresh_from_db();self.assertFalse(tc.current(self.encounter))
        usage = list(Usage.objects.values_list('id','actual_cost','charged_cost'))
        with patch('core.jobs.httpx.post') as provider:
            self.assertEqual(self.apply(job, notes, transcript_mode='replace', transcript_reconfirm='on').status_code, 302)
        provider.assert_not_called()
        self.encounter.refresh_from_db();self.assertTrue(tc.current(self.encounter))
        self.assertNotEqual(tc.current(self.encounter), token)
        self.assertEqual(TranscriptReview.objects.filter(job=job).count(), 2)
        self.assertEqual(list(Usage.objects.values_list('id','actual_cost','charged_cost')), usage)

    def test_employee_cannot_supply_clinician_context_and_failure_rolls_back(self):
        job = self.transcript()
        Membership.objects.filter(user=self.user, department=self.encounter.department).update(role='employee')
        previous = self.encounter.notes
        self.assertEqual(self.apply(job).status_code, 403)
        self.encounter.refresh_from_db();self.assertEqual(self.encounter.notes, previous)
        self.assertFalse(TranscriptReview.objects.exists())
        self.assertEqual(self.apply(job, transcript_patient_context='').status_code, 302)
        self.encounter.refresh_from_db();self.assertFalse(tc.current(self.encounter))
        self.assertFalse(TranscriptReview.objects.get().patient_context)

    def test_database_append_only_guard_cannot_be_bypassed_by_queryset(self):
        job = self.transcript();self.assertEqual(self.apply(job).status_code, 302)
        review = TranscriptReview.objects.get(job=job)
        with self.assertRaises(ValidationError):review.save()
        for operation in [lambda:TranscriptReview.objects.filter(pk=review.pk).update(reviewed_text='Lost evidence'),
                          lambda:TranscriptReview.objects.filter(pk=review.pk).delete()]:
            with self.assertRaises(DatabaseError), transaction.atomic():operation()
        review.refresh_from_db();self.assertEqual(review.reviewed_text, 'Takes warfarin 5 mg daily.')

    def test_database_refuses_review_of_a_different_source_span_before_retaining_it(self):
        job=self.transcript()
        source=self.encounter.notes
        for changes in [dict(reviewed_text='Another synthetic value'),dict(span_end=len(source)+1),
                        dict(source_revision=self.encounter.revision+1)]:
            values={'encounter':self.encounter,'job':job,'reviewed_by':self.user,
                'source_revision':self.encounter.revision,'source_sha256':tc.digest(source),
                'raw_sha256':tc.digest('Takes warfarin 5 mg daily.'),'result_sha256':job.result_sha256,
                'reviewed_sha256':tc.digest(source),'reviewed_text':source,'span_start':0,'span_end':len(source)}
            values.update(changes)
            with self.subTest(changes=changes),self.assertRaises(DatabaseError),transaction.atomic():
                TranscriptReview.objects.create(**values)
        self.assertFalse(TranscriptReview.objects.exists())
        self.assertEqual(self.apply(job).status_code,302)

    @override_settings(**SONNET_55_DEFAULT)
    def test_job_admission_only_uses_current_server_review_on_default_sonnet_route(self):
        job = self.transcript();self.assertEqual(self.apply(job).status_code, 302)
        self.encounter.refresh_from_db()
        config = g.configuration(form('Medications'))
        self.assertEqual(config['ordinary_field_contract'], r.CONTRACT)
        admitted = tc.admitted({**config,'patient_default_context':'forged'}, self.encounter)
        self.assertEqual(admitted['patient_default_context'], tc.current(self.encounter))
        self.assertNotIn('patient_default_context', tc.admitted({'patient_default_context':'forged'}, self.encounter))
        fields = form('Medications'); request = r.compile_request(self.encounter.notes, fields, admitted)
        payload=json.loads(request['wire']['messages'][0]['content'])
        self.assertEqual(payload['patient_default_context']['default_subject'], 'encounter_patient')
        self.assertNotIn(tc.current(self.encounter), json.dumps(payload))


@override_settings(**dict(V3,**SONNET_55_DEFAULT))
class VoiceToWord(V3Fixture, TestCase):
    """Actual HTTP/worker/Word export with mocked providers; live calls are a separate gate."""
    def test_reviewed_subjectless_transcript_reaches_default_doctor_generation_and_export(self):
        doc=Document();doc.add_paragraph('Medications: ____________');doc.add_paragraph('Symptoms: ____________')
        buffer=io.BytesIO();doc.save(buffer)
        template,version=self.uploaded(data=buffer.getvalue(),name='Synthetic subjectless voice')
        self.confirm(template,version)
        self.assertEqual(self.new_encounter(version,notes='').status_code,302)
        encounter=Encounter.objects.get(template_version=version)
        url=f'/encounters/{encounter.pk}/'
        with patch.dict(os.environ,BoundaryFixture.VOICE_ENV):
            quote=self.client.get(url).context['usage_quote']
            response=self.client.post(url+'voice/',{'audio':SimpleUploadedFile('synthetic.wav',wav_bytes(tag='scope')),
                'consent':'on','usage_quote':quote})
            self.assertEqual(response.status_code,302)
            with patch('core.jobs.httpx.post',return_value=deepgram('Takes warfarin 5 mg daily. Reports cough.',
                'synthetic-reviewed-voice',duration=1)):self.assertTrue(run_one())
        voice=Job.objects.get(kind='transcription',encounter=encounter)
        voice.refresh_from_db();self.assertEqual(voice.state,'succeeded',voice.error)
        raw_bytes=private_path(voice.result_path).read_bytes()
        self.assertContains(self.client.get(url),'Statements with no subject in this reviewed transcript are about this patient.')
        encounter.refresh_from_db()
        self.assertEqual(self.client.post(url+'notes/',{'revision':encounter.revision,'transcript_job':voice.pk,
            'corrected_transcript':'Takes warfarin 5 mg daily. Reports cough.','transcript_mode':'replace',
            'transcript_reviewed':'on','transcript_patient_context':'on'}).status_code,302)
        encounter.refresh_from_db()
        page=self.client.get(url);self.assertContains(page,'data-patient-context-current')
        self.assertEqual(self.client.post(url+'generate/',{'revision':encounter.revision,
            'usage_quote':page.context['usage_quote']}).status_code,302)
        generation=Job.objects.get(kind='generation',encounter=encounter)
        self.assertEqual(generation.config['patient_default_context'],tc.current(encounter))
        sent=[]
        def reply(**request):
            sent.append(request);payload=json.loads(request['messages'][0]['content'])
            self.assertEqual(payload['patient_default_context']['default_subject'],'encounter_patient')
            picks={f['key']:['Takes warfarin 5 mg daily.' if f['label']=='Medications' else 'Reports cough.']
                   for f in version.fields}
            body=json.dumps(relation_wire(encounter.notes,version.fields,picks))
            return SimpleNamespace(model=g.SONNET_5_5,stop_reason='end_turn',content=[SimpleNamespace(type='text',text=body)],
                id='msg_synthetic_reviewed_voice',_request_id='req_synthetic_reviewed_voice',
                usage=SimpleNamespace(input_tokens=100,output_tokens=20,cache_read_input_tokens=0,cache_creation_input_tokens=0),
                model_dump=lambda **kwargs:{'id':'msg_synthetic_reviewed_voice','text':body})
        with patch('anthropic.Anthropic',return_value=SimpleNamespace(messages=SimpleNamespace(create=reply))):
            self.assertTrue(run_one())
        generation.refresh_from_db();self.assertEqual(generation.state,'succeeded',generation.error)
        self.assertEqual(len(sent),1)
        encounter.refresh_from_db()
        self.assertEqual([encounter.answers[f['key']]['status'] for f in version.fields],['supported','supported'])
        self.assertEqual(Usage.objects.filter(job=voice).count(),1)
        self.assertEqual(Usage.objects.filter(job=generation).count(),1)
        payload={'revision':encounter.revision}
        for field in version.fields:
            payload['answer_'+field['key']]=encounter.answers[field['key']]['text']
            payload['reviewed_'+field['key']]='on'
        self.assertEqual(self.client.post(url+'save/',payload).status_code,302)
        encounter.refresh_from_db();self.assertTrue(tc.current(encounter))
        self.assertEqual(self.client.post(url+'finalize/',{'revision':encounter.revision,'reviewed':'on'}).status_code,302)
        encounter.refresh_from_db();self.assertEqual(encounter.status,'finalized')
        export=self.client.get(url+'export/');self.assertEqual(export.status_code,200)
        text=part_text(b''.join(export.streaming_content))
        self.assertIn('Takes warfarin 5 mg daily.',text);self.assertIn('Reports cough.',text)
        self.assertEqual(private_path(voice.result_path).read_bytes(),raw_bytes)
