"""A heading governs the lines it introduces; a finding written under it is never a ready answer.

⚠️ THE DEFECT THIS PINS, REPRODUCED ON 02c29f99 BEFORE ANY CHANGE. With a *Review of systems* question on
the form, the notes

    Review of systems: negative for the following:
    Symptoms: chest pain

came back with Symptoms SUPPORTED as "chest pain": pre-filled, needing only the review tick, and the
finalized export said the patient had chest pain. The heading parsed as the Review-of-systems answer, so
every line looked like a self-contained labelled answer (`scoped_lines`), and nothing downstream (save,
finalize, export) re-checks a supported answer. `Family history: the following conditions:` above
`Diabetes: yes` attributed the family's diabetes to the patient the same way.

The EXPECTED answers below were written from what each note says clinically, before the repair existed.
They do not depend on how the repair works. Everything here is offline: fixture generation or
provider-SHAPED selections built in the test (not retained provider output), and no provider call.
"""
from io import BytesIO
import json

from django.test import TestCase
from docx import Document

from core import generation
from core.generation import Selection, Selections, validate_selections
from core.jobs import run_one
from core.tests.test_workflow import DocumentWorkflow

FIXTURE = {'mode': 'fixture', 'prompt_version': 'legacy-offline-test'}
HEADING = 'Review of systems: negative for the following:'


def form(*labels):
    return [{'key': f'q{index}', 'label': label, 'kind': 'text', 'locator': {'cell': index}}
            for index, label in enumerate(labels)]


def fixture_answers(labels, notes):
    """What the offline generator returns, by question label."""
    fields = form(*labels)
    answers = generation.generate(notes, fields, FIXTURE)[0]
    return {field['label']: answers[field['key']] for field in fields}


def provider_answers(labels, notes, picks):
    """A provider-SHAPED response (built here, not retained): each question quotes its own full line,
    `conflict: false`, which is what the retained live WM response does for every one of its 90 answers."""
    fields = form(*labels)
    chosen = Selections(fields=[Selection(key=field['key'], quotes=picks.get(field['label'], []), conflict=False)
                                for field in fields])
    answers = validate_selections(notes, fields, chosen, evidence_contract='legacy')
    return {field['label']: answers[field['key']] for field in fields}


class ExpectationsMixin:
    """The two outcomes a clinician can be shown, stated independently of the implementation."""

    def assert_ready(self, answer, text):
        self.assertEqual((answer['status'], answer['text']), ('supported', text))

    def assert_needs_decision(self, answer, candidate, context):
        """Not a ready answer, not a one-click offer, and the candidate is kept with what governs it."""
        self.assertEqual(answer['status'], 'conflict', answer)
        self.assertEqual(answer['text'], '', 'a governed finding is never pre-filled')
        self.assertNotIn('suggestion', answer, 'a governed finding is never offered for one click')
        held = answer.get('held_candidate')
        self.assertIsNotNone(held, 'the candidate must be preserved for the clinician, not discarded')
        self.assertEqual(held['text'], candidate)
        quotes = [entry['quote'] for entry in answer['evidence']] + [entry['quote'] for entry in held['evidence']]
        self.assertTrue(any(context in quote for quote in quotes),
                        f'the governing line {context!r} must be shown with the candidate: {quotes}')


class TheReproducedDefect(ExpectationsMixin, TestCase):
    """E1-E8 of the milestone brief, on the offline generator and on provider-shaped selections."""

    def both(self, labels, notes, picks):
        yield 'fixture', fixture_answers(labels, notes)
        yield 'provider-shaped', provider_answers(labels, notes, picks)

    def test_e1_a_mapped_denial_heading_never_prefills_the_finding_below_it(self):
        notes = HEADING + '\nSymptoms: chest pain'
        for path, answers in self.both(['Review of systems', 'Symptoms'], notes,
                                       {'Review of systems': [HEADING], 'Symptoms': ['Symptoms: chest pain']}):
            with self.subTest(path=path):
                self.assert_needs_decision(answers['Symptoms'], 'chest pain', HEADING)
                self.assert_needs_decision(answers['Review of systems'], 'negative for the following:', HEADING)

    def test_e2_an_ordinary_labelled_line_before_the_heading_stays_ready_and_does_not_hide_it(self):
        notes = 'Indication: screening colonoscopy\n' + HEADING + '\nSymptoms: chest pain'
        for path, answers in self.both(['Indication', 'Review of systems', 'Symptoms'], notes,
                                       {'Indication': ['Indication: screening colonoscopy'],
                                        'Review of systems': [HEADING], 'Symptoms': ['Symptoms: chest pain']}):
            with self.subTest(path=path):
                self.assert_ready(answers['Indication'], 'screening colonoscopy')
                self.assert_needs_decision(answers['Symptoms'], 'chest pain', HEADING)

    def test_e3_an_unmapped_heading_keeps_the_candidate_instead_of_discarding_it(self):
        notes = 'Indication: screening colonoscopy\nROS: negative for the following:\nSymptoms: chest pain'
        for path, answers in self.both(['Indication', 'Symptoms'], notes,
                                       {'Indication': ['Indication: screening colonoscopy'],
                                        'Symptoms': ['Symptoms: chest pain']}):
            with self.subTest(path=path):
                self.assert_needs_decision(answers['Symptoms'], 'chest pain', 'ROS: negative for the following:')
                # The existing whole-note rule for an unlabelled line is NOT relaxed: Indication stays a
                # decision too, but its candidate is kept rather than lost.
                self.assertEqual(answers['Indication']['status'], 'conflict')
                self.assertEqual(answers['Indication'].get('held_candidate', {}).get('text'), 'screening colonoscopy')

    def test_e4_mixed_findings_only_the_governed_paragraph_needs_a_decision(self):
        notes = ('Chief complaint: right calf pain for three weeks\n' + HEADING + '\nSymptoms: chest pain, dyspnea\n\n'
                 'Examination: mild right ankle swelling\nPlan: duplex ultrasound')
        labels = ['Chief complaint', 'Review of systems', 'Symptoms', 'Examination', 'Plan']
        picks = {'Chief complaint': ['Chief complaint: right calf pain for three weeks'], 'Review of systems': [HEADING],
                 'Symptoms': ['Symptoms: chest pain, dyspnea'], 'Examination': ['Examination: mild right ankle swelling'],
                 'Plan': ['Plan: duplex ultrasound']}
        for path, answers in self.both(labels, notes, picks):
            with self.subTest(path=path):
                self.assert_ready(answers['Chief complaint'], 'right calf pain for three weeks')
                self.assert_needs_decision(answers['Symptoms'], 'chest pain, dyspnea', HEADING)
                self.assert_ready(answers['Examination'], 'mild right ankle swelling')
                self.assert_ready(answers['Plan'], 'duplex ultrasound')

    def test_e5_a_blank_line_after_the_heading_does_not_release_the_list_below_it(self):
        notes = HEADING + '\n\nSymptoms: chest pain'
        for path, answers in self.both(['Review of systems', 'Symptoms'], notes,
                                       {'Review of systems': [HEADING], 'Symptoms': ['Symptoms: chest pain']}):
            with self.subTest(path=path):
                self.assert_needs_decision(answers['Symptoms'], 'chest pain', HEADING)

    def test_e6_a_heading_with_nothing_after_it_governs_the_lines_above_it(self):
        closing = 'Review of systems: all of the above denied:'
        notes = 'Symptoms: chest pain\n' + closing
        for path, answers in self.both(['Symptoms', 'Review of systems'], notes,
                                       {'Symptoms': ['Symptoms: chest pain'], 'Review of systems': [closing]}):
            with self.subTest(path=path):
                self.assert_needs_decision(answers['Symptoms'], 'chest pain', closing)

    def test_e7_an_attribution_heading_never_gives_the_patient_the_familys_finding(self):
        heading = 'Family history: the following conditions:'
        notes = heading + '\nDiabetes: yes'
        for path, answers in self.both(['Family history', 'Diabetes'], notes,
                                       {'Family history': [heading], 'Diabetes': ['Diabetes: yes']}):
            with self.subTest(path=path):
                self.assert_needs_decision(answers['Diabetes'], 'yes', heading)

    def test_e8_a_line_initial_correction_keeps_the_candidate_without_offering_it(self):
        notes = 'Symptoms: no chest pain\nCorrection: severe crushing chest pain.'
        for path, answers in self.both(['Symptoms'], notes, {'Symptoms': ['Symptoms: no chest pain']}):
            with self.subTest(path=path):
                self.assert_needs_decision(answers['Symptoms'], 'no chest pain', 'Correction: severe crushing chest pain.')


class ValidAnswersStillWork(ExpectationsMixin, TestCase):
    """POSITIVE CONTROLS: without them, every assertion above could pass by blanking everything."""

    def test_p1_a_fully_labelled_affirmative_note_is_ready_unchanged(self):
        notes = ('Symptoms: intermittent right calf ache\nExamination: mild right ankle swelling; no skin changes\n'
                 'Plan: duplex ultrasound of both legs')
        answers = fixture_answers(['Symptoms', 'Examination', 'Plan'], notes)
        self.assert_ready(answers['Symptoms'], 'intermittent right calf ache')
        self.assert_ready(answers['Examination'], 'mild right ankle swelling; no skin changes')
        self.assert_ready(answers['Plan'], 'duplex ultrasound of both legs')

    def test_p2_a_denial_written_as_the_answer_itself_is_ready(self):
        notes = 'Symptoms: no chest pain\nReview of systems: negative for chest pain and dyspnea'
        answers = fixture_answers(['Symptoms', 'Review of systems'], notes)
        self.assert_ready(answers['Symptoms'], 'no chest pain')
        self.assert_ready(answers['Review of systems'], 'negative for chest pain and dyspnea')

    def test_p3_anatomical_and_backward_pointers_are_answers_not_headings(self):
        notes = 'Findings: No reflux below the knee.\nImpression: No DVT in the veins listed above.'
        answers = fixture_answers(['Findings', 'Impression'], notes)
        self.assert_ready(answers['Findings'], 'No reflux below the knee.')
        self.assert_ready(answers['Impression'], 'No DVT in the veins listed above.')

    def test_p4_a_colon_inside_a_value_is_not_a_heading(self):
        notes = 'Plan: review at 10:30\nFollow-up: in 2 weeks: sooner if worse'
        answers = fixture_answers(['Plan', 'Follow-up'], notes)
        self.assert_ready(answers['Plan'], 'review at 10:30')
        self.assert_ready(answers['Follow-up'], 'in 2 weeks: sooner if worse')


class TheDecisionIsEnforcedThroughFinalizationAndExport(DocumentWorkflow):
    """E1 on a real Word form, authenticated, through generate -> review -> finalize -> export."""
    NOTES = HEADING + '\nSymptoms: chest pain'
    TYPED = 'denies chest pain'

    def make_ros_docx(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Review of systems'
        table.cell(1, 0).text = 'Symptoms'
        buffer = BytesIO(); document.save(buffer)
        return buffer.getvalue()

    def generated(self):
        template, mapping = self.upload(self.make_ros_docx(), 'docx')
        encounter = self.create_encounter(mapping, self.NOTES)
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/generate/').status_code, 302)
        self.assertTrue(run_one())
        encounter.refresh_from_db()
        by_label = {f['label'].rstrip(':'): f['key'] for f in mapping.fields}
        self.ros, self.symptoms = by_label['Review of systems'], by_label['Symptoms']
        return encounter

    def save(self, encounter, **extra):
        payload = {'revision': encounter.revision,
                   'answer_' + self.ros: '', 'answer_' + self.symptoms: '',
                   'reviewed_' + self.ros: 'on', 'reviewed_' + self.symptoms: 'on'}
        payload.update(extra)
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/save/', payload).status_code, 302)
        encounter.refresh_from_db()

    def finalize_attempt(self, encounter):
        return self.client.post(f'/encounters/{encounter.pk}/finalize/',
                                {'reviewed': 'on', 'revision': encounter.revision}, follow=True)

    def test_the_governed_finding_is_generated_as_a_decision_not_an_answer(self):
        encounter = self.generated()
        answer = encounter.answers[self.symptoms]
        self.assertEqual((answer['status'], answer['text']), ('conflict', ''))
        self.assertNotIn('suggestion', answer)
        self.assertEqual(answer['held_candidate']['text'], 'chest pain')

    def test_ticking_reviewed_alone_cannot_finalize_and_nothing_is_exported(self):
        encounter = self.generated()
        self.save(encounter)
        response = self.finalize_attempt(encounter)
        encounter.refresh_from_db()
        self.assertNotEqual(encounter.status, 'finalized')
        self.assertContains(response, 'Symptoms')
        self.assertEqual(self.client.get(f'/encounters/{encounter.pk}/export/').status_code, 404)

    def test_the_held_candidate_cannot_be_accepted_by_a_crafted_post(self):
        encounter = self.generated()
        self.save(encounter, **{'accept_' + self.symptoms: 'on', 'accept_' + self.ros: 'on'})
        self.assertEqual(encounter.answers[self.symptoms]['text'], '')
        self.assertEqual(encounter.answers[self.symptoms]['status'], 'conflict')
        self.assertFalse(encounter.answers[self.symptoms].get('accepted_suggestion'))

    def test_a_typed_decision_finalizes_and_the_export_carries_it_not_the_candidate(self):
        encounter = self.generated()
        self.save(encounter, **{'answer_' + self.symptoms: self.TYPED,
                                'answer_' + self.ros: 'negative for chest pain'})
        exported = self.finalize(encounter)
        cells = [cell.text for row in Document(BytesIO(exported)).tables[0].rows for cell in row.cells]
        self.assertEqual(cells, ['Review of systems', 'negative for chest pain', 'Symptoms', self.TYPED])
