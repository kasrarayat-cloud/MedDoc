"""One bounded, populated synthetic rehearsal of the installed 0018 schema's forward upgrade.

This does not access an installed database, run an installer, qualify a provider,
or authorize schema reversal. Historical input rows and files are retained as
their own evidence, separately from today's authenticated read controls.
"""
import hashlib
import io
import json
import os
from datetime import date
from decimal import Decimal
from pathlib import Path
import tempfile
from unittest.mock import patch

from django.contrib.auth.hashers import make_password
from django.core.management import call_command
from django.db import connection
from django.db.migrations.executor import MigrationExecutor
from django.db.migrations.recorder import MigrationRecorder
from django.test import TransactionTestCase, override_settings
from django.utils import timezone
from docx import Document

from core import models
from core.security import private_path
from .test_workflow import TEST_SETTINGS

BASELINE = ('core', '0018_servicehealth')


def fingerprint(value):
    raw = json.dumps(value, sort_keys=True, default=str, separators=(',', ':')).encode()
    return hashlib.sha256(raw).hexdigest()


def historical_snapshot(apps):
    result = {}
    for model in apps.get_app_config('core').get_models(include_auto_created=True):
        columns = [f.attname for f in model._meta.local_fields]
        result[model._meta.label_lower] = list(model.objects.order_by('pk').values(*columns))
    return result


@override_settings(**TEST_SETTINGS)
class InstalledBaselineUpgradePreservation(TransactionTestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix='quillmd-synthetic-upgrade-')
        self.addCleanup(self.directory.cleanup)
        self.media = override_settings(MEDIA_ROOT=Path(self.directory.name) / 'private')
        self.media.enable(); self.addCleanup(self.media.disable)
        deny = patch('socket.socket.connect', side_effect=AssertionError('Network denied'))
        deny.start(); self.addCleanup(deny.stop)
        self.capture = None
        if value := os.environ.get('QUILLMD_WORD_ARTIFACTS'):
            self.capture = Path(value) / 'installed-baseline-upgrade'
            self.capture.mkdir(mode=0o700, parents=True, exist_ok=False)

    def tearDown(self):
        # Restore the forward graph even after a failed assertion, for orderly own-test-DB disposal.
        executor = MigrationExecutor(connection)
        executor.migrate(executor.loader.graph.leaf_nodes())
        super().tearDown()

    def evidence(self, name, value):
        if self.capture:
            (self.capture / name).write_text(json.dumps(value, indent=2, sort_keys=True, default=str))

    def file(self, key, data):
        path = private_path(key)
        path.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
        path.write_bytes(data); path.chmod(0o600)
        return hashlib.sha256(data).hexdigest()

    def seed_old_release(self, apps):
        def create(model_name, **values):
            return apps.get_model('core', model_name).objects.create(**values)
        now = timezone.now()
        user = create('User', username='upgrade-clinician', email='upgrade@synthetic.invalid',
                      password=make_password('SYNTHETIC-ONLY-upgrade'), verified_at=now)
        outsider = create('User', username='upgrade-outsider', email='upgrade-outsider@synthetic.invalid',
                          password=make_password('SYNTHETIC-ONLY-outsider'), verified_at=now)
        organization = create('Organization', name='Synthetic historical upgrade organization')
        department = create('Department', organization=organization, name='Historical authorized clinic',
                            organization_pays=True)
        other = create('Department', organization=organization, name='Historical other clinic')
        create('Membership', user=user, department=department, role='clinician')
        create('Membership', user=outsider, department=other, role='clinician')
        patient = create('Patient', department=department, display_name='SYNTHETIC HISTORICAL PATIENT',
                         reference='SYNTHETIC-UPGRADE', created_by=user)
        document = Document(); table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Symptoms'
        original_stream = io.BytesIO(); document.save(original_stream)
        original = original_stream.getvalue()
        table.cell(0, 1).text = 'Synthetic knee pain.'
        final_stream = io.BytesIO(); document.save(final_stream)
        exported = final_stream.getvalue()
        original_key, final_key, receipt_key = ('originals/historical.docx', 'exports/historical.docx',
                                                'receipts/historical.json')
        source_hash = self.file(original_key, original)
        export_hash = self.file(final_key, exported)
        receipt = b'{"scope":"synthetic historical metering fixture","units":{"input":100,"output":20}}\n'
        receipt_hash = self.file(receipt_key, receipt)
        template = create('Template', department=department, name='Synthetic historical original', format='docx',
                          original_name='historical.docx', original_path=original_key, sha256=source_hash, created_by=user)
        fields = [{'key': 'symptoms', 'label': 'Symptoms', 'kind': 'text', 'locator': {'cell': 0}}]
        version = create('TemplateVersion', template=template, number=1, fields=fields, created_by=user)
        answers = {'symptoms': {'text': 'Synthetic knee pain.', 'reviewed': True,
                               'reviewed_by': str(user.pk), 'status': 'supported', 'evidence': []}}
        final = create('Encounter', department=department, patient=patient, template_version=version,
                       created_by=user, notes='Symptoms: Synthetic knee pain.', service_date=date(2026, 9, 20),
                       answers=answers, status='finalized', revision=3, finalized_at=now, finalized_by=user,
                       export_path=final_key, export_sha256=export_hash)
        draft = create('Encounter', department=department, patient=patient, template_version=version,
                       created_by=user, notes='Symptoms: Synthetic resumed draft.', service_date=date(2026, 9, 21),
                       answers={}, revision=2)
        done = create('Job', encounter=final, user=user, state='succeeded', request_key='historical-completed',
                      revision=3, provider='synthetic:historical-fixture', payer='org:' + str(organization.pk),
                      config={'scope': 'synthetic historical metering; no real invoice'}, result_path=receipt_key,
                      result_sha256=receipt_hash, finished_at=now)
        queued = create('Job', encounter=draft, user=user, request_key='historical-reservation', revision=2,
                        config={'scope': 'synthetic historical queued job'}, payer='org:' + str(organization.pk))
        promo = create('Wallet', user=user, kind='promo', balance=Decimal('.90'))
        paid = create('Wallet', user=user, kind='paid', balance=Decimal('5'))
        allocation = create('Wallet', organization=organization, kind='paid', balance=Decimal('9.80'), reserved=Decimal('.40'))
        for wallet, amount, after, kind, key, job in (
            (promo, '1', '1', 'welcome', 'upgrade-promo-credit', None),
            (promo, '-.10', '.90', 'charge', 'upgrade-promo-charge', done),
            (paid, '5', '5', 'payment', 'upgrade-test-payment', None),
            (allocation, '10', '10', 'allocation', 'upgrade-org-allocation', None),
            (allocation, '-.20', '9.80', 'charge', 'upgrade-org-charge', done),
        ):
            create('LedgerEntry', wallet=wallet, amount=Decimal(amount), balance_after=Decimal(after), kind=kind,
                   key=key, job=job, description='Synthetic historical billing fixture')
        create('Reservation', job=queued, wallet=allocation, amount=Decimal('.40'))
        metering = dict(provider='synthetic:historical-fixture', actual_cost=Decimal('.20'),
                        units={'input': 100, 'output': 20}, rates={'scope': 'synthetic fixture'}, simulated=True)
        create('ProviderReceipt', job=done, result_path=receipt_key, result_sha256=receipt_hash, **metering)
        usage = create('Usage', job=done, user=user, department=department, organization=organization,
                       payer='org:' + str(organization.pk), charged_cost=Decimal('.30'), **metering)
        create('CostReconciliation', usage=usage, evidence_kind='unclassified', reconciled_cost=Decimal('.22'),
               evidence_reference='SYNTHETIC unresolved fixture; not authentic invoice evidence',
               recorded_by=user, charge_delta=Decimal('.03'), ledger_status='pending_funds')
        create('Payment', user=user, wallet=paid, amount_cents=500, stripe_session='synthetic-test-checkout',
               stripe_intent='synthetic-test-intent', status='paid', request_key='synthetic-payment-key')
        create('WebhookEvent', event_id='synthetic-test-event', kind='checkout.session.completed',
               payload_hash='1' * 64, object_data={'id': 'synthetic-test-checkout', 'amount_total': 500},
               semantic_hash='2' * 64, processed_at=now)
        release = create('CodeSet', system='ICD-10-CM', version='SYNTHETIC FY2026', effective_from=date(2025, 10, 1),
                         effective_until=date(2026, 10, 1), source_url='https://synthetic.invalid/fixture', sha256='3' * 64)
        code = create('Code', release=release, code='SYNTHETIC', description='Synthetic historical coding fixture')
        create('CodeSuggestion', encounter=final, terminology=code, explanation='Synthetic historical reviewed code',
               evidence='Symptoms: Synthetic knee pain.', accepted_by=user, accepted_at=now)
        create('AuditEvent', actor=user, organization=organization, department=department, action='synthetic.baseline',
               target=str(final.pk), details={'scope': 'Synthetic historical migration fixture'})
        return {'user': user.pk, 'outsider': outsider.pk, 'template': template.pk, 'version': version.pk,
                'final': final.pk, 'draft': draft.pk, 'files': {original_key: original, final_key: exported,
                                                               receipt_key: receipt}}

    def test_populated_0018_forward_upgrade_preserves_rows_files_balances_and_authorized_download(self):
        executor = MigrationExecutor(connection)
        executor.migrate([BASELINE])
        old_apps = executor.loader.project_state([BASELINE]).apps
        seeded = self.seed_old_release(old_apps)
        before = historical_snapshot(old_apps)
        self.evidence('baseline-rows.json', before)
        executor = MigrationExecutor(connection)
        targets = executor.loader.graph.leaf_nodes()
        plan = executor.migration_plan(targets)
        forward = [(m.app_label, m.name) for m, backwards in plan if not backwards and m.app_label == 'core']
        self.assertEqual(len(forward), 20)
        self.assertIn(('core', '0029_merge_batch_disposition'), forward)
        self.assertIn(('core', '0038_reinspection_confirmation_guard'), forward)
        source = Path(__file__).resolve().parents[1] / 'migrations'
        hashes = {name: hashlib.sha256((source / (name + '.py')).read_bytes()).hexdigest()
                  for _, name in forward}
        self.evidence('migration-plan.json', {'baseline': BASELINE, 'leaf': targets, 'forward': forward, 'sha256': hashes})
        executor.migrate(targets)
        # Also exercise normal post_migrate state publication without running an installed helper.
        call_command('migrate', interactive=False, verbosity=0)
        after = historical_snapshot(old_apps)
        self.evidence('upgraded-old-columns.json', after)
        self.assertEqual(after, before)
        self.assertEqual(fingerprint(after), fingerprint(before))
        for key, expected in seeded['files'].items():
            self.assertEqual(private_path(key).read_bytes(), expected)
            if self.capture:
                (self.capture / Path(key).name).write_bytes(expected)
        version = models.TemplateVersion.objects.get(pk=seeded['version'])
        self.assertEqual((version.contract, version.corrections, version.assistance, version.import_manifest),
                         ('', {}, {}, {}))
        self.assertIsNone(version.generation_constraints)
        self.assertFalse(version.needs_confirmation)
        self.assertIsNone(models.Payment.objects.get().livemode)
        self.assertIsNone(models.WebhookEvent.objects.get().livemode)
        coding = models.CodeSuggestion.objects.get()
        self.assertEqual((coding.release_sha256, coding.review_required), ('', False))
        self.client.force_login(models.User.objects.get(pk=seeded['user']))
        for route, expected in (
            (f"/templates/{seeded['template']}/original/", seeded['files']['originals/historical.docx']),
            (f"/encounters/{seeded['final']}/export/", seeded['files']['exports/historical.docx']),
        ):
            response = self.client.get(route)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(b''.join(response.streaming_content), expected)
        self.client.force_login(models.User.objects.get(pk=seeded['outsider']))
        for route in (f"/templates/{seeded['template']}/original/", f"/encounters/{seeded['final']}/export/"):
            self.assertEqual(self.client.get(route).status_code, 404)
        for label in ('core.wallet', 'core.ledgerentry', 'core.reservation', 'core.usage', 'core.providerreceipt',
                      'core.payment', 'core.webhookevent', 'core.costreconciliation', 'core.encounter',
                      'core.template', 'core.templateversion', 'core.codeset', 'core.code', 'core.codesuggestion'):
            self.assertEqual(historical_snapshot(old_apps)[label], before[label], label)
        applied = [row for row in MigrationRecorder(connection).applied_migrations() if row[0] == 'core']
        report = {'scope': 'synthetic own-database forward rehearsal; no installed mutation or provider calls',
                  'baseline': BASELINE, 'target': targets, 'applied_core_count': len(applied),
                  'old_columns_sha256': fingerprint(before), 'row_counts': {k: len(v) for k, v in before.items()},
                  'files': {k: {'bytes': len(v), 'sha256': hashlib.sha256(v).hexdigest()} for k, v in seeded['files'].items()},
                  'authorized_downloads': 2, 'denied_cross_department_downloads': 2,
                  'historical_payment_mode': 'unknown; NULL preserved',
                  'historical_invoice_status': 'unclassified synthetic evidence; no reconciliation inferred',
                  'migration_sha256': hashes, 'passed': True}
        self.evidence('verification.json', report)
