"""Self-service synthetic doctor pilot: what a new doctor is told, and the safeguards the pilot relies on.

Isolated synthetic database, fixture AI, no provider request. Each test names the guard it would
catch if it were removed.
"""
from io import BytesIO
import hashlib
import os
from pathlib import Path
import tempfile
import time
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.messages import get_messages
from django.core import signing
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from docx import Document
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from core.generation import configuration
from core.jobs import enqueue, run_one
from core.models import Department, Encounter, Job, Membership, Organization, Patient, Template, TemplateVersion, User
from core.security import private_path
from core.views import UPLOAD_REFUSAL_HINTS, VOICE_DOCUMENT_CHANGED, VOICE_QUOTE_EXPIRED, VOICE_QUOTE_SECONDS, mpeg_frame_sync
from documents import DocumentError
from core.tests.audio_fixtures import wav_bytes
from core.tests.test_billing_boundary import LIVE, BoundaryFixture, deepgram
from core.tests.test_billing_policy import STORAGES
from .test_security import SyntheticFixture
from .test_workflow import DocumentWorkflow
from .test_workflow import TEST_SETTINGS


def texts(response):
    return [str(message) for message in get_messages(response.wsgi_request)]


def docx_with_content_control():
    """A blank-looking form carrying one Word content control, the structure QuillMD refuses."""
    document = Document(); document.add_paragraph('Symptoms: ______')
    buffer = BytesIO(); document.save(buffer)
    source = ZipFile(BytesIO(buffer.getvalue()))
    xml = source.read('word/document.xml').decode()
    control = ('<w:sdt><w:sdtPr><w:alias w:val="Choice"/></w:sdtPr><w:sdtContent><w:p><w:r><w:t>Choose</w:t></w:r></w:p>'
               '</w:sdtContent></w:sdt>')
    xml = xml.replace('<w:sectPr', control + '<w:sectPr', 1)
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item, xml if item.filename == 'word/document.xml' else source.read(item.filename))
    return output.getvalue()



def fillable_pdf(link=False, name='Plan', tooltip='Plan'):
    """One multi-line Helvetica text field, the shape the upload page calls supported."""
    buffer = BytesIO(); document = canvas.Canvas(buffer)
    document.drawString(30, 780, 'Plan:')
    document.acroForm.textfield(name=name, tooltip=tooltip, x=30, y=600, width=420, height=140,
                                fontName='Helvetica', fontSize=10, fieldFlags='multiline')
    if link:
        document.linkURL('https://clinic.invalid/', (30, 60, 200, 80))
    document.showPage(); document.save()
    return buffer.getvalue()

@override_settings(**TEST_SETTINGS)
class NewDoctorPages(SyntheticFixture, TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(); self.addCleanup(directory.cleanup)
        override = override_settings(MEDIA_ROOT=Path(directory.name)); override.enable(); self.addCleanup(override.disable)
        # Exactly what views.verify creates for a self-service account: a workspace and nothing in it.
        self.doctor = User.objects.create_user(username='new-doctor', email='new-doctor@synthetic.invalid', verified_at=timezone.now())
        organization = Organization.objects.create(name='Synthetic workspace')
        self.department = Department.objects.create(organization=organization, name='Personal practice')
        Membership.objects.create(user=self.doctor, department=self.department, role='owner')
        self.client.force_login(self.doctor)

    def test_an_empty_workspace_points_to_a_supported_form_and_a_patient(self):
        page = self.client.get('/')
        self.assertContains(page, 'Upload a supported blank form')
        self.assertContains(page, '<a href="/patients/">Add a patient</a>')
        self.assertNotContains(page, 'Upload a Word or PDF form')

    def test_the_only_department_is_preselected_on_the_patient_and_upload_forms(self):
        for url in ('/patients/', '/templates/upload/'):
            with self.subTest(url=url):
                self.assertContains(self.client.get(url), f'<option value="{self.department.pk}" selected>')
        second = Department.objects.create(organization=self.department.organization, name='Second practice')
        Membership.objects.create(user=self.doctor, department=second, role='owner')
        self.assertNotContains(self.client.get('/patients/'), f'<option value="{self.department.pk}" selected>')

    def test_the_upload_page_explains_supported_and_refused_forms_before_upload(self):
        page = self.client.get('/templates/upload/')
        self.assertContains(page, 'data-supported-forms')
        for words in ('two-column question/answer tables', 'a label and an underlined Tab', 'not inside a table', 'Typed underscores', 'Fillable PDF', 'standard Helvetica font', 'not Auto',
                      'qualified WM IOV form', 'exactly as you received it', 'Refused at upload',
                      '<span class="mono">{{Question}}</span> placeholders', 'web or email links or linked images', 'with links, attachments or actions'):
            self.assertContains(page, words)
        self.assertNotContains(page, 'reported before export')

    def test_a_refused_content_control_says_what_the_doctor_can_do(self):
        response = self.client.post('/templates/upload/', {'name': 'Synthetic dropdown form', 'department': self.department.pk,
            'synthetic_only': 'on', 'file': SimpleUploadedFile('dropdowns.docx', docx_with_content_control())})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Template.objects.exists())
        self.assertContains(response, 'need the advanced document adapter')   # the engine's refusal is unchanged
        self.assertContains(response, 'The only supported form with dropdowns is the qualified WM IOV file')

    def marked_tab_form(self, where, mark):
        """The committed underlined-Tab fixture with one invisible mark in its first answer line."""
        from documents.test_tab_form import ORIGINAL
        with ZipFile(BytesIO(ORIGINAL.read_bytes())) as source:
            xml = source.read('word/document.xml').decode()
            index = xml.index('<w:tab/>')
            paragraph = xml.rindex('<w:p>', 0, index) if '<w:p>' in xml[:index] else xml.rindex('<w:p ', 0, index)
            run = xml.index('<w:r>', paragraph) if '<w:r>' in xml[paragraph:index] else xml.index('<w:r ', paragraph)
            at = run if where == 'paragraph' else xml.index('>', run) + 1
            xml = xml[:at] + mark + xml[at:]
            output = BytesIO()
            with ZipFile(output, 'w', ZIP_DEFLATED) as target:
                for item in source.infolist():
                    target.writestr(item, xml if item.filename == 'word/document.xml' else source.read(item.filename))
        return output.getvalue()

    def test_an_underlined_line_with_an_invisible_mark_gets_its_own_hint_not_the_content_control_one(self):
        # A spelling mark sits beside the runs; Word writes a page-break marker INSIDE the first run of
        # a line that starts a page. The doctor can see neither, so the hint names both.
        for where, mark in (('paragraph', '<w:proofErr w:type="spellStart"/>'), ('run', '<w:lastRenderedPageBreak/>')):
            with self.subTest(mark=mark):
                response = self.client.post('/templates/upload/', {'name': 'Synthetic marked form', 'department': self.department.pk,
                    'synthetic_only': 'on', 'file': SimpleUploadedFile('marked.docx', self.marked_tab_form(where, mark))})
                self.assertContains(response, 'underlined answer area crosses an unsupported structure')
                self.assertContains(response, 'must be an ordinary paragraph, not inside a table cell')
                self.assertContains(response, 'a page-break marker when the line starts a page')
                self.assertNotContains(response, 'The only supported form with dropdowns')
                self.assertFalse(Template.objects.exists())

    def test_a_linked_word_or_pdf_form_is_told_to_remove_the_link(self):
        from docx.opc.constants import RELATIONSHIP_TYPE
        document = Document(); document.add_paragraph('Symptoms: {{Symptoms}}')
        document.part.relate_to('https://clinic.invalid/', RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        buffer = BytesIO(); document.save(buffer)
        for name, data, refusal, hint in (
                ('linked.docx', buffer.getvalue(), 'External links and linked resources are not accepted',
                 'Remove every web or email link and linked image from the form'),
                ('linked.pdf', fillable_pdf(link=True), 'PDF actions, attachments, XFA, and external resources',
                 'Remove every link, attachment, document open action')):
            with self.subTest(name=name):
                response = self.client.post('/templates/upload/', {'name': 'Synthetic linked form', 'department': self.department.pk,
                    'synthetic_only': 'on', 'file': SimpleUploadedFile(name, data)})
                self.assertContains(response, refusal)
                self.assertContains(response, hint)
                self.assertFalse(Template.objects.exists())
        # CONTROL: the same PDF without its link is accepted, so the link is what was refused.
        response = self.client.post('/templates/upload/', {'name': 'Synthetic plain form', 'department': self.department.pk,
            'synthetic_only': 'on', 'file': SimpleUploadedFile('plain.pdf', fillable_pdf())})
        self.assertEqual(response.status_code, 302)

    def test_every_hint_is_keyed_on_exactly_one_real_engine_refusal(self):
        """If the engine's wording changes, a hint silently stops appearing; this keeps them aligned."""
        import re as regex
        root = Path(__file__).resolve().parents[2] / 'documents'
        refusals = [m for name in ('engine.py', 'tab_form.py')
                    for m in regex.findall(r"fail\(\s*(?:\"([^\"]+)\"|'([^']+)')", (root / name).read_text())]
        messages = [a or b for a, b in refusals]
        self.assertGreater(len(messages), 20)
        for marker, _ in UPLOAD_REFUSAL_HINTS:
            with self.subTest(marker=marker):
                self.assertEqual(sum(marker in message for message in messages), 1, marker)

    def test_an_encrypted_or_old_word_file_is_named_as_such(self):
        response = self.client.post('/templates/upload/', {'name': 'Synthetic protected form', 'department': self.department.pk,
            'synthetic_only': 'on', 'file': SimpleUploadedFile('protected.docx', b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + bytes(512))})
        self.assertContains(response, 'password-protected, encrypted, or an old Word .doc file')
        self.assertFalse(Template.objects.exists())

    def test_new_document_without_a_patient_links_to_patients(self):
        self.make_scope('someone-else')          # another organization's patient must not hide the notice
        page = self.client.get('/encounters/new/')
        self.assertContains(page, 'data-no-patient')
        self.assertContains(page, 'Add a synthetic patient on Patients')
        self.assertNotContains(page, 'The selected department determines')
        Patient.objects.create(department=self.department, display_name='SYNTHETIC ONE', reference='S1', created_by=self.doctor)
        self.assertNotContains(self.client.get('/encounters/new/'), 'data-no-patient')

    def test_the_qualified_form_page_lists_prewritten_statements_before_any_draft(self):
        blob = b'synthetic qualified original'
        template = Template.objects.create(department=self.department, name='Synthetic qualified', format='docx', original_name='q.docx',
                                           original_path='originals/q.docx', sha256=hashlib.sha256(blob).hexdigest(), created_by=self.doctor)
        requirement = {'key': 'claim_1', 'label': 'Prewritten statement: Right impression', 'kind': 'assertion', 'default_variant': 'selected',
                       'variants': [{'id': 'selected', 'label': 'Documented right finding', 'text_template': 'Right {sdt_1}{right_reflux_connector}{sdt_2}',
                                     'required_control_keys': ['sdt_1', 'sdt_2'], 'empty_control_keys': []},
                                    {'id': 'not_documented', 'label': 'Not documented', 'text_template': '   Right findings were not documented.',
                                     'required_control_keys': [], 'empty_control_keys': ['sdt_2']}]}
        identity = {'key': 'claim_2', 'label': 'Prewritten statement: Interpreting physician', 'kind': 'identity', 'default_variant': 'trusted',
                    'variants': [{'id': 'trusted', 'label': 'Signer', 'text_template': 'PHYSICIAN: {signer_name}'}]}
        TemplateVersion.objects.create(template=template, number=1, created_by=self.doctor, document_backend={'backend': 'synthetic-qualified'},
            fields=[{'key': 'sdt_1', 'label': 'Right vein', 'kind': 'choice'}, {'key': 'sdt_2', 'label': 'Right finding', 'kind': 'choice'}],
            review_requirements=[requirement, identity])
        page = self.client.get(f'/templates/{template.pk}/')
        self.assertContains(page, 'data-statement-guide')
        self.assertContains(page, 'Prewritten statement: Right impression')
        self.assertContains(page, 'Right [Right vein] [is] [Right finding]')
        self.assertContains(page, 'Right findings were not documented.')
        self.assertContains(page, 'Needs answers to: Right vein, Right finding.')
        self.assertContains(page, 'Needs these left unanswered: Right finding.')
        self.assertNotContains(page, 'Prewritten statement: Interpreting physician')   # identity is never written in notes
        # An ordinary form shows no statement guide.
        ordinary = self.make_scope('ordinary-page', self.department.organization)[4]
        Membership.objects.create(user=self.doctor, department=ordinary.department, role='clinician')
        self.assertNotContains(self.client.get(f'/templates/{ordinary.pk}/'), 'data-statement-guide')


@override_settings(**TEST_SETTINGS)
class PilotSafeguards(SyntheticFixture, TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(); self.addCleanup(directory.cleanup)
        override = override_settings(MEDIA_ROOT=Path(directory.name)); override.enable(); self.addCleanup(override.disable)
        self.scope = self.make_scope('pilot')
        self.user, _, self.department, self.patient, self.template, self.encounter, blob = self.scope
        path = private_path(self.template.original_path); path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(blob)
        self.client.force_login(self.user)

    def test_downloads_refuse_a_changed_original_or_export(self):
        private_path(self.template.original_path).write_bytes(b'tampered original')
        self.assertNotEqual(self.client.get(f'/templates/{self.template.pk}/original/').status_code, 200)
        export = private_path('exports/synthetic-export.docx'); export.parent.mkdir(parents=True, exist_ok=True); export.write_bytes(b'export')
        Encounter.objects.filter(pk=self.encounter.pk).update(status='finalized', export_path='exports/synthetic-export.docx',
            export_sha256=hashlib.sha256(b'export').hexdigest(), finalized_at=timezone.now(), finalized_by=self.user)
        export.write_bytes(b'tampered export')
        self.assertEqual(self.client.get(f'/encounters/{self.encounter.pk}/export/').status_code, 409)

    def test_the_qualified_form_never_offers_or_accepts_leave_blank(self):
        version = TemplateVersion.objects.create(template=self.template, number=2, fields=self.encounter.template_version.fields,
                                                 document_backend={'backend': 'synthetic-qualified'}, review_requirements=[], created_by=self.user)
        self.encounter.template_version = version; self.encounter.answers = {}
        self.encounter.save(update_fields=['template_version', 'answers'])
        with patch('core.views.wm_resolve', side_effect=DocumentError('Synthetic boundary; no Word runtime.')):
            page = self.client.get(f'/encounters/{self.encounter.pk}/')
            self.assertNotContains(page, 'name="blank_')
            revision = self.encounter.revision
            response = self.client.post(f'/encounters/{self.encounter.pk}/save/', {'revision': revision, 'answer_symptoms': '', 'blank_symptoms': 'on'})
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, 'cannot be marked', status_code=400)
        self.encounter.refresh_from_db()
        self.assertEqual((self.encounter.revision, self.encounter.answers), (revision, {}))

    def test_every_creation_form_requires_its_synthetic_attestation(self):
        """Without the tick nothing is created; with it the same request succeeds (the positive control)."""
        from documents.test_tab_form import ORIGINAL
        def attempt(model, url, data, attested, files=None):
            before = model.objects.count()
            self.client.post(url, {**data, **(attested if files is None else {**attested, **files()})} if attested else {**data, **(files() if files else {})})
            return model.objects.count() - before
        cases = [
            (Encounter, '/encounters/new/', {'patient': self.patient.pk, 'template_version': self.encounter.template_version.pk,
                                             'notes': 'Symptoms: fictional', 'service_date': timezone.localdate().isoformat()}, None),
            (Patient, '/patients/', {'department': self.department.pk, 'display_name': 'SYNTHETIC NEW', 'reference': 'SN-1'}, None),
            (Template, '/templates/upload/', {'name': 'Synthetic attested form', 'department': self.department.pk},
             lambda: {'file': SimpleUploadedFile('Colonoscopy_Operative_Template.docx', ORIGINAL.read_bytes())}),
        ]
        for model, url, data, files in cases:
            with self.subTest(form=url):
                self.assertEqual(attempt(model, url, data, None, files), 0)
                self.assertEqual(attempt(model, url, data, {'synthetic_only': 'on'}, files), 1)
        self.client.logout()
        account = {'first_name': 'Un', 'last_name': 'Attested', 'password1': 'synthetic-Passw0rd-long', 'password2': 'synthetic-Passw0rd-long'}
        self.assertEqual(attempt(User, '/register/', {**account, 'email': 'unattested@synthetic.invalid'}, None), 0)
        self.assertEqual(attempt(User, '/register/', {**account, 'email': 'attested@synthetic.invalid'}, {'synthetic_only': 'on'}), 1)

    def test_a_patient_and_form_from_different_departments_cannot_make_a_document(self):
        other = self.make_scope('pilot-sibling', self.scope[1])
        Membership.objects.create(user=self.user, department=other[2], role='clinician')
        encounters = Encounter.objects.count()
        response = self.client.post('/encounters/new/', {'patient': self.patient.pk, 'template_version': other[5].template_version.pk,
            'notes': 'Symptoms: fictional', 'service_date': timezone.localdate().isoformat(), 'synthetic_only': 'on'})
        self.assertContains(response, 'Patient and template must belong to the same department.')
        self.assertEqual(Encounter.objects.count(), encounters)

    def test_content_security_policy_allows_recording_playback_and_nothing_wider(self):
        policy = self.client.get('/')['Content-Security-Policy']
        self.assertIn("media-src 'self' blob:;", policy)
        self.assertIn("default-src 'self';", policy)
        self.assertNotIn('blob:', policy.replace("media-src 'self' blob:;", ''))


@override_settings(AI_MODE='fixture', ADMIN_SITE=False, STORAGES=STORAGES)
class PilotVoice(BoundaryFixture, TestCase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.user)
        self.encounter.answers = {}; self.encounter.save(update_fields=['answers'])
        self.url = f'/encounters/{self.encounter.pk}/'

    def quote_signed_ago(self, seconds):
        with patch.dict(os.environ, self.VOICE_ENV), patch('django.core.signing.time.time', return_value=time.time() - seconds):
            return signing.dumps({'user': str(self.user.pk), 'encounter': str(self.encounter.pk), 'revision': self.encounter.revision,
                                  'config': configuration()}, salt='usage-quote')

    def send(self, quote, name='synthetic.wav', content=None):
        upload = SimpleUploadedFile(name, content if content is not None else wav_bytes(seconds=1, tag=name))
        with patch.dict(os.environ, self.VOICE_ENV):
            return self.client.post(self.url + 'voice/', {'audio': upload, 'consent': 'on', 'usage_quote': quote})

    def uploads(self):
        folder = private_path('audio')
        return sorted(folder.iterdir()) if folder.exists() else []

    def test_a_recording_sent_twenty_nine_minutes_after_the_page_loaded_is_still_accepted(self):
        self.assertEqual(VOICE_QUOTE_SECONDS, 1800)   # the page promises 30 minutes
        response = self.send(self.quote_signed_ago(29 * 60))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Job.objects.filter(encounter=self.encounter, kind='transcription').count(), 1, texts(response))
        # Generation keeps its own shorter window: this widening is for the voice path only.
        with patch.dict(os.environ, self.VOICE_ENV):
            refused = self.client.post(self.url + 'generate/', {'usage_quote': self.quote_signed_ago(20 * 60), 'revision': self.encounter.revision})
        self.assertIn('Usage quote expired or changed. Reload to review the current maximum charge.', texts(refused))

    def test_a_document_change_is_named_as_the_reason_not_an_expired_price(self):
        quote = self.quote_signed_ago(0)
        # Answers saved in another window move the revision without touching the notes.
        saved = self.client.post(self.url + 'save/', {'revision': self.encounter.revision, 'answer_symptoms': 'Fictional answer only'})
        self.assertEqual(saved.status_code, 302)
        self.encounter.refresh_from_db()
        response = self.send(quote)
        self.assertIn(VOICE_DOCUMENT_CHANGED, texts(response))
        self.assertNotIn(VOICE_QUOTE_EXPIRED, texts(response))
        self.assertFalse(Job.objects.filter(encounter=self.encounter).exists())
        self.assertEqual(self.uploads(), [])

    def test_the_upload_check_uses_the_measuring_readers_mpeg_rule_for_every_two_byte_prefix(self):
        for first in range(256):
            for second in range(256):
                expected = first == 0xFF and second & 0xE0 == 0xE0     # core/audio.py's dispatch rule
                self.assertEqual(mpeg_frame_sync(bytes([first, second, 0, 0])), expected, (first, second))

    def test_an_expired_voice_quote_says_the_recording_was_not_kept_and_keeps_nothing(self):
        response = self.send(self.quote_signed_ago(31 * 60))
        self.assertIn(VOICE_QUOTE_EXPIRED, texts(response))
        self.assertFalse(Job.objects.filter(encounter=self.encounter).exists())
        self.assertEqual(self.uploads(), [])

    def test_unreadable_audio_names_the_supported_formats(self):
        response = self.send(self.quote_signed_ago(0), 'noise.wav', b'NOT AUDIO AT ALL, ONLY TEXT')
        self.assertTrue(any('Upload a .wav, .webm, .mp3 or .m4a recording' in text for text in texts(response)), texts(response))
        self.assertEqual(self.uploads(), [])

    def test_mpeg_frames_the_measuring_reader_accepts_pass_the_upload_check(self):
        # \xff\xfa (MPEG-1 Layer III with CRC) was refused by the view although core/audio.py reads it.
        reached = []
        def stop_at_enqueue(*args, **kwargs):
            reached.append(True); raise ValidationError('synthetic stop at enqueue')
        with patch('core.views.enqueue', side_effect=stop_at_enqueue):
            response = self.send(self.quote_signed_ago(0), 'frames.mp3', b'\xff\xfa\x90\x00' + bytes(64))
        self.assertEqual(reached, [True], texts(response))

    def test_a_partly_processed_silent_recording_shows_its_coverage_note_not_no_speech(self):
        response = self.send(self.quote_signed_ago(0), 'partial.wav', wav_bytes(seconds=10, tag='partial'))
        self.assertEqual(response.status_code, 302)
        with patch.dict(os.environ, self.VOICE_ENV), patch('core.jobs.httpx.post', return_value=deepgram('', 'synthetic-partial', duration=0)):
            self.assertTrue(run_one())
        page = self.client.get(self.url)
        self.assertContains(page, 'Transcription incomplete')
        self.assertNotContains(page, 'data-empty-transcript')

    def test_an_empty_transcript_is_explained_instead_of_an_unexplained_empty_box(self):
        response = self.send(self.quote_signed_ago(0), 'silence.wav')
        self.assertEqual(response.status_code, 302)
        with patch.dict(os.environ, self.VOICE_ENV), patch('core.jobs.httpx.post', return_value=deepgram('', 'synthetic-silence', duration=1)):
            self.assertTrue(run_one())
        job = Job.objects.get(encounter=self.encounter, kind='transcription')
        self.assertEqual(job.state, 'succeeded', job.error)
        page = self.client.get(self.url)
        self.assertContains(page, 'data-empty-transcript')
        self.assertContains(page, 'No speech was recognised in this recording')
        # Negative control: a transcript with speech never carries the notice.
        self.client.post(self.url + 'notes/', {'revision': self.encounter.revision, 'notes': 'Symptoms: fictional'})
        self.encounter.refresh_from_db()
        self.send(self.quote_signed_ago(0), 'speech.wav')
        with patch.dict(os.environ, self.VOICE_ENV), patch('core.jobs.httpx.post', return_value=deepgram('Fictional words spoken.', 'synthetic-speech', duration=1)):
            self.assertTrue(run_one())
        spoken = self.client.get(self.url)
        self.assertContains(spoken, 'Fictional words spoken.')
        # Both transcripts now await review; only the silent one carries the notice.
        self.assertContains(spoken, 'data-empty-transcript', count=1)


@override_settings(ADMIN_SITE=False, STORAGES=STORAGES, **LIVE)
class PilotProviderResults(BoundaryFixture, TestCase):
    def test_a_provider_result_is_delivered_only_inside_its_department(self):
        job = self.live_job('synthetic-pilot-result')
        self.assertEqual(job.state, 'succeeded', job.error)
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(f'/jobs/{job.pk}/result/').status_code, 200)   # positive control: the owner can
        for outsider in (self.make_scope('pilot-outsider')[0], self.make_scope('pilot-sibling', self.scope[1])[0]):
            with self.subTest(outsider=outsider.email):   # another organization, and a sibling department
                self.client.force_login(outsider)
                self.assertEqual(self.client.get(f'/jobs/{job.pk}/result/').status_code, 404)


class BrowserLineBreaksReachTheDocumentOnce(DocumentWorkflow):
    """A browser posts a textarea's line breaks as CRLF. Stored as posted, a two-line answer could
    never finalize on a fillable PDF (a carriage return is not printable) and became a blank line
    between its lines in a Word table cell."""
    TYPED = 'duplex ultrasound of both legs\r\nreview in two weeks'
    STORED = 'duplex ultrasound of both legs\nreview in two weeks'

    def test_a_two_line_pdf_answer_finalizes_with_one_line_break(self):
        template, mapping = self.upload(fillable_pdf(), 'pdf')
        encounter = self.create_encounter(mapping, 'Plan: duplex ultrasound.')
        self.generate(encounter)
        self.review(encounter, override=self.TYPED)
        self.assertEqual(encounter.answers['Plan']['text'], self.STORED)
        exported = self.finalize(encounter)
        self.assertEqual(PdfReader(BytesIO(exported)).get_form_text_fields()['Plan'], self.STORED)

    def test_a_two_line_table_answer_exports_without_an_extra_blank_line(self):
        template, mapping = self.upload(self.make_docx(), 'docx')
        encounter = self.create_encounter(mapping, 'Symptoms: No chest pain.\nExamination: Normal gait.')
        self.generate(encounter)
        self.review(encounter, override=self.TYPED)
        self.assertTrue(all(answer['text'] == self.STORED for answer in encounter.answers.values()))
        exported = self.finalize(encounter)
        self.assertEqual(Document(BytesIO(exported)).tables[0].cell(0, 1).text, self.STORED)

    def test_a_character_a_pdf_cannot_hold_is_refused_naming_the_question(self):
        template, mapping = self.upload(fillable_pdf(), 'pdf')
        encounter = self.create_encounter(mapping, 'Plan: duplex ultrasound.')
        self.generate(encounter)
        for typed, where in (('review in two weeks \u2014 sooner if worse', '(\u201c\u2014\u201d, U+2014 at character 21)'),
                             ('review\u00a0in two weeks', '(U+00A0 at character 7)')):
            with self.subTest(where=where):
                self.review(encounter, override=typed)
                response = self.client.post(f'/encounters/{encounter.pk}/finalize/', {'reviewed': 'on', 'revision': encounter.revision}, follow=True)
                self.assertContains(response, 'The answer for Plan has a character a PDF form cannot hold ' + where)
                encounter.refresh_from_db()
                self.assertEqual(encounter.status, 'review')

    def test_the_refusal_names_a_renamed_question_as_the_doctor_sees_it(self):
        """The engine re-inspects the original, whose field is still `Text1`; the doctor renamed it."""
        import json
        template, mapping = self.upload(fillable_pdf(name='Text1', tooltip=None), 'pdf')
        self.assertEqual(mapping.fields[0]['label'], 'Text1')
        response = self.client.post(f'/templates/{template.pk}/', {'labels': json.dumps({'Text1': 'Follow-up plan'}), 'mapping_version': str(mapping.pk)})
        self.assertEqual(response.status_code, 302)
        renamed = template.versions.get(number=2)
        encounter = self.create_encounter(renamed, 'Follow-up plan: duplex ultrasound.')
        self.generate(encounter)
        self.review(encounter, override='review in two weeks \u2014 sooner if worse')
        response = self.client.post(f'/encounters/{encounter.pk}/finalize/', {'reviewed': 'on', 'revision': encounter.revision}, follow=True)
        self.assertContains(response, 'The answer for Follow-up plan has a character a PDF form cannot hold')
        self.assertNotContains(response, 'The answer for Text1')


class ALabelledNoteWithAnatomyFinalizes(DocumentWorkflow):
    """Round 5, end to end: a vein note's ordinary `below the knee` must not stop a labelled note."""

    def test_no_reflux_below_the_knee_is_an_answer_and_the_document_finalizes(self):
        template, mapping = self.upload(self.make_docx(), 'docx')
        encounter = self.create_encounter(mapping, 'Symptoms: No reflux below the knee.\nExamination: Normal venous study.')
        self.generate(encounter)          # asserts every answer came back supported
        self.review(encounter)
        exported = self.finalize(encounter)
        self.assertEqual(Document(BytesIO(exported)).tables[0].cell(0, 1).text, 'Normal venous study.')
