"""Exact saved-state approval, with independently valid source-bound answers."""
from copy import deepcopy
import hashlib
from pathlib import Path
import tempfile
from unittest.mock import patch

from docx import Document
from django.http import QueryDict
from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from django.utils import timezone

from documents import inspect_template, DocumentError
from documents.semantic_preparation_roles import default_definitions
from core import document_review as review, semantic_generation as semantic
from core.models import AuditEvent, Department, Encounter, Membership, Organization, Patient, Template, TemplateVersion, User, Job
from core.security import private_path
from core.tests.test_semantic_generation import fact, assignments, FixtureSender
from core.tests.test_workflow import TEST_SETTINGS

@override_settings(**TEST_SETTINGS)
class DocumentApprovalTests(TestCase):
    def setUp(self):
        directory=tempfile.TemporaryDirectory(); self.addCleanup(directory.cleanup)
        settings=override_settings(MEDIA_ROOT=Path(directory.name));settings.enable();self.addCleanup(settings.disable)
        self.actor=User.objects.create_user(username='approval-fixture',email='approval@synthetic.invalid',verified_at=timezone.now(),first_name='Synthetic',last_name='Clinician')
        org=Organization.objects.create(name='Synthetic approval practice')
        self.department=Department.objects.create(organization=org,name='X-ray')
        Membership.objects.create(user=self.actor,department=self.department,role='clinician')
        patient=Patient.objects.create(department=self.department,created_by=self.actor,display_name='Synthetic review patient')
        document=Document();table=document.add_table(rows=2,cols=2)
        table.cell(0,0).text='History';table.cell(1,0).text='Drug allergies'
        self.original=private_path('fixture.docx');document.save(self.original)
        fields=inspect_template(self.original)['fields'];self.fields=fields
        self.definitions=default_definitions(fields)
        template=Template.objects.create(department=self.department,created_by=self.actor,name='Synthetic review',format='docx',original_name='fixture.docx',original_path='fixture.docx',sha256=hashlib.sha256(self.original.read_bytes()).hexdigest())
        version=TemplateVersion.objects.create(template=template,number=1,fields=fields,created_by=self.actor)
        self.notes='Right knee stiffness. She takes a multivitamin.'
        facts=[fact(self.notes,fields,'symptom','Right knee stiffness',kind='history'),fact(self.notes,fields,'medicine','a multivitamin',kind='history')]
        self.answers=[]
        for identity in ('symptom','medicine'):
            selected=assignments(fields,{fields[0]['key']:[identity]})
            self.answers.append(semantic.run(self.notes,fields,self.definitions,FixtureSender(facts,selected),scope_contract=semantic.DURATION_SCOPE_CONTRACT))
        self.encounter=Encounter.objects.create(department=self.department,patient=patient,created_by=self.actor,template_version=version,notes=self.notes,service_date=timezone.localdate(),answers=self.answers[0],status='review')
        # This approval-boundary fixture supplies already approved definitions;
        # it does not stand in for live fresh preparation acceptance.
        definitions=patch('core.semantic_preparation.validated_semantic_definitions',return_value=self.definitions)
        definitions.start();self.addCleanup(definitions.stop)
        self.client.force_login(self.actor)
        self.route=f'/encounters/{self.encounter.pk}/review/'

    def post(self,intent='approve'):
        self.encounter.refresh_from_db()
        post={'intent':intent,'review_token':review.review_token(self.encounter,self.actor,self.definitions)}
        for field in self.fields:
            key=field['key'];post['answer_'+key]=self.encounter.answers[key]['text'];post['decision_'+key]='keep'
        return post

    def assert_unapproved(self):
        self.encounter.refresh_from_db()
        self.assertNotEqual(self.encounter.status,'finalized')
        self.assertFalse(self.encounter.export_path)
        self.assertFalse(AuditEvent.objects.filter(action='encounter.document_approve').exists())

    def test_one_approval_keeps_optional_missing_and_exports_exact_cells(self):
        before=deepcopy(self.encounter.answers)
        self.assertEqual(self.client.post(self.route,self.post()).status_code,302)
        self.encounter.refresh_from_db()
        self.assertEqual(self.encounter.status,'finalized')
        self.assertEqual(self.encounter.answers,before)
        self.assertFalse(any(a.get('reviewed') for a in self.encounter.answers.values()))
        missing=self.encounter.answers[self.fields[1]['key']]
        self.assertEqual(missing['status'],'missing');self.assertNotIn('blank_decision',missing)
        receipt=self.encounter.template_review['document_approval']
        self.assertEqual(receipt['answers_sha256'],semantic.digest(before))
        self.assertEqual(receipt['actor'],str(self.actor.pk))
        self.assertEqual(receipt['approved_revision'],self.encounter.revision)
        doc=Document(private_path(self.encounter.export_path))
        self.assertEqual([r.cells[1].text for r in doc.tables[0].rows],['Right knee stiffness',''])
        self.assertEqual([r.cells[0].text for r in doc.tables[0].rows],['History','Drug allergies'])
        self.assertEqual(AuditEvent.objects.filter(action='encounter.document_approve').count(),1)
        finalized=deepcopy(receipt);sha=self.encounter.export_sha256
        self.assertEqual(self.client.post(self.route,{}).status_code,302)
        self.encounter.refresh_from_db()
        self.assertEqual(self.encounter.template_review['document_approval'],finalized)
        self.assertEqual(self.encounter.export_sha256,sha)
        self.assertEqual(AuditEvent.objects.filter(action='encounter.document_approve').count(),1)

    def test_approved_soap_name_persists_through_review_finalization_and_download(self):
        from io import BytesIO
        from core.tests.test_document_identity import soap_definitions
        document=Document();table=document.add_table(rows=4,cols=2)
        labels=['Subjective','Objective','Assessment','Plan']
        for row,label in zip(table.rows,labels):row.cells[0].text=label
        self.original=private_path('soap.docx')
        document.save(self.original)
        original_bytes=self.original.read_bytes()
        self.fields=inspect_template(self.original)['fields']
        self.definitions.clear();self.definitions.update(soap_definitions(self.fields))
        template=Template.objects.create(department=self.department,created_by=self.actor,name='SOAP fixture',
            format='docx',original_name='soap.docx',original_path='soap.docx',sha256=hashlib.sha256(original_bytes).hexdigest())
        version=TemplateVersion.objects.create(template=template,number=1,fields=self.fields,created_by=self.actor)
        self.encounter.template_version=version
        patient=self.encounter.patient;patient.display_name='Synthetic Avery Sen';patient.save(update_fields=['display_name'])
        notes='Synthetic Avery Sen reports heartburn and constipation.'
        facts=[fact(notes,self.fields,'name','Synthetic Avery Sen',kind='identity',name_quote='Synthetic Avery Sen'),
               fact(notes,self.fields,'symptoms','heartburn and constipation',kind='history')]
        generated=semantic.run(notes,self.fields,self.definitions,FixtureSender(facts,
            assignments(self.fields,{'cell_0':['name','symptoms']})),
            scope_contract=semantic.DOCUMENT_IDENTITY_SCOPE_CONTRACT)
        self.encounter.notes=notes;self.encounter.answers=generated;self.encounter.save()
        self.encounter.refresh_from_db()
        self.assertEqual(self.encounter.answers,generated)
        self.assertEqual(generated['cell_0']['text'],'Synthetic Avery Sen; Heartburn and constipation')
        response=self.client.get(f'/encounters/{self.encounter.pk}/')
        self.assertContains(response,generated['cell_0']['text'])
        self.assertEqual(response.content.decode().count('name="intent" value="approve"'),1)
        self.assertEqual(self.client.post(self.route,self.post()).status_code,302)
        self.encounter.refresh_from_db()
        self.assertEqual(self.encounter.answers,generated)
        self.assertEqual(self.encounter.status,'finalized')
        receipt=self.encounter.template_review['document_approval']
        self.assertEqual(receipt['answers_sha256'],semantic.digest(generated))
        self.assertEqual(receipt['approved_revision'],self.encounter.revision)
        download=self.client.get(f'/encounters/{self.encounter.pk}/export/')
        self.assertEqual(download.status_code,200)
        downloaded=b''.join(download.streaming_content)
        self.assertEqual(hashlib.sha256(downloaded).hexdigest(),self.encounter.export_sha256)
        word=Document(BytesIO(downloaded))
        self.assertEqual([r.cells[0].text for r in word.tables[0].rows],labels)
        self.assertEqual([r.cells[1].text for r in word.tables[0].rows],
                         [generated[f['key']]['text'] for f in self.fields])
        self.assertEqual(len(word.tables),1)
        self.assertEqual(self.original.read_bytes(),original_bytes)
        self.assertFalse(any(a.get('reviewed') or a.get('semantic_edit_history') for a in generated.values()))

    def test_same_revision_new_valid_answers_and_same_text_metadata_refuse_old_review(self):
        for newer in [self.answers[1],{**self.answers[0],self.fields[0]['key']:{**self.answers[0][self.fields[0]['key']], 'reviewed_by':str(self.actor.pk)}}]:
            self.encounter.answers=deepcopy(self.answers[0]);self.encounter.save()
            post=self.post()
            for field in self.fields:semantic.validate_stored_answer(self.notes,field,self.definitions[field['key']],newer[field['key']])
            self.encounter.answers=deepcopy(newer);self.encounter.save()
            self.assertEqual(self.client.post(self.route,post).status_code,409)
            self.assert_unapproved();self.assertEqual(self.encounter.answers,newer)

    def test_revision_actor_template_source_and_signature_are_bound(self):
        old=self.post()
        self.encounter.revision+=1;self.encounter.save()
        self.assertEqual(self.client.post(self.route,old).status_code,409);self.assert_unapproved()
        tampered=self.post();tampered['review_token']+='bad'
        self.assertEqual(self.client.post(self.route,tampered).status_code,409);self.assert_unapproved()
        other=User.objects.create_user(username='other',email='other@synthetic.invalid',verified_at=timezone.now())
        Membership.objects.create(user=other,department=self.department,role='clinician')
        post=self.post();self.client.force_login(other)
        self.assertEqual(self.client.post(self.route,post).status_code,409);self.assert_unapproved()

    def test_pending_jobs_and_old_finalization_route_refuse(self):
        for state in review.PENDING:
            job=Job.objects.create(encounter=self.encounter,user=self.actor,request_key='review-'+state,state=state)
            self.assertEqual(self.client.post(self.route,self.post()).status_code,409);self.assert_unapproved()
            job.delete()
        self.client.post(f'/encounters/{self.encounter.pk}/finalize/',{'revision':0,'reviewed':'on'})
        self.assert_unapproved()
        self.assertEqual(self.client.post(f'/encounters/{self.encounter.pk}/save/',{'revision':0}).status_code,409)

    def test_missing_duplicate_controls_or_implicit_submit_cannot_approve(self):
        for key in ['intent','review_token','answer_'+self.fields[0]['key'],'decision_'+self.fields[0]['key']]:
            post=self.post();post.pop(key)
            self.assertEqual(self.client.post(self.route,post).status_code,409);self.assert_unapproved()
        post=self.post();post['answer_'+self.fields[0]['key']]=['Right knee stiffness','different']
        self.assertEqual(self.client.post(self.route,post).status_code,409);self.assert_unapproved()
        post=self.post('save')
        self.assertEqual(self.client.post(self.route,post).status_code,302);self.assert_unapproved()

    def test_conflict_needs_actual_decision_and_keeps_history_without_per_field_review(self):
        key=self.fields[0]['key'];answer=self.encounter.answers[key]
        answer.update(status='conflict',text='',candidate_text='Right knee stiffness')
        self.encounter.save()
        self.assertEqual(self.client.post(self.route,self.post()).status_code,409);self.assert_unapproved()
        post=self.post();post['decision_'+key]='blank'
        self.assertEqual(self.client.post(self.route,post).status_code,302)
        self.encounter.refresh_from_db();answer=self.encounter.answers[key]
        self.assertEqual(answer['status'],'left_blank');self.assertFalse(answer.get('reviewed'))
        self.assertEqual(answer['blank_decision']['by'],str(self.actor.pk))
        self.assertEqual(answer['semantic_edit_history'][-1]['previous']['candidate_text'],'Right knee stiffness')

    def test_export_failure_does_not_save_edits_or_approval(self):
        old=deepcopy(self.encounter.answers);post=self.post()
        post['answer_'+self.fields[0]['key']]='Documented clinician correction'
        with patch('documents.fill_template',side_effect=DocumentError('Export failed')):
            self.assertEqual(self.client.post(self.route,post).status_code,409)
        self.assert_unapproved();self.assertEqual(self.encounter.answers,old);self.assertEqual(self.encounter.revision,0)

    def test_compact_ui_has_no_routine_checkboxes_and_one_approval(self):
        response=self.client.get(f'/encounters/{self.encounter.pk}/')
        self.assertEqual(response.status_code,200)
        html=response.content.decode()
        self.assertIn('Not documented',html)
        self.assertNotIn('name="reviewed_cell_',html)
        self.assertNotIn('I reviewed this answer',html)
        self.assertEqual(html.count('name="intent" value="approve"'),1)
        self.assertNotIn('document_approval',self.encounter.template_review)

    def test_signature_identity_and_required_omission_cannot_be_waived(self):
        from types import SimpleNamespace
        from core.semantic_review import save_answer
        fields=[{'key':'signature','label':'Physician signature','kind':'text','locator':{'cell':0}}]
        definitions=default_definitions(fields)
        empty=semantic.run('',fields,definitions,FixtureSender([],assignments(fields)))
        encounter=SimpleNamespace(notes='',department=self.department,template_version=SimpleNamespace(fields=fields))
        review.approval_blockers(encounter,empty,definitions,self.actor)
        for name,allowed in [('Synthetic Clinician',True),('Another Doctor',False)]:
            answer=save_answer(encounter,fields[0],empty['signature'],name,self.actor,False,False,1,definitions,document_level=True)
            self.assertNotIn('reviewed',answer)
            if allowed:review.approval_blockers(encounter,{'signature':answer},definitions,self.actor)
            else:
                with self.assertRaises(ValidationError):review.approval_blockers(encounter,{'signature':answer},definitions,self.actor)
        fields[0]['required']=True
        with self.assertRaises(ValidationError):review.approval_blockers(encounter,empty,definitions,self.actor)

    def test_unchanged_browser_newlines_preserve_original_bytes_and_history(self):
        from core.semantic_review import save_answer
        key=self.fields[0]['key']
        for value in ['\nRight knee stiffness','Right knee stiffness\r\nNo fever']:
            self.encounter.answers[key]=save_answer(self.encounter,self.fields[0],self.answers[0][key],value,self.actor,False,False,1,self.definitions,document_level=True)
            self.encounter.save();before=deepcopy(self.encounter.answers);post=self.post('save')
            post['answer_'+key]=value.replace('\r\n','\n')
            self.assertEqual(self.client.post(self.route,post).status_code,302)
            self.encounter.refresh_from_db();self.assertEqual(self.encounter.answers,before)


from django.db import transaction, close_old_connections, connections
from django.test import TransactionTestCase, Client
import threading

@override_settings(**TEST_SETTINGS)
class DocumentApprovalConcurrencyTests(TransactionTestCase):
    setUp=DocumentApprovalTests.setUp
    post=DocumentApprovalTests.post
    assert_unapproved=DocumentApprovalTests.assert_unapproved

    def test_concurrent_source_edit_wins_before_waiting_approval(self):
        post=self.post();client=Client();client.force_login(self.actor)
        started=threading.Event();results=[]
        def approve():
            close_old_connections();started.set()
            try:results.append(client.post(self.route,post).status_code)
            finally:connections.close_all()
        with transaction.atomic():
            current=Encounter.objects.select_for_update().get(pk=self.encounter.pk)
            thread=threading.Thread(target=approve);thread.start();self.assertTrue(started.wait(5))
            current.notes+=' A source correction.';current.revision+=1;current.save()
        thread.join(15);self.assertFalse(thread.is_alive());self.assertEqual(results,[409]);self.assert_unapproved()
        self.assertTrue(self.encounter.notes.endswith('A source correction.'))

    def test_parallel_duplicate_approval_creates_exactly_one_receipt_and_export(self):
        post=self.post();clients=[Client(),Client()]
        for client in clients:client.force_login(self.actor)
        barrier=threading.Barrier(2);results=[]
        def approve(client):
            close_old_connections()
            try:
                barrier.wait(5);results.append(client.post(self.route,post).status_code)
            finally:connections.close_all()
        threads=[threading.Thread(target=approve,args=(client,)) for client in clients]
        for thread in threads:thread.start()
        for thread in threads:thread.join(15);self.assertFalse(thread.is_alive())
        self.assertEqual(sorted(results),[302,302])
        self.assertEqual(AuditEvent.objects.filter(action='encounter.document_approve').count(),1)
        self.assertEqual(AuditEvent.objects.filter(action='encounter.finalize').count(),1)
        self.assertEqual(len(list(private_path('exports').glob('*.docx'))),1)
