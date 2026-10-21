"""Prospective tagged units: exact canonical values, no inference or old repair."""
from copy import deepcopy
import json
from pathlib import Path

from anthropic import transform_schema
from django.test import SimpleTestCase
from docx import Document

from core import preparation, preparation_choice_transport as codec, word_contract
from core.security import private_path
from core.tests import test_preparation_schema_transport as legacy
from core.tests import test_preparation_stream_transport as sdk_tests
from core.tests import test_preparation_compact_components as compact
from core.tests.test_provider_schema import expand
from documents import DocumentError
from documents.component_contracts import CONFIG_KEY, CHOICE_CONTRACT
from documents.preparation import TYPED_PROTOCOL, digest, inventory


class PreparationUnitPolicyTests(SimpleTestCase):
    setUp = legacy.PreparationSchemaTransportTests.setUp

    def live(self, transport=codec.UNIT_POLICY_TRANSPORT):
        return {**self.config, 'mode': 'live', 'word_preparation_protocol': TYPED_PROTOCOL,
                CONFIG_KEY: CHOICE_CONTRACT, preparation.SCHEMA_TRANSPORT_KEY: transport}

    def canonical(self, role='proposer', proposal=None, *, inv=None):
        fixture = {**self.live(), 'mode': 'fixture'}
        fixture.pop(preparation.SCHEMA_TRANSPORT_KEY)
        fixture.pop(preparation.stream_transport.CONFIG_KEY, None)
        fixture.pop(preparation.REASONING_KEY, None)
        return preparation.provider_role(self.inventory if inv is None else inv,
                                        role, fixture, 1, proposal)[0]

    def evidence(self, wire, role='proposer', proposal=None, *, inv=None):
        inv = self.inventory if inv is None else inv
        envelope = preparation.provider_request(inv, role, self.live(), proposal)
        return {'response_transport': codec.UNIT_POLICY_TRANSPORT, 'request_envelope': envelope,
                'request_envelope_sha256': digest(envelope), 'response': {'stop_reason': 'end_turn',
                'content': [{'type': 'text', 'text': json.dumps(wire, ensure_ascii=False)}]}}

    def decode(self, wire, role='proposer', proposal=None, *, evidence=None, inv=None):
        inv = self.inventory if inv is None else inv
        return preparation.decode_role_response(inv, role, self.live(), wire, proposal,
            self.evidence(wire, role, proposal, inv=inv) if evidence is None else evidence,
            digest(preparation.role_payload(inv, role, proposal, config=self.live())))

    @staticmethod
    def replace_rows(wire, rows):
        next(codec.contracts(wire))['components'] = json.dumps(rows, ensure_ascii=False)
        return wire

    def rows(self, wire):
        return json.loads(next(codec.contracts(wire))['components'])

    def test_eleven_positions_restore_all_twelve_keys_and_three_scope_values(self):
        self.assertEqual(codec.UNIT_POLICY_COLUMNS, ('id', 'question_quote', 'label', 'meaning',
            'value_type', 'scope', 'unit_policy', 'required', 'encounter_scope', 'choices', 'option_quote'))
        proposal = self.canonical()
        for role in ('proposer', 'checker'):
            value = self.canonical(role, proposal)
            before = deepcopy(value)
            wire = codec.encode_unit_policy_rows(value)
            self.assertEqual(len(self.rows(wire)[0]), 11)
            self.assertEqual(self.rows(wire)[0][6], ['none'])
            decoded = self.decode(wire, role, proposal)
            self.assertEqual(decoded, before)
            self.assertEqual(value, before)
            self.assertEqual(set(next(codec.components(decoded))), set(codec.COMPONENT_COLUMNS))
            word_contract.validate_result(self.inventory, decoded, role, self.live())

    def test_all_disjoint_policies_roundtrip_without_normalizing_invalid_pairs(self):
        for unit, requirement in ((None, 'none'), (None, 'documented'), ('µmol/L', 'printed'),
                                  ('null', 'printed'), ('"é"\\unit', 'printed'), ('x' * 80, 'printed')):
            value = self.canonical()
            next(codec.components(value)).update(unit=unit, unit_requirement=requirement)
            self.assertEqual(codec.decode(codec.encode_unit_policy_rows(value), codec.UNIT_POLICY_TRANSPORT), value)
        invalid_pairs = [('minutes', 'documented'), ('minutes', 'none'), (None, 'printed'),
                         ('', 'printed'), ('  ', 'printed'), ('x' * 81, 'printed'), (False, 'printed'),
                         (0, 'printed'), ('minutes', None), (None, 'unknown')]
        for unit, requirement in invalid_pairs:
            value = self.canonical()
            next(codec.components(value)).update(unit=unit, unit_requirement=requirement)
            before = deepcopy(value)
            with self.subTest(unit=unit, requirement=requirement), self.assertRaises(DocumentError):
                codec.encode_unit_policy_rows(value)
            self.assertEqual(value, before)

    def test_tag_type_arity_case_and_literal_types_refuse(self):
        wire = codec.encode_unit_policy_rows(self.canonical())
        cases = [None, False, {}, 'none', 'null', [], [''], ['None'], ['DOCUMENTED'], ['unknown'],
                 [None], [False], [0], ['none', None], ['documented', 'minutes'], ['printed'],
                 ['printed', None], ['printed', False], ['printed', 1], ['printed', []], ['printed', {}],
                 ['printed', ''], ['printed', ' '], ['printed', 'x' * 81], ['printed', 'm', 'extra']]
        for policy in cases:
            rows = self.rows(wire)
            rows[0][6] = policy
            with self.subTest(policy=policy), self.assertRaises(DocumentError):
                codec.decode(self.replace_rows(deepcopy(wire), rows), codec.UNIT_POLICY_TRANSPORT)

    def test_malformed_rows_nonfinite_duplicates_and_nested_transport_refuse(self):
        wire = codec.encode_unit_policy_rows(self.canonical())
        row = self.rows(wire)[0]
        for rows in (None, False, {}, '[]', [None], [{}], [row[:-1]], [row + [None]], [json.dumps(row)]):
            with self.subTest(rows=repr(rows)[:40]), self.assertRaises(DocumentError):
                codec.decode(self.replace_rows(deepcopy(wire), rows), codec.UNIT_POLICY_TRANSPORT)
        for encoded in (None, [], '', 'null', '[', 'NaN', '[Infinity]', '[NaN]',
                        '[[{"a":1,"a":2}]]', '[[{"a":1,"\\u0061":2}]]',
                        '[' * 1200 + ']' * 1200, ' ' * (codec.MAX_ENCODED_BYTES + 1)):
            changed = deepcopy(wire)
            next(codec.contracts(changed))['components'] = encoded
            with self.subTest(encoded=repr(encoded)[:40]), self.assertRaises(DocumentError):
                codec.decode(changed, codec.UNIT_POLICY_TRANSPORT)

    def test_every_other_position_retains_exact_types_and_scope_cardinality(self):
        wire = codec.encode_unit_policy_rows(self.canonical())
        cases = []
        for index in (0, 1, 3, 4, 8):
            cases.extend((index, value) for value in (None, False, 1, [], {}))
        for index in (2, 10):
            cases.extend((index, value) for value in (False, 1, [], {}))
        cases.extend((7, value) for value in (None, 0, 1, 'true', [], {}))
        cases.extend((9, value) for value in (None, '[]', {}, [None], [False], [1]))
        cases.extend((5, value) for value in (None, {}, 'patient', [], ['patient'],
            ['patient', 'any'], ['patient', 'any', 'any', 'extra'], ['patient', None, 'any']))
        for index, value in cases:
            rows = self.rows(wire)
            rows[0][index] = value
            with self.subTest(index=index, value=value), self.assertRaises(DocumentError):
                codec.decode(self.replace_rows(deepcopy(wire), rows), codec.UNIT_POLICY_TRANSPORT)

    def test_source_absent_units_required_bounds_choices_and_signing_still_refuse(self):
        cases = [lambda c: c.update(question_quote='not present in the original'),
                 lambda c: c.update(unit='arbitrary_unprinted_unit', unit_requirement='printed'),
                 lambda c: c.update(required=False), lambda c: c.update(value_type='unknown'),
                 lambda c: c.update(choices=['invented']), lambda c: c.update(option_quote='invented'),
                 lambda c: c['scope'].update(subject='unknown')]
        for mutate in cases:
            value = self.canonical()
            mutate(next(codec.components(value)))
            decoded = codec.decode(codec.encode_unit_policy_rows(value), codec.UNIT_POLICY_TRANSPORT)
            self.assertEqual(decoded, value)
            with self.assertRaises(ValueError):
                word_contract.validate_result(self.inventory, decoded, 'proposer', self.live())
        for count in (0, 31):
            value = self.canonical()
            contract = next(codec.contracts(value))
            component = contract['components'][0]
            contract['components'] = [{**deepcopy(component), 'id': 'c' + str(i)} for i in range(count)]
            decoded = codec.decode(codec.encode_unit_policy_rows(value), codec.UNIT_POLICY_TRANSPORT)
            self.assertEqual(len(next(codec.contracts(decoded))['components']), count)
            with self.assertRaises(ValueError):
                word_contract.validate_result(self.inventory, decoded, 'proposer', self.live())
        value = self.canonical()
        next(codec.contracts(value)).update(rendering='protected', components=[])
        with self.assertRaises(ValueError):
            word_contract.validate_result(self.inventory,
                codec.decode(codec.encode_unit_policy_rows(value), codec.UNIT_POLICY_TRANSPORT), 'proposer', self.live())

    def test_arbitrary_literal_units_are_not_keyword_allowlisted_or_inferred(self):
        for number, unit in enumerate(('kPa', 'µmol/L', '"é"\\unit', 'null'), 1):
            doc = Document()
            doc.add_table(rows=1, cols=2).cell(0, 0).text = 'Describe measurement in ' + unit + '.'
            path = Path(self.root) / f'unit-{number}.docx'
            doc.save(path)
            inv = inventory(path, protocol=TYPED_PROTOCOL)
            value = self.canonical(inv=inv)
            definition = value['answer_definitions']['fields'][0]['definition']
            definition.update(value_type='text', answer_mode='narrative', unit=None, choices=[])
            component = next(codec.components(value))
            component.update(unit=unit, unit_requirement='printed')
            decoded = self.decode(codec.encode_unit_policy_rows(value), inv=inv)
            self.assertEqual(decoded, value)
            word_contract.validate_result(inv, decoded, 'proposer', self.live())
            component.update(unit=None, unit_requirement='documented')
            self.assertEqual(codec.decode(codec.encode_unit_policy_rows(value), codec.UNIT_POLICY_TRANSPORT), value)
        self.assertIn('mentioned only in an example or exclusion', codec.UNIT_POLICY_INSTRUCTIONS)
        # Instructions require review; the codec does not parse clinical semantics.
        self.assertIn('mark unresolved on ambiguity', codec.UNIT_POLICY_INSTRUCTIONS)

    def test_new_profile_only_scope_guidance_preserves_all_explicit_scopes(self):
        instructions = codec.instructions(codec.UNIT_POLICY_TRANSPORT)
        self.assertIn('eligibility filters on EVERY requested clause', instructions)
        self.assertIn('outer temporality:any', instructions)
        self.assertIn('Do not invent time categories', instructions)
        for transport in (codec.ARRAY_TRANSPORT, codec.QUOTE_TRANSPORT, codec.TRANSPORT, codec.COMPACT_TRANSPORT):
            self.assertNotIn('eligibility filters on EVERY requested clause', codec.instructions(transport))
        for temporality in ('past', 'current', 'future', 'any'):
            value = self.canonical()
            value['answer_definitions']['fields'][0]['definition']['scope']['temporality'] = temporality
            next(codec.components(value))['scope']['temporality'] = temporality
            self.assertEqual(codec.decode(codec.encode_unit_policy_rows(value), codec.UNIT_POLICY_TRANSPORT), value)

    def test_full_twenty_sixty_component_choices_and_nulls_roundtrip(self):
        inv, parts = compact.compound_form(self.root)
        proposal = compact.compound_canonical(inv, parts, self.live())
        for role in ('proposer', 'checker'):
            value = compact.compound_canonical(inv, parts, self.live(), role, proposal)
            decoded = self.decode(codec.encode_unit_policy_rows(value), role, proposal, inv=inv)
            self.assertEqual(decoded, value)
            self.assertEqual(sum(len(c['components']) for c in codec.contracts(decoded)), 60)
            self.assertEqual(next(codec.components(decoded))['choices'], ['open', 'closed'])
            self.assertEqual(list(codec.contracts(decoded))[1]['components'][0]['choices'], ['null', 'active'])
            word_contract.validate_result(inv, decoded, role, self.live())

    def test_schema_and_evidence_remain_bound_and_no_cross_version_fallback(self):
        value = self.canonical()
        wire = codec.encode_unit_policy_rows(value)
        canonical_schema = preparation.role_schema('proposer', TYPED_PROTOCOL, self.inventory, self.live())
        schema = codec.wire_schema(canonical_schema, codec.UNIT_POLICY_TRANSPORT)
        description = schema['properties']['answer_definitions']['properties']['fields']['items']['properties'][
            'definition']['properties']['component_contract']['properties']['components']['description']
        self.assertIn('EXACTLY eleven explicit values', description)
        self.assertIn('Reference-only RESTORED canonical object schema after decode, NOT a wire row', description)
        self.assertIn('NEVER emit separate unit or unit_requirement positions', description)
        envelope = self.evidence(wire)['request_envelope']
        self.assertEqual(expand(envelope['output_config']['format']['schema']), transform_schema(schema))
        self.assertIn(codec.UNIT_POLICY_INSTRUCTIONS, envelope['system'])
        self.assertNotIn(codec.COMPACT_INSTRUCTIONS, envelope['system'])
        self.assertEqual(preparation.role_limits(self.inventory, self.live()),
                         preparation.role_limits(self.inventory, self.live(codec.COMPACT_TRANSPORT)))
        for mutate in (lambda e: e.update(response_transport=codec.COMPACT_TRANSPORT),
                       lambda e: e.pop('request_envelope_sha256'),
                       lambda e: e['response']['content'][0].update(text='{}')):
            evidence = self.evidence(wire)
            mutate(evidence)
            with self.assertRaises(DocumentError):
                self.decode(wire, evidence=evidence)
        v5 = compact.encoded_rows(value)
        self.assertEqual(codec.decode(v5, codec.COMPACT_TRANSPORT), value)
        with self.assertRaises(DocumentError):
            codec.decode(v5, codec.UNIT_POLICY_TRANSPORT)
        with self.assertRaises(DocumentError):
            codec.decode(wire, codec.COMPACT_TRANSPORT)

    def test_large_profile_keeps_two_hundred_four_eighty_without_changing_limits(self):
        inv, _ = compact.compound_form(self.root)
        old = self.live(codec.COMPACT_TRANSPORT)
        self.assertGreater(preparation.role_limits(inv, old)[0], preparation.MAX_TYPED_OUTPUT)
        self.assertEqual(preparation.role_limits(inv, old), preparation.role_limits(inv, self.live()))
        self.assertEqual(preparation.timing_limits(inv, old), (200, 480))
        self.assertEqual(preparation.timing_limits(inv, self.live()), (200, 480))


class PreparationUnitPolicySDKTests(SimpleTestCase):
    setUp = legacy.PreparationSchemaTransportTests.setUp
    receipt = sdk_tests.PreparationStreamTransportTests.receipt
    sdk_client = sdk_tests.PreparationStreamTransportTests.sdk_client
    invoke = sdk_tests.PreparationStreamTransportTests.invoke
    trace = sdk_tests.PreparationStreamTransportTests.trace
    canonical = PreparationUnitPolicyTests.canonical

    def live(self, *, choice=False):
        return {**self.config, 'mode': 'live', 'model': preparation.REASONING_MODEL,
                'word_preparation_protocol': TYPED_PROTOCOL, CONFIG_KEY: CHOICE_CONTRACT,
                preparation.SCHEMA_TRANSPORT_KEY: codec.UNIT_POLICY_TRANSPORT,
                preparation.REASONING_KEY: preparation.REASONING_PROFILE,
                preparation.stream_transport.CONFIG_KEY: preparation.stream_transport.TRANSPORT}

    def test_real_sdk_both_roles_exact_request_raw_usage_and_replay(self):
        proposal = self.canonical()
        for role in ('proposer', 'checker'):
            canonical = self.canonical(role, proposal)
            wire = codec.encode_unit_policy_rows(canonical)
            body = sdk_tests.Body(sdk_tests.events(wire))
            result, receipt, sent, client, factory, reserve = self.invoke(body, role=role, proposal=proposal, choice=True)
            self.assertEqual(result, canonical)
            self.assertEqual(receipt.state, 'succeeded')
            self.assertEqual(len(sent), 1)
            self.assertEqual(sent[0], preparation.provider_request(self.inventory, role, self.live(), proposal))
            self.assertEqual(sent[0]['thinking'], {'type': 'between_tools'})
            self.assertEqual(sent[0]['output_config']['effort'], 'high')
            self.assertEqual(factory.call_args.kwargs['max_retries'], 0)
            self.assertEqual(reserve.call_count, 1)
            self.assertTrue(body.closed)
            self.assertTrue(client.is_closed())
            raw = json.loads(private_path(receipt.response_path).read_bytes())
            self.assertEqual(raw['response_transport'], codec.UNIT_POLICY_TRANSPORT)
            self.assertEqual(raw['request_envelope_sha256'], digest(sent[0]))
            self.assertEqual(json.loads(''.join(b.get('text', '') for b in raw['response']['content'])), wire)
            self.assertEqual(preparation.read_retained_role(receipt, 'live'), canonical)
            self.assertIsNotNone(receipt.computed_cost)

    def test_bad_policy_preserves_known_cost_and_raw_once_without_retry(self):
        wire = codec.encode_unit_policy_rows(self.canonical())
        contract = next(codec.contracts(wire))
        rows = json.loads(contract['components'])
        rows[0][6] = ['documented', 'minutes']
        contract['components'] = json.dumps(rows)
        body = sdk_tests.Body(sdk_tests.events(wire))
        result, receipt, sent, client, factory, reserve = self.invoke(body, choice=True)
        self.assertIsInstance(result, DocumentError)
        self.assertEqual(receipt.state, 'failed')
        self.assertIsNotNone(receipt.computed_cost)
        self.assertEqual(len(sent), 1)
        self.assertEqual(reserve.call_count, 1)
        self.assertEqual(factory.call_args.kwargs['max_retries'], 0)
        raw = json.loads(private_path(receipt.response_path).read_bytes())
        self.assertEqual(json.loads(''.join(b.get('text', '') for b in raw['response']['content'])), wire)
        with self.assertRaises(DocumentError):
            preparation.read_retained_role(receipt, 'live')
