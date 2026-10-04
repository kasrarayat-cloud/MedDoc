"""The actual qualified fixture wire for an unmapped completed document status.

The fixture only records deterministic relations. The shared validators still
require a valid whole relation and refuse questions, actors and source changes.
"""
from io import BytesIO
from unittest.mock import patch

from django.test import SimpleTestCase, override_settings
from docx import Document

from core import generation as g, relations_fixture
from documents import word_v3
from .test_workflow import TEST_SETTINGS


QUALIFIED = dict(TEST_SETTINGS, AI_PROVIDER='anthropic', AI_MODEL=g.SONNET_5_5,
    AI_INPUT_RATE='2', AI_OUTPUT_RATE='10', AI_CACHED_RATE='0.20',
    AI_THINKING='between_tools', AI_EFFORT='medium', AI_INFERENCE_GEO='',
    V3_FIELD_RELATIONS=False, V3_FIELD_RELATIONS_SWITCH=None,
    V3_HISTORY_CONTEXT=False, V3_GENERATION_BATCHES=False)


@override_settings(**QUALIFIED)
class DocumentStatusFixture(SimpleTestCase):
    def setUp(self):
        doc = Document()
        doc.add_heading('Clinic record', level=2)
        doc.add_paragraph('Consent form signed and dated: ______________')
        doc.add_paragraph('Plan: ______________')
        out = BytesIO(); doc.save(out)
        self.original = out.getvalue()
        self.fields = [f for f in word_v3.inspect(self.original)['fields'] if f['label'] == 'Plan']
        self.assertEqual(len(self.fields), 1)
        self.field = self.fields[0]

    def answer(self, notes):
        emitted = []
        original = relations_fixture.body
        def capture(*args, **kwargs):
            result = original(*args, **kwargs)
            emitted.append(result)
            return result
        config = g.configuration()
        self.assertEqual(config['ordinary_field_contract'], 'ordinary-field-relations-v4')
        self.assertEqual(config['request_profile']['profile'], 'claude-sonnet-5-5-ordinary-v1')
        with patch.object(relations_fixture, 'body', capture):
            answers, cost, *_ = g.generate(notes, self.fields, config)
        self.assertEqual(cost, 0)
        self.assertEqual(len(emitted), 1)
        return answers[self.field['key']], emitted[0]['fields'][0]

    def check(self, notes, supported):
        answer, wire = self.answer(notes)
        self.assertEqual((answer['status'], answer['text']),
            ('supported', 'Rest.') if supported else ('conflict', ''), (notes, answer, wire))
        pick = 'Plan: Rest.'
        stored = answer if supported else {**answer, 'status': 'supported', 'text': 'Rest.',
            'evidence': [{'quote': pick, 'start': notes.index(pick), 'end': notes.index(pick) + len(pick)}]}
        page = g.as_currently_governed(g.governance_for(notes, self.fields), notes, {**self.field, **stored})
        self.assertEqual((page['status'], page['text']),
            ('supported', 'Rest.') if supported else ('conflict', ''))
        self.assertEqual(g.governed_decisions(notes, self.fields, {self.field['key']: stored}),
            [] if supported else ['Plan'])
        return answer, wire

    def test_unchanged_r105_fixture_emits_a_complete_usable_relation(self):
        notes = 'Consent form signed and dated: Yes\nPlan: Rest.'
        answer, wire = self.check(notes, True)
        self.assertEqual(wire, {'key': self.field['key'], 'unit_ids': [2], 'conflict': False,
            'context': {'answer': [2], 'qualifier': [], 'independent': [1], 'contrary': [], 'unresolved': []}})
        self.assertEqual(answer['evidence'], [{'quote': 'Plan: Rest.', 'start': 35, 'end': 46}])
        filled = word_v3.fill(self.original, self.fields, {self.field['key']: answer['text']})
        paragraphs = [p.text for p in Document(BytesIO(filled)).paragraphs]
        self.assertEqual(paragraphs[1], 'Consent form signed and dated: ______________')
        self.assertIn('Rest.', paragraphs[2])

    def test_closed_affirmative_artifact_captions_have_the_same_wire_ownership(self):
        for status in ('Consent form signed: Yes', 'Document signed and dated: Yes'):
            with self.subTest(status=status):
                _, wire = self.check(status + '\nPlan: Rest.', True)
                self.assertEqual(wire['context']['independent'], [1])

    def test_question_conditional_negated_incomplete_and_invalidated_status_remain_unresolved(self):
        statuses = ('Consent form signed and dated: Yes?', 'Consent form signed and dated?: Yes',
            'Consent form signed and dated: No', 'Consent form signed and dated: not yet',
            'Consent form signed and dated: awaiting signature', 'Consent form will be signed and dated: Yes',
            'If consent form signed and dated: Yes', 'Consent form signed and dated: Yes, entered in error',
            'The patient signed and dated: Yes', 'Her father signed and dated: Yes',
            'The manager signed and dated: Yes', 'Consent form signed by her father: Yes')
        for status in statuses:
            with self.subTest(status=status):
                _, wire = self.check(status + '\nPlan: Rest.', False)
                self.assertEqual(wire['context']['unresolved'], [1])

    def test_document_status_never_resets_a_preceding_foreign_clinical_actor(self):
        for lead in ('The doctor underwent colonoscopy.', 'Her father has ankle pain.'):
            with self.subTest(lead=lead):
                _, wire = self.check(lead + '\nConsent form signed and dated: Yes\nPlan: Rest.', False)
                self.assertEqual(wire['context']['unresolved'], [1, 2])

    def test_other_source_governors_remain_in_the_whole_wire(self):
        for tail in ('Correction: the plan was entered in error.', 'The patient will not rest.',
                     'If the patient rests, call the doctor.'):
            with self.subTest(tail=tail):
                _, wire = self.check('Consent form signed and dated: Yes\nPlan: Rest.\n' + tail, False)
                self.assertEqual(wire['context']['independent'], [1])
                self.assertEqual(wire['context']['unresolved'], [3])

    def test_invalid_profile_fixture_cannot_manufacture_a_missing_relation(self):
        notes = 'Consent form signed and dated: Yes\nPlan: Rest.'
        with override_settings(AI_MODEL='gpt-4.1-mini', V3_FIELD_RELATIONS_SWITCH=False):
            config = g.configuration()
            self.assertIn('invalid', config['request_profile'])
            self.assertNotIn('ordinary_field_contract', config)
            answer = g.generate(notes, self.fields, config)[0][self.field['key']]
        self.assertEqual((answer['status'], answer['text']), ('conflict', ''))
        self.assertNotIn('ordinary_relation', answer)
