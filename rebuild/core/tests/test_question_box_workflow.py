"""Synthetic ordinary prose through the real worker/review/export path; provider selections are mocked."""
from decimal import Decimal
import hashlib
import io
from unittest.mock import patch
from zipfile import ZipFile

from django.test import TestCase, override_settings
from docx import Document
from lxml import etree

from core import generation as g
from core.jobs import run_one
from core.models import Encounter, Job, Usage
from core.security import private_path
from documents import word_v3 as w
from documents.test_question_boxes import box, data
from .test_word_v3_workflow import V3, V3Fixture


def selected_answers(notes, fields, selected_key):
    asked, context, original = g.v3_generation_fields(fields)
    units, _ = g.source_units(notes, asked, context_labels=context)
    selections = g.UnitSelections(fields=[g.UnitSelection(key=f['key'],
        unit_ids=[u['id'] for u in units] if f['key'] == selected_key else [], conflict=False)
        for f in asked])
    answers = g.validate_selections(notes, asked, selections, context_labels=context,
                                    evidence_contract=g.TAB_CONTRACT)
    return g.v3_adopted(original, answers, notes, asked, context)


@override_settings(**V3)
class QuestionBoxWorkflow(V3Fixture, TestCase):
    target_slot = 'b/tbl[0]/tr[0]/tc[0]'
    ordinary_notes = 'The patient reports weekly contact with the supervisor.'
    ordinary_correction = 'The patient reports weekly supervised contact.'
    ordinary_status = 'supported'
    long_correction = ('SYNTHETIC long clinician review of supervised contact. ' * 90).strip()

    def source(self):
        doc = Document()
        doc.add_paragraph('How is the relationship condition monitored?')
        box(doc)
        doc.add_paragraph('Printed instructions and signature area remain unchanged.')
        return data(doc)

    def capture(self, mode, template, version, encounter, original, exported):
        pass

    def check_mapping(self, version):
        pass

    def test_long_mapping_review_save_add_and_oversized_post_preserve_versions(self):
        doc = Document()
        labels = ['What ' + 'q' * (length - 6) + '?' for length in (334, 409)]
        for label in labels:
            doc.add_paragraph(label)
            box(doc)
        box(doc)
        original = data(doc)
        template, version = self.uploaded(data=original, name='SYNTHETIC long mapping metadata')
        before = version.fields
        page = self.client.get(f'/templates/{template.pk}/')
        self.assertContains(page, 'maxlength="2000"')
        for label in labels: self.assertContains(page, label)
        payload = {'action': 'v3_save', 'mapping_version': version.pk,
                   'kind_' + version.fields[0]['key']: 'text'}
        self.assertEqual(self.client.post(f'/templates/{template.pk}/', payload).status_code, 302)
        latest = self.latest(template)
        self.assertEqual([f['label'] for f in latest.fields], labels)
        oversized = {'action': 'v3_save', 'mapping_version': latest.pk,
                     'label_' + latest.fields[0]['key']: 'x' * 1001}
        self.assertEqual(self.client.post(f'/templates/{template.pk}/', oversized).status_code, 400)
        add = {'action': 'v3_add', 'mapping_version': latest.pk, 'kind': 'long_text',
               'slot': 'word/document.xml#b/tbl[2]/tr[0]/tc[0]', 'label': 'x' * 1001}
        self.assertEqual(self.client.post(f'/templates/{template.pk}/', add).status_code, 400)
        self.assertEqual(self.latest(template).pk, latest.pk)
        add['label'] = labels[1]
        self.assertEqual(self.client.post(f'/templates/{template.pk}/', add).status_code, 302)
        final = self.latest(template)
        self.assertEqual(final.fields[-1]['label'], labels[1])
        self.assertEqual(final.fields[-1]['origin'], 'manual')
        version.refresh_from_db()
        self.assertEqual(version.fields, before)
        self.assertEqual(private_path(template.original_path).read_bytes(), original)

    def test_missing_ordinary_long_worker_correction_finalization_and_export(self):
        original = self.source()
        template, version = self.uploaded(data=original, name='SYNTHETIC separate answer box')
        self.check_mapping(version)
        field = next(f for f in version.fields if f['locator']['slot'] == self.target_slot)
        download = self.client.get(f'/templates/{template.pk}/original/')
        self.assertEqual(b''.join(download.streaming_content), original)
        page = self.client.get(f'/templates/{template.pk}/')
        self.assertContains(page, field['label'])
        self.confirm(template, version)
        for mode in ('missing', 'ordinary', 'long'):
            with self.subTest(mode=mode):
                notes = 'Administrative visit record.' if mode == 'missing' else self.ordinary_notes
                self.assertEqual(self.new_encounter(version, notes=notes).status_code, 302)
                encounter = Encounter.objects.latest('created_at')
                answers = selected_answers(notes, version.fields, None if mode == 'missing' else field['key'])
                self.assertEqual(answers[field['key']]['status'], 'missing' if mode == 'missing' else self.ordinary_status)
                mocked = (answers, Decimal('0'), {}, {}, 'synthetic-unit-selection', '', True)
                with patch('core.jobs.generation.generate', return_value=mocked) as provider:
                    self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/generate/').status_code, 302)
                    self.assertTrue(run_one())
                provider.assert_called_once()
                encounter.refresh_from_db()
                self.assertEqual(Job.objects.get(encounter=encounter).state, 'succeeded')
                usage = Usage.objects.get(job__encounter=encounter)
                self.assertTrue(usage.simulated)
                self.assertEqual(usage.charged_cost, 0)
                self.assertEqual(encounter.answers[field['key']]['text'], notes if mode != 'missing' and self.ordinary_status == 'supported' else '')
                self.assertEqual(self.client.get(f'/encounters/{encounter.pk}/').status_code, 200)
                value = {'missing': '', 'ordinary': self.ordinary_correction,
                         'long': self.long_correction}[mode]
                payload = {'revision': encounter.revision}
                for f in version.fields:
                    payload['reviewed_' + f['key']] = 'on'
                    payload['answer_' + f['key']] = value if f['key'] == field['key'] else ''
                response = self.client.post(f'/encounters/{encounter.pk}/save/', payload)
                self.assertEqual(response.status_code, 302)
                encounter.refresh_from_db()
                if value:
                    self.assertEqual(encounter.answers[field['key']]['status'], 'edited')
                    self.assertEqual(encounter.answers[field['key']]['edited_by'], str(self.user.pk))
                response = self.client.post(f'/encounters/{encounter.pk}/finalize/',
                    {'revision': encounter.revision, 'reviewed': 'on'}, follow=True)
                self.assertEqual(response.status_code, 200)
                encounter.refresh_from_db()
                self.assertEqual(encounter.status, 'finalized')
                exported = private_path(encounter.export_path).read_bytes()
                download = self.client.get(f'/encounters/{encounter.pk}/export/')
                self.assertEqual(b''.join(download.streaming_content), exported)
                self.assertEqual(hashlib.sha256(exported).hexdigest(), encounter.export_sha256)
                self.assertEqual(private_path(template.original_path).read_bytes(), original)
                old = w.SlotIndex(w.open_package(original))
                new = w.SlotIndex(w.open_package(exported))
                target = 'word/document.xml#' + self.target_slot
                expected = value or w.text_of(old.by_id[target].el).strip()
                self.assertEqual(w.text_of(new.by_id[target].el).strip(), expected)
                for slot in old.slots:
                    if slot.part == 'word/document.xml' and slot.id != target and \
                            not slot.id.startswith(target + '/') and not target.startswith(slot.id + '/'):
                        self.assertEqual(etree.tostring(slot.el), etree.tostring(new.by_id[slot.id].el), slot.id)
                for slot in old.roots:
                    if slot.part != 'word/document.xml' or not target.startswith(slot.id + '/'):
                        self.assertEqual(etree.tostring(slot.el), etree.tostring(new.by_id[slot.id].el), slot.id)
                with ZipFile(io.BytesIO(original)) as before, ZipFile(io.BytesIO(exported)) as after:
                    self.assertEqual(before.namelist(), after.namelist())
                    for part in before.namelist():
                        if part != 'word/document.xml':
                            self.assertEqual(before.read(part), after.read(part), part)
                self.capture(mode, template, version, encounter, original, exported)


class PairedCellWorkflow(QuestionBoxWorkflow):
    target_slot = 'b/tbl[0]/tr[1]/tc[1]'
    ordinary_notes = 'The patient reports knee pain.'
    ordinary_correction = 'The patient reports improving knee pain.'

    def source(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        for row, label in enumerate(('Date of Birth', 'Symptoms:')):
            table.cell(row, 0).text = label
            table.cell(row, 0).add_paragraph()
        return data(doc)

    def check_mapping(self, version):
        self.assertEqual(len(version.fields), 2)
        self.assertTrue(all(f['locator']['target']['type'] == 'cell' for f in version.fields))


@override_settings(**V3)
class NestedCheckboxWorkflow(V3Fixture, TestCase):
    target_slot = 'b/tbl[0]/tr[0]/tc[0]/tbl[0]/tr[0]/tc[0]/p[0]/glyphs'
    question = 'Which symptoms occur? Tick all that apply.'
    wanted_options = ['Tightness', 'Nausea']

    def source(self):
        from documents.test_question_boxes import NestedCheckboxOwnership
        return data(NestedCheckboxOwnership().form())

    def capture(self, mode, original, version, encounter, exported):
        pass

    def test_nested_choice_upload_review_worker_correction_finalize_download(self):
        original = self.source()
        template, version = self.uploaded(data=original, name='SYNTHETIC nested choice ownership')
        target = next(f for f in version.fields if f['locator']['slot'] == self.target_slot)
        self.assertEqual((target['label'], target['kind']), (self.question, 'multi_choice'))
        self.assertEqual(b''.join(self.client.get(f'/templates/{template.pk}/original/').streaming_content), original)
        self.assertContains(self.client.get(f'/templates/{template.pk}/'), self.question)
        self.confirm(template, version)
        for mode in ('missing', 'ordinary'):
            notes = 'Administrative visit record.' if mode == 'missing' else 'The patient reports symptoms during a headache.'
            self.assertEqual(self.new_encounter(version, notes=notes).status_code, 302)
            encounter = Encounter.objects.latest('created_at')
            answers = selected_answers(notes, version.fields, None if mode == 'missing' else target['key'])
            self.assertEqual(answers[target['key']]['status'], 'missing' if mode == 'missing' else 'conflict')
            # Ordinary option semantics still require a clinician decision. This is
            # structural/application evidence, never automatic clinical success.
            with patch('core.jobs.generation.generate', return_value=(answers, Decimal('0'), {}, {},
                       'synthetic-unit-selection', '', True)) as provider:
                self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/generate/').status_code, 302)
                self.assertTrue(run_one())
            provider.assert_called_once(); encounter.refresh_from_db()
            self.assertEqual(Job.objects.get(encounter=encounter).state, 'succeeded')
            usage = Usage.objects.get(job__encounter=encounter)
            self.assertTrue(usage.simulated);self.assertEqual(usage.charged_cost, 0)
            self.assertContains(self.client.get(f'/encounters/{encounter.pk}/'), 'data-v3-options')
            payload = {'revision': encounter.revision}
            for field in version.fields:
                payload['answer_' + field['key']] = self.wanted_options if mode == 'ordinary' and field is target else ''
                payload['reviewed_' + field['key']] = 'on'
            self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/save/', payload).status_code, 302)
            encounter.refresh_from_db()
            if mode == 'ordinary':
                self.assertEqual(encounter.answers[target['key']]['status'], 'edited')
                self.assertEqual(encounter.answers[target['key']]['text'], '\n'.join(self.wanted_options))
            self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/finalize/',
                {'revision': encounter.revision, 'reviewed': 'on'}).status_code, 302)
            encounter.refresh_from_db();self.assertEqual(encounter.status, 'finalized')
            response = self.client.get(f'/encounters/{encounter.pk}/export/')
            self.assertEqual(response.status_code, 200)
            exported = b''.join(response.streaming_content)
            self.assertEqual(hashlib.sha256(exported).hexdigest(), encounter.export_sha256)
            self.assertEqual(private_path(template.original_path).read_bytes(), original)
            before = w.SlotIndex(w.open_package(original));after = w.SlotIndex(w.open_package(exported))
            expected = ([option in self.wanted_options for option in target['options']] if mode == 'ordinary'
                        else w.option_states(before, target))
            self.assertEqual(w.option_states(after, target), expected)
            changed_paragraphs = {box['part']+'#'+box['slot'].rsplit('/glyphs',1)[0]
                                  for box in target['locator']['target']['boxes']}
            for slot in before.slots:
                if slot.kind == 'p' and (mode == 'missing' or slot.id not in changed_paragraphs):
                    self.assertEqual(etree.tostring(slot.el), etree.tostring(after.by_id[slot.id].el))
            self.capture(mode, original, version, encounter, exported)


@override_settings(**V3)
class NestedWriteInWorkflow(V3Fixture, TestCase):
    target_slot = 'b/tbl[0]/tr[0]/tc[0]/tbl[0]/tr[0]/tc[1]/p[0]'
    ordinary_notes = 'The patient reports pressure around the forehead.'
    long_correction = ('SYNTHETIC retained pressure description with literal ☐ notation. ' * 45).strip()

    def source(self):
        from documents.test_question_boxes import NestedWriteIns
        return data(NestedWriteIns().form(heading='', subsection=False, label='Other:'))

    def capture(self, mode, original, version, encounter, exported, generated):
        pass

    def test_disjoint_write_in_upload_review_worker_finalize_download(self):
        original = self.source()
        template, version = self.uploaded(data=original, name='SYNTHETIC separate option write-in')
        field = next((f for f in version.fields if f['locator']['slot'] == self.target_slot
                      and f['locator']['target']['type'] == 'span'), None)
        self.assertIsNotNone(field, 'The independently expected printed write-in is absent')
        group = next(f for f in version.fields if f['locator']['target']['type'] == 'checkboxes'
                     and any(box['slot'] == self.target_slot + '/glyphs' for box in f['locator']['target']['boxes']))
        page = self.client.get(f'/templates/{template.pk}/')
        self.assertContains(page, field['label'])
        self.assertEqual(b''.join(self.client.get(f'/templates/{template.pk}/original/').streaming_content), original)
        # The editable mapping review round trip must preserve both original targets.
        self.assertEqual(self.client.post(f'/templates/{template.pk}/',
            {'action': 'v3_save', 'mapping_version': version.pk}).status_code, 302)
        reviewed = self.latest(template)
        self.assertEqual(reviewed.fields, version.fields)
        version = reviewed;self.confirm(template, version)
        for mode in ('missing', 'ordinary', 'long'):
            notes = 'Administrative visit record.' if mode == 'missing' else self.ordinary_notes
            self.assertEqual(self.new_encounter(version, notes=notes).status_code, 302)
            encounter = Encounter.objects.latest('created_at')
            answers = selected_answers(notes, version.fields, None if mode == 'missing' else field['key'])
            generated = answers[field['key']].copy()
            self.assertEqual(answers[group['key']]['status'], 'missing')
            with patch('core.jobs.generation.generate', return_value=(answers, Decimal('0'), {}, {},
                       'synthetic-write-in-selection', '', True)) as provider:
                self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/generate/').status_code, 302)
                self.assertTrue(run_one())
            provider.assert_called_once();encounter.refresh_from_db()
            self.assertEqual(Job.objects.get(encounter=encounter).state, 'succeeded')
            usage = Usage.objects.get(job__encounter=encounter)
            self.assertTrue(usage.simulated);self.assertEqual(usage.charged_cost, 0)
            self.assertEqual(self.client.get(f'/encounters/{encounter.pk}/').status_code, 200)
            value = '' if mode == 'missing' else self.ordinary_notes if mode == 'ordinary' else self.long_correction
            payload = {'revision': encounter.revision}
            for current in version.fields:
                payload['reviewed_' + current['key']] = 'on'
                payload['answer_' + current['key']] = (value if current['key'] == field['key'] else
                    ['Other'] if mode != 'missing' and current['key'] == group['key'] else '')
            self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/save/', payload).status_code, 302)
            encounter.refresh_from_db()
            self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/finalize/',
                {'revision': encounter.revision, 'reviewed': 'on'}).status_code, 302)
            encounter.refresh_from_db();self.assertEqual(encounter.status, 'finalized')
            response = self.client.get(f'/encounters/{encounter.pk}/export/')
            self.assertEqual(response.status_code, 200)
            exported = b''.join(response.streaming_content)
            self.assertEqual(hashlib.sha256(exported).hexdigest(), encounter.export_sha256)
            self.assertEqual(private_path(template.original_path).read_bytes(), original)
            old = w.SlotIndex(w.open_package(original));new = w.SlotIndex(w.open_package(exported))
            identity = 'word/document.xml#' + self.target_slot
            source = w.paragraph_text(old.by_id[identity].el)[0];target = field['locator']['target']
            expected = source if not value else source[:target['start']] + value + source[target['end']:]
            if value: expected = expected.replace(w.BOX, w.CHECKED, 1)
            self.assertEqual(w.paragraph_text(new.by_id[identity].el)[0], expected)
            allowed = {box['part']+'#'+box['slot'].rsplit('/glyphs', 1)[0]
                       for box in group['locator']['target']['boxes']} | {identity}
            for slot in old.slots:
                if slot.kind == 'p' and (mode == 'missing' or slot.id not in allowed):
                    self.assertEqual(etree.tostring(slot.el), etree.tostring(new.by_id[slot.id].el), slot.id)
            with ZipFile(io.BytesIO(original)) as before, ZipFile(io.BytesIO(exported)) as after:
                for name in before.namelist():
                    if mode == 'missing' or name != 'word/document.xml':
                        self.assertEqual(before.read(name), after.read(name), name)
            self.capture(mode, original, version, encounter, exported, generated)
