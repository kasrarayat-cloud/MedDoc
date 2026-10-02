"""Global rules follow the exact original and Word writes; synthetic cases only."""
import copy,hashlib,io,tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from docx import Document
from django.core.exceptions import ValidationError
from django.test import SimpleTestCase,TestCase,override_settings
from core import generation_batch as b,generation as g,v3_mapping,batch_jobs
from core.models import TemplateVersion,Encounter,Usage,Job
from documents import word_v3 as w
from .test_generation_batch import manifest,response,BatchFixture,SETTINGS
from .test_word_v3_workflow import V3Fixture,V3,part_text


def original(*,single=False,rows=103):
    doc=Document();doc.add_paragraph('Select at most one active group.')
    table=doc.add_table(rows=rows,cols=2)
    for i,label in enumerate(('Group A','Group B')):
        table.cell(i,0).text=label
        table.cell(i,1).text=('☒ Active' if i==0 else '☐ Active')+('' if single else ' ☐ Reviewed')
    for i in range(2,rows):table.cell(i,0).text=f'Note {i:03d}'
    out=io.BytesIO();doc.save(out);return out.getvalue()


def rules(fields,sha):
    # Independent fixture expectation: first two authored physical rows. Single
    # boxes retain the printed label Active rather than inventing Group labels.
    selected={label:fields[i] for i,label in enumerate(('Group A','Group B'))}
    for i,f in enumerate(selected.values()):
        assert f'/tr[{i}]/tc[1]/' in f['locator']['slot']
    rule=manifest(fields,sha)
    rule['instruction_references']=[{'id':'r','part':'word/document.xml','slot':'b/p[0]','sha256':hashlib.sha256(b'Select at most one active group.').hexdigest()}]
    rule['rules']=[{'kind':'max_active_groups','source_id':'r','maximum':1,'groups':[
        [{'key':selected[label]['key'],'values':['checked' if selected[label]['kind']=='checkbox' else 'Active']}]
        for label in ('Group A','Group B')]}]
    return rule,selected


def decisions(fields,selected,first='',second='Active',*,reviewed=True):
    answers={f['key']:{'text':'','status':'left_blank' if reviewed else 'missing','reviewed':reviewed} for f in fields}
    for label,text in [('Group A',first),('Group B',second)]:
        answers[selected[label]['key']]={'text':text,'status':'supported' if text else ('left_blank' if reviewed else 'missing'),'reviewed':reviewed}
    return answers


@override_settings(**SETTINGS)
class EffectiveRules(SimpleTestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.media=override_settings(MEDIA_ROOT=Path(self.temp.name));self.media.enable();self.addCleanup(self.media.disable)
        deny=patch('socket.socket.connect',side_effect=AssertionError('Network forbidden'));deny.start();self.addCleanup(deny.stop)

    def fixture(self,*,single=False):
        data=original(single=single);fields=w.inspect(data)['fields'];sha=hashlib.sha256(data).hexdigest()
        Path(self.temp.name,'source.docx').write_bytes(data)
        rule,selected=rules(fields,sha)
        version=SimpleNamespace(pk='synthetic',contract='v3',review_requirements=[],fields=fields,
            template=SimpleNamespace(sha256=sha,original_path='source.docx'),generation_constraints=rule)
        return data,version,selected

    def test_blank_preserves_precheck_and_complete_membership_uses_physical_export(self):
        data,version,selected=self.fixture()
        for first,second,invalid in [('', 'Active',True),('Active','Active',True),('','',False),('Reviewed','Active',False)]:
            answers=decisions(version.fields,selected,first,second)
            if invalid:
                with self.assertRaisesMessage(ValidationError,'maximum selection'):b.check_answers(version,answers)
            else:b.check_answers(version,answers)
            physical=part_text(w.fill(data,version.fields,v3_mapping.fill_answers(version.fields,answers)))
            self.assertEqual(physical.count('☒ Active'),2 if invalid else 1)
        # Metadata cannot erase a real printed check or invent one.
        selected['Group A'].pop('printed_ticked',None)
        version.generation_constraints['fields_sha256']=b.digest(version.fields)
        with self.assertRaisesMessage(ValidationError,'maximum selection'):
            b.check_answers(version,decisions(version.fields,selected))
        self.assertEqual(Path(self.temp.name,'source.docx').read_bytes(),data)

    def test_multichoice_active_membership_is_not_whole_newline_equality(self):
        data,version,selected=self.fixture()
        for f in selected.values():f['kind']='multi_choice'
        version.generation_constraints['fields_sha256']=b.digest(version.fields)
        answers=decisions(version.fields,selected,'Active\nReviewed','Active\nReviewed')
        with self.assertRaisesMessage(ValidationError,'maximum selection'):b.check_answers(version,answers)
        self.assertEqual(part_text(w.fill(data,version.fields,v3_mapping.fill_answers(version.fields,answers))).count('☒ Active'),2)

    def test_reviewed_blank_checkbox_clears_but_unreviewed_preserves(self):
        data,version,selected=self.fixture(single=True)
        self.assertEqual({f['kind'] for f in selected.values()},{'checkbox'})
        answers=decisions(version.fields,selected,'','checked')
        b.check_answers(version,answers)
        self.assertEqual(part_text(w.fill(data,version.fields,v3_mapping.fill_answers(version.fields,answers))).count('☒ Active'),1)
        answers[selected['Group A']['key']].update(status='missing',reviewed=False)
        with self.assertRaisesMessage(ValidationError,'maximum selection'):b.check_answers(version,answers)

    def test_unknown_constrained_state_locator_and_source_fail_closed(self):
        data,version,selected=self.fixture()
        answers=decisions(version.fields,selected,'','')
        # This malformed/unknown control is not counted as an unchecked box.
        selected['Group A']['locator']['target']['boxes'][0]['glyph']=99
        version.generation_constraints['fields_sha256']=b.digest(version.fields)
        with self.assertRaisesMessage(ValidationError,'unreadable'):b.check_answers(version,answers)
        selected['Group A']['locator']['target']['boxes'][0]['glyph']=0
        selected['Group A']['locator']['fingerprint']='f'*64
        version.generation_constraints['fields_sha256']=b.digest(version.fields)
        with self.assertRaisesMessage(ValidationError,'could not be verified'):b.check_answers(version,answers)
        data,version,selected=self.fixture();answers=decisions(version.fields,selected,'','')
        version.template.sha256='b'*64;version.generation_constraints['source_sha256']='b'*64
        with self.assertRaisesMessage(ValidationError,'original form changed'):b.check_answers(version,answers)

    def test_unread_annotation_holds_only_omitted_constrained_control(self):
        _,version,selected=self.fixture();selected['Group A']['printed_unread']=[0]
        version.generation_constraints['fields_sha256']=b.digest(version.fields)
        with self.assertRaisesMessage(ValidationError,'unreadable'):
            b.check_answers(version,decisions(version.fields,selected,'',''))
        b.check_answers(version,decisions(version.fields,selected,'Reviewed','Active'))
        version.generation_constraints['rules'][0]['groups']=version.generation_constraints['rules'][0]['groups'][1:]
        b.check_answers(version,decisions(version.fields,selected,'','Active'))

    def test_predicate_options_and_preserved_unique_priority_are_checked(self):
        _,version,selected=self.fixture()
        version.generation_constraints['rules'][0]['groups'][0][0]['values']=['not an original option']
        with self.assertRaisesMessage(ValidationError,'exact option'):b.constraints(version)
        version.generation_constraints['rules']=[{'kind':'unique_values','source_id':'r','keys':[f['key'] for f in selected.values()],'allowed_values':['Active','Reviewed']}]
        with self.assertRaisesMessage(ValidationError,'priority'):
            b.check_answers(version,decisions(version.fields,selected,'','Active'))
        b.check_answers(version,decisions(version.fields,selected,'Reviewed','Active'))


    def test_merged_header_copies_must_agree_on_preserved_state(self):
        from documents.test_word_v3 import _sdt
        from docx.oxml.ns import qn
        doc=Document();section=doc.sections[0];section.different_first_page_header_footer=True
        for header,value in ((section.header,'Active'),(section.first_page_header,'Reviewed')):
            p=header.paragraphs[0];p.text='Activity '
            control=_sdt('dropdown',options=['Active','Reviewed'],placeholder=value)
            pr=control.find(qn('w:sdtPr'));pr.remove(pr.find(qn('w:showingPlcHdr')))
            p._p.append(control)
        stream=io.BytesIO();doc.save(stream);data=stream.getvalue()
        fields=w.inspect(data)['fields'];control=next(f for f in fields if f['kind']=='choice')
        self.assertEqual(len(control['locator']['target'].get('copies',[])),1)
        sha=hashlib.sha256(data).hexdigest();Path(self.temp.name,'source.docx').write_bytes(data)
        m=manifest(fields,sha);m['instruction_references']=[{'id':'r','part':'word/document.xml','slot':'b/p[0]','sha256':'2'*64}]
        m['rules']=[{'kind':'max_active_groups','source_id':'r','maximum':0,'groups':[[{'key':control['key'],'values':['Active']}]]}]
        version=SimpleNamespace(fields=fields,template=SimpleNamespace(sha256=sha,original_path='source.docx'),generation_constraints=m)
        answers={f['key']:{'text':'','reviewed':True} for f in fields}
        with self.assertRaisesMessage(ValidationError,'unreadable'):b.check_answers(version,answers)
        answers[control['key']]['text']='Reviewed';b.check_answers(version,answers)
        answers[control['key']]['text']='Active'
        with self.assertRaisesMessage(ValidationError,'maximum selection'):b.check_answers(version,answers)


@override_settings(**SETTINGS)
class EffectiveWorkflow(BatchFixture,TestCase):
    def setup_original(self):
        data=original();template,first=self.uploaded(data=data,name='Synthetic preserved option states')
        rule,selected=rules(first.fields,template.sha256)
        version=TemplateVersion.objects.create(template=template,number=2,fields=first.fields,contract='v3',generation_constraints=rule,created_by=self.user)
        self.confirm(template,version);self.new_encounter(version,'Group B: Active.')
        self.encounter=Encounter.objects.latest('created_at');return data,version,selected

    def test_worker_checks_original_effective_state_before_atomic_adoption(self):
        _,version,selected=self.setup_original();batch=self.admit();calls,call=self.fake(batch)
        with patch('anthropic.Anthropic') as provider:
            provider.return_value.messages.create.side_effect=call
            for _ in range(3):batch_jobs.run_one()
            self.assertEqual(len(calls),2)
        batch.refresh_from_db();self.encounter.refresh_from_db()
        self.assertEqual(batch.state,'failed');self.assertEqual(self.encounter.answers,{})
        self.assertFalse(batch_jobs.pending(self.encounter))
        self.assertEqual(Usage.objects.filter(job__generation_part__batch=batch).count(),2)

    def test_page_save_forced_finalize_hold_then_explicit_override_exports_one(self):
        data,version,selected=self.setup_original();e=self.encounter
        bad=decisions(version.fields,selected)
        def post(answers):
            out={'revision':e.revision}
            for key,a in answers.items():
                out['answer_'+key]=a['text'];out['reviewed_'+key]='on'
                if not a['text']:out['blank_'+key]='on'
            return out
        with patch('anthropic.Anthropic',side_effect=AssertionError('No provider allowed')):
            self.assertEqual(self.client.post(f'/encounters/{e.pk}/save/',post(bad)).status_code,409)
            # Old/forced supported state cannot bypass page or finalization.
            for a in bad.values():a['edited_by']=str(self.user.pk)
            e.answers=bad;e.status='review';e.save(update_fields=['answers','status'])
            page=self.client.get(f'/encounters/{e.pk}/');self.assertIn('maximum selection',page.context['batch_constraint_error'])
            self.client.post(f'/encounters/{e.pk}/finalize/',{'revision':e.revision,'reviewed':'on'});e.refresh_from_db()
            self.assertNotEqual(e.status,'finalized');self.assertFalse(e.export_path)
            good=decisions(version.fields,selected,'Reviewed','Active')
            self.assertEqual(self.client.post(f'/encounters/{e.pk}/save/',post(good)).status_code,302);e.refresh_from_db()
            self.client.post(f'/encounters/{e.pk}/finalize/',{'revision':e.revision,'reviewed':'on'});e.refresh_from_db()
            self.assertEqual(e.status,'finalized')
            download=self.client.get(f'/encounters/{e.pk}/export/');self.assertEqual(download.status_code,200)
            out=b''.join(download.streaming_content);self.assertEqual(part_text(out).count('☒ Active'),1)
            self.assertEqual(hashlib.sha256(out).hexdigest(),e.export_sha256)
        self.assertEqual(Job.objects.count(),0);self.assertEqual(Usage.objects.count(),0)
