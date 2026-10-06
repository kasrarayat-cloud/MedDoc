"""Retain the incomplete eight-row source; qualify a separately named signing/export control."""
import hashlib
import importlib.util
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from docx import Document

from core.billing import welcome_credit
from core.jobs import run_one
from core.models import Department, Encounter, Job, Membership, Organization, Patient, Reservation, Template, Usage, User
from core.security import private_path
from core.tests.test_workflow import TEST_SETTINGS

PREPARER = Path(__file__).resolve().parents[2] / 'prepare-word-authority-browser.py'
spec = importlib.util.spec_from_file_location('word_authority_preparer', PREPARER)
preparer = importlib.util.module_from_spec(spec)
spec.loader.exec_module(preparer)


@override_settings(**TEST_SETTINGS)
class WordAuthorityFixture(TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.override = override_settings(MEDIA_ROOT=Path(self.directory.name) / 'private')
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.inputs = Path(self.directory.name) / 'inputs'
        self.original, self.control = preparer.prepare(self.inputs)
        self.user = User.objects.create_user(username='word-authority', email='doctor@synthetic.invalid', verified_at=timezone.now())
        dept = Department.objects.create(organization=Organization.objects.create(name='Synthetic authority'), name='X-ray')
        Membership.objects.create(user=self.user, department=dept, role='clinician')
        self.patient = Patient.objects.create(department=dept, display_name='Synthetic patient', created_by=self.user)
        welcome_credit(self.user)
        self.client.force_login(self.user)

    def uploaded_encounter(self, case):
        data = (self.inputs / case['form']).read_bytes()
        response = self.client.post('/templates/upload/', {'name': case['purpose'], 'department': self.patient.department_id,
            'synthetic_only': 'on', 'file': SimpleUploadedFile(case['form'], data)})
        self.assertEqual(response.status_code, 302)
        template = Template.objects.get(name=case['purpose'])
        version = template.versions.get(number=1)
        self.assertEqual([field['label'] for field in version.fields], case['rows'])
        self.assertEqual(template.sha256, case['original_sha256'])
        response = self.client.post('/encounters/new/', {'patient': self.patient.pk, 'template_version': version.pk,
            'notes': case['notes'], 'service_date': '2026-09-30', 'synthetic_only': 'on'})
        self.assertEqual(response.status_code, 302)
        return template, version, Encounter.objects.get(template_version=version), data

    def retain(self, name, template, version, encounter, original, extra=None):
        if not os.environ.get('QUILLMD_WORD_ARTIFACTS'):
            return
        path = Path(os.environ['QUILLMD_WORD_ARTIFACTS']) / 'authority-http' / name
        path.mkdir(parents=True, exist_ok=True, mode=0o700)
        (path / 'original.docx').write_bytes(original)
        evidence = {'original_sha256': template.sha256, 'mapping': version.fields,
            'import_manifest': version.import_manifest, 'engine': version.contract or 'legacy',
            'encounter': str(encounter.pk), 'status': encounter.status, 'notes': encounter.notes, 'answers': encounter.answers,
            'jobs': list(Job.objects.filter(encounter=encounter).values('state', 'kind', 'config')),
            'reservations': Reservation.objects.filter(job__encounter=encounter).count(),
            'usage': list(Usage.objects.filter(job__encounter=encounter).values('simulated', 'actual_cost', 'charged_cost')),
            'provider_requests': 0, **(extra or {})}
        (path / 'evidence.json').write_text(json.dumps(evidence, default=str, indent=2))
        if encounter.export_path:
            (path / 'export.docx').write_bytes(private_path(encounter.export_path).read_bytes())

    def test_preparer_retains_original_denominator_and_never_overwrites_evidence(self):
        self.assertEqual(len(self.original['rows']), 8)
        self.assertIn('Consent form signed', self.original['rows'])
        self.assertEqual(self.control['rows'], [row for row in self.original['rows'] if row != 'Consent form signed'])
        self.assertEqual(self.control['signing'], self.original['signing'])
        self.assertIn('Initial assessment', self.control['ordinary'])
        before = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in self.inputs.iterdir()}
        with self.assertRaises(FileExistsError):
            preparer.prepare(self.inputs)
        self.assertEqual(before, {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in self.inputs.iterdir()})

    def test_retained_eight_row_human_consent_refuses_without_job_reservation_or_charge(self):
        template, version, encounter, original = self.uploaded_encounter(self.original)
        self.assertTrue(version.import_manifest['coverage']['block_generation'])
        refusals = []
        for _ in range(2):
            response = self.client.post(f'/encounters/{encounter.pk}/generate/', follow=True)
            self.assertContains(response, 'This import has unmapped answer areas. No generation or charge is allowed')
            refusals.append(response.status_code)
        encounter.refresh_from_db()
        self.assertEqual(encounter.status, 'draft')
        self.assertEqual(encounter.answers, {})
        self.assertEqual(Job.objects.filter(encounter=encounter).count(), 0)
        self.assertEqual(Reservation.objects.filter(job__encounter=encounter).count(), 0)
        self.assertEqual(Usage.objects.filter(job__encounter=encounter).count(), 0)
        self.assertFalse(encounter.export_path)
        self.assertEqual(private_path(template.original_path).read_bytes(), original)
        self.retain('retained-eight-row-refusal', template, version, encounter, original,
            {'repeat_response_statuses': refusals, 'supported_coverage': False, 'export_assertions_evaluated': False})

    def test_separate_seven_row_control_review_finalize_export_preserves_source_and_signing_cells(self):
        template, version, encounter, original = self.uploaded_encounter(self.control)
        self.assertFalse(version.import_manifest['coverage']['block_generation'])
        self.client.post(f'/encounters/{encounter.pk}/generate/')
        self.assertTrue(run_one())
        encounter.refresh_from_db()
        self.assertEqual(Job.objects.get(encounter=encounter).state, 'succeeded')
        keys = {field['label']: field['key'] for field in version.fields}
        for label, value in self.control['ordinary'].items():
            self.assertEqual(encounter.answers[keys[label]]['status'], 'supported')
            self.assertEqual(encounter.answers[keys[label]]['text'], value)
            self.assertTrue(encounter.answers[keys[label]]['evidence'])
        for label in self.control['signing']:
            self.assertFalse(encounter.answers[keys[label]]['text'])
        payload = {'revision': encounter.revision}
        for key, answer in encounter.answers.items():
            payload['answer_' + key] = answer['text']
            payload['reviewed_' + key] = 'on'
        self.client.post(f'/encounters/{encounter.pk}/save/', payload)
        encounter.refresh_from_db()
        self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'})
        encounter.refresh_from_db()
        self.assertEqual(encounter.status, 'finalized')
        response = self.client.get(f'/encounters/{encounter.pk}/export/')
        self.assertEqual(response.status_code, 200)
        exported = b''.join(response.streaming_content)
        document = Document(BytesIO(exported))
        before = Document(BytesIO(original))
        rows = {row.cells[0].text: row.cells[1] for row in document.tables[0].rows}
        original_rows = {row.cells[0].text: row.cells[1] for row in before.tables[0].rows}
        for label, value in self.control['ordinary'].items():
            self.assertEqual(rows[label].text, value)
        for label, attack in self.control['signing'].items():
            self.assertEqual(rows[label]._tc.xml, original_rows[label]._tc.xml)
            self.assertNotIn(attack, '\n'.join(row.cells[1].text for row in document.tables[0].rows))
        self.assertEqual([p.text for p in document.paragraphs], [p.text for p in before.paragraphs])
        self.assertEqual(private_path(template.original_path).read_bytes(), original)
        usage = Usage.objects.get(job__encounter=encounter)
        self.assertTrue(usage.simulated)
        self.assertEqual(usage.actual_cost, 0)
        self.assertEqual(usage.charged_cost, 0)
        self.retain('separate-seven-row-control', template, version, encounter, original,
            {'ordinary_supported': 5, 'signing_cells_unchanged': 2, 'export_sha256': hashlib.sha256(exported).hexdigest()})
        other = User.objects.create_user(username='other-authority', verified_at=timezone.now())
        other_dept = Department.objects.create(organization=self.patient.department.organization, name='Ultrasound')
        Membership.objects.create(user=other, department=other_dept, role='clinician')
        self.client.force_login(other)
        for path in (f'/templates/{template.pk}/original/', f'/api/encounters/{encounter.pk}/', f'/encounters/{encounter.pk}/export/'):
            self.assertEqual(self.client.get(path).status_code, 404)
