"""Source-only preparation security, cost and approval binding tests; no database."""
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
import tempfile
from types import SimpleNamespace
from unittest.mock import patch,Mock
from django.test import SimpleTestCase,override_settings
from django.core.exceptions import ValidationError
from docx import Document
from documents.preparation import inventory,digest,validate_role
from documents.sdk_service import WM_IOV_SHA256
from core import preparation as prep
from core.template_source import source_for_version,approval_binding

class PreparationPureTests(SimpleTestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.settings=override_settings(MEDIA_ROOT=self.root,WORD_PREPARATION_ENABLED=True)
        self.settings.enable();self.addCleanup(self.settings.disable)
        d=Document();d.add_paragraph('Assessment: ____');d.add_paragraph('Patient signature: ____')
        d.save(self.root/'original.docx');self.inv=inventory(self.root/'original.docx')

    def config(self,**kw):
        return {'mode':'fixture','provider':'anthropic','model':'configured-label','input_rate':'2',
                'output_rate':'10','cached_rate':'0.2',**kw}

    def packet(self):
        proposal,_=prep.provider_role(self.inv,'proposer',self.config(),1)
        checker,_=prep.provider_role(self.inv,'checker',self.config(),1,proposal)
        p=SimpleNamespace(template_id='template1',template=SimpleNamespace(sha256=self.inv['source_sha256']),
            inventory=deepcopy(self.inv),proposal=proposal,checker=checker,corrections={},revision=0,
            candidate_sha256=self.inv['source_sha256'],candidate_info={'fields':[{'key':'x','label':'Assessment',
            'kind':'text','locator':{'paragraph':1}}],'document_backend':{},'review_requirements':[]},changes=[])
        return p

    def test_fixture_two_roles_bind_all_anchors_and_complete_original_stories(self):
        p=self.packet();validate_role(self.inv,p.proposal,'proposer');validate_role(self.inv,p.checker,'checker')
        self.assertEqual(len(p.proposal['anchors']),2)
        self.assertEqual(len(p.checker['stories']),len(self.inv['stories']))

    def test_ordinary_budget_quote_cannot_overrun_point_ten(self):
        config=self.config(mode='live');payload=prep.role_payload(self.inv,'proposer')
        self.assertEqual(prep.role_limits(self.inv),(2500,Decimal('0.10')))
        self.assertGreater(prep.maximum_cost({'text':'X'*100000},config),Decimal('0.10'))
        self.assertGreater(prep.maximum_cost(payload,config),0)

    def test_name_or_anchor_count_cannot_obtain_wm_exception(self):
        for alteration in ['name','anchors','source','backend']:
            value=deepcopy(self.inv)
            if alteration=='name':value['name']='WM IOV Template'
            if alteration=='anchors':value['anchors']*=45
            if alteration=='source':value['source_sha256']=WM_IOV_SHA256
            if alteration=='backend':value['document_backend']={'backend':'wm-iov-foss-v2','source_sha256':WM_IOV_SHA256}
            value['inventory_sha256']=digest({k:v for k,v in value.items() if k!='inventory_sha256'})
            self.assertEqual(prep.role_limits(value),(2500,Decimal('0.10')))

    def test_wm_cost_classification_requires_every_pinned_shape_property(self):
        # Classifier fixture only: genuine WM bytes are checked separately by
        # the actual networkless-worker qualification, not asserted here.
        value=deepcopy(self.inv);value['source_sha256']=WM_IOV_SHA256
        value['anchors']=[{**value['anchors'][0],'kind':'existing'} for _ in range(90)]
        value['fields']=[{} for _ in range(90)];value['stories']=[value['stories'][0] for _ in range(16)]
        value['document_backend']={'backend':'wm-iov-foss-v2','source_sha256':WM_IOV_SHA256}
        value['inventory_sha256']=digest({k:v for k,v in value.items() if k!='inventory_sha256'})
        self.assertEqual(prep.role_limits(value),(6000,Decimal('0.16')))
        for changed in ['source','anchors','fields','stories','backend','backend_source','kind']:
            bad=deepcopy(value)
            if changed=='source':bad['source_sha256']='0'*64
            elif changed in {'anchors','fields','stories'}:bad[changed].pop()
            elif changed=='backend':bad['document_backend']['backend']='other'
            elif changed=='backend_source':bad['document_backend']['source_sha256']='0'*64
            elif changed=='kind':bad['anchors'][0]['kind']='literal_blank'
            bad['inventory_sha256']=digest({k:v for k,v in bad.items() if k!='inventory_sha256'})
            with self.subTest(changed=changed):self.assertEqual(prep.role_limits(bad),(2500,Decimal('0.10')))

    def test_eighth_attempt_is_wm_only_and_ninth_never_creates(self):
        for attempts,sha,allowed in [(7,WM_IOV_SHA256,True),(7,self.inv['source_sha256'],False),(8,WM_IOV_SHA256,False)]:
            manager=Mock();manager.exclude.return_value=manager;manager.count.return_value=attempts
            manager.values_list.return_value.distinct.return_value.count.return_value=1
            template=SimpleNamespace(sha256=sha)
            with self.subTest(attempts=attempts,sha=sha),override_settings(WORD_PREPARATION_MAX_RUNS=8),\
                    patch.object(prep.TemplatePreparation,'objects',manager),patch.object(prep,'configuration',return_value=self.config()):
                if allowed:
                    self.assertIs(prep.create_preparation(template,'doctor'),manager.create.return_value)
                    manager.create.assert_called_once()
                else:
                    with self.assertRaises(ValidationError):prep.create_preparation(template,'doctor')
                    manager.create.assert_not_called()

    def test_final_two_sends_require_wm_and_fifteenth_has_no_effect(self):
        import time
        from documents.service import DocumentError
        for sends,sha,limits in [(12,self.inv['source_sha256'],(2500,Decimal('0.10'))),
                                  (13,WM_IOV_SHA256,(2500,Decimal('0.10'))),
                                  (14,WM_IOV_SHA256,(6000,Decimal('0.16')))]:
            manager=Mock();manager.exclude.return_value=manager;manager.count.return_value=sends
            p=self.packet();p.template.sha256=sha;p.config=self.config()
            with self.subTest(sends=sends,sha=sha),override_settings(WORD_PREPARATION_MAX_SENDS=14),\
                    patch.object(prep.PreparationReceipt,'objects',manager),patch.object(prep,'role_limits',return_value=limits),\
                    patch.object(prep,'provider_role') as provider:
                with self.assertRaises(DocumentError) as refused:prep.call_role(p,'proposer',time.monotonic()+10)
                self.assertEqual(refused.exception.code,'budget_limit');manager.create.assert_not_called();provider.assert_not_called()

    def test_approval_seal_catches_proposal_checker_correction_placement_revision_and_candidate_drift(self):
        p=self.packet();sealed=approval_binding(p)
        for name in ['proposal','checker','corrections','revision','candidate_sha256','candidate_info','changes']:
            changed=deepcopy(p)
            if name=='revision':setattr(changed,name,1)
            elif name=='candidate_sha256':setattr(changed,name,'0'*64)
            elif name=='changes':setattr(changed,name,[{'anchor_id':'a1'}])
            else:getattr(changed,name)['tampered']=True
            self.assertNotEqual(approval_binding(changed),sealed)

    def version(self):
        p=self.packet();p.template.original_path='original.docx';p.template.format='docx'
        p.state='approved';p.approved_by_id='doctor1';p.approved_at='now';p.candidate_path='candidate.docx'
        (self.root/p.candidate_path).write_bytes((self.root/'original.docx').read_bytes())
        p.approval_sha256=approval_binding(p)
        return SimpleNamespace(template=p.template,template_id='template1',preparation=p,preparation_id='prep1',
            prepared_path=p.candidate_path,prepared_sha256=p.candidate_sha256,fields=p.candidate_info['fields'],
            document_backend={},review_requirements=[])

    def test_approved_source_is_distinct_and_both_original_and_candidate_are_checked(self):
        v=self.version();self.assertEqual(source_for_version(v),self.root/'candidate.docx')
        (self.root/'original.docx').write_bytes(b'changed')
        with self.assertRaises((ValidationError,ValueError)):source_for_version(v)

    def test_cross_template_unapproved_stale_and_altered_candidate_refuse(self):
        for kind in ['other','unapproved','stale','bytes','placement']:
            v=self.version()
            if kind=='other':v.template_id='other'
            if kind=='unapproved':v.preparation.state='review'
            if kind=='stale':v.preparation.revision=1
            if kind=='bytes':(self.root/'candidate.docx').write_bytes(b'bad')
            if kind=='placement':v.fields=[{**v.fields[0],'locator':{'paragraph':999}}]
            with self.subTest(kind=kind),self.assertRaises((ValidationError,ValueError)):source_for_version(v)

    def test_legacy_version_uses_original_and_cannot_claim_unsigned_prepared_source(self):
        p=self.packet();template=SimpleNamespace(original_path='original.docx',format='docx',sha256=self.inv['source_sha256'])
        v=SimpleNamespace(template=template,preparation_id=None,prepared_path='',prepared_sha256='')
        self.assertEqual(source_for_version(v),self.root/'original.docx')
        v.prepared_path='candidate.docx'
        with self.assertRaises(ValidationError):source_for_version(v)

    def test_fixture_returns_actual_fixture_model_separately_from_configured_label(self):
        _,evidence=prep.provider_role(self.inv,'proposer',self.config(),1)
        self.assertTrue(evidence['simulated']);self.assertEqual(evidence['cost'],'0')
        self.assertNotEqual(evidence['model'],self.config()['model'])

    def test_live_transport_zero_retries_and_actual_response_identity_are_retained(self):
        response=SimpleNamespace(model='actual-returned-model',id='response1',_request_id='request1',
            usage=SimpleNamespace(input_tokens=100,output_tokens=20,cache_read_input_tokens=4,cache_creation_input_tokens=0),
            content=[SimpleNamespace(type='text',text='{}')],model_dump=lambda **kw:{'model':'actual-returned-model','usage':{'input_tokens':100}})
        client=Mock();client.messages.create.return_value=response
        with patch('anthropic.Anthropic',return_value=client) as constructor:
            raw,evidence=prep.provider_role(self.inv,'proposer',self.config(mode='live'),7)
        constructor.assert_called_once_with(timeout=7,max_retries=0)
        self.assertEqual(client.messages.create.call_count,1)
        self.assertEqual(evidence['model'],'actual-returned-model');self.assertEqual(evidence['request_id'],'request1')
        self.assertEqual(evidence['resource_id'],'response1');self.assertFalse(evidence['simulated'])
        self.assertEqual(Decimal(evidence['cost']),Decimal('0.0004008'))

    def test_official_transport_schema_transform_preserves_local_exact_count_guards(self):
        from anthropic import transform_schema
        for role in ['proposer','checker']:
            original=prep.role_schema(role);before=deepcopy(original);wire=transform_schema(original)
            self.assertEqual(original,before)
            for key in ['anchors','stories']:
                self.assertIn('maxItems',original['properties'][key]);self.assertNotIn('maxItems',wire['properties'][key])
                self.assertEqual(wire['properties'][key]['items']['required'],original['properties'][key]['items']['required'])
                self.assertFalse(wire['properties'][key]['items']['additionalProperties'])
            result,_=prep.provider_role(self.inv,role,self.config(),1)
            result['anchors'].pop()
            with self.assertRaises(ValueError):validate_role(self.inv,result,role)

    def test_failure_diagnostic_is_closed_and_does_not_export_exception_text(self):
        import anthropic,httpx
        response=httpx.Response(400,request=httpx.Request('POST','https://offline.invalid'),headers={'request-id':'req_local_1'})
        error=anthropic.BadRequestError('private-key-and-document-text',response=response,body={'error':{'type':'invalid_request_error','message':'private-note'},'secret':'private-key'})
        evidence=prep.safe_provider_failure(error,'proposer','provider_role',12.5,False,False)
        self.assertEqual(evidence['error_class'],'BadRequestError');self.assertEqual(evidence['http_status'],400)
        self.assertEqual(evidence['provider_request_id'],'req_local_1');self.assertEqual(evidence['provider_error_type'],'invalid_request_error')
        self.assertNotIn('private',str(evidence));self.assertFalse(evidence['automatic_retry'])
        class HostileError(Exception):
            status_code=True;request_id='private\nheader';body={'error':{'type':'private-code'}}
        bad=prep.safe_provider_failure(HostileError('private'),'anything','anything',0,False,False)
        self.assertIsNone(bad['http_status']);self.assertEqual(bad['provider_request_id'],'')
        self.assertEqual(bad['error_class'],'other_exception');self.assertEqual(bad['provider_error_type'],'unestablished')

    def test_role_failure_retains_separate_evidence_without_model_usage_or_retry_claim(self):
        import anthropic,httpx,time,json
        response=httpx.Response(400,request=httpx.Request('POST','https://offline.invalid'),headers={'request-id':'req_local_2'})
        error=anthropic.BadRequestError('private-message',response=response,body={'error':{'type':'invalid_request_error'}})
        p=self.packet();p.config=self.config(mode='live')
        for after_response in [False,True]:
            receipt=SimpleNamespace(response_path='',response_sha256='',returned_model='',provider_request_id='',provider_resource_id='',
                computed_cost=None,usage={},failure={},failure_path='',failure_sha256='',save=Mock())
            manager=Mock();manager.exclude.return_value=manager;manager.count.return_value=0;manager.values_list.return_value=[];manager.create.return_value=receipt
            def fail(*args, **kwargs):
                if after_response:args[-1]({'model':'retained-real-response-model','request_id':'response_req','resource_id':'response_id','response':{'raw':'retained'},'cost':None})
                raise error
            with patch.object(prep.PreparationReceipt,'objects',manager),patch('core.preparation.provider_role',side_effect=fail) as provider:
                with self.assertRaises(anthropic.BadRequestError) as caught:prep.call_role(p,'proposer',time.monotonic()+30)
            self.assertIs(caught.exception,error);self.assertEqual(provider.call_count,1)
            self.assertEqual(receipt.state,'failed' if after_response else 'unknown')
            self.assertIsNone(receipt.computed_cost);self.assertEqual(receipt.usage,{})
            self.assertEqual(receipt.returned_model,'retained-real-response-model' if after_response else '')
            saved=json.loads((self.root/receipt.failure_path).read_bytes());self.assertEqual(saved,receipt.failure)
            self.assertEqual(saved['http_status'],400);self.assertNotIn('private-message',str(saved))
            if after_response:self.assertNotEqual(receipt.failure_path,receipt.response_path)
            else:self.assertEqual(receipt.response_path,'')

    def test_missing_usage_retains_raw_response_and_unknown_cost(self):
        response=SimpleNamespace(model='actual-model',id='response1',content=[],usage=None,
                                  model_dump=lambda **kw:{'model':'actual-model','usage':None})
        client=Mock();client.messages.create.return_value=response
        with patch('anthropic.Anthropic',return_value=client):
            _,evidence=prep.provider_role(self.inv,'proposer',self.config(mode='live'),1)
        self.assertIsNone(evidence['cost']);self.assertEqual(evidence['response']['usage'],None)

    def test_raw_response_is_observed_before_invalid_usage_or_content_properties(self):
        class Response:
            model='genuine-model';id='response1';_request_id='request1'
            def model_dump(self,**kw):return {'model':self.model,'usage':{'broken':True},'content':[{'type':'text','text':'bad'}]}
            @property
            def usage(self):raise ValueError('invalid usage property')
        client=Mock();client.messages.create.return_value=Response();seen=[]
        with patch('anthropic.Anthropic',return_value=client),self.assertRaises(ValueError):
            prep.provider_role(self.inv,'proposer',self.config(mode='live'),1,on_response=seen.append)
        self.assertEqual(len(seen),1);self.assertEqual(seen[0]['model'],'genuine-model')
        self.assertEqual(seen[0]['request_id'],'request1');self.assertIsNone(seen[0]['cost'])

    def test_candidate_keeps_exact_layout_reminder_and_blocks_unresolved_warning(self):
        from documents import inspect_template
        advisory="DOCX answers expand and can change pagination. Review the exported document's layout."
        for extra in [None,'An unmapped clinical answer remains.',advisory+' Changed']:
            p=self.packet();p.template.original_path='original.docx';p.candidate_history=[];p.save=Mock()
            def inspected(path):
                info=inspect_template(path)
                info['warnings']=[advisory]+([extra] if extra else [])
                return info
            with patch('core.preparation.inspect_template',side_effect=inspected):
                if extra:
                    with self.assertRaises(prep.DocumentError):prep.build_candidate(p,{})
                else:
                    prep.build_candidate(p,{})
                    self.assertEqual(p.candidate_info['warnings'],[advisory])
                    self.assertTrue((self.root/p.candidate_path).is_file())
        self.assertNotIn('An unmapped clinical answer remains.',prep.LAYOUT_ADVISORIES)

    def test_manual_row_warning_is_admitted_only_for_exact_candidate_row_proof(self):
        from documents.engine import UNMAPPED_NOTES
        from documents import inspect_template
        d=Document();d.add_paragraph('Name: {{Name}}');table=d.add_table(rows=1,cols=3)
        table.cell(0,0).text='Assessment';d.save(self.root/'original.docx')
        self.inv=inventory(self.root/'original.docx')
        choices={a['id']:a['choices'][0]['id'] for a in self.inv['anchors'] if a['kind']=='choose_cell'}
        for wrong in [False,True]:
            p=self.packet();p.template.original_path='original.docx';p.candidate_history=[];p.save=Mock()
            def inspected(path):
                info=inspect_template(path)
                if wrong:info['warnings']=[UNMAPPED_NOTES['multi_column'].format(count=2)]
                return info
            with patch('core.preparation.inspect_template',side_effect=inspected):
                if wrong:
                    with self.assertRaises(prep.DocumentError):prep.build_candidate(p,choices)
                else:
                    prep.build_candidate(p,choices)
                    self.assertEqual(p.candidate_info['reviewed_multi_column_rows'],[{'part':'word/document.xml','row':0}])
                    self.assertIn(UNMAPPED_NOTES['multi_column'].format(count=1),p.candidate_info['warnings'])

    def test_prepared_source_expected_hash_refuses_before_any_parser_process(self):
        from documents.service import _run,DocumentError
        with patch('documents.service.subprocess.run') as process:
            with self.assertRaises(DocumentError):_run(self.root/'original.docx','fill',expected_sha256='0'*64)
        process.assert_not_called()


class WMContinuationTests(SimpleTestCase):
    setUp=PreparationPureTests.setUp
    config=PreparationPureTests.config
    packet=PreparationPureTests.packet
    # Classifier fixtures; the genuine unchanged WM mapping is independently
    # checked with its retained baseline and real networkless processor.
    def wm_packet(self):
        import hashlib,json
        from django.utils import timezone
        p=self.packet();p.pk='prep1';p.config=self.config();p.template.department='department1'
        p.template.original_path='original.docx';p.template.sha256=WM_IOV_SHA256
        raw=(Path(__file__).resolve().parents[3]/'WM IOV Template.docx').read_bytes()
        self.assertEqual(hashlib.sha256(raw).hexdigest(),WM_IOV_SHA256)
        (self.root/'original.docx').write_bytes(raw);(self.root/'candidate.docx').write_bytes(raw)
        inv=p.inventory;inv['source_sha256']=WM_IOV_SHA256
        inv['fields']=[{'key':'q'+str(i),'label':'Question '+str(i),'kind':'text','locator':{'index':i}} for i in range(90)]
        inv['questions']=[{'id':'q'+str(i),'text':'Question '+str(i),'protected_signing':False} for i in range(90)]
        inv['anchors']=[{'id':'a'+str(i),'question_id':'q'+str(i),'kind':'existing','locator':None,'protected_signing':False} for i in range(90)]
        inv['stories']=[{**inv['stories'][0],'part':'word/story'+str(i)+'.xml'} for i in range(16)]
        inv['review_requirements']=[{'id':'r'+str(i)} for i in range(14)]
        inv['document_backend']={'backend':'wm-iov-foss-v2','source_sha256':WM_IOV_SHA256,'mapping_digest':prep.WM_MAPPING_DIGEST}
        inv['inventory_sha256']=digest({k:v for k,v in inv.items() if k!='inventory_sha256'})
        p.proposal,_=prep.provider_role(inv,'proposer',self.config(),1)
        p.checker,_=prep.provider_role(inv,'checker',self.config(),1,p.proposal)
        p.checker['anchors'][0]['verdict']='unresolved'
        p.candidate_path='candidate.docx';p.candidate_sha256=WM_IOV_SHA256;p.candidate_info={}
        p.state='blocked';p.error_code='guided_layout';p.started_at=p.finished_at=timezone.now()
        p.correction_history=[];p.candidate_history=[{'path':p.candidate_path,'sha256':WM_IOV_SHA256,'revision':0,'placements':{}}]
        p.save=Mock();directory=self.root/'preparation-receipts';directory.mkdir(mode=0o700,exist_ok=True)
        receipts=[]
        for i,role in enumerate(['proposer','checker']):
            value=p.proposal if role=='proposer' else p.checker
            evidence={'model':'fixture-returned-'+role,'request_id':'','resource_id':'','simulated':True,'response':value}
            raw=json.dumps(evidence).encode();relative='preparation-receipts/'+str(i)*48+'.json'
            target=self.root/relative;target.write_bytes(raw);target.chmod(0o600)
            receipts.append(SimpleNamespace(pk=role,role=role,state='succeeded',finished_at=timezone.now(),
                response_path=relative,response_sha256=hashlib.sha256(raw).hexdigest(),returned_model=evidence['model'],
                provider_request_id='',provider_resource_id='',simulated=True,computed_cost=Decimal('0'),
                provider=p.config['provider'],requested_model=p.config['model'],
                request_sha256=digest(prep.role_payload(inv,role,p.proposal if role=='checker' else None))))
        p.receipts=Mock();p.receipts.order_by.return_value=receipts
        info={'fields':deepcopy(inv['fields']),'review_requirements':deepcopy(inv['review_requirements']),
              'warnings':list(prep.WM_LAYOUT_ADVISORIES),'sha256':WM_IOV_SHA256,'format':'docx',**inv['document_backend']}
        return p,info,receipts

    def test_exact_wm_warning_proof_keeps_full_visible_baseline_and_rejects_substitution(self):
        p,info,_=self.wm_packet()
        result=prep.candidate_inspection(p,info,{'reviewed_multi_column_rows':[]},{})
        self.assertEqual(result['warnings'],list(prep.WM_LAYOUT_ADVISORIES))
        self.assertEqual(result['fields'],p.inventory['fields'])
        self.assertEqual(result['qualified_warning_baseline']['inventory_sha256'],p.inventory['inventory_sha256'])
        for change in ['missing','extra','text','order','backend','mapping','source','format','fields','requirements','candidate','generic']:
            bad=deepcopy(info);q=deepcopy(p)
            if change=='missing':bad['warnings'].pop()
            elif change=='extra':bad['warnings'].append('Unmapped answer area.')
            elif change=='text':bad['warnings'][0]+=' changed'
            elif change=='order':bad['warnings'].reverse()
            elif change=='backend':bad['backend']='other'
            elif change=='mapping':bad['mapping_digest']='0'*64
            elif change=='source':bad['sha256']='0'*64
            elif change=='format':bad['format']='pdf'
            elif change=='fields':bad['fields'][0]['locator']['index']=True
            elif change=='requirements':bad['review_requirements'][0]['id']='changed'
            elif change=='candidate':q.candidate_sha256='0'*64
            elif change=='generic':q.inventory['document_backend']={};q.inventory['inventory_sha256']=digest({k:v for k,v in q.inventory.items() if k!='inventory_sha256'})
            with self.subTest(change=change),self.assertRaises((prep.DocumentError,ValueError)):
                prep.candidate_inspection(q,bad,{'reviewed_multi_column_rows':[]},{})

    def continuation(self,p,info):
        manager=Mock();manager.select_for_update.return_value.select_related.return_value.get.return_value=p
        actor=SimpleNamespace(pk='doctor1')
        with patch.object(prep.TemplatePreparation,'objects',manager),patch.object(prep,'membership'),\
                patch.object(prep,'inspect_template',return_value=info),patch.object(prep,'inspect_preparation',return_value=p.inventory),\
                patch.object(prep,'audit'),patch.object(prep,'provider_role',side_effect=AssertionError('No provider call')):
            return prep.continue_wm_review.__wrapped__(p.pk,actor,p.revision,p.candidate_sha256)

    def test_continuation_preserves_blocked_history_raw_roles_and_returns_guided_only(self):
        p,info,receipts=self.wm_packet()
        original={k:deepcopy(getattr(p,k)) for k in ['proposal','checker','error_code','started_at','finished_at','changes','candidate_path','candidate_sha256']}
        old_history=deepcopy(p.candidate_history);raw=[(self.root/r.response_path).read_bytes() for r in receipts]
        self.assertIs(self.continuation(p,info),p)
        self.assertEqual(p.state,'guided');self.assertEqual(p.revision,1);self.assertEqual(p.corrections,{})
        for key,value in original.items():self.assertEqual(getattr(p,key),value)
        self.assertEqual(p.candidate_history[:-1],old_history)
        self.assertEqual(p.candidate_history[-1]['prior_error_code'],'guided_layout')
        self.assertEqual(raw,[(self.root/r.response_path).read_bytes() for r in receipts])
        p.save.assert_called_once_with(update_fields=['candidate_info','candidate_history','revision','state'])
        self.assertFalse(prep.can_continue_wm_review(p))
        with self.assertRaises(ValidationError):self.continuation(p,info)

    def test_stale_revision_wrong_hash_unfinished_role_and_source_drift_refuse_without_save(self):
        for change in ['stale','hash','state','error','candidate_info','history','membership','receipt','request','raw','inventory','source']:
            p,info,receipts=self.wm_packet();manager=Mock()
            manager.select_for_update.return_value.select_related.return_value.get.return_value=p
            expected_revision=p.revision;expected_sha=p.candidate_sha256
            if change=='stale':expected_revision+=1
            elif change=='hash':expected_sha='0'*64
            elif change=='state':p.state='review'
            elif change=='error':p.error_code='provider_unknown'
            elif change=='candidate_info':p.candidate_info={'fields':[]}
            elif change=='history':p.candidate_history=[]
            elif change=='receipt':receipts[0].state='failed'
            elif change=='request':receipts[1].request_sha256='0'*64
            elif change=='raw':(self.root/receipts[0].response_path).write_bytes(b'{}')
            elif change=='source':(self.root/p.candidate_path).write_bytes(b'PK\x03\x04changed')
            membership=Mock(side_effect=ValidationError('Denied') if change=='membership' else None)
            rederived=deepcopy(p.inventory)
            if change=='inventory':rederived['issues'].append({'code':'changed'})
            with self.subTest(change=change),patch.object(prep.TemplatePreparation,'objects',manager),\
                    patch.object(prep,'membership',membership),patch.object(prep,'inspect_template',return_value=info),\
                    patch.object(prep,'inspect_preparation',return_value=rederived),patch.object(prep,'audit') as audit,\
                    patch.object(prep,'provider_role') as provider:
                with self.assertRaises((ValidationError,prep.DocumentError)):
                    prep.continue_wm_review.__wrapped__(p.pk,SimpleNamespace(pk='doctor1'),expected_revision,expected_sha)
                p.save.assert_not_called();audit.assert_not_called();provider.assert_not_called()

    def test_retained_role_reader_refuses_unsafe_file_and_incomplete_provider_result(self):
        import json,hashlib
        for change in ['fifo','symlink','mode','hardlink','path','hash','model','request','truncated','raw_simulated','raw_simulated_type','raw_simulated_missing']:
            p,_,receipts=self.wm_packet();r=receipts[0];path=self.root/r.response_path
            if change=='fifo':path.unlink();__import__('os').mkfifo(path)
            elif change=='symlink':path.unlink();path.symlink_to(self.root/receipts[1].response_path)
            elif change=='mode':path.chmod(0o644)
            elif change=='hardlink':__import__('os').link(path,self.root/'linked.json')
            elif change=='path':r.response_path='../original.docx'
            elif change=='hash':r.response_sha256='0'*64
            else:
                value=json.loads(path.read_bytes())
                if change=='model':value['model']='other'
                elif change=='request':value['request_id']='other'
                elif change=='raw_simulated':value['simulated']=False
                elif change=='raw_simulated_type':value['simulated']=1
                elif change=='raw_simulated_missing':value.pop('simulated')
                else:r.simulated=False;value['simulated']=False;value['response']={'model':r.returned_model,'id':'','stop_reason':'max_tokens','content':[]}
                raw=json.dumps(value).encode();path.write_bytes(raw);r.response_sha256=hashlib.sha256(raw).hexdigest()
            with self.subTest(change=change),self.assertRaises(ValidationError):prep.read_retained_role(r,'live' if change=='truncated' else 'fixture')
            if path.is_fifo() or path.is_symlink():path.unlink()
            if change=='hardlink':(self.root/'linked.json').unlink()

    def test_live_and_fixture_receipts_require_typed_mode_match_before_raw_evidence_reads(self):
        import json,hashlib
        for mode,simulated,allowed in [('fixture',True,True),('live',False,True),('live',True,False),
                ('fixture',False,False),('live',1,False),('fixture',0,False),('live','false',False),
                ('fixture',None,False),('unknown',True,False)]:
            p,info,receipts=self.wm_packet();p.config['mode']=mode
            for r in receipts:
                r.simulated=simulated
                # Both forms have otherwise valid complete original role JSON;
                # a simulated envelope cannot establish a live provider result.
                if simulated is False:
                    role_result=p.proposal if r.role=='proposer' else p.checker
                    value={'model':r.returned_model,'request_id':'','resource_id':'',
                           'simulated':False,'response':{'model':r.returned_model,'id':'','stop_reason':'end_turn',
                                       'content':[{'type':'text','text':json.dumps(role_result)}]}}
                    raw=json.dumps(value).encode();(self.root/r.response_path).write_bytes(raw)
                    r.response_sha256=hashlib.sha256(raw).hexdigest()
            with self.subTest(mode=mode,simulated=simulated):
                if allowed:self.assertEqual(self.continuation(p,info).state,'guided')
                else:
                    with self.assertRaises(ValidationError):self.continuation(p,info)
                    p.save.assert_not_called()
                    with patch.object(prep,'_read_retained_role') as read:
                        with self.assertRaises(ValidationError):prep.read_retained_role(receipts[0],mode)
                        read.assert_not_called()

    def test_live_envelope_cannot_disagree_with_its_raw_simulation_provenance(self):
        import json,hashlib
        for simulated in [True,1,None,'false']:
            p,info,receipts=self.wm_packet();p.config['mode']='live'
            for r in receipts:
                r.simulated=False;role_result=p.proposal if r.role=='proposer' else p.checker
                value={'model':r.returned_model,'request_id':'','resource_id':'','simulated':simulated,
                       'response':{'model':r.returned_model,'id':'','stop_reason':'end_turn',
                                   'content':[{'type':'text','text':json.dumps(role_result)}]}}
                raw=json.dumps(value).encode();(self.root/r.response_path).write_bytes(raw)
                r.response_sha256=hashlib.sha256(raw).hexdigest()
            with self.subTest(raw_simulated=simulated),self.assertRaises(ValidationError):self.continuation(p,info)
            p.save.assert_not_called()

    @override_settings(STORAGES={'default':{'BACKEND':'django.core.files.storage.FileSystemStorage'},
                                'staticfiles':{'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}})
    def test_continuation_post_only_and_blocked_form_ui_has_no_correction_controls(self):
        from django.test import RequestFactory
        from django.template.loader import render_to_string
        from core import preparation_views
        request=RequestFactory().get('/templates/unused/preparation/continue-review/')
        with patch.object(preparation_views,'selected') as selected:
            self.assertEqual(preparation_views.continue_review(request,'unused').status_code,405);selected.assert_not_called()
        p,_,_=self.wm_packet();p.template.pk='00000000-0000-0000-0000-000000000001';p.template.name='Blank form'
        html=render_to_string('template_preparation.html',{'user':SimpleNamespace(username='synthetic-doctor',email='',is_staff=False),'preparation':p,'template':p.template,'can_correct':False,
            'can_continue_wm_review':False,'anchor_rows':p.inventory['anchors'],'preview_token':'a'*48})
        self.assertNotIn('Save doctor corrections',html);self.assertNotIn('Recheck retained layout · no AI request',html)
        html=render_to_string('template_preparation.html',{'user':SimpleNamespace(username='synthetic-doctor',email='',is_staff=False),'preparation':p,'template':p.template,'can_correct':False,
            'can_continue_wm_review':True,'preview_token':'a'*48})
        self.assertIn('Recheck retained layout · no AI request',html);self.assertNotIn('Save doctor corrections',html)
