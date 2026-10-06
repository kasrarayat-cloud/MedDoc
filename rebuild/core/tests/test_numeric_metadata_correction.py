"""An explicit different date attribute does not correct a completed finding.

Actual retained procedure mapping/source; whole valid wire, page and finalization.
Global/uncertain corrections, other actors, dates and incomplete proof still hold.
"""
import copy
import json
from pathlib import Path
from io import BytesIO

from django.test import SimpleTestCase, override_settings
from docx import Document

from core import generation as g, v3_field_relations as r
from documents import word_v3
from .test_v3_field_relations import relation_wire
from .test_v3_supersession_floor import SONNET_55_DEFAULT, raw
from .test_transcript_context import proof


@override_settings(**SONNET_55_DEFAULT)
class NumericMetadataCorrection(SimpleTestCase):
    def setUp(self):
        fixture=Path(__file__).resolve().parents[3]/'docs/rebuild/v1/voice-retained-20261003/raw-asr-response-fixtures.json'
        case=next(c for c in json.loads(fixture.read_text())['cases']if c['attempt']=='005')
        self.fields=case['catalogue']
        self.field=next(f for f in self.fields if f['label']=='Polyps')

    def answers(self,notes,fields,picks,reviewed=True,relations=None):
        wire=relation_wire(notes,fields,picks)
        if relations:
            for item in wire['fields']:
                if item['key']in relations:item['context']=relations[item['key']]
        request=raw(notes,fields,wire,g.SONNET_5_5)
        if reviewed:request['patient_default_context']=proof(notes)
        return r.validated(request)

    def check(self,notes,supported,pick='No polyps were seen.',reviewed=True):
        answer=self.answers(notes,self.fields,{self.field['key']:[pick]},reviewed)[self.field['key']]
        self.assertEqual((answer['status'],answer['text']),('supported',pick)if supported else('conflict',''),(notes,answer))
        forced=answer if supported else{**answer,'status':'supported','text':pick,
            'evidence':[{'quote':pick,'start':notes.index(pick),'end':notes.index(pick)+len(pick)}]}
        page=g.as_currently_governed(g.governance_for(notes,self.fields),notes,{**self.field,**forced})
        self.assertEqual((page['status'],page['text']),('supported',pick)if supported else('conflict',''))
        self.assertEqual(g.governed_decisions(notes,self.fields,{self.field['key']:forced}),[]if supported else['Polyps'])
        return answer

    def test_retained_numeric_iso_and_newline_date_metadata_are_apart(self):
        for note in ('No polyps were seen. Correction: procedure date was 09/21/2026.',
                     'No polyps were seen. Correction: procedure date was 2026-09-21.',
                     'No polyps were seen.\nCorrection: procedure date was 09/21/2026.'):
            with self.subTest(notes=note):
                answer=self.check(note,True)
                self.assertEqual(answer['evidence'],[{'quote':'No polyps were seen.','start':0,'end':20}])

    def test_invalid_ambiguous_uncertain_negative_and_foreign_date_changes_stay_held(self):
        endings=('procedure date was 09/21/2026?', 'procedure date was 02/30/2026.',
            'procedure date was 09/10/2026.', 'procedure date was 09/21.',
            'procedure date was not 09/21/2026.', 'procedure date was possibly 09/21/2026.',
            'if procedure date was 09/21/2026.', 'procedure date will be 09/21/2026.',
            'procedure date was yesterday.', 'procedure time was 09:21.',
            "her father's procedure date was 09/21/2026.", 'my procedure date was 09/21/2026.',
            'the procedure was entered in error.', 'all results were entered in error.',
            'the conclusion was entered in error.', 'those findings were incorrect.',
            'procedure date was 09/21/2026. Correction: all findings were wrong.')
        for end in endings:
            with self.subTest(end=end):self.check('No polyps were seen. Correction: '+end,False)
        for notes,pick in (
            ('The doctor underwent colonoscopy. No polyps were seen. Correction: procedure date was 09/21/2026.','No polyps were seen.'),
            ('Her father had no polyps. Correction: procedure date was 09/21/2026.','Her father had no polyps.'),
            ('No polyps were seen on 09/20/2026. Correction: procedure date was 09/21/2026.','No polyps were seen on 09/20/2026.'),
            ('No polyps were seen. Polyps were seen. Correction: procedure date was 09/21/2026.','No polyps were seen.')):
            with self.subTest(notes=notes):self.check(notes,False,pick)

    def test_date_destination_calendar_only_and_mixed_evidence_do_not_borrow_the_scope(self):
        doc=Document();doc.add_paragraph('Procedure date: __________');doc.add_paragraph('Polyps: __________')
        out=BytesIO();doc.save(out)
        fields=word_v3.inspect(out.getvalue())['fields']
        date=next(f for f in fields if f['label']=='Procedure date')
        notes='Procedure date: 09/20/2026\nCorrection: procedure date was 09/21/2026.'
        answer=self.answers(notes,fields,{date['key']:['Procedure date: 09/20/2026']})[date['key']]
        self.assertEqual((answer['status'],answer['text']),('conflict',''))
        notes='No polyps were seen. Correction: procedure date was 09/21/2026.'
        correction='Correction: procedure date was 09/21/2026.'
        for picks in ([correction],['No polyps were seen.',correction]):
            with self.subTest(picks=picks):
                answer=self.answers(notes,self.fields,{self.field['key']:picks})[self.field['key']]
                self.assertEqual((answer['status'],answer['text']),('conflict',''))

    def test_missing_qualifier_incomplete_relation_and_unknown_patient_proof_remain_held(self):
        notes='No polyps were seen. Correction: procedure date was 09/21/2026.'
        self.check(notes,False,reviewed=False) # separate source-ownership limitation, not a date exception
        complete=self.check(notes,True)
        for relation in (None,'not-a-signed-relation'):
            stored={**complete,'ordinary_relation':relation}
            self.assertEqual(g.governed_decisions(notes,self.fields,{self.field['key']:stored}),['Polyps'])
        wire=relation_wire(notes,self.fields,{self.field['key']:['No polyps were seen.']})
        item=next(i for i in wire['fields']if i['key']==self.field['key'])
        item['context']['independent'].remove(2);item['context']['qualifier'].append(2)
        request=raw(notes,self.fields,wire,g.SONNET_5_5);request['patient_default_context']=proof(notes)
        answer=r.validated(request)[self.field['key']]
        self.assertEqual((answer['status'],answer['text']),('conflict',''))

    def test_scope_is_per_answer_and_original_global_governors_are_preserved(self):
        notes='No polyps were seen. Correction: procedure date was 09/21/2026.'
        answer=self.check(notes,True)
        for reverse in (False,True):
            governance=g.governance_for(notes,self.fields)
            original=copy.deepcopy((governance[4],governance[5],governance[7]))
            variants=[answer,{**answer,'ordinary_relation':None}]
            if reverse:variants.reverse()
            for stored in variants:
                context=g.page_context(governance,notes,'Polyps',stored['evidence'],answer_text=stored['text'],
                    ordinary_relation=stored.get('ordinary_relation'),field_key=self.field['key'])
                self.assertEqual(context is None,bool(stored.get('ordinary_relation')))
            self.assertEqual((governance[4],governance[5],governance[7]),original)
        self.check(notes+'\nCorrection: all prior observations were entered in error.',False)

    def test_original_elliptical_blood_loss_ambiguity_c08_remains_held(self):
        field=next(f for f in self.fields if f['label']=='Estimated blood loss')
        first='Estimated blood loss was zero milliliters.'
        source=first+' Less than one milliliter. Correction: procedure date was 09/21/2026.'
        answer=self.answers(source,self.fields,{field['key']:[first]})[field['key']]
        self.assertEqual((answer['status'],answer['text']),('conflict',''))

    def test_explicit_patient_typed_prose_is_useful_without_transcript_context_or_labels(self):
        notes='The patient had no polyps.'
        self.check(notes,True,pick=notes,reviewed=False)

    def test_backward_source_dependencies_survive_a_separate_date_correction(self):
        for source in ('The earlier findings were wrong.', 'The earlier findings were incorrect.',
                'This finding was wrong.', 'The earlier conclusion was wrong.',
                'The clinician invalidated the earlier report.', 'The earlier report was invalidated.',
                'The prior assessment was superseded.', 'The earlier report cannot be relied upon.',
                'Do not use the earlier findings.', 'The clinician doubted the preceding findings.',
                'The earlier findings may be wrong.', 'Were the earlier findings wrong?',
                'The earlier findings were probably not wrong.', 'The earlier findings may not be wrong.',
                'If the earlier findings were not wrong, review them.',
                'The earlier report was for the wrong patient.'):
            with self.subTest(source=source):
                self.check('No polyps were seen. '+source+' Correction: procedure date was 09/21/2026.',False)

    def test_closed_noninvalidating_status_and_other_date_targets_keep_their_polarity(self):
        for source in ('The earlier findings were not wrong.', 'The earlier findings were correct.',
                'The clinician did not invalidate the earlier report.',
                'The clinician validated the earlier report.',
                "Her father's procedure date was wrong.", "Dr Synthetic's appointment date was wrong.",
                'The earlier procedure date was wrong.'):
            with self.subTest(source=source):
                self.check('No polyps were seen. '+source+' Correction: procedure date was 09/21/2026.',True)
        for separator in (' ', '\n'):
            with self.subTest(separator=separator):
                self.check('No polyps were seen.'+separator+'Correction: examination date was 09/21/2026.',True)

    def test_unknown_backward_reference_and_wider_quotes_are_not_assumed_independent(self):
        for source in ('The clinician removed the earlier dressing.',
                '"The earlier findings were wrong."', 'Was the earlier procedure date wrong?'):
            with self.subTest(source=source):
                self.check('No polyps were seen. '+source+' Correction: procedure date was 09/21/2026.',False)

    def test_exact_root_iso_examination_subject_and_negative_iso_controls(self):
        for separator in (' ', '\n'):
            with self.subTest(separator=separator):
                self.check('No polyps were seen.'+separator+'Correction: examination date was 2026-09-22.',True)
        for change in ('examination date was 2026-02-30.', 'examination date was 2026-09-22?',
                'examination date was possibly 2026-09-22.', 'examination date was not 2026-09-22.',
                "her father's examination date was 2026-09-22.", 'my examination date was 2026-09-22.',
                'examination date will be 2026-09-22.'):
            with self.subTest(change=change):
                self.check('No polyps were seen. Correction: '+change,False)
        self.check('No polyps were seen. The earlier findings were wrong. '
                   'Correction: examination date was 2026-09-22.',False)

    def test_demonstrative_date_subject_cannot_discharge_a_literal_pointer(self):
        for change in ('that procedure date was 09/21/2026.',
                'this procedure date was 09/21/2026.', 'these procedure dates were 09/21/2026.',
                'those procedure dates were 09/21/2026.',
                'that examination date was 2026-09-22.', 'this examination date was 2026-09-22.'):
            with self.subTest(change=change):
                self.check('No polyps were seen. Correction: '+change,False)
        for change in ('procedure date was 09/21/2026.', 'the procedure date was 09/21/2026.',
                'examination date was 2026-09-22.', 'the examination date was 2026-09-22.'):
            with self.subTest(change=change):
                self.check('No polyps were seen. Correction: '+change,True)
