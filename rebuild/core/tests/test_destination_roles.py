"""Destination authority through the clinical workflow: whether a place may receive an answer at all.

A model never creates permission to write a signature, perform an attestation or change an identity field; the
permission comes from the ORIGINAL's structure (core.destination_roles). Synthetic blank forms only; every provider is
mocked and no request leaves the process. RF../PC.. ids are the destination-authority matrix's red-first exploits and
positive controls (v1-product-20260930 a1 mapping-matrix redfirst.json).
"""
import copy
import hashlib
import io
import itertools
import json
from types import SimpleNamespace
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from core import destination_roles, generation, mapping_assistance as mapping, v3_mapping
from core.jobs import enqueue, run_one
from core.models import Encounter, Job, MappingConfirmation, Reservation, Template, TemplateVersion
from core.security import private_path
from core.tests.test_mapping_assistance import SETTINGS as MAPPING_SETTINGS
from core.tests.test_word_v3_workflow import V3, V3Fixture, part_text, import_metadata
from core.tests.test_workflow import TEST_SETTINGS
from documents import DocumentError, fill_template, word_v3

BODY = 'word/document.xml#'
LIVE = {'AI_MODE': 'live', 'AI_PROVIDER': 'anthropic', 'AI_MODEL': 'claude-sonnet-4-6', 'AI_INPUT_RATE': '3',
        'AI_OUTPUT_RATE': '15', 'AI_CACHED_RATE': '.3'}


def saved(doc):
    out = io.BytesIO(); doc.save(out); return out.getvalue()


def box(doc, text=None, pict=False):
    table = doc.add_table(rows=1, cols=1)
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); borders.append(edge)
    table.cell(0, 0)._tc.get_or_add_tcPr().append(borders)
    if text:
        table.cell(0, 0).paragraphs[0].add_run(text)
    if pict:
        table.cell(0, 0).paragraphs[0].add_run()._r.append(OxmlElement('w:pict'))
    return table


def lines(*texts, heading='Clinic record'):
    doc = Document(); doc.add_heading(heading, level=2)
    for text in texts:
        doc.add_paragraph(text)
    return saved(doc)


def legacy_rows(*labels):
    doc = Document(); doc.add_paragraph('Clinic record')
    table = doc.add_table(rows=len(labels), cols=2)
    for n, label in enumerate(labels):
        table.cell(n, 0).text = label
    return saved(doc)


def forged(data, slot, kind='text', label='Forged answer'):
    """A field a stored mapping could hold from before this change: exactly the place the engine's locate finds, with
    only the role check bypassed (never the placement)."""
    with patch.object(word_v3, 'destination_roles_of', return_value=[]):
        return word_v3.locate(data, [{'slot': BODY + slot, 'kind': kind, 'label': label}])[0]


def provider_capture(sent):
    """A mocked Anthropic client: records every request and answers every asked question with nothing."""
    def create(**request):
        sent.append(request)
        payload = json.loads(request['messages'][0]['content'])
        units = 'source_units' in payload
        body = {'fields': [{'key': q['key'], ('unit_ids' if units else 'quotes'): [], 'conflict': False}
                           for q in payload['questions']]}
        return SimpleNamespace(usage=SimpleNamespace(input_tokens=1, output_tokens=1, cache_read_input_tokens=0,
                                                     cache_creation_input_tokens=0),
                               stop_reason='end_turn', model='claude-sonnet-4-6', id='msg_synthetic',
                               content=[SimpleNamespace(type='text', text=json.dumps(body))], _request_id='req_synthetic')
    client = SimpleNamespace(messages=SimpleNamespace(create=create))
    return patch('anthropic.Anthropic', return_value=client)


@override_settings(**V3)
class V3Gates(V3Fixture, TestCase):
    SIGNATURE_FORM = ('Signature: ______________', 'Plan: ______')

    def upload_lines(self, *texts, name='Synthetic signing line'):
        data = lines(*texts)
        template, version = self.uploaded(data=data, name=name)
        return data, template, version

    def stored(self, template, fields):
        """A stored mapping version (and its confirmation) as a release before this change could have left it."""
        latest = self.latest(template)
        version = TemplateVersion.objects.create(template=template, number=latest.number + 1, fields=fields,
                                                 contract='v3', created_by=self.user, **import_metadata(latest, fields))
        MappingConfirmation.objects.create(version=version, confirmed_by=self.user)
        return version

    def encounter(self, version, notes):
        self.assertEqual(self.new_encounter(version, notes=notes).status_code, 302)
        return Encounter.objects.latest('created_at')

    def test_rf01_remove_then_add_text_on_a_signature_line_is_refused_signature_kind_is_allowed(self):
        data, template, version = self.upload_lines(*self.SIGNATURE_FORM)
        signature = next(f for f in version.fields if f['kind'] == 'signature')
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_save', 'mapping_version': version.pk,
                                                                  'remove_' + signature['key']: 'on'})
        self.assertEqual(response.status_code, 302)
        latest = self.latest(template)
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_add', 'mapping_version': latest.pk,
                                    'slot': BODY + 'b/p[1]', 'kind': 'text', 'label': 'Signature'})
        self.assertEqual(response.status_code, 400)
        self.assertIn('signing, initials or consent area', response.context['v3_error'])
        self.assertEqual(self.latest(template).pk, latest.pk)
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_add', 'mapping_version': latest.pk,
                                    'slot': BODY + 'b/p[1]', 'kind': 'signature', 'label': 'Signature'})
        self.assertEqual(response.status_code, 302)
        self.confirm(template)

    def test_rf18_a_manual_answer_in_a_signing_question_box_is_refused(self):
        doc = Document(); doc.add_heading('Consultant review', level=2); doc.add_paragraph('Signed by the consultant:')
        box(doc); doc.add_paragraph('Plan: ______')
        template, version = self.uploaded(data=saved(doc), name='Synthetic signing box')
        for slot in ('b/tbl[0]/tr[0]/tc[0]', 'b/tbl[0]/tr[0]/tc[0]/p[0]'):
            with self.subTest(slot=slot):
                response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_add', 'mapping_version': version.pk,
                                            'slot': BODY + slot, 'kind': 'long_text', 'label': 'Consultant review'})
                self.assertEqual(response.status_code, 400)
        self.assertEqual(template.versions.count(), 1)

    def test_rf19_a_stored_mapping_on_a_signature_place_is_held_everywhere(self):
        """not confirmed, not asked, save 400, finalize refused, fill protected_signing_area; Plan unaffected."""
        data, template, version = self.upload_lines(*self.SIGNATURE_FORM)
        plan = next(f for f in version.fields if f['label'] == 'Plan')
        bad = forged(data, 'b/p[1]', label='Signature')
        fields = [plan, bad]
        # confirmation of such a mapping is refused (a fresh version, unconfirmed)
        latest = TemplateVersion.objects.create(template=template, number=2, fields=fields, contract='v3', created_by=self.user,
                                                **import_metadata(version, fields))
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_confirm', 'mapping_version': latest.pk,
                                                                   'checked_every_field': 'on'})
        self.assertEqual(response.status_code, 400)
        self.assertIn('Signature', response.context['v3_error'])
        self.assertIsNone(TemplateVersion.objects.get(pk=latest.pk).confirmed)
        # a stored confirmed version from before this change
        stored = self.stored(template, fields)
        self.assertEqual(destination_roles.v3_held(stored), {bad['key']: word_v3.ROLE_REFUSALS['signing']})
        e = self.encounter(stored, 'Signature: Dr Forged Signer\nPlan: Rest and fluids.')
        enqueue(Encounter.objects.get(pk=e.pk), self.user); self.assertTrue(run_one())
        e.refresh_from_db()
        self.assertEqual((e.answers[bad['key']]['status'], e.answers[bad['key']]['text']), ('missing', ''))
        self.assertEqual(e.answers[bad['key']]['reason'], word_v3.ROLE_REFUSALS['signing'])
        self.assertEqual(e.answers[plan['key']]['text'], 'Rest and fluids.')
        # never asked: the live request carries no question for it (its label stays a context label)
        sent = []
        with override_settings(**LIVE), provider_capture(sent):
            generation.generate(e.notes, destination_roles.generation_fields(stored), generation.configuration())
        payload = json.loads(sent[0]['messages'][0]['content'])
        self.assertEqual([q['key'] for q in payload['questions']], [plan['key']])
        self.assertIn('Signature', payload['context_labels'])
        # save: any text there is refused, nothing saved
        post = {'revision': e.revision, 'answer_' + plan['key']: 'Rest and fluids.', 'reviewed_' + plan['key']: 'on',
                'answer_' + bad['key']: 'Dr Forged Signer', 'reviewed_' + bad['key']: 'on'}
        response = self.client.post(f'/encounters/{e.pk}/save/', post)
        self.assertEqual(response.status_code, 400)
        self.assertIn('Signature', response.content.decode())
        before = Encounter.objects.get(pk=e.pk)
        self.assertEqual(before.revision, e.revision)
        # finalize: refused, no export (an answer stored by an earlier release)
        answers = dict(before.answers)
        answers[bad['key']] = {'text': 'Dr Forged Signer', 'status': 'edited', 'reviewed': True, 'reviewed_by': str(self.user.pk), 'evidence': []}
        answers[plan['key']] = {**answers[plan['key']], 'reviewed': True, 'reviewed_by': str(self.user.pk)}
        Encounter.objects.filter(pk=e.pk).update(answers=answers)
        response = self.client.post(f'/encounters/{e.pk}/finalize/', {'revision': before.revision, 'reviewed': 'on'}, follow=True)
        self.assertIn('never takes a typed or generated answer', response.content.decode())
        self.assertEqual(Encounter.objects.get(pk=e.pk).status, 'review')
        # the fill itself refuses by the original's roles, never trusting the mapping
        private_path('exports').mkdir(parents=True, exist_ok=True)
        with self.assertRaises(DocumentError) as refused:
            fill_template(private_path(template.original_path), fields, {bad['key']: 'Dr Forged Signer'},
                          private_path('exports/forged.docx'))
        self.assertEqual(refused.exception.code, 'protected_signing_area')

    def test_pc01_name_beside_a_signature_is_generated_and_exported(self):
        data, template, version = self.upload_lines('Name: ______    Signature: ______', 'Plan: ______')
        name = next(f for f in version.fields if f['label'] == 'Name')
        self.confirm(template, version)
        e = self.encounter(version, 'Name: Synthetic Example\nPlan: Rest.')
        enqueue(Encounter.objects.get(pk=e.pk), self.user); self.assertTrue(run_one())
        e.refresh_from_db()
        self.assertEqual((e.answers[name['key']]['status'], e.answers[name['key']]['text']), ('supported', 'Synthetic Example'))
        post = {'revision': e.revision}
        for f in version.fields:
            if f['kind'] in v3_mapping.SIGNING_KINDS:
                continue
            post['answer_' + f['key']] = e.answers[f['key']]['text']; post['reviewed_' + f['key']] = 'on'
        self.assertEqual(self.client.post(f'/encounters/{e.pk}/save/', post).status_code, 302)
        e.refresh_from_db()
        self.client.post(f'/encounters/{e.pk}/finalize/', {'revision': e.revision, 'reviewed': 'on'})
        e.refresh_from_db()
        self.assertEqual(e.status, 'finalized')
        self.assertIn('Name: Synthetic Example', part_text(private_path(e.export_path).read_bytes()))


@override_settings(**dict(TEST_SETTINGS, QUILLMD_WORD_ENGINE='legacy'))
class LegacyDefaultPath(V3Fixture, TestCase):
    """The retained ambiguous form refuses; distinct supported legacy forms exercise human review."""
    LABELS = ('Patient signature', 'Consultant signature', 'Initials', 'Consent form signed', 'Initial assessment',
              'Electronically signed by', 'Plan')
    NOTES = ('Patient signature: John Example\nConsultant signature: Dr Example\nInitials: JE\n'
             'Consent form signed: Yes\nInitial assessment: Stable\nElectronically signed by: Dr Other\nPlan: Rest.')

    def setUp(self):
        super().setUp()
        self.data = legacy_rows(*self.LABELS)
        self.template, self.version = self.uploaded(data=self.data, name='Synthetic legacy roles')
        self.by = self.by_label(self.version)
        self.assertEqual(self.version.contract, '')
        original, imported = self.template, self.version
        self.assertTrue(imported.import_manifest['coverage']['block_generation'])
        self.assertEqual(self.new_encounter(imported, notes=self.NOTES).status_code, 302)
        with self.assertRaisesMessage(ValidationError, 'No generation or charge'):
            enqueue(Encounter.objects.latest('created_at'), self.user)
        # A synthetic record representing an old, unbound import keeps the
        # exact original/map. Its missing manifest cannot authorize this area.
        historical = Template.objects.create(department=self.department, name='Synthetic historical unbound roles',
            format=original.format, original_name=original.original_name, original_path=original.original_path,
            sha256=original.sha256, created_by=self.user)
        version = TemplateVersion.objects.create(template=historical, number=1, fields=copy.deepcopy(imported.fields),
                                                 created_by=self.user)
        self.assertEqual(version.fields, imported.fields)
        self.assertEqual(self.new_encounter(version, notes=self.NOTES).status_code, 302)
        with self.assertRaisesMessage(ValidationError, 'human-owned or unresolved answer area'):
            enqueue(Encounter.objects.latest('created_at'), self.user)
        self.assertEqual(Job.objects.count(), 0); self.assertEqual(Reservation.objects.count(), 0)
        self.assertEqual(private_path(original.original_path).read_bytes(), self.data)
        # Supplemental supported source, not a correction or pass of the
        # original Consent-form case. Known signing fields stay human-only.
        labels = tuple('Signed consent was received on' if label == 'Consent form signed' else label for label in self.LABELS)
        self.supported_data = legacy_rows(*labels)
        self.template, self.version = self.uploaded(data=self.supported_data, name='Synthetic supported legacy roles')
        self.by = self.by_label(self.version)
        self.assertFalse(self.version.import_manifest['coverage'].get('block_generation', False))
        self.notes = self.NOTES.replace('Consent form signed: Yes', 'Signed consent was received on: 3 October 2026')

    def generated(self):
        self.assertEqual(self.new_encounter(self.version, notes=self.notes).status_code, 302)
        e = Encounter.objects.latest('created_at')
        enqueue(Encounter.objects.get(pk=e.pk), self.user); self.assertTrue(run_one())
        e.refresh_from_db()
        return e

    def test_rf16_rf17_signature_initials_and_signer_rows_are_never_generated_checklist_rows_are(self):
        e = self.generated()
        for label in ('Patient signature', 'Consultant signature', 'Initials', 'Electronically signed by'):
            with self.subTest(label=label):
                answer = e.answers[self.by[label]['key']]
                self.assertEqual((answer['status'], answer['text']), ('missing', ''))
                self.assertTrue(answer['reason'])
        for label, text in (('Signed consent was received on', '3 October 2026'), ('Initial assessment', 'Stable'), ('Plan', 'Rest.')):
            with self.subTest(label=label):   # PC09 / PC07: checklist states and adjectives stay ordinary answers
                self.assertEqual(e.answers[self.by[label]['key']]['text'], text)
        sent = []
        with override_settings(**LIVE), provider_capture(sent):
            generation.generate(e.notes, destination_roles.generation_fields(self.version), generation.configuration())
        asked = [q['label'] for q in json.loads(sent[0]['messages'][0]['content'])['questions']]
        self.assertEqual(asked, ['Signed consent was received on', 'Initial assessment', 'Plan'])

    def review_all(self, e, override=()):
        post = {'revision': e.revision}
        for f in self.version.fields:
            post['answer_' + f['key']] = dict(override).get(f['label'], e.answers.get(f['key'], {}).get('text', ''))
            post['reviewed_' + f['key']] = 'on'
        return self.client.post(f'/encounters/{e.pk}/save/', post)

    def test_generated_or_accepted_text_on_a_role_row_is_refused_typed_text_exports(self):
        e = self.generated()
        key = self.by['Patient signature']['key']
        # a draft from before this change could hold generated text there
        Encounter.objects.filter(pk=e.pk).update(answers={**e.answers, key: {'text': 'John Example', 'status': 'supported',
                                                                              'evidence': [], 'reviewed': False}})
        e.refresh_from_db()
        response = self.review_all(e)
        self.assertEqual(response.status_code, 400)
        self.assertIn('Patient signature', response.content.decode())
        self.assertEqual(Encounter.objects.get(pk=e.pk).revision, e.revision)
        # the clinician's own typed entry is theirs, and it exports (legacy: provenance, not a blanket refusal)
        response = self.review_all(e, override={'Patient signature': 'J. Example (signed on paper)'})
        self.assertEqual(response.status_code, 302)
        e.refresh_from_db()
        self.assertEqual(e.answers[key]['status'], 'edited')
        self.client.post(f'/encounters/{e.pk}/finalize/', {'revision': e.revision, 'reviewed': 'on'})
        e.refresh_from_db()
        self.assertEqual(e.status, 'finalized')
        self.assertIn('J. Example (signed on paper)', part_text(private_path(e.export_path).read_bytes()))

    def test_finalize_refuses_generated_text_on_a_role_row(self):
        e = self.generated()
        key = self.by['Initials']['key']
        answers = {k: {**v, 'reviewed': True, 'reviewed_by': str(self.user.pk)} for k, v in e.answers.items()}
        answers[key] = {'text': 'JE', 'status': 'supported', 'evidence': [], 'reviewed': True, 'reviewed_by': str(self.user.pk)}
        Encounter.objects.filter(pk=e.pk).update(answers=answers, status='review')
        response = self.client.post(f'/encounters/{e.pk}/finalize/', {'revision': e.revision, 'reviewed': 'on'}, follow=True)
        self.assertIn('never fills from notes', response.content.decode())
        self.assertEqual(Encounter.objects.get(pk=e.pk).status, 'review')

    def test_pc14_the_signer_row_keeps_the_finalizers_name_rule(self):
        e = self.generated()
        key = self.by['Electronically signed by']['key']
        response = self.review_all(e, override={'Electronically signed by': 'Dr Somebody Else'})
        self.assertEqual(response.status_code, 302)
        e.refresh_from_db()
        response = self.client.post(f'/encounters/{e.pk}/finalize/', {'revision': e.revision, 'reviewed': 'on'}, follow=True)
        self.assertIn('must match your signed-in clinician name', response.content.decode())
        name = ' '.join((self.user.get_full_name() or self.user.email).split())
        e.refresh_from_db()
        response = self.review_all(e, override={'Electronically signed by': name})
        e.refresh_from_db()
        self.client.post(f'/encounters/{e.pk}/finalize/', {'revision': e.revision, 'reviewed': 'on'})
        e.refresh_from_db()
        self.assertEqual(e.status, 'finalized')
        self.assertEqual(e.answers[key]['text'], name)


@override_settings(**TEST_SETTINGS)
class LegacyIdentityReadsThePrintedLabel(V3Fixture, TestCase):
    def test_l05_participant_signature_stays_blank_after_mapping_rename(self):
        """The retained FV-01 source no longer appoints a finalizer after a mapped rename."""
        template, version = self.uploaded(data=legacy_rows("Participant's signature", 'Plan'), name='Synthetic L05')
        by = self.by_label(version)
        response = self.client.post(f'/templates/{template.pk}/', {
            'mapping_version': version.pk, 'labels': json.dumps({by["Participant's signature"]['key']: 'Signature',
                                                                  by['Plan']['key']: 'Plan'})})
        latest = self.latest(template)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.by_label(latest)['Signature']['key'], by["Participant's signature"]['key'])
        self.assertEqual(self.new_encounter(latest, notes='Plan: Rest.').status_code, 302)
        e = Encounter.objects.latest('created_at')
        key = by["Participant's signature"]['key']
        post = {'revision': e.revision, 'answer_' + key: '', 'blank_' + key: 'on', 'reviewed_' + key: 'on',
                'answer_' + by['Plan']['key']: 'Rest.', 'reviewed_' + by['Plan']['key']: 'on'}
        # A mapped rename cannot appoint the finalizer in another person's area.
        self.assertEqual(self.client.post(f'/encounters/{e.pk}/save/', post).status_code, 302)
        e.refresh_from_db()
        self.client.post(f'/encounters/{e.pk}/finalize/', {'revision': e.revision, 'reviewed': 'on'})
        e.refresh_from_db(); self.assertEqual(e.status, 'finalized')
        self.assertEqual(Document(private_path(e.export_path)).tables[0].cell(0, 1).text, '')


@override_settings(**MAPPING_SETTINGS)
class ModelProposals(V3Fixture, TestCase):
    def setUp(self):
        super().setUp()
        self.config = generation.configuration()

    def version_of(self, data, name, drop=()):
        template, version = self.uploaded(data=data, name=name)
        if drop:
            fields = [f for f in version.fields if f['label'] not in drop]
            self.assertEqual(len(fields), len(version.fields) - len(drop))
            version = TemplateVersion.objects.create(template=template, number=2, fields=fields, contract='v3',
                                                     created_by=self.user, **import_metadata(version, fields))
        return template, version

    def compiled(self, version):
        compiled = mapping.compile_request(version, self.config)
        return compiled, json.loads(compiled['request']['messages'][0]['content'])

    @staticmethod
    def change(payload, slot, text, kind='text', fid=-1, glyph=-1, role=0):
        sid = next(n for n, row in enumerate(payload['slots']) if row['slot'] == BODY + slot)
        source = payload['slots'][sid][('text', 'heading', 'row_label')[role]]
        ref = payload['text_refs'].index([sid, role, 0])
        start = source.index(text)
        return [fid, sid, mapping.wire.KINDS.index(kind), ref, start, start + len(text), glyph]

    @staticmethod
    def keep_all(payload):
        body = mapping.wire.fixture(payload)
        return body

    def test_rf07_rf08_rf09_a_model_addition_is_never_adopted_it_is_a_review_item(self):
        cases = []
        data = lines('Signature: ______________', 'Plan: ______')
        cases.append(('RF07', data, ('Signature',), 'b/p[1]', 'Signature', 'text'))
        for fill in (None, 'pict', 'Use black ink'):
            doc = Document(); doc.add_heading('Consultant review', level=2); doc.add_paragraph('Signed by the consultant:')
            box(doc, text=fill if fill not in (None, 'pict') else None, pict=fill == 'pict'); doc.add_paragraph('Plan: ______')
            cases.append(('RF08-%s' % fill, saved(doc), (), 'b/tbl[0]/tr[0]/tc[0]/p[0]', 'Consultant review', 'long_text'))
        doc = Document(); doc.add_heading('Clinic record', level=2)
        t = doc.add_table(rows=2, cols=2); t.cell(0, 0).text = 'Name'; t.cell(0, 1).text = 'Signature'; t.cell(1, 1).text = 'Use black ink'
        cases.append(('RF09', saved(doc), (), 'b/tbl[0]/tr[1]/tc[1]/p[0]', 'Use black ink', 'text'))
        for n, (rid, data, drop, slot, label, kind) in enumerate(cases):
            with self.subTest(rid=rid):
                template, version = self.version_of(data, 'Synthetic %s' % rid, drop)
                compiled, payload = self.compiled(version)
                body = self.keep_all(payload)
                role = 1 if label == 'Consultant review' else 0
                body['x'] = [self.change(payload, slot, label, kind, role=role)]
                result = mapping.validate_proposal(version, compiled, body)
                self.assertEqual(result['fields'], version.fields, 'nothing is adopted')
                self.assertEqual(result['review']['changes'], [])
                proposed = [row for row in result['review']['unresolved'] if row.get('proposal')]
                self.assertEqual([(row['slot'], row['proposal']['label']) for row in proposed], [(BODY + slot, label)])
                self.assertIn('never adopted automatically', proposed[0]['reason'])
                if rid != 'RF09':
                    self.assertIn('only a Signature or Initials field can be added', proposed[0]['reason'])

    def test_rf11_one_invalid_change_is_held_and_a_valid_relabel_is_adopted(self):
        doc = Document(); doc.add_paragraph('Section 3: Public protection')
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = 'Does a formal offender supervision scheme currently apply to this person?'
        t.cell(1, 0).text = 'Date the lead agency was last informed'
        template, version = self.version_of(saved(doc), 'Synthetic RF11')
        compiled, payload = self.compiled(version)
        body = self.keep_all(payload)
        date = next(f for f in version.fields if f['kind'] == 'date')
        other = next(f for f in version.fields if f['key'] != date['key'])
        fid_date = next(n for n, f in enumerate(payload['fields']) if f['key'] == date['key'])
        fid_other = next(n for n, f in enumerate(payload['fields']) if f['key'] == other['key'])
        for code in ('k',):
            body['f'][code] = [n for n in body['f'][code] if n not in (fid_date, fid_other)]
        body['f']['r'] = [fid_other, fid_date]
        date_slot = date['locator']['slot']
        other_slot = other['locator']['slot']
        # invalid: at its own cell, the label quotes ANOTHER slot's text (the retained evidence/002 change-0 shape)
        question = next(n for n, row in enumerate(payload['slots']) if row['slot'] == BODY + 'b/tbl[0]/tr[0]/tc[0]/p[0]')
        invalid = [fid_other, next(n for n, row in enumerate(payload['slots']) if row['slot'] == BODY + other_slot),
                   0, payload['text_refs'].index([question, 0, 0]), 0, 22, -1]
        valid = self.change(payload, date_slot, 'Date the lead agency was last informed', kind='date', fid=fid_date, role=2)
        body['x'] = [invalid, valid]
        result = mapping.validate_proposal(version, compiled, body)
        adopted = next(f for f in result['fields'] if f['key'] == date['key'])
        self.assertEqual((adopted['label'], adopted['locator'], adopted['origin']),
                         ('Date the lead agency was last informed', date['locator'], 'model_assisted'))
        self.assertEqual(next(f for f in result['fields'] if f['key'] == other['key']), other)
        self.assertEqual(result['review']['changes'], [{'key': date['key'], 'replacement': date['key']}])
        self.assertEqual([r['wire'] for r in result['review']['rejected_changes']], [invalid])

    def test_rf12_the_outcome_is_the_same_for_every_order_of_the_changes(self):
        doc = Document(); doc.add_heading('Clinic record', level=2); doc.add_paragraph('Pulse findings: ______')
        doc.add_paragraph('Allergies'); doc.add_paragraph('Smoker ☐ Yes ☐ No'); doc.add_paragraph('Plan: ______')
        template, version = self.version_of(saved(doc), 'Synthetic permutations', drop=('Smoker',))
        compiled, payload = self.compiled(version)
        pulse = next(f for f in version.fields if f['label'] == 'Pulse findings')
        fid = next(n for n, f in enumerate(payload['fields']) if f['key'] == pulse['key'])
        base = self.keep_all(payload); base['f']['k'].remove(fid); base['f']['r'] = [fid]
        changes = [self.change(payload, 'b/p[1]', 'Pulse findings', fid=fid, kind='long_text'),
                   self.change(payload, 'b/p[2]', 'Allergies'), self.change(payload, 'b/p[2]', 'Allergies', kind='long_text'),
                   self.change(payload, 'b/p[3]', 'Yes', kind='checkbox', glyph=0),
                   self.change(payload, 'b/p[3]', 'No', kind='checkbox', glyph=1),
                   [-1, 0, 0, 399999, 0, 1, -1]]
        outcomes = set()
        # The original never changes across the 720 orders: its role report is read in the isolated parser once per
        # distinct mapping and reused (the reading itself is exercised by every other test).
        reports, real = {}, destination_roles.report
        def report(version, fields=None):
            key = mapping.canonical(fields if fields is not None else version.fields)
            if key not in reports:reports[key] = real(version, fields)
            return copy.deepcopy(reports[key])
        with patch.object(destination_roles, 'report', side_effect=report):
            for order in itertools.permutations(changes):
                body = copy.deepcopy(base); body['x'] = [list(c) for c in order]
                result = mapping.validate_proposal(version, compiled, body)
                outcomes.add(mapping.canonical({'fields': result['fields'], 'changes': result['review']['changes'],
                                                'unresolved': result['review']['unresolved'],
                                                'rejected': result['review']['rejected_changes']}))
        self.assertEqual(len(outcomes), 1, 'the order of the changes decided the outcome')
        result = json.loads(outcomes.pop())
        self.assertEqual(result['changes'], [{'key': pulse['key'], 'replacement': pulse['key']}])
        boxes = [r for r in result['rejected'] if r['slot'] == BODY + 'b/p[3]']
        self.assertEqual(len(boxes), 2)
        self.assertEqual(sorted(boxes[0]['with']), [boxes[1]['id']])   # RF13: one printed option group, named together
        allergies = [r for r in result['rejected'] if r['slot'] == BODY + 'b/p[2]']
        self.assertEqual([len(r['with']) for r in allergies], [1, 1])

    def test_dependent_changes_are_held_naming_each_other_and_an_independent_one_is_adopted(self):
        """D6 (the same question), D3 (an addition landing on a field another change replaces): every member is held
        with its reason; an independent relabel beside them is still adopted."""
        data = lines('Pulse findings: ______', 'Pulse findings (left): ______', 'Plan: ______')
        template, version = self.version_of(data, 'Synthetic dependence')
        compiled, payload = self.compiled(version)
        first, second, plan = [next(f for f in version.fields if f['label'] == label)
                               for label in ('Pulse findings', 'Pulse findings (left)', 'Plan')]
        fid = {f['key']: n for n, f in enumerate(payload['fields'])}
        def body_with(changes, replaced):
            body = self.keep_all(payload)
            body['f']['k'] = [n for n in body['f']['k'] if n not in replaced]; body['f']['r'] = sorted(replaced)
            body['x'] = changes
            return body
        again = self.change(payload, 'b/p[2]', 'Pulse findings', fid=fid[second['key']])
        # a relabel that asks a KEPT field's question: held, depending on that field
        result = mapping.validate_proposal(version, compiled, body_with([again], {fid[second['key']]}))
        self.assertEqual(result['review']['changes'], [])
        self.assertIn(mapping.QUESTION_REASON.format(label='Pulse findings'),
                      ' '.join(r['reason'] for r in result['review']['unresolved']))
        # two relabels asking one question: both held, each naming the other; the independent Plan relabel adopted
        own = self.change(payload, 'b/p[1]', 'Pulse findings', fid=fid[first['key']])
        independent = self.change(payload, 'b/p[3]', 'Plan', fid=fid[plan['key']], kind='long_text')
        result = mapping.validate_proposal(version, compiled, body_with(
            [own, again, independent], {fid[first['key']], fid[second['key']], fid[plan['key']]}))
        self.assertEqual(result['review']['changes'], [{'key': plan['key'], 'replacement': plan['key']}])
        held = {r['key']: r for r in result['review']['rejected_changes']}
        self.assertEqual(held[first['key']]['with'], [held[second['key']]['id']])
        self.assertEqual(held[second['key']]['with'], [held[first['key']]['id']])
        # D3: an addition on the place another change replaces holds that replacement too
        addition = self.change(payload, 'b/p[1]', 'Pulse findings', kind='long_text')
        result = mapping.validate_proposal(version, compiled, body_with([own, addition], {fid[first['key']]}))
        self.assertEqual(result['review']['changes'], [])
        self.assertTrue(all(r['with'] for r in result['review']['rejected_changes']))

    def test_rf14_a_replacement_keeps_its_place(self):
        data = lines('Pulse findings: ______', 'Plan: ______')
        template, version = self.version_of(data, 'Synthetic RF14')
        compiled, payload = self.compiled(version)
        pulse = next(f for f in version.fields if f['label'] == 'Pulse findings')
        fid = next(n for n, f in enumerate(payload['fields']) if f['key'] == pulse['key'])
        body = self.keep_all(payload); body['f']['k'].remove(fid); body['f']['r'] = [fid]
        body['x'] = [self.change(payload, 'b/p[1]', 'Pulse findings', fid=fid, kind='long_text')]
        result = mapping.validate_proposal(version, compiled, body)
        field = next(f for f in result['fields'] if f['key'] == pulse['key'])
        self.assertEqual((field['kind'], field['locator']), ('long_text', pulse['locator']))
        out = private_path('exports/rf14.docx'); out.parent.mkdir(parents=True, exist_ok=True)
        fill_template(private_path(template.original_path), result['fields'], {pulse['key']: 'weak radial pulse'}, out)
        self.assertIn('Pulse findings: weak radial pulse', part_text(out.read_bytes()))
        self.assertNotIn('______', part_text(out.read_bytes()).split('Pulse findings: weak radial pulse')[0][-10:])

    def test_a_replacement_needs_an_automatic_ordinary_field_with_no_role(self):
        data = lines('Signature: ______________', 'Plan: ______')
        template, version = self.version_of(data, 'Synthetic replacement rules')
        signature = next(f for f in version.fields if f['kind'] == 'signature')
        plan = next(f for f in version.fields if f['label'] == 'Plan')
        manual = {**plan, 'origin': 'manual'}
        # (labelled 'Answer': a label the identity check reads as the finalizer's is refused earlier, by NOT_ORDINARY)
        automatic = {k: v for k, v in forged(data, 'b/p[1]', label='Answer').items() if k != 'origin'}
        cases = (('signature kind', [signature, plan], signature, 'Signature', 'text', mapping.SIGNING_REASON),
                 ('manual field', [signature, manual], manual, 'Plan', 'text', mapping.NOT_AUTOMATIC),
                 ('kind to checkbox', [signature, plan], plan, 'Plan', 'checkbox', mapping.KIND_CHANGE),
                 ('an automatic-looking field on a role place', [plan, automatic], automatic, 'Signature', 'text',
                  mapping.ROLE_PLACE))
        for n, (name, fields, prior, label, kind, reason) in enumerate(cases):
            with self.subTest(case=name):
                stored = TemplateVersion.objects.create(template=template, number=n + 2, fields=fields, contract='v3',
                                                        created_by=self.user, **import_metadata(version, fields))
                prior = prior or fields[1]
                compiled, payload = self.compiled(stored)
                fid = next(i for i, f in enumerate(payload['fields']) if f['key'] == prior['key'])
                body = self.keep_all(payload); body['f']['k'].remove(fid); body['f']['r'] = [fid]
                slot = prior['locator']['slot']
                body['x'] = [self.change(payload, slot, label, kind=kind, fid=fid)]
                result = mapping.validate_proposal(stored, compiled, body)
                self.assertEqual(result['fields'], stored.fields)
                self.assertIn(reason, ' '.join(row['reason'] for row in result['review']['unresolved']))

    def test_rf15_resolve_applies_removals_first_and_revalidates(self):
        data = lines('Allergies: ______', 'Plan: ______')
        template, version = self.version_of(data, 'Synthetic RF15')
        allergies = next(f for f in version.fields if f['label'] == 'Allergies')
        slot = BODY + allergies['locator']['slot']
        assisted = TemplateVersion.objects.create(template=template, number=2, fields=version.fields, contract='v3',
            **import_metadata(version, version.fields),
            created_by=self.user, assistance={'unresolved': [
                {'slot': slot, 'reason': 'Proposed addition', 'key': '', 'text': ''},
                {'slot': slot, 'reason': 'Existing field requires clinician review', 'key': allergies['key'], 'text': ''}]})
        url = f'/templates/{template.pk}/assistance/'
        # 'mapped' for an area whose only field the same submission removes: refused
        response = self.client.post(url, {'action': 'resolve', 'mapping_version': assisted.pk,
                                          'resolution_0': 'mapped', 'resolution_1': 'not_answer'})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.latest(template).pk, assisted.pk)
        # a key-less 'not an answer' on an area a field still writes to: refused
        response = self.client.post(url, {'action': 'resolve', 'mapping_version': assisted.pk,
                                          'resolution_0': 'not_answer', 'resolution_1': 'mapped'})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.latest(template).pk, assisted.pk)
        # consistent decisions are saved
        response = self.client.post(url, {'action': 'resolve', 'mapping_version': assisted.pk,
                                          'resolution_0': 'mapped', 'resolution_1': 'mapped'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.latest(template).assistance['unresolved'], [])


# ─── refutation round 1 (e555540d): each test is red on e555540d and names the finding it closes ─────────────────
def framed(h1, inside, prompt=None):
    doc = Document()
    if h1:
        doc.add_heading(h1, level=1)
    if prompt:
        doc.add_paragraph(prompt)
    else:
        doc.add_heading('Clinical details', level=2)
    cell = box(doc).cell(0, 0)
    for n, text in enumerate(inside):
        (cell.paragraphs[0] if n == 0 else cell.add_paragraph()).add_run(text)
    doc.add_heading('Management', level=2); doc.add_paragraph('Plan: ______')
    return saved(doc)


class Walk:
    """generate -> review every row -> finalize, on the fixture generator (no provider)."""

    def generated(self, version, notes):
        self.assertEqual(self.new_encounter(version, notes=notes).status_code, 302)
        e = Encounter.objects.latest('created_at')
        enqueue(Encounter.objects.get(pk=e.pk), self.user); self.assertTrue(run_one())
        e.refresh_from_db()
        return e

    def reviewed(self, e, version, typed=(), skip_signing=False):
        post = {'revision': e.revision}
        for f in version.fields:
            if skip_signing and f.get('kind') in v3_mapping.SIGNING_KINDS:
                continue
            post['answer_' + f['key']] = dict(typed).get(f['label'], e.answers.get(f['key'], {}).get('text', ''))
            post['reviewed_' + f['key']] = 'on'
        response = self.client.post(f'/encounters/{e.pk}/save/', post)
        e.refresh_from_db()
        return response

    def finalized(self, e):
        response = self.client.post(f'/encounters/{e.pk}/finalize/', {'revision': e.revision, 'reviewed': 'on'}, follow=True)
        e.refresh_from_db()
        return response


@override_settings(**TEST_SETTINGS)
class LegacyRound1(Walk, V3Fixture, TestCase):
    def test_l1f1_l1f5_signing_rows_are_never_generated_or_exported(self):
        """L1:F1 (HIGH) and L1:F5: e555540d generated 'Dr Example' / 'JE' into these legacy rows and exported them."""
        for label, value in (('Signed (Consultant)', 'Dr Example'), ('Signed (print name)', 'Dr Example'),
                             ('Electronically signed (print name)', 'Dr Example'), ('Signed (GP)', 'Dr Example'),
                             ('Signed (patient)', 'John Example'), ('Signed for and on behalf of', 'Dr Example'),
                             ('Countersigned by consultant', 'Dr Example'), ('Signature and date', 'Dr Example'),
                             ('Initial here', 'JE'), ('Initial each page', 'JE')):
            with self.subTest(label=label):
                template, version = self.uploaded(data=legacy_rows(label, 'Plan'), name='Synthetic L1 ' + label)
                by = self.by_label(version)
                e = self.generated(version, f'{label}: {value}\nPlan: Rest.')
                answer = e.answers[by[label]['key']]
                self.assertEqual((answer['status'], answer['text']), ('missing', ''))
                self.assertEqual(self.reviewed(e, version).status_code, 302)
                self.finalized(e)
                self.assertEqual(e.status, 'finalized')
                self.assertNotIn(value, part_text(private_path(e.export_path).read_bytes()))

    def test_r2_06_a_numbered_initial_reading_is_generated(self):
        """R2-06: 'Initial 12-lead ECG' was held as an initials row and never generated."""
        template, version = self.uploaded(data=legacy_rows('Initial 12-lead ECG', 'Plan'), name='Synthetic R2-06')
        e = self.generated(version, 'Initial 12-lead ECG: Sinus rhythm\nPlan: Rest.')
        answer = e.answers[self.by_label(version)['Initial 12-lead ECG']['key']]
        self.assertEqual((answer['status'], answer['text']), ('supported', 'Sinus rhythm'))

    def test_r2_07_unknown_parenthetical_signer_is_human_owned_not_the_finalizer(self):
        """October 3 supersedes the old assumption that any name qualifier identifies the finalizer."""
        for label in ('Electronically signed by (print name)', 'Signed by (name)', 'Signature (consultant)'):
            with self.subTest(label=label):
                template, version = self.uploaded(data=legacy_rows(label, 'Plan'), name='Synthetic R2-07 ' + label)
                key = self.by_label(version)[label]['key']
                e = self.generated(version, label + ': Dr Somebody Else\nPlan: Rest.')
                self.assertEqual((e.answers[key]['status'], e.answers[key]['text']), ('missing', ''))
                self.assertEqual(self.reviewed(e, version).status_code, 302)
                self.finalized(e)
                self.assertEqual(e.status, 'finalized')
                self.assertEqual(Document(private_path(e.export_path)).tables[0].cell(0, 1).text, '')

    def test_r2_08_an_established_clinician_signature_cannot_gain_unknown_authority(self):
        for n, label in enumerate(('Signature (consultant)', 'Signature (print name)')):
            with self.subTest(label=label):
                template, version = self.uploaded(data=legacy_rows('Signature', 'Plan'), name='Synthetic R2-08 %d' % n)
                by = self.by_label(version)
                response = self.client.post(f'/templates/{template.pk}/', {
                    'mapping_version': version.pk,
                    'labels': json.dumps({by['Signature']['key']: label, by['Plan']['key']: 'Plan'})})
                self.assertContains(response, 'A signature question keeps a signature label')
                self.assertEqual(self.latest(template).pk, version.pk)

    def test_another_persons_or_unknown_parenthetical_is_not_the_finalizers(self):
        from core.tab_review import another_signers_line, clinician_identity
        for label in ('Signed (patient)', 'Signed by (witness)', 'Signature (guardian)',
                      'Signature (participant)', 'Signature (mother)', 'Signature (father)', 'Signature (donor)',
                      'Signed (Consultant)', 'Signature (GP)', 'Signed (print name)', 'Signature (registrar)'):
            with self.subTest(label=label):
                self.assertFalse(clinician_identity({'label': label}))
                self.assertTrue(another_signers_line({'label': label}))


@override_settings(**V3)
class V3Round1(Walk, V3Fixture, TestCase):
    def confirm_status(self, template, version):
        return self.client.post(f'/templates/{template.pk}/', {'action': 'v3_confirm', 'mapping_version': version.pk,
                                                               'checked_every_field': 'on'}).status_code

    def test_r2_01_a_frame_holding_ordinary_questions_is_confirmed_generated_and_exported(self):
        """R2-01 (HIGH): the frame under 'Initial referral' was a signing box; confirmation listed all three fields."""
        template, version = self.uploaded(data=framed('Initial referral', ('Presenting complaint: ______',
                                          'Current medication: ______', 'Allergies: ______')), name='Synthetic R2-01')
        self.assertEqual(self.confirm_status(template, version), 302)
        e = self.generated(version, 'Presenting complaint: Chest pain\nCurrent medication: Aspirin\n'
                                    'Allergies: None known\nPlan: Rest.')
        self.assertEqual(self.reviewed(e, version).status_code, 302)
        self.finalized(e)
        self.assertEqual(e.status, 'finalized')
        self.assertIn('Presenting complaint: Chest pain', part_text(private_path(e.export_path).read_bytes()))

    def test_r2_02_a_printed_name_beside_a_signature_is_confirmed(self):
        for n, text in enumerate(('Printed name:    Signature: ______', 'Consultant:    Signature: ______    Date: ______')):
            with self.subTest(text=text):
                template, version = self.uploaded(data=lines(text, 'Plan: ______'), name='Synthetic R2-02 %d' % n)
                self.assertEqual(self.confirm_status(template, version), 302)

    def test_r2_03_the_clinician_may_answer_an_initial_presentation_box(self):
        doc = Document(); doc.add_heading('History', level=2); doc.add_paragraph('Describe the initial presentation?')
        box(doc); doc.add_paragraph('Plan: ______')
        template, version = self.uploaded(data=saved(doc), name='Synthetic R2-03')
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_add', 'mapping_version': version.pk,
                                    'slot': BODY + 'b/tbl[0]/tr[0]/tc[0]', 'kind': 'long_text',
                                    'label': 'Describe the initial presentation?'})
        self.assertEqual(response.status_code, 302)

    def test_r2_04_an_initial_reading_can_be_replaced_by_its_text_answer(self):
        template, version = self.uploaded(data=lines('Initial GCS: ______', 'Plan: ______'), name='Synthetic R2-04')
        gcs = self.by_label(version)['Initial GCS']
        self.assertEqual(gcs['kind'], 'text')  # ordinary initial clinical measurement is now recovered
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_save', 'mapping_version': version.pk,
                                                                  'remove_' + gcs['key']: 'on'})
        self.assertEqual(response.status_code, 302)
        latest = self.latest(template)
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_add', 'mapping_version': latest.pk,
                                    'slot': BODY + 'b/p[1]', 'kind': 'text', 'label': 'Initial GCS'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.confirm_status(template, self.latest(template)), 302)
        reviewed = self.latest(template)
        e = self.generated(reviewed, 'Initial GCS: 13\nPlan: Rest.')
        self.assertEqual(e.answers[self.by_label(reviewed)['Initial GCS']['key']]['text'], '13')
        self.assertEqual(self.reviewed(e, reviewed).status_code, 302)
        self.finalized(e); self.assertEqual(e.status, 'finalized')
        self.assertIn('13', part_text(private_path(e.export_path).read_bytes()))

    def test_r2_05_an_unlocked_control_inside_a_locked_frame_is_confirmed(self):
        doc = Document(); doc.add_heading('Clinic record', level=2)
        frame = OxmlElement('w:sdt'); pr = OxmlElement('w:sdtPr'); lock = OxmlElement('w:lock')
        lock.set(qn('w:val'), 'contentLocked'); pr.append(lock); pr.append(OxmlElement('w:richText')); frame.append(pr)
        content = OxmlElement('w:sdtContent'); frame.append(content)
        p = doc.add_paragraph('Hospital number: ')
        sdt = OxmlElement('w:sdt'); spr = OxmlElement('w:sdtPr'); alias = OxmlElement('w:alias')
        alias.set(qn('w:val'), 'Hospital number'); spr.append(alias); spr.append(OxmlElement('w:showingPlcHdr'))
        spr.append(OxmlElement('w:text')); sdt.append(spr); body = OxmlElement('w:sdtContent')
        run = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = 'Click or tap here to enter text.'; run.append(t)
        body.append(run); sdt.append(body); p._p.append(sdt); content.append(p._p)
        doc.element.body.find(qn('w:sectPr')).addprevious(frame)
        doc.add_paragraph('Plan: ______')
        template, version = self.uploaded(data=saved(doc), name='Synthetic R2-05')
        self.assertEqual(self.confirm_status(template, version), 302)

    def test_l1f2_a_signature_and_date_line_is_never_generated(self):
        """L1:F2 (HIGH): 'Signature and date: ____' was a date field; generation wrote 'Dr Example 30/09/2026'."""
        template, version = self.uploaded(data=lines('Signature and date: ______________', 'Plan: ______'),
                                          name='Synthetic L1F2')
        self.confirm(template, version)
        e = self.generated(version, 'Signature and date: Dr Example 30/09/2026\nPlan: Rest.')
        field = self.by_label(version)['Signature and date']
        self.assertEqual(field['kind'], 'signature')
        self.assertEqual((e.answers[field['key']]['status'], e.answers[field['key']]['text']), ('missing', ''))

    def test_l1f6_the_line_under_a_signature_caption_is_listed_at_confirmation(self):
        """L1:F6(a): the engine's 'Answer' field on the writing line under 'Consultant signature:' sits on the signing
        place; confirmation lists it, and after its removal the mapping confirms."""
        doc = Document(); doc.add_heading('Clinic record', level=2); doc.add_paragraph('Consultant signature:')
        doc.add_paragraph('____________________'); doc.add_paragraph('Plan: ______')
        template, version = self.uploaded(data=saved(doc), name='Synthetic L1F6')
        answer = next(f for f in version.fields if f['locator']['slot'] == 'b/p[2]')
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_confirm', 'mapping_version': version.pk,
                                                                   'checked_every_field': 'on'})
        self.assertEqual(response.status_code, 400)
        self.assertIn(answer['label'], response.context['v3_error'])
        self.client.post(f'/templates/{template.pk}/', {'action': 'v3_save', 'mapping_version': version.pk,
                                                        'remove_' + answer['key']: 'on'})
        self.assertEqual(self.confirm_status(template, self.latest(template)), 302)

    def test_l1f7_declared_residual_a_typed_initial_in_a_log_tables_initials_column_exports(self):
        """L1:F7 (pre-existing on 43fcf214) is a declared residual again: a log table's column carries no role (refutation
        of 258650f2, L1:R1-03: a column read from its printed header refused ordinary clinical columns). Only the
        clinician's typed rows reach it; table_rows is never generated."""
        doc = Document(); doc.add_heading('Medication log', level=2)
        t = doc.add_table(rows=3, cols=3); t.style = 'Table Grid'
        for c, h in enumerate(('Drug', 'Dose', 'Nurse initials')):
            t.cell(0, c).text = h
        doc.add_paragraph('Plan: ______')
        template, version = self.uploaded(data=saved(doc), name='Synthetic L1F7')
        self.confirm(template, version)
        e = self.generated(version, 'Plan: Rest.')
        log = next(f for f in version.fields if f['kind'] == 'table_rows')
        response = self.reviewed(e, version, typed={log['label']: 'Paracetamol | 1 g | JE'}, skip_signing=True)
        self.assertEqual(response.status_code, 302)
        self.finalized(e)
        self.assertEqual(e.status, 'finalized')
        self.assertIn('JE', part_text(private_path(e.export_path).read_bytes()))


@override_settings(**MAPPING_SETTINGS)
class ModelProposalsRound1(V3Fixture, TestCase):
    def setUp(self):
        super().setUp()
        self.config = generation.configuration()

    def proposal(self, version, replaced=(), uncertain=(), dismissed=(), changes=()):
        compiled = mapping.compile_request(version, self.config)
        payload = json.loads(compiled['request']['messages'][0]['content'])
        body = mapping.wire.fixture(payload)
        fids = {f['key']: n for n, f in enumerate(payload['fields'])}
        moved = {fids[k] for k in list(replaced) + list(uncertain)}
        body['f']['k'] = [n for n in body['f']['k'] if n not in moved]
        body['f']['r'] = sorted(fids[k] for k in replaced); body['f']['u'] = sorted(fids[k] for k in uncertain)
        sids = {row['slot']: n for n, row in enumerate(payload['slots'])}
        for slot in dismissed:
            for verdict in ('c', 'u', 'n'):
                body['s'][verdict] = [n for n in body['s'][verdict] if n != sids[slot]]
            body['s']['n'] = sorted(body['s']['n'] + [sids[slot]])
        body['x'] = [change(payload, fids, sids) for change in changes]
        return mapping.validate_proposal(version, compiled, body)

    @staticmethod
    def relabel(key, slot, kind, start, end, role=0):
        def build(payload, fids, sids):
            sid = sids[BODY + slot]
            return [fids[key], sid, mapping.wire.KINDS.index(kind), payload['text_refs'].index([sid, role, 0]), start, end, -1]
        return build

    def test_l3f1_a_replacement_label_is_printed_wording_not_spaces_or_a_blank_line(self):
        """L3:F1: ' ', ':', ': ' and '______' were adopted as the label (question_key(' ') == ''), unasking the question."""
        template, version = self.uploaded(data=lines('Pulse findings: ______', 'Plan: ______'), name='Synthetic L3F1')
        pulse = self.by_label(version)['Pulse findings']
        for name, (start, end) in (('space', (15, 16)), ('colon', (14, 15)), ('colon and space', (14, 16)),
                                   ('blank line', (16, 22))):
            with self.subTest(name=name):
                result = self.proposal(version, replaced=[pulse['key']],
                                       changes=[self.relabel(pulse['key'], 'b/p[1]', pulse['kind'], start, end)])
                self.assertEqual(result['review']['changes'], [])
                self.assertEqual(next(f for f in result['fields'] if f['key'] == pulse['key']), pulse)
                self.assertIn(mapping.LABEL_SHAPE, ' '.join(r['reason'] for r in result['review']['unresolved']))
        result = self.proposal(version, replaced=[pulse['key']],
                               changes=[self.relabel(pulse['key'], 'b/p[1]', pulse['kind'], 0, 15)])
        self.assertEqual(next(f for f in result['fields'] if f['key'] == pulse['key'])['label'], 'Pulse findings')

    def test_l3f2_the_finalizers_signature_label_is_never_relabelled(self):
        """L3:F2: 'Signature/Date' (identity) relabelled to 'Date' was adopted; generation then filled it."""
        template, version = self.uploaded(data=lines('Signature/Date: ______', 'Plan: ______'), name='Synthetic L3F2')
        signed = self.by_label(version)['Signature/Date']
        fields = [{**f, 'kind': 'date'} if f['key'] == signed['key'] else f for f in version.fields]
        stored = TemplateVersion.objects.create(template=template, number=2, contract='v3', created_by=self.user,
            fields=fields, **import_metadata(version, fields))  # before this role repair; the signing kind stays wrong
        at = 'Signature/Date: ______'.index('Date')
        result = self.proposal(stored, replaced=[signed['key']],
                               changes=[self.relabel(signed['key'], 'b/p[1]', 'date', at, at + 4)])
        self.assertEqual(result['review']['changes'], [])
        self.assertIn(mapping.NOT_ORDINARY, ' '.join(r['reason'] for r in result['review']['unresolved']))

    def test_l3f2_an_identity_label_on_a_place_with_no_role_is_never_relabelled(self):
        """Where no role holds the place, only the identity check does: a stored automatic field labelled as the
        finalizer's signature on the 'Pulse findings' line is never relabelled away from it by a model."""
        template, version = self.uploaded(data=lines('Pulse findings: ______', 'Plan: ______'), name='Synthetic L3F2b')
        pulse = self.by_label(version)['Pulse findings']
        fields = [{**f, 'label': 'Signed by'} if f['key'] == pulse['key'] else f for f in version.fields]
        stored = TemplateVersion.objects.create(template=template, number=2, contract='v3', created_by=self.user,
            fields=fields, **import_metadata(version, fields))
        result = self.proposal(stored, replaced=[pulse['key']],
                               changes=[self.relabel(pulse['key'], 'b/p[1]', pulse['kind'], 0, 14)])
        self.assertEqual(result['review']['changes'], [])
        self.assertEqual(next(f for f in result['fields'] if f['key'] == pulse['key'])['label'], 'Signed by')

    def test_l3f3_two_answers_on_one_line_are_independent(self):
        """L3:F3: 'Weight: ____ kg  Height: ____ cm' - dependence compared paragraphs, so a Height relabel was held for a
        Weight verdict and two disjoint relabels held each other."""
        text = 'Weight: ______ kg  Height: ______ cm'
        template, version = self.uploaded(data=lines(text, 'Plan: ______'), name='Synthetic L3F3')
        weight, height = version.fields[0], version.fields[1]
        at = text.index('Height')
        height_change = self.relabel(height['key'], 'b/p[1]', height['kind'], at, at + 6)
        weight_change = self.relabel(weight['key'], 'b/p[1]', weight['kind'], 0, 6)
        result = self.proposal(version, replaced=[height['key']], uncertain=[weight['key']], changes=[height_change])
        self.assertEqual(result['review']['changes'], [{'key': height['key'], 'replacement': height['key']}])
        adopted = next(f for f in result['fields'] if f['key'] == height['key'])
        self.assertEqual((adopted['label'], adopted['locator']), ('Height', height['locator']))
        result = self.proposal(version, replaced=[weight['key'], height['key']], changes=[height_change, weight_change])
        self.assertEqual(sorted(c['key'] for c in result['review']['changes']), sorted([weight['key'], height['key']]))

    def test_l3f4_a_change_inside_a_cell_the_proposal_dismisses_is_held(self):
        """L3:F4: the proposal marked the date cell not an answer, and a relabel naming the cell's paragraph was adopted."""
        doc = Document(); doc.add_heading('Section 3: Public protection', level=2)
        t = doc.add_table(rows=2, cols=2); t.style = 'Table Grid'
        t.cell(0, 0).text = 'Name of the lead agency'; t.cell(1, 0).text = 'Date the lead agency was informed'
        doc.add_paragraph('Plan: ______')
        template, version = self.uploaded(data=saved(doc), name='Synthetic L3F4')
        date = next(f for f in version.fields if f['kind'] == 'date')
        cell = date['locator']['slot']
        result = self.proposal(version, replaced=[date['key']], dismissed=[BODY + cell],
                               changes=[self.relabel(date['key'], cell + '/p[0]', 'date', 0, 20, role=2)])
        self.assertEqual(result['review']['changes'], [])
        self.assertIn('not an answer', ' '.join(r['reason'] for r in result['review']['unresolved']))
        # The printed line INSIDE a covered cell called 'not an answer' is not the cell (the retained evidence/002
        # response does exactly this): a relabel of the cell's field is still adopted.
        result = self.proposal(version, replaced=[date['key']], dismissed=[BODY + cell + '/p[0]'],
                               changes=[self.relabel(date['key'], cell, 'date', 0, 20, role=2)])
        self.assertEqual(result['review']['changes'], [{'key': date['key'], 'replacement': date['key']}])


# ─── salvage of a32243eb (b2): the two regressions vs 43fcf214 and the legacy usefulness narrowing ─────────────────
OTHER_SIGNERS_DATED = (('Signature (participant, dd/mm/yyyy)', 'John Example 30/09/2026'),
                       ('Signed (mother, DOB)', 'Jane Example 01/01/1990'))


@override_settings(**TEST_SETTINGS)
class SalvageLegacy(Walk, V3Fixture, TestCase):
    def test_rr1_10_another_named_signers_dated_line_is_never_generated(self):
        """RR1-10 (MEDIUM, regression vs 43fcf214): a32243eb generated and exported these; 43fcf214 held them (identity).
        The patient's and the witness's dated lines are held too (another_signers_line), and there a clinician's typed
        value finalizes and exports. The participant's and the mother's are identity rows again (L1:R1-01)."""
        for n, (label, value) in enumerate((('Signature (patient, dd/mm/yyyy)', 'John Example 30/09/2026'),
                                            ('Signed by (witness, DOB)', 'Jane Example 01/01/1990'))):
            with self.subTest(label=label):
                template, version = self.uploaded(data=legacy_rows(label, 'Plan'), name='Synthetic RR1-10 %d' % n)
                field = self.by_label(version)[label]
                e = self.generated(version, f'{label}: {value}\nPlan: Rest.')
                answer = e.answers[field['key']]
                self.assertEqual((answer['status'], answer['text']), ('missing', ''))
                self.assertEqual(answer['reason'], destination_roles.ROLE_NOT_GENERATED)
                self.assertEqual(self.reviewed(e, version).status_code, 302)
                self.finalized(e)
                self.assertEqual(e.status, 'finalized')
                self.assertNotIn(value, part_text(private_path(e.export_path).read_bytes()))
                typed = self.generated(version, 'Plan: Rest.')
                self.assertEqual(self.reviewed(typed, version, typed={label: value}).status_code, 302)
                self.finalized(typed)
                self.assertEqual(typed.status, 'finalized')
                self.assertIn(value, part_text(private_path(typed.export_path).read_bytes()))

    def test_pc07_a_dash_qualified_initial_reading_is_generated_as_on_43fcf214(self):
        """PC07 / PC18 usefulness: a32243eb held 'Troponin — Initial' (never generated); 43fcf214 generated it."""
        template, version = self.uploaded(data=legacy_rows('Troponin — Initial', 'Initial', 'Plan'),
                                          name='Synthetic PC07 dash')
        by = self.by_label(version)
        e = self.generated(version, 'Troponin — Initial: 12 ng/L\nInitial: JE\nPlan: Rest.')
        troponin, initial = e.answers[by['Troponin — Initial']['key']], e.answers[by['Initial']['key']]
        self.assertEqual((troponin['status'], troponin['text']), ('supported', '12 ng/L'))
        self.assertEqual((initial['status'], initial['text']), ('missing', ''))    # a bare 'Initial' stays a place
        self.assertEqual(self.reviewed(e, version).status_code, 302)
        self.finalized(e)
        self.assertEqual(e.status, 'finalized')
        exported = part_text(private_path(e.export_path).read_bytes())
        self.assertIn('12 ng/L', exported)
        self.assertNotIn('JE', exported)

    def test_rr1_10_the_legacy_request_does_not_ask_it(self):
        """Layer 1 (core.destination_roles.legacy_reason): the row is never asked; its label is a context label."""
        template, version = self.uploaded(data=legacy_rows('Signed by (witness, DOB)', 'Plan'), name='Synthetic RR1-10 ask')
        asked = {f['label']: f for f in destination_roles.generation_fields(version)}
        self.assertEqual(asked['Signed by (witness, DOB)']['destination_role'], {'reason': destination_roles.ROLE_NOT_GENERATED})
        self.assertNotIn('destination_role', asked['Plan'])
        questions, context, _ = generation.v3_generation_fields(list(asked.values()))
        self.assertEqual([f['label'] for f in questions], ['Plan'])
        self.assertIn('Signed by (witness, DOB)', context)

    def test_rr1_10_the_validator_holds_it_on_any_path(self):
        """Layer 2 (the generation validator, where 43fcf214 held identity): a field that reaches generation unmarked
        (a caller that did not prepare it) is still answered 'missing' with the reason."""
        template, version = self.uploaded(data=legacy_rows('Signature (patient, dd/mm/yyyy)', 'Plan'), name='Synthetic RR1-10 v')
        fields = [dict(f) for f in version.fields]
        answers = generation.generate('Signature (patient, dd/mm/yyyy): John Example 30/09/2026\nPlan: Rest.', fields,
                                      generation.configuration(fields))[0]
        by = self.by_label(version)
        self.assertEqual((answers[by['Signature (patient, dd/mm/yyyy)']['key']]['status'],
                          answers[by['Signature (patient, dd/mm/yyyy)']['key']]['reason']),
                         ('missing', destination_roles.ROLE_NOT_GENERATED))
        self.assertEqual(answers[by['Plan']['key']]['status'], 'supported')

    def test_rr1_10_other_and_unknown_signers_keep_human_ownership(self):
        from core.tab_review import another_signers_line, clinician_identity, signature_label
        for label in ('Signature (patient, dd/mm/yyyy)', 'Signed by (witness, DOB)', 'Signature (patient)',
                      'Signed by (witness)', 'Signature (participant, dd/mm/yyyy)', 'Signed (mother, DOB)',
                      'Signature (father, born)', 'Signed (Consultant)', 'Signature (GP)',
                      'Signature (patient, date)', 'Patient signature'):
            with self.subTest(label=label):
                self.assertFalse(clinician_identity({'label': label}))
                self.assertTrue(another_signers_line({'label': label}))
        for label in ('Signature date', 'Date of signature', 'Consent obtained by (name and grade)', 'Babinski sign'):
            with self.subTest(label=label):
                self.assertFalse(another_signers_line({'label': label}))
        self.assertFalse(signature_label('Signed (Consultant)'))


@override_settings(**V3)
class SalvageV3(Walk, V3Fixture, TestCase):
    def test_rr1_10_another_named_signers_dated_line_is_never_generated_on_v3(self):
        for n, (label, value) in enumerate(OTHER_SIGNERS_DATED):
            with self.subTest(label=label):
                template, version = self.uploaded(data=lines(label + ': ______________', 'Plan: ______'),
                                                  name='Synthetic RR1-10 v3 %d' % n)
                field = self.by_label(version)[label]
                self.confirm(template, version)          # not a role place: confirmable, as on 43fcf214
                e = self.generated(version, f'{label}: {value}\nPlan: Rest.')
                answer = e.answers[field['key']]
                self.assertEqual((answer['status'], answer['text']), ('missing', ''))
                self.assertEqual(answer['reason'], destination_roles.ROLE_NOT_GENERATED)
                self.assertEqual(self.reviewed(e, version).status_code, 302)
                self.finalized(e)
                self.assertEqual(e.status, 'finalized')
                self.assertNotIn(value, part_text(private_path(e.export_path).read_bytes()))


def control_line(before, alias, after):
    """'<before><inline plain-text content control><after>' as one paragraph."""
    doc = Document(); doc.add_heading('Clinic record', level=2)
    p = doc.add_paragraph(before)
    sdt = OxmlElement('w:sdt'); pr = OxmlElement('w:sdtPr'); a = OxmlElement('w:alias'); a.set(qn('w:val'), alias)
    pr.append(a); pr.append(OxmlElement('w:showingPlcHdr')); pr.append(OxmlElement('w:text')); sdt.append(pr)
    content = OxmlElement('w:sdtContent'); r = OxmlElement('w:r'); t = OxmlElement('w:t')
    t.text = 'Click or tap here to enter text.'; r.append(t); content.append(r); sdt.append(content); p._p.append(sdt)
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = after; r.append(t); p._p.append(r)
    doc.add_paragraph('Plan: ______')
    return saved(doc)


@override_settings(**MAPPING_SETTINGS)
class SalvageResolveOverlap(V3Fixture, TestCase):
    """RR3-03 (MEDIUM, regression vs 43fcf214): validate_proposal V1 excused the engine's own pre-existing overlap and
    adopted a relabel; the resolve step's V1 did not, so the version could never be confirmed. One rule now
    (mapping_assistance.model_overlaps)."""

    def setUp(self):
        super().setUp()
        self.config = generation.configuration()

    def assisted(self, template, version, replace, text):
        compiled = mapping.compile_request(version, self.config)
        payload = json.loads(compiled['request']['messages'][0]['content'])
        body = mapping.wire.fixture(payload)
        fid = next(i for i, f in enumerate(payload['fields']) if f['key'] == replace['key'])
        body['f']['k'] = [n for n in body['f']['k'] if n != fid]; body['f']['r'] = [fid]
        body['x'] = [ModelProposals.change(payload, replace['locator']['slot'], text, kind=replace['kind'], fid=fid)]
        proposal = mapping.validate_proposal(version, compiled, body)
        new = TemplateVersion.objects.create(template=template, number=version.number + 1, fields=proposal['fields'],
            **import_metadata(version, proposal['fields']),
            contract='v3', review_requirements=version.review_requirements, created_by=self.user,
            assistance={**proposal['review'], 'job_id': 'synthetic', 'base_version': str(version.pk), 'simulated': True})
        return proposal, new

    def resolve(self, template, new):
        mapped = v3_mapping.mapped_slots(new.fields)
        post = {'action': 'resolve', 'mapping_version': new.pk}
        for n, row in enumerate(new.assistance['unresolved']):
            post['resolution_%d' % n] = 'mapped' if row.get('key') or v3_mapping.is_mapped(row['slot'], mapped) \
                else 'not_answer'
        return self.client.post(f'/templates/{template.pk}/assistance/', post)

    def test_rr3_03_a_relabel_beside_an_inline_control_resolves_and_confirms(self):
        template, version = self.uploaded(data=control_line('Hospital number: ', 'Hospital number', '    Weight: ______ kg'),
                                          name='Synthetic RR3-03')
        weight = self.by_label(version)['Weight']
        other = self.by_label(version)['Hospital number']
        self.assertTrue(mapping.overlaps(weight, other))      # the engine's own overlap (compared by paragraph)
        proposal, new = self.assisted(template, version, weight, 'Weight')
        self.assertEqual(proposal['review']['changes'], [{'key': weight['key'], 'replacement': weight['key']}])
        self.assertEqual(self.resolve(template, new).status_code, 302)
        latest = self.latest(template)
        self.assertEqual(latest.number, new.number + 1)
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_confirm', 'mapping_version': latest.pk,
                                                                   'checked_every_field': 'on'})
        self.assertEqual(response.status_code, 302)

    def test_a_new_overlap_of_a_model_field_is_still_refused_at_resolve(self):
        """The control: an overlap the base mapping did NOT have (here a model field moved onto its neighbour's place)."""
        template, version = self.uploaded(data=lines('Weight: ______ kg  Height: ______ cm', 'Plan: ______'),
                                          name='Synthetic RR3-03 control')
        weight, height = version.fields[0], version.fields[1]
        self.assertFalse(mapping.overlaps(weight, height))
        moved = {**weight, 'locator': copy.deepcopy(height['locator']), 'origin': 'model_assisted'}
        fields = [moved if f['key'] == weight['key'] else f for f in version.fields]
        new = TemplateVersion.objects.create(template=template, number=2, fields=fields, contract='v3',
            **import_metadata(version, fields),
            created_by=self.user, assistance={'unresolved': [{'slot': BODY + weight['locator']['slot'], 'key': '',
                                                             'reason': 'Unresolved source area', 'text': ''}],
                                              'base_version': str(version.pk)})
        response = self.client.post(f'/templates/{template.pk}/assistance/',
                                    {'action': 'resolve', 'mapping_version': new.pk, 'resolution_0': 'mapped'})
        self.assertEqual(response.status_code, 400)
        self.assertIn('write to the same place', response.content.decode())
        self.assertEqual(self.latest(template).pk, new.pk)

    def test_rr1_10_another_persons_signature_line_is_never_relabelled_by_a_model(self):
        template, version = self.uploaded(data=lines('Signed (mother, DOB): ______________', 'Plan: ______'),
                                          name='Synthetic RR1-10 relabel')
        signed = self.by_label(version)['Signed (mother, DOB)']
        self.assertEqual(signed['kind'], 'date')
        compiled = mapping.compile_request(version, self.config)
        payload = json.loads(compiled['request']['messages'][0]['content'])
        body = mapping.wire.fixture(payload)
        fid = next(i for i, f in enumerate(payload['fields']) if f['key'] == signed['key'])
        body['f']['k'] = [n for n in body['f']['k'] if n != fid]; body['f']['r'] = [fid]
        # 'DOB' is printed wording of that line and no identity label: adopted, generation would write a date there.
        body['x'] = [ModelProposals.change(payload, signed['locator']['slot'], 'DOB', kind='date', fid=fid)]
        result = mapping.validate_proposal(version, compiled, body)
        self.assertEqual(result['review']['changes'], [])
        self.assertIn(mapping.NOT_ORDINARY, ' '.join(r['reason'] for r in result['review']['unresolved']))

    def test_one_rule_proposal_and_resolve_read_model_overlaps(self):
        template, version = self.uploaded(data=control_line('Hospital number: ', 'Hospital number', '    Weight: ______ kg'),
                                          name='Synthetic RR3-03 rule')
        weight, control = self.by_label(version)['Weight'], self.by_label(version)['Hospital number']
        relabelled = {**weight, 'origin': 'model_assisted'}
        # the engine's own overlap, same places: excused
        self.assertEqual(mapping.model_overlaps([relabelled, control], version.fields), [])
        # no base recorded: nothing is excused
        self.assertEqual(len(mapping.model_overlaps([relabelled, control], [])), 1)
        # a place that moved is judged as new, even where the prior places overlapped too
        moved = {**relabelled, 'locator': copy.deepcopy(control['locator'])}
        self.assertEqual(len(mapping.model_overlaps([moved, control], version.fields)), 1)


# ─── b2 repair round (refutation of 258650f2): each regression vs 43fcf214 is removed by restoring 43fcf214's behaviour
# for its input class - its cause dropped or narrowed, no wording reader added. Each test is red on 258650f2.
def pdf_form(fields):
    from reportlab.pdfgen import canvas
    buffer = io.BytesIO(); c = canvas.Canvas(buffer); y = 780
    for printed, name in fields:
        c.drawString(30, y, printed)
        c.acroForm.textfield(name=name, x=200, y=y - 8, width=300, height=22, fontName='Helvetica', fontSize=10)
        y -= 60
    c.showPage(); c.save(); return buffer.getvalue()


def log_table(headers, group=None, rows=3):
    """A bordered log table: a header row (or a two-level header, `group` = (text, first, last) merged over columns
    first..last) and `rows` empty rows."""
    doc = Document(); doc.add_paragraph('Synthetic results log')
    extra = 1 if group else 0
    t = doc.add_table(rows=rows + 1 + extra, cols=len(headers)); t.style = 'Table Grid'
    if group:
        text, start, end = group
        for c in range(len(headers)):
            if c < start or c > end:
                t.cell(0, c).text = headers[c]
        merged = t.cell(0, start).merge(t.cell(0, end)); merged.text = text
        for c in range(len(headers)):
            if start <= c <= end:
                t.cell(1, c).text = headers[c]
            else:
                t.cell(0, c).merge(t.cell(1, c))
    else:
        for c, h in enumerate(headers):
            t.cell(0, c).text = h
    return saved(doc)


FINALIZER = 'Synthetic Alpha Clinician'
IDENTITY_REFUSAL = 'must match your signed-in clinician name'


def messages_of(response):
    return ' '.join(str(m) for m in (response.context or {}).get('messages') or [])


@override_settings(**TEST_SETTINGS)
class RepairLegacyIdentity(Walk, V3Fixture, TestCase):
    """L1:R1-01, L1:R1-02, L2:R2-01 / L3:F1 (legacy half): identity and the rename guard as on 43fcf214."""

    def page_row(self, e, key):
        return next(r for r in self.client.get(f'/encounters/{e.pk}/').context['answer_rows'] if r['key'] == key)

    def test_r1_01_a_qualified_role_is_not_assigned_to_the_finalizer(self):
        """Unknown relationship-specific roles remain human-owned; a finalizer identity cannot infer their authority."""
        for n, label in enumerate(("Signature (participant's consultant)", 'Signed by (doctor recruiting the participant)',
                                   'Signature (donor coordinator)', "Signature (father's GP)", 'Signature (consultant)')):
            with self.subTest(label=label):
                template, version = self.uploaded(data=legacy_rows(label, 'Plan'), name='Synthetic R1-01 %d' % n)
                key = self.by_label(version)[label]['key']
                e = self.generated(version, f'{label}: Dr Somebody Else\nPlan: Rest.')
                self.assertEqual((e.answers[key]['status'], e.answers[key]['text'], e.answers[key]['reason']),
                                 ('missing', '', destination_roles.ROLE_NOT_GENERATED))
                row = self.page_row(e, key)
                self.assertEqual((row['identity'], row['blank_eligible']), (False, True))
                self.assertEqual(self.reviewed(e, version).status_code, 302)
                self.finalized(e)
                self.assertEqual(e.status, 'finalized')
                self.assertEqual(Document(private_path(e.export_path)).tables[0].cell(0, 1).text, '')

    def test_rr1_10_a_dated_signer_remains_human_owned_after_mapping_rename(self):
        for n, (label, value) in enumerate(OTHER_SIGNERS_DATED + (('Signature (father, born)', 'Dr Somebody Else'),)):
            with self.subTest(label=label):
                template, version = self.uploaded(data=legacy_rows(label, 'Plan'), name='Synthetic RR1-10 id %d' % n)
                key = self.by_label(version)[label]['key']
                e = self.generated(version, f'{label}: {value}\nPlan: Rest.')
                self.assertEqual((e.answers[key]['status'], e.answers[key]['text'], e.answers[key]['reason']),
                                 ('missing', '', destination_roles.ROLE_NOT_GENERATED))
                response = self.client.post(f'/templates/{template.pk}/', {
                    'mapping_version': version.pk, 'labels': json.dumps({key: 'Signature', self.by_label(version)['Plan']['key']: 'Plan'})})
                self.assertEqual(response.status_code, 302)
                renamed = self.latest(template)
                newer = self.generated(renamed, 'Signature: ' + FINALIZER + '\nPlan: Rest.')
                self.assertEqual((newer.answers[key]['status'], newer.answers[key]['text']), ('missing', ''))
                row = self.page_row(newer, key)
                self.assertEqual((row['identity'], row['blank_eligible']), (False, True))
                self.assertEqual(self.reviewed(newer, renamed).status_code, 302)
                self.finalized(newer)
                self.assertEqual(newer.status, 'finalized')
                self.assertEqual(Document(private_path(newer.export_path)).tables[0].cell(0, 1).text, '')

    def test_r1_02_a_field_renamed_to_a_signature_label_is_the_finalizers(self):
        """L1:R1-02 (regression vs 43fcf214): identity read only the PRINTED label, so a Word row 'Doctor' or a PDF field
        'Text1' renamed to a signature label finalized another clinician's typed name (and could be left blank)."""
        from django.core.files.uploadedfile import SimpleUploadedFile
        response = self.client.post('/templates/upload/', {'name': 'Synthetic R1-02 pdf', 'department': self.department.pk,
                                    'synthetic_only': 'on', 'file': SimpleUploadedFile('form.pdf', pdf_form(
                                        [('Signature of doctor:', 'Text1'), ('Plan:', 'Text2')]))})
        self.assertEqual(response.status_code, 302)
        pdf = Template.objects.get(name='Synthetic R1-02 pdf')
        word, _ = self.uploaded(data=legacy_rows('Doctor', 'Plan'), name='Synthetic R1-02 word')
        for template, old, new in ((pdf, 'Text1', 'Signature'), (word, 'Doctor', "Doctor's signature")):
            with self.subTest(label=new):
                version = self.latest(template)
                labels = {f['key']: (new if f['label'] == old else f['label']) for f in version.fields}
                response = self.client.post(f'/templates/{template.pk}/', {'mapping_version': version.pk,
                                                                          'labels': json.dumps(labels)})
                self.assertEqual(response.status_code, 302)
                renamed = self.latest(template)
                key = self.by_label(renamed)[new]['key']
                e = self.generated(renamed, f'{new}: Dr Somebody Else\nPlan: Rest.')
                self.assertEqual(e.answers[key]['status'], 'missing')
                row = self.page_row(e, key)
                self.assertEqual((row['identity'], row['blank_eligible']), (True, False))
                post = {'revision': e.revision}
                for f in renamed.fields:
                    post['answer_' + f['key']] = '' if f['key'] == key else 'Rest.'
                    post['reviewed_' + f['key']] = 'on'
                post['blank_' + key] = 'on'
                response = self.client.post(f'/encounters/{e.pk}/save/', post)
                self.assertEqual(response.status_code, 400)
                self.assertIn('cannot be marked', response.content.decode())
                self.assertEqual(self.reviewed(e, renamed, typed={new: 'Dr Somebody Else'}).status_code, 302)
                self.assertIn(IDENTITY_REFUSAL, messages_of(self.finalized(e)))
                self.assertEqual(e.status, 'review')


    def test_the_review_page_says_why_a_signing_row_was_not_generated(self):
        """Demo rehearsal gap: a held 'Patient signature' row showed only 'No supporting passage was attached'."""
        template, version = self.uploaded(data=legacy_rows('Patient signature', 'Plan'), name='Synthetic held reason')
        by = self.by_label(version)
        e = self.generated(version, 'Patient signature: Synthetic Person\nPlan: Rest.')
        page = self.client.get(f'/encounters/{e.pk}/')
        rows = {r['key']: r for r in page.context['answer_rows']}
        self.assertEqual(rows[by['Patient signature']['key']]['destination_reason'], destination_roles.ROLE_NOT_GENERATED)
        self.assertEqual(rows[by['Plan']['key']]['destination_reason'], '')
        self.assertContains(page, 'data-destination-role', count=1)
        self.assertContains(page, 'never generated from notes')

    def test_fv_01_original_human_roles_cannot_appoint_finalizer_identity_after_rename(self):
        for n, old in enumerate(('Initials','Signature and date','Signed off by consultant')):
            with self.subTest(old=old):
                template, version = self.uploaded(data=legacy_rows(old, 'Plan'), name='Synthetic FV-01 %d' % n)
                labels = {f['key']: ('Signature' if f['label'] == old else f['label']) for f in version.fields}
                response = self.client.post(f'/templates/{template.pk}/', {'mapping_version': version.pk,
                                                                          'labels': json.dumps(labels)})
                self.assertEqual(response.status_code, 302)
                renamed = self.latest(template);key = self.by_label(renamed)['Signature']['key']
                e = self.generated(renamed, 'Signature: ' + FINALIZER + '\nPlan: Rest.')
                self.assertEqual((e.answers[key]['status'],e.answers[key]['text']), ('missing',''))
                row = self.page_row(e, key)
                self.assertEqual((row['identity'], row['blank_eligible']), (False, True))
                self.assertEqual(self.reviewed(e,renamed).status_code,302)
                self.finalized(e);self.assertEqual(e.status,'finalized')
                self.assertEqual(Document(private_path(e.export_path)).tables[0].cell(0,1).text,'')



@override_settings(**V3)
class RepairV3Rename(Walk, V3Fixture, TestCase):
    def test_r2_01_other_signer_dated_line_keeps_original_human_role_after_relabel(self):
        """Retained L2:R2-01 / L3:F1 sources stay blank through actual export after mapping relabel."""
        for n, (label, renamed) in enumerate((('Signature (participant, dd/mm/yyyy)', 'Participant'),
                                              ('Signature (participant, dd/mm/yyyy)', 'Participant signature and date'),
                                              ('Signed (mother, DOB)', 'Mother DOB'), ('Signed (mother, DOB)', 'DOB'),
                                              ('Signed (mother, DOB)', "Mother's DOB"),
                                              ('Signed (mother, DOB)', "Mother's signature (DOB)"),
                                              ('Signature (father, born)', 'Father born'))):
            with self.subTest(label=label, renamed=renamed):
                template, version = self.uploaded(data=lines(label + ': ______________', 'Plan: ______'),
                                                  name='Synthetic R2-01 %d' % n)
                field = self.by_label(version)[label]
                response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_save', 'mapping_version': version.pk,
                                                                          'label_' + field['key']: renamed})
                self.assertEqual(response.status_code, 302)
                latest = self.latest(template)
                mapped = next(f for f in latest.fields if f['key'] == field['key'])
                self.assertEqual(mapped['locator'], field['locator'])
                self.assertEqual(mapped['kind'], field['kind'])
                self.assertFalse(destination_roles.v3_held(latest))
                self.confirm(template, latest)
                e = self.generated(latest, f'{renamed}: Dr Somebody Else\nPlan: Rest.')
                self.assertEqual((e.answers[field['key']]['status'], e.answers[field['key']]['text']), ('missing', ''))
                page = self.client.get(f'/encounters/{e.pk}/')
                row = next(r for r in page.context['answer_rows'] if r['key'] == field['key'])
                self.assertFalse(row['identity'])
                self.assertEqual(self.reviewed(e, latest).status_code, 302)
                self.finalized(e); self.assertEqual(e.status, 'finalized')
                exported = part_text(private_path(e.export_path).read_bytes())
                self.assertIn(label + ': ______________', exported)
                self.assertNotIn('Dr Somebody Else', exported)


@override_settings(**V3)
class RepairV3Usefulness(Walk, V3Fixture, TestCase):
    def rows_post(self, e, version, texts):
        post = {'revision': e.revision}
        for f in version.fields:
            if f.get('kind') in v3_mapping.SIGNING_KINDS:
                continue
            post['answer_' + f['key']] = texts.get(f['label'], (e.answers.get(f['key']) or {}).get('text', ''))
            post['reviewed_' + f['key']] = 'on'
        response = self.client.post(f'/encounters/{e.pk}/save/', post)
        e.refresh_from_db()
        return response

    def test_r1_03_a_log_tables_clinical_columns_take_typed_rows(self):
        """L1:R1-03 (regression vs 43fcf214): a log-table column whose printed header the signing reader reads as an
        initials place ('Visit type (initial/follow-up)', 'Dose (initial)', 'Initial' under 'Troponin (ng/L)') refused
        the clinician's typed rows at save and fill. 43fcf214 filled them."""
        cases = ((['Date', 'Visit type (initial/follow-up)', 'Clinician notes'], None, ['01/09/2026', 'Follow-up', 'Stable']),
                 (['Drug', 'Dose (initial)', 'Dose (maintenance)'], None, ['Levothyroxine', '25 micrograms', '50 micrograms']),
                 (['Date', 'Initial', 'Repeat'], ('Troponin (ng/L)', 1, 2), ['01/09/2026', '12', '14']),
                 (['Date', 'Troponin initial', 'Troponin 3 h'], None, ['01/09/2026', '12', '14']),
                 (['Problem', 'Initial and ongoing management'], None, ['Sepsis', 'IV antibiotics']))
        for n, (headers, group, row) in enumerate(cases):
            with self.subTest(headers=headers):
                template, version = self.uploaded(data=log_table(headers, group=group), name='Synthetic R1-03 %d' % n)
                table = next(f for f in version.fields if f['kind'] == 'table_rows')
                version = self.confirm(template)
                e = self.generated(version, 'Clinic review.')
                self.assertEqual(self.rows_post(e, version, {table['label']: ' | '.join(row)}).status_code, 302)
                self.finalized(e)
                self.assertEqual(e.status, 'finalized')
                exported = part_text(private_path(e.export_path).read_bytes())
                for value in row:
                    self.assertIn(value, exported)

    def test_r1_05_a_date_pair_that_does_not_close_the_label_keeps_43fcf214s_kind_and_is_generated(self):
        """L1:R1-05 (regression vs 43fcf214): the signature/initials-and-date pair arms read 'Date signed / date received',
        'Care plan signed and date of next review', 'Patient initials and date of birth' and 'Date of diagnosis / initial
        stage' as signing places, so they were never generated (and a replacement field there was refused)."""
        for n, (label, kind, value) in enumerate((('Date signed / date received', 'date', '01/09/2026'),
                                                  ('Date of diagnosis / initial stage', 'date', '12/03/2026'),
                                                  ('Care plan signed and date of next review', 'long_text', 'Yes, review 01/12/2026'),
                                                  ('Patient initials and date of birth', 'date', '01/01/1970'))):
            with self.subTest(label=label):
                template, version = self.uploaded(data=lines(label + ': ______________', 'Plan: ______________'),
                                                  name='Synthetic R1-05 %d' % n)
                field = self.by_label(version)[label]
                self.assertEqual(field['kind'], kind)
                if label == 'Care plan signed and date of next review':
                    self.assertTrue(field['write_protected'])
                    self.assertEqual(self.client.post(f'/templates/{template.pk}/', {'action': 'v3_confirm',
                        'mapping_version': version.pk, 'checked_every_field': 'on'}).status_code, 400)
                    issues = version.import_manifest['coverage']['selected_engine_issues']
                    response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_coverage',
                        'mapping_version': version.pk, **{'exclude_' + i['id']: 'on' for i in issues}})
                    self.assertEqual(response.status_code, 302)
                    version = self.confirm(template)
                    e = self.generated(version, f'{label}: {value}\nPlan: Rest.')
                    self.assertEqual(e.answers[field['key']]['text'], '')
                    self.assertEqual(self.rows_post(e, version, {}).status_code, 302)
                    self.finalized(e); self.assertEqual(e.status, 'finalized')
                    self.assertIn(label + ': ______________', part_text(private_path(e.export_path).read_bytes()))
                    continue  # explicit partial coverage; not a recovered ordinary-field pass
                version = self.confirm(template)
                e = self.generated(version, f'{label}: {value}\nPlan: Rest.')
                self.assertEqual(e.answers[field['key']]['status'], 'supported')
                self.assertEqual(self.rows_post(e, version, {}).status_code, 302)
                self.finalized(e)
                self.assertEqual(e.status, 'finalized')

    def test_r1_05_ambiguous_signed_form_stays_human_after_remove_add(self):
        """Retained L1:R1-05 source requires an exact unchanged exclusion; mapping clicks cannot authorize it."""
        template, version = self.uploaded(data=lines('Consent form signed and dated: ______________', 'Plan: ______________'),
                                          name='Synthetic R1-05 consent')
        field = self.by_label(version)['Consent form signed and dated']
        self.assertEqual(field['kind'], 'signature')
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_save', 'mapping_version': version.pk,
                                                                  'remove_' + field['key']: 'on'})
        self.assertEqual(response.status_code, 302)
        latest = self.latest(template)
        self.assertNotIn(field['key'], {f['key'] for f in latest.fields})
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_add', 'mapping_version': latest.pk,
                                    'slot': BODY + 'b/p[1]', 'kind': 'text', 'label': 'Consent form signed and dated'})
        self.assertEqual(response.status_code, 400)
        issues = latest.import_manifest['coverage']['selected_engine_issues']
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_coverage',
            'mapping_version': latest.pk, **{'exclude_' + i['id']: 'on' for i in issues}})
        self.assertEqual(response.status_code, 302)
        reviewed = self.confirm(template)
        e = self.generated(reviewed, 'Consent form signed and dated: Yes\nPlan: Rest.')
        self.assertEqual(self.rows_post(e, reviewed, {}).status_code, 302)
        self.finalized(e); self.assertEqual(e.status, 'finalized')
        self.assertIn('Consent form signed and dated: ______________', part_text(private_path(e.export_path).read_bytes()))


@override_settings(**TEST_SETTINGS)
class RepairLegacyUsefulness(Walk, V3Fixture, TestCase):
    def test_r1_04_readings_the_kind_rule_or_a_closing_pair_does_not_make_a_signing_place_are_generated(self):
        """L1:R1-04 (regression vs 43fcf214) and PC05: these legacy rows were held as signing places and never generated.
        43fcf214 generated and exported each value."""
        for n, (label, value) in enumerate((('Care plan signed and date of next review', 'Yes, review 01/12/2026'),
                                            ('Consent form signed and dated', 'Yes'),
                                            ('Date of referral / initial contact', '02/09/2026'),
                                            ('Date signed / date received', '01/09/2026 / 03/09/2026'),
                                            ('Patient initials and date of birth', 'J.E. 01/01/1970'),
                                            ('Tinel sign both wrists', 'Positive'),
                                            ('Sign and symptom review', 'No new symptoms'))):
            with self.subTest(label=label):
                template, version = self.uploaded(data=legacy_rows(label, 'Plan'), name='Synthetic R1-04 %d' % n)
                key = self.by_label(version)[label]['key']
                if label in ('Care plan signed and date of next review', 'Consent form signed and dated', 'Sign and symptom review'):
                    before = (Job.objects.count(), Reservation.objects.count())
                    self.assertTrue(version.import_manifest['coverage']['block_generation'])
                    self.assertEqual(self.new_encounter(version, f'{label}: {value}\nPlan: Rest.').status_code, 302)
                    e = Encounter.objects.latest('created_at')
                    with self.assertRaisesMessage(ValidationError, 'No generation or charge'):
                        enqueue(e, self.user)
                    self.assertEqual((Job.objects.count(), Reservation.objects.count()), before)
                    self.assertEqual(e.answers, {}); self.assertFalse(e.export_path)
                    continue  # retained source is explicitly unsupported on this legacy route
                e = self.generated(version, f'{label}: {value}\nPlan: Rest.')
                self.assertEqual((e.answers[key]['status'], e.answers[key]['text']), ('supported', value))
                self.assertEqual(self.reviewed(e, version).status_code, 302)
                self.finalized(e)
                self.assertEqual(e.status, 'finalized')
                self.assertIn(value, part_text(private_path(e.export_path).read_bytes()))


@override_settings(**V3)
class RoleReportReuse(V3Fixture, TestCase):
    """Integrated-candidate refutation R2-1: every v3 review save re-ran the isolated parser (5-11x slower than
    43fcf214). A verified reading of the same, unchanged original is reused across requests of one process."""

    def setUp(self):
        super().setUp()
        destination_roles._PROCESS_REPORTS.clear()

    def counted(self):
        import documents.service as service
        real = service.destination_roles
        calls = []

        def counting(*args, **kwargs):
            calls.append(args[0])
            return real(*args, **kwargs)
        return calls, patch.object(service, 'destination_roles', counting)

    def test_a_second_request_reuses_the_reading_of_an_unchanged_original(self):
        template, version = self.uploaded(data=lines('Plan: ______', 'Signature: ______'), name='Synthetic reuse')
        calls, counting = self.counted()
        with counting:
            first = destination_roles.v3_held(TemplateVersion.objects.get(pk=version.pk))
            second = destination_roles.v3_held(TemplateVersion.objects.get(pk=version.pk))   # a new request's object
        self.assertEqual(first, second)
        self.assertEqual(len(calls), 1)

    def test_a_changed_mapping_is_read_again(self):
        template, version = self.uploaded(data=lines('Plan: ______', 'Signature: ______'), name='Synthetic reread')
        calls, counting = self.counted()
        with counting:
            destination_roles.v3_held(TemplateVersion.objects.get(pk=version.pk))
            fields = [dict(f) for f in version.fields]
            fields[0]['label'] = fields[0]['label'] + ' (renamed)'
            destination_roles.v3_held(TemplateVersion.objects.get(pk=version.pk), fields)
        self.assertEqual(len(calls), 2)

    def test_an_original_rewritten_in_place_with_its_old_size_and_mtime_is_refused(self):
        """Cache refutation F2: inode, size and mtime are not the content; a hit re-hashes the file."""
        import os
        template, version = self.uploaded(data=lines('Plan: ______', 'Signature: ______'), name='Synthetic rewritten')
        destination_roles.v3_held(TemplateVersion.objects.get(pk=version.pk))
        path = private_path(template.original_path)
        original, stat = open(path, 'rb').read(), os.stat(path)
        forged = original[:-1] + bytes([original[-1] ^ 0x01])
        try:
            with open(path, 'r+b') as stream:
                stream.write(forged)
            os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns))
            self.assertEqual((os.stat(path).st_ino, os.stat(path).st_size, os.stat(path).st_mtime_ns),
                             (stat.st_ino, stat.st_size, stat.st_mtime_ns))
            with self.assertRaises((DocumentError, ValidationError)):
                destination_roles.v3_held(TemplateVersion.objects.get(pk=version.pk))
        finally:
            with open(path, 'r+b') as stream:
                stream.write(original)

    def test_an_original_present_but_unreadable_is_refused_after_a_cached_reading(self):
        """Cache refutation F1: a chmod changes no identity the old key read; a hit re-reads the file."""
        import os
        template, version = self.uploaded(data=lines('Plan: ______', 'Signature: ______'), name='Synthetic chmod')
        destination_roles.v3_held(TemplateVersion.objects.get(pk=version.pk))
        path = private_path(template.original_path)
        mode = os.stat(path).st_mode
        os.chmod(path, 0)
        try:
            if os.access(path, os.R_OK):
                self.skipTest('this account can read a mode-000 file')
            with self.assertRaises((DocumentError, ValidationError, OSError)):
                destination_roles.v3_held(TemplateVersion.objects.get(pk=version.pk))
        finally:
            os.chmod(path, mode)

    def test_a_caller_mutating_a_reading_does_not_change_the_next_one(self):
        template, version = self.uploaded(data=lines('Plan: ______', 'Signature: ______'), name='Synthetic poison')
        first = destination_roles.report(TemplateVersion.objects.get(pk=version.pk))
        expected = copy.deepcopy(first)
        first.clear()
        hit = destination_roles.report(TemplateVersion.objects.get(pk=version.pk))
        hit['violations'] = {'forged': {'reason': 'x'}}
        self.assertEqual(destination_roles.report(TemplateVersion.objects.get(pk=version.pk)), expected)

    def test_an_integrity_failure_is_never_retained(self):
        template, version = self.uploaded(data=lines('Plan: ______', 'Signature: ______'), name='Synthetic not kept')
        import documents.service as service
        real = service.destination_roles
        with patch.object(service, 'destination_roles', lambda *a, **k: {**real(*a, **k), 'sha256': '0' * 64}):
            for _ in range(2):
                with self.assertRaises(ValidationError):
                    destination_roles.v3_held(TemplateVersion.objects.get(pk=version.pk))
        self.assertEqual(destination_roles._PROCESS_REPORTS, {})

    def test_a_reading_larger_than_the_entry_limit_is_not_retained(self):
        template, version = self.uploaded(data=lines('Plan: ______', 'Signature: ______'), name='Synthetic large')
        calls, counting = self.counted()
        with counting, patch.object(destination_roles, '_PROCESS_ENTRY_LIMIT', 10):
            destination_roles.v3_held(TemplateVersion.objects.get(pk=version.pk))
            destination_roles.v3_held(TemplateVersion.objects.get(pk=version.pk))
        self.assertEqual(len(calls), 2)

    def test_an_unreadable_original_is_still_refused_after_a_cached_reading(self):
        import os
        template, version = self.uploaded(data=lines('Plan: ______', 'Signature: ______'), name='Synthetic unreadable')
        destination_roles.v3_held(TemplateVersion.objects.get(pk=version.pk))
        path = private_path(template.original_path)
        os.rename(path, str(path) + '.aside')
        try:
            with self.assertRaises((DocumentError, ValidationError, OSError)):
                destination_roles.v3_held(TemplateVersion.objects.get(pk=version.pk))
        finally:
            os.rename(str(path) + '.aside', path)


@override_settings(**V3)
class MappingReviewCoverage(Walk, V3Fixture, TestCase):
    @override_settings(AI_PROVIDER='anthropic', AI_MODEL='claude-sonnet-5-5', AI_THINKING='between_tools',
                       AI_EFFORT='medium', AI_INPUT_RATE='2', AI_OUTPUT_RATE='10', AI_CACHED_RATE='.2',
                       V3_FIELD_RELATIONS=True)
    def test_reserved_manager_statement_is_not_asked_or_written_through_upload_and_export(self):
        original = legacy_rows('This section is for the line manager to add comments', 'Plan')
        template, version = self.uploaded(data=original, name='Synthetic manager source reservation')
        self.assertEqual(version.contract, 'v3')
        manager, plan = version.fields
        self.assertTrue(manager.get('write_protected'))
        issues = version.import_manifest['coverage']['selected_engine_issues']
        self.assertTrue(any(i['code'] == 'human_authored_destination' for i in issues))
        # Retain the original first mapping and source; the explicit reviewed
        # exclusion is recovery, never an automatic whole-document success.
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_coverage',
            'mapping_version': version.pk, **{'exclude_' + i['id']: 'on' for i in issues}})
        self.assertEqual(response.status_code, 302)
        reviewed = self.latest(template); self.confirm(template, reviewed)
        notes = 'This section is for the line manager to add comments: Dr Note Person\nPlan: Rest.'
        e = self.generated(reviewed, notes)
        import os
        from pathlib import Path
        if location := os.environ.get('QUILLMD_WORD_ARTIFACTS'):
            capture = Path(location) / 'manager-reserved-statement'
            capture.mkdir(mode=0o700, parents=True)
            (capture / 'original.docx').write_bytes(original)
            (capture / 'first.json').write_text(json.dumps({'notes': notes, 'fields': reviewed.fields,
                'import_manifest': reviewed.import_manifest, 'answers': e.answers,
                'configuration': Job.objects.get(encounter=e, kind='generation').config}, indent=2, default=str))
        self.assertEqual((e.answers[manager['key']]['status'], e.answers[manager['key']]['text']), ('missing', ''))
        self.assertEqual(e.answers[plan['key']]['text'], 'Rest.')
        self.assertEqual(self.reviewed(e, reviewed).status_code, 302)
        self.finalized(e); self.assertEqual(e.status, 'finalized')
        exported = Document(private_path(e.export_path))
        self.assertEqual(exported.tables[0].cell(0, 1).text, '')
        self.assertEqual(exported.tables[0].cell(1, 1).text, 'Rest.')
        self.assertEqual(private_path(template.original_path).read_bytes(), original)

    def test_rf08_mapping_review_reads_bound_roles_before_exclusions_but_admission_stays_held(self):
        from core import import_binding
        doc = Document(); doc.add_heading('Consultant review', level=2)
        doc.add_paragraph('Signed by the consultant:'); box(doc); doc.add_paragraph('Plan: ______')
        original = saved(doc)
        template, version = self.uploaded(data=original, name='Synthetic RF08 review before exclusions')
        issue = version.import_manifest['coverage']['selected_engine_issues'][0]
        self.assertEqual((issue['code'], issue['slot']), ('unmapped_answer', 'b/tbl[0]/tr[0]/tc[0]'))
        with override_settings(**MAPPING_SETTINGS):
            self.config = generation.configuration()
        compiled, payload = ModelProposals.compiled(self, version)
        body = ModelProposals.keep_all(payload)
        body['x'] = [ModelProposals.change(payload, 'b/tbl[0]/tr[0]/tc[0]/p[0]',
                                         'Consultant review', 'long_text', role=1)]
        result = mapping.validate_proposal(version, compiled, body)
        self.assertEqual(result['fields'], version.fields)
        self.assertEqual(result['review']['changes'], [])
        self.assertTrue(any(row.get('proposal') for row in result['review']['unresolved']))
        self.assertTrue(destination_roles.report(version, mapping_review=True)['roles'])
        # This allowance is role inspection only: all immutable identities
        # are still checked, including cached readings of the exact original.
        for key in ('source_sha256', 'engine', 'engine_version', 'engine_code_sha256', 'mapping_sha256'):
            with self.subTest(tampered=key):
                altered = copy.copy(version); altered.import_manifest = copy.deepcopy(version.import_manifest)
                altered.import_manifest[key] = 'changed'
                with self.assertRaisesMessage(ValidationError, 'does not match this original'):
                    destination_roles.report(altered, mapping_review=True)
        with self.assertRaisesMessage(ValidationError, 'No generation or charge'):
            destination_roles.report(version)
        MappingConfirmation.objects.create(version=version, confirmed_by=self.user)
        self.assertEqual(self.new_encounter(version, 'Plan: Rest.').status_code, 302)
        blocked = Encounter.objects.latest('created_at')
        with self.assertRaisesMessage(ValidationError, 'No generation or charge'):
            enqueue(blocked, self.user)
        self.assertEqual(Job.objects.count(), 0); self.assertEqual(Reservation.objects.count(), 0)
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_coverage',
            'mapping_version': version.pk, 'exclude_' + issue['id']: 'on'})
        self.assertEqual(response.status_code, 302)
        reviewed = self.latest(template); self.assertTrue(import_binding.coverage_reviewed(reviewed))
        self.confirm(template, reviewed)
        e = self.generated(reviewed, 'Plan: Rest.')
        self.assertEqual(self.reviewed(e, reviewed).status_code, 302)
        self.finalized(e); self.assertEqual(e.status, 'finalized')
        exported = Document(private_path(e.export_path))
        self.assertEqual(exported.tables[0].cell(0, 0).text, '')
        self.assertIn('Rest.', part_text(private_path(e.export_path).read_bytes()))
        self.assertEqual(hashlib.sha256(private_path(template.original_path).read_bytes()).hexdigest(),
                         hashlib.sha256(original).hexdigest())
