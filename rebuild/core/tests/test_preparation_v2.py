"""Product role protocol and historical evidence compatibility; no live calls."""
from copy import deepcopy
from contextlib import contextmanager, nullcontext
from decimal import Decimal
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
from typing import Annotated
from unittest.mock import Mock, patch
from docx import Document
from django.template.loader import render_to_string
from django.test import SimpleTestCase

from documents.preparation import inventory, TYPED_PROTOCOL, validate_role, digest
from core import preparation


class TypedRoleProductTests(SimpleTestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.path = Path(temporary.name) / 'original.docx'
        document = Document()
        document.add_paragraph('Temperature: ____ °C')
        document.add_paragraph('Fever: ____ (Yes / No)')
        document.add_paragraph('Patient signature: ____')
        document.save(self.path)
        self.config = {'mode': 'fixture', 'provider': 'anthropic', 'model': 'configured-label',
                       'input_rate': '3', 'output_rate': '15', 'cached_rate': '0.3'}

    def test_both_new_fixture_roles_include_metadata_and_are_clearly_simulated(self):
        value = inventory(self.path, protocol=TYPED_PROTOCOL)
        proposal, proposed_evidence = preparation.provider_role(value, 'proposer', self.config, 1)
        checked, checked_evidence = preparation.provider_role(value, 'checker', self.config, 1, proposal)
        for role, response, evidence in [('proposer', proposal, proposed_evidence), ('checker', checked, checked_evidence)]:
            validate_role(value, response, role)
            self.assertTrue(evidence['simulated'])
            self.assertEqual(evidence['model'], 'synthetic-blank-preparation-' + role)
            self.assertEqual(evidence['cost'], '0')
            self.assertEqual(response['anchors'][0]['metadata']['unit'], '°C')
            self.assertEqual(response['anchors'][1]['metadata']['choices'], ['Yes', 'No'])

    def test_historical_protocol_payload_schema_and_results_are_not_upgraded(self):
        value = inventory(self.path)
        before = deepcopy(value)
        proposal, _ = preparation.provider_role(value, 'proposer', self.config, 1)
        checker, _ = preparation.provider_role(value, 'checker', self.config, 1, proposal)
        self.assertFalse(any('metadata' in row for row in proposal['anchors'] + checker['anchors']))
        self.assertNotIn('metadata', preparation.role_schema('proposer')['properties']['anchors']['items']['properties'])
        self.assertEqual(preparation.role_payload(value, 'checker', proposal),
                         {'original': preparation.model_inventory(before), 'role': 'checker', 'proposal': proposal})
        validate_role(value, proposal, 'proposer')

    def test_short_schema_and_quote_fit_existing_per_role_bound_without_calling_api(self):
        value = inventory(self.path, protocol=TYPED_PROTOCOL)
        max_output, cap = preparation.role_limits(value)
        self.assertEqual((max_output, cap), (2500, Decimal('0.10')))
        proposal, _ = preparation.provider_role(value, 'proposer', self.config, 1)
        live = {**self.config, 'mode': 'live'}
        for role in ['proposer', 'checker']:
            payload = preparation.role_payload(value, role, proposal if role == 'checker' else None)
            self.assertLessEqual(preparation.maximum_cost(payload, live, max_output), cap)
        self.assertIn('metadata', preparation.role_schema('checker', TYPED_PROTOCOL)['properties']['anchors']['items']['required'])

    def test_pinned_sdk_transforms_both_typed_roles_without_weakening_source_checks(self):
        from anthropic import transform_schema
        value = inventory(self.path, protocol=TYPED_PROTOCOL)
        for role in ['proposer', 'checker']:
            schema = preparation.role_schema(role, TYPED_PROTOCOL)
            before = deepcopy(schema)
            # Preserve the actual pre-send counterexample without invoking a
            # client: the pinned SDK cannot transform JSON Schema type arrays.
            broken = deepcopy(schema)
            broken['properties']['anchors']['items']['properties']['metadata']['properties']['unit'] = {
                'type': ['string', 'null'], 'maxLength': 80}
            with self.assertRaises(AssertionError):
                transform_schema(broken)
            wire = transform_schema(schema)
            self.assertEqual(schema, before)
            metadata = wire['properties']['anchors']['items']['properties']['metadata']
            self.assertEqual(metadata['properties']['unit']['anyOf'][1], {'type': 'null'})
            self.assertEqual(metadata['properties']['unit']['anyOf'][0]['type'], 'string')
            # The SDK documents unsupported length limits in the wire schema;
            # local exact source checks remain mandatory after the response.
            self.assertIn('maxLength: 80', metadata['properties']['unit']['anyOf'][0]['description'])
            self.assertFalse(metadata['additionalProperties'])
            self.assertEqual(metadata['required'], before['properties']['anchors']['items']['properties']['metadata']['required'])
            proposal, _ = preparation.provider_role(value, 'proposer', self.config, 1)
            result, _ = preparation.provider_role(value, role, self.config, 1, proposal)
            validate_role(value, result, role)
            for unsupported_unit in [None, '°F', 'x'*81, False]:
                changed = deepcopy(result)
                changed['anchors'][0]['metadata']['unit'] = unsupported_unit
                with self.subTest(role=role, unsupported_unit=unsupported_unit), self.assertRaises(ValueError):
                    validate_role(value, changed, role)

    def test_nullable_unit_schema_matches_independent_bounded_string_contract(self):
        from pydantic import StringConstraints, TypeAdapter, ValidationError
        # Pydantic independently generates the same nullable 80-character
        # string schema; both declarations have the old union's meaning.
        adapter = TypeAdapter(Annotated[str, StringConstraints(max_length=80)] | None)
        for role in ['proposer', 'checker']:
            schema = preparation.role_schema(role, TYPED_PROTOCOL)
            unit = schema['properties']['anchors']['items']['properties']['metadata']['properties']['unit']
            self.assertEqual(unit, adapter.json_schema())
        for valid in [None, '', '°C', 'x'*80, '°'*80]:
            self.assertEqual(adapter.validate_python(valid, strict=True), valid)
        for invalid in ['x'*81, '°'*81, False, 0, 1.2, [], {}]:
            with self.subTest(invalid=invalid), self.assertRaises(ValidationError):
                adapter.validate_python(invalid, strict=True)

    def test_typed_live_branch_reaches_stubbed_transport_for_both_roles(self):
        import json
        value = inventory(self.path, protocol=TYPED_PROTOCOL)
        proposal, _ = preparation.provider_role(value, 'proposer', self.config, 1)
        for role in ['proposer', 'checker']:
            fixture, _ = preparation.provider_role(value, role, self.config, 1, proposal)
            response = SimpleNamespace(model='offline-stub-returned-model', id='offline-response', _request_id='offline-request',
                usage=SimpleNamespace(input_tokens=100, output_tokens=20, cache_read_input_tokens=0, cache_creation_input_tokens=0),
                content=[SimpleNamespace(type='text', text=json.dumps(fixture))],
                model_dump=lambda **kw: {'model': 'offline-stub-returned-model', 'usage': {'input_tokens': 100, 'output_tokens': 20}})
            client = Mock(); client.messages.create.return_value = response
            observed = Mock()
            with patch('anthropic.Anthropic', return_value=client) as constructor:
                raw, evidence = preparation.provider_role(value, role, {**self.config, 'mode': 'live'}, 7, proposal, observed)
            constructor.assert_called_once_with(timeout=7, max_retries=0)
            client.messages.create.assert_called_once()
            arguments = client.messages.create.call_args.kwargs
            self.assertEqual(arguments['model'], self.config['model'])
            wire = arguments['output_config']['format']['schema']
            self.assertEqual(wire['properties']['anchors']['items']['properties']['metadata']['properties']['unit']['anyOf'][1], {'type': 'null'})
            self.assertEqual(json.loads(arguments['messages'][0]['content']), preparation.role_payload(value, role, proposal))
            validate_role(value, json.loads(raw), role)
            observed.assert_called_once()
            self.assertEqual(evidence['model'], 'offline-stub-returned-model')

    def underscore_source(self):
        document=Document()
        for text in ['Synthetic underscore form',
                'Choose one listed option for Fever. Preserve °C. Patient signature must remain blank.',
                'Patient name: ____________________','Date of birth: __________',
                'Temperature: ______ °C','Fever: Yes / No ______','Patient signature: ____________________']:
            document.add_paragraph(text)
        document.save(self.path)
        return inventory(self.path,protocol=TYPED_PROTOCOL)

    def test_canonical_text_and_dob_uncertainty_refuse_the_retained_failure_counterexample(self):
        # Actual failed response is retained unchanged under its SHA-256:
        # 6ff0708d23b08c2a5d994caa78bd323b8b06a7f4507f25b91a918fa105e773f7.
        # Reproduce its two wrong uncertainty values against equivalent source;
        # this is not a rewrite or successful replay of that paid response.
        value=self.underscore_source()
        self.assertEqual(value['questions'][0]['answer_constraints'],{'version':1,'data_type':'text'})
        self.assertEqual(value['questions'][1]['answer_constraints'],{'version':1,'data_type':'text'})
        for role in ['proposer','checker']:
            result,_=preparation.provider_role(value,role,self.config,1)
            validate_role(value,result,role)
            for index in [0,1]:
                metadata=result['anchors'][index]['metadata']
                self.assertEqual(metadata,{'label':'Patient name' if index==0 else 'Date of birth',
                    'data_type':'text','unit':None,'choices':[],'destination':'a'+str(index+1),
                    'uncertainty':'constraints_unstated'})
                wrong=deepcopy(result);wrong['anchors'][index]['metadata']['uncertainty']='none'
                with self.subTest(role=role,index=index),self.assertRaisesMessage(ValueError,
                        'The proposed field metadata differs from the original question, constraints or protection.'):
                    validate_role(value,wrong,role)
            self.assertEqual(result['anchors'][2]['metadata']['data_type'],'number')
            self.assertEqual(result['anchors'][2]['metadata']['unit'],'°C')
            self.assertEqual(result['anchors'][2]['metadata']['uncertainty'],'none')
            self.assertEqual(result['anchors'][3]['metadata']['choices'],['Yes','No'])
            self.assertEqual(result['anchors'][3]['metadata']['uncertainty'],'none')
            self.assertEqual(result['anchors'][4]['metadata']['uncertainty'],'signing_protected')
            for index,change in [(1,{'data_type':'date'}),(2,{'unit':'°F'}),
                    (3,{'choices':['Yes','No','Unknown']}),(4,{'uncertainty':'constraints_unstated'})]:
                wrong=deepcopy(result);wrong['anchors'][index]['metadata'].update(change)
                with self.subTest(role=role,index=index,change=change),self.assertRaises(ValueError):
                    validate_role(value,wrong,role)

    def test_actual_sdk_requests_explain_text_uncertainty_to_both_roles_and_preserve_v1(self):
        import httpx
        from anthropic import Anthropic
        value=self.underscore_source()
        proposal,_=preparation.provider_role(value,'proposer',self.config,1)
        historical=inventory(self.path)
        for source in [value,historical]:
            for role in ['proposer','checker']:
                fixture,_=preparation.provider_role(source,role,self.config,1)
                sent=[]
                def handle(request):
                    self.assertEqual(request.url.host,'synthetic-provider.invalid')
                    self.assertEqual(request.url.path,'/v1/messages')
                    sent.append(json.loads(request.content))
                    return httpx.Response(200,headers={'request-id':'offline-request'},json={
                        'id':'offline-response','type':'message','role':'assistant','model':'offline-returned-model',
                        'content':[{'type':'text','text':json.dumps(fixture)}],'stop_reason':'end_turn',
                        'stop_sequence':None,'usage':{'input_tokens':100,'output_tokens':20}})
                with httpx.Client(transport=httpx.MockTransport(handle)) as http:
                    client=Anthropic(api_key='synthetic-offline-only',base_url='https://synthetic-provider.invalid',
                        http_client=http,timeout=7,max_retries=0)
                    with patch('anthropic.Anthropic',return_value=client) as constructor:
                        raw,evidence=preparation.provider_role(source,role,{**self.config,'mode':'live'},7,
                            proposal if role=='checker' and source is value else None)
                constructor.assert_called_once_with(timeout=7,max_retries=0)
                self.assertEqual(len(sent),1)
                body=sent[0];payload=json.loads(body['messages'][0]['content'])
                self.assertEqual(body['model'],self.config['model'])
                self.assertEqual(payload['original'],preparation.model_inventory(source))
                self.assertEqual(payload['role'],role)
                validate_role(source,json.loads(raw),role)
                self.assertFalse(evidence['simulated'])
                self.assertEqual(evidence['model'],'offline-returned-model')
                if source is value:
                    prompt=body['system']
                    self.assertIn('version 1 and data_type text',prompt)
                    self.assertIn('does not justify uncertainty none',prompt)
                    self.assertIn('Date of birth label without a printed format remains text with constraints_unstated',prompt)
                    self.assertLess(prompt.index('(1) Signing'),prompt.index('(2) Ambiguous'))
                    self.assertLess(prompt.index('(2) Ambiguous'),prompt.index('(3) If the original inventory has issues'))
                    self.assertLess(prompt.index('(3) If the original inventory has issues'),prompt.index('(4) Otherwise every source text'))
                    self.assertIn('use uncertainty none',prompt)
                    self.assertEqual(payload['original']['questions'][1]['answer_constraints'],
                        {'version':1,'data_type':'text'})
                    self.assertNotIn('format',payload['original']['questions'][1]['answer_constraints'])
                    self.assertTrue(payload['original']['anchors'][4]['protected_signing'])
                    if role=='checker':self.assertEqual(payload['proposal'],proposal)
                else:
                    self.assertNotIn('Apply uncertainty in this exact priority order',body['system'])
                    self.assertFalse(any('metadata' in row for row in json.loads(raw)['anchors']))

    def test_metadata_display_escapes_untrusted_original_and_shows_uncertainty(self):
        metadata = {'label': '<script>original text</script>', 'data_type': 'number', 'unit': '°C',
                    'choices': [], 'destination': 'a1', 'uncertainty_label': 'Printed constraint verified'}
        html = render_to_string('template_preparation_metadata.html', {'metadata': metadata})
        self.assertIn('&lt;script&gt;', html)
        self.assertNotIn('<script>', html)
        self.assertIn('Unit: °C', html)
        self.assertIn('Printed constraint verified', html)


class WMOriginalStoryContextProductTests(SimpleTestCase):
    """Real pinned SDK transport, mocked network; no paid/model acceptance claim."""
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        from documents.test_preparation_wm_context import original_xml_inventory
        cls.value, _, _, _ = original_xml_inventory()

    def setUp(self):
        self.config = {'mode': 'fixture', 'provider': 'anthropic', 'model': 'configured-label',
                       'input_rate': '3', 'output_rate': '15', 'cached_rate': '0.3'}

    def test_only_new_exact_original_preparations_receive_prospective_context_mode(self):
        original_config = deepcopy(self.config)
        for sha, expected in [(preparation.WM_IOV_SHA256, True), ('0'*64, False)]:
            with self.subTest(sha=sha), patch.object(preparation, 'feature_enabled', return_value=True), \
                    patch.object(preparation, 'configuration', return_value=original_config), \
                    patch('core.preparation_campaign.active_campaign_config', side_effect=lambda config, _: config), \
                    patch('core.testing_authority.identity', return_value={'mode': 'standing_authority'}), \
                    patch('core.testing_authority.create_preparation') as create:
                preparation.create_preparation(SimpleNamespace(sha256=sha), 'synthetic-doctor')
            config = create.call_args.args[2]
            self.assertEqual('word_source_review_context' in config, expected)
            if expected:
                self.assertEqual(config['word_source_review_context'], preparation.WM_SOURCE_REVIEW_CONTEXT)
            self.assertEqual(original_config, self.config)

    def test_old_payload_digest_and_both_retained_role_results_are_unchanged_by_new_context(self):
        before = deepcopy(self.value)
        proposal, _ = preparation.provider_role(self.value, 'proposer', self.config, 1)
        checker, _ = preparation.provider_role(self.value, 'checker', self.config, 1, proposal)
        for role, response in [('proposer', proposal), ('checker', checker)]:
            # Historical unresolved findings remain original evidence.
            response['stories'][0]['verdict'] = 'unresolved'
            response['stories'][7]['verdict'] = 'unresolved'
            proposal_arg = proposal if role == 'checker' else None
            expected = {'original': preparation.model_inventory(before), 'role': role}
            if role == 'checker':
                expected['proposal'] = proposal
            actual = preparation.role_payload(self.value, role, proposal_arg)
            self.assertEqual(actual, expected)
            self.assertEqual(digest(actual), digest(expected))
            new = preparation.role_payload(self.value, role, proposal_arg,
                                           source_review_context=preparation.WM_SOURCE_REVIEW_CONTEXT)
            self.assertNotEqual(digest(new), digest(actual))
            validate_role(self.value, response, role)
            self.assertEqual([s['story_id'] for s in response['stories'] if s['verdict'] == 'unresolved'], ['s1', 's8'])
        self.assertEqual(self.value, before)

    def test_both_actual_sdk_requests_receive_original_context_and_unknown_results_stay_unresolved(self):
        import httpx
        from anthropic import Anthropic
        sent = []
        config = {**self.config, 'mode': 'live',
                  'word_source_review_context': preparation.WM_SOURCE_REVIEW_CONTEXT}
        proposer, _ = preparation.provider_role(self.value, 'proposer', self.config, 1)
        for role in ['proposer', 'checker']:
            output, _ = preparation.provider_role(self.value, role, self.config, 1)
            output['stories'][0]['verdict'] = 'unresolved'
            output['stories'][7]['verdict'] = 'unresolved'
            def handle(request):
                self.assertEqual(request.url.host, 'synthetic-provider.invalid')
                sent.append(json.loads(request.content))
                return httpx.Response(200, headers={'request-id': 'offline-wm-context-request'}, json={
                    'id': 'offline-wm-context-response', 'type': 'message', 'role': 'assistant',
                    'model': 'offline-returned-model', 'stop_reason': 'end_turn', 'stop_sequence': None,
                    'content': [{'type': 'text', 'text': json.dumps(output)}],
                    'usage': {'input_tokens': 100, 'output_tokens': 100}})
            with httpx.Client(transport=httpx.MockTransport(handle)) as http:
                client = Anthropic(api_key='synthetic-not-a-credential', base_url='https://synthetic-provider.invalid',
                                   http_client=http, timeout=7, max_retries=0)
                with patch('anthropic.Anthropic', return_value=client) as constructor:
                    raw, evidence = preparation.provider_role(self.value, role, config, 7,
                                                               proposer if role == 'checker' else None)
            constructor.assert_called_once_with(timeout=7, max_retries=0)
            wire = sent[-1]
            payload = json.loads(wire['messages'][0]['content'])
            context = payload['original']['source_review_context']
            self.assertEqual(len(context['body']['answer_destinations']), 90)
            self.assertEqual(len(context['body']['preserved_review_requirements']), 14)
            self.assertEqual(len(context['glossary']['definitions']), 28)
            self.assertIn('retain NEVER means the doctor has approved', wire['system'])
            self.assertIn('Keep unresolved for any remaining omission', wire['system'])
            self.assertIn('actual finalization', wire['system'])
            if role == 'checker':
                self.assertEqual(payload['proposal'], proposer)
            response = json.loads(raw)
            self.assertEqual(response, output)
            validate_role(self.value, response, role)
            self.assertEqual(evidence['model'], 'offline-returned-model')
            self.assertFalse(evidence['simulated'])  # Mock transport; never counted as a live model check.
        self.assertEqual(len(sent), 2)

    def test_explicit_invalid_modes_and_wrong_original_fail_before_provider_or_reservation(self):
        for mode in [None, '', False, {}, 'universal-word']:
            with self.subTest(mode=mode), patch('anthropic.Anthropic') as constructor, self.assertRaises(ValueError):
                preparation.provider_role(self.value, 'proposer',
                    {**self.config, 'mode': 'live', 'word_source_review_context': mode}, 7)
            constructor.assert_not_called()
        bad = deepcopy(self.value)
        bad['source_sha256'] = '0'*64
        bad['inventory_sha256'] = digest({k: v for k, v in bad.items() if k != 'inventory_sha256'})
        packet = SimpleNamespace(inventory=bad, config={**self.config, 'mode': 'live',
            'word_source_review_context': preparation.WM_SOURCE_REVIEW_CONTEXT})
        with patch('core.testing_authority.reserve_role') as reserve, \
                patch('anthropic.Anthropic') as constructor, self.assertRaises(ValueError):
            preparation.call_role(packet, 'proposer', 999999999)
        reserve.assert_not_called()
        constructor.assert_not_called()


class TypedWMDeadlineTests(SimpleTestCase):
    """Deadline classifier/queue controls; synthetic SDK transport, no live API."""
    setUp=TypedRoleProductTests.setUp

    def wm_inventory(self):
        # A classifier fixture, not a claim that these synthetic questions were
        # discovered in WM. Actual WM source/mapping qualification is separate.
        value=inventory(self.path,protocol=TYPED_PROTOCOL)
        value['source_sha256']=preparation.WM_IOV_SHA256
        value['fields']=[{'key':'field_'+str(i),'label':'Synthetic question '+str(i),
            'kind':'text','locator':{'index':i}} for i in range(90)]
        value['questions']=[{'id':'q'+str(i),'text':'Synthetic question '+str(i),
            'protected_signing':False} for i in range(90)]
        value['anchors']=[{'id':'a'+str(i),'question_id':'q'+str(i),'kind':'existing',
            'field_key':'field_'+str(i),'locator':None,'protected_signing':False} for i in range(90)]
        value['stories']=[{**value['stories'][0],'part':'word/story'+str(i)+'.xml'} for i in range(16)]
        value['review_requirements']=[{'key':'claim_'+str(i)} for i in range(14)]
        value['document_backend']={'backend':'wm-iov-foss-v2','source_sha256':preparation.WM_IOV_SHA256,
            'mapping_digest':preparation.WM_MAPPING_DIGEST}
        value['issues']=[]
        self.seal(value)
        return value

    def seal(self,value):
        value['inventory_sha256']=digest({k:v for k,v in value.items() if k!='inventory_sha256'})

    def test_only_complete_source_qualified_typed_wm_gets_the_longer_deadline(self):
        value=self.wm_inventory()
        before=deepcopy(value)
        self.assertEqual(preparation.timing_limits(value),(120,360))
        self.assertEqual(preparation.role_limits(value),(16000,Decimal('0.30')))
        self.assertEqual(value,before)
        self.assertEqual(preparation.timing_limits(inventory(self.path,protocol=TYPED_PROTOCOL)),(60,180))
        for changed in ['protocol','source','backend','backend_source','mapping','backend_extra',
                'fields','anchors','questions','stories','review_requirements','kind','issues']:
            bad=deepcopy(value)
            if changed=='protocol':bad['protocol']='blank-word-preparation-v1'
            elif changed=='source':bad['source_sha256']='0'*64
            elif changed=='backend':bad['document_backend']['backend']='other'
            elif changed=='backend_source':bad['document_backend']['source_sha256']='0'*64
            elif changed=='mapping':bad['document_backend']['mapping_digest']='0'*64
            elif changed=='backend_extra':bad['document_backend']['unapproved']=True
            elif changed=='kind':bad['anchors'][0]['kind']='literal_blank'
            elif changed=='issues':bad['issues']=[{'code':'unsupported_structure','message':'Synthetic issue'}]
            else:bad[changed].pop()
            self.seal(bad)
            with self.subTest(changed=changed):
                self.assertEqual(preparation.timing_limits(bad),(60,180))
                old_limit=(6000 if changed=='protocol' else 12000 if changed in {
                    'mapping','backend_extra','questions','review_requirements','issues'} else 2500)
                self.assertEqual(preparation.role_limits(bad)[0],old_limit)
        historical=deepcopy(value);historical['protocol']='blank-word-preparation-v1';self.seal(historical)
        self.assertEqual(preparation.role_limits(historical),(6000,Decimal('0.16')))
        bad=deepcopy(value);bad['source_sha256']='0'*64
        with self.assertRaises(ValueError):preparation.timing_limits(bad)

    def receipt(self):
        return SimpleNamespace(response_path='',computed_cost=None,quote_maximum=Decimal('0.220432'),
            state='running',save=Mock())

    def test_actual_sdk_timeout_is_scoped_clamped_and_never_retries(self):
        import httpx
        from anthropic import Anthropic,transform_schema
        prompts=[]
        wm=self.wm_inventory();ordinary=inventory(self.path,protocol=TYPED_PROTOCOL)
        for value,remaining,expected_timeout in [(wm,300,120),(wm,43,43),(ordinary,170,60)]:
            for role in ['proposer','checker']:
                fixture,_=preparation.provider_role(value,role,self.config,1)
                before=deepcopy(value);sent=[];clients=[]
                def handle(request):
                    self.assertEqual(request.url.host,'synthetic-provider.invalid')
                    self.assertEqual(request.extensions['timeout']['read'],expected_timeout)
                    sent.append(json.loads(request.content))
                    return httpx.Response(200,headers={'request-id':'offline-request'},json={
                        'id':'offline-response','type':'message','role':'assistant','model':'offline-returned-model',
                        'content':[{'type':'text','text':json.dumps(fixture)}],'stop_reason':'end_turn',
                        'stop_sequence':None,'usage':{'input_tokens':100,'output_tokens':20}})
                def client(**kwargs):
                    self.assertEqual(kwargs,{'timeout':expected_timeout,'max_retries':0})
                    result=Anthropic(api_key='synthetic-offline-only',base_url='https://synthetic-provider.invalid',
                        http_client=http,**kwargs);clients.append(result);return result
                packet=SimpleNamespace(inventory=value,config={**self.config,'mode':'live'})
                with httpx.Client(transport=httpx.MockTransport(handle)) as http, \
                        patch('anthropic.Anthropic',side_effect=client), \
                        patch('core.testing_authority.identity',return_value=True), \
                        patch('core.testing_authority.reserve_role',return_value=self.receipt()), \
                        patch('core.preparation_campaign.provider_slot',return_value=nullcontext()), \
                        patch.object(preparation,'retain_response',return_value=('offline.json','a'*64)), \
                        patch.object(preparation.time,'monotonic',return_value=100):
                    result=preparation.call_role(packet,role,100+remaining)
                self.assertEqual(result,fixture);self.assertEqual(len(sent),1)
                self.assertEqual(clients[0].max_retries,0)
                self.assertEqual(sent[0]['max_tokens'],16000 if value is wm else 2500)
                self.assertEqual(sent[0]['output_config']['format']['schema'],
                    transform_schema(preparation.role_schema(role,TYPED_PROTOCOL)))
                prompts.append((role,sent[0]['system']))
                self.assertEqual(packet.config,{**self.config,'mode':'live'})
                self.assertEqual(value,before)
        for role in ['proposer','checker']:
            self.assertEqual(len({prompt for item,prompt in prompts if item==role}),1)

    def test_complete_wm_output_still_requires_every_source_metadata_row_and_story(self):
        value=self.wm_inventory()
        for role in ['proposer','checker']:
            response,_=preparation.provider_role(value,role,self.config,1)
            self.assertEqual(len(response['anchors']),90);self.assertEqual(len(response['stories']),16)
            self.assertEqual(validate_role(value,response,role),response)
            for mutation in ['missing_anchor','duplicate_anchor','missing_story','duplicate_story',
                    'label','destination','invented_unit','invented_choice','invented_type','lost_uncertainty']:
                bad=deepcopy(response)
                if mutation=='missing_anchor':bad['anchors'].pop()
                elif mutation=='duplicate_anchor':bad['anchors'][1]=deepcopy(bad['anchors'][0])
                elif mutation=='missing_story':bad['stories'].pop()
                elif mutation=='duplicate_story':bad['stories'][1]=deepcopy(bad['stories'][0])
                else:
                    key,changed={'label':('label','Different question'),'destination':('destination','other'),
                        'invented_unit':('unit','°C'),'invented_choice':('choices',['Yes','No']),
                        'invented_type':('data_type','date'),'lost_uncertainty':('uncertainty','none')}[mutation]
                    bad['anchors'][0]['metadata'][key]=changed
                with self.subTest(role=role,mutation=mutation),self.assertRaises(ValueError):
                    validate_role(value,bad,role)

    def test_actual_sdk_truncated_output_retains_known_response_usage_and_failure(self):
        import httpx
        from anthropic import Anthropic
        from django.core.exceptions import ValidationError
        value=self.wm_inventory();fixture,_=preparation.provider_role(value,'proposer',self.config,1)
        text=json.dumps(fixture,separators=(',',':'))
        # A synthetic equivalent of the retained live failure: all90 complete
        # anchors, four complete stories, then an unterminated fifth digest.
        needle='"content_sha256":"'+fixture['stories'][4]['content_sha256']
        fifth=text.index('"story_id":"'+fixture['stories'][4]['story_id']+'"')
        cut=text.index(needle,fifth)+len('"content_sha256":"')+14
        truncated=text[:cut]
        with self.assertRaises(json.JSONDecodeError):json.loads(truncated)
        receipt=self.receipt();packet=SimpleNamespace(inventory=value,config={**self.config,'mode':'live'})
        receipt.quote_maximum=preparation.maximum_cost(preparation.role_payload(value,'proposer'),packet.config,16000)
        quote=receipt.quote_maximum;sent=[];retained=[]
        def handle(request):
            self.assertEqual(json.loads(request.content)['max_tokens'],16000);sent.append(request)
            return httpx.Response(200,headers={'request-id':'offline-truncated-request'},json={
                'id':'offline-truncated-response','type':'message','role':'assistant','model':'offline-returned-model',
                'content':[{'type':'text','text':truncated}],'stop_reason':'max_tokens','stop_sequence':None,
                'usage':{'input_tokens':20566,'output_tokens':16000,
                         'output_tokens_details':{'thinking_tokens':1937}}})
        def retain(*args):
            retained.append(deepcopy(args[-1]));return 'offline-retained.json','c'*64
        with httpx.Client(transport=httpx.MockTransport(handle)) as http:
            client=Anthropic(api_key='synthetic-offline-only',base_url='https://synthetic-provider.invalid',
                http_client=http,timeout=120,max_retries=0)
            with patch('anthropic.Anthropic',return_value=client), \
                    patch('core.testing_authority.identity',return_value=True), \
                    patch('core.testing_authority.reserve_role',side_effect=[receipt,ValidationError('Role already reserved')]) as reserve, \
                    patch('core.preparation_campaign.provider_slot',return_value=nullcontext()), \
                    patch.object(preparation,'retain_response',side_effect=retain), \
                    patch.object(preparation.time,'monotonic',return_value=100):
                with self.assertRaises(preparation.DocumentError) as incomplete:
                    preparation.call_role(packet,'proposer',460)
                self.assertEqual(incomplete.exception.code,'output_limit')
                frozen=(receipt.state,receipt.quote_maximum,receipt.computed_cost,deepcopy(receipt.failure))
                with self.assertRaises(ValidationError):preparation.call_role(packet,'proposer',460)
        self.assertEqual(len(sent),1);self.assertEqual(reserve.call_args_list[0].args[-1],quote)
        self.assertEqual(receipt.state,'failed');self.assertEqual(receipt.quote_maximum,quote)
        self.assertEqual(receipt.computed_cost,(Decimal(20566)*3+Decimal(16000)*15)/1000000)
        self.assertEqual(receipt.returned_model,'offline-returned-model')
        self.assertEqual(receipt.provider_request_id,'offline-truncated-request')
        self.assertEqual(receipt.provider_resource_id,'offline-truncated-response')
        self.assertEqual(receipt.error_code,'output_limit')
        self.assertEqual(receipt.failure['error_class'],'DocumentError')
        self.assertEqual(receipt.failure['provider_stop_reason'],'max_tokens')
        self.assertTrue(receipt.failure['provider_invoked']);self.assertTrue(receipt.failure['response_retained'])
        self.assertTrue(receipt.failure['usage_established']);self.assertFalse(receipt.failure['automatic_retry'])
        self.assertEqual(retained[0]['response']['stop_reason'],'max_tokens')
        self.assertEqual(retained[0]['response']['content'][0]['text'],truncated)
        self.assertEqual(retained[-1],receipt.failure)
        self.assertEqual((receipt.state,receipt.quote_maximum,receipt.computed_cost,receipt.failure),frozen)

    def test_reservation_and_slot_delay_recheck_the_actual_send_window(self):
        import httpx
        from anthropic import Anthropic
        from django.core.exceptions import ValidationError
        value=inventory(self.path,protocol=TYPED_PROTOCOL)
        fixture,_=preparation.provider_role(value,'proposer',self.config,1)
        for delayed_at in ['reservation','slot']:
            for delay in [1,2,5]:
                with self.subTest(delayed_at=delayed_at,delay=delay):
                    clock={'now':100.0};receipt=self.receipt();sent=[];retained=[]
                    def reserve(*args):
                        if delayed_at=='reservation':clock['now']+=delay
                        return receipt
                    @contextmanager
                    def slot(*args):
                        if delayed_at=='slot':clock['now']+=delay
                        yield
                    def handle(request):
                        sent.append(request)
                        self.assertEqual(clock['now'],101)
                        self.assertEqual(request.extensions['timeout']['read'],3)
                        return httpx.Response(200,headers={'request-id':'offline-request'},json={
                            'id':'offline-response','type':'message','role':'assistant','model':'offline-returned-model',
                            'content':[{'type':'text','text':json.dumps(fixture)}],'stop_reason':'end_turn',
                            'stop_sequence':None,'usage':{'input_tokens':100,'output_tokens':20}})
                    def retain(*args):
                        retained.append(deepcopy(args[-1]));return 'offline-evidence.json','a'*64
                    packet=SimpleNamespace(inventory=value,config={**self.config,'mode':'live'})
                    with httpx.Client(transport=httpx.MockTransport(handle)) as http:
                        def client(**kwargs):
                            self.assertEqual(kwargs,{'timeout':3,'max_retries':0})
                            return Anthropic(api_key='synthetic-offline-only',
                                base_url='https://synthetic-provider.invalid',http_client=http,**kwargs)
                        with patch('anthropic.Anthropic',side_effect=client) as constructed, \
                                patch('core.testing_authority.identity',return_value=True), \
                                patch('core.testing_authority.reserve_role',side_effect=reserve) as reserved, \
                                patch('core.preparation_campaign.provider_slot',side_effect=slot), \
                                patch.object(preparation,'retain_response',side_effect=retain), \
                                patch.object(preparation.time,'monotonic',side_effect=lambda:clock['now']):
                            if delay==1:
                                self.assertEqual(preparation.call_role(packet,'proposer',104),fixture)
                                self.assertEqual(receipt.state,'succeeded');self.assertEqual(len(sent),1)
                            else:
                                with self.assertRaises(preparation.DocumentError) as error:
                                    preparation.call_role(packet,'proposer',104)
                                self.assertEqual(error.exception.code,'deadline')
                                self.assertEqual(sent,[]);constructed.assert_not_called()
                                self.assertEqual(receipt.state,'failed')
                                self.assertEqual(receipt.error_code,'deadline')
                                self.assertFalse(receipt.failure['provider_invoked'])
                                self.assertFalse(receipt.failure['response_retained'])
                                self.assertFalse(receipt.failure['usage_established'])
                                self.assertFalse(receipt.failure['automatic_retry'])
                                self.assertEqual(receipt.response_path,'');self.assertIsNone(receipt.computed_cost)
                                self.assertEqual(retained,[receipt.failure])
                                frozen=(receipt.state,receipt.quote_maximum,deepcopy(receipt.failure))
                                reserved.side_effect=ValidationError('Role already reserved')
                                with self.assertRaises(ValidationError):
                                    preparation.call_role(packet,'proposer',clock['now']+180)
                                self.assertEqual((receipt.state,receipt.quote_maximum,receipt.failure),frozen)
                                constructed.assert_not_called()
                    self.assertEqual(receipt.quote_maximum,Decimal('0.220432'))

    def test_reserved_deadline_failure_stops_worker_before_checker_and_candidate(self):
        value=inventory(self.path,protocol=TYPED_PROTOCOL);clock={'now':100.0};receipt=self.receipt()
        packet=SimpleNamespace(inventory={},config={**self.config,'mode':'live',
            'word_preparation_protocol':TYPED_PROTOCOL},
            template=SimpleNamespace(sha256=value['source_sha256'],department='synthetic-department',original_path='original.docx'),
            requested_by='synthetic-doctor',state='queued',pk='synthetic-prep',save=Mock(),error_code='')
        manager=Mock();manager.select_for_update.return_value.filter.return_value.order_by.return_value.first.return_value=packet
        manager.filter.return_value.exists.return_value=False
        def inspect(*args,**kwargs):clock['now']=276;return deepcopy(value)
        def reserve(*args):clock['now']+=5;return receipt
        with patch.object(preparation,'feature_enabled',return_value=True), \
                patch.object(preparation.transaction,'atomic',side_effect=nullcontext), \
                patch.object(preparation.TemplatePreparation,'objects',manager), \
                patch.object(preparation,'membership'),patch.object(preparation,'private_path',return_value=self.path), \
                patch.object(preparation,'inspect_preparation',side_effect=inspect), \
                patch('core.testing_authority.identity',return_value=True), \
                patch('core.testing_authority.reserve_role',side_effect=reserve) as reserved, \
                patch('core.preparation_campaign.provider_slot',return_value=nullcontext()), \
                patch.object(preparation,'retain_response',return_value=('offline-failure.json','b'*64)), \
                patch.object(preparation,'provider_role') as provider, \
                patch.object(preparation,'build_candidate') as build,patch.object(preparation,'audit'), \
                patch.object(preparation.time,'monotonic',side_effect=lambda:clock['now']):
            self.assertTrue(preparation.run_one())
        self.assertEqual(reserved.call_count,1);self.assertEqual(reserved.call_args.args[1],'proposer')
        provider.assert_not_called();build.assert_not_called()
        self.assertEqual(packet.state,'blocked');self.assertEqual(packet.error_code,'deadline')
        self.assertEqual(receipt.state,'failed');self.assertFalse(receipt.failure['provider_invoked'])
        self.assertEqual(receipt.quote_maximum,Decimal('0.220432'));self.assertIsNone(receipt.computed_cost)

    def worker_fixture(self,value,inspection_seconds,role_seconds,candidate_seconds):
        clock={'now':100.0};packet=SimpleNamespace(inventory={},config={'word_preparation_protocol':value['protocol']},
            template=SimpleNamespace(sha256=value['source_sha256'],department='synthetic-department',original_path='original.docx'),
            requested_by='synthetic-doctor',state='queued',pk='synthetic-prep',save=Mock(),error_code='')
        manager=Mock();manager.select_for_update.return_value.filter.return_value.order_by.return_value.first.return_value=packet
        manager.filter.return_value.exists.return_value=False
        def inspect(*args,**kwargs):clock['now']+=inspection_seconds;return deepcopy(value)
        calls=[]
        def role(prep,which,deadline,proposal=None):
            calls.append((which,deadline));clock['now']+=role_seconds
            return preparation.provider_role(value,which,self.config,1)[0]
        def candidate(*args):clock['now']+=candidate_seconds
        with patch.object(preparation,'feature_enabled',return_value=True), \
                patch.object(preparation.transaction,'atomic',side_effect=nullcontext), \
                patch.object(preparation.TemplatePreparation,'objects',manager), \
                patch.object(preparation,'membership'),patch.object(preparation,'private_path',return_value=self.path), \
                patch.object(preparation,'inspect_template'),patch.object(preparation,'inspect_preparation',side_effect=inspect), \
                patch.object(preparation,'call_role',side_effect=role), \
                patch.object(preparation,'build_candidate',side_effect=candidate) as build, \
                patch.object(preparation,'audit'),patch.object(preparation.time,'monotonic',side_effect=lambda:clock['now']):
            self.assertTrue(preparation.run_one())
        return packet,calls,build

    def test_worker_total_clock_includes_inventory_and_preserves_default_profile(self):
        for value,total in [(self.wm_inventory(),360),(inventory(self.path,protocol=TYPED_PROTOCOL),180),
                (inventory(self.path),180)]:
            before=deepcopy(value)
            packet,calls,build=self.worker_fixture(value,20,20,10)
            self.assertEqual(calls,[('proposer',100+total),('checker',100+total)])
            self.assertIn(packet.state,['review','guided']);build.assert_called_once()
            self.assertEqual(value,before)
            self.assertEqual(packet.config,{'word_preparation_protocol':value['protocol']})

    def test_inventory_time_cannot_buy_a_fresh_total_deadline_or_skip_candidate_reserve(self):
        packet,calls,build=self.worker_fixture(self.wm_inventory(),110,100,1)
        self.assertEqual(calls,[('proposer',460),('checker',460)])
        self.assertEqual(packet.state,'blocked');self.assertEqual(packet.error_code,'deadline')
        build.assert_not_called()  # Only50s remain; the existing60s reserve is mandatory.

    def test_candidate_completion_is_checked_against_the_original_total_deadline(self):
        packet,_,build=self.worker_fixture(self.wm_inventory(),20,100,141)
        build.assert_called_once()
        self.assertEqual(packet.state,'blocked');self.assertEqual(packet.error_code,'deadline')

    def test_timed_out_role_retains_unknown_quote_and_cannot_buy_a_replay(self):
        import httpx
        from anthropic import Anthropic,APITimeoutError
        from django.core.exceptions import ValidationError
        value=self.wm_inventory();receipt=self.receipt();sent=[]
        def handle(request):
            sent.append(request);raise httpx.ReadTimeout('synthetic timeout',request=request)
        packet=SimpleNamespace(inventory=value,config={**self.config,'mode':'live'})
        with httpx.Client(transport=httpx.MockTransport(handle)) as http:
            client=Anthropic(api_key='synthetic-offline-only',base_url='https://synthetic-provider.invalid',
                http_client=http,timeout=120,max_retries=0)
            with patch('anthropic.Anthropic',return_value=client), \
                    patch('core.testing_authority.identity',return_value=True), \
                    patch('core.testing_authority.reserve_role',side_effect=[receipt,ValidationError('Role already reserved')]), \
                    patch('core.preparation_campaign.provider_slot',return_value=nullcontext()), \
                    patch.object(preparation,'retain_response',return_value=('offline-failure.json','b'*64)), \
                    patch.object(preparation.time,'monotonic',return_value=100):
                with self.assertRaises(APITimeoutError):preparation.call_role(packet,'proposer',460)
                frozen=(receipt.state,receipt.quote_maximum,receipt.computed_cost,deepcopy(receipt.failure))
                with self.assertRaises(ValidationError):preparation.call_role(packet,'proposer',460)
        self.assertEqual(len(sent),1)
        self.assertEqual(receipt.state,'unknown');self.assertIsNone(receipt.computed_cost)
        self.assertEqual(receipt.quote_maximum,Decimal('0.220432'))
        self.assertFalse(receipt.failure['automatic_retry']);self.assertFalse(receipt.failure['response_retained'])
        self.assertTrue(receipt.failure['provider_invoked'])
        self.assertEqual(receipt.failure['error_class'],'APITimeoutError')
        self.assertEqual((receipt.state,receipt.quote_maximum,receipt.computed_cost,receipt.failure),frozen)

