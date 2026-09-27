"""Local parser syntax and actual fixture workflows; no live provider quality claim."""
from decimal import Decimal
import hashlib
import json
import os
from pathlib import Path
from unittest.mock import patch

from django.test import SimpleTestCase, TestCase, override_settings
from docx import Document

from core import generation as g, v3_mapping, v3_syntax
from core.jobs import run_one
from core.models import Encounter, Job, Usage
from core.security import private_path
from documents import word_v3 as w
from documents.test_question_boxes import box, data
from .test_question_box_workflow import selected_answers
from .test_v3_ordinary_notes import field, run
from .test_word_v3_workflow import V3, V3Fixture


A35 = ('What structured activities does the patient partake in including employment, study, or other interests? '
       'If the patient has no structured activity, how do they occupy their time and will this pose a problem '
       'from a risk perspective in the long term? Have they been supported in disclosure of their offending '
       'history (where appropriate?)')
ACTIVITY = 'The patient attends an art class every Monday.'
POSITIVES = (
    ('What activities does the patient attend?', ACTIVITY),
    ('History', 'The patient fell from a ladder yesterday and hurt his left wrist.'),
    ('Allergies', 'The patient developed a rash after amoxicillin.'),
    ('Social history', 'The patient smokes ten cigarettes a day.'),
    ('Mobility', 'The patient walks two blocks daily.'),
    ('How does the patient sleep?', 'The patient sleeps through the night.'),
)
FRAGMENTS = ('Warfarin 5 mg daily.', 'Right knee.', 'The patient attending art classes.',
             'The patient who attends art classes.', 'The patient a violin every Monday.',
             'The patient medications and allergies.', 'The patient attends an art class and.',
             'The patient attends an art class because.', 'The patient attends an art class with.')


class SentenceSyntax(SimpleTestCase):
    def test_materially_varied_complete_sentences_keep_exact_text_and_offsets(self):
        for label, quote in POSITIVES:
            with self.subTest(quote=quote):
                fs = [field('v3_1', label, 'long_text')]
                answer = run(quote, {label: [quote]}, fields=fs)[label]
                self.assertEqual((answer['status'], answer['text']), ('supported', quote))
                self.assertEqual(answer['evidence'], [{'quote': quote, 'start': 0, 'end': len(quote)}])
                self.assertEqual(g.governed_decisions(quote, fs, {'v3_1': {**answer, 'reviewed': True}}), [])

    def test_fragments_are_not_released_on_generation_or_stored_reader(self):
        fs = [field('v3_1', 'History', 'long_text')]
        for quote in FRAGMENTS:
            with self.subTest(quote=quote):
                answer = run(quote, {'History': [quote]}, fields=fs)['History']
                self.assertEqual((answer['status'], answer['text']), ('conflict', ''))
                old = {'text': quote, 'status': 'supported', 'reviewed': True,
                       'evidence': [{'quote': quote, 'start': 0, 'end': len(quote)}]}
                self.assertEqual(g.governed_decisions(quote, fs, {'v3_1': old}), ['History'])

    def test_no_actor_context_correction_side_or_shape_guard_is_replaced(self):
        cases = (
            ('Family history:\n' + ACTIVITY, ACTIVITY, 'Activities'),
            ('Her father attends an art class every Monday.', 'Her father attends an art class every Monday.', 'Activities'),
            ('The patient asks whether her father attends art classes.', 'The patient asks whether her father attends art classes.', 'Activities'),
            ('The patient hopes to attend art classes.', 'The patient hopes to attend art classes.', 'Activities'),
            ('The patient might attend art classes.', 'The patient might attend art classes.', 'Activities'),
            ('The patient could attend art classes.', 'The patient could attend art classes.', 'Activities'),
            ('The patient and her father attend art classes.', 'The patient and her father attend art classes.', 'Activities'),
            ('Her father enjoys music. He attends an art class every Monday.', 'He attends an art class every Monday.', 'Activities'),
            (ACTIVITY + ' Correction: the art class ended last month.', ACTIVITY, 'Activities'),
            ('If symptoms improve, the patient attends an art class.', 'If symptoms improve, the patient attends an art class.', 'Activities'),
            ('Does the patient attend an art class?', 'Does the patient attend an art class?', 'Activities'),
            ('The patient attends an art class?', 'The patient attends an art class?', 'Activities'),
            ('The patient attends an art class...', 'The patient attends an art class...', 'Activities'),
            ('The patient limps because his right knee hurts.', 'The patient limps because his right knee hurts.', 'Left knee'),
            ('Activities: the patient attends an art class.', 'Activities: the patient attends an art class.', 'Allergies'),
        )
        for notes, quote, label in cases:
            with self.subTest(notes=notes, label=label):
                fs = [field('v3_1', label, 'long_text')]
                if label != 'Activities': fs.append(field('v3_2', 'Activities', 'long_text'))
                answer = run(notes, {label: [quote]}, fields=fs)[label]
                self.assertEqual((answer['status'], answer['text']), ('conflict', ''))

    def test_dependency_failure_and_bounds_preserve_old_finite_path(self):
        with patch.object(v3_syntax, '_parser', return_value=None):
            self.assertEqual(v3_syntax.complete_ids([{'id': 1, 'quote': ACTIVITY}]), set())
            self.assertEqual(run(ACTIVITY, {'History': [ACTIVITY]}, fields=[field('v3_1', 'History')])['History']['status'], 'conflict')
            old = 'Patient presents with knee pain.'
            self.assertEqual(run(old, {'History': [old]}, fields=[field('v3_1', 'History')])['History']['status'], 'supported')
        self.assertEqual(v3_syntax.complete_ids([{'id': 1, 'quote': ACTIVITY * 30}]), set())
        units = [{'id': i, 'quote': ACTIVITY} for i in range(1, 66)]
        self.assertEqual(v3_syntax.complete_ids(units), set(range(1, 65)))

    def test_different_package_version_or_model_bytes_fail_closed(self):
        with patch.object(v3_syntax.importlib.metadata, 'version', return_value='0.0.0'):
            self.assertIsNone(v3_syntax._parser.__wrapped__())
        with patch.object(v3_syntax.hashlib, 'sha256') as digest:
            digest.return_value.hexdigest.return_value = '0' * 64
            self.assertIsNone(v3_syntax._parser.__wrapped__())

    def test_negative_activity_is_preserved_not_converted_to_positive(self):
        quote = 'The patient never attends an art class.'
        label = 'What activities does the patient attend?'
        answer = run(quote, {label: [quote]}, fields=[field('v3_1', label, 'long_text')])[label]
        self.assertEqual((answer['status'], answer['text']), ('supported', quote))

    def test_compound_review_preserves_printed_text_without_inventing_coverage(self):
        parts = v3_mapping.question_parts({'kind': 'long_text', 'label': A35})
        self.assertEqual(len(parts), 3)
        self.assertEqual(''.join(parts), A35)
        self.assertEqual(v3_mapping.question_parts({'kind': 'long_text', 'label': 'What activities?'}), [])
        for text in ('First？ Second？ trailing instruction.', 'First? Second (where appropriate?)'):
            self.assertEqual(''.join(v3_mapping.question_parts({'kind': 'long_text', 'label': text})), text)


@override_settings(**V3)
class SentenceWorkflow(V3Fixture, TestCase):
    def exercise(self, question, notes, *, correction=None, source=None, target_slot='b/tbl[0]/tr[0]/tc[0]'):
        if source is None:
            doc = Document(); doc.add_paragraph(question); box(doc)
            source = data(doc)
        template, version = self.uploaded(data=source, name='SYNTHETIC sentence ' + question[:80])
        target = next(f for f in version.fields if f['locator']['slot'] == target_slot)
        self.assertEqual(target['label'], question)
        self.confirm(template, version)
        self.assertEqual(self.new_encounter(version, notes=notes).status_code, 302)
        encounter = Encounter.objects.latest('created_at')
        answers = selected_answers(notes, version.fields, target['key'])
        self.assertEqual((answers[target['key']]['status'], answers[target['key']]['text']), ('supported', notes))
        mocked = (answers, Decimal('0'), {}, {}, 'synthetic-unit-selection', '', True)
        with patch('core.jobs.generation.generate', return_value=mocked) as provider:
            self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/generate/').status_code, 302)
            self.assertTrue(run_one())
        provider.assert_called_once()
        encounter.refresh_from_db()
        self.assertEqual(Job.objects.get(encounter=encounter).state, 'succeeded')
        self.assertTrue(Usage.objects.get(job__encounter=encounter).simulated)
        page = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertContains(page, notes)
        if question == A35:
            self.assertContains(page, 'Coverage has not been verified.')
            self.assertContains(page, 'including any parts left unanswered')
            for part in v3_mapping.question_parts(target): self.assertContains(page, part)
        else:
            self.assertNotContains(page, 'data-question-coverage')
        # A source-supported draft is still subject to the existing clinician review gate.
        self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'})
        encounter.refresh_from_db(); self.assertNotEqual(encounter.status, 'finalized')
        value = notes if correction is None else correction
        payload = {'revision': encounter.revision}
        for f in version.fields:
            payload['answer_' + f['key']] = value if f['key'] == target['key'] else ''
            payload['reviewed_' + f['key']] = 'on'
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/save/', payload).status_code, 302)
        encounter.refresh_from_db()
        answer = encounter.answers[target['key']]
        self.assertEqual(answer['status'], 'supported' if correction is None else 'edited')
        self.assertEqual(answer['reviewed_by'], str(self.user.pk))
        self.assertEqual(answer['text'], value)
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/finalize/',
            {'revision': encounter.revision, 'reviewed': 'on'}).status_code, 302)
        encounter.refresh_from_db(); self.assertEqual(encounter.status, 'finalized')
        response = self.client.get(f'/encounters/{encounter.pk}/export/')
        self.assertEqual(response.status_code, 200)
        exported = b''.join(response.streaming_content)
        self.assertEqual(hashlib.sha256(exported).hexdigest(), encounter.export_sha256)
        self.assertEqual(private_path(template.original_path).read_bytes(), source)
        target_id = 'word/document.xml#' + target_slot
        self.assertEqual(w.text_of(w.SlotIndex(w.open_package(exported)).by_id[target_id].el).strip(), value)
        if (capture := os.environ.get('QUILLMD_SYNTAX_ARTIFACTS')):
            folder = Path(capture) / hashlib.sha256(question.encode()).hexdigest()[:12]
            folder.mkdir(parents=True, exist_ok=False)
            (folder / 'draft-page.html').write_bytes(page.content)
            (folder / 'original.docx').write_bytes(source)
            (folder / 'exported.docx').write_bytes(exported)
            (folder / 'result.json').write_text(json.dumps({'question': question, 'notes': notes,
                'source_answer': answers[target['key']], 'final_answer': answer,
                'assisted_completion': correction is not None,
                'export_sha256': encounter.export_sha256, 'simulated': True}, indent=2))
        return encounter, source, exported

    def test_forced_old_supported_fragment_or_foreign_subject_stays_held_on_page_save_finalize(self):
        doc = Document(); doc.add_paragraph('What activities does the patient attend?'); box(doc)
        template, version = self.uploaded(data=data(doc), name='SYNTHETIC stored syntax guards')
        self.confirm(template, version)
        target = version.fields[0]
        for notes in ('Her father attends an art class every Monday.', 'The patient attending art classes.',
                      'The patient hopes to attend art classes.'):
            with self.subTest(notes=notes):
                self.new_encounter(version, notes=notes)
                e = Encounter.objects.latest('created_at')
                e.answers = {target['key']: {'text': notes, 'status': 'supported', 'reviewed': True,
                    'reviewed_by': str(self.user.pk), 'evidence': [{'quote': notes, 'start': 0, 'end': len(notes)}]}}
                e.save()
                page = self.client.get(f'/encounters/{e.pk}/')
                self.assertContains(page, 'data-held-candidate')
                self.assertContains(page, 'Source statement requiring review:')
                self.client.post(f'/encounters/{e.pk}/finalize/', {'revision': e.revision, 'reviewed': 'on'})
                e.refresh_from_db(); self.assertNotEqual(e.status, 'finalized')
                self.client.post(f'/encounters/{e.pk}/save/', {'revision': e.revision,
                    'answer_' + target['key']: '', 'reviewed_' + target['key']: 'on',
                    'held_shown_' + target['key']: '1', 'blank_shown_' + target['key']: '1'})
                e.refresh_from_db()
                self.client.post(f'/encounters/{e.pk}/finalize/', {'revision': e.revision, 'reviewed': 'on'})
                e.refresh_from_db(); self.assertNotEqual(e.status, 'finalized')

    def test_varied_complete_answers_need_review_but_no_answer_edit(self):
        # Full English questions keep the reusable question-box importer in scope.
        for question, notes in (
            POSITIVES[0], ('How does the patient sleep?', POSITIVES[5][1]),
            ('What distance does the patient walk?', POSITIVES[4][1]),
            ('How did the injury happen?', POSITIVES[1][1]),
        ):
            with self.subTest(question=question): self.exercise(question, notes)

    def test_compound_activity_supported_but_missing_disclosure_is_assisted(self):
        value = ACTIVITY + ' Support with disclosure of offending history is not documented in the source notes.'
        self.exercise(A35, ACTIVITY, correction=value)
