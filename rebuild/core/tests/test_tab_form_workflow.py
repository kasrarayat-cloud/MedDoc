"""Isolated authenticated original-form workflow; manual review creates no jobs."""

import hashlib
from io import BytesIO
import json
from pathlib import Path
import tempfile
from zipfile import ZipFile

from html.parser import HTMLParser

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from docx import Document

from core.billing import welcome_credit
from core.models import AuditEvent, Department, Encounter, Job, LedgerEntry, Membership, Organization, Patient, Template, Usage, User
from core.security import private_path
from documents.test_tab_form import ORIGINAL, SOURCE_SHA256
from .test_workflow import TEST_SETTINGS


class RenderedForm(HTMLParser):
    """What a browser would submit from one rendered form: enabled fields, ticked boxes only."""
    def __init__(self, action):
        super().__init__(); self.action = action; self.inside = False; self.data = {}; self.textarea = None; self.select = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "form": self.inside = attrs.get("action") == self.action
        if not self.inside or "disabled" in attrs or not attrs.get("name"): return
        if tag == "input":
            if attrs.get("type") == "checkbox":
                if "checked" in attrs: self.data[attrs["name"]] = "on"
            else: self.data[attrs["name"]] = attrs.get("value", "")
        elif tag == "textarea": self.textarea = attrs["name"]; self.data[self.textarea] = ""
        elif tag == "select": self.select = attrs["name"]
        elif tag == "option" and self.select and "selected" in attrs: self.data[self.select] = attrs.get("value", "")

    def handle_data(self, data):
        if self.textarea: self.data[self.textarea] += data

    def handle_endtag(self, tag):
        if tag == "form": self.inside = False
        if tag == "textarea": self.textarea = None
        if tag == "select": self.select = None


def rendered_post(html, action):
    form = RenderedForm(action); form.feed(html); return form.data


@override_settings(**TEST_SETTINGS)
class TabFormWorkflowTests(TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="quillmd-tab-http-")
        self.addCleanup(self.directory.cleanup)
        self.override = override_settings(MEDIA_ROOT=Path(self.directory.name) / "private")
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.user = User.objects.create_user(username="tab-alpha", email="tab-alpha@synthetic.invalid",
                                             first_name="Synthetic", last_name="Alpha Clinician",
                                             verified_at=timezone.now())
        self.other = User.objects.create_user(username="tab-beta", email="tab-beta@synthetic.invalid",
                                              first_name="Synthetic", last_name="Beta Clinician",
                                              verified_at=timezone.now())
        organization = Organization.objects.create(name="Synthetic tab-form practice")
        self.department = Department.objects.create(organization=organization, name="Synthetic endoscopy")
        for clinician in (self.user, self.other):
            Membership.objects.create(user=clinician, department=self.department, role="clinician")
        welcome_credit(self.user)
        self.ledger_before = LedgerEntry.objects.count()
        self.client.force_login(self.user)
        self.patient = Patient.objects.create(department=self.department, display_name="SYNTHETIC TAB PATIENT",
                                               reference="SYNTHETIC-TAB", created_by=self.user)
        self.original = ORIGINAL.read_bytes()
        self.assertEqual(hashlib.sha256(self.original).hexdigest(), SOURCE_SHA256)
        response = self.client.post("/templates/upload/", {"name": "Synthetic original colonoscopy form",
            "department": self.department.pk, "synthetic_only": "on",
            "file": SimpleUploadedFile("Colonoscopy_Operative_Template.docx", self.original)})
        self.assertEqual(response.status_code, 302, getattr(response, "context", None))
        self.template = Template.objects.get(name="Synthetic original colonoscopy form")
        self.version = self.template.versions.get(number=1)
        self.assertEqual(len(self.version.fields), 24)
        self.signature_key = next(f["key"] for f in self.version.fields if f.get("answer_policy") == "clinician_identity")

    def encounter(self, version=None):
        version = version or self.version
        response = self.client.post("/encounters/new/", {"patient": self.patient.pk,
            "template_version": version.pk, "notes": "PATIENT NAME: SYNTHETIC TAB PATIENT",
            "service_date": timezone.localdate().isoformat(), "synthetic_only": "on"})
        self.assertEqual(response.status_code, 302)
        return Encounter.objects.get(template_version=version)

    def save(self, encounter, signature):
        values = {field["key"]: f"Synthetic documented answer {index:02d}"
                  for index, field in enumerate(encounter.template_version.fields)}
        values[self.signature_key] = signature
        payload = {"revision": encounter.revision}
        for key, value in values.items():
            payload["answer_" + key] = value
            payload["reviewed_" + key] = "on"
        response = self.client.post(f"/encounters/{encounter.pk}/save/", payload)
        self.assertEqual(response.status_code, 302)
        encounter.refresh_from_db()
        self.assertTrue(all(answer["reviewed"] for answer in encounter.answers.values()))
        return values

    def finish(self, encounter):
        response = self.client.post(f"/encounters/{encounter.pk}/finalize/",
                                   {"revision": encounter.revision, "reviewed": "on"}, follow=True)
        self.assertEqual(response.status_code, 200)
        encounter.refresh_from_db()
        return response

    def assert_no_billable_work(self):
        self.assertEqual(Job.objects.count(), 0)
        self.assertEqual(Usage.objects.count(), 0)
        self.assertEqual(LedgerEntry.objects.count(), self.ledger_before)

    def test_mapping_manual_review_and_authenticated_original_format_download(self):
        mapping = self.client.get(f"/templates/{self.template.pk}/")
        self.assertContains(mapping, "Confirm that these consecutive answer lines belong to one question.")
        original_fields = self.version.fields
        labels = {field["key"]: field["label"] for field in original_fields}
        labels[original_fields[0]["key"]] = "Patient display name"
        response = self.client.post(f"/templates/{self.template.pk}/",
                                   {"mapping_version": self.version.pk, "labels": json.dumps(labels)})
        self.assertEqual(response.status_code, 302)
        revised = self.template.versions.get(number=2)
        self.version.refresh_from_db()
        self.assertEqual(self.version.fields, original_fields)
        self.assertEqual(revised.fields[0]["locator"], original_fields[0]["locator"])
        encounter = self.encounter(revised)
        page = self.client.get(f"/encounters/{encounter.pk}/")
        self.assertContains(page, 'id="source-notes-edit"', count=1)
        self.assertContains(page, 'data-generate-form', count=1)
        for field in revised.fields:
            self.assertContains(page, f'name="answer_{field["key"]}"', count=1)
        self.assertEqual(encounter.answers, {})
        values = self.save(encounter, self.user.get_full_name())
        self.finish(encounter)
        self.assertEqual(encounter.status, "finalized")
        self.assertEqual(encounter.finalized_by_id, self.user.pk)
        response = self.client.get(f"/encounters/{encounter.pk}/export/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])
        output = b"".join(response.streaming_content)
        self.assertEqual(hashlib.sha256(output).hexdigest(), encounter.export_sha256)
        document = Document(BytesIO(output))
        self.assertEqual(len(document.paragraphs), 31)
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("PATIENT NAME:", text)
        self.assertNotIn("Patient display name", text)
        for value in values.values():
            self.assertEqual(text.count(value), 1)
        with ZipFile(BytesIO(self.original)) as original, ZipFile(BytesIO(output)) as export:
            self.assertEqual(original.namelist(), export.namelist())
            for name in original.namelist():
                if name != "word/document.xml":
                    self.assertEqual(original.read(name), export.read(name), name)
        self.assertEqual(private_path(self.template.original_path).read_bytes(), self.original)
        original_download = self.client.get(f"/templates/{self.template.pk}/original/")
        self.assertEqual(b"".join(original_download.streaming_content), self.original)
        self.assertEqual(AuditEvent.objects.filter(action="encounter.finalize", target=str(encounter.pk)).count(), 1)
        self.assert_no_billable_work()

    def test_other_person_signature_is_refused_before_export(self):
        encounter = self.encounter()
        self.save(encounter, self.other.get_full_name())
        response = self.finish(encounter)
        self.assertContains(response, "signature must match your signed-in clinician name")
        self.assertEqual(encounter.status, "review")
        self.assertFalse(encounter.export_path)
        self.assertFalse(AuditEvent.objects.filter(action="encounter.finalize", target=str(encounter.pk)).exists())
        self.assert_no_billable_work()

    def test_named_finalizer_cannot_inherit_another_clinicians_signature_review(self):
        encounter = self.encounter()
        self.save(encounter, self.other.get_full_name())
        self.assertEqual(encounter.answers[self.signature_key]["reviewed_by"], str(self.user.pk))
        self.client.force_login(self.other)
        self.finish(encounter)
        self.assertEqual(encounter.status, "review")
        self.assertFalse(encounter.export_path)
        self.assertFalse(AuditEvent.objects.filter(action="encounter.finalize", target=str(encounter.pk)).exists())
        # Explicit review by the named finalizer repairs the refusal without new generation.
        self.save(encounter, self.other.get_full_name())
        self.finish(encounter)
        self.assertEqual(encounter.status, "finalized")
        self.assertEqual(encounter.finalized_by_id, self.other.pk)
        self.assert_no_billable_work()

    def blank_key(self):
        return next(f["key"] for f in self.version.fields if f["key"] != self.signature_key)

    def save_with_blank(self, encounter, signature):
        """Every row reviewed; one ordinary row decided "Leave blank — reviewed"."""
        values = self.save(encounter, signature)
        key = self.blank_key()
        payload = {"revision": encounter.revision, "answer_" + key: "", "blank_" + key: "on", "blank_shown_" + key: "1"}
        for other, value in values.items():
            payload["reviewed_" + other] = "on"
            if other != key: payload["answer_" + other] = value
        self.assertEqual(self.client.post(f"/encounters/{encounter.pk}/save/", payload).status_code, 302)
        encounter.refresh_from_db()
        self.assertEqual(encounter.answers[key]["status"], "left_blank")

    def test_second_clinician_saving_the_rendered_page_inherits_no_review(self):
        """A plain Save of the page as another clinician sees it re-stamps nobody's review as theirs."""
        encounter = self.encounter()
        self.save_with_blank(encounter, self.other.get_full_name())
        blank = self.blank_key()
        action = f"/encounters/{encounter.pk}/save/"
        # Control: the clinician who reviewed sees their own ticks, so the parser reads real boxes.
        own_page = self.client.get(f"/encounters/{encounter.pk}/")
        own = rendered_post(own_page.content.decode(), action)
        self.assertEqual({k for k in own if k.startswith("reviewed_")},
                         {"reviewed_" + f["key"] for f in encounter.template_version.fields})
        self.assertNotContains(own_page, "data-reviewed-by-other")
        self.client.force_login(self.other)
        page = self.client.get(f"/encounters/{encounter.pk}/")
        payload = rendered_post(page.content.decode(), action)
        self.assertFalse([k for k in payload if k.startswith("reviewed_")])
        self.assertEqual(payload.get("blank_" + blank), "on")          # the colleague's decision is shown, not re-made
        self.assertContains(page, "data-reviewed-by-other")
        self.assertContains(page, "data-signature-rule", count=1)
        self.assertContains(page, "exactly as <strong>Synthetic Beta Clinician</strong>")
        self.assertContains(page, "finalizing as <strong>Synthetic Beta Clinician</strong>")
        self.assertEqual(self.client.post(action, payload).status_code, 302)
        encounter.refresh_from_db()
        self.assertFalse(any(answer.get("reviewed") for answer in encounter.answers.values()))
        self.assertEqual((encounter.answers[blank]["status"], encounter.answers[blank]["blank_decision"]["by"]),
                         ("left_blank", str(self.user.pk)))
        withdrawn = AuditEvent.objects.filter(action="encounter.answers_saved", target=str(encounter.pk)).latest("created_at")
        self.assertEqual(len(withdrawn.details["withdrawn_reviews"]), len(encounter.template_version.fields))
        self.assertEqual(encounter.answers[self.signature_key]["text"], self.other.get_full_name())
        self.finish(encounter)
        self.assertEqual(encounter.status, "review")
        self.assertFalse(encounter.export_path)
        self.assertFalse(AuditEvent.objects.filter(action="encounter.finalize", target=str(encounter.pk)).exists())
        # The named clinician's own explicit review is what allows finalization.
        self.save_with_blank(encounter, self.other.get_full_name())
        self.finish(encounter)
        self.assertEqual((encounter.status, encounter.finalized_by_id), ("finalized", self.other.pk))
        self.assert_no_billable_work()

    def test_an_employee_save_then_a_colleagues_plain_save_cannot_pass_a_blank_row_to_finalization(self):
        encounter = self.encounter()
        self.save_with_blank(encounter, "")
        blank = self.blank_key()
        employee = User.objects.create_user(username="tab-employee", email="tab-employee@synthetic.invalid", verified_at=timezone.now())
        Membership.objects.create(user=employee, department=self.department, role="employee")
        action = f"/encounters/{encounter.pk}/save/"
        self.client.force_login(employee)
        self.assertEqual(self.client.post(action, rendered_post(self.client.get(f"/encounters/{encounter.pk}/").content.decode(), action)).status_code, 302)
        encounter.refresh_from_db()
        self.assertEqual((encounter.answers[blank]["status"], encounter.answers[blank]["reviewed"]), ("left_blank", False))
        self.client.force_login(self.other)
        page = self.client.get(f"/encounters/{encounter.pk}/")
        self.assertContains(page, "left blank — review needed")
        payload = rendered_post(page.content.decode(), action)
        for field in encounter.template_version.fields:          # the colleague ticks every row except the blank one
            if field["key"] != blank: payload["reviewed_" + field["key"]] = "on"
        self.assertEqual(self.client.post(action, payload).status_code, 302)
        encounter.refresh_from_db()
        self.assertFalse(encounter.answers[blank]["reviewed"])
        self.finish(encounter)
        self.assertEqual(encounter.status, "review")
        payload["reviewed_" + blank] = "on"; payload["revision"] = encounter.revision
        self.assertEqual(self.client.post(action, payload).status_code, 302)
        encounter.refresh_from_db()
        self.assertTrue(encounter.answers[blank]["reviewed"])
        self.finish(encounter)
        self.assertEqual((encounter.status, encounter.finalized_by_id), ("finalized", self.other.pk))

    def last_withdrawn(self, encounter):
        return AuditEvent.objects.filter(action="encounter.answers_saved", target=str(encounter.pk)).latest("created_at").details["withdrawn_reviews"]

    def test_the_withdrawal_audit_names_exactly_the_reviews_removed(self):
        encounter = self.encounter()
        keys = [f["key"] for f in self.version.fields]
        action = f"/encounters/{encounter.pk}/save/"
        payload = {"revision": encounter.revision, "reviewed_" + keys[0]: "on", "reviewed_" + keys[1]: "on"}
        payload.update({"answer_" + k: "Synthetic documented answer" for k in keys if k != self.signature_key})
        self.assertEqual(self.client.post(action, payload).status_code, 302)
        encounter.refresh_from_db()
        self.assertEqual(self.last_withdrawn(encounter), [])                     # a first review withdraws nothing
        own = rendered_post(self.client.get(f"/encounters/{encounter.pk}/").content.decode(), action)
        self.assertEqual(self.client.post(action, own).status_code, 302)
        encounter.refresh_from_db()
        self.assertEqual(self.last_withdrawn(encounter), [])                     # the reviewer's own re-save neither
        self.client.force_login(self.other)
        theirs = rendered_post(self.client.get(f"/encounters/{encounter.pk}/").content.decode(), action)
        theirs["reviewed_" + keys[2]] = "on"                                     # a row nobody had reviewed
        self.assertEqual(self.client.post(action, theirs).status_code, 302)
        encounter.refresh_from_db()
        self.assertEqual(self.last_withdrawn(encounter), sorted(keys[:2]))
        self.assertEqual(encounter.answers[keys[2]]["reviewed_by"], str(self.other.pk))

    def test_a_new_blank_decision_is_the_clinicians_own_review_without_a_second_tick(self):
        encounter = self.encounter()
        values = self.save(encounter, "")
        key = self.blank_key()
        payload = {"revision": encounter.revision, "answer_" + key: "", "blank_" + key: "on", "blank_shown_" + key: "1"}
        for other, value in values.items():
            if other != key: payload.update({"answer_" + other: value, "reviewed_" + other: "on"})
        self.assertEqual(self.client.post(f"/encounters/{encounter.pk}/save/", payload).status_code, 302)
        encounter.refresh_from_db()
        self.assertEqual((encounter.answers[key]["status"], encounter.answers[key]["reviewed"], encounter.answers[key]["reviewed_by"]),
                         ("left_blank", True, str(self.user.pk)))
        self.assertContains(self.client.get(f"/encounters/{encounter.pk}/"), "left blank — reviewed")
        self.finish(encounter)
        self.assertEqual(encounter.status, "finalized")

    def test_finalize_refuses_each_unmet_guard_on_a_form_that_can_finalize(self):
        encounter = self.encounter()
        self.save(encounter, "")
        url = f"/encounters/{encounter.pk}/finalize/"
        def refused(data, text):
            response = self.client.post(url, {"revision": encounter.revision, "reviewed": "on", **data}, follow=True)
            encounter.refresh_from_db()
            self.assertContains(response, text)
            self.assertEqual(encounter.status, "review"); self.assertFalse(encounter.export_path)
        refused({"reviewed": ""}, "Explicit clinical review confirmation is required.")
        refused({"revision": encounter.revision - 1}, "Document changed; reload before finalizing.")
        for state in ("queued", "running", "reconciliation"):
            with self.subTest(state=state):
                job = Job.objects.create(encounter=encounter, user=self.user, request_key="synthetic-guard-" + state, state=state)
                refused({}, "Resolve pending jobs before finalizing.")
                job.delete()
        original = private_path(self.template.original_path)
        original.write_bytes(self.original + b"tampered")
        refused({}, "Original integrity check failed.")
        original.write_bytes(self.original)
        self.finish(encounter)                                           # positive control: every guard met
        self.assertEqual(encounter.status, "finalized")
        self.assertTrue(encounter.export_path)

    def test_explicit_blank_signature_can_finalize_and_remains_blank(self):
        encounter = self.encounter()
        self.save(encounter, "")
        self.finish(encounter)
        self.assertEqual(encounter.status, "finalized")
        response = self.client.get(f"/encounters/{encounter.pk}/export/")
        document = Document(BytesIO(b"".join(response.streaming_content)))
        self.assertEqual(document.paragraphs[29].text, "Electronically signed by: \t")
        self.assert_no_billable_work()
