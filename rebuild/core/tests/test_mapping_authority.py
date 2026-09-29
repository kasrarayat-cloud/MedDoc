"""ONE structure-only authority (documents.word_v3.held_places) consulted by every writer and reader of a v3 mapping,
and model changes adopted by INDEPENDENCE (never by wire order). Synthetic BLANK forms and the retained offline f1
response only; native transports are never constructed and nothing leaves this process.

Covers the verified findings at 48aa5e8d (mapping-leaks): R1, R2, CAND-R1, CAND-R2-AND-LINES, MANUAL-PATH-LEAKS,
ENGINE-DETERMINISTIC-GAPS (a)-(f), ORDER-DEPENDENT-ADOPTION, the retained evidence/002 replay and its mixed
valid+invalid variant, the f5 negative controls, an f4 exam-grid multi-change independence case, the over-hold controls
and the protection of a mapping stored before the rule.
"""
import copy
import hashlib
import importlib.util
import io
import itertools
import json
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, TestCase, override_settings
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from core import generation, mapping_assistance as mapping, mapping_wire as wire, v3_mapping
from core.models import Encounter, Template, TemplateVersion
from core.security import private_path
from core.tab_review import clinician_identity, signature_label
from core.tests.test_mapping_assistance import SETTINGS
from core.tests.test_mapping_places import bordered_box, box_form
from core.tests.test_word_v3_workflow import V3, V3Fixture
from documents import word_v3
from documents.service import DocumentError, fill_template, inspect_template

BODY = 'word/document.xml#'
REPOSITORY = Path(__file__).resolve().parents[3]
FORMS = REPOSITORY / 'docs' / 'rebuild' / 'v1' / 'sonnet5-campaign' / 'diagnostics' / 'mapping' / 'build_forms.py'
RETAINED = Path(__file__).resolve().parent / 'fixtures' / 'mapping_live_f1'


def saved(doc):
    out = io.BytesIO(); doc.save(out); return out.getvalue()


def frozen(name):
    """A frozen, byte-deterministic diagnostic twin (no engine, no network): `name` is its builder in build_forms.py."""
    spec = importlib.util.spec_from_file_location('sonnet5_mapping_build_forms_authority', FORMS)
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
    return getattr(module, name)()


def r1_form(question, variant):
    doc = Document(); doc.add_paragraph('Participant ID: __________')
    doc.add_paragraph('Consultant review', style='Heading 2'); doc.add_paragraph(question)
    table = bordered_box(doc)
    if variant == 'ink':
        table.cell(0, 0).text = 'Use black ink'
    if variant == 'pict':
        run = OxmlElement('w:r'); run.append(OxmlElement('w:pict')); table.cell(0, 0).paragraphs[0]._p.append(run)
    return saved(doc)


def r2_form(variant, header='Signature'):
    doc = Document(); doc.add_paragraph('Clinic record', style='Heading 2')
    table = doc.add_table(rows=2, cols=2); table.cell(0, 0).text = 'Name'; table.cell(0, 1).text = header
    if variant == 'ink':
        table.cell(1, 1).text = 'Use black ink'
    if variant == 'pict':
        run = OxmlElement('w:r'); run.append(OxmlElement('w:pict')); table.cell(1, 1).paragraphs[0]._p.append(run)
    return saved(doc)


def initials_row(printed):
    doc = Document(); doc.add_paragraph('Plan: ______')
    table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = 'Clinician initials'
    if printed:
        table.cell(0, 1).text = printed
    return saved(doc)


def lines(*texts):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    return saved(doc)


def dependent_form():
    """b/p[5] 'Allergies' with no engine field once dropped: two additions compete for it (ORDER-DEPENDENT)."""
    return lines('Visit summary', 'Plan: ______', 'History: ______', 'Medication: ______', 'Signature: ______',
                 'Allergies')


class ModelPath:
    """The model-assisted mapping path, offline: inspect a synthetic original, compile the real request, and run the
    real validate_proposal on a hand-made compact response."""

    def setUp(self):
        deny = patch('anthropic.Anthropic', side_effect=AssertionError('Unmocked provider construction is forbidden.'))
        deny.start(); self.addCleanup(deny.stop)
        self.directory = tempfile.TemporaryDirectory(prefix='mapping-authority-')
        self.addCleanup(self.directory.cleanup)
        override = override_settings(MEDIA_ROOT=Path(self.directory.name)); override.enable(); self.addCleanup(override.disable)
        self.config = generation.configuration()
        getattr(v3_mapping, '_HELD_CACHE', {}).clear()   # (absent before the authority: red, not an error)

    def version(self, name, data, drop=(), fields=None):
        path = private_path(f'originals/{name}.docx'); path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(data)
        info = inspect_template(path, engine='v3')
        fields = copy.deepcopy(info['fields'] if fields is None else fields)
        fields = [f for f in fields if f['label'] not in drop and f['kind'] not in drop]
        return SimpleNamespace(template=SimpleNamespace(format='docx', original_path=f'originals/{name}.docx',
                                                        sha256=info['sha256']),
                               contract='v3', document_backend={}, fields=fields)

    def compile(self, version):
        compiled = mapping.compile_request(version, self.config)
        return compiled, json.loads(compiled['request']['messages'][0]['content'])

    @staticmethod
    def sid(payload, slot):
        return next(n for n, row in enumerate(payload['slots']) if row['slot'] in (BODY + slot, slot))

    @staticmethod
    def quote(payload, sid, text, roles=(0, 1, 2)):
        for position, (owner, role, _) in enumerate(payload['text_refs']):
            if owner == sid and role in roles:
                source = payload['slots'][sid][('text', 'heading', 'row_label')[role]]
                if text in source:
                    return [position, source.index(text), source.index(text) + len(text)]
        raise AssertionError(f'{text!r} is not printed at slot {sid}')

    def respond(self, payload, changes, replace=(), verdicts=None, dismiss=()):
        # A coherent response: the slot a change writes at is covered, unless the case names it in `dismiss` (a proposal
        # that also calls it not an answer contradicts itself - R2-F6).
        mapped = {slot for f in payload['fields'] for slot in f['slots']}
        targets = {c[1] for c in changes} - set(dismiss)
        covered = lambda n, row: n in targets or (n not in dismiss and v3_mapping.is_mapped(row['slot'], mapped))
        verdicts = verdicts or {}
        groups = {'k': [], 'r': list(replace), 'n': [], 'u': []}
        for n in range(len(payload['fields'])):
            if n not in replace:
                groups[verdicts.get(n, 'k')].append(n)
        return {'f': groups,
                's': {'c': [n for n, row in enumerate(payload['slots']) if covered(n, row)],
                      'n': [n for n, row in enumerate(payload['slots']) if not covered(n, row)],
                      'u': []},
                'x': [list(c) for c in changes], 'e': False}

    def addition(self, payload, slot, kind, label, roles=(0, 1, 2)):
        sid = self.sid(payload, slot)
        return [-1, sid, wire.KINDS.index(kind), *self.quote(payload, sid, label, roles), -1]

    @staticmethod
    def adopted(result):
        return [f for f in result['fields'] if f.get('origin') == 'model_assisted']

    def assertNothingAdopted(self, version, result, change):
        self.assertEqual(self.adopted(result), [], result['review'].get('rejected_changes'))
        self.assertEqual(result['fields'], version.fields)
        self.assertEqual(result['review']['changes'], [])
        rows = [r for r in result['review']['rejected_changes'] if r['wire'] == change]
        self.assertEqual(len(rows), 1, result['review']['rejected_changes'])
        self.assertTrue(result['review']['unresolved'])
        return rows[0]

    def outcome(self, version, compiled, response):
        result = mapping.validate_proposal(version, compiled, copy.deepcopy(response))
        review = {k: v for k, v in result['review'].items() if k != 'wire_response'}
        return mapping.canonical({'fields': result['fields'], 'review': review})

    def assertOrderIndependent(self, version, compiled, response):
        outcomes = set()
        for order in itertools.permutations(range(len(response['x']))):
            permuted = {**response, 'x': [response['x'][n] for n in order]}
            outcomes.add(self.outcome(version, compiled, permuted))
        self.assertEqual(len(outcomes), 1)
        return json.loads(outcomes.pop())

    def fills_into(self, version, field, answer='SYNTHETIC-ANSWER'):
        """What the engine's fill writes for a stored mapping holding `field` (refuses with DocumentError)."""
        out = Path(self.directory.name) / ('fill-%s.docx' % hashlib.sha256(json.dumps(field, sort_keys=True).encode()).hexdigest()[:10])
        return fill_template(private_path(version.template.original_path), [field], {field['key']: answer}, out)


@override_settings(**SETTINGS)
class R1R2AndLines(ModelPath, SimpleTestCase):
    """R1 / CAND-R1: a signing or consent question box is held whatever it holds; R2 / CAND-R2-AND-LINES: a signing
    column cell, a signature line, f5's signature line and initials box, and an initials row are held with the engine's
    own signing field present or dropped."""

    def test_R1_every_question_box_variant_is_held_and_never_filled(self):
        for question in ('Signed by the consultant:', 'I confirm the patient gave consent to proceed:',
                         'Record that the patient consents to proceed:'):
            for variant in ('empty', 'pict', 'ink'):
                version = self.version('r1', r1_form(question, variant))
                compiled, payload = self.compile(version)
                for slot, kind in (('b/tbl[0]/tr[0]/tc[0]/p[0]', 'long_text'), ('b/tbl[0]/tr[0]/tc[0]', 'long_text'),
                                   ('b/tbl[0]/tr[0]/tc[0]/p[0]', 'text')):
                    with self.subTest(question=question, variant=variant, slot=slot, kind=kind):
                        label = 'Use black ink' if variant == 'ink' and kind == 'text' else 'Consultant review'
                        roles = (0,) if label == 'Use black ink' else (1,)
                        change = self.addition(payload, slot, kind, label, roles)
                        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
                        self.assertNothingAdopted(version, result, change)
                # A stored mapping holding a text field in that box is refused by fill, whatever the box holds.
                index = word_v3.SlotIndex(word_v3.open_package(r1_form(question, variant)))
                field = {'key': 'stored', 'label': 'Consultant review', 'kind': 'long_text', 'context': '',
                         'locator': {'contract': 'v3', 'part': 'word/document.xml', 'slot': 'b/tbl[0]/tr[0]/tc[0]',
                                     'target': {'type': 'cell'},
                                     'fingerprint': word_v3.fingerprint(index.by_id[BODY + 'b/tbl[0]/tr[0]/tc[0]'].el)}}
                with self.subTest(question=question, variant=variant, fill=True), \
                        self.assertRaises(DocumentError) as refused:
                    self.fills_into(version, field)
                self.assertEqual(refused.exception.code, 'held_place')

    def test_R2_a_signing_column_cell_is_held_with_the_engine_field_kept_or_dropped(self):
        for header in ('Signature', 'Countersigned'):
            for variant in ('empty', 'pict', 'ink'):
                for drop in ((), ('signature',)):
                    with self.subTest(header=header, variant=variant, drop=drop):
                        version = self.version('r2', r2_form(variant, header), drop=drop)
                        compiled, payload = self.compile(version)
                        label, roles = ('Use black ink', (0,)) if variant == 'ink' else ('Clinic record', (1,))
                        change = self.addition(payload, 'b/tbl[0]/tr[1]/tc[1]/p[0]', 'text', label, roles)
                        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
                        self.assertNothingAdopted(version, result, change)

    def test_CAND_R2_lines_f5_signing_places_and_initials_rows_with_the_engine_field_dropped(self):
        cases = [('line', lines('Plan: ______', 'Signature: ______________'), 'b/p[1]', 'Signature: ______________', (0,), ('signature',)),
                 ('countersigned', lines('Plan: ______', 'Countersigned by: ______________'), 'b/p[1]', 'Countersigned by', (0,), ('signature',)),
                 ('signature-and-date', lines('Plan: ______', 'Consultant signature and date: ______'), 'b/p[1]',
                  'Consultant signature and date', (0,), ('signature', 'date')),
                 ('f5-line', frozen('f5_negative_controls'), 'b/p[5]', 'Signature of examining clinician', (0,), ('signature',)),
                 ('f5-initials', frozen('f5_negative_controls'), 'b/tbl[1]/tr[0]/tc[1]', 'Clinician initials', (2,), ('initials',)),
                 ('initial-here', initials_row('Initial here'), 'b/tbl[0]/tr[0]/tc[1]/p[0]', 'Initial here', (0,), ('initials',)),
                 ('black-ink', initials_row('Use black ink'), 'b/tbl[0]/tr[0]/tc[1]/p[0]', 'Use black ink', (0,), ('initials',))]
        for name, data, slot, label, roles, drop in cases:
            with self.subTest(name=name):
                version = self.version(name, data, drop=drop)
                compiled, payload = self.compile(version)
                change = self.addition(payload, slot, 'text', label, roles)
                result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
                self.assertNothingAdopted(version, result, change)

    def test_the_f5_declaration_heading_and_consent_declaration_are_held_and_its_guidance_is_not(self):
        # ENGINE-DETERMINISTIC-GAPS (d), (e): f5.C2 is write-protected by a new inspection; f5.C6 takes no model answer.
        version = self.version('f5', frozen('f5_negative_controls'))
        declaration = next(f for f in version.fields if f['locator']['slot'] == 'b/p[4]/ff[0]'
                           or any(b.get('slot') == 'b/p[4]/ff[0]' for b in f['locator']['target'].get('boxes', [])))
        self.assertTrue(declaration['write_protected'])
        self.assertNotIn(declaration['key'], [f['key'] for f in v3_mapping.answerable(version.fields)])
        compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/p[3]', 'text', 'Declaration')
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        self.assertNothingAdopted(version, result, change)
        # Control: the authority does not hold ordinary guidance (f5.C1 is a not-an-answer control graded elsewhere).
        guidance = self.addition(payload, 'b/p[1]', 'text', 'The reviewing clinician completes this page')
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [guidance]))
        self.assertEqual([f['locator']['slot'] for f in self.adopted(result)], ['b/p[1]'])

    def test_over_hold_controls_stay_answerable_on_the_model_path(self):
        data = lines('Vital signs: ______', 'Signs and symptoms: ______', 'Initial assessment: ______',
                     'Initial presentation: ______', 'Consent obtained: ______')
        for label, slot in (('Vital signs', 'b/p[0]'), ('Signs and symptoms', 'b/p[1]'), ('Initial assessment', 'b/p[2]'),
                            ('Initial presentation', 'b/p[3]'), ('Consent obtained', 'b/p[4]')):
            with self.subTest(label=label):
                version = self.version('over', data, drop=(label,))
                compiled, payload = self.compile(version)
                result = mapping.validate_proposal(version, compiled, self.respond(payload, [self.addition(payload, slot, 'text', label)]))
                self.assertEqual([f['label'] for f in self.adopted(result)], [label])


@override_settings(**SETTINGS)
class Independence(ModelPath, SimpleTestCase):
    """INDEPENDENCE-DESIGN: dependent changes are held together (never first-come), independent ones are adopted, and
    every permutation of the changes gives the same mapping and records."""

    def test_ORDER_DEPENDENT_two_additions_at_one_place_are_both_held_in_either_order(self):
        version = self.version('dep', dependent_form())
        compiled, payload = self.compile(version)
        text = self.addition(payload, 'b/p[5]', 'text', 'Allergies')
        long = self.addition(payload, 'b/p[5]', 'long_text', 'Allergies')
        result = self.assertOrderIndependent(version, compiled, self.respond(payload, [text, long]))
        self.assertEqual(result['fields'], version.fields)
        self.assertEqual(result['review']['changes'], [])
        rows = {tuple(r['wire']): r for r in result['review']['rejected_changes']}
        self.assertEqual(set(rows), {tuple(text), tuple(long)})
        for mine, other in ((text, long), (long, text)):
            self.assertEqual(rows[tuple(mine)]['with'], [mapping.change_id(other)])
            self.assertIn('D2', rows[tuple(mine)]['reason'])
        # Control: either addition alone is adopted.
        for change in (text, long):
            result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
            self.assertEqual([f['kind'] for f in self.adopted(result)], [wire.KINDS[change[2]]])

    def test_an_independent_valid_change_is_adopted_beside_a_held_one_in_every_order(self):
        version = self.version('mixed', lines('Plan: ______', 'History', 'Signature: ______'), drop=('Signature',))
        compiled, payload = self.compile(version)
        good = self.addition(payload, 'b/p[1]', 'long_text', 'History')
        harmful = self.addition(payload, 'b/p[2]', 'text', 'Signature: ______')
        result = self.assertOrderIndependent(version, compiled, self.respond(payload, [good, harmful]))
        self.assertEqual([(f['label'], f['locator']['slot']) for f in result['fields'] if f.get('origin') == 'model_assisted'],
                         [('History', 'b/p[1]')])
        self.assertEqual([r['wire'] for r in result['review']['rejected_changes']], [harmful])

    def smoker(self):
        doc = Document(); doc.add_paragraph('Smoker', style='Heading 2'); doc.add_paragraph('Smoker \u2610 Yes \u2610 No')
        doc.add_paragraph('Plan: ______')
        return self.version('smoker', saved(doc), drop=('Smoker', 'Plan'))

    def box(self, payload, glyph, text, roles):
        sid = self.sid(payload, 'b/p[1]/glyphs')
        return [-1, sid, wire.KINDS.index('checkbox'), *self.quote(payload, sid, text, roles), glyph]

    def assertHeldTogether(self, version, compiled, changes, dependence):
        result = self.assertOrderIndependent(version, compiled, self.respond(json.loads(compiled['request']['messages'][0]['content']), changes))
        self.assertEqual(result['fields'], version.fields)
        rows = {tuple(r['wire']): r for r in result['review']['rejected_changes']}
        self.assertEqual(set(rows), {tuple(c) for c in changes})
        for change in changes:
            self.assertEqual(rows[tuple(change)]['with'], sorted(mapping.change_id(c) for c in changes if c != change))
            self.assertIn(dependence, rows[tuple(change)]['reason'])
        for change in changes:              # control: each alone is adopted
            alone = mapping.validate_proposal(version, compiled, self.respond(json.loads(compiled['request']['messages'][0]['content']), [change]))
            self.assertEqual(len(self.adopted(alone)), 1, alone['review']['rejected_changes'])

    def test_D2_alone_two_changes_writing_at_one_place(self):
        version = self.smoker(); compiled, payload = self.compile(version)
        self.assertHeldTogether(version, compiled, [self.addition(payload, 'b/p[2]', 'text', 'Plan', (0,)),
                                                    self.addition(payload, 'b/p[2]', 'long_text', 'Smoker', (1,))], 'D2')

    def test_D4_alone_one_printed_question_claimed_by_two_boxes(self):
        version = self.smoker(); compiled, payload = self.compile(version)
        changes = [self.box(payload, 0, 'Smoker', (0,)), self.box(payload, 1, 'Smok', (0,))]
        self.assertHeldTogether(version, compiled, changes, 'D4')

    def test_D6_alone_one_question_asked_twice(self):
        version = self.smoker(); compiled, payload = self.compile(version)
        changes = [self.box(payload, 0, 'Smoker', (0,)), self.box(payload, 1, 'Smoker', (1,))]
        self.assertHeldTogether(version, compiled, changes, 'D6')

    def test_D1_two_replacements_of_one_field_name_each_other(self):
        version = self.version('d1', lines('Plan: ______', 'History: ______'))
        compiled, payload = self.compile(version)
        plan = next(n for n, f in enumerate(payload['fields']) if f['label'] == 'Plan')
        sid = self.sid(payload, 'b/p[0]')
        first = [plan, sid, wire.KINDS.index('long_text'), *self.quote(payload, sid, 'Plan', (0,)), -1]
        second = [plan, sid, wire.KINDS.index('text'), *self.quote(payload, sid, 'Plan', (0,)), -1]
        result = self.assertOrderIndependent(version, compiled, self.respond(payload, [first, second], replace=[plan]))
        self.assertEqual(result['fields'], version.fields)
        rows = {tuple(r['wire']): r for r in result['review']['rejected_changes']}
        self.assertEqual(rows[tuple(first)]['with'], [mapping.change_id(second)])
        self.assertIn('D1', rows[tuple(first)]['reason'])

    def test_V4_a_located_field_with_a_key_outside_the_contract_is_held(self):
        version = self.smoker(); compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/p[2]', 'text', 'Plan', (0,))
        real = mapping.locate_fields

        def extra(path, targets):
            located = real(path, targets)
            located['fields'][0]['layout_candidate'] = 'synthetic'
            return located
        with patch.object(mapping, 'locate_fields', side_effect=extra):
            result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        row = self.assertNothingAdopted(version, result, change)
        self.assertIn('outside the mapping contract', row['reason'])

    def test_D3_a_change_on_a_field_the_proposal_would_drop_is_held_for_the_clinician(self):
        version = self.version('removal', lines('Plan: ______', 'History: ______'))
        compiled, payload = self.compile(version)
        plan = next(n for n, f in enumerate(payload['fields']) if f['label'] == 'Plan')
        change = self.addition(payload, 'b/p[0]', 'long_text', 'Plan')
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change], verdicts={plan: 'n'}))
        row = self.assertNothingAdopted(version, result, change)
        self.assertEqual(row['depends_on'], [payload['fields'][plan]['key']])
        self.assertIn('only the clinician can remove', row['reason'])

    def test_D5_a_change_to_a_place_another_fields_relation_reads_is_held(self):
        version = self.version('relation', lines('Plan: ______', 'History: ______'))
        reader = next(f for f in version.fields if f['label'] == 'History')
        reader['source_instructions'] = [{'part': 'word/document.xml', 'slot': 'b/p[0]', 'fingerprint': '0' * 64,
                                          'text': 'Plan: ______'}]
        compiled, payload = self.compile(version)
        prior = next(n for n, f in enumerate(payload['fields']) if f['label'] == 'Plan')
        sid = self.sid(payload, 'b/p[0]')
        change = [prior, sid, wire.KINDS.index('long_text'), payload['text_refs'].index([prior, 4, 0]), 0, 4, -1]
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change], replace=[prior]))
        row = self.assertNothingAdopted(version, result, change)
        self.assertEqual(row['depends_on'], [reader['key']])
        self.assertIn('D5', row['reason'])

    def test_V2_the_resulting_mapping_is_read_again_by_the_authority(self):
        # Defence in depth: were the engine's locate ever to return a field inside the signature part of a mixed caption
        # line (a ranged hold, which the pre-locate place check leaves to the engine), the whole-mapping reading holds it.
        data = lines('Plan', 'Name: ______ Signature: ______ Date: ______')
        version = self.version('v2', data, drop=('signature', 'Plan'))
        compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/p[0]', 'text', 'Plan')
        entry = word_v3.held_places(word_v3.open_package(data))[BODY + 'b/p[1]']
        start = entry.ranges[0][0] + 12
        index = word_v3.SlotIndex(word_v3.open_package(data))
        forged = {'key': 'm3_forged00000', 'label': 'Plan', 'kind': 'text', 'context': '', 'origin': 'manual',
                  'locator': {'contract': 'v3', 'part': 'word/document.xml', 'slot': 'b/p[1]',
                              'target': {'type': 'span', 'start': start, 'end': start, 'shape': 'after_label'},
                              'fingerprint': word_v3.fingerprint(index.by_id[BODY + 'b/p[1]'].el)}}
        with patch.object(mapping, 'locate_fields', return_value={'sha256': version.template.sha256, 'fields': [forged]}):
            result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        row = self.assertNothingAdopted(version, result, change)
        self.assertEqual((row['reason'], row['check']), (mapping.SIGNING_REASON, 'mapping'))

    def test_V4_an_adopted_field_carries_only_the_locator_keys_and_originals_stay_byte_identical(self):
        version = self.version('keys', lines('Plan: ______', 'History'))
        compiled, payload = self.compile(version)
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [self.addition(payload, 'b/p[1]', 'long_text', 'History')]))
        added = self.adopted(result)
        self.assertEqual(len(added), 1)
        self.assertLessEqual(set(added[0]), mapping.FIELD_KEYS)
        kept = [f for f in result['fields'] if f.get('origin') != 'model_assisted']
        self.assertEqual([mapping.canonical(f) for f in kept], [mapping.canonical(f) for f in version.fields])

    def test_f4_exam_grid_multiple_changes_are_decided_independently_in_every_order(self):
        # D10: two independent relabels of pulse cells are adopted; two additions competing for one comment line are
        # held together; an option quoted from the shared header line is refused alone. 24 orders, one outcome.
        version = self.version('f4', frozen('f4_exam_grid'))
        compiled, payload = self.compile(version)
        by_label = {f['label']: n for n, f in enumerate(payload['fields'])}
        radial = by_label['Radial — Left']; femoral = by_label['Femoral — Left']

        def relabel(fid, slot, text):
            sid = self.sid(payload, slot)
            return [fid, sid, wire.KINDS.index('text'), *self.quote(payload, sid, text, (2,)), -1]
        first = relabel(radial, 'b/tbl[0]/tr[1]/tc[1]', 'Radial')
        second = relabel(femoral, 'b/tbl[0]/tr[2]/tc[1]', 'Femoral')
        compete_a = self.addition(payload, 'b/tbl[1]/tr[1]/tc[0]/p[0]', 'text', 'Pulse findings')
        compete_b = self.addition(payload, 'b/tbl[1]/tr[1]/tc[0]/p[0]', 'long_text', 'Pulse findings')
        response = self.respond(payload, [first, second, compete_a, compete_b], replace=[radial, femoral])
        result = self.assertOrderIndependent(version, compiled, response)
        adopted = sorted(f['label'] for f in result['fields'] if f.get('origin') == 'model_assisted')
        self.assertEqual(adopted, ['Femoral', 'Radial'])
        held = {tuple(r['wire']) for r in result['review']['rejected_changes']}
        self.assertEqual(held, {tuple(compete_a), tuple(compete_b)})
        # Control: each competitor alone is adopted (the hold is the dependence, not the place).
        for change in (compete_a, compete_b):
            alone = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
            self.assertEqual([f['label'] for f in self.adopted(alone)], ['Pulse findings'])

    def test_f4_an_option_quoted_from_the_shared_header_line_is_refused_alone(self):
        # CHECK.md M1: one such change refused the WHOLE proposal at the candidate; now only itself, in either order.
        version = self.version('f4m1', frozen('f4_exam_grid'))
        compiled, payload = self.compile(version)
        radial = next(n for n, f in enumerate(payload['fields']) if f['label'] == 'Radial — Left')
        sid = self.sid(payload, 'b/tbl[0]/tr[1]/tc[1]')
        relabel = [radial, sid, wire.KINDS.index('text'), *self.quote(payload, sid, 'Radial', (2,)), -1]
        line = self.sid(payload, 'b/tbl[1]/tr[2]/tc[0]/p[3]'); header = self.sid(payload, 'b/tbl[1]/tr[2]/tc[0]/p[2]')
        option = self.quote(payload, header, 'Normal', (0,))
        quoted = [-1, line, wire.KINDS.index('choice'), *self.quote(payload, line, 'Skin colour', (0,)), -1, *option]
        result = self.assertOrderIndependent(version, compiled, self.respond(payload, [relabel, quoted], replace=[radial]))
        self.assertEqual([f['label'] for f in result['fields'] if f.get('origin') == 'model_assisted'], ['Radial'])
        rows = result['review']['rejected_changes']
        self.assertEqual([(r['wire'], r['check']) for r in rows], [(quoted, 'wire')])
        self.assertIn('Source text belongs to another destination.', rows[0]['reason'])


@override_settings(**SETTINGS)
class RetainedEvidence(ModelPath, SimpleTestCase):
    """evidence/002, replayed offline: the retained response and its mixed valid+invalid variant."""

    def setUp(self):
        super().setUp()
        self.request = json.loads((RETAINED / 'request.json').read_text())
        self.response = json.loads(json.loads((RETAINED / 'provider-result.json').read_text())['content'][0]['text'])
        self.body = self.request['body_json']
        self.payload = json.loads(self.body['messages'][0]['content'])
        self.version_ = self.version('f1', frozen('f1_guidance_confirmation'))
        self.assertEqual(self.version_.template.sha256, self.payload['source_sha256'])
        self.compiled = {'request': self.body, 'request_sha256': mapping.digest(self.body),
                         'profile_sha256': self.payload['profile_sha256']}

    def run_(self, x):
        return mapping.validate_proposal(self.version_, self.compiled, {**copy.deepcopy(self.response), 'x': x})

    def test_the_retained_response_adopts_nothing_and_holds_both_changes_keyed(self):
        result = self.run_(self.response['x'])
        self.assertEqual(result['fields'], self.version_.fields)
        self.assertEqual({r['key'] for r in result['review']['rejected_changes']}, {'v3_1', 'v3_2'})
        self.assertEqual({r['key'] for r in result['review']['unresolved'] if r['key']}, {'v3_1', 'v3_2'})

    def test_the_mixed_variant_adopts_the_valid_fix_and_holds_the_invalid_change_in_either_order(self):
        bad, valid = [0, 3, 0, 3, 0, 79, -1], [1, 8, 2, 18, 0, 38, -1]
        outcomes = set()
        for x in ([bad, valid], [valid, bad]):
            result = self.run_(x)
            adopted = self.adopted(result)
            self.assertEqual([(f['label'], f['kind'], f['locator']['slot']) for f in adopted],
                             [('Date the lead agency was last informed', 'date', 'b/tbl[0]/tr[1]/tc[1]')])
            self.assertEqual([r['wire'] for r in result['review']['rejected_changes']], [bad])
            self.assertIn('Source text belongs to another destination.', result['review']['rejected_changes'][0]['reason'])
            self.assertIn(next(f for f in self.version_.fields if f['key'] == 'v3_1'), result['fields'])
            outcomes.add(mapping.canonical({'fields': result['fields'],
                                            'review': {k: v for k, v in result['review'].items() if k != 'wire_response'}}))
        self.assertEqual(len(outcomes), 1)

    def test_an_envelope_failure_still_refuses_the_whole_response(self):
        broken = copy.deepcopy(self.response); broken['f']['r'] = [0]
        with self.assertRaises(ValidationError):
            mapping.validate_proposal(self.version_, self.compiled, broken)


@override_settings(**SETTINGS)
class RefutedModelPath(ModelPath, SimpleTestCase):
    """The refutation round against be5fba45, on the model path: an answer place a caption governs (F1), a signing
    column the engine misread (F2), a declaration in every shape (F3) and a declaration block's empty line (R2-F2) are
    held; two selections on one option line (R2-F1) are dependent; one question in two contexts is two questions
    (R2-F5); a change the same proposal dismisses is held (R2-F6)."""

    def assertHeldChange(self, version, compiled, payload, change, **respond):
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change], **respond))
        row = self.assertNothingAdopted(version, result, change)
        return result, row

    def test_F1_F2_F3_additions_on_governed_answer_places_are_held(self):
        doc = Document(); doc.add_paragraph('Authorisation', style='Heading 2'); doc.add_paragraph('Signature of consultant:')
        doc.add_paragraph('_' * 30)
        version = self.version('f1m2', saved(doc), drop=('signature', 'Answer'))
        compiled, payload = self.compile(version)
        _, row = self.assertHeldChange(version, compiled, payload, self.addition(payload, 'b/p[2]', 'text', 'Authorisation', (1,)))
        self.assertIn(mapping.SIGNING_REASON, row['reason'])
        doc = Document(); doc.add_paragraph('Sign-off')
        table = doc.add_table(rows=3, cols=4)
        for r, row_ in enumerate((('', 'Name', 'Signature', 'Date'), ('Nurse', '', '', ''), ('Doctor', '', '', ''))):
            for c, value in enumerate(row_):
                if value:
                    table.cell(r, c).text = value
        data = saved(doc)
        engine = inspect_template_bytes(self, 'f2m1', data)
        version = self.version('f2m1', data, fields=[f for f in engine if f['locator']['slot'] != 'b/tbl[0]/tr[1]/tc[2]'])
        compiled, payload = self.compile(version)
        self.assertHeldChange(version, compiled, payload, self.addition(payload, 'b/tbl[0]/tr[1]/tc[2]', 'text', 'Nurse', (2,)))
        for line, label in (('I ____________________ (full name) consent to the procedure described above.', 'full name'),
                            ('To the best of my knowledge I confirm that the information above is correct: ____________',
                             'To the best of my knowledge')):
            with self.subTest(line=line):
                data = lines('Patient declaration', line)
                version = self.version('f3m', data, fields=[])
                compiled, payload = self.compile(version)
                self.assertHeldChange(version, compiled, payload, self.addition(payload, 'b/p[1]', 'text', label, (0,)))

    def test_R2_F2_an_addition_in_the_f5_declaration_block_after_its_table_is_held(self):
        version = self.version('f5r2', frozen('f5_negative_controls'))
        compiled, payload = self.compile(version)
        _, row = self.assertHeldChange(version, compiled, payload,
                                       self.addition(payload, 'b/p[6]', 'long_text', 'Declaration', (1,)))
        self.assertIn(mapping.HELD_REASONS['declaration'], row['reason'])

    def smoker_line(self):
        doc = Document(); doc.add_paragraph('Smoker'); doc.add_paragraph('Smoker \u2610 Yes \u2610 No')
        doc.add_paragraph('Plan: ______')
        data = saved(doc)
        fields = [f for f in inspect_template_bytes(self, 'smoker-line', data) if not f['locator']['slot'].startswith('b/p[1]')]
        return self.version('smoker-line', data, fields=fields)

    def box_at(self, payload, glyph, text):
        sid = self.sid(payload, 'b/p[1]/glyphs')
        return [-1, sid, wire.KINDS.index('checkbox'), *self.quote(payload, sid, text, (0,)), glyph]

    def test_R2_F1_two_selections_on_one_option_line_are_held_together_and_a_second_beside_a_kept_one(self):
        version = self.smoker_line(); compiled, payload = self.compile(version)
        yes, no = self.box_at(payload, 0, 'Yes'), self.box_at(payload, 1, 'No')
        result = self.assertOrderIndependent(version, compiled, self.respond(payload, [yes, no]))
        self.assertEqual(result['fields'], version.fields)
        rows = {tuple(r['wire']): r for r in result['review']['rejected_changes']}
        self.assertEqual(set(rows), {tuple(yes), tuple(no)})
        self.assertIn('D6 they select among the options printed on one line', rows[tuple(yes)]['reason'])
        for change in (yes, no):                   # control: each alone is adopted
            alone = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
            self.assertEqual(len(self.adopted(alone)), 1, alone['review']['rejected_changes'])
        # A second selection beside a KEPT one on the same line is held for the clinician.
        first = self.adopted(mapping.validate_proposal(version, compiled, self.respond(payload, [yes])))[0]
        kept = SimpleNamespace(**{**vars(version), 'fields': version.fields + [first]})
        compiled, payload = self.compile(kept)
        result = mapping.validate_proposal(kept, compiled, self.respond(payload, [self.box_at(payload, 1, 'No')]))
        self.assertEqual(result['fields'], kept.fields)
        row = result['review']['rejected_changes'][0]
        self.assertEqual(row['depends_on'], [first['key']])
        self.assertIn('already selects', row['reason'])

    def test_R2_F5_one_question_asked_again_is_held_alone_as_in_a_pair(self):
        # The f4 grid prints 'Describe any abnormality' in two rows: the engine's two fields already ask it. A model
        # change asking it again is held for the clinician whether it arrives alone or beside another (be5fba45 held
        # the pair and adopted each one alone).
        version = self.version('f4r2f5', frozen('f4_exam_grid'))
        compiled, payload = self.compile(version)
        asking = sorted(f['key'] for f in version.fields if f['label'] == 'Describe any abnormality')
        self.assertEqual(len(asking), 2)
        first = self.addition(payload, 'b/tbl[1]/tr[1]/tc[0]/p[1]', 'long_text', 'Describe any abnormality')
        second = self.addition(payload, 'b/tbl[1]/tr[5]/tc[0]/p[1]', 'long_text', 'Describe any abnormality')
        result = self.assertOrderIndependent(version, compiled, self.respond(payload, [first, second]))
        self.assertEqual(result['fields'], version.fields)
        for change in (first, second):
            alone = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
            row = self.assertNothingAdopted(version, alone, change)
            self.assertEqual(row['depends_on'], asking)
            self.assertIn('D6 it asks the question', row['reason'])
        # Control: a question the form does not already ask in that context is adopted.
        other = self.addition(payload, 'b/tbl[1]/tr[1]/tc[0]/p[0]', 'long_text', 'Pulse findings')
        self.assertEqual(len(self.adopted(mapping.validate_proposal(version, compiled, self.respond(payload, [other])))), 1)

    def test_R2_F6_a_change_the_same_proposal_dismisses_is_held(self):
        version = self.version('r2f6', lines('Plan: ______', 'History'))
        compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/p[1]', 'long_text', 'History')
        result, row = self.assertHeldChange(version, compiled, payload, change, dismiss=(self.sid(payload, 'b/p[1]'),))
        self.assertIn('marks its destination', row['reason'])
        self.assertTrue(any(r['slot'].endswith('#b/p[1]') for r in result['review']['unresolved']))
        coherent = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        self.assertEqual([f['label'] for f in self.adopted(coherent)], ['History'])


@override_settings(**SETTINGS)
class RoundThreeModelPath(ModelPath, SimpleTestCase):
    """Round 3 on the model path: a kept field another change only PROPOSES to replace is still kept (RR2-1), and the
    options of one printed question on several lines of a cell or on consecutive lines are one question's (RR2-2)."""

    def pulse(self):
        version = self.version('rr21', lines('Pulse findings: ______', 'Pulse findings', 'Plan: ______'))
        compiled, payload = self.compile(version)
        pulse = next(n for n, f in enumerate(payload['fields']) if f['label'] == 'Pulse findings')
        key = payload['fields'][pulse]['key'] if 'key' in payload['fields'][pulse] else \
            next(f['key'] for f in version.fields if f['label'] == 'Pulse findings')
        sid = self.sid(payload, 'b/p[0]')
        addition = self.addition(payload, 'b/p[1]', 'long_text', 'Pulse findings', (0,))
        same = [pulse, sid, wire.KINDS.index('text'), *self.quote(payload, sid, 'Pulse findings', (0,)), -1]
        relabel = [pulse, sid, wire.KINDS.index('text'), *self.quote(payload, sid, 'Pulse', (0,)), -1]
        failing = [pulse, sid, wire.KINDS.index('checkbox'), *self.quote(payload, sid, 'Pulse findings', (0,)), -1]
        return version, compiled, payload, pulse, key, addition, same, relabel, failing

    def test_RR2_1_a_kept_field_is_passed_over_only_when_its_replacement_is_adopted(self):
        version, compiled, payload, pulse, key, addition, same, relabel, failing = self.pulse()
        alone = mapping.validate_proposal(version, compiled, self.respond(payload, [addition]))
        self.assertEqual(self.assertNothingAdopted(version, alone, addition)['depends_on'], [key])
        # A replacement of that field that is NOT adopted - it fails alone, or it is held with a rival - leaves the
        # field kept, so the addition asking its question again is held exactly as alone (every order).
        for changes in ([addition, failing], [addition, same, relabel]):
            with self.subTest(changes=len(changes)):
                result = self.assertOrderIndependent(version, compiled, self.respond(payload, changes, replace=[pulse]))
                self.assertEqual(result['fields'], version.fields)
                row = next(r for r in result['review']['rejected_changes'] if r['wire'] == addition)
                self.assertEqual(row['depends_on'], [key])
                self.assertIn('D6 it asks the question', row['reason'])
                self.assertTrue(any(r['slot'].endswith('#b/p[1]') for r in result['review']['unresolved']))
        # Controls: each replacement alone is adopted; with the relabel ADOPTED no kept field asks the question.
        for change in (same, relabel):
            one = mapping.validate_proposal(version, compiled, self.respond(payload, [change], replace=[pulse]))
            self.assertEqual(len(self.adopted(one)), 1, one['review']['rejected_changes'])
        both = mapping.validate_proposal(version, compiled, self.respond(payload, [addition, relabel], replace=[pulse]))
        self.assertEqual(sorted(f['label'] for f in self.adopted(both)), ['Pulse', 'Pulse findings'])

    def box(self, payload, slot, text):
        sid = self.sid(payload, slot)
        return [-1, sid, wire.KINDS.index('checkbox'), *self.quote(payload, sid, text, (0,)), 0]

    def test_RR2_2_options_of_one_question_on_two_lines_of_a_cell_or_on_consecutive_lines_are_held_together(self):
        doc = Document(); doc.add_paragraph('Smoking history')
        table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = 'Current smoker'
        table.cell(0, 1).paragraphs[0].text = '\u2610 Yes'; table.cell(0, 1).add_paragraph('\u2610 No')
        doc.add_paragraph('Plan: ______')
        data = saved(doc)
        cell = self.version('rr22-cell', data, fields=[f for f in inspect_template_bytes(self, 'rr22-cell', data)
                                                        if '/tc[1]' not in f['locator']['slot']])
        data = lines('Smoker', '\u2610 Yes', '\u2610 No', 'Plan: ______')
        stacked = self.version('rr22-lines', data, fields=[f for f in inspect_template_bytes(self, 'rr22-lines', data)
                                                            if f['label'] == 'Plan'])
        for version, yes_slot, no_slot in ((cell, 'b/tbl[0]/tr[0]/tc[1]/p[0]/glyphs', 'b/tbl[0]/tr[0]/tc[1]/p[1]/glyphs'),
                                           (stacked, 'b/p[1]/glyphs', 'b/p[2]/glyphs')):
            with self.subTest(slot=yes_slot):
                compiled, payload = self.compile(version)
                yes, no = self.box(payload, yes_slot, 'Yes'), self.box(payload, no_slot, 'No')
                result = self.assertOrderIndependent(version, compiled, self.respond(payload, [yes, no]))
                self.assertEqual(result['fields'], version.fields)
                rows = {tuple(r['wire']): r for r in result['review']['rejected_changes']}
                self.assertEqual(set(rows), {tuple(yes), tuple(no)})
                self.assertIn('D6 they select among the options', rows[tuple(yes)]['reason'])
                for change in (yes, no):          # control: each alone is adopted
                    one = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
                    self.assertEqual(len(self.adopted(one)), 1, one['review']['rejected_changes'])
                # A second selection beside a KEPT one of the same question is held for the clinician.
                first = self.adopted(mapping.validate_proposal(version, compiled, self.respond(payload, [yes])))[0]
                kept = SimpleNamespace(**{**vars(version), 'fields': version.fields + [first]})
                compiled, payload = self.compile(kept)
                result = mapping.validate_proposal(kept, compiled, self.respond(payload, [self.box(payload, no_slot, 'No')]))
                self.assertEqual(result['fields'], kept.fields)
                self.assertEqual(result['review']['rejected_changes'][0]['depends_on'], [first['key']])


def inspect_template_bytes(case, name, data):
    path = private_path(f'originals/{name}.docx'); path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(data)
    return inspect_template(path, engine='v3')['fields']


class IdentityRule(SimpleTestCase):
    """ENGINE-DETERMINISTIC-GAPS (f), F5, R3-L1: the identity rule recognises the CLINICIAN as the signer - a clinical
    role with up to two qualifying words ('Signature of examining clinician', 'Consultant signature') - and never another
    person named in the same shape: that stays a held signature place, completed by hand."""

    def test_signer_phrases_are_identity_and_date_questions_are_not(self):
        for label in ('Signature of examining clinician', 'Consultant signature', 'Signed by', 'Electronically signed by',
                      'Signature of the physician', "Doctor's signature", "Supervising consultant's signature",
                      'Signed by the reviewing doctor'):
            with self.subTest(label=label):
                self.assertTrue(signature_label(label))
                self.assertTrue(clinician_identity({'label': label}))
        for label in ('Signature date', 'Date of signature', 'Time of signature', 'Place of signature', 'Signature (date)',
                      'Patient signature', 'Signature of patient', 'Witness signature', 'Signed by (witness)', 'Date signed'):
            with self.subTest(label=label):
                self.assertFalse(signature_label(label))

    def test_F5_R3_L1_another_persons_signature_is_never_the_clinicians_identity(self):
        actor = SimpleNamespace(pk=7, email='alpha@example.invalid', get_full_name=lambda: 'Synthetic Alpha Clinician')
        from core.tab_review import validate_identity
        for label in ("Participant's signature", 'Signature of participant', 'Signature of person giving consent',
                      "Mother's signature", 'Translator signature', 'Donor signature', 'Research subject signature',
                      'Chaperone signature', 'Second checker signature', 'Supervisor signature', 'Trainee signature',
                      'Signature of assessor', 'Signed by the student', 'Nurse signature'):
            with self.subTest(label=label):
                self.assertFalse(signature_label(label))
                self.assertFalse(clinician_identity({'label': label}))
                # validate_identity does not demand the finalizer's own name there (no false refusal).
                validate_identity([{'key': 'k', 'label': label}],
                                  {'k': {'text': 'Synthetic Mother', 'reviewed': True, 'reviewed_by': '7'}}, actor)


@override_settings(**V3, QUILLMD_MAPPING_ASSISTANCE=False)
class EveryWriterAndReader(V3Fixture, TestCase):
    """MANUAL-PATH-LEAKS and the stored-mapping protection, through the application's own views: a clinician's manual
    field, a correction, confirmation, generation, review save and finalize all consult the ONE authority."""

    def setUp(self):
        super().setUp()
        getattr(v3_mapping, '_HELD_CACHE', {}).clear()   # (absent before the authority: red, not an error)

    def post(self, template, data, version=None):
        return self.client.post(f'/templates/{template.pk}/', {'mapping_version': (version or self.latest(template)).pk, **data})

    def test_MANUAL_a_text_field_in_a_signing_box_is_refused_and_a_signing_kind_is_not(self):
        template, version = self.uploaded(data=r1_form('Signed by the consultant:', 'ink'), name='Synthetic R1')
        response = self.post(template, {'action': 'v3_add', 'slot': BODY + 'b/tbl[0]/tr[0]/tc[0]/p[0]', 'kind': 'long_text',
                                        'label': 'Consultant review'})
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, 'held place', status_code=400)
        self.assertEqual(self.latest(template).pk, version.pk)
        response = self.post(template, {'action': 'v3_add', 'slot': BODY + 'b/tbl[0]/tr[0]/tc[0]/p[0]', 'kind': 'signature',
                                        'label': 'Consultant review'})
        self.assertEqual(response.status_code, 302)

    def test_MANUAL_remove_the_signature_field_then_add_it_back_as_text_is_refused(self):
        for data, slot, label in ((r2_form('empty'), 'b/tbl[0]/tr[1]/tc[1]', 'Signature'),
                                  (lines('Plan: ______', 'Signature: ______________'), 'b/p[1]', 'Signature')):
            with self.subTest(slot=slot):
                template, version = self.uploaded(data=data, name='Synthetic manual ' + slot)
                signature = next(f for f in version.fields if f['kind'] == 'signature')
                self.assertEqual(self.post(template, {'action': 'v3_save', 'remove_' + signature['key']: 'on'}).status_code, 302)
                removed = self.latest(template)
                self.assertNotIn(signature['key'], [f['key'] for f in removed.fields])
                response = self.post(template, {'action': 'v3_add', 'slot': BODY + slot, 'kind': 'text', 'label': label})
                self.assertEqual(response.status_code, 400)
                self.assertEqual(self.latest(template).pk, removed.pk)

    def stored(self, data, name, field_slot, target, label='Consultant review', kind='long_text'):
        """A mapping STORED before the rule: the engine's fields plus an answerable field on a held place."""
        template, version = self.uploaded(data=data, name=name)
        index = word_v3.SlotIndex(word_v3.open_package(data))
        field = {'key': 'm3_stored0000', 'label': label, 'kind': kind, 'context': '', 'origin': 'manual',
                 'locator': {'contract': 'v3', 'part': 'word/document.xml', 'slot': field_slot, 'target': target,
                             'fingerprint': word_v3.fingerprint(index.by_id[BODY + field_slot].el)}}
        stored = TemplateVersion.objects.create(template=template, number=version.number + 1, fields=version.fields + [field],
                                                contract='v3', created_by=self.user)
        return template, stored, field

    def test_a_mapping_stored_before_the_rule_is_never_confirmed_generated_saved_into_or_exported(self):
        template, stored, field = self.stored(r1_form('Signed by the consultant:', 'pict'), 'Synthetic stored R1',
                                              'b/tbl[0]/tr[0]/tc[0]', {'type': 'cell'})
        # Confirmation is refused through the view.
        response = self.post(template, {'action': 'v3_confirm', 'checked_every_field': 'on'}, stored)
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, 'Consultant review', status_code=400)
        self.assertIsNone(TemplateVersion.objects.get(pk=stored.pk).confirmed)
        # A confirmation made before the rule does not help: generation, save and finalize still refuse.
        from core.models import MappingConfirmation
        MappingConfirmation.objects.create(version=stored, confirmed_by=self.user)
        stored.refresh_from_db()
        refusal = v3_mapping.generation_refusal(stored)
        self.assertTrue(refusal.startswith("This form's mapping puts an answer"), refusal)
        self.assertEqual(self.new_encounter(stored, notes='Participant ID: SYNTHETIC-1').status_code, 302)
        encounter = Encounter.objects.order_by('-created_at').first()
        payload = {'revision': encounter.revision, 'answer_' + field['key']: 'SYNTHETIC-ANSWER'}
        response = self.client.post(f'/encounters/{encounter.pk}/save/', payload)
        self.assertEqual(response.status_code, 400)
        encounter.refresh_from_db(); self.assertEqual(encounter.answers, {})
        ordinary = next(f for f in stored.fields if f['label'] == 'Participant ID')
        response = self.client.post(f'/encounters/{encounter.pk}/save/',
                                    {'revision': encounter.revision, 'answer_' + ordinary['key']: 'SYNTHETIC-1',
                                     **{'reviewed_' + f['key']: 'on' for f in stored.fields}})
        self.assertEqual(response.status_code, 302)
        encounter.refresh_from_db()
        response = self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'},
                                    follow=True)
        encounter.refresh_from_db()
        self.assertNotEqual(encounter.status, 'finalized')
        self.assertEqual(encounter.export_path, '')
        self.assertContains(response, 'Nothing was finalized')
        # And the engine's fill refuses the write itself, whatever calls it.
        with self.assertRaises(DocumentError) as refused:
            fill_template(private_path(template.original_path), stored.fields, {field['key']: 'SYNTHETIC-ANSWER'},
                          Path(self.directory.name) / 'stored-out.docx')
        self.assertEqual(refused.exception.code, 'held_place')

    def test_a_correction_never_turns_a_held_field_into_another_answer_kind(self):
        template, stored, field = self.stored(r2_form('pict'), 'Synthetic stored R2', 'b/tbl[0]/tr[1]/tc[1]', {'type': 'cell'},
                                              label='Clinic record', kind='text')
        response = self.post(template, {'action': 'v3_save', 'kind_' + field['key']: 'long_text'}, stored)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.latest(template).pk, stored.pk)

    def test_ENGINE_GAPS_a_new_upload_protects_consent_and_first_person_fields_and_confirms(self):
        # (a) consent statement lines; (b)/(c) Countersigned and 'signature and date' are signing kinds.
        data = lines('I confirm the patient gave consent to proceed: ______', 'Record that the patient consents to proceed: ______',
                     'Countersigned by: ______', 'Consultant signature and date: ______', 'Consent obtained: ______',
                     'Plan: ______')
        template, version = self.uploaded(data=data, name='Synthetic gaps')
        by = {f['label']: f for f in version.fields}
        self.assertTrue(by['I confirm the patient gave consent to proceed']['write_protected'])
        self.assertTrue(by['Record that the patient consents to proceed']['write_protected'])
        self.assertEqual(by['Countersigned by']['kind'], 'signature')
        self.assertEqual(by['Consultant signature and date']['kind'], 'signature')
        self.assertFalse(by['Consent obtained'].get('write_protected'))
        self.assertFalse(by['Plan'].get('write_protected'))
        asked, _, _ = generation.v3_generation_fields(version.fields)
        self.assertEqual(sorted(f['label'] for f in asked), ['Consent obtained', 'Plan'])
        self.confirm(template)
        self.assertEqual(v3_mapping.generation_refusal(self.latest(template)), '')

    def test_a_protected_declaration_checkbox_reviewed_blank_exports_exactly_as_printed(self):
        data = frozen('f5_negative_controls')
        template, version = self.uploaded(data=data, name='Synthetic f5')
        confirmed = self.confirm(template)
        self.assertEqual(self.new_encounter(confirmed, notes='Plan: rest').status_code, 302)
        encounter = Encounter.objects.order_by('-created_at').first()
        payload = {'revision': encounter.revision}
        for field in confirmed.fields:
            if field['kind'] in ('signature', 'initials'):
                continue
            payload['reviewed_' + field['key']] = 'on'
            if 'knee' in field['label']:
                payload['answer_' + field['key']] = 'SYNTHETIC finding'
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/save/', payload).status_code, 302)
        encounter.refresh_from_db()
        self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'})
        encounter.refresh_from_db()
        self.assertEqual(encounter.status, 'finalized')
        exported = b''.join(self.client.get(f'/encounters/{encounter.pk}/export/').streaming_content)
        before = word_v3.SlotIndex(word_v3.open_package(data)).by_id[BODY + 'b/p[4]'].el
        after = word_v3.SlotIndex(word_v3.open_package(exported)).by_id[BODY + 'b/p[4]'].el
        from lxml import etree
        self.assertEqual(etree.tostring(before, method='c14n'), etree.tostring(after, method='c14n'))

    def test_the_admin_grid_keeps_its_drug_dose_and_time_columns(self):
        doc = Document(); doc.add_paragraph('Medicines given', style='Heading 2')
        table = doc.add_table(rows=3, cols=4)
        for column, header in enumerate(('Drug', 'Dose', 'Time', 'Signature')):
            table.cell(0, column).text = header
        template, version = self.uploaded(data=saved(doc), name='Synthetic admin grid')
        grid = next(f for f in version.fields if f['kind'] == 'table_rows')
        self.assertNotIn('write_protected', grid)
        self.assertEqual(v3_mapping.held_columns(version), {grid['key']: [3]})
        self.confirm(template)
        self.assertEqual(self.new_encounter(self.latest(template), notes='Plan: rest').status_code, 302)
        encounter = Encounter.objects.order_by('-created_at').first()
        signed = self.client.post(f'/encounters/{encounter.pk}/save/', {'revision': encounter.revision,
                                  'answer_' + grid['key']: 'Drug A | 5 mg | 08:00 | SYNTHETIC-ANSWER'})
        self.assertEqual(signed.status_code, 400)
        ordinary = self.client.post(f'/encounters/{encounter.pk}/save/', {'revision': encounter.revision,
                                    'answer_' + grid['key']: 'Drug A | 5 mg | 08:00'})
        self.assertEqual(ordinary.status_code, 302)


@override_settings(**SETTINGS)
class ResolveAfterRemoval(V3Fixture, TestCase):
    """R2-F3: the clinician's area decisions are judged against the mapping they save."""

    def setUp(self):
        super().setUp()
        deny = patch('anthropic.Anthropic', side_effect=AssertionError('Unmocked provider construction is forbidden.'))
        deny.start(); self.addCleanup(deny.stop)
        getattr(v3_mapping, '_HELD_CACHE', {}).clear()

    def test_R2_F3_a_record_is_never_resolved_mapped_onto_a_field_another_decision_removes(self):
        template, version = self.uploaded(data=lines('Plan: ______', 'History: ______'), name='Synthetic resolve')
        compiled = mapping.compile_request(version, generation.configuration())
        payload = json.loads(compiled['request']['messages'][0]['content'])
        plan = next(n for n, f in enumerate(payload['fields']) if f['label'] == 'Plan')
        sid = next(n for n, row in enumerate(payload['slots']) if row['slot'].endswith('#b/p[0]'))
        mapped = {slot for f in payload['fields'] for slot in f['slots']}
        response = {'f': {'k': [n for n in range(len(payload['fields'])) if n != plan], 'r': [], 'n': [plan], 'u': []},
                    's': {'c': [n for n, r in enumerate(payload['slots']) if v3_mapping.is_mapped(r['slot'], mapped)],
                          'n': [n for n, r in enumerate(payload['slots']) if not v3_mapping.is_mapped(r['slot'], mapped)], 'u': []},
                    'x': [[-1, sid, wire.KINDS.index('long_text'), payload['text_refs'].index([sid, 0, 0]), 0, 4, -1]], 'e': False}
        result = mapping.validate_proposal(version, compiled, response)
        proposed = TemplateVersion.objects.create(template=template, number=2, fields=result['fields'], contract='v3',
                                                  created_by=self.user,
                                                  assistance={**result['review'], 'job_id': 'synthetic', 'base_version': str(version.pk)})
        records = proposed.assistance['unresolved']
        self.assertEqual(sorted(bool(r.get('key')) for r in records), [False, True])
        # The addition's record 'mapped' while the field that maps its place is removed: refused, nothing saved.
        body = {'action': 'resolve', 'mapping_version': str(proposed.pk)}
        for n, row in enumerate(records):
            body['resolution_%d' % n] = 'not_answer' if row.get('key') else 'mapped'
        refused = self.client.post(f'/templates/{template.pk}/assistance/', body)
        self.assertEqual(refused.status_code, 400)
        self.assertEqual(self.latest(template).pk, proposed.pk)
        # Control: keeping the field, the same 'mapped' decision is accepted.
        for n, row in enumerate(records):
            body['resolution_%d' % n] = 'mapped'
        self.assertEqual(self.client.post(f'/templates/{template.pk}/assistance/', body).status_code, 302)
        self.assertEqual(self.latest(template).number, 3)


@override_settings(**SETTINGS)
class ResolveNotAnAnswer(V3Fixture, TestCase):
    """RR2-3: 'not an answer' is judged against the mapping the decisions save, symmetric with 'mapped' (R2-F3)."""

    def setUp(self):
        super().setUp()
        deny = patch('anthropic.Anthropic', side_effect=AssertionError('Unmocked provider construction is forbidden.'))
        deny.start(); self.addCleanup(deny.stop)
        getattr(v3_mapping, '_HELD_CACHE', {}).clear()

    def test_RR2_3_an_area_an_adopted_field_still_maps_is_never_recorded_not_an_answer(self):
        template, version = self.uploaded(data=lines('Plan: ______', 'History'), name='Synthetic not-an-answer')
        compiled = mapping.compile_request(version, generation.configuration())
        payload = json.loads(compiled['request']['messages'][0]['content'])
        sid = next(n for n, row in enumerate(payload['slots']) if row['slot'].endswith('#b/p[1]'))
        mapped = {slot for f in payload['fields'] for slot in f['slots']}
        others = [n for n in range(len(payload['slots'])) if n != sid]
        response = {'f': {'k': list(range(len(payload['fields']))), 'r': [], 'n': [], 'u': []},
                    's': {'c': [n for n in others if v3_mapping.is_mapped(payload['slots'][n]['slot'], mapped)],
                          'n': [n for n in others if not v3_mapping.is_mapped(payload['slots'][n]['slot'], mapped)],
                          'u': [sid]},
                    'x': [[-1, sid, wire.KINDS.index('long_text'), payload['text_refs'].index([sid, 0, 0]), 0, 7, -1]],
                    'e': False}
        result = mapping.validate_proposal(version, compiled, response)
        self.assertEqual([f['label'] for f in result['fields'] if f.get('origin') == 'model_assisted'], ['History'])
        proposed = TemplateVersion.objects.create(template=template, number=2, fields=result['fields'], contract='v3',
                                                  created_by=self.user,
                                                  assistance={**result['review'], 'job_id': 'synthetic', 'base_version': str(version.pk)})
        records = proposed.assistance['unresolved']
        self.assertEqual([(bool(r.get('key')), r['slot'].endswith('#b/p[1]')) for r in records], [(False, True)])
        adopted = next(f['key'] for f in result['fields'] if f.get('origin') == 'model_assisted')
        body = {'action': 'resolve', 'mapping_version': str(proposed.pk), 'resolution_0': 'not_answer'}
        self.assertEqual(self.client.post(f'/templates/{template.pk}/assistance/', body).status_code, 302)
        saved_version = self.latest(template)
        self.assertEqual(saved_version.number, 3)
        self.assertNotIn(adopted, [f['key'] for f in saved_version.fields])
        self.assertEqual([f['key'] for f in saved_version.corrections['removed']], [adopted])
        self.assertEqual([f['key'] for f in saved_version.fields], [f['key'] for f in version.fields])

    def test_RR2_3_control_mapped_keeps_the_adopted_field(self):
        template, version = self.uploaded(data=lines('Plan: ______', 'History'), name='Synthetic mapped control')
        compiled = mapping.compile_request(version, generation.configuration())
        payload = json.loads(compiled['request']['messages'][0]['content'])
        sid = next(n for n, row in enumerate(payload['slots']) if row['slot'].endswith('#b/p[1]'))
        mapped = {slot for f in payload['fields'] for slot in f['slots']}
        others = [n for n in range(len(payload['slots'])) if n != sid]
        response = {'f': {'k': list(range(len(payload['fields']))), 'r': [], 'n': [], 'u': []},
                    's': {'c': [n for n in others if v3_mapping.is_mapped(payload['slots'][n]['slot'], mapped)],
                          'n': [n for n in others if not v3_mapping.is_mapped(payload['slots'][n]['slot'], mapped)],
                          'u': [sid]},
                    'x': [[-1, sid, wire.KINDS.index('long_text'), payload['text_refs'].index([sid, 0, 0]), 0, 7, -1]],
                    'e': False}
        result = mapping.validate_proposal(version, compiled, response)
        proposed = TemplateVersion.objects.create(template=template, number=2, fields=result['fields'], contract='v3',
                                                  created_by=self.user,
                                                  assistance={**result['review'], 'job_id': 'synthetic', 'base_version': str(version.pk)})
        body = {'action': 'resolve', 'mapping_version': str(proposed.pk), 'resolution_0': 'mapped'}
        self.assertEqual(self.client.post(f'/templates/{template.pk}/assistance/', body).status_code, 302)
        self.assertEqual(self.latest(template).fields, result['fields'])


@override_settings(**{**V3, 'QUILLMD_WORD_ENGINE': 'legacy'})
class LegacyFill(V3Fixture, TestCase):
    """The legacy (non-v3) Word fill consults the same authority where it writes into the customer's file."""

    def test_legacy_fill_refuses_a_held_row_and_keeps_the_identity_and_ordinary_rows(self):
        doc = Document(); table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = 'Plan'; table.cell(1, 0).text = 'Countersigned by'; table.cell(2, 0).text = 'Electronically signed by'
        with tempfile.TemporaryDirectory(prefix='legacy-authority-') as directory:
            path = Path(directory) / 'legacy.docx'; path.write_bytes(saved(doc))
            fields = inspect_template(path, engine='legacy')['fields']
            by = {f['label']: f for f in fields}
            self.assertEqual(set(by), {'Plan', 'Countersigned by', 'Electronically signed by'})
            fill_template(path, fields, {by['Plan']['key']: 'Rest.'}, Path(directory) / 'plan.docx')
            fill_template(path, fields, {by['Electronically signed by']['key']: 'Synthetic Alpha Clinician'},
                          Path(directory) / 'identity.docx', identity_keys=[by['Electronically signed by']['key']])
            with self.assertRaises(DocumentError) as refused:
                fill_template(path, fields, {by['Electronically signed by']['key']: 'Synthetic Alpha Clinician'},
                              Path(directory) / 'no-identity.docx')
            self.assertEqual(refused.exception.code, 'held_place')
            with self.assertRaises(DocumentError) as refused:
                fill_template(path, fields, {by['Countersigned by']['key']: 'SYNTHETIC-ANSWER'}, Path(directory) / 'held.docx',
                              identity_keys=[by['Electronically signed by']['key']])
            self.assertEqual(refused.exception.code, 'held_place')
            self.assertFalse((Path(directory) / 'held.docx').exists())

    def test_F5_legacy_fill_never_writes_the_clinicians_name_into_another_persons_signature_row(self):
        for label in ("Participant's signature", 'Signature of person giving consent', 'Consultant signature'):
            with self.subTest(label=label):
                doc = Document(); doc.add_paragraph('Research consent record')
                table = doc.add_table(rows=2, cols=2); table.cell(0, 0).text = 'Participant name'; table.cell(1, 0).text = label
                with tempfile.TemporaryDirectory(prefix='legacy-identity-') as directory:
                    path = Path(directory) / 'legacy.docx'; path.write_bytes(saved(doc))
                    fields = inspect_template(path, engine='legacy')['fields']
                    row = next(f for f in fields if f['label'] == label)
                    keys = [f['key'] for f in fields if clinician_identity(f)]
                    if label == 'Consultant signature':
                        self.assertEqual(keys, [row['key']])
                        fill_template(path, fields, {row['key']: 'Synthetic Alpha Clinician'}, Path(directory) / 'out.docx',
                                      identity_keys=keys)
                        continue
                    self.assertEqual(keys, [])
                    with self.assertRaises(DocumentError) as refused:
                        fill_template(path, fields, {row['key']: 'Synthetic Alpha Clinician'}, Path(directory) / 'out.docx',
                                      identity_keys=keys)
                    self.assertEqual(refused.exception.code, 'held_place')

    def test_RR5_a_parenthetical_naming_another_signer_is_never_the_clinicians_identity(self):
        for label in ('Signature (participant)', 'Signature (mother)', 'Signature (donor)', 'Signed by (research participant)',
                      'Signature (person giving consent)', 'Signature (next-of-kin)', 'Signed by (MD)'):
            with self.subTest(label=label):
                doc = Document(); doc.add_paragraph('Research consent record')
                table = doc.add_table(rows=2, cols=2); table.cell(0, 0).text = 'Participant name'; table.cell(1, 0).text = label
                with tempfile.TemporaryDirectory(prefix='legacy-identity-') as directory:
                    path = Path(directory) / 'legacy.docx'; path.write_bytes(saved(doc))
                    fields = inspect_template(path, engine='legacy')['fields']
                    row = next(f for f in fields if f['label'] == label)
                    keys = [f['key'] for f in fields if clinician_identity(f)]
                    if label == 'Signed by (MD)':                  # control: the clinician's own row
                        self.assertEqual(keys, [row['key']])
                        fill_template(path, fields, {row['key']: 'Synthetic Alpha Clinician'}, Path(directory) / 'out.docx',
                                      identity_keys=keys)
                        continue
                    self.assertFalse(signature_label(label))
                    self.assertEqual(keys, [])
                    with self.assertRaises(DocumentError) as refused:
                        fill_template(path, fields, {row['key']: 'Synthetic Alpha Clinician'}, Path(directory) / 'out.docx',
                                      identity_keys=keys)
                    self.assertEqual(refused.exception.code, 'held_place')

    def test_F6_legacy_tab_line_under_a_signing_heading_is_held(self):
        for heading in ('Patient signature', 'Sign & date'):
            with self.subTest(heading=heading):
                doc = Document(); doc.add_paragraph('Consent form'); doc.add_paragraph(heading, style='Heading 2')
                line = OxmlElement('w:p'); run = OxmlElement('w:r'); properties = OxmlElement('w:rPr')
                underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); properties.append(underline)
                run.append(properties); run.append(OxmlElement('w:tab')); line.append(run)
                doc.element.body.insert(len(doc.element.body) - 1, line)
                with tempfile.TemporaryDirectory(prefix='legacy-heading-') as directory:
                    path = Path(directory) / 'legacy.docx'; path.write_bytes(saved(doc))
                    fields = inspect_template(path, engine='legacy')['fields']
                    self.assertTrue(fields)
                    with self.assertRaises(DocumentError) as refused:
                        fill_template(path, fields, {fields[0]['key']: 'SYNTHETIC-ANSWER'}, Path(directory) / 'out.docx',
                                      identity_keys=[f['key'] for f in fields if clinician_identity(f)])
                    self.assertEqual(refused.exception.code, 'held_place')
