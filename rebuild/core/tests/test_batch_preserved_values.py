"""Synthetic constraints checked against actual preserved/written Word values; no DB."""
import hashlib
import io
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, override_settings
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from core import generation_batch as b, v3_mapping
from documents import word_v3 as w
from documents.test_word_v3 import _sdt, ff_text_field, ff_field


def raw(doc):
    stream = io.BytesIO(); doc.save(stream); return stream.getvalue()


def legacy(p, value='', default=''):
    ff_text_field(p, 'Priority')
    begin = next(p._p.iter(qn('w:fldChar')))
    if default:
        node = OxmlElement('w:default'); node.set(qn('w:val'), default)
        begin.find(qn('w:ffData')).find(qn('w:textInput')).append(node)
    if value:
        r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = value; r.append(t)
        p._p[-1].addprevious(r)


def dropdown(p, value, *, placeholder=False):
    control = _sdt('dropdown', options=['1', '2', '3'], placeholder=value)
    if not placeholder:
        pr = control.find(qn('w:sdtPr')); pr.remove(pr.find(qn('w:showingPlcHdr')))
    p._p.append(control)


class PreservedConstraintValues(SimpleTestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        media = override_settings(MEDIA_ROOT=Path(self.temp.name)); media.enable(); self.addCleanup(media.disable)
        deny = patch('socket.socket.connect', side_effect=AssertionError('Network forbidden'))
        deny.start(); self.addCleanup(deny.stop)

    def fixture(self, doc, *, fields=None):
        data = raw(doc); fields = fields if fields is not None else w.inspect(data)['fields']
        self.assertEqual(len(fields), 2)
        Path(self.temp.name, 'original.docx').write_bytes(data)
        sha = hashlib.sha256(data).hexdigest()
        manifest = {'contract': b.CONSTRAINTS, 'source_sha256': sha, 'fields_sha256': b.digest(fields),
                    'coverage': 'complete', 'provenance': 'synthetic_fixture', 'evidence_sha256': '1' * 64,
                    'instruction_references': [{'id': 'rule', 'part': 'word/document.xml', 'slot': 'b/p[0]',
                                               'sha256': hashlib.sha256(b'Priorities must be unique.').hexdigest()}],
                    'rules': [{'kind': 'unique_values', 'source_id': 'rule', 'keys': [f['key'] for f in fields],
                               'allowed_values': ['1', '2', '3']}]}
        version = SimpleNamespace(fields=fields, template=SimpleNamespace(sha256=sha, original_path='original.docx'),
                                  generation_constraints=manifest)
        return data, version

    def answers(self, version, first='', second=''):
        return {f['key']: {'text': value, 'reviewed': True, 'status': 'supported' if value else 'left_blank'}
                for f, value in zip(version.fields, (first, second))}

    def values(self, data, version, answers):
        # Native XML results in the actual export, independently of constraint readers.
        output = w.fill(data, version.fields, v3_mapping.fill_answers(version.fields, answers))
        root = w.open_package(output).tree('word/document.xml')
        return [w.text_of(p) for p in root.iter(qn('w:p'))]

    def test_preserved_native_scalar_duplicate_refuses_and_explicit_change_succeeds(self):
        doc = Document(); doc.add_paragraph('Priorities must be unique.')
        legacy(doc.add_paragraph('Priority A: '), '1')
        legacy(doc.add_paragraph('Priority B: '))
        data, version = self.fixture(doc)
        bad = self.answers(version, '', '1')
        with self.assertRaisesMessage(ValidationError, 'priority'): b.check_answers(version, bad)
        self.assertEqual(self.values(data, version, bad)[1:], ['Priority A: 1', 'Priority B: 1'])
        good = self.answers(version, '2', '1'); b.check_answers(version, good)
        self.assertEqual(self.values(data, version, good)[1:], ['Priority A: 2', 'Priority B: 1'])
        self.assertEqual(Path(self.temp.name, 'original.docx').read_bytes(), data)

    def test_unknown_displayed_dropdown_is_unreadable_and_explicit_change_succeeds(self):
        doc = Document(); doc.add_paragraph('Priorities must be unique.')
        dropdown(doc.add_paragraph('Priority A: '), '4')
        dropdown(doc.add_paragraph('Priority B: '), '2')
        data, version = self.fixture(doc)
        with self.assertRaisesMessage(ValidationError, 'unreadable'): b.check_answers(version, self.answers(version))
        self.assertEqual(self.values(data, version, self.answers(version))[1:], ['Priority A: 4', 'Priority B: 2'])
        good = self.answers(version, '3', ''); b.check_answers(version, good)
        self.assertEqual(self.values(data, version, good)[1:], ['Priority A: 3', 'Priority B: 2'])

    def test_placeholder_and_known_displayed_dropdown_remain_usable(self):
        doc = Document(); doc.add_paragraph('Priorities must be unique.')
        dropdown(doc.add_paragraph('Priority A: '), 'Choose an item.', placeholder=True)
        dropdown(doc.add_paragraph('Priority B: '), '2')
        _, version = self.fixture(doc)
        b.check_answers(version, self.answers(version))
        b.check_answers(version, self.answers(version, '1', ''))
        with self.assertRaisesMessage(ValidationError, 'priority'): b.check_answers(version, self.answers(version, '2', ''))

    def test_legacy_dropdown_invalid_index_is_unknown_and_replacement_is_readable(self):
        doc = Document(); doc.add_paragraph('Priorities must be unique.')
        for label, index in [('Priority A: ', '9'), ('Priority B: ', '1')]:
            p = doc.add_paragraph(label); ff_field(p, 'dropdown', options=['1', '2', '3'])
            result = OxmlElement('w:result'); result.set(qn('w:val'), index)
            next(p._p.iter(qn('w:ddList'))).append(result)
        _, version = self.fixture(doc)
        with self.assertRaisesMessage(ValidationError, 'unreadable'): b.check_answers(version, self.answers(version))
        b.check_answers(version, self.answers(version, '3', ''))

    def test_legacy_default_without_visible_result_requires_explicit_answer(self):
        doc = Document(); doc.add_paragraph('Priorities must be unique.')
        legacy(doc.add_paragraph('Priority A: '), default='1')
        legacy(doc.add_paragraph('Priority B: '), '2')
        _, version = self.fixture(doc)
        with self.assertRaisesMessage(ValidationError, 'unreadable'): b.check_answers(version, self.answers(version))
        b.check_answers(version, self.answers(version, '1', ''))

    def test_preserved_sdt_date_scalar_is_counted(self):
        doc = Document(); doc.add_paragraph('Priorities must be unique.')
        for label, value in [('Priority A: ', '1'), ('Priority B: ', '2')]:
            control = _sdt('rich', placeholder=value)
            control.find(qn('w:sdtPr')).append(OxmlElement('w:date'))
            doc.add_paragraph(label)._p.append(control)
        _, version = self.fixture(doc)
        self.assertEqual({f['kind'] for f in version.fields}, {'date'})
        b.check_answers(version, self.answers(version))
        with self.assertRaisesMessage(ValidationError, 'priority'): b.check_answers(version, self.answers(version, '', '1'))

    def test_plain_cells_and_separate_spans_read_written_values_and_blanks(self):
        for shape in ('cell', 'span'):
            with self.subTest(shape=shape):
                doc = Document(); doc.add_paragraph('Priorities must be unique.')
                if shape == 'cell':
                    table = doc.add_table(rows=2, cols=2)
                    for index in range(2):table.cell(index, 0).text = f'Priority {index}'
                else:
                    for index in range(2):doc.add_paragraph(f'Priority {index}: ____________')
                _, version = self.fixture(doc)
                b.check_answers(version, self.answers(version))
                b.check_answers(version, self.answers(version, '1', '2'))
                with self.assertRaisesMessage(ValidationError, 'priority'): b.check_answers(version, self.answers(version, '1', '1'))

    def test_exported_scalar_unknowns_and_missing_values_cannot_count_as_blank(self):
        doc = Document(); doc.add_paragraph('Priorities must be unique.')
        legacy(doc.add_paragraph('Priority A: '), '1'); legacy(doc.add_paragraph('Priority B: '), '2')
        _, version = self.fixture(doc); answers = self.answers(version); manifest = b.constraints(version)
        for returned in ({}, {version.fields[0]['key']: '1', version.fields[1]['key']: None}):
            with self.assertRaises(ValidationError):
                b.validate_constraints(manifest, answers, version.fields, read_options=lambda *_: returned)
        with self.assertRaisesMessage(ValidationError, 'required'): b.validate_constraints(manifest, answers, version.fields)

    def test_preserved_scalar_applies_to_maximum_active_groups_too(self):
        doc = Document(); doc.add_paragraph('Priorities must be unique.')
        legacy(doc.add_paragraph('Priority A: '), '1'); legacy(doc.add_paragraph('Priority B: '), '1')
        _, version = self.fixture(doc)
        version.generation_constraints['rules'] = [{'kind': 'max_active_groups', 'source_id': 'rule', 'maximum': 1,
            'groups': [[{'key': f['key'], 'values': ['1']}] for f in version.fields]}]
        with self.assertRaisesMessage(ValidationError, 'maximum selection'): b.check_answers(version, self.answers(version))
        b.check_answers(version, self.answers(version, '2', ''))

    def test_scalar_replicas_must_agree_and_explicit_write_updates_both(self):
        doc = Document(); doc.add_paragraph('Priorities must be unique.')
        section = doc.sections[0]; section.different_first_page_header_footer = True
        for header, value in ((section.header, '1'), (section.first_page_header, '2')):
            p = header.paragraphs[0]; p.text = 'Priority A: '
            control = _sdt('rich', placeholder=value)
            control.find(qn('w:sdtPr')).append(OxmlElement('w:date')); p._p.append(control)
        legacy(doc.add_paragraph('Priority B: '), '3')
        data, version = self.fixture(doc)
        copied = next(f for f in version.fields if f['locator']['target'].get('copies'))
        self.assertEqual(len(copied['locator']['target']['copies']), 1)
        answers = {f['key']: {'text': '', 'reviewed': True} for f in version.fields}
        with self.assertRaisesMessage(ValidationError, 'unreadable'): b.check_answers(version, answers)
        answers[copied['key']]['text'] = '1'; b.check_answers(version, answers)
        output = w.fill(data, version.fields, v3_mapping.fill_answers(version.fields, answers))
        package = w.open_package(output)
        for part in {copied['locator']['part'], copied['locator']['target']['copies'][0]['part']}:
            self.assertIn('Priority A: 1', w.text_of(package.tree(part)))

    def test_shifted_neighbouring_spans_refuse_ambiguous_readback(self):
        doc = Document(); doc.add_paragraph('Priorities must be unique.')
        doc.add_paragraph('Priority A: ____________  Priority B: ____________')
        _, version = self.fixture(doc)
        # Both writes change the other's original bounding text. Until that
        # relationship is qualified, unknown must not be counted as empty.
        with self.assertRaisesMessage(ValidationError, 'unreadable'):
            b.check_answers(version, self.answers(version, '1', '2'))

    def placeholder_fixture(self, kind, flag, *, missing_content=False, duplicate=False):
        doc = Document(); doc.add_paragraph('Priorities must be unique.')
        for index, value in enumerate(('1', '2')):
            control = _sdt('dropdown' if kind == 'choice' else 'rich', options=['1', '2', '3'], placeholder=value)
            pr = control.find(qn('w:sdtPr'))
            for node in pr.findall(qn('w:showingPlcHdr')): pr.remove(node)
            if kind == 'scalar':pr.append(OxmlElement('w:date'))
            if index == 0:
                marker = OxmlElement('w:showingPlcHdr')
                if flag is not None:marker.set(qn('w:val'), flag)
                pr.append(marker)
                if duplicate:pr.append(OxmlElement('w:showingPlcHdr'))
                if missing_content:control.remove(control.find(qn('w:sdtContent')))
            doc.add_paragraph(f'Priority {index}: ')._p.append(control)
        return self.fixture(doc)

    def test_false_placeholder_flags_preserve_value_and_true_flags_remain_placeholders(self):
        # OnOffType lexical values, independently listed from the XML contract.
        for kind in ('scalar', 'choice'):
            for flag in ('0', 'false', 'off'):
                with self.subTest(kind=kind, flag=flag):
                    _, version = self.placeholder_fixture(kind, flag)
                    with self.assertRaisesMessage(ValidationError, 'priority'):
                        b.check_answers(version, self.answers(version, '', '1'))
                    b.check_answers(version, self.answers(version, '3', '1'))
            for flag in (None, '1', 'true', 'on'):
                with self.subTest(kind=kind, flag=flag):
                    _, version = self.placeholder_fixture(kind, flag)
                    b.check_answers(version, self.answers(version, '', '1'))

    def test_malformed_placeholder_and_missing_content_are_unknown(self):
        for kind in ('scalar', 'choice'):
            for flag, missing, duplicate in [('bogus', False, False), ('', False, False),
                                              ('0', False, True), ('0', True, False)]:
                with self.subTest(kind=kind, flag=flag, missing=missing, duplicate=duplicate):
                    _, version = self.placeholder_fixture(kind, flag, missing_content=missing, duplicate=duplicate)
                    with self.assertRaisesMessage(ValidationError, 'unreadable'):
                        b.check_answers(version, self.answers(version))

    def test_malformed_legacy_boundaries_never_become_known_empty(self):
        for mode in ('missing', 'duplicate', 'text_before', 'begin_text', 'separator_text', 'end_text'):
            with self.subTest(mode=mode):
                doc = Document(); doc.add_paragraph('Priorities must be unique.')
                p = doc.add_paragraph('Priority A: '); legacy(p, '1')
                separator = next(n for n in p._p.iter(qn('w:fldChar')) if n.get(qn('w:fldCharType')) == 'separate')
                if mode == 'missing':p._p.remove(separator.getparent())
                elif mode == 'duplicate':
                    r = OxmlElement('w:r'); n = OxmlElement('w:fldChar'); n.set(qn('w:fldCharType'), 'separate')
                    r.append(n); separator.getparent().addnext(r)
                elif mode == 'text_before':
                    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = '3'; r.append(t)
                    separator.getparent().addprevious(r)
                else:
                    boundary = {'begin_text': 'begin', 'separator_text': 'separate', 'end_text': 'end'}[mode]
                    marker = next(n for n in p._p.iter(qn('w:fldChar')) if n.get(qn('w:fldCharType')) == boundary)
                    t = OxmlElement('w:t'); t.text = '3'; marker.getparent().append(t)
                legacy(doc.add_paragraph('Priority B: '), '2')
                _, version = self.fixture(doc)
                with self.assertRaisesMessage(ValidationError, 'unreadable'):
                    b.check_answers(version, self.answers(version, '', '2'))
                # The unchanged writer cannot remove text in an unresolved
                # result area by adding a separator. Explicit text still holds.
                with self.assertRaisesMessage(ValidationError, 'unreadable'):
                    b.check_answers(version, self.answers(version, '3', '2'))
