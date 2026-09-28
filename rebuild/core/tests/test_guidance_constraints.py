"""Printed coloured values cannot disappear from global constraints after appending."""
import copy
import hashlib
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, override_settings
from docx import Document
from docx.shared import RGBColor

from core import generation_batch as b
from documents import word_v3 as w
from documents.test_question_boxes import data


class GuidanceConstraints(SimpleTestCase):
    def test_actual_sandbox_constraint_reader_refuses_preserved_scalar_and_option_values(self):
        for option in (False, True):
            for legacy in (False, True):
                with self.subTest(option=option, legacy=legacy), tempfile.TemporaryDirectory() as directory:
                    doc = Document(); instruction = 'Select no Yes options.' if option else 'Priorities must be unique.'
                    doc.add_paragraph(instruction); table = doc.add_table(rows=2, cols=2)
                    if option:
                        table.cell(0, 0).text = 'Question'; table.cell(0, 1).text = 'Yes / No'
                        table.cell(1, 0).text = 'Do you smoke?'
                        table.cell(1, 1).paragraphs[0].add_run('Yes').font.color.rgb = RGBColor.from_string('808080')
                    else:
                        table.cell(0, 0).text = 'Priority A'; table.cell(1, 0).text = 'Priority B'
                        table.cell(0, 1).paragraphs[0].add_run('1').font.color.rgb = RGBColor.from_string('0000FF')
                    original = data(doc); fields = w.inspect(original)['fields']; self.assertEqual(len(fields), 1 if option else 2)
                    retained = next(f for f in fields if f['locator']['target']['type'] == 'cell_append')
                    if legacy:
                        retained['locator']['target'].update(type='cell_replace', all=True); retained.pop('source_instructions')
                    sha = hashlib.sha256(original).hexdigest(); Path(directory, 'original.docx').write_bytes(original)
                    rule = ({'kind': 'max_active_groups', 'source_id': 'r', 'maximum': 0,
                             'groups': [[{'key': retained['key'], 'values': ['Yes']}]]} if option else
                            {'kind': 'unique_values', 'source_id': 'r', 'keys': [f['key'] for f in fields],
                             'allowed_values': ['1', '2']})
                    manifest = {'contract': b.CONSTRAINTS, 'source_sha256': sha, 'fields_sha256': b.digest(fields),
                                'coverage': 'complete', 'provenance': 'synthetic_fixture', 'evidence_sha256': '1' * 64,
                                'instruction_references': [{'id': 'r', 'part': 'word/document.xml', 'slot': 'b/p[0]',
                                                           'sha256': hashlib.sha256(instruction.encode()).hexdigest()}], 'rules': [rule]}
                    version = SimpleNamespace(fields=fields, template=SimpleNamespace(sha256=sha, original_path='original.docx'),
                                              generation_constraints=manifest)
                    answers = {f['key']: {'text': ('No' if option else '2') if f is retained else '1',
                                        'reviewed': True, 'status': 'supported'} for f in fields}
                    before = copy.deepcopy(fields)
                    with override_settings(MEDIA_ROOT=Path(directory)), patch('socket.socket.connect', side_effect=AssertionError('Network forbidden')):
                        with self.assertRaisesMessage(ValidationError, 'unreadable'):
                            b.check_answers(version, answers)
                    self.assertEqual(fields, before)
                    self.assertEqual(Path(directory, 'original.docx').read_bytes(), original)
