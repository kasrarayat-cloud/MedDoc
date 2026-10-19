"""Compact clinician review: explicit consent, complete definitions and no stale approval."""
from copy import deepcopy
from unittest.mock import patch
from django.core.exceptions import ValidationError
from django.test import TestCase, SimpleTestCase, override_settings
from django.urls import reverse
from core import preparation, template_review, word_contract


@override_settings(WORD_PREPARATION_ENABLED=True, AI_MODE='fixture', ADMIN_SITE=False,
    TESTING_AUTHORITY_MODE='', TESTING_AUTHORITY_ID='', TESTING_AUTHORITY_SHA256='',
    STORAGES={'default':{'BACKEND':'django.core.files.storage.FileSystemStorage'},
              'staticfiles':{'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class CompactReviewTests(TestCase):
    def setUp(self):
        from core.tests.test_word_answer_contract import OrdinaryClinicianWordTests
        OrdinaryClinicianWordTests.setUp(self)

    def fresh(self, ambiguous=False, second_unresolved=False):
        from core.tests.test_word_answer_contract import OrdinaryClinicianWordTests
        original = preparation.provider_role
        def role(*args, **kwargs):
            result, receipt = original(*args, **kwargs)
            if ambiguous:
                row = result['answer_definitions']['fields'][0]
                row['definition'].update(value_type='person_name', answer_mode='atomic')
                row['definition']['scope']['subject'] = 'patient'
                if args[1] == 'checker': row.update(state='unresolved', reason='The form does not identify whose name belongs here.')
            if second_unresolved and args[1] == 'checker':
                result['answer_definitions']['fields'][1].update(state='unresolved', reason='Clarify the intended answer.')
            return result, receipt
        with patch.object(preparation, 'provider_role', side_effect=role):
            return OrdinaryClinicianWordTests.fresh_upload(self, 'guided' if ambiguous or second_unresolved else 'review')

    def page(self, template):
        return self.client.get(reverse('template_preparation', args=[template.pk]))

    def post(self, response):
        from core.tests.test_semantic_preparation import DefinitionPreparationPureTests
        context = response.context
        values = {r['field']['key']: r['definition'] for r in context['definition_rows']}
        post = DefinitionPreparationPureTests.post(values)
        del post['reviewed_field']
        for key, definition in values.items(): post['field_'+key+'_answer_mode'] = definition['answer_mode']
        post.update(compact_review='on', revision=str(context['preparation'].revision),
            candidate_sha256=context['preparation'].candidate_sha256, review_summary='on',
            review_token_save=context['review_token_save'], review_token_approve=context['review_token_approve'])
        return post

    def save(self, template, post):
        return self.client.post(reverse('template_preparation_correct', args=[template.pk]), post)

    def approve(self, template, post):
        return self.client.post(reverse('template_preparation_approve', args=[template.pk]), post)

    def test_one_unchecked_consent_can_explicitly_approve_all_complete_definitions_once(self):
        template, prep = self.fresh(); page = self.page(template)
        self.assertContains(page, 'Review detected fields')
        self.assertContains(page, 'name="review_summary" required>')
        self.assertNotContains(page, 'name="reviewed_field"')
        self.assertNotContains(page, 'name="reviewed_definition"')
        self.assertContains(page, '<summary>Advanced')
        post = self.post(page); no_consent = post.copy(); del no_consent['review_summary']
        self.assertContains(self.approve(template, no_consent), 'confirm the review summary', status_code=409)
        self.assertFalse(template.versions.exists())
        self.assertEqual(self.approve(template, post).status_code, 302)
        self.assertEqual(template.versions.count(), 1)
        self.assertEqual(word_contract.require_generation(template.versions.get()), word_contract.proposed(prep))
        self.assertEqual(self.approve(template, post).status_code, 409)
        self.assertEqual(template.versions.count(), 1)

    def test_unsaved_edits_cannot_approve_even_without_javascript(self):
        template, prep = self.fresh(); post = self.post(self.page(template))
        key = prep.inventory['fields'][0]['key']
        post['field_'+key+'_meaning'] += ' Keep only the requested answer.'
        self.assertContains(self.approve(template, post), 'Save your edited definitions', status_code=409)
        self.assertFalse(template.versions.exists())
        self.assertEqual(self.save(template, post).status_code, 302)
        self.assertEqual(self.approve(template, self.post(self.page(template))).status_code, 302)
        self.assertEqual(word_contract.require_generation(template.versions.get())[key]['meaning'], post['field_'+key+'_meaning'])

    def test_scope_choice_is_required_even_when_ai_proposed_patient_and_keeps_raw_roles(self):
        template, prep = self.fresh(ambiguous=True); raw = deepcopy((prep.proposal, prep.checker))
        page = self.page(template); post = self.post(page); key = prep.inventory['fields'][0]['key']
        self.assertContains(page, 'Whose name belongs here?')
        self.assertContains(page, '<option value="">Choose a person</option>')
        self.assertNotContains(page, 'data-review-approve')
        self.assertContains(self.save(template, post), 'Choose a person', status_code=409)
        post['name_subject_'+key] = 'patient'
        self.assertEqual(self.save(template, post).status_code, 302)
        prep.refresh_from_db(); self.assertEqual((prep.proposal, prep.checker), raw)
        self.assertEqual(len(prep.correction_history), 1)
        self.assertEqual(word_contract.proposed(prep)[key]['scope']['subject'], 'patient')
        page = self.page(template); self.assertNotContains(page, 'Whose name belongs here?')
        self.assertEqual(self.approve(template, self.post(page)).status_code, 302)

    def test_name_choice_preserves_explicit_advanced_meaning_and_other_requires_role(self):
        template, prep = self.fresh(ambiguous=True); post = self.post(self.page(template)); key=prep.inventory['fields'][0]['key']
        post['name_subject_'+key] = 'other'
        self.assertContains(self.save(template, post), 'Describe whose name', status_code=409)
        post['name_other_'+key] = 'accompanying parent'
        post['field_'+key+'_meaning'] = 'The accompanying parent’s full legal name, including all surnames.'
        self.assertEqual(self.save(template, post).status_code, 302)
        prep.refresh_from_db(); saved=word_contract.proposed(prep)[key]
        self.assertEqual(saved['meaning'], post['field_'+key+'_meaning'])
        self.assertEqual(saved['scope']['subject'], 'other')

    def test_every_unresolved_question_needs_a_decision_and_accept_cannot_disguise_edits(self):
        template, prep = self.fresh(ambiguous=True, second_unresolved=True)
        post = self.post(self.page(template)); keys=[f['key'] for f in prep.inventory['fields']]
        post['name_subject_'+keys[0]]='patient'
        self.assertContains(self.save(template, post), 'Resolve the question', status_code=409)
        post['decision_'+keys[1]]='accept'; post['field_'+keys[1]+'_subject']='clinician'
        self.assertContains(self.save(template, post), 'Resolve the question', status_code=409)
        post['decision_'+keys[1]]='edit'
        self.assertContains(self.save(template, post), 'update its meaning', status_code=409)
        post['field_'+keys[1]+'_meaning']='The clinician’s documented reason for attending.'
        self.assertEqual(self.save(template, post).status_code,302)

    def test_hidden_controls_must_round_trip_and_destinations_are_not_editable(self):
        template, prep = self.fresh(); page=self.page(template); post=self.post(page)
        key=prep.inventory['fields'][0]['key']
        for suffix in ('unit','minimum','scale_maximum','choices','answer_mode'):
            missing=post.copy();del missing['field_'+key+'_'+suffix]
            self.assertContains(self.save(template,missing), 'missing or repeated',status_code=409)
        repeated=post.copy();repeated.appendlist('field_'+key+'_unit','Celsius')
        self.assertContains(self.client.post(reverse('template_preparation_correct',args=[template.pk]),repeated.urlencode(),content_type='application/x-www-form-urlencoded'), 'missing or repeated',status_code=409)
        post['field_'+key+'_locator_sha256']='f'*64
        self.assertEqual(self.save(template,post).status_code,302)
        prep.refresh_from_db()
        self.assertEqual(word_contract.proposed(prep)[key]['permitted_destination'], page.context['definition_rows'][0]['definition']['permitted_destination'])

    def test_token_binds_actor_operation_revision_and_full_source_snapshot(self):
        template,prep=self.fresh();post=self.post(self.page(template));items=word_contract.review_rows(prep)
        for change in ('actor','operation','revision','candidate','locator','meaning'):
            edited=deepcopy(prep);actor=deepcopy(self.user);body=post.copy()
            if change=='actor':actor.pk='00000000-0000-0000-0000-000000000000'
            if change=='operation':body['review_token_approve']=body['review_token_save']
            if change=='revision':edited.revision+=1
            if change=='candidate':edited.candidate_sha256='f'*64
            changed_items=deepcopy(items)
            if change=='locator':changed_items[0]['field']['locator']['cell']=99
            if change=='meaning':changed_items[0]['definition']['meaning']='Changed meaning.'
            with self.subTest(change=change),self.assertRaises(ValidationError):
                template_review.submitted(edited,changed_items,actor,body,'approve')

    def test_label_spacing_uses_actual_word_boundaries_without_changing_mapping(self):
        from docx import Document
        from documents import inspect_template
        doc=Document(); table=doc.add_table(rows=0,cols=2)
        for parts in [('HbA1c',),('eGFR',),('McDonald',),('Person full name',),('Intensity','0 to 10')]:
            cell=table.add_row().cells[0];cell.text=parts[0]
            for text in parts[1:]:cell.add_paragraph(text)
        path=self.root/'display.docx';doc.save(path);result=inspect_template(path)
        self.assertEqual(result['display_labels']['Intensity0 to 10'],'Intensity 0 to 10')
        for text in ('HbA1c','eGFR','McDonald'):self.assertEqual(result['display_labels'][text],text)
        self.assertIn('Intensity0 to 10',[f['label'] for f in result['fields']])

    def test_optional_display_inspection_cannot_strand_a_retained_source_or_skip_integrity(self):
        from documents import DocumentError
        from django.core.cache import cache
        cache.delete('template-review-labels-v1:'+self.template.sha256)
        with patch('documents.inspect_template', side_effect=DocumentError('Current inspector refusal', 'overflow_risk')):
            self.assertEqual(template_review.label_displays(self.template), {})
            self.template.sha256='f'*64
            with self.assertRaises(ValidationError): template_review.label_displays(self.template)

    def test_read_only_pending_definitions_do_not_claim_approval(self):
        from core.models import Membership
        template,prep=self.fresh()
        Membership.objects.filter(user=self.user).update(role='employee')
        page=self.page(template)
        self.assertContains(page,'Review detected fields')
        self.assertNotContains(page,'Approved fields')
        self.assertNotContains(page,'data-review-approve')
        self.assertNotContains(page,'data-review-save')


@override_settings(WORD_PREPARATION_ENABLED=True, AI_MODE='fixture', ADMIN_SITE=False,
    TESTING_AUTHORITY_MODE='', TESTING_AUTHORITY_ID='', TESTING_AUTHORITY_SHA256='',
    STORAGES={'default':{'BACKEND':'django.core.files.storage.FileSystemStorage'},
              'staticfiles':{'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class SeparateDefinitionRetryTests(TestCase):
    def setUp(self):
        from core.tests.test_semantic_preparation import DefinitionPreparationDatabaseTests
        DefinitionPreparationDatabaseTests.setUp(self)
        from core import semantic_preparation
        from core.tests.test_semantic_preparation import historical_preparation
        self.prep=historical_preparation(self.source,self.user)
        self.assertTrue(semantic_preparation.run_one());self.prep.refresh_from_db()
        self.client.force_login(self.user)

    def definition_post(self, page):
        from core.tests.test_semantic_preparation import DefinitionPreparationPureTests
        values={r['field']['key']:r['definition'] for r in page.context['definition_rows']}
        post=DefinitionPreparationPureTests.post(values)
        del post['reviewed_field']
        post.update(compact_review='on',revision=str(self.prep.revision),review_summary='on',
            inventory_sha256=self.prep.inventory['inventory_sha256'],
            review_token_save=page.context['review_token_save'],
            review_token_approve=page.context['review_token_approve'])
        return post

    def test_documented_patient_name_round_trips_one_summary_and_saved_immutable_approval(self):
        page=self.client.get(reverse('template_definitions',args=[self.prep.pk]))
        self.assertContains(page,'Patient name in this answer')
        self.assertContains(page,'No additional requirement')
        self.assertContains(page,'name="review_summary" required>',count=1)
        self.assertNotContains(page,'name="reviewed_field"')
        key=self.source.fields[2]['key']
        self.assertContains(page,'for="patient-name-'+key+'"')
        post=self.definition_post(page)
        post['field_'+key+'_patient_name_requirement']='documented'
        post['field_'+key+'_meaning']='The documented patient name and reported symptom history.'
        approve=reverse('template_definitions_approve',args=[self.prep.pk])
        self.assertContains(self.client.post(approve,post),'Save your edited definitions',status_code=409)
        self.assertEqual(self.template.versions.count(),1)
        self.assertEqual(self.client.post(reverse('template_definitions_correct',args=[self.prep.pk]),post).status_code,302)
        self.prep.refresh_from_db()
        page=self.client.get(reverse('template_definitions',args=[self.prep.pk]))
        self.assertContains(page,'Patient name in this answer: Include documented patient name')
        self.assertContains(self.client.post(approve,post),'changed since this review opened',status_code=409)
        self.assertEqual(self.client.post(approve,self.definition_post(page)).status_code,302)
        approved=self.template.versions.order_by('-number').first()
        self.assertEqual(approved.semantic_definitions[key]['patient_name_requirement'],'documented')
        self.assertEqual(approved.fields,self.source.fields)
        self.assertEqual(self.prep.ai_stage_receipts.count(),2)

    def test_patient_name_control_missing_repeated_and_incompatible_scope_are_refused(self):
        page=self.client.get(reverse('template_definitions',args=[self.prep.pk]))
        post=self.definition_post(page);key=self.source.fields[2]['key']
        control='field_'+key+'_patient_name_requirement'
        url=reverse('template_definitions_correct',args=[self.prep.pk])
        missing=post.copy();del missing[control]
        self.assertContains(self.client.post(url,missing),'missing or repeated',status_code=409)
        repeated=post.copy();repeated.appendlist(control,'documented')
        self.assertContains(self.client.post(url,repeated.urlencode(),content_type='application/x-www-form-urlencoded'),
            'missing or repeated',status_code=409)
        post[control]='documented';post['field_'+key+'_subject']='clinician'
        post['field_'+key+'_meaning']='The clinician name and history.'
        self.assertContains(self.client.post(url,post),'existing patient name or narrative answer',status_code=409)
        self.prep.refresh_from_db()
        self.assertEqual(self.prep.revision,0)
        self.assertFalse(self.prep.corrections)

    def test_patient_name_requirement_is_in_review_snapshot_and_historical_controls_stay_absent(self):
        from django.template.loader import render_to_string
        page=self.client.get(reverse('template_definitions',args=[self.prep.pk]))
        items=page.context['definition_rows'];post=self.definition_post(page)
        changed=deepcopy(items);changed[2]['definition']['patient_name_requirement']='documented'
        with self.assertRaisesMessage(ValidationError,'changed since this review opened'):
            template_review.submitted(self.prep,changed,self.user,post,'approve')
        for version in (1,2):
            row=deepcopy(items[2]);row['definition'].pop('patient_name_requirement')
            row['definition']['version']=version
            if version==1:row['definition'].pop('answer_mode')
            displayed=template_review.rows([row])[0]
            html=render_to_string('template_review_field.html',{'row':displayed,'editable':True,'definition_version':version})
            self.assertNotIn('patient_name_requirement',html)
            self.assertNotIn('Patient name in this answer',html)

    def test_error_form_keeps_edits_and_issues_a_retryable_retained_baseline_token(self):
        from core.tests.test_semantic_preparation import DefinitionPreparationPureTests
        from core import semantic_preparation
        page=self.client.get(reverse('template_definitions',args=[self.prep.pk]))
        values={r['field']['key']:r['definition'] for r in page.context['definition_rows']}
        post=DefinitionPreparationPureTests.post(values)
        for key,value in values.items():
            if value['version']==2:post['field_'+key+'_answer_mode']=value['answer_mode']
        post.update(compact_review='on',revision=str(self.prep.revision),review_token_save=page.context['review_token_save'])
        key=next(iter(values));meaning='Only the documented patient name, preserving all surnames.'
        post['field_'+key+'_meaning']=meaning
        url=reverse('template_definitions_correct',args=[self.prep.pk])
        error=self.client.post(url,post)
        self.assertEqual(error.status_code,409)
        self.assertContains(error,meaning,status_code=409)
        from django.core import signing
        self.assertEqual(signing.loads(error.context['review_token_save'],salt='template-review-v1'),
                         signing.loads(page.context['review_token_save'],salt='template-review-v1'))
        post['review_token_save']=error.context['review_token_save'];post['review_summary']='on'
        self.assertEqual(self.client.post(url,post).status_code,302)
        self.prep.refresh_from_db()
        self.assertEqual(semantic_preparation.corrected_definitions(self.prep)[key]['meaning'],meaning)
        self.assertEqual(self.template.versions.count(),1)  # Saving is not approval.


class SummaryConstraintTests(SimpleTestCase):
    def test_browser_serialization_preserves_multiline_choices_meaning_and_unit_bytes(self):
        from types import SimpleNamespace
        from django.http import QueryDict
        from documents.semantic_preparation_roles import destination, validate_definitions
        field={'key':'status','label':'Allergy status','kind':'text','locator':{'cell':0}}
        definition={'version':2,'value_type':'choice','answer_mode':'atomic',
            'meaning':'\nPatient allergy status.\nRetain the documented choice.', 'unit':' documented status ',
            'choices':['No\nknown allergy','Known allergy'],
            'scope':{'subject':'patient','temporality':'current','polarity':'any'},
            'constraints':dict.fromkeys(('minimum','maximum','scale_minimum','scale_maximum')),
            'permitted_destination':destination(field)}
        validate_definitions([field],{'status':definition})
        prep=SimpleNamespace(pk='prep',revision=0,state='review',config={},inventory={},changes=[],
            candidate_sha256='a'*64,template=SimpleNamespace(pk='template',sha256='a'*64))
        user=SimpleNamespace(pk='doctor');items=[{'field':field,'definition':definition,'state':'pass'}]
        post=QueryDict(mutable=True)
        for key,value in {'meaning':definition['meaning'].replace('\n','\r\n'),'choices':'\r\n'.join(definition['choices']).replace('No\nknown','No\r\nknown'),
            'value_type':'choice','unit':definition['unit'],'answer_mode':'atomic',**definition['scope'],**definition['constraints']}.items():
            post['field_status_'+key]=value or ''
        post.update(review_summary='on',**template_review.tokens(prep,items,user.pk))
        values,_,_=template_review.submitted(prep,items,user,post,'approve')
        self.assertEqual(values,{'status':definition})
        post['field_status_meaning']='Patient allergy status.\r\nPreserve the full answer.'
        values,_,_=template_review.submitted(prep,items,user,post,'save')
        self.assertEqual(values['status']['meaning'],'Patient allergy status.\nPreserve the full answer.')
        self.assertEqual(values['status']['choices'],definition['choices'])
        self.assertEqual(values['status']['unit'],definition['unit'])
        post['field_status_choices']='No known allergy\r\nKnown allergy'
        with self.assertRaisesMessage(ValidationError,'multiline choices'):
            template_review.submitted(prep,items,user,post,'save')

    def test_zero_scale_narrower_limits_and_celsius_limits_remain_visible(self):
        definition={'version':2,'value_type':'number','unit':None,'scope':{'subject':'patient'},
            'constraints':{'minimum':'2','maximum':'8','scale_minimum':'0','scale_maximum':'10'}}
        item={'field':{'key':'score','label':'Score','locator':{'cell':0}},'definition':definition}
        row=template_review.rows([item])[0]
        self.assertEqual(row['type_summary'],'Number from 0 to 10')
        self.assertEqual(row['constraints_summary'],'Minimum 2 · Maximum 8 · Scale 0 to 10')
        definition.update(unit='Celsius');definition['constraints']={'minimum':'32','maximum':'42','scale_minimum':None,'scale_maximum':None}
        row=template_review.rows([item])[0]
        self.assertEqual(row['type_summary'],'Temperature in Celsius')
        self.assertEqual(row['constraints_summary'],'Minimum 32 · Maximum 42')

    def test_generic_person_label_requires_explicit_initial_choice_but_qualified_names_do_not(self):
        from types import SimpleNamespace
        prep=SimpleNamespace(state='review',corrections={},candidate_sha256='x')
        rows=[{'field':{'key':'name','label':'Person full name'},'definition':{'value_type':'person_name'},'state':'pass'}]
        self.assertEqual(template_review.review_items(prep,rows)[0]['state'],'unresolved')
        for label in ('Full name (please print)', 'Name of person', 'Fullname', 'Person full name (clinician to complete)', 'Name (for clinician use)', 'Clinician to complete: Name'):
            rows[0]['field']['label']=label
            self.assertEqual(template_review.review_items(prep,rows)[0]['state'],'unresolved')
        rows[0]['field']['label']='Patient full name'
        self.assertEqual(template_review.review_items(prep,rows)[0]['state'],'pass')
        rows[0]['field']['label']='Person full name';prep.state='approved'
        self.assertEqual(template_review.review_items(prep,rows)[0]['state'],'pass')

    def test_timing_polarity_and_whole_degree_restrictions_are_visible(self):
        definition={'version':2,'value_type':'integer','unit':'Celsius','constraints':{},
                    'scope':{'subject':'patient','temporality':'past','polarity':'negative'}}
        item={'field':{'key':'temperature','label':'Temperature','locator':{'cell':0}},'definition':definition}
        row=template_review.rows([item])[0]
        self.assertEqual(row['type_summary'],'Temperature in Celsius · whole degrees')
        self.assertEqual(row['scope_summary'],'Past information · Negative statements only')
