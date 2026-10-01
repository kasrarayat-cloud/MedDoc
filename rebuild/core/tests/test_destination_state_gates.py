"""roles-2 through every gate: the three-state destination (ordinary / protected / uncertain) on the legacy default path
and on contract v3 - generation, review save, finalization, export, mapping confirmation, the complete-form plan and
the review page. Synthetic blank forms only; fixture generation (no provider, no network).

REG = the 18 clinical labels f5fe720f no longer generated; HOLD = the 17 must-hold labels; L1:/L2:/L3: = recorded
residual shapes (docs/rebuild/v1/destination-roles-20260930/residuals.json).
"""
import json
from unittest.mock import patch

from django.test import TestCase, override_settings
from docx import Document

from core import destination_roles, generation, v3_mapping
from core.models import Encounter, TemplateVersion
from core.security import private_path
from core.tests.test_destination_roles import BODY, LIVE, Walk, legacy_rows, lines, provider_capture, saved
from core.tests.test_word_v3_workflow import V3, V3Fixture, part_text
from core.tests.test_workflow import TEST_SETTINGS
from documents import word_v3

#: The 14 labels roles-2 restores (ordinary on evidence) with a synthetic value each.
RESTORED = (('Troponin (initial)', '12 ng/L'), ('Lactate (initial)', '2.1 mmol/L'), ('Dose (initial)', '25 micrograms'),
            ('Troponin initial', '14 ng/L'), ('Troponin — Initial (ng/L)', '16'), ('Initial (baseline) weight', '72 kg'),
            ('Initial and ongoing management', 'IV antibiotics'), ('Visit type (initial/follow-up)', 'Follow-up'),
            ('Initial / repeat', 'Repeat'), ('Gene expression signature', 'Luminal A'),
            ('Consent form signed by patient', 'Yes'), ('Drug chart signed by prescriber', 'Yes'),
            ('Patient signature obtained', 'Yes'), ('Witness signature obtained', 'No'))
#: The 4 of the 18 that stay held, and why (destination_state): uncertain, or a true signing/initials place.
STILL_HELD = (('Signed off by consultant', 'Dr Example', destination_roles.UNCERTAIN_NOT_GENERATED),
              ('Signed on admission', 'Yes', destination_roles.UNCERTAIN_NOT_GENERATED),
              ('Initial', 'JE', destination_roles.ROLE_NOT_GENERATED),
              ('Patient initials', 'JE', destination_roles.ROLE_NOT_GENERATED))


def offered(answer):
    """What generation offered for an asked field: its supported text, or the candidate it held for the clinician's
    decision (generation's own governance, unchanged here) - never offered for a field it was not asked."""
    return answer.get('text') or (answer.get('held_candidate') or {}).get('text', '')


def notes_for(pairs, plan=True):
    return '\n'.join(f'{label}: {value}' for label, value in pairs) + ('\nPlan: Rest.' if plan else '')


@override_settings(**TEST_SETTINGS)
class LegacyThreeStates(Walk, V3Fixture, TestCase):
    def test_reg_the_restored_labels_are_generated_finalized_and_exported(self):
        template, version = self.uploaded(data=legacy_rows(*[x[0] for x in RESTORED + STILL_HELD], 'Plan'),
                                          name='Synthetic REG legacy')
        by = self.by_label(version)
        e = self.generated(version, notes_for(RESTORED + tuple((l, v) for l, v, _ in STILL_HELD)))
        for label, value in RESTORED:
            with self.subTest(label=label):
                self.assertEqual(offered(e.answers[by[label]['key']]), value)
        for label, _, reason in STILL_HELD:
            with self.subTest(label=label):
                answer = e.answers[by[label]['key']]
                self.assertEqual((answer['status'], answer['text'], answer.get('reason')), ('missing', '', reason))
        # the review page says why each held row was not generated
        page = self.client.get(f'/encounters/{e.pk}/').content.decode()
        self.assertEqual(page.count('data-destination-role'), len(STILL_HELD))
        self.assertIn(destination_roles.UNCERTAIN_NOT_GENERATED[:40], page)
        held = {label: value for label, value in RESTORED if e.answers[by[label]['key']]['status'] == 'conflict'}
        self.assertEqual(self.reviewed(e, version, typed=held).status_code, 302)
        self.finalized(e)
        self.assertEqual(e.status, 'finalized')
        exported = part_text(private_path(e.export_path).read_bytes())
        for label, value in RESTORED:
            self.assertIn(value, exported)
        for label, value, _ in STILL_HELD:
            self.assertNotIn(label + value, exported)

    def test_reg_the_provider_is_never_asked_a_held_row(self):
        template, version = self.uploaded(data=legacy_rows(*[x[0] for x in RESTORED + STILL_HELD], 'Plan'),
                                          name='Synthetic REG asked')
        e = self.generated(version, 'Plan: Rest.')
        sent = []
        with override_settings(**LIVE), provider_capture(sent):
            generation.generate(e.notes, destination_roles.generation_fields(version), generation.configuration())
        asked = [q['label'] for q in json.loads(sent[0]['messages'][0]['content'])['questions']]
        self.assertEqual(asked, [l for l, _ in RESTORED] + ['Plan'])

    def test_uncertain_generated_text_is_refused_typed_text_exports(self):
        template, version = self.uploaded(data=legacy_rows('Signed off by consultant', 'Plan'), name='Synthetic U legacy')
        key = self.by_label(version)['Signed off by consultant']['key']
        e = self.generated(version, 'Plan: Rest.')
        # a draft from before roles-2 could hold generated text there
        Encounter.objects.filter(pk=e.pk).update(answers={**e.answers, key: {'text': 'Dr Example', 'status': 'supported',
                                                                              'evidence': [], 'reviewed': False}})
        e.refresh_from_db()
        response = self.reviewed(e, version)
        self.assertEqual(response.status_code, 400)
        self.assertIn('does not show whether each is a signing or initials place', response.content.decode())
        response = self.reviewed(e, version, typed={'Signed off by consultant': 'Yes - 30/09/2026'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(e.answers[key]['status'], 'edited')
        self.finalized(e)
        self.assertEqual(e.status, 'finalized')
        self.assertIn('Yes - 30/09/2026', part_text(private_path(e.export_path).read_bytes()))

    def test_finalize_refuses_generated_text_on_an_uncertain_row(self):
        template, version = self.uploaded(data=legacy_rows('Signed on admission', 'Plan'), name='Synthetic U final')
        key = self.by_label(version)['Signed on admission']['key']
        e = self.generated(version, 'Plan: Rest.')
        answers = {k: {**v, 'reviewed': True, 'reviewed_by': str(self.user.pk)} for k, v in e.answers.items()}
        answers[key] = {'text': 'Yes', 'status': 'supported', 'evidence': [], 'reviewed': True,
                        'reviewed_by': str(self.user.pk)}
        Encounter.objects.filter(pk=e.pk).update(answers=answers, status='review')
        response = self.finalized(e)
        self.assertIn('never fills', response.content.decode())
        self.assertEqual(e.status, 'review')

    def test_b2_n7_a_pre_roles2_draft_of_a_restored_label_saves(self):
        """B2-N7: a draft f5fe720f... no - 43fcf214 generated holds 'supported' text for a restored label; f5fe720f
        refused it at save. roles-2 restores the label, so the draft saves and exports."""
        template, version = self.uploaded(data=legacy_rows('Initial and ongoing management', 'Plan'), name='Synthetic N7')
        key = self.by_label(version)['Initial and ongoing management']['key']
        e = self.generated(version, 'Plan: Rest.')
        Encounter.objects.filter(pk=e.pk).update(answers={**e.answers, key: {'text': 'IV antibiotics', 'status': 'supported',
                                                                              'evidence': [], 'reviewed': False}})
        e.refresh_from_db()
        self.assertEqual(self.reviewed(e, version).status_code, 302)

    def test_l1_rr1_05_l3_f3_01_a_wrapped_label_cell_is_read_the_same_way_at_every_gate(self):
        """L1:RR1-05 / L3:F3-01: 'Consultant' + 'signature' (two paragraphs) - the engine's label joins them
        ('Consultantsignature'); generation, save and finalization all read the printed words from the original."""
        doc = Document(); doc.add_paragraph('Clinic record')
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).paragraphs[0].add_run('Consultant'); table.cell(0, 0).add_paragraph('signature')
        table.cell(1, 0).text = 'Plan'
        template, version = self.uploaded(data=saved(doc), name='Synthetic RR1-05')
        field = version.fields[0]
        self.assertEqual(field['label'], 'Consultantsignature')
        e = self.generated(version, 'Consultantsignature: Dr Somebody Else\nPlan: Rest.')
        self.assertEqual((e.answers[field['key']]['status'], e.answers[field['key']]['text']), ('missing', ''))
        Encounter.objects.filter(pk=e.pk).update(answers={**e.answers, field['key']: {
            'text': 'Dr Somebody Else', 'status': 'supported', 'evidence': [], 'reviewed': False}})
        e.refresh_from_db()
        self.assertEqual(self.reviewed(e, version).status_code, 400)

    def test_l2_rr3_05_the_finalizers_own_line_naming_the_patient_keeps_the_name_rule(self):
        for label in ('Signed by (doctor responsible for patient)', "Signature (patient's consultant)"):
            with self.subTest(label=label):
                template, version = self.uploaded(data=legacy_rows(label, 'Plan'), name='Synthetic RR3-05 ' + label[:12])
                e = self.generated(version, 'Plan: Rest.')
                self.assertEqual(self.reviewed(e, version, typed={label: 'Dr Somebody Else'}).status_code, 302)
                response = self.finalized(e)
                self.assertIn('must match your signed-in clinician name', response.content.decode())
                self.assertEqual(e.status, 'review')


@override_settings(**V3)
class V3ThreeStates(Walk, V3Fixture, TestCase):
    def rows_post(self, e, version, texts):
        post = {'revision': e.revision}
        for f in version.fields:
            if f.get('kind') in v3_mapping.SIGNING_KINDS:
                continue
            post['answer_' + f['key']] = texts.get(f['label'], (e.answers.get(f['key']) or {}).get('text', ''))
            post['reviewed_' + f['key']] = 'on'
        response = self.client.post(f'/encounters/{e.pk}/save/', post)
        e.refresh_from_db()
        return response

    def test_reg_the_restored_labels_are_generated_finalized_and_exported_on_v3(self):
        template, version = self.uploaded(data=lines(*[f'{l}: ______' for l, _ in RESTORED], 'Plan: ______'),
                                          name='Synthetic REG v3')
        by = self.by_label(version)
        for label, _ in RESTORED:
            self.assertNotIn(by[label]['kind'], v3_mapping.SIGNING_KINDS, label)
        version = self.confirm(template)
        e = self.generated(version, notes_for(RESTORED))
        for label, value in RESTORED:
            with self.subTest(label=label):
                self.assertEqual(offered(e.answers[by[label]['key']]), value)
        held = {label: value for label, value in RESTORED if e.answers[by[label]['key']]['status'] == 'conflict'}
        self.assertEqual(self.rows_post(e, version, held).status_code, 302)
        self.finalized(e)
        self.assertEqual(e.status, 'finalized')
        exported = part_text(private_path(e.export_path).read_bytes())
        for label, value in RESTORED:
            self.assertIn(value, exported)

    def test_an_uncertain_place_is_shown_held_and_typed_by_the_clinician(self):
        template, version = self.uploaded(data=lines('Signed off by consultant: ______', 'Plan: ______'),
                                          name='Synthetic U v3')
        field = self.by_label(version)['Signed off by consultant']
        self.assertEqual(field['kind'], 'signature')
        self.assertEqual(self.client.post(f'/templates/{template.pk}/', {'action': 'v3_confirm', 'mapping_version': version.pk,
                         'checked_every_field': 'on'}).status_code, 302)
        e = self.generated(version, 'Signed off by consultant: Dr Example\nPlan: Rest.')
        page = self.client.get(f'/encounters/{e.pk}/').content.decode()
        self.assertIn('data-destination-role', page)                       # shown with what needs review
        self.assertIn('answer-fields page, then enter it yourself', page)
        # the clinician replaces the engine's signature field with a text field: allowed (uncertain is no role)
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_save', 'mapping_version': version.pk,
                                                                  'remove_' + field['key']: 'on'})
        self.assertEqual(response.status_code, 302)
        latest = self.latest(template)
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_add', 'mapping_version': latest.pk,
                                    'slot': BODY + 'b/p[1]', 'kind': 'text', 'label': 'Signed off by consultant'})
        self.assertEqual(response.status_code, 302)
        version = self.confirm(template)                                   # confirmed: held, not refused
        added = self.by_label(version)['Signed off by consultant']
        self.assertIn(added['key'], destination_roles.v3_uncertain(version))
        self.assertNotIn(added['key'], destination_roles.v3_protected(version))
        self.assertIn(added['key'], destination_roles.v3_held(version))     # a complete-form plan refuses it
        e = self.generated(version, 'Signed off by consultant: Dr Example\nPlan: Rest.')
        answer = e.answers[added['key']]
        self.assertEqual((answer['status'], answer['text']), ('missing', ''))  # never generated
        # generated text from an earlier draft is refused at save; the clinician's own entry exports
        Encounter.objects.filter(pk=e.pk).update(answers={**e.answers, added['key']: {
            'text': 'Dr Example', 'status': 'supported', 'evidence': [], 'reviewed': False}})
        e.refresh_from_db()
        response = self.rows_post(e, version, {})
        self.assertEqual(response.status_code, 400)
        self.assertIn('does not show whether each is a signing or initials place', response.content.decode())
        self.assertEqual(self.rows_post(e, version, {'Signed off by consultant': 'Yes (on paper)'}).status_code, 302)
        self.finalized(e)
        self.assertEqual(e.status, 'finalized')
        self.assertIn('Yes (on paper)', part_text(private_path(e.export_path).read_bytes()))

    def test_finalize_refuses_generated_text_on_an_uncertain_v3_place(self):
        template, version = self.uploaded(data=lines('Reviewed and signed: ______', 'Plan: ______'), name='Synthetic U v3 f')
        field = self.by_label(version)['Reviewed and signed']
        self.client.post(f'/templates/{template.pk}/', {'action': 'v3_save', 'mapping_version': version.pk,
                                                        'remove_' + field['key']: 'on'})
        self.client.post(f'/templates/{template.pk}/', {'action': 'v3_add', 'mapping_version': self.latest(template).pk,
                         'slot': BODY + 'b/p[1]', 'kind': 'text', 'label': 'Reviewed and signed'})
        version = self.confirm(template)
        added = self.by_label(version)['Reviewed and signed']
        e = self.generated(version, 'Plan: Rest.')
        answers = {k: {**v, 'reviewed': True, 'reviewed_by': str(self.user.pk)} for k, v in e.answers.items()}
        answers[added['key']] = {'text': 'Dr Example', 'status': 'supported', 'evidence': [], 'reviewed': True,
                                 'reviewed_by': str(self.user.pk)}
        Encounter.objects.filter(pk=e.pk).update(answers=answers, status='review')
        response = self.finalized(e)
        self.assertIn('never fills', response.content.decode())
        self.assertEqual(e.status, 'review')

    def test_a_stored_mapping_on_a_newly_protected_line_is_refused_at_confirmation_and_export(self):
        """L1:F2: an f5fe720f mapping holds a text field on the signature line under a declaration; roles-2 refuses it."""
        data = lines('Please complete in block capitals.', '______________________', 'Signature of clinician',
                     '______________________', 'Name and grade')
        template, version = self.uploaded(data=data, name='Synthetic F2 stored')
        line = next(f for f in version.fields if f['locator']['slot'] == 'b/p[2]')
        self.assertEqual(line['kind'], 'signature')                        # the engine's own field: never generated
        stored = TemplateVersion.objects.create(template=template, number=2, contract='v3', created_by=self.user,
                                                fields=[{**f, 'kind': 'text', 'label': 'Please complete in block capitals.'}
                                                        if f['key'] == line['key'] else f for f in version.fields])
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_confirm', 'mapping_version': stored.pk,
                                                                   'checked_every_field': 'on'})
        self.assertEqual(response.status_code, 400)
        with self.assertRaises(word_v3.V3Failure) as refused:
            word_v3.fill(data, stored.fields, {line['key']: 'Dr Somebody Else'})
        self.assertEqual(refused.exception.code, 'protected_signing_area')

    def test_b2_n6_an_initial_reading_is_generated_without_a_replacement(self):
        template, version = self.uploaded(data=lines('Initial and ongoing management: ______', 'Plan: ______'),
                                          name='Synthetic N6')
        field = self.by_label(version)['Initial and ongoing management']
        self.assertNotIn(field['kind'], v3_mapping.SIGNING_KINDS)
        version = self.confirm(template)
        e = self.generated(version, 'Initial and ongoing management: IV antibiotics\nPlan: Rest.')
        self.assertEqual(e.answers[field['key']]['text'], 'IV antibiotics')

    def test_the_reading_is_named_and_reused_per_template(self):
        template, version = self.uploaded(data=lines('Name: ______', 'Plan: ______'), name='Synthetic reading')
        self.assertEqual(destination_roles.report(version)['reading'], word_v3.ROLE_READING)


from core import mapping_assistance as mapping  # noqa: E402
from core.tests.test_mapping_assistance import SETTINGS as MAPPING_SETTINGS  # noqa: E402


@override_settings(**MAPPING_SETTINGS)
class ModelProposalTestGaps(V3Fixture, TestCase):
    """L2:RR3-06 / L2:RR3-08 / L3:F3-03: the model-path arms no test pinned (mutants MF and MG survived the focused
    suites), and a model proposal never moves a place out of protected or uncertain (roles-2)."""

    def setUp(self):
        super().setUp()
        self.config = generation.configuration()

    def proposal(self, version, replaced=(), dismissed=(), changes=()):
        compiled = mapping.compile_request(version, self.config)
        payload = json.loads(compiled['request']['messages'][0]['content'])
        body = mapping.wire.fixture(payload)
        fids = {f['key']: n for n, f in enumerate(payload['fields'])}
        body['f']['k'] = [n for n in body['f']['k'] if n not in {fids[k] for k in replaced}]
        body['f']['r'] = sorted(fids[k] for k in replaced)
        sids = {row['slot']: n for n, row in enumerate(payload['slots'])}
        for slot in dismissed:
            for verdict in ('c', 'u', 'n'):
                body['s'][verdict] = [n for n in body['s'][verdict] if n != sids[slot]]
            body['s']['n'] = sorted(body['s']['n'] + [sids[slot]])
        body['x'] = [change(payload, fids, sids) for change in changes]
        return mapping.validate_proposal(version, compiled, body)

    @staticmethod
    def relabel(key, slot, kind, start, end, role=0):
        def build(payload, fids, sids):
            sid = sids[BODY + slot]
            return [fids[key], sid, mapping.wire.KINDS.index(kind), payload['text_refs'].index([sid, role, 0]), start, end, -1]
        return build

    def test_rr3_08_mutant_mf_a_label_with_no_question_is_no_label(self):
        """'?' cleans to a non-empty label whose question key is empty: only LABEL_SHAPE's question_key arm refuses it
        (mutant MF keeps 'if not label' alone and adopts it). A punctuation-only '(' still passes both arms (residual
        L2:RR3-06, mapping_assistance.py - handed off, not owned here); this test pins what the code refuses today."""
        text = 'Pulse findings? ______'
        template, version = self.uploaded(data=lines(text, 'Plan: ______'), name='Synthetic MF')
        pulse = version.fields[0]
        at = text.index('?')
        result = self.proposal(version, replaced=[pulse['key']],
                               changes=[self.relabel(pulse['key'], 'b/p[1]', pulse['kind'], at, at + 1)])
        self.assertEqual(result['review']['changes'], [])
        self.assertIn(mapping.LABEL_SHAPE, ' '.join(r['reason'] for r in result['review']['unresolved']))

    def test_rr3_08_mutant_mg_a_change_on_a_dismissed_destination_names_the_dismissal(self):
        doc = Document(); doc.add_heading('Section 3: Public protection', level=2)
        t = doc.add_table(rows=2, cols=2); t.style = 'Table Grid'
        t.cell(0, 0).text = 'Name of the lead agency'; t.cell(1, 0).text = 'Date the lead agency was informed'
        doc.add_paragraph('Plan: ______')
        template, version = self.uploaded(data=saved(doc), name='Synthetic RR3-08')
        date = next(f for f in version.fields if f['kind'] == 'date')
        cell = date['locator']['slot']
        result = self.proposal(version, replaced=[date['key']], dismissed=[BODY + cell],
                               changes=[self.relabel(date['key'], cell + '/p[0]', 'date', 0, 20, role=2)])
        self.assertEqual(result['review']['changes'], [])
        self.assertIn(mapping.DISMISSED_REASON.format(slot=BODY + cell),
                      ' '.join(r['reason'] for r in result['review']['unresolved']))

    def test_a_model_relabel_never_moves_an_uncertain_or_protected_place(self):
        """The engine's field on an uncertain place ('Sign') and on a protected one ('Patient signature') is a
        Signature field: a model replacement relabelling it as ordinary text is never adopted."""
        template, version = self.uploaded(data=lines('Sign: ______', 'Patient signature: ______', 'Plan: ______'),
                                          name='Synthetic model roles2')
        for label in ('Sign', 'Patient signature'):
            with self.subTest(label=label):
                field = self.by_label(version)[label]
                self.assertEqual(field['kind'], 'signature')
                result = self.proposal(version, replaced=[field['key']],
                                       changes=[self.relabel(field['key'], field['locator']['slot'], 'text', 0, len(label))])
                self.assertEqual(result['review']['changes'], [])
                self.assertEqual(next(f for f in result['fields'] if f['key'] == field['key']), field)
