"""Existing imports recover through a fresh, linked inspection, never a changed provenance stamp."""
import copy
import hashlib
import io
import json
import os
from pathlib import Path
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from django.utils import timezone

from core import import_binding
from core.jobs import enqueue
from core.models import AuditEvent, Department, Encounter, Job, MappingConfirmation, Membership, Reservation, Template, TemplateVersion, User
from core.security import private_path
from .test_word_v3_workflow import V3Fixture, build_legacy_table_form, part_text
from .test_workflow import TEST_SETTINGS

# Actual generic V3 fingerprint from the independently pinned 5bf30245 source,
# before the original-preserving D01 layout change. Only the FIRST fixture import
# represents that older code; all review, admission and export routes run normally.
OLD_PROCESSOR = 'ce385128ac1133cf9879388ff01100333fc77ab7bd4a33b346aef4f44e007a40'
SETTINGS = dict(TEST_SETTINGS, QUILLMD_WORD_ENGINE='v3', QUILLMD_ROLE_SUGGESTIONS=False)


@override_settings(**SETTINGS)
class TemplateReinspection(V3Fixture, TestCase):
    def old_form(self):
        with patch('documents.service.engine_fingerprint', return_value=OLD_PROCESSOR):
            return self.uploaded(data=build_legacy_table_form(), name='Synthetic preserved form')

    def recheck(self, template, version, **extra):
        return self.client.post(f'/templates/{template.pk}/', {'action': 'reinspect_original',
            'mapping_version': version.pk, 'review_reinspection': 'on', **extra})

    def fresh(self, template, version):
        response = self.recheck(template, version)
        self.assertEqual(response.status_code, 302, response.content.decode())
        fresh = Template.objects.get(pk=response.url.rstrip('/').split('/')[-1])
        return fresh, fresh.versions.get(number=1)

    def archive(self, label, **records):
        """Seal real original/export bytes before assertions or temporary MEDIA_ROOT teardown."""
        root = os.environ.get('QUILLMD_WORD_ARTIFACTS')
        if not root:
            return
        target = Path(root) / self._testMethodName / label
        target.mkdir(parents=True, exist_ok=False)
        hashes = {}
        for name, data in records.items():
            data = data if isinstance(data, bytes) else (json.dumps(data, sort_keys=True, indent=2, default=str) + '\n').encode()
            (target / name).write_bytes(data)
            hashes[name] = hashlib.sha256(data).hexdigest()
        (target / 'hashes.json').write_text(json.dumps(hashes, indent=2) + '\n')

    def finish(self, version):
        self.assertEqual(self.new_encounter(version, notes='Allergies: None documented\nPlan: Rest').status_code, 302)
        encounter = Encounter.objects.get(template_version=version)
        payload = {'revision': encounter.revision}
        for field in version.fields:
            payload['answer_' + field['key']] = {'Allergies': 'None documented', 'Plan': 'Rest'}[field['label']]
            payload['reviewed_' + field['key']] = 'on'
        response = self.client.post(f'/encounters/{encounter.pk}/save/', payload)
        self.assertEqual(response.status_code, 302, response.content.decode())
        encounter.refresh_from_db()
        response = self.client.post(f'/encounters/{encounter.pk}/finalize/',
            {'revision': encounter.revision, 'reviewed': 'on'}, follow=True)
        encounter.refresh_from_db()
        exported = private_path(encounter.export_path).read_bytes() if encounter.export_path else b''
        self.archive('final-route', **{'original.docx': private_path(version.template.original_path).read_bytes(),
            'export.docx': exported, 'response.html': response.content,
            'state.json': {'status': encounter.status, 'export_sha256': encounter.export_sha256,
                'mapping': version.fields, 'import_manifest': version.import_manifest, 'answers': encounter.answers}})
        self.assertEqual(encounter.status, 'finalized', [str(m) for m in response.context['messages']])
        self.assertEqual(hashlib.sha256(exported).hexdigest(), encounter.export_sha256)
        self.assertIn('None documented', part_text(exported)); self.assertIn('Rest', part_text(exported))
        return encounter, exported

    def test_old_import_corrections_stay_stale_and_fresh_original_needs_its_own_review(self):
        old, version = self.old_form()
        original = private_path(old.original_path).read_bytes()
        prior = copy.deepcopy(version.import_manifest)
        plan = self.by_label(version)['Plan']
        response = self.client.post(f'/templates/{old.pk}/', {'action': 'v3_save', 'mapping_version': version.pk,
            'label_' + plan['key']: 'Reviewed plan'})
        self.assertEqual(response.status_code, 302)
        corrected = self.latest(old)
        self.assertEqual(corrected.import_manifest['engine_code_sha256'], OLD_PROCESSOR)
        refusal = self.client.post(f'/templates/{old.pk}/', {'action': 'v3_confirm', 'mapping_version': corrected.pk,
            'checked_every_field': 'on'})
        self.assertEqual(refusal.status_code, 400)
        self.assertContains(self.client.get(f'/templates/{old.pk}/'), 'Recheck answer locations')
        fresh, current = self.fresh(old, corrected)
        self.assertNotEqual(fresh.pk, old.pk); self.assertNotEqual(fresh.original_path, old.original_path)
        self.assertEqual(private_path(fresh.original_path).read_bytes(), original)
        self.assertEqual(fresh.sha256, old.sha256)
        self.assertNotEqual(current.import_manifest['engine_code_sha256'], OLD_PROCESSOR)
        self.assertEqual(set(self.by_label(current)), {'Allergies', 'Plan'}, 'Old correction is not silently copied')
        self.assertEqual(current.import_manifest['reinspection']['source_version_id'], str(corrected.pk))
        self.assertTrue(current.needs_confirmation)
        self.assertEqual(self.new_encounter(current).status_code, 200); self.assertFalse(Encounter.objects.exists())
        current = self.confirm(fresh, current)
        encounter, exported = self.finish(current)
        self.assertEqual(encounter.template_version_id, current.pk)
        version.refresh_from_db(); self.assertEqual(version.import_manifest, prior)
        self.assertEqual(old.versions.count(), 2); self.assertEqual(private_path(old.original_path).read_bytes(), original)
        event = AuditEvent.objects.get(action='template.reinspected', target=str(fresh.pk))
        self.assertEqual(event.details['source_version_id'], str(corrected.pk))
        self.assertEqual(event.details['source_sha256'], old.sha256)
        self.assertContains(self.client.get(f'/templates/{old.pk}/'), str(fresh.pk))
        self.assertContains(self.client.get(f'/templates/{fresh.pk}/'), str(old.pk))
        self.assertEqual(Job.objects.count(), 0); self.assertEqual(Reservation.objects.count(), 0)

    @override_settings(QUILLMD_WORD_ENGINE='legacy')
    def test_new_legacy_reinspection_requires_review_before_encounter_job_or_finalize(self):
        old, version = self.uploaded(data=build_legacy_table_form(), name='Synthetic legacy source')
        self.assertFalse(version.needs_confirmation, 'Historical ordinary legacy contract is preserved')
        fresh, current = self.fresh(old, version)
        self.assertEqual(current.contract, ''); self.assertTrue(current.needs_confirmation)
        self.assertEqual(self.new_encounter(current).status_code, 200)
        encounter = Encounter.objects.create(department=self.department, patient=self.patient, template_version=current,
            notes='Plan: Rest', created_by=self.user, service_date=timezone.localdate())
        with self.assertRaises(ValidationError): enqueue(encounter, self.user)
        self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'})
        encounter.refresh_from_db(); self.assertNotEqual(encounter.status, 'finalized')
        self.assertFalse(Job.objects.exists()); self.assertFalse(Reservation.objects.exists())
        response = self.client.post(f'/templates/{fresh.pk}/', {'action': 'confirm_reinspection',
            'mapping_version': current.pk, 'checked_every_field': 'on'})
        self.assertEqual(response.status_code, 302, response.content.decode())
        current.refresh_from_db(); self.assertFalse(current.needs_confirmation)
        import_binding.verify(current)
        self.assertEqual(self.new_encounter(current).status_code, 302)

    def test_retry_reopens_current_fresh_inspection_without_another_original(self):
        old, version = self.old_form(); fresh, current = self.fresh(old, version)
        originals = set(Path(self.directory.name).glob('private/originals/*'))
        retry = self.recheck(old, version)
        self.assertEqual(retry.status_code, 302); self.assertIn(str(fresh.pk), retry.url)
        self.assertEqual(Template.objects.count(), 2)
        self.assertEqual(set(Path(self.directory.name).glob('private/originals/*')), originals)
        self.assertEqual(AuditEvent.objects.filter(action='template.reinspected').count(), 1)

    def test_acknowledgement_current_version_and_clinician_authority_are_required(self):
        old, version = self.old_form()
        self.assertEqual(self.recheck(old, version, review_reinspection='').status_code, 400)
        self.assertEqual(self.recheck(old, version, mapping_version='00000000-0000-0000-0000-000000000000').status_code, 400)
        for user in (self.manager, self.employee):
            self.client.force_login(user); self.assertEqual(self.recheck(old, version).status_code, 403)
        stranger = User.objects.create_user(username='foreign-synthetic', email='foreign@synthetic.invalid', verified_at=timezone.now())
        department = Department.objects.create(organization=self.department.organization, name='Other department')
        Membership.objects.create(user=stranger, department=department, role='clinician')
        self.client.force_login(stranger)
        self.assertEqual(self.recheck(old, version).status_code, 404)
        self.assertEqual(self.client.get(f'/templates/{old.pk}/original/').status_code, 404)
        self.assertEqual(Template.objects.count(), 1); self.assertFalse(MappingConfirmation.objects.exists())

    def test_corrupt_preserved_original_is_not_reinspected_or_allocated(self):
        old, version = self.old_form()
        path = private_path(old.original_path); original = path.read_bytes(); path.write_bytes(original + b'changed')
        before = set(path.parent.iterdir()); response = self.recheck(old, version)
        self.assertEqual(response.status_code, 400); self.assertEqual(Template.objects.count(), 1)
        self.assertEqual(set(path.parent.iterdir()), before)
        self.assertIn(b'Original integrity check failed', response.content)

    def test_historical_download_reads_stored_bytes_without_reinspection_or_current_engine(self):
        old, version = self.uploaded(data=build_legacy_table_form(), name='Synthetic finalized predecessor')
        version = self.confirm(old, version); encounter, exported = self.finish(version)
        before = copy.deepcopy(encounter.answers); fresh, current = self.fresh(old, version)
        with patch('documents.service.engine_fingerprint', side_effect=AssertionError('Downloads must not require current engine')), \
             patch('documents.inspect_template', side_effect=AssertionError('Downloads must not reimport')):
            response = self.client.get(f'/encounters/{encounter.pk}/export/')
            self.assertEqual(b''.join(response.streaming_content), exported)
            response = self.client.get(f'/templates/{old.pk}/original/')
            self.assertEqual(b''.join(response.streaming_content), private_path(old.original_path).read_bytes())
        encounter.refresh_from_db(); self.assertEqual(encounter.template_version_id, version.pk); self.assertEqual(encounter.answers, before)
        stranger = User.objects.create_user(username='download-stranger', email='stranger@synthetic.invalid', verified_at=timezone.now())
        self.client.force_login(stranger); self.assertEqual(self.client.get(f'/encounters/{encounter.pk}/export/').status_code, 404)

    def test_forged_current_v3_destination_cannot_be_confirmed_after_fresh_reinspection(self):
        old, version = self.old_form(); fresh, current = self.fresh(old, version)
        fields = copy.deepcopy(current.fields)
        fields[0]['locator']['slot'] = 'b/tbl[99]/tr[0]/tc[1]'
        invalid = TemplateVersion.objects.create(template=fresh, number=2, fields=fields, contract=current.contract,
            import_manifest=import_binding.for_mapping(current.import_manifest, fields), created_by=self.user)
        response = self.client.post(f'/templates/{fresh.pk}/', {'action': 'v3_confirm',
            'mapping_version': invalid.pk, 'checked_every_field': 'on'})
        self.archive('invalid-physical-destination', **{'original.docx': private_path(fresh.original_path).read_bytes(),
            'fields.json': fields, 'response.html': response.content,
            'result.json': {'status': response.status_code, 'confirmed': bool(TemplateVersion.objects.get(pk=invalid.pk).confirmed)}})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(MappingConfirmation.objects.exists()); self.assertFalse(Job.objects.exists()); self.assertFalse(Reservation.objects.exists())

    def test_valid_manual_recovery_and_ordinary_label_corrections_can_be_reviewed(self):
        old, version = self.old_form(); fresh, current = self.fresh(old, version)
        allergies = self.by_label(current)['Allergies']; plan = self.by_label(current)['Plan']
        response = self.client.post(f'/templates/{fresh.pk}/', {'action': 'v3_save', 'mapping_version': current.pk,
            'remove_' + allergies['key']: 'on', 'label_' + plan['key']: 'Care plan'})
        self.assertEqual(response.status_code, 302)
        current = self.latest(fresh)
        response = self.client.post(f'/templates/{fresh.pk}/', {'action': 'v3_add', 'mapping_version': current.pk,
            'slot': allergies['locator']['part'] + '#' + allergies['locator']['slot'], 'kind': 'text', 'label': 'Known allergies'})
        self.assertEqual(response.status_code, 302, response.content.decode())
        current = self.latest(fresh); manual = self.by_label(current)['Known allergies']
        self.assertEqual(manual['origin'], 'manual')
        current = self.confirm(fresh, current); self.assertFalse(current.needs_confirmation)
        self.assertEqual(self.new_encounter(current).status_code, 302)
        self.assertEqual(current.import_manifest['reinspection']['source_version_id'], str(version.pk))

    def test_known_destination_swap_and_changed_fingerprint_are_not_ordinary_corrections(self):
        template, current = self.uploaded(data=build_legacy_table_form(), name='Synthetic current physical mapping')
        for number, change in enumerate(('swap', 'fingerprint'), 2):
            fields = copy.deepcopy(current.fields)
            if change == 'swap': fields[0]['locator'] = copy.deepcopy(fields[1]['locator'])
            else: fields[0]['locator']['fingerprint'] = '0' * 64
            invalid = TemplateVersion.objects.create(template=template, number=number, fields=fields, contract=current.contract,
                import_manifest=import_binding.for_mapping(current.import_manifest, fields), created_by=self.user)
            response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_confirm',
                'mapping_version': invalid.pk, 'checked_every_field': 'on'})
            self.assertEqual(response.status_code, 400, change)
        self.assertFalse(MappingConfirmation.objects.exists()); self.assertFalse(Job.objects.exists()); self.assertFalse(Reservation.objects.exists())

    def test_original_legacy_field_checkbox_and_dropdown_anchors_remain_reviewable(self):
        from docx import Document
        from documents.test_word_v3 import ff_field
        document = Document(); document.add_paragraph('Synthetic clinical form')
        paragraph = document.add_paragraph('Activity: ')
        ff_field(paragraph, 'FORMDROPDOWN', options=('Rest', 'Exercise'), name='Activity')
        paragraph = document.add_paragraph('Fever: ')
        ff_field(paragraph, 'FORMCHECKBOX', checked=False, name='Fever')
        stream = io.BytesIO(); document.save(stream)
        template, current = self.uploaded(data=stream.getvalue(), name='Synthetic legacy field locations')
        self.assertEqual({field['kind'] for field in current.fields}, {'choice', 'checkbox'})
        current = self.confirm(template, current)
        self.assertFalse(current.needs_confirmation)
        self.assertEqual(self.new_encounter(current).status_code, 302)

    @override_settings(QUILLMD_WORD_ENGINE='legacy')
    def test_legacy_review_refuses_changed_physical_destination_and_missing_acknowledgement(self):
        old, version = self.uploaded(data=build_legacy_table_form(), name='Synthetic legacy binding')
        fresh, current = self.fresh(old, version)
        self.assertEqual(self.client.post(f'/templates/{fresh.pk}/', {'action': 'confirm_reinspection', 'mapping_version': current.pk}).status_code, 400)
        fields = copy.deepcopy(current.fields); fields[0]['locator']['index'] = 999
        invalid = TemplateVersion.objects.create(template=fresh, number=2, fields=fields,
            import_manifest=import_binding.for_mapping(current.import_manifest, fields), created_by=self.user)
        response = self.client.post(f'/templates/{fresh.pk}/', {'action': 'confirm_reinspection',
            'mapping_version': invalid.pk, 'checked_every_field': 'on'})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(MappingConfirmation.objects.exists()); self.assertFalse(Job.objects.exists()); self.assertFalse(Reservation.objects.exists())
