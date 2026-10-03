"""Explicit source review cannot be replaced by a generic confirmation or copied writer."""
import copy
import io
from pathlib import Path
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import TestCase,override_settings
from docx import Document

from core import import_binding
from core.jobs import enqueue,run_one
from core.models import Encounter,Job,Reservation,TemplateVersion
from core.security import private_path
from documents.test_source_coverage import numbered
from .test_word_v3_workflow import V3Fixture
from .test_workflow import TEST_SETTINGS


@override_settings(**dict(TEST_SETTINGS,QUILLMD_WORD_ENGINE='v3',QUILLMD_ROLE_SUGGESTIONS=False))
class SourceReviewWorkflow(V3Fixture,TestCase):
    def source(self,shape):
        doc=Document()
        if shape=='list':
            numbered(doc,'Gender',90,'decimal')
            for item in ('Male','Female','Other'):numbered(doc,item,91,'lowerLetter')
        elif shape=='unbalanced':doc.add_paragraph('Site: [left/right')
        else:doc.add_paragraph('The care received was right for [me/us].')
        doc.add_paragraph('Reason for visit: ______')
        out=io.BytesIO();doc.save(out);return doc,out.getvalue()

    def test_exact_review_preserves_printed_sources_and_ordinary_field_workflow(self):
        for shape in ('list','alternatives','unbalanced'):
            with self.subTest(shape=shape):
                original,raw=self.source(shape)
                template,version=self.uploaded(data=raw,name='Synthetic exact source review '+shape)
                issues=version.import_manifest['coverage']['selected_engine_issues']
                self.assertTrue(issues)
                with self.assertRaises(ValidationError):import_binding.verify(version)
                response=self.client.post(f'/templates/{template.pk}/',dict(action='v3_confirm',mapping_version=version.pk,
                                                                          checked_every_field='on'))
                self.assertEqual(response.status_code,400)
                self.assertEqual(self.new_encounter(version).status_code,200)
                self.assertFalse(Encounter.objects.filter(template_version=version).exists())
                self.assertEqual(Job.objects.filter(encounter__template_version=version).count(),0)
                response=self.client.post(f'/templates/{template.pk}/',dict(action='v3_coverage',mapping_version=version.pk,
                                            **{'exclude_'+issue['id']:'on' for issue in issues}))
                self.assertEqual(response.status_code,302,response.content.decode())
                reviewed=self.latest(template);self.assertNotEqual(version.pk,reviewed.pk)
                self.assertEqual(version.import_manifest,reviewed.import_manifest)
                self.assertFalse(reviewed.import_manifest['coverage']['automatic_complete'])
                import_binding.verify(reviewed);self.confirm(template,reviewed)
                self.assertEqual(self.new_encounter(reviewed,'Reason for visit: Synthetic review').status_code,302)
                encounter=Encounter.objects.get(template_version=reviewed);job=enqueue(encounter,self.user)
                self.assertTrue(run_one());job.refresh_from_db();self.assertEqual(job.state,'succeeded',job.error)
                encounter.refresh_from_db();key=reviewed.fields[0]['key']
                response=self.client.post(f'/encounters/{encounter.pk}/save/',{'revision':encounter.revision,
                                    'answer_'+key:'Synthetic review','reviewed_'+key:'on'})
                self.assertEqual(response.status_code,302);encounter.refresh_from_db()
                response=self.client.post(f'/encounters/{encounter.pk}/finalize/',{'revision':encounter.revision,'reviewed':'on'})
                self.assertEqual(response.status_code,302);encounter.refresh_from_db();self.assertEqual(encounter.status,'finalized')
                export=Document(private_path(encounter.export_path))
                self.assertEqual([p.text for p in export.paragraphs[:-1]],[p.text for p in original.paragraphs[:-1]])
                self.assertEqual(export.paragraphs[-1].text,'Reason for visit: Synthetic review')

    def test_precise_exclusion_cannot_be_copied_into_by_a_forged_manual_field(self):
        _,raw=self.source('alternatives');template,version=self.uploaded(data=raw)
        issue=version.import_manifest['coverage']['selected_engine_issues'][0]
        version.corrections={'coverage_review':[{'id':issue['id'],'disposition':'leave_unchanged',
                                 'source_sha256':template.sha256,'reviewed_by':str(self.user.pk)}]}
        self.assertTrue(import_binding.coverage_reviewed(version))
        fields=copy.deepcopy(version.fields)
        fields.append(dict(key='manual_alternative',label='Forged free text',kind='text',locator={
            'part':issue['part'],'slot':issue['slot'],'target':dict(type='span',start=issue['start'],end=issue['end'])}))
        version.fields=fields;version.import_manifest=import_binding.for_mapping(version.import_manifest,fields)
        self.assertFalse(import_binding.coverage_reviewed(version))
        with self.assertRaises(ValidationError):import_binding.verify(version)

    def test_unqualified_fixed_header_refuses_generic_review_before_job_or_reservation(self):
        from documents.test_source_coverage import SourceCoverage
        doc=Document();SourceCoverage().fixed_header(doc);doc.add_paragraph('Reason: ______')
        out=io.BytesIO();doc.save(out);template,version=self.uploaded(data=out.getvalue())
        self.assertTrue(version.import_manifest['coverage']['block_generation'])
        with self.assertRaises(ValidationError):import_binding.verify(version)
        response=self.client.post(f'/templates/{template.pk}/',{'action':'v3_confirm',
                                 'mapping_version':version.pk,'checked_every_field':'on'})
        self.assertEqual(response.status_code,400)
        self.assertEqual(self.new_encounter(version).status_code,200)
        self.assertEqual(Job.objects.count(),0);self.assertEqual(Reservation.objects.count(),0)

    def test_measured_document_quarantine_refuses_even_precise_exclusion_before_money(self):
        _,raw=self.source('alternatives');template,version=self.uploaded(data=raw)
        manifest=copy.deepcopy(version.import_manifest)
        quarantined=TemplateVersion.objects.create(template=template,number=2,fields=version.fields,
                    contract=version.contract,created_by=self.user,import_manifest=manifest,
                    corrections={'coverage_review':[{'id':issue['id'],'disposition':'leave_unchanged',
                                  'source_sha256':template.sha256,'reviewed_by':str(self.user.pk)}
                                 for issue in manifest['coverage']['selected_engine_issues']]})
        self.assertEqual(quarantined.import_manifest,version.import_manifest)
        from documents.word_v3 import KNOWN_LAYOUT_FAILURES
        declaration={'failed_anchors':[version.fields[0]['locator']],
                     'evidence':['synthetic measured-failure control']}
        with patch.dict(KNOWN_LAYOUT_FAILURES,{template.sha256:declaration}):
            self.assertFalse(import_binding.coverage_reviewed(quarantined))
            with self.assertRaisesMessage(ValidationError,'retained rendered answer failures'):
                import_binding.verify(quarantined)
            response=self.client.post(f'/templates/{template.pk}/',{'action':'v3_confirm',
                                     'mapping_version':quarantined.pk,'checked_every_field':'on'})
            self.assertEqual(response.status_code,400)
            self.assertEqual(self.new_encounter(quarantined).status_code,200)
            self.assertEqual(Job.objects.count(),0);self.assertEqual(Reservation.objects.count(),0)

    def test_spaced_date_source_has_valid_binding_and_genuine_export(self):
        from documents.service import fill_template
        doc=Document();doc.add_paragraph('Date: [DD / MM / YYYY]')
        out=io.BytesIO();doc.save(out);template,version=self.uploaded(data=out.getvalue())
        self.assertFalse(version.import_manifest['coverage'].get('block_generation'))
        self.confirm(template,version);import_binding.verify(version)
        field=version.fields[0];self.assertEqual(field['kind'],'date')
        target=Path(self.directory.name)/'spaced-date-export.docx'
        fill_template(private_path(template.original_path),version.fields,{field['key']:'03/14/2026'},target)
        self.assertEqual(Document(target).paragraphs[0].text,'Date: 03/14/2026')
