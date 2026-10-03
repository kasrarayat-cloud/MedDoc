"""Signer ownership and withdrawn offers under the October 3 continuation."""
from copy import deepcopy
from types import SimpleNamespace
from unittest.mock import patch

from django.test import SimpleTestCase, TestCase, override_settings

from core import destination_roles, generation
from core.jobs import enqueue, run_one
from core.models import AuditEvent, Encounter, Job, LedgerEntry, Usage
from core.tab_review import another_signers_line, clinician_identity, validate_identity
from core.tests.test_destination_roles import legacy_rows
from core.tests.test_role_suggestions import ACTION, Legacy, LIVE, NOTES, ON, labelled_provider
from core.tests.test_v3_supersession_floor import SONNET_55_DEFAULT


class SignerOwnership(SimpleTestCase):
    def test_unknown_and_other_person_roles_are_never_the_finalizers_identity(self):
        for label in ('Signature (participant)', 'Signed (mother)', 'Signature (unfamiliar role)',
                      'Signature (participant, dd/mm/yyyy)', 'Signature (patient, date)',
                      'Patient signature', 'Witness signature', 'Manager signature',
                      'Signed (Consultant)', 'Signed (print name)'):
            with self.subTest(label=label):
                field = {'key': 'role', 'label': label}
                self.assertTrue(another_signers_line(field))
                self.assertFalse(clinician_identity(field))
                self.assertFalse(clinician_identity({**field, 'answer_policy': 'clinician_identity'}))
                mapped = {'key': 'role', 'label': 'Signature', 'answer_policy': 'clinician_identity'}
                printed = destination_roles.printed(mapped, {'role': {'printed_label': label, 'role': 'signature'}})
                self.assertEqual(printed['label'], label)
                self.assertFalse(clinician_identity(printed))
                self.assertEqual(destination_roles.legacy_reason(field, {'printed_label': label, 'role': 'signature'}),
                                 destination_roles.ROLE_NOT_GENERATED)

    def test_original_structural_role_cannot_become_the_finalizer_by_rename_or_old_policy(self):
        for label, role in (('Signature of participant','signature'),('Signature of witness','signature'),
                            ('Signature of patient','signature'),('Signed off by manager','signature'),
                            ('Initials','initials')):
            with self.subTest(label=label):
                field={'key':'k','label':'Signature','answer_policy':'clinician_identity'}
                row={'printed_label':label,'role':role}
                source=destination_roles.printed(field,{'k':row})
                self.assertEqual(source['label'],label)
                self.assertFalse(clinician_identity(source))
                self.assertEqual(destination_roles.legacy_reason(field,row),destination_roles.ROLE_NOT_GENERATED)
        mapped={'key':'k','label':'Signature','answer_policy':'clinician_identity'}
        self.assertTrue(clinician_identity(destination_roles.printed(mapped,{'k':{'printed_label':'Doctor','role':''}})))

    def test_explicit_clinician_identity_and_ordinary_date_questions_keep_their_workflows(self):
        actor = SimpleNamespace(pk=1, email='synthetic@invalid', get_full_name=lambda: 'Synthetic Clinician')
        for label in ('Signature', 'Electronically signed by', 'Doctor signature', 'Signature (clinician)'):
            self.assertTrue(clinician_identity({'key': 'k', 'label': label}))
            validate_identity([{'key': 'k', 'label': label}], {'k': {
                'text': 'Synthetic Clinician', 'reviewed': True, 'reviewed_by': '1'}}, actor)
        for label in ('Signature date', 'Date of signature', 'Signature (date and time)'):
            self.assertFalse(clinician_identity({'label': label}))
            self.assertFalse(another_signers_line({'label': label}))


@override_settings(**ON)
class WithdrawnSignerOffers(Legacy, TestCase):
    @override_settings(**dict(SONNET_55_DEFAULT, AI_MODE='fixture', QUILLMD_AI_MODE='fixture',
                              AI_THINKING='between_tools', V3_FIELD_RELATIONS=True,
                              QUILLMD_WORD_ENGINE='auto', QUILLMD_WORD_AUTO_PREFER_V3=True))
    def test_initial_and_signed_receipt_answers_reach_actual_bound_http_export_without_correction(self):
        import hashlib
        from docx import Document
        from core.security import private_path
        desired = {'Initial diagnosis': 'Sprain', 'Signed consent was received on': '3 October 2026', 'Plan': 'Rest.'}
        self.template, self.version = self.uploaded(data=legacy_rows(*desired), name='Synthetic ordinary source roles')
        self.assertEqual(self.version.contract, 'v3')
        self.assertFalse(self.version.import_manifest['coverage']['block_generation'])
        original = private_path(self.template.original_path).read_bytes()
        self.version = self.confirm(self.template, self.version); self.by = self.by_label(self.version)
        notes = '\n\n'.join(label + ': ' + value for label, value in desired.items())
        e = self.generated(notes)
        job=Job.objects.get(encounter=e,kind='generation')
        self.assertEqual(job.state,'succeeded',job.error)
        for label, value in desired.items():
            self.assertEqual(e.answers[self.key(label)]['text'], value)
            self.assertEqual(e.answers[self.key(label)]['status'], 'supported')
        first = deepcopy(e.answers)
        self.assertEqual(self.review_all(e).status_code, 302)
        self.assertEqual({k: a['text'] for k, a in e.answers.items()}, {k: a['text'] for k, a in first.items()})
        # The stored import determines export even when the new-import default changes.
        with override_settings(QUILLMD_WORD_ENGINE='legacy'): self.finalized(e)
        self.assertEqual(e.status, 'finalized')
        exported = private_path(e.export_path).read_bytes()
        import io
        document = Document(io.BytesIO(exported))
        for n, (label, value) in enumerate(desired.items()):
            self.assertEqual(document.tables[0].cell(n, 0).text, label)
            self.assertEqual(document.tables[0].cell(n, 1).text, value)
        self.assertEqual(private_path(self.template.original_path).read_bytes(), original)
        self.assertEqual(hashlib.sha256(exported).hexdigest(), e.export_sha256)
        download = self.client.get(f'/encounters/{e.pk}/export/')
        self.assertEqual(download.status_code, 200)
        self.assertEqual(b''.join(download.streaming_content), exported)

    def test_flag_one_admits_no_role_questions_and_preserves_useful_ordinary_answer(self):
        for flag in ('0', '1'):
            with self.subTest(flag=flag):
                e, job, payload = self.live(flag=flag)
                self.assertEqual(job.state, 'succeeded', job.error)
                self.assertNotIn('role_suggestion_contract', job.config)
                self.assertEqual([q['label'] for q in payload['questions']], ['Plan'])
                self.assertEqual(e.answers[self.key('Plan')]['text'], 'Rest.')
                self.assertFalse(any(a.get('role_suggestion') for a in e.answers.values()))
                self.assertFalse(destination_roles.suggestions_enabled())
                self.assertFalse(any((f.get('destination_role') or {}).get('suggest') for f in
                                     destination_roles.generation_fields(self.version, suggest=True)))

    def test_old_queued_contract_is_kept_but_never_asks_signer(self):
        e = self.encounter(); sent = []
        # Simulates the immutable admission an earlier release already made.
        # This changes neither a stored job nor its settlement after creation.
        with override_settings(**dict(LIVE, QUILLMD_ROLE_SUGGESTIONS='1')):
            with patch.object(generation, 'admitted_with_role_suggestions',
                              side_effect=lambda config, version: {**config,
                                  'role_suggestion_contract': generation.ROLE_SUGGESTION_CONTRACT}):
                job = enqueue(Encounter.objects.get(pk=e.pk), self.user)
            recorded = deepcopy(job.config)
            with labelled_provider(sent): self.assertTrue(run_one())
        job.refresh_from_db(); e.refresh_from_db()
        self.assertEqual(job.config, recorded)
        self.assertEqual(job.state, 'succeeded', job.error)
        import json
        payload = json.loads(sent[0]['messages'][0]['content'])
        self.assertEqual([q['label'] for q in payload['questions']], ['Plan'])
        self.assertEqual(e.answers[self.key('Plan')]['text'], 'Rest.')
        self.assertFalse(any(a.get('role_suggestion') for a in e.answers.values()))
        self.assertEqual(Usage.objects.filter(job=job).count(), 1)

    def test_old_stored_offer_is_unavailable_without_answer_accounting_or_audit_rewrites(self):
        e = self.generated(); key = self.key('Troponin (initial)')
        start = NOTES.index('14 ng/L')
        e.answers[key]['role_suggestion'] = {'text': '14 ng/L', 'evidence': [
            {'quote': '14 ng/L', 'start': start, 'end': start + len('14 ng/L')}]}
        e.save(update_fields=['answers'])
        before = (deepcopy(e.answers), e.revision, e.status,
                  list(LedgerEntry.objects.values()), list(Usage.objects.values()), list(Job.objects.values()))
        roles = destination_roles.legacy_roles(self.version)
        self.assertEqual(destination_roles.current_suggestions(self.version.fields, e.answers, roles, e.notes), {})
        self.assertNotIn('data-role-suggestion', self.client.get(f'/encounters/{e.pk}/').content.decode())
        self.assertEqual(self.accept(e).status_code, 404)
        e.refresh_from_db()
        self.assertEqual((e.answers, e.revision, e.status, list(LedgerEntry.objects.values()),
                          list(Usage.objects.values()), list(Job.objects.values())), before)
        self.assertFalse(AuditEvent.objects.filter(action=ACTION).exists())

    def test_unknown_printed_signer_stays_human_only_after_mapping_rename(self):
        import json
        for label in ('Signature (participant)', 'Patient signature', 'Manager signature',
                      'Signature of participant', 'Signature of witness', 'Signature of patient', 'Signed off by manager'):
            with self.subTest(label=label):
                self.template, self.version = self.uploaded(data=legacy_rows(label, 'Plan'),
                                                            name='Synthetic printed role ' + label)
                self.by = self.by_label(self.version); key = self.key(label)
                labels = {f['key']: ('Signature' if f['key'] == key else f['label']) for f in self.version.fields}
                self.assertEqual(self.client.post(f'/templates/{self.template.pk}/', {
                    'mapping_version': self.version.pk, 'labels': json.dumps(labels)}).status_code, 302)
                self.version = self.latest(self.template)
                e = self.generated('Signature: Synthetic finalizer\nPlan: Rest.')
                self.assertEqual((e.answers[key]['status'], e.answers[key]['text']), ('missing', ''))
                row = next(r for r in self.client.get(f'/encounters/{e.pk}/').context['answer_rows'] if r['key'] == key)
                self.assertFalse(row['identity'])
                self.assertTrue(row['blank_eligible'])
                self.assertEqual(self.review_all(e).status_code, 302)
                self.finalized(e)
                self.assertEqual(e.status, 'finalized')
                # The finalized document retains the signer's area blank.
                from core.security import private_path
                from docx import Document
                self.assertEqual(Document(private_path(e.export_path)).tables[0].cell(0, 1).text, '')
