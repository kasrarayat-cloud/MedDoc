"""Approved name destinations, source roles and unchanged Word destinations.

Provider fixtures are intentionally independent of the local name requirement:
a passing checker must not make an omitted or ambiguous identity approvable.
"""
from copy import deepcopy
from pathlib import Path
import tempfile

from django.core.exceptions import ValidationError
from django.test import SimpleTestCase
from docx import Document

from core import semantic_generation as semantic
from core.tests.test_semantic_generation import FixtureSender, assignments, fact, alex_case, NOTES
from documents import inspect_template, fill_template
from documents.semantic_preparation_roles import default_definitions


def soap_definitions(fields):
    definitions = default_definitions(fields)
    for key, definition in definitions.items():
        definition.update(version=3, answer_mode='narrative', patient_name_requirement='none')
    definitions[fields[0]['key']].update(patient_name_requirement='documented',
        meaning='The complete documented patient name followed by the reported symptoms.')
    return definitions


class DocumentIdentityTests(SimpleTestCase):
    fields = [{'key': 'subjective', 'label': 'Subjective', 'kind': 'text', 'locator': {'cell': 0}}]
    notes = 'Synthetic Avery Sen reports heartburn and constipation.'

    def case(self, notes=None):
        notes = notes or self.notes
        fields = deepcopy(self.fields)
        definitions = soap_definitions(fields)
        facts = [fact(notes, fields, 'identity', 'Synthetic Avery Sen', kind='identity', name_quote='Synthetic Avery Sen'),
                 fact(notes, fields, 'symptom', 'heartburn', kind='history')]
        return notes, fields, definitions, facts

    def generate(self, notes, fields, definitions, facts, selected, **kwargs):
        sender = FixtureSender(facts, assignments(fields, selected), **kwargs)
        answers = semantic.run(notes, fields, definitions, sender,
            scope_contract=semantic.DOCUMENT_IDENTITY_SCOPE_CONTRACT)
        for field in fields:
            semantic.validate_stored_answer(notes, field, definitions[field['key']], answers[field['key']])
        return answers, sender

    def test_explicit_name_is_in_actual_answer_not_only_its_evidence(self):
        notes, fields, definitions, facts = self.case()
        answers, sender = self.generate(notes, fields, definitions, facts, {'subjective':['identity','symptom']})
        answer = answers['subjective']
        self.assertEqual(answer['text'], 'Synthetic Avery Sen; Heartburn')
        self.assertEqual(answer['patient_name_check']['verdict'], 'pass')
        self.assertFalse(answer['reviewed'])
        self.assertEqual(len(sender.calls), 3)
        for call in sender.calls:
            self.assertEqual(call[2]['approved_patient_name_destinations'], ['subjective'])

    def test_checker_pass_does_not_waive_omission_and_repair_is_bounded(self):
        notes, fields, definitions, facts = self.case()
        missing = assignments(fields, {'subjective':['symptom']})
        for repaired, expected in [(missing, 'ambiguous'),
                (assignments(fields, {'subjective':['identity','symptom']}), 'supported')]:
            with self.subTest(expected=expected):
                answers, sender = self.generate(notes, fields, definitions, facts,
                    {'subjective':['symptom']}, repaired=repaired)
                answer = answers['subjective']
                self.assertEqual(answer['status'], expected)
                self.assertEqual(answer['independent_check']['verdict'], 'pass')
                self.assertEqual(len(sender.calls), 5)
                self.assertEqual(sender.calls[3][2]['local_patient_name_findings']['subjective']['verdict'], 'omission')
                if expected == 'ambiguous':
                    self.assertEqual(answer['text'], '')
                    self.assertEqual(answer['candidate_text'], 'Heartburn')
                    self.assertEqual(answer['patient_name_check']['verdict'], 'omission')

    def test_patient_referrer_and_relative_are_separate_people(self):
        notes = self.notes + ' Referring clinician: Dr. Priya Chen. The patient’s mother, Synthetic Elin Sen, has diabetes.'
        notes, fields, definitions, facts = self.case(notes)
        facts += [fact(notes, fields, 'referrer', 'Priya Chen', kind='identity',
                       subject='clinician', name_quote='Priya Chen'),
                  fact(notes, fields, 'relative', 'Synthetic Elin Sen', kind='identity',
                       subject='other', name_quote='Synthetic Elin Sen')]
        answers, _ = self.generate(notes, fields, definitions, facts, {'subjective':['identity','symptom']})
        self.assertEqual(answers['subjective']['text'], 'Synthetic Avery Sen; Heartburn')
        self.assertEqual(answers['subjective']['patient_name_check']['fact_ids'], ['identity'])
        for wrong in ('referrer', 'relative'):
            with self.subTest(wrong=wrong):
                selected = {'subjective':[wrong,'symptom']}
                answers, _ = self.generate(notes, fields, definitions, facts, selected,
                                          repaired=assignments(fields, selected))
                self.assertNotEqual(answers['subjective']['status'], 'supported')
                self.assertEqual(answers['subjective']['text'], '')

    def test_missing_name_is_held_without_inventing_registry_name(self):
        notes = 'The patient reports heartburn.'
        fields = deepcopy(self.fields); definitions = soap_definitions(fields)
        facts = [fact(notes, fields, 'symptom', 'heartburn', kind='history')]
        answers, sender = self.generate(notes, fields, definitions, facts, {'subjective':['symptom']})
        self.assertEqual(answers['subjective']['patient_name_check']['verdict'], 'missing')
        self.assertEqual(answers['subjective']['status'], 'ambiguous')
        self.assertEqual(len(sender.calls), 3)

    def test_clipped_names_and_clinician_subject_labels_cannot_pass(self):
        cases = [('Alex Morgan reports heartburn.', 'Alex'),
                 ('Alex Morgan reports heartburn.', 'Morgan'),
                 ('Nurse Alex Morgan reports heartburn. The patient reports constipation.', 'Alex Morgan'),
                 ('Dr. Alex Morgan reports that the patient has heartburn.', 'Alex Morgan'),
                 ('Patient name: Alex Morgan.', 'Alex'),
                 ('Patient name: Alex Morgan.', 'Morgan')]
        for notes, name in cases:
            with self.subTest(notes=notes, name=name):
                fields = deepcopy(self.fields); definitions = soap_definitions(fields)
                facts = [fact(notes, fields, 'identity', name, kind='identity', name_quote=name)]
                answers, sender = self.generate(notes, fields, definitions, facts, {'subjective':['identity']})
                self.assertEqual(answers['subjective']['status'], 'ambiguous')
                self.assertEqual(answers['subjective']['text'], '')
                self.assertEqual(len(sender.calls), 3)

    def test_explicit_patient_name_slot_accepts_complete_unicode_name(self):
        notes = 'Patient name: Zoë van Dijk.'
        fields = deepcopy(self.fields); definitions = soap_definitions(fields)
        facts = [fact(notes, fields, 'identity', 'Zoë van Dijk', kind='identity', name_quote='Zoë van Dijk')]
        answers, _ = self.generate(notes, fields, definitions, facts, {'subjective':['identity']})
        self.assertEqual(answers['subjective']['text'], 'Zoë van Dijk')

    def test_competing_or_uncertain_identity_is_held(self):
        notes = self.notes + ' Patient name: Synthetic Avery Ren.'
        notes, fields, definitions, facts = self.case(notes)
        facts.append(fact(notes, fields, 'alternative', 'Synthetic Avery Ren', kind='identity', name_quote='Synthetic Avery Ren'))
        answers, _ = self.generate(notes, fields, definitions, facts, {'subjective':['identity','symptom']})
        self.assertEqual(answers['subjective']['patient_name_check']['verdict'], 'ambiguous')
        self.assertEqual(answers['subjective']['status'], 'ambiguous')

    def test_omitted_competing_source_name_does_not_manufacture_uniqueness(self):
        notes = 'Patient name: Alex Morgan. Patient name: Jamie Rivers. The patient reports heartburn.'
        fields = deepcopy(self.fields); definitions = soap_definitions(fields)
        facts = [fact(notes, fields, 'identity', 'Alex Morgan', kind='identity', name_quote='Alex Morgan'),
                 fact(notes, fields, 'symptom', 'heartburn', kind='history')]
        answers, _ = self.generate(notes, fields, definitions, facts, {'subjective':['identity','symptom']})
        check = answers['subjective']['patient_name_check']
        self.assertEqual(answers['subjective']['status'], 'ambiguous')
        self.assertEqual(check['unrepresented_source_names'][0]['quote'], 'Jamie Rivers')
        notes, fields, definitions, facts = self.case()
        facts[0]['subject']['kind'] = 'uncertain'
        answers, _ = self.generate(notes, fields, definitions, facts, {'subjective':['identity','symptom']})
        self.assertEqual(answers['subjective']['status'], 'ambiguous')

    def test_unapproved_requirement_and_forged_persisted_name_refuse(self):
        notes, fields, definitions, facts = self.case()
        with self.assertRaisesMessage(ValidationError, 'document-identity generation contract'):
            semantic.run(notes, fields, definitions, FixtureSender(facts, assignments(fields)))
        answers, _ = self.generate(notes, fields, definitions, facts, {'subjective':['identity','symptom']})
        for change in ('text', 'finding', 'selected'):
            changed = deepcopy(answers['subjective'])
            if change == 'text': changed['text'] = 'Registry Replacement; heartburn'
            elif change == 'finding': changed['patient_name_check']['fact_ids'] = ['symptom']
            else: changed['fact_ids'] = ['symptom']
            with self.subTest(change=change), self.assertRaises(ValidationError):
                semantic.validate_stored_answer(notes, fields[0], definitions['subjective'], changed)

    def test_historical_generic_subjective_is_not_reinterpreted(self):
        notes, fields, _, facts = self.case()
        definitions = default_definitions(fields)
        answers = semantic.run(notes, fields, definitions,
            FixtureSender(facts, assignments(fields, {'subjective':['symptom']})),
            scope_contract=semantic.CONSENT_CONTEXT_SCOPE_CONTRACT)
        self.assertEqual(answers['subjective']['text'], 'Heartburn')
        self.assertNotIn('patient_name_check', answers['subjective'])
        semantic.validate_stored_answer(notes, fields[0], definitions['subjective'], answers['subjective'])

    def test_custom_form_name_and_other_fields_keep_their_existing_destinations(self):
        fields, definitions, facts, selected = alex_case()
        for definition in definitions.values():
            definition.update(version=3, answer_mode='narrative' if definition['value_type']=='text' else 'atomic',
                              patient_name_requirement='none')
        definitions['cell_0'].update(value_type='person_name', answer_mode='atomic', patient_name_requirement='documented')
        answers, _ = self.generate(NOTES, fields, definitions, facts,
                                  {row['key']:row['fact_ids'] for row in selected['fields']})
        self.assertEqual(answers['cell_0']['text'], 'Synthetic Alex Morgan')
        self.assertEqual(answers['cell_4']['text'], '3')
        self.assertIn('No head injury', answers['cell_2']['text'])
        with tempfile.TemporaryDirectory() as directory:
            original = Path(directory)/'custom.docx'
            document = Document(); table = document.add_table(rows=len(fields), cols=2)
            for row, field in zip(table.rows, fields): row.cells[0].text = field['label']
            document.save(original)
            detected = inspect_template(original)['fields']
            self.assertEqual([f['locator'] for f in detected], [f['locator'] for f in fields])
            rendered = Path(directory)/'filled.docx'
            fill_template(original, detected, {key:a['text'] for key,a in answers.items()}, rendered)
            exported = Document(rendered)
            self.assertEqual([r.cells[0].text for r in exported.tables[0].rows], [f['label'] for f in fields])
            self.assertEqual([r.cells[1].text for r in exported.tables[0].rows], [answers[f['key']]['text'] for f in fields])
