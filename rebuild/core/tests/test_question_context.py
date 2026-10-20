"""Real Word question boundaries, immutable destinations and historical readers."""
from copy import deepcopy
import hashlib
from pathlib import Path
import tempfile
from types import SimpleNamespace
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, TestCase, override_settings
from docx import Document

from core import preparation, word_contract
from core.template_source import approval_binding, physical_source_for_version
from documents.component_contracts import CONFIG_KEY as COMPONENT_KEY, ENCOUNTER_CONTRACT, literal_part
from documents.engine import DocumentFailure, docx_inspect, structure
from documents.preparation import TYPED_PROTOCOL, SIGNING_CONTEXT, candidate, digest, inventory, validate_inventory
from documents.preparation_service import inspect_preparation, prepare_candidate
from documents.question_context import CONFIG_KEY, CONTEXT, project_fields
from documents.semantic_preparation_roles import MEANING_CONTRACT_KEY, COMPOUND_ANSWER_CONTRACT
from documents.service import fill_template
from documents.table_constraints import with_table_constraints


class QuestionContextTests(SimpleTestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.path = self.root / 'blank.docx'

    def save_questions(self, questions, *, columns=2):
        document = Document()
        table = document.add_table(rows=(len(questions) + columns // 2 - 1) // (columns // 2), cols=columns)
        for index, question in enumerate(questions):
            cell = table.rows[index // (columns // 2)].cells[2 * (index % (columns // 2))]
            paragraphs = question if isinstance(question, list) else [question]
            cell.text = paragraphs[0]
            for paragraph in paragraphs[1:]:
                cell.add_paragraph(paragraph)
        document.save(self.path)
        return document

    def current(self):
        return inventory(self.path, protocol=TYPED_PROTOCOL, signing_context=SIGNING_CONTEXT,
                         question_context=CONTEXT)

    def config(self):
        return {'mode': 'fixture', 'provider': 'anthropic', 'model': 'offline-only',
                'input_rate': '2', 'output_rate': '10', 'cached_rate': '0.2',
                word_contract.CONFIG_KEY: word_contract.GRAPH_PROFILE,
                COMPONENT_KEY: ENCOUNTER_CONTRACT, 'semantic_definition_version': 4,
                MEANING_CONTRACT_KEY: COMPOUND_ANSWER_CONTRACT,
                CONFIG_KEY: CONTEXT}

    def test_only_real_paragraph_tab_and_break_boundaries_enable_complete_quotes(self):
        document = self.save_questions([['Symptom history', 'Describe onset and frequency.'],
            'Objective\tExclude assessment.', 'Family history\nRecord documented conditions.'])
        # A styled run boundary must remain inside the same word.
        cell = document.tables[0].rows[0].cells[0]
        cell.paragraphs[0].clear()
        cell.paragraphs[0].add_run('Symptom his').bold = True
        cell.paragraphs[0].add_run('tory')
        document.save(self.path)
        old = inventory(self.path, protocol=TYPED_PROTOCOL, signing_context=SIGNING_CONTEXT)
        current = self.current()
        self.assertEqual([f['label'] for f in docx_inspect(self.path)[3]], [f['label'] for f in old['fields']])
        self.assertNotIn('question_context', old)
        self.assertEqual(current['fields'][0]['label'], 'Symptom history\nDescribe onset and frequency.')
        for field, quote in zip(current['fields'], ['Symptom history', 'Objective', 'Family history']):
            self.assertTrue(literal_part(field['label'], quote))
        self.assertFalse(literal_part(old['fields'][0]['label'], 'Symptom history'))
        for fragment in ['his', 'tory', 'Symptom his', 'toryDescribe']:
            self.assertFalse(literal_part(current['fields'][0]['label'], fragment))
        self.assertEqual(structure(current['fields']), structure(old['fields']))
        self.assertEqual(current['parts'], old['parts'])
        self.assertEqual(current['stories'], old['stories'])
        self.assertNotEqual(current['inventory_sha256'], old['inventory_sha256'])

    def test_same_paragraph_concatenation_and_styled_runs_do_not_create_boundaries(self):
        document = self.save_questions(['HistoryDescribe onset.', 'brain'])
        cell = document.tables[0].rows[1].cells[0]
        cell.paragraphs[0].clear()
        cell.paragraphs[0].add_run('b')
        cell.paragraphs[0].add_run('rain').italic = True
        document.save(self.path)
        fields = self.current()['fields']
        self.assertEqual(fields[0]['label'], 'HistoryDescribe onset.')
        self.assertEqual(fields[1]['label'], 'brain')
        self.assertFalse(literal_part(fields[0]['label'], 'History'))
        self.assertFalse(literal_part(fields[1]['label'], 'rain'))

    def test_alternating_pairs_and_duplicate_titles_use_exact_preceding_cells(self):
        self.save_questions([['History', 'First instructions.'], ['History', 'Different instructions.'],
            ['Objective', 'Third instructions.'], ['Plan', 'Fourth instructions.']], columns=4)
        value = self.current()
        self.assertEqual([f['label'] for f in value['fields']], ['History\nFirst instructions.',
            'History\nDifferent instructions.', 'Objective\nThird instructions.', 'Plan\nFourth instructions.'])
        for kind in ['locator', 'label', 'key', 'missing', 'duplicate']:
            changed = deepcopy(value['fields'])
            if kind == 'locator': changed[0]['locator'] = deepcopy(changed[1]['locator'])
            if kind == 'label': changed[0]['label'] = changed[1]['label']
            if kind == 'key': changed[0]['key'] = changed[1]['key']
            if kind == 'missing': changed.pop()
            if kind == 'duplicate': changed.append(deepcopy(changed[0]))
            with self.subTest(kind=kind), self.assertRaises(DocumentFailure):
                project_fields(self.path, changed, CONTEXT)

    def test_physical_constraints_remain_identical_and_actual_word_export_rederives_them(self):
        self.save_questions([['Weight (kg)', 'Record measured weight.'], ['History', 'Describe current symptoms.']])
        old = inventory(self.path, protocol=TYPED_PROTOCOL)
        value = self.current()
        self.assertEqual(structure(old['fields']), structure(value['fields']))
        self.assertEqual(with_table_constraints(self.path, old['fields']), old['fields'])
        self.assertEqual(with_table_constraints(self.path, value['fields']), value['fields'])
        prepared = self.root / 'prepared.docx'
        result = prepare_candidate(self.path, value, [a['id'] for a in value['anchors']], prepared)
        self.assertEqual(result['prepared_fields'], value['fields'])
        output = self.root / 'filled.docx'
        fill_template(prepared, result['prepared_fields'], {'cell_0': '68', 'cell_1': 'Reported cough.'}, output)
        filled = Document(output)
        self.assertEqual(filled.tables[0].rows[0].cells[0].text, 'Weight (kg)\nRecord measured weight.')
        self.assertEqual(filled.tables[0].rows[0].cells[1].text, '68')
        self.assertEqual(filled.tables[0].rows[1].cells[1].text, 'Reported cough.')
        with self.assertRaises(ValueError):
            fill_template(prepared, result['prepared_fields'], {'cell_0': '68 lb'}, self.root / 'wrong.docx')

    def test_boundary_sensitive_constraints_stop_before_provider_and_signatures_stay_protected(self):
        self.save_questions([['Weight (kg)', '°F']])
        old = inventory(self.path, protocol=TYPED_PROTOCOL)
        self.assertEqual(old['fields'][0]['answer_constraints']['unit'], 'kg')
        with patch('anthropic.Anthropic') as client, self.assertRaisesMessage(DocumentFailure, 'constraints ambiguous'):
            self.current()
        client.assert_not_called()
        self.save_questions([['Patient', 'signature'], ['Clinical vital sign', 'Record measurement.'],
            ['Clinician', 'sign below']])
        value = self.current()
        self.assertEqual([q['protected_signing'] for q in value['questions']], [True, False, True])
        self.assertEqual(value['questions'][1]['text'], 'Clinical vital sign\nRecord measurement.')
        self.assertFalse(any(i['code'] == 'guided_signing' for i in value['issues']))

    def test_candidate_rederives_source_even_when_forged_inventory_is_redigested(self):
        self.save_questions([['History', 'Describe symptoms.'], ['Objective', 'Exclude assessment.']])
        value = self.current()
        original = self.path.read_bytes()
        output = self.root / 'pure-candidate.docx'
        candidate(self.path, output, value, [a['id'] for a in value['anchors']])
        self.assertEqual(output.read_bytes(), original)
        for kind in ['question', 'label', 'locator', 'source', 'part']:
            changed = deepcopy(value)
            if kind == 'question': changed['questions'][0]['text'] = changed['questions'][1]['text']
            if kind == 'label': changed['fields'][0]['label'] = changed['fields'][1]['label']
            if kind == 'locator': changed['fields'][0]['locator'] = changed['fields'][1]['locator']
            if kind == 'source': changed['source_sha256'] = '0' * 64
            if kind == 'part': changed['parts']['word/document.xml'] = '0' * 64
            changed['inventory_sha256'] = digest({k:v for k,v in changed.items() if k != 'inventory_sha256'})
            with self.subTest(kind=kind), self.assertRaisesMessage(DocumentFailure, 'stale inventory'):
                candidate(self.path, self.root / (kind + '.docx'), changed, [a['id'] for a in changed['anchors']])
        self.assertEqual(self.path.read_bytes(), original)

    def test_unknown_or_unbound_context_cannot_send_or_reserve_and_absence_stays_exact(self):
        self.save_questions([['History', 'Describe symptoms.']])
        old = inventory(self.path, protocol=TYPED_PROTOCOL)
        self.assertEqual(old, inventory(self.path, protocol=TYPED_PROTOCOL))
        current = self.current()
        for unknown in [None, False, 1, '', 'future-source-questions']:
            changed = deepcopy(current); changed['question_context'] = unknown
            changed['inventory_sha256'] = digest({k:v for k,v in changed.items() if k != 'inventory_sha256'})
            with self.subTest(unknown=unknown), self.assertRaises(DocumentFailure): validate_inventory(changed)
            with patch('anthropic.Anthropic') as client, patch('core.testing_authority.reserve_role') as reserve:
                with self.assertRaises(ValueError):
                    preparation.provider_role(current, 'proposer', {**self.config(), CONFIG_KEY: unknown}, 1)
                client.assert_not_called(); reserve.assert_not_called()
        for source, config in [(current, {}), (old, self.config())]:
            with self.assertRaises(ValueError): preparation.role_payload(source, 'proposer', config=config)
        with self.assertRaises(DocumentFailure): inventory(self.path, question_context=CONTEXT)
        with self.assertRaises(ValueError): inspect_preparation(self.path, protocol=TYPED_PROTOCOL, question_context='future')
        self.assertEqual(inspect_preparation(self.path, protocol=TYPED_PROTOCOL,
            signing_context=SIGNING_CONTEXT, question_context=CONTEXT), current)

    @override_settings(WORD_PREPARATION_ENABLED=True, RELEASE_SHA='offline-only')
    def test_new_queue_binds_context_without_mutating_configuration_or_downgrading_unknown(self):
        config = self.config(); config.pop(CONFIG_KEY)
        before = deepcopy(config)
        with patch.object(preparation, 'configuration', return_value=config), \
                patch('core.preparation_campaign.active_campaign_config', side_effect=lambda value, _: value), \
                patch('core.testing_authority.identity', return_value=True), \
                patch('core.testing_authority.create_preparation') as create:
            preparation.create_preparation(SimpleNamespace(sha256='a' * 64), SimpleNamespace())
        self.assertEqual(create.call_args.args[2][CONFIG_KEY], CONTEXT)
        self.assertEqual(config, before)
        for unknown in [None, False, 'future-question-context']:
            with self.subTest(unknown=unknown), patch.object(preparation, 'configuration',
                    return_value={**config, CONFIG_KEY: unknown}), \
                    patch('core.preparation_campaign.active_campaign_config', side_effect=lambda value, _: value), \
                    patch('core.testing_authority.create_preparation') as create:
                with self.assertRaises(ValueError):
                    preparation.create_preparation(SimpleNamespace(sha256='a' * 64), SimpleNamespace())
                create.assert_not_called()

    def test_approved_source_binds_policy_labels_and_unchanged_original_bytes(self):
        self.save_questions([['History', 'Describe current symptoms.'], ['Objective', 'Describe examination findings.']])
        value = self.current(); config = self.config()
        proposal, _ = preparation.provider_role(value, 'proposer', config, 1)
        checker, _ = preparation.provider_role(value, 'checker', config, 1, proposal)
        prepared = self.root / 'prepared.docx'
        result = prepare_candidate(self.path, value, [a['id'] for a in value['anchors']], prepared)
        template = SimpleNamespace(sha256=value['source_sha256'], original_path='blank.docx', format='docx')
        prep = SimpleNamespace(template_id='template', template=template, inventory=value, config=config,
            proposal=proposal, checker=checker, corrections={}, correction_history=[], revision=0,
            candidate_sha256=hashlib.sha256(prepared.read_bytes()).hexdigest(), candidate_path='prepared.docx',
            candidate_info={'fields':result['prepared_fields'], 'document_backend':{}, 'review_requirements':[]},
            changes=[], state='approved', approved_by_id='doctor', approved_at='synthetic-review-time')
        word_contract.seal(prep, [f['key'] for f in value['fields']], SimpleNamespace(pk='doctor'))
        prep.approval_sha256 = approval_binding(prep)
        version = SimpleNamespace(template=template, template_id='template', preparation=prep, preparation_id='prep',
            fields=deepcopy(value['fields']), document_backend={}, review_requirements=[],
            prepared_path='prepared.docx', prepared_sha256=prep.candidate_sha256)
        with override_settings(MEDIA_ROOT=self.root):
            self.assertEqual(physical_source_for_version(version), prepared)
            self.assertEqual(word_contract.definitions(version), word_contract.proposed(prep))
            # Historical physical-only approvals permit display-label edits.
            # Existing inline semantic approvals already require exact sealed
            # labels through definitions(); the new policy adds no old gate.
            legacy = deepcopy(version)
            legacy.preparation.config = {}
            legacy.preparation.inventory = inventory(self.path, protocol=TYPED_PROTOCOL)
            legacy.preparation.candidate_info = {'fields': deepcopy(legacy.preparation.inventory['fields']),
                'document_backend': {}, 'review_requirements': []}
            legacy.fields = deepcopy(legacy.preparation.inventory['fields'])
            legacy.fields[0]['label'] = 'Clinician display label'
            legacy.preparation.approval_sha256 = approval_binding(legacy.preparation)
            self.assertEqual(physical_source_for_version(legacy), prepared)
            for kind in ['config', 'inventory', 'version_label', 'candidate_label', 'source_bytes']:
                changed = deepcopy(version)
                if kind == 'config': changed.preparation.config.pop(CONFIG_KEY)
                if kind == 'inventory': changed.preparation.inventory.pop('question_context')
                if kind == 'version_label': changed.fields[0]['label'] = 'HistoryDescribe current symptoms.'
                if kind == 'candidate_label': changed.preparation.candidate_info['fields'][0]['label'] = 'Other question'
                if kind == 'source_bytes': self.path.write_bytes(self.path.read_bytes() + b'changed')
                with self.subTest(kind=kind), self.assertRaises((ValidationError, ValueError)):
                    physical_source_for_version(changed)


@override_settings(WORD_PREPARATION_ENABLED=True, AI_MODE='fixture', ADMIN_SITE=False,
    TESTING_AUTHORITY_MODE='', TESTING_AUTHORITY_ID='', TESTING_AUTHORITY_SHA256='',
    STORAGES={'default':{'BACKEND':'django.core.files.storage.FileSystemStorage'},
              'staticfiles':{'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class QuestionContextApprovalTests(TestCase):
    def setUp(self):
        from core.tests import test_word_answer_contract
        test_word_answer_contract.OrdinaryClinicianWordTests.setUp(self)

    def test_actual_upload_roles_separate_synthetic_approval_reload_and_word_fill_keep_source_boundaries(self):
        from core.tests import test_word_answer_contract
        from core.models import TemplateVersion
        from core.security import private_path
        original = Document(self.root / 'blank.docx')
        for row in original.tables[0].rows:
            # Real paragraphs, without editing the fixture's original title.
            row.cells[0].add_paragraph('Retain the documented answer to this question.')
        original.save(self.root / 'blank.docx')
        raw = (self.root / 'blank.docx').read_bytes()
        template, prep = test_word_answer_contract.OrdinaryClinicianWordTests.fresh_upload(self)
        self.assertEqual(prep.config[CONFIG_KEY], CONTEXT)
        self.assertEqual(prep.inventory['question_context'], CONTEXT)
        self.assertTrue(all('\n' in f['label'] for f in prep.inventory['fields']))
        for response in (prep.proposal, prep.checker):
            self.assertEqual(response['inventory_sha256'], prep.inventory['inventory_sha256'])
        version = test_word_answer_contract.OrdinaryClinicianWordTests.approve(self, template, prep)
        version = TemplateVersion.objects.select_related('preparation', 'template').get(pk=version.pk)
        self.assertEqual(version.fields, prep.inventory['fields'])
        self.assertTrue(word_contract.definitions(version))
        self.assertEqual(private_path(template.original_path).read_bytes(), raw)
        source = physical_source_for_version(version)
        self.assertEqual(source.read_bytes(), raw)
        filled = self.root / 'actual-export.docx'
        fill_template(source, version.fields, {}, filled)
        actual = Document(filled)
        self.assertEqual([r.cells[0].text for r in actual.tables[0].rows],
                         [r.cells[0].text for r in original.tables[0].rows])
        self.assertTrue(all(not r.cells[1].text for r in actual.tables[0].rows))
