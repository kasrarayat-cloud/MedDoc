"""Complete private workflows using synthetic files and the real document sandbox."""

from datetime import timedelta
from decimal import Decimal
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import tempfile
from zipfile import ZipFile

from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from docx import Document
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from core.billing import welcome_credit
from core.jobs import run_one
from core.models import AuditEvent, Department, Encounter, Invitation, Job, LedgerEntry, Membership, Organization, Patient, Template, TemplateVersion, Usage, User, Verification, Wallet
from core.security import digest, private_path


TEST_SETTINGS = {
    "AI_MODE": "fixture",
    "ADMIN_SITE": False,
    "EMAIL_BACKEND": "django.core.mail.backends.locmem.EmailBackend",
    "STORAGES": {"default": {"BACKEND": "django.core.files.storage.FileSystemStorage"}, "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"}},
}


@override_settings(**TEST_SETTINGS)
class DocumentWorkflow(TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.override = override_settings(MEDIA_ROOT=Path(self.directory.name) / "private")
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.addCleanup(self.directory.cleanup)
        self.user = User.objects.create_user(username="workflow-clinician", email="workflow@synthetic.invalid", verified_at=timezone.now())
        self.organization = Organization.objects.create(name="Synthetic workflow practice")
        self.department = Department.objects.create(organization=self.organization, name="Synthetic X-ray")
        Membership.objects.create(user=self.user, department=self.department, role="clinician")
        welcome_credit(self.user)
        self.client.force_login(self.user)
        response = self.client.post("/patients/", {"department": self.department.pk, "display_name": "SYNTHETIC WORKFLOW PATIENT", "reference": "SYNTHETIC-WORKFLOW", "synthetic_only": "on"})
        self.assertEqual(response.status_code, 302)
        self.patient = Patient.objects.get(reference="SYNTHETIC-WORKFLOW")

    def make_docx(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Symptoms: ").bold = True
        paragraph.add_run("{{Symptoms}}")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Examination"
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def make_pdf(self):
        buffer = BytesIO()
        document = canvas.Canvas(buffer)
        document.drawString(30, 780, "Symptoms: documented findings")
        document.acroForm.textfield(name="Symptoms", tooltip="Symptoms", x=30, y=730, width=300, height=30, fontName="Helvetica", fontSize=10)
        document.showPage()
        document.save()
        return buffer.getvalue()

    def upload(self, data, suffix):
        response = self.client.post("/templates/upload/", {"name": f"Synthetic {suffix} form", "department": self.department.pk, "synthetic_only": "on", "file": SimpleUploadedFile(f"synthetic.{suffix}", data)})
        self.assertEqual(response.status_code, 302, getattr(response, "context", None))
        template = Template.objects.get(name=f"Synthetic {suffix} form")
        self.assertEqual(private_path(template.original_path).read_bytes(), data)
        self.assertEqual(template.sha256, hashlib.sha256(data).hexdigest())
        return template, template.versions.get(number=1)

    def create_encounter(self, version, notes):
        response = self.client.post("/encounters/new/", {"patient": self.patient.pk, "template_version": version.pk, "notes": notes, "service_date": timezone.localdate().isoformat(), "synthetic_only": "on"})
        self.assertEqual(response.status_code, 302, getattr(response, "context", None))
        encounter = Encounter.objects.get(template_version=version)
        self.assertEqual(encounter.status, "draft")
        self.assertEqual(encounter.answers, {})
        return encounter

    def generate(self, encounter):
        response = self.client.post(f"/encounters/{encounter.pk}/generate/")
        self.assertEqual(response.status_code, 302)
        self.assertTrue(run_one())
        encounter.refresh_from_db()
        self.assertEqual(encounter.status, "review")
        self.assertEqual(Job.objects.get(encounter=encounter).state, "succeeded")
        usage = Usage.objects.get(job__encounter=encounter)
        self.assertTrue(usage.simulated)
        self.assertEqual(usage.actual_cost, 0)
        self.assertEqual(usage.charged_cost, 0)
        for answer in encounter.answers.values():
            self.assertEqual(answer["status"], "supported")
            self.assertFalse(answer["reviewed"])
            for evidence in answer["evidence"]:
                self.assertEqual(encounter.notes[evidence["start"]:evidence["end"]], evidence["quote"])

    def review(self, encounter, override=None):
        payload = {"revision": encounter.revision}
        for key, answer in encounter.answers.items():
            payload["answer_" + key] = override if override is not None else answer["text"]
            payload["reviewed_" + key] = "on"
        response = self.client.post(f"/encounters/{encounter.pk}/save/", payload)
        self.assertEqual(response.status_code, 302)
        encounter.refresh_from_db()
        self.assertTrue(all(answer["reviewed"] for answer in encounter.answers.values()))
        return payload

    def finalize(self, encounter):
        response = self.client.post(f"/encounters/{encounter.pk}/finalize/", {"reviewed": "on", "revision": encounter.revision})
        self.assertEqual(response.status_code, 302)
        encounter.refresh_from_db()
        self.assertEqual(encounter.status, "finalized")
        export = self.client.get(f"/encounters/{encounter.pk}/export/")
        self.assertEqual(export.status_code, 200)
        exported = b"".join(export.streaming_content)
        self.assertEqual(hashlib.sha256(exported).hexdigest(), encounter.export_sha256)
        return exported

    def test_word_upload_generate_review_export_and_finalization_guards(self):
        original = self.make_docx()
        template, mapping = self.upload(original, "docx")
        encounter = self.create_encounter(mapping, "Symptoms: No chest pain.\nExamination: Normal gait.")
        self.generate(encounter)
        # Finalization before per-field review must leave the document unfinalized.
        self.client.post(f"/encounters/{encounter.pk}/finalize/", {"reviewed": "on", "revision": encounter.revision})
        encounter.refresh_from_db()
        self.assertEqual(encounter.status, "review")
        stale_payload = self.review(encounter)
        response = self.client.post(f"/encounters/{encounter.pk}/save/", {**stale_payload, "answer_" + mapping.fields[0]["key"]: "Stale concurrent edit"})
        self.assertEqual(response.status_code, 409)
        exported = self.finalize(encounter)
        document = Document(BytesIO(exported))
        self.assertEqual(document.paragraphs[0].text, "Symptoms: No chest pain.")
        self.assertTrue(document.paragraphs[0].runs[0].bold)
        self.assertEqual(document.tables[0].cell(0, 0).text, "Examination")
        self.assertEqual(document.tables[0].cell(0, 1).text, "Normal gait.")
        self.assertEqual(private_path(template.original_path).read_bytes(), original)
        response = self.client.get(f"/templates/{template.pk}/original/")
        self.assertEqual(b"".join(response.streaming_content), original)
        self.assertEqual(self.client.post(f"/encounters/{encounter.pk}/save/", {"revision": encounter.revision}).status_code, 409)
        self.assertEqual(self.client.post(f"/encounters/{encounter.pk}/notes/", {"revision": encounter.revision, "notes": "Forbidden post-finalization source edit"}).status_code, 409)
        self.client.post(f"/encounters/{encounter.pk}/generate/")
        self.assertEqual(Job.objects.filter(encounter=encounter).count(), 1)
        saved_path = encounter.export_path
        self.client.post(f"/encounters/{encounter.pk}/finalize/", {"reviewed": "on", "revision": encounter.revision})
        encounter.refresh_from_db()
        self.assertEqual(encounter.export_path, saved_path)
        self.assertEqual(AuditEvent.objects.filter(action="encounter.finalize", target=str(encounter.pk)).count(), 1)

    def test_pdf_upload_generate_review_export_preserves_original_and_question(self):
        original = self.make_pdf()
        template, mapping = self.upload(original, "pdf")
        encounter = self.create_encounter(mapping, "Symptoms: No chest pain.")
        self.generate(encounter)
        self.review(encounter)
        exported = self.finalize(encounter)
        pdf = PdfReader(BytesIO(exported))
        self.assertEqual(pdf.get_form_text_fields()["Symptoms"], "No chest pain.")
        self.assertIn("Symptoms: documented findings", pdf.pages[0].extract_text())
        self.assertEqual(private_path(template.original_path).read_bytes(), original)

    def test_mapping_versions_do_not_rewrite_existing_encounter_or_original(self):
        original = self.make_docx()
        template, mapping = self.upload(original, "docx")
        encounter = self.create_encounter(mapping, "Symptoms: No chest pain.\nExamination: Normal gait.")
        labels = {field["key"]: "Revised " + field["label"] for field in mapping.fields}
        response = self.client.post(f"/templates/{template.pk}/", {"labels": json.dumps(labels), "mapping_version": str(mapping.pk)})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(template.versions.count(), 2)
        mapping.refresh_from_db()
        encounter.refresh_from_db()
        self.assertEqual(encounter.template_version_id, mapping.pk)
        self.assertTrue(all(not field["label"].startswith("Revised") for field in mapping.fields))
        self.assertEqual(private_path(template.original_path).read_bytes(), original)
        self.generate(encounter)
        self.review(encounter)
        self.finalize(encounter)

    def test_pdf_overflow_blocks_export_without_truncation(self):
        template, mapping = self.upload(self.make_pdf(), "pdf")
        encounter = self.create_encounter(mapping, "Symptoms: No chest pain.")
        self.generate(encounter)
        long_answer = "W" * 1000
        self.review(encounter, override=long_answer)
        response = self.client.post(f"/encounters/{encounter.pk}/finalize/", {"reviewed": "on", "revision": encounter.revision})
        self.assertEqual(response.status_code, 302)
        encounter.refresh_from_db()
        self.assertEqual(encounter.status, "review")
        self.assertEqual(encounter.answers["Symptoms"]["text"], long_answer)
        self.assertEqual(encounter.export_path, "")
        self.assertEqual(self.client.get(f"/encounters/{encounter.pk}/export/").status_code, 404)
        self.assertEqual(list((Path(self.directory.name) / "private" / "exports").glob("*")), [])

    def test_malicious_uploads_show_error_and_leave_no_private_or_database_artifact(self):
        malicious = BytesIO(self.make_docx())
        with ZipFile(malicious, "a") as package:
            package.writestr("../escape.txt", "synthetic traversal")
        for name, payload in [("fake.pdf", b"<script>synthetic</script>"), ("traversal.docx", malicious.getvalue()), ("corrupt.pdf", b"%PDF-1.7\nmalformed")]:
            with self.subTest(name=name):
                response = self.client.post("/templates/upload/", {"name": "Unsafe form", "department": self.department.pk, "synthetic_only": "on", "file": SimpleUploadedFile(name, payload)})
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.context["form"].errors)
                self.assertEqual(Template.objects.count(), 0)
                self.assertEqual(list((Path(self.directory.name) / "private" / "originals").glob("*")), [])


@override_settings(**TEST_SETTINGS)
class VerifiedWelcomeWorkflow(TestCase):
    def registration_payload(self, email):
        return {"email": email, "first_name": "Synthetic", "last_name": "Tester", "password1": "Synthetic-Only-4821!private", "password2": "Synthetic-Only-4821!private", "synthetic_only": "on"}

    def assert_one_personal_credit(self, user):
        self.assertEqual(LedgerEntry.objects.filter(wallet__user=user, kind="welcome").count(), 1)
        self.assertEqual(Wallet.objects.get(user=user, kind="promo").balance, Decimal("10"))
        self.assertEqual(Wallet.objects.get(user=user, kind="paid").balance, 0)
        self.assertFalse(Wallet.objects.filter(user=user, organization__isnull=False).exists())

    def test_provisioned_account_verification_grants_once_and_duplicate_is_rejected(self):
        # Existing, controlled accounts retain verification; no public signup.
        user = User.objects.create_user(username="new", email="new@synthetic.invalid",
            password="Synthetic-Only-4821!private", first_name="Synthetic", last_name="Tester")
        response = self.client.post("/verification/resend/", {"email": user.email})
        self.assertEqual(response.status_code, 200, getattr(response, "context", None))
        self.assertIsNone(user.verified_at)
        self.assertFalse(Wallet.objects.filter(user=user).exists())
        self.assertEqual(len(mail.outbox), 1)
        raw = re.search(r"/verify/([^/]+)/", mail.outbox[0].body).group(1)
        response = self.client.post(f"/verify/{raw}/")
        self.assertEqual(response.status_code, 302)
        user.refresh_from_db()
        self.assertIsNotNone(user.verified_at)
        self.assert_one_personal_credit(user)
        self.assertEqual(self.client.post(f"/verify/{raw}/").status_code, 404)
        welcome_credit(user)
        self.assert_one_personal_credit(user)

    def test_invited_employee_has_individual_verified_login_and_one_personal_credit(self):
        owner = User.objects.create_user(username="inviter", email="inviter@synthetic.invalid", verified_at=timezone.now())
        org = Organization.objects.create(name="Synthetic invited practice")
        dept = Department.objects.create(organization=org, name="X-ray")
        Membership.objects.create(user=owner, department=dept, role="owner")
        raw = "synthetic-invitation-token-not-a-real-credential"
        invite = Invitation.objects.create(email="invited@synthetic.invalid", department=dept, role="employee", token_hash=digest(raw), expires_at=timezone.now() + timedelta(days=1), invited_by=owner)
        response = self.client.post(f"/invite/{raw}/", self.registration_payload(invite.email))
        self.assertEqual(response.status_code, 302, getattr(response, "context", None))
        employee = User.objects.get(email=invite.email)
        self.assertIsNotNone(employee.verified_at)
        self.assertTrue(employee.check_password("Synthetic-Only-4821!private"))
        self.assertEqual(Membership.objects.get(user=employee, department=dept).role, "employee")
        self.assert_one_personal_credit(employee)
        self.assertEqual(self.client.post(f"/invite/{raw}/", self.registration_payload(invite.email)).status_code, 404)
        self.assert_one_personal_credit(employee)
