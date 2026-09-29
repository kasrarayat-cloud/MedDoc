"""ONE place rule for every model-proposed mapping change (addition or replacement). Synthetic BLANK forms only.

A model change is never adopted - it is held as an unresolved record - when its target
  (a) is, contains or lies inside a place the deterministic engine holds: a write-protected field, a detector
      protected region, a signing/initials area or a signing/consent question box, a printed authorship or
      attestation boundary, or wording the engine's kind rule reads as a signature/initials label, each printed
      part read alone (never one joined caption);
  (b) lies in a text box, shape, frame, header or footer, unless it is exactly the replaced field's own destination;
  (c) lies inside a mapped destination, unless it is exactly the replaced field's own destination.
A replacement inherits the prior field's write protection. An integrity failure while locating refuses the whole
response. The per-slot reference presentation never makes a form that fitted the request bound refuse.
Native transports are never constructed; nothing leaves this process.
"""
import copy
import io
import json
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, override_settings
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from core import generation, mapping_assistance as mapping, mapping_wire as wire, v3_mapping
from core.security import private_path
from core.tests.test_mapping_assistance import SETTINGS
from documents import word_v3
from documents.service import DocumentError, inspect_template, fill_template
from documents.test_signing_areas import signing_form

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
WP = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
WPS = 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape'
BODY = 'word/document.xml#'


def saved(doc):
    out = io.BytesIO(); doc.save(out); return out.getvalue()


def bordered_box(doc):
    table = doc.add_table(rows=1, cols=1)
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); borders.append(edge)
    table.cell(0, 0)._tc.get_or_add_tcPr().append(borders)
    return table


def box_form(heading, question):
    """A heading, a printed question and its own empty one-cell bordered box (the engine's question-box shape)."""
    doc = Document(); doc.add_paragraph('Participant ID: __________')
    doc.add_paragraph(heading, style='Heading 2'); doc.add_paragraph(question); bordered_box(doc)
    return saved(doc)


def textbox_form():
    """'Diagnosis: ____' holding an anchored floating text box that prints 'For office use only'."""
    doc = Document(); doc.add_paragraph('Patient reference: __________')
    paragraph = doc.add_paragraph('Diagnosis: ______________')
    paragraph._p.append(etree.fromstring(
        f'<w:r xmlns:w="{W}" xmlns:wp="{WP}" xmlns:a="{A}" xmlns:wps="{WPS}"><w:drawing>'
        '<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" behindDoc="0" locked="0" '
        'layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>5000000</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>8500000</wp:posOffset></wp:positionV>'
        '<wp:extent cx="1800000" cy="600000"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        '<wp:docPr id="10" name="Office box"/><wp:cNvGraphicFramePr/><a:graphic>'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp>'
        '<wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1800000" cy="600000"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr><wps:txbx><w:txbxContent><w:p><w:r>'
        '<w:t>For office use only</w:t></w:r></w:p></w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp>'
        '</a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'))
    return saved(doc)


def header_form():
    doc = Document(); doc.add_paragraph('Plan: ______________')
    doc.sections[0].header.paragraphs[0].text = 'Clinic reference: ______________'
    return saved(doc)


def frame_form():
    """'Office code: ____' printed in a positioned text frame (w:framePr)."""
    doc = Document(); doc.add_paragraph('Plan: ______________')
    paragraph = doc.add_paragraph('Office code: ______________')
    frame = OxmlElement('w:framePr')
    for name, value in (('w', '2000'), ('hAnchor', 'page'), ('vAnchor', 'page'), ('x', '8000'), ('y', '14000')):
        frame.set(qn('w:' + name), value)
    paragraph._p.get_or_add_pPr().insert(0, frame)
    return saved(doc)


def statement_form():
    doc = Document(); doc.add_paragraph('Participant ID: __________')
    doc.add_paragraph('This section is for the patient to add any comments or concerns about their care')
    bordered_box(doc)
    return saved(doc)


def role_form():
    doc = Document(); doc.add_paragraph('Visit summary', style='Heading 1')
    doc.add_paragraph('Plan: ______________')
    doc.add_paragraph('For the patient:')
    doc.add_paragraph('How have you felt since the last visit? ______________')
    doc.add_paragraph('')
    return saved(doc)


def attestation_form():
    doc = Document(); doc.add_paragraph('Plan: ______________')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'I confirm that the patient has read the information form'
    return saved(doc)


def simple_form():
    doc = Document(); table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Plan'; table.cell(1, 0).text = 'History'
    doc.add_paragraph('Instructions: ' + 'Safe blank source wording. ' * 20)
    doc.add_paragraph('Signature: ____________')
    return saved(doc)


def big_form(count=750, length=130):
    """MP-2's form: 750 instruction paragraphs of about 130 characters and three answer lines."""
    doc = Document(); doc.add_paragraph('Synthetic blank intake form')
    words = ('please read each instruction carefully before completing the section below and return the form to the '
             'front desk when finished ').split()
    for k in range(count):
        text = f'Instruction {k}: '; i = k
        while len(text) < length:
            text += words[i % len(words)] + ' '; i += 1
        doc.add_paragraph(text.strip())
    for label in ('Reason for visit: ____________________', 'Comments: ____________________', 'Plan: ____________________'):
        doc.add_paragraph(label)
    return saved(doc)


@override_settings(**SETTINGS)
class PlaceRule(SimpleTestCase):
    def setUp(self):
        deny = patch('anthropic.Anthropic', side_effect=AssertionError('Unmocked provider construction is forbidden in offline tests.'))
        deny.start(); self.addCleanup(deny.stop)
        self.directory = tempfile.TemporaryDirectory(prefix='mapping-places-')
        self.addCleanup(self.directory.cleanup)
        override = override_settings(MEDIA_ROOT=Path(self.directory.name)); override.enable(); self.addCleanup(override.disable)
        self.config = generation.configuration()

    # ─── helpers ───────────────────────────────────────────────────────────────────────────────────────────────
    def version(self, name, data, fields=None, drop=()):
        path = private_path(f'originals/{name}.docx'); path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(data)
        info = inspect_template(path, engine='v3')
        fields = copy.deepcopy(info['fields'] if fields is None else fields)
        fields = [f for f in fields if f['label'] not in drop]
        return SimpleNamespace(template=SimpleNamespace(format='docx', original_path=f'originals/{name}.docx', sha256=info['sha256']),
                               contract='v3', document_backend={}, fields=fields)

    def compile(self, version):
        compiled = mapping.compile_request(version, self.config)
        return compiled, json.loads(compiled['request']['messages'][0]['content'])

    @staticmethod
    def sid(payload, slot):
        return next(n for n, row in enumerate(payload['slots']) if row['slot'] == BODY + slot or row['slot'] == slot)

    @staticmethod
    def quote(payload, sid, text, roles=(0, 1, 2)):
        """[reference, start, end] of `text` in slot `sid`'s own source (text, heading, row_label)."""
        for position, (owner, role, _) in enumerate(payload['text_refs']):
            if owner == sid and role in roles:
                source = payload['slots'][sid][('text', 'heading', 'row_label')[role]]
                if text in source:
                    return [position, source.index(text), source.index(text) + len(text)]
        raise AssertionError(f'{text!r} is not printed at slot {sid}')

    @staticmethod
    def field_id(payload, key):
        return next(n for n, f in enumerate(payload['fields']) if f['key'] == key)

    def respond(self, payload, changes, replace=()):
        # A coherent response: the slot a change writes at is covered (a proposal that also calls it not an answer
        # contradicts itself and is held - see test_mapping_authority's self-contradiction case).
        mapped = {slot for f in payload['fields'] for slot in f['slots']}
        targets = {c[1] for c in changes}
        covered = lambda n, row: n in targets or v3_mapping.is_mapped(row['slot'], mapped)
        body = {'f': {'k': [n for n in range(len(payload['fields'])) if n not in replace], 'r': list(replace), 'n': [], 'u': []},
                's': {'c': [n for n, row in enumerate(payload['slots']) if covered(n, row)],
                      'n': [n for n, row in enumerate(payload['slots']) if not covered(n, row)], 'u': []},
                'x': [list(c) for c in changes], 'e': False}
        return body

    def addition(self, payload, slot, kind, label, roles=(0, 1, 2)):
        sid = self.sid(payload, slot)
        return [-1, sid, wire.KINDS.index(kind), *self.quote(payload, sid, label, roles), -1]

    def replacement(self, payload, key, slot, kind, label=None):
        fid = self.field_id(payload, key); sid = self.sid(payload, slot)
        if label is None:
            reference = [payload['text_refs'].index([fid, 4, 0]), 0, len(payload['fields'][fid]['label'])]
        else:
            reference = self.quote(payload, sid, label)
        return [fid, sid, wire.KINDS.index(kind), *reference, -1]

    def assertHeld(self, version, result, change, reason, *, key=''):
        """Nothing adopted, every prior field kept exactly, the change recorded and held for the clinician."""
        review = result['review']
        self.assertEqual(review['changes'], [])
        self.assertEqual(result['fields'], version.fields)
        self.assertFalse([f for f in result['fields'] if f.get('origin') == 'model_assisted'])
        self.assertEqual([f['key'] for f in v3_mapping.answerable(result['fields'])],
                         [f['key'] for f in v3_mapping.answerable(version.fields)])
        rows = [r for r in review['rejected_changes'] if r['wire'] == change]
        self.assertEqual(len(rows), 1, review['rejected_changes'])
        self.assertIn(reason, rows[0]['reason'])
        records = [r for r in review['unresolved'] if reason in r['reason'] and r['key'] == key]
        self.assertTrue(records, review['unresolved'])
        if key:
            # The KEYED record sits on the field's own destination: 'mapped' keeps the field, the proposal is named.
            prior = next(f for f in version.fields if f['key'] == key)
            self.assertEqual(records[0]['slot'], next(iter(sorted(v3_mapping.mapped_slots([prior])))))
            self.assertEqual(records[0]['proposed_slot'], rows[0]['slot'])
        else:
            self.assertEqual(records[0]['slot'], rows[0]['slot'])
        return rows[0]

    def adopted(self, result):
        return [f for f in result['fields'] if f.get('origin') == 'model_assisted']

    # ─── (a) the engine's held places ──────────────────────────────────────────────────────────────────────────
    def test_F1_an_addition_at_a_signing_or_consent_question_box_is_held(self):
        # The first-person consent question is held as the declaration it is (the same reader as a printed line).
        cases = [('Consultant review', 'Signed by the consultant:', 'Consultant review', mapping.HELD_REASONS['consent']),
                 ('Consent', 'I confirm the risks were explained and the patient consents to proceed:', 'Consent',
                  mapping.HELD_REASONS['declaration']),
                 ('Patient initials', 'Please initial below to confirm you have read the leaflet:', 'Patient',
                  mapping.HELD_REASONS['consent'])]
        for number, (heading, question, label, reason) in enumerate(cases):
            version = self.version(f'box{number}', box_form(heading, question))
            compiled, payload = self.compile(version)
            for slot in ('b/tbl[0]/tr[0]/tc[0]', 'b/tbl[0]/tr[0]/tc[0]/p[0]'):
                with self.subTest(heading=heading, slot=slot):
                    change = self.addition(payload, slot, 'long_text', label, roles=(1,))
                    result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
                    self.assertHeld(version, result, change, reason)

    def test_F1_control_an_ordinary_question_box_still_takes_a_model_addition(self):
        # Not vacuous: the same shape with an ordinary question is adopted when the engine's own field is missing.
        version = self.version('ordinary', box_form('Wound review', 'Please describe the wound?'),
                               drop=('Please describe the wound?',))
        compiled, payload = self.compile(version)
        # The box's own refs carry only its heading; the question paragraph is another slot.
        change = self.addition(payload, 'b/tbl[0]/tr[0]/tc[0]', 'long_text', 'Wound review', roles=(1,))
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        self.assertEqual([(f['label'], f['kind']) for f in self.adopted(result)], [('Wound review', 'long_text')])
        self.assertEqual(result['review']['rejected_changes'], [])

    def test_F3_a_protected_signing_block_caption_is_held_part_by_part(self):
        version = self.version('signing', signing_form(table=False))
        compiled, payload = self.compile(version)
        sid = self.sid(payload, 'b/p[2]')
        self.assertEqual(payload['slots'][sid]['text'], 'Name of Participant\t\tDate\t\tSignature')
        for kind, label in (('date', 'Date'), ('text', 'Name of Participant'), ('signature', 'Signature')):
            with self.subTest(label=label):
                change = self.addition(payload, 'b/p[2]', kind, label)
                result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
                self.assertHeld(version, result, change, 'cannot become generated answers.')
        # The engine's classifier reads each printed part alone: the date part never excuses the signature part.
        self.assertEqual(word_v3.label_parts('Name of Participant\t\tDate\t\tSignature'),
                         ['Name of Participant', 'Date', 'Signature'])
        self.assertTrue(mapping.signing_label('Name of Participant\t\tDate\t\tSignature'))
        self.assertFalse(mapping.signing_label('Name of Participant\t\tDate'))

    def test_F4_a_nominal_replacement_of_a_write_protected_statement_box_is_held(self):
        version = self.version('statement', statement_form())
        compiled, payload = self.compile(version)
        prior = next(f for f in version.fields if f.get('write_protected'))
        slot = next(iter(v3_mapping.mapped_slots([prior])))
        change = self.replacement(payload, prior['key'], slot, prior['kind'])
        result = mapping.validate_proposal(version, compiled,
                                           self.respond(payload, [change], replace=[self.field_id(payload, prior['key'])]))
        self.assertHeld(version, result, change, 'protected', key=prior['key'])
        self.assertIn(prior, result['fields'])

    def test_F4_a_replacement_inherits_the_prior_fields_write_protection(self):
        # The place itself is ordinary (the engine does not hold it); the stored field carries the protection.
        version = self.version('inherit', simple_form())
        prior = next(f for f in version.fields if f['label'] == 'Plan')
        prior['write_protected'] = True; prior['review_reason'] = 'Synthetic protected reason.'
        compiled, payload = self.compile(version)
        slot = next(iter(v3_mapping.mapped_slots([prior])))
        change = self.replacement(payload, prior['key'], slot, prior['kind'])
        result = mapping.validate_proposal(version, compiled,
                                           self.respond(payload, [change], replace=[self.field_id(payload, prior['key'])]))
        replaced = self.adopted(result)
        self.assertEqual(len(replaced), 1)
        self.assertEqual(result['review']['changes'], [{'key': prior['key'], 'replacement': replaced[0]['key']}])
        self.assertIs(replaced[0].get('write_protected'), True)
        self.assertEqual(replaced[0].get('review_reason'), 'Synthetic protected reason.')
        self.assertNotIn(replaced[0]['key'], [f['key'] for f in v3_mapping.answerable(result['fields'])])
        # The printed answer policy travels with it too; an unprotected prior gives an unprotected replacement.
        version = self.version('inherit-policy', simple_form())
        prior = next(f for f in version.fields if f['label'] == 'Plan')
        policy = {'source_author': 'patient', 'source': {'part': 'word/document.xml', 'slot': 'b/p[0]', 'fingerprint': '0' * 64,
                                                         'text': 'For the patient:'}}
        prior['answer_policy'] = copy.deepcopy(policy)
        compiled, payload = self.compile(version)
        change = self.replacement(payload, prior['key'], next(iter(v3_mapping.mapped_slots([prior]))), prior['kind'])
        result = mapping.validate_proposal(version, compiled,
                                           self.respond(payload, [change], replace=[self.field_id(payload, prior['key'])]))
        replaced = self.adopted(result)
        self.assertEqual([f.get('answer_policy') for f in replaced], [policy])
        self.assertNotIn('write_protected', replaced[0])

    def test_printed_authorship_and_attestation_places_are_held(self):
        version = self.version('role', role_form())
        compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/p[4]', 'long_text', 'Visit summary', roles=(1,))
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        self.assertHeld(version, result, change, 'authorship or attestation')
        version = self.version('attest', attestation_form())
        compiled, payload = self.compile(version)
        prior = next(f for f in version.fields if f.get('answer_policy'))
        change = self.replacement(payload, prior['key'], 'b/tbl[0]/tr[0]/tc[1]', prior['kind'])
        result = mapping.validate_proposal(version, compiled,
                                           self.respond(payload, [change], replace=[self.field_id(payload, prior['key'])]))
        self.assertHeld(version, result, change, 'authorship or attestation', key=prior['key'])

    def test_every_way_the_engine_holds_a_cell_is_read(self):
        # A signature COLUMN: the engine's own signature area (its label comes from the column header, so neither the
        # cell's text nor its row label names it).
        doc = Document(); doc.add_paragraph('Clinic record', style='Heading 2')
        table = doc.add_table(rows=2, cols=2); table.cell(0, 0).text = 'Name'; table.cell(0, 1).text = 'Signature'
        version = self.version('column', saved(doc), drop=('Signature',))
        compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/tbl[0]/tr[1]/tc[1]', 'text', 'Clinic record', roles=(1,))
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        self.assertHeld(version, result, change, mapping.SIGNING_REASON)
        # A signing ROW LABEL over a cell the engine made no field for (a picture keeps the cell from being empty).
        doc = Document(); doc.add_paragraph('Clinic record', style='Heading 2')
        table = doc.add_table(rows=2, cols=2); table.cell(0, 0).text = 'Plan'; table.cell(1, 0).text = 'Signature'
        run = OxmlElement('w:r'); run.append(OxmlElement('w:pict')); table.cell(1, 1).paragraphs[0]._p.append(run)
        version = self.version('rowlabel', saved(doc))
        self.assertFalse([f for f in version.fields if f['label'] == 'Signature'])
        compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/tbl[0]/tr[1]/tc[1]/p[0]', 'text', 'Clinic record', roles=(1,))
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        self.assertHeld(version, result, change, mapping.SIGNING_REASON)
        # A protected REGION: printed text in a patient's statement box. A NEW inspection write-protects the engine's
        # own field there (the ONE authority); a mapping stored before that rule still carries it answerable, and a
        # model change of it is held all the same.
        doc = Document(); doc.add_paragraph('Participant ID: __________')
        doc.add_paragraph('This section is for the patient to add any comments or concerns about their care')
        bordered_box(doc).cell(0, 0).text = 'Comments:'
        version = self.version('region', saved(doc))
        prior = next(f for f in version.fields if f['label'] == 'Comments')
        self.assertIs(prior.get('write_protected'), True)
        prior.pop('write_protected'); prior.pop('review_reason', None)
        compiled, payload = self.compile(version)
        change = self.replacement(payload, prior['key'], next(iter(v3_mapping.mapped_slots([prior]))), prior['kind'])
        result = mapping.validate_proposal(version, compiled,
                                           self.respond(payload, [change], replace=[self.field_id(payload, prior['key'])]))
        self.assertHeld(version, result, change, 'protected', key=prior['key'])

    def test_a_held_place_is_refused_before_the_engine_locates_anything(self):
        # A signing question box labelled from a harmless heading: only the place itself says it is held, so the
        # refusal must come before the engine is asked to locate anything there.
        version = self.version('early', box_form('Consultant review', 'Signed by the consultant:'))
        compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/tbl[0]/tr[0]/tc[0]', 'long_text', 'Consultant review', roles=(1,))
        self.assertFalse(mapping.signing_label('Consultant review'))
        calls = []; real = mapping.locate_fields
        with patch.object(mapping, 'locate_fields', side_effect=lambda *args: calls.append(args) or real(*args)):
            result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        self.assertEqual(calls, [])
        self.assertHeld(version, result, change, 'Signing or consent areas cannot become generated answers.')

    def test_signing_wording_is_read_whole_from_what_the_label_quotes_and_from_the_label(self):
        # The label 'Patient' quotes the heading 'Patient initials': the WHOLE quoted wording names an initials area.
        doc = Document(); doc.add_paragraph('Patient initials', style='Heading 2'); doc.add_paragraph('Notes: ______')
        version = self.version('quoted', saved(doc), drop=('Notes',))
        compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/p[1]', 'long_text', 'Patient', roles=(1,))
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        self.assertHeld(version, result, change, mapping.SIGNING_REASON)
        control = self.addition(payload, 'b/p[1]', 'long_text', 'Notes', roles=(0,))
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [control]))
        self.assertEqual([f['label'] for f in self.adopted(result)], ['Notes'])
        # 'Signature date' is a date line to the engine, but a text answer labelled 'Signature' there is a signature.
        doc = Document(); doc.add_paragraph('Signature date: ______')
        version = self.version('label', saved(doc), drop=('Signature date',))
        compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/p[0]', 'text', 'Signature')
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        self.assertHeld(version, result, change, mapping.SIGNING_REASON)
        control = self.addition(payload, 'b/p[0]', 'date', 'Signature date')
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [control]))
        self.assertEqual([f['kind'] for f in self.adopted(result)], ['date'])

    def test_a_signing_field_is_never_replaced_by_another_kind_even_where_the_engine_holds_nothing(self):
        version = self.version('manual-signature', simple_form())
        prior = next(f for f in version.fields if f['label'] == 'Plan')
        prior['kind'] = 'signature'                      # a clinician's own kind for this cell
        compiled, payload = self.compile(version)
        slot = next(iter(v3_mapping.mapped_slots([prior])))
        change = self.replacement(payload, prior['key'], slot, 'long_text')
        result = mapping.validate_proposal(version, compiled,
                                           self.respond(payload, [change], replace=[self.field_id(payload, prior['key'])]))
        self.assertHeld(version, result, change, mapping.SIGNING_REASON, key=prior['key'])

    # ─── (b) anchored containers ───────────────────────────────────────────────────────────────────────────────
    def test_F2_a_replacement_into_an_anchored_text_box_is_held_and_nothing_is_written_there(self):
        version = self.version('textbox', textbox_form())
        compiled, payload = self.compile(version)
        prior = next(f for f in version.fields if f['label'] == 'Diagnosis')
        change = self.replacement(payload, prior['key'], 'b/p[1]/txbx[0]/p[0]', 'text')
        result = mapping.validate_proposal(version, compiled,
                                           self.respond(payload, [change], replace=[self.field_id(payload, prior['key'])]))
        self.assertHeld(version, result, change, 'text box', key=prior['key'])
        self.assertFalse(any('txbx' in slot for slot in v3_mapping.mapped_slots(result['fields'])))
        # An addition there is held the same way; the exact own destination is still replaceable.
        change = self.addition(payload, 'b/p[1]/txbx[0]/p[0]', 'text', 'For office use only')
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        self.assertHeld(version, result, change, 'text box')
        change = self.replacement(payload, prior['key'], 'b/p[1]', 'text')
        result = mapping.validate_proposal(version, compiled,
                                           self.respond(payload, [change], replace=[self.field_id(payload, prior['key'])]))
        self.assertEqual([f['locator']['slot'] for f in self.adopted(result)], ['b/p[1]'])

    def test_headers_footers_and_frames_take_only_a_replacement_at_the_fields_own_place(self):
        for name, data, label, slot in (('header', header_form(), 'Clinic reference', 'word/header1.xml#hf/p[0]'),
                                        ('frame', frame_form(), 'Office code', 'b/p[1]')):
            with self.subTest(name):
                version = self.version(name + '-add', data, drop=(label,))
                compiled, payload = self.compile(version)
                change = self.addition(payload, slot, 'text', label)
                result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
                self.assertHeld(version, result, change, 'text box')
                version = self.version(name + '-keep', data)
                compiled, payload = self.compile(version)
                prior = next(f for f in version.fields if f['label'] == label)
                change = self.replacement(payload, prior['key'], slot, 'text')
                result = mapping.validate_proposal(version, compiled,
                                                   self.respond(payload, [change], replace=[self.field_id(payload, prior['key'])]))
                self.assertEqual(len(self.adopted(result)), 1, result['review']['rejected_changes'])

    # ─── (c) inside a mapped destination ───────────────────────────────────────────────────────────────────────
    def test_a_replacement_inside_its_own_destination_but_not_at_it_is_held(self):
        version = self.version('inside', simple_form())
        compiled, payload = self.compile(version)
        prior = next(f for f in version.fields if f['label'] == 'Plan')
        slot = next(iter(v3_mapping.mapped_slots([prior])))
        change = self.replacement(payload, prior['key'], slot + '/p[0]', prior['kind'])
        result = mapping.validate_proposal(version, compiled,
                                           self.respond(payload, [change], replace=[self.field_id(payload, prior['key'])]))
        self.assertHeld(version, result, change, 'inside a mapped destination', key=prior['key'])
        change = self.replacement(payload, prior['key'], slot, prior['kind'])
        result = mapping.validate_proposal(version, compiled,
                                           self.respond(payload, [change], replace=[self.field_id(payload, prior['key'])]))
        self.assertEqual(len(self.adopted(result)), 1)

    def test_a_box_replacement_resolved_to_its_own_line_is_not_held_as_inside(self):
        # Control for (c): the model names the printed box line; the engine resolves it to the paragraph that holds
        # the field's own box, which is the field's own place, not a place inside another destination.
        doc = Document(); doc.add_paragraph('Plan: ______'); doc.add_paragraph('Smoker \u2610')
        version = self.version('glyph', saved(doc))
        compiled, payload = self.compile(version)
        prior = next(f for f in version.fields if f['label'] == 'Smoker')
        for slot in ('b/p[1]/glyphs', 'b/p[1]'):
            with self.subTest(slot=slot):
                change = self.replacement(payload, prior['key'], slot, 'checkbox', 'Smoker')
                result = mapping.validate_proposal(version, compiled,
                                                   self.respond(payload, [change], replace=[self.field_id(payload, prior['key'])]))
                self.assertEqual([(f['label'], f['kind']) for f in self.adopted(result)], [('Smoker', 'checkbox')],
                                 result['review']['rejected_changes'])

    # ─── F5 ────────────────────────────────────────────────────────────────────────────────────────────────────
    def test_F5_an_unreadable_original_while_locating_refuses_the_whole_response(self):
        version = self.version('f5', simple_form(), drop=('Plan',))
        compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/tbl[0]/tr[0]/tc[1]', 'long_text', 'Plan')
        real = mapping.locate_fields
        path = private_path(version.template.original_path); original = path.read_bytes()

        def swapped(where, targets):
            path.write_bytes(b'not a synthetic docx any more')
            try:
                return real(where, targets)
            finally:
                path.write_bytes(original)
        with patch.object(mapping, 'locate_fields', side_effect=swapped), self.assertRaises(mapping.Refusal):
            mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        # Control: the same change, the original intact, is adopted.
        result = mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
        self.assertEqual([f['label'] for f in self.adopted(result)], ['Plan'])
        # Control: the engine refusing THIS place (a cell with printed text) is one change's record, never a refusal.
        printed = self.addition(payload, 'b/tbl[0]/tr[0]/tc[0]', 'text', 'Plan')
        try:
            result = mapping.validate_proposal(version, compiled, self.respond(payload, [printed, change]))
        except mapping.Refusal as refusal:
            self.fail(f'one unlocatable place refused the whole response: {refusal}')
        self.assertEqual([f['label'] for f in self.adopted(result)], ['Plan'])
        self.assertIn('printed text', next(r for r in result['review']['rejected_changes'] if r['wire'] == printed)['reason'])

    def test_F5_the_engines_place_facts_are_bound_to_the_original(self):
        version = self.version('f5b', simple_form(), drop=('Plan',))
        compiled, payload = self.compile(version)
        change = self.addition(payload, 'b/tbl[0]/tr[0]/tc[1]', 'long_text', 'Plan')
        real = mapping.held_places
        for name, effect in (('another original', lambda path, *rest: {**real(path, *rest), 'sha256': '0' * 64}),
                             ('unreadable', DocumentError('The isolated document processor failed.', 'parser_failed'))):
            with self.subTest(name):
                with patch.object(mapping, 'held_places', side_effect=effect), self.assertRaises(Exception) as caught:
                    mapping.validate_proposal(version, compiled, self.respond(payload, [change]))
                self.assertIsInstance(caught.exception, mapping.Refusal)

    # ─── MP-2 ──────────────────────────────────────────────────────────────────────────────────────────────────
    def test_MP2_a_form_that_fitted_the_bound_still_fits_and_small_forms_get_the_compact_presentation(self):
        version = self.version('big', big_form())
        try:
            compiled, payload = self.compile(version)
        except ValidationError as refused:
            self.fail(f'a form that fitted the bound before the presentation is refused: {refused}')
        size = len(mapping.canonical(compiled['request']).encode())
        self.assertLessEqual(size, mapping.MAX_PAYLOAD_BYTES)
        self.assertEqual(compiled['input_token_bound'], size + 12000)
        # Over the bound with the presentation, the request is the unpresented v3 shape: byte for byte what fitted
        # before the presentation existed (no per-slot lists, the v3 prompt).
        self.assertEqual(compiled['presentation'], 'none')
        self.assertEqual(compiled['request']['system'], mapping.SYSTEM_V3)
        self.assertFalse(any('refs' in row for row in payload['slots']) or any('label_ref' in f for f in payload['fields']))
        # The presented request of this form would not fit; the bound itself is unchanged.
        self.assertEqual(mapping.MAX_PAYLOAD_BYTES, 400_000)
        presented = copy.deepcopy(payload); wire.present(presented)
        request = generation.anthropic_request(self.config['model'], compiled['profile'], mapping.SYSTEM, presented, wire.Proposal)
        self.assertGreater(len(mapping.canonical(request).encode()), mapping.MAX_PAYLOAD_BYTES)
        # The fallback expands a response exactly as the presented request would.
        answer = wire.fixture(payload)
        self.assertEqual(wire.expand(payload, copy.deepcopy(answer)), wire.expand(presented, copy.deepcopy(answer)))
        # A small form is presented compactly: each slot lists its own reference IDs; empty lists are omitted.
        version = self.version('small', simple_form())
        compiled, payload = self.compile(version)
        self.assertEqual(compiled['presentation'], 'slot_refs')
        self.assertEqual(compiled['request']['system'], mapping.SYSTEM)
        for sid, row in enumerate(payload['slots']):
            own = [n for n, (owner, role, _) in enumerate(payload['text_refs']) if owner == sid and role != 4]
            self.assertEqual(row.get('refs'), own or None)
            self.assertIsInstance(row.get('refs', []), list)
