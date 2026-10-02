"""Offline synthetic mapping operation proofs. Native transports are always mocked."""
import copy
import hashlib
import io
import json
import os
import tempfile
from decimal import Decimal
from datetime import timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from django.core.exceptions import ValidationError
from django.db import DatabaseError, transaction
from django.test import TestCase, override_settings
from core.tests.surfaces import SurfaceClient
from django.utils import timezone
from docx import Document
from core import generation, mapping_assistance as mapping, mapping_jobs, jobs
from core.billing import welcome_credit
from core.models import AuditEvent, Department, Encounter, Job, LedgerEntry, Membership, Organization, Patient, PlatformCost, ProviderReceipt, Reservation, Template, TemplateVersion, Usage, User, Wallet
from core.security import private_path
from documents.service import inspect_template, fill_template
from core.tests.test_workflow import TEST_SETTINGS

SETTINGS={**TEST_SETTINGS,'QUILLMD_WORD_ENGINE':'v3','QUILLMD_MAPPING_ASSISTANCE':True,
          'QUILLMD_MAPPING_ASSISTANCE_LIVE':False,'AI_MODE':'fixture','AI_PROVIDER':'anthropic',
          'AI_MODEL':'claude-sonnet-5','AI_THINKING':'disabled','AI_EFFORT':'medium',
          'AI_INPUT_RATE':'2','AI_OUTPUT_RATE':'10','AI_CACHED_RATE':'0.2','AI_INFERENCE_GEO':''}

@override_settings(**SETTINGS)
class MappingAssistance(TestCase):
    client_class = SurfaceClient
    def setUp(self):
        deny=patch('anthropic.Anthropic',side_effect=AssertionError('Unmocked provider construction is forbidden in offline tests.'))
        deny.start();self.addCleanup(deny.stop)
        self.directory=tempfile.TemporaryDirectory(prefix='mapping-offline-')
        self.addCleanup(self.directory.cleanup)
        override=override_settings(MEDIA_ROOT=Path(self.directory.name));override.enable();self.addCleanup(override.disable)
        self.user=User.objects.create_user(username='mapping',email='mapping@synthetic.invalid',verified_at=timezone.now())
        self.org=Organization.objects.create(name='Synthetic mapping')
        self.department=Department.objects.create(organization=self.org,name='Synthetic mapping')
        Membership.objects.create(user=self.user,department=self.department,role='clinician')
        welcome_credit(self.user)
        doc=Document();table=doc.add_table(rows=2,cols=2)
        table.cell(0,0).text='Plan';table.cell(1,0).text='History'
        doc.add_paragraph('Instructions: '+'Safe blank source wording. '*20)
        doc.add_paragraph('Signature: ____________')
        buffer=io.BytesIO();doc.save(buffer);self.original=buffer.getvalue()
        path=private_path('originals/mapping.docx');path.parent.mkdir(parents=True);path.write_bytes(self.original)
        info=inspect_template(path,engine='v3')
        self.template=Template.objects.create(department=self.department,name='Synthetic mapping',format='docx',
            original_name='mapping.docx',original_path='originals/mapping.docx',sha256=info['sha256'],created_by=self.user)
        self.version=TemplateVersion.objects.create(template=self.template,number=1,fields=info['fields'],contract='v3',created_by=self.user)
        self.config=generation.configuration();self.compiled=mapping.compile_request(self.version,self.config)
        self.client.force_login(self.user)

    def body(self):
        return mapping_jobs.send(self.compiled,{**self.config,'mode':'fixture'})[0]

    def enqueue(self):
        return mapping_jobs.enqueue(self.version,self.user,self.config,self.compiled['request_sha256'])

    def artifact(self,name,payload):
        if location:=os.getenv('QUILLMD_MAPPING_ARTIFACT_DIR'):
            path=Path(location);path.mkdir(parents=True,exist_ok=True)
            (path/name).write_bytes(payload)

    def test_existing_context_inherits_only_with_exact_resolved_locator(self):
        prior = self.version.fields[0]
        long_context = 'Known reviewed source context ' + 'complete words ' * 30
        prior['context'] = long_context
        self.compiled = mapping.compile_request(self.version, self.config)
        payload = json.loads(self.compiled['request']['messages'][0]['content'])
        slot = next(iter(mapping.v3_mapping.mapped_slots([prior])))
        sid = next(n for n, row in enumerate(payload['slots']) if row['slot'] == slot)
        body = self.body(); body['f']['k'].remove(0); body['f']['r'] = [0]
        body['x'] = [[0, sid, 0, payload['text_refs'].index([0, 4, 0]), 0, len(prior['label']), -1]]
        result = mapping.validate_proposal(self.version, self.compiled, body)
        replacement = next(f for f in result['fields'] if f['label'] == prior['label'])
        self.assertEqual(replacement['context'], long_context)
        self.assertEqual(replacement['locator'], prior['locator'])
        self.assertEqual((replacement['key'], replacement['origin']), (prior['key'], 'model_assisted'))
        # ⚠️ A REPLACEMENT NEVER MOVES A FIELD: naming a slot inside its place keeps the prior locator and context
        # exactly (it used to re-locate after the printed text - M11: 'Pulse findings: ______ weak radial pulse').
        descendant = copy.deepcopy(body); descendant['x'][0][1] = sid + 1
        kept = mapping.validate_proposal(self.version, self.compiled, descendant)
        kept_field = next(f for f in kept['fields'] if f['key'] == prior['key'])
        self.assertEqual((kept_field['locator'], kept_field['context']), (prior['locator'], long_context))
        # Another field's slot is ONE invalid change: held with its reason; the fields stay exactly as stored.
        foreign_slot = payload['fields'][1]['slots'][0]
        foreign = copy.deepcopy(body)
        foreign['x'][0][1] = next(n for n, row in enumerate(payload['slots']) if row['slot'] == foreign_slot)
        held = mapping.validate_proposal(self.version, self.compiled, foreign)
        self.assertEqual(held['fields'], self.version.fields)
        self.assertEqual(held['review']['changes'], [])
        self.assertIn('Replacement target belongs to another field destination.',
                      ' '.join(row['reason'] for row in held['review']['unresolved']))

    def test_complete_long_labels_are_priced_in_actual_request_and_old_profile_refuses(self):
        original_fields = copy.deepcopy(self.version.fields)
        for length in (334, 409, 1000):
            self.version.fields = [{**original_fields[0], 'label': 'Q' * length}]
            compiled = mapping.compile_request(self.version, self.config)
            payload = json.loads(compiled['request']['messages'][0]['content'])
            self.assertEqual(payload['fields'][0]['label'], 'Q' * length)
            self.assertEqual(compiled['input_token_bound'], len(mapping.canonical(compiled['request']).encode()) + 12000)
            self.assertEqual(compiled['request']['max_tokens'], 20000)
            cost = (Decimal(compiled['input_token_bound']) * Decimal('2') + Decimal(20000) * Decimal('10')) / 1000000
            self.assertGreaterEqual(Decimal(compiled['maximum_cost']), cost)
        self.version.refresh_from_db()
        self.assertEqual(self.version.fields, original_fields)
        with self.assertRaises(ValidationError):
            mapping.request_profile({**self.config, 'mapping_contract': 'blank-mapping-v2',
                'mapping_request_profile': self.compiled['profile'],
                'mapping_profile_sha256': self.compiled['profile_sha256']})

    def test_full_outline_and_fixed_schema_preserve_long_source_and_untrusted_words(self):
        request=self.compiled['request'];payload=json.loads(request['messages'][0]['content'])
        self.assertTrue(any(len(row['text'])>120 for row in payload['slots']))
        self.assertIn('UNTRUSTED DATA',request['system'])
        self.assertNotIn('Safe blank source wording',json.dumps(request['output_config']['format']['schema']))
        self.assertEqual(request['model'],'claude-sonnet-5')
        self.assertEqual(request['thinking'],{'type':'disabled'})
        self.assertEqual(request['output_config']['effort'],'medium')
        self.assertNotIn('temperature',request)
        with patch.object(mapping,'MAX_PAYLOAD_BYTES',1):
            with self.assertRaises(ValidationError):mapping.compile_request(self.version,self.config)

    def test_duplicated_unknown_or_extra_schema_members_refuse(self):
        body=self.body()
        mutations=[{**body,'f':{**body['f'],'u':[body['f']['k'][0]]}}, {**body,'xml':'execute this'}]
        for candidate in mutations:
            with self.subTest(candidate=list(candidate)):
                with self.assertRaises((ValidationError,ValueError)):mapping.validate_proposal(self.version,self.compiled,candidate)

    def test_a_missing_field_or_slot_is_held_for_review_not_refused(self):
        """Campaign attempt 24 (a model leaving slots unclassified) refused the whole proposal; an omitted field keeps its
        deterministic mapping and an omitted slot becomes an unresolved review item. Nothing is adopted from it."""
        body=self.body()
        for name,candidate in [('field',{**body,'f':{**body['f'],'k':body['f']['k'][:-1]}}),
                               ('slot',{**body,'s':{**body['s'],'u':body['s']['u'][:-1]}})]:
            with self.subTest(omitted=name):
                result=mapping.validate_proposal(self.version,self.compiled,candidate)
                self.assertEqual(result['fields'],self.version.fields)
                self.assertEqual(result['review'].get('changes') or [],[])
                if name=='slot':
                    dropped=json.loads(self.compiled['request']['messages'][0]['content'])['slots'][body['s']['u'][-1]]['slot']
                    self.assertIn(dropped,[row['slot'] for row in result['review']['unresolved']])

    def target(self,slot,label,kind='text',fid=-1,options=()):
        payload=json.loads(self.compiled['request']['messages'][0]['content'])
        sid=next(n for n,row in enumerate(payload['slots']) if row['slot']==slot)
        def reference(text):
            for n,(owner,role,index) in enumerate(payload['text_refs']):
                if owner!=sid or role==4:continue
                row=payload['slots'][owner]
                source=row['options'][index] if role==3 else row.get(('text','heading','row_label')[role],'')
                if text in source:return [n,source.index(text),source.index(text)+len(text)]
            raise AssertionError('Synthetic target text missing from source')
        return [fid,sid,mapping.wire.KINDS.index(kind),*reference(label),-1]+[n for option in options for n in reference(option)]

    def test_wrong_slots_and_references_refuse_and_overlap_stays_visible(self):
        """A wrong slot or text reference is ONE held change (M05), never a refused response and never adopted."""
        body=self.body();prior=self.version.fields[0]
        slot=next(iter(mapping.v3_mapping.mapped_slots([prior])))
        target=self.target(slot,prior['label'])
        for index,value in [(1,1199),(3,399999)]:
            altered=target.copy();altered[index]=value;body['x']=[altered]
            with self.subTest(index=index):
                result=mapping.validate_proposal(self.version,self.compiled,body)
                self.assertEqual(result['fields'],self.version.fields)
                self.assertEqual([row['wire'] for row in result['review']['rejected_changes']],[altered])
                self.assertTrue(any(row.get('change')==result['review']['rejected_changes'][0]['id']
                                    for row in result['review']['unresolved']))
        body['x']=[target]
        result=mapping.validate_proposal(self.version,self.compiled,body)
        self.assertEqual(result['fields'],self.version.fields)
        self.assertTrue(result['review']['unresolved'])

    def test_source_mutation_refuses_before_enqueue(self):
        private_path(self.template.original_path).write_bytes(b'corrupt synthetic file')
        with self.assertRaises((ValidationError,ValueError)):self.enqueue()
        self.assertFalse(Job.objects.exists())

    def test_worker_uses_template_scope_receipt_ledger_and_review_without_patients(self):
        job=self.enqueue();self.assertEqual(self.enqueue().pk,job.pk)
        self.assertIsNone(job.encounter_id);self.assertEqual(job.department,self.department)
        with patch('anthropic.Anthropic') as provider:
            self.assertTrue(jobs.run_one());provider.assert_not_called()
        job.refresh_from_db();self.assertEqual(job.state,'succeeded',job.error)
        self.assertTrue(ProviderReceipt.objects.get(job=job).simulated)
        self.assertEqual(Usage.objects.get(job=job).actual_cost,Decimal(0))
        self.assertFalse(Patient.objects.exists())
        self.assertFalse(Reservation.objects.filter(job=job,state='held').exists())
        latest=self.template.versions.latest('number')
        self.assertTrue(latest.needs_confirmation)
        self.assertTrue(latest.assistance['unresolved'])
        self.artifact('mapping-review.html',self.client.get(f'/templates/{self.template.pk}/assistance/').content)
        self.artifact('existing-mapping-reference.html',self.client.get(f'/templates/{self.template.pk}/').content)
        response=self.client.post(f'/templates/{self.template.pk}/',{'action':'v3_confirm','mapping_version':latest.pk,'checked_every_field':'on'})
        self.assertEqual(response.status_code,400)
        self.assertEqual(private_path(self.template.original_path).read_bytes(),self.original)

    def test_a_charged_mapping_result_is_downloadable_under_the_generation_rule(self):
        # RR2-L1 (docs/rebuild/v1/residuals-assessment-20260929.md). A generation's stored provider output is downloadable
        # by an active member of the job's department while the job is charged (jobs.deliverable). A mapping job has no
        # encounter, and job_result scoped by encounter__department, so its charged output answered 404 even to its
        # requester. RED on d5367521: 404. The rule is now the same: the department is the template's.
        job=self.enqueue();jobs.run_one();job.refresh_from_db()
        self.assertEqual(job.state,'succeeded',job.error);self.assertTrue(jobs.deliverable(job))
        response=self.client.get(f'/jobs/{job.pk}/result/')
        self.assertEqual(response.status_code,200)
        self.assertEqual(hashlib.sha256(b''.join(response.streaming_content)).hexdigest(),job.result_sha256)
        page=self.client.get(f'/templates/{self.template.pk}/assistance/')
        self.assertContains(page,f'/jobs/{job.pk}/result/')
        event=AuditEvent.objects.get(action='job.result_download',target=str(job.pk))
        self.assertEqual((event.actor,event.department),(self.user,self.department))
        # Any active member of the template's department, as for a generation; nobody outside it.
        colleague=User.objects.create_user(username='mapping-colleague',email='mapping-colleague@synthetic.invalid',verified_at=timezone.now())
        Membership.objects.create(user=colleague,department=self.department,role='employee')
        other=Department.objects.create(organization=self.org,name='Synthetic sibling department')
        outsider=User.objects.create_user(username='mapping-outsider',email='mapping-outsider@synthetic.invalid',verified_at=timezone.now())
        Membership.objects.create(user=outsider,department=other,role='owner')
        for person,status in ((colleague,200),(outsider,404)):
            with self.subTest(person=person.email):
                self.client.force_login(person)
                self.assertEqual(self.client.get(f'/jobs/{job.pk}/result/').status_code,status)
        Membership.objects.filter(user=colleague).update(active=False)
        self.client.force_login(colleague)
        self.assertEqual(self.client.get(f'/jobs/{job.pk}/result/').status_code,404)

    def test_an_uncharged_mapping_result_is_never_delivered(self):
        # Control for RR2-L1: an absorbed platform failure's stored output stays undelivered, as a generation's does.
        with override_settings(AI_MODE='live',QUILLMD_MAPPING_ASSISTANCE_LIVE=True):
            self.config=generation.configuration();self.compiled=mapping.compile_request(self.version,self.config)
            job=self.enqueue()
            with patch('anthropic.Anthropic') as client:
                client.return_value.messages.create.return_value=self.response(stop='max_tokens')
                jobs.run_one()
        job.refresh_from_db();self.assertEqual(job.state,'failed',job.error)
        self.assertTrue(job.result_path);self.assertFalse(jobs.deliverable(job))
        self.assertEqual(self.client.get(f'/jobs/{job.pk}/result/').status_code,404)
        self.assertNotContains(self.client.get(f'/templates/{self.template.pk}/assistance/'),f'/jobs/{job.pk}/result/')

    def test_scope_and_source_config_are_immutable_in_database(self):
        job=self.enqueue()
        for change in ({'mapping_version':None},{'kind':'generation'},{'config':{}},{'user_id':None}):
            with self.subTest(change=change),self.assertRaises(DatabaseError),transaction.atomic():
                Job.objects.filter(pk=job.pk).update(**change)
        with self.assertRaises(DatabaseError),transaction.atomic():
            Template.objects.filter(pk=self.template.pk).update(sha256='0'*64)

    def test_disabled_flag_cancel_and_revoked_membership_release_without_send(self):
        job=self.enqueue()
        with override_settings(QUILLMD_MAPPING_ASSISTANCE=False),patch('core.mapping_jobs.send') as send:
            self.assertTrue(jobs.run_one());send.assert_not_called()
        job.refresh_from_db();self.assertEqual(job.state,'failed')
        self.assertFalse(Reservation.objects.filter(job=job,state='held').exists())
        job=self.enqueue();jobs.cancel(job,self.user);job.refresh_from_db();self.assertEqual(job.state,'cancelled')
        job=self.enqueue();Membership.objects.filter(user=self.user).update(active=False)
        with patch('core.mapping_jobs.send') as send:jobs.run_one();send.assert_not_called()
        job.refresh_from_db();self.assertEqual(job.state,'failed')

    def test_cross_department_quote_and_hidden_default_route(self):
        with override_settings(QUILLMD_MAPPING_ASSISTANCE=False):
            self.assertEqual(self.client.get(f'/templates/{self.template.pk}/assistance/').status_code,404)
        response=self.client.get(f'/templates/{self.template.pk}/assistance/')
        self.assertEqual(response.status_code,200)
        approval=response.context['approval']
        outsider=User.objects.create_user(username='outside',email='outside@synthetic.invalid',verified_at=timezone.now())
        self.client.force_login(outsider)
        response=self.client.post(f'/templates/{self.template.pk}/assistance/',{'action':'queue','approval':approval,'blank_only':'on'})
        self.assertEqual(response.status_code,404)
        self.assertFalse(Job.objects.exists())

    def test_manual_resolution_records_assistance_and_old_version_is_preserved(self):
        self.enqueue();jobs.run_one();version=self.template.versions.latest('number')
        body={'action':'resolve','mapping_version':str(version.pk)}
        # An area a field already writes to is 'mapped'; 'not an answer' there would contradict the mapping (RF15).
        mapped=mapping.v3_mapping.mapped_slots(version.fields)
        body.update({'resolution_'+str(n):'mapped' if mapping.v3_mapping.is_mapped(row['slot'],mapped) else 'not_answer'
                     for n,row in enumerate(version.assistance['unresolved'])})
        response=self.client.post(f'/templates/{self.template.pk}/assistance/',body)
        self.assertEqual(response.status_code,302)
        latest=self.template.versions.latest('number')
        self.assertTrue(latest.assistance['assisted']);self.assertEqual(latest.assistance['unresolved'],[])
        version.refresh_from_db();self.assertTrue(version.assistance['unresolved'])

    def test_native_envelope_refusal_incomplete_and_wrong_model_never_parse(self):
        for model,stop in [('claude-sonnet-4-6','end_turn'),('claude-sonnet-5','max_tokens'),('claude-sonnet-5','refusal')]:
            with self.subTest(model=model,stop=stop),self.assertRaises(ValidationError):
                mapping_jobs.parse_response(SimpleNamespace(model=model,stop_reason=stop,content=[]),False,'claude-sonnet-5')
        response=self.response()
        response.content=[SimpleNamespace(type='text',text='{"fields": [], "fields": []}')]
        with self.assertRaises(ValidationError):mapping_jobs.parse_response(response,False,'claude-sonnet-5')

    @override_settings(AI_MODE='live')
    def test_live_permission_off_refuses_before_reservation_or_provider(self):
        config=generation.configuration();compiled=mapping.compile_request(self.version,config)
        with patch('anthropic.Anthropic') as provider,self.assertRaises(ValidationError):
            mapping_jobs.enqueue(self.version,self.user,config,compiled['request_sha256'])
        provider.assert_not_called();self.assertFalse(Job.objects.exists())

    def test_missing_field_can_be_added_then_reviewed_finalized_and_exported(self):
        missing=next(f for f in self.version.fields if f['label']=='Plan')
        fields=[f for f in self.version.fields if f['key']!=missing['key']]
        self.version=TemplateVersion.objects.create(template=self.template,number=2,fields=fields,contract='v3',created_by=self.user)
        self.compiled=mapping.compile_request(self.version,self.config)
        body=self.body()
        body['s']={'c':[],'n':list(range(len(json.loads(self.compiled['request']['messages'][0]['content'])['slots']))),'u':[]}
        body['x']=[self.target(next(iter(mapping.v3_mapping.mapped_slots([missing]))),'Plan',missing['kind'])]
        job=self.enqueue()
        with patch('core.mapping_jobs.send',return_value=(body,Decimal(0),{},{},'synthetic-mapping-fixture','','',True)):
            jobs.run_one()
        job.refresh_from_db();self.assertEqual(job.state,'succeeded',job.error)
        version=self.template.versions.latest('number')
        # ⚠️ A MODEL ADDITION IS NEVER ADOPTED AUTOMATICALLY: it is an explicit review item naming what was proposed,
        # and it blocks confirmation until the clinician adds it (one add on the mapping page) or dismisses it.
        self.assertEqual(len(version.fields),len(fields))
        proposed=[row for row in version.assistance['unresolved'] if row.get('proposal')]
        self.assertEqual([(row['slot'],row['proposal']['label'],row['proposal']['kind']) for row in proposed],
                         [(next(iter(mapping.v3_mapping.mapped_slots([missing]))),'Plan',missing['kind'])])
        self.assertEqual(self.client.post(f'/templates/{self.template.pk}/',{'action':'v3_confirm','mapping_version':version.pk,'checked_every_field':'on'}).status_code,400)
        row=proposed[0]
        response=self.client.post(f'/templates/{self.template.pk}/',{'action':'v3_add','mapping_version':version.pk,
                                  'slot':row['slot'],'kind':row['proposal']['kind'],'label':row['proposal']['label']})
        self.assertEqual(response.status_code,302)
        version=self.template.versions.latest('number')
        self.assertEqual(len(version.fields),len(fields)+1)
        self.assertEqual(next(f for f in version.fields if f['label']=='Plan')['origin'],'manual')
        mapped=mapping.v3_mapping.mapped_slots(version.fields)
        body={'action':'resolve','mapping_version':str(version.pk)}
        body.update({'resolution_'+str(n):'mapped' if mapping.v3_mapping.is_mapped(r['slot'],mapped) else 'not_answer'
                     for n,r in enumerate(version.assistance['unresolved'])})
        self.assertEqual(self.client.post(f'/templates/{self.template.pk}/assistance/',body).status_code,302)
        version=self.template.versions.latest('number')
        self.assertEqual(version.assistance['unresolved'],[])
        self.assertEqual(self.client.post(f'/templates/{self.template.pk}/',{'action':'v3_confirm','mapping_version':version.pk,'checked_every_field':'on'}).status_code,302)
        patient=Patient.objects.create(department=self.department,display_name='SYNTHETIC MAPPING WORKFLOW',reference='mapping-flow',created_by=self.user)
        response=self.client.post('/encounters/new/',{'patient':patient.pk,'template_version':version.pk,'notes':'Plan: rest.','service_date':timezone.localdate().isoformat(),'synthetic_only':'on'})
        self.assertEqual(response.status_code,302)
        encounter=Encounter.objects.get()
        payload={'revision':encounter.revision}
        for field in version.fields:
            if field['kind'] in ('signature','initials'):continue
            payload['answer_'+field['key']]='Rest and return if symptoms worsen.' if field['label']=='Plan' else ''
            payload['reviewed_'+field['key']]='on'
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/save/',payload).status_code,302)
        encounter.refresh_from_db()
        response=self.client.post(f'/encounters/{encounter.pk}/finalize/',{'revision':encounter.revision,'reviewed':'on'})
        self.assertEqual(response.status_code,302)
        encounter.refresh_from_db();self.assertEqual(encounter.status,'finalized')
        from core.tests.test_word_v3_workflow import part_text
        exported=private_path(encounter.export_path).read_bytes()
        self.assertIn('Rest and return if symptoms worsen.',part_text(exported))
        self.assertEqual(hashlib.sha256(exported).hexdigest(),encounter.export_sha256)
        self.assertEqual(private_path(self.template.original_path).read_bytes(),self.original)
        self.assertEqual(self.client.get(f'/encounters/{encounter.pk}/export/').status_code,200)
        self.artifact('workflow-original.docx',self.original)
        self.artifact('workflow-export.docx',exported)
        added=next(f for f in version.fields if f['label']=='Plan')
        for name,text in [('missing',''),('long','Synthetic detailed plan. '*80)]:
            output=Path(self.directory.name)/(name+'.docx')
            fill_template(private_path(self.template.original_path),version.fields,{added['key']:text},output)
            document=Document(output)
            self.assertEqual(document.tables[0].cell(0,1).text,text)
            self.assertEqual(document.tables[0].cell(0,0).text,'Plan')
            self.assertIn('Signature:',document.paragraphs[-1].text)
            self.artifact('workflow-'+name+'.docx',output.read_bytes())

    def response(self,body=None,stop='end_turn',model='claude-sonnet-5'):
        body=body if body is not None else self.body()
        self.response_number=getattr(self,'response_number',0)+1
        data={'model':model,'stop_reason':stop,'content':[{'type':'text','text':json.dumps(body)}]}
        return SimpleNamespace(model=model,stop_reason=stop,content=[SimpleNamespace(**data['content'][0])],
            id=f'msg_synthetic_mapping_{self.response_number}',_request_id=f'req_synthetic_mapping_{self.response_number}',
            usage=SimpleNamespace(input_tokens=100,output_tokens=30,cache_read_input_tokens=0,cache_creation_input_tokens=0),
            model_dump=lambda **kwargs:data)

    @override_settings(AI_MODE='live',QUILLMD_MAPPING_ASSISTANCE_LIVE=True)
    def test_mocked_native_receipt_charge_and_financial_views_keep_exact_model(self):
        self.config=generation.configuration();self.compiled=mapping.compile_request(self.version,self.config)
        job=self.enqueue();response=self.response()
        with patch('anthropic.Anthropic') as client:
            client.return_value.messages.create.return_value=response
            jobs.run_one()
        client.assert_called_once_with(timeout=120,max_retries=0)
        sent=client.return_value.messages.create.call_args.kwargs
        self.assertEqual(sent['model'],'claude-sonnet-5');self.assertEqual(sent['thinking'],{'type':'disabled'})
        self.assertEqual(sent['max_tokens'],20000)
        self.assertEqual(job.config['request_profile']['max_output_tokens'],8000)
        self.assertEqual(job.config['mapping_request_profile']['max_output_tokens'],20000)
        job.refresh_from_db();self.assertEqual(job.state,'succeeded',job.error)
        usage=Usage.objects.get(job=job)
        self.assertEqual(usage.units['response_model'],'claude-sonnet-5')
        self.assertEqual(usage.actual_cost,Decimal('0.0005'))
        self.assertEqual(usage.charged_cost,Decimal('0.00075'))
        self.assertEqual(self.client.get(f'/billing/jobs/{job.pk}/').status_code,200)
        from core.management.commands.audit_billing import audit_report
        report=audit_report();self.assertTrue(all(report['rows'][0]['checks'].values()),report['rows'][0]['checks'])

    @override_settings(AI_MODE='live',QUILLMD_MAPPING_ASSISTANCE_LIVE=True)
    def test_mocked_incomplete_or_wrong_model_is_receipted_absorbed_and_never_adopted(self):
        self.config=generation.configuration();self.compiled=mapping.compile_request(self.version,self.config)
        for stop,model in [('max_tokens','claude-sonnet-5'),('end_turn','claude-sonnet-4-6')]:
            with self.subTest(stop=stop,model=model):
                job=self.enqueue()
                with patch('anthropic.Anthropic') as client:
                    client.return_value.messages.create.return_value=self.response(stop=stop,model=model)
                    jobs.run_one()
                job.refresh_from_db();self.assertEqual(job.state,'failed',job.error)
                self.assertTrue(ProviderReceipt.objects.filter(job=job).exists())
                self.assertTrue(PlatformCost.objects.filter(job=job,kind='absorbed_failure').exists())
                self.assertEqual(sum(LedgerEntry.objects.filter(job=job).values_list('amount',flat=True)),Decimal(0))
                self.assertEqual(self.template.versions.count(),1)

    @override_settings(AI_MODE='live',QUILLMD_MAPPING_ASSISTANCE_LIVE=True)
    def test_a_reversal_that_fails_holds_the_mapping_job_with_its_key_and_charge(self):
        # RS-3 (docs/rebuild/v1/residuals-assessment-20260929.md), the mapping worker's copy of the rule. RED on 61c5cc8f:
        # 'failed' with the charge kept and the request key freed, so the same mapping could be bought again.
        self.config=generation.configuration();self.compiled=mapping.compile_request(self.version,self.config)
        job=self.enqueue();key=job.request_key
        with patch('anthropic.Anthropic') as client,patch('core.mapping_jobs.absorb_failed_job',side_effect=RuntimeError('synthetic reversal failure')):
            client.return_value.messages.create.return_value=self.response(stop='max_tokens')
            self.assertTrue(jobs.run_one())
        job.refresh_from_db()
        self.assertEqual((job.state,job.request_key),('reconciliation',key))
        self.assertEqual(job.error,jobs.ABSORPTION_FAILED)
        usage=Usage.objects.get(job=job)
        self.assertEqual(-sum(LedgerEntry.objects.filter(job=job,kind='usage').values_list('amount',flat=True)),usage.charged_cost)
        self.assertFalse(LedgerEntry.objects.filter(job=job,kind='absorbed').exists())
        self.assertFalse(PlatformCost.objects.filter(job=job,kind='absorbed_failure').exists())
        self.assertEqual(AuditEvent.objects.get(action='usage.absorb_failed',target=str(job.pk)).details['error_class'],'RuntimeError')
        self.assertEqual(AuditEvent.objects.get(action='mapping.reconciliation',target=str(job.pk)).details['billing'],'platform_failure')
        self.assertEqual(self.enqueue().pk,job.pk)  # the same request is not bought twice
        self.assertEqual(self.template.versions.count(),1)

    @override_settings(AI_MODE='live',QUILLMD_MAPPING_ASSISTANCE_LIVE=True)
    def test_definite_constructor_failure_releases_without_unknown_outcome(self):
        self.config=generation.configuration();self.compiled=mapping.compile_request(self.version,self.config)
        job=self.enqueue()
        with patch('anthropic.Anthropic',side_effect=ValueError('synthetic local setup failure')) as constructor:
            jobs.run_one()
        constructor.assert_called_once()
        job.refresh_from_db();self.assertEqual(job.state,'failed',job.error)
        self.assertFalse(Reservation.objects.filter(job=job,state='held').exists())
        self.assertFalse(ProviderReceipt.objects.filter(job=job).exists())
        self.assertFalse(Usage.objects.filter(job=job).exists())
        self.assertEqual(sum(LedgerEntry.objects.filter(job=job).values_list('amount',flat=True)),Decimal(0))
        self.assertEqual(self.template.versions.count(),1)

    @override_settings(AI_MODE='live',QUILLMD_MAPPING_ASSISTANCE_LIVE=True)
    def test_no_provider_credential_is_not_processed_and_releases_the_hold(self):
        """MONEY-2, the mapping path. Reproduce: live mapping, every ANTHROPIC_* variable unset, a HOME with no profile,
        the REAL anthropic SDK, every connect refused. RED on bf8f7b78: the SDK raised TypeError while building the
        request inside messages.create, so the job went to 'reconciliation' ('Provider outcome or settlement is
        uncertain') with its hold still held. Control: the same real SDK with ANTHROPIC_API_KEY set sends once."""
        import httpx
        from core.tests.test_money_closeout import (REAL_ANTHROPIC, SYNTHETIC_KEY, message_json, no_provider_credential,
                                                   over_mock_transport, synthetic_credential)
        self.config=generation.configuration();self.compiled=mapping.compile_request(self.version,self.config)
        job=self.enqueue();connects=[]
        def refuse(*args,**kwargs):
            connects.append(args);raise OSError('synthetic test: the network is denied')
        with no_provider_credential(),patch('anthropic.Anthropic',REAL_ANTHROPIC),patch('socket.socket.connect',side_effect=refuse):
            self.assertTrue(jobs.run_one())
        job.refresh_from_db();self.assertEqual(job.state,'failed',job.error)
        self.assertTrue(job.error.startswith('not_processed: '),job.error)
        self.assertEqual(list(Reservation.objects.filter(job=job).values_list('state',flat=True)),['released'])
        self.assertFalse(ProviderReceipt.objects.filter(job=job).exists());self.assertFalse(Usage.objects.filter(job=job).exists())
        self.assertEqual(connects,[])
        self.assertTrue(all(w.reserved==0 for w in Wallet.objects.filter(user=self.user)))
        # Positive control: the same real SDK and a fresh request, with the credential present, leaves exactly once.
        fresh=self.enqueue();self.assertNotEqual(fresh.pk,job.pk);sent=[]
        def handler(request):
            sent.append(request)
            return httpx.Response(200,json=message_json(json.dumps(self.body()),ident='msg_synthetic_mapping'),
                                  headers={'request-id':'req_synthetic_mapping'})
        with synthetic_credential(),patch('socket.socket.connect',side_effect=refuse),over_mock_transport(handler):
            self.assertTrue(jobs.run_one())
        fresh.refresh_from_db();self.assertEqual(fresh.state,'succeeded',fresh.error)
        self.assertEqual(len(sent),1);self.assertEqual(sent[0].headers.get('x-api-key'),SYNTHETIC_KEY)
        self.assertEqual(connects,[])

    @override_settings(AI_MODE='live',QUILLMD_MAPPING_ASSISTANCE_LIVE=True)
    def test_unknown_provider_outcome_keeps_hold_and_same_attempt_without_retry(self):
        self.config=generation.configuration();self.compiled=mapping.compile_request(self.version,self.config)
        job=self.enqueue()
        with patch('anthropic.Anthropic') as client:
            client.return_value.messages.create.side_effect=TimeoutError('synthetic unknown outcome')
            jobs.run_one();self.assertFalse(jobs.run_one())
        self.assertEqual(client.return_value.messages.create.call_count,1)
        job.refresh_from_db();self.assertEqual(job.state,'reconciliation')
        self.assertTrue(Reservation.objects.filter(job=job,state='held').exists())
        self.assertEqual(self.enqueue().pk,job.pk)
        self.assertFalse(Usage.objects.exists());self.assertFalse(ProviderReceipt.objects.exists())

    def test_manual_edit_cannot_clear_unresolved_assistance(self):
        self.enqueue();jobs.run_one();version=self.template.versions.latest('number')
        first=version.fields[0]
        response=self.client.post(f'/templates/{self.template.pk}/',{'action':'v3_save','mapping_version':version.pk,
            'label_'+first['key']:first['label']+' clarified'})
        self.assertEqual(response.status_code,302)
        latest=self.template.versions.latest('number')
        self.assertTrue(latest.assistance['unresolved']);self.assertTrue(latest.assistance['assisted'])

    def test_queued_http_cancellation_and_stale_operation_keep_financial_state(self):
        job=self.enqueue()
        response=self.client.post(f'/templates/{self.template.pk}/assistance/',{'action':'cancel','job':job.pk})
        self.assertEqual(response.status_code,302)
        job.refresh_from_db();self.assertEqual(job.state,'cancelled')
        self.assertFalse(ProviderReceipt.objects.filter(job=job).exists())
        job=self.enqueue();job.state='running';job.started_at=timezone.now()-timedelta(minutes=6)
        job.save(update_fields=['state','started_at'])
        jobs.mark_stale();job.refresh_from_db();self.assertEqual(job.state,'reconciliation')
        self.assertEqual(self.enqueue().pk,job.pk)

    def test_reverse_scope_migration_refuses_to_erase_mapping_operations(self):
        import importlib
        from django.db import connection
        reverse=importlib.import_module('core.migrations.0022_mapping_job_scope').REVERSE
        self.enqueue()
        with self.assertRaises(DatabaseError),transaction.atomic():
            with connection.cursor() as cursor:cursor.execute(reverse)
        self.assertEqual(Job.objects.filter(kind='mapping').count(),1)

    @override_settings(AI_MODE='live',QUILLMD_MAPPING_ASSISTANCE_LIVE=True)
    def test_mapping_is_included_in_operator_drain_state(self):
        self.config=generation.configuration();self.compiled=mapping.compile_request(self.version,self.config)
        self.enqueue()
        from core.admin_console import live_generation_queue_row
        row=live_generation_queue_row(True)
        self.assertIn('1 queued, 0 running, 0 in reconciliation',row['value'])

    def test_administrator_financial_scope_and_unknown_resolution_are_usable(self):
        from django_otp.plugins.otp_totp.models import TOTPDevice
        job=self.enqueue();jobs.run_one()
        self.user.is_staff=True;self.user.save(update_fields=['is_staff'])
        device=TOTPDevice.objects.create(user=self.user,name='Synthetic mapping authenticator',confirmed=True)
        session=self.client.session;session['otp_device_id']=device.persistent_id;session.save()
        self.assertEqual(self.client.get(f'/console/jobs/{job.pk}/').status_code,200)
        version=self.template.versions.latest('number')
        compiled=mapping.compile_request(version,self.config)
        pending=mapping_jobs.enqueue(version,self.user,self.config,compiled['request_sha256'])
        pending.state='reconciliation';pending.save(update_fields=['state'])
        response=self.client.post(f'/console/jobs/{pending.pk}/resolve/',{'outcome':'release_no_cost','evidence_note':'Synthetic no-request fixture resolution.'})
        self.assertEqual(response.status_code,302)
        pending.refresh_from_db();self.assertEqual(pending.state,'failed')
        self.assertFalse(Reservation.objects.filter(job=pending,state='held').exists())

    def test_upload_quote_queue_uses_existing_mapping_review_route(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        response=self.client.post('/templates/upload/',{'department':self.department.pk,'name':'Synthetic upload mapping',
            'synthetic_only':'on','file':SimpleUploadedFile('synthetic.docx',self.original)})
        self.assertEqual(response.status_code,302)
        template=Template.objects.get(name='Synthetic upload mapping')
        page=self.client.get(f'/templates/{template.pk}/assistance/')
        self.assertEqual(page.status_code,200)
        response=self.client.post(f'/templates/{template.pk}/assistance/',{'action':'queue','approval':page.context['approval'],'blank_only':'on'})
        self.assertEqual(response.status_code,302)
        with patch('anthropic.Anthropic') as provider:jobs.run_one();provider.assert_not_called()
        latest=template.versions.latest('number')
        self.assertTrue(latest.assistance['simulated']);self.assertTrue(latest.needs_confirmation)

    def test_mapping_quote_prices_its_exact_profile_without_rewriting_generation(self):
        before=copy.deepcopy(self.config)
        self.assertEqual(self.compiled['request']['max_tokens'],20000)
        self.assertEqual(self.compiled['profile']['timeout_seconds'],120)
        self.assertEqual(self.compiled['profile_sha256'],mapping.digest(self.compiled['profile']))
        payload=json.loads(self.compiled['request']['messages'][0]['content'])
        self.assertEqual(payload['profile_sha256'],self.compiled['profile_sha256'])
        expected=(Decimal(self.compiled['input_token_bound'])*2+Decimal(20000)*10)/1000000
        self.assertEqual(Decimal(self.compiled['maximum_cost']),expected)
        self.assertEqual(self.config,before)
        self.assertEqual(self.config['request_profile']['max_output_tokens'],8000)
        with override_settings(AI_INFERENCE_GEO='us'):
            config=generation.configuration();compiled=mapping.compile_request(self.version,config)
            expected=(Decimal(compiled['input_token_bound'])*2+Decimal(20000)*10)/1000000*generation.US_INFERENCE_MULTIPLIER
            self.assertEqual(Decimal(compiled['maximum_cost']),expected.quantize(generation.COST_QUANTUM,rounding='ROUND_CEILING'))

    def test_old_queued_contract_refuses_before_new_compiler_and_releases(self):
        from core.billing import reserve
        config={**self.config,'mapping_contract':'blank-mapping-v1','source_sha256':self.template.sha256,
                'department_id':str(self.department.pk),'mapping_version_id':str(self.version.pk)}
        job=Job.objects.create(kind='mapping',mapping_version=self.version,user=self.user,config=config,request_key='synthetic-old-contract')
        reserve(job,Decimal(0))
        with patch('core.mapping_assistance.compile_request') as compile_new,patch('core.mapping_jobs.send') as send:
            jobs.run_one();compile_new.assert_not_called();send.assert_not_called()
        job.refresh_from_db();self.assertEqual(job.state,'failed')
        self.assertEqual(job.config,config)
        self.assertFalse(Reservation.objects.filter(job=job,state='held').exists())
        self.assertFalse(ProviderReceipt.objects.filter(job=job).exists())

    def test_incomplete_capacity_is_visible_and_prevents_unassisted_confirmation(self):
        body=self.body();body['e']=True
        body['s']={'c':[],'n':list(range(len(json.loads(self.compiled['request']['messages'][0]['content'])['slots']))),'u':[]}
        job=self.enqueue()
        with patch('core.mapping_jobs.send',return_value=(body,Decimal(0),{},{},'synthetic-mapping-fixture','','',True)):
            jobs.run_one()
        latest=self.template.versions.latest('number')
        self.assertTrue(latest.assistance['incomplete'])
        self.assertEqual(len(latest.assistance['areas']),len(body['s']['n']))
        self.assertTrue(any('capacity' in area['reason'] for area in latest.assistance['unresolved']))
        self.assertEqual(self.client.post(f'/templates/{self.template.pk}/',{'action':'v3_confirm','mapping_version':latest.pk,'checked_every_field':'on'}).status_code,400)

    def test_signature_downgrade_remains_visible_and_preserves_existing_field(self):
        body=self.body();prior=next(f for f in self.version.fields if f['kind']=='signature')
        fid=self.version.fields.index(prior);body['f']['k'].remove(fid);body['f']['r']=[fid]
        body['x']=[self.target(next(iter(mapping.v3_mapping.mapped_slots([prior]))),prior['label'],'text',fid)]
        result=mapping.validate_proposal(self.version,self.compiled,body)
        self.assertIn(prior,result['fields'])
        self.assertTrue(any('Signature' in row['reason'] for row in result['review']['unresolved']))

    def test_tampered_mapping_profile_or_hash_refuses_unsent(self):
        job=self.enqueue()
        for update in [{'mapping_profile_sha256':'0'*64},{'mapping_request_profile':{**job.config['mapping_request_profile'],'max_output_tokens':20001}},
                       {'mapping_request_profile':{**job.config['mapping_request_profile'],'timeout_seconds':150}}]:
            with self.subTest(update=update),self.assertRaises(jobs.NoProviderRequest):
                mapping_jobs.permission({**job.config,**update},queued=True)

    def test_minimum_and_maximum_metadata_quote_match_request_and_byte_admission(self):
        # Pure synthetic source metadata substitutes only the outline for this compiler boundary.
        for count,field_count in [(0,0),(1200,200)]:
            rows=[{'slot':f'word/document.xml#body/p[{n}]','text':'Plan','heading':'','row_label':'','options':[]} for n in range(count)]
            fields=[{**self.version.fields[0],'key':f'synthetic_{n}'} for n in range(field_count)]
            version=SimpleNamespace(template=self.template,contract='v3',document_backend='',fields=fields)
            with self.subTest(count=count),patch.object(mapping,'assistance_outline',return_value={'sha256':self.template.sha256,'slots':rows}):
                compiled=mapping.compile_request(version,self.config)
                self.assertEqual(compiled['request']['max_tokens'],20000)
                self.assertEqual(compiled['input_token_bound'],len(mapping.canonical(compiled['request']).encode())+12000)
                expected=(Decimal(compiled['input_token_bound'])*2+Decimal(20000)*10)/1000000
                self.assertEqual(Decimal(compiled['maximum_cost']),expected)
                size=len(mapping.canonical(compiled['request']).encode())
                with patch.object(mapping,'MAX_PAYLOAD_BYTES',size):mapping.compile_request(version,self.config)
                with patch.object(mapping,'MAX_PAYLOAD_BYTES',size-1),self.assertRaises(ValidationError):mapping.compile_request(version,self.config)
        with patch.object(mapping,'MAX_SLOTS',1199),patch.object(mapping,'assistance_outline',return_value={'sha256':self.template.sha256,'slots':rows}),self.assertRaises(ValidationError):
            mapping.compile_request(version,self.config)
