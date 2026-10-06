"""Persisted manual printed-option recovery through the ordinary doctor routes."""
import io

from django.test import TestCase, override_settings
from docx import Document

from core.models import Encounter, Job
from core.jobs import enqueue, run_one
from core.security import private_path
from core.tests.test_word_v3_workflow import V3, V3Fixture
from documents.test_paragraph_option_recovery import source


@override_settings(**V3)
class ParagraphOptionWorkflow(V3Fixture, TestCase):
    @override_settings(AI_PROVIDER='anthropic', AI_MODEL='claude-sonnet-5-5', AI_THINKING='between_tools',
                       AI_EFFORT='medium', AI_INPUT_RATE='2', AI_OUTPUT_RATE='10', AI_CACHED_RATE='.2',
                       V3_FIELD_RELATIONS=True)
    def test_mapping_review_owns_complete_list_separately_from_narrative_and_downloads_reviewed_choices(self):
        doc = Document(io.BytesIO(source()))
        # Same printed-list + separate response-box family as D30. The source
        # list starts unassigned as selections; the box has its own narrative.
        doc.paragraphs[-1]._p.getparent().remove(doc.paragraphs[-1]._p)
        box = doc.add_table(rows=1, cols=1)
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        borders = OxmlElement('w:tcBorders')
        for edge in ('top', 'bottom', 'left', 'right'):
            node = OxmlElement('w:' + edge); node.set(qn('w:val'), 'single'); borders.append(node)
        box.cell(0, 0)._tc.get_or_add_tcPr().append(borders)
        doc.add_paragraph('Plan: ______')
        stream = io.BytesIO(); doc.save(stream); original = stream.getvalue()
        template, first = self.uploaded(data=original, name='Synthetic complete printed list')
        narrative = next(f for f in first.fields if f['locator']['slot'] == 'b/tbl[0]/tr[0]/tc[0]')
        self.assertNotIn('options', narrative)
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_add', 'mapping_version': first.pk,
            'slot': 'word/document.xml#b/p[1]', 'label': 'Accommodation arrangements', 'kind': 'multi_choice'})
        self.assertEqual(response.status_code, 302, response.content.decode())
        mapped = self.latest(template)
        choices = next(f for f in mapped.fields if f['kind'] == 'multi_choice')
        self.assertEqual(choices['origin'], 'manual')
        self.assertEqual(choices['options'], ['Own flat', 'With family', 'With partner', 'Other'])
        self.assertEqual(len(choices['locator']['target']['paragraphs']), 4)
        self.assertIn(choices['key'], mapped.corrections['added'])
        # The same group selected through another source alias cannot make a
        # duplicate writer; the append-only mapping version does not change.
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_add', 'mapping_version': mapped.pk,
            'slot': 'word/document.xml#b/p[2]', 'label': 'Duplicate list', 'kind': 'choice'})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.latest(template).pk, mapped.pk)
        self.confirm(template, mapped)
        self.assertEqual(self.new_encounter(mapped, 'The patient lives in an own flat with their partner.').status_code, 302)
        encounter = Encounter.objects.latest('created_at'); job = enqueue(encounter, self.user)
        self.assertTrue(run_one()); job.refresh_from_db(); self.assertEqual(job.state, 'succeeded', job.error)
        encounter.refresh_from_db()
        # Explicitly corrected choices are useful recovery, never a first-pass
        # generation success or a correction to the frozen corpus denominator.
        post = {'revision': encounter.revision}
        for field in mapped.fields:
            key = field['key']; post['reviewed_' + key] = 'on'
            post['answer_' + key] = 'Own flat\nWith partner' if key == choices['key'] else ''
            if key != choices['key']: post['blank_' + key] = 'on'
        response = self.client.post(f'/encounters/{encounter.pk}/save/', post)
        self.assertEqual(response.status_code, 302, response.content.decode()); encounter.refresh_from_db()
        self.assertEqual(encounter.answers[choices['key']]['status'], 'edited')
        response = self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'}, follow=True)
        encounter.refresh_from_db(); self.assertEqual(encounter.status, 'finalized')
        raw = private_path(encounter.export_path).read_bytes(); actual = Document(io.BytesIO(raw))
        self.assertEqual([p.text for p in actual.paragraphs[2:6]], ['☒ Own flat', 'With family', '☒ With partner', 'Other'])
        self.assertEqual(actual.tables[0].cell(0, 0).text, '')
        self.assertEqual(private_path(template.original_path).read_bytes(), original)
        download = self.client.get(f'/encounters/{encounter.pk}/export/')
        self.assertEqual(download.status_code, 200); self.assertEqual(b''.join(download.streaming_content), raw)
