"""Bounded ordinary context repairs. Retained synthetic notes; mocked unit selections, no provider."""
from django.test import SimpleTestCase
from core import generation as g
from .test_v3_ordinary_notes import FIELDS, field, run


class IndependentPrecedingContext(SimpleTestCase):
    def assert_shared_supported(self, notes, fields, label, quote, expected):
        answer = run(notes, {label: [quote]}, fields=fields)[label]
        self.assertEqual((answer['status'], answer['text']), ('supported', expected))
        governance = g.governance_for(notes, fields)
        self.assertIsNone(g.governed_supported(governance, notes, answer, label))
        key = next(f['key'] for f in fields if f['label'] == label)
        self.assertEqual(g.governed_decisions(notes, fields, {key: {**answer, 'reviewed': True}}), [])
        self.assertEqual(answer['evidence'][0]['quote'], quote)
        self.assertEqual(notes[answer['evidence'][0]['start']:answer['evidence'][0]['end']], quote)

    def test_retained_patient_sentence_after_physician_answer(self):
        fields = [field('r', 'Referring physician'), field('v', 'Reason for visit'),
                  field('s', 'Smoker?', 'choice', ['Yes', 'No'])]
        quote = 'Patient presents with right knee pain.'
        notes = 'Referring physician: Dr Example\n' + quote + '\nSmoker: yes'
        self.assert_shared_supported(notes, fields, 'Reason for visit', quote, quote)
        answers = run(notes, {'Referring physician': ['Referring physician: Dr Example'],
                             'Reason for visit': [quote], 'Smoker?': ['Smoker: yes']}, fields=fields)
        self.assertEqual({label: a['text'] for label, a in answers.items()},
                         {'Referring physician': 'Dr Example', 'Reason for visit': quote, 'Smoker?': 'Yes'})
        self.assertEqual({a['status'] for a in answers.values()}, {'supported'})

    def test_retained_explicit_plan_after_plain_answers(self):
        quote = 'Plan is physiotherapy twice weekly and review in six weeks.'
        for prefix in ('Patient presents with right knee pain for three weeks. Allergies: penicillin. ',
                       'Patient presents with right knee pain for three weeks.\nAllergies: penicillin.\nSmoker: No.\n'):
            with self.subTest(prefix=prefix):
                self.assert_shared_supported(prefix + quote, FIELDS, 'Plan', quote, quote[len('Plan is '):])

    def test_prior_scope_and_anaphoric_subject_are_not_independent(self):
        fields = [field('r', 'Referring physician'), field('a', 'Allergies'), field('v', 'Reason for visit'), field('p', 'Plan', 'long_text')]
        for preceding, quote, label in (
            ('Her sister was seen in oncology last month.\nAllergies: penicillin.\n', 'Plan is chemotherapy.', 'Plan'),
            ('Referring physician: Dr Example\n', 'He reports right knee pain.', 'Reason for visit'),
            ('Family history:\nAllergies: penicillin.\n', 'Plan is chemotherapy.', 'Plan'),
            ('If the wound looks infected:\nAllergies: penicillin.\n', 'Plan is cephalexin.', 'Plan'),
        ):
            with self.subTest(notes=preceding + quote):
                answer = run(preceding + quote, {label: [quote]}, fields=fields)[label]
                self.assertEqual((answer['status'], answer['text']), ('conflict', ''))

    def test_corrections_consent_duplicate_and_wrong_destinations_stay_held(self):
        quote = 'Patient presents with right knee pain.'
        for label, kind in (('Lt knee', 'long_text'), ('Allergies', 'text'), ('Consent obtained', 'text')):
            fields = [field('r', 'Referring physician'), field('v', label, kind)]
            notes = 'Referring physician: Dr Example\n' + quote
            self.assertEqual(run(notes, {label: [quote]}, fields=fields)[label]['status'], 'conflict')
        fields = [field('r', 'Referring physician'), field('v', 'Reason for visit'), field('w', 'Reason for visit')]
        notes = 'Referring physician: Dr Example\n' + quote
        self.assertEqual(run(notes, {'Reason for visit': [quote]}, fields=fields)['Reason for visit']['status'], 'conflict')
        fields = [field('r', 'Referring physician'), field('v', 'Reason for visit')]
        for tail in ('\nSorry, I meant the left knee.', '\n\nActually, he has left knee pain.'):
            self.assertEqual(run(notes + tail, {'Reason for visit': [quote]}, fields=fields)['Reason for visit']['status'], 'conflict')

    def test_a_patient_presentation_cannot_gain_a_plan_destination(self):
        quote = 'Patient presents with right knee pain.'
        for label in ('Plan', 'Treatment plan', 'Management plan'):
            fields = [field('r', 'Referring physician'), field('p', label, 'long_text'),
                      field('s', 'Smoker', 'choice', ['Yes', 'No'])]
            notes = 'Referring physician: Dr Example\n' + quote + '\nSmoker: yes'
            answer = run(notes, {label: [quote]}, fields=fields)[label]
            self.assertEqual((answer['status'], answer['text']), ('conflict', ''), label)

    def test_a_complete_explicit_plan_never_drops_following_context(self):
        quote = 'Plan is physiotherapy twice weekly and review in six weeks.'
        for tail in ('\nIt was cancelled.', '\nSorry, only once weekly.', '\nThe patient declined therapy.'):
            notes = 'Allergies: penicillin.\n' + quote + tail
            self.assertEqual(run(notes, {'Plan': [quote]})['Plan']['status'], 'conflict')


class IndependentSourceBlocks(SimpleTestCase):
    def test_the_retained_complete_ordinary_note_has_three_exact_answers(self):
        quotes = ['Patient presents with right knee pain for three weeks.', 'Allergies: penicillin.',
                  'Plan is physiotherapy twice weekly and review in six weeks.']
        fields = [field('c', 'Chief complaint'), field('a', 'Allergies'), field('p', 'Plan', 'long_text')]
        for separator in (' ', '\n'):
            notes = separator.join(quotes)
            answers = run(notes, dict(zip(('Chief complaint', 'Allergies', 'Plan'), ([q] for q in quotes))), fields=fields)
            for label, expected in zip(('Chief complaint', 'Allergies', 'Plan'), (quotes[0], 'penicillin.', quotes[2][8:])):
                self.assertEqual((answers[label]['status'], answers[label]['text']), ('supported', expected))
            stored = {f['key']: {**answers[f['label']], 'reviewed': True} for f in fields}
            self.assertEqual(g.governed_decisions(notes, fields, stored), [])

    def test_explicit_other_topic_does_not_erase_a_correction_or_inherited_subject(self):
        for notes, picks in (
            ('Allergies: penicillin.\nPlan is amoxicillin, as the penicillin allergy was disproven.',
             {'Allergies': ['Allergies: penicillin.']}),
            ('For the patient\'s mother only:\nAllergies: penicillin.\nPlan: return in two weeks.',
             {'Allergies': ['Allergies: penicillin.'], 'Plan': ['Plan: return in two weeks.']}),
            ('If infection develops:\nAllergies: penicillin.\nPlan: amoxicillin.',
             {'Allergies': ['Allergies: penicillin.'], 'Plan': ['Plan: amoxicillin.']}),
            ('Consent obtained: yes.\nPlan is physiotherapy twice weekly.',
             {'Consent obtained': ['Consent obtained: yes.']}),
        ):
            for label, answer in run(notes, picks).items():
                if label in picks:self.assertEqual((answer['status'], answer['text']), ('conflict', ''), notes)

    def test_a_continuation_stays_in_its_source_block_and_cannot_be_omitted(self):
        notes = 'Allergies: penicillin.\nIt was disproven on testing.\nPlan is physiotherapy twice weekly.'
        answer = run(notes, {'Allergies': ['Allergies: penicillin.']})['Allergies']
        self.assertEqual((answer['status'], answer['text']), ('conflict', ''))


class ConditionalPlanScopes(SimpleTestCase):
    def test_complete_plan_keeps_its_conditional_return_instruction(self):
        fields = [field('h', 'Chief complaint'), field('a', 'Allergies'), field('p', 'Plan', 'long_text')]
        head = 'Patient presents with a cough for five days.'
        plan = 'Plan: amoxicillin 500 mg three times daily for five days.'
        follow = 'Return if symptoms worsen.'
        for notes, plan_quotes, expected in (
            ('\n'.join((head, 'Allergies: none known.', plan, follow)), [plan, follow], plan[6:] + '\n' + follow),
            ('\n'.join((head, 'Allergies: none known.', plan + ' ' + follow)), [plan + ' ' + follow], plan[6:] + ' ' + follow),
        ):
            picks = {'Chief complaint': [head], 'Allergies': ['Allergies: none known.'], 'Plan': plan_quotes}
            answers = run(notes, picks, fields=fields)
            self.assertEqual((answers['Plan']['status'], answers['Plan']['text']), ('supported', expected))
            self.assertEqual(answers['Chief complaint']['text'], head)
            self.assertEqual(answers['Allergies']['text'], 'none known.')
            stored = {f['key']: {**answers[f['label']], 'reviewed': True} for f in fields}
            self.assertEqual(g.governed_decisions(notes, fields, stored), [])

    def test_conditional_plan_does_not_authorize_omission_or_conditional_findings(self):
        fields = [field('h', 'Chief complaint'), field('a', 'Allergies'), field('p', 'Plan', 'long_text')]
        plan = 'Plan: amoxicillin 500 mg three times daily for five days.'
        notes = 'Allergies: none known.\n' + plan + '\nReturn if symptoms worsen.'
        self.assertEqual(run(notes, {'Plan': [plan]}, fields=fields)['Plan']['status'], 'conflict')
        for prefix in ('If infection develops:\n', 'For the patient\'s mother only:\n'):
            governed = prefix + notes
            self.assertEqual(run(governed, {'Allergies': ['Allergies: none known.'], 'Plan': [plan, 'Return if symptoms worsen.']},
                                 fields=fields)['Plan']['status'], 'conflict')
        condition = 'If the wound looks infected, start amoxicillin.'
        self.assertEqual(run(condition, {'Chief complaint': [condition]}, fields=fields)['Chief complaint']['status'], 'conflict')


from io import BytesIO
from pathlib import Path
import hashlib
import json
from unittest.mock import patch
from docx import Document
from django.test import TestCase, override_settings
from core.jobs import enqueue, run_one
from core.models import Encounter
from core.security import private_path
from .test_v3_ordinary_notes import LIVE, unit_response
from .test_word_v3_workflow import V3EncounterFixture, part_text


def complete_note_form(path):
    doc = Document()
    doc.add_paragraph('Synthetic ordinary clinical form')
    table = doc.add_table(rows=3, cols=2)
    for cell, label in zip((table.cell(i, 0) for i in range(3)), ('Chief complaint', 'Allergies', 'Plan')):
        cell.text = label
    out = BytesIO(); doc.save(out)
    Path(path).write_bytes(out.getvalue())
    return out.getvalue()


@override_settings(**LIVE)
class CompleteOrdinaryNoteWorkflow(V3EncounterFixture, TestCase):
    def setUp(self):
        with patch('core.tests.test_word_v3_workflow.build_intake_form', side_effect=complete_note_form):
            super().setUp()
        self.assertEqual([f['label'] for f in self.version.fields], ['Chief complaint', 'Allergies', 'Plan'])

    def walk(self, notes, picks, expected):
        response = self.client.post(f'/encounters/{self.encounter.pk}/notes/',
                                   {'revision': self.encounter.revision, 'notes': notes})
        self.assertEqual(response.status_code, 302)
        with patch('anthropic.Anthropic') as client:
            client.return_value.messages.create.side_effect = lambda **kw: unit_response(json.loads(kw['messages'][0]['content']), picks)
            job = enqueue(Encounter.objects.get(pk=self.encounter.pk), self.user)
            self.assertTrue(run_one())
            self.assertEqual(client.return_value.messages.create.call_count, 1)
        job.refresh_from_db(); self.assertEqual(job.state, 'succeeded', job.error)
        self.encounter.refresh_from_db()
        answers = self.encounter.answers
        self.assertEqual({label: (answers[self.key(label)]['status'], answers[self.key(label)]['text']) for label in expected},
                         {label: ('supported', value) for label, value in expected.items()})
        page = self.client.get(f'/encounters/{self.encounter.pk}/')
        self.assertEqual({r['label']: r['status'] for r in page.context['answer_rows']}, dict.fromkeys(expected, 'supported'))
        self.assertEqual(self.save(expected).status_code, 302)
        self.encounter.refresh_from_db()
        self.assertEqual({label: self.encounter.answers[self.key(label)]['status'] for label in expected}, dict.fromkeys(expected, 'supported'))
        response = self.finalize()
        self.encounter.refresh_from_db()
        self.assertEqual(self.encounter.status, 'finalized', [str(m) for m in response.context['messages']])
        exported = private_path(self.encounter.export_path).read_bytes()
        body = part_text(exported)
        for value in expected.values():
            # Word's runs encode a newline as a break; every original clause must remain.
            for line in value.splitlines():self.assertIn(line, body)
        self.assertEqual(hashlib.sha256(exported).hexdigest(), self.encounter.export_sha256)
        self.assertEqual(private_path(self.template.original_path).read_bytes(), self.original)

    def test_complete_ordinary_prose_reaches_the_original_word_form(self):
        head = 'Patient presents with right knee pain for three weeks.'
        allergy = 'Allergies: penicillin.'
        plan = 'Plan is physiotherapy twice weekly and review in six weeks.'
        self.walk(' '.join((head, allergy, plan)), {'Chief complaint': [head], 'Allergies': [allergy], 'Plan': [plan]},
                  {'Chief complaint': head, 'Allergies': 'penicillin.', 'Plan': plan[8:]})

    def test_complete_conditional_plan_reaches_the_original_word_form(self):
        head = 'Patient presents with a cough for five days.'
        allergy = 'Allergies: none known.'
        plan = 'Plan: amoxicillin 500 mg three times daily for five days.'
        follow = 'Return if symptoms worsen.'
        self.walk('\n'.join((head, allergy, plan, follow)),
                  {'Chief complaint': [head], 'Allergies': [allergy], 'Plan': [plan, follow]},
                  {'Chief complaint': head, 'Allergies': 'none known.', 'Plan': plan[6:] + '\n' + follow})
