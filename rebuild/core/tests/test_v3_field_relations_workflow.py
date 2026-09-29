"""Authenticated synthetic workflow; SDK responses simulated, no external calls."""
import copy
import hashlib
import io
import json
import os
import uuid
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.test import TestCase, override_settings
from docx import Document

from core import generation as g
from core.jobs import run_one
from core.models import AuditEvent, Encounter, Job, Reservation, Usage
from core.security import private_path
from documents import word_v3 as w
from .test_v3_field_relations import relation_wire
from .test_v3_ordinary_input import ordinary_input_form
from .test_v3_ordinary_notes import LIVE
from .test_word_v3_workflow import V3Fixture


def sdk_response(body):
    identity = str(uuid.uuid4())
    return SimpleNamespace(id='msg_synthetic_' + identity, _request_id=identity, model=g.SONNET_5,
        stop_reason='end_turn', content=[SimpleNamespace(type='text', text=json.dumps(body))],
        usage=SimpleNamespace(input_tokens=100, output_tokens=100, cache_read_input_tokens=0,
                              cache_creation_input_tokens=0),
        model_dump=lambda **kw: {'id': 'msg_synthetic_' + identity})


def other_form():
    doc = Document()
    for question in ('What does your headache usually feel like?', 'Which symptoms occur before headache begins?'):
        doc.add_heading(question, level=2)
        doc.add_paragraph('Other: ____________')
    output = io.BytesIO(); doc.save(output)
    return output.getvalue()


@override_settings(**dict(LIVE, V3_FIELD_RELATIONS=True, V3_HISTORY_CONTEXT=False, V3_GENERATION_BATCHES=False))
class OrdinaryRelationWorkflow(V3Fixture, TestCase):
    def setUp(self):
        super().setUp()
        deny = patch('socket.socket.connect', side_effect=AssertionError('External network denied'))
        deny.start(); self.addCleanup(deny.stop)

    def walk(self, source, notes, target_label=None, target_context=None, expected=None, selected=None):
        selected = notes if selected is None else selected
        template, version = self.uploaded(data=source, name='Synthetic ordinary relation ' + str(Encounter.objects.count()))
        self.confirm(template, version)
        target = next(f for f in version.fields if
            (target_label is None or g.question_key(f['label']) == g.question_key(target_label)) and
            (target_context is None or target_context in f.get('context', '')))
        self.assertEqual(self.new_encounter(version, notes=notes).status_code, 302)
        encounter = Encounter.objects.latest('created_at')
        page = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertEqual(page.status_code, 200)
        quote = page.context['usage_quote']
        captured = []
        def respond(**wire):
            captured.append(wire)
            payload = json.loads(wire['messages'][0]['content'])
            self.assertEqual(payload['catalogue'], version.fields)
            choices = {target['key']: [target['key'] + (':1' if expected == 'No' else ':0')]} if target['kind'] == 'choice' else {}
            return sdk_response(relation_wire(notes, version.fields, {target['key']: [selected]}, choices=choices))
        with patch('anthropic.Anthropic') as provider:
            provider.return_value.messages.create.side_effect = respond
            self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/generate/',
                {'usage_quote': quote}).status_code, 302)
            self.assertTrue(run_one())
            self.assertEqual(provider.return_value.messages.create.call_count, 1)
        job = Job.objects.get(encounter=encounter)
        self.assertEqual(job.state, 'succeeded', job.error)
        reservation = Reservation.objects.get(job=job)
        usage = Usage.objects.get(job=job)
        self.assertNotEqual(reservation.state, 'held')
        self.assertGreater(usage.charged_cost, 0)
        self.assertLessEqual(usage.charged_cost, reservation.amount)
        encounter.refresh_from_db(); generated = copy.deepcopy(encounter.answers)
        answer = generated[target['key']]
        expected = selected if expected is None else expected
        self.assertEqual((answer['status'], answer['text']), ('supported', expected))
        start = notes.index(selected)
        self.assertEqual(answer['evidence'], [{'quote': selected, 'start': start, 'end': start + len(selected)}])
        page = self.client.get(f'/encounters/{encounter.pk}/')
        rows = {row['key']: row for row in page.context['answer_rows']}
        self.assertEqual((rows[target['key']]['status'], rows[target['key']]['text']), ('supported', expected))
        payload = {'revision': encounter.revision}
        for f in version.fields:
            value = generated[f['key']]
            self.assertIn(value['status'], ('supported', 'missing'))
            payload['answer_' + f['key']] = value['text']
            payload['reviewed_' + f['key']] = 'on'
            if not value['text']: payload['blank_' + f['key']] = 'on'
        for requirement in version.review_requirements:
            payload['requirement_' + requirement['key']] = 'on'
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/save/', payload).status_code, 302)
        encounter.refresh_from_db()
        self.assertEqual(encounter.answers[target['key']]['status'], 'supported')
        self.assertEqual(encounter.answers[target['key']]['text'], expected)
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/finalize/',
            {'revision': encounter.revision, 'reviewed': 'on'}).status_code, 302)
        encounter.refresh_from_db(); self.assertEqual(encounter.status, 'finalized')
        response = self.client.get(f'/encounters/{encounter.pk}/export/')
        self.assertEqual(response.status_code, 200)
        exported = b''.join(response.streaming_content)
        self.assertEqual(hashlib.sha256(exported).hexdigest(), encounter.export_sha256)
        self.assertEqual(private_path(template.original_path).read_bytes(), source)
        self.assertTrue(AuditEvent.objects.filter(action='encounter.finalize', target=str(encounter.pk), actor=self.user).exists())
        self.assertTrue(AuditEvent.objects.filter(action='encounter.export', target=str(encounter.pk), actor=self.user).exists())
        from core.models import User
        from django.utils import timezone
        outsider = User.objects.create_user(username='outsider_' + str(encounter.pk),
            email=str(encounter.pk) + '@synthetic.invalid', verified_at=timezone.now())
        self.client.force_login(outsider)
        self.assertIn(self.client.get(f'/encounters/{encounter.pk}/export/').status_code, (403, 404))
        self.client.force_login(self.user)
        if target['kind'] == 'text':
            slots = w.SlotIndex(w.open_package(exported))
            slot = slots.by_id[target['locator']['part'] + '#' + target['locator']['slot']]
            self.assertIn(expected, w.text_of(slot.el))
        if (folder := os.environ.get('QUILLMD_ORDINARY_ARTIFACTS')):
            out = Path(folder); out.mkdir(parents=True, exist_ok=False)
            (out / 'original.docx').write_bytes(source)
            (out / 'exported.docx').write_bytes(exported)
            (out / 'request.json').write_text(json.dumps(captured[0], indent=2))
            (out / 'result.json').write_text(json.dumps({'synthetic_provider': True, 'notes': notes,
                'target': target, 'generated': generated, 'final': encounter.answers,
                'export_sha256': encounter.export_sha256, 'reservation': str(reservation.amount),
                'charged': str(usage.charged_cost), 'no_answer_correction': True}, indent=2))
        return encounter

    def test_repeated_other_full_context_no_answer_correction_authenticated_export(self):
        source = Path(os.environ['QUILLMD_ORDINARY_ORIGINAL']).read_bytes() if os.environ.get('QUILLMD_ORDINARY_ORIGINAL') else other_form()
        self.walk(source, 'The patient reports that the headache usually feels like intermittent buzzing.',
                  target_label='Other:', target_context='feel like')

    def test_same_side_ordinary_sentence_and_predicate_choice_export_unchanged(self):
        for label, note, expected in (
            ('Right knee', 'The patient has right knee pain.', None),
            ('Smoker', 'The patient does not smoke.', 'No'),
            ('Smoker', 'The patient smokes.', 'Yes')):
            with self.subTest(label=label):
                source = ordinary_input_form(Path(self.directory.name) / (label + '.docx'), [label])
                self.walk(source, note, target_label=label, expected=expected)

    def test_d31_mappa_direct_opposite_is_held_through_the_worker_and_agreement_exports(self):
        # The D31 question, the selected positive statement and every other unit
        # declared "independent" by the simulated provider, as the refuters ran it.
        label = 'Please confirm that this patient is subject to Multi-Agency Public Protection Arrangements (MAPPA):'
        full = 'The patient is subject to Multi-Agency Public Protection Arrangements (MAPPA).'
        source = ordinary_input_form(Path(self.directory.name) / 'd31-mappa.docx', [label])
        for other in ('The patient is subject to MAPPA.', 'Patients subject to MAPPA are reviewed monthly.'):
            with self.subTest(other=other):
                self.walk(source, full + ' ' + other, target_label=label, selected=full)
        template, version = self.uploaded(data=source, name='Synthetic D31 opposite'); self.confirm(template, version)
        target = next(f for f in version.fields if g.question_key(f['label']) == g.question_key(label))
        for other in ('Subject to MAPPA: no.', 'MAPPA: No.', 'Is the patient subject to MAPPA? No.',
                      'Our client is not subject to MAPPA.', 'The patient is not subject to MAPPA at present.',
                      'Today the patient was not subject to MAPPA.', 'The patient is subject to no MAPPA arrangements.'):
            notes = full + ' ' + other
            with self.subTest(other=other):
                self.assertEqual(self.new_encounter(version, notes=notes).status_code, 302)
                encounter = Encounter.objects.latest('created_at')
                quote = self.client.get(f'/encounters/{encounter.pk}/').context['usage_quote']
                body = relation_wire(notes, version.fields, {target['key']: [full]})
                with patch('anthropic.Anthropic') as provider:
                    provider.return_value.messages.create.side_effect = lambda **wire: sdk_response(body)
                    self.client.post(f'/encounters/{encounter.pk}/generate/', {'usage_quote': quote})
                    self.assertTrue(run_one())
                self.assertEqual(Job.objects.get(encounter=encounter).state, 'succeeded')
                encounter.refresh_from_db()
                self.assertEqual((encounter.answers[target['key']]['status'], encounter.answers[target['key']]['text']),
                                 ('conflict', ''))
                # A supported copy stored before this hold existed is held on review, save and finalize.
                encounter.answers = {**encounter.answers, target['key']: {'status': 'supported', 'text': full,
                    'reviewed': True, 'evidence': [{'quote': full, 'start': 0, 'end': len(full)}]}}
                encounter.save()
                self.assertContains(self.client.get(f'/encounters/{encounter.pk}/'), 'data-held-candidate')
                for save in (False, True):
                    if save:
                        self.client.post(f'/encounters/{encounter.pk}/save/', {'revision': encounter.revision,
                            'answer_' + target['key']: full, 'reviewed_' + target['key']: 'on'})
                        encounter.refresh_from_db()
                    self.client.post(f'/encounters/{encounter.pk}/finalize/',
                                     {'revision': encounter.revision, 'reviewed': 'on'})
                    encounter.refresh_from_db()
                    self.assertNotEqual(encounter.status, 'finalized')
                    self.assertFalse(encounter.export_path)

    def assert_worker_holds(self, version, target, selected, notes):
        """Generated through enqueue and the worker with the provider patched, and a
        stored supported copy held on review, save and finalize."""
        self.assertEqual(self.new_encounter(version, notes=notes).status_code, 302)
        encounter = Encounter.objects.latest('created_at')
        quote = self.client.get(f'/encounters/{encounter.pk}/').context['usage_quote']
        body = relation_wire(notes, version.fields, {target['key']: [selected]})
        with patch('anthropic.Anthropic') as provider:
            provider.return_value.messages.create.side_effect = lambda **wire: sdk_response(body)
            self.client.post(f'/encounters/{encounter.pk}/generate/', {'usage_quote': quote})
            self.assertTrue(run_one())
        self.assertEqual(Job.objects.get(encounter=encounter).state, 'succeeded')
        encounter.refresh_from_db()
        self.assertEqual((encounter.answers[target['key']]['status'], encounter.answers[target['key']]['text']),
                         ('conflict', ''))
        start = notes.index(selected)
        encounter.answers = {**encounter.answers, target['key']: {'status': 'supported', 'text': selected,
            'reviewed': True, 'evidence': [{'quote': selected, 'start': start, 'end': start + len(selected)}]}}
        encounter.save()
        self.assertContains(self.client.get(f'/encounters/{encounter.pk}/'), 'data-held-candidate')
        for save in (False, True):
            if save:
                self.client.post(f'/encounters/{encounter.pk}/save/', {'revision': encounter.revision,
                    'answer_' + target['key']: selected, 'reviewed_' + target['key']: 'on'})
                encounter.refresh_from_db()
            self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'})
            encounter.refresh_from_db()
            self.assertNotEqual(encounter.status, 'finalized')
            self.assertFalse(encounter.export_path)

    def test_refuted_round_two_opposites_are_held_through_the_worker(self):
        # R2-1..R2-7 with the refuters' inputs: the D31 question and every other unit
        # declared "independent" by the simulated provider.
        label = 'Please confirm that this patient is subject to Multi-Agency Public Protection Arrangements (MAPPA):'
        full = 'The patient is subject to Multi-Agency Public Protection Arrangements (MAPPA).'
        forms = {}
        for key, question in (('d31', label), ('risk', 'Risk to self:'), ('allergy', 'Allergies:')):
            source = ordinary_input_form(Path(self.directory.name) / (key + '-round-two.docx'), [question])
            template, version = self.uploaded(data=source, name='Synthetic round two ' + key)
            self.confirm(template, version)
            forms[key] = (version, next(f for f in version.fields if g.question_key(f['label']) == g.question_key(question)))
        for key, selected, other in (
                ('d31', 'The patient is on MAPPA.', 'The patient is not on MAPPA.'),
                ('d31', full, 'Subject to MAPPA. None.'),
                ('d31', full, 'The patient was not subject to MAPPA as of this week.'),
                ('d31', 'The patient is subject to MAPPA.', 'The patient is not subject to MAPPA arrangements.'),
                ('d31', full, 'MAPPA status: none.'),
                ('risk', 'The patient is not suicidal.', 'She reports suicidal ideation with a plan.'),
                ('allergy', 'The patient is allergic to penicillin.', 'The patient is non allergic to penicillin.')):
            with self.subTest(selected=selected, other=other):
                version, target = forms[key]
                self.assert_worker_holds(version, target, selected, selected + ' ' + other)

    def generate_with(self, version, notes, text):
        """A new encounter generated through the quote, enqueue and the worker, the provider patched to return `text`
        (a keyed answer) exactly once."""
        self.assertEqual(self.new_encounter(version, notes=notes).status_code, 302)
        encounter = Encounter.objects.latest('created_at')
        quote = self.client.get(f'/encounters/{encounter.pk}/').context['usage_quote']
        with patch('anthropic.Anthropic') as provider:
            provider.return_value.messages.create.side_effect = lambda **wire: sdk_response(json.loads(text))
            self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/generate/',
                                              {'usage_quote': quote}).status_code, 302)
            self.assertTrue(run_one())
            self.assertEqual(provider.return_value.messages.create.call_count, 1)
        encounter.refresh_from_db()
        return encounter, Job.objects.get(encounter=encounter)

    def test_a_reported_contradiction_holds_its_fields_through_the_worker_review_finalize_and_export(self):
        # Live attempt 020 (N13): the model returned both fields with conflict true and their selected statements
        # marked contrary. The job failed as a platform failure; it must succeed with each field HELD, the review
        # save (no clinical edit: held fields left blank, reviewed) and finalization must complete, and the export
        # must fill nothing from the contradicted statements.
        from .test_v3_field_relations_contract import ATTEMPT_020_TEXT, frozen_case
        notes, frozen_fields = frozen_case('N13')
        source = ordinary_input_form(Path(self.directory.name) / 'n13-shape.docx',
                                     [f['label'] for f in frozen_fields])
        template, version = self.uploaded(data=source, name='Synthetic N13 shape'); self.confirm(template, version)
        keys = [f['key'] for f in version.fields]
        self.assertEqual([g.question_key(f['label']) for f in version.fields],
                         [g.question_key(f['label']) for f in frozen_fields])
        text = ATTEMPT_020_TEXT.replace('"v3_1"', json.dumps(keys[0])).replace('"v3_2"', json.dumps(keys[1]))
        asked, context, _ = g.v3_generation_fields(version.fields)
        self.assertEqual([u['quote'] for u in g.source_units(notes, asked, context_labels=context)[0]],
                         ['The patient is suicidal.', 'Suicidal ideation: none.', 'The patient is on methadone.',
                          'The patient is not on diazepam or methadone.'])
        encounter, job = self.generate_with(version, notes, text)
        self.assertEqual(job.state, 'succeeded', job.error)                    # (before: failed, 'disagree')
        self.assertNotEqual(Reservation.objects.get(job=job).state, 'held')
        self.assertEqual({k: (v['status'], v['text']) for k, v in encounter.answers.items()},
                         {keys[0]: ('conflict', ''), keys[1]: ('conflict', '')})
        self.assertTrue(all(v.get('ordinary_relation') for v in encounter.answers.values()))
        page = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertEqual(page.status_code, 200)
        self.assertEqual({row['key']: row['status'] for row in page.context['answer_rows']},
                         {keys[0]: 'conflict', keys[1]: 'conflict'})
        payload = {'revision': encounter.revision}
        for key in keys:
            payload.update({'answer_' + key: '', 'reviewed_' + key: 'on', 'blank_' + key: 'on'})
        for requirement in version.review_requirements:
            payload['requirement_' + requirement['key']] = 'on'
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/save/', payload).status_code, 302)
        encounter.refresh_from_db()
        self.assertEqual({k: v.get('text', '') for k, v in encounter.answers.items()}, {keys[0]: '', keys[1]: ''})
        self.assertNotIn('supported', {v['status'] for v in encounter.answers.values()})
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/finalize/',
            {'revision': encounter.revision, 'reviewed': 'on'}).status_code, 302)
        encounter.refresh_from_db(); self.assertEqual(encounter.status, 'finalized')
        response = self.client.get(f'/encounters/{encounter.pk}/export/')
        self.assertEqual(response.status_code, 200)
        exported = b''.join(response.streaming_content)
        slots = w.SlotIndex(w.open_package(exported))
        for f in version.fields:
            written = w.text_of(slots.by_id[f['locator']['part'] + '#' + f['locator']['slot']].el)
            for statement in ('suicidal.', 'none.', 'methadone'):
                self.assertNotIn(statement, written)

    def test_a_selection_the_model_calls_independent_is_held_and_its_record_refused_on_every_reader(self):
        # The same disagreement with conflict FALSE: nothing but the record's disagreement holds the field. A stored
        # `supported` copy that keeps that record is held on the page and refused by a stale save and finalization.
        source = ordinary_input_form(Path(self.directory.name) / 'disagreement.docx', ['Allergies', 'Activity'])
        template, version = self.uploaded(data=source, name='Synthetic disagreement'); self.confirm(template, version)
        allergies, activity = (next(f['key'] for f in version.fields if g.question_key(f['label']) == g.question_key(name))
                               for name in ('Allergies', 'Activity'))
        first, second = 'The patient is not allergic to penicillin.', 'The patient walks every day.'
        notes = first + '\n\n' + second
        body = relation_wire(notes, version.fields, {allergies: [first], activity: [second]},
                             overrides={(allergies, first): 'independent'})
        encounter, job = self.generate_with(version, notes, json.dumps(body))
        self.assertEqual(job.state, 'succeeded', job.error)
        self.assertEqual((encounter.answers[allergies]['status'], encounter.answers[allergies]['text']), ('conflict', ''))
        self.assertEqual((encounter.answers[activity]['status'], encounter.answers[activity]['text']), ('supported', second))
        held = encounter.answers[allergies]
        forged = {**{k: v for k, v in held.items() if k != 'held_candidate'}, 'status': 'supported', 'text': first,
                  'reviewed': True, 'evidence': [{'quote': first, 'start': 0, 'end': len(first)}]}
        encounter.answers = {**encounter.answers, allergies: forged}
        encounter.save()
        self.assertContains(self.client.get(f'/encounters/{encounter.pk}/'), 'data-held-candidate')
        for save in (False, True):
            if save:
                self.client.post(f'/encounters/{encounter.pk}/save/', {'revision': encounter.revision,
                    'answer_' + allergies: first, 'reviewed_' + allergies: 'on',
                    'answer_' + activity: second, 'reviewed_' + activity: 'on'})
                encounter.refresh_from_db()
            self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'})
            encounter.refresh_from_db()
            self.assertNotEqual(encounter.status, 'finalized')
            self.assertFalse(encounter.export_path)

    def test_forced_historical_wrong_side_cannot_pass_review_or_stale_save_or_finalize(self):
        source = ordinary_input_form(Path(self.directory.name) / 'wrong-side.docx', ['Right knee'])
        template, version = self.uploaded(data=source, name='Historical side guard'); self.confirm(template, version)
        notes = 'The patient has left knee pain.'
        self.new_encounter(version, notes=notes); encounter = Encounter.objects.latest('created_at')
        f = version.fields[0]
        encounter.answers = {f['key']: {'status': 'supported', 'text': notes, 'reviewed': True,
            'evidence': [{'quote': notes, 'start': 0, 'end': len(notes)}]}}
        encounter.save()
        page = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertContains(page, 'data-held-candidate')
        for save in (False, True):
            if save:
                self.client.post(f'/encounters/{encounter.pk}/save/', {'revision': encounter.revision,
                    'answer_' + f['key']: notes, 'reviewed_' + f['key']: 'on'})
                encounter.refresh_from_db()
            self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'})
            encounter.refresh_from_db(); self.assertNotEqual(encounter.status, 'finalized')
            self.assertFalse(encounter.export_path)
