"""'Leave blank — reviewed': an explicit clinician decision, distinct from an untouched empty field,
a missing answer and an accepted suggestion. Fixture generation only (AI_MODE=fixture)."""
from io import BytesIO
from zipfile import ZipFile

from django.utils import timezone
from docx import Document

from core.jobs import run_one
from core.models import AuditEvent, Encounter, Membership, User
from core.tests.test_workflow import DocumentWorkflow


def exported_text(data):
    xml = ZipFile(BytesIO(data)).read('word/document.xml').decode()
    return xml


class LeaveBlankReviewed(DocumentWorkflow):
    # This pointer genuinely leaves source scope unresolved under current v7;
    # a complete harmless synthetic disclaimer no longer does.
    NOTES = ('See the findings listed in the report.\n'
             'Symptoms: no chest pain\n'
             'Examination: normal gait')

    def generated(self, data=None):
        _template, mapping = self.upload(data or self.make_docx(), 'docx')
        encounter = self.create_encounter(mapping, self.NOTES)
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/generate/').status_code, 302)
        self.assertTrue(run_one())
        encounter.refresh_from_db()
        self.keys = {f['label'].rstrip(':'): f['key'] for f in mapping.fields}
        return encounter

    def save(self, encounter, payload):
        encounter.refresh_from_db()
        response = self.client.post(f'/encounters/{encounter.pk}/save/', {'revision': encounter.revision, **payload})
        encounter.refresh_from_db()
        return response

    def finalize_page(self, encounter):
        encounter.refresh_from_db()
        response = self.client.post(f'/encounters/{encounter.pk}/finalize/', {'reviewed': 'on', 'revision': encounter.revision}, follow=True)
        encounter.refresh_from_db()
        return response.content.decode()

    def reviewed_all(self):
        return {'reviewed_' + key: 'on' for key in self.keys.values()}

    def test_an_untouched_conflict_still_blocks_finalization(self):
        encounter = self.generated()
        self.assertEqual({encounter.answers[k]['status'] for k in self.keys.values()}, {'conflict'})
        self.assertEqual(self.save(encounter, self.reviewed_all()).status_code, 302)
        page = self.finalize_page(encounter)
        self.assertEqual(encounter.status, 'review')
        self.assertIn('Still in conflict: Examination, Symptoms.', page)
        self.assertIn('Leave blank — reviewed', page)

    def test_an_explicit_blank_resolves_a_conflict_keeps_its_history_and_exports_blank(self):
        encounter = self.generated()
        s, x = self.keys['Symptoms'], self.keys['Examination']
        before = encounter.answers[s]
        self.save(encounter, {**self.reviewed_all(), 'blank_' + s: 'on', 'blank_shown_' + s: '1',
                              'answer_' + x: 'normal gait on examination'})
        answer = encounter.answers[s]
        self.assertEqual(answer['status'], 'left_blank')
        self.assertEqual(answer['text'], '')
        self.assertIs(answer['reviewed'], True)
        self.assertFalse(answer['accepted_suggestion'])
        self.assertEqual(answer['blank_decision']['by'], str(self.user.pk))
        self.assertEqual(answer['blank_decision']['revision'], encounter.revision)
        self.assertEqual(answer['blank_decision']['from_status'], 'conflict')
        # the original conflict, candidate and evidence are retained, not rewritten
        self.assertEqual(answer['suggestion'], before['suggestion'])
        self.assertEqual(answer['evidence'], before['evidence'])
        self.assertEqual(answer['suggestion_evidence'], before['suggestion_evidence'])
        self.assertEqual([h['decision'] for h in answer['blank_history']], ['left_blank'])
        event = AuditEvent.objects.filter(action='encounter.answers_saved', target=str(encounter.pk)).latest('created_at')
        self.assertEqual(event.details['left_blank'], [s])
        self.finalize_page(encounter)
        self.assertEqual(encounter.status, 'finalized')
        exported = b''.join(self.client.get(f'/encounters/{encounter.pk}/export/').streaming_content)
        text = exported_text(exported)
        self.assertNotIn('no chest pain', text, 'an unaccepted candidate must not reach the export')
        self.assertIn('normal gait on examination', text)
        page = self.client.get(f'/encounters/{encounter.pk}/').content.decode()
        self.assertIn('Left blank after clinician review by ' + self.user.email, page)

    def test_reopen_and_ordinary_saves_carry_the_decision_and_withdrawal_is_recorded(self):
        encounter = self.generated()
        s, x = self.keys['Symptoms'], self.keys['Examination']
        self.save(encounter, {**self.reviewed_all(), 'blank_' + s: 'on', 'blank_shown_' + s: '1', 'accept_' + x: 'on'})
        decision = encounter.answers[s]['blank_decision']
        page = self.client.get(f'/encounters/{encounter.pk}/').content.decode()
        self.assertIn(f'id="blank-{s}" name="blank_{s}" type="checkbox" checked', page)
        self.assertIn('left blank — reviewed', page)
        # A save that does not carry the control (an employee's disabled box, an API client) keeps it.
        self.save(encounter, self.reviewed_all())
        self.assertEqual(encounter.answers[s]['status'], 'left_blank')
        self.assertEqual(encounter.answers[s]['blank_decision'], decision, 'the decision is not re-attributed')
        # Unticking the enabled control withdraws it, back to the status it resolved.
        self.save(encounter, {**self.reviewed_all(), 'blank_shown_' + s: '1'})
        self.assertEqual(encounter.answers[s]['status'], 'conflict')
        self.assertNotIn('blank_decision', encounter.answers[s])
        self.assertEqual([h['decision'] for h in encounter.answers[s]['blank_history']], ['left_blank', 'withdrawn'])
        # Deciding again, then typing an answer later, withdraws it too; history keeps every step.
        self.save(encounter, {**self.reviewed_all(), 'blank_' + s: 'on', 'blank_shown_' + s: '1'})
        self.save(encounter, {**self.reviewed_all(), 'answer_' + s: 'reports palpitations', 'blank_shown_' + s: '1'})
        self.assertEqual(encounter.answers[s]['status'], 'edited')
        self.assertEqual(encounter.answers[s]['text'], 'reports palpitations')
        self.assertEqual([h['decision'] for h in encounter.answers[s]['blank_history']],
                         ['left_blank', 'withdrawn', 'left_blank', 'withdrawn'])

    def test_a_blank_on_an_untouched_missing_field_is_distinct_from_missing(self):
        data = self.make_docx_rows(['Symptoms', 'Allergies'])
        encounter = self.generated(data)
        a = self.keys['Allergies']
        self.assertEqual(encounter.answers[a]['status'], 'missing')
        self.save(encounter, {**self.reviewed_all(), 'accept_' + self.keys['Symptoms']: 'on', 'blank_' + a: 'on', 'blank_shown_' + a: '1'})
        self.assertEqual(encounter.answers[a]['status'], 'left_blank')
        self.assertEqual(encounter.answers[a]['blank_decision']['from_status'], 'missing')

    def test_contradictory_decisions_are_refused_and_nothing_is_saved(self):
        encounter = self.generated()
        s = self.keys['Symptoms']
        before = (encounter.answers, encounter.revision)
        for extra in ({'answer_' + s: 'typed text'}, {'accept_' + s: 'on'}):
            response = self.save(encounter, {**self.reviewed_all(), 'blank_' + s: 'on', **extra})
            self.assertEqual(response.status_code, 400)
            self.assertIn('Choose one decision', response.content.decode())
            self.assertEqual((encounter.answers, encounter.revision), before)

    def test_a_signature_field_never_offers_or_accepts_the_blank_decision(self):
        encounter = self.generated(self.make_docx_rows(['Symptoms', 'Electronically signed by']))
        sig = self.keys['Electronically signed by']
        page = self.client.get(f'/encounters/{encounter.pk}/').content.decode()
        self.assertNotIn(f'name="blank_{sig}"', page)
        self.assertIn(f'name="blank_{self.keys["Symptoms"]}"', page)
        before = encounter.answers
        response = self.save(encounter, {**self.reviewed_all(), 'blank_' + sig: 'on'})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(encounter.answers, before)

    def test_only_an_owner_or_clinician_may_decide_or_withdraw(self):
        encounter = self.generated()
        s = self.keys['Symptoms']
        employee = User.objects.create_user(username='employee', email='employee@synthetic.invalid', verified_at=timezone.now())
        Membership.objects.create(user=employee, department=self.department, role='employee')
        self.client.force_login(employee)
        before = encounter.answers
        self.assertEqual(self.save(encounter, {'blank_' + s: 'on'}).status_code, 403)
        self.assertEqual(encounter.answers, before)
        self.client.force_login(self.user)
        self.save(encounter, {**self.reviewed_all(), 'blank_' + s: 'on', 'blank_shown_' + s: '1'})
        self.client.force_login(employee)
        self.assertEqual(self.save(encounter, {'blank_shown_' + s: '1'}).status_code, 403)
        self.assertEqual(encounter.answers[s]['status'], 'left_blank')

    def test_stale_edits_and_finalized_documents_are_still_refused(self):
        encounter = self.generated()
        s = self.keys['Symptoms']
        stale = encounter.revision
        self.save(encounter, self.reviewed_all())
        response = self.client.post(f'/encounters/{encounter.pk}/save/', {'revision': stale, 'blank_' + s: 'on'})
        self.assertEqual(response.status_code, 409)
        self.save(encounter, {**self.reviewed_all(), 'blank_' + s: 'on', 'blank_shown_' + s: '1',
                              'blank_' + self.keys['Examination']: 'on'})
        self.finalize_page(encounter)
        self.assertEqual(encounter.status, 'finalized')
        response = self.client.post(f'/encounters/{encounter.pk}/save/', {'revision': encounter.revision, 'blank_shown_' + s: '1'})
        self.assertEqual(response.status_code, 409)
        encounter.refresh_from_db()
        self.assertEqual(encounter.answers[s]['status'], 'left_blank')

    def test_withdrawing_a_blank_on_a_cleared_supported_answer_reads_as_an_edit(self):
        _template, mapping = self.upload(self.make_docx(), 'docx')
        encounter = self.create_encounter(mapping, 'Symptoms: No chest pain.\nExamination: Normal gait.')
        self.client.post(f'/encounters/{encounter.pk}/generate/')
        self.assertTrue(run_one())
        encounter.refresh_from_db()
        self.keys = {f['label'].rstrip(':'): f['key'] for f in mapping.fields}
        s = self.keys['Symptoms']
        self.assertEqual(encounter.answers[s]['status'], 'supported')
        self.save(encounter, {**self.reviewed_all(), 'answer_' + self.keys['Examination']: encounter.answers[self.keys['Examination']]['text'],
                              'blank_' + s: 'on', 'blank_shown_' + s: '1'})
        self.assertEqual(encounter.answers[s]['blank_decision']['from_status'], 'supported')
        self.save(encounter, {**self.reviewed_all(), 'answer_' + self.keys['Examination']: encounter.answers[self.keys['Examination']]['text'],
                              'blank_shown_' + s: '1'})
        self.assertEqual(encounter.answers[s]['text'], '')
        self.assertEqual(encounter.answers[s]['status'], 'edited', 'an emptied answer must never read as supported')

    def test_the_rendered_form_withdraws_by_typing_or_accepting(self):
        encounter = self.generated()
        s, x = self.keys['Symptoms'], self.keys['Examination']
        self.save(encounter, {**self.reviewed_all(), 'blank_' + s: 'on', 'blank_shown_' + s: '1',
                              'blank_' + x: 'on', 'blank_shown_' + x: '1'})
        # Exactly what the page posts: the box still ticked as rendered, plus the new decision.
        response = self.save(encounter, {**self.reviewed_all(), 'blank_' + s: 'on', 'blank_shown_' + s: '1',
                                         'answer_' + s: 'reports palpitations',
                                         'blank_' + x: 'on', 'blank_shown_' + x: '1', 'accept_' + x: 'on'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual((encounter.answers[s]['status'], encounter.answers[s]['text']), ('edited', 'reports palpitations'))
        self.assertEqual(encounter.answers[x]['text'], 'normal gait')
        self.assertTrue(encounter.answers[x]['accepted_suggestion'])
        for key in (s, x):
            self.assertEqual([h['decision'] for h in encounter.answers[key]['blank_history']], ['left_blank', 'withdrawn'])

    def test_whitespace_is_not_an_answer_and_cannot_erase_or_resolve_anything(self):
        encounter = self.generated()
        s, x = self.keys['Symptoms'], self.keys['Examination']
        # A space does not clear a generated conflict.
        self.save(encounter, {**self.reviewed_all(), 'answer_' + x: '   '})
        self.assertEqual((encounter.answers[x]['status'], encounter.answers[x]['text']), ('conflict', ''))
        self.save(encounter, {**self.reviewed_all(), 'blank_' + s: 'on', 'blank_shown_' + s: '1'})
        decision = encounter.answers[s]['blank_decision']
        # An employee's space neither erases the decision nor slips past the authority check.
        employee = User.objects.create_user(username='space-employee', email='space@synthetic.invalid', verified_at=timezone.now())
        Membership.objects.create(user=employee, department=self.department, role='employee')
        self.client.force_login(employee)
        self.assertEqual(self.save(encounter, {'answer_' + s: '   '}).status_code, 302)
        self.assertEqual(encounter.answers[s]['status'], 'left_blank')
        self.assertEqual(encounter.answers[s]['blank_decision'], decision)
        self.assertEqual([h['decision'] for h in encounter.answers[s]['blank_history']], ['left_blank'])

    def make_docx_rows(self, labels):
        document = Document()
        document.add_paragraph('Fictional software test form.')
        table = document.add_table(rows=0, cols=2)
        for label in labels:
            cells = table.add_row().cells
            cells[0].text = label
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()
