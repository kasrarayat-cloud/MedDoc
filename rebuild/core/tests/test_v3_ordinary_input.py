"""Bounded retained ordinary-input repairs; mock source selections, zero provider calls."""
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from django.test import SimpleTestCase, TestCase, override_settings

from core import generation as g
from .test_v3_ordinary_notes import FIELDS, LIVE, field, run
from .test_v3_ordinary_context import (CompleteOrdinaryNoteWorkflow, medication_review_form,
                                     retained_note)
from .test_word_v3_workflow import V3EncounterFixture, part_text
from core.security import private_path


class InlinePrintedOptions(SimpleTestCase):
    def test_actual_inline_option_uses_the_same_source_sentence_boundaries(self):
        quote = 'Smoker: No.'
        for head in ('Patient presents with right knee pain.', '45-year-old man presents with knee pain.'):
            notes = head + ' ' + quote
            answer = run(notes, {'Smoker': [quote]})['Smoker']
            self.assertEqual((answer['status'], answer['text']), ('supported', 'No'))
            governance = g.governance_for(notes, FIELDS)
            self.assertIsNone(g.governed_supported(governance, notes, answer, 'Smoker'))
            self.assertEqual(answer['evidence'], [{'quote': quote, 'start': len(head) + 1, 'end': len(notes)}])

    def test_inline_options_keep_uncertainty_wrong_destination_and_governing_context(self):
        for notes, quote in (
            ('Patient presents with knee pain. Smoker: No?', 'Smoker: No?'),
            ('Patient presents with knee pain. Smoker: No...', 'Smoker: No...'),
            ('Patient presents with knee pain. Allergies: No.', 'Allergies: No.'),
            ('For the patient\'s mother only:\nPatient presents with knee pain. Smoker: No.', 'Smoker: No.'),
            ('Patient presents with knee pain. Smoker: No. Actually, yes.', 'Smoker: No.'),
            ('Patient presents with knee pain. Smoker: No. The patient smokes daily.', 'Smoker: No.'),
            ('Patient presents with knee pain. Smoker: No. He says that answer is false.', 'Smoker: No.'),
        ):
            with self.subTest(notes=notes):
                answer = run(notes, {'Smoker': [quote]})['Smoker']
                self.assertEqual((answer['status'], answer['text']), ('conflict', ''))
        fields = [field('s1', 'Smoker', 'choice', ['Yes', 'No']), field('s2', 'Smoker?', 'choice', ['Yes', 'No'])]
        notes = 'Patient presents with knee pain. Smoker: No.'
        self.assertTrue(all(a['status'] == 'conflict' for a in run(notes, {'Smoker': ['Smoker: No.'], 'Smoker?': ['Smoker: No.']}, fields=fields).values()))

    def test_locator_rejects_a_partial_value_or_offset_not_in_the_original_source(self):
        labels = g.recognized_labels(FIELDS)
        notes = 'Patient presents with knee pain. Smoker: No.'
        start = notes.index('Smoker:')
        for entry in ({'quote': 'No', 'start': notes.index('No')},
                      {'quote': 'Smoker: No.', 'start': start - 1},
                      {'quote': 'No.', 'start': notes.index('No') - 1}):
            self.assertFalse(g.v3_labelled_source(notes, labels, 'smoker', [entry]))


@override_settings(**LIVE)
class CompleteInlineMedicationReview(V3EncounterFixture, TestCase):
    walk = CompleteOrdinaryNoteWorkflow.walk

    def setUp(self):
        with patch('core.tests.test_word_v3_workflow.build_intake_form', side_effect=medication_review_form):
            super().setUp()

    def test_retained_note9_as_one_paragraph_completes_original_word_export(self):
        case = retained_note('r1-note9')
        # Layout variant only; the unchanged original remains separately measured.
        notes = case['notes'].replace('\n', ' ')
        expected = {'Chief complaint': 'Patient presents for medication review.',
                    'Medications': 'Current medications are atorvastatin 20 mg and metformin 1000 mg.',
                    'Allergies': 'codeine.', 'Smoker': 'No', 'Examination': 'not performed today.',
                    'Assessment': 'stable type 2 diabetes and hyperlipidemia.', 'Plan': 'continue current medications.'}
        self.walk(notes, case['picks'], expected)
        body = part_text(private_path(self.encounter.export_path).read_bytes())
        self.assertIn('Smoker  ☐ Yes  ☒ No', body)


class AncillarySmokerContext(SimpleTestCase):
    def test_retained_visit_and_social_tails_keep_exact_smoker_answer(self):
        for quote, tail, expected in (('Smoker: yes', 'Patient seen today.', 'Yes'),
                                      ('Smoker: No.', 'The patient drinks socially.', 'No')):
            notes = quote + '\n' + tail
            answer = run(notes, {'Smoker': [quote]})['Smoker']
            self.assertEqual((answer['status'], answer['text']), ('supported', expected))
            reading = g.governance_for(notes, FIELDS)
            self.assertEqual(g.governed_decisions(notes, FIELDS, {'v3_4': {**answer, 'reviewed': True}}), [])
            relation = reading[11]['preceding']['ancillary_relations']['smoker']
            self.assertEqual(relation, {1: {2}})
            self.assertEqual([unit['quote'] for unit in reading[11]['preceding']['units']], [quote, tail])

    def test_ancillary_relation_never_resets_an_unknown_or_contradicted_question(self):
        for label, tail in (('Seen recently?', 'Patient seen today.'),
                            ('Alcohol use', 'The patient drinks socially.'),
                            ('Medications', 'Patient seen today.'),
                            ('Consent obtained', 'Patient seen today.')):
            fields = [field('a', label, 'choice', ['Yes', 'No'])]
            quote = label.rstrip('?') + ': No.'
            notes = quote + '\n' + tail
            answer = run(notes, {label: [quote]}, fields=fields)[label]
            self.assertEqual(answer['status'], 'conflict', label)
            forced = {'status': 'supported', 'text': 'No', 'reviewed': True,
                      'evidence': [{'quote': quote, 'start': 0, 'end': len(quote)}]}
            self.assertEqual(g.governed_decisions(notes, fields, {'a': forced}), [label])

    def test_no_extra_qualifier_correction_negation_or_family_scope_is_discarded(self):
        quote = 'Smoker: No.'
        for before, tail in (
            ('', 'The patient smokes daily.'),
            ('', 'The patient drinks socially and smokes daily.'),
            ('', 'The patient drinks socially, but the earlier answer is wrong.'),
            ('', 'Patient seen today for smoking cessation.'),
            ('', 'Patient not seen today.'),
            ('', 'Patient seen today. Actually, smoker yes.'),
            ('', 'The patient drinks socially.\nShe smokes daily.'),
            ('For the patient\'s mother only:\n', 'The patient drinks socially.'),
            ('If the patient attends:\n', 'Patient seen today.'),
        ):
            notes = before + quote + '\n' + tail
            answer = run(notes, {'Smoker': [quote]})['Smoker']
            self.assertEqual((answer['status'], answer['text']), ('conflict', ''), notes)
            start = len(before)
            forced = {'status': 'supported', 'text': 'No', 'reviewed': True,
                      'evidence': [{'quote': quote, 'start': start, 'end': start + len(quote)}]}
            self.assertIn('Smoker', g.governed_decisions(notes, FIELDS, {'v3_4': forced}), notes)


class AgeSexPresentation(SimpleTestCase):
    fields = [field('r', 'Referring physician'), field('h', 'Chief complaint'),
              field('s', 'Smoker', 'choice', ['Yes', 'No'])]

    def test_retained_explicit_presentations_keep_full_age_and_time_wording(self):
        for case_name in ('r1-note1', 'r1-note4'):
            case = retained_note(case_name)
            quote = case['picks']['Chief complaint'][0]
            notes = 'Referring physician: Dr Example\n' + quote + '\nSmoker: No.'
            answers = run(notes, {'Chief complaint': [quote]}, fields=self.fields)
            answer = answers['Chief complaint']
            self.assertEqual((answer['status'], answer['text']), ('supported', quote), notes)
            self.assertEqual(answer['evidence'], [{'quote': quote, 'start': notes.index(quote), 'end': notes.index(quote) + len(quote)}])
            self.assertEqual(g.governed_decisions(notes, self.fields, {'h': {**answer, 'reviewed': True}}), [])

    def test_age_does_not_reset_family_conditional_correction_or_destination(self):
        quote = retained_note('r1-note1')['picks']['Chief complaint'][0]
        for prefix, suffix in (('For the patient\'s father only:\n', ''),
                               ('If the patient attends:\n', ''),
                               ('Her sister was seen yesterday.\n', ''),
                               ('', '\nActually, the pain is in the left knee.'),
                               ('', '\nThe patient denies that symptom.')):
            notes = prefix + 'Referring physician: Dr Example\n' + quote + suffix
            answer = run(notes, {'Chief complaint': [quote]}, fields=self.fields)['Chief complaint']
            self.assertEqual((answer['status'], answer['text']), ('conflict', ''), notes)
        notes = 'Referring physician: Dr Example\n' + quote
        for label in ('Allergies', 'Lt knee', 'Plan', 'Consent obtained'):
            fields = [field('r', 'Referring physician'), field('x', label, 'long_text')]
            self.assertEqual(run(notes, {label: [quote]}, fields=fields)[label]['status'], 'conflict', label)

    def test_age_subject_does_not_discard_an_unselected_pronoun_continuation(self):
        quote = retained_note('r1-note1')['picks']['Chief complaint'][0]
        notes = 'Referring physician: Dr Example\n' + quote + '\nHe reports the opposite side.'
        answer = run(notes, {'Chief complaint': [quote]}, fields=self.fields)['Chief complaint']
        self.assertEqual((answer['status'], answer['text']), ('conflict', ''))


def ordinary_input_form(path, labels):
    from docx import Document
    doc = Document()
    doc.add_paragraph('Synthetic ordinary input form')
    text_labels = [label for label in labels if label != 'Smoker']
    if text_labels:
        table = doc.add_table(rows=len(text_labels), cols=2)
        for index, label in enumerate(text_labels):
            table.cell(index, 0).text = label
    if 'Smoker' in labels:
        doc.add_paragraph('Smoker  ☐ Yes  ☐ No')
    out = BytesIO(); doc.save(out)
    Path(path).write_bytes(out.getvalue())
    return out.getvalue()


@override_settings(**LIVE)
class CompleteAncillarySmokerWorkflow(V3EncounterFixture, TestCase):
    walk = CompleteOrdinaryNoteWorkflow.walk

    def setUp(self):
        with patch('core.tests.test_word_v3_workflow.build_intake_form',
                   side_effect=lambda path: ordinary_input_form(path, ['Smoker'])):
            super().setUp()

    def test_retained_social_tail_completes_original_checkbox_export(self):
        self.walk('Smoker: No.\nThe patient drinks socially.', {'Smoker': ['Smoker: No.']}, {'Smoker': 'No'})
        self.assertIn('Smoker  ☐ Yes  ☒ No', part_text(private_path(self.encounter.export_path).read_bytes()))

    def test_retained_visit_tail_completes_original_checkbox_export(self):
        self.walk('Smoker: yes\nPatient seen today.', {'Smoker': ['Smoker: yes']}, {'Smoker': 'Yes'})
        self.assertIn('Smoker  ☒ Yes  ☐ No', part_text(private_path(self.encounter.export_path).read_bytes()))


@override_settings(**LIVE)
class CompleteAgePresentationWorkflow(V3EncounterFixture, TestCase):
    walk = CompleteOrdinaryNoteWorkflow.walk

    def setUp(self):
        with patch('core.tests.test_word_v3_workflow.build_intake_form',
                   side_effect=lambda path: ordinary_input_form(path, ['Referring physician', 'Chief complaint', 'Smoker'])):
            super().setUp()

    def test_retained_age_sentence_keeps_its_time_qualification_through_export(self):
        quote = retained_note('r1-note1')['picks']['Chief complaint'][0]
        notes = 'Referring physician: Dr Example\n' + quote + '\nSmoker: No.'
        self.walk(notes, {'Referring physician': ['Referring physician: Dr Example'],
                          'Chief complaint': [quote], 'Smoker': ['Smoker: No.']},
                  {'Referring physician': 'Dr Example', 'Chief complaint': quote, 'Smoker': 'No'})


class InlineOwnershipRefutation(SimpleTestCase):
    def test_existing_supported_inline_option_under_unknown_context_is_governed(self):
        for prefix in ('Father: MI at 50.', 'Family hx: father had an MI.', 'Seen today.'):
            label = 'Consent obtained' if prefix == 'Seen today.' else 'Smoker'
            key = 'v3_5' if label == 'Consent obtained' else 'v3_4'
            quote = label + ': yes.'
            notes = prefix + ' ' + quote
            answer = run(notes, {label: [quote]})[label]
            self.assertEqual((answer['status'], answer['text']), ('conflict', ''))
            stored = {'status': 'supported', 'text': 'Yes', 'reviewed': True,
                      'evidence': [{'quote': quote, 'start': len(prefix) + 1, 'end': len(notes)}]}
            governance = g.governance_for(notes, FIELDS)
            shown = g.as_currently_governed(governance, notes, {**stored, 'label': label})
            self.assertEqual((shown['status'], shown['text']), ('conflict', ''))
            self.assertIn(label, g.governed_decisions(notes, FIELDS, {key: stored}))

    def test_physical_line_boundaries_match_source_units_without_changing_the_older_separate_line_contract(self):
        for separator in ('\n', '\r\n', '\r', '\u2028'):
            notes = 'Father: MI at 50.' + separator + 'Smoker: yes.'
            answer = run(notes, {'Smoker': ['Smoker: yes.']})['Smoker']
            self.assertEqual((answer['status'], answer['text']), ('supported', 'Yes'))
            reading = g.governance_for(notes, FIELDS)[11]['preceding']
            self.assertEqual(reading['inline_options'], set())


class CompleteValueEvidenceContext(SimpleTestCase):
    fields = [*FIELDS, field('recent', 'Seen recently', 'choice', ['Yes', 'No'])]

    def test_complete_value_keeps_the_same_context_as_its_labelled_source_unit(self):
        for notes, label, held in (
            ('Father: MI at 50. Smoker: No.', 'Smoker', True),
            ('Family hx: father had an MI. Smoker: No.', 'Smoker', True),
            ('Smoker: No.\nThe patient drinks socially but smokes daily.', 'Smoker', True),
            ('Seen recently: No.\nPatient seen today.', 'Seen recently', True),
            ('Patient presents with knee pain. Smoker: No.', 'Smoker', False),
            ('Smoker: No.\nThe patient drinks socially.', 'Smoker', False),
        ):
            for quote in (label + ': No.', 'No.'):
                with self.subTest(notes=notes, quote=quote):
                    start = notes.index(quote)
                    evidence = [{'quote': quote, 'start': start, 'end': start + len(quote)}]
                    self.assertTrue(g.v3_labelled_source(notes, g.recognized_labels(self.fields),
                                                        g.question_key(label), evidence))
                    answer = {'status': 'supported', 'text': 'No', 'reviewed': True, 'evidence': evidence}
                    governance = g.governance_for(notes, self.fields)
                    page = g.as_currently_governed(governance, notes, {**answer, 'label': label})
                    self.assertEqual((page['status'], page['text']), ('conflict', '') if held else ('supported', 'No'))
                    key = next(f['key'] for f in self.fields if f['label'] == label)
                    self.assertEqual(g.governed_decisions(notes, self.fields, {key: answer}), [label] if held else [])
                    offer = {'text': 'No', 'status': 'edited', 'suggestion': 'No',
                             'suggestion_evidence': evidence, 'accepted_suggestion': True, 'reviewed': True}
                    self.assertEqual(g.governed_decisions(notes, self.fields, {key: offer}), [label] if held else [])

    def test_binding_requires_the_complete_quote_and_exact_offsets(self):
        notes = 'Father: MI at 50. Smoker: No.'
        reading = g.governance_for(notes, self.fields)[11]['preceding']
        start = notes.index('No.')
        good = {'quote': 'No.', 'start': start, 'end': start + 3}
        bound = g.v3_evidence_units(reading, [good])
        self.assertEqual([u['quote'] for u in bound.values()], ['Smoker: No.'])
        for bad in ({**good, 'quote': 'No', 'end': start + 2}, {**good, 'start': start - 1},
                    {**good, 'end': start + 2}, {**good, 'quote': 'Yes'}):
            self.assertEqual(g.v3_evidence_units(reading, [bad]), {})
        self.assertFalse(g.v3_labelled_source(notes, g.recognized_labels(self.fields), 'seen recently', [good]))
        self.assertIn('Seen recently', g.governed_decisions(notes, self.fields, {'recent': {
            'status': 'supported', 'text': 'No', 'reviewed': True, 'evidence': [good]}}))


@override_settings(**LIVE)
class CompleteValueEvidenceWorkflow(V3EncounterFixture, TestCase):
    def setUp(self):
        with patch('core.tests.test_word_v3_workflow.build_intake_form',
                   side_effect=lambda path: ordinary_input_form(path, ['Smoker'])):
            super().setUp()

    def stored_value(self, notes):
        start = notes.index('No.')
        self.encounter.notes = notes
        self.encounter.answers = {self.key('Smoker'): {'status': 'supported', 'text': 'No', 'reviewed': True,
            'evidence': [{'quote': 'No.', 'start': start, 'end': start + 3}]}}
        self.encounter.save(update_fields=['notes', 'answers'])

    def test_old_supported_value_cannot_skip_page_save_or_finalization_context(self):
        self.stored_value('Father: MI at 50. Smoker: No.')
        page = self.client.get(f'/encounters/{self.encounter.pk}/')
        row = page.context['answer_rows'][0]
        self.assertEqual((row['status'], row['text']), ('conflict', ''))
        self.finalize(); self.encounter.refresh_from_db()
        self.assertNotEqual(self.encounter.status, 'finalized')
        self.assertFalse(self.encounter.export_path)
        self.assertEqual(self.save({'Smoker': 'No'}).status_code, 302)
        self.encounter.refresh_from_db()
        # An unchanged stale page is not a new clinician decision. Its stored
        # status stays supported, but the current page and finalization hold it.
        self.assertEqual(self.encounter.answers[self.key('Smoker')]['status'], 'supported')
        self.finalize(); self.encounter.refresh_from_db()
        self.assertNotEqual(self.encounter.status, 'finalized')
        self.assertFalse(self.encounter.export_path)
        self.assertEqual(self.save({'Smoker': ''}, extra={'held_shown_' + self.key('Smoker'): '1'}).status_code, 302)
        self.encounter.refresh_from_db()
        self.assertEqual(self.encounter.answers[self.key('Smoker')]['status'], 'conflict')
        self.finalize(); self.encounter.refresh_from_db()
        self.assertNotEqual(self.encounter.status, 'finalized')
        self.assertFalse(self.encounter.export_path)

    def test_patient_value_evidence_still_saves_finalizes_and_exports_checked_option(self):
        self.stored_value('Patient presents with knee pain. Smoker: No.')
        page = self.client.get(f'/encounters/{self.encounter.pk}/')
        self.assertEqual(page.context['answer_rows'][0]['status'], 'supported')
        self.assertEqual(self.save({'Smoker': 'No'}).status_code, 302)
        self.finalize(); self.encounter.refresh_from_db()
        self.assertEqual(self.encounter.status, 'finalized')
        self.assertIn('Smoker  ☐ Yes  ☒ No', part_text(private_path(self.encounter.export_path).read_bytes()))
