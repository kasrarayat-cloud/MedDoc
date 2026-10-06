"""Persisted doctor recovery of complete source-bound flowing ruled narrative."""
import io
import os
import json
import hashlib
from pathlib import Path

from django.test import TestCase, override_settings
from docx import Document

from core.models import Encounter
from core.jobs import enqueue, run_one
from core.security import private_path
from core.tests.test_word_v3_workflow import V3, V3Fixture
from documents.test_ruled_paragraph_recovery import source
from documents import word_v3 as w


@override_settings(**V3)
class RuledParagraphWorkflow(V3Fixture, TestCase):
    @override_settings(AI_PROVIDER='anthropic', AI_MODEL='claude-sonnet-5-5', AI_THINKING='between_tools',
                       AI_EFFORT='medium', AI_INPUT_RATE='2', AI_OUTPUT_RATE='10', AI_CACHED_RATE='.2',
                       V3_FIELD_RELATIONS=True)
    def test_manual_complete_area_review_generation_correction_finalization_and_faithful_download(self):
        original = source(); template, first = self.uploaded(data=original, name='Synthetic ruled narrative')
        artifact_root = os.environ.get('QUILLMD_WORD_ARTIFACTS')
        artifacts = Path(artifact_root) / ('ruled-narrative-' + str(template.pk)) if artifact_root else None
        if artifacts:
            artifacts.mkdir(parents=True, mode=0o700); (artifacts / 'original.docx').write_bytes(original)
        def retain(phase, version, encounter=None, job=None):
            if artifacts is None: return
            record = {'phase': phase, 'original_sha256': hashlib.sha256(original).hexdigest(),
                      'mapping_version': str(version.pk), 'fields': version.fields,
                      'import_manifest': version.import_manifest, 'corrections': version.corrections}
            if encounter:
                record.update(notes=encounter.notes, answers=encounter.answers, status=encounter.status)
                if encounter.export_path:
                    export = private_path(encounter.export_path).read_bytes()
                    (artifacts / 'export.docx').write_bytes(export); record['export_sha256'] = hashlib.sha256(export).hexdigest()
            if job: record.update(job_state=job.state, job_error=job.error, job_config=job.config)
            (artifacts / (phase + '.json')).write_text(json.dumps(record, indent=2, default=str) + '\n')
        retain('first-import', first)
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_add', 'mapping_version': first.pk,
            'slot': 'word/document.xml#b/p[4]', 'label': 'Impression', 'kind': 'long_text'})
        # The unchanged original now has an automatic full-area owner. Keep the
        # first duplicate refusal, then exercise explicit manual recovery by
        # removing that owner through the actual mapping-correction route.
        self.assertEqual(response.status_code, 400, response.content.decode())
        self.assertContains(response, 'That place already belongs to a field', status_code=400)
        self.assertEqual(self.latest(template).pk, first.pk)
        automatic = next(f for f in first.fields if f['locator'].get('slot') == 'b/p[4]')
        self.assertEqual(automatic.get('origin', 'automatic'), 'automatic')
        self.assertEqual(automatic['locator']['target']['type'], 'ruled_paragraphs')
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_save', 'mapping_version': first.pk,
            'remove_' + automatic['key']: 'on'})
        self.assertEqual(response.status_code, 302, response.content.decode())
        without_automatic = self.latest(template)
        retain('automatic-owner-explicitly-removed', without_automatic)
        self.assertNotIn(automatic['key'], {f['key'] for f in without_automatic.fields})
        self.assertEqual(first.fields, template.versions.get(pk=first.pk).fields)
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_add', 'mapping_version': without_automatic.pk,
            'slot': 'word/document.xml#b/p[4]', 'label': 'Impression', 'kind': 'long_text'})
        self.assertEqual(response.status_code, 302, response.content.decode())
        mapped = self.latest(template); narrative = next(f for f in mapped.fields if f['label'] == 'Impression')
        retain('manual-mapping', mapped)
        self.assertEqual(narrative['origin'], 'manual'); self.assertEqual(len(narrative['locator']['extra']), 13)
        self.assertIn(narrative['key'], mapped.corrections['added'])
        response = self.client.post(f'/templates/{template.pk}/', {'action': 'v3_add', 'mapping_version': mapped.pk,
            'slot': 'word/document.xml#b/p[6]', 'label': 'Partial suffix', 'kind': 'long_text'})
        self.assertEqual(response.status_code, 400); self.assertEqual(self.latest(template).pk, mapped.pk)
        # Fixed bracket guidance may remain an explicit source exclusion. That
        # recorded recovery stays partial; no answer is written into guidance.
        issues = mapped.import_manifest.get('coverage', {}).get('selected_engine_issues', [])
        if mapped.import_manifest.get('coverage', {}).get('block_generation'):
            self.assertTrue(issues)
            post = {'action': 'v3_coverage', 'mapping_version': mapped.pk,
                    **{'exclude_' + issue['id']: 'on' for issue in issues}}
            response = self.client.post(f'/templates/{template.pk}/', post)
            self.assertEqual(response.status_code, 302, response.content.decode()); mapped = self.latest(template)
        self.confirm(template, mapped)
        self.assertEqual(self.new_encounter(mapped, 'The patient is stable. Impression: outpatient review.').status_code, 302)
        encounter = Encounter.objects.latest('created_at'); job = enqueue(encounter, self.user)
        worked = run_one(); job.refresh_from_db(); encounter.refresh_from_db()
        retain('generated', mapped, encounter, job)
        self.assertTrue(worked); self.assertEqual(job.state, 'succeeded', job.error)
        encounter.refresh_from_db()
        answer = '\n'.join('Reviewed synthetic sentence %02d.' % n for n in range(19))
        post = {'revision': encounter.revision}
        for field in mapped.fields:
            key = field['key']; post['reviewed_' + key] = 'on'; post['answer_' + key] = answer if key == narrative['key'] else ''
            if key != narrative['key']: post['blank_' + key] = 'on'
        response = self.client.post(f'/encounters/{encounter.pk}/save/', post)
        self.assertEqual(response.status_code, 302, response.content.decode()); encounter.refresh_from_db()
        retain('reviewed', mapped, encounter, job)
        self.assertEqual(encounter.answers[narrative['key']]['status'], 'edited')
        self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'}, follow=True)
        encounter.refresh_from_db(); retain('finalized', mapped, encounter, job)
        self.assertEqual(encounter.status, 'finalized')
        raw = private_path(encounter.export_path).read_bytes()
        self.assertEqual(w.effective_values(original, [narrative], {narrative['key']: answer})[narrative['key']], answer)
        before, after = Document(io.BytesIO(original)), Document(io.BytesIO(raw))
        self.assertEqual(len(before.paragraphs), len(after.paragraphs))
        for number in (0, 1, 2, 3, 5, 7, 29, 31):
            self.assertEqual(before.paragraphs[number]._p.xml, after.paragraphs[number]._p.xml)
        self.assertEqual(before.sections[0].header.paragraphs[0].text, after.sections[0].header.paragraphs[0].text)
        self.assertEqual(private_path(template.original_path).read_bytes(), original)
        download = self.client.get(f'/encounters/{encounter.pk}/export/')
        self.assertEqual(download.status_code, 200); self.assertEqual(b''.join(download.streaming_content), raw)
