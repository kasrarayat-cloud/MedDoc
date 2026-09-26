"""Real private HTTP/model/Word sandbox workflows with synthetic clinical data.

Requires the separately built qualified SDK runtime. Missing runtime is an
explicit skip; it is never replaced by a mocked document engine.
"""
from copy import deepcopy
from decimal import Decimal
import hashlib
import json
import os
from pathlib import Path
import tempfile
from unittest import skipUnless
from unittest.mock import patch
from zipfile import ZipFile
from io import BytesIO
from html.parser import HTMLParser
import html
import re

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone

from core.billing import welcome_credit
from core.generation import configuration
from core.jobs import enqueue, run_one
from core.models import Department, Encounter, Job, Membership, Organization, Patient, Template, User
from core.security import private_path
from core.wm_review import BACKEND, ORIGINAL_SHA256
from documents import DocumentError, fill_template, resolve_template
from .test_workflow import TEST_SETTINGS
from core.tests.audio_fixtures import wav_bytes


def word_runtime_is_usable():
    """Every condition sdk_service.configuration() needs, so a partly-configured host SKIPS rather than fails.

    This predicate named two of the THREE variables the service requires, and said nothing about
    ownership. With two of three set, these eight tests ran and failed for an environment reason —
    reporting a defect in the document workflow when the only thing wrong was the host. A skip that
    is narrower than the thing it guards is a check that cries wolf; this one now lists exactly what
    is missing, so whoever meets it can fix it or accept it knowingly.
    """
    missing = [name for name in ('QUILLMD_WORD_RUNTIME', 'QUILLMD_WORD_MANIFEST',
                                 'QUILLMD_WORD_MANIFEST_SHA256') if not os.getenv(name)]
    if missing:
        return False, 'not configured here: ' + ', '.join(missing)
    runtime = Path(os.getenv('QUILLMD_WORD_RUNTIME'))
    if not runtime.exists():
        return False, 'the configured runtime does not exist: %s' % runtime
    # The service refuses a processor the running account can modify, which is the point of the
    # separate service identity. Running these as the owner proves nothing about the shipped path.
    try:
        owned_by_us = runtime.stat().st_uid == os.getuid()
    except OSError as error:
        return False, 'the runtime cannot be stat()ed: %s' % error
    if owned_by_us:
        return False, ('this account owns the Word processor, and the service refuses a processor it '
                       'can modify — run these where the runtime belongs to another identity')
    return True, ''


WORD_RUNTIME_USABLE, WORD_RUNTIME_REASON = word_runtime_is_usable()
if os.getenv('QUILLMD_REQUIRE_WORD_TESTS') == '1':
    if not WORD_RUNTIME_USABLE:
        raise RuntimeError('Required Word integration tests cannot run: ' + WORD_RUNTIME_REASON)
    if os.getenv('QUILLMD_WORD_QUALIFICATION') == '1':
        raise RuntimeError('Required Word integration tests must use the ordinary ownership boundary.')


@skipUnless(WORD_RUNTIME_USABLE,
            'Live Word document workflow NOT EXERCISED — ' + (WORD_RUNTIME_REASON or 'unknown reason') +
            '. No claim is made about this path from this run.')
@override_settings(**TEST_SETTINGS)
class WMDocumentWorkflow(TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.state = patch.dict(os.environ, {'QUILLMD_STATE_DIR': self.directory.name})
        self.state.start(); self.addCleanup(self.state.stop)
        self.override = override_settings(MEDIA_ROOT=Path(self.directory.name) / 'private')
        self.override.enable(); self.addCleanup(self.override.disable)
        self.user = User.objects.create_user(username='wm-clinician', email='wm@synthetic.invalid',
                                             first_name='SYNTHETIC', last_name='Actual Clinician',
                                             verified_at=timezone.now())
        self.organization = Organization.objects.create(name='Synthetic WM organization')
        self.department = Department.objects.create(organization=self.organization, name='Synthetic X-ray')
        Membership.objects.create(user=self.user, department=self.department, role='clinician')
        welcome_credit(self.user)
        self.client.force_login(self.user)
        self.patient = Patient.objects.create(department=self.department, display_name='SYNTHETIC WM PATIENT',
                                               reference='WM-SYNTHETIC', created_by=self.user)
        original_path = Path(__file__).resolve().parents[3] / 'WM IOV Template.docx'
        self.original = original_path.read_bytes()
        self.assertEqual(hashlib.sha256(self.original).hexdigest(), ORIGINAL_SHA256)
        response = self.client.post('/templates/upload/', {'name': 'Synthetic WM workflow',
            'department': self.department.pk, 'synthetic_only': 'on',
            'file': SimpleUploadedFile('WM IOV Template.docx', self.original)})
        self.assertEqual(response.status_code, 302, getattr(response, 'context', None))
        self.template = Template.objects.get(name='Synthetic WM workflow')
        self.version = self.template.versions.get(number=1)
        self.assertEqual(self.version.document_backend['backend'], BACKEND)
        self.assertEqual((len(self.version.fields), len(self.version.review_requirements)), (90, 14))

    def prepare(self, *, positive=False, missing_measurement=False, long_answer=False, select_variants=True, heading=None):
        values = {}
        for field in self.version.fields:
            if field['kind'] == 'choice':
                value = next(v for v in field['options'] if v not in ('reflux positive.', 'Asymptomatic'))
            elif field['kind'] == 'measurement':
                value = '3.2 mm' if field['locator']['unit'] == 'mm' else '400 ms'
            else:
                value = 'SYNTHETIC ' + field['key']
            values[field['key']] = value
        if positive:
            values['sdt_-213273323'] = 'reflux positive.'
            values['measurement_1_2'] = '650 ms'
            values['sdt_1706206184'] = ', however, there has been a recurrence.'
        if missing_measurement: values['measurement_15_1'] = ''
        if long_answer: values['patient_name'] = 'SYNTHETIC 😀 ' + ('documented ' * 1800) + 'END-OF-COMPLETE-ANSWER'
        if heading:
            # A question whose own value introduces the lines after it, and a diameter below it the export can be read by.
            values[heading] = 'none of the following:'
            values['measurement_15_1'] = '7.7 mm'
        variants = {row['key']: row['default_variant'] for row in self.version.review_requirements}
        for key in ('claim_23', 'claim_27'): variants[key] = 'original'
        for key in ('claim_30', 'claim_31', 'claim_34', 'claim_35'): variants[key] = 'selected'
        variants.update(claim_19='technique_only', claim_32='positive' if positive else 'negative', claim_36='negative')
        answers = {key: {'text': value, 'status': 'supported' if value else 'missing',
                         'reviewed_by': '', 'revision': 0, 'evidence': []} for key, value in values.items()}
        context = {'reviewer_id': str(self.user.pk), 'revision': 0, 'notes': '',
                   'signer_name': self.user.get_full_name(), 'finalized_at': timezone.now().isoformat()}
        resolved = resolve_template(private_path(self.template.original_path), self.version.document_backend,
                                    answers, claim_variants=variants, trusted_context=context)['review_requirements']
        lines = [field['label'] + ': ' + values[field['key']] for field in self.version.fields if values[field['key']]]
        lines += [row['label'] + ': ' + row['text'] for row in resolved if row['kind'] == 'assertion']
        notes = '\n'.join(lines)
        response = self.client.post('/encounters/new/', {'patient': self.patient.pk,
            'template_version': self.version.pk, 'notes': notes, 'service_date': timezone.localdate().isoformat(), 'synthetic_only': 'on'})
        self.assertEqual(response.status_code, 302)
        encounter = Encounter.objects.get(template_version=self.version)
        self.client.post(f'/encounters/{encounter.pk}/generate/')
        self.assertTrue(run_one())
        encounter.refresh_from_db()
        self.assertEqual(Job.objects.get(encounter=encounter).state, 'succeeded')
        order = [field['key'] for field in self.version.fields]
        self.governed = ({key for key in order[order.index(heading):] if values[key]} if heading else set())
        for key, value in values.items():
            if key in self.governed:
                answer = encounter.answers[key]
                self.assertEqual((answer['status'], answer['text'], answer['held_candidate']['text']), ('conflict', '', value), key)
                self.assertNotIn('suggestion', answer)
            else:
                self.assertEqual(encounter.answers[key]['text'], value)
        # First select alternatives without claiming to have seen new wording.
        self.values, self.variants = values, variants
        if select_variants:self.save(encounter, reviewed=False)
        return encounter

    def payload(self, encounter, *, reviewed=True):
        response = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertEqual(response.status_code, 200)
        if reviewed:self.assertFalse(response.context['wm_review_error'], response.context['wm_review_error'])
        payload = {'revision': encounter.revision}
        for field in self.version.fields:
            key, value = field['key'], self.values[field['key']]
            payload['answer_' + key] = value
            payload['source_' + key] = field['label'] + ': ' + value if value else ''
            if reviewed: payload['reviewed_' + key] = 'on'
        for row in response.context['wm_requirements']:
            key = row['key']
            payload['claim_variant_' + key] = self.variants[key]
            payload['claim_token_' + key] = row.get('review_token', '')
            payload['claim_source_' + key] = row['label'] + ': ' + row['text'] if row['kind'] == 'assertion' else ''
            if reviewed: payload['claim_reviewed_' + key] = 'on'
        return payload

    def save(self, encounter, *, reviewed=True, payload=None):
        response = self.client.post(f'/encounters/{encounter.pk}/save/', payload or self.payload(encounter, reviewed=reviewed))
        self.assertEqual(response.status_code, 302)
        encounter.refresh_from_db()
        return response

    def finalize(self, encounter):
        response = self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'})
        self.assertEqual(response.status_code, 302)
        encounter.refresh_from_db()
        self.assertEqual(encounter.status, 'finalized', encounter.template_review.get('assessment', {}).get('blockers'))
        response = self.client.get(f'/encounters/{encounter.pk}/export/')
        self.assertEqual(response.status_code, 200)
        payload = b''.join(response.streaming_content)
        self.assertEqual(hashlib.sha256(payload).hexdigest(), encounter.export_sha256)
        self.assertEqual(private_path(self.template.original_path).read_bytes(), self.original)
        return payload

    def test_upload_generate_review_original_export_actual_identity_and_blank_measurement(self):
        encounter = self.prepare(missing_measurement=True)
        self.save(encounter)
        self.assertTrue(encounter.template_review['assessment']['complete'])
        self.assertEqual(encounter.answers['measurement_15_1']['text'], '')
        exported = self.finalize(encounter)
        with ZipFile(BytesIO(exported)) as zipped:
            document = zipped.read('word/document.xml').decode()
        self.assertIn(self.user.get_full_name(), document)
        self.assertIn('No saphenous reflux in right lower extremity.', document)
        self.assertNotIn('Evaluation Only', document)
        self.assertEqual(encounter.finalized_by_id, self.user.pk)

    def test_a_question_whose_value_is_a_heading_holds_what_follows_through_review_finalize_and_export(self):
        """The older-form case followed through on the qualified form (2026-09-23): a location written as a heading held
        every answer after it (measured on 2d900862: all READY). Held: a tick on the page as rendered cannot finalize;
        the clinician's own answers, typed on the page, finalize and are exported."""
        encounter = self.prepare(heading='location_11', select_variants=False)
        self.assertIn('measurement_15_1', self.governed)
        self.assertEqual(encounter.answers['measurement_1_1']['text'], '3.2 mm')        # above it: ready
        heading = 'Right tributary row 11: anatomical location: none of the following:'
        page = self.client.get(f'/encounters/{encounter.pk}/')
        rendered = page.content.decode()
        sections = {}
        for part in rendered.split('<section class="answer">')[1:]:
            for field in self.version.fields:
                if f'. {field["label"]}</h3>' in part.split('</h3>', 1)[0] + '</h3>':
                    sections[field['key']] = part.split('</section>')[0]
        for key in self.governed:
            section = sections[key]
            # Held on the page: its own badge, the heading among its GOVERNING lines (the notes box shows every line
            # anyway), and not the "conflicting values" message the empty held box used to produce.
            self.assertIn('data-held-candidate', section, key)
            self.assertIn('class="badge conflict"', section, key)
            quotes = [html.unescape(q) for q in re.findall(r'data-governing-line><strong>[^<]*</strong> <q>(.*?)</q>', section, re.S)]
            self.assertIn(heading, quotes, key)
            self.assertNotIn('The notes contain conflicting values for this question', section, key)

        class SaveForm(HTMLParser):
            """What a browser posts for this page AS RENDERED: every control of the save form, unchanged."""
            def __init__(self):
                super().__init__(); self.values = {}; self.inside = False; self.text = None; self.select = None
            def handle_starttag(self, tag, attrs):
                attrs = dict(attrs)
                if tag == 'form':
                    self.inside = (attrs.get('action') or '').endswith('/save/')
                if not self.inside or not attrs.get('name') and tag != 'option':
                    return
                if tag == 'input':
                    if attrs.get('type') != 'checkbox' or 'checked' in attrs:
                        self.values[attrs['name']] = attrs.get('value', 'on' if attrs.get('type') == 'checkbox' else '')
                elif tag == 'textarea':
                    self.text = attrs['name']; self.values[self.text] = ''
                elif tag == 'select':
                    self.select = attrs['name']; self.values.setdefault(self.select, '')
                elif tag == 'option' and self.select and 'selected' in attrs:
                    self.values[self.select] = attrs.get('value', '')
            def handle_endtag(self, tag):
                if tag == 'textarea': self.text = None
                elif tag == 'select': self.select = None
                elif tag == 'form': self.inside = False
            def handle_data(self, data):
                if self.text: self.values[self.text] += data
        form = SaveForm(); form.feed(rendered)
        payload = {name: html.unescape(value).lstrip('\n') for name, value in form.values.items()}
        for key in self.governed:
            self.assertEqual(payload.get('answer_' + key, ''), '', key)                 # the held box renders empty
            self.assertEqual(payload.get('held_shown_' + key), '1', key)
        shown = {name: value for name, value in payload.items() if name.startswith('held_shown_')}
        # The page as rendered, and the clinician ticks every box (the statements' sources pasted as the form asks).
        for field in self.version.fields:
            payload['reviewed_' + field['key']] = 'on'
        for row in page.context['wm_requirements']:
            payload['claim_reviewed_' + row['key']] = 'on'
            if row['kind'] == 'assertion' and not payload.get('claim_source_' + row['key']):
                payload['claim_source_' + row['key']] = row['label'] + ': ' + row['text']
        self.save(encounter, payload=payload)
        self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'})
        encounter.refresh_from_db()
        self.assertNotEqual(encounter.status, 'finalized')
        self.assertFalse(encounter.export_path)
        # The clinician types each answer on the page (the form accepts only the value under its own label).
        payload = self.payload(encounter, reviewed=False)
        payload.update(shown)
        self.save(encounter, payload=payload)
        for key in self.governed:
            self.assertEqual((encounter.answers[key]['status'], encounter.answers[key]['text']), ('edited', self.values[key]), key)
        self.save(encounter)
        exported = self.finalize(encounter)
        # The adapter writes a measurement's number into its own cell (the unit is the column's), and the export splits
        # the table by leg, so the value is looked for in every exported table cell (CI measured both).
        from docx import Document
        cells = [cell.text.strip() for table in Document(BytesIO(exported)).tables for row in table.rows for cell in row.cells]
        self.assertIn('7.7', cells)

    def test_positive_reflux_approved_statement_and_long_answer_are_not_truncated(self):
        encounter = self.prepare(positive=True, long_answer=True)
        self.save(encounter)
        exported = self.finalize(encounter)
        with ZipFile(BytesIO(exported)) as zipped:
            document = zipped.read('word/document.xml').decode()
        self.assertIn('END-OF-COMPLETE-ANSWER', document)
        self.assertNotIn('No saphenous reflux in right lower extremity.', document)
        self.assertIn('reflux positive.', document)

    def test_initial_positive_conflict_page_submits_rendered_hidden_identity_variants(self):
        encounter = self.prepare(positive=True, select_variants=False)
        response = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertTrue(response.context['wm_review_error'])
        # The adapter's own conflict strings, named on the page: which statement, which answer.
        from core.wm_conflicts import GENERAL
        named = [row for row in response.context['wm_conflicts']
                 if any('saphenous reflux finding' in statement[0] for statement in row['statements'])]
        self.assertTrue(named, response.context['wm_conflicts'])
        for row in named:
            self.assertNotEqual(row['explanation'], GENERAL, row)
            self.assertTrue(row['answers'], row)
            # A measurement is named, never quoted.
            self.assertTrue(all(answer[1] == '' for answer in row['answers'] if '(ms)' in answer[0]), row)
        # The positive reflux the default no-reflux wording contradicts is named: its reflux time over 500 ms, or its
        # impression. (The adapter reports the rule it meets first; CI measured the reflux-time one.)
        self.assertTrue(any('reflux time' in answer[0]
                            or ('impression: reflux or absence' in answer[0] and answer[1] == 'reflux positive.')
                            for row in named for answer in row['answers']), named)
        self.assertContains(response, 'data-wm-conflict')
        class HiddenVariants(HTMLParser):
            def __init__(self):
                super().__init__(); self.values = {}
            def handle_starttag(self, tag, attrs):
                attrs = dict(attrs)
                if tag == 'input' and attrs.get('type') == 'hidden' and attrs.get('name', '').startswith('claim_variant_'):
                    self.values[attrs['name']] = attrs.get('value', '')
        rendered = HiddenVariants(); rendered.feed(response.content.decode())
        self.assertEqual(rendered.values, {'claim_variant_claim_11': 'trusted',
                                          'claim_variant_claim_39': 'original',
                                          'claim_variant_claim_45': 'trusted',
                                          'claim_variant_claim_46': 'trusted'})
        payload = self.payload(encounter, reviewed=False)
        payload.update(rendered.values)
        self.save(encounter, payload=payload)
        self.assertEqual(encounter.template_review['schema'], 1)
        self.assertEqual(encounter.template_review['decisions']['claim_32']['variant'], 'positive')
        self.assertFalse(encounter.template_review['assessment']['complete'])
        response = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertFalse(response.context['wm_review_error'], response.context['wm_review_error'])

    def test_changed_statement_or_forged_visible_proof_requires_new_review(self):
        encounter = self.prepare()
        payload = self.payload(encounter)
        payload['claim_variant_claim_19'] = 'original'
        self.save(encounter, payload=payload)
        self.assertFalse(encounter.template_review['decisions']['claim_19']['reviewed'])
        self.assertFalse(encounter.template_review['assessment']['complete'])
        self.variants['claim_19'] = 'original'
        payload = self.payload(encounter)
        payload['claim_token_claim_9'] = 'forged-proof'
        self.save(encounter, payload=payload)
        self.assertFalse(encounter.template_review['decisions']['claim_9']['reviewed'])
        self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'})
        encounter.refresh_from_db(); self.assertNotEqual(encounter.status, 'finalized')
        payload = self.payload(encounter)
        with patch('documents.resolve_template', side_effect=DocumentError('Synthetic busy-processor failure.')):
            self.save(encounter, payload=payload)
        self.assertEqual(encounter.template_review['schema'], 0)
        self.assertTrue(all(not row['reviewed'] for row in encounter.template_review['decisions'].values()))

    def test_unsupported_choice_stays_visible_and_cannot_export(self):
        encounter = self.prepare()
        payload = self.payload(encounter)
        payload['answer_sdt_-213273323'] = 'Unsupported synthetic option'
        self.save(encounter, payload=payload)
        self.assertEqual(encounter.answers['sdt_-213273323']['text'], 'Unsupported synthetic option')
        response = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertContains(response, 'Unsupported current value: Unsupported synthetic option')
        self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': encounter.revision, 'reviewed': 'on'})
        encounter.refresh_from_db(); self.assertFalse(encounter.export_path)

    def test_source_edit_and_simulated_voice_adoption_clear_reviews(self):
        encounter = self.prepare(); self.save(encounter)
        self.client.post(f'/encounters/{encounter.pk}/notes/', {'revision': encounter.revision,
                          'notes': encounter.notes + '\nSynthetic correction: revised source'})
        encounter.refresh_from_db()
        self.assertEqual(encounter.template_review, {}); self.assertEqual(encounter.answers, {})
        # Explicitly simulated transcription: this verifies adoption/invalidation,
        # not microphone capture or a live transcription provider.
        audio = private_path('audio/synthetic.wav'); audio.parent.mkdir(parents=True)
        audio.write_bytes(wav_bytes(tag='wm-voice')); audio.chmod(0o600)
        config = {**configuration(), 'voice_rate_minute': '0', 'voice_rate_second': '', 'voice_max_cost': '0'}
        encounter.template_review = {'schema': 0, 'synthetic_test_marker': True}
        encounter.save(update_fields=['template_review'])
        enqueue(encounter, self.user, kind='transcription', input_path='audio/synthetic.wav', config=config)
        with patch('core.jobs.transcribe', return_value=('SYNTHETIC voice correction', Decimal('0'),
                   {'duration_seconds': '1'}, {'usd_per_minute': '0'}, 'synthetic-transcription-fixture', '', True)) as provider:
            self.assertTrue(run_one())
        provider.assert_called_once()
        encounter.refresh_from_db()
        self.assertIn('SYNTHETIC voice correction', encounter.notes)
        self.assertEqual(encounter.template_review, {}); self.assertEqual(encounter.answers, {})

    def test_employee_read_scope_review_denial_and_other_department_isolation(self):
        encounter = self.prepare()
        employee = User.objects.create_user(username='wm-employee', email='wm-employee@synthetic.invalid', verified_at=timezone.now())
        Membership.objects.create(user=employee, department=self.department, role='employee')
        self.client.force_login(employee)
        self.assertEqual(self.client.get(f'/encounters/{encounter.pk}/').status_code, 200)
        before = encounter.revision
        response = self.client.post(f'/encounters/{encounter.pk}/save/', {'revision': before, 'claim_reviewed_claim_9': 'on'})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.client.post(f'/encounters/{encounter.pk}/finalize/', {'revision': before, 'reviewed': 'on'}).status_code, 403)
        encounter.refresh_from_db(); self.assertEqual(encounter.revision, before)
        other = Department.objects.create(organization=self.organization, name='Synthetic ultrasound')
        Membership.objects.filter(user=employee).update(department=other)
        for suffix in ('', 'export/', 'save/', 'finalize/'):
            url = f'/encounters/{encounter.pk}/' + suffix
            response = self.client.post(url, {'revision': before}) if suffix in ('save/', 'finalize/') else self.client.get(url)
            self.assertEqual(response.status_code, 404)

    def test_mapping_versions_copy_backend_and_requirements_but_reject_label_changes(self):
        labels = {row['key']: row['label'] for row in self.version.fields}
        response = self.client.post(f'/templates/{self.template.pk}/', {'labels': json.dumps(labels), 'mapping_version': str(self.version.pk)})
        self.assertEqual(response.status_code, 302)
        second = self.template.versions.get(number=2)
        self.assertEqual(second.document_backend, self.version.document_backend)
        self.assertEqual(second.review_requirements, self.version.review_requirements)
        labels['patient_name'] = 'Changed qualified label'
        response = self.client.post(f'/templates/{self.template.pk}/', {'labels': json.dumps(labels), 'mapping_version': str(second.pk)})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'fixed question labels')
        self.assertEqual(self.template.versions.count(), 2)


class RenderedSaveForm(HTMLParser):
    """What a browser posts for a page AS RENDERED: every control of the save form, unchanged."""
    def __init__(self):
        super().__init__(); self.values = {}; self.inside = False; self.text = None; self.select = None
    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'form':
            self.inside = (attrs.get('action') or '').endswith('/save/')
        if not self.inside or not attrs.get('name') and tag != 'option':
            return
        if tag == 'input':
            if attrs.get('type') != 'checkbox' or 'checked' in attrs:
                self.values[attrs['name']] = attrs.get('value', 'on' if attrs.get('type') == 'checkbox' else '')
        elif tag == 'textarea':
            self.text = attrs['name']; self.values[self.text] = ''
        elif tag == 'select':
            self.select = attrs['name']; self.values.setdefault(self.select, '')
        elif tag == 'option' and self.select and 'selected' in attrs:
            self.values[self.select] = attrs.get('value', '')
    def handle_endtag(self, tag):
        if tag == 'textarea': self.text = None
        elif tag == 'select': self.select = None
        elif tag == 'form': self.inside = False
    def handle_data(self, data):
        if self.text: self.values[self.text] += data

    @classmethod
    def of(cls, page):
        form = cls(); form.feed(page.content.decode())
        return {name: html.unescape(value).lstrip('\n') for name, value in form.values.items()}


@skipUnless(WORD_RUNTIME_USABLE,
            'Live Word document workflow NOT EXERCISED — ' + (WORD_RUNTIME_REASON or 'unknown reason') +
            '. No claim is made about this path from this run.')
@override_settings(**TEST_SETTINGS)
class WMNarrativeWorkflow(TestCase):
    """Ordinary narrative notes through the real application and the real Word engine (docs/rebuild/wm-narrative-design.md):
    a RETAINED provider-shaped response (the frozen benchmark's careful set, replayed in fixture mode; not fresh live
    generation) → generation → the page as rendered → the clinician's decisions → finalization → original export."""
    setUp = WMDocumentWorkflow.setUp

    def benchmark(self, name):
        from core.tests.wm_narrative_benchmark import load
        return load(name)

    def generate(self, case_id):
        case = {row['id']: row for row in self.benchmark('cases.json')['cases']}[case_id]
        response = self.benchmark('responses-careful.json')['cases'][case_id]
        replay = Path(self.directory.name) / 'replay'
        replay.mkdir(exist_ok=True)
        (replay / (hashlib.sha256(case['notes'].encode()).hexdigest() + '.json')).write_text(json.dumps(response))
        self.narrative = self.settings(WM_NARRATIVE=True, AI_REPLAY_DIR=str(replay))
        self.narrative.enable(); self.addCleanup(self.narrative.disable)
        created = self.client.post('/encounters/new/', {'patient': self.patient.pk, 'template_version': self.version.pk,
            'notes': case['notes'], 'service_date': timezone.localdate().isoformat(), 'synthetic_only': 'on'})
        self.assertEqual(created.status_code, 302)
        encounter = Encounter.objects.get(template_version=self.version)
        self.client.post(f'/encounters/{encounter.pk}/generate/')
        self.assertTrue(run_one())
        encounter.refresh_from_db()
        job = Job.objects.get(encounter=encounter)
        self.assertEqual((job.state, job.config.get('narrative_contract'), job.provider),
                         ('succeeded', 'wm-narrative-v1', 'synthetic-narrative-replay'))
        return case, encounter

    def test_the_benchmark_mapping_is_the_engines_own_inspection(self):
        mapping = self.benchmark('mapping.json')
        self.assertEqual(self.version.fields, mapping['fields'])
        self.assertEqual(self.version.review_requirements, mapping['review_requirements'])
        self.assertEqual(self.version.document_backend, mapping['backend'])

    def test_narrative_notes_are_reviewed_and_exported_into_the_unchanged_original(self):
        case, encounter = self.generate('C12')
        answers = encounter.answers
        for key, expected in case['expected_fields'].items():
            if expected['outcome'] == 'supported':
                self.assertEqual((answers[key]['status'], answers[key]['text']), ('supported', expected['value']), key)
                self.assertTrue(all(entry['kind'] == 'unit' and case['notes'][entry['start']:entry['end']] == entry['quote']
                                    for entry in answers[key]['evidence']), key)
        self.assertEqual(set(encounter.template_review), {'proposals'})
        page = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertContains(page, 'data-wm-narrative')
        self.assertContains(page, 'data-statement-proposed')
        # Only the two statements nothing proposed start on their defaults, and the previous-treatment default needs
        # answers the notes do not give: the page names exactly those conflicts before anything is reviewed.
        named = {statement[0] for conflict in page.context['wm_conflicts'] for statement in conflict['statements']}
        self.assertTrue(named and all('previous treatment' in name for name in named), page.context['wm_conflicts'])
        payload = RenderedSaveForm.of(page)
        # The page as rendered carries every generated answer and its sentences.
        for field in self.version.fields:
            self.assertEqual(payload.get('answer_' + field['key'], ''), answers[field['key']]['text'], field['key'])
        # The notes say nothing about previous treatment: the clinician chooses "not documented", resting on those
        # empty answers. No statement line was added to the notes.
        for key in ('claim_23', 'claim_27'):
            payload['claim_variant_' + key] = 'not_documented'
            payload['claim_basis_' + key] = 'answers'
        self.client.post(f'/encounters/{encounter.pk}/save/', payload)
        encounter.refresh_from_db()
        # The chosen wordings are shown before they can be reviewed: review everything on the page as it now renders.
        page = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertFalse(page.context['wm_review_error'], page.context['wm_review_error'])
        payload = RenderedSaveForm.of(page)
        self.assertEqual((payload.get('claim_basis_claim_23'), payload.get('claim_variant_claim_23')), ('answers', 'not_documented'))
        for field in self.version.fields:
            payload['reviewed_' + field['key']] = 'on'
        for row in page.context['wm_requirements']:
            payload['claim_reviewed_' + row['key']] = 'on'
        self.client.post(f'/encounters/{encounter.pk}/save/', payload)
        encounter.refresh_from_db()
        self.assertTrue(encounter.template_review['assessment']['complete'], encounter.template_review['assessment']['blockers'])
        self.assertEqual(encounter.notes, case['notes'])
        exported = WMDocumentWorkflow.finalize(self, encounter)
        from docx import Document
        document = Document(BytesIO(exported))
        cells = [cell.text.strip() for table in document.tables for row in table.rows for cell in row.cells]
        for value in ('6.0', '1600', '6.1', '1300'):
            self.assertIn(value, cells)
        with ZipFile(BytesIO(exported)) as zipped:
            xml = zipped.read('word/document.xml').decode()
        self.assertIn('Saphenous reflux is present in the right lower extremity.', xml)
        self.assertIn('Previous treatment history was not documented for the right lower extremity.', xml)
        self.assertIn('Clinical correlation with the above findings is required.', xml)
        self.assertNotIn('No saphenous reflux in right lower extremity.', xml)
        self.assertEqual(xml.count('<w:sdt>') + xml.count('<w:sdt '), 28)

    def test_the_engine_refuses_a_narrative_answer_its_cited_sentences_do_not_support(self):
        """The second boundary, driven directly with a forged request: a narrative answer on a side its sentence does
        not name, and one whose value its sentence does not contain, are refused at fill whatever the parent sent. The
        positive control passes the answer check and stops later, at the unreviewed statements."""
        notes = 'The right great saphenous vein measures 7.4 mm at the saphenofemoral junction.'
        unit = {'quote': notes, 'start': 0, 'end': len(notes), 'kind': 'unit'}
        variants = {row['key']: row['default_variant'] for row in self.version.review_requirements}
        variants.update(claim_9='original', claim_19='not_documented', claim_23='not_documented', claim_27='not_documented',
                        claim_30='not_documented', claim_31='not_documented', claim_34='not_documented',
                        claim_35='not_documented')
        context = {'reviewer_id': str(self.user.pk), 'revision': 3, 'notes': notes,
                   'signer_name': self.user.get_full_name(), 'finalized_at': timezone.now().isoformat()}
        def attempt(key, text):
            answers = {field['key']: {'text': '', 'status': 'missing', 'reviewed_by': str(self.user.pk), 'revision': 3,
                                      'evidence': []} for field in self.version.fields}
            answers[key] = {'text': text, 'status': 'supported', 'reviewed_by': str(self.user.pk), 'revision': 3,
                            'evidence': [unit]}
            with self.assertRaises(DocumentError) as refused:
                fill_template(private_path(self.template.original_path), self.version.fields, answers,
                              Path(self.directory.name) / ('forged-' + key + text.replace(' ', '') + '.docx'),
                              document_backend=self.version.document_backend, reviewed_requirements={},
                              claim_variants=variants, trusted_context=context)
            return refused.exception.code
        self.assertEqual(attempt('measurement_15_1', '7.4 mm'), 'unsupported_answer')    # the LEFT row, a right sentence
        self.assertEqual(attempt('measurement_1_1', '6.9 mm'), 'unsupported_answer')     # a number it does not contain
        self.assertEqual(attempt('measurement_1_1', '7.4 mm'), 'review_required')        # control: the answer passes

    def test_the_engine_reads_characters_as_the_parent_does(self):
        """Refutation round 2 (parity lens): the engine and the parent read a side, a quantity and a text answer over the
        same classes of character, so the export never refuses what review accepted, and neither reads a side the
        other does not. Each note is judged by the parent first, then the engine gets the same answer."""
        from core import generation, wm_narrative
        variants = {row['key']: row['default_variant'] for row in self.version.review_requirements}
        variants.update(claim_9='original', claim_19='not_documented', claim_23='not_documented', claim_27='not_documented',
                        claim_30='not_documented', claim_31='not_documented', claim_34='not_documented',
                        claim_35='not_documented')
        fields = {field['key']: field for field in self.version.fields}
        labels = [row['label'] for row in self.version.review_requirements]

        def both(notes, key, text):
            unit = next(u for u in wm_narrative.units(notes, self.version.fields, labels) if u['end'] - u['start'] > 1)
            unit = {'quote': unit['quote'], 'start': unit['start'], 'end': unit['end'], 'kind': 'unit'}
            recognized = generation.recognized_labels(self.version.fields, labels)
            parent = wm_narrative.support(fields[key], text, [unit], notes, recognized,
                                          [field['label'] for field in self.version.fields] + labels)
            context = {'reviewer_id': str(self.user.pk), 'revision': 3, 'notes': notes,
                       'signer_name': self.user.get_full_name(), 'finalized_at': timezone.now().isoformat()}
            answers = {field['key']: {'text': '', 'status': 'missing', 'reviewed_by': str(self.user.pk), 'revision': 3,
                                      'evidence': []} for field in self.version.fields}
            answers[key] = {'text': text, 'status': 'supported', 'reviewed_by': str(self.user.pk), 'revision': 3,
                            'evidence': [unit]}
            with self.assertRaises(DocumentError) as refused:
                fill_template(private_path(self.template.original_path), self.version.fields, answers,
                              Path(self.directory.name) / 'parity.docx', document_backend=self.version.document_backend,
                              reviewed_requirements={}, claim_variants=variants, trusted_context=context)
            return parent, refused.exception.code

        # `review_required`: the answer check passed and the fill stopped later, at the unreviewed statements.
        self.assertEqual(both('The RİGHT great saphenous vein measures 7.4 mm at the saphenofemoral junction.',
                              'measurement_1_1', '7.4 mm'), (None, 'review_required'))
        # A number the hard wrap put before its unit's line is one quantity in one unit, for both (round 17).
        for eol in ('\n', '\r\n'):
            self.assertEqual(both('The right great saphenous vein measures 7.4' + eol + 'mm at the saphenofemoral junction.',
                                  'measurement_1_1', '7.4 mm'), (None, 'review_required'))
        self.assertEqual(both('The great saphenous vein at the saphenofemoral junction measures 7.4 mm and the legs are '
                              'bothersome.', 'measurement_15_1', '7.4 mm'), ('side_unknown', 'unsupported_answer'))
        self.assertEqual(both('İBRAHİM YILMAZ presents for a venous study.', 'patient_name', 'Ibrahim Yilmaz'),
                         (None, 'review_required'))
        self.assertEqual(both('KARL GROSS presents for a venous study.', 'patient_name', 'Karl Gro\u00df'),
                         ('unsupported', 'unsupported_answer'))
        # A bound read after the value, in place (round 3: the engine no longer copies the rest of the text per number).
        self.assertEqual(both('The right great saphenous vein measures 7.4 mm or more at the saphenofemoral junction.',
                              'measurement_1_1', '7.4 mm'), ('unsupported', 'unsupported_answer'))
        self.assertEqual(both('The right great saphenous vein measures 7.4 mm (approx.) at the saphenofemoral junction.',
                              'measurement_1_1', '7.4 mm'), ('unsupported', 'unsupported_answer'))
        self.assertEqual(both('The right great saphenous vein measures 7.4 mm' + ' ' * 4000 + 'at the saphenofemoral junction.',
                              'measurement_1_1', '7.4 mm'), (None, 'review_required'))

    def generate_from(self, notes, response):
        """A retained response for notes of the test's own, replayed through the real generation."""
        replay = Path(self.directory.name) / 'replay'
        replay.mkdir(exist_ok=True)
        (replay / (hashlib.sha256(notes.encode()).hexdigest() + '.json')).write_text(json.dumps(response))
        self.narrative = self.settings(WM_NARRATIVE=True, AI_REPLAY_DIR=str(replay))
        self.narrative.enable(); self.addCleanup(self.narrative.disable)
        created = self.client.post('/encounters/new/', {'patient': self.patient.pk, 'template_version': self.version.pk,
            'notes': notes, 'service_date': timezone.localdate().isoformat(), 'synthetic_only': 'on'})
        self.assertEqual(created.status_code, 302)
        encounter = Encounter.objects.get(template_version=self.version)
        self.client.post(f'/encounters/{encounter.pk}/generate/')
        self.assertTrue(run_one())
        encounter.refresh_from_db()
        return encounter

    def test_the_first_page_reviews_the_proposed_wordings_and_keeps_that_review(self):
        """Refutation round 3 (integrity lens): the first page resolved the proposed statements but reviewed the
        defaults, showed an error and the original wordings, and discarded the clinician's statement reviews."""
        from core import wm_narrative
        notes = ('SYNTHETIC PATIENT P had a bilateral detailed venous study.\n'
                 'The right great saphenous vein measures 7.4 mm at the saphenofemoral junction.\n'
                 'The left great saphenous vein measures 5.0 mm at the saphenofemoral junction.\n'
                 'Clinical correlation with the above findings is required.')
        labels = [row['label'] for row in self.version.review_requirements]
        ids = [unit['id'] for unit in wm_narrative.units(notes, self.version.fields, labels)]
        picks = {'sdt_1475334564': ('Bilateral Detailed', [ids[0]]), 'measurement_1_1': ('7.4 mm', [ids[1]]),
                 'measurement_15_1': ('5.0 mm', [ids[2]])}
        rows = [{'key': field['key'], 'value': picks.get(field['key'], ('', []))[0], 'unit_ids': picks.get(field['key'], ('', []))[1],
                 'conflict': False, 'subject': 'patient', 'certainty': 'definite',
                 'timing': 'prior' if wm_narrative.previous_treatment(field) else 'current'} for field in self.version.fields]
        statements = [{'key': 'claim_9', 'variant': 'exam_choice', 'unit_ids': [ids[0]]},
                      {'key': 'claim_39', 'variant': 'original', 'unit_ids': [ids[3]]}]
        statements += [{'key': key, 'variant': 'not_documented', 'unit_ids': []}
                       for key in ('claim_19', 'claim_23', 'claim_27', 'claim_30', 'claim_31', 'claim_32', 'claim_34',
                                   'claim_35', 'claim_36')]
        encounter = self.generate_from(notes, {'fields': rows, 'statements': statements})
        self.assertEqual(encounter.template_review['proposals']['claim_9']['variant'], 'exam_choice')
        page = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertFalse(page.context['wm_review_error'], page.context['wm_review_error'])
        shown = {row['key']: row for row in page.context['wm_requirements']}
        self.assertEqual(shown['claim_9']['selected_variant'], 'exam_choice')
        self.assertIn('Bilateral Detailed', shown['claim_9']['text'])
        payload = RenderedSaveForm.of(page)
        for field in self.version.fields:
            payload['reviewed_' + field['key']] = 'on'
        for row in page.context['wm_requirements']:
            payload['claim_reviewed_' + row['key']] = 'on'
        self.client.post(f'/encounters/{encounter.pk}/save/', payload)
        encounter.refresh_from_db()
        reviewed = {key: row['reviewed'] for key, row in encounter.template_review['decisions'].items()}
        self.assertTrue(reviewed and all(reviewed.values()), reviewed)

    def test_review_is_never_complete_where_finalization_refuses(self):
        """Refutation round 3 (integrity lens): a save changing which answers cite a correction left review complete
        while finalization refused the answers the correction now held. Review applies finalization's recheck."""
        case, encounter = self.generate('C12')
        answers = dict(encounter.answers)
        held = {key for key, answer in answers.items() if answer.get('status') == 'supported'}
        from core import generation, wm_review
        # A correction the notes gain nobody cites: finalization's recheck holds the measurements it could correct.
        encounter.notes = case['notes'] + '\n\nCorrection: the knee measurement is 6.1 mm.'
        governed = generation.governed_decisions(encounter.notes, self.version.fields, answers,
                                                 [row['label'] for row in self.version.review_requirements])
        self.assertTrue(governed, 'the control: finalization refuses these answers')
        assessment = wm_review._answer_assessment(encounter, self.user, self.version.fields, self.version.review_requirements)
        refused = {field['label'] for field in self.version.fields
                   if 'narrative_held' in assessment[field['key']]['blockers']}
        self.assertEqual(refused, set(governed) & {field['label'] for field in self.version.fields if field['key'] in held})

    # ── round 17e's review: provenance through the page, the save, finalization and export ──────────────────────────
    PROVENANCE_NOTES = ('SYNTHETIC PATIENT P had a bilateral detailed venous study.\n'
                        'The right great saphenous vein measures 7.4 mm at the saphenofemoral junction.\n'
                        "Her mother's right great saphenous vein measures 7.4 mm at the saphenofemoral junction.\n"
                        'The left leg is more swollen, and the great saphenous vein measures 5.1 mm at the saphenofemoral junction.\n'
                        'There is no saphenous reflux in the right lower extremity.\n'
                        'There is no saphenous reflux in the left lower extremity.\n'
                        'Clinical correlation with the above findings is required.')
    PROVENANCE_PICKS = {'sdt_1475334564': ('Bilateral Detailed', [0]), 'measurement_1_1': ('7.4 mm', [1]),
                        'measurement_15_1': ('5.1 mm', [3])}
    PROVENANCE_STATEMENTS = ([('claim_9', 'exam_choice', [0]), ('claim_32', 'negative', [4]), ('claim_36', 'negative', [5]),
                              ('claim_39', 'original', [6])]
                             + [(key, 'not_documented', []) for key in ('claim_19', 'claim_23', 'claim_27', 'claim_30',
                                                                        'claim_31', 'claim_34', 'claim_35')])

    def replayed(self, notes, picks, statements, qualifiers=None):
        """A careless model's response for notes of the test's own (the patient's, definite, current), picks and
        statements by unit index, replayed through the real generation."""
        from core import wm_narrative
        labels = [row['label'] for row in self.version.review_requirements]
        ids = [unit['id'] for unit in wm_narrative.units(notes, self.version.fields, labels)]
        rows = []
        for field in self.version.fields:
            value, indexes = picks.get(field['key'], ('', []))
            row = {'key': field['key'], 'value': value, 'unit_ids': [ids[i] for i in indexes], 'conflict': False,
                   'subject': 'patient', 'certainty': 'definite',
                   'timing': 'prior' if wm_narrative.previous_treatment(field) else 'current'}
            row.update((qualifiers or {}).get(field['key'], {}))
            rows.append(row)
        return self.generate_from(notes, {'fields': rows, 'statements': [
            {'key': key, 'variant': variant, 'unit_ids': [ids[i] for i in indexes]} for key, variant, indexes in statements]})

    def post_page(self, encounter, edits=(), sources=(), review=True, statements=None, extra=()):
        page = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertEqual(page.status_code, 200)
        payload = RenderedSaveForm.of(page)
        if review:
            payload.update({'reviewed_' + field['key']: 'on' for field in self.version.fields})
        if review if statements is None else statements:
            payload.update({'claim_reviewed_' + row['key']: 'on' for row in page.context['wm_requirements']})
        payload.update({'answer_' + key: value for key, value in dict(edits).items()})
        payload.update({'source_' + key: value for key, value in dict(sources).items()})
        payload.update(dict(extra))
        self.client.post(f'/encounters/{encounter.pk}/save/', payload)
        encounter.refresh_from_db()
        return page

    def member(self, name, role):
        user = User.objects.create_user(username=name, email=name + '@synthetic.invalid', verified_at=timezone.now())
        Membership.objects.create(user=user, department=self.department, role=role)
        return user

    def test_a_held_statement_proposal_opens_the_page(self):
        """Round 17e's review (provenance lens): a proposal `threshold_wording` held is `{'held': reason}`, and the page
        raised on it (`KeyError: 'variant'`) for every owner and clinician."""
        notes = ('SYNTHETIC PATIENT P had a bilateral detailed venous study.\n'
                 'The deep and superficial venous systems were examined.\n'
                 'The right great saphenous vein measures 7.4 mm at the saphenofemoral junction, with possible reflux of '
                 'approximately 550 ms.\n'
                 'The left great saphenous vein measures 5.0 mm at the saphenofemoral junction.\n'
                 'Clinical correlation with the above findings is required.')
        picks = {'sdt_1475334564': ('Bilateral Detailed', [0]), 'sdt_429781496': ('deep and superficial', [1]),
                 'measurement_1_1': ('7.4 mm', [2]), 'measurement_1_2': ('550 ms', [2]), 'measurement_15_1': ('5.0 mm', [3])}
        statements = ([('claim_9', 'exam_choice', [0]), ('claim_39', 'original', [4]), ('claim_19', 'technique_only', [1])]
                      + [(key, 'not_documented', []) for key in ('claim_23', 'claim_27', 'claim_30', 'claim_31', 'claim_32',
                                                                 'claim_34', 'claim_35', 'claim_36')])
        encounter = self.replayed(notes, picks, statements, {'measurement_1_2': {'certainty': 'uncertain'}})
        self.assertIn('held', encounter.template_review['proposals']['claim_19'])      # (the control: it is held)
        page = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertEqual(page.status_code, 200)
        self.assertIn('data-statement-held', page.content.decode())

    def test_an_answer_cited_to_other_sentences_is_the_clinicians_and_stays_so_through_export(self):
        """Round 17e's review (provenance lens): `7.4 mm`, re-cited by the clinician to the mother's sentence with its
        text unchanged, stayed `supported` — shown as read from the notes, finalized and exported as source-supported."""
        encounter = self.replayed(self.PROVENANCE_NOTES, self.PROVENANCE_PICKS, self.PROVENANCE_STATEMENTS)
        self.assertEqual(encounter.answers['measurement_1_1']['status'], 'supported')     # (the control)
        mother = "Her mother's right great saphenous vein measures 7.4 mm at the saphenofemoral junction."
        self.post_page(encounter, edits={'measurement_15_1': '5.1 mm'}, sources={'measurement_1_1': mother})
        answer = encounter.answers['measurement_1_1']
        self.assertEqual((answer['text'], answer['status'], answer['edited_by']), ('7.4 mm', 'edited', str(self.user.pk)))
        self.assertEqual([entry['quote'] for entry in answer['evidence']], [mother])
        page = self.client.get(f'/encounters/{encounter.pk}/').content.decode()
        section = page.split('Right GSV at the SFJ: diameter (mm)</h3>')[1].split('</section>')[0]
        self.assertIn('entered by ' + self.user.email + ', not read from the notes', section)
        self.assertNotIn('data-narrative-reading', section)          # (no "read from the notes as" beside it)
        # (an answer the clinician left as generated keeps generation's evidence and status)
        self.post_page(encounter)
        self.assertEqual(encounter.answers['measurement_1_1']['status'], 'edited')
        self.assertEqual(encounter.answers['sdt_1475334564']['status'], 'supported')

    def test_an_employees_answer_names_the_employee_and_keeps_the_proposals(self):
        """Round 17e's review (provenance lens): an employee's choice for a held value was stored with no author and read
        as the reviewing clinician's; and the employee's save dropped generation's statement proposals."""
        encounter = self.replayed(self.PROVENANCE_NOTES, self.PROVENANCE_PICKS, self.PROVENANCE_STATEMENTS)
        self.assertNotEqual(encounter.answers['measurement_15_1']['status'], 'supported')   # (held: the legs compared)
        proposals = encounter.template_review['proposals']
        employee = self.member('wm-employee', 'employee')
        self.client.force_login(employee)
        self.post_page(encounter, edits={'measurement_15_1': '5.1 mm'}, review=False)
        answer = encounter.answers['measurement_15_1']
        self.assertEqual((answer['status'], answer['edited_by'], answer['reviewed']), ('edited', str(employee.pk), False))
        self.assertEqual(encounter.template_review, {'proposals': proposals})
        self.client.force_login(self.user)
        page = self.client.get(f'/encounters/{encounter.pk}/').content.decode()
        self.assertIn('entered by wm-employee@synthetic.invalid, not read from the notes', page)
        self.post_page(encounter)
        answer = encounter.answers['measurement_15_1']
        self.assertEqual((answer['status'], answer['edited_by'], answer['reviewed_by']),
                         ('edited', str(employee.pk), str(self.user.pk)))

    def test_generation_is_refused_over_a_clinicians_decisions(self):
        """Round 17e's review (provenance lens): a POST to generate, which the page never offers once answers exist,
        replaced a clinician's reviewed decisions with a new draft."""
        encounter = self.replayed(self.PROVENANCE_NOTES, self.PROVENANCE_PICKS, self.PROVENANCE_STATEMENTS)
        self.post_page(encounter, edits={'measurement_15_1': '5.1 mm'})
        before = (dict(encounter.answers), Job.objects.filter(encounter=encounter).count())
        employee = self.member('wm-employee-2', 'employee')
        self.client.force_login(employee)
        self.client.post(f'/encounters/{encounter.pk}/generate/', {'revision': encounter.revision})
        self.assertFalse(run_one())
        encounter.refresh_from_db()
        self.assertEqual((encounter.answers, Job.objects.filter(encounter=encounter).count()), before)

    def test_generation_is_refused_over_reviewed_statement_decisions(self):
        """The focused pass of round 17e's review: the refusal read only the answers, so statement reviews alone — the 14
        wordings ticked, no answer — were replaced by a second, billed draft."""
        encounter = self.replayed(self.PROVENANCE_NOTES, self.PROVENANCE_PICKS, self.PROVENANCE_STATEMENTS)
        self.post_page(encounter, review=False, statements=True)
        self.assertTrue(any(d.get('reviewed') for d in encounter.template_review.get('decisions', {}).values()))   # (the control)
        self.assertFalse(any(a.get('reviewed') for a in encounter.answers.values()))
        before = (dict(encounter.answers), dict(encounter.template_review), Job.objects.filter(encounter=encounter).count())
        self.client.post(f'/encounters/{encounter.pk}/generate/', {'revision': encounter.revision})
        self.assertFalse(run_one())
        encounter.refresh_from_db()
        self.assertEqual((encounter.answers, encounter.template_review, Job.objects.filter(encounter=encounter).count()), before)

    def test_generation_is_refused_over_an_answer_a_person_entered(self):
        """The focused pass's mutation proofs: the refusal's reading of the answers was reached by no test — each one's
        clinician had also ticked the statements. An answer typed with nothing ticked is a person's decision too."""
        encounter = self.replayed(self.PROVENANCE_NOTES, self.PROVENANCE_PICKS, self.PROVENANCE_STATEMENTS)
        self.post_page(encounter, edits={'measurement_15_1': '5.1 mm'}, review=False)
        self.assertEqual(encounter.answers['measurement_15_1']['status'], 'edited')          # (the controls: typed,
        self.assertFalse(any(a.get('reviewed') for a in encounter.answers.values()))          # nothing ticked)
        self.assertFalse(any(d.get('reviewed') for d in (encounter.template_review or {}).get('decisions', {}).values()))
        before = (dict(encounter.answers), Job.objects.filter(encounter=encounter).count())
        self.client.post(f'/encounters/{encounter.pk}/generate/', {'revision': encounter.revision})
        self.assertFalse(run_one())
        encounter.refresh_from_db()
        self.assertEqual((encounter.answers, Job.objects.filter(encounter=encounter).count()), before)

    # ── round 18's review: clinician decisions and provenance through the page ─────────────────────────────────────
    def generation_is_refused_after(self, extra, edits, by=None):
        encounter = self.replayed(self.PROVENANCE_NOTES, self.PROVENANCE_PICKS, self.PROVENANCE_STATEMENTS)
        clinician = self.user
        if by is not None:
            self.client.force_login(by)
        self.post_page(encounter, edits=edits, review=False, extra=extra)
        self.client.force_login(clinician)
        self.assertFalse(any(a.get('reviewed') for a in encounter.answers.values()))     # (nothing ticked)
        before = (dict(encounter.answers), dict(encounter.template_review), Job.objects.filter(encounter=encounter).count())
        self.client.post(f'/encounters/{encounter.pk}/generate/', {'revision': encounter.revision})
        self.assertFalse(run_one())
        encounter.refresh_from_db()
        self.assertEqual((encounter.answers, encounter.template_review, Job.objects.filter(encounter=encounter).count()), before)
        return encounter

    def test_round_18_generation_is_refused_over_an_unticked_statement_choice(self):
        """Round 18 (clinician-decision and provenance lenses): a statement wording chosen and saved without its tick was
        replaced by a second generation."""
        encounter = self.generation_is_refused_after({'claim_variant_claim_32': 'not_assessed'}, {})
        self.assertEqual(encounter.template_review['decisions']['claim_32']['variant'], 'not_assessed')

    def test_round_18_generation_is_refused_over_a_cleared_answer(self):
        """Round 18: an answer a person cleared (`missing`, with its `edited_by`) was re-filled by a second generation.
        Cleared by an employee, whose save records no statement decisions: the clearing alone refuses (the mutation proofs
        found a clinician's save refused by its decisions whatever the answers said)."""
        encounter = self.generation_is_refused_after({}, {'measurement_1_1': ''}, by=self.member('wm-employee-r18g', 'employee'))
        self.assertFalse((encounter.template_review or {}).get('decisions'))
        self.assertEqual((encounter.answers['measurement_1_1']['status'],
                          bool(encounter.answers['measurement_1_1'].get('edited_by'))), ('missing', True))

    def test_round_18_an_employees_save_keeps_the_clinicians_statement_decisions(self):
        """Round 18 (clinician-decision and provenance lenses): an employee's ordinary save emptied the review, and the
        clinician's page fell back to the mapping's defaults (`No saphenous reflux …` the clinician had rejected)."""
        encounter = self.replayed(self.PROVENANCE_NOTES, self.PROVENANCE_PICKS, self.PROVENANCE_STATEMENTS)
        self.post_page(encounter, review=False, extra={'claim_variant_claim_32': 'not_assessed'})   # (the clinician decides,
        self.post_page(encounter)                                                         #  sees the wording, ticks all)
        self.assertEqual(encounter.template_review['decisions']['claim_32'], {**encounter.template_review['decisions']['claim_32'],
                                                                              'variant': 'not_assessed', 'reviewed': True})
        clinician = self.user
        employee = self.member('wm-employee-r18', 'employee')
        self.client.force_login(employee)
        self.post_page(encounter, review=False)                                           # (the employee's plain save)
        self.assertIn('claim_32', encounter.template_review.get('decisions') or {})
        decision = encounter.template_review['decisions']['claim_32']
        self.assertEqual((decision['variant'], decision['reviewed']), ('not_assessed', False))
        from core.models import AuditEvent
        withdrawn = AuditEvent.objects.filter(action='encounter.answers_saved', target=str(encounter.pk)).latest('created_at')
        self.assertIn('claim_32', withdrawn.details['withdrawn_statement_reviews'])
        self.client.force_login(clinician)
        page = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertEqual(RenderedSaveForm.of(page)['claim_variant_claim_32'], 'not_assessed')
        # (and it is never finalized as it stands: the kept decisions are unreviewed and incomplete)
        self.assertEqual(encounter.template_review.get('schema'), 0)
        # (and an employee's statement ticks are refused before anything is read — the role guard — so the kept decisions
        # stay unreviewed: the save's own `reviewed=False` is a second guard, equivalent under the first, which the
        # mutation proofs could not separate)
        self.client.force_login(employee)
        revision = encounter.revision
        self.post_page(encounter, review=False, statements=True)
        self.assertEqual(encounter.revision, revision)
        self.assertFalse(any(d.get('reviewed') for d in encounter.template_review['decisions'].values()))
        self.assertEqual(encounter.template_review['decisions']['claim_32']['variant'], 'not_assessed')

    def test_round_18_a_cleared_answer_names_who_cleared_it(self):
        """Round 18 (provenance lens): a value the model read and a person deleted showed as a plain `missing`, as if the
        notes gave nothing."""
        encounter = self.replayed(self.PROVENANCE_NOTES, self.PROVENANCE_PICKS, self.PROVENANCE_STATEMENTS)
        employee = self.member('wm-employee-r18c', 'employee')
        self.client.force_login(employee)
        self.post_page(encounter, edits={'measurement_1_1': ''}, review=False)
        self.client.force_login(self.user)
        page = self.client.get(f'/encounters/{encounter.pk}/')
        self.assertContains(page, 'data-cleared-by')
        self.assertContains(page, 'cleared by wm-employee-r18c@synthetic.invalid')

    def test_round_18_review_refuses_text_the_engine_would_rewrite_in_a_cell(self):
        """Round 18 (provenance lens): the engine writes a table cell through a replace that reads `&p`, `&l`, `&m`, `&b`
        and `&&` as breaks — `thigh&mid calf` exported as `thigh`, a page break and `id calf`."""
        from core import wm_review
        encounter = self.replayed(self.PROVENANCE_NOTES, self.PROVENANCE_PICKS, self.PROVENANCE_STATEMENTS)
        self.post_page(encounter, edits={'location_11': 'thigh&mid calf'}, review=False)
        fields, requirements = wm_review._mapping(encounter)
        assessed = wm_review._answer_assessment(encounter, self.user, fields, requirements)
        self.assertIn('text_the_form_rewrites', assessed['location_11']['blockers'])
        self.post_page(encounter, edits={'location_11': 'thigh and mid calf'}, review=False)
        assessed = wm_review._answer_assessment(encounter, self.user, fields, requirements)
        self.assertNotIn('text_the_form_rewrites', assessed['location_11']['blockers'])

    def test_the_save_keeps_the_copy_the_next_posted_line_runs_on_to(self):
        """Round 17e's review (false-rejection lens): one line written under both legs with the side only on the line
        after it — the save kept the right leg's copy for a left answer, and an unchanged re-save moved a correct right
        answer onto the left leg's copy."""
        line = 'The great saphenous vein measures 5.0 mm at the saphenofemoral junction without reflux in the'
        notes = 'VENOUS DUPLEX\n\n' + line + '\nright leg.\n\n' + line + '\nleft leg.'
        encounter = self.replayed(notes, {}, [])
        for key, side in (('measurement_15_1', 'left leg.'), ('measurement_1_1', 'right leg.')):
            self.post_page(encounter, edits={key: '5.0 mm'}, sources={key: line + '\n' + side}, review=False)
            cited = [entry['start'] for entry in encounter.answers[key]['evidence']]
            self.assertEqual(cited, [notes.index(line + '\n' + side), notes.index(side, notes.index(line + '\n' + side))])
        self.post_page(encounter, review=False)          # (an unchanged re-save moves neither)
        self.assertEqual(encounter.answers['measurement_1_1']['evidence'][0]['start'], notes.index(line + '\nright leg.'))
        self.assertEqual(encounter.answers['measurement_15_1']['evidence'][0]['start'], notes.index(line + '\nleft leg.'))
