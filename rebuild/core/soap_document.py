"""Frozen, blank built-in SOAP original. Pure construction: no application state or provider call."""
from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

LABELS = ('Subjective', 'Objective', 'Assessment', 'Plan')
ORIGINAL_NAME = 'QuillMD-SOAP-v1.docx'


def original_bytes():
    """Deterministic blank DOCX; ordinary answer cells grow onto later pages without fixed row heights."""
    from docx import Document
    document = Document()
    document.add_heading('SOAP note', level=1)
    table = document.add_table(rows=len(LABELS), cols=2)
    table.style = 'Table Grid'
    for index, label in enumerate(LABELS):
        table.cell(index, 0).text = label
    properties = document.core_properties
    properties.title = 'SOAP note'
    properties.author = 'QuillMD'
    properties.last_modified_by = 'QuillMD'
    properties.created = properties.modified = datetime(2026, 10, 3, tzinfo=timezone.utc)
    source = BytesIO()
    document.save(source)
    output = BytesIO()
    # DOCX writers timestamp ZIP entries at save time. Bind this blank original to one reproducible version instead.
    with ZipFile(source) as package, ZipFile(output, 'w', compression=ZIP_DEFLATED) as stable:
        for name in package.namelist():
            entry = ZipInfo(name, date_time=(2026, 10, 3, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            stable.writestr(entry, package.read(name))
    return output.getvalue()
