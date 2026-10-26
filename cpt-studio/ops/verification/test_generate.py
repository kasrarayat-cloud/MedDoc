"""Tests for the verification documents, their manifest and generate.py.

    cd cpt-studio && .venv/bin/python -m unittest discover -s ops/verification -p 'test_*.py' -v
"""
import copy
import hashlib
import importlib.util
import io
import json
import re
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import docx
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent


def _load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module  # dataclasses resolve their module while the class is built
    spec.loader.exec_module(module)
    return module


generate = _load('cpt_verification_generate', HERE / 'generate.py')
profiles = _load('cpt_verification_profiles', ROOT / 'studio' / 'gi' / 'profiles.py')

MANIFEST_PATH = HERE / 'manifest.json'
MANIFEST = json.loads(MANIFEST_PATH.read_text(encoding='utf-8'))
DOCS = {d['id']: d for d in MANIFEST['documents']}
CODE_TOKEN = re.compile(r'(?<![\w.-])(?:\d{5}|\d{4}[FTU])(?![\w-])')
SERVICE_STATUSES = {'performed', 'incomplete', 'planned', 'historical', 'cancelled', 'negated', 'uncertain'}
CONTEXT_FACTS = {'service_date', 'setting', 'billing_context', 'patient_status'}
REQUIRED_FACTS = {fact for p in profiles.PROFILES.values() for fact in p.required_facts}


def normalized(text):
    return re.sub(r'\s+', ' ', text).strip()


def source_text(doc):
    return (HERE / doc['source']).read_text(encoding='ascii')


def docx_lines(data):
    document = docx.Document(io.BytesIO(data))
    lines = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            lines.append(docx.text.paragraph.Paragraph(child, document).text)
        elif tag == 'tbl':
            for row in docx.table.Table(child, document).rows:
                lines.append(' | '.join(cell.text for cell in row.cells))
    return lines


class DocumentContentTests(unittest.TestCase):
    def test_ten_documents_in_the_requested_delivery_formats(self):
        formats = [DOCS[f'gi-{n:02d}']['delivery']['format'] for n in range(1, 11)]
        self.assertEqual(formats, ['docx', 'pdf', 'scanned-pdf', 'paste', 'txt', 'pdf', 'paste', 'docx', 'pdf', 'txt'])

    def test_document_one_is_the_leads_draft_verbatim(self):
        data = (HERE / DOCS['gi-01']['source']).read_bytes()
        self.assertEqual(hashlib.sha256(data).hexdigest(),
                         '585b1d0a02f6fa9d78657f6b897430f89fd85c67106c6d5c0e4366bae3048265')

    def test_every_source_matches_its_manifest_hash_and_carries_the_banner(self):
        for doc in MANIFEST['documents']:
            with self.subTest(doc=doc['id']):
                data = (HERE / doc['source']).read_bytes()
                self.assertEqual(hashlib.sha256(data).hexdigest(), doc['source_sha256'])
                text = data.decode('ascii')
                self.assertEqual(text.split('\n', 1)[0], generate.BANNER)
                self.assertNotIn('\r', text)
                self.assertTrue(text.endswith('\n'))

    def test_sources_contain_no_code_shaped_tokens(self):
        # Report checks look for code-shaped tokens; the documents must not supply any.
        for doc in MANIFEST['documents']:
            with self.subTest(doc=doc['id']):
                self.assertEqual(CODE_TOKEN.findall(source_text(doc)), [])

    def test_every_named_patient_and_clinician_is_marked_fictional(self):
        for doc in MANIFEST['documents']:
            with self.subTest(doc=doc['id']):
                text = source_text(doc)
                for line in text.splitlines():
                    if line.startswith(('Patient:', 'Endoscopist:', 'Clinician:')):
                        self.assertIn('(fictional)', line)

    def test_missing_information_note_documents_no_date_setting_status_or_basis(self):
        body = source_text(DOCS['gi-07']).split('\n', 1)[1].lower()
        self.assertIsNone(re.search(r'\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{2,4}', body))
        for cue in ('office', 'clinic', 'hospital', 'telehealth', 'established', 'new patient', 'minutes',
                    'medical decision', 'mdm', 'today', 'visit', 'outpatient', 'inpatient'):
            self.assertIsNone(re.search(rf'\b{cue}\b', body), cue)

    def test_planned_and_historical_services_are_written_as_such(self):
        self.assertIn('is scheduled for 2026-11-04', source_text(DOCS['gi-01']))
        self.assertIn('No endoscopy or other procedure was performed today.', source_text(DOCS['gi-01']))
        five = source_text(DOCS['gi-05'])
        self.assertIn('A colonoscopy in May 2021', five)
        self.assertIn('No colonoscopy or other procedure was performed today.', five)
        self.assertIn('The cecum was not reached.', source_text(DOCS['gi-06']))


class ManifestTests(unittest.TestCase):
    def test_author_expectations_are_never_adjudicated(self):
        self.assertIn('not independent adjudication', MANIFEST['author_expectations'].lower())
        for doc in MANIFEST['documents']:
            with self.subTest(doc=doc['id']):
                expected = doc['author_expected_cpt']
                self.assertIs(expected['adjudicated'], False)
                self.assertTrue(expected['note'])

    def test_family_slugs_exist_in_the_gi_profiles(self):
        for doc in MANIFEST['documents']:
            routing = doc['expected_routing']
            slugs = set(routing['families']) | set(routing['not_performed'])
            slugs |= {item['family'] for item in doc['must_not_code'] + doc['incomplete_services']}
            if doc['review_plan']['reject']:
                slugs.add(doc['review_plan']['reject']['family'])
            with self.subTest(doc=doc['id']):
                self.assertLessEqual(slugs, set(profiles.PROFILES))

    def test_routing_expectations_use_the_contract_vocabulary(self):
        for doc in MANIFEST['documents']:
            for slug, family in doc['expected_routing']['families'].items():
                with self.subTest(doc=doc['id'], family=slug):
                    self.assertIn(family['presence'], {'required', 'optional'})
                    self.assertLessEqual(set(family['status_in']), SERVICE_STATUSES)
                    self.assertLessEqual(set(family.get('facts', {})), REQUIRED_FACTS)
                    self.assertLessEqual(set(family.get('missing_includes', []) + family.get('missing_excludes', [])),
                                         REQUIRED_FACTS)

    def test_a_family_expected_performed_is_never_also_listed_not_performed(self):
        for doc in MANIFEST['documents']:
            routing = doc['expected_routing']
            for slug in routing['not_performed']:
                family = routing['families'].get(slug)
                if family:
                    with self.subTest(doc=doc['id'], family=slug):
                        self.assertFalse({'performed', 'incomplete'} & set(family['status_in']))

    def test_must_not_code_families_are_not_performed_or_are_held_for_clarification(self):
        for doc in MANIFEST['documents']:
            for item in doc['must_not_code']:
                family = doc['expected_routing']['families'].get(item['family'])
                with self.subTest(doc=doc['id'], family=item['family']):
                    self.assertTrue(item['reason'])
                    if family and 'performed' in family['status_in']:
                        self.assertTrue(doc['expected_clarifications'])

    def test_context_and_clarification_facts_use_known_names(self):
        for doc in MANIFEST['documents']:
            with self.subTest(doc=doc['id']):
                self.assertLessEqual(set(doc['expected_context']), CONTEXT_FACTS)
                facts = {item['fact'] for item in doc['expected_clarifications']}
                self.assertLessEqual(facts, CONTEXT_FACTS | REQUIRED_FACTS)

    def test_missing_information_note_expects_the_four_clarifications(self):
        facts = {item['fact'] for item in DOCS['gi-07']['expected_clarifications']}
        self.assertEqual(facts, {'service_date', 'setting', 'patient_status', 'em_basis'})
        self.assertEqual(DOCS['gi-07']['upload']['service_date'], None)

    def test_uploads_never_supply_a_service_date(self):
        for doc in MANIFEST['documents']:
            with self.subTest(doc=doc['id']):
                self.assertIsNone(doc['upload']['service_date'])

    def test_designated_rejections_name_a_profiled_family(self):
        rejecting = {d['id']: d['review_plan']['reject'] for d in MANIFEST['documents'] if d['review_plan']['reject']}
        self.assertEqual(set(rejecting), {'gi-03', 'gi-04', 'gi-08'})
        for plan in rejecting.values():
            self.assertTrue(plan['prefer_codes'] and plan['reason'])


class ClassifyAndWrapTests(unittest.TestCase):
    def test_classify_recognises_banner_title_headings_paragraphs_and_tables(self):
        text = (f'{generate.BANNER}\n\nPractice (fictional)\nReport\n\nPatient: A (fictional)\n\nFindings\n'
                'Normal mucosa.\n\nSpecimens\nJar | Site\nA | Rectum\n')
        kinds = [kind for kind, _ in generate.classify(text)]
        self.assertEqual(kinds, ['banner', 'blank', 'title', 'para', 'blank', 'para', 'blank', 'heading', 'para',
                                 'blank', 'heading', 'table'])
        self.assertEqual(generate.classify(text)[-1][1], [['Jar', 'Site'], ['A', 'Rectum']])

    def test_a_line_with_a_colon_or_full_stop_is_never_a_heading(self):
        text = f'{generate.BANNER}\n\nPlan: continue.\nNext line\n\nNot a heading.\nNext line\n'
        self.assertNotIn('heading', [kind for kind, _ in generate.classify(text)])
        self.assertNotIn('title', [kind for kind, _ in generate.classify(text)])

    def test_ragged_table_is_refused(self):
        with self.assertRaises(generate.ManifestError):
            generate.classify(f'{generate.BANNER}\nA | B\nC | D | E\n')

    def test_wrap_fits_the_width_and_keeps_an_overlong_word_whole(self):
        lines = generate.wrap('aaa bbb ccc dddddddddd e', 7, len)
        self.assertEqual(lines, ['aaa bbb', 'ccc', 'dddddddddd', 'e'])
        self.assertEqual(' '.join(lines), 'aaa bbb ccc dddddddddd e')


class ManifestValidationTests(unittest.TestCase):
    def write_manifest(self, folder, mutate):
        manifest = copy.deepcopy(MANIFEST)
        mutate(manifest)
        (folder / 'documents').mkdir(exist_ok=True)
        for doc in MANIFEST['documents']:
            (folder / doc['source']).write_bytes((HERE / doc['source']).read_bytes())
        path = folder / 'manifest.json'
        path.write_text(json.dumps(manifest), encoding='utf-8')
        return path

    def assert_refused(self, mutate, message, *, edit_source=None):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / 'verification'
            folder.mkdir()
            path = self.write_manifest(folder, mutate)
            if edit_source:
                edit_source(folder)
            with self.assertRaisesRegex(generate.ManifestError, message):
                generate.generate(Path(tmp) / 'out', path)

    def test_duplicate_id_is_refused(self):
        self.assert_refused(lambda m: m['documents'][1].update(id='gi-01'), 'duplicate document id')

    def test_unknown_format_is_refused(self):
        self.assert_refused(lambda m: m['documents'][0]['delivery'].update(format='rtf'), 'unknown delivery format')

    def test_extension_must_match_the_format(self):
        self.assert_refused(lambda m: m['documents'][1]['delivery'].update(file='gi-02.docx'), 'must be a unique bare name')

    def test_path_in_delivery_file_is_refused(self):
        self.assert_refused(lambda m: m['documents'][1]['delivery'].update(file='../gi-02.pdf'), 'must be a unique bare name')

    def test_source_outside_the_manifest_directory_is_refused(self):
        self.assert_refused(lambda m: m['documents'][0].update(source='../outside.txt'), 'outside the manifest directory')

    def test_changed_source_is_refused_until_the_manifest_hash_is_updated(self):
        def edit(folder):
            path = folder / MANIFEST['documents'][4]['source']
            path.write_bytes(path.read_bytes().replace(b'2026-10-20', b'2026-10-19', 1))
        self.assert_refused(lambda m: None, 'does not match the manifest', edit_source=edit)

    def test_missing_banner_is_refused(self):
        def edit(folder):
            path = folder / MANIFEST['documents'][6]['source']
            path.write_bytes(path.read_bytes().replace(b'SYNTHETIC TEST DOCUMENT', b'TEST DOCUMENT', 1))

        def rehash(manifest):
            data = (HERE / MANIFEST['documents'][6]['source']).read_bytes().replace(
                b'SYNTHETIC TEST DOCUMENT', b'TEST DOCUMENT', 1)
            manifest['documents'][6]['source_sha256'] = hashlib.sha256(data).hexdigest()
        self.assert_refused(rehash, 'first line must be', edit_source=edit)

    def test_crlf_source_is_refused(self):
        def edit(folder):
            path = folder / MANIFEST['documents'][6]['source']
            path.write_bytes(path.read_bytes().replace(b'\n', b'\r\n'))

        def rehash(manifest):
            data = (HERE / MANIFEST['documents'][6]['source']).read_bytes().replace(b'\n', b'\r\n')
            manifest['documents'][6]['source_sha256'] = hashlib.sha256(data).hexdigest()
        self.assert_refused(rehash, 'LF line endings', edit_source=edit)

    def test_output_inside_the_manifest_directory_is_refused(self):
        with self.assertRaisesRegex(generate.ManifestError, 'refusing to write outputs inside'):
            generate.generate(HERE / 'out', MANIFEST_PATH)

    def test_unknown_only_id_is_refused(self):
        with tempfile.TemporaryDirectory() as tmp, self.assertRaisesRegex(generate.ManifestError, 'unknown document ids'):
            generate.generate(Path(tmp), MANIFEST_PATH, {'gi-99'})


class GeneratedOutputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        base = Path(cls.tmp.name)
        cls.first = base / 'first'
        cls.second = base / 'second'
        cls.results = {row['id']: row for row in generate.generate(cls.first, MANIFEST_PATH)}
        generate.generate(cls.second, MANIFEST_PATH)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def output(self, doc_id):
        return (self.first / DOCS[doc_id]['delivery']['file']).read_bytes()

    def test_outputs_are_byte_identical_across_runs(self):
        names = sorted(p.name for p in self.first.iterdir())
        self.assertEqual(names, sorted(p.name for p in self.second.iterdir()))
        for name in names:
            with self.subTest(file=name):
                self.assertEqual((self.first / name).read_bytes(), (self.second / name).read_bytes())

    def test_index_lists_every_output_with_its_real_hash(self):
        index = json.loads((self.first / 'generated.json').read_text(encoding='utf-8'))
        self.assertEqual(index['manifest_sha256'], hashlib.sha256(MANIFEST_PATH.read_bytes()).hexdigest())
        self.assertEqual(len(index['documents']), 10)
        for row in index['documents']:
            with self.subTest(doc=row['id']):
                data = (self.first / row['file']).read_bytes()
                self.assertEqual(row['sha256'], hashlib.sha256(data).hexdigest())
                self.assertEqual(row['size'], len(data))

    def test_text_and_paste_outputs_are_the_source_bytes(self):
        for doc in MANIFEST['documents']:
            if doc['delivery']['format'] in {'txt', 'paste'}:
                with self.subTest(doc=doc['id']):
                    self.assertEqual(self.output(doc['id']), (HERE / doc['source']).read_bytes())

    def test_pdf_text_layer_reproduces_the_source(self):
        for doc in MANIFEST['documents']:
            if doc['delivery']['format'] == 'pdf':
                with self.subTest(doc=doc['id']):
                    reader = PdfReader(io.BytesIO(self.output(doc['id'])))
                    text = '\n'.join(page.extract_text() for page in reader.pages)
                    self.assertEqual(normalized(text), normalized(source_text(doc)))
                    self.assertEqual(reader.metadata['/Subject'], generate.BANNER)

    def test_scanned_pdf_is_image_only(self):
        reader = PdfReader(io.BytesIO(self.output('gi-03')))
        self.assertGreaterEqual(len(reader.pages), 1)
        for page in reader.pages:
            resources = page['/Resources']
            self.assertNotIn('/Font', resources)
            self.assertEqual(page.extract_text(), '')
            images = list(resources['/XObject'].values())
            self.assertEqual(len(images), 1)
            image = images[0].get_object()
            self.assertEqual((image['/Width'], image['/Height']), (2550, 3300))
            self.assertEqual(image['/ColorSpace'], '/DeviceGray')
            self.assertNotIn(b'BT', page.get_contents().get_data())
        self.assertEqual([round(float(v)) for v in reader.pages[0].mediabox], [0, 0, 612, 792])

    def test_docx_keeps_every_source_line_in_order_with_headings_and_a_table(self):
        for doc_id, tables in (('gi-01', 0), ('gi-08', 1)):
            with self.subTest(doc=doc_id):
                data = self.output(doc_id)
                expected = [line for line in source_text(DOCS[doc_id]).split('\n') if line.strip()]
                self.assertEqual(docx_lines(data), expected)
                document = docx.Document(io.BytesIO(data))
                self.assertEqual(len(document.tables), tables)
                headings = [p.text for p in document.paragraphs if p.style.name.startswith('Heading')]
                self.assertIn('Plan', headings)
                self.assertEqual(document.core_properties.author, generate.AUTHOR)
        table = docx.Document(io.BytesIO(self.output('gi-08'))).tables[0]
        self.assertEqual([c.text for c in table.rows[0].cells], ['Jar', 'Site', 'Method'])

    def test_docx_members_carry_fixed_timestamps(self):
        with zipfile.ZipFile(io.BytesIO(self.output('gi-01'))) as archive:
            self.assertEqual({info.date_time for info in archive.infolist()}, {generate.ZIP_TIME})
            self.assertEqual(archive.infolist()[0].filename, '[Content_Types].xml')


if __name__ == '__main__':
    unittest.main()
