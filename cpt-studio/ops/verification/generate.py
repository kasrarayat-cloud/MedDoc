"""Render the synthetic verification documents into the formats the browser verifier uploads.

    .venv/bin/python ops/verification/generate.py OUT_DIR [--manifest PATH] [--only ID,ID]

Each manifest document is rendered from its source text (``documents/*.txt``):

  txt          the source bytes unchanged, uploaded as a .txt file
  paste        the source bytes unchanged, pasted into the text field
  pdf          a PDF with a real text layer (reportlab, Helvetica)
  scanned-pdf  an image-only PDF (Pillow): every page is one 300 dpi grayscale
               image with no font and no text layer, so extraction has to OCR it
  docx         a Word document: headings, paragraphs, and a table for runs of
               ' | '-delimited lines (so extracted cells re-join to the source)

Output is deterministic: no wall-clock time, random identifier or host font
reaches a file, so the same sources give byte-identical outputs. The sha256 of
every output is printed and written to OUT_DIR/generated.json.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

GENERATOR_VERSION = 'cpt-studio-verification-generate@1'
BANNER = 'SYNTHETIC TEST DOCUMENT - FICTIONAL PATIENT - NOT FOR CLINICAL USE'
AUTHOR = 'CPT Studio synthetic verification (fictional)'
FORMATS = {'txt': '.txt', 'paste': '.paste.txt', 'pdf': '.pdf', 'scanned-pdf': '.pdf', 'docx': '.docx'}
FIXED_TIME = datetime(2026, 10, 23, 0, 0, 0)
ZIP_TIME = (2026, 10, 23, 0, 0, 0)
HERE = Path(__file__).resolve().parent
TABLE_SEPARATOR = ' | '
MAX_HEADING_CHARS = 72

PAGE_W, PAGE_H = letter
MARGIN_PT = 54
# kind -> (font, size pt, leading pt)
PDF_STYLES = {'banner': ('Helvetica-Bold', 9, 14), 'title': ('Helvetica-Bold', 13, 19),
              'heading': ('Helvetica-Bold', 11, 17), 'para': ('Helvetica', 10, 13.5),
              'row': ('Helvetica', 10, 13.5)}
SCAN_DPI = 300
SCAN_JPEG_QUALITY = 90  # grayscale JPEG, as an office scanner writes it
SCAN_MARGIN_PX = 225
# kind -> (font size px, line height px, stroke for a bold face)
SCAN_STYLES = {'banner': (34, 50, 1), 'title': (52, 72, 1), 'heading': (44, 62, 1), 'para': (40, 54, 0),
               'row': (40, 54, 0)}


class ManifestError(Exception):
    """The manifest or a source document does not satisfy the generator's contract."""


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


# --- structure ----------------------------------------------------------------

def classify(text: str) -> list[tuple[str, object]]:
    """Source lines -> blocks: banner | title | heading | para | table | blank.

    A heading is a short line between a blank line and a non-blank line with no
    colon and no closing full stop; the first heading is the document title.
    Consecutive lines containing ' | ' form one table (a list of cell lists).
    """
    lines = text.split('\n')
    if lines and lines[-1] == '':
        lines.pop()
    blocks: list[tuple[str, object]] = []
    seen_title = False
    for index, line in enumerate(lines):
        if index == 0:
            blocks.append(('banner', line))
        elif not line.strip():
            blocks.append(('blank', ''))
        elif TABLE_SEPARATOR in line:
            cells = line.split(TABLE_SEPARATOR)
            if blocks and blocks[-1][0] == 'table':
                blocks[-1][1].append(cells)
            else:
                blocks.append(('table', [cells]))
        elif _is_heading(lines, index):
            blocks.append(('heading' if seen_title else 'title', line))
            seen_title = True
        else:
            blocks.append(('para', line))
    for kind, payload in blocks:
        if kind == 'table' and len({len(row) for row in payload}) != 1:
            raise ManifestError(f'table rows have different cell counts: {payload!r}')
    return blocks


def _is_heading(lines: list[str], index: int) -> bool:
    line = lines[index]
    before = lines[index - 1] if index > 0 else ''
    after = lines[index + 1] if index + 1 < len(lines) else ''
    return (not before.strip() and bool(after.strip()) and len(line) <= MAX_HEADING_CHARS
            and ':' not in line and not line.rstrip().endswith('.') and not line[:1].isdigit())


def wrap(text: str, width: float, measure) -> list[str]:
    """Greedy word wrap on single spaces; a word wider than the line stays whole."""
    words = text.split(' ')
    lines, current = [], ''
    for word in words:
        candidate = f'{current} {word}' if current else word
        if not current or measure(candidate) <= width:
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def layout_lines(blocks) -> list[tuple[str, str]]:
    """Blocks -> (kind, text) visual lines before wrapping; tables become their source lines."""
    out = []
    for kind, payload in blocks:
        if kind == 'table':
            out.extend(('row', TABLE_SEPARATOR.join(cells)) for cells in payload)
        else:
            out.append((kind, payload))
    return out


# --- renderers ----------------------------------------------------------------

def render_pdf(text: str, title: str) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter, invariant=1, pageCompression=1)
    _pdf_metadata(pdf, title)
    width = PAGE_W - 2 * MARGIN_PT
    y = PAGE_H - MARGIN_PT
    for kind, line in layout_lines(classify(text)):
        if kind == 'blank':
            y -= PDF_STYLES['para'][2] * 0.6
            continue
        font, size, leading = PDF_STYLES[kind]
        for piece in wrap(line, width, lambda s, f=font, z=size: stringWidth(s, f, z)):
            if y - leading < MARGIN_PT:
                pdf.showPage()
                y = PAGE_H - MARGIN_PT
            y -= leading
            pdf.setFont(font, size)
            pdf.drawString(MARGIN_PT, y, piece)
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def render_scanned_pdf(text: str, title: str) -> bytes:
    # Pillow's writer, not reportlab: a reportlab page always carries a font
    # resource and an empty text object, which a scanner's page never has.
    pages = _scan_pages(text)
    stamp = FIXED_TIME.timetuple()
    buffer = io.BytesIO()
    pages[0].save(buffer, 'PDF', save_all=True, append_images=pages[1:], resolution=SCAN_DPI,
                  quality=SCAN_JPEG_QUALITY, title=title, author=AUTHOR, subject=BANNER,
                  creator=GENERATOR_VERSION, producer=GENERATOR_VERSION, creationDate=stamp, modDate=stamp)
    return buffer.getvalue()


def _scan_pages(text: str) -> list[Image.Image]:
    # Pillow's bundled scalable face, so no host font file decides the output.
    fonts = {kind: ImageFont.load_default(size=size) for kind, (size, _, _) in SCAN_STYLES.items()}
    width_px, height_px = int(PAGE_W / 72 * SCAN_DPI), int(PAGE_H / 72 * SCAN_DPI)
    usable = width_px - 2 * SCAN_MARGIN_PX
    pages: list[Image.Image] = []

    def new_page():
        image = Image.new('L', (width_px, height_px), 255)
        pages.append(image)
        return ImageDraw.Draw(image), SCAN_MARGIN_PX

    draw, y = new_page()
    for kind, line in layout_lines(classify(text)):
        if kind == 'blank':
            y += SCAN_STYLES['para'][1] // 2
            continue
        _, height, stroke = SCAN_STYLES[kind]
        font = fonts[kind]
        for piece in wrap(line, usable, font.getlength):
            if y + height > height_px - SCAN_MARGIN_PX:
                draw, y = new_page()
            draw.text((SCAN_MARGIN_PX, y), piece, font=font, fill=0, stroke_width=stroke, stroke_fill=0)
            y += height
    return pages


def _pdf_metadata(pdf, title: str) -> None:
    pdf.setTitle(title)
    pdf.setAuthor(AUTHOR)
    pdf.setSubject(BANNER)
    pdf.setCreator(GENERATOR_VERSION)


def render_docx(text: str, title: str) -> bytes:
    document = Document()
    props = document.core_properties
    props.title, props.subject, props.comments = title, BANNER, 'Synthetic verification document; fictional patient.'
    props.author = props.last_modified_by = AUTHOR
    props.created = props.modified = props.last_printed = FIXED_TIME
    props.revision = 1
    for kind, payload in classify(text):
        if kind == 'banner':
            document.add_paragraph().add_run(payload).bold = True
        elif kind == 'title':
            document.add_heading(payload, level=1)
        elif kind == 'heading':
            document.add_heading(payload, level=2)
        elif kind == 'para':
            document.add_paragraph(payload)
        elif kind == 'table':
            _docx_table(document, payload)
    buffer = io.BytesIO()
    document.save(buffer)
    return normalize_zip(buffer.getvalue())


def _docx_table(document, rows) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Table Grid'
    for r, cells in enumerate(rows):
        for c, value in enumerate(cells):
            cell = table.cell(r, c)
            cell.text = ''
            run = cell.paragraphs[0].add_run(value)
            run.bold = r == 0


def normalize_zip(data: bytes) -> bytes:
    """Rewrite a zip with fixed member timestamps and attributes, keeping member order."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', compression=zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            member = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
            member.compress_type = zipfile.ZIP_DEFLATED
            member.create_system = 3
            member.external_attr = 0o644 << 16
            target.writestr(member, source.read(info.filename))
    return out.getvalue()


RENDERERS = {
    'txt': lambda data, text, title: data,
    'paste': lambda data, text, title: data,
    'pdf': lambda data, text, title: render_pdf(text, title),
    'scanned-pdf': lambda data, text, title: render_scanned_pdf(text, title),
    'docx': lambda data, text, title: render_docx(text, title),
}


# --- manifest -----------------------------------------------------------------

def load_manifest(path: Path) -> dict:
    try:
        manifest = json.loads(path.read_text(encoding='utf-8'))
    except (OSError, ValueError) as exc:
        raise ManifestError(f'cannot read manifest {path}: {exc}') from exc
    documents = manifest.get('documents')
    if not isinstance(documents, list) or not documents:
        raise ManifestError('manifest has no documents')
    ids, files = set(), set()
    for doc in documents:
        doc_id = doc.get('id')
        delivery = doc.get('delivery') or {}
        fmt, name = delivery.get('format'), delivery.get('file')
        if not doc_id or doc_id in ids:
            raise ManifestError(f'missing or duplicate document id: {doc_id!r}')
        if fmt not in FORMATS:
            raise ManifestError(f'{doc_id}: unknown delivery format {fmt!r}')
        if not name or Path(name).name != name or not name.endswith(FORMATS[fmt]) or name in files:
            raise ManifestError(f'{doc_id}: delivery file {name!r} must be a unique bare name ending {FORMATS[fmt]}')
        ids.add(doc_id)
        files.add(name)
    return manifest


def read_source(manifest_dir: Path, doc: dict) -> tuple[bytes, str]:
    source = (manifest_dir / doc['source']).resolve()
    if manifest_dir not in source.parents:
        raise ManifestError(f'{doc["id"]}: source {doc["source"]!r} is outside the manifest directory')
    data = source.read_bytes()
    if sha256_bytes(data) != doc.get('source_sha256'):
        raise ManifestError(f'{doc["id"]}: {doc["source"]} sha256 {sha256_bytes(data)} does not match the '
                            f'manifest ({doc.get("source_sha256")}); update the manifest deliberately')
    try:
        text = data.decode('ascii')
    except UnicodeDecodeError as exc:
        raise ManifestError(f'{doc["id"]}: source must be ASCII ({exc})') from exc
    if '\r' in text or not text.endswith('\n'):
        raise ManifestError(f'{doc["id"]}: source must use LF line endings and end with a newline')
    if text.split('\n', 1)[0] != BANNER:
        raise ManifestError(f'{doc["id"]}: first line must be {BANNER!r}')
    return data, text


def generate(out_dir: Path, manifest_path: Path, only: set[str] | None = None) -> list[dict]:
    manifest_path = manifest_path.resolve()
    manifest = load_manifest(manifest_path)
    out_dir = out_dir.resolve()
    if out_dir == manifest_path.parent or manifest_path.parent in out_dir.parents:
        raise ManifestError(f'refusing to write outputs inside {manifest_path.parent}')
    selected = [d for d in manifest['documents'] if only is None or d['id'] in only]
    unknown = (only or set()) - {d['id'] for d in manifest['documents']}
    if unknown:
        raise ManifestError(f'unknown document ids: {sorted(unknown)}')
    out_dir.mkdir(parents=True, exist_ok=True)
    results = []
    for doc in selected:
        data, text = read_source(manifest_path.parent, doc)
        fmt, name = doc['delivery']['format'], doc['delivery']['file']
        rendered = RENDERERS[fmt](data, text, f'{doc["id"]} {doc["title"]} (synthetic)')
        (out_dir / name).write_bytes(rendered)
        results.append({'id': doc['id'], 'format': fmt, 'file': name, 'size': len(rendered),
                        'sha256': sha256_bytes(rendered), 'source': doc['source'],
                        'source_sha256': doc['source_sha256']})
    index = {'generator': GENERATOR_VERSION, 'manifest_sha256': sha256_bytes(manifest_path.read_bytes()),
             'documents': results}
    (out_dir / 'generated.json').write_text(json.dumps(index, indent=2) + '\n', encoding='utf-8')
    return results


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    parser.add_argument('out_dir', type=Path)
    parser.add_argument('--manifest', type=Path, default=HERE / 'manifest.json')
    parser.add_argument('--only', help='comma-separated document ids')
    args = parser.parse_args(argv)
    only = {item.strip() for item in args.only.split(',') if item.strip()} if args.only else None
    try:
        results = generate(args.out_dir, args.manifest, only)
    except ManifestError as exc:
        print(f'generate: {exc}', file=sys.stderr)
        return 2
    for row in results:
        print(f'{row["sha256"]}  {row["size"]:>8}  {row["format"]:<11}  {row["id"]}  {row["file"]}')
    print(f'{len(results)} file(s) written to {args.out_dir.resolve()} (index: generated.json)')
    return 0


if __name__ == '__main__':
    sys.exit(main())
