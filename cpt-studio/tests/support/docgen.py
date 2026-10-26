"""Synthetic document builders for tests. Synthetic text only — never patient data.

make_pdf          real text layer (reportlab), deterministic bytes
make_scanned_pdf  page images only, no text layer (Pillow)
make_docx         paragraphs and tables (python-docx), optional header text
make_encrypted_pdf  RC4-128 (no crypto library needed); '' opens without a password
concat_pdfs       pages of several PDFs in order (e.g. a text page then a scanned one)
make_docx_from_body / rewrite_zip  raw WordprocessingML and package surgery
"""
from __future__ import annotations

import io
import textwrap
import zipfile

import docx
import pypdf
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas

PDF_LINE_CHARS = 90
PDF_LINE_HEIGHT_PT = 14
PDF_TOP_PT = 720
PDF_BOTTOM_PT = 72
SCAN_LINE_CHARS = 50

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
BODY_NAMESPACES = {
    'w': W_NS,
    'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
    'wps': 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
    'v': 'urn:schemas-microsoft-com:vml',
    'w14': 'http://schemas.microsoft.com/office/word/2010/wordml',
    'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}


def _wrapped(text: str, width: int) -> list[str]:
    lines = []
    for paragraph in text.split('\n'):
        lines.extend(textwrap.wrap(paragraph, width) or [''])
    return lines


def make_pdf(pages: list[str], *, image_px: tuple[int, int] | None = None,
             page_size_pt: tuple[float, float] = letter) -> bytes:
    """One page per string (Helvetica 11 pt, wrapped at 90 characters); '' is a page with no text.

    ``image_px`` adds a blank image of that many pixels (width, height) below
    the text of every page.
    """
    buffer = io.BytesIO()
    canvas = Canvas(buffer, pagesize=page_size_pt, invariant=1)
    for text in pages:
        lines = _wrapped(text, PDF_LINE_CHARS)
        if PDF_TOP_PT - PDF_LINE_HEIGHT_PT * len(lines) < PDF_BOTTOM_PT:
            raise ValueError('The page text does not fit on one page.')
        canvas.setFont('Helvetica', 11)
        for number, line in enumerate(lines):
            if line:
                canvas.drawString(72, PDF_TOP_PT - PDF_LINE_HEIGHT_PT * number, line)
        if image_px:
            canvas.drawImage(ImageReader(Image.new('L', image_px, 255)), 72, PDF_BOTTOM_PT, width=288, height=216)
        canvas.showPage()
    canvas.save()
    return buffer.getvalue()


def make_scanned_pdf(pages: list[str], dpi: int = 200) -> bytes:
    """Each page's text drawn into a greyscale letter-size image; the PDF holds only the images."""
    images = [_page_image(text, dpi) for text in pages]
    buffer = io.BytesIO()
    images[0].save(buffer, format='PDF', save_all=True, append_images=images[1:], resolution=float(dpi))
    return buffer.getvalue()


def _page_image(text: str, dpi: int) -> Image.Image:
    image = Image.new('L', (int(8.5 * dpi), int(11 * dpi)), 255)
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=dpi // 5)
    for number, line in enumerate(_wrapped(text, SCAN_LINE_CHARS)):
        draw.text((dpi, dpi + number * (dpi // 4)), line, fill=0, font=font)
    return image


def make_docx(blocks: list[str | list[list[str]]], header_text: str = '') -> bytes:
    """Strings become paragraphs; a list of rows (lists of cell strings) becomes a table."""
    document = docx.Document()
    for block in blocks:
        if isinstance(block, str):
            document.add_paragraph(block)
            continue
        table = document.add_table(rows=len(block), cols=max(len(row) for row in block))
        for row, values in zip(table.rows, block):
            for cell, value in zip(row.cells, values):
                cell.text = value
    if header_text:
        document.sections[0].header.paragraphs[0].text = header_text
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_encrypted_pdf(text: str, password: str, owner_password: str = 'synthetic-owner-password') -> bytes:
    """A one-page PDF encrypted with RC4-128."""
    writer = pypdf.PdfWriter(clone_from=pypdf.PdfReader(io.BytesIO(make_pdf([text]))))
    writer.encrypt(user_password=password, owner_password=owner_password, algorithm='RC4-128')
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def concat_pdfs(documents: list[bytes]) -> bytes:
    writer = pypdf.PdfWriter()
    for data in documents:
        writer.append(pypdf.PdfReader(io.BytesIO(data)))
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def rewrite_zip(data: bytes, replacements: dict[str, bytes | None]) -> bytes:
    """Copy a zip package, replacing (bytes), removing (None) or adding members."""
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            if name not in replacements:
                target.writestr(name, source.read(name))
        for name, content in replacements.items():
            if content is not None:
                target.writestr(name, content)
    return buffer.getvalue()


def make_docx_from_body(body_xml: str, header_text: str = '') -> bytes:
    """A DOCX whose body is ``body_xml``: WordprocessingML elements using the
    prefixes in BODY_NAMESPACES (the section properties are kept)."""
    data = make_docx([], header_text=header_text)
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read('word/document.xml'))
    body = root.find(f'{{{W_NS}}}body')
    declarations = ' '.join(f'xmlns:{prefix}="{uri}"' for prefix, uri in BODY_NAMESPACES.items())
    fragment = etree.fromstring(f'<w:fragment {declarations}>{body_xml}</w:fragment>')
    for position, element in enumerate(list(fragment)):  # insert() moves the element out of the fragment
        body.insert(position, element)
    return rewrite_zip(data, {'word/document.xml': etree.tostring(root, xml_declaration=True, encoding='UTF-8',
                                                                  standalone=True)})
