"""Regressions for refutation round 3 over GI section structure (round "3b", studio/gi/sections.py).

Each class reproduces one confirmed round-3 finding (routing:N or contract:N in the round-3 triage) from its
exact input, adds the wordings its verifier listed, and pins positive controls, in both directions:

- over-correction: an ordinary, clearly performed procedure is excluded or held;
- under-correction: a service that was not performed still becomes a ready code.

Routing runs for real, for the document's own service date and through the JSON round trip the pipeline
applies; the real validation engine assesses provider-shaped selections against an empty in-memory reference.
Synthetic text only.
"""
import json
import time
import unittest
from datetime import date

from studio.gi.context import extract_context
from studio.gi.detect import route, status_for_span
from studio.gi.sections import layout
from studio.reference.store import InMemoryReferenceLookup
from studio.types import (EvidenceSpan, NormalizedResponse, PageSpan, ProviderCode, RoutingResult, from_jsonable,
                          to_jsonable)
from studio.validation.engine import assess

HEADER = 'SYNTHETIC TEST DOCUMENT - FICTIONAL PATIENT\nDate of service: 2026-10-21\n'
NOT_TODAY = ('planned', 'historical', 'cancelled', 'negated')
POSITIVE = ('performed', 'incomplete')


def span(text, phrase, occurrence=0):
    start = -1
    for _ in range(occurrence + 1):
        start = text.index(phrase, start + 1)
    return start, start + len(phrase)


def service_date(text):
    fact = extract_context(text)['service_date']
    return date.fromisoformat(fact.value) if fact.origin == 'document' and fact.value else None


def routed(text):
    """Routing as the pipeline stores it: for the document's own date, through JSON and back."""
    result = route(text, service_date(text))
    return from_jsonable(RoutingResult, json.loads(json.dumps(to_jsonable(result))))


def read(text, phrase, occurrence=0):
    return status_for_span(text, *span(text, phrase, occurrence), service_date(text))


def family(text, slug):
    finding = routed(text).finding(slug)
    return finding.status if finding else None


def fact_value(text, slug, name):
    finding = routed(text).finding(slug)
    return next((f.value for f in (finding.facts if finding else ()) if f.name == name), None)


def run_engine(text, picks):
    """picks: (code, evidence phrase). The proposals by code."""
    codes = []
    for index, (code, phrase) in enumerate(picks):
        start, end = span(text, phrase)
        codes.append(ProviderCode(role='selected', system='cpt', code=code, display='UNVERIFIED',
                                  evidence=(EvidenceSpan(0, start, end, text[start:end], True),), index=index))
    result = assess(text=text, pages=(PageSpan(1, 0, len(text)),), context=extract_context(text), routing=routed(text),
                    normalized=NormalizedResponse('replay', 'test', codes=tuple(codes)),
                    reference=InMemoryReferenceLookup((), {}), service_date=service_date(text),
                    provider_status='succeeded', extraction_status='ok')
    return {p.code: p for p in result.proposals}


def ready(proposal):
    return (proposal.bucket, proposal.status) == ('selected', 'ready')


def clarified(proposal):
    return [c.fact for c in proposal.clarifications]


def starts_a_sentence(text, phrase):
    start = text.index(phrase)
    return layout(text).sentence_for(start, start + 1)[0] == start


# --- routing:1 / contract:1 / routing:11: one wrap rule, pulled both ways ------------------------------------

class Routing1AnsweredFieldTests(unittest.TestCase):
    """routing:1 (a regression from round 2b status:16) - a line that answers a field ("Schatzki ring present: no")
    never runs on into the next line, so its "no" never negates the act written there."""

    B2_BIOPSY = 'Biopsies were taken from the mid esophagus to evaluate for eosinophilic esophagitis.'
    B2 = (HEADER + 'Indication: Dysphagia.\nFindings:\nEsophagus: normal\nSchatzki ring present: no\n' + B2_BIOPSY + '\n'
          'The gastroscope was advanced to the second part of the duodenum.\nProcedure: EGD with biopsy.\n')
    POLYP = 'A 6 mm sigmoid polyp was removed with a cold snare.'
    B3 = (HEADER + 'Procedure: Colonoscopy with cold snare polypectomy and tattoo.\nFindings: The colonoscope was advanced '
          'to the cecum.\nTattoo placed today: no\n' + POLYP + '\n')
    B1 = (HEADER + 'Procedure:\nAnticoagulation held today: no\nColonoscopy with cold snare polypectomy was performed.\n'
          'Findings: The colonoscope was advanced to the cecum. A 6 mm polyp was removed with a cold snare.\n')
    C = (HEADER + 'Procedure: Colonoscopy with cold snare polypectomy.\nTechnique: Sedation was offered but the patient '
         'refused\nColonoscopy was then performed without sedation and the colonoscope was advanced to the cecum.\n'
         'Findings: ' + POLYP + '\n')

    def test_the_claimed_inputs(self):
        self.assertTrue(starts_a_sentence(self.B2, self.B2_BIOPSY))
        self.assertEqual(read(self.B2, self.B2_BIOPSY)[0], 'performed')
        self.assertTrue(ready(run_engine(self.B2, [('43239', self.B2_BIOPSY)])['43239']))
        self.assertEqual(read(self.B3, self.POLYP)[0], 'performed')
        self.assertTrue(ready(run_engine(self.B3, [('45385', self.POLYP)])['45385']))
        self.assertEqual(family(self.B1, 'gi.colonoscopy'), 'performed')
        self.assertTrue(ready(run_engine(self.B1, [('45385', 'A 6 mm polyp was removed with a cold snare.')])['45385']))

    def test_a_refusal_that_ends_its_clause_does_not_deny_the_next_sentence(self):
        """case (c): "the patient refused" ends its sentence when the next line has a verb of its own."""
        self.assertTrue(starts_a_sentence(self.C, 'Colonoscopy was then performed'))
        self.assertEqual(family(self.C, 'gi.colonoscopy'), 'performed')
        self.assertEqual(fact_value(self.C, 'gi.colonoscopy', 'extent'), 'cecum')
        self.assertTrue(ready(run_engine(self.C, [('45385', self.POLYP)])['45385']))

    def test_the_verifier_variants(self):
        egd = (HEADER + 'Procedure: EGD with biopsy.\nIndication: Dysphagia.\nFindings:\n{field}\nBiopsies were taken '
               'from the distal esophagus.\nThe gastroscope was advanced to the second part of the duodenum.\n')
        for field in ('Hiatal hernia present: no', 'Hiatal hernia: no', 'Hiatal hernia: No', 'Hiatal hernia - no',
                      'Hiatal hernia \u2013 no', 'Hiatal hernia? no', 'Anticoagulants held: not', 'Prior EGD: never',
                      'Sedation given: patient declined'):
            text = egd.format(field=field)
            with self.subTest(field=field):
                self.assertTrue(starts_a_sentence(text, 'Biopsies were taken'))
                self.assertTrue(ready(run_engine(text, [('43239', 'Biopsies were taken from the distal esophagus.')])
                                      ['43239']))
                self.assertEqual(fact_value(text, 'gi.egd', 'location'), 'esophagus')
        for field in ('Active bleeding: no', 'Retained stool: no', 'Mass seen: no'):
            text = self.B3.replace('Tattoo placed today: no', field)
            with self.subTest(field=field):
                self.assertTrue(ready(run_engine(text, [('45385', self.POLYP)])['45385']))

    def test_controls_are_unchanged(self):
        egd = (HEADER + 'Procedure: EGD with biopsy.\nFindings:\n{field}\nBiopsies were taken from the distal esophagus.\n'
               'The gastroscope was advanced to the second part of the duodenum.\n')
        for field in ('Hernia: no', 'Bleeding: no', 'Hiatal hernia: none', 'Hiatal hernia: absent', 'Hiatal hernia: no.'):
            text = egd.format(field=field)
            with self.subTest(field=field):
                self.assertTrue(ready(run_engine(text, [('43239', 'Biopsies were taken from the distal esophagus.')])
                                      ['43239']))

    def test_a_field_whose_value_is_cut_prose_still_runs_on(self):
        """A value that is a clause ("because of time constraints, no") is prose cut on "no", not an answer."""
        text = (HEADER + 'Procedure: Colonoscopy\nFindings: The colonoscope was advanced to the cecum.\nComment: because of '
                'time constraints, no\nEUS of the pancreatic cyst was performed today.\n')
        self.assertFalse(starts_a_sentence(text, 'EUS of the pancreatic cyst'))
        self.assertEqual(family(text, 'gi.eus'), 'negated')
        text = text.replace('Comment: because of time constraints, no', 'Comment: patient had no')
        self.assertFalse(starts_a_sentence(text, 'EUS of the pancreatic cyst'))
        self.assertEqual(family(text, 'gi.eus'), 'negated')
        text = (HEADER + 'Procedure: Colonoscopy\nFindings: The cecum was reached; the colonoscopy was performed - no\n'
                'EGD was done.\n')
        self.assertFalse(starts_a_sentence(text, 'EGD was done.'))
        self.assertNotIn(family(text, 'gi.egd'), POSITIVE)


OFFICE = (HEADER + 'Place of service: Office (POS 11)\nEstablished patient office visit.\nChief complaint: Rectal '
          'bleeding.\nMedical decision making: moderate.\nTotal time on the date of the encounter was 32 minutes.\n')


class Contract1AnsweredRowTests(unittest.TestCase):
    """contract:1 (a new defect from round 2b status:16) - an answered field ending in "No", or an option row, never
    runs on into the next line: today's procedure is not negated and the same-day E/M hold is kept."""

    def test_the_claimed_field_lines(self):
        snare = 'Cold snare polypectomy of the 6 mm sigmoid polyp was performed.'
        text = (HEADER + 'Procedure: Colonoscopy\nFindings: The colonoscope was advanced to the cecum. A 6 mm sigmoid polyp '
                'was seen.\nBiopsies obtained: No\n' + snare + '\n')
        self.assertTrue(starts_a_sentence(text, snare))
        self.assertTrue(ready(run_engine(text, [('45385', snare)])['45385']))
        flex = 'Flexible sigmoidoscopy to 30 cm with cold forceps biopsy was performed today.'
        text = OFFICE + 'Procedure note\nAnal fissure on external inspection: no\n' + flex + '\n'
        proposals = run_engine(text, [('99214', 'Established patient office visit.'), ('45331', flex)])
        self.assertTrue(ready(proposals['45331']))
        self.assertEqual(proposals['99214'].status, 'needs_clarification')
        self.assertIn('separate_em', clarified(proposals['99214']))

    def test_the_listed_variants(self):
        flex = (HEADER + 'Procedure: Flexible sigmoidoscopy.\nFindings: The sigmoidoscope was advanced to 40 cm.\nPolyps '
                'found: no\nBiopsy of the sigmoid mucosa was taken.\n')
        self.assertTrue(ready(run_engine(flex, [('45331', 'Biopsy of the sigmoid mucosa was taken.')])['45331']))
        polyp = 'A 5 mm cecal polyp was removed with a cold snare.'
        text = (HEADER + 'Procedure: Colonoscopy\nFindings: The colonoscope was advanced to the cecum.\nTattoo placed: no\n'
                + polyp + '\n')
        self.assertTrue(ready(run_engine(text, [('45385', polyp)])['45385']))

    def test_option_rows(self):
        for rows in ('Biopsy taken \u2610 Yes \u2612 No\nCold snare polypectomy \u2612 Yes \u2610 No',
                     'Biopsy taken [ ] Yes [x] No\nCold snare polypectomy [x] Yes [ ] No'):
            text = HEADER + 'Procedure: Colonoscopy\nFindings: The colonoscope was advanced to the cecum.\n' + rows + '\n'
            with self.subTest(rows=rows):
                self.assertTrue(starts_a_sentence(text, 'Cold snare polypectomy'))
                self.assertTrue(ready(run_engine(text, [('45385', 'Cold snare polypectomy')])['45385']))
        # printed radio rows with the answer only in the form block: the row is held as at 57b4b3747, never excluded
        text = (HEADER + 'Procedure: Colonoscopy\nFindings: The colonoscope was advanced to the cecum.\nBiopsy taken      Yes'
                '      No\nCold snare polypectomy      Yes      No\n')
        self.assertTrue(starts_a_sentence(text, 'Cold snare polypectomy'))
        self.assertNotIn(read(text, 'Cold snare polypectomy')[0], NOT_TODAY)
        self.assertNotEqual(run_engine(text, [('45385', 'Cold snare polypectomy')])['45385'].bucket, 'excluded')


class Routing11WrapAfterAVerbTests(unittest.TestCase):
    """routing:11 - a layout break after a verb that still lacks its object ("did not perform", "did not undergo",
    "will schedule", "we recommend") no longer splits the negation or the plan from the service on the next line:
    the two lines read as the same text on one line."""

    def note(self, cut):
        return (HEADER + 'Procedure: Colonoscopy with cold snare polypectomy.\nFindings: The colonoscope was advanced to the '
                'cecum. A 6 mm sigmoid polyp was removed with a cold snare. The patient asked about the pancreatic cyst; '
                f'because of time constraints {cut}\nEUS of the pancreatic cyst.\n')

    NOT_DONE_CUTS = ('we did not perform', 'she did not undergo', 'we will schedule', 'we recommend', "we didn't perform",
                     'we cannot perform', 'we were unable to perform', 'the patient chose to forgo',
                     'the patient elected not to undergo', 'the patient declined', 'we will', 'it is unclear whether',
                     'we discussed EGD versus')

    def test_the_claimed_and_listed_cuts(self):
        for cut in self.NOT_DONE_CUTS:
            text = self.note(cut)
            with self.subTest(cut=cut):
                self.assertFalse(starts_a_sentence(text, 'EUS of the pancreatic cyst.'))
                self.assertNotIn(family(text, 'gi.eus'), POSITIVE)
                self.assertFalse(ready(run_engine(text, [('43259', 'EUS of the pancreatic cyst.')])['43259']))
                self.assertEqual(family(text, 'gi.colonoscopy'), 'performed')

    def test_a_wrapped_line_reads_as_the_same_text_on_one_line(self):
        """The join adds no wording of its own: whatever the cue words make of the one-line text ("we decided to
        defer EUS" has none), the wrapped text reads the same."""
        for cut in self.NOT_DONE_CUTS + ('we decided to defer', 'the patient asked about'):
            text = self.note(cut)
            one_line = text.replace(cut + '\nEUS', cut + ' EUS')
            with self.subTest(cut=cut):
                self.assertEqual(read(text, 'EUS of the pancreatic cyst.')[0],
                                 read(one_line, 'EUS of the pancreatic cyst.')[0])
                self.assertEqual(family(text, 'gi.eus'), family(one_line, 'gi.eus'))

    def test_a_line_that_ends_its_sentence_stays_apart(self):
        """The last word is a noun or a predicate, or the next line has a verb of its own: the next line reads as it
        does when the line before ends in a period."""
        base = HEADER + 'Procedure: EGD with biopsy.\nFindings: The gastroscope was advanced to the second part of the duodenum.\n'
        for prev, nxt in (('The polyp was removed with the snare', 'EGD was performed.'),
                          ('Retroflexion not performed', 'EGD with biopsy of the antrum.'),
                          ('Repeat colonoscopy recommended', 'EGD with biopsy of the antrum.'),
                          ('Patient referred to surgery', 'EGD with biopsy of the antrum.'),
                          ('Sedation was offered but the patient declined', 'EGD was performed without sedation.'),
                          ('Sedation was offered but the patient declined', 'EGD performed without sedation.'),
                          ('No change compared to prior', 'EGD with biopsy of the antrum.'),
                          ('Colonoscopy was done in May', 'EGD with biopsy of the antrum.'),
                          ('Procedure started at 8 AM', 'EGD with biopsy of the antrum.'),
                          ('Patient with known AS', 'EGD with biopsy of the antrum.'),
                          ('LA grade I esophagitis', 'EGD with biopsy of the antrum.')):
            text = base + prev + '\n' + nxt + '\n'
            with self.subTest(prev=prev, nxt=nxt):
                self.assertTrue(starts_a_sentence(text, nxt))
                self.assertEqual(read(text, nxt), read(base + prev + '.\n' + nxt + '\n', nxt))
                self.assertEqual(read(text, nxt)[0], 'performed')

    def test_a_letterhead_title_and_field_lines_stay_apart(self):
        text = ('SYNTHETIC TEST DOCUMENT - FICTIONAL PATIENT\nHarborview Digestive Health (fictional practice)\n'
                'Established patient visit - office progress note\nDate of service: 2026-10-21\nPlace of service: Office '
                '(POS 11)\nEstablished patient office visit.\n')
        for line in ('Established patient visit', 'Date of service', 'Place of service', 'Established patient office'):
            with self.subTest(line=line):
                self.assertTrue(starts_a_sentence(text, line))

    def test_the_rule_reads_only_the_ends_of_long_lines(self):
        """A 50,000-character line ending on every shape the rule reads lays out in bounded time."""
        for tail in ('Hiatal hernia - no', 'did not perform', 'x: no', 'Biopsy taken \u2610 Yes \u2612 No'):
            text = HEADER + 'Findings: ' + 'quickly ' * 6250 + tail + '\nEUS of the cyst.\n'
            started = time.monotonic()
            layout(text)
            with self.subTest(tail=tail):
                self.assertLess(time.monotonic() - started, 2)


# --- routing:2 / routing:12: colon-less labels and plan headings ---------------------------------------------

def word_document(paragraphs):
    """A real .docx through the real extractor: paragraphs are (text, style or None)."""
    import io

    import docx
    from studio.ingest import Limits
    from studio.ingest import docx as docx_ingest
    document = docx.Document()
    for text, style in paragraphs:
        if style and style.startswith('Heading'):
            document.add_heading(text, level=int(style.split()[-1]))
        else:
            document.add_paragraph(text, style=style)
    buffer = io.BytesIO()
    document.save(buffer)
    result = docx_ingest.extract_docx(buffer.getvalue(), Limits())
    assert result.status == 'ok', result.messages
    return result.text


EMR_LINE = 'Colonoscopy with EMR of the ascending colon polyp at a tertiary center.'


def emr_note(heading):
    return (HEADER + 'Procedure: Colonoscopy\nFindings: The colonoscope was advanced to the cecum. A 30 mm laterally '
            'spreading polyp in the ascending colon was tattooed and not removed.\n' + heading + '\n' + EMR_LINE + '\n')


class Routing2ListItemsAndStatementsTests(unittest.TestCase):
    """routing:2 (a regression from round 2b status:7) - a numbered findings item or a statement ("3. Further workup
    recommended", "Previous treatments failed") is content, never a colon-less plan or history heading: the
    performed polypectomy after it stays a findings sentence and 45385 is ready."""

    POLYP_ITEM = '6 mm sigmoid polyp removed with cold snare.'

    def numbered(self, item):
        return (HEADER + 'Procedure: Colonoscopy with biopsy and cold snare polypectomy.\nFindings:\n1. The colonoscope was '
                'advanced to the cecum.\n2. A 3 cm ascending colon mass, biopsied and tattooed.\n3. ' + item + '\n4. '
                + self.POLYP_ITEM + '\n')

    def prose(self, line):
        return (HEADER + 'Procedure: Colonoscopy with biopsy and cold snare polypectomy.\nFindings: The colonoscope was '
                'advanced to the cecum. A 3 cm mass in the ascending colon was biopsied.\n' + line + '\nA 6 mm sigmoid '
                'polyp was removed with a cold snare.\n')

    STATEMENTS = ('Further workup recommended', 'Further evaluation recommended', 'Surgical management recommended',
                  'Further imaging recommended', 'Follow-up imaging advised', 'Previous treatments failed',
                  'Prior therapy ineffective', 'Prior radiation therapy changes', 'Genetic testing recommended',
                  'History of AS')

    def test_the_claimed_numbered_item_and_its_variants(self):
        for item in self.STATEMENTS + ('Further evaluation', 'Upcoming procedures'):
            text = self.numbered(item)
            with self.subTest(item=item):
                self.assertEqual(layout(text).section_at(text.index(self.POLYP_ITEM)).kind, 'findings')
                self.assertTrue(ready(run_engine(text, [('45385', self.POLYP_ITEM)])['45385']))

    def test_the_claimed_prose_line_and_its_variants(self):
        for line in self.STATEMENTS:
            text = self.prose(line)
            with self.subTest(line=line):
                self.assertEqual(layout(text).section_at(text.index('A 6 mm sigmoid')).kind, 'findings')
                self.assertTrue(ready(run_engine(text, [('45385', 'A 6 mm sigmoid polyp was removed with a cold snare.')])
                                      ['45385']))

    def test_a_word_numbered_list_whose_numbers_extraction_drops(self):
        text = word_document([('SYNTHETIC TEST DOCUMENT - FICTIONAL PATIENT', None), ('Date of service: 2026-10-21', None),
                              ('Procedure: Colonoscopy with biopsy and cold snare polypectomy.', None),
                              ('Findings', 'Heading 2'),
                              ('The colonoscope was advanced to the cecum', 'List Number'),
                              ('A 3 cm ascending colon mass, biopsied and tattooed', 'List Number'),
                              ('Surgical evaluation recommended', 'List Number'),
                              ('A 6 mm sigmoid polyp removed with cold snare', 'List Number')])
        self.assertIn('\nSurgical evaluation recommended\n', text)
        self.assertEqual(layout(text).section_at(text.index('A 6 mm sigmoid')).kind, 'findings')
        self.assertTrue(ready(run_engine(text, [('45385', 'A 6 mm sigmoid polyp removed with cold snare')])['45385']))

    def test_an_impression_list(self):
        text = (HEADER + 'Procedure: Colonoscopy with biopsy and cold snare polypectomy.\nFindings: The colonoscope was '
                'advanced to the cecum. A 6 mm sigmoid polyp was removed with a cold snare.\nImpression:\n1. Ascending '
                'colon mass, biopsied.\n2. Genetic testing recommended\n3. Sigmoid polyp, removed with cold snare\n')
        self.assertEqual(layout(text).section_at(text.index('Sigmoid polyp')).kind, 'impression')
        self.assertNotEqual(run_engine(text, [('45385', 'Sigmoid polyp, removed with cold snare')])['45385'].bucket,
                            'excluded')

    def test_headings_still_open(self):
        """A known heading still opens with a number; an unnumbered label naming a section still opens one."""
        for heading, kind in (('2. Plan', 'plan'), ('3. Recommendations', 'plan'), ('B. Further management', 'plan'),
                              ('Upcoming procedures', 'plan'), ('## Upcoming procedures', 'plan'),
                              ('**Upcoming procedures**', 'plan'), ('Upcoming procedures (pending authorization)', 'plan'),
                              ('Recommended interventions', 'plan'), ('Summary and recommendations', 'plan'),
                              ('Earlier endoscopic work-up', 'history'), ('Previous endoscopic studies', 'history')):
            text = emr_note(heading)
            with self.subTest(heading=heading):
                self.assertEqual(layout(text).section_at(text.index(EMR_LINE)).kind, kind)
                self.assertEqual(run_engine(text, [('45390', EMR_LINE)])['45390'].bucket, 'excluded')

    def test_list_items_and_findings_lines_stay_content(self):
        for heading in ('3. Upcoming procedures', 'b) Further evaluation', '* Upcoming procedures',
                        'Previous anastomosis', 'Prior polypectomy site', 'Previous stent procedures'):
            text = emr_note(heading)
            with self.subTest(heading=heading):
                self.assertEqual(layout(text).section_at(text.index(EMR_LINE)).kind, 'findings')


class Routing12PlanHeadingTests(unittest.TestCase):
    """routing:12 - colon-less plan headings outside the list ("Pending procedures", "Referrals", "Orders", "To do")
    no longer continue the findings section: the planned EMR is excluded and no EMR fact is read."""

    HEADINGS = ('Pending procedures', 'Pending procedure', 'Pending orders', 'Referrals', 'Referral', 'Orders', 'To do',
                'To-do', 'TO DO', 'To do:', 'Referrals:', 'Orders:')

    def assert_planned_emr_excluded(self, text):
        self.assertEqual(layout(text).section_at(text.index(EMR_LINE)).kind, 'plan')
        self.assertEqual(read(text, EMR_LINE)[0], 'planned')
        self.assertNotIn('EMR', fact_value(text, 'gi.colonoscopy', 'interventions') or '')
        self.assertEqual(run_engine(text, [('45390', EMR_LINE)])['45390'].bucket, 'excluded')

    def test_the_claimed_headings(self):
        for heading in self.HEADINGS:
            with self.subTest(heading=heading):
                self.assert_planned_emr_excluded(emr_note(heading))
        for heading in ('Orders:', 'To do:', 'Referral:'):  # with the planned line inline
            with self.subTest(inline=heading):
                self.assert_planned_emr_excluded(emr_note(heading).replace(heading + '\n', heading + ' '))

    def test_a_word_heading_paragraph(self):
        for heading in ('Pending procedures', 'To do'):
            text = word_document([('SYNTHETIC TEST DOCUMENT - FICTIONAL PATIENT', None),
                                  ('Date of service: 2026-10-21', None), ('Procedure: Colonoscopy', None),
                                  ('Findings: The colonoscope was advanced to the cecum. A 30 mm laterally spreading '
                                   'polyp in the ascending colon was tattooed and not removed.', None),
                                  (heading, 'Heading 2'), (EMR_LINE, None)])
            with self.subTest(heading=heading):
                self.assertIn('\n' + heading + '\n', text)
                self.assert_planned_emr_excluded(text)

    def test_content_that_mentions_the_words_is_not_a_heading(self):
        for line in ('Referrals were sent to surgery.', 'Orders placed for CBC.', 'Things to do'):
            text = emr_note(line)
            with self.subTest(line=line):
                self.assertEqual(layout(text).section_at(text.index(EMR_LINE)).kind, 'findings')


# --- contract:8: a bare label with a participle --------------------------------------------------------------

class Contract8PerformedTodayHeadingTests(unittest.TestCase):
    """contract:8 - "Procedures performed today:", "Procedures done today:" and "Interventions performed:" open a
    procedure section like "Procedures performed:" does, so a checked procedure under them no longer inherits the
    previous "Assessment and plan" (excluded as planned, with the E/M released)."""

    FLEX = '☒ Flexible sigmoidoscopy to 30 cm'
    PICKS = [('99214', 'Established patient office visit.'), ('45330', 'Flexible sigmoidoscopy to 30 cm')]
    LABELS = ('Procedures performed today:', 'Procedures done today:', 'Interventions performed:',
              'Procedure performed today:', 'Interventions performed today:')

    def office(self, label, items):
        return (OFFICE.replace('moderate.', 'Moderate.') + 'Assessment and plan: Grade II internal hemorrhoids.\n'
                + label + '\n' + items + '\n')

    def test_the_claimed_labels(self):
        for label in self.LABELS + ('Procedures performed:',):
            text = self.office(label, self.FLEX)
            with self.subTest(label=label):
                self.assertEqual(layout(text).section_at(text.index(self.FLEX)).kind, 'procedure')
                self.assertEqual(family(text, 'gi.flex_sig'), 'performed')
                self.assertEqual(fact_value(text, 'gi.flex_sig', 'extent'), '30 cm')
                proposals = run_engine(text, self.PICKS)
                self.assertTrue(ready(proposals['45330']))
                self.assertEqual(proposals['99214'].status, 'needs_clarification')
                self.assertIn('separate_em', clarified(proposals['99214']))

    def test_other_procedures_listed_under_the_label_hold_the_visit(self):
        items = '☒ Anoscopy\n☒ Rubber band ligation of an internal hemorrhoid x3'
        text = OFFICE + 'Procedures performed today:\n' + items + '\n'
        other = routed(text).finding('gi.other')
        self.assertIsNotNone(other)
        self.assertEqual({m.span.text for m in other.mentions},
                         {'Anoscopy', 'Rubber band ligation of an internal hemorrhoid x3'})
        visit = run_engine(text, [('99214', 'Established patient office visit.')])['99214']
        self.assertEqual(visit.status, 'needs_clarification')

    def test_labels_for_procedures_still_to_be_done_are_plan_headings(self):
        line = 'EGD to the second portion of the duodenum with esophageal balloon dilation.'
        for label in ('Procedures to be done:', 'Procedures to be performed:', 'Procedures that will be performed:',
                      'Procedure to be done'):
            text = (HEADER + 'Procedure: Colonoscopy\nFindings: The colonoscope was advanced to the cecum. A 30 mm '
                    'laterally spreading polyp in the ascending colon was tattooed and not removed.\n' + label + '\n'
                    + line + '\n')
            with self.subTest(label=label):
                self.assertEqual(layout(text).section_at(text.index(line)).kind, 'plan')
                self.assertNotIn(family(text, 'gi.egd'), POSITIVE)
                self.assertEqual(run_engine(text, [('43249', line)])['43249'].bucket, 'excluded')

    def test_a_sentence_ending_in_a_colon_is_still_content(self):
        for line in ('The following procedures were performed today:', 'Findings included:'):
            text = (HEADER + 'Procedure: Colonoscopy\nFindings: The colonoscope was advanced to the cecum.\n' + line +
                    '\nA 6 mm sigmoid polyp was removed with a cold snare.\n')
            with self.subTest(line=line):
                self.assertEqual(layout(text).section_at(text.index('A 6 mm sigmoid')).kind, 'findings')



if __name__ == '__main__':
    unittest.main()
