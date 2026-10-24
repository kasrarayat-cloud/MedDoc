"""Regression tests for the DOCX ingestion repairs (refutation round 1: ingest:5, 6, 9, 10, 11).

Every document is synthetic and built here with python-docx or from raw
WordprocessingML (tests.support.docgen); nothing is read from disk. Each class
names the finding it pins and carries positive controls.
"""
from __future__ import annotations

import io
import unittest
import zipfile

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from studio.ingest import Limits, extract_document
from studio.ingest import docx as docx_ingest
from tests.support import docgen

DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
HIDDEN_ONE = 'Not extracted from the document body: 1 run of hidden text.'


def extract(data: bytes):
    return docx_ingest.extract_docx(data, Limits())


def saved(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def body(xml: str) -> bytes:
    return docgen.make_docx_from_body(xml)


# --- ingest:5 — text hidden through a style --------------------------------------------------

HEADER = ['SYNTHETIC TEST DOCUMENT - FICTIONAL PATIENT - NOT FOR CLINICAL USE', 'Procedure report',
          'Date of service: 2026-10-21']
VISIBLE = 'Procedure: EGD with cold forceps biopsy of the gastric antrum.'
SAMPLE = 'Colonoscopy with snare polypectomy performed.'


def report() -> docx.Document:
    document = docx.Document()
    for line in HEADER + [VISIBLE]:
        document.add_paragraph(line)
    return document


def char_style(document, name: str, hidden: bool | None, based_on: str | None = None, font: str | None = None):
    style = document.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
    if based_on:
        style.base_style = document.styles[based_on]
    if hidden is not None:
        style.font.hidden = hidden
    if font:
        style.font.name = font
    return style


class DocxStyleHiddenTextTests(unittest.TestCase):
    """ingest:5 — w:vanish from a character, paragraph, table or default style hides the text."""

    def assert_hidden(self, data: bytes, sample: str = SAMPLE, count: int = 1):
        result = extract(data)
        self.assertEqual(result.status, 'ok')
        self.assertNotIn(sample, result.text)
        self.assertIn(VISIBLE, result.text)
        noun = 'run' if count == 1 else 'runs'
        self.assertEqual(result.messages, (f'Not extracted from the document body: {count} {noun} of hidden text.',))

    def test_character_style_marked_hidden(self):
        document = report()
        char_style(document, 'TemplateInstruction', hidden=True)
        document.add_paragraph().add_run(SAMPLE, style='TemplateInstruction')
        self.assert_hidden(saved(document))

    def test_paragraph_style_marked_hidden(self):
        document = report()
        document.styles.add_style('HiddenGuidance', WD_STYLE_TYPE.PARAGRAPH).font.hidden = True
        document.add_paragraph(SAMPLE, style='HiddenGuidance')
        self.assert_hidden(saved(document))

    def test_hidden_through_a_based_on_chain(self):
        document = report()
        char_style(document, 'InstrBase', hidden=True)
        char_style(document, 'InstrChild', hidden=None, based_on='InstrBase')
        document.add_paragraph().add_run(SAMPLE, style='InstrChild')
        self.assert_hidden(saved(document))
        paragraph = report()
        base = paragraph.styles.add_style('GuideBase', WD_STYLE_TYPE.PARAGRAPH)
        base.font.hidden = True
        paragraph.styles.add_style('GuideChild', WD_STYLE_TYPE.PARAGRAPH).base_style = base
        paragraph.add_paragraph(SAMPLE, style='GuideChild')
        self.assert_hidden(saved(paragraph))

    def test_runs_nested_in_hyperlinks_and_controls_inherit_the_paragraph_style(self):
        document = report()
        document.styles.add_style('HiddenGuidance', WD_STYLE_TYPE.PARAGRAPH).font.hidden = True
        paragraph = document.add_paragraph(style='HiddenGuidance')
        sdt = etree.fromstring(
            f'<w:sdt xmlns:w="{docgen.W_NS}"><w:sdtContent><w:hyperlink w:anchor="x"><w:r><w:t>{SAMPLE}</w:t>'
            '</w:r></w:hyperlink></w:sdtContent></w:sdt>')
        paragraph._p.append(sdt)
        self.assert_hidden(saved(document))

    def test_default_paragraph_style_and_document_defaults(self):
        document = docx.Document()
        document.styles['Normal'].font.hidden = True
        document.add_paragraph(SAMPLE)
        result = extract(saved(document))
        self.assertEqual((result.status, result.messages),
                         ('failed', ('The Word document is empty.', HIDDEN_ONE)))
        defaults = docx.Document()
        rpr = defaults.styles.element.find(f'{qn("w:docDefaults")}/{qn("w:rPrDefault")}/{qn("w:rPr")}')
        rpr.append(OxmlElement('w:vanish'))
        defaults.add_paragraph(SAMPLE)
        self.assertEqual(extract(saved(defaults)).messages, ('The Word document is empty.', HIDDEN_ONE))

    def test_table_style_marked_hidden(self):
        document = report()
        document.styles.add_style('HiddenTable', WD_STYLE_TYPE.TABLE).font.hidden = True
        table = document.add_table(rows=1, cols=2)
        table.style = 'HiddenTable'
        table.cell(0, 0).text, table.cell(0, 1).text = 'Procedure', SAMPLE
        result = extract(saved(document))
        self.assertNotIn(SAMPLE, result.text)
        self.assertEqual(result.messages, ('Not extracted from the document body: 2 runs of hidden text.',))

    def test_controls_visible_text_is_kept(self):
        plain = report()
        char_style(plain, 'PlainNote', hidden=None)
        plain.add_paragraph().add_run(SAMPLE, style='PlainNote')
        unhidden = report()
        char_style(unhidden, 'TemplateInstruction', hidden=True)
        unhidden.add_paragraph().add_run(SAMPLE, style='TemplateInstruction').font.hidden = False
        child = report()
        base = child.styles.add_style('GuideBase', WD_STYLE_TYPE.PARAGRAPH)
        base.font.hidden = True
        visible_child = child.styles.add_style('GuideChild', WD_STYLE_TYPE.PARAGRAPH)
        visible_child.base_style = base
        visible_child.font.hidden = False
        child.add_paragraph(SAMPLE, style='GuideChild')
        for name, document in (('plain style', plain), ('direct val=0 over a hidden style', unhidden),
                               ('child style un-hides its base', child)):
            with self.subTest(name):
                result = extract(saved(document))
                self.assertEqual((result.status, result.messages), ('ok', ()))
                self.assertTrue(result.text.endswith(SAMPLE))

    def test_direct_vanish_still_counts_once(self):
        document = report()
        run = document.add_paragraph().add_run(SAMPLE)
        run.font.hidden = True
        self.assert_hidden(saved(document))

    def test_template_sample_paragraph_through_the_child_process(self):
        sample = ('Colonoscopy: The colonoscope was advanced to the cecum. A 7 mm sessile polyp in the sigmoid colon '
                  'was removed with a cold snare.')
        document = report()
        document.styles.add_style('TemplateSample', WD_STYLE_TYPE.PARAGRAPH).font.hidden = True
        document.add_paragraph(sample, style='TemplateSample')
        result = extract_document(saved(document), 'egd.docx', DOCX, ocr=None)
        self.assertEqual((result.status, result.messages), ('ok', (HIDDEN_ONE,)))
        self.assertNotIn('cold snare', result.text)


# --- ingest:6 — showingPlcHdr is an on/off flag ------------------------------------------------

PROCEDURE = 'Colonoscopy with cold snare polypectomy (sigmoid, 7 mm)'
FINDINGS = 'Findings: three polyps removed with cold snare.'
PLACEHOLDER_ONE = 'Not extracted from the document body: 1 empty content control (placeholder text).'


def sdt_pr(flag: str) -> str:
    return f'<w:sdtPr><w:alias w:val="Procedure"/><w:tag w:val="proc"/><w:id w:val="1234"/>{flag}</w:sdtPr>'


def inline_control(flag: str, content: str = PROCEDURE) -> str:
    return (f'<w:p><w:r><w:t xml:space="preserve">Procedure performed: </w:t></w:r><w:sdt>{sdt_pr(flag)}'
            f'<w:sdtContent><w:r><w:t>{content}</w:t></w:r></w:sdtContent></w:sdt></w:p>')


class DocxPlaceholderFlagTests(unittest.TestCase):
    """ingest:6 — a control whose showingPlcHdr is off holds real content."""

    def test_flag_off_keeps_inline_content(self):
        for value in ('false', '0', 'off', ' false '):
            with self.subTest(value=value):
                result = extract(body(inline_control(f'<w:showingPlcHdr w:val="{value}"/>')))
                self.assertEqual((result.status, result.text, result.messages),
                                 ('ok', f'Procedure performed: {PROCEDURE}', ()))

    def test_flag_off_keeps_a_block_control(self):
        flag = '<w:showingPlcHdr w:val="0"/>'
        xml = (f'<w:sdt>{sdt_pr(flag)}<w:sdtContent><w:p><w:r><w:t>{FINDINGS}</w:t></w:r></w:p></w:sdtContent>'
               '</w:sdt>')
        result = extract(body(xml))
        self.assertEqual((result.status, result.text, result.messages), ('ok', FINDINGS, ()))

    def test_flag_off_keeps_a_table_row_control(self):
        row = (f'<w:tr><w:tc><w:p><w:r><w:t>Procedure</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>{PROCEDURE}</w:t>'
               '</w:r></w:p></w:tc></w:tr>')
        flag = '<w:showingPlcHdr w:val="false"/>'
        xml = ('<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr><w:tblGrid><w:gridCol w:w="2000"/>'
               f'<w:gridCol w:w="6000"/></w:tblGrid><w:sdt>{sdt_pr(flag)}<w:sdtContent>{row}</w:sdtContent></w:sdt>'
               '</w:tbl>')
        result = extract(body(xml))
        self.assertEqual((result.status, result.text, result.messages), ('ok', f'Procedure | {PROCEDURE}', ()))

    def test_true_placeholders_are_still_skipped(self):
        for flag in ('<w:showingPlcHdr/>', '<w:showingPlcHdr w:val="true"/>', '<w:showingPlcHdr w:val="1"/>',
                     '<w:showingPlcHdr w:val="on"/>'):
            with self.subTest(flag=flag):
                result = extract(body(inline_control(flag, 'Click or tap here to enter text.')))
                self.assertEqual((result.text, result.messages), ('Procedure performed: ', (PLACEHOLDER_ONE,)))


# --- ingest:9 — comments and processing instructions inside w:t --------------------------------

PARAS = ['SYNTHETIC TEST DOCUMENT - FICTIONAL PATIENT - NOT FOR CLINICAL USE', 'Date of service: 2026-10-20',
         'Procedure: colonoscopy with cold snare polypectomy of a 7 mm sigmoid polyp.',
         'Procedure: EGD with cold forceps biopsy of the gastric antrum.',
         'Impression: sigmoid polyp removed; gastric antrum biopsied.']


def patched_document(edits: list[tuple[str, str]], before_root: str = '') -> bytes:
    data = docgen.make_docx(PARAS)
    xml = zipfile.ZipFile(io.BytesIO(data)).read('word/document.xml').decode('utf-8')
    for old, new in edits:
        assert xml.count(old) == 1, f'expected exactly one {old!r}'
        xml = xml.replace(old, new)
        assert new in xml
    if before_root:
        end = xml.index('?>') + 2
        xml = xml[:end] + before_root + xml[end:]
    return docgen.rewrite_zip(data, {'word/document.xml': xml.encode('utf-8')})


class DocxTextNodeTests(unittest.TestCase):
    """ingest:9 — w:t is all of its character data; comments and PIs are not content."""

    def test_comment_or_pi_inside_a_run_keeps_the_whole_sentence(self):
        cases = {'comment mid-run': [('colonoscopy with', 'colonoscopy <!-- merge field: procedure_detail -->with')],
                 'PI mid-run': [('EGD with', 'EGD <?tpl field="biopsy"?>with')],
                 'comment first': [('>Impression: sigmoid', '><!-- tpl:impression -->Impression: sigmoid')],
                 'two comments': [('colonoscopy with', 'colonoscopy <!--a-->wi<!--b-->th')]}
        for label, edits in cases.items():
            with self.subTest(label):
                result = extract(patched_document(edits))
                self.assertEqual((result.status, result.messages), ('ok', ()))
                self.assertEqual(result.text, '\n'.join(PARAS))

    def test_a_document_type_declaration_is_refused(self):
        for prefix in ('<!DOCTYPE w:document [<!ENTITY dx "cold snare polypectomy">]>',
                       '<!DOCTYPE w:document []>'):
            with self.subTest(prefix=prefix):
                edits = [('cold snare polypectomy', '&dx;')] if 'ENTITY' in prefix else []
                result = extract(patched_document(edits, before_root=prefix))
                self.assertEqual((result.status, result.text, result.messages),
                                 ('failed', '', ('The document declares a DTD, which Word documents do not use; it '
                                                 'was not read.',)))

    def test_header_text_counts_comment_tails(self):
        data = docgen.make_docx(['Body.'], header_text='Synthetic Clinic Letterhead')
        package = zipfile.ZipFile(io.BytesIO(data))
        header = next(name for name in package.namelist() if name.startswith('word/header'))
        xml = package.read(header).decode('utf-8')
        assert xml.count('Synthetic Clinic') == 1
        data = docgen.rewrite_zip(data, {header: xml.replace('Synthetic Clinic', 'Synthetic <!--x-->Clinic').encode()})
        self.assertEqual(extract(data).messages, ('Header/footer text is not extracted (25 characters).',))


# --- ingest:10 — Markup Compatibility: the first understood Choice, else the Fallback ---------------

COMPLETED = 'Colonoscopy completed to cecum with snare polypectomy.'
CANCELLED = 'Colonoscopy cancelled before sedation.'
UNKNOWN = 'xmlns:q7="urn:example:synthetic:never-understood"'
SKIPPED_ONE = ('Not extracted from the document body: 1 alternative content block for newer Word versions (its '
               'backward-compatible version was read instead).')


def t(text: str) -> str:
    return f'<w:t xml:space="preserve">{text}</w:t>'


class DocxAlternateContentTests(unittest.TestCase):
    """ingest:10 — a Choice whose Requires namespaces are not understood is not read."""

    def placements(self, choice: str, fallback: str) -> dict[str, str]:
        return {
            'paragraph level': (f'<w:p><w:r>{t("Procedure status: ")}</w:r><mc:AlternateContent {UNKNOWN}>'
                                f'<mc:Choice Requires="q7"><w:r>{t(choice)}</w:r></mc:Choice>'
                                f'<mc:Fallback><w:r>{t(fallback)}</w:r></mc:Fallback></mc:AlternateContent></w:p>'),
            'run level': (f'<w:p><w:r>{t("Procedure status: ")}</w:r><w:r><mc:AlternateContent {UNKNOWN}>'
                          f'<mc:Choice Requires="q7">{t(choice)}</mc:Choice><mc:Fallback>{t(fallback)}</mc:Fallback>'
                          '</mc:AlternateContent></w:r></w:p>'),
            'body level': (f'<w:p><w:r>{t("Procedure status: ")}</w:r></w:p><mc:AlternateContent {UNKNOWN}>'
                           f'<mc:Choice Requires="q7"><w:p><w:r>{t(choice)}</w:r></w:p></mc:Choice>'
                           f'<mc:Fallback><w:p><w:r>{t(fallback)}</w:r></w:p></mc:Fallback></mc:AlternateContent>'),
        }

    def test_unknown_choice_reads_the_fallback(self):
        for choice, fallback in ((COMPLETED, CANCELLED), (CANCELLED, COMPLETED)):
            for label, xml in self.placements(choice, fallback).items():
                with self.subTest(label, shown=fallback):
                    result = extract(body(xml))
                    self.assertEqual(result.status, 'ok')
                    self.assertIn(fallback, result.text)
                    self.assertNotIn(choice, result.text)
                    self.assertEqual(result.messages, (SKIPPED_ONE,))

    def test_unknown_choice_without_fallback_is_partial(self):
        xml = (f'<w:p><w:r>{t("Procedure status: ")}</w:r><mc:AlternateContent {UNKNOWN}><mc:Choice Requires="q7">'
               f'<w:r>{t(COMPLETED)}</w:r></mc:Choice></mc:AlternateContent></w:p>')
        result = extract(body(xml))
        self.assertEqual((result.status, result.text), ('partial', 'Procedure status: '))
        self.assertEqual(result.messages, ('The document has 1 block of content for newer Word versions that could '
                                           'not be read and has no backward-compatible version.',))
        # The parent process accepts a partial Word result from the sandboxed child.
        self.assertEqual(extract_document(body(xml), 'status.docx', DOCX, ocr=None), result)

    def test_understood_choice_is_read_once(self):
        box = '<w:txbxContent><w:p><w:r><w:t>Text box: cecum reached.</w:t></w:r></w:p></w:txbxContent>'
        xml = ('<w:p><w:r><w:t>Anchor paragraph.</w:t></w:r><w:r><mc:AlternateContent><mc:Choice Requires="wps">'
               f'<w:drawing><wp:anchor><a:graphic><a:graphicData><wps:wsp><wps:txbx>{box}</wps:txbx></wps:wsp>'
               '</a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape>'
               f'<v:textbox>{box}</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>')
        result = extract(body(xml))
        self.assertEqual((result.text, result.messages), ('Anchor paragraph.\nText box: cecum reached.', ()))
        w14 = (f'<w:p><mc:AlternateContent><mc:Choice Requires="w14"><w:r>{t(COMPLETED)}</w:r></mc:Choice>'
               f'<mc:Fallback><w:r>{t(CANCELLED)}</w:r></mc:Fallback></mc:AlternateContent></w:p>')
        self.assertEqual(extract(body(w14)).text, COMPLETED)

    def test_table_rows_wrapped_in_alternate_content_are_read(self):
        def row(text):
            cells = ''.join(f'<w:tc><w:p><w:r>{t(value)}</w:r></w:p></w:tc>' for value in ('Procedure', text))
            return f'<w:tr>{cells}</w:tr>'

        grid = ('<w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr><w:tblGrid><w:gridCol w:w="2000"/>'
                '<w:gridCol w:w="6000"/></w:tblGrid>')
        understood = (f'<w:tbl>{grid}<mc:AlternateContent><mc:Choice Requires="w14">{row(COMPLETED)}</mc:Choice>'
                      f'<mc:Fallback>{row(CANCELLED)}</mc:Fallback></mc:AlternateContent></w:tbl>')
        self.assertEqual(extract(body(understood)).text, f'Procedure | {COMPLETED}')
        unknown = (f'<w:tbl>{grid}<mc:AlternateContent {UNKNOWN}><mc:Choice Requires="q7">{row(COMPLETED)}'
                   f'</mc:Choice><mc:Fallback>{row(CANCELLED)}</mc:Fallback></mc:AlternateContent></w:tbl>')
        result = extract(body(unknown))
        self.assertEqual((result.text, result.messages), (f'Procedure | {CANCELLED}', (SKIPPED_ONE,)))

    def test_emoji_fallback_is_kept(self):
        xml = (f'<w:p><w:r>{t("Prep quality: ")}</w:r><w:r><mc:AlternateContent '
               'xmlns:w16se="http://schemas.microsoft.com/office/word/2015/wordml/symex"><mc:Choice Requires="w16se">'
               '<w16se:symEx w16se:font="Segoe UI Emoji" w16se:char="2705"/></mc:Choice>'
               f'<mc:Fallback>{t("✅")}</mc:Fallback></mc:AlternateContent></w:r><w:r>{t(" adequate (BBPS 8).")}</w:r>'
               '</w:p>')
        result = extract(body(xml))
        self.assertEqual((result.text, result.messages), ('Prep quality: ✅ adequate (BBPS 8).', ()))


# --- ingest:11 — glyphs typed in a symbol font -------------------------------------------------

def symbol_document(segments) -> bytes:
    """segments: (text, font) pairs, or ('SYM', font, hex) for a w:sym element."""
    document = docx.Document()
    paragraph = document.add_paragraph()
    for segment in segments:
        if segment[0] == 'SYM':
            sym = OxmlElement('w:sym')
            sym.set(qn('w:font'), segment[1])
            sym.set(qn('w:char'), segment[2])
            paragraph.add_run()._r.append(sym)
            continue
        text, font = segment
        run = paragraph.add_run(text)
        if font:
            run.font.name = font
    return saved(document)


TECHNIQUE = 'Polypectomy technique: '


class DocxSymbolFontTests(unittest.TestCase):
    """ingest:11 — symbol-font glyphs are mapped to the Unicode symbols they show, or counted."""

    def test_checkbox_glyphs_are_mapped(self):
        cases = {
            'Wingdings w:t': [('þ', 'Wingdings'), (' cold snare  ', None), ('¨', 'Wingdings'), (' hot snare', None)],
            'Wingdings private-use w:t': [('', 'Wingdings'), (' cold snare  ', None), ('', 'Wingdings'),
                                          (' hot snare', None)],
            'Wingdings 2 w:t': [('R', 'Wingdings 2'), (' cold snare  ', None), ('£', 'Wingdings 2'),
                                (' hot snare', None)],
            'w:sym': [('SYM', 'Wingdings', 'F0FE'), (' cold snare  ', None), ('SYM', 'Wingdings', 'F0A8'),
                      (' hot snare', None)],
        }
        # Round 2 (ingest:0): every recognised box is written as ☐ (unchecked) or ☒ (checked).
        for label, segments in cases.items():
            with self.subTest(label):
                result = extract(symbol_document([(TECHNIQUE, None)] + segments))
                self.assertEqual((result.status, result.messages), ('ok', ()))
                self.assertEqual(result.text, f'{TECHNIQUE}☒ cold snare  ☐ hot snare')

    def test_unmapped_symbol_glyphs_are_counted_not_shown_as_letters(self):
        result = extract(symbol_document([('Mood: ', None), ('J', 'Wingdings'), (' and ', None),
                                          ('SYM', 'Webdings', 'F061')]))
        self.assertEqual(result.text, 'Mood:  and ')
        self.assertEqual(result.messages, ('Not extracted from the document body: 2 symbol-font characters.',))

    def test_symbol_font_is_read_through_its_encoding(self):
        result = extract(symbol_document([('Polyp ', None), ('³', 'Symbol'), (' 10 mm; ', None), ('m', 'Symbol'),
                                          ('g dose ', None), ('±', 'Symbol'), (' 5', None)]))
        self.assertEqual((result.text, result.messages), ('Polyp ≥ 10 mm; μg dose ± 5', ()))

    def test_symbol_font_from_a_character_style(self):
        document = docx.Document()
        char_style(document, 'Boxes', hidden=None, font='Wingdings')
        paragraph = document.add_paragraph(TECHNIQUE)
        paragraph.add_run('þ', style='Boxes')
        paragraph.add_run(' cold snare')
        result = extract(saved(document))
        self.assertEqual((result.text, result.messages), (f'{TECHNIQUE}☒ cold snare', ()))

    def test_ordinary_fonts_and_symbols_are_unchanged(self):
        result = extract(symbol_document([('Prep ', None), ('þ ¨ ☑ ☐ ≥ µ', 'Calibri'), (' fine.', 'Arial')]))
        self.assertEqual((result.text, result.messages), ('Prep þ ¨ ☑ ☐ ≥ µ fine.', ()))

    def test_private_use_glyph_in_an_ordinary_font_is_counted(self):
        result = extract(symbol_document([('Checked: ', None), ('', None)]))
        self.assertEqual((result.text, result.messages),
                         ('Checked: ', ('Not extracted from the document body: 1 symbol-font character.',)))


if __name__ == '__main__':
    unittest.main()
