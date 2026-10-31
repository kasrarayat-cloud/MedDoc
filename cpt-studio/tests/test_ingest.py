"""Document ingestion (ARCHITECTURE 5.1). Synthetic documents only; no database.

The sandbox tests run the real child process. OCR is exercised three ways:
forced unavailable (ocr=None), with executable stand-ins for pdftoppm and
tesseract (always runs, checks the plumbing), and with the real tools when
this host has them.
"""
from __future__ import annotations

import importlib.util
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import textwrap
import time
import unittest
import zipfile
from unittest import mock

import docx
import pypdf

from studio.ingest import (DETECT, Limits, canonicalize, engine_versions, extract_document, extract_paste,
                           find_ocr_tools, ocr_available)
from studio.ingest import _child as child_module
from studio.ingest import docx as docx_ingest
from studio.ingest import extract as extract_module
from studio.ingest.pdf import extract_pdf
from studio.ingest.text import extract_txt, join_pages, page_phrase
from studio.types import ExtractionResult, PageSpan
from tests.support import docgen

PDF = 'application/pdf'
DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
TXT = 'text/plain'
MIB = 1024 * 1024
OCR_NOTE = 'OCR can misread characters — check evidence against the original.'
NO_TEXT = 'No text could be read from the PDF.'
W = docgen.W_NS


def page_texts(result: ExtractionResult) -> list[str]:
    """Each page's slice of the text, after checking the pages tile it exactly."""
    slices = [result.text[span.start:span.end] for span in result.pages]
    assert [span.page for span in result.pages] == list(range(1, len(slices) + 1))
    assert '\n\n'.join(slices) == result.text, 'pages must tile the text with the separator'
    return slices


def replace_once(text: str, old: str, new: str) -> str:
    assert text.count(old) == 1, f'expected exactly one {old!r}'
    return text.replace(old, new)


def gone(pid: int, within_s: float = 5.0) -> bool:
    """True once ``pid`` no longer runs (absent, or a zombie awaiting its reaper)."""
    deadline = time.monotonic() + within_s
    while time.monotonic() < deadline:
        try:
            with open(f'/proc/{pid}/stat') as handle:
                if handle.read().rsplit(')', 1)[1].split()[0] == 'Z':
                    return True
        except (FileNotFoundError, ProcessLookupError):
            return True
        time.sleep(0.05)
    return False


# --- canonical text ---------------------------------------------------------------

class CanonicalizeTests(unittest.TestCase):
    def test_line_endings_become_lf(self):
        self.assertEqual(canonicalize('a\r\nb\rc\nd\r\n\r\n'), 'a\nb\nc\nd\n\n')

    def test_nul_is_removed(self):
        self.assertEqual(canonicalize('co\x00lon\x00'), 'colon')

    def test_unicode_nfc(self):
        self.assertEqual(canonicalize('café'), 'café')
        self.assertEqual(canonicalize('Å'), 'Å')
        self.assertEqual(canonicalize('가'), '가')

    def test_nothing_else_changes(self):
        text = ('\tTabs, trailing spaces   \n\n\n\x0cform feed\x0b  nbsp  line sep ​zwsp '
                'ﬁ ligature ① circled １ fullwidth')
        self.assertEqual(canonicalize(text), text)

    def test_nul_inside_a_line_ending_and_a_combining_sequence(self):
        self.assertEqual(canonicalize('a\r\x00\nb'), 'a\n\nb')  # line endings are handled before NUL removal
        self.assertEqual(canonicalize('e\x00́'), 'é')  # NFC runs after NUL removal

    def test_idempotent(self):
        for text in ('a\r\n\r\x00\nb', 'e\x00́\r', 'ᄀ\x00ᅡ', 'plain'):
            once = canonicalize(text)
            self.assertEqual(canonicalize(once), once)


class PageHelperTests(unittest.TestCase):
    def test_page_phrase(self):
        self.assertEqual(page_phrase([2]), 'Page 2')
        self.assertEqual(page_phrase([1, 2]), 'Pages 1-2')
        self.assertEqual(page_phrase([1, 2, 3, 5, 7, 8]), 'Pages 1-3, 5, 7-8')

    def test_join_pages_spans_slice_back_including_empty_pages(self):
        text, spans = join_pages(['first', '', 'third'])
        self.assertEqual(text, 'first\n\n\n\nthird')
        self.assertEqual(spans, (PageSpan(1, 0, 5), PageSpan(2, 7, 7), PageSpan(3, 9, 14)))

    def test_join_pages_refuses_non_canonical_text(self):
        with self.assertRaises(ValueError):
            join_pages(['line\r'])


class LimitsTests(unittest.TestCase):
    def test_defaults(self):
        self.assertEqual(Limits(), Limits(timeout_s=240, max_pages=40, max_chars=400_000, ocr_dpi=300,
                                          ocr_page_timeout_s=60, min_text_chars_per_page=20, max_bytes=15 * MIB))

    def test_rejects_non_positive_non_finite_and_wrong_types(self):
        for bad in ({'timeout_s': 0}, {'timeout_s': -1}, {'timeout_s': float('nan')}, {'timeout_s': float('inf')},
                    {'max_pages': 1.5}, {'max_pages': True}, {'max_chars': '400000'}, {'ocr_dpi': 0}):
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                Limits(**bad)


# --- TXT and paste ----------------------------------------------------------------

class TxtTests(unittest.TestCase):
    def test_utf8_text_is_one_page(self):
        result = extract_txt('Synthetic GI note: EGD to the duodenum.'.encode(), Limits())
        self.assertEqual((result.status, result.method, result.messages), ('ok', 'txt', ()))
        self.assertEqual(result.text, 'Synthetic GI note: EGD to the duodenum.')
        self.assertEqual(result.pages, (PageSpan(1, 0, len(result.text)),))

    def test_bom_is_stripped_and_text_canonicalised(self):
        result = extract_txt(b'\xef\xbb\xbfLine one\r\nLine\x00 two\rcafe\xcc\x81', Limits())
        self.assertEqual(result.text, 'Line one\nLine two\ncafé')

    def test_invalid_utf8_fails(self):
        for data in (b'caf\xe9 au lait', 'Synthetic note'.encode('utf-16'), b'\xed\xa0\x80 surrogate'):
            with self.subTest(data=data[:6]):
                result = extract_txt(data, Limits())
                self.assertEqual((result.status, result.text, result.pages), ('failed', '', ()))
                self.assertEqual(result.messages, ('The text file is not valid UTF-8.',))

    def test_empty_whitespace_and_bom_only_fail(self):
        for data in (b'', b' \n\t\r\n', b'\xef\xbb\xbf', b'\x00\x00'):
            with self.subTest(data=data):
                self.assertEqual(extract_txt(data, Limits()).messages, ('The text file is empty.',))

    def test_too_long_fails_without_truncation(self):
        result = extract_txt(b'x' * 11, Limits(max_chars=10))
        self.assertEqual((result.status, result.text), ('failed', ''))
        self.assertEqual(result.messages, ('The text file has 11 characters; the limit is 10.',))
        self.assertEqual(extract_txt(b'x' * 10, Limits(max_chars=10)).status, 'ok')

    def test_through_the_child_process(self):
        result = extract_document('Ward note\r\nERCP planned next week.'.encode(), 'note.txt', 'text/plain; charset=utf-8')
        self.assertEqual((result.status, result.method, result.text), ('ok', 'txt', 'Ward note\nERCP planned next week.'))
        failure = extract_document(b'\xff\xfe\x00bad', 'note.txt', TXT)
        self.assertEqual(failure.messages, ('The text file is not valid UTF-8.',))


class PasteTests(unittest.TestCase):
    def test_pasted_text_is_canonical_and_one_page(self):
        result = extract_paste('Office visit.\r\nEstablished patient.\x00')
        self.assertEqual((result.status, result.method), ('ok', 'paste'))
        self.assertEqual(result.text, 'Office visit.\nEstablished patient.')
        self.assertEqual(result.pages, (PageSpan(1, 0, len(result.text)),))

    def test_addendum_method(self):
        result = extract_paste('Addendum: cecum was reached.', method='addendum')
        self.assertEqual((result.status, result.method), ('ok', 'addendum'))
        self.assertEqual(extract_paste('  ', method='addendum').messages, ('The addendum is empty.',))

    def test_empty_too_long_and_invalid_fail(self):
        self.assertEqual(extract_paste('').messages, ('The pasted text is empty.',))
        self.assertEqual(extract_paste('\x00 \r\n').messages, ('The pasted text is empty.',))
        self.assertEqual(extract_paste('abcdef', limits=Limits(max_chars=5)).messages,
                         ('The pasted text has 6 characters; the limit is 5.',))
        self.assertEqual(extract_paste('lone \ud800 surrogate').messages,
                         ('The pasted text contains characters that are not valid Unicode text.',))

    def test_programming_errors_raise(self):
        with self.assertRaises(ValueError):
            extract_paste('text', method='upload')
        with self.assertRaises(TypeError):
            extract_paste(b'bytes')


# --- DOCX -------------------------------------------------------------------------

class DocxTests(unittest.TestCase):
    def extract(self, data, limits=None):
        return docx_ingest.extract_docx(data, limits or Limits())

    def body(self, xml, **kwargs):
        return self.extract(docgen.make_docx_from_body(xml), **kwargs)

    def test_paragraphs_and_tables_in_document_order(self):
        data = docgen.make_docx(['Procedure: colonoscopy', [['Finding', 'Location'], ['Polyp 5 mm', 'Sigmoid']],
                                 'Impression: polyp removed.'])
        result = self.extract(data)
        self.assertEqual((result.status, result.method, result.messages), ('ok', 'docx', ()))
        self.assertEqual(result.text, 'Procedure: colonoscopy\nFinding | Location\nPolyp 5 mm | Sigmoid\n'
                                      'Impression: polyp removed.')
        self.assertEqual(result.pages, (PageSpan(1, 0, len(result.text)),))

    def test_empty_cells_keep_their_place_and_empty_rows_are_dropped(self):
        result = self.extract(docgen.make_docx([[['Biopsy', ''], ['', ''], ['', 'Rectum']]]))
        self.assertEqual(result.text, 'Biopsy | \n | Rectum')

    def test_line_breaks_and_tabs_inside_paragraphs(self):
        self.assertEqual(self.extract(docgen.make_docx(['Line one\nLine two\tTabbed'])).text,
                         'Line one\nLine two\tTabbed')

    def test_merged_cells_are_read_once(self):
        document = docx.Document()
        table = document.add_table(rows=3, cols=3)
        for row, values in zip(table.rows, [['Procedure', 'Site', 'Size'], ['Polypectomy', 'Sigmoid', '5 mm'],
                                            ['', 'Rectum', '3 mm']]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        table.cell(1, 0).merge(table.cell(2, 0))
        table.cell(0, 1).merge(table.cell(0, 2))
        # Control: python-docx's row API repeats a merged cell in every row it spans.
        self.assertEqual(table.rows[2].cells[0].text.strip(), 'Polypectomy')
        buffer = io.BytesIO()
        document.save(buffer)
        result = self.extract(buffer.getvalue())
        self.assertEqual(result.text, 'Procedure | Site Size\nPolypectomy | Sigmoid | 5 mm\n | Rectum | 3 mm')

    def test_header_and_footer_text_is_reported_not_extracted(self):
        result = self.extract(docgen.make_docx(['Body text about the EGD.'], header_text='Synthetic Clinic Letterhead'))
        self.assertEqual(result.text, 'Body text about the EGD.')
        self.assertEqual(result.messages, ('Header/footer text is not extracted (25 characters).',))

    def test_footnote_and_comment_text_is_reported_not_extracted(self):
        document = docx.Document()
        paragraph = document.add_paragraph('Polyp removed.')
        document.add_comment(paragraph.runs, text='Reviewer note text', author='Synthetic reviewer')
        buffer = io.BytesIO()
        document.save(buffer)
        data = buffer.getvalue()
        package = zipfile.ZipFile(io.BytesIO(data))
        rels = replace_once(package.read('word/_rels/document.xml.rels').decode(), '</Relationships>',
                            '<Relationship Id="rIdFootnotes" Target="footnotes.xml" Type="http://schemas.'
                            'openxmlformats.org/officeDocument/2006/relationships/footnotes"/></Relationships>')
        types = replace_once(package.read('[Content_Types].xml').decode(), '</Types>',
                             '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-'
                             'officedocument.wordprocessingml.footnotes+xml"/></Types>')
        footnotes = (f'<w:footnotes xmlns:w="{W}"><w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/>'
                     '</w:r></w:p></w:footnote><w:footnote w:id="1"><w:p><w:r><w:t>See pathology.</w:t></w:r></w:p>'
                     '</w:footnote></w:footnotes>')
        data = docgen.rewrite_zip(data, {'word/_rels/document.xml.rels': rels.encode(),
                                         '[Content_Types].xml': types.encode(), 'word/footnotes.xml': footnotes.encode()})
        result = self.extract(data)
        self.assertEqual(result.text, 'Polyp removed.')
        self.assertEqual(result.messages, ('Footnote/endnote text is not extracted (13 characters).',
                                           'Comment text is not extracted (16 characters).'))

    def test_content_controls_are_read_and_placeholders_skipped(self):
        xml = ('<w:sdt><w:sdtPr><w:alias w:val="Indication"/></w:sdtPr><w:sdtContent>'
               '<w:p><w:r><w:t>Indication: screening colonoscopy.</w:t></w:r></w:p></w:sdtContent></w:sdt>'
               '<w:p><w:r><w:t xml:space="preserve">Prep quality: </w:t></w:r><w:sdt><w:sdtPr><w:dropDownList/>'
               '</w:sdtPr><w:sdtContent><w:r><w:t>Adequate</w:t></w:r></w:sdtContent></w:sdt></w:p>'
               '<w:p><w:r><w:t xml:space="preserve">Biopsy site: </w:t></w:r><w:sdt><w:sdtPr><w:showingPlcHdr/>'
               '</w:sdtPr><w:sdtContent><w:r><w:t>Click or tap here to enter text.</w:t></w:r></w:sdtContent>'
               '</w:sdt></w:p>')
        data = docgen.make_docx_from_body(xml)
        # Control: python-docx's paragraph text drops the content-control value.
        self.assertEqual(docx.Document(io.BytesIO(data)).paragraphs[0].text, 'Prep quality: ')
        result = self.extract(data)
        self.assertEqual(result.text, 'Indication: screening colonoscopy.\nPrep quality: Adequate\nBiopsy site: ')
        self.assertEqual(result.messages,
                         ('Not extracted from the document body: 1 empty content control (placeholder text).',))

    def test_directly_hidden_runs_are_skipped(self):
        xml = ('<w:p><w:r><w:t xml:space="preserve">Polyp removed. </w:t></w:r><w:r><w:rPr><w:vanish/></w:rPr>'
               '<w:t>Hidden template note: biopsy performed.</w:t></w:r><w:r><w:rPr><w:vanish w:val="0"/></w:rPr>'
               '<w:t>Shown again.</w:t></w:r></w:p>')
        result = self.body(xml)
        self.assertEqual(result.text, 'Polyp removed. Shown again.')
        self.assertEqual(result.messages, ('Not extracted from the document body: 1 run of hidden text.',))

    def test_tracked_changes_read_as_accepted(self):
        xml = ('<w:p><w:r><w:t xml:space="preserve">Biopsy </w:t></w:r><w:del w:id="1" w:author="s">'
               '<w:r><w:delText xml:space="preserve">not </w:delText></w:r></w:del><w:ins w:id="2" w:author="s">'
               '<w:r><w:t xml:space="preserve">was </w:t></w:r></w:ins><w:r><w:t>taken.</w:t></w:r></w:p>'
               '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Kept row</w:t></w:r></w:p></w:tc></w:tr>'
               '<w:tr><w:trPr><w:del w:id="3" w:author="s"/></w:trPr><w:tc><w:p><w:r><w:t>Deleted row</w:t></w:r>'
               '</w:p></w:tc></w:tr></w:tbl>')
        result = self.body(xml)
        self.assertEqual(result.text, 'Biopsy was taken.\nKept row')
        self.assertEqual(result.messages,
                         ('The document has tracked changes; inserted text is included and deleted text is not.',))

    def test_text_box_alternate_content_is_read_once_after_its_paragraph(self):
        box = '<w:txbxContent><w:p><w:r><w:t>Text box: cecum reached.</w:t></w:r></w:p></w:txbxContent>'
        xml = ('<w:p><w:r><w:t>Anchor paragraph.</w:t></w:r><w:r><mc:AlternateContent><mc:Choice Requires="wps">'
               f'<w:drawing><wp:anchor><a:graphic><a:graphicData><wps:wsp><wps:txbx>{box}</wps:txbx></wps:wsp>'
               '</a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape>'
               f'<v:textbox>{box}</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>'
               '<w:p><w:r><w:t>Next paragraph.</w:t></w:r></w:p>')
        self.assertEqual(self.body(xml).text, 'Anchor paragraph.\nText box: cecum reached.\nNext paragraph.')

    def test_fields_show_their_result_not_their_code(self):
        xml = ('<w:p><w:r><w:t xml:space="preserve">Date: </w:t></w:r><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
               '<w:r><w:instrText xml:space="preserve"> DATE \\@ "yyyy-MM-dd" </w:instrText></w:r>'
               '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>2026-10-21</w:t></w:r>'
               '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
               '<w:p><w:hyperlink w:anchor="findings"><w:r><w:t>Linked text</w:t></w:r></w:hyperlink>'
               '<w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple></w:p>')
        self.assertEqual(self.body(xml).text, 'Date: 2026-10-21\nLinked text1')

    def test_content_without_text_is_counted_not_dropped_silently(self):
        xml = ('<w:p><w:r><w:t xml:space="preserve">Consent signed </w:t></w:r><w:r><w:sym w:font="Wingdings" '
               'w:char="F0FE"/></w:r><w:r><w:drawing><wp:inline><a:graphic><a:graphicData><pic:pic/>'
               '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>'
               '<w:p><w:r><w:fldChar w:fldCharType="begin"><w:ffData><w:name w:val="Check1"/><w:checkBox>'
               '<w:default w:val="1"/></w:checkBox></w:ffData></w:fldChar></w:r><w:r><w:instrText> FORMCHECKBOX '
               '</w:instrText></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r><w:r>'
               '<w:t xml:space="preserve"> Biopsy taken</w:t></w:r></w:p>')
        result = self.body(xml)
        # The Wingdings check box (F0FE) is read as the checked box it shows (ingest:11), and the legacy
        # FORMCHECKBOX, checked by default, is written as its state in place of the field (round 2, ingest:1).
        self.assertEqual(result.text, 'Consent signed ☒\n☒ Biopsy taken')
        self.assertEqual(result.messages, ('Not extracted from the document body: 1 image.',))

    def test_empty_document_fails_and_keeps_its_notes(self):
        self.assertEqual(self.extract(docgen.make_docx([])).messages, ('The Word document is empty.',))
        result = self.extract(docgen.make_docx(['', '  '], header_text='Letterhead only'))
        self.assertEqual((result.status, result.text), ('failed', ''))
        self.assertEqual(result.messages, ('The Word document is empty.',
                                           'Header/footer text is not extracted (14 characters).'))

    def test_too_long_fails_without_truncation(self):
        result = self.extract(docgen.make_docx(['x' * 11]), Limits(max_chars=10))
        self.assertEqual(result.messages, ('The Word document has 11 characters; the limit is 10.',))

    def test_damaged_packages_fail(self):
        whole = docgen.make_docx(['Synthetic text.'])
        for data in (b'PK\x03\x04' + b'\x00' * 200, whole[:300], b'not a zip at all',
                     docgen.rewrite_zip(whole, {'word/document.xml': None})):
            with self.subTest(data=data[:12]):
                result = self.extract(data)
                self.assertEqual((result.status, result.messages), ('failed', (docx_ingest.DAMAGED,)))

    def test_ole_documents_are_refused(self):
        result = self.extract(docx_ingest.OLE_SIGNATURE + b'\x00' * 512)
        self.assertEqual(result.messages, (docx_ingest.OLE,))

    def test_macro_and_template_packages_are_refused(self):
        whole = docgen.make_docx(['Synthetic text.'])
        types = zipfile.ZipFile(io.BytesIO(whole)).read('[Content_Types].xml').decode()
        main = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml'
        cases = {
            'application/vnd.ms-word.document.macroEnabled.main+xml': docx_ingest.MACROS,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml':
                'The file is a Word template, not a Word document.',
            'application/vnd.ms-excel.sheet.main+xml': 'The file is not a Word document.',
        }
        for content_type, message in cases.items():
            with self.subTest(content_type=content_type):
                data = docgen.rewrite_zip(whole, {'[Content_Types].xml': replace_once(types, main, content_type).encode()})
                self.assertEqual(self.extract(data).messages, (message,))
        with_project = docgen.rewrite_zip(whole, {'word/vbaProject.bin': b'\x00' * 64})
        self.assertEqual(self.extract(with_project).messages, (docx_ingest.MACROS,))

    def test_package_size_guards(self):
        whole = docgen.make_docx(['Synthetic text.'])
        padded = docgen.rewrite_zip(whole, {'word/media/padding.bin': b'\x00' * (2 * MIB)})
        with mock.patch.object(docx_ingest, 'MAX_PACKAGE_BYTES', MIB):
            self.assertEqual(self.extract(padded).messages,
                             ('The DOCX package expands to more than 1 MB; it was not opened.',))
        with mock.patch.object(docx_ingest, 'MAX_PACKAGE_MEMBERS', 3):
            self.assertEqual(self.extract(whole).messages, ('The DOCX package has more than 3 parts; it was not opened.',))

    def test_through_the_child_process_and_deterministic(self):
        data = docgen.make_docx(['EGD with biopsy.', [['Site', 'Antrum']]], header_text='Synthetic header')
        first = extract_document(data, 'egd.docx', DOCX)
        self.assertEqual((first.status, first.text), ('ok', 'EGD with biopsy.\nSite | Antrum'))
        self.assertEqual(first.messages, ('Header/footer text is not extracted (15 characters).',))
        self.assertEqual(first, extract_document(data, 'egd.docx', DOCX))


# --- PDF text layer -----------------------------------------------------------------

class PdfTests(unittest.TestCase):
    PAGES = ['Colonoscopy report. Date of service: 2026-10-21. Indication: screening.',
             'Findings: 5 mm sigmoid polyp removed with cold snare. Cecum reached.',
             'Impression: one polyp removed. Plan: surveillance colonoscopy in 7 years.']

    def test_text_layer_pages_have_exact_spans(self):
        result = extract_document(docgen.make_pdf(self.PAGES), 'report.pdf', PDF, ocr=None)
        self.assertEqual((result.status, result.method, result.messages, result.ocr_pages),
                         ('ok', 'pdf-text', (), ()))
        self.assertEqual([text.strip() for text in page_texts(result)], self.PAGES)
        self.assertEqual(result.engine, {'ingest': 'cpt-studio-ingest@4', 'pypdf': '6.21.4'})

    def test_each_page_marker_lands_in_its_own_span(self):
        pages = [f'Marker page {n}: synthetic findings text for this page only.' for n in range(1, 7)]
        result = extract_document(docgen.make_pdf(pages), 'six.pdf', PDF, ocr=None)
        for number, text in enumerate(page_texts(result), 1):
            self.assertEqual([n for n in range(1, 7) if f'Marker page {n}:' in text], [number])

    def test_password_protected_pdf_fails(self):
        result = extract_document(docgen.make_encrypted_pdf('Synthetic encrypted note text.', 'secret'), 'e.pdf', PDF)
        self.assertEqual((result.status, result.text, result.messages), ('failed', '', ('The PDF is password-protected.',)))

    def test_aes_encrypted_pdf_is_reported_by_what_this_server_can_do(self):
        # An AES-256 (V5/R6) security handler: pypdf needs AES just to test the empty password.
        handler = (b' /Encrypt << /Filter /Standard /V 5 /R 6 /Length 256 /P -4 /O <' + b'00' * 48 + b'> /U <'
                   + b'00' * 48 + b'> /OE <' + b'00' * 32 + b'> /UE <' + b'00' * 32 + b'> /Perms <' + b'00' * 16
                   + b'> /CF << /StdCF << /CFM /AESV3 /AuthEvent /DocOpen /Length 32 >> >> /StmF /StdCF /StrF /StdCF >>')
        plain = docgen.make_pdf(['Synthetic AES-encrypted note text.'])
        data = plain.replace(b'trailer\n<<', b'trailer\n<<' + handler)
        self.assertEqual(data.count(b'/AESV3'), 1)
        result = extract_document(data, 'aes.pdf', PDF, ocr=None)
        expected = ('The PDF is password-protected.' if importlib.util.find_spec('cryptography')
                    else 'The PDF is encrypted with AES, which this server cannot decrypt.')
        self.assertEqual((result.status, result.messages), ('failed', (expected,)))

    def test_owner_password_only_pdf_opens(self):
        result = extract_document(docgen.make_encrypted_pdf('Synthetic owner-restricted note text.', ''), 'o.pdf', PDF)
        self.assertEqual((result.status, result.text.strip()), ('ok', 'Synthetic owner-restricted note text.'))

    def test_damaged_and_truncated_pdfs_fail(self):
        whole = docgen.make_pdf(self.PAGES)
        cases = {b'%PDF-1.7\nthis is not a pdf body\n%%EOF\n': 'The PDF could not be read; the file is damaged or '
                                                                'is not a valid PDF.',
                 b'random bytes, no header': 'The PDF could not be read; the file is damaged or is not a valid PDF.',
                 whole[:len(whole) // 2]: 'The PDF is incomplete: it has no end-of-file marker. Upload the complete '
                                          'file.'}
        for data, message in cases.items():
            with self.subTest(data=data[:16]):
                result = extract_document(data, 'bad.pdf', PDF, ocr=None)
                self.assertEqual((result.status, result.messages), ('failed', (message,)))

    def test_too_many_pages_fails_without_partial_text(self):
        result = extract_document(docgen.make_pdf(self.PAGES), 'r.pdf', PDF, limits=Limits(max_pages=2), ocr=None)
        self.assertEqual((result.status, result.text, result.messages),
                         ('failed', '', ('The PDF has 3 pages; the limit is 2.',)))

    def test_pdf_without_pages_fails(self):
        writer = pypdf.PdfWriter()
        buffer = io.BytesIO()
        writer.write(buffer)
        self.assertEqual(extract_document(buffer.getvalue(), 'n.pdf', PDF).messages, ('The PDF has no pages.',))

    def test_too_much_text_fails_without_truncation(self):
        data = docgen.make_pdf(['A' * 30, 'B' * 30])
        # Over the limit on the text layer alone, before any OCR.
        self.assertEqual(extract_document(data, 'l.pdf', PDF, limits=Limits(max_chars=50), ocr=None).messages,
                         ('The extracted text has 62 characters; the limit is 50.',))
        # Over the limit only once the page separator is counted.
        self.assertEqual(extract_document(data, 'l.pdf', PDF, limits=Limits(max_chars=63), ocr=None).messages,
                         ('The extracted text has 64 characters; the limit is 63.',))
        self.assertEqual(extract_document(data, 'l.pdf', PDF, limits=Limits(max_chars=64), ocr=None).status, 'ok')

    def test_sparse_text_layer_page_counts_as_no_text_layer(self):
        data = docgen.make_pdf([self.PAGES[0], 'Page 2 of 2'])
        result = extract_document(data, 'r.pdf', PDF, ocr=None)
        self.assertEqual((result.status, result.method, result.ocr_pages), ('partial', 'pdf-text', ()))
        self.assertEqual(result.messages, ('Page 2 has no usable text layer and OCR is not available on this server.',))
        self.assertEqual([text.strip() for text in page_texts(result)], [self.PAGES[0], 'Page 2 of 2'])
        self.assertEqual(result.engine['ocr'], 'not available')

    def test_invalid_characters_in_a_text_layer_are_replaced_and_reported(self):
        data = docgen.make_pdf([self.PAGES[0]])
        with mock.patch.object(pypdf.PageObject, 'extract_text', return_value='Polyp \ud800 removed with cold snare.'), \
                tempfile.TemporaryDirectory() as workdir:
            result = extract_pdf(data, Limits(), None, time.monotonic() + 60, workdir)
        self.assertEqual((result.status, result.text), ('ok', 'Polyp � removed with cold snare.'))
        self.assertEqual(result.messages, ('Page 1 contains characters the PDF does not encode as valid text; '
                                           'they are shown as �.',))

    def test_text_page_with_a_large_image_is_named(self):
        large = extract_document(docgen.make_pdf([self.PAGES[0]], image_px=(1200, 1000)), 'i.pdf', PDF, ocr=None)
        self.assertEqual((large.status, large.method, large.text.strip()), ('ok', 'pdf-text', self.PAGES[0]))
        self.assertEqual(large.messages, ('Page 1 combines a text layer with a large image; only the text layer was '
                                          'read, so check the text against the original if the image holds text.',))
        small = extract_document(docgen.make_pdf([self.PAGES[0]], image_px=(300, 100)), 'i.pdf', PDF, ocr=None)
        self.assertEqual((small.status, small.messages), ('ok', ()))

    def test_deterministic(self):
        data = docgen.make_pdf(self.PAGES)
        self.assertEqual(extract_document(data, 'r.pdf', PDF, ocr=None), extract_document(data, 'r.pdf', PDF, ocr=None))


# --- scanned PDFs and OCR ---------------------------------------------------------

PDFTOPPM_SCRIPT = '''#!/bin/sh
if [ "$1" = "-v" ]; then echo "pdftoppm version 0.0-fake" >&2; exit 0; fi
page=""
while [ $# -gt 2 ]; do
  if [ "$1" = "-f" ]; then page="$2"; fi
  shift
done
@RENDER@
'''
TESSERACT_SCRIPT = '''#!/bin/sh
case "$1" in
  --version) echo "tesseract 0.0-fake"; exit 0 ;;
  --list-langs) echo 'List of available languages in "/fake/" (1):'; echo "@LANGUAGES@"; exit 0 ;;
esac
page=$(cat "$1")
@RECOGNISE@
'''


class FakeOcr:
    """Executable stand-ins for pdftoppm and tesseract: no recognition, synthetic output.

    pdftoppm writes 'page N' as the "image"; tesseract runs ``recognise``, a
    shell fragment in which $page holds that text.
    """
    RENDER = 'printf "page %s" "$page" > "$2.png"'
    RECOGNISE = 'echo "Synthetic OCR text for $page: cecum reached, cold snare polypectomy."'

    def __init__(self, recognise=RECOGNISE, render=RENDER, languages='eng'):
        self.directory = tempfile.mkdtemp(prefix='fake-ocr-')
        self._write('pdftoppm', PDFTOPPM_SCRIPT.replace('@RENDER@', render))
        self._write('tesseract', TESSERACT_SCRIPT.replace('@RECOGNISE@', recognise).replace('@LANGUAGES@', languages))
        self.tools = find_ocr_tools(self.directory)

    def _write(self, name, script):
        path = os.path.join(self.directory, name)
        with open(path, 'w') as handle:
            handle.write(script)
        os.chmod(path, 0o700)

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        shutil.rmtree(self.directory)


def ocr_text(page: int) -> str:
    return f'Synthetic OCR text for page {page}: cecum reached, cold snare polypectomy.\n'


class OcrDetectionTests(unittest.TestCase):
    def test_finds_tools_and_reports_their_versions(self):
        with FakeOcr() as fake:
            self.assertIsNotNone(fake.tools)
            self.assertEqual((fake.tools.tesseract_version, fake.tools.pdftoppm_version),
                             ('tesseract 0.0-fake', 'pdftoppm version 0.0-fake'))
            self.assertEqual(fake.tools.tesseract, os.path.join(fake.directory, 'tesseract'))

    def test_missing_language_or_tool_means_unavailable(self):
        with FakeOcr(languages='deu') as fake:
            self.assertIsNone(fake.tools)
        with FakeOcr() as fake:
            os.unlink(os.path.join(fake.directory, 'pdftoppm'))
            self.assertIsNone(find_ocr_tools(fake.directory))

    def test_availability_and_engine_versions_agree(self):
        versions = engine_versions()
        self.assertIsInstance(ocr_available(), bool)
        self.assertEqual((versions['pypdf'], versions['python-docx']), ('6.21.4', '1.2.0'))
        self.assertEqual('ocr' not in versions, ocr_available())
        if not ocr_available():
            self.assertEqual(versions['ocr'], 'not available')


class OcrUnavailableTests(unittest.TestCase):
    """OCR forced off, which is also this host's real state when the tools are absent."""

    def test_scanned_only_pdf_fails_naming_the_page(self):
        result = extract_document(docgen.make_scanned_pdf(['Synthetic scanned colonoscopy report.']), 's.pdf', PDF,
                                  ocr=None)
        self.assertEqual((result.status, result.method, result.text, result.ocr_pages), ('failed', 'pdf-text', '', ()))
        self.assertEqual(result.messages, ('Page 1 has no usable text layer and OCR is not available on this server.',
                                           NO_TEXT))
        self.assertEqual(result.engine['ocr'], 'not available')

    def test_several_scanned_pages_are_listed_together(self):
        result = extract_document(docgen.make_scanned_pdf(['One', 'Two', 'Three']), 's.pdf', PDF, ocr=None)
        self.assertEqual(result.messages[0],
                         'Pages 1-3 have no usable text layer and OCR is not available on this server.')

    def test_text_page_plus_scanned_page_is_partial(self):
        text_page = 'Office visit, established patient. Total time 25 minutes.'
        data = docgen.concat_pdfs([docgen.make_pdf([text_page]), docgen.make_scanned_pdf(['Scanned page text.'])])
        result = extract_document(data, 'mixed.pdf', PDF, ocr=None)
        self.assertEqual((result.status, result.method, result.ocr_pages), ('partial', 'pdf-text', ()))
        self.assertEqual(result.messages, ('Page 2 has no usable text layer and OCR is not available on this server.',))
        self.assertEqual(page_texts(result), [text_page + '\n', ''])


class OcrPlumbingTests(unittest.TestCase):
    """The OCR path end to end through the child, with stand-in tools."""

    def test_scanned_pdf_is_read_with_ocr(self):
        with FakeOcr() as fake:
            result = extract_document(docgen.make_scanned_pdf(['one', 'two']), 's.pdf', PDF, ocr=fake.tools)
        self.assertEqual((result.status, result.method, result.ocr_pages), ('ok', 'pdf-ocr', (1, 2)))
        self.assertEqual(page_texts(result), [ocr_text(1), ocr_text(2)])
        self.assertEqual(result.messages, (f'Pages 1-2 were read with OCR; {OCR_NOTE}',))
        self.assertEqual((result.engine['tesseract'], result.engine['pdftoppm']),
                         ('tesseract 0.0-fake', 'pdftoppm version 0.0-fake'))

    def test_text_and_scanned_pages_are_mixed(self):
        text_page = 'EGD report: duodenum reached, biopsies of the antrum taken.'
        data = docgen.concat_pdfs([docgen.make_pdf([text_page]), docgen.make_scanned_pdf(['scanned'])])
        with FakeOcr() as fake:
            result = extract_document(data, 'm.pdf', PDF, ocr=fake.tools)
        self.assertEqual((result.status, result.method, result.ocr_pages), ('ok', 'pdf-mixed', (2,)))
        self.assertEqual(page_texts(result), [text_page + '\n', ocr_text(2)])
        self.assertEqual(result.messages, (f'Page 2 was read with OCR; {OCR_NOTE}',))

    def test_a_failed_page_makes_the_result_partial(self):
        recognise = 'if [ "$page" = "page 2" ]; then exit 3; fi\n' + FakeOcr.RECOGNISE
        with FakeOcr(recognise=recognise) as fake:
            result = extract_document(docgen.make_scanned_pdf(['one', 'two']), 's.pdf', PDF, ocr=fake.tools)
        self.assertEqual((result.status, result.method, result.ocr_pages), ('partial', 'pdf-mixed', (1,)))
        self.assertEqual(page_texts(result), [ocr_text(1), ''])
        self.assertEqual(result.messages, (f'Page 1 was read with OCR; {OCR_NOTE}',
                                           'Page 2 has no usable text layer and OCR could not read it '
                                           '(tesseract exited with status 3).'))

    def test_page_timeout(self):
        with FakeOcr(recognise='sleep 5\n' + FakeOcr.RECOGNISE) as fake:
            started = time.monotonic()
            result = extract_document(docgen.make_scanned_pdf(['one']), 's.pdf', PDF,
                                      limits=Limits(ocr_page_timeout_s=1), ocr=fake.tools)
        self.assertLess(time.monotonic() - started, 4)
        self.assertEqual(result.messages, ('Page 1 has no usable text layer and OCR could not read it (OCR did not '
                                           'finish within the 1-second page limit).', NO_TEXT))

    def test_tool_failures_are_specific(self):
        cases = [(':', FakeOcr.RECOGNISE, 'pdftoppm produced no image'),
                 (FakeOcr.RENDER, "printf '\\377\\376'", 'tesseract returned text that is not valid UTF-8'),
                 (FakeOcr.RENDER, 'kill -TERM $$', 'tesseract was stopped by SIGTERM')]
        for render, recognise, problem in cases:
            with self.subTest(problem=problem), FakeOcr(recognise=recognise, render=render) as fake:
                result = extract_document(docgen.make_scanned_pdf(['one']), 's.pdf', PDF, ocr=fake.tools)
                self.assertEqual(result.messages,
                                 (f'Page 1 has no usable text layer and OCR could not read it ({problem}).', NO_TEXT))

    def test_oversized_page_is_not_rendered(self):
        data = docgen.make_pdf([''], page_size_pt=(30000, 30000))
        with FakeOcr() as fake:
            result = extract_document(data, 'huge.pdf', PDF, ocr=fake.tools)
        self.assertEqual(result.messages, ('Page 1 has no usable text layer and OCR could not read it (the page is too '
                                           'large to render: 417 × 417 inches).', NO_TEXT))

    def test_blank_scanned_page_is_reported(self):
        text_page = 'Flexible sigmoidoscopy to 40 cm. No polyps seen.'
        data = docgen.concat_pdfs([docgen.make_pdf([text_page]), docgen.make_scanned_pdf([''])])
        with FakeOcr(recognise=':') as fake:
            result = extract_document(data, 'b.pdf', PDF, ocr=fake.tools)
            only_blank = extract_document(docgen.make_scanned_pdf(['']), 'b.pdf', PDF, ocr=fake.tools)
        self.assertEqual((result.status, result.method, result.ocr_pages), ('ok', 'pdf-mixed', (2,)))
        self.assertEqual(result.messages, (f'Page 2 was read with OCR; {OCR_NOTE}',
                                           'Page 2 has no readable text (OCR found none).'))
        self.assertEqual((only_blank.status, only_blank.method), ('failed', 'pdf-ocr'))
        self.assertEqual(only_blank.messages[-1], NO_TEXT)

    def test_pages_left_when_the_time_budget_runs_out_are_named(self):
        data = docgen.concat_pdfs([docgen.make_pdf([self.text_page()]), docgen.make_scanned_pdf(['a', 'b'])])
        with FakeOcr() as fake, tempfile.TemporaryDirectory() as workdir:
            result = extract_pdf(data, Limits(), fake.tools, time.monotonic() - 1, workdir)
        self.assertEqual((result.status, result.method, result.ocr_pages), ('partial', 'pdf-text', ()))
        self.assertEqual(result.messages, ('Pages 2-3 have no usable text layer and were not read: the 240-second '
                                           'extraction time limit was reached.',))

    @staticmethod
    def text_page():
        return 'ERCP report: bile duct cannulated, sphincterotomy performed.'


@unittest.skipUnless(ocr_available(), 'tesseract and pdftoppm are not installed on this host')
class RealOcrTests(unittest.TestCase):
    def test_scanned_pdf_is_read_with_real_ocr(self):
        data = docgen.make_scanned_pdf(['Synthetic colonoscopy report', 'Cecum reached. Polyp removed.'])
        result = extract_document(data, 'scan.pdf', PDF, ocr=DETECT)
        self.assertEqual((result.status, result.method, result.ocr_pages), ('ok', 'pdf-ocr', (1, 2)))
        first, second = (text.lower() for text in page_texts(result))
        self.assertIn('colonoscopy', first)
        self.assertIn('polyp', second)
        self.assertEqual(result.messages, (f'Pages 1-2 were read with OCR; {OCR_NOTE}',))


# --- the sandbox ------------------------------------------------------------------

LIMIT_PROBE = textwrap.dedent('''
    import json, os, resource
    names = ('RLIMIT_AS', 'RLIMIT_CPU', 'RLIMIT_FSIZE', 'RLIMIT_NPROC', 'RLIMIT_CORE')
    print(json.dumps({'limits': {name: resource.getrlimit(getattr(resource, name)) for name in names},
                      'env': sorted(os.environ), 'cwd': os.getcwd(), 'own_group': os.getpgid(0) == os.getpid()}))
''')
SPAWNER = textwrap.dedent('''
    import subprocess, time
    child = subprocess.Popen(['sleep', '30'])
    print(child.pid, flush=True)
    time.sleep(30)
''')


class SandboxTests(unittest.TestCase):
    def child(self, code, data=b'Synthetic text.', limits=None):
        with mock.patch.object(extract_module, 'CHILD_ARGV', (sys.executable, '-c', code)):
            return extract_document(data, 'note.txt', TXT, limits=limits or Limits(timeout_s=30), ocr=None)

    def test_resource_limits_environment_and_session(self):
        with mock.patch.dict(os.environ, {'CPT_STUDIO_PROBE_SECRET': 'must-not-leak'}), \
                mock.patch.object(extract_module, 'CHILD_ARGV', (sys.executable, '-c', LIMIT_PROBE)):
            run = extract_module._run_child(b'', Limits(timeout_s=7))
        self.assertEqual((run.outcome, run.returncode), ('exited', 0))
        probe = json.loads(run.stdout)
        limits = probe['limits']
        self.assertEqual(limits['RLIMIT_AS'], [1536 * MIB] * 2)
        self.assertEqual(limits['RLIMIT_CPU'], [8, 13])
        self.assertEqual(limits['RLIMIT_FSIZE'], [256 * MIB] * 2)
        self.assertEqual(limits['RLIMIT_CORE'], [0, 0])
        nproc_soft, nproc_hard = limits['RLIMIT_NPROC']
        self.assertEqual(nproc_soft, nproc_hard)
        self.assertLessEqual(nproc_soft, extract_module._user_task_count() + 2 * extract_module.NPROC_HEADROOM)
        self.assertTrue(set(probe['env']) <= {'PATH', 'LC_ALL', 'OMP_THREAD_LIMIT', 'TESSDATA_PREFIX',
                                              'PYTHONDONTWRITEBYTECODE', 'PYTHONNOUSERSITE', 'PYTHONUTF8'})
        self.assertEqual(probe['cwd'], str(extract_module.SERVICE_DIR))
        self.assertTrue(probe['own_group'])

    def test_memory_limit_is_enforced(self):
        result = self.child('bytearray(2 * 1024 ** 3)')
        self.assertEqual(result.messages, ('Extraction failed: the extraction process exited with status 1.',))

    def test_timeout_kills_the_whole_process_group(self):
        started = time.monotonic()
        with mock.patch.object(extract_module, 'CHILD_ARGV', (sys.executable, '-c', SPAWNER)):
            run = extract_module._run_child(b'', Limits(timeout_s=1.5))
        self.assertLess(time.monotonic() - started, 5)
        self.assertEqual((run.outcome, run.returncode), ('timeout', -9))
        self.assertTrue(gone(int(run.stdout.split()[0])), 'the grandchild must not outlive the timeout')

    def test_timeout_message(self):
        self.assertEqual(self.child('import time; time.sleep(30)', limits=Limits(timeout_s=1.5)).messages,
                         ('Extraction stopped after 1.5 seconds.',))
        real = extract_document(docgen.make_pdf(['Synthetic text.']), 'r.pdf', PDF, limits=Limits(timeout_s=0.001))
        self.assertEqual((real.status, real.text, real.messages),
                         ('failed', '', ('Extraction stopped after 0.001 seconds.',)))

    def test_crashes_and_exit_statuses(self):
        cases = {'import os, signal; os.kill(os.getpid(), signal.SIGTERM)':
                 'Extraction failed: the extraction process crashed (SIGTERM).',
                 'import os, signal; os.kill(os.getpid(), signal.SIGKILL)':
                 'Extraction stopped: the extraction process was killed, usually because the document needed more '
                 'memory than allowed.',
                 # Not dumpable first: SIGXCPU dumps core, and a piped core_pattern ignores RLIMIT_CORE.
                 'import ctypes, os, signal; ctypes.CDLL(None).prctl(4, 0, 0, 0, 0); '
                 'os.kill(os.getpid(), signal.SIGXCPU)':
                 'Extraction stopped: the document needed more CPU time than allowed.',
                 'import sys; sys.exit(3)': 'Extraction failed: the extraction process exited with status 3.'}
        for code, message in cases.items():
            with self.subTest(code=code):
                result = self.child(code)
                self.assertEqual((result.status, result.method, result.text, result.messages),
                                 ('failed', 'txt', '', (message,)))

    def test_oversized_output_is_cut_off(self):
        result = self.child('import sys; sys.stdout.write("x" * 3_000_000)', limits=Limits(max_chars=1000))
        self.assertEqual(result.messages, ('Extraction stopped: the extractor produced more output than allowed.',))

    def test_unreadable_or_inconsistent_answers_are_refused(self):
        good = {'status': 'ok', 'method': 'txt', 'text': 'abc', 'pages': [{'page': 1, 'start': 0, 'end': 3}],
                'messages': [], 'ocr_pages': [], 'engine': {'ingest': 'x'}}
        self.assertEqual(self.child(f'print({json.dumps(json.dumps(good))})').status, 'ok')  # control
        variants = [
            'not json',
            json.dumps(good | {'pages': [{'page': 1, 'start': 0, 'end': 2}]}),  # span short of the text
            json.dumps(good | {'text': 'a\r\nb', 'pages': [{'page': 1, 'start': 0, 'end': 4}]}),  # not canonical
            json.dumps(good | {'method': 'docx'}),  # wrong method for the file type
            json.dumps(good | {'status': 'partial', 'messages': ['x']}),  # a text file is never partial
            json.dumps(good | {'status': 'failed', 'text': '', 'pages': [], 'messages': []}),  # a failure says why
            json.dumps(good | {'ocr_pages': [1]}),  # OCR claimed for a text file
            json.dumps(good | {'pages': [{'page': True, 'start': 0, 'end': 3}]}),  # a bool is not a page number
            json.dumps({key: value for key, value in good.items() if key != 'engine'}),
            '[' * 100_000,  # nesting deep enough to exhaust the JSON parser's recursion
        ]
        for answer in variants:
            with self.subTest(answer=answer[:60]):
                result = self.child(f'print({json.dumps(answer)})')
                self.assertEqual(result.messages, ('Extraction failed: the extractor returned an unreadable or '
                                                   'inconsistent result.',))

    def test_malformed_request_to_the_real_child(self):
        run = extract_module._run_child(b'no header line', Limits())
        self.assertEqual((run.outcome, run.returncode, run.stdout), ('exited', 2, b''))

    def test_child_that_cannot_start(self):
        with mock.patch.object(extract_module, 'CHILD_ARGV', ('/nonexistent/cpt-studio-python',)):
            result = extract_document(b'Synthetic text.', 'note.txt', TXT)
        self.assertEqual((result.status, result.messages),
                         ('failed', ('Extraction could not start on this server (No such file or directory).',)))
        with mock.patch.object(extract_module.tempfile, 'mkdtemp', side_effect=OSError(28, 'No space left on device')):
            result = extract_document(b'Synthetic text.', 'note.txt', TXT)
        self.assertEqual(result.messages, ('Extraction could not start on this server (No space left on device).',))

    def test_child_makes_itself_non_dumpable(self):
        code = ('import ctypes; from studio.ingest._child import disable_core_dumps; libc = ctypes.CDLL(None); '
                'print(libc.prctl(3, 0, 0, 0, 0)); disable_core_dumps(); print(libc.prctl(3, 0, 0, 0, 0))')
        done = subprocess.run([sys.executable, '-c', code], cwd=extract_module.SERVICE_DIR,
                              env=extract_module._child_env(), capture_output=True, text=True, timeout=60)
        self.assertEqual(done.stdout.split(), ['1', '0'], done.stderr)  # PR_GET_DUMPABLE before and after

    def test_child_does_not_import_django(self):
        code = ('import sys, studio.ingest._child, studio.ingest.pdf, studio.ingest.docx; '
                'print(sorted(m for m in sys.modules if m.split(".")[0] == "django"))')
        done = subprocess.run([sys.executable, '-c', code], cwd=extract_module.SERVICE_DIR,
                              env=extract_module._child_env(), capture_output=True, text=True, timeout=60)
        self.assertEqual((done.returncode, done.stdout.strip()), (0, '[]'), done.stderr)

    def test_unexpected_errors_fail_the_document_not_the_child(self):
        for error, message in ((MemoryError(), 'The document needs more memory to read than the extraction limit '
                                               'allows.'),
                               (RuntimeError('boom'), 'The document could not be read: the extractor stopped '
                                                      'unexpectedly (RuntimeError).')):
            with self.subTest(error=error), mock.patch.object(child_module, 'parse', side_effect=error):
                result = child_module.guarded_parse(b'x', 'docx', Limits(), None, 0.0, '/nonexistent')
                self.assertEqual((result.status, result.method, result.messages), ('failed', 'docx', (message,)))


class EntryPointTests(unittest.TestCase):
    def test_size_and_name_checks_run_before_any_child(self):
        with mock.patch.object(extract_module, '_run_child', side_effect=AssertionError('no child expected')):
            self.assertEqual(extract_document(b'', 'n.txt', TXT).messages, ('The file is empty.',))
            self.assertEqual(extract_document(b'x' * 101, 'n.txt', TXT, limits=Limits(max_bytes=100)).messages,
                             ('The file is larger than 100 bytes.',))
            self.assertEqual(extract_document(b'x' * (MIB + 1), 'n.txt', TXT, limits=Limits(max_bytes=MIB)).messages,
                             ('The file is larger than 1 MB.',))
            mismatch = extract_document(docgen.make_pdf(['Synthetic text.']), 'note.docx', PDF)
            self.assertEqual((mismatch.method, mismatch.messages),
                             ('pdf-text', ('The file name ends in .docx but the file is PDF.',)))

    def test_programming_errors_raise(self):
        with self.assertRaises(ValueError):
            extract_document(b'x', 'n.rtf', 'application/rtf')
        with self.assertRaises(TypeError):
            extract_document(b'x', 'n.txt', TXT, ocr='off')
        with self.assertRaises(TypeError):
            extract_document('text, not bytes', 'n.txt', TXT)


if __name__ == '__main__':
    unittest.main()
