"""Regressions for refutation round 3, FORMS lens, Word side: check boxes recognised by structure.

- ingest:4 / contract:6: a w14:checkbox content control is written ☒ (w14:checked on) or ☐, whatever symbols its
  checked and unchecked states are set to;
- ingest:5: a list paragraph whose numbering bullet is a box glyph is written with ☐ or ☒ at its start.

Each document is synthetic, built here from python-docx with raw WordprocessingML for the controls and the
numbering part, extracted with the sandboxed extract_document, routed for real and assessed by the real engine.
"""
import unittest

from studio.ingest import extract_document
from tests.support import docgen
from tests.test_repairs_status_r2_routing import BOXED, read, run_engine
from tests.test_repairs_r3_forms_routing import outcome

DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
W14 = 'http://schemas.microsoft.com/office/word/2010/wordml'


def docx_result(data: bytes):
    return extract_document(data, 'synthetic.docx', DOCX, ocr=None)


def docx_text(data: bytes) -> str:
    result = docx_result(data)
    if result.status != 'ok':
        raise AssertionError((result.status, result.messages))
    return result.text


def _escape(text: str) -> str:
    return text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def paragraph(text: str, ppr: str = '') -> str:
    return f'<w:p>{ppr}<w:r><w:t xml:space="preserve">{_escape(text)}</w:t></w:r></w:p>'


def control(checked: bool, unchecked_state=('2610', 'MS Gothic'), checked_state=('2612', 'MS Gothic'),
            run_text=None, run_font=None, checked_xml=None, hidden=False, placeholder=False, sym=False) -> str:
    """A Word check-box content control as Word saves it: the states it is set to in w14:checkbox and the current
    state's glyph as its content (``run_text`` overrides the glyph; ``sym`` writes it as a w:sym)."""
    state = checked_state if checked else unchecked_state
    glyph = chr(int(state[0], 16)) if run_text is None else run_text
    font = run_font or state[1]
    checked_element = (f'<w14:checked w14:val="{1 if checked else 0}"/>' if checked_xml is None else checked_xml)
    flags = ('<w:showingPlcHdr/>' if placeholder else '')
    vanish = '<w:vanish/>' if hidden else ''
    if sym:
        content = (f'<w:r><w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}"/>{vanish}</w:rPr>'
                   f'<w:sym w:font="{font}" w:char="F0{state[0][-2:]}"/></w:r>')
    else:
        content = (f'<w:r><w:rPr><w:rFonts w:ascii="{font}" w:eastAsia="{font}" w:hAnsi="{font}"/>{vanish}</w:rPr>'
                   f'<w:t>{_escape(glyph)}</w:t></w:r>')
    return (f'<w:sdt><w:sdtPr><w:id w:val="-1"/>{flags}<w14:checkbox>{checked_element}'
            f'<w14:checkedState w14:val="{checked_state[0]}" w14:font="{checked_state[1]}"/>'
            f'<w14:uncheckedState w14:val="{unchecked_state[0]}" w14:font="{unchecked_state[1]}"/></w14:checkbox>'
            f'</w:sdtPr><w:sdtContent>{content}</w:sdtContent></w:sdt>')


def note_paragraphs(lines=None) -> str:
    return ''.join(paragraph(line) for line in (lines if lines is not None else BOXED.strip('\n').split('\n')))


def box_table_docx(rows, lines=None, after=(), **box) -> bytes:
    """``lines`` (default BOXED) as paragraphs, then a table: each cell is text or ('box', checked)."""
    def cell(value):
        if isinstance(value, tuple):
            return f'<w:tc><w:p>{control(value[1], **box)}</w:p></w:tc>'
        return f'<w:tc>{paragraph(value)}</w:tc>'
    table = '<w:tbl>' + ''.join('<w:tr>' + ''.join(cell(v) for v in row) + '</w:tr>' for row in rows) + '</w:tbl>'
    return docgen.make_docx_from_body(note_paragraphs(lines) + table + ''.join(paragraph(p) for p in after))


def paragraphs_docx(extra, lines=None) -> bytes:
    return docgen.make_docx_from_body(note_paragraphs(lines) + ''.join(paragraph(p) for p in extra))


def boxed_paragraphs_docx(items, **box) -> bytes:
    """One paragraph per (checked, label): the control, then a space and the label."""
    body = ''.join(f'<w:p>{control(checked, **box)}<w:r><w:t xml:space="preserve"> {_escape(label)}</w:t></w:r></w:p>'
                   for checked, label in items)
    return docgen.make_docx_from_body(note_paragraphs() + body)


# Symbol sets a check-box control may be set to: (unchecked state, checked state).
SYMBOL_SETS = {
    'default MS Gothic': (('2610', 'MS Gothic'), ('2612', 'MS Gothic')),
    'Segoe UI Symbol 25A2 / 2611': (('25A2', 'Segoe UI Symbol'), ('2611', 'Segoe UI Symbol')),
    'Segoe UI Symbol 2B1C / 2714': (('2B1C', 'Segoe UI Symbol'), ('2714', 'Segoe UI Symbol')),
    'square 25A2 / 25A0': (('25A2', 'Segoe UI Symbol'), ('25A0', 'Segoe UI Symbol')),
    'circle 25CB / 25CF': (('25CB', 'Segoe UI Symbol'), ('25CF', 'Segoe UI Symbol')),
    'space / X': (('0020', 'Calibri'), ('0058', 'Calibri')),
    'Wingdings 2 A3 / 50': (('00A3', 'Wingdings 2'), ('0050', 'Wingdings 2')),
}


class Ingest4CheckBoxControlTests(unittest.TestCase):
    """ingest:4 / contract:6 - a check-box content control is written from w14:checked, not from its glyph."""

    def test_every_symbol_set_in_a_paragraph(self):
        for name, (unchecked, checked) in SYMBOL_SETS.items():
            with self.subTest(name):
                text = docx_text(boxed_paragraphs_docx([(False, 'Biopsy'), (True, 'Cold snare polypectomy')],
                                                       unchecked_state=unchecked, checked_state=checked))
                self.assertTrue(text.endswith('\n☐ Biopsy\n☒ Cold snare polypectomy'), text[-80:])
                self.assertEqual(read(text, 'Biopsy')[0], 'negated')
                proposals, _ = run_engine(text, [('45380', 'Biopsy'), ('45385', 'Cold snare polypectomy')])
                self.assertEqual((outcome(proposals['45380']), outcome(proposals['45385'])), ('excluded', 'ready'))

    def test_every_symbol_set_in_a_table(self):
        for name, (unchecked, checked) in SYMBOL_SETS.items():
            with self.subTest(name):
                text = docx_text(box_table_docx([['Biopsy', ('box', False)], ['Cold snare polypectomy', ('box', True)]],
                                                unchecked_state=unchecked, checked_state=checked))
                self.assertTrue(text.endswith('\nBiopsy | ☐\nCold snare polypectomy | ☒'), text[-80:])
                proposals, _ = run_engine(text, [('45380', 'Biopsy'), ('45385', 'Cold snare polypectomy')])
                self.assertEqual((outcome(proposals['45380']), outcome(proposals['45385'])), ('excluded', 'ready'))

    def test_a_symbol_font_glyph_written_as_w_sym(self):
        unchecked, checked = SYMBOL_SETS['Wingdings 2 A3 / 50']
        text = docx_text(boxed_paragraphs_docx([(False, 'Biopsy'), (True, 'Cold snare polypectomy')],
                                               unchecked_state=unchecked, checked_state=checked, sym=True))
        self.assertTrue(text.endswith('\n☐ Biopsy\n☒ Cold snare polypectomy'), text[-80:])

    def test_w14_checked_is_on_off(self):
        for checked_xml, glyph in (('<w14:checked/>', '☒'), ('<w14:checked w14:val="true"/>', '☒'),
                                   ('<w14:checked w14:val="on"/>', '☒'), ('<w14:checked w14:val="false"/>', '☐'),
                                   ('<w14:checked w14:val="off"/>', '☐'), ('', '☐')):
            with self.subTest(checked_xml=checked_xml):
                text = docx_text(boxed_paragraphs_docx([(False, 'Biopsy')], checked_xml=checked_xml))
                self.assertTrue(text.endswith(f'\n{glyph} Biopsy'), text[-40:])

    def test_a_hidden_or_placeholder_control_writes_nothing(self):
        hidden = docx_result(boxed_paragraphs_docx([(False, 'Biopsy')], hidden=True))
        self.assertTrue(hidden.text.endswith('\n Biopsy'), hidden.text[-40:])
        self.assertIn('run of hidden text', ' '.join(hidden.messages))
        placeholder = docx_result(boxed_paragraphs_docx([(False, 'Biopsy')], placeholder=True))
        self.assertTrue(placeholder.text.endswith('\n Biopsy'), placeholder.text[-40:])
        self.assertIn('placeholder', ' '.join(placeholder.messages))

    def test_a_control_around_a_whole_paragraph(self):
        """Block level: the control's content is a paragraph; the glyph is written in its place."""
        def block_control(checked):
            return ('<w:sdt><w:sdtPr><w14:checkbox><w14:checked w14:val="%d"/><w14:checkedState w14:val="2714" '
                    'w14:font="Segoe UI Symbol"/><w14:uncheckedState w14:val="2B1C" w14:font="Segoe UI Symbol"/>'
                    '</w14:checkbox></w:sdtPr><w:sdtContent><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:sdtContent>'
                    '</w:sdt>' % (checked, '\u2714' if checked else '\u2b1c'))
        rows = ''.join(f'<w:tr><w:tc>{paragraph(label)}</w:tc><w:tc>{block_control(checked)}</w:tc></w:tr>'
                       for label, checked in (('Biopsy', 0), ('Cold snare polypectomy', 1)))
        text = docx_text(docgen.make_docx_from_body(note_paragraphs() + '<w:tbl>' + rows + '</w:tbl>'))
        self.assertTrue(text.endswith('\nBiopsy | ☐\nCold snare polypectomy | ☒'), text[-80:])

    def test_a_control_that_is_not_a_check_box_is_read_as_before(self):
        body = (note_paragraphs() + '<w:p><w:sdt><w:sdtPr><w:id w:val="-2"/><w:text/></w:sdtPr><w:sdtContent><w:r>'
                '<w:t>Biopsy</w:t></w:r></w:sdtContent></w:sdt></w:p>')
        self.assertTrue(docx_text(docgen.make_docx_from_body(body)).endswith('\nBiopsy'))


# --- ingest:5 - list bullets drawn as boxes ------------------------------------------------------------------------

NUMBERING = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">{abstracts}{nums}'
             '</w:numbering>')


def level(ilvl, fmt, text, font=None, vanish=False) -> str:
    fonts = f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:hint="default"/>' if font else ''
    rpr = f'<w:rPr>{fonts}{"<w:vanish/>" if vanish else ""}</w:rPr>' if fonts or vanish else ''
    return (f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{_escape(text)}"/>'
            f'<w:lvlJc w:val="left"/>{rpr}</w:lvl>')


def numbering(abstracts: dict, nums: dict) -> bytes:
    """abstracts: id -> level xml (ilvl 0); nums: numId -> (abstractNumId, override level xml or '')."""
    a = ''.join(f'<w:abstractNum w:abstractNumId="{i}"><w:multiLevelType w:val="hybridMultilevel"/>{lvl}'
                f'</w:abstractNum>' for i, lvl in abstracts.items())
    n = ''.join(f'<w:num w:numId="{num}"><w:abstractNumId w:val="{a_id}"/>'
                + (f'<w:lvlOverride w:ilvl="0">{override}</w:lvlOverride>' if override else '') + '</w:num>'
                for num, (a_id, override) in nums.items())
    return NUMBERING.format(abstracts=a, nums=n).encode('utf-8')


def numpr(num_id, ilvl=0) -> str:
    return f'<w:pPr><w:numPr><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr></w:pPr>'


def list_docx(items, abstracts, nums, styles=None) -> bytes:
    """items: (text, pPr xml)."""
    body = note_paragraphs() + ''.join(paragraph(text, ppr) for text, ppr in items)
    data = docgen.make_docx_from_body(body)
    replacements = {'word/numbering.xml': numbering(abstracts, nums)}
    if styles is not None:
        import io
        import zipfile
        original = zipfile.ZipFile(io.BytesIO(data)).read('word/styles.xml').decode('utf-8')
        replacements['word/styles.xml'] = original.replace('</w:styles>', styles + '</w:styles>').encode('utf-8')
    return docgen.rewrite_zip(data, replacements)


WINGDINGS_BOXES = {0: level(0, 'bullet', '\uf0a8', 'Wingdings'), 1: level(0, 'bullet', '\uf0fe', 'Wingdings')}


class Ingest5BoxBulletTests(unittest.TestCase):
    """ingest:5 - a list paragraph whose bullet is a box glyph starts with ☐ or ☒."""

    def assert_boxed(self, data):
        text = docx_text(data)
        self.assertTrue(text.endswith('\n☒ Cold snare polypectomy\n☐ Biopsy'), text[-80:])
        self.assertEqual(read(text, 'Biopsy')[0], 'negated')
        proposals, _ = run_engine(text, [('45380', 'Biopsy'), ('45385', 'Cold snare polypectomy')])
        self.assertEqual((outcome(proposals['45380']), outcome(proposals['45385'])), ('excluded', 'ready'))

    def test_wingdings_bullets_through_a_direct_numbering(self):
        self.assert_boxed(list_docx([('Cold snare polypectomy', numpr(2)), ('Biopsy', numpr(1))], WINGDINGS_BOXES,
                                    {1: (0, ''), 2: (1, '')}))

    def test_wingdings_bullets_through_a_paragraph_style(self):
        styles = ''.join(f'<w:style w:type="paragraph" w:styleId="{sid}"><w:name w:val="{sid}"/>'
                         f'<w:pPr><w:numPr><w:numId w:val="{num}"/></w:numPr></w:pPr></w:style>'
                         for sid, num in (('CheckOff', 1), ('CheckOn', 2)))
        style = '<w:pPr><w:pStyle w:val="{}"/></w:pPr>'
        self.assert_boxed(list_docx([('Cold snare polypectomy', style.format('CheckOn')),
                                     ('Biopsy', style.format('CheckOff'))], WINGDINGS_BOXES, {1: (0, ''), 2: (1, '')},
                                    styles=styles))

    def test_a_level_override(self):
        plain = {0: level(0, 'bullet', '\uf0b7', 'Symbol')}
        self.assert_boxed(list_docx([('Cold snare polypectomy', numpr(2)), ('Biopsy', numpr(1))], plain,
                                    {1: (0, level(0, 'bullet', '\uf0a8', 'Wingdings')),
                                     2: (0, level(0, 'bullet', '\uf0fe', 'Wingdings'))}))

    def test_a_numbering_style_link(self):
        """An abstract numbering that links to a numbering style takes its levels from that style's numbering."""
        styles = ('<w:style w:type="numbering" w:styleId="CheckList"><w:name w:val="Check List"/><w:pPr><w:numPr>'
                  '<w:numId w:val="7"/></w:numPr></w:pPr></w:style>')
        abstracts = {5: '<w:numStyleLink w:val="CheckList"/>', 6: level(0, 'bullet', '\uf0a8', 'Wingdings'),
                     1: level(0, 'bullet', '\uf0fe', 'Wingdings')}
        self.assert_boxed(list_docx([('Cold snare polypectomy', numpr(2)), ('Biopsy', numpr(3))], abstracts,
                                    {3: (5, ''), 7: (6, ''), 2: (1, '')}, styles=styles))

    def test_literal_box_characters(self):
        abstracts = {0: level(0, 'bullet', '☐', 'Segoe UI Symbol'), 1: level(0, 'bullet', '☑', 'Segoe UI Symbol')}
        self.assert_boxed(list_docx([('Cold snare polypectomy', numpr(2)), ('Biopsy', numpr(1))], abstracts,
                                    {1: (0, ''), 2: (1, '')}))

    def test_ordinary_lists_are_unchanged(self):
        abstracts = {0: level(0, 'bullet', '\uf0b7', 'Symbol'), 1: level(0, 'decimal', '%1.'),
                     2: level(0, 'bullet', 'o', 'Courier New'), 3: level(0, 'bullet', '\uf0a8', 'Wingdings', vanish=True)}
        text = docx_text(list_docx([('Biopsy', numpr(1)), ('Cold snare polypectomy', numpr(2)), ('Hemoclip', numpr(3)),
                                    ('Tattoo', numpr(0)), ('Retroflexion', numpr(4))],
                                   abstracts, {1: (0, ''), 2: (1, ''), 3: (2, ''), 4: (3, '')}))
        self.assertTrue(text.endswith('\nBiopsy\nCold snare polypectomy\nHemoclip\nTattoo\nRetroflexion'), text[-80:])
        builtin = docgen.make_docx([])
        import docx as python_docx
        import io
        document = python_docx.Document(io.BytesIO(builtin))
        document.add_paragraph('Biopsy', style='List Bullet')
        document.add_paragraph('Cold snare polypectomy', style='List Number')
        buffer = io.BytesIO()
        document.save(buffer)
        self.assertEqual(docx_text(buffer.getvalue()), 'Biopsy\nCold snare polypectomy')


if __name__ == '__main__':
    unittest.main()
